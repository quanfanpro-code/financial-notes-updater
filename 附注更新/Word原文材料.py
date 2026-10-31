# coding: utf-8
"""保留Word单元格和既有位置关联，只呈现原文，不判断业务含义。"""
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl.utils.cell import get_column_letter

from .表格 import file_hash


def word_source_material(link_path, positions=None):
    from .分步流程 import _load_link
    link = _load_link(link_path)
    ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    with ZipFile(link['word_path']) as package:
        document = ET.fromstring(package.read('word/document.xml'))
        styles = ({style.get(ns + 'styleId'): style for style in
                   ET.fromstring(package.read('word/styles.xml')).findall(ns + 'style')}
                  if 'word/styles.xml' in package.namelist() else {})
    links = {row['table_id']: row for row in link['context'].get('tables', [])
             if row.get('table_id')}
    table_number = 0

    def paragraph(node):
        return ''.join(part.text or '' if part.tag == ns + 't' else
                       '\t' if part.tag == ns + 'tab' else '\n'
                       for part in node.iter() if part.tag in (ns + 't', ns + 'tab', ns + 'br', ns + 'cr'))

    def outline_level(node):
        value = node.find(ns + 'pPr/' + ns + 'outlineLvl')
        style = node.find(ns + 'pPr/' + ns + 'pStyle')
        identifier = style.get(ns + 'val') if style is not None else None
        seen = set()
        while value is None and identifier in styles and identifier not in seen:
            seen.add(identifier)
            style = styles[identifier]
            value = style.find(ns + 'pPr/' + ns + 'outlineLvl')
            parent = style.find(ns + 'basedOn')
            identifier = parent.get(ns + 'val') if parent is not None else None
        return int(value.get(ns + 'val')) if value is not None else None

    def blocks(parent, path):
        result = []
        for index, node in enumerate(parent):
            where = path + '/' + str(index)
            if node.tag == ns + 'p':
                block = {'kind': 'paragraph', 'word_position': where, 'text': paragraph(node)}
                level = outline_level(node)
                if level is not None:
                    block['word_outline_level'] = level
                result.append(block)
            elif node.tag == ns + 'tbl':
                result.append(table(node, where))
            elif node.tag in (ns + 'sdt', ns + 'sdtContent'):
                result.extend(blocks(node, where))
        return result

    def table(node, where):
        nonlocal table_number
        table_number += 1
        number = table_number
        linked = links.get(where, {})
        rows = []
        vertical = {}
        for row_index, row in enumerate(node.findall(ns + 'tr')):
            before = row.find(ns + 'trPr/' + ns + 'gridBefore')
            after = row.find(ns + 'trPr/' + ns + 'gridAfter')
            start = int(before.get(ns + 'val')) if before is not None else 0
            column = start
            cells = []
            for cell_index, cell in enumerate(row.findall(ns + 'tc')):
                span = cell.find(ns + 'tcPr/' + ns + 'gridSpan')
                width = int(span.get(ns + 'val')) if span is not None else 1
                merge = cell.find(ns + 'tcPr/' + ns + 'vMerge')
                value = {'word_row': row_index + 1, 'word_column': column + 1,
                         'column_span': width, 'row_span': 1,
                         'content': blocks(cell, where + '/row/' + str(row_index) + '/cell/' + str(cell_index))}
                if merge is not None and merge.get(ns + 'val') != 'restart':
                    anchor = vertical.get(column)
                    value['vertical_merge'] = 'continue'
                    value['merge_anchor'] = ([anchor['word_row'], anchor['word_column']]
                                             if anchor is not None else None)
                    if anchor is not None:
                        anchor['row_span'] += 1
                elif merge is not None:
                    value['vertical_merge'] = 'restart'
                    vertical[column] = value
                else:
                    vertical.pop(column, None)
                if linked.get('sheet'):
                    coordinate = get_column_letter(linked['first_column'] + column) + str(linked['first_row'] + row_index)
                    value['source_excel'] = linked['sheet'] + '!' + coordinate
                    if positions is not None:
                        value['current_excel'] = positions.get(value['source_excel'])
                cells.append(value)
                column += width
            rows.append({'word_row': row_index + 1, 'grid_before': start,
                         'grid_after': int(after.get(ns + 'val')) if after is not None else 0,
                         'cells': cells})
        return {'kind': 'table', 'table_number': number, 'table_id': where,
                'range_name': linked.get('range_name'), 'sheet': linked.get('sheet'),
                'grid_columns': len(node.findall(ns + 'tblGrid/' + ns + 'gridCol')),
                'rows': rows}

    result = {'word_source': {'path': link['word_path'], 'sha256': link['word_hash']},
              'source_excel': {'path': link['a_path'], 'sha256': link['a_hash']},
              'blocks': blocks(document.find(ns + 'body'), '/word/document.xml#body')}
    if file_hash(link['word_path']) != link['word_hash'] or file_hash(link['a_path']) != link['a_hash']:
        raise ValueError('读取期间Word或已有Excel位置关联发生变化')
    return result


def word_table_context(material, areas):
    """按原生提纲保留本节、祖先共用说明及明确读取的表，不判断业务适用性。"""
    from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries
    blocks = material['blocks']; tables = {}
    for index, block in enumerate(blocks):
        if block['kind'] != 'table':
            continue
        cells = [cell for row in block['rows'] for cell in row['cells'] if cell.get('source_excel')]
        if not cells:
            continue
        points = [(*coordinate_to_tuple(cell['source_excel'].rsplit('!', 1)[1]), cell) for cell in cells]
        top = min(r for r, c, cell in points); left = min(c for r, c, cell in points)
        bottom = max(r + cell['row_span'] - 1 for r, c, cell in points)
        right = max(c + cell['column_span'] - 1 for r, c, cell in points)
        tables[index] = {'table_id': block['table_id'], 'sheet': block['sheet'],
                         'range': f'{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}'}
    targets = []
    for index, table in tables.items():
        l, t, r, b = range_boundaries(table['range'])
        for area in areas:
            al, at, ar, ab = range_boundaries(area['range'])
            if table['sheet'] == area['sheet'] and al <= r and ar >= l and at <= b and ab >= t:
                targets.append(index); break
    headings = [(i, block['word_outline_level']) for i, block in enumerate(blocks)
                if block['kind'] == 'paragraph' and isinstance(block.get('word_outline_level'), int)
                and 0 <= block['word_outline_level'] < 9]
    # shortcut: 没有原生提纲或明确关联表时保留全文，取得可靠范围后再缩。
    if not targets or not headings:
        return None
    keep = set(range(headings[0][0] + 1))
    report_level = min(level for index, level in headings)
    for index, level in headings:
        if level == report_level:
            stop = next((i for i, depth in headings if i > index), len(blocks))
            keep.update(range(index, stop))
    for target in targets:
        ancestors = []
        for index, level in headings:
            if index >= target:
                break
            while ancestors and ancestors[-1][1] >= level:
                ancestors.pop()
            ancestors.append((index, level))
        if not ancestors:
            return None
        for index, level in ancestors:
            stop = next((i for i, depth in headings if i > index), len(blocks))
            keep.update(range(index, stop))
        index, level = ancestors[-1]
        stop = next((i for i, depth in headings if i > index and depth <= level), len(blocks))
        keep.update(range(index, stop))
    lines = []; ranges = []
    for index in sorted(keep):
        block = blocks[index]
        if block['kind'] == 'paragraph':
            lines.append('原段落 ' + block['word_position'] +
                         ('；原提纲级别=' + str(block['word_outline_level'])
                          if 'word_outline_level' in block else '') + '\n' + block['text'])
        elif index in tables:
            table = tables[index]; ranges.append(table)
            lines.append('[原Word表' + str(block['table_number']) + '：' + block['table_id'] +
                         '；原表区=' + table['sheet'] + '!' + table['range'] + ']')
        else:
            lines.append(word_text_material({**material, 'blocks': [block]}))
    return ({'title': '原Word正文顺序（表格引用指向实际表区，各段适用范围由模型判断）',
             'text': '\n'.join(lines), 'source_path': material['word_source']['path'],
             'source_hash': material['word_source']['sha256'],
             'reading_basis': '按原文件提纲保留本节全部原文、报告层及祖先开头说明、明确读取的表；是否适用由模型判断。'}, ranges)


def word_request_material(material):
    """只将重复字段名和段落位置前缀共享，原文、合并和坐标值不变。"""
    from copy import deepcopy
    result = deepcopy(material)
    fields = ['word_column', 'column_span', 'row_span',
              'source_excel', 'current_excel', 'content']
    body = '/word/document.xml#body'
    for index, table in enumerate(result['blocks']):
        if table['kind'] == 'paragraph' and set(table) <= {'kind', 'word_position', 'text'}:
            result['blocks'][index] = [table['word_position'].removeprefix(body), table['text']]
            continue
        if table['kind'] != 'table':
            continue
        for row in table['rows']:
            packed = []
            for cell in row['cells']:
                content = []
                for block in cell['content']:
                    if block['kind'] == 'paragraph' and set(block) <= {'kind', 'word_position', 'text'}:
                        content.append([block['kind'], block['word_position'].removeprefix(table['table_id']), block['text']])
                    else:
                        content.append(block)
                values = [cell.get(key) for key in fields]
                for key in ('source_excel', 'current_excel'):
                    at = fields.index(key)
                    if values[at] and table.get('sheet') and values[at].startswith(table['sheet'] + '!'):
                        values[at] = values[at].removeprefix(table['sheet'] + '!')
                values[-1] = content
                values.append({key: value for key, value in cell.items() if key not in fields and key != 'word_row'})
                packed.append(values)
            row['cells'] = packed
    result['cell_fields'] = fields + ['attributes']
    result['content_fields'] = ['kind', 'word_position_relative_to_table_id', 'text']
    result['paragraph_fields'] = ['word_position_relative_to_body', 'text']
    result['paragraph_position_base'] = body
    result['excel_coordinates_relative_to_table_sheet'] = True
    return result


def word_text_material(material):
    """按原顺序呈现全文与二维网格，位置关联独立列示，不拼接表头含义。"""
    import json
    from .表格文字 import grid_text
    lines = ['原Word完整文字及表格结构。表内行列号只在各自表内有效，不是Excel列号。',
             '并入表示原文件明确的合并覆盖；原格空白与省略位置分开保留。',
             '原提取Excel位置和当前Excel位置仅用于定位，不参与金标准或身份证判断。']

    def blocks(items):
        for block in items:
            if block['kind'] == 'paragraph':
                level = block.get('word_outline_level')
                prefix = '#' * (level + 1) + ' ' if isinstance(level, int) and 0 <= level < 9 else ''
                lines.extend(['\n原段落位置：' + block['word_position'] +
                              ('；原Word提纲级别=' + str(level) if level is not None else ''),
                              prefix + block['text']])
                continue
            if block['kind'] != 'table':
                lines.append(json.dumps(block, ensure_ascii=False))
                continue
            table = block
            cells = [cell for row in table['rows'] for cell in row['cells']]
            values = {}
            merges = []
            for cell in cells:
                if cell.get('vertical_merge') == 'continue':
                    continue
                values[cell['word_row'], cell['word_column']] = '\n'.join(
                    part['text'] if part['kind'] == 'paragraph' else '〔此格内含嵌套表〕'
                    for part in cell['content'])
                if cell['row_span'] > 1 or cell['column_span'] > 1:
                    merges.append((cell['word_row'], cell['word_column'],
                                   cell['word_row'] + cell['row_span'] - 1,
                                   cell['word_column'] + cell['column_span'] - 1))
            width = max([table['grid_columns']] + [c['word_column'] + c['column_span'] - 1 for c in cells])
            lines.append('\nWord表号：' + str(table['table_number']) + '；原表标识：' + table['table_id'] +
                         '；原范围名称：' + str(table.get('range_name')))
            lines.append(grid_text([r['word_row'] for r in table['rows']],
                                   {c: '第' + str(c) + '列' for c in range(1, width + 1)}, values, merges,
                                   lambda r, c: '第' + str(r) + '行第' + str(c) + '列'))
            lines.append('原始格与Excel位置关联：')
            for row in table['rows']:
                if row['grid_before'] or row['grid_after']:
                    lines.append('第' + str(row['word_row']) + '行：行前省略' + str(row['grid_before']) +
                                 '列，行后省略' + str(row['grid_after']) + '列。')
                for cell in row['cells']:
                    lines.append('第' + str(cell['word_row']) + '行第' + str(cell['word_column']) + '列：原提取Excel=' +
                                 str(cell.get('source_excel')) + '；当前Excel=' + str(cell.get('current_excel')) +
                                 '；行跨度=' + str(cell['row_span']) + '；列跨度=' + str(cell['column_span']))
                    if cell.get('vertical_merge') == 'continue':
                        lines.append('纵向合并延续起点：' + str(cell.get('merge_anchor')) + '；该格原段落：' +
                                     json.dumps([part['text'] for part in cell['content'] if part['kind']=='paragraph'], ensure_ascii=False))
                    for part in cell['content']:
                        if part['kind'] == 'table':
                            blocks([part])
    blocks(material['blocks'])
    lines.append('\n原始文件与位置关联依据：' + json.dumps(
        {key: value for key, value in material.items() if key != 'blocks'}, ensure_ascii=False))
    return '\n'.join(lines)
