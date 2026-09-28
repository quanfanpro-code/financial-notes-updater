"""独立五样式评测工具。人工映射仅在评测阶段读取，不向识别引擎提供答案。"""
from __future__ import annotations
import argparse
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import re
from openpyxl import load_workbook

DEFAULT_SOURCE = Path(r"D:\BaiduSyncdisk\数字资产\审计数据自动搬运系统")
DIMENSION_ALIASES = {
    "期间角色":"period_role", "币种":"currency", "单位":"unit", "倍率":"scale",
    "报表口径":"scope", "核算主体":"entity", "汇总层级":"aggregation_role",
    "字段":"metric", "field_path":"metric", "measure":"metric",
    "项目":"project", "project_name":"project",
    "客户名称":"counterparty", "单位名称":"counterparty",
    "被投资单位":"investee", "被投资企业":"investee",
    "账龄":"age_bucket", "账龄区间":"age_bucket", "aging_bucket":"age_bucket",
    "asset_category":"asset_class", "资产类别":"asset_class",
    "计量模式":"measurement_model", "变动类型":"movement_type",
}
TRACE_KEYS = {"source_row_path","source_column_path","source_value_type","source_subtable",
              "raw_row_path","raw_column_path","row_path","column_path","sheet","cell",
              "address","来源汇总表","来源科目"}
ROLE_ALIASES = {"期末":"closing","年末":"closing","期初":"opening","年初":"opening",
                "本期":"current","本年":"current","上期":"prior","上年":"prior"}
VALUE_ALIASES = {
    "currency":{"人民币":"CNY","人民币元":"CNY","RMB":"CNY","美元":"USD","港币":"HKD","港元":"HKD"},
    "scope":{"合并":"consolidated","合并报表":"consolidated","母公司":"parent","母公司报表":"parent"},
    "aggregation_role":{"明细":"detail","合计":"total","小计":"subtotal"},
}


def _hash(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _jsonl(path):
    for number, line in enumerate(Path(path).read_text(encoding="utf-8-sig").splitlines(),1):
        if line.strip():
            try:
                yield json.loads(line)
            except json.JSONDecodeError as error:
                raise ValueError(f"{path} 第 {number} 行不是有效 JSON") from error


def _paths(style, source_root):
    if isinstance(style,bool) or style not in range(1,6):
        raise ValueError("样式号必须为 1 至 5")
    root = Path(source_root)
    folder = root / "work" / "excel-mapping-v1" / f"style{style}-manual-review-v2"
    paths = sorted(folder.glob("*人工映射.jsonl")) if style == 2 else [folder / f"样式{style}-合并映射.jsonl"]
    if not paths or any(not p.is_file() for p in paths):
        raise ValueError(f"未找到样式 {style} 的人工映射")
    workbook = root / "附注样式示例" / f"附注样式{style}.xlsx"
    gold = root / "语义金标准" / "财务报表附注语义金标准-v1.jsonl"
    if not workbook.is_file() or not gold.is_file():
        raise ValueError("缺少原始工作簿或正式金标准")
    return workbook, gold, paths


def _cell(value):
    address = str(value).replace("$","").upper().strip()
    if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*",address):
        raise ValueError("无效的单元格坐标：" + address)
    return address


def _position(item):
    sheet = item.get("sheet",item.get("worksheet"))
    if not isinstance(sheet,str) or not sheet:
        raise ValueError("评测记录缺少工作表名称")
    return sheet,_cell(item.get("cell",""))


def read_benchmark(style: int, source_root=DEFAULT_SOURCE) -> list[dict]:
    """读取人工真值和原始单元格；公式另保留缓存，绝不重新计算或覆盖原文件。"""
    workbook, gold_path, paths = _paths(style,source_root)
    gold = {item["id"]:item for item in _jsonl(gold_path)}
    raw = load_workbook(workbook,data_only=False,read_only=False)
    cached = load_workbook(workbook,data_only=True,read_only=False)
    records = []
    seen = set()
    source_hash = _hash(workbook)
    try:
        for path in paths:
            for item in _jsonl(path):
                slot_id = item.get("target_slot_id")
                if not slot_id:
                    continue
                if slot_id not in gold or gold[slot_id].get("status") != "active":
                    raise ValueError("人工映射引用非 active 金标准：" + str(slot_id))
                sheet,cell = _position(item)
                if (sheet,cell) in seen:
                    raise ValueError(f"人工真值重复：{sheet}!{cell}")
                seen.add((sheet,cell))
                if sheet not in raw.sheetnames:
                    raise ValueError("原工作簿没有工作表：" + sheet)
                source = raw[sheet][cell]
                is_formula = source.data_type == "f"
                actual = cached[sheet][cell].value if is_formula else source.value
                records.append({
                    "sheet":sheet,"cell":cell,"slot_id":slot_id,
                    "dimensions":item.get("dimensions") or {},
                    "scope":gold[slot_id].get("scope"),
                    "value_type":(item.get("target_slot") or {}).get("value_type",gold[slot_id].get("value_type")),
                    "value":actual,"raw_value":source.value,"is_formula":is_formula,
                    "value_available":not is_formula or actual is not None,
                    "source_hash":source_hash,"mapping_file":str(path),
                })
    finally:
        raw.close()
        cached.close()
    if not records:
        raise ValueError("该样式没有业务格人工映射")
    return records


def _normal_dimensions(dimensions):
    if not isinstance(dimensions,dict):
        raise ValueError("维度必须为对象")
    result = {}
    for key,value in dimensions.items():
        if key in TRACE_KEYS:
            continue
        key = DIMENSION_ALIASES.get(key,key)
        if isinstance(value,str):
            value = value.strip()
            if key == "period_role":
                value = ROLE_ALIASES.get(value,value)
            value = VALUE_ALIASES.get(key,{}).get(value,value)
            if key == "scale":
                try:
                    value = str(Decimal(value).normalize())
                except InvalidOperation:
                    pass
        elif key == "scale" and isinstance(value,(int,float)) and not isinstance(value,bool):
            value = str(Decimal(str(value)).normalize())
        if key in result and result[key] != value:
            # 不用后一个同义键覆盖前一个冲突值，保留到复核明细。
            result[key] = {"alias_conflict":[result[key],value]}
        else:
            result[key] = value
    return result


def _differences(expected,predicted):
    left = _normal_dimensions(expected.get("dimensions") or {})
    right = _normal_dimensions(predicted.get("dimensions") or {})
    expected_scope = expected.get("scope")
    if expected_scope:
        left["scope"] = VALUE_ALIASES["scope"].get(expected_scope,expected_scope)
        scope = predicted.get("scope",right.get("scope"))
        right["scope"] = VALUE_ALIASES["scope"].get(scope,scope)
    # 无法自动确认的维度单列复核，不因 ID 一致而假定完整语义一致。
    return [{"dimension":key,"expected":left.get(key),"predicted":right.get(key),
             "expected_present":key in left,"predicted_present":key in right}
            for key in sorted(set(left)|set(right))
            if (key in left) != (key in right) or left.get(key) != right.get(key)]


def _value(value):
    if isinstance(value,bool) or value is None:
        return type(value).__name__,value
    if isinstance(value,(int,float,Decimal)):
        return "number",Decimal(str(value))
    if isinstance(value,str):
        number = value.strip()
        if re.fullmatch(r"-?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?",number):
            try:
                return "number",Decimal(number.replace(",",""))
            except InvalidOperation:
                pass
    return "literal",value


def _verified_value_source(prediction,workbook,source_hash,benchmark):
    """重算仅允许缓存变化；原始单元格内容和合并结构必须保持一致。"""
    original_hash=prediction.get("original_hash")
    supplied=prediction.get("source_hash")
    value_source=str(workbook)
    if original_hash:
        if original_hash!=source_hash:
            raise ValueError("识别记录绑定的原始文件与人工基准不一致")
        working=Path(prediction.get("source_path",""))
        if not working.is_file() or not supplied or _hash(working)!=supplied:
            raise ValueError("重算副本缺失或哈希校验失败")
        original=load_workbook(workbook,data_only=False)
        current=load_workbook(working,data_only=False)
        cached=load_workbook(working,data_only=True)
        try:
            if original.sheetnames!=current.sheetnames:
                raise ValueError("重算副本工作表与原始资料不一致")
            for name in original.sheetnames:
                left=original[name];right=current[name]
                raw_left={cell.coordinate:(cell.data_type,cell.value) for cell in left._cells.values() if cell.value is not None}
                raw_right={cell.coordinate:(cell.data_type,cell.value) for cell in right._cells.values() if cell.value is not None}
                if raw_left!=raw_right or {str(r) for r in left.merged_cells.ranges}!={str(r) for r in right.merged_cells.ranges}:
                    raise ValueError("计算副本改动了业务单元格内容，不能套用原始真值")
            for target in benchmark:
                if target["is_formula"]:
                    target["value"]=cached[target["sheet"]][target["cell"]].value
                    target["value_available"]=target["value"] is not None
        finally:
            original.close();current.close();cached.close()
        value_source=str(working)
    elif supplied and supplied!=source_hash:
        raise ValueError("识别结果与评测原始工作簿的哈希不一致，不能套用该真值")
    return value_source,bool(supplied)


def compare_predictions(prediction_result: dict, style: int, source_root=DEFAULT_SOURCE) -> dict:
    """逐工作表/坐标核对；正确只代表 ID 正确，完整语义结论始终另行核验。"""
    if not isinstance(prediction_result,dict):
        raise ValueError("识别结果必须是 JSON 对象")
    benchmark = read_benchmark(style,source_root)
    source_hash = benchmark[0]["source_hash"]
    workbook,_,_=_paths(style,source_root)
    value_source,hash_verified=_verified_value_source(prediction_result,workbook,source_hash,benchmark)
    expected = {_position(item):item for item in benchmark}
    groups = defaultdict(list)
    invalid = []
    for category in ("mappings","excluded","unresolved"):
        items = prediction_result.get(category,[])
        if not isinstance(items,list):
            raise ValueError(category + " 必须是数组")
        for item in items:
            try:
                if not isinstance(item,dict):
                    raise ValueError("记录必须为对象")
                position = _position(item)
                if category == "mappings" and not isinstance(item.get("slot_id",item.get("target_slot_id")),str):
                    raise ValueError("映射缺少金标准 ID")
                groups[position].append((category,item))
            except (ValueError,TypeError) as error:
                invalid.append({"category":category,"record":item,"reason":str(error)})
    details = {key:[] for key in ("correct","wrong","missing","extra","duplicate","dimension_review","value_mismatch","value_unverifiable")}
    values_checked = 0
    for position,items in groups.items():
        if len(items)>1:
            details["duplicate"].append({"sheet":position[0],"cell":position[1],"records":[{"classification":c,**item} for c,item in items]})
    for position,target in expected.items():
        basic = {"sheet":position[0],"cell":position[1],"expected_slot_id":target["slot_id"]}
        items = groups.get(position,[])
        if len(items)>1:
            continue
        if not items or items[0][0] != "mappings":
            details["missing"].append({**basic,"classification":items[0][0] if items else "absent"})
            continue
        item = items[0][1]
        slot_id = item.get("slot_id",item.get("target_slot_id"))
        if slot_id != target["slot_id"]:
            details["wrong"].append({**basic,"predicted_slot_id":slot_id})
        else:
            details["correct"].append(basic)
            differences = _differences(target,item)
            if differences:
                details["dimension_review"].append({**basic,"differences":differences})
        if "value" in item:
            if target["value_available"]:
                values_checked += 1
                if _value(target["value"]) != _value(item["value"]):
                    details["value_mismatch"].append({**basic,"expected":target["value"],"predicted":item["value"]})
            else:
                details["value_unverifiable"].append({**basic,"reason":"源公式无可核验缓存，不能以空值作为计算结果"})
    for position,items in groups.items():
        if position not in expected:
            for category,item in items:
                if category == "mappings":
                    details["extra"].append({"sheet":position[0],"cell":position[1],"predicted_slot_id":item.get("slot_id",item.get("target_slot_id"))})
    counts = {key:len(items) for key,items in details.items()}
    counts.update(expected=len(expected),prediction_records=sum(len(v) for v in groups.values()),
                  duplicate_records=sum(len(v["records"]) for v in details["duplicate"]),
                  invalid=len(invalid),values_checked=values_checked)
    return {
        "style":style,"source_hash":source_hash,"counts":counts,"details":{**details,"invalid":invalid},
        "slot_accuracy":counts["correct"]/len(expected),"full_semantic_verified":False,
        "source_hash_verified":hash_verified,"value_source":value_source,
        "limitations":[
            "correct 和 slot_accuracy 仅统计单元格金标准 ID；维度、期间、主体、单位及取数须另外复核。",
            "dimension_review 包含核心同义键规范化后的不一致及尚不能自动证明等价的差异。",
            "未返回 value 时不进行模型数值抄录核验；业务程序仍须从来源单元格重新取数。",
            "本结果不证明新样式泛化；五种布局必须实行留一式独立验证。",
        ],
    }


def benchmark_summary(source_root=DEFAULT_SOURCE) -> dict:
    """只读核实五样式有效映射数量、结构排除数及文件哈希，不调用模型。"""
    rows = []
    for style in range(1,6):
        workbook,gold,paths = _paths(style,source_root)
        records = read_benchmark(style,source_root)
        excluded = sum(not item.get("target_slot_id") for path in paths for item in _jsonl(path))
        rows.append({"style":style,"workbook":str(workbook),"workbook_sha256":_hash(workbook),
                     "mapping_count":len(records),"excluded_structure_records":excluded,
                     "mapping_files":[{"path":str(p),"sha256":_hash(p)} for p in paths],
                     "all_ids_active":True})
    return {"source_root":str(Path(source_root)),"gold_sha256":_hash(Path(source_root)/"语义金标准"/"财务报表附注语义金标准-v1.jsonl"),
            "styles":rows,"total_mapping_count":sum(r["mapping_count"] for r in rows),
            "real_model_evaluated":False,
            "note":"仅核实人工基准，不表示这些单元格已被模型正确识别。"}


def main(argv=None):
    parser = argparse.ArgumentParser(description="独立比较新表识别结果与人工五样式基准")
    parser.add_argument("--source-root",default=str(DEFAULT_SOURCE))
    parser.add_argument("--summary",action="store_true")
    parser.add_argument("--prediction")
    parser.add_argument("--style",type=int)
    parser.add_argument("--output",required=True,help="全新结果 JSON 路径")
    args = parser.parse_args(argv)
    output = Path(args.output)
    if output.exists():
        parser.error("结果文件已存在，请使用新的输出文件名；不会覆盖已有文件")
    if args.summary:
        result = benchmark_summary(args.source_root)
    elif args.prediction and args.style:
        prediction = json.loads(Path(args.prediction).read_text(encoding="utf-8-sig"))
        result = compare_predictions(prediction,args.style,args.source_root)
    else:
        parser.error("选择 --summary，或同时填写 --prediction 和 --style")
    text = json.dumps(result,ensure_ascii=False,indent=2,default=str)
    if "\ufffd" in text:
        raise ValueError("输出包含异常替换字符")
    output.parent.mkdir(parents=True,exist_ok=True)
    with output.open("x",encoding="utf-8-sig") as handle:
        handle.write(text)
    if json.loads(output.read_text(encoding="utf-8-sig")) != json.loads(text):
        raise ValueError("评测结果回读不一致")
    print(str(output))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
