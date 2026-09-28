# coding: utf-8
"""说明原文独立于Excel原格；范围只用于确认说明适用性，不决定财务语义。"""
from copy import deepcopy
import json
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries, coordinate_to_tuple
from .表格 import _definitions, _destinations, _word_table_ranges
from .统一语义 import _require, definition_hash


def context_catalog(context, source_path=None):
    """原文生成稳定引用；命名区域沿用现有Word链接，移动行列后按当前范围取证。"""
    _require(isinstance(context,(list,tuple)),'语义说明须为列表')
    ranges={}
    if source_path and any(isinstance(item,dict) and item.get('range_name') for item in context):
        book=load_workbook(source_path,data_only=False)
        try:
            names={}
            for definition in _definitions(book):
                for sheet,area in _destinations(definition):names.setdefault(definition.name,[]).append({'sheet':sheet,'range':area})
            for sheet,items in (_word_table_ranges(context,book,names) or {}).items():
                for item in items:
                    left,top,right,bottom=range_boundaries(item['range'])
                    ranges[(sheet,item['table_id'],item['range_name'])]=[top,left,bottom,right]
        finally:book.close()
    catalog={}
    for item in context:
        if isinstance(item,str):
            if not item.strip():continue
            record={'kind':'user_context','content':item,'scope':None}
        else:
            _require(isinstance(item,dict),'语义说明须为原文或原文记录')
            if not any(item.get(key) for key in ('title','chapter_context','unit_text','text')):continue
            scope=None
            if item.get('sheet'):
                bounds=ranges.get((item['sheet'],item.get('table_id'),item.get('range_name')))
                if bounds is None:
                    keys=('first_row','first_column','row_count','column_count')
                    _require(all(type(item.get(key)) is int and item[key]>0 for key in keys),'局部说明缺少明确的适用范围')
                    row,column,rows,columns=(item[key] for key in keys)
                    bounds=[row,column,row+rows-1,column+columns-1]
                scope={'sheet':item['sheet'],'bounds':bounds}
            record={'kind':'document_context','content':deepcopy(item),'scope':scope}
        catalog['@context:'+definition_hash(item)]=record
    return catalog


def applicable_context(catalog,sheet,bounds=None):
    """整包可展示相交说明；每个财务格仍独立核实范围。"""
    result={}
    for key,record in catalog.items():
        scope=record['scope']
        if scope:
            if scope['sheet']!=sheet:continue
            area=scope['bounds']
            if bounds and (area[0]>bounds[2] or area[2]<bounds[0] or area[1]>bounds[3] or area[3]<bounds[1]):continue
        result[key]=record
    return result


def resolve_context(reference,catalog,location):
    sheet,separator,address=location.rpartition('!')
    _require(separator,'说明依据缺少实际所属单元格')
    row,column=coordinate_to_tuple(address)
    allowed=applicable_context(catalog,sheet,[row,column,row,column])
    _require(reference in allowed,'说明依据不存在、已改变或不适用于此格：'+reference)
    return deepcopy(allowed[reference])


def append_context_sheet(book,context,source_path):
    catalog=context_catalog(context,source_path)
    if not catalog:return
    sheet=book.create_sheet('语义说明依据')
    sheet.append(['依据引用','来源','适用工作表','适用范围','说明原文'])
    for reference,row in catalog.items():
        scope=row['scope'] or {}
        sheet.append([reference,'用户提供的业务说明' if row['kind']=='user_context' else '文档说明',
                      scope.get('sheet','全部'),json.dumps(scope.get('bounds','全部'),ensure_ascii=False),
                      row['content'] if isinstance(row['content'],str) else json.dumps(row['content'],ensure_ascii=False)])
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for column,width in {'A':55,'B':24,'C':25,'D':24,'E':100}.items():sheet.column_dimensions[column].width=width
    for row in sheet:
        for cell in row:
            if isinstance(cell.value,str):cell.data_type='s'
