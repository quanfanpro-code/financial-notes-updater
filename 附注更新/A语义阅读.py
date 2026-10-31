# coding: utf-8
"""模型按完整原文逐格识别金标准身份，程序保存各格结果。"""
from copy import deepcopy
from collections import Counter, deque
from pathlib import Path
import json
import re
import time

from openpyxl.utils.cell import range_boundaries, get_column_letter
from openpyxl.formula.tokenizer import TokenizerError
from .统一语义 import _require, definition_hash, validate_facts, required_cell_dimensions, _dimension_value, json_dates
from .统一证据 import context_catalog, applicable_context
from .表格 import read_workbook, materialize_candidate_blanks, file_hash, Tokenizer, _REFERENCE
from .表格编排 import _stage_receipt






# 用户指定的A/B格子对应不以整份报表口径为前提，标准原文件保持不变。



def addresses(area):
    area=area.replace('、',',').replace('，',',')
    if ',' in area:
        return [address for part in area.split(',') for address in addresses(part.strip())]
    left, top, right, bottom = range_boundaries(area)
    _require(all(type(n) is int for n in (left,top,right,bottom)) and 1<=left<=right and 1<=top<=bottom,
             '范围须为有效的单格或矩形')
    _require((right-left+1)*(bottom-top+1)<=500000, '范围过大')
    return [f'{get_column_letter(c)}{r}' for r in range(top,bottom+1) for c in range(left,right+1)]


def data_cells(table, cells):
    selected=[]
    for area in table['data_ranges']:
        for address in addresses(area):
            if address in cells:
                _require(address not in selected, '财务范围重复')
                selected.append(address)
    _require(set(selected)<=set(addresses(table['range'])), '财务范围越出本表')
    return selected


def expand_scopes(answer, table, cells, standard):
    """把模型明确的指标及各项取值与范围组合，不推断业务含义。"""
    owned=set(data_cells(table,cells));result={a:{'dimensions':{}} for a in owned}
    def scope(ranges):
        selected=set()
        for area in ranges:
            current=owned if area=='*' else set(addresses(area))
            _require(current<=set(addresses(table['range'])),'取值范围越出本表')
            current &= owned
            _require(current and not current&selected,'取值范围没有财务目标或重复')
            selected.update(current)
        return selected
    for kind,key,category in (('metrics','metric_id','metric'),('dimension_bindings','dimension_id','dimension')):
        for row in answer.get(kind,[]):
            _require(row['id'] in standard[category],'不存在的金标准编号')
            for address in scope(row['ranges']):
                other='dimension_id' if key=='metric_id' else 'metric_id'
                _require(not result[address].get(other) and result[address].get(key,row['id'])==row['id'],'同一数据位置有两个金标准编号')
                result[address][key]=row['id']
    for row in answer['dimensions']:
        key=row['id'];definition=standard['dimension'].get(key)
        refs=[row['evidence']] if isinstance(row['evidence'],str) else row['evidence']
        _require(definition is not None,'不存在的维度')
        value=deepcopy(row['value'])
        if not standard.get('_complete_identity') and definition['value_type']=='selection' and isinstance(value,dict):
            value.setdefault('domain',definition['domain'])
            value.setdefault('completeness','unknown' if value.get('mode')=='unknown' else 'complete')
        for address in scope(row['ranges']):
            previous=result[address]['dimensions'].get(key)
            _require(previous is None or previous['value']==value,'同一数据位置的维度取值冲突')
            result[address]['dimensions'][key]={'value':value,'evidence':list(dict.fromkeys((previous['evidence'] if previous else [])+refs))}
            if 'label_references' in row or previous and 'label_references' in previous:
                result[address]['dimensions'][key]['label_references']=list(dict.fromkeys(
                    (previous.get('label_references',[]) if previous else [])+row.get('label_references',[])))
    for row in answer.get('unresolved',[]):
        for address in scope([row['range']]):result[address].update(metric_id=None,reason=row['reason'])
    return result


def formula_source_material(snapshot, selected):
    """追溯静态引用及其行列文字；不求值，也不据公式直接发证。"""
    compact=snapshot.get('semantic_formula_sources',3)>=3
    direct=snapshot.get('semantic_formula_sources',3)==4
    full=snapshot.get('semantic_formula_sources',3)>=5
    sheets={s['name']:s for s in snapshot['sheets']};queue=deque();seen=set()
    for sheet in snapshot['sheets']:
        queue.extend((sheet['name'],a) for a,c in sheet['cells'].items()
                     if a in selected.get(sheet['name'],set()) and c.get('formula'))
    result={'引用':{},'原文':{},'未展开':[]};indexes={}
    if compact:result.update(引用范围={},共同行列文字={},合并标题范围={})
    if direct:
        result['未递归数值计算']=[]
        result['说明']='保留所选格的公式、直接引用及行列文字；文字标题公式继续追溯。未递归数值计算中每个格的原式见原文同名项，不展开其下层、不计算金额。原文中的空对象仍是已读数值或空白位置，数值及空白状态不用于判断身份。'
        result['行列文字规则']='每格按格址读取共同行列文字：同一行、本列上方非公式文字，以及合并标题范围中覆盖本行且位于本列或左侧的标题、本列上方跨列标题及这些标题所在行的非公式文字。各范围由该格行列及合并标题范围确定，不把其他行或其他列的数字当作标签。'
    def context_cells(name,address):
        sheet=sheets[name];cell=sheet['cells'][address];row,col=cell['row'],cell['column']
        if name not in indexes:
            rows={};columns={}
            for a,c in sheet['cells'].items():
                if isinstance(c.get('value'),str) or c.get('formula'):
                    rows.setdefault(c['row'],[]).append(a);columns.setdefault(c['column'],[]).append(a)
            indexes[name]=(rows,columns)
        rows,columns=indexes[name]
        owned=set(rows.get(row,[]));merged_anchors=set()
        owned.update(a for a in columns.get(col,[]) if sheet['cells'][a]['row']<row and
                     (not sheet['cells'][a].get('formula') or isinstance(sheet['cells'][a].get('cached_value'),str)))
        if compact:result['合并标题范围'].setdefault(name,sheet['merges'])
        for area in sheet['merges']:
            left,top,right,bottom=range_boundaries(area)
            if (top<=row<=bottom and left<=col) or (left<=col<=right and bottom<row):
                anchor=f'{get_column_letter(left)}{top}';owned.add(anchor);merged_anchors.add(anchor)
                owned.update(a for a in rows.get(top,[]) if not sheet['cells'][a].get('formula'))
        text={}
        for a in sorted(owned):
            c=sheet['cells'].get(a,{})
            if c.get('formula'):
                if a!=address and (a in merged_anchors or (c['row']==row and c['column']<col) or isinstance(c.get('cached_value'),str)):
                    text[name+'!'+a]={'公式':c['formula']}
                    if not direct or isinstance(c.get('cached_value'),str) and (a in merged_anchors or c['row']==row):queue.append((name,a))
            elif isinstance(c.get('value'),str) and c['value'].strip():text[name+'!'+a]=c['value']
        # 共用文字按原位置只保存一次；每格保留适用范围，避免公式链平方复制全文。
        if not compact:return text
        result['共同行列文字'].update(text)
        scope={'同行':f'{name}!{row}:{row}',
               '合并标题范围引用':name}
        if row>1:scope['同列上方']=f'{name}!{get_column_letter(col)}1:{get_column_letter(col)}{row-1}'
        return scope
    while queue:
        name,address=queue.popleft();location=name+'!'+address
        if location in seen:continue
        seen.add(location);cell=sheets[name]['cells'].get(address,{})
        if not cell.get('formula'):continue
        try:tokens=Tokenizer(cell['formula']).items
        except (TokenizerError,ValueError,IndexError):
            result['未展开'].append({'位置':location,'引用':cell['formula']});continue
        if full and any(t.type=='FUNC' and t.subtype=='OPEN' and t.value.upper() in {'INDIRECT(','OFFSET('} for t in tokens):
            result['未展开'].append({'位置':location,'引用':cell['formula']});continue
        refs=[]
        for token in tokens:
            if token.type!='OPERAND' or token.subtype!='RANGE':continue
            match=_REFERENCE.fullmatch(token.value)
            target=(match['sheet'] or name).strip("'").replace("''","'") if match else None
            if not match or target not in sheets:
                result['未展开'].append({'位置':location,'引用':token.value});continue
            area=match['a']+(':'+match['b'] if match['b'] else '')
            left,top,right,bottom=range_boundaries(area.replace('$','').upper())
            # 只遍历真实读取网格，整列或越界引用不凭空生成单元格。
            cells=sheets[target]['cells'];single=match['a'].replace('$','').upper()
            sources=([(single,cells[single])] if single in cells else []) if not match['b'] else [
                (a,c) for a,c in cells.items() if left<=c['column']<=right and top<=c['row']<=bottom]
            if not sources:
                result['未展开'].append({'位置':location,'引用':token.value});continue
            if compact and match['b']:
                scope_ref=target+'!'+area.replace('$','').upper()
                result['引用范围'].setdefault(scope_ref,[target+'!'+a for a,c in sources])
                refs.append({'范围引用':scope_ref})
            for a,c in sources:
                ref=target+'!'+a
                if not compact or not match['b']:refs.append(ref)
                if ref not in result['原文']:
                    row_context=context_cells(target,a)
                    record={} if direct else {'行列文字':row_context}
                    if c.get('formula'):record['公式']=c['formula']
                    elif isinstance(c.get('value'),str):
                        if compact:record['文字引用']=ref
                        else:record['文字']=c['value']
                    elif not direct:record['数据位置']='只确认业务含义，不使用数值及空白状态'
                    result['原文'][ref]=record
                    if direct and c.get('formula') and not isinstance(c.get('cached_value'),str):
                        result['未递归数值计算'].append(ref)
                if not direct or isinstance(c.get('cached_value'),str):queue.append((target,a))
        result['引用'][location]=list({json.dumps(ref,ensure_ascii=False):ref for ref in refs}.values())
    return result


def validate_material_references(source):
    """只检查已传材料中的引用闭合；不能据此宣称财务上下文已经完整。"""
    proof=source.get('公式引用依据',{})
    _require(not proof.get('未展开'),'公式引用未展开，停止本次识别请求：'+json.dumps(proof.get('未展开'),ensure_ascii=False))
    _require(not proof.get('未递归数值计算'),'公式下层引用未展开，不能提交历史残缺材料')
    references=proof.get('引用',{});original=proof.get('原文',{});shared=proof.get('共同行列文字',{})
    formulas=set(source.get('公式原文',{}))|{k for k,v in original.items() if v.get('公式')}
    formulas.update(k for k,v in shared.items() if isinstance(v,dict) and v.get('公式'))
    _require(formulas<=set(references),'公式缺少引用展开：'+str(sorted(formulas-set(references))))
    for items in references.values():
        for ref in items:
            if isinstance(ref,dict):
                key=ref['范围引用']
                _require(key in proof.get('引用范围',{}),'公式缺少引用范围内容：'+key)
                targets=proof['引用范围'][key]
            else:targets=[ref]
            _require(set(targets)<=set(original),'公式缺少引用原文：'+str(sorted(set(targets)-set(original))))
    for row in original.values():
        if '文字引用' in row:_require(row['文字引用'] in shared,'公式缺少共同原文：'+row['文字引用'])
    def check(value):
        if isinstance(value,dict):
            if '原文引用' in value:_require(value['原文引用'] in source.get('共同原文',{}),'缺少共同原文：'+str(value['原文引用']))
            for child in value.values():check(child)
        elif isinstance(value,list):
            for child in value:check(child)
    check(source.get('context',{}))


def source_material(snapshot, context, tables=None, semantic_only=False):
    """按原行列呈现，不抹掉空白、合并关系或适用说明。"""
    catalog=context_catalog(context,snapshot['path'])
    aliases={f'@C{i}':ref for i,ref in enumerate(catalog)}
    selected={s['name']:set(s['cells']) for s in snapshot['sheets']}
    if tables is not None:
        selected={}
        for table in tables:
            selected.setdefault(table['sheet'],set()).update(addresses(table['range']))
            if table.get('read_full_rows'):
                first,last=range_boundaries(table['range'])[1::2]
                source=next(s for s in snapshot['sheets'] if s['name']==table['sheet'])
                selected[table['sheet']].update(address for address,cell in source['cells'].items() if first<=cell['row']<=last)
            for area in table.get('context_ranges',[]):selected[table['sheet']].update(addresses(area))
            if snapshot.get('semantic_identity_version')==2:
                for ref in table.get('context_references',[]):selected.setdefault(ref['sheet'],set()).update(addresses(ref['range']))
    sheets=[]
    targets={}
    if semantic_only:
        for table in tables or []:targets.setdefault(table['sheet'],set()).update(a for area in table['data_ranges'] for a in addresses(area))
    for sheet in snapshot['sheets']:
        owned=selected.get(sheet['name'],set());lines={}
        for address,cell in sheet['cells'].items():
            if address not in owned:continue
            value=cell.get('cached_value') if cell.get('formula') else cell['value']
            if address in targets.get(sheet['name'],set()) and (value in (None,'') or isinstance(value,(int,float)) or cell.get('formula')):
                value='[待确认业务含义]'
            if isinstance(value,str) and ('\n' in value or '\r' in value):
                value=json.dumps(value,ensure_ascii=False)
            lines.setdefault(cell['row'],[]).append(f'{get_column_letter(cell["column"])}={"[空白]" if value is None or value=="" else value}')
        if lines:sheets.append({'sheet':sheet['name'],'rows':'\n'.join(str(r)+': '+' | '.join(parts) for r,parts in sorted(lines.items())),
                                'merges':[a for a in sheet['merges'] if set(addresses(a))&owned]})
    context_rows={}
    for alias,ref in aliases.items():
        record=catalog[ref];scope=record['scope']
        if snapshot.get('semantic_identity_version')!=2 and tables is not None and scope and not any(t['sheet']==scope['sheet'] and
            ref in applicable_context(catalog,t['sheet'],[range_boundaries(t['range'])[1],range_boundaries(t['range'])[0],range_boundaries(t['range'])[3],range_boundaries(t['range'])[2]]) for t in tables):continue
        content=record['content']
        if isinstance(content,dict):content={k:content[k] for k in ('title','chapter_context','unit_text','text') if content.get(k)}
        context_rows[alias]={'scope':scope,'text':content}
    texts=[]
    def collect(value):
        if isinstance(value,str) and len(value)>40:texts.append(value)
        elif isinstance(value,dict):
            for item in value.values():collect(item)
        elif isinstance(value,list):
            for item in value:collect(item)
    collect(context_rows)
    common={text:'T'+str(i) for i,(text,count) in enumerate(Counter(texts).items()) if count>1}
    def compact(value):
        if isinstance(value,str) and value in common:return {'原文引用':common[value]}
        if isinstance(value,dict):return {k:compact(v) for k,v in value.items()}
        if isinstance(value,list):return [compact(v) for v in value]
        return value
    material={'sheets':sheets,'context':compact(context_rows)}
    if tables is not None:material['本次读取范围']=deepcopy(tables)
    material['原有表区']={name:areas for name,areas in snapshot.get('names',{}).items()
        if any(area['range'] and area['sheet'] in selected and
               set(addresses(area['range']))&selected[area['sheet']] for area in areas)}
    if snapshot.get('semantic_formula_sources'):
        material['公式原文']={sheet['name']+'!'+address:cell['formula'] for sheet in snapshot['sheets']
            for address,cell in sheet['cells'].items() if address in selected.get(sheet['name'],set()) and cell.get('formula')}
        material['公式用途']='仅用于理解构成和引用来源。与文字说明不一致时核对原文及直接来源，保留冲突依据；不计算新金额，不按数值相等判断身份。'
        if snapshot['semantic_formula_sources']>=2 and material['公式原文']:
            formula_selected=selected
            if snapshot['semantic_formula_sources']==4 and any(table.get('data_ranges') for table in tables or []):
                formula_selected={}
                for table in tables:
                    formula_selected.setdefault(table['sheet'],set()).update(a for area in table.get('data_ranges',[]) for a in addresses(area))
                for sheet in snapshot['sheets']:
                    formula_selected.setdefault(sheet['name'],set()).update(a for a,c in sheet['cells'].items()
                        if a in selected.get(sheet['name'],set()) and c.get('formula') and c.get('cached_value') is None)
            material['公式引用依据']=formula_source_material(snapshot,formula_selected)
    if semantic_only:material['数值处理']='仅身份确认阶段隐藏已确定数据范围内的数值及空白状态；全部位置都要确认含义，原件及最终保存的原值不变。'
    if common:material['共同原文']={key:text for text,key in common.items()}
    def text_parts(value):
        if isinstance(value,str):yield value
        elif isinstance(value,dict):
            for child in value.values():yield from text_parts(child)
        elif isinstance(value,list):
            for child in value:yield from text_parts(child)
    aliases['_context_sources']=[{'text':text,'ref':aliases[key],'scope':record['scope']}
        for key,record in context_rows.items() for text in text_parts(record['text'])]
    def contains(value,text):
        if isinstance(value,str):return value==text
        if isinstance(value,dict):return any(contains(v,text) for v in value.values())
        if isinstance(value,list):return any(contains(v,text) for v in value)
        return False
    for text,key in common.items():
        aliases[key]=[{'ref':aliases[alias],'scope':record['scope']} for alias,record in context_rows.items() if contains(record['text'],text)]
    # 允许模型引用表内唯一原文；这里只定位逐字一致的文字，不推断其财务含义。
    for table in tables or []:
        cells=next(s for s in snapshot['sheets'] if s['name']==table['sheet'])['cells']
        literals={}
        for address in dict.fromkeys(a for area in [table['range'],*table.get('context_ranges',[])] for a in addresses(area)):
            cell=cells.get(address,{})
            value=cell.get('cached_value') if cell.get('formula') else cell.get('value')
            if isinstance(value,str) and value.strip() and ('semantic_candidates' in snapshot or not cell.get('formula')):
                literals.setdefault(value,[]).append(address)
        left,top,right,bottom=range_boundaries(table['range'])
        for value,refs in literals.items():
            if re.fullmatch(r'\$?[A-Z]+\$?[1-9]\d*',value) or value.startswith(('@','T')):continue
            row_counts=Counter(cells[ref]['row'] for ref in refs)
            for ref in refs:
                row=cells[ref]['row']
                if row_counts[row]!=1 and len(row_counts)!=1:continue
                # 同一行重复的原文保留全部出处，不猜其中哪个才是模型所指。
                bounds=[top,left,bottom,right] if len(row_counts)==1 else [row,left,row,right]
                aliases.setdefault(value,[]).append({'ref':ref,'scope':{'sheet':table['sheet'],'bounds':bounds}})
    if snapshot.get('semantic_identity_version')==2:
        shown={name+'!'+a for name,items in selected.items() for a in items}
        shown.update(material.get('公式引用依据',{}).get('原文',{}))
        shown.update(material.get('公式引用依据',{}).get('共同行列文字',{}))
        evidence={s['name']+'!'+a:deepcopy(c) for s in snapshot['sheets'] for a,c in s['cells'].items() if s['name']+'!'+a in shown}
        aliases['_all_evidence']=evidence
        # 同文仅保留全部原位置，不以行列距离推定它的业务管辖范围。
        for key in list(aliases):
            if isinstance(aliases[key],list) and not key.startswith('_') and not re.fullmatch(r'T\d+',key):aliases.pop(key)
        for ref,c in evidence.items():
            value=c.get('cached_value') if c.get('formula') else c.get('value')
            if isinstance(value,str) and value.strip():aliases.setdefault(value,[]).append({'ref':ref,'scope':None})
    aliases['_material_positions']=selected
    return material,aliases


def validate_reading(answer, snapshot):
    _require(isinstance(answer.get('tables'),list) and answer['tables'],'没有返回子表')
    sheets={s['name']:s for s in snapshot['sheets']};seen=set();covered={}
    for table in answer['tables']:
        _require(table['id'] not in seen and table['sheet'] in sheets and table['meaning'] and table['subject'],'子表缺少唯一编号或含义')
        seen.add(table['id']);sheet=sheets[table['sheet']]
        area=set(addresses(table['range']));left,top,right,bottom=range_boundaries(table['range'])
        _require(right<=sheet['max_column'] and bottom<=sheet['max_row'],'子表越出原文件')
        _require(not area&covered.setdefault(table['sheet'],set()),'子表范围重叠')
        covered[table['sheet']].update(area)
        for target in table['data_ranges']:_require(set(addresses(target))<=area,'数据范围越出子表：'+table['sheet']+'!'+target+' 不在 '+table['range'])
        data_cells(table,sheet['cells'])
        if snapshot.get('semantic_identity_version')==2:
            for ref in table.get('context_references',[]):
                _require(ref['sheet'] in sheets,'共同原文引用工作表不存在')
                left,top,right,bottom=range_boundaries(ref['range'])
                _require(top>0 and left>0 and right<=sheets[ref['sheet']]['max_column'] and bottom<=sheets[ref['sheet']]['max_row'],'共同原文引用越出真实网格')
    return deepcopy(answer)


def reading_questions(reading, snapshot):
    questions=[]
    for table in reading['tables']:
        packet=table_packet(table,snapshot,[]);owned=set(data_cells(table,packet['cells']))
        outside=[a for a,c in packet['cells'].items() if a not in owned and type(c['value']) in (int,float)
                 and ('semantic_candidates' not in snapshot or table['sheet']+'!'+a in snapshot['semantic_candidates'])]
        text=[a for a in owned if isinstance(packet['cells'][a]['value'],str) and
              packet['cells'][a]['value'].strip() not in {'','-','—','–','/'} and not packet['cells'][a].get('formula')]
        if outside or text:questions.append({'id':table['id'],'数据范围外的数值':outside,'须分清业务事实与身份标签的文字':sorted(text)})
    return questions


def reading_check_payload(reading, questions, snapshot, context):
    ids={q['id'] for q in questions};tables=[t for t in reading['tables'] if t['id'] in ids]
    source,_=source_material(snapshot,context,tables)
    return {'task':'a_semantic_read_check','reading':tables,'questions':questions,'source':source}


def apply_reading_check(reading, answer, questions, snapshot):
    wanted={q['id'] for q in questions};rows=answer['tables']
    _require(len(rows)==len(wanted) and {t['id'] for t in rows}==wanted,'局部阅读复查漏表或重复')
    index={t['id']:t for t in reading['tables']}
    for row in rows:
        original=index[row['id']]
        _require(all(row[k]==original[k] for k in ('sheet','range')),'局部复查改变了整份表格编排')
        index[row['id']]=row
    return validate_reading({**reading,'tables':[index[t['id']] for t in reading['tables']]},snapshot)


def metric_catalog(standard, include_basis=True):
    return [{'id':m['id'],'meaning':m['definition']['economic_object']+'：'+m['definition']['measure'],
             'time_kind':m['definition']['time_kind'],
             **({'basis':m['definition']['measurement_basis'].split('。',1)[0]} if include_basis else {})}
            for m in standard['metric'].values() if m['status']=='active']


def validate_selection(answer, reading, standard):
    wanted={t['id'] for t in reading['tables']};rows=answer['tables']
    _require(len(rows)==len(wanted) and {t['id'] for t in rows}==wanted,'标准检索遗漏或重复子表')
    for row in rows:
        _require(isinstance(row['metric_ids'],list),'指标候选须为列表')
        _require(all(m in standard['metric'] and standard['metric'][m]['status']=='active' for m in row['metric_ids']),'检索引用不存在的标准')
    return deepcopy(answer)


def selection_payload(reading,snapshot,standard,include_basis=True):
    labels=[]
    for table in reading['tables']:
        packet=table_packet(table,snapshot,[])
        labels.append({'id':table['id'],'sheet':table['sheet'],'text':{a:value for a,c in packet['cells'].items()
            if isinstance(value:=c.get('cached_value') if c.get('formula') else c.get('value'),str) and value.strip()
            and ('semantic_candidates' in snapshot or not c.get('formula'))}})
    return {'task':'a_semantic_select','reading':reading,'catalog':metric_catalog(standard,include_basis),'原表文字':labels}


def standard_material(ids, standard):
    if not standard.get('_historical_material'):
        # 带上标准明文引用的其他定义；是否适用仍由本次识别决定。
        selected=list(dict.fromkeys(ids))
        for key in selected:
            for side in standard['metric'][key]['definition'].get('ratio',{}).values():
                if side['metric_id'] not in selected:selected.append(side['metric_id'])
            for reference in re.findall(r'(?<![A-Za-z0-9_.])metric\.[A-Za-z0-9_.]+',json.dumps(standard['metric'][key]['definition'],ensure_ascii=False)):
                reference=reference.rstrip('.')
                if reference not in selected and standard['metric'].get(reference,{}).get('status')=='active':
                    selected.append(reference)
        metrics=[deepcopy(standard['metric'][key]) for key in selected]
        dimensions={d['id'] for m in metrics for d in m['dimension_refs']}
        return {'metrics':metrics,'dimensions':[deepcopy(standard['dimension'][key]) for key in sorted(dimensions)],
                '共同识别原则':'范围以指标经济对象、计量基础和维度职责为总体。原文覆盖完整业务总体且无附加筛选时用all，不要求列出所有类别证明全体；不能把合计改成当前展示的明细名单。存在父项限定时保留父项，明确账龄或类别本身已是确定范围。保留其中项、交集、并集及排除条件，不添加原文没有的限制；不凭金额、当前行数、位置或常见做法推断。只记已知含义，真正缺少的条件才列未决。'}
    # 仅供旧回执审计重放；新请求不使用历史删段材料。
    metrics=[standard['metric'][m] for m in dict.fromkeys(ids)]
    dimensions={d['id'] for m in metrics for d in m['dimension_refs']}
    repeated=Counter(part for d in standard['dimension'].values() for part in d['meaning'].split(' ') if len(part)>100)
    definitions=[]
    for key in sorted(dimensions):
        original=standard['dimension'][key]
        row={k:v for k,v in original.items() if k in {'id','value_type','domain','allowed_values'}}
        # 各维度重复的通用识别说明统一表达，指标定义和该维度独有的业务职责保持原文。
        row['meaning']=' '.join(p for p in original['meaning'].split(' ') if repeated[p]<10)
        definitions.append(row)
    return {'metrics':[{'id':m['id'],'definition':{k:v for k,v in m['definition'].items() if k!='ratio'},
                        'dimension_refs':m['dimension_refs']} for m in metrics], 'dimensions':definitions,
            '共同识别原则':'范围以指标经济对象、计量基础和维度职责为总体。原文覆盖完整业务总体且无附加筛选时用all，不要求列出所有类别证明全体；不能把合计改成当前展示的明细名单。存在父项限定时保留父项，明确账龄或类别本身已是确定范围。保留其中项、交集、并集及排除条件，不添加原文没有的限制；不凭金额、当前行数、位置或常见做法推断。只记已知含义，真正缺少的条件才列未决。'}


def cell_requirements(ids,standard):
    metrics={key:standard['metric'][key] for key in dict.fromkeys(ids)}
    if standard.get('_complete_identity'):
        return {'complete_fact_dimensions':{key:sorted(required_cell_dimensions(m,standard)) for key,m in metrics.items()},
                'missing_values':'全部必需维度照实列缺项，保留已确认槽位；本次比较排除另存，不补取值。',
                'ratio_definitions':{key:deepcopy(m['definition']['ratio']) for key,m in metrics.items() if 'ratio' in m['definition']}}
    return {'complete_fact_dimensions':{key:sorted(required_cell_dimensions(m,standard)) for key,m in metrics.items()},
            'report_scope':{'required_for_cell':False,'used_for_matching':False},
            'missing_values':'实际取值未提供不否定已确认槽位；省略无据取值，不填unknown字符串或假成员。',
            'ratio_definitions':{key:deepcopy(m['definition']['ratio']) for key,m in metrics.items() if 'ratio' in m['definition']}}


def checked_records(answer, table, packet, standard, aliases):
    from .统一识别 import validate_response
    expanded=expand_scopes(answer,table,packet['cells'],standard);records={}
    label_refs=[a for a,c in packet['cells'].items() if isinstance(c['value'],str) and c['value'].strip() and not c.get('formula')]
    _require(label_refs,'表格缺少文字依据')
    for address in packet['candidate_cells']:
        if address not in expanded:
            records[address]={'kind':'non_business','category':'label','reason':'模型整体阅读明确的数据范围以外的标题、表头或标签', 'evidence_cells':label_refs[:1]}
            continue
        row=expanded[address];metric=row.get('metric_id')
        if metric is None:
            records[address]={'kind':'unresolved','reason':row.get('reason','缺少对应标准'),'evidence_cells':label_refs[:1]};continue
        _require(metric in standard['metric'], '未找到指标：'+str(metric))
        allowed={d['id'] for d in standard['metric'][metric]['dimension_refs']}
        required=required_cell_dimensions(standard['metric'][metric],standard)
        _require(set(row.get('dimensions',{}))<=set(standard['dimension']), '模型使用了不存在的维度ID')
        if standard.get('_complete_identity'):
            _require(set(row.get('dimensions',{}))<=allowed,'已声明维度不属于该指标；原答复保留，须明确修正，不能静默删除')
        dims={k:v['value'] for k,v in row.get('dimensions',{}).items() if k in allowed}
        try:
            evidence={k:normalize_evidence(v['evidence'],table['sheet'],aliases,address) for k,v in row.get('dimensions',{}).items() if k in allowed}
        except (ValueError,KeyError,TypeError,IndexError) as error:
            records[address]={'kind':'unresolved','reason':'本格原文依据不能解析：'+str(error),'evidence_cells':label_refs[:1]}
            continue
        record={'kind':'metric_fact','metric_id':metric,'dimensions':dims,'dimension_evidence':evidence,'reason':row.get('reason',table['meaning'])}
        if required-set(dims) or any(isinstance(v,dict) and v.get('completeness','complete')!='complete' for v in dims.values()):
            record.update(kind='partial_metric_fact',missing_dimensions=sorted(required-set(dims)),evidence_cells=label_refs[:1])
        records[address]=record
    records=validate_response({'cells':records},packet,standard,isolate_errors=True)
    return apply_role_corrections(records,{'corrections':{'non_business':answer.get('non_business',[])}},table,packet,standard,aliases)


def normalize_evidence(refs, sheet, aliases, address=None):
    if isinstance(refs,str):refs=[refs]
    result=[]
    for ref in refs:
        if '_all_evidence' in aliases and not ref.startswith('@') and ref not in aliases:
            prefix,separator,area=ref.rpartition('!')
            target=prefix.strip("'") if separator else sheet
            try:locations=[target+'!'+a for a in addresses((area if separator else ref).replace('$',''))]
            except (ValueError,TypeError):locations=[]
            if locations:
                _require(set(locations)<=set(aliases['_all_evidence']),'依据引用未在本次完整材料中提供：'+ref)
                result.extend(locations);continue
        literal=ref.removeprefix('原文引用：').removeprefix('原文引用:').strip()
        if re.fullmatch(r'@T\d+',literal) and literal[1:] in aliases:literal=literal[1:]
        if literal in aliases:ref=literal
        items=aliases.get(ref)
        if items is None and literal and not re.fullmatch(r'\$?[A-Z]+\$?[1-9]\d*(?::\$?[A-Z]+\$?[1-9]\d*)?',literal) and not literal.startswith('@'):
            items=[item for item in aliases.get('_context_sources',[]) if literal in item['text']]
            items.extend(item for text,refs in aliases.items() if isinstance(refs,list) and not text.startswith(('@','_'))
                and not re.fullmatch(r'T\d+',text) and literal in text for item in refs)
        if isinstance(items,list) and items:
            _require(address is not None,'共同原文引用需要明确适用位置')
            column,row,_,_=range_boundaries(address)
            matching=[]
            for item in items:
                scope=item['scope']
                if '_all_evidence' in aliases or not scope or (scope['sheet']==sheet and scope['bounds'][0]<=row<=scope['bounds'][2] and scope['bounds'][1]<=column<=scope['bounds'][3]):matching.append(item['ref'])
            _require(matching,'共同原文说明不适用于当前位置：'+ref)
            result.extend(matching);continue
        if ref.startswith('@'):result.append(aliases.get(ref,ref));continue
        if '!' in ref:
            prefix,ref=ref.rsplit('!',1)
            _require(prefix.strip("'")==sheet,'依据引用了其他工作表')
        # 证据地址可带括注；只读取明确的地址，括注不作为财务判断或原文证据。
        annotated=re.fullmatch(r'(\$?[A-Z]+\$?[1-9]\d*(?::\$?[A-Z]+\$?[1-9]\d*)?)(?:[（(][^\r\n]*[）)]|[=：].+)',ref)
        if annotated:ref=annotated.group(1)
        result.extend(addresses(ref.replace('$','')))
    return list(dict.fromkeys(result))


def table_packet(table, snapshot, context):
    sheet=next(s for s in snapshot['sheets'] if s['name']==table['sheet'])
    area=set(addresses(table['range']))
    shown=area|{a for span in table.get('context_ranges',[]) for a in addresses(span)}
    result={'source_path':snapshot['path'],'source_hash':snapshot['sha256'],'sheet':sheet['name'],
            'cells':{a:c for a,c in sheet['cells'].items() if a in shown},
            'candidate_cells':[a for a in sheet['cells'] if a in area and
                ('semantic_candidates' not in snapshot or sheet['name']+'!'+a in snapshot['semantic_candidates'])], 'document_context':context}
    if snapshot.get('semantic_identity_version')==2:
        result['source_evidence']=snapshot.get('_source_evidence',{})
        result['identity_collection_version']=1
    return result


def binding_payload(tables, selection, snapshot, context, standard, semantic_only=True, reading_notes=None):
    material,aliases=source_material(snapshot,context,tables,semantic_only=semantic_only)
    selected={r['id']:r for r in selection['tables']}
    ids=[m for t in tables for m in selected[t['id']]['metric_ids']]
    # 检索已读取完整行名，只提供实际选中的定义，不再次扩成整个同类指标。
    payload={'task':'a_semantic_bind','reading':tables,'source':material,'standards':standard_material(ids,standard),
             'cell_requirements':cell_requirements(ids,standard)}
    if standard.get('_complete_identity'):
        payload['identity_collection_version']=1
        snapshot['_source_evidence']=deepcopy(aliases['_all_evidence'])
        snapshot.setdefault('_all_source_evidence',{}).update(deepcopy(aliases['_all_evidence']))
    if snapshot.get('semantic_slot_bindings'):
        payload['槽位与取值分层']='先确认金标准槽位及维度定义，再记录原文实际取值。名称或金额未填写不否定已知槽位；缺少具体取值时保留已有维度，集合用unknown或partial如实记录，不填假名称，不改成all或合计，不放进unresolved删除已知指标。核查仍须确认指标及已知维度都有原文依据。'
    vocabulary={}
    for target in snapshot.get('semantic_known_targets',[]):
        metric=target.get('metric_id')
        if target.get('semantic_key') is None or metric not in ids:continue
        allowed={r['id'] for r in standard['metric'][metric]['dimension_refs']}
        for key,value in target.get('dimensions',{}).items():
            definition=standard['dimension'].get(key,{})
            if key not in allowed or key in {'entity','report_scope','period','currency','unit_scale'}:continue
            if definition.get('value_type') not in {'selection','text'} or definition.get('allowed_values'):continue
            item={'metric_id':metric,'dimension_id':key,'value':deepcopy(value)}
            vocabulary[definition_hash(item)]=item
    if vocabulary:payload['已确认维度取值']=list(vocabulary.values())
    locations={t['sheet']+'!'+a for t in tables for span in t['data_ranges'] for a in addresses(span)}
    templates=[r for r in snapshot.get('semantic_style_templates',[]) if r['location'] in locations and r['metric_id'] in ids]
    if templates:
        payload['style_templates']=deepcopy(templates)
        payload['style_reuse_instruction']='核对本表上下文和所列差异。一致部分用style_reuse:[{template_id,evidence:[本次原文位置],reason:一致依据}]沿用；不要重写这些格的指标和constant_dimensions。只在dimensions补当前主体、实际日期、币种、单位及context_bindings所需业务范围。其他格及不一致部分按现有metrics、dimensions识别。模板不是当前原文证据，禁止沿用历史年份或主体。'
    if reading_notes is not None:payload['reading_notes']=deepcopy(reading_notes)
    return payload,aliases


def source_identity_records(row, table, packet, standard, aliases):
    """保存模型明确的原文含义，不以标准必填项是否齐全判定原文身份。"""
    from .统一识别 import validate_identity_review
    result={};owned=set(data_cells(table,packet['cells']))
    for declaration in row.get('source_identities',[]):
        meaning=declaration.get('meaning');difference=declaration.get('standard_difference','')
        _require(isinstance(meaning,str) and meaning.strip(),'原文身份证缺少实际含义')
        _require(isinstance(difference,str),'标准差别说明须为文字')
        candidate=declaration.get('closest_metric_id')
        _require(candidate is None or candidate in standard['metric'] and standard['metric'][candidate]['status']=='active','最近标准不存在或未启用')
        selected=[]
        for span in declaration['ranges']:
            area=owned if span=='*' else set(addresses(span))
            _require(area<=set(addresses(table['range'])),'原文身份证范围越表')
            selected.extend(sorted(area&owned))
        _require(selected and len(selected)==len(set(selected)) and not set(selected)&set(result),'原文身份证没有目标或范围重复')
        for address in selected:
            evidence=normalize_evidence(declaration['evidence'],table['sheet'],aliases,address)
            _require(evidence and not set(evidence)&owned,'原文身份证须引用解释含义的文字，不能引用目标数据')
            validate_identity_review({'checked_cells':[address],'issues':[{'cells':[address],'reason':meaning,'evidence_cells':evidence}]},packet)
            result[address]={'source_reference':{'file_hash':packet['source_hash'],'location':table['sheet']+'!'+address},
                'table_id':table['id'],'meaning':meaning,'evidence_cells':evidence,
                'closest_metric_id':candidate,'standard_difference':difference}
    return result


def validate_bindings(answer, tables, snapshot, context, standard, candidate=False, require_identity=False):
    _require(isinstance(answer.get('tables'),list),'身份结果须按表返回')
    answer=deepcopy(answer)
    if require_identity:
        _require(not any(t.get('source_identities') for t in answer['tables']),
                 '识别成功必须完整对应金标准，请通过metrics和dimensions修正绑定')
        _require(all(m.get('id') in standard['metric'] for row in answer['tables'] for m in row.get('metrics',[])),
                 '指标ID须引用当前金标准，不能填写表号或自造编码；未查明的范围保留unresolved，不填无效指标')
    index={t['id']:t for t in tables}
    if any(row.get('style_reuse') for row in answer['tables']):
        from .B样式复用 import apply_reuse
        answer['tables']=[apply_reuse(row,index[row['id']],snapshot.get('semantic_style_templates',[])) for row in answer['tables']]
    for row in answer['tables']:
        if row.get('id') not in index:
            scopes=[a for m in row.get('metrics',[]) for a in m['ranges']]+[u['range'] for u in row.get('unresolved',[])]
            _require(scopes and '*' not in scopes,'无法从明确范围定位返回的原表')
            declared=set().union(*(set(addresses(a)) for a in scopes))
            matches=[t for t in tables if row.get('id')==t['sheet'] and declared<=set(addresses(t['range']))]
            _require(len(matches)==1,'返回表编号无唯一对应原表')
            row['id']=matches[0]['id']
        row.setdefault('dimensions',[])
        row.setdefault('metrics',[])
        for metric in row.get('metrics',[]):
            for dimension in metric.pop('dimensions',[]):
                _require('*' not in metric['ranges'] and '*' not in dimension['ranges'],'嵌套维度须声明明确范围')
                metric_scope=set().union(*(set(addresses(a)) for a in metric['ranges']))
                dimension_scope=set().union(*(set(addresses(a)) for a in dimension['ranges']))
                _require(dimension_scope<=metric_scope,'嵌套维度范围超出对应指标')
                row['dimensions'].append(dimension)
    _require(len(answer['tables'])==len(index) and {t['id'] for t in answer['tables']}==set(index),'身份结果漏表或重复')
    if any('source_identities' in row for row in answer['tables']):
        _,aliases=source_material(snapshot,context,tables)
        for row in answer['tables']:
            table=index[row['id']]
            source_identity_records(row,table,table_packet(table,snapshot,context),standard,aliases)
    if candidate:
        try:answer=distribute_shared_dimensions(answer,tables,snapshot)
        except (ValueError,KeyError,TypeError):return deepcopy(answer)
        return answer
    answer=distribute_shared_dimensions(answer,tables,snapshot)
    _,aliases=source_material(snapshot,context,tables)
    for row in answer['tables']:
        table=index[row['id']]
        checked_records(row,table,table_packet(table,snapshot,context),standard,aliases)
    return deepcopy(answer)


def distribute_shared_dimensions(answer,tables,snapshot):
    """只分配模型明确声明的跨表范围；星号仍只代表所在表。"""
    result=deepcopy(answer);index={t['id']:t for t in tables}
    owned={t['id']:set(data_cells(t,next(s for s in snapshot['sheets'] if s['name']==t['sheet'])['cells'])) for t in tables}
    rows={r['id']:r for r in result['tables']}
    for row in result['tables']:row['dimensions']=[]
    for row in answer['tables']:
        origin=index[row['id']]
        siblings=[t for t in tables if t['sheet']==origin['sheet']]
        permitted=set().union(*(set(addresses(t['range'])) for t in siblings))
        for dimension in row['dimensions']:
            scope=set().union(*(owned[row['id']] if area=='*' else set(addresses(area)) for area in dimension['ranges']))
            _require(scope<=permitted,'共同信息范围超出本次已明确的数据范围')
            _require(any(scope&owned[t['id']] for t in siblings),'共同信息没有财务目标')
            for target in siblings:
                selected=scope&owned[target['id']]
                if selected:
                    ranges=dimension['ranges'] if scope<=set(addresses(origin['range'])) else sorted(selected)
                    rows[target['id']]['dimensions'].append({**deepcopy(dimension),'ranges':ranges})
    return result


def review_for_table(row, table, packet, aliases):
    from .统一识别 import validate_identity_review
    checked=packet['candidate_cells']
    if 'checked_ranges' in row:
        area={a for span in row['checked_ranges'] for a in addresses(span)}
        _require(area<=set(addresses(table['range'])),'核查声明越出本表')
        checked=[a for a in checked if a in area]
    issues=[]
    for issue in row.get('issues',[]):
        area=set(addresses(issue['range']))
        _require(area<=set(addresses(table['range'])),'核查范围越出本表')
        cells=[a for a in packet['candidate_cells'] if a in area]
        _require(cells,'核查问题没有对应原格')
        issues.append({'cells':cells,'reason':issue['reason'],'evidence_cells':normalize_evidence(issue['evidence'],table['sheet'],aliases,cells[0])})
    return validate_identity_review({'checked_cells':checked,'issues':issues},packet)


def validate_review(answer, tables, snapshot, context, focus=None, require_identity=False):
    index={t['id']:t for t in tables}
    if answer and 'tables' not in answer and all(k in index and isinstance(v,dict) and v.get('id',k)==k for k,v in answer.items()):
        answer={'tables':[dict(v,id=k) for k,v in answer.items()]}
    _require(isinstance(answer.get('tables'),list),'核查须按表返回')
    if require_identity:
        _require(not any(row.get('corrections',{}).get('source_identities') for row in answer['tables'] if isinstance(row,dict)),
                 '识别成功必须完整对应金标准，核查不能改回文字身份')
    valid=[t for t in answer['tables'] if isinstance(t,dict) and isinstance(t.get('id'),str)]
    _require(valid,'没有可恢复的完整子表核查记录')
    incomplete=len(valid)!=len(answer['tables']) or {t['id'] for t in valid}!=set(index)
    answer={**deepcopy(answer),'tables':deepcopy(valid)}
    if incomplete:answer['incomplete']=True
    _require(len(answer['tables'])==len({t['id'] for t in answer['tables']}) and
             {t['id'] for t in answer['tables']}<=set(index),'核查表重复或越界')
    _,aliases=source_material(snapshot,context,tables)
    for row in answer['tables']:
        for item in row.get('corrections',{}).get('non_business',[]):
            _require(isinstance(item,dict) and {'range','category','reason','evidence'}<=set(item),
                     'non_business须明确range、category、reason和evidence，不能由程序猜范围或类别')
            _require(isinstance(item['range'],str) and bool(item['range'].strip()) and item['category'] in {'label','annotation'},
                     'non_business须用明确范围及label或annotation类别')
        for item in row.get('corrections',{}).get('remove_dimensions',[]):
            _require(isinstance(item,str) and bool(item.strip()) or
                     isinstance(item,dict) and set(item)=={'id','ranges'} and isinstance(item['id'],str)
                     and isinstance(item['ranges'],list) and bool(item['ranges']),
                     '维度移除须为维度ID或明确的{id,ranges}，不能猜范围')
        if focus is not None:
            wanted={a for span in focus[row['id']] for a in addresses(span)}
            if 'checked_ranges' in row:
                _require({a for span in row['checked_ranges'] for a in addresses(span)}==wanted,'核查回复范围与请求不符')
            row['checked_ranges']=deepcopy(focus[row['id']])
            corrections=row.get('corrections',{})
            _require(not corrections.get('dimensions'),'局部核查须用有范围的修正')
            for item in corrections.get('remove_dimensions',[]):
                _require(isinstance(item,dict),'局部移除须逐项提供{id,ranges}，不能移除整表维度')
                _require({a for span in item['ranges'] for a in addresses(span)}<=wanted,'维度移除越出本次核查范围')
            for key in ('metrics','patches','non_business','source_identities'):
                for change in corrections.get(key,[]):
                    spans=[change['range']] if key=='non_business' else change['ranges']
                    _require({a for span in spans for a in addresses(span)}<=wanted,'更正越出本次核查范围')
            _require({a for span in corrections.get('resolved_ranges',[]) for a in addresses(span)}<=wanted,'解除未决越出本次核查范围')
        corrections=row.get('corrections',{});patches=corrections.get('patches',[])
        metric_patches=[p for p in patches if (p.get('id')=='metric' and isinstance(p.get('value'),str) and p['value'].startswith('metric.'))
                        or (isinstance(p.get('id'),str) and p['id'].startswith('metric.') and 'value' not in p)]
        if metric_patches:
            metrics=corrections.get('metrics',[])+[{'ranges':p['ranges'],'id':p['value'] if p['id']=='metric' else p['id']} for p in metric_patches]
            corrections['metrics']=list({json.dumps(m,sort_keys=True):m for m in metrics}.values())
            corrections['patches']=[p for p in patches if p not in metric_patches]
        table=index[row['id']];review_for_table(row,table,table_packet(table,snapshot,context),aliases)
    return deepcopy(answer)


def validate_review_result(answer,binding,tables,snapshot,context,standard,payload):
    answer=validate_review(answer,tables,snapshot,context,payload.get('review_focus'),require_identity=True)
    _require(all(isinstance(row.get('issues'),list) for row in answer['tables']),'核查必须明确返回issues数组，只有表号不等于已经核查')
    merged=validate_bindings(reviewed_binding(binding,answer,tables,snapshot),tables,snapshot,context,standard)
    index={t['id']:t for t in tables};rows={r['id']:r for r in merged['tables']}
    _,aliases=source_material(snapshot,context,tables)
    if standard.get('_complete_identity'):
        from .A合集 import validate_relations
        validate_relations(binding,answer,tables,aliases)
    supplied={m['id'] for m in payload['standards']['metrics']}
    for review in answer['tables']:
        table=index[review['id']];packet=table_packet(table,snapshot,context)
        _require(all(m['id'] in supplied for m in review.get('corrections',{}).get('metrics',[])),
                 '指标更正须依据本次实际发送的完整定义，不能仅凭编码选标准')
        checked=review_for_table(review,table,packet,aliases)
        blocked={a for issue in checked['issues'] for a in issue['cells']}
        row=rows[review['id']]
        declared={a for issue in row.get('unresolved',[]) for a in addresses(issue['range'])}
        records=checked_records(row,table,packet,standard,aliases)
        for address in checked['checked_cells']:
            record=records[address]
            if address in blocked:continue
            if record['kind']=='invalid_response' or record['kind']=='unresolved' and address not in declared:
                raise ValueError('核查后仍未修正：'+review['id']+'!'+address+'；'+record['reason'])
    return deepcopy(answer)


def reviewed_binding(binding, review, tables=None, snapshot=None):
    """核查沿用原结果，只采纳有明确取值及范围的维度更正。"""
    result=deepcopy(binding);index={r['id']:r for r in result['tables']}
    for row in review.get('tables',[]):
        corrections=row.get('corrections',{})
        _require(set(corrections)<= {'dimensions','metrics','non_business','remove_dimensions','patches','resolved_ranges','source_identities','relations'},'核查更正字段不支持')
        replacements=corrections.get('dimensions',[])
        ids={d['id'] for d in replacements}|{d for d in corrections.get('remove_dimensions',[]) if isinstance(d,str)}
        target=index[row['id']]
        for change in corrections.get('source_identities',[]):
            table=next(t for t in tables if t['id']==row['id'])
            owned=set(data_cells(table,next(s for s in snapshot['sheets'] if s['name']==table['sheet'])['cells']))
            def declared(item):
                area=set().union(*(owned if span=='*' else set(addresses(span)) for span in item['ranges']))
                _require(area<=set(addresses(table['range'])),'原文身份更正越表')
                return area&owned
            changed=declared(change)
            _require(changed,'原文身份更正没有财务目标')
            kept=[{**old,'ranges':sorted(remaining)} for old in target.get('source_identities',[]) if (remaining:=declared(old)-changed)]
            target['source_identities']=kept+[deepcopy(change)]
        if corrections.get('non_business'):
            target.setdefault('non_business',[]).extend(deepcopy(corrections['non_business']))
        target['dimensions']=[d for d in target['dimensions'] if d['id'] not in ids]+deepcopy(replacements)
        for change in corrections.get('metrics',[]):
            table=next(t for t in tables if t['id']==row['id'])
            owned=set(data_cells(table,next(s for s in snapshot['sheets'] if s['name']==table['sheet'])['cells']))
            changed=set(a for area in change['ranges'] for a in addresses(area))
            _require(changed and changed<=set(addresses(table['range'])),'指标修正越出本表')
            changed &= owned
            _require(changed,'指标修正没有财务目标')
            kept=[]
            for original in target['metrics']:
                remaining=(set(a for area in original['ranges'] for a in addresses(area))&owned)-changed
                if remaining:kept.append({**original,'ranges':sorted(remaining)})
            target['metrics']=kept+[deepcopy(change)]
        changes=[(change,False) for change in corrections.get('patches',[])]+[
            (change,True) for change in corrections.get('remove_dimensions',[]) if isinstance(change,dict)]
        for change,remove in changes:
            table=next(t for t in tables if t['id']==row['id'])
            owned=set(data_cells(table,next(s for s in snapshot['sheets'] if s['name']==table['sheet'])['cells']))
            def scope(item):
                declared=set().union(*(owned if a=='*' else set(addresses(a)) for a in item['ranges']))
                _require(declared<=set(addresses(table['range'])),'局部修正越出本表')
                return declared&owned
            changed=scope(change)
            _require(changed,'局部修正没有财务目标')
            kept=[]
            for original in target['dimensions']:
                if original['id']!=change['id']:kept.append(original);continue
                remaining=scope(original)-changed
                if remaining:kept.append({**original,'ranges':sorted(remaining)})
            target['dimensions']=kept+([] if remove else [deepcopy(change)])
        resolved=set(a for area in corrections.get('resolved_ranges',[]) for a in addresses(area))
        if resolved:
            table=next(t for t in tables if t['id']==row['id'])
            owned=set(data_cells(table,next(s for s in snapshot['sheets'] if s['name']==table['sheet'])['cells']))
            _require(resolved<=set(addresses(table['range'])),'解除未决范围越出本表')
            resolved &= owned
            _require(resolved,'解除未决范围没有财务目标')
            target['unresolved']=[{**item,'range':a} for item in target.get('unresolved',[]) for a in addresses(item['range']) if a not in resolved]
    return result


def duplicate_identity_questions(binding,tables,snapshot,standard):
    """同列不同原文却有相同身份时请模型查明，不自动判错或猜测财务范围。"""
    if standard.get('_complete_identity'):return []
    questions=[];index={t['id']:t for t in tables}
    for row in binding['tables']:
        table=index[row['id']];cells=table_packet(table,snapshot,[])['cells']
        try:expanded=expand_scopes(row,table,cells,standard)
        except (ValueError,KeyError,TypeError):continue
        source_only={a for item in row.get('source_identities',[]) for span in item['ranges']
                     for a in (expanded if span=='*' else addresses(span))}
        groups={}
        for address,item in sorted(expanded.items()):
            if address in source_only:continue
            metric=item.get('metric_id')
            if not metric:continue
            allowed={d['id'] for d in standard['metric'][metric]['dimension_refs']}
            values={k:v['value'] for k,v in item['dimensions'].items() if k in allowed}
            key=(cells[address]['column'],metric,json.dumps(values,sort_keys=True,ensure_ascii=False))
            groups.setdefault(key,[]).append(address)
        for group in groups.values():
            if len(group)<2:continue
            labels={a:'；'.join(c['value'] for c in cells.values() if c['row']==cells[a]['row'] and c['column']<cells[a]['column']
                and isinstance(c['value'],str) and c['value'].strip() and not c.get('formula')) for a in group}
            if len(set(labels.values())-{''})<2:continue
            questions.append({'id':row['id'],'ranges':sorted(group),'check':'same_identity_different_rows',
                'reason':'这些同列位置目前得到完全相同的身份，但行原文不同：'+json.dumps(labels,ensure_ascii=False)+
                '。请逐行保留各自对象/项目/类别及合计的总体；不能把多个明细的合集共同赋给每条明细。共同公司、期间等不能覆盖这些差异。只有原文确实表达相同范围才可沿用，并说明依据。'})
    return questions


def semantic_questions(binding,tables,snapshot,standard):
    """指出已发生过的合计范围疑点；只提问，财务结论仍由模型依据原文给出。"""
    if standard.get('_complete_identity'):return []
    questions=duplicate_identity_questions(binding,tables,snapshot,standard);index={t['id']:t for t in tables}
    for row in binding['tables']:
        table=index[row['id']];cells=next(s for s in snapshot['sheets'] if s['name']==table['sheet'])['cells']
        parent_checked=set()
        owned=set(data_cells(table,cells))
        objects=' '.join(standard['metric'][m['id']]['definition']['economic_object'] for m in row['metrics'] if m['id'] in standard['metric'])
        total_rows={c['row'] for a,c in cells.items() if a in addresses(table['range']) and isinstance(c['value'],str) and ('合计' in c['value'] or '小计' in c['value'])}
        for item in row['dimensions']:
            definition=standard['dimension'].get(item['id'],{});domain=definition.get('domain','');value=item['value']
            if not domain or domain not in objects or not isinstance(value,dict):continue
            reason=None;question_ranges=item['ranges']
            if value.get('mode') in {'all','members'} and table['subject']!=domain and item['id'] not in parent_checked:
                parent_checked.add(item['id']);question_ranges=table['data_ranges']
                examples=list(dict.fromkeys(str(name) for d in row['dimensions'] if d['id']==item['id'] and isinstance(d['value'],dict) for name in d['value'].get('members',[])))[:3]
                reason=f'本维度总体为指标计量基础下的“{domain}”，当前表为“{table["subject"]}”。请核对本表全部该维度取值是否保留实际科目/父项：all不能省掉局部父项，成员名称脱离表题也须完整（当前示例：{examples}）。指标本身已经保证的限定不必重复，不添加原文没有的条件。'
            elif value.get('mode')=='members':
                area=set(a for r in item['ranges'] for a in (owned if r=='*' else addresses(r)))
                if any(cells[a]['row'] in total_rows for a in area&owned):
                    reason='这些成员名单覆盖了原表合计/小计行。请核对合计代表的父项或排名条件，不能用目前列出的公司或明细名单代替；只有原文确实限定列举成员才保留名单。'
            if reason:questions.append({'id':row['id'],'dimension':item['id'],'ranges':question_ranges,'reason':reason})
    return questions


def apply_role_corrections(first, review, table, packet, standard, aliases):
    from .统一识别 import validate_response
    for row in review.get('corrections',{}).get('non_business',[]):
        area=addresses(row['range'])
        _require(set(area)<=set(packet['candidate_cells']),'业务角色修正超出原表')
        _require(row['category'] in {'label','annotation'},'只允许有据的标签或注释修正')
        record={'kind':'non_business','category':row['category'],'reason':row['reason'],
                'evidence_cells':normalize_evidence(row['evidence'],table['sheet'],aliases,area[0])}
        first.update(validate_response({'cells':{a:record for a in area}},
                     {**packet,'candidate_cells':area},standard))
    return first


def shared_binding(binding):
    """只合并完全相同且范围不重叠的声明，原文措辞不做同义改写。"""
    result=deepcopy(binding)
    for row in result['tables']:
        for key in ('metrics','dimensions'):
            grouped={}
            for item in row[key]:
                signature=json.dumps({k:v for k,v in item.items() if k!='ranges'},sort_keys=True,ensure_ascii=False)
                grouped.setdefault(signature,[]).append(item)
            merged=[]
            for group in grouped.values():
                ranges=[a for item in group for a in ([item['ranges']] if isinstance(item['ranges'],str) else item['ranges'])]
                try:flat=[a for area in ranges for a in addresses(area)]
                except (ValueError,TypeError):merged.extend(group);continue
                if len(flat)!=len(set(flat)):merged.extend(group);continue
                merged.append({**group[0],'ranges':ranges})
            row[key]=merged
    return result


REVIEW_EVIDENCE_RULE='核查只引用text_evidence列出的当前表实际文字格，或source.context内适用的说明编号。reading_notes及其他表的文字不能代替当前表证据；缺少具体事实保持未提供，不据此编造取值。'


def review_text_evidence(tables,snapshot,context):
    return {t['id']:{t['sheet']+'!'+a:value
        for a,c in table_packet(t,snapshot,context)['cells'].items()
        if isinstance(value:=c.get('cached_value') if c.get('formula') else c.get('value'),str) and value.strip()}
        for t in tables}


def review_material(payload, binding, tables, snapshot, context, standard, prior_reviews=()):
    _,aliases=source_material(snapshot,context,tables);index={t['id']:t for t in tables};issues=[]
    catalog=context_catalog(context,snapshot['path'])
    for row in binding['tables']:
        table=index[row['id']]
        source_known=source_identity_records(row,table,table_packet(table,snapshot,context),standard,aliases)
        try:records=checked_records(row,table,table_packet(table,snapshot,context),standard,aliases)
        except (ValueError,KeyError,TypeError) as error:
            reason=str(error)
            if reason=='范围须为有效的单格或矩形':
                reason='某个身份条目的ranges或evidence包含不能解析的位置引用；不是说reading中的表范围有错。详见下面各条维度问题。'
            issues.append({'id':row['id'],'reason':reason})
            # 首个冲突不能遮住其他已知错误；逐项检查仅补报无效项，不重复列必需维度缺项。
            for dimension in row['dimensions']:
                try:
                    partial=checked_records({**row,'dimensions':[dimension]},table,table_packet(table,snapshot,context),standard,aliases)
                    failures={r['reason'] for r in partial.values() if r['kind']=='invalid_response'}
                except (ValueError,KeyError,TypeError) as detail:failures={str(detail)}
                for reason in failures:
                    if '@context:' in reason and not standard.get('_complete_identity'):reason='所引用的@C说明不适用于这些位置，请依据当前表的表头、行名或全局说明重新引用。'
                    issues.append({'id':row['id'],'dimension':dimension['id'],'ranges':dimension['ranges'],'reason':reason})
            continue
        expanded=expand_scopes(row,table,table_packet(table,snapshot,context)['cells'],standard)
        for dimension in row['dimensions']:
            key=dimension['id'];definition=standard['dimension'].get(key)
            declared={a for span in dimension['ranges'] for a in (expanded if span=='*' else addresses(span))}
            relevant=sorted(a for a in declared&set(expanded) if expanded[a].get('metric_id') in standard['metric'] and
                key in {d['id'] for d in standard['metric'][expanded[a]['metric_id']]['dimension_refs']})
            if definition is None or not relevant:continue
            value=expanded[relevant[0]]['dimensions'][key]['value'];failures=[]
            try:_dimension_value(value,definition)
            except (ValueError,TypeError,KeyError) as error:failures.append(str(error))
            try:normalize_evidence(dimension['evidence'],table['sheet'],aliases,relevant[0])
            except (ValueError,TypeError,KeyError,IndexError) as error:failures.append('原文依据：'+str(error))
            for reason in failures:
                issues.append({'id':row['id'],'dimension':key,'ranges':relevant,'reason':reason,'provided':deepcopy(value),
                    'expected':{k:deepcopy(definition[k]) for k in ('value_type','domain','allowed_values') if k in definition}})
        reasons={}
        for address,record in records.items():
            if address in source_known and record['kind'] in {'partial_metric_fact','unresolved'}:continue
            if record['kind'] in {'invalid_response','unresolved'}:
                reason=record['reason']
                if reason=='缺少对应标准':reason='请核对原文及已查标准，修正指标选择；报告标准缺项须提供原文与标准对照证据。'
            elif record['kind']=='partial_metric_fact':
                if snapshot.get('semantic_slot_bindings'):continue
                reason='金标准绑定未完成，请结合行列和共同说明检查这些维度：'+','.join(record.get('missing_dimensions',[]))+'；不得用文字身份替代，也不得补造条件。'
            else:continue
            for alias,ref in aliases.items():
                if not isinstance(ref,str):continue
                if ref in reason:
                    scope=catalog[ref]['scope']
                    if scope and not standard.get('_complete_identity'):
                        top,left,bottom,right=scope['bounds']
                        area=f'{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}'
                        reason=f'所引说明{alias}只适用于{scope["sheet"]}!{area}，不适用于当前问题位置；请换用适用的原文依据。'
                    else:reason=reason.replace(ref,alias)
            reasons.setdefault(reason,[]).append(address)
        issues.extend({'id':row['id'],'cells':cells,'reason':reason} for reason,cells in reasons.items())
    grouped={}
    for issue in issues:
        key=json.dumps({k:v for k,v in issue.items() if k not in {'ranges','cells'}},sort_keys=True,ensure_ascii=False)
        target=grouped.setdefault(key,{k:v for k,v in issue.items() if k not in {'ranges','cells'}})
        for field in ('ranges','cells'):
            if field in issue:target[field]=list(dict.fromkeys(target.get(field,[])+issue[field]))
    issues=list(grouped.values())
    result={**payload,'task':'a_semantic_review','recognized':shared_binding(binding),'共同声明合并':True,
            '返回骨架':{'tables':[{'id':t['id'],'corrections':{'patches':[],'resolved_ranges':[]},'issues':[]} for t in tables]}}
    # 绑定可纠正初选；核查必须看到实际采用的完整定义，不能只沿用初选目录。
    metrics=[m['id'] for m in payload.get('standards',{}).get('metrics',[])]
    metrics.extend(m['id'] for row in binding['tables'] for m in row.get('metrics',[])
                   if m['id'] in standard['metric'] and standard['metric'][m['id']]['status']=='active')
    selected=set(metrics)
    related=[key for key,m in standard['metric'].items() if m['status']=='active' and key not in selected
             and (ratio:=m['definition'].get('ratio'))
             and {ratio['numerator']['metric_id'],ratio['denominator']['metric_id']}<=selected]
    if related:result['review_candidates']=related;metrics.extend(related)
    if metrics:result['standards']=standard_material(metrics,standard)
    result['cell_requirements']=cell_requirements([m['id'] for m in result['standards']['metrics']],standard)
    if issues:result['program_issues']=issues
    questions=semantic_questions(binding,tables,snapshot,standard)
    if questions:result['semantic_questions']=questions
    previous=[{'id':row['id'],**issue} for answer in prior_reviews for row in answer.get('tables',[]) for issue in row.get('issues',[])]
    if previous:result['previous_issues']=previous
    if 'semantic_candidates' in snapshot and snapshot.get('semantic_formula_sources',0)>=4:
        result['text_evidence']=review_text_evidence(tables,snapshot,context)
        result['evidence_rule']=REVIEW_EVIDENCE_RULE
    if standard.get('_complete_identity'):
        result['identity_collection_version']=1
        result.pop('text_evidence',None)
        result['evidence_rule']='只引用source实际提供的原文及共同引用；跨表与同文的业务适用性由识别核查确认，程序仅检验引用内容存在。'
    return result


def assemble(snapshot, context, standard, reading, batches, b_scope=None, b_outline=None, require_identity=False):
    from .统一识别 import apply_identity_review, workbook_coverage
    result={'facts':[],'owners':{},'excluded':[],'unresolved':[],'source_identities':[]}
    confirm_slots=bool(snapshot.get('semantic_slot_bindings'))
    if confirm_slots:result['slot_bindings']=[]
    tables={t['id']:t for t in reading['tables']};processed=set()
    for batch in batches:
        binding,review=batch[:2]
        review=carry_role_corrections(review,batch[2] if len(batch)>2 else [])
        _,aliases=source_material(snapshot,context,[tables[r['id']] for r in binding['tables']])
        binding=reviewed_binding(binding,review,list(tables.values()),snapshot)
        if require_identity:
            _require(not any(row.get('source_identities') for row in binding['tables']),
                     '识别成功必须完整对应金标准，核查不能改回文字身份')
        reviews={r['id']:r for r in review.get('tables',[])}
        for row in binding['tables']:
            table=tables[row['id']];packet=table_packet(table,snapshot,context)
            try:first=checked_records(row,table,packet,standard,aliases)
            except (ValueError,KeyError,TypeError) as error:
                first={a:{'kind':'unresolved','reason':'身份声明尚未修正：'+str(error),'evidence_cells':[]} for a in packet['candidate_cells']}
            first=apply_role_corrections(first,reviews.get(row['id'],{}),table,packet,standard,aliases)
            checked=review_for_table(reviews[row['id']],table,packet,aliases) if row['id'] in reviews else {'checked_cells':[],'issues':[]}
            part=apply_identity_review(first,checked,standard,confirm_slots=confirm_slots)
            source_known=source_identity_records(row,table,packet,standard,aliases)
            # 明确存在标准差别时，不能因旧metrics恰好填齐字段就保留为精确事实。
            retained=[]
            for fact in part['facts']:
                address=fact['source_reference']['location'].rpartition('!')[2]
                if address in source_known:
                    part['unresolved'].append({'cell':address,'reason':source_known[address]['standard_difference']})
                else:retained.append(fact)
            part['facts']=retained
            blocked={a for issue in checked['issues'] for a in issue['cells']}
            blocked.update(a for issue in row.get('unresolved',[]) for a in addresses(issue['range']))
            excluded={item['cell'] for item in part['excluded']}
            _require(not excluded&set(source_known),'原文身份证与非业务分类冲突')
            for address,identity in source_known.items():
                if address in checked['checked_cells'] and address not in blocked:
                    result['source_identities'].append(identity)
            _require(not {(table['sheet'],a) for a in packet['candidate_cells']}&processed,'重复采纳子表')
            processed.update((table['sheet'],a) for a in packet['candidate_cells'])
            result['facts'].extend(part['facts']);result['owners'].update(part['owners'])
            if confirm_slots:result['slot_bindings'].extend({'sheet':table['sheet'],**r} for r in part['slot_bindings'])
            for key in ('excluded','unresolved'):
                result[key].extend({'sheet':table['sheet'],**r} for r in part[key])
    table_areas={}
    for table in tables.values():table_areas.setdefault(table['sheet'],set()).update(addresses(table['range']))
    for sheet in snapshot['sheets']:
        physical=set(a for t in sheet.get('word_table_ranges',[]) for a in addresses(t['range']))
        for address,cell in sheet['cells'].items():
            if b_scope is not None and sheet['name']+'!'+address not in snapshot['semantic_candidates']:continue
            if (sheet['name'],address) in processed:continue
            if address not in table_areas.get(sheet['name'],set()) and address not in physical and (cell['value'] is None or not str(cell['value']).strip()):
                result['excluded'].append({'sheet':sheet['name'],'cell':address,'category':'empty_padding',
                    'reason':'整体阅读所列各表及原Word物理表以外的空白间隔','evidence_cells':[],
                    'recognition_provenance':{'method':'a_semantic_padding','record_id':definition_hash(reading)}})
            else:result['unresolved'].append({'sheet':sheet['name'],'cell':address,'reason':'此处尚未完成语义识别及核查'})
    result['outline']={'complete':all(any(t['sheet']==s['name'] and a in addresses(t['range']) for t in tables.values())
                                    for s in snapshot['sheets'] for a,c in s['cells'].items() if c['value'] is not None and str(c['value']).strip()),
                       'subjects':list(dict.fromkeys(t['subject'] for t in tables.values())),'subtables':reading['tables']}
    if b_scope is not None:
        result['outline']=deepcopy(b_outline)
        result['recognition_scope']=deepcopy(b_scope)
    result['coverage']=workbook_coverage(result,snapshot);result['complete']=result['coverage']['complete']
    validate_facts(result['facts'],standard,result['owners'])
    # 旧回执可重读；新识别的完成依据只有准确对应金标准的事实。
    facts={f['source_reference']['location'] for f in result['facts']}
    facts.update(r['sheet']+'!'+r['cell'] for r in result.get('slot_bindings',[]))
    source_known={f['source_reference']['location'] for f in result['source_identities']}
    _require(not facts&source_known,'同一格不能同时声明完全对应和存在标准差别')
    remaining=[u for u in result['unresolved'] if u['sheet']+'!'+u['cell'] not in source_known]
    if require_identity:remaining=result['unresolved']
    result['b_identity_coverage' if b_scope is not None else 'a_identity_coverage']={'financial_cells':len(facts|source_known),'unresolved_cells':len(remaining),
        'unresolved':remaining,'complete':not remaining and result['outline']['complete']}
    if b_scope is not None:
        from .表格编排 import required_coverage
        result['recognition_scope']['coverage']=required_coverage(result,standard)
    return result


def carry_role_corrections(review, earlier):
    result=deepcopy(review)
    for answer in earlier:
        for old in answer.get('tables',[]):
            if isinstance(old.get('issues'),list) and not any(r['id']==old['id'] for r in result.get('tables',[])):
                result.setdefault('tables',[]).append({'id':old['id'],'checked_ranges':[],'corrections':{},'issues':[]})
    for row in result.get('tables',[]):
        history=[old for answer in [*earlier,result] for old in answer.get('tables',[]) if old['id']==row['id']]
        # 旧回执未限定范围且明确返回问题数组，表示全表核查；后续局部核查不能缩小它。
        if any('checked_ranges' not in old and isinstance(old.get('issues'),list) for old in history):
            row.pop('checked_ranges',None)
        else:
            row['checked_ranges']=list(dict.fromkeys(span for old in history for span in old.get('checked_ranges',[])))
        previous=[fix for answer in earlier for old in answer.get('tables',[]) if old['id']==row['id'] for fix in old.get('corrections',{}).get('non_business',[])]
        if previous:
            corrections=row.setdefault('corrections',{})
            corrections['non_business']=previous+corrections.get('non_business',[])
        # 旧问题须明确解除，不能因后轮漏报而消失；只保留尚未解除的原范围。
        issues=[]
        for answer in [*earlier,result]:
            for old in answer.get('tables',[]):
                if old['id']!=row['id']:continue
                cleared={a for area in old.get('corrections',{}).get('resolved_ranges',[]) for a in addresses(area)}
                issues=[{**issue,'range':a} for issue in issues for a in addresses(issue['range']) if a not in cleared]
                issues.extend(deepcopy(old.get('issues',[])))
        row['issues']=list({json.dumps(issue,sort_keys=True,ensure_ascii=False):issue for issue in issues}.values())
    return result


def focused_review_payload(payload, focus, standard):
    """保留整表原文，只发送本次明确核查范围的已有身份及实际定义。"""
    result=deepcopy(payload);tables={t['id']:t for t in result['reading']}
    _require(set(focus)==set(tables),'局部核查须明确每张表的范围')
    _require(any(focus.values()),'本次核查范围为空')
    metrics=[]
    for row in result['recognized']['tables']:
        wanted={a for span in focus[row['id']] for a in addresses(span)}
        _require(wanted<=set(addresses(tables[row['id']]['range'])),'本次核查范围越表')
        data={a for span in tables[row['id']]['data_ranges'] for a in addresses(span)}
        def clip(spans, financial=True):
            declared=set().union(*(data if span=='*' else set(addresses(span)) for span in spans))
            return sorted(declared&wanted&(data if financial else wanted))
        for key in ('metrics','dimensions','source_identities'):
            if key not in row:continue
            row[key]=[{**item,'ranges':selected} for item in row.get(key,[]) if (selected:=clip(item['ranges']))]
        for key in ('unresolved','non_business'):
            if key in row:row[key]=[{**item,'range':a} for item in row[key] for a in clip([item['range']],financial=key!='non_business')]
        metrics.extend(item['id'] for item in row['metrics'])
        metrics.extend(item['closest_metric_id'] for item in row.get('source_identities',[]) if item.get('closest_metric_id'))
    metrics.extend(result.get('review_candidates',[]))
    _require(all(m in standard['metric'] and standard['metric'][m]['status']=='active' for m in metrics),'核查候选指标无效')
    result['standards']=standard_material(metrics,standard)
    if 'cell_requirements' in result:result['cell_requirements']=cell_requirements(metrics,standard)
    result['review_focus']=deepcopy(focus)
    for key in ('program_issues','semantic_questions','previous_issues'):
        relevant=[]
        for issue in result.pop(key,[]):
            # 没有位置的具体问题仍须保留；不能把无法解析引用本身的错误也筛掉。
            if not isinstance(issue,dict) or 'id' not in issue:
                relevant.append(issue);continue
            wanted={a for span in focus.get(issue['id'],[]) for a in addresses(span)}
            if not wanted:continue
            field=next((name for name in ('cells','ranges','range') if name in issue),None)
            if field is None:
                relevant.append(issue);continue
            spans=[issue[field]] if field=='range' else issue[field]
            declared={a for span in spans for a in (wanted if span=='*' else addresses(span))}
            selected=sorted(declared&wanted)
            if field=='range':relevant.extend({**issue,'range':a} for a in selected)
            elif selected:relevant.append({**issue,field:selected})
        if relevant:result[key]=relevant
    return result


def scoped_read_tables(snapshot, b_scope, b_outline):
    from .表格编排 import scoped_subtables
    detailed={d['subtable_id'] for d in b_scope['details']}
    tables=[deepcopy(r) for t in scoped_subtables(b_outline,b_scope) if t['id'] in detailed for r in t['ranges']]
    if snapshot.get('semantic_formula_sources',0)>=4:
        from .B范围定位 import location_payload
        previews={s['sheet']:s['preview_range'] for s in location_payload(snapshot,[],[])['sheets']}
        # 定位读过的同表公共说明继续交给语义系统，不能在切入子表时丢失。
        for table in tables:
            table['context_ranges']=list(dict.fromkeys([*table.get('context_ranges',[]),previews[table['sheet']]]))
    return tables


def semantic_read_input(snapshot, context, b_scope=None, b_outline=None):
    """A读整份，B读已定位范围；后续选标准、绑定和核查共用。"""
    tables=None
    if b_scope is not None:
        snapshot['semantic_candidates']=set(b_scope['candidate_locations'])
        snapshot['semantic_known_targets']=deepcopy(b_scope['targets'])
        tables=scoped_read_tables(snapshot,b_scope,b_outline)
    material,_=source_material(snapshot,context,tables)
    payload={'task':'a_semantic_read','source':material}
    if b_scope is not None:
        payload['candidate_locations']=b_scope['candidate_locations']
        payload['located_reading']=[{'id':d['subtable_id'],'structure':d['structure']} for d in b_scope['details']]
        if snapshot.get('semantic_formula_sources',0)>=4:
            # 初读只定编排；逐表绑定和核查仍由source_material提供完整公式依据。
            if snapshot['semantic_formula_sources']==4:material.pop('公式引用依据',None)
            payload['reading_ranges']=tables
            payload['range_rule']='每个tables.range须在reading_ranges列出的range内；公共表头或说明只填context_ranges，不并入或扩大tables.range。表范围须覆盖全部candidate_locations；data_ranges只列其中的实际财务格，其他候选仍须由后续共同语义流程确认标签身份并核查，不能丢弃候选。不能据预览追加候选。'
    return payload


def validate_scoped_reading(answer, snapshot, b_scope=None, b_outline=None):
    if b_scope is not None and not b_scope['candidate_locations']:
        _require(answer=={'tables':[],'notes':'未找到必要候选，原A需要继续保留。'},'无候选不能保存身份阅读结果')
        return deepcopy(answer)
    reading=validate_reading(answer,snapshot)
    if b_scope is not None:
        from .表格编排 import in_ranges
        tables=scoped_read_tables(snapshot,b_scope,b_outline)
        allowed=[{'sheet':t['sheet'],'range':area} for t in tables for area in [t['range'],*t.get('context_ranges',[])]]
        for table in reading['tables']:
            _require(all(in_ranges(table['sheet'],a,tables) for a in addresses(table['range'])), 'B阅读超出已定位范围：'+table['sheet']+'!'+table['range'])
            _require(all(in_ranges(table['sheet'],a,allowed) for span in table.get('context_ranges',[]) for a in addresses(span)), 'B公共说明超出已展示范围：'+table['sheet']+'!'+str(table.get('context_ranges')))
        actual={t['sheet']+'!'+a for t in reading['tables'] for r in t['data_ranges'] for a in addresses(r)}
        expected=set(b_scope['candidate_locations'])
        if snapshot.get('semantic_formula_sources',0)>=4:
            covered={t['sheet']+'!'+a for t in reading['tables'] for a in addresses(t['range'])}
            _require(expected<=covered and actual<=expected,
                     'B阅读遗漏必要候选或增加了无关数据格：遗漏候选='+str(sorted(expected-covered))+'；多出数据='+str(sorted(actual-expected)))
        else:
            _require(actual==expected, 'B阅读遗漏必要候选或增加了无关数据格')
    return reading


def batch_origin(batch, mapping, cache):
    """跨版本复用保留旧回执来源；业务材料仍由当前标准逐项重建核对。"""
    reference=batch.get('reused_from')
    if not reference:return mapping
    key=(reference['path'],reference['sha256'])
    if key not in cache:
        _require(file_hash(reference['path'])==reference['sha256'],'复用的旧语义映射发生变化')
        old=json.loads(Path(reference['path']).read_text(encoding='utf-8-sig'),object_pairs_hook=json_dates)
        _require(old['source_hash']==mapping['source_hash'] and old.get('source_context',[])==mapping.get('source_context',[]),
                 '复用的原件或上下文不一致')
        _require(file_hash(old['gold_path'])==old['gold_hash'],'复用的历史标准发生变化')
        cache[key]=old
    old=cache[key]
    fields=('binding','prior_reviews','review')
    proof=old.get('semantic_reading') or old.get('a_semantic_reading',{})
    _require(any(all(candidate.get(k,[])==batch.get(k,[]) for k in fields) for candidate in proof.get('batches',[])),
             '复用回执不属于所引用的旧映射')
    return batch_origin(next(candidate for candidate in proof['batches']
                             if all(candidate.get(k,[])==batch.get(k,[]) for k in fields)),old,cache)


def semantic_groups(tables, saved_groups):
    """沿用已核查分组；其余按已读出的科目共用材料，子表身份保持独立。"""
    saved_ids={t['id'] for group in saved_groups.values() for t in group}
    subjects={}
    for table in tables:
        if table['id'] not in saved_ids:subjects.setdefault(table['subject'],[]).append(table)
    groups=[];grouped=set()
    for table in tables:
        if table['id'] in grouped:continue
        group=saved_groups.get(table['id'],subjects.get(table['subject'],[]))
        groups.append(group);grouped.update(t['id'] for t in group)
    return groups


def checked_reading(mapping):
    """核实全部子表都有完整核查记录，不把中断进度当作全表起点。"""
    proof=mapping.get('semantic_reading') or mapping.get('a_semantic_reading',{})
    if not proof.get('reading'):return None
    _,reading=_stage_receipt(proof['reading'][0],mapping,'a_semantic_read')
    reviewed=set();origins={}
    for batch in proof.get('batches',[]):
        origin=batch_origin(batch,mapping,origins)
        for ref in batch.get('prior_reviews',[])+batch.get('review',[]):
            _,answer=_stage_receipt(ref,origin,'a_semantic_review')
            if answer.get('incomplete') or any(not isinstance(r,dict) or not isinstance(r.get('issues'),list) for r in answer['tables']):return None
            reviewed.update(r['id'] for r in answer['tables'] if isinstance(r,dict) and isinstance(r.get('issues'),list))
    return reading if {t['id'] for t in reading['tables']}<=reviewed else None






def scope_material(snapshot, context):
    """范围认识保留所有行列文字；公式身份留给必要候选的完整识别。"""
    material,_=source_material({**snapshot,'semantic_formula_sources':0},context)
    for sheet in material['sheets']:
        grouped=[];start=end=None;previous=None
        for line in sheet['rows'].splitlines():
            number,body=line.split(': ',1);number=int(number)
            if body==previous and number==end+1:
                end=number
                continue
            if start is not None:grouped.append(f'{start}'+(f'-{end}' if end!=start else '')+': '+previous)
            start=end=number;previous=body
        if start is not None:grouped.append(f'{start}'+(f'-{end}' if end!=start else '')+': '+previous)
        sheet['rows']='\n'.join(grouped)
    material['行号说明']='连续重复行写为起行-止行，该范围内每一行的列值完全相同；没有删除任何行。仅认识范围，不据此确认公式格身份。'
    return material


def _compact_coordinate_ranges(addresses):
    """只合并同列连续坐标，展开后保留全部原位置。"""
    columns = {}
    for address in addresses:
        column, row, _, _ = range_boundaries(address)
        columns.setdefault(column, []).append(row)
    result = []
    for column, rows in sorted(columns.items()):
        rows = sorted(rows); start = end = rows[0]; prefix = get_column_letter(column)
        for row in rows[1:] + [None]:
            if row is not None and row == end + 1: end = row; continue
            result.append(f'{prefix}{start}' if start == end else f'{prefix}{start}:{prefix}{end}')
            if row is not None: start = end = row
    return result


def compact_formula_material(snapshot, material):
    """逐格保留实际公式，只合并重复的来源文字。"""
    proof=material.get('公式引用依据',{})
    locations=set(material.get('公式原文',{}))|set(proof.get('原文',{}))|set(proof.get('共同行列文字',{}))|set(proof.get('引用',{}))
    if not locations:return material
    formulas={};sheets=[]
    for sheet in snapshot['sheets']:
        cells={a:c for a,c in sheet['cells'].items() if sheet['name']+'!'+a in locations}
        if not cells:continue
        sheets.append({**sheet,'cells':cells})
        for address,cell in cells.items():
            if not cell.get('formula'):continue
            formulas[sheet['name']+'!'+address]=cell['formula']
    unresolved = {}
    for item in proof.get('未展开', []):
        sheet, address = item['位置'].rsplit('!', 1)
        attributes = {k: v for k, v in item.items() if k != '位置'}
        key = (sheet, json.dumps(attributes, ensure_ascii=False, sort_keys=True))
        group = unresolved.setdefault(key, {**attributes, '位置': []})
        group['位置'].append(address)
    for (sheet, _), group in unresolved.items():
        group['位置'] = [sheet + '!' + address for address in _compact_coordinate_ranges(group['位置'])]
    sources=scope_material({**snapshot,'sheets':sheets},[])
    result={k:v for k,v in material.items() if k not in {'公式原文','公式引用依据'}}
    result['公式及来源']={'公式原文':formulas,'来源行列原文':sources['sheets'],
        '行号说明':sources['行号说明'],'未展开':list(unresolved.values()),
        '公式写法':'逐格列出实际公式，不需要展开相对偏移；未展开位置中的连续范围包含范围内全部格。所有递归引用来源及共用文字保留在来源行列原文中。'}
    return result


DIRECT_PROMPT = '表中指定的各格分别是什么意思，和金标准里的哪个槽位或维度最接近？ 按本格承载的内容选标准：名称或类别格可选相应维度，不套用同行金额格的指标。请逐格答完作答位置列表，不只给示例。'

CELL_IDENTITY_OUTPUT = '逐格给出location、metric_id或dimension_id、dimensions（按金标准原始维度ID）和template_period_role，值依本格原文；标准编号分别用metric_id或dimension_id写明，不合写字段名。dimensions填维度ID到实际值的对象，不只列名称。适用于本格的共同属性也各格写明，未知照实写。不把金额列的期间、币种或单位套给名称格；本格明示的限定照实保留。period用kind=instant及date，或kind=duration及start/end，日期为YYYY-MM-DD。集合值写在该集合维度原始ID下，mode用（all/members/predicate/complement/unknown），按实际给members、predicate或base/excluded。期间角色closing/opening/current/prior分别为期末/期初/本期/上期。同一金标准维度下，本格与已有参照含义相同时，直接复用参照的完整实际值对象，不另换mode或改写members、predicate等表达；本方原文不同或未知时照实保留；空白业务格照常识别，项目未写只将范围记未知。原文依据可另列，不展开共同声明或解释同义句，不交A/B位置取数或写回安排。文字、列表、表格或对象均可。'


def _with_declared_context(mapping, ranges):
    """定向范围沿用模型已指定的共同阅读引用，不按内容判断适用性。"""
    declared=[table for group in mapping.get('正常入口范围认识',{}).get('ranges',[])
              for table in group.get('tables',[group])]
    result=[]
    for area in ranges:
        refs=list(area.get('context_references',[]));owned=set(addresses(area['range']))
        for table in declared:
            if table['sheet']==area['sheet'] and owned.intersection(addresses(table['range'])):
                refs.extend(table.get('context_references',[]))
        result.append({**area,'context_references':list({json.dumps(ref,sort_keys=True):ref for ref in refs}.values())})
    return result


def recognize_list_relations(engine, a, b, config, destination):
    """旧名单入口接回B逐格识别，只更新明确指定的身份字段。"""
    dimension=config['list_dimension']
    _require(dimension in engine.standard['dimension'] and
             engine.standard['dimension'][dimension]['value_type']=='selection','名单维度须来自金标准')
    field_scope=deepcopy(config.get('b_field_scope') or {
        area['sheet']+'!'+area['range']:[dimension] for area in config['b_source_ranges']})
    metadata={key:deepcopy(config[key]) for key in
        ('update_purpose','comparison_requirements','a_required_ranges','b_source_ranges') if key in config}
    metadata.update(carrier='B',normal_entry=True,field_scope=field_scope,
        additional_context=str(config.get('b_additional_context','') or ''))
    if config.get('b_source_images'):metadata['source_images']=deepcopy(config['b_source_images'])
    result=recognize_a(engine,b['source_path'],b.get('source_context',[]),artifact_metadata=metadata,
        previous_result=b['mapping_path'],target_mapping=a,allowed_metrics=config.get('b_allowed_metrics'),
        source_snapshot=deepcopy(b['机械记录']))
    destination=Path(destination);destination.mkdir(parents=True,exist_ok=True)
    answer_path=destination/'B识别答复.txt'
    answer_path.write_text(result.get('答复接入来源',{}).get('本次原答复',result['识别答复']),encoding='utf-8-sig')
    from .A合集 import direct_recognition_status
    status=direct_recognition_status(result)
    return {'status':status,'message':('B指定格识别字段已完整接入并保存。' if status=='complete' else 'B指定格原答复和已接入字段已保存；识别尚未完整接入，未重答。'),
        'files':{'b_mapping':result['mapping_path'],'recognition_answer':str(answer_path)},'usage':dict(engine.usage)}


def recognize_a(engine, source, context, progress=None, artifact_metadata=None, previous_result=None, *, b_scope=None, b_outline=None, target_mapping=None, sheets=None, allowed_metrics=None, target_table=None, source_snapshot=None):
    """A/B各进行一次直接识别；旧分轮辅助函数仅供历史成果读取。"""
    metadata = artifact_metadata or {}
    carrier = metadata.get('carrier', 'B' if target_mapping is not None or b_scope is not None else 'A')
    field_scope=metadata.get('field_scope')
    period_labels_only=bool(field_scope and
        'period只补明确的原表标签引用，不改日期' in metadata.get('additional_context','') and
        all(set(fields)<={'period','template_period_role'} for fields in field_scope.values()))
    snapshot = source_snapshot if source_snapshot is not None else read_workbook(source, context=context)
    materialize_candidate_blanks(snapshot)
    snapshot['semantic_formula_sources'] = 5
    snapshot['semantic_identity_version'] = 2
    previous=None;execution=None;original_collections={}
    if previous_result:
        stored=json.loads(Path(previous_result).read_text(encoding='utf-8-sig'),object_pairs_hook=json_dates)
        if stored.get('schema')=='附注更新执行记录-v1':
            from .A合集 import resume_direct_output, direct_digest
            _require(carrier=='A','更新执行记录仅用于A副本接续')
            if not field_scope:
                requested=([target_table['sheet']+'!'+cell for area in target_table['data_ranges'] for cell in addresses(area)]
                    if target_table else list(stored.get('template_layout',{}).get('added',{})))
                field_scope={loc:['metric_id','dimension_id','template_period_role',*engine.standard['dimension']] for loc in requested}
            previous,execution,original_collections=resume_direct_output(previous_result,snapshot,context,engine.standard,field_scope)
            if not field_scope:
                destination=engine.work_dir/('直接识别_'+str(time.time_ns()))
                destination.mkdir(parents=True)
                path=destination/'A识别结果.json'
                previous.update(usage=dict(engine.usage),mapping_path=str(path.resolve()))
                previous['封存校验']=direct_digest(previous)
                engine._save(path,previous)
                return previous
        elif field_scope:
            from .统一勾稽 import load_fact_mapping
            previous=load_fact_mapping(previous_result,engine.standard)
        else:
            previous=stored
    selected = None
    if sheets is not None:
        _require(set(sheets) <= {s['name'] for s in snapshot['sheets']}, '指定工作表不存在')
        selected = [{'sheet': s['name'], 'range': f'A1:{get_column_letter(s["max_column"])}{s["max_row"]}'}
                    for s in snapshot['sheets'] if s['name'] in sheets]
    if carrier=='B' and metadata.get('b_source_ranges'):
        selected=[{**area,'data_ranges':[],'read_full_rows':True} for area in metadata['b_source_ranges']]
    if target_table is not None:
        target_table = deepcopy(target_table)
        sheet = next(s for s in snapshot['sheets'] if s['name'] == target_table['sheet'])
        targets = [a for span in target_table['data_ranges'] for a in addresses(span)]
        _require(bool(targets) and len(targets) == len(set(targets)), '固定目标为空或重复')
        _require(set(targets) <= set(addresses(target_table['range'])) <= set(addresses(
            f'A1:{get_column_letter(sheet["max_column"])}{sheet["max_row"]}')), '固定目标越出原表范围')
        selected = [target_table]
    if field_scope and not selected:
        # 多表字段补交也保留完整二维网格；作答位置仍只由原字段范围限定。
        requested_sheets={area.rsplit('!',1)[0].strip("'") for area in field_scope if '!' in area}
        if not requested_sheets and len(snapshot['sheets'])==1:requested_sheets={snapshot['sheets'][0]['name']}
        selected=[{'sheet':s['name'],'range':f'A1:{get_column_letter(s["max_column"])}{s["max_row"]}',
                   'data_ranges':[]} for s in snapshot['sheets'] if s['name'] in requested_sheets]
    if selected and previous and previous.get('source_hash')==snapshot['sha256']:
        selected=_with_declared_context(previous,selected)
    if period_labels_only and carrier=='A':
        requested={sheet+'!'+cell for span in field_scope for sheet,area in [span.rsplit('!',1)] for cell in addresses(area)}
        named=[area for areas in snapshot.get('names',{}).values() for area in areas
               if any(area['sheet']+'!'+cell in requested for cell in addresses(area['range']))]
        if named:selected=named
    reading_context=context; word_ranges=[]
    if carrier=='A' and target_table is not None and metadata.get('a_link'):
        from .Word原文材料 import word_source_material, word_table_context
        areas=[*selected, *(ref for area in selected for ref in area.get('context_references',[]))]
        scoped=word_table_context(word_source_material(metadata['a_link']),areas)
        if scoped is not None:
            section,word_ranges=scoped; linked={row['table_id'] for row in word_ranges}
            reading_context=[item for item in context if not isinstance(item,dict) or
                (item.get('title')!='原Word正文顺序（表格引用指向实际表区，各段适用范围由模型判断）'
                 and (not item.get('table_id') or item['table_id'] in linked))]+[section]
            selected=[{**area,'context_references':list({json.dumps(ref,sort_keys=True):ref for ref in
                [*area.get('context_references',[]),*word_ranges]}.values())} for area in selected]
    material_aliases={}
    if carrier=='B' and target_mapping is None:
        material=scope_material(snapshot,reading_context)
    else:
        direct_formula_input=carrier=='B'
        reading_snapshot=({**snapshot,'semantic_formula_sources':4}
                          if direct_formula_input else snapshot)
        material, material_aliases = source_material(reading_snapshot, reading_context, selected)
        if carrier=='B':
            proof=material.get('公式引用依据',{})
            material=compact_formula_material(reading_snapshot,material)
            if direct_formula_input and '公式及来源' in material:
                material['公式及来源']['未递归数值计算']=deepcopy(proof.get('未递归数值计算',[]))
                material['公式及来源']['说明']=proof.get('说明','')
                material['公式及来源']['公式写法']=material['公式及来源']['公式写法'].replace(
                    '所有递归引用来源及共用文字保留在来源行列原文中。',
                    '保留直接引用和标题公式链，不展开数值公式下层，具体范围见本段说明。')
    word_context=carrier=='A' and any(item.get('title')=='原Word正文顺序（表格引用指向实际表区，各段适用范围由模型判断）' for item in reading_context if isinstance(item,dict))
    if selected and carrier!='B' and not word_context and not all(area.get('context_references') for area in selected):
        # 已有完整Word正文或模型声明的阅读范围时，不再附上其他整表。
        context_sheets={area['sheet'] for area in selected}
        requested={}
        for span in field_scope or []:
            sheet,area=span.rsplit('!',1);requested.setdefault(sheet,set()).update(addresses(area))
        if not requested:
            for area in selected:requested.setdefault(area['sheet'],set()).update(addresses(area['range']))
        references=formula_source_material({**snapshot,'semantic_formula_sources':4},requested)
        context_sheets.update(location.rsplit('!',1)[0]
            for field in ('原文','共同行列文字') for location in references.get(field,{}))
        complete_ranges=[{'sheet':sheet['name'],'range':f'A1:{get_column_letter(sheet["max_column"])}{sheet["max_row"]}'}
                         for sheet in snapshot['sheets'] if sheet['name'] in context_sheets]
        complete,_=source_material({**snapshot,'semantic_formula_sources':1},context,complete_ranges)
        full={key:complete[key] for key in ('sheets','公式原文','原有表区')
              if key in complete and complete[key]!=material.get(key)}
        if full:
            material['完整工作表原文']={**full,'说明':'目标及直接引用工作表的其余完整原文和公式；不扩大作答范围。'}
    gold_text = Path(engine.gold_path).read_text(encoding='utf-8-sig')
    # 只去除旧转换及历史样例，指标与维度定义的原行保持完整。
    definitions = '\n'.join(line for line in gold_text.splitlines() if line.strip()
        and (lambda row: row.get('record_type') == 'dimension' or
             row.get('record_type') == 'metric' and row.get('status') == 'active')(json.loads(line)))
    if allowed_metrics is not None:
        related = standard_material([key for key in allowed_metrics if key in engine.standard['metric']], engine.standard)
        ids = {r['id'] for kind in ('metrics', 'dimensions') for r in related[kind]}
        ids.update(key for key in allowed_metrics if key in engine.standard['dimension'])
        if field_scope:ids.update(key for fields in field_scope.values() for key in fields)
        definitions = '\n'.join(line for line in definitions.splitlines() if json.loads(line)['id'] in ids)
    material['载体']=carrier
    material['源文件']={'path':snapshot['path'],'sha256':snapshot['sha256']}
    materials = {'原文.json': material, '金标准.jsonl': definitions}
    if selected:
        from .表格文字 import workbook_table_text, workbook_table_ranges
        grid_areas=[*selected,*word_ranges]
        whole=material.get('完整工作表原文',{})
        full_sheets={row['sheet'] for row in whole.get('sheets',[])}
        grid_areas.extend({'sheet':s['name'],'range':f'A1:{get_column_letter(s["max_column"])}{s["max_row"]}'}
                          for s in snapshot['sheets'] if s['name'] in full_sheets)
        materials['本次原表网格.txt']=workbook_table_text(snapshot,grid_areas)
    if metadata.get('update_purpose'):materials['本次明确更新要求.json']=deepcopy(metadata['update_purpose'])
    if metadata.get('comparison_requirements'):
        materials['用户明确的比较排除.json']=metadata['comparison_requirements']
    if metadata.get('既有另一方身份参照',{}).get('source_ranges'):
        materials['既有另一方身份证表达.json']=deepcopy(metadata['既有另一方身份参照'])
    if target_table is not None:
        materials['原文.json']['固定目标'] = target_table
    if metadata.get('target_tables'):
        materials['原文.json']['范围认识声明的候选表区']=metadata['target_tables']
    if carrier == 'B':
        if target_mapping is None:
            materials.pop('金标准.jsonl')
        else:
            _require(target_mapping['gold_hash']==engine.gold_hash,'A需要与本次识别没有使用同一金标准')
            materials['A需要.json'] = {'载体':'A','目标表区':deepcopy(metadata.get('a_required_ranges',[])),
                **direct_reference(target_mapping,compact=True,ranges=metadata.get('a_required_ranges'))}
            saved=materials['A需要.json']
            materials['A需要.json']={key:saved[key] for key in
                ('载体','目标表区','source_hash','gold_hash','outline') if key in saved}
            materials['A需要.json']['需要']=[row for row in saved.get('需要',[])
                if allowed_metrics is None or any(row.get(key) in ids for key in ('metric_id','dimension_id'))]
    if previous_result:
        if execution:
            materials['上次搬运记录.json']={key:deepcopy(execution[key]) for key in
                ('output_path','output_hash','collections','update_purpose','comparison_requirements',
                 'template_layout','operations','retained_values')}
            if selected:
                owned={area['sheet']+'!'+cell for area in selected for cell in addresses(area['range'])}
                layout=materials['上次搬运记录.json']['template_layout']
                layout['positions']={old:new for old,new in layout['positions'].items() if new in owned}
                layout['added']={loc:row for loc,row in layout.get('added',{}).items() if loc in owned}
                materials['上次搬运记录.json']['本次位置范围']=deepcopy(selected)
            materials['上次搬运记录.json']['完整记录']={'path':str(previous_result),'sha256':file_hash(previous_result)}
            for side,saved in original_collections.items():
                from .A合集 import _records
                known_ranges=([{'sheet':loc.rsplit('!',1)[0],'range':loc.rsplit('!',1)[1]}
                    for loc,row in _records(saved) if any(row.get(key) in allowed_metrics for key in ('metric_id','dimension_id'))] if allowed_metrics else None)
                materials['上次'+side+'原判断.json']=direct_reference(saved,compact=True,ranges=known_ranges)
                source_sheets={area['sheet'] for area in known_ranges or []}
                complete_ranges=([{'sheet':s['name'],'range':f'A1:{get_column_letter(s["max_column"])}{s["max_row"]}'}
                    for s in saved['机械记录']['sheets'] if s['name'] in source_sheets] if side=='B' and source_sheets else None)
                # 原工作表及逐格公式完整保留，不重复嵌入递归来源副本。
                materials['上次'+side+'原文.json'],_=source_material(
                    {**saved['机械记录'],'semantic_formula_sources':1},saved['source_context'],complete_ranges)
        if previous is not None and (previous.get('source_hash') == snapshot['sha256'] and previous.get('gold_hash') == engine.gold_hash
                and previous.get('source_context') == context):
            if previous.get('recognition_method') != 'direct':
                materials['已有结果.json'] = direct_reference(previous)
    if field_scope:
        _require(previous is not None and previous.get('recognition_method')=='direct'
                 and previous.get('source_hash')==snapshot['sha256'] and previous.get('gold_hash')==engine.gold_hash,
                 '字段纠正须使用同一原件及金标准的已有直接识别合集')
        materials['本次字段范围.json']={'载体':carrier,'字段范围':field_scope}
        materials['已有身份.json']=direct_reference(previous,compact=True,
            ranges=[{'sheet':area.rsplit('!',1)[0],'range':area.rsplit('!',1)[1]} for area in field_scope]
            if all('!' in area for area in field_scope) else None)
        if period_labels_only:
            from .A合集 import _records
            selected_locations={sheet+'!'+cell for area in field_scope
                                for sheet,span in [area.rsplit('!',1)] for cell in addresses(span)}
            materials['已有期间只读参照.json']={location:deepcopy(row['dimensions']['period'])
                for location,row in _records(previous)
                if location in selected_locations and 'period' in row['dimensions']}
    request_materials=({key:value for key,value in materials.items() if key!='A需要.json'}
                       if period_labels_only else materials)
    if field_scope and all(set(fields)<={'period','template_period_role','entity','report_scope','currency','unit_scale'} for fields in field_scope.values()) and not period_labels_only and not metadata.get('source_targets'):
        request_materials={key:value for key,value in materials.items() if key!='A需要.json'}
    source_targets=metadata.get('source_targets')
    if metadata.get('normal_entry') and 'A需要.json' in request_materials:
        request_materials['B原文.json']=request_materials.pop('原文.json')
    if '本次原表网格.txt' in request_materials:
        from .表格文字 import TABLE_STRUCTURE_READING
    if metadata.get('source_images'):
        _require(all(image.get('source_hash')==snapshot['sha256'] for image in metadata['source_images']),
                 '原表图片不属于本次原件，未发送识别')
        request_materials['原表图片.json']=deepcopy(metadata['source_images'])
    if target_table is not None:
        request_materials['本次作答位置.json']={'载体':carrier,'位置':[target_table['sheet']+'!'+cell
            for area in target_table['data_ranges'] for cell in addresses(area)]}
    if carrier=='B' and target_mapping is None:
        prompt='读懂B整份文件的科目、子表、位置和共同说明，保存范围认识，包括空白模板；本次不逐格识别。'
    else:
        prompt=DIRECT_PROMPT+' 本次识别'+carrier+'，只回答材料列出的作答位置和字段，依据本方完整原文及相关金标准。'
        if '本次原表网格.txt' in request_materials:
            prompt+=TABLE_STRUCTURE_READING
        if metadata.get('additional_context'):
            prompt+='\n'+metadata['additional_context']
        prompt+='\n'+CELL_IDENTITY_OUTPUT
    if '本次原表网格.txt' in request_materials and '_material_positions' in material_aliases:
        covered={}
        for _,sheet,span in workbook_table_ranges(snapshot,grid_areas):
            covered.setdefault(sheet,set()).update(addresses(span))
        extra=[]
        for sheet in snapshot['sheets']:
            positions=material_aliases['_material_positions'].get(sheet['name'],set())-covered.get(sheet['name'],set())
            if positions:
                extra.append({**sheet,'cells':{a:c for a,c in sheet['cells'].items() if a in positions}})
        outside,_=source_material({**snapshot,'sheets':extra,'semantic_identity_version':0,'semantic_formula_sources':0},[])
        request_materials=dict(request_materials)
        for key in ('原文.json','B原文.json'):
            if key not in request_materials:continue
            reading_material={k:v for k,v in request_materials[key].items() if k not in ('sheets','完整工作表原文')}
            if outside['sheets']:reading_material['sheets']=outside['sheets']
            request_materials[key]=reading_material
    _require(len(prompt)<=2000,'识别提示词超过2000字符；未发送，原材料保留。')
    answer = engine.request_text(prompt, request_materials)
    _require(file_hash(source) == snapshot['sha256'] and file_hash(engine.gold_path) == engine.gold_hash,
             '识别期间原件或金标准已变化；原答复已经保留')
    record = {**metadata, 'schema_version': 2, 'recognition_method': 'direct', 'carrier': carrier,
              'source_path': snapshot['path'], 'source_hash': snapshot['sha256'], 'source_context': context,
              'gold_path': engine.gold_path, 'gold_hash': engine.gold_hash,
              '原文材料': material, '金标准原文': definitions,
              '识别答复': answer, '答复指纹': definition_hash(answer),
              'usage': dict(engine.usage), 'recognition_scope': {'mode': 'b_outline' if carrier == 'B' and target_mapping is None else 'direct'}}
    if target_table is not None:
        record['target_table'] = target_table
    if target_mapping is not None and 'A需要.json' in materials:
        record['A需要'] = materials['A需要.json']
    destination = engine.work_dir / ('直接识别_' + str(time.time_ns()))
    path = destination / (carrier + '识别结果.json')
    if source_targets:
        from .A合集 import save_direct_answer
        base={**previous,'target_table':target_table,'target_tables':deepcopy(metadata.get('target_tables',[])),
              'source_targets':deepcopy(source_targets),
              'b_source_ranges':deepcopy(metadata.get('b_source_ranges',[])),
              'a_required_ranges':deepcopy(metadata.get('a_required_ranges',[]))}
        identity_scope=field_scope or {area['sheet']+'!'+area['range']:
            ['metric_id','dimension_id','template_period_role',*engine.standard['dimension']]
            for area in metadata.get('b_source_ranges',[])}
        _require(identity_scope,'B身份识别缺少本次已定位的B范围；不能改为保存A/B位置对应')
        result=save_direct_answer({**base,'usage':dict(engine.usage)},answer,path,
            field_scope=identity_scope,withdrawn=metadata.get('withdrawn'),snapshot=previous['机械记录'],
            standard=engine.standard,target_mapping=target_mapping)
        result['本次关系答复']=answer
        return result
    if record['recognition_scope']['mode'] != 'b_outline':
        from .A合集 import save_direct_answer
        if field_scope:
            base=deepcopy(previous)
            base.update(usage=dict(engine.usage),**metadata)
            if execution:base['原文材料']=material
            if target_mapping is not None:base['A需要']=materials['A需要.json']
            base['source_context']=list({json.dumps(row,ensure_ascii=False,sort_keys=True):row
                for row in [*context,*previous['source_context']]}.values())
            base['金标准原文']='\n'.join(dict.fromkeys((previous['金标准原文']+'\n'+definitions).splitlines()))
            base['封存定义原文']=base['金标准原文'].splitlines()
            record=save_direct_answer(base,answer,path,field_scope=field_scope,
                withdrawn=metadata.get('withdrawn'),snapshot=previous['机械记录'],standard=engine.standard,
                target_mapping=target_mapping)
        else:
            record = save_direct_answer(record, answer, path, snapshot=snapshot, standard=engine.standard)
    else:
        destination.mkdir(parents=True)
        record['mapping_path'] = str(path)
        engine._save(path, record)
    if progress:
        progress(record)
    return record


def direct_reference(record, *, compact=False, ranges=None):
    """只搬运已保存的结果，不重建、补齐或解释旧身份。"""
    if compact:
        from .A合集 import _records
        owned={area['sheet']+'!'+cell for area in ranges or [] for cell in addresses(area['range'])}
        selected=lambda loc:not ranges or loc in owned
        needs=[{'location':loc,**{key:row[key] for key in ('metric_id','dimension_id') if row.get(key)},'dimensions':deepcopy(row['dimensions']),
                **({'missing_dimensions':deepcopy(row['missing_dimensions'])} if row.get('missing_dimensions') else {})}
               for loc,row in _records(record) if selected(loc)]
        def without_original(value):
            if isinstance(value,dict):return {k:without_original(v) for k,v in value.items() if k!='original'}
            if isinstance(value,list):return [without_original(v) for v in value]
            return deepcopy(value)
        def kept_ranges(spans,sheet):
            kept=[]
            for span in spans:
                cells=list(addresses(span));chosen=[cell for cell in cells if selected(sheet+'!'+cell)]
                kept.extend([span] if chosen==cells else chosen)
            return kept
        tables=[]
        for table in record.get('outline',{}).get('subtables',[]):
            copied={'sheet':table['sheet'],'metrics':[],'dimensions':[]}
            for kind in ('metrics','dimensions'):
                for row in table.get(kind,[]):
                    if record.get('逐格字段版本')==1 and (kind!='dimensions' or row.get('逐格原字段') is not True):continue
                    if not (kind=='metrics' and 'template_period_role' in row or
                            kind=='dimensions' and row.get('id')=='template_period_role'):continue
                    declared=[part for span in row['ranges'] for part in
                              (table['data_ranges'] if span=='*' else [span])]
                    spans=kept_ranges(declared,table['sheet'])
                    if not spans:continue
                    item={**without_original(row),'ranges':spans}
                    # 已声明的坐标映射只裁切范围；未分配的角色列表保持原样。
                    if isinstance(item.get('template_period_role'),dict):
                        item['template_period_role']={part:role for span,role in item['template_period_role'].items()
                            for part in ([cell for area in spans for cell in addresses(area)
                                if re.fullmatch(span+r'\d+',cell)] if re.fullmatch(r'[A-Z]+',span)
                                else kept_ranges([span],table['sheet']))}
                    copied[kind].append(item)
            if copied['metrics'] or copied['dimensions']:tables.append(copied)
        return {'source_hash':record['source_hash'],'gold_hash':record['gold_hash'],'需要':needs,
                'outline':{'subtables':tables},
                'unresolved':[without_original(row) for row in record.get('unresolved',[]) if
                    not row.get('sheet') or not row.get('cell') or selected(row['sheet']+'!'+row['cell'])],
                'source_identities':[deepcopy(row) for row in record.get('source_identities',[]) if selected(row['source_reference']['location'])]}
    fields = ('carrier', 'source_hash', 'gold_hash', 'recognition_method', '识别答复',
              'outline', 'facts', 'slot_bindings', 'source_identities', 'unresolved',
              '已确认关系', '完整性')
    return {key: deepcopy(record[key]) for key in fields if key in record}


def supplied_review_standards(original, supplied, standard):
    """补查可增加相关定义，实际发送的每项仍须与当前金标准逐字一致。"""
    ids=[m['id'] for m in supplied['metrics']]
    _require({m['id'] for m in original['metrics']}<=set(ids),'补查遗漏原有候选标准')
    _require(all(m in standard['metric'] and standard['metric'][m]['status']=='active' for m in ids),'补查引用无效标准')
    expected=standard_material(ids,standard)
    _require(expected==supplied,'补查定义与正式金标准不符')
    return expected


def verify_mapping(mapping, standard, *, rebuild=False):
    """从保存的原始请求和回执重建结果，不能只信任最终JSON中的完成声明。"""
    from .语义 import _stable_content
    from .统一语义 import read_standard_v2
    historical=deepcopy(standard) if standard['sha256']==mapping['gold_hash'] else read_standard_v2(mapping['gold_path'])
    historical['_historical_material']=mapping.get('formula_source_version',0)<5
    historical['_historical_report_scope']=mapping.get('cell_identity_version')!=1
    _require(historical['sha256']==mapping['gold_hash'],'原识别标准已经改变')
    context=mapping.get('source_context',[])
    snapshot=read_workbook(mapping['source_path'],context=context);materialize_candidate_blanks(snapshot)
    snapshot['semantic_formula_sources']=mapping.get('formula_source_version',0)
    snapshot['semantic_slot_bindings']=mapping.get('slot_binding_version')==1
    _require(snapshot['sha256']==mapping['source_hash'],'识别原件发生变化')
    scoped=bool(mapping.get('semantic_reading'))
    b_scope=mapping['recognition_scope'] if scoped else None
    b_outline=mapping['outline'] if scoped else None
    proof=mapping['semantic_reading' if scoped else 'a_semantic_reading']
    read_input=semantic_read_input(snapshot,context,b_scope,b_outline)
    if scoped and mapping.get('style_reuse_version')==1:
        from .B样式复用 import prepare_reuse
        prepare_reuse(snapshot,b_scope,b_outline,historical)
    if mapping.get('dimension_vocabulary_version')!=1:snapshot.pop('semantic_known_targets',None)
    receipt_mapping=mapping;origin_cache={}
    def read(proofs,task_name,payload):
        _require(len(proofs)==1,'每个阶段只采纳一个完整回执')
        task,answer=_stage_receipt(proofs[0],receipt_mapping,task_name)
        # 提示词原文已包含在实际任务哈希中；以后改进措辞不能使已保存的合法成果失效。
        _require(isinstance(task.get('system'),str) and bool(task['system'].strip()) and
                 task['payload']==_stable_content(payload),'语义结果与实际请求材料不符')
        return answer
    no_candidates=scoped and not b_scope['candidate_locations']
    if no_candidates:
        _require(not any(proof.values()),'无候选不能伪造语义回执')
        reading={'tables':[],'notes':'未找到必要候选，原A需要继续保留。'}
    else:reading=validate_scoped_reading(read(proof['reading'],'a_semantic_read',read_input),snapshot,b_scope,b_outline)
    if proof.get('reading_check'):
        questions=reading_questions(reading,snapshot)
        answer=read(proof['reading_check'],'a_semantic_read_check',
                    reading_check_payload(reading,questions,snapshot,context))
        reading=apply_reading_check(reading,answer,questions,snapshot)
        validate_scoped_reading(reading,snapshot,b_scope,b_outline)
    batches=[]
    if proof['selection']:
        old_task,_=_stage_receipt(proof['selection'][0],mapping,'a_semantic_select')
        select_input=selection_payload(reading,snapshot,historical,include_basis=any('basis' in m for m in old_task['payload']['catalog']))
        if '原表文字' not in old_task['payload']:select_input.pop('原表文字')
        selection=validate_selection(read(proof['selection'],'a_semantic_select',select_input),reading,historical)
        index={t['id']:t for t in reading['tables']}
        selected={r['id']:r for r in selection['tables']}
        for reference in proof.get('selection_additions',[]):
            task,_=_stage_receipt(reference,mapping,'a_semantic_select')
            partial=task['payload']['reading']
            _require(all(t==index.get(t['id']) for t in partial['tables']),'补检索改变了原表理解')
            select_input=selection_payload(partial,snapshot,historical,include_basis=any('basis' in m for m in task['payload']['catalog']))
            if 'selection_questions' in task['payload']:select_input['selection_questions']=task['payload']['selection_questions']
            addition=validate_selection(read([reference],'a_semantic_select',select_input),partial,historical)
            selected.update((r['id'],r) for r in addition['tables'])
        selection={'tables':[selected[t['id']] for t in reading['tables']]}
        for item in proof['batches']:
            receipt_mapping=batch_origin(item,mapping,origin_cache)
            _require(len(item['binding'])==1,'缺少唯一身份确认回执')
            task,_=_stage_receipt(item['binding'][0],receipt_mapping,'a_semantic_bind')
            tables=task['payload']['reading']
            _require(all(t==index.get(t['id']) for t in tables),'身份确认没有沿用原表理解')
            payload,_=binding_payload(tables,selection,snapshot,context,historical,semantic_only='数值处理' in task['payload']['source'],
                                      reading_notes=reading.get('notes','') if 'reading_notes' in task['payload'] else None)
            # 按该次回执实际提供的候选重建，补充同对象候选后不能使已合法保存的旧结果失效。
            supplied=[m['id'] for m in task['payload']['standards']['metrics']]
            payload['standards']=standard_material(supplied,historical)
            if 'cell_requirements' in task['payload']:payload['cell_requirements']=cell_requirements(supplied,historical)
            else:payload.pop('cell_requirements',None)
            binding=validate_bindings(read(item['binding'],'a_semantic_bind',payload),tables,snapshot,context,historical,candidate=True)
            earlier=[]
            seen={t['id']:{a:set(supplied) for a in table_packet(t,snapshot,context)['candidate_cells']} for t in tables}
            def record_definitions(answer,review_task):
                ids={m['id'] for m in review_task['payload']['standards']['metrics']}
                _,review_aliases=source_material(snapshot,context,tables)
                for row in answer['tables']:
                    table=next(t for t in tables if t['id']==row['id'])
                    checked=review_for_table(row,table,table_packet(table,snapshot,context),review_aliases)
                    for address in checked['checked_cells']:seen[row['id']][address].update(ids)
            for reference in item.get('prior_reviews',[]):
                old_task,_=_stage_receipt(reference,receipt_mapping,'a_semantic_review')
                prior_payload={**payload,'task':'a_semantic_review','recognized':binding}
                if 'text_evidence' in old_task['payload']:
                    prior_payload.update(text_evidence=review_text_evidence(tables,snapshot,context),evidence_rule=REVIEW_EVIDENCE_RULE)
                if 'reading_notes' in old_task['payload']:prior_payload['reading_notes']=deepcopy(reading.get('notes',''))
                if old_task['payload'].get('共同声明合并'):
                    prior_payload.update(recognized=shared_binding(binding),共同声明合并=True)
                focus=old_task['payload'].get('review_focus')
                if 'review_candidates' in old_task['payload']:prior_payload['review_candidates']=old_task['payload']['review_candidates']
                if focus is not None:prior_payload=focused_review_payload(prior_payload,focus,historical)
                prior_payload['standards']=supplied_review_standards(prior_payload['standards'],old_task['payload']['standards'],historical)
                if 'cell_requirements' in old_task['payload']:
                    prior_payload['cell_requirements']=cell_requirements([m['id'] for m in prior_payload['standards']['metrics']],historical)
                else:prior_payload.pop('cell_requirements',None)
                for key in ('program_issues','semantic_questions','previous_issues','返回骨架'):
                    if key in old_task['payload']:prior_payload[key]=old_task['payload'][key]
                old_review=validate_review(read([reference],'a_semantic_review',prior_payload),tables,snapshot,context,focus)
                record_definitions(old_review,old_task)
                binding=reviewed_binding(binding,old_review,tables,snapshot);earlier.append(old_review)
            review_payload={**payload,'task':'a_semantic_review','recognized':binding}
            if item['review']:
                old_task,_=_stage_receipt(item['review'][0],receipt_mapping,'a_semantic_review')
                if 'text_evidence' in old_task['payload']:
                    review_payload.update(text_evidence=review_text_evidence(tables,snapshot,context),evidence_rule=REVIEW_EVIDENCE_RULE)
                if 'reading_notes' in old_task['payload']:review_payload['reading_notes']=deepcopy(reading.get('notes',''))
                if old_task['payload'].get('共同声明合并'):
                    review_payload.update(recognized=shared_binding(binding),共同声明合并=True)
                focus=old_task['payload'].get('review_focus')
                if 'review_candidates' in old_task['payload']:review_payload['review_candidates']=old_task['payload']['review_candidates']
                if focus is not None:review_payload=focused_review_payload(review_payload,focus,historical)
                review_payload['standards']=supplied_review_standards(review_payload['standards'],old_task['payload']['standards'],historical)
                if 'cell_requirements' in old_task['payload']:
                    review_payload['cell_requirements']=cell_requirements([m['id'] for m in review_payload['standards']['metrics']],historical)
                else:review_payload.pop('cell_requirements',None)
                if 'program_issues' in old_task['payload']:
                    # 诊断文案只是辅助提示；原始文案由任务哈希保护，业务原文、标准、身份及结果仍重新核验。
                    review_payload['program_issues']=old_task['payload']['program_issues']
                if 'semantic_questions' in old_task['payload']:review_payload['semantic_questions']=old_task['payload']['semantic_questions']
                if '返回骨架' in old_task['payload']:review_payload['返回骨架']=old_task['payload']['返回骨架']
                if 'previous_issues' in old_task['payload']:review_payload['previous_issues']=old_task['payload']['previous_issues']
            review=validate_review(read(item['review'],'a_semantic_review',
                review_payload),tables,snapshot,context,review_payload.get('review_focus')) if item['review'] else {'tables':[]}
            if item['review']:record_definitions(review,old_task)
            final_binding=reviewed_binding(binding,review,tables,snapshot)
            _,review_aliases=source_material(snapshot,context,tables)
            for row in final_binding['tables']:
                table=next(t for t in tables if t['id']==row['id'])
                try:records=checked_records(row,table,table_packet(table,snapshot,context),historical,review_aliases)
                except (ValueError,KeyError,TypeError):continue
                for address,record in records.items():
                    if record['kind'] in {'metric_fact','partial_metric_fact'}:
                        fact=record.get('fact',record.get('partial_semantics'))
                        _require(fact['metric_id'] in seen[row['id']][address],'本格身份缺少实际发送的完整标准：'+row['id']+'!'+address)
            batches.append((binding,review,earlier))
    else:_require(not proof['batches'],'尚未检索标准不能已有身份结果')
    expected=assemble(snapshot,context,historical,reading,batches,b_scope,b_outline,require_identity=mapping.get('identity_version')==1)
    if rebuild:return expected
    for key in ('facts','owners','excluded','unresolved','outline','coverage','complete'):
        _require(mapping.get(key)==expected[key],'保存结果与原语义回执不符：'+key)
    _require(mapping.get('slot_bindings',[])==expected.get('slot_bindings',[]),'保存结果与原语义回执不符：slot_bindings')
    for key in ('source_identities','b_identity_coverage' if scoped else 'a_identity_coverage'):
        if key in mapping or expected['source_identities']:
            _require(mapping.get(key)==expected[key],'保存结果与原语义回执不符：'+key)
