"""只展开模型声明的表区和行列语义，不从位置或数值推断业务。"""
from copy import deepcopy
import re

from openpyxl.utils.cell import range_boundaries, get_column_letter


def expand_structure(payload, result):
    if not isinstance(result, dict) or set(result) != {'dimension_values', 'regions', 'cells'}:
        raise ValueError('结构语义回执须包含dimension_values、regions、cells')
    if not isinstance(result['regions'], list) or not isinstance(result['cells'], dict):
        raise ValueError('结构语义表区或单格例外格式无效')
    owned = set(payload['candidate_cells'])
    if set(result['cells']) - owned:
        raise ValueError('单格例外超出候选范围')
    expanded = {}; covered = set()
    for region in result['regions']:
        if not isinstance(region, dict) or set(region) != {'range', 'record', 'rows', 'columns'}:
            raise ValueError('表区须明确范围、共同记录和行列差异')
        area = region['range']
        if not isinstance(area, str) or not re.fullmatch(r'[A-Z]{1,3}[1-9]\d*(?::[A-Z]{1,3}[1-9]\d*)?', area):
            raise ValueError('表区地址格式无效')
        left, top, right, bottom = range_boundaries(area)
        if left > right or top > bottom or (right-left+1)*(bottom-top+1) > len(owned):
            raise ValueError('表区超出候选范围')
        addresses = {f'{get_column_letter(c)}{r}' for r in range(top, bottom+1) for c in range(left, right+1)}
        if addresses - owned or addresses & covered:
            raise ValueError('表区超出候选范围或重复覆盖')
        covered.update(addresses)
        base = region['record']
        if not isinstance(base, dict): raise ValueError('表区共同记录无效')
        axes = []
        for key, allowed in [('rows', set(range(top, bottom+1))),
                             ('columns', {get_column_letter(c) for c in range(left, right+1)})]:
            parts = region[key]; lookup = {}
            if not isinstance(parts, list): raise ValueError('行列差异须为数组')
            for part in parts:
                if not isinstance(part, dict) or set(part) != {'indices', 'record'}:
                    raise ValueError('行列差异格式无效')
                indices = part['indices']; patch = part['record']
                if (not isinstance(indices, list) or not indices or
                    any(type(i) is not (int if key == 'rows' else str) for i in indices) or
                    len(set(indices)) != len(indices) or set(indices)-allowed or set(indices)&set(lookup)):
                    raise ValueError('行列差异越界或重复')
                if not isinstance(patch, dict) or set(patch)-{'metric_id', 'dimension_refs', 'missing_dimensions'}:
                    raise ValueError('行列差异只能补充指标和维度引用')
                for index in indices: lookup[index] = patch
            axes.append(lookup)
        for row in range(top, bottom+1):
            for col in range(left, right+1):
                letter = get_column_letter(col); address = f'{letter}{row}'
                if address in result['cells']: continue
                record = deepcopy(base)
                for patch in (axes[0].get(row, {}), axes[1].get(letter, {})):
                    for key, value in patch.items():
                        if key in {'dimension_refs', 'missing_dimensions'}:
                            current = record.get(key, [])
                            if not isinstance(current, list) or not isinstance(value, list):
                                raise ValueError('维度引用和缺失维度须为数组')
                            record[key] = current + deepcopy(value)
                        elif key in record:
                            raise ValueError('指标在共同记录与行列差异中重复声明')
                        else: record[key] = deepcopy(value)
                expanded[address] = record
    expanded.update(deepcopy(result['cells']))
    # 漏格继续由既有逐格校验器记录未决；共享维度仍逐条校验证据及重复引用。
    return {'dimension_values': deepcopy(result['dimension_values']), 'cells': expanded}
