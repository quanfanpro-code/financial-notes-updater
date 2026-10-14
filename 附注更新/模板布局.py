# coding: utf-8
"""复用既有行列移动器，保持槽位及原位置来源，不按金额推断对应。"""
from copy import deepcopy
from openpyxl import load_workbook
from .统一结构 import _shape, _position, _location_text, structural_context
from .表格 import _apply_operation, _definitions, same_excel_value, read_workbook, materialize_candidate_blanks, file_hash
from .统一语义 import _require, definition_hash


def relocate_mapping(mapping, layout):
    view=deepcopy(mapping);positions=layout['positions']
    original={r['source_reference']['location']:r for r in mapping['facts']}
    context=structural_context(mapping,layout)
    context_references={'@context:'+definition_hash(before):'@context:'+definition_hash(after)
                        for before,after in zip(mapping.get('source_context',[]),context)}
    removed_evidence={}
    def preserve_evidence(reference):
        if reference.startswith('@context:'):return context_references.get(reference,reference)
        moved=positions.get(reference,reference)
        if moved is not None:return moved
        if reference not in removed_evidence:
            _require(file_hash(mapping['source_path'])==mapping['source_hash'],'移除位置的原始识别依据已变化')
            book=load_workbook(mapping['source_path'],data_only=False)
            try:
                sheet,address=reference.rsplit('!',1);value=book[sheet][address].value
                _require(value is not None and str(value).strip(),'移除位置的原始识别依据为空')
            finally:book.close()
            removed_evidence[reference]=f'原模板已移除位置的识别依据（仅用于接续原识别）：{mapping["source_hash"]} / {reference} / {value}'
        return '@context:'+definition_hash(removed_evidence[reference])
    def move_reference(reference, destination, template=None):
        if reference.startswith('@context:'):return preserve_evidence(reference)
        if template and positions.get(reference) is not None:
            before=_position(template);after=_position(destination);prior=_position(reference)
            ref=list(_position(positions[reference]))
            if prior[0]==before[0]:
                # 名称跟随新行、列标题跟随新列，其余位置先按全表变化移动。
                if prior[1]==before[1]:ref[1]=after[1]
                if prior[2]==before[2]:ref[2]=after[2]
                return _location_text(ref)
        return preserve_evidence(reference)
    def move_fact(row, destination, template=None):
        row=deepcopy(row)
        row['id']=destination;row['source_reference']['location']=destination
        if row['record_type']=='metric_fact':
            row['owner_reference']=destination
            row['dimension_evidence']={key:[move_reference(ref,destination,template) for ref in refs] for key,refs in row['dimension_evidence'].items()}
        if template:row.update(raw_value=None,raw_value_state='blank')
        return row
    view['facts']=[move_fact(row,positions[location]) for location,row in original.items() if positions[location] is not None]
    for destination,item in layout['added'].items():
        if item['template'] in original:
            view['facts'].append(move_fact(original[item['template']],destination,item['template']))
    view['owners']={r['owner_reference']:deepcopy(r['dimensions']) for r in view['facts'] if r['record_type']=='metric_fact'}
    for kind in ('excluded','unresolved','slot_bindings'):
        moved=[]
        for row in view.get(kind,[]):
            location=row['sheet']+'!'+row['cell'];destination=positions.get(location,location)
            if destination is not None:
                if row.get('partial_semantics'):
                    known=row['partial_semantics']
                    known['dimension_evidence']={key:[move_reference(ref,destination) for ref in refs] for key,refs in known['dimension_evidence'].items()}
                    known['evidence_cells']=[move_reference(ref,destination) for ref in known['evidence_cells']]
                row['sheet'],row['cell']=destination.rsplit('!',1);moved.append(row)
        if kind=='slot_bindings':
            original_slots={r['sheet']+'!'+r['cell']:r for r in mapping.get(kind,[])}
            for destination,item in layout['added'].items():
                if item['template'] not in original_slots:continue
                row=deepcopy(original_slots[item['template']]);known=row['partial_semantics']
                known['dimension_evidence']={key:[move_reference(ref,destination,item['template']) for ref in refs] for key,refs in known['dimension_evidence'].items()}
                known['evidence_cells']=[move_reference(ref,destination,item['template']) for ref in known['evidence_cells']]
                row['sheet'],row['cell']=destination.rsplit('!',1);moved.append(row)
        view[kind]=moved
    view['source_context']=context+list(removed_evidence.values())
    view['_template_layout']=layout
    return view


def project_layout(mapping,source,standard,operations,source_locations=None):
    a=load_workbook(mapping['source_path'],data_only=False);b=load_workbook(source['source_path'],data_only=False)
    try:
        # 局部更新仍会移动后方整表，未保存节点的间隔空白也必须移动。
        snapshot=read_workbook(mapping['source_path'],context=mapping.get('source_context',[]))
        materialize_candidate_blanks(snapshot)
        for sheet in snapshot['sheets']:
            for address in sheet['cells']:a[sheet['name']][address]
        for kind in ('excluded','unresolved','slot_bindings'):
            for row in mapping.get(kind,[]):a[row['sheet']][row['cell']]
        selected=source
        if source_locations is not None:
            known={r['source_reference']['location'] for r in source['facts']}
            known.update(r['sheet']+'!'+r['cell'] for r in source.get('slot_bindings',[]))
            _require(source_locations and set(source_locations)<=known,'本次名单须引用已确认B来源')
            selected={**source,'facts':[r for r in source['facts'] if r['source_reference']['location'] in source_locations],
                      'slot_bindings':[r for r in source.get('slot_bindings',[]) if r['sheet']+'!'+r['cell'] in source_locations]}
        positions,added,removed=_shape(a,{'operations':operations},mapping,selected,b,standard)
    finally:a.close();b.close()
    layout={'operations':deepcopy(operations),'positions':{key:_location_text(value) if value is not None else None for key,value in positions.items()},
            'added':{_location_text(key):value for key,value in added.items()},'removed':removed}
    return relocate_mapping(mapping,layout)


def verify_output(parent, output, plan):
    """按既有写入器在内存重放，核对布局、文字、数值及样式。"""
    expected=load_workbook(parent['source_path'],data_only=False);actual=load_workbook(output,data_only=False)
    try:
        for op in plan['operations']:_apply_operation(expected,op)
        for row in [*plan.get('labels',[]),*plan['updates']]:
            cell=expected[row['sheet']][row['cell']];cell.value=row['value']
            if row.get('number_format_update'):cell.number_format=row['number_format_update']['new_format']
            if isinstance(row['value'],str):cell.data_type='s'
        _require(expected.sheetnames==actual.sheetnames,'模板更新改变了工作表清单')
        _require([(r.name,r.attr_text) for r in _definitions(expected)]==[(r.name,r.attr_text) for r in _definitions(actual)],'模板更新后的Word命名区域不符')
        for before,after in zip(expected,actual):
            _require(set(map(str,before.merged_cells.ranges))==set(map(str,after.merged_cells.ranges)),'模板更新后合并范围不符')
            for position in set(before._cells)|set(after._cells):
                left=before.cell(*position);right=after.cell(*position)
                _require(same_excel_value(right.value,left.value) and
                         (tuple(right._style) if right.has_style else None)==(tuple(left._style) if left.has_style else None),
                         '模板更新有清单外的数值或格式改动：'+before.title+'!'+left.coordinate)
    finally:expected.close();actual.close()
