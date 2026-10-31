"""以独立成果文件连接提取、A/B识别、勾稽和回写，支持复用与下一轮更新。"""
from __future__ import annotations
import copy
from collections import Counter
import hashlib
import json
import threading
import uuid
from datetime import datetime
from pathlib import Path
from .流程 import ROOT,bridge,save_json,trace_workbook,check_cancel
from .设置 import validate_settings
from .语义 import SemanticEngine,SemanticCancelled,SemanticProgressError,_normal_dimensions,recognition_progress
from .表格 import read_workbook,file_hash,write_updated_workbook,recalculate_copy,materialize_business_blanks,cell_at
from .金标准 import ensure_compatible,validate_mapping_scope,normalize_mapping_metric,read_standard
from .统一语义 import json_dates

MAP_SCHEMA="附注语义映射-v1"
LINK_SCHEMA="附注Word关联-v1"
PLAN_SCHEMA="附注更新清单-v1"
STEPS={"extract_a":"提取附注表格","recognize_a":"识别附注表格","recognize_b":"识别更新数据","match":"建立语义对应","write":"生成更新文件"}
CARRIER_NAMES={"A":"附注表格","B":"更新数据表"}
FILE_PREFIXES={"A":"附注表格","B":"更新数据"}

def _read(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"),object_pairs_hook=json_dates)

def _json(value):
    return json.dumps(value,ensure_ascii=False,sort_keys=True,separators=(",",":"),default=str)

def _path(config,key,label):
    value=str(config.get(key) or "").strip()
    if not value or not Path(value).is_file():raise ValueError("请选择存在的"+label)
    return Path(value).resolve()

def _check(path,digest,label):
    path=Path(path)
    if not path.is_file() or file_hash(path)!=digest:
        raise ValueError(label+"已变化或不存在，请重新生成相应成果")
    return path

def _gold(config,record=None):
    path=Path(config.get("gold_path") or (record or {}).get("gold_path") or "")
    if not path.is_file():raise ValueError("请选择正式语义金标准")
    digest=file_hash(path)
    if record:ensure_compatible(record,path)
    return path.resolve(),digest

def _load_link(path,a_path=None):
    item=_read(path)
    if item.get("schema")!=LINK_SCHEMA:raise ValueError("这不是本程序保存的附注表格与Word位置关联文件")
    _check(item["word_path"],item["word_hash"],"关联的Word原件")
    _check(item["linked_word"],item["linked_word_hash"],"带连接的Word")
    _check(item["a_path"],item["a_hash"],"关联的附注表格")
    if a_path and file_hash(a_path)!=item["a_hash"]:raise ValueError("附注表格与Word位置关联文件不属于当前附注表格")
    return item

def _engine(config,gold,identity,log,cancel):
    settings=validate_settings(config)
    identity=dict(identity)
    if identity.get("carrier")=="B":
        settings["layout_mode"]="cells_v1"
        identity["layout_mode"]="cells_v1"
    key=hashlib.sha256(_json(identity).encode("utf-8")).hexdigest()[:24]
    cache=Path(config.get("output_dir") or ROOT/"输出")/"语义缓存"/key
    return SemanticEngine(settings,str(gold),str(cache),log,cancel)

def _gold_snapshot(source,digest):
    """按内容保存一份识别依据，避免外部改写选择文件破坏既有映射。"""
    import shutil
    folder=ROOT/"金标准"/"识别快照";folder.mkdir(parents=True,exist_ok=True)
    target=folder/(digest+".jsonl")
    if target.exists():return _check(target,digest,"识别用金标准快照")
    temporary=folder/(digest+"_"+uuid.uuid4().hex+".tmp")
    shutil.copyfile(source,temporary)
    _check(temporary,digest,"新金标准快照")
    try:temporary.rename(target)
    except FileExistsError:
        # 并发识别相同版本时保留临时副本，核验已经建立的正式快照。
        _check(target,digest,"并发建立的金标准快照")
    return target

def _normalize_mapping(mapping,slot):
    """成果读写共用兼容规则；只在内存中规范已证明等价的表示。"""
    normalized=normalize_mapping_metric(slot,mapping)
    if "dimensions" in normalized:
        normalized["dimensions"]=_normal_dimensions(normalized["dimensions"],slot)
    return normalized


def _record(snapshot,result,carrier,gold,original_path,context,a_link=None):
    item=copy.deepcopy(result)
    definitions=read_standard(gold)
    normalized=[]
    for mapping in item.get("mappings",[]):
        slot=definitions.get(mapping.get("slot_id"))
        if not slot or slot.get("status")!="active":
            raise ValueError("语义映射引用了不在当前active金标准中的ID："+str(mapping.get("slot_id")))
        normalized.append(_normalize_mapping(mapping,slot))
    item["mappings"]=normalized
    item.update(schema=MAP_SCHEMA,carrier=carrier,source_path=snapshot["path"],
        source_hash=snapshot["sha256"],original_path=str(original_path),
        original_hash=file_hash(original_path),gold_path=str(gold),gold_hash=file_hash(gold),
        context=context,recalculated=bool(snapshot.get("recalculated")),
        recalculation=copy.deepcopy(snapshot.get("recalculation",{})))
    if a_link:item["a_link"]=str(a_link)
    item["confirmed_business_ranges"]={s["name"]:copy.deepcopy(s.get("confirmed_business_ranges",[])) for s in snapshot["sheets"]}
    return item


def _recognition_context(link):
    """正文按原顺序只带一次，已提取表用原表区引用连接。"""
    from zipfile import ZipFile
    from xml.etree import ElementTree as ET
    from openpyxl.utils import get_column_letter
    word=_check(link['word_path'],link['word_hash'],'关联的Word原件')
    with ZipFile(word) as package:root=ET.fromstring(package.read('word/document.xml'))
    ns='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    context=[{k:v for k,v in item.items() if k not in
              ('chapter_context','preceding_blocks','unit_text_candidates','context_note')}
             for item in link['context'].get('tables',[])]
    extracted={item['table_id']:item for item in context if item.get('table_id')}
    paragraphs=[]
    def read(node):
        if node.tag==ns+'tbl':
            for row in node.findall(ns+'tr'):
                cells=[' / '.join(''.join(t.text or '' for t in p.iter(ns+'t')).strip()
                                  for p in cell.iter(ns+'p')) for cell in row.findall(ns+'tc')]
                if any(cells):paragraphs.append(' | '.join(cells))
            return
        if node.tag==ns+'p':
            text=''.join(part.text or '' for part in node.iter(ns+'t')).strip()
            if text:paragraphs.append(text)
            return
        for child in node:read(child)
    for index,node in enumerate(root.find(ns+'body')):
        item=extracted.get('/word/document.xml#body/'+str(index))
        if node.tag==ns+'tbl' and item and item.get('sheet'):
            left=item['first_column'];top=item['first_row']
            area=(f'{get_column_letter(left)}{top}:'
                  f'{get_column_letter(left+item["column_count"]-1)}{top+item["row_count"]-1}')
            paragraphs.append('[原文此处为表格：'+item['sheet']+'!'+area+']')
            continue
        read(node)
    if paragraphs:
        context.append({'title':'原Word正文顺序（表格引用指向实际表区，各段适用范围由模型判断）',
                        'text':'\n'.join(paragraphs),'source_path':str(word),'source_hash':link['word_hash']})
    return context


def _mapping_review(record,snapshot,path):
    """给使用者查看的逐格清单；正式后续输入仍为语义映射文件。"""
    from openpyxl import Workbook
    from .表格 import actual_value
    gold={r.get("id"):r for line in Path(record["gold_path"]).read_text(encoding="utf-8-sig").splitlines() if line.strip() for r in [json.loads(line)]}
    book=Workbook();sheet=book.active;sheet.title="业务格语义"
    sheet.append(["工作表","单元格","源内容或公式","可核实的值","金标准ID","指标槽位名称","标准附注","标准表","期间","币种","单位","倍率","实际维度（完整记录）","判断依据"])
    for m in record["mappings"]:
        c=cell_at(snapshot,m["sheet"],m["cell"]);slot=gold.get(m["slot_id"],{});d=m.get("dimensions",{})
        try:value=actual_value(snapshot,m["sheet"],m["cell"])
        except ValueError:value="公式结果尚未核实"
        raw=c.get("value")
        sheet.append([m["sheet"],m["cell"],"空白" if raw is None else raw,"空白" if value is None else value,
            m["slot_id"],m.get("semantic_field"),(slot.get("note") or {}).get("name"),(slot.get("table") or {}).get("name"),
            _json(d.get("period")) if isinstance(d.get("period"),dict) else d.get("period"),
            d.get("currency"),d.get("unit"),d.get("scale"),_json(d),m.get("reason")])
    for title,key in (("待核实","unresolved"),("非业务格","excluded"),("本次范围外","out_of_scope")):
        ws=book.create_sheet(title);ws.append(["工作表","单元格","分类","依据或待核实原因"])
        for item in record.get(key,[]):ws.append([item.get("sheet"),item.get("cell"),item.get("category"),item.get("reason")])
    for ws in book:
        ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
        for row in ws:
            for c in row:
                if isinstance(c.value,str):c.data_type="s"
        for col,width in {"A":23,"B":13,"C":30,"D":24,"E":26,"F":42,"G":24,"H":30,"I":24,"J":14,"K":14,"L":14,"M":70,"N":65}.items():
            ws.column_dimensions[col].width=width
    with Path(path).open("xb") as stream:book.save(stream)
    book.close()
    return str(path)



def _standard_review(record,snapshot,run,structure_path):
    """保留未确认项的来源证据，交给人工或高级AI处理；不自动认定缺标准。"""
    import shutil
    from .表格 import actual_value
    carrier=record["carrier"]
    gold_copy=run/"待处理时使用的金标准.jsonl"
    shutil.copyfile(record["gold_path"],gold_copy)
    _check(gold_copy,record["gold_hash"],"待处理材料中的金标准副本")
    items=[]
    for unresolved in record.get("unresolved",[]):
        item=copy.deepcopy(unresolved)
        try:
            source=cell_at(snapshot,item.get("sheet"),item.get("cell"))
            item["source_cell"]=copy.deepcopy(source)
            try:
                item["verified_value"]=actual_value(snapshot,item["sheet"],item["cell"])
                item["value_verified"]=True
            except ValueError as error:
                item.update(value_verified=False,value_note=str(error))
        except (ValueError,KeyError,TypeError):
            item["source_note"]="涉及范围或结构判断，请按实际结构文件查看，未凭空补造单元格"
        items.append(item)
    instructions=run/"待处理语义说明.md"
    instructions.write_text(
        "# 处理未确认语义\n\n"
        "本目录保存本次识别的语义映射、实际表格结构、核对表和当时使用的金标准副本。"
        "JSON文件中的路径用于本机追溯；将材料交给外部高级AI时，应同时提供这些文件。\n\n"
        "先根据原表的表头、所在附注、口径和实际维度解释每个未确认项。"
        "依次判断：现有active标准是否已覆盖；是否只是开放维度的新取值；是否缺少文件未注明的信息；"
        "是否仍有判断分歧；最后才判断需要新增或修订金标准。不可按新客户、新期间、新排版重复建标准。\n\n"
        "如果现有标准已经覆盖，请给出具体ID、维度和单元格依据，补充必要事实后重新识别。"
        "若确需改标准，请用本目录金标准副本作为起点，保留全部旧ID和字段，"
        "将完整修订结果另存为新的JSONL文件。新增定义按原格式提供语义、口径、维度、来源依据和引用关系；"
        "修改或停用旧定义须说明原因。不得用金额、坐标或某个样式替代语义定义。\n\n"
        "人工或高级AI完成语义判断后，仅在确需修订标准时，才在程序设置的金标准管理中选择修订文件，"
        "检查新增、修订、停用清单，再保存并使用新版本。程序检查结构和引用，不能代替语义判断。"
        "新版本保存后重新识别未完成的附注表格或更新数据表；已有映射所用定义未变时可以继续复用。\n",
        encoding="utf-8-sig")
    package={
        "schema":"附注待处理语义-v1","carrier":carrier,"source_path":record["source_path"],
        "source_hash":record["source_hash"],"original_path":record["original_path"],"original_hash":record["original_hash"],
        "gold_path":str(gold_copy),"gold_hash":record["gold_hash"],"original_gold_path":record["gold_path"],
        "mapping_path":str(run/(FILE_PREFIXES[carrier]+"语义映射.json")),"structure_path":str(structure_path),
        "review_path":record["review_path"],"instructions_path":str(instructions),
        "additional_context":record.get("additional_context",""),"context":record["context"],
        "items":items,"proposed_changes":[],"notice":"未确认不等于金标准缺项，须先检查现有定义和开放维度"}
    return save_json(run/(FILE_PREFIXES[carrier]+"待处理语义与金标准补充材料.json"),package)

def _snapshot(record):
    source=_check(record["source_path"],record["source_hash"],"语义映射对应的工作簿")
    if record.get("original_path"):
        _check(record["original_path"],record["original_hash"],"识别时的原始资料")
    snapshot=read_workbook(source,context=record.get("context",[]))
    basis=record.get("recognition_basis",{})
    if basis.get("candidate_blanks"):
        if basis.get("layout_mode")!="cells_v1" or basis["candidate_blanks"]!="source_grid_v1":
            raise ValueError("不支持这份映射的空白候选识别方式")
        from .表格 import materialize_candidate_blanks
        materialize_candidate_blanks(snapshot)
    for sheet,areas in record.get("confirmed_business_ranges",{}).items():
        for area in areas:
            if area.get("source_hash")!=record["source_hash"]:
                raise ValueError("业务空白的来源证据与工作簿不一致")
            table={**copy.deepcopy(area),"range":area["table_range"]}
            if "layout_batch_id" in area or "layout_review" in area:
                actual_sheet=next((s for s in snapshot["sheets"] if s["name"]==sheet),None)
                if actual_sheet is None:raise ValueError("逐格布局引用未知工作表")
                SemanticEngine._restore_cell_table(snapshot,actual_sheet,table)
            else:
                materialize_business_blanks(snapshot,sheet,table)
    if record.get("recalculated"):
        proof=record.get("recalculation",{})
        if proof.get("method")!="Excel.CalculateFullRebuild" or proof.get("source_sha256")!=record.get("original_hash"):
            raise ValueError("公式重算缺少可核验的原文件身份")
        snapshot.update(recalculated=True,recalculation=proof)
    return snapshot

def _load_mapping(path,carrier=None,complete=True):
    record=_read(path)
    if record.get("schema")!=MAP_SCHEMA:raise ValueError("请选择本程序保存的语义映射文件")
    if carrier and record.get("carrier")!=carrier:raise ValueError("选择的语义映射不是"+CARRIER_NAMES[carrier]+"的映射")
    if complete and (not record.get("complete") or record.get("unresolved")):
        raise ValueError("语义映射还有未确认项，请先完成该表识别")
    if not isinstance(record.get("mappings"),list) or (complete and not record["mappings"]):
        raise ValueError("语义映射没有已确认的业务格")
    gold=Path(record.get("gold_path") or "")
    _check(gold,record.get("gold_hash"),"映射所用金标准")
    active={row.get("id"):row for line in gold.read_text(encoding="utf-8-sig").splitlines() if line.strip()
        for row in [json.loads(line)] if row.get("status")=="active" and row.get("id")}
    snapshot=_snapshot(record)
    selection=record.get("scope_selection")
    if not selection and (record.get("out_of_scope") or record.get("recognition_basis",{}).get("selection_hash") not in (None,SemanticEngine._selection_identity(None))):
        raise ValueError("语义映射缺少实际识别范围来源，不能只保留章外结论")
    if selection:
        if selection.get("source_hash")!=record["source_hash"] or selection.get("target_hash")!=hashlib.sha256(_json(selection.get("target_meanings")).encode("utf-8")).hexdigest():
            raise ValueError("语义映射的识别范围来源或目标含义校验失败")
        from .识别范围 import read_verified_scope
        verified_scope=read_verified_scope(selection,snapshot,selection.get("target_meanings"))
        if verified_scope is not None:
            basis=record.get("recognition_basis",{})
            if basis.get("layout_mode")!="cells_v1" or basis.get("candidate_blanks")!="source_grid_v1":
                raise ValueError("外部范围确认与映射的完整候选格依据不一致")
        expected={(s["name"],a) for s in snapshot["sheets"] for a,c in s["cells"].items() if verified_scope is not None or not c.get("derived_blank")}
        selected=[(sheet,a) for sheet,addresses in selection.get("selected_cells",{}).items() for a in addresses]
        omitted=[(area["sheet"],a) for area in selection.get("out_of_scope",[]) for a in area.get("cells",[])]
        if set(selected+omitted)!=expected or len(set(selected+omitted))!=len(selected+omitted):
            raise ValueError("语义映射的识别范围存在遗漏或重复")
        if any((m["sheet"],m["cell"]) in set(omitted) for m in record["mappings"]):
            raise ValueError("范围外数据被错误地纳入业务语义映射")
        if verified_scope is not None:
            confirmed_omitted=[dict(sheet=area["sheet"],cell=a,category="out_of_scope",reason=area["reason"],evidence_cells=area["evidence_cells"]) for area in selection["out_of_scope"] for a in area["cells"]]
            if record.get("out_of_scope",[])!=confirmed_omitted:
                raise ValueError("映射的章外清单与实际外部确认不一致")
            positions=[(item.get("sheet"),item.get("cell")) for category in ("mappings","excluded","unresolved","out_of_scope") for item in record.get(category,[])]
            if set(positions)!=expected or len(positions)!=len(set(positions)):
                raise ValueError("外部范围映射存在遗漏、重复或包外格")
    seen=set()
    for item in record["mappings"]:
        pos=(item.get("sheet"),item.get("cell"))
        if pos in seen:raise ValueError("语义映射包含重复业务格")
        seen.add(pos);cell_at(snapshot,*pos)
        slot=item.get("slot_id")
        if slot not in active:raise ValueError("语义映射引用了不在当前active金标准中的ID："+str(slot))
        validate_mapping_scope(active[slot],item)
    if any(m.get("review_method")=="external_semantic_review" or "external_review" in m for m in record["mappings"]):
        from .外部复核 import validate_external_reviews
        validate_external_reviews(record,snapshot)
    # 先完成原件、金标准与原记录校验，再规范内存副本，不改写旧成果文件。
    record["mappings"]=[_normalize_mapping(item,active[item["slot_id"]]) for item in record["mappings"]]
    return record,snapshot


def _selected_sources(config,ar,br):
    """界面当前选择与落盘成果必须指向相同资料，空路径允许纯成果文件复用。"""
    for key,record,label in (("a_path",ar,"A"),("b_path",br,"B")):
        if str(config.get(key) or "").strip():
            source=_path(config,key,CARRIER_NAMES[label])
            if file_hash(source) not in {record.get("source_hash"),record.get("original_hash")}:
                raise ValueError("当前选择的"+CARRIER_NAMES[label]+"与已保存语义映射不是同一份资料，请重新识别或选择对应映射")


def _input_records(paths):
    return {str(Path(path).resolve()):file_hash(path) for path in paths if path}

def _verify_inputs(inputs):
    for path,digest in inputs.items():_check(path,digest,"该步骤依赖的文件")

def _extract(config,run,log,cancel):
    word=_path(config,"word_path","原始附注 Word")
    if word.suffix.lower()!=".docx":raise ValueError("原始附注 Word需要为.docx")
    before=file_hash(word);a=run/"附注表格.xlsx";linked=run/"原始附注_带连接.docx";context=run/"Word表格上下文.json"
    log("机械提取Word表格及位置关联；参与更新的范围由后续已有范围识别确定")
    result=bridge({"action":"extract","word":str(word),"excel":str(a),"linked_word":str(linked),"context":str(context)},run,cancel)
    _check(word,before,"原始附注 Word")
    link={"schema":LINK_SCHEMA,"word_path":str(word),"word_hash":before,"linked_word":str(linked),
        "linked_word_hash":file_hash(linked),"a_path":str(a),"a_hash":file_hash(a),"context":_read(context)}
    link["scope"]=link["context"].get("scope") or {}
    link_path=save_json(run/"附注表格与Word位置关联.json",link)
    return {"status":"complete","message":"原始表格及位置关联已保存；下一步由模型确定处理范围并识别语义。",
        "files":{"a_path":str(a),"a_link":link_path},"table_count":result["table_count"]}


def _apply_reviewed_existing(result,learned,engine,snapshot):
    """独立复核已确认现有槽位时接入原成果；不改旧结果、原数据或金标准。"""
    proposed=learned.get("reviewed_existing_mappings") or []
    if not proposed:return result
    report=_check(learned.get("report_path",""),learned.get("report_sha256"),"已有语义复核报告")
    saved=_read(report)
    if _json(saved)!=_json({k:v for k,v in learned.items() if k not in {"report_path","report_sha256"}}):
        raise ValueError("已有语义复核意见与已保存报告不一致")
    if learned.get("source_hash")!=snapshot["sha256"] or learned.get("gold_hash")!=engine.gold_hash:
        raise ValueError("已有语义复核的原表或金标准版本不一致")
    pending={(m.get("sheet"),m.get("cell")) for m in result.get("unresolved",[])}
    accepted=[];seen=set();sheets={s["name"]:s for s in snapshot["sheets"]}
    for raw in proposed:
        pos=(raw.get("sheet"),raw.get("cell"))
        if pos not in pending or pos in seen:raise ValueError("已有语义复核引用非待确认格或出现重复格")
        seen.add(pos)
        review=raw.get("semantic_review") or {}
        second=review.get("independent_review") or {}
        if (raw.get("reviewed") is not True or raw.get("review_method")!="existing_slot_proposal_and_independent_review"
                or not review.get("proposal") or second.get("accepted") is not True or not str(second.get("reason") or "").strip()):
            raise ValueError("已有语义缺少独立复核通过的真实依据")
        sheet=sheets.get(pos[0])
        if sheet is None:raise ValueError("已有语义复核引用未知工作表")
        areas={_json(a):a for a in sheet.get("confirmed_business_ranges",[])+result.get("confirmed_business_ranges",{}).get(pos[0],[])
               if a.get("source_hash")==snapshot["sha256"] and a.get("table_range")==raw.get("table_range")}
        if len(areas)!=1:raise ValueError("已有语义复核缺少唯一原业务范围")
        area=next(iter(areas.values()));table={**area,"range":area["table_range"]}
        engine._validate_layout({"tables":[table]},sheet,[])
        from .语义 import _inside,_range
        if not _inside(pos[1],_range(table["range"])):raise ValueError("已有语义复核引用原业务范围外单元格")
        for opinion in (review["proposal"],second):
            evidence=opinion.get("evidence_cells")
            if not isinstance(evidence,list) or not evidence or any(a not in sheet["cells"] for a in evidence):
                raise ValueError("已有语义复核缺少可回读的原格证据")
        allowed={s["id"] for s in engine._candidate_slots(table)}
        checked=engine._validate_classification({"mappings":[raw],"excluded":[],"unresolved":[]},[pos[1]],allowed,sheet,table)[pos[1]][1]
        checked.update(reviewed=True,review_method=raw["review_method"],semantic_review=copy.deepcopy(review))
        accepted.append(checked)
    updated=copy.deepcopy(result)
    updated["mappings"].extend(accepted)
    updated["unresolved"]=[m for m in updated["unresolved"] if (m["sheet"],m["cell"]) not in seen]
    expected={(s["name"],a) for s in snapshot["sheets"] for a in s["cells"]}
    positions=[(i["sheet"],i["cell"]) for k in ("mappings","excluded","unresolved","out_of_scope") for i in updated.get(k,[])]
    if set(positions)!=expected or len(positions)!=len(set(positions)):
        raise ValueError("已有语义复核后存在遗漏或重复，未采用")
    updated["complete"]=not updated["unresolved"]
    updated.setdefault("existing_semantic_reviews",[]).append({"path":str(report),"sha256":learned["report_sha256"],"accepted_cells":len(accepted)})
    # 新证据另存；后续新增金标准接续时仍可核实完整原结果，不改写旧模型回执。
    updated.pop("evidence_path",None);updated.pop("evidence_hash",None)
    proof=engine.work_dir/("识别结果_已有语义复核_"+uuid.uuid4().hex+".json")
    engine._save(proof,updated)
    updated.update(evidence_path=str(proof.resolve()),evidence_hash=file_hash(proof))
    return updated


def _check_b_mapping_mode(config, path):
    """范围映射不能冒充完整来源识别成果；原文件保持不变。"""
    if not config.get("scope_prefilter", False) and path:
        record = _read(path)
        if record.get("scope_selection") or record.get("scope_selection_path"):
            raise ValueError("所选更新数据映射曾按附注筛选，不能作为完整更新数据映射。请在第 3 步清空旧映射选择后重新识别；需要沿用原筛选方式时，先勾选“先按附注内容筛选来源（可选）”。")


def _recognize_v2(config, carrier, run, log, cancel, progress=None):
    from .统一识别 import UnifiedSemanticEngine
    key = 'a_path' if carrier == 'A' else 'b_path'
    original = _path(config, key, CARRIER_NAMES[carrier])
    gold, digest = _gold(config)
    context = []
    metadata = {'carrier': carrier, 'original_path': str(original), 'original_hash': file_hash(original)}
    for name in ('update_purpose','comparison_requirements'):
        if config.get(name):metadata[name]=copy.deepcopy(config[name])
    if carrier=='B':
        for name in ('b_source_ranges','a_required_ranges'):
            if config.get(name):metadata[name]=copy.deepcopy(config[name])
    for name in ('field_scope','withdrawn','source_targets','source_candidates','arranged_targets','source_images'):
        if config.get(carrier.lower()+'_'+name):
            metadata[name]=copy.deepcopy(config[carrier.lower()+'_'+name])
    if carrier == 'A' and config.get('a_link'):
        link_path = _path(config, 'a_link', '附注表格与Word位置关联文件')
        context.extend(_recognition_context(_load_link(link_path, original)))
        metadata['a_link'] = str(link_path)
    additional = str(config.get('a_context_text' if carrier == 'A' else 'b_context_text', '') or '').strip()
    metadata['additional_context'] = '\n'.join(filter(None, (additional,
        str(config.get(carrier.lower()+'_additional_context', '') or '').strip())))
    if additional:
        context.append(additional)
    identity = hashlib.sha256(_json({'source': file_hash(original), 'gold': digest, 'context': context, 'carrier': carrier}).encode('utf-8')).hexdigest()
    engine = UnifiedSemanticEngine(validate_settings(config), gold,
        Path(config.get('output_dir') or ROOT / '输出') / '直接识别缓存' / identity, log, cancel)
    target = None
    if carrier == 'B' and config.get('a_mapping'):
        from .统一勾稽 import load_fact_mapping
        target = load_fact_mapping(_path(config, 'a_mapping', 'A识别结果'), engine.standard)
        if target.get('carrier') != 'A':
            raise ValueError('B查找需要必须来自A结果')
        if config.get('a_path'):
            _check(config['a_path'], target.get('source_hash'), 'A原件')
    if carrier=='A' and config.get('b_mapping') and config.get('b_source_ranges'):
        from .统一勾稽 import load_fact_mapping
        from .A语义阅读 import direct_reference
        try:
            reference_path=_path(config,'b_mapping','已保存B识别结果')
            reference=load_fact_mapping(reference_path,engine.standard)
            if reference.get('carrier')!='B':raise ValueError('身份表达参照不是B结果')
            if config.get('b_path'):_check(config['b_path'],reference['source_hash'],'B原件')
            saved=direct_reference(reference,compact=True,ranges=config['b_source_ranges'])
            metadata['既有另一方身份参照']={'载体':'B','path':str(reference_path),
                'sha256':file_hash(reference_path),'source_ranges':copy.deepcopy(config['b_source_ranges']),
                '原样身份':{key:saved[key] for key in ('source_hash','gold_hash','需要','outline') if key in saved}}
        except (ValueError,KeyError,TypeError,OSError) as error:
            log('已明确原表的B身份参照不能复用，A继续按本方原文识别：'+str(error))
    mapping_key = 'a_mapping' if carrier == 'A' else 'b_mapping'
    previous_result=config.get(mapping_key)
    if carrier=='A' and config.get('a_prime_mapping'):
        selected=_read(previous_result) if previous_result else {}
        execution=_read(config['a_prime_mapping'])
        if (selected.get('source_hash')!=metadata['original_hash'] and
                execution.get('output_hash')==metadata['original_hash']):
            previous_result=config['a_prime_mapping']
    continuing=bool(previous_result and _read(previous_result).get('schema')=='附注更新执行记录-v1')
    if carrier=='B' and config.get('list_dimension'):
        from .A语义阅读 import recognize_list_relations
        from .统一勾稽 import load_fact_mapping
        if target is None:raise ValueError('B名单字段识别须沿用已有A合集')
        previous=load_fact_mapping(_path(config,'b_mapping','B识别结果'),engine.standard)
        log('逐格识别已授权的B名单身份字段，其他已有结果复用')
        return recognize_list_relations(engine,target,previous,config,run)
    if not continuing and (carrier=='A' or target is not None) and not any(config.get(carrier.lower()+'_'+key) for key in ('target_table','field_scope','source_targets','source_ranges','allowed_metrics')):
        from .正常入口识别 import recognize_document
        return recognize_document(engine,original,context,metadata,target,run,log,progress,
            previous_result=previous_result,scope_answers=config.get(carrier.lower()+'_scope_answers'))
    log(carrier + '：直接识别并保存原答复')
    result = engine.recognize_workbook(original, context=context, artifact_metadata=metadata,
        target_mapping=target, previous_result=previous_result,
        allowed_metrics=config.get(carrier.lower() + '_allowed_metrics'),
        target_table=config.get(carrier.lower() + '_target_table'))
    path = Path(result['mapping_path'])
    answer_path = run / (carrier + '识别答复.txt')
    answer_path.write_text(result.get('本次逐格原答复',result.get('本次关系答复',result.get('答复接入来源',{}).get('本次原答复',result['识别答复']))), encoding='utf-8-sig')
    from .A合集 import direct_recognition_status
    status=direct_recognition_status(result) if result.get('collection_version')==2 else 'saved'
    return {'status': status,
            'message': carrier + ('本次识别字段已完整接入并保存。' if status=='complete' else '原答复和已接入字段已保存；识别尚未完整接入，缺交或未确定字段保留，未重答。'),
            'files': {mapping_key: str(path), 'recognition_answer': str(answer_path)}, 'usage': result['usage']}


def _recognize(config, carrier, run, log, cancel, progress=None):
    return _recognize_v2(config, carrier, run, log, cancel, progress)

def _match_v2(config,run,log,cancel,ap,bp):
    """匹配入口保存搬运安排，写入阶段不再补做对应。"""
    return _prepare_update_v2(config,run,log,cancel,ap,bp)


def _prepare_update_v2(config,run,log,cancel,ap,bp):
    """新版独立映射先接入同一个勾稽入口；范围及结构未验收时明确保留待处理。"""
    from .统一语义 import read_standard_v2
    from .统一勾稽 import load_fact_mapping
    raw=_read(ap)
    if raw.get('collection_version')==2:
        from .A合集 import _standard
        gold=Path(raw['gold_path']);digest=raw['gold_hash'];standard=_standard(raw)
    else:
        gold,digest=_gold(config)
        standard=read_standard_v2(gold)
    ar=load_fact_mapping(ap,standard); br=load_fact_mapping(bp,standard)
    if any(record.get('collection_version')!=2 for record in (ar,br)):
        raise ValueError('匹配只使用A、B两个已保存身份证合集；历史映射保留，不执行旧位置取数流程')
    if ar.get('collection_version')==2 or br.get('collection_version')==2:
        from .模板更新编排 import direct_update_plan
        for key,record in (('a_path',ar),('b_path',br)):
            if config.get(key):_check(config[key],record['source_hash'],CARRIER_NAMES['A' if key=='a_path' else 'B'])
            _check(record['source_path'],record['source_hash'],'语义原件')
        plan=direct_update_plan(ar,br,config.get('update_purpose'),config.get('validation_scope'),
                                comparison_requirements=config.get('comparison_requirements'))
        paths=[ap,bp,ar['source_path'],br['source_path']]
        for record in (ar,br):
            for reference in record.get('副本接续',{}).get('执行记录',[]):
                _check(reference['path'],reference['sha256'],'A副本接续记录')
                paths.append(reference['path'])
            evidence=record.get('calculation_evidence')
            if evidence:
                paths.extend([evidence['path'],_read(evidence['path'])['workbook_path']])
        selected_link=config.get('a_link',ar.get('a_link'))
        if selected_link:
            link=_load_link(selected_link,ar['source_path'])
            paths.extend([selected_link,link['word_path'],link['linked_word']])
        plan.update(schema='附注更新清单-v2',inputs=_input_records(paths),a_mapping=str(ap),b_mapping=str(bp),
                    gold_path=str(gold),a_link=selected_link,word_operations=copy.deepcopy(plan['operations']))
        if config.get('validation_scope'):plan['validation_scope']=copy.deepcopy(config['validation_scope'])
        path=save_json(run/'更新清单.json',plan)
        save_json(run/'更新清单校验.json',{'path':path,'sha256':file_hash(path)})
        return {'status':'complete','message':'两合集的对应、搬运及名单安排已保存；未明确内容原样保留。',
                'files':{'update_plan':path},'usage':{}}


def _match(config,run,log,cancel):
    ap=_path(config,"a_mapping","附注表格语义映射文件");bp=_path(config,"b_mapping","更新数据语义映射文件")
    versions=[_read(path).get('schema_version') for path in (ap,bp)]
    if versions!=[2,2]:
        raise ValueError('匹配只使用A、B两个已保存身份证合集；历史映射保留，不执行旧对应流程')
    return _match_v2(config,run,log,cancel,ap,bp)

def _write_v2(config,run,log,cancel,path,plan):
    """完整安排只读取其绑定材料，不重新识别或重算语义对应。"""
    # 旧同名标记曾接受模型额外编排，不能直接执行其保存结果。
    if plan.get('matching_method')!='unique_exact_identity_v4':
        raise ValueError('此清单不属于当前唯一身份证匹配结果；请从已保存的两个合集重新匹配，识别答复无需重做')
    records=[]
    for key in ('a_mapping','b_mapping'):
        mapping_path=str(Path(plan[key]).resolve())
        record=_read(mapping_path)
        record.update(mapping_path=mapping_path,mapping_hash=plan['inputs'][mapping_path])
        records.append(record)
    return _write_direct(config,run,log,cancel,path,plan,*records)


def _write_direct(config,run,log,cancel,path,plan,ar,br):
    from .A合集 import direct_output_record, validate_direct_collection
    _verify_inputs(plan['inputs'])
    for mapping in (ar,br):validate_direct_collection(mapping,for_matching=True)
    plan_hash=file_hash(path)
    for key,record in (('a_path',ar),('b_path',br),('a_mapping',ar),('b_mapping',br)):
        if config.get(key):_check(config[key],record['mapping_hash'] if key.endswith('mapping') else record['source_hash'],'当前选择与更新安排')
    working=plan.get('working_a') or ar
    if plan.get('structure'):
        working=_read(plan['structure']['mapping_path'])
    link_path=config.get('a_link',plan.get('a_link'))
    link=None
    if link_path:
        if not plan.get('a_link') or Path(link_path).resolve()!=Path(plan['a_link']).resolve():
            raise ValueError('所选Word位置关联与保存的安排不对应')
        link=_load_link(link_path,ar['source_path'])
    check_cancel(cancel)
    prime=run/('更新后附注表格'+Path(working['source_path']).suffix)
    write_updated_workbook(working['source_path'],prime,plan)
    files={'a_prime':str(prime),'update_plan':str(path)}
    if link:
        word=run/'更新后附注.docx'
        chapter=(link.get('scope') or link.get('context',{}).get('scope') or {}).get('chapter_title')
        scope={'chapter_title':chapter} if chapter else {}
        result=bridge({'action':'update','word':link['linked_word'],'excel':str(prime),'output':str(word),
                       'operations':plan.get('word_operations',plan.get('operations',[])),
                       'caption_updates':plan.get('caption_updates',[]),**scope},run,cancel)
        files['word']=str(word)
        context=bridge({'action':'read','word':str(word),**scope},run,cancel)
        new_link={'schema':LINK_SCHEMA,'word_path':str(word),'word_hash':file_hash(word),'linked_word':str(word),
                  'linked_word_hash':file_hash(word),'a_path':str(prime),'a_hash':file_hash(prime),'context':context,'scope':context.get('scope',{})}
        files['word_link']=save_json(run/'更新后附注表格与Word位置关联.json',new_link)
    record=direct_output_record(ar,br,prime,plan,str(path))
    if files.get('word_link'):record['a_link']=files['word_link']
    files['a_prime_mapping']=save_json(run/'更新执行记录.json',record)
    if ar.get('collection_version')!=2:
        trace=run/'本次更新执行记录.xlsx';trace_workbook(trace,plan);files['trace']=str(trace)
    _verify_inputs(plan['inputs'])
    _check(path,plan_hash,'本次更新清单')
    return {'status':'complete','message':'已按保存的安排生成副本；未安排内容继续保留。',
            'files':files,'updated_cells':sum(not row.get('unprovided_reason') for row in plan.get('updates',[])),
            'unprovided_cells':len(plan.get('unprovided_fields',[])),
            'retained_blank_cells':len(plan.get('retained_blanks',[])),
            'unused_source_cells':len(plan.get('unused_sources',[])),'usage':{}}


def _write(config,run,log,cancel):
    path=_path(config,"update_plan","更新清单")
    seal=path.with_name("更新清单校验.json")
    if not seal.is_file() or _read(seal).get("sha256")!=file_hash(path):
        raise ValueError("更新清单缺少校验记录或已被改动，请重新生成清单")
    plan=_read(path)
    if plan.get('schema')=='附注更新清单-v2':return _write_v2(config,run,log,cancel,path,plan)
    if plan.get("schema")!=PLAN_SCHEMA:
        raise ValueError("所选文件没有完整搬运安排；历史候选记录只供查看，不在写入时启动旧流程")
    return _write_v2(config,run,log,cancel,path,plan)


def run_step(step,config,log=None,cancel=None,progress=None):
    log=log or (lambda message:None);cancel=cancel or threading.Event();run=None
    latest_progress={}
    def publish(event):
        latest_progress.clear();latest_progress.update(copy.deepcopy(event))
        if progress is not None:progress(copy.deepcopy(event))
    try:
        check_cancel(cancel)
        if step not in STEPS:raise ValueError("未知流程步骤")
        output=Path(config.get("output_dir") or ROOT/"输出").resolve()
        output.mkdir(parents=True,exist_ok=True)
        run=output/(STEPS[step]+"_"+datetime.now().strftime("%Y%m%d_%H%M%S")+"_"+uuid.uuid4().hex[:6])
        run.mkdir()
        if step=="extract_a":result=_extract(config,run,log,cancel)
        elif step=="recognize_a":result=_recognize(config,"A",run,log,cancel,progress=publish)
        elif step=="recognize_b":result=_recognize(config,"B",run,log,cancel,progress=publish)
        elif step=="match":result=_match(config,run,log,cancel)
        else:result=_write(config,run,log,cancel)
        result.update(step=step,output_dir=str(run))
    except SemanticCancelled as error:
        result={"status":"cancelled","step":step,"message":str(error),"files":copy.deepcopy(latest_progress.get("files",{})),"output_dir":str(run or config.get("output_dir") or ROOT/"输出")}
    except Exception as error:
        message=f"{type(error).__name__}：{error}";secret=str(config.get("api_key") or "")
        if secret:message=message.replace(secret,"[密钥已隐藏]")
        result={"status":"failed","step":step,"message":message,"files":copy.deepcopy(latest_progress.get("files",{})),"output_dir":str(run or config.get("output_dir") or ROOT/"输出")}
    if run:save_json(run/"步骤结果.json",result)
    log(result["message"])
    return result
