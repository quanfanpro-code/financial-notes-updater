"""可追溯 Excel 搬运；结构变更仅接受可验证的普通 A1 引用。"""
from __future__ import annotations
import ast
import copy
import datetime
import hashlib
import json
import re
import unicodedata
from collections import defaultdict
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.workbook.properties import CalcProperties
from openpyxl.cell.cell import MergedCell
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import get_column_letter, column_index_from_string, range_boundaries, quote_sheetname

_RECALCULATED = {}
_CURRENCY = {"RMB":"CNY", "人民币":"CNY", "人民币元":"CNY", "美元":"USD", "港元":"HKD", "港币":"HKD", "欧元":"EUR"}
_REFERENCE = re.compile(r"^(?:(?P<sheet>'(?:[^']|'')+'|[^!]+)!)?(?P<a>\$?[A-Z]{1,3}\$?[1-9]\d*)(?::(?P<b>\$?[A-Z]{1,3}\$?[1-9]\d*))?$",re.I)

class WorkbookError(ValueError):
    """不能安全读取或更新。"""

def _close_workbook(workbook):
    """Workbook.close 不释放 keep_vba 的内存归档，须在缓冲区回收前关闭。"""
    if workbook is None:return
    try:workbook.close()
    finally:
        archive=getattr(workbook,"vba_archive",None)
        if archive is not None:archive.close()

def file_hash(path):
    h=hashlib.sha256()
    with Path(path).open("rb") as f:
        for block in iter(lambda:f.read(1024*1024),b""): h.update(block)
    return h.hexdigest()

def _definitions(w):
    yield from w.defined_names.values()
    for s in w: yield from s.defined_names.values()

def _destinations(d):
    try:return list(d.destinations)
    except (TypeError,ValueError,AttributeError):return []

def _external(w):
    result=["外部工作簿链接"] if getattr(w,"_external_links",[]) else []
    for s in w:
        for c in s._cells.values():
            if c.data_type=="f" and (re.search(r"\[[^]]+\][^+*/(),;]*!",str(c.value)) or "|" in str(c.value) or re.search(r"\b(?:WEBSERVICE|RTD|CUBE\w*)\s*\(",str(c.value),re.I)):
                result.append(f"{s.title}!{c.coordinate}")
    return result

def _word_omissions(context, workbook):
    """桥接声明的空位在 Word 中没有单元格，不能作为财务空白识别。"""
    tables=context.get("tables",[]) if isinstance(context,dict) else context
    if not isinstance(tables,(list,tuple)):return {}
    result={}
    for table in tables:
        if not isinstance(table,dict) or "structural_omissions" not in table:continue
        entries=table["structural_omissions"]
        if not isinstance(entries,list):raise WorkbookError("Word 结构省略信息必须为列表")
        for item in entries:
            if not isinstance(item,dict):raise WorkbookError("Word 结构省略信息无效")
            sheet=item.get("sheet");area=str(item.get("range","")).replace("$","").upper()
            reason=item.get("reason")
            if not isinstance(sheet,str) or sheet not in workbook.sheetnames:
                raise WorkbookError("Word 结构省略信息引用未知工作表")
            if table.get("sheet") not in (None,sheet):
                raise WorkbookError("Word 结构省略位置与所属表格不一致")
            if not isinstance(reason,str) or not reason.strip():
                raise WorkbookError("Word 结构省略信息缺少原因")
            try:left,top,right,bottom=range_boundaries(area)
            except (ValueError,TypeError):raise WorkbookError("Word 结构省略范围无效") from None
            if not all(type(v) is int for v in (left,top,right,bottom)) or not (1<=left<=right<=16384 and 1<=top<=bottom<=1048576):
                raise WorkbookError("Word 结构省略范围无效")
            record={"sheet":sheet,"range":area,"reason":reason}
            records=result.setdefault(sheet,[])
            if record not in records:records.append(record)
    return result

def _word_table_ranges(context, workbook, names):
    """Word 表边界只是载体事实；当前命名区域优先于提取时的位置。"""
    tables=context.get("tables",[]) if isinstance(context,dict) else context
    if not isinstance(tables,(list,tuple)):return None
    tables=[t for t in tables if isinstance(t,dict) and t.get("table_id") and t.get("range_name") and t.get("sheet")]
    if not tables:return None
    result={}
    defined_names={d.name for d in _definitions(workbook)}
    for table in tables:
        sheet=table["sheet"];name=table["range_name"]
        if sheet not in workbook.sheetnames:raise WorkbookError("Word 表格范围引用未知工作表")
        if name in defined_names:
            areas={e["range"] for e in names.get(name,[]) if e["sheet"]==sheet}
            if len(areas)!=1:raise WorkbookError("Word 表格命名区域缺失、歧义或与工作表不一致")
            area=next(iter(areas)).replace("$","").upper()
            try:left,top,right,bottom=range_boundaries(area)
            except (ValueError,TypeError):raise WorkbookError("Word 表格命名区域不是有效矩形") from None
        else:
            top=table.get("first_row");left=table.get("first_column")
            rows=table.get("row_count");cols=table.get("column_count")
            if not all(type(v) is int and v>0 for v in (top,left,rows,cols)):
                raise WorkbookError("Word 表格缺少当前命名区域及有效的原始位置")
            right=left+cols-1;bottom=top+rows-1
        if not all(type(v) is int for v in (left,top,right,bottom)) or not (1<=left<=right<=16384 and 1<=top<=bottom<=1048576):
            raise WorkbookError("Word 表格物理范围无效")
        if (right-left+1)*(bottom-top+1)>500000:raise WorkbookError("Word 表格物理范围超过50万格，请核实位置关联")
        area=f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}"
        record={"table_id":table["table_id"],"range_name":name,"range":area}
        if record not in result.setdefault(sheet,[]):result[sheet].append(record)
    return result


def read_workbook(path,context=None):
    path=Path(path).resolve()
    before=file_hash(path)
    w=load_workbook(path,data_only=False,keep_vba=path.suffix.lower()==".xlsm",keep_links=True)
    cache=None
    out={"path":str(path),"sha256":before,"sheets":[],"names":{},"context":context if context is not None else []}
    try:
        omissions_by_sheet=_word_omissions(out["context"],w)
        cache=load_workbook(path,data_only=True,keep_links=True)
        for d in _definitions(w):
            for sheet,address in _destinations(d):out["names"].setdefault(d.name,[]).append({"sheet":sheet,"range":address})
        word_ranges=_word_table_ranges(out["context"],w,out["names"])
        for s in w:
            physical_boxes=[range_boundaries(item["range"]) for item in (word_ranges or {}).get(s.title,[])]
            omissions=omissions_by_sheet.get(s.title,[])
            omission_boxes=[range_boundaries(item["range"]) for item in omissions]
            def omitted(c):return any(l<=c.column<=r and t<=c.row<=b for l,t,r,b in omission_boxes)
            for m in s.merged_cells.ranges:
                if any(l<=m.max_col and r>=m.min_col and t<=m.max_row and b>=m.min_row for l,t,r,b in omission_boxes):
                    raise WorkbookError(f"{s.title}!{m} 与 Word 结构省略范围相交")
            for c in s._cells.values():
                if (c.value is not None or c.data_type=="f") and omitted(c):
                    raise WorkbookError(f"{s.title}!{c.coordinate} 在 Word 结构省略位置存在实际值或公式")
            anchors=[(c.row,c.column) for c in s._cells.values() if not isinstance(c,MergedCell) and c.value is not None]
            boxes=[]
            for entries in out["names"].values():
                for entry in entries:
                    if entry["sheet"]==s.title:
                        try:
                            box=range_boundaries(entry["range"])
                            if all(isinstance(x,int) for x in box):boxes.append(box)
                        except ValueError:pass
            for m in s.merged_cells.ranges:
                if s.cell(m.min_row,m.min_col).value is not None:boxes.append((m.min_col,m.min_row,m.max_col,m.max_row))
            max_row=max([r for r,c in anchors]+[b[3] for b in boxes+physical_boxes]+[1])
            max_col=max([c for r,c in anchors]+[b[2] for b in boxes+physical_boxes]+[1])
            hidden=[(d.min or column_index_from_string(k),d.max or column_index_from_string(k)) for k,d in s.column_dimensions.items() if d.hidden]
            def visible(c):return not isinstance(c,MergedCell) and not omitted(c)
            cells={}
            for c in list(s._cells.values()):
                if not visible(c):continue
                if c.value is None and (not c.has_style or c.row>max_row or c.column>max_col):continue
                formula=str(c.value) if c.data_type=="f" else None
                cells[c.coordinate]={"row":c.row,"column":c.column,"value":c.value,"formula":formula,"cached_value":cache[s.title][c.coordinate].value if formula else c.value,"number_format":c.number_format,"style_id":c.style_id}
            # 仅展开确定的表格范围，不按污染后的 worksheet.max_row 扫描。
            for left,top,right,bottom in boxes:
                if (right-left+1)*(bottom-top+1)>500000:raise WorkbookError("命名业务区域超过50万格，请核实是否包含无用空白格式")
                for row in range(top,bottom+1):
                    for col in range(left,right+1):
                        c=s.cell(row,col)
                        if visible(c) and c.coordinate not in cells:
                            cells[c.coordinate]={"row":row,"column":col,"value":c.value,"formula":None,"cached_value":c.value,"number_format":c.number_format,"style_id":c.style_id}
            for c in cells.values():
                c["hidden_row"]=bool(s.row_dimensions.get(c["row"]) and s.row_dimensions[c["row"]].hidden)
                c["hidden_column"]=any(lo<=c["column"]<=hi for lo,hi in hidden)
                c["hidden_sheet"]=s.sheet_state!="visible"
                c["hidden"]=c["hidden_row"] or c["hidden_column"] or c["hidden_sheet"]
            out["sheets"].append({"name":s.title,"hidden":s.sheet_state!="visible","hidden_rows":[r for r,d in s.row_dimensions.items() if d.hidden],"hidden_columns":hidden,"max_row":max_row,"max_column":max_col,"merges":[str(m) for m in s.merged_cells.ranges],"cells":cells})
            if omissions:out["sheets"][-1]["structural_omissions"]=copy.deepcopy(omissions)
            if word_ranges is not None:out["sheets"][-1]["word_table_ranges"]=copy.deepcopy(word_ranges.get(s.title,[]))
        out["external_references"]=_external(w)
        if out["sha256"] in _RECALCULATED:out.update(recalculated=True,recalculation=_RECALCULATED[out["sha256"]])
        if file_hash(path)!=before:raise WorkbookError("来源文件在读取期间发生变化，请重新读取")
        return out
    finally:
        try:_close_workbook(cache)
        finally:_close_workbook(w)


def materialize_candidate_blanks(snapshot):
    """按来源真实边界补入尚待分类的空位置，不推定业务范围或空白含义。"""
    source_hash=snapshot.get("sha256")
    if not isinstance(source_hash,str) or not re.fullmatch(r"[0-9a-fA-F]{64}",source_hash):
        raise WorkbookError("候选空白缺少可核验的来源工作簿哈希")
    prepared=[]
    for sheet in snapshot.get("sheets",[]):
        bottom=sheet.get("max_row");right=sheet.get("max_column")
        if type(bottom) is not int or type(right) is not int or not (1<=bottom<=1048576 and 1<=right<=16384):
            raise WorkbookError("候选空白的实际内容边界无效")
        if bottom*right>500000:
            raise WorkbookError("候选范围超过50万格，请核实来源实际边界；未完成空白覆盖")
        merges=[range_boundaries(area) for area in sheet.get("merges",[])]
        omissions=[range_boundaries(item["range"]) for item in sheet.get("structural_omissions",[])]
        physical=[range_boundaries(item["range"]) for item in sheet.get("word_table_ranges",[])]
        prepared.append((sheet,bottom,right,merges,omissions,physical))
    added={}
    for sheet,bottom,right,merges,omissions,physical in prepared:
        cells=sheet["cells"]
        if not cells and not physical:continue  # 空工作表默认的 A1 边界不证明存在业务表。
        bounds=f"A1:{get_column_letter(right)}{bottom}"
        hidden_rows=set(sheet.get("hidden_rows",[]));hidden_columns=sheet.get("hidden_columns",[])
        for row in range(1,bottom+1):
            for column in range(1,right+1):
                address=f"{get_column_letter(column)}{row}"
                if address in cells:continue
                if any(l<=column<=r and t<=row<=b for l,t,r,b in omissions):continue
                if "word_table_ranges" in sheet and not any(l<=column<=r and t<=row<=b for l,t,r,b in physical):continue
                if any(l<=column<=r and t<=row<=b and (column,row)!=(l,t) for l,t,r,b in merges):continue
                row_hidden=row in hidden_rows
                column_hidden=any(l<=column<=r for l,r in hidden_columns)
                sheet_hidden=bool(sheet.get("hidden"))
                cells[address]={"row":row,"column":column,"value":None,"formula":None,"cached_value":None,
                    "number_format":"General","style_id":0,"hidden_row":row_hidden,"hidden_column":column_hidden,
                    "hidden_sheet":sheet_hidden,"hidden":row_hidden or column_hidden or sheet_hidden,
                    "derived_blank":True,"blank_evidence":{"kind":"source_grid","source_hash":source_hash,
                        "sheet":sheet["name"],"cell":address,"bounds":bounds,
                        "reason":"原工作簿实际内容边界内未列出的空位置；无值、无公式，尚未判断业务含义"}}
                added.setdefault(sheet["name"],[]).append(address)
    return added


def materialize_business_blanks(snapshot, sheet_name, table):
    """仅把已核验业务范围内真实缺失的格补为候选，不替代语义分类。"""
    sheet = next((s for s in snapshot["sheets"] if s["name"] == sheet_name), None)
    if sheet is None:
        raise WorkbookError("补格引用未知工作表")
    area = str(table.get("range", "")).replace("$", "").upper()
    try:
        left, top, right, bottom = range_boundaries(area)
    except (ValueError, TypeError):
        raise WorkbookError("业务补格范围无效") from None
    if not all(type(v) is int for v in (left, top, right, bottom)) or not (1 <= left <= right <= sheet["max_column"] and 1 <= top <= bottom <= sheet["max_row"]):
        raise WorkbookError("业务补格范围超出已读取的实际边界")
    if not str(table.get("table_semantic", "")).strip() or not table.get("note_ids") or table.get("scope") not in {"consolidated", "parent", "standalone"}:
        raise WorkbookError("补格须有已确认的业务表含义、科目和报表口径")
    if any(address not in sheet["cells"] for address in table.get("header_cells", [])):
        raise WorkbookError("补格业务范围的表头证据不存在")
    merges = [range_boundaries(value) for value in sheet.get("merges", [])]
    omissions = [range_boundaries(item["range"]) for item in sheet.get("structural_omissions", [])]
    physical_boxes = [range_boundaries(item["range"]) for item in sheet.get("word_table_ranges", [])]
    proof = {"source_hash": snapshot["sha256"], "table_range": area,
             "table_semantic": table["table_semantic"], "note_ids": list(table["note_ids"]),
             "scope": table["scope"], "header_cells": list(table.get("header_cells", []))}
    records = sheet.setdefault("confirmed_business_ranges", [])
    if proof not in records:
        records.append(proof)
    hidden_rows = set(sheet.get("hidden_rows", []))
    hidden_columns = sheet.get("hidden_columns", [])
    added = []
    for row in range(top, bottom + 1):
        for column in range(left, right + 1):
            address = f"{get_column_letter(column)}{row}"
            if address in sheet["cells"] or any(l <= column <= r and t <= row <= b for l, t, r, b in omissions):
                continue
            if "word_table_ranges" in sheet and not any(l <= column <= r and t <= row <= b for l, t, r, b in physical_boxes):
                continue
            if any(l <= column <= r and t <= row <= b and (column, row) != (l, t) for l, t, r, b in merges):
                continue
            row_hidden = row in hidden_rows
            column_hidden = any(lo <= column <= hi for lo, hi in hidden_columns)
            sheet["cells"][address] = {"row": row, "column": column, "value": None,
                "formula": None, "cached_value": None, "number_format": "General", "style_id": 0,
                "hidden_row": row_hidden, "hidden_column": column_hidden, "hidden_sheet": bool(sheet.get("hidden")),
                "hidden": row_hidden or column_hidden or bool(sheet.get("hidden")),
                "derived_blank": True, "blank_evidence": copy.deepcopy(proof)}
            added.append(address)
    return added

def cell_at(snapshot,sheet,cell):
    address=str(cell).replace("$","").upper()
    for s in snapshot["sheets"]:
        if s["name"]==sheet:
            if address in s["cells"]:return s["cells"][address]
            raise WorkbookError(f"{sheet}!{address} 不在已读取候选格中")
    raise WorkbookError("未知工作表："+str(sheet))

def _number(value):
    if isinstance(value,bool):raise WorkbookError("布尔值不能当金额或数量")
    if isinstance(value,(int,float,Decimal)):result=Decimal(str(value))
    elif isinstance(value,str):
        text=unicodedata.normalize("NFKC",value).strip()
        if text.startswith("(") and text.endswith(")"):text="-"+text[1:-1]
        if not re.fullmatch(r"[+-]?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?(?:[eE][+-]?\d+)?",text):raise WorkbookError("来源不是可靠数值，不把横线或文字当0")
        result=Decimal(text.replace(",",""))
    else:raise WorkbookError("来源不是可靠数值")
    if not result.is_finite():raise WorkbookError("数值不是有限数")
    return result

def _plain(v):return (int(v) if v==v.to_integral_value() else float(v)) if isinstance(v,Decimal) else v

def actual_value(snapshot,sheet,cell,_stack=None):
    item=cell_at(snapshot,sheet,cell)
    if not item.get("formula"):return item.get("value")
    formula=str(item["formula"])
    if re.search(r"\[[^]]+\][^+*/(),;]*!",formula) or "|" in formula or re.search(r"\b(?:WEBSERVICE|RTD|CUBE\w*)\s*\(",formula,re.I):raise WorkbookError("公式包含无法验证的外部引用")
    value=item.get("cached_value")
    if snapshot.get("recalculated") is True and value is not None and not (isinstance(value,str) and value.startswith("#")):return value
    if "^" in formula:raise WorkbookError("指数公式须由Excel重算，以保留Excel运算顺序")
    stack=set() if _stack is None else _stack
    identity=(sheet,str(cell).replace("$","").upper())
    if identity in stack:raise WorkbookError("公式循环引用")
    stack.add(identity)
    try:
        pieces=[]
        for t in Tokenizer(formula).items:
            if t.type=="OPERAND" and t.subtype=="RANGE":
                m=_REFERENCE.fullmatch(t.value)
                if not m:raise WorkbookError("名称公式须由Excel重算副本后使用")
                name=(m["sheet"] or sheet).strip("'").replace("''","'")
                if name!=sheet:raise WorkbookError("跨表公式须由Excel重算副本后使用")
                a=m["a"].replace("$","");b=(m["b"] or "").replace("$","")
                pieces.append(f'R("{a}","{b}")' if b else f'C("{a}")')
            elif t.type=="FUNC":
                if t.subtype=="OPEN" and t.value.upper()!="SUM(":raise WorkbookError("复杂公式须由Excel重算副本后使用")
                pieces.append(t.value.upper())
            elif t.type=="OPERAND" and t.subtype!="NUMBER":raise WorkbookError("非数值公式须由Excel重算")
            elif t.type=="OPERATOR-POSTFIX" and t.value=="%":pieces.append("/100")
            elif t.type in {"OPERATOR-INFIX","OPERATOR-PREFIX","PAREN","SEP","WSPACE","OPERAND"}:pieces.append(t.value.replace("^","**").replace(";",","))
            else:raise WorkbookError("公式操作须由Excel重算")
        def ref(address,for_sum=False):
            v=actual_value(snapshot,sheet,address,stack)
            if v is None or v=="" or (for_sum and isinstance(v,(str,bool))):return Decimal(0)
            return _number(v)
        def calc(n,for_sum=False):
            if isinstance(n,ast.Expression):return calc(n.body,for_sum)
            if isinstance(n,ast.Constant) and isinstance(n.value,(int,float)) and not isinstance(n.value,bool):return Decimal(str(n.value))
            if isinstance(n,ast.UnaryOp) and isinstance(n.op,(ast.UAdd,ast.USub)):
                v=calc(n.operand,for_sum);return -v if isinstance(n.op,ast.USub) else v
            if isinstance(n,ast.BinOp):
                a,b=calc(n.left),calc(n.right)
                if isinstance(n.op,ast.Add):return a+b
                if isinstance(n.op,ast.Sub):return a-b
                if isinstance(n.op,ast.Mult):return a*b
                if isinstance(n.op,ast.Div):return a/b
                if isinstance(n.op,ast.Pow) and b==b.to_integral_value() and abs(b)<=12:return a**int(b)
            if isinstance(n,ast.Call) and isinstance(n.func,ast.Name) and not n.keywords:
                if n.func.id=="C" and len(n.args)==1 and isinstance(n.args[0],ast.Constant):return ref(n.args[0].value,for_sum)
                if n.func.id=="R" and len(n.args)==2 and all(isinstance(x,ast.Constant) for x in n.args):
                    l,t,r,b=range_boundaries(n.args[0].value+":"+n.args[1].value)
                    if (r-l+1)*(b-t+1)>100000:raise WorkbookError("大范围公式须由Excel重算")
                    return [ref(f"{get_column_letter(c)}{row}",True) for row in range(t,b+1) for c in range(l,r+1)]
                if n.func.id=="SUM":
                    total=Decimal(0)
                    for arg in n.args:
                        v=calc(arg,True);total+=sum(v,Decimal(0)) if isinstance(v,list) else v
                    return total
            raise WorkbookError("公式不能安全独立核算")
        return _plain(calc(ast.parse("".join(pieces),mode="eval")))
    except WorkbookError:raise
    except (SyntaxError,ArithmeticError,ValueError,TypeError) as exc:raise WorkbookError("公式无法可靠计算："+str(exc)) from exc
    finally:stack.remove(identity)

def _new_output(source,output):
    if source==output or output.exists():raise WorkbookError("输出须为尚不存在的新文件，不能覆盖输入或已有结果")
    if source.suffix.lower()!=output.suffix.lower():raise WorkbookError("输出须保留工作簿原扩展名")

def recalculate_copy(source,output):
    import pythoncom
    import win32com.client
    source,output=Path(source).resolve(),Path(output).resolve();_new_output(source,output)
    before=file_hash(source);w=load_workbook(source,data_only=False,keep_links=True)
    try:external=_external(w)
    finally:_close_workbook(w)
    if external:raise WorkbookError("拒绝采用含外部引用的重算结果："+"、".join(external[:10]))
    output.parent.mkdir(parents=True,exist_ok=True);pythoncom.CoInitialize();excel=book=None
    try:
        excel=win32com.client.DispatchEx("Excel.Application");excel.Visible=False;excel.DisplayAlerts=False;excel.EnableEvents=False;excel.AskToUpdateLinks=False;excel.AutomationSecurity=3
        book=excel.Workbooks.Open(str(source),UpdateLinks=0,ReadOnly=True,IgnoreReadOnlyRecommended=True,AddToMru=False)
        if book.LinkSources(1) or book.Connections.Count:raise WorkbookError("来源有外部链接或数据连接，无法证明重算有效")
        excel.CalculateFullRebuild();book.SaveCopyAs(str(output))
        if file_hash(source)!=before:raise WorkbookError("来源在重算期间发生变化")
        if not output.is_file():raise WorkbookError("Excel未生成重算副本")
        _RECALCULATED[file_hash(output)]={"source_sha256":before,"external_references":[],"method":"Excel.CalculateFullRebuild"}
        return str(output)
    finally:
        try:
            if book is not None:book.Close(SaveChanges=False)
        finally:
            if excel is not None:excel.Quit()
            pythoncom.CoUninitialize()

def _normalize(v):
    if isinstance(v,str):return re.sub(r"\s+"," ",unicodedata.normalize("NFKC",v)).strip()
    if isinstance(v,dict):return {str(k):_normalize(x) for k,x in sorted(v.items())}
    if isinstance(v,list):return [_normalize(x) for x in v]
    return v

def semantic_key(mapping):
    slot=mapping.get("slot_id") or mapping.get("target_slot_id")
    if not slot:raise WorkbookError("映射没有金标准ID")
    aliases={"期间角色":"period_role","币种":"currency","单位":"unit","倍率":"scale","汇总层级":"aggregation_level"}
    excluded={"source_row_path","source_column_path","source_value_type","source_workbook","source_sheet","raw_row_path","raw_column_path","row_path","column_path","field_path","sheet","cell","address","unit","scale","entity","主体"}
    dimensions={}
    for original,v in (mapping.get("dimensions") or {}).items():
        key=aliases.get(original,original)
        if key in excluded:continue
        v=_normalize(v)
        if key=="currency" and isinstance(v,str):v=_CURRENCY.get(v,v.upper())
        dimensions[key]=v
    scope=mapping.get("scope") or ("consolidated" if slot.startswith("C-") else "parent" if slot.startswith("P-") else "")
    kind=mapping.get("value_type") or (mapping.get("dimensions") or {}).get("source_value_type") or (mapping.get("target_slot") or {}).get("value_type")
    if not kind:raise WorkbookError("映射缺少实际值类型")
    # semantic_field用于展示；身份由金标准槽位与实际业务维度共同确定。
    # 不同指标须遵守槽位定义，不能借任意metric把缺失指标塞进现有槽位。
    return json.dumps({"slot_id":slot,"scope":scope,"dimensions":dimensions,"value_type":kind},ensure_ascii=False,sort_keys=True,separators=(",",":"),default=str)

def _location(m):return m.get("sheet") or m.get("worksheet"),m["cell"]

def _units(m):
    d=m.get("dimensions") or {}
    return str(d.get("unit",d.get("单位",""))).strip(),_number(d.get("scale",d.get("倍率",1)))

def _convert(value,source,target,source_cell,target_cell):
    """沿用旧调用入口，只搬运原值，不检查或改变数值及显示格式。"""
    return value

def _monetary_number_format(value):
    if not isinstance(value,str) or value in {"General","@",""} or "%" in value or is_date_format(value):return False
    visible=re.sub(r'"[^\"]*"|\\.|\[[^]]*\]',"",value)
    return bool(re.search(r"[0#?]",visible))


def _format_identity(mapping):
    dimensions=mapping.get("dimensions") or {}
    if mapping.get("value_type")!="monetary" or mapping.get("reviewed") is not True:return None
    if not dimensions.get("currency") or not dimensions.get("unit") or "scale" not in dimensions:return None
    unit,scale=_units(mapping)
    if scale<=0:return None
    return (json.loads(semantic_key(mapping))["scope"],_CURRENCY.get(str(dimensions["currency"]),str(dimensions["currency"]).upper()),
        unit,scale,_normalize(dimensions.get("period")),_normalize(dimensions.get("entity")),_normalize(dimensions.get("reporting_entity")))


def _in_format_table(address,area):
    col,row,_,_=range_boundaries(address);left,top,right,bottom=range_boundaries(area)
    return left<=col<=right and top<=row<=bottom


def _number_format_update(snapshot,target,mappings,target_cell,value):
    """仅沿用原表同列已确认金额格的一致格式，不从来源表借格式。"""
    if target_cell.get("formula") or target_cell.get("value") not in (None,"") or target_cell.get("number_format") not in {"General","@",""}:return None
    if isinstance(value,bool) or not isinstance(value,(int,float,Decimal)):return None
    identity=_format_identity(target);area=target.get("table_range")
    if identity is None or not area:return None
    sheet,address=_location(target);column=range_boundaries(address)[0]
    if not _in_format_table(address,area):return None
    original=next(s for s in snapshot["sheets"] if s["name"]==sheet)
    if "word_table_ranges" in original:
        left,top,right,bottom=range_boundaries(area)
        containing=[r for r in original["word_table_ranges"] if _in_format_table(f"{get_column_letter(left)}{top}",r["range"])
                    and _in_format_table(f"{get_column_letter(right)}{bottom}",r["range"])]
        if len(containing)!=1:return None
    counts=defaultdict(int)
    for m in mappings:counts[_location(m)]+=1
    sources=[]
    for m in mappings:
        location=_location(m)
        if location==(sheet,address) or location[0]!=sheet or counts[location]!=1:continue
        if m.get("table_range")!=area or range_boundaries(location[1])[0]!=column or not _in_format_table(location[1],area):continue
        if _format_identity(m)!=identity:continue
        cell=cell_at(snapshot,*location);fmt=cell.get("number_format")
        if cell.get("formula") or not _monetary_number_format(fmt):continue
        try:numeric=_number(cell.get("value"))
        except WorkbookError:continue
        sources.append({"sheet":sheet,"cell":location[1],"number_format":fmt,"original_value":_plain(numeric),
            "mapping":{k:copy.deepcopy(m[k]) for k in ("sheet","cell","slot_id","scope","dimensions","value_type","reviewed","table_range") if k in m}})
    formats={item["number_format"] for item in sources}
    if len(formats)!=1:return None
    return {"old_format":target_cell.get("number_format"),"new_format":next(iter(formats)),"source_a_hash":snapshot["sha256"],
        "table_range":area,"target_mapping":{k:copy.deepcopy(target[k]) for k in ("sheet","cell","slot_id","scope","dimensions","value_type","reviewed","table_range") if k in target},
        "sources":sources,"reason":"原空白金额格沿用同一已确认业务表、同列、同期间币种单位的原数值格唯一显示格式"}


def _validate_number_format_update(workbook,entry,source_hash):
    proof=entry.get("number_format_update")
    if proof is None:return
    sheet,address=entry.get("sheet") or entry.get("worksheet"),entry["cell"]
    try:
        target=proof["target_mapping"];cell=workbook[sheet][address];area=proof["table_range"]
        if proof["source_a_hash"]!=source_hash or _location(target)!=(sheet,address) or target.get("table_range")!=area:raise ValueError()
        if target.get("slot_id")!=entry.get("slot_id") or target.get("dimensions")!=entry.get("dimensions"):raise ValueError()
        if cell.value not in (None,"") or cell.data_type=="f" or cell.number_format!=proof["old_format"] or cell.number_format not in {"General","@",""}:raise ValueError()
        if not _monetary_number_format(proof["new_format"]) or not _in_format_table(address,area):raise ValueError()
        if isinstance(entry.get("value"),bool) or not isinstance(entry.get("value"),(int,float,Decimal)):raise ValueError()
        identity=_format_identity(target)
        if identity is None or not proof["sources"]:raise ValueError()
        seen=set()
        for source in proof["sources"]:
            donor=source["mapping"];location=_location(donor)
            if location in seen or location==(sheet,address) or location!=(source["sheet"],source["cell"]):raise ValueError()
            seen.add(location)
            if location[0]!=sheet or range_boundaries(location[1])[0]!=range_boundaries(address)[0] or donor.get("table_range")!=area or not _in_format_table(location[1],area):raise ValueError()
            if _format_identity(donor)!=identity:raise ValueError()
            original=workbook[location[0]][location[1]]
            if original.data_type=="f" or original.number_format!=source["number_format"] or original.number_format!=proof["new_format"]:raise ValueError()
            if _number(original.value)!=_number(source["original_value"]):raise ValueError()
    except (KeyError,TypeError,ValueError,IndexError) as error:
        raise WorkbookError(f"金额显示格式依据与原表不一致：{sheet}!{address}") from error


def build_update_plan(a,b,a_mappings,b_mappings):
    plan={"source_a_hash":a["sha256"],"source_b_hash":b["sha256"],"updates":[],"issues":[],"unplaced_sources":[],"retained_blanks":[],"operations":[],"labels":[]}
    candidates=defaultdict(list);consumed=set();occupied=set()
    for index,m in enumerate(b_mappings):
        try:candidates[semantic_key(m)].append((index,m))
        except WorkbookError as exc:
            sheet,cell=_location(m);plan["issues"].append({"kind":"invalid_source_mapping","reason":str(exc),"sheet":sheet,"cell":cell})
    target_counts=defaultdict(int)
    for m in a_mappings:target_counts[_location(m)]+=1
    for target in a_mappings:
        sheet,cell=_location(target);issue_kind="unresolved_target"
        try:
            key=semantic_key(target);matches=candidates.get(key,[])
            if target_counts[(sheet,cell)]>1:raise WorkbookError("目标格有重复语义映射，未写入")
            tc=cell_at(a,sheet,cell)
            if not matches:
                if not tc.get("formula") and tc.get("value") in (None,""):
                    plan["retained_blanks"].append({"sheet":sheet,"cell":cell,"slot_id":target.get("slot_id") or target.get("target_slot_id"),
                        "dimensions":copy.deepcopy(target.get("dimensions") or {}),"semantic_field":target.get("semantic_field",""),
                        "semantic_key":key,"old_value":tc.get("value"),"source_a_hash":a["sha256"],"source_b_hash":b["sha256"],
                        "reason":"更新数据未提供该语义，目标原为空白，保持原空白；未作为已更新数据"})
                    continue
                issue_kind="missing_source"
                raise WorkbookError("更新数据表没有相同主体、期间、对象及业务含义的数据")
            values=[];sources=[];blank_sources=[]
            for index,source in matches:
                source_sheet,address=_location(source);raw=actual_value(b,source_sheet,address)
                source_record={"sheet":source_sheet,"cell":address,"value":raw}
                if raw in (None,""):
                    blank_sources.append(source_record)
                    continue
                values.append(_convert(raw,source,target,cell_at(b,source_sheet,address),tc))
                sources.append(source_record)
            if not values:
                if not tc.get("formula") and tc.get("value") in (None,""):
                    plan["retained_blanks"].append({"sheet":sheet,"cell":cell,"slot_id":target.get("slot_id") or target.get("target_slot_id"),
                        "dimensions":copy.deepcopy(target.get("dimensions") or {}),"semantic_field":target.get("semantic_field",""),
                        "semantic_key":key,"old_value":tc.get("value"),"source_a_hash":a["sha256"],"source_b_hash":b["sha256"],"blank_sources":blank_sources,
                        "reason":"同语义来源尚未填写，目标原为空白，保持原空白；未作为已更新数据"})
                    consumed.update(i for i,m in matches)
                    continue
                issue_kind="source_unknown_blank"
                raise WorkbookError("相同语义的更新数据尚未填写，不能据此清空目标原金额或认定为零")
            if any(type(v)!=type(values[0]) or v!=values[0] for v in values[1:]):raise WorkbookError("相同业务身份的多个来源数值冲突，未自动选取或求和")
            plan["updates"].append({"sheet":sheet,"cell":cell,"value":_plain(values[0]),"old_value":tc.get("value"),"slot_id":target.get("slot_id") or target.get("target_slot_id"),"dimensions":copy.deepcopy(target.get("dimensions") or {}),"semantic_field":target.get("semantic_field",""),"sources":sources,"blank_sources":blank_sources})
            format_update = _number_format_update(a, target, a_mappings, tc, _plain(values[0]))
            if format_update:
                plan["updates"][-1]["number_format_update"] = format_update
            consumed.update(i for i,m in matches)
        except (WorkbookError,KeyError,TypeError,ValueError) as exc:plan["issues"].append({"kind":issue_kind,"reason":str(exc),"sheet":sheet,"cell":cell})
    plan["unplaced_sources"]=[copy.deepcopy(m) for i,m in enumerate(b_mappings) if i not in consumed]
    return plan

def write_updated_workbook(source,output,plan):
    source,output=Path(source).resolve(),Path(output).resolve();_new_output(source,output)
    before=file_hash(source)
    if plan.get("source_a_hash") and plan["source_a_hash"]!=before:raise WorkbookError("附注表格在识别后发生变化，禁止执行过期计划")
    w=load_workbook(source,data_only=False,keep_vba=source.suffix.lower()==".xlsm",keep_links=True);expected=[]
    try:
        for operation in plan.get("operations",[]):_apply_operation(w,operation)
        for entry in plan.get("updates",[]):_validate_number_format_update(w,entry,before)
        seen=set()
        for entry in [*plan.get("labels",[]),*plan.get("updates",[])]:
            sheet,address=entry.get("sheet") or entry.get("worksheet"),entry["cell"]
            if sheet not in w:raise WorkbookError("计划引用未知工作表")
            if (sheet,address) in seen:raise WorkbookError("计划对同一格安排了多次写入")
            seen.add((sheet,address));c=w[sheet][address]
            if isinstance(c,MergedCell):raise WorkbookError("不能写入合并区域非首格")
            prior_style=copy.copy(c._style)
            value=_plain(entry.get("value"));c.value=value
            if entry.get("number_format_update"):
                c.number_format=entry["number_format_update"]["new_format"]
                allowed_style=copy.copy(prior_style)
                if allowed_style is None:allowed_style=copy.copy(c._style)
                allowed_style.numFmtId=c._style.numFmtId
                if c._style!=allowed_style:raise WorkbookError(f"金额格发生了数字格式以外的样式变化：{sheet}!{address}")
            if isinstance(value,str):c.data_type="s"
            expected.append((sheet,address,value,copy.copy(c._style)))
        if file_hash(source)!=before:raise WorkbookError("源文件在写入前发生变化，已停止")
        output.parent.mkdir(parents=True,exist_ok=True)
        if w.calculation is None:w.calculation=CalcProperties()
        w.calculation.fullCalcOnLoad=True;w.calculation.forceFullCalc=True
        with output.open("xb") as f:w.save(f)
    finally:_close_workbook(w)
    check=load_workbook(output,data_only=False,keep_vba=output.suffix.lower()==".xlsm",keep_links=True)
    try:
        for sheet,address,value,style in expected:
            c=check[sheet][address];equal=c.value==value or (value=="" and c.value is None)
            same_style = (tuple(c._style) if c.has_style else None) == (tuple(style) if style is not None and any(style) else None)
            if not equal or not same_style or (isinstance(value,str) and c.data_type=="f"):raise WorkbookError(f"输出回读验证失败：{sheet}!{address}")
        if file_hash(source)!=before:raise WorkbookError("来源文件发生变化，输出不可视为通过验证")
    finally:_close_workbook(check)
    return str(output)

def _name_range(w,name):
    choices=[d for d in _definitions(w) if d.name==name]
    if len(choices)!=1 or len(_destinations(choices[0]))!=1:raise WorkbookError("结构变更须指向唯一连续命名区："+str(name))
    sheet,address=_destinations(choices[0])[0];box=range_boundaries(address)
    if not all(isinstance(x,int) for x in box):raise WorkbookError("不能对整行或整列名称作结构变更")
    return choices[0],w[sheet],box

def _interval(lo,hi,pos,count,action):
    if action=="insert":
        if pos<=lo:return lo+count,hi+count
        if pos<=hi:return lo,hi+count
        return lo,hi
    last=pos+count-1
    if hi<pos:return lo,hi
    if lo>last:return lo-count,hi-count
    size=hi-lo+1-max(0,min(hi,last)-max(lo,pos)+1)
    if size<=0:raise WorkbookError("删除会移除其他公式或名称引用的完整区域")
    start=lo if lo<pos else pos
    return start,start+size-1

def _transform(address,target_sheet,formula_sheet,op,box):
    m=_REFERENCE.fullmatch(address)
    if not m:
        if any(c in address for c in "[!:"):raise WorkbookError("不能安全调整公式引用："+address)
        return address
    name=(m["sheet"] or formula_sheet).strip("'").replace("''","'")
    if name!=target_sheet:return address
    a,b=m["a"],m["b"];left,top,right,bottom=range_boundaries(a+(":"+b if b else ""))
    pos=(box[1] if op["axis"]=="row" else box[0])+op["index"];count=op.get("count",1)
    if op["axis"]=="row":top,bottom=_interval(top,bottom,pos,count,op["action"])
    else:
        if bottom<box[1] or top>box[3] or right<pos or left>box[2]:return address
        if top<box[1] or bottom>box[3] or right>box[2]:raise WorkbookError("局部插删列会拆散跨表公式范围")
        left,right=_interval(left,right,pos,count,op["action"])
    def fmt(old,row,col):
        x=re.fullmatch(r"(\$?)[A-Z]+(\$?)\d+",old,re.I)
        return x[1]+get_column_letter(col)+x[2]+str(row)
    result=fmt(a,top,left)+(":"+fmt(b,bottom,right) if b else "")
    return (m["sheet"]+"!" if m["sheet"] else "")+result

def _apply_operation(w,op):
    definition,s,box=_name_range(w,op.get("range_name"))
    axis,action=op.get("axis"),op.get("action");count,index=op.get("count",1),op.get("index")
    if axis not in {"row","column"} or action not in {"insert","delete"} or type(count) is not int or count<1 or type(index) is not int:raise WorkbookError("结构操作参数无效")
    length=box[3]-box[1]+1 if axis=="row" else box[2]-box[0]+1
    if index<0 or index>length or (action=="delete" and (index+count>length or count>=length)):raise WorkbookError("结构操作超出命名区范围")
    if any(ws.tables or ws._charts or ws._pivots or ws.data_validations.count or len(ws.conditional_formatting) for ws in w):raise WorkbookError("结构变更涉及表、图表、透视表、数据验证或条件格式依赖，未执行未经维护的引用变更")
    pos=(box[1] if axis=="row" else box[0])+index
    template=op.get("template_index",max(0,min(index-1,length-1)))
    if type(template) is not int or not 0<=template<length:raise WorkbookError("样式模板行列超出范围")
    template_pos=(box[1] if axis=="row" else box[0])+template
    styles={}
    for i in (range(box[0],box[2]+1) if axis=="row" else range(box[1],box[3]+1)):
        c=s.cell(template_pos,i) if axis=="row" else s.cell(i,template_pos);styles[i]=copy.copy(c._style)
    row_style=copy.copy(s.row_dimensions.get(template_pos)) if axis=="row" else None
    merges=[]
    for m in list(s.merged_cells.ranges):
        lo,hi=(m.min_row,m.max_row) if axis=="row" else (m.min_col,m.max_col)
        overlap_rows=not (m.max_row<box[1] or m.min_row>box[3]);overlap_cols=not (m.max_col<box[0] or m.min_col>box[2])
        if action=="insert" and lo<=template_pos<=hi and (overlap_cols if axis=="row" else overlap_rows):raise WorkbookError("结构样式模板涉及合并区域，请选取普通行列模板")
        affected=axis=="row" or (overlap_rows and not (m.max_col<pos or m.min_col>box[2]))
        if not affected:merges.append(str(m));continue
        if axis=="column" and (m.min_row<box[1] or m.max_row>box[3]):raise WorkbookError("局部插删列跨越合并区域")
        if (action=="insert" and lo<pos<=hi) or (action=="delete" and not (hi<pos or lo>=pos+count)):raise WorkbookError("结构操作跨越或删除合并区域，未执行")
        merges.append(_transform(str(m),s.title,s.title,op,box))
    formulas=[]
    for ws in w:
        for c in list(ws._cells.values()):
            if c.data_type!="f":continue
            if not isinstance(c.value,str):raise WorkbookError("结构变更涉及数组或动态溢出公式，未执行未经维护的依赖变更")
            text=str(c.value)
            if re.search(r"\b(?:INDIRECT|OFFSET)\s*\(",text,re.I):raise WorkbookError("动态地址公式不能证明结构变更后的引用安全")
            tokens=Tokenizer(text).items
            for token in tokens:
                if token.type=="OPERAND" and token.subtype=="RANGE":token.value=_transform(token.value,s.title,ws.title,op,box)
            formulas.append((c,"="+"".join(t.value for t in tokens)))
    names=[]
    for d in _definitions(w):
        destinations=_destinations(d);changed=[]
        if not destinations and any(x in str(d.attr_text) for x in ("!","(")):raise WorkbookError("定义名称含公式，结构变更须明确处理该依赖")
        for sheet,address in destinations:
            if sheet!=s.title:changed.append(f"{quote_sheetname(sheet)}!{address}");continue
            left,top,right,bottom=range_boundaries(address)
            if not all(isinstance(x,int) for x in (left,top,right,bottom)):raise WorkbookError("结构变更遇到整行或整列定义名称")
            if d is definition:
                if axis=="row":bottom+=count if action=="insert" else -count
                else:right+=count if action=="insert" else -count
            elif axis=="row":top,bottom=_interval(top,bottom,pos,count,action)
            elif not (bottom<box[1] or top>box[3] or right<pos or left>box[2]):
                if top<box[1] or bottom>box[3] or right>box[2]:raise WorkbookError("局部列变更会拆散其他命名区")
                left,right=_interval(left,right,pos,count,action)
            changed.append(f"{quote_sheetname(sheet)}!${get_column_letter(left)}${top}:${get_column_letter(right)}${bottom}")
        if changed:names.append((d,",".join(changed)))
    if axis=="column" and action=="insert":
        for c in s._cells.values():
            if box[1]<=c.row<=box[3] and box[2]<c.column<=box[2]+count and (c.value is not None or c.has_style):raise WorkbookError("表格右侧有内容或样式，插列会覆盖，未执行")
    # 打印区域保持包含原表；整行移动同步调整范围和重复标题。
    print_ranges=[]
    for area in s._print_area.ranges:
        left,top,right,bottom=area.bounds
        if axis=="row":top,bottom=_interval(top,bottom,pos,count,action)
        elif action=="insert" and not (bottom<box[1] or top>box[3]):right=max(right,box[2]+count)
        print_ranges.append(f"${get_column_letter(left)}${top}:${get_column_letter(right)}${bottom}")
    title_rows=s.print_title_rows
    if title_rows and axis=="row":
        first,last=[int(v.replace("$","")) for v in title_rows.split(":")];first,last=_interval(first,last,pos,count,action);title_rows=f"{first}:{last}"
    for c,text in formulas:c.value=text
    for m in list(s.merged_cells.ranges):s.unmerge_cells(str(m))
    selected=[(key,c) for key,c in s._cells.items() if (c.row>=pos if axis=="row" else box[1]<=c.row<=box[3] and pos<=c.column<=box[2])]
    selected.sort(key=lambda pair:pair[0][0 if axis=="row" else 1],reverse=action=="insert")
    shift=count if action=="insert" else -count
    for (row,col),c in selected:
        s._cells.pop((row,col),None);coordinate=row if axis=="row" else col
        if action=="delete" and pos<=coordinate<pos+count:continue
        if axis=="row":row+=shift
        else:col+=shift
        c.row,c.column=row,col;s._cells[(row,col)]=c
    if axis=="row":
        dimensions=[(row,d) for row,d in s.row_dimensions.items() if row>=pos]
        for row,d in sorted(dimensions,reverse=action=="insert"):
            del s.row_dimensions[row]
            if action=="delete" and row<pos+count:continue
            d.index=row+shift;s.row_dimensions[row+shift]=d
    for m in merges:s.merge_cells(m)
    if action=="insert":
        for offset in range(count):
            for i,style in styles.items():
                c=s.cell(pos+offset,i) if axis=="row" else s.cell(i,pos+offset);c._style=copy.copy(style)
            if row_style is not None:
                d=copy.copy(row_style);d.index=pos+offset;s.row_dimensions[pos+offset]=d
    for d,text in names:d.attr_text=text
    if print_ranges:s.print_area=print_ranges
    if title_rows:s.print_title_rows=title_rows
