# coding: utf-8
"""A识别出口的原答复封存和只读校验；不重建、修正或比较财务身份。"""
from copy import deepcopy
from pathlib import Path
import hashlib
import json
import re

from .统一语义 import _require, definition_hash, identity_completeness, json_date, json_dates

FIELDS=('facts','slot_bindings','owners','unresolved','excluded','source_identities','outline','coverage',
        'complete','a_identity_coverage','source_hash','gold_hash','source_context','recognition_scope',
        'source_path','original_path','original_hash','gold_path','carrier','schema_version',
        'identity_version','cell_identity_version','formula_source_version','slot_binding_version',
        'a_semantic_reading','封存定义原文','原判断与核查','字段来源','来源缺口','完整性',
        '已确认关系','关系缺口','机械记录','原文证据','未通过答复','collection_version')


def _digest(text):
    return hashlib.sha256(text.encode('utf-8')).hexdigest()


def _at(value, path):
    for key in path:value=value[key]
    return value


def _inside(address, ranges, table):
    from .A语义阅读 import addresses
    owned={a for span in table['data_ranges'] for a in addresses(span)} if '*' in ranges else set()
    return any(address in (owned if span=='*' else addresses(span)) for span in ranges)


def _records(mapping):
    for row in mapping['facts']:
        if row['record_type']=='metric_fact' or (mapping.get('逐格字段版本')==1 and row['record_type']=='dimension_binding' and 'dimensions' in row):
            yield row['source_reference']['location'],row
    for row in mapping.get('slot_bindings',[]):yield row['sheet']+'!'+row['cell'],row['partial_semantics']


def _standard(mapping):
    result={'metric':{},'dimension':{},'sha256':mapping['gold_hash'],'_complete_identity':True}
    for line in mapping['封存定义原文']:
        row=json.loads(line);result[row['record_type']][row['id']]=row
    return result


def _documents(mapping):
    result={}
    for key,doc in mapping['原判断与核查'].items():
        _require(_digest(doc['task_text'])==doc['task_sha256'] and _digest(doc['receipt_text'])==doc['receipt_sha256'], '封存原任务或回执文字变化')
        task=json.loads(doc['task_text'].lstrip('\ufeff'));receipt=json.loads(doc['receipt_text'].lstrip('\ufeff'))
        _require(key==task['task_id']==receipt['task_id']==definition_hash({k:v for k,v in task.items() if k!='task_id'}),'封存任务指纹不符')
        _require(receipt['result_hash']==definition_hash(receipt['result']),'封存原回答指纹不符')
        _require(task['source']==mapping['source_hash'],'封存任务原件不符')
        result[key]=(task,receipt['result'])
    return result


def validate_relations(binding,review,tables,aliases):
    from .A语义阅读 import addresses,normalize_evidence
    index={t['id']:t for t in tables};original={r['id']:r for r in binding['tables']}
    for row in review['tables']:
        relations=original[row['id']].get('relations',[])
        for i in row.get('checked_relations',[]):_require(type(i) is int and 0<=i<len(relations),'关系核查序号没有原判断')
        fixes=row.get('corrections',{}).get('relations',[])
        _require(len({r['index'] for r in fixes})==len(fixes),'同一关系有冲突的多项修正')
        for fix in fixes:_require(type(fix['index']) is int and 0<=fix['index']<len(relations),'关系修正没有原判断')
        for relation in [*relations,*[r['value'] for r in fixes]]:
            _require(set(relation)=={'content','record_ranges','exceptions','evidence'} and isinstance(relation['content'],str) and relation['content'].strip(),'关系缺少原有内容、范围、例外或证据')
            _require(bool(relation['record_ranges']),'关系未明确本方记录范围')
            for exception in relation['exceptions']:
                _require(isinstance(exception.get('content'),str) and exception['content'].strip(),'关系例外缺少原判断内容')
            for scope in relation['record_ranges']+relation['exceptions']:
                _require(scope['table_id'] in index and scope['ranges'],'关系范围没有原表')
                table=index[scope['table_id']];owned={a for span in table['data_ranges'] for a in addresses(span)}
                for span in scope['ranges']:_require(span=='*' or set(addresses(span))<=owned,'关系范围越出已确认财务目标')
            table=index[row['id']]
            _require(bool(normalize_evidence(relation['evidence'],table['sheet'],aliases,next(iter(addresses(table['data_ranges'][0]))))),'关系没有实际原文证据')


def _relation_blocked(relation, review, tables):
    from .A语义阅读 import addresses
    index={t['id']:t for t in tables}
    for scope in relation['record_ranges']+relation['exceptions']:
        table=index[scope['table_id']]
        for row in review['tables']:
            if row['id']!=scope['table_id']:continue
            if any(_inside(a,scope['ranges'],table) for issue in row.get('issues',[]) for a in addresses(issue['range'])):
                return True
    return False


def _plain_answer(text, standard, sheet_names=(), metric_context=None, default_sheet=None, field_scope=None):
    """读取文字或表格中明示的坐标、编号和维度名；业务文字原样保留。"""
    from .A语义阅读 import addresses
    names={key:key for key in standard['dimension']}
    for key,row in standard['dimension'].items():
        label=row.get('domain')
        if label and sum(d.get('domain')==label for d in standard['dimension'].values())==1:
            names[label]=key
            if row.get('value_type')=='selection':names[label+'范围']=key
    names.update({k:'metric_id' for k in ('metric_id','标准槽位编号','标准编号','槽位编号','所选槽位','所选金标准槽位','槽位','标准','计量指标')})
    if 'entity' in standard['dimension']:names['主体']='entity'
    if 'report_scope' in standard['dimension']:names.update({'报表范围':'report_scope','报表口径':'report_scope','本次对应范围':'report_scope'})
    if 'unit_scale' in standard['dimension']:names.update({'金额倍率':'unit_scale','倍率':'unit_scale'})
    if 'currency' in standard['dimension']:names.update({'币种及倍率':'currency','金额币种':'currency'})
    if 'period' in standard['dimension']:
        names.update({label:'period' for label in ('期间时点','时点','报告期间','实际期间','实际日期','期间')})
    for label,key in re.findall(r'([\u4e00-\u9fff]+)范围中[，,]\s*`(\w+)`',text):
        if key in standard['dimension']:names[label]=key
    range_names={key for key in re.findall(r'`?(\w+)`?\s*按(?:各行实际范围取值|上列各格分别取值)',text) if key in standard['dimension']}
    if len(range_names)==1:names.update({label:next(iter(range_names)) for label in ('范围','指标取值范围')})
    shared_field_aliases={}
    for declaration in re.finditer(r'[“"]([^”"]+)[”"]取值同时归属\s*([^，,。；;\n]+)',text):
        keys=[key for key in re.findall(r'`?(\w+)`?',declaration[2]) if key in standard['dimension']]
        if len(keys)>1:
            names[declaration[1]]=keys[0]
            shared_field_aliases[keys[0]]=(keys[1:],declaration[0])
    for declaration in re.finditer(r'((?:`?\w+`?\s*(?:与|和|、)\s*)+`?\w+`?)\s*采用同一范围值',text):
        keys=[key for key in re.findall(r'\w+',declaration[1]) if key in standard['dimension']]
        if len(keys)>1:shared_field_aliases[keys[0]]=(keys[1:],declaration[0])
    if len(shared_field_aliases)==1 and '项目范围' in names:names['范围']=names['项目范围']
    names['template_period_role']='template_period_role';names['period_role']='template_period_role'
    fields='|'.join((r'(?<![\u4e00-\u9fff])' if names[k]=='period' and not k.isascii() or k=='范围' else r'(?<!的)' if k=='主体' else '')+re.escape(k)
                    for k in sorted(names,key=len,reverse=True))
    marker=re.compile(r'(?<![A-Za-z0-9_.])('+fields+r')(?>`?)(?:\s*[（(][^（）()\n]*[）)])?\s*(?:(?:[:：=＝]|(?:均|依次|分别)?为(?!准)|是|取(?:值)?|[，,]?\s*实际取值(?:均)?(?:为|[:：=＝]))\s*|(?=`[^`]+`|“[^”]+”|按(?:上列格子|所在行分别取)|分别与同一行[A-Z]+列|同上|与[^；;。]+相同))')
    annotated_field=re.compile(r'[\u4e00-\u9fff]+\s*[（(](?:维度\s*)?`?('+fields+r')`?[）)]\s*(?:为|是|[:：=])')
    known_sheets='|'.join(re.escape(name) for name in sorted(sheet_names,key=len,reverse=True))
    sheet_pattern=known_sheets or r"[^\s:：;；,，、|()（）'`]+"
    coordinate=re.compile(r"(?<![A-Za-z0-9_.@])(?:(?:'[^']+'|"+sheet_pattern+r")!)?\$?[A-Z]+\$?[1-9]\d*(?::\$?[A-Z]+\$?[1-9]\d*)?(?![A-Za-z0-9_.])")
    metrics=re.compile(r'(?<![A-Za-z0-9_.])('+ '|'.join(re.escape(k) for k in sorted(standard['metric'],key=len,reverse=True))+r')(?![A-Za-z0-9_.])')
    declined=re.compile(r'(?:不据此标作|不作为|不标作|不识别为|不得把[^。；;]*?套成)\s*`?(?:'+ '|'.join(re.escape(k) for k in standard['metric'])+r')')
    result={'metrics':[],'dimensions':[],'unresolved':[]};scope=['*'];headers=None;previous_table_cells=None
    stated_data=re.findall(r'(?:数据格|财务格)(?:为|是)\s*((?:'+coordinate.pattern+r'[、，,\s]*)+)',text.replace('`',''))
    if len(stated_data)==1:result['data_ranges']=coordinate.findall(stated_data[0])
    for block in re.split(r'\n(?=#{1,6}\s)',text.replace(chr(96),'').replace('**','')):
        declaration=re.search(r'对应指标依次(?:是|为)([^。\n]+)',block)
        columns=re.findall(r'(?:^|[：；])\s*([A-Z]+)[^A-Z；。\n]*列→A',block)
        ids=metrics.findall(declaration[1]) if declaration else []
        if not columns or len(columns)!=len(ids):continue
        column_ids=dict(zip(columns,ids))
        for line in block.splitlines():
            if '←' not in line:continue
            for area in coordinate.findall(line.split('←',1)[1]):
                source,sep,span=area.rpartition('!')
                if not sep:continue
                for cell in addresses(span):
                    key=column_ids.get(re.match(r'[A-Z]+',cell)[0])
                    if key:result['metrics'].append({'id':key,'ranges':[source+'!'+cell],
                        'original':{'column_declaration':declaration[0],'row_declaration':line}})
    sheet=default_sheet or (sheet_names[0] if len(sheet_names)==1 else None)
    requested={}
    if field_scope:
        from .A语义阅读 import addresses
        for area,keys in field_scope.items():
            source,sep,span=area.rpartition('!')
            for cell in addresses(span if sep else area):
                requested[(source if sep else sheet)+'!'+cell]=set(keys)
    for block in re.split(r'\n(?=#{1,6}\s)',text.replace('`','').replace('**','')):
        heading=block.splitlines()[0] if block.splitlines() else ''
        areas=coordinate.findall(heading)
        if not areas:areas=next((coordinate.findall(line) for line in block.splitlines()
                                if any(name+'!' in line for name in sheet_names)),[])
        owners={area.rsplit('!',1)[0] for area in areas if '!' in area}
        if not requested or len(owners)!=1:continue
        owner=next(iter(owners))
        owned={owner+'!'+cell for area in areas for cell in addresses(area.rsplit('!',1)[-1])}&requested.keys()
        for role,columns in re.findall(r'B\s*(期末|期初)(?:组|列)\s*([A-Z]+(?::[A-Z]+)+)',block):
            locations=[loc for loc in owned if re.match(r'[A-Z]+',loc.rsplit('!',1)[1])[0] in columns.split(':')
                       and 'template_period_role' in requested[loc]]
            if locations:result['dimensions'].append({'id':'template_period_role','ranges':locations,
                'value':{'期末':'closing','期初':'opening'}[role],'original':block.splitlines()[0]+'；B'+role+'组 '+columns})
        alias_pairs=re.findall(metrics.pattern+r'[（(]([^（）()]+)[）)]',block)
        aliases={label:key for key,label in alias_pairs if len({k for k,v in alias_pairs if v==label})==1}
        declaration=re.search(r'列映射为[:：]([^\n]+)',block)
        declared_columns={}
        if declaration:
            for part in re.split(r'[、，；;。]',declaration[1]):
                pair=re.fullmatch(r'\s*([A-Z]+)[^A-Z→]*→\s*(.+?)\s*',part)
                if pair and pair[2] in aliases:declared_columns[pair[1]]=aliases[pair[2]]
        group_ids=set();group_columns={}
        for original in block.splitlines():
            if not original.startswith('|'):continue
            values=[v.strip() for v in original.strip('|').split('|')]
            if not values or not re.fullmatch(r'\d+(?:[/／、–—-]\d+)*',values[0]):continue
            ids=set(metrics.findall(original))
            if ids:
                group_ids=ids;group_columns={col:key for col,key in declared_columns.items() if key in ids}
                for part in re.split(r'[；;]',original):
                    if '→' not in part:continue
                    before,after=part.rsplit('→',1);keys=metrics.findall(after)
                    if len(keys)==1:
                        group_columns.update((col,keys[0]) for col in re.findall(r'(?<![A-Za-z0-9_.])([A-Z]+)(?![A-Za-z0-9_.])',before))
            elif '同一组槽位' not in original:continue
            rows={n for first,last in re.findall(r'(\d+)(?:[–—-](\d+))?',values[0]) for n in range(int(first),int(last or first)+1)}
            locations=[loc for loc in owned if int(re.search(r'\d+$',loc)[0]) in rows]
            for loc in locations:
                col=re.match(r'[A-Z]+',loc.rsplit('!',1)[1])[0];key=group_columns.get(col)
                # 同组行可引用已识别的同列编号，不按名称近似或剩余编号猜配。
                known={(metric_context or {}).get(other) for other in locations if re.match(r'[A-Z]+',other.rsplit('!',1)[1])[0]==col}&group_ids
                if key is None and len(known)==1:key=next(iter(known))
                if key and 'metric_id' in requested[loc]:result['metrics'].append({'id':key,'ranges':[loc],'original':original})
    metric_group=None;last_metric=None;last_selection=None;column_periods={};column_roles={};named_periods={};ordered_values={}
    row_roles=[];section_roles={};section_lines=[];declared_fields=set();section_duration=None;row_texts={};row_field=None
    ratio_names={};row_metric_columns={};section_fields=[];common_selection=None;row_objects={}
    pending_row_fields=[]
    selection_defaults={};member_field=None;member_shared={};member_enums=[];paragraph_scope=None;period_links=[]
    column_directions={};last_column_roles={};sheet_columns={};item_groups=[];group_rules=[];group_direction=None;labeled_cells=[]
    selection_aliases={};period_boundaries={};column_defaults={};column_label_refs={};column_table_scope=None
    metric_aliases={m[1]:m[2] for m in re.finditer(r'([\u4e00-\u9fff]+)使用\s*`?('+metrics.pattern+r')',text)}
    heading_metric=None;bullet_scopes=[];field_references=[];range_heading=None
    row_column_groups={tuple(re.findall(r'[A-Z]+',value)) for value in re.findall(
        r'(?:对该行|适用于该行的)\s*([A-Z]+(?:[、，,]\s*[A-Z]+)+)\s*(?:两格|各格|金额格)',text.replace('**',''))}
    header_dates={label:begin+'至'+end for label,begin,end in re.findall(
        r'(本年|上年|本期|上期)(?:数|发生额)?为\s*`?(\d{4}-\d{2}-\d{2})`?\s*至\s*`?(\d{4}-\d{2}-\d{2})',text)}
    def column_metrics(value):
        value=value.replace('`','')
        columns={}
        for part in re.split(r'[；;。]',value):
            ids=metrics.findall(part)
            if len(ids)!=1:continue
            if declined.search(part) or re.search(re.escape(ids[0])+r'\s*(?:不适用|不采用)',part):continue
            before,after=part.split(ids[0],1)
            declared=r'(?<![A-Z0-9])([A-Z]+(?:[、，,/]\s*[A-Z]+)*)\s*(?:列|[:：=]|(?=\s*(?:用|使用|适用|采用)))'
            matches=list(re.finditer(declared,before)) or list(re.finditer(declared,after))
            for match in matches:
                for column in re.findall(r'[A-Z]+',match[1]):columns[column]=ids[0]
        # 两段“分别”以原文唯一的列名末尾简称相接，不按财务科目推定编号。
        labels={}
        for item in re.finditer(r'([A-Z]+(?:[、，,/][A-Z]+)*)列分别为([^，,。；;]+)',value):
            cols=re.findall(r'[A-Z]+',item[1]);names=item[2].split('、')
            if len(cols)==len(names):labels.update(zip(cols,names))
        labels.update((col,label.strip()) for col,label in re.findall(r'([A-Z]+)列为([^；;。]+)',value))
        linked=[]
        for item in re.finditer(r'([\u4e00-\u9fff]+)格适用\s*(.*?)(?=[；;。]|[，,]\s*[\u4e00-\u9fff]+格适用|$)',value):
            ids=set(metrics.findall(item[2]))
            if len(ids)==1:linked.append((item[1],next(iter(ids))))
        for column,label in labels.items():
            ids={key for name,key in linked if label.endswith(name)}
            if len(ids)==1:
                key=next(iter(ids))
                if column in columns and columns[column]!=key:columns.pop(column)
                else:columns[column]=key
        for item in re.finditer(r'([\u4e00-\u9fff]+(?:、[\u4e00-\u9fff]+)+)分别选\s*([^；;。]+)',value):
            names=item[1].split('、');ids=metrics.findall(item[2])
            if len(names)!=len(ids):continue
            for name,key in zip(names,ids):
                matched=[col for col,label in labels.items() if label.endswith(name)]
                if len(matched)==1 and sum(labels[matched[0]].endswith(n) for n in names)==1:columns[matched[0]]=key
        same_columns=re.findall(r'([A-Z]+)列槽位相同',value)
        if same_columns and len(set(columns.values()))==1:
            columns.update((column,next(iter(columns.values()))) for column in same_columns)
        return columns
    date_pattern=r'\d{4}-\d{2}-\d{2}(?:\s*至\s*\d{4}-\d{2}-\d{2})?'
    common_periods={};table_period_rows=[];section_metric_start=0
    for declaration in re.finditer(r'期间角色与日期[:：]([^。\n]+)',text.replace(chr(96),'').replace('**','')):
        for value in re.finditer(r'(本期|上期)\s*[=＝]\s*\d{4}年度[（(]('+date_pattern+r')[）)]',declaration[1]):
            common_periods[value[1]]=re.sub(r'\s+','',value[2])
    lines=[]
    for source_line in text.splitlines():
        sentences=[]
        for sentence in re.split(r'(?<=。)',source_line):
            if not sentence.strip():continue
            if sentences and (not coordinate.search(sentence) or
                re.match(r'\s*(?:对应原表(?:合并表头|期间标签|标签)|(?:原表)?期间标签(?:坐标[:：]|位于))',sentence) or
                re.match(r'\s*(?:Word|A|B)?对应',sentence) and '同为' in sentence):sentences[-1]+=sentence
            else:sentences.append(sentence)
        column_relation=field_scope and all(set(keys)<={'period','template_period_role'} for keys in field_scope.values()) and '分别为' in source_line and ('金额列为' in source_line or '为列标题' in source_line)
        common_line=re.match(r'^\s*(?:[AB]\s*)?各格共同维度\s*[:：]',source_line.replace('**',''))
        source_bullet=bool(re.match(r'\s*-\s*B\s+`?'+coordinate.pattern,source_line))
        separate=not source_bullet and not common_line and not column_relation and len(sentences)>1 and not source_line.lstrip().startswith('|') and all(
            coordinate.search(s) and s.count('`')%2==0 and
            s.count('“')==s.count('”') for s in sentences)
        lines.extend(sentences if separate else [source_line])
    column_scope=re.compile(r'([A-Z]+)列共同适用\s*[:：]?$')
    for number,original in enumerate(lines):
        table_literals={};table_cell_ends=[]
        table_scope=None;table_metric_text=None
        if paragraph_scope is not None:
            scope=paragraph_scope;paragraph_scope=None
        source_line=original.strip()
        line=re.sub(r'\*\*([^*]+)\*\*',lambda m:'`'+m[1]+'`' if
            re.search(r'[:：=＝为]\s*$',source_line[:m.start()]) and '`' not in m[1] else m[1],source_line).replace('＝','=')
        line=annotated_field.sub(lambda m:' '+m[1]+'=',line)
        if re.search(r'(?:实际日期(?:为|[:：])|\bperiod`?\s*[=：:])\s*`?'+date_pattern,line):
            line=re.sub(r'^期间\s*[:：]\s*(?!`?\d{4})','',line)
        if shared_field_aliases:
            line=re.sub(r'((?:项目)?范围)同上',r'\1：同上',line)
            line=re.sub(r'((?:项目)?范围)(?:是|限于)',r'\1：',line)
        # 未指明位置的正反例说明不能扩成全表选定编号，逐项声明照常接入。
        mixed_metric_explanation=not coordinate.search(line) and bool(declined.search(line))
        if mixed_metric_explanation:
            result.setdefault('unscoped_metric_explanations',[]).append(original)
            line=metrics.sub('',line)
        line=re.sub(r'\bmetric_id`?\s+(?=未知|unknown)', 'metric_id=',line)
        line=re.sub(r'币种\s+`?([A-Z]{3})`?(?=$|[，,；;。])',r'币种=`\1`',line)
        line=re.sub(r'((?:实际)?期间为)\d{4}年度[（(]('+date_pattern+r')[）)]',r'\1\2',line)
        if re.match(r'^(?:上述B来源均为|以上[一二三四五六七八九十百0-9]+格的共同维度值)',line):
            preceding=[]
            for prior in reversed(lines[:number]):
                if not prior.strip():
                    if preceding:break
                    continue
                preceding.append(prior)
            declared=qualify(coordinate.findall('\n'.join(reversed(preceding))))
            if declared:scope=list(dict.fromkeys(declared))
        if 'template_period_role' in line:
            line=re.sub(r'([A-Z]+)列[“"][^”"]+[”"](?:和|及)([A-Z]+)列[“"][^”"]+[”"]',r'\1、\2列',line)
        line=re.sub(r'(?:两表|各表)均为([^，,。]+?)(合并报表|单体报表)',r'entity=\1；report_scope=\2',line)
        line=re.sub(r'币种\s*([^，,。]+[（(][A-Z]{3}[）)])',r'币种=\1',line)
        line=re.sub(r'金额倍率\s+(?=\d)',r'金额倍率=',line)
        for dated in re.finditer(r'([A-Z]+)列为\s*`?\d{4}年度[（(](\d{4}-\d{2}-\d{2}至\d{4}-\d{2}-\d{2})[）)]',line):
            column_periods[dated[1]]=dated[2]
        line=re.sub(r'(\d{4}-\d{2}-\d{2})[—–](\d{4}-\d{2}-\d{2})',r'\1至\2',line)
        line=re.sub(r'列\s*([A-Z]+(?:[、，,]\s*[A-Z]+)*)(?=为|均为)',r'\1列',line)
        if re.search(r'适用于\s*`?'+coordinate.pattern,line):
            line=re.sub(r'实际日期(?=\s*(?:为|[:：=]))','period',line)
        line=re.sub(r'([A-Z]+)列(?=[、，,]\s*[A-Z]+列)',r'\1',line)
        from openpyxl.utils.cell import column_index_from_string, get_column_letter
        line=re.sub(r'(?<![A-Z])([A-Z]+)\s*[:：]\s*([A-Z]+)(?=列)',
            lambda m:'、'.join(get_column_letter(c) for c in range(column_index_from_string(m[1]),column_index_from_string(m[2])+1)),line)
        line=re.sub(r'(template_period_role\s*[:：=]\s*`?)(?:本期|上期)(期初|期末)(?:余额)?',r'\1\2',line)
        if not coordinate.search(line) and re.match(r'^[#*\s-]*(?:共同(?:范围|实际取值|取值|维度)|适用于全部财务格的维度)',line):
            scope=['*']
        if re.search(r'(?:这些格|本表|本范围)[^。；;]{0,12}(?:没有|无)(?:可据以选择的|可用的|可适用的)?金标准槽位',line):
            result['unresolved'].append({'id':'metric_id','ranges':scope[:],'reason':'原答复明确无适用槽位','original':original})
            continue
        line=re.sub(r'([A-Z]+(?:[、，,][A-Z]+)*)列第(\d+):(\d+)行',
            lambda m:'、'.join(col+m[2]+':'+col+m[3] for col in re.findall(r'[A-Z]+',m[1])),line)
        line=re.sub(r'(?<![A-Z])([A-Z]+)(\d+):([A-Z]+)(\d+)中([A-Z]+)列',
            lambda m:m[5]+m[2]+':'+m[5]+m[4] if column_index_from_string(m[1])<=column_index_from_string(m[5])<=column_index_from_string(m[3]) else m[0],line)
        line=re.sub(r'→\s*指定数据范围\s*','适用于 ',line)
        line=re.sub(r'。(?=对应原表(?:合并表头|期间标签|标签))','，',line)
        reference=re.search(r'`?('+coordinate.pattern+r')`?\s+所在[^。；;]*?为\s*`?('+coordinate.pattern+r')',line)
        if reference:line=line[:reference.start(1)]+line[reference.end(1):]
        line=re.sub(r'(\d{4})年(\d{1,2})月(\d{1,2})日',lambda m:f'{m[1]}-{int(m[2]):02d}-{int(m[3]):02d}',line)
        line=re.sub(r'(\d{4})-(\d{2})-(\d{2})至(\d{1,2})月(\d{1,2})日',lambda m:f'{m[1]}-{m[2]}-{m[3]}至{m[1]}-{int(m[4]):02d}-{int(m[5]):02d}',line)
        line=re.sub(r'`(\d{4}-\d{2}-\d{2})`\s*至\s*`(\d{4}-\d{2}-\d{2})`',r'`\1至\2`',line)
        if not line or line.startswith('```'):continue
        column_text=line.replace('`','')
        for declared in re.finditer(r'[“"](本期|上期|期初|期末)(?:发生额|余额)?[”"]对应(?:期间)?\s*('+date_pattern+r')',column_text):
            named_periods[declared[1]]=declared[2]
        for common in re.finditer(r'(本期|上期)(?:=|＝|为|对应)\s*(?:\d{4}年度[（(])?\s*('+date_pattern+r')',column_text):
            named_periods[common[1]]=re.sub(r'\s+','',common[2])
        period_pair=re.search(r'[“"](本期|上期|期初|期末)[／/](本期|上期|期初|期末)[”"][^。]*?分别保留期间[：:]\s*('+date_pattern+r')[、，,]\s*('+date_pattern+r')',column_text)
        if period_pair:named_periods.update({period_pair[1]:period_pair[3],period_pair[2]:period_pair[4]})
        pair=re.fullmatch(r'\s*-?\s*(\w+)\s*[/／]\s*(\w+)\s*[:：]\s*([^/／]+)[/／]([^/／]+)',column_text)
        if pair and pair[1] in standard['dimension'] and pair[2] in standard['dimension']:
            line=f'{pair[1]}={pair[3]}；{pair[2]}={pair[4]}';column_text=line
        for stated in re.finditer(r'([A-Z]+)至([A-Z]+)列(?:依次)?(?:是|为)([^。]+)',column_text):
            from openpyxl.utils.cell import column_index_from_string, get_column_letter
            columns=[get_column_letter(i) for i in range(column_index_from_string(stated[1]),column_index_from_string(stated[2])+1)]
            values=re.split(r'[、，,]',stated[3])
            if len(columns)==len(values):
                for column,value in zip(columns,values):
                    role=re.match(r'\s*(期末|期初|本期|上期)',value)
                    if role:column_roles[column]=role[1]
        column_period_declared=False
        declarations=list(re.finditer(r'([A-Z]+(?:[、，,]\s*[A-Z]+)*)\s*(?:列(?:角色)?(?:为)?|各格的)\s*[:：]?\s*',column_text))
        for i,declaration in enumerate(declarations):
            clause=column_text[declaration.end():declarations[i+1].start() if i+1<len(declarations) else len(column_text)]
            stated_role=re.search(r'template_period_role\s*[:：=]\s*(期末|期初|本期|上期|closing|opening|current|prior)',clause)
            if not stated_role:stated_role=re.search(r'(?:角色|各行|各格)(?:[（(][^（）()]*[）)])?(?:为)?\s*(期末|期初|本期|上期)',clause)
            stated_date=re.search(r'(?:period\s*(?:[:：=]|均为|为)|实际(?:日期|期间|时点)(?:为)?|[、，,]\s*(?:期间为|日期))\s*('+date_pattern+r'|\d{4}年度)',clause)
            if not stated_date:stated_date=re.match(r'(?:期末|期初|本期|上期)[，,]\s*日期\s*('+date_pattern+r')',clause)
            if not stated_date:stated_date=re.match(r'('+date_pattern+r')',clause)
            if stated_date and not stated_role:
                stated_role=(re.match(r'(期末|期初|本期|上期)(?=$|[，,；;\s])',clause)
                    or re.match(r'\s*[（(](期末|期初|本期|上期)[）)]',clause[stated_date.end():]))
            if stated_role or stated_date:
                column_period_declared=True
                role={'closing':'期末','opening':'期初','current':'本期','prior':'上期'}.get(stated_role[1],stated_role[1]) if stated_role else None
                for column in re.findall(r'[A-Z]+',declaration[1]):
                    if role:column_roles[column]=role
                    if stated_date:column_periods[column]=stated_date[1]
                    reference=re.search(r'对应\s*('+coordinate.pattern+r')',clause)
                    if not reference:reference=re.search(r'[，,]\s*('+coordinate.pattern+r')[，,]\s*原文',clause)
                    if reference:column_label_refs[column]=reference[1]
                    else:column_label_refs.pop(column,None)
                if role and stated_date:named_periods[role]=stated_date[1]
        for clause in re.split(r'[；;。]',column_text):
            columns=[col for group in re.findall(r'([A-Z]+(?:[、，,]\s*[A-Z]+)*)列',clause) for col in re.findall(r'[A-Z]+',group)]
            role=re.search(r'template_period_role\s*[:：=]\s*(期末|期初|本期|上期)',clause)
            date=re.search(r'(?:period\s*(?:为|[:：=])|时点(?:为)?|期间(?:为)?)\s*('+date_pattern+r')',clause)
            if columns and role and date:
                column_roles.update((column,role[1]) for column in columns)
                column_periods.update((column,date[1]) for column in columns)
        for declared in re.finditer(r'([A-Z]+(?:[、，,][A-Z]+)*)列[^。；;]*?[（(](\w+)\s*=\s*([^（）()]+)[）)]',column_text):
            if declared[2] in standard['dimension']:
                for column in re.findall(r'[A-Z]+',declared[1]):column_defaults[(column,declared[2])]=(declared[3],original)
        # 同句分别声明多组列时，先记下原话中的列归属，避免两个期间都落到共同范围。
        column_declarations=list(re.finditer(r'([A-Z]+(?:[、，,]\s*[A-Z]+)*)\s*列(?:共同适用|为)?\s*(?:`?period\s*[:：=]\s*`?('
            +date_pattern+r')`?[^；;。]*?)?`?template_period_role\s*[:：=]\s*`?(期末|期初|本期|上期|closing|opening|current|prior)\b',line))
        for declared in column_declarations:
            role={'closing':'期末','opening':'期初','current':'本期','prior':'上期'}.get(declared[3],declared[3])
            date=re.sub(r'\s+','',declared[2]) if declared[2] else None
            if date:named_periods[role]=date
            for column in re.findall(r'[A-Z]+',declared[1]):
                if date:column_periods[column]=date
                column_roles[column]=role
        for declared in re.finditer(r'([A-Z]+)列[“「"]([^”」"]+)[”」"]\s*[=：:]\s*`?(?:template_period_role\s*[:：=]\s*)?(期初|期末|本期|上期|closing|opening|current|prior)`?',line):
            role={'closing':'期末','opening':'期初','current':'本期','prior':'上期'}.get(declared[3],declared[3])
            column_roles[declared[1]]=role
            date=re.match(r'\s*[，,]\s*(?:时点|期间)\s*`?('+date_pattern+r')',line[declared.end():])
            if date:column_periods[declared[1]]=date[1];named_periods[role]=date[1]
        for declared in reversed(column_declarations):line=line[:declared.start()]+line[declared.end():]
        if re.match(r'A(?:所需|需要)',line) and not coordinate.search(line) and re.search(r'未找到|未列',line):continue
        if not coordinate.search(line) and re.search(r'(?:而非|不是)A(?:对应|所需|需要)',line):continue
        field_heading=re.match(r'^\s*(?:[-*]\s*)?\*\*`?([A-Za-z_]+)`?\*\*',original)
        field_heading=field_heading[1] if field_heading and field_heading[1] in standard['dimension'] else None
        declared=[key for key in standard['dimension'] if re.search(r'(?<![\w.])'+re.escape(key)+r'(?![\w.])',line)]
        if not coordinate.search(line) and line.rstrip().endswith(('：',':')) and len(declared)==1:
            field_heading=declared[0]
        field_first=re.match(r'^\s*[-*]?\s*`?(\w+)`?\s*[:：]\s*(.+)$',line)
        if field_first and field_first[1] in standard['dimension']:
            stated=re.match(r'(.+?)\s*(?:均为|为)\s*(.+)$',field_first[2])
            if stated and coordinate.search(stated[1]) and not coordinate.sub('',stated[1]).strip(' `、，,；;'):
                line=stated[1]+'：'+field_first[1]+'='+stated[2]
        section_heading=bool(re.match(r'^\s*(?:[-*]\s*)?\*\*[^*]+\*\*(?:\s*[（(].*)?$',original))
        if section_heading:
            selections=[key for key in declared if standard['dimension'][key]['value_type']=='selection']
            range_heading=selections[0] if len(selections)==1 else None
        if line.startswith('#') or field_heading:
            range_heading=None
            heading_metric=metric_aliases.get(line.strip(' #`'))
            if line.startswith('#'):
                scope=['*'];metric_group=None;bullet_scopes=[];section_metric_start=len(result['metrics'])
            row_roles=[];section_roles={};section_lines=[];declared_fields=set();section_duration=None;row_texts={}
            row_metric_columns={};row_objects={}
            member_field=None;member_shared={};member_enums=[]
            item_groups=[];group_rules=[];group_direction=None;labeled_cells=[]
            fields=[key for key in standard['dimension'] if re.search(r'(?<![\w.])'+re.escape(key)+r'(?![\w.])',line)]
            section_fields=fields
            row_field=fields[0] if len(fields)==1 else None
            common_selection=fields[0] if len(fields)==1 and standard['dimension'][fields[0]]['value_type']=='selection' else None
            denominators=[key for key in fields if key.endswith('_denominator_selection')]
            if len(denominators)==1:ratio_names['分母']=denominators[0]
        if not coordinate.search(line) and re.search(r'以下各(?:坐标|行|格)均为\s*`?mode\s*=\s*members',line):
            fields=[key for key in standard['dimension'] if re.search(r'(?<![\w.])'+re.escape(key)+r'(?![\w.])',line)
                    and standard['dimension'][key].get('value_type')=='selection']
            if len(fields)==1:
                common_selection=fields[0]
                selection_defaults[common_selection]={'domain':standard['dimension'][common_selection]['domain'],'mode':'members'}
                continue
        by_row=re.search(r'\x60?(\w+)\x60?\s*[:：=]\s*按各行[^。]*分别识别',line)
        if by_row and standard['dimension'].get(by_row[1],{}).get('value_type')=='selection':
            common_selection=by_row[1];continue
        common_mode=re.search(r'(\w+)`?\s*均为\s*`?(members|predicate)\b',line)
        group_mode=re.search(r'(\w+)`?\s*[、，,]\s*`?mode\s*=\s*members',line)
        if group_mode and not coordinate.search(line) and standard['dimension'].get(group_mode[1],{}).get('value_type')=='selection':
            common_selection=group_mode[1]
            continue
        if common_mode and not coordinate.search(line) and standard['dimension'].get(common_mode[1],{}).get('value_type')=='selection':
            common_selection=common_mode[1]
            selection_defaults[common_selection]={'mode':common_mode[2]}
            if not re.search(r'[:：]\s*\S',line[common_mode.end():]):continue
        if ratio_names.get('分母') in section_fields:
            line=re.sub(r'(?<=[；;])\s*分母(?=为)',lambda m:' '+ratio_names['分母'],line)
        shared=re.search(r'(?:均使用金标准维度|以下仅列[^。；;]*?的|以下均为)\s+`?(\w+)',line)
        if shared and standard['dimension'].get(shared[1],{}).get('value_type')=='selection':
            common_selection=shared[1]
            if re.search(r'选取方式均为\s*`?members\b',line):
                selection_defaults[common_selection]={'mode':'members'}
        first=next(marker.finditer(line),None)
        if first and len(coordinate.findall(line[:first.start()].rsplit('。',1)[-1]))==1:
            line=re.sub(r'([；;]\s*`?(?:'+coordinate.pattern+r')`?\s*)为(?=[“`])',
                        lambda m:m[1]+'`'+first[1]+'`为',line)
        if '自然年度' in text:
            line=re.sub(r'(\d{4})年度',lambda m:f'{m[1]}-01-01至{m[1]}-12-31',line)
        declared_ratio=set(metrics.findall(line))
        if len(declared_ratio)==1:
            ratio=standard['metric'][next(iter(declared_ratio))].get('definition',{}).get('ratio',{})
            for label,side in (('分子','numerator'),('分母','denominator')):
                keys={v.get('key') for v in ratio.get(side,{}).get('dimension_bindings',{}).values()
                      if v.get('op')=='dimension' and standard['dimension'].get(v.get('key'),{}).get('value_type')=='selection'}
                if len(keys)==1:
                    key=next(iter(keys));line=line.replace(label+'范围同上',key+'：同上')
                    if not re.search(r'(?<![\w.])'+re.escape(key)+r'(?![\w.])',line):
                        line=re.sub(r'(?<![A-Za-z_])'+label+r'(?:范围)?(?=为|[:：])',key,line)
        if last_metric and last_selection:
            line=line.replace('同一编号及选择范围',last_metric+'；'+last_selection+'：同上')
        if last_selection and re.search(r'同(?:一|变动)?(?:槽位|指标)',line):
            line=re.sub(r'(?:该)?选择维度(?:取值)?\s*(?=[:：=])',last_selection,line)
        if last_selection:
            line=re.sub(r'(?<=[：:])\s*同一总体范围(?=[；;。])',last_selection+'：同上',line)
        if range_heading:
            line=re.sub(r'(?<![\w])(?:范围|对象)(?=为|[:：=])',range_heading,line)
        explicit=list(dict.fromkeys(key for key in metrics.findall(line) if key in standard['metric']))
        refs=[]
        if len(explicit)==1:
            refs=[r['id'] for r in standard['metric'][explicit[0]].get('dimension_refs',[])]
            for label,suffix in (('活动','_activity'),('方向','_direction'),('列报口径','_presentation'),
                                 ('项目范围','_selection'),('选择维度','_selection')):
                keys=[key for key in refs if key.endswith(suffix)]
                if len(keys)==1:
                    line=re.sub(r'(?<![\w])'+label+r'(?:取值)?\s*(?=[:：=]|(?:分别|依次)?为)',keys[0],line)
            changes=[key for key in refs if key.endswith('_change_selection')]
            if len(changes)==1:
                line=re.sub(r'(?<![\w])变动事项(?:范围)?\s*(?=[:：=]|(?:均|分别|依次)?为)',changes[0],line)
                objects=[key for key in refs if key.endswith('_selection') and key!=changes[0]]
                if len(objects)==1:
                    line=re.sub(r'对象(?:范围)?\s*(?=[:：=]|(?:分别|依次)?为)',' '+objects[0]+' ',line)
        for label,key in re.findall(r'(分子|分母)(?:维度)?\s*`(\w+)`',line):
            if key in standard['dimension']:ratio_names[label]=key
        combined=re.search(r'分子和分母的选择均为\s*(.+)',line)
        if combined and set(ratio_names)=={'分子','分母'}:
            line=line[:combined.start()]+'；'.join(key+'='+combined[1] for key in ratio_names.values())
        paired_fields=re.search(r'分子\s*`?(\w+)`?\s*与分母\s*`?(\w+)`?\s*均为\s*(.+)',line)
        if paired_fields and all(key in standard['dimension'] for key in paired_fields.group(1,2)):
            line=line[:paired_fields.start()]+'；'.join(key+'='+paired_fields[3] for key in paired_fields.group(1,2))
        mentioned=[name for name in sheet_names if any(s in line for s in
            ('`'+name+'`',"'"+name+"'!",'`'+name+'!','「'+name+'」','《'+name+'》','“'+name+'”','"'+name+'"'))
            or line.startswith('#') and line.endswith(name)]
        explicit_sheets=list(dict.fromkeys(a.rsplit('!',1)[0].strip("'") for a in coordinate.findall(line)
                         if '!' in a and a.rsplit('!',1)[0].strip("'") in sheet_names))
        mentioned.extend(name for name in sheet_names if re.fullmatch(
            re.escape(name)+r'(?:[（(][^（）()]*[）)])?',line.strip(' #`')))
        heading_sheet=next((name for name in sheet_names if line.startswith((name+'：',name+':'))),None)
        mentioned=list(dict.fromkeys(mentioned))
        previous_sheet=sheet
        first_coordinate=coordinate.search(line)
        if field_scope and not explicit_sheets:
            stated={cell for area in coordinate.findall(line) for cell in addresses(area.rsplit('!',1)[-1].replace('$',''))}
            owners={location.rsplit('!',1)[0] for location in requested if location.rsplit('!',1)[1] in stated}
            if len(owners)==1:sheet=next(iter(owners))
        # 引用来源位于答复值内，不改变本段目标坐标的所属表。
        if len(explicit_sheets)==1 and first_coordinate and ('!' in first_coordinate[0] or
                field_scope and all(set(keys)<={'period','template_period_role'} for keys in field_scope.values())):
            sheet=explicit_sheets[0]
        elif heading_sheet:sheet=heading_sheet
        elif not first_coordinate and len(mentioned)==1:sheet=mentioned[0]
        if previous_sheet and sheet!=previous_sheet:
            column_periods={};column_roles={};named_periods={};row_objects={};column_directions={};last_column_roles={};column_defaults={}
            metric_group=None;period_boundaries={};column_label_refs={}
        def qualify(areas):
            return [sheet+'!'+a if sheet and '!' not in a and a!='*' else a for a in areas]
        if line.startswith('#') and requested:
            # 列声明仅展开到本段明示表区内的请求位置，不沿用上一张表的范围。
            stated={area.rsplit('!',1)[0]+'!'+cell for area in qualify(coordinate.findall(line)) if '!' in area
                    for cell in addresses(area.rsplit('!',1)[-1])}
            column_table_scope=[location for location in requested if location in stated] or None
        for declaration in re.finditer(r'适用于\s*('+coordinate.pattern+r')[^；。\n]*?实际日期\s*('+date_pattern+r')',line.replace(chr(96),'')):
            result['dimensions'].append({'id':'period','ranges':qualify([declaration[1]]),'value':declaration[2],'original':original})
        source_item=re.match(r'^-\s*B\s+`?('+coordinate.pattern+r')`?\s+[^：:]+[：:]',line)
        source_values=list(re.finditer(r'(本期|上期|期初|期末)\s*(?:值来源\s*)?`?('+coordinate.pattern+r')',line))
        if source_item and source_values:
            ids=metrics.findall(line)
            selected=ids[0] if len(ids)==1 else last_metric if re.search(r'同一指标',line) else None
            member=re.search(r'成员取值[“"]([^”"]+)[”"]',line)
            value={'mode':'members','members':[member[1]]} if member else 'all' if re.search(r'mode\s*=\s*all',line) else None
            locations=qualify([item[2] for item in source_values])
            if selected:
                last_metric=selected;result['metrics'].append({'id':selected,'ranges':locations,'original':original})
            if common_selection and value is not None:
                result['dimensions'].append({'id':common_selection,'ranges':locations,'value':value,'original':original})
            for item in source_values:
                located=qualify([item[2]])
                result['dimensions'].append({'id':'template_period_role','ranges':located,'value':{'本期':'current','上期':'prior','期初':'opening','期末':'closing'}[item[1]],'original':original})
                if item[1] in named_periods:result['dimensions'].append({'id':'period','ranges':located,'value':named_periods[item[1]],'original':original})
            label=qualify([source_item[1]])[0]
            if label not in locations:result.setdefault('non_business',[]).append({'range':label,'category':'label','reason':'答复明示的行名称位置；金额来源另列','original':original})
            continue
        indent=len(original)-len(original.lstrip())
        while bullet_scopes and indent<=bullet_scopes[-1][0]:bullet_scopes.pop()
        nested_scope=bool(bullet_scopes)
        if nested_scope:scope=bullet_scopes[-1][1][:]
        declared_table=re.match(r'^(?:范围\d*|目标范围)\s*[:：｜|]\s*('+coordinate.pattern+r')(?=$|[。；;])',line.replace(chr(96),''))
        if declared_table:column_table_scope=qualify([declared_table[1]])
        selected_source=re.match(r'^\s*[-*]?\s*(?:选择|采用|选用)的\s*B来源\s*[:：]\s*('+coordinate.pattern+r')',line.replace('`',''))
        if selected_source:
            scope=qualify([selected_source[1]]);metric_group=None
            continue
        bullet_heading=re.fullmatch(r'\s*[-*]\s+\*\*([^*]+)\*\*\s*',original)
        if bullet_heading and coordinate.search(bullet_heading[1]) and not marker.search(line) and not metrics.search(line):
            scope=qualify(coordinate.findall(bullet_heading[1]));bullet_scopes.append((indent,scope[:]))
            metric_group=None
            continue
        if nested_scope and column_period_declared:
            for declaration in declarations:
                for column in re.findall(r'[A-Z]+',declaration[1]):
                    located=[(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell
                        for area in scope for cell in addresses(area.rsplit('!',1)[-1])
                        if re.match(r'[A-Z]+',cell)[0]==column]
                    if not located:continue
                    for key,values in (('period',column_periods),('template_period_role',column_roles)):
                        if column in values:result['dimensions'].append({'id':key,'ranges':located,'value':values[column],'original':original})
        label_ref=re.search(r'(?:原表)?期间标签(?:引用|坐标)?\s*(?:为|是|[:：=])\s*`?('+coordinate.pattern+r')',line)
        if label_ref:
            head=re.split(r'[:：](?![A-Z]+\d)',line[:label_ref.start()],maxsplit=1)[0]
            areas=qualify(coordinate.findall(head))
            if areas:period_links.append((qualify([label_ref[1]]),areas))
        # 名单答复可用表头和金额列一次声明期间引用，不要求逐格重述。
        plain_columns=line.replace('`','')
        if field_scope and all(set(keys)<={'period','template_period_role'} for keys in field_scope.values()) and not re.search(date_pattern,plain_columns):
            table_sheet=(first_coordinate[0].rsplit('!',1)[0].strip("'")
                         if first_coordinate and '!' in first_coordinate[0] else sheet)
            header_pairs=re.search(r'('+coordinate.pattern+r')[、，,]\s*('+coordinate.pattern+r')\s*分别为(本期|上期|期初|期末)[、，,](本期|上期|期初|期末)',plain_columns)
            amount_columns=re.search(r'金额列为\s*([^。；;]+)',plain_columns)
            header_band=re.search(r'('+coordinate.pattern+r')\s*为列标题',plain_columns)
            column_pair=re.search(r'([A-Z]+)[、，,]([A-Z]+)\s*列分别为(本期|上期|期初|期末)[、，,](本期|上期|期初|期末)',plain_columns)
            links=[]
            if table_sheet in sheet_names and header_pairs and amount_columns:
                for ref in header_pairs.group(1,2):
                    column=re.match(r'[A-Z]+',ref.rsplit('!',1)[-1])[0]
                    areas=[area for area in coordinate.findall(amount_columns[1])
                           if all(re.match(r'[A-Z]+',cell)[0]==column for cell in addresses(area.rsplit('!',1)[-1]))]
                    links.append((ref,areas))
            elif table_sheet in sheet_names and header_band and column_pair:
                header_cells=addresses(header_band[1].rsplit('!',1)[-1]);header_rows={re.search(r'\d+$',cell)[0] for cell in header_cells}
                row_areas=re.findall(r'('+coordinate.pattern+r')\s*(?:是|为)\s*(?:明细及占位行|合计)',plain_columns)
                if len(header_rows)==1 and row_areas:
                    rows=sorted({int(re.search(r'\d+$',cell)[0]) for area in row_areas for cell in addresses(area.rsplit('!',1)[-1])})
                    for column in column_pair.group(1,2):
                        ref=column+next(iter(header_rows))
                        if ref in header_cells:links.append((ref,[column+str(row) for row in rows]))
            if any(areas for ref,areas in links):
                for ref,areas in links:
                    qualify_column=lambda area:area if '!' in area else table_sheet+'!'+area
                    if areas:period_links.append(([qualify_column(ref)],[qualify_column(area) for area in areas]))
                continue
        label=re.search(r'(?<=[：:。；;])\s*(?:原表)?(?:期间)?标签(?:坐标[:：]|位于)\s*`?('
            +coordinate.pattern+r')`?\s*(?:[，,]\s*(?:为|原文)\s*)?[“"「]((?:期初|期末|本期|上期)[^”"」]*)[”"」]',line)
        if label:
            areas=qualify(coordinate.findall(line[:label.start()]))
            if areas:
                refs=qualify([label[1]]);period_links.append((refs,areas))
                result['dimensions'].append({'id':'template_period_role','ranges':areas,
                    'value':{'期初':'opening','期末':'closing','本期':'current','上期':'prior'}[label[2][:2]],'original':original})
                date=re.search(r'(?:本期|上期)?期间(?:同)?为\s*`?('+date_pattern+r')',line)
                if date:
                    result['dimensions'].append({'id':'period','ranges':areas,
                        'value':re.sub(r'\s+','',date[1]),'evidence':refs,'original':original})
                # 标签是引用位置，不是新数据目标；同段其他字段继续按原文接入。
                line=line[:label.start()]+line[label.end():]
                if not marker.search(line):continue
        candidate_heading=re.match(r'^[AB]?候选[:：]\s*(.+)$',line)
        if candidate_heading and not coordinate.sub('',candidate_heading[1]).strip(' `、，,。'):
            scope=qualify(coordinate.findall(candidate_heading[1]));continue
        if field_heading and not coordinate.search(line) and line.rstrip().endswith(('：',':')) and 'domain' not in line:continue
        head,sep,body=line.partition('：')
        located=qualify(coordinate.findall(head)) if sep else []
        if requested and located and not marker.search(body) and not metrics.search(body):
            cells=[area.rsplit('!',1)[0]+'!'+cell for area in located if '!' in area
                   for cell in addresses(area.rsplit('!',1)[1])]
            keys={key for cell in cells for key in requested.get(cell,())}
            remainder=re.sub(r'[（(][^（）()]*[）)]','',coordinate.sub('',head)).strip(' -*`、，,；;')
            if cells and not remainder and len(keys)==1 and all(requested.get(cell)==keys for cell in cells):
                key=next(iter(keys))
                if body.strip() and standard['dimension'].get(key,{}).get('value_type')=='text':
                    # 请求已唯一指定开放文本字段，只搬运本段原话，不推断字段或改写含义。
                    result['dimensions'].append({'id':key,'ranges':cells,'value':body.strip(),'original':original})
                    continue
        linked_labels=False;previous_areas=[]
        if '标签' in line or '标识' in line:
            role_labels={}
            for labelled in re.finditer('('+coordinate.pattern+r')\s*[“"](本期|上期|期末|期初)(?:发生额|余额)?[”"]',line.replace('`','')):
                role_labels.setdefault(labelled[2],[]).extend(qualify([labelled[1]]))
            for stated in re.finditer('('+coordinate.pattern+r')\s*[（(](本期|上期|期末|期初)[）)]',line.replace('`','')):
                if len(role_labels.get(stated[2],[]))==1:
                    period_links.append((role_labels[stated[2]],qualify([stated[1]])))
        period_request=bool(field_scope and all(set(keys)<={'period','template_period_role'} for keys in field_scope.values()))
        if period_request:
            for link in re.finditer(r'('+coordinate.pattern+r')\s*[“\"「][^”\"」]+[”\"」]\s*对应\s*((?:'+coordinate.pattern+r'[、，,\s]*)+)',line):
                period_links.append((qualify([link[1]]),qualify(coordinate.findall(link[2]))))
                linked_labels=True
        header_rows=re.findall(r'第(\d+)行',body) if sep else []
        header_columns=re.findall(r'([A-Z]+)列[“"「][^”"」]+[”"」]',body) if sep else []
        if period_request and located and len(header_rows)==len(located) and header_columns and '列标签' in body:
            # 按答复列出的段、表头行和列逐项展开，不从表格位置猜标签。
            for area,number in zip(located,header_rows):
                name,span=area.rsplit('!',1)
                for column in header_columns:
                    cells=[name+'!'+cell for cell in addresses(span) if re.match(r'[A-Z]+',cell)[0]==column]
                    if cells:period_links.append(([name+'!'+column+number],cells))
            continue
        for clause in re.split(r'[；;](?!\s*角色(?:为|是))|[。]|[，,](?=\s*(?:`?(?:'+coordinate.pattern+r')`?[、\s]*)+对应)',line):
            column_prefix=re.match(r'\s*对应的\s*([A-Z]+)\s*列数据格',clause)
            if column_prefix:
                clause=clause[column_prefix.end():]
                if clause.lstrip().startswith('对应'):
                    # 答复明说沿用前句行范围，交给已有坐标展开；不从原表推定。
                    clause+='适用于对应的 '+column_prefix[1]+'列各范围'
            before,sep,after=clause.partition('适用于')
            corresponding=re.fullmatch(r'\s*对应的\s*([A-Z]+)\s*列各范围\s*',after) if sep else None
            if corresponding and previous_areas:
                columns=[re.fullmatch(r'([A-Z]+)\d+(?::\1\d+)?',area.rsplit('!',1)[-1]) for area in previous_areas]
                if all(columns) and len({m[1] for m in columns})==1:
                    after='、'.join((area.rsplit('!',1)[0]+'!' if '!' in area else '')+
                        re.sub(r'[A-Z]+(?=\d)',corresponding[1],area.rsplit('!',1)[-1]) for area in previous_areas)
            if sep and coordinate.search(after) and not before.rstrip().endswith('不') and (re.search(r'期间标签|标签(?:原文)?|[“\"「](?:期初|期末|本期|上期|本年|上年)',before) or period_request and re.search(r'[“\"「][^”\"」]+[”\"」]',before)):
                anchor=re.search(r'标签(?:位于|坐标)?\s*`?('+coordinate.pattern+r')',before)
                refs=qualify([anchor[1]] if anchor else coordinate.findall(before))
                areas=qualify(coordinate.findall(after));label_text=before
                paired=before.rstrip().endswith(('分别','依次'))
            else:
                before,sep,after=re.sub(r'适用(?:原表标签|标签)?(?!于)','对应',clause).partition('对应')
                if not sep or before.rstrip().endswith('不') or not (re.search(r'[“\"「](?:期初|期末|本期|上期|本年|上年)',after) or period_request and re.search(r'[“\"「][^”\"」]+[”\"」]',after)):continue
                areas=qualify(coordinate.findall(before));refs=qualify(coordinate.findall(after));label_text=after
                paired=before.rstrip().endswith(('分别','依次'))
            if not refs or not areas or paired and len(refs)!=len(areas):continue
            previous_areas=areas[:]
            pairs=[([ref],[area]) for ref,area in zip(refs,areas)] if paired else [(refs,areas)]
            columns=re.search(r'分别适用于\s*((?:[A-Z]+\s*列[、和及与\s]*)+)指定范围',clause)
            if columns and len(refs)==len(re.findall(r'[A-Z]+',columns[1])):
                pairs=[([ref],[area.rsplit('!',1)[0]+'!'+cell for area in areas
                              for cell in addresses(area.rsplit('!',1)[-1]) if re.match(r'[A-Z]+',cell)[0]==column])
                       for ref,column in zip(refs,re.findall(r'[A-Z]+',columns[1]))]
            labels=set(re.findall(r'角色(?:为|是)(期初|期末|本期|上期)',clause))
            if not labels:labels=set(re.findall(r'[“\"「](期初|期末|本期|上期)',label_text))
            for refs,areas in pairs:
                period_links.append((refs,areas))
                if len(labels)==1:
                    result['dimensions'].append({'id':'template_period_role','ranges':areas,
                        'value':{'期初':'opening','期末':'closing','本期':'current','上期':'prior'}[next(iter(labels))],
                        'original':original})
            linked_labels=True
            if marker.search(clause):
                without_refs=coordinate.sub(lambda m:'' if qualify([m[0]])[0] in set(refs)-set(areas) else m[0],clause)
                line=line.replace(clause,without_refs)
        if linked_labels and not marker.search(line):continue
        if common_selection and known_sheets:
            for stated in reversed(list(re.finditer(r'('+known_sheets+r')!(\d+)(?::(\d+))?(?!\d)',line))):
                locations=[]
                for metric in result['metrics']:
                    for area in qualify(metric['ranges']):
                        if area=='*' or area.rsplit('!',1)[0]!=stated[1]:continue
                        locations.extend(stated[1]+'!'+cell for cell in addresses(area.rsplit('!',1)[-1])
                            if int(stated[2])<=int(re.search(r'\d+$',cell)[0])<=int(stated[3] or stated[2]))
                if locations:line=line[:stated.start()]+'、'.join(dict.fromkeys(locations))+line[stated.end():]
        plain=line.replace('`','')
        alias_declaration=re.search(r'以下“([^”]+)”均为\s*(\w+)',plain)
        if alias_declaration and standard['dimension'].get(alias_declaration[2],{}).get('value_type')=='selection':
            key=alias_declaration[2]
            declared_member=re.search(r'mode\s*=\s*members[、，,]\s*domain\s*=\s*([^、，,；;。]+)',plain)
            if declared_member:
                # 仅声明了成员方式，未给逐项成员；保留已知部分并撤下旧的含糊取值。
                selection_aliases[alias_declaration[1]]=(key,{'mode':'unknown','domain':declared_member[1],'known_scope':{'mode':'members'}})
            for alias in re.finditer(r'“([^”]+)”[^；;。]*?mode\s*=\s*all[、，,]\s*domain\s*=\s*([^、，,；;。]+)',plain):
                selection_aliases[alias[1]]=(key,{'mode':'all','domain':alias[2]})
            continue
        for alias,(key,value) in selection_aliases.items():
            for match in re.finditer(r'('+coordinate.pattern+r')\s*为(?:全体|各)?'+re.escape(alias),plain):
                result['dimensions'].append({'id':key,'ranges':qualify([match[1]]),'value':deepcopy(value),'original':original})
        enum_declaration=re.search(r'各[^；;，,]+格的\s*(\w+)\s*=\s*',plain)
        if enum_declaration and standard['dimension'].get(enum_declaration[1],{}).get('allowed_values'):
            key=enum_declaration[1];located=[]
            for value in standard['dimension'][key]['allowed_values']:
                for match in re.finditer(re.escape(value)+r'('+coordinate.pattern+r')',plain):
                    areas=qualify([match[1]]);located.extend(areas)
                    result['dimensions'].append({'id':key,'ranges':areas,'value':value,'original':original})
            if located:
                common=re.search(r'两者的\s*(\w+)\s*均为\s*domain\s*=\s*([^、，,；;。]+)[、，,]\s*mode\s*=\s*all',plain)
                if common and standard['dimension'].get(common[1],{}).get('value_type')=='selection':
                    result['dimensions'].append({'id':common[1],'ranges':located,
                        'value':{'domain':common[2],'mode':'all'},'original':original})
                continue
        selection_declaration=re.search(r'(\w+)的domain为“([^”]+)”',plain)
        if selection_declaration and standard['dimension'].get(selection_declaration[1],{}).get('value_type')=='selection':
            common_selection=selection_declaration[1]
            selection_defaults[common_selection]={'domain':selection_declaration[2]}
            if re.search(r'mode\s*=\s*members',plain):selection_defaults[common_selection]['mode']='members'
            continue
        if common_selection:
            for clause in re.split(r'[；;。]',plain):
                member=re.search(r'mode\s*=\s*members[，,]\s*成员“([^”]+)”',clause)
                if member:
                    located=qualify(coordinate.findall(clause[:member.start()]))
                    if located:
                        result['dimensions'].append({'id':common_selection,'ranges':located,
                            'value':{**selection_defaults.get(common_selection,{}),'mode':'members','members':[member[1]]},'original':original})
            sequence=re.search(r'范围：依次为“([^”]+)”',plain)
            if sequence:
                line=line[:line.index('范围：')] + common_selection+'依次为'+sequence[1]
            elif not line.startswith('|') and coordinate.search(line) and re.search(r'未识别出[^。]*成员|不是[^。]*成员',plain):
                for clause in re.split(r'[；;。]',plain):
                    located=qualify(coordinate.findall(clause))
                    if located and re.search(r'未识别出[^。]*成员|不是[^。]*成员',clause):
                        if '该说明行' in clause:
                            from .A语义阅读 import addresses
                            excluded={(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell for area in located
                                      for cell in addresses(area.rsplit('!',1)[-1])}
                            for row in result['dimensions']:
                                source=row.get('original','')
                                if row['id']==common_selection and isinstance(source,str) and '范围：依次为' in source and '除说明行外' in source:
                                    row['ranges']=[(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell for area in row['ranges']
                                        for cell in addresses(area.rsplit('!',1)[-1])
                                        if (area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell not in excluded]
                            result['dimensions']=[row for row in result['dimensions'] if row['ranges']]
                        result['dimensions'].append({'id':common_selection,'ranges':located,'value':'未知','original':original})
        shared_selection=re.search(r'(?<![A-Za-z0-9_])('+ '|'.join(re.escape(k) for k in standard['dimension'])+
            r')均为domain\s*=\s*([^,，；;]+)[,，]\s*mode\s*=\s*members',plain)
        if shared_selection and standard['dimension'].get(shared_selection[1],{}).get('value_type')=='selection':
            common_selection=shared_selection[1]
            selection_defaults[common_selection]={'domain':shared_selection[2],'mode':'members'}
            # 这里只声明后续成员列表的共同字段，不能当作缺成员值的完整身份。
            line=line.replace(re.search(r'`?'+re.escape(shared_selection[1])+r'`?均为`?domain[^。]+',line)[0],'')
            plain=line.replace('`','')
        member_list=re.search(r'members\s*=\s*\[([^\]]+)\]',plain)
        if common_selection and member_list:
            from .A语义阅读 import addresses
            located=qualify(coordinate.findall(plain.split('：',1)[0]))
            if located:
                result['dimensions'].append({'id':common_selection,'ranges':located,'value':{
                    **selection_defaults.get(common_selection,{}),'mode':'members',
                    'members':[v.strip(' "\'') for v in re.split(r'[,，、]',member_list[1])]},'original':original})
                if metric_group is not None:metric_group['ranges']=list(dict.fromkeys(metric_group['ranges']+located))
                for area in located:
                    for cell in addresses(area.rsplit('!',1)[-1]):
                        col=re.match(r'[A-Z]+',cell)[0];location=(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell
                        if col in column_periods:result['dimensions'].append({'id':'period','ranges':[location],'value':column_periods[col],'original':original})
                        if col in column_roles:result['dimensions'].append({'id':'template_period_role','ranges':[location],'value':column_roles[col],'original':original})
                continue
        quoted_members=re.search(r'(?:mode\s*[:：]\s*members[，,]\s*并列成员为|同为)((?:“[^”]+”\s*)+)',plain)
        previous_members=next((r['value'] for r in reversed(result['dimensions']) if r['id']==common_selection),None)
        if common_selection and quoted_members and ('mode' in quoted_members[0] or
                isinstance(previous_members,dict) and previous_members.get('mode')=='members'):
            located=qualify(coordinate.findall(plain.split('：',1)[0]))
            if located:
                result['dimensions'].append({'id':common_selection,'ranges':located,
                    'value':{'mode':'members','members':re.findall(r'“([^”]+)”',quoted_members[1])},'original':original})
                continue
        above_ranges={}
        above=re.match(r'^\s*-\s*上述各项的\s*(\w+)\s',plain)
        if above and above[1] in standard['dimension']:
            preceding=[]
            for statement in reversed(lines[:number]):
                if not statement.lstrip().startswith('-') or not metrics.search(statement):break
                preceding.append(statement)
            areas=list(dict.fromkeys(area for row in result['metrics']
                if row.get('original') in preceding for area in row['ranges']))
            if areas:above_ranges[above[1]]=areas
        if metric_context and re.search(r'分子|分母',plain):
            located=qualify(coordinate.findall(plain.split('：',1)[0]))
            from .A语义阅读 import addresses
            ratio_cells=[area.rsplit('!',1)[0]+'!'+cell for area in located if '!' in area
                         for cell in addresses(area.rsplit('!',1)[1])]
            ratios=[standard['metric'].get(metric_context.get(cell),{}).get('definition',{}).get('ratio',{}) for cell in ratio_cells]
            roles={}
            for label,part in (('分子','numerator'),('分母','denominator')):
                keys=[{v['key'] for v in ratio.get(part,{}).get('dimension_bindings',{}).values()
                       if v.get('op')=='dimension' and standard['dimension'].get(v.get('key'),{}).get('value_type')=='selection'} for ratio in ratios]
                if keys and len(keys[0])==1 and all(k==keys[0] for k in keys):roles[label]=next(iter(keys[0]))
            if len(roles)==2:
                handled=False
                both=re.search(r'分子[、和]分母均为([^。；;]+)',plain)
                if both:
                    result['dimensions'].extend({'id':key,'ranges':ratio_cells,'value':both[1],'original':original} for key in roles.values())
                    handled=True
                for label,key in roles.items():
                    literal=re.search(label+r'(?:均为|可识别为)\s*(\{[^{}]+\})',plain)
                    if literal:
                        try:value=json.loads(literal[1])
                        except ValueError:continue
                        result['dimensions'].append({'id':key,'ranges':ratio_cells,'value':value,'original':original})
                        handled=True
                reference=re.search(r'分子[^；;。]*(?:分别同上|对应[^；;。]*行)([\d–—\-、]+)',plain)
                if reference:
                    rows={r for start,end in re.findall(r'(\d+)[–—-](\d+)',reference[1]) for r in range(int(start),int(end)+1)}
                    key=roles['分子']
                    for cell in ratio_cells:
                        row=int(re.search(r'\d+$',cell)[0]);value=row_objects.get(key,{}).get(row)
                        if row in rows and value is not None:
                            result['dimensions'].append({'id':key,'ranges':[cell],'value':deepcopy(value),'original':original})
                            handled=True
                if re.search(r'分子范围未知',plain):
                    result['dimensions'].append({'id':roles['分子'],'ranges':ratio_cells,'value':'未知','original':original})
                    handled=True
                if handled:continue
        assigned=re.search(r'(?:资产|对象)范围(?:分别|依次)为[：:]?\s*([^。]+)',plain)
        if assigned and explicit:
            choices=set.intersection(*({r['id'] for r in standard['metric'][key]['dimension_refs']
                if standard['dimension'][r['id']].get('value_type')=='selection'} for key in explicit))
            choices-={names[m[1]] for m in marker.finditer(line)}
            if len(choices)==1:
                key=next(iter(choices));located=qualify(coordinate.findall(plain[:assigned.start()]))
                from .A语义阅读 import addresses
                for rows,value in re.findall(r'行([\d、]+)选([^；;]+)',assigned[1]):
                    selected=[area.rsplit('!',1)[0]+'!'+cell for area in located if '!' in area
                              for cell in addresses(area.rsplit('!',1)[1]) if re.search(r'\d+$',cell)[0] in rows.split('、')]
                    if selected:result['dimensions'].append({'id':key,'ranges':selected,'value':value,'original':original})
        for match in re.finditer('('+coordinate.pattern+r')\s+([\u4e00-\u9fff]+)(?=[、，,:：;；。]|$)',plain):
            labeled_cells.append((qualify([match[1]])[0],match[2],original))
        literal_rules=[]
        for clause in re.split(r'[；;。]',plain):
            declared=re.match(r'^\s*(?:[-*]\s*)?(?:上述)?([\u4e00-\u9fff]+)格用\s*(.+)$',clause)
            if not declared:continue
            metric=metrics.match(declared[2])
            if metric and len(set(metrics.findall(declared[2])))==1:
                literal_rules.append((declared[1],metric[1]))
        resolved_labels={}
        for location,label,source in labeled_cells if literal_rules else ():
            choices={key for name,key in literal_rules if name==label}
            if len(choices)!=1:continue
            result['metrics'].append({'id':next(iter(choices)),'ranges':[location],
                'original':{'original':source,'binding_reference':original}})
            resolved_labels.setdefault(source,set()).add(location)
        for row in result['unresolved'][:] if resolved_labels else ():
            if row.get('reason')!='这段答复未明确可机械读取的字段归属' or not isinstance(row.get('original'),str):continue
            if re.search(r'未知|未明确|未确定|不明|无法确认|不能确认|待定',row['original']):continue
            resolved=resolved_labels.get(row['original'],set())
            if not resolved:continue
            remaining=[area for area in row['ranges'] if area not in resolved]
            if remaining:row['ranges']=remaining
            else:result['unresolved'].remove(row)
        if re.match(r'^[-*]\s',plain):
            item_groups=[]
            for clause in re.split(r'[；;。]',plain):
                before,sep,description=clause.partition('：')
                if not sep:before,sep,description=clause.partition(':')
                located=coordinate.findall(before)
                if sep and located and not coordinate.sub('',before).strip(' -、，,'):
                    item_groups.append((qualify(located),description,original))
        if '明确的事项按行名识别' in plain and len(explicit)==1:
            change_fields=[r['id'] for r in standard['metric'][explicit[0]]['dimension_refs'] if r['id'].endswith('_change_selection')]
            labels=re.search(r'行(\d+[^。]+)[。]',plain)
            if len(change_fields)==1 and labels:
                from .A语义阅读 import addresses
                located=qualify(coordinate.findall(plain))
                for row_number,label in re.findall(r'(?:^|、)(\d+)([^、]+)',labels[1]):
                    group=[(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell for area in located
                           for cell in addresses(area.rsplit('!',1)[-1]) if re.search(r'\d+$',cell)[0]==row_number]
                    if group:result['dimensions'].append({'id':change_fields[0],'ranges':group,'value':label,'original':original})
        shared_members=re.search(r'共同适用的选取维度均为\s*(\w+)，业务域为“([^”]+)”，模式为\s*members',plain)
        if shared_members and standard['dimension'].get(shared_members[1],{}).get('value_type')=='selection':
            selection_defaults[shared_members[1]]={'domain':shared_members[2],'mode':'members'}
            continue
        member_heading=re.search(r'以下各格的\s*(\w+)\.members\s*取值为|以下各格除\s*(\w+)\s*外',plain)
        if member_heading:
            key=next(k for k in member_heading.groups() if k)
            if key in selection_defaults:
                member_field=key;member_shared={};member_enums=[]
                for field,value in re.findall(r'(\w+)\s*均为“([^”]+)”',plain):
                    if field in standard['dimension']:member_shared[field]=value
                member_enums=[field for field in standard['dimension'] if field not in member_shared and
                    re.search(r'(?<![\w.])'+re.escape(field)+r'(?![\w.])',plain) and standard['dimension'][field].get('allowed_values')]
                continue
        member_entry=re.match(r'^\s*-\s*((?:'+coordinate.pattern+r'[\s、，,]*)+)[：:]\s*(.+)$',plain)
        if member_field and member_entry:
            located=qualify(coordinate.findall(member_entry[1]));value=member_entry[2].rstrip('。')
            if metric_group is not None:metric_group['ranges']=list(dict.fromkeys(metric_group['ranges']+located))
            parts=value.split('；',1);enum=[key for key in member_enums if parts[0] in standard['dimension'][key]['allowed_values']]
            extra=dict(member_shared)
            if len(parts)==2 and len(enum)==1:extra[enum[0]]=parts[0];value=parts[1]
            result['dimensions'].append({'id':member_field,'ranges':located,
                'value':{**selection_defaults[member_field],'members':[value]},'original':original})
            result['dimensions'].extend({'id':key,'ranges':located,'value':value,'original':original} for key,value in extra.items())
            continue
        paired_rows=re.findall(r'(\d+)/(\d+)行(?:对象为|为|是)([^；;。]+)',plain)
        row_span=re.search(r'第([\d–—\-、，,]+)行',plain)
        column_ids=[(cols,key) for cols,key in re.findall(r'([A-Z]+(?:[、，,/][A-Z]+)*)列为\s*(metric\.[\w.]+)',plain) if key in standard['metric']]
        if row_span and column_ids and not nested_scope:
            paragraph_scope=scope
            row_numbers={row for start,end in re.findall(r'(\d+)[–—-](\d+)',row_span[1]) for row in range(int(start),int(end)+1)}
            columns=list(dict.fromkeys(col for cols,_ in column_ids for col in re.findall(r'[A-Z]+',cols)))
            # 有行范围的本段不能退成全组；未明示的列也不能按原表补出。
            scope=qualify([col+str(row) for row in sorted(row_numbers) for col in columns])
            period_pairs=[(column,match[2]) for match in re.finditer(
                r'([A-Z]+(?:[、，,/][A-Z]+)*)列(?:期间)?(?:为|是)(期初|期末|本期|上期)',plain)
                for column in re.findall(r'[A-Z]+',match[1])]
            for pair in re.finditer(r'([A-Z]+(?:[、，,/][A-Z]+)+)列[^，,；;。]*[，,]\s*期间分别为'
                    r'((?:期初|期末|本期|上期)(?:[、，,/](?:期初|期末|本期|上期))+)',plain):
                period_columns=re.findall(r'[A-Z]+',pair[1])
                period_labels=re.findall(r'期初|期末|本期|上期',pair[2])
                if len(period_columns)==len(period_labels):period_pairs.extend(zip(period_columns,period_labels))
            for column,label in period_pairs:
                result['dimensions'].append({'id':'template_period_role',
                    'ranges':qualify([column+str(row) for row in sorted(row_numbers)]),
                    'value':{'期初':'opening','期末':'closing','本期':'current','上期':'prior'}[label],
                    'original':original})
        if paired_rows and row_span and column_ids:
            row_numbers={row for start,end in re.findall(r'(\d+)[–—-](\d+)',row_span[1]) for row in range(int(start),int(end)+1)}
            choices=set.intersection(*({r['id'] for r in standard['metric'][key]['dimension_refs']
                if standard['dimension'][r['id']].get('value_type')=='selection'} for _,key in column_ids))
            if len(choices)==1:
                columns=list(dict.fromkeys(col for cols,_ in column_ids for col in re.findall(r'[A-Z]+',cols)))
                for left,right,value in paired_rows:
                    if {int(left),int(right)}<=row_numbers:
                        result['dimensions'].append({'id':next(iter(choices)),'ranges':qualify([col+row for row in (left,right) for col in columns]),
                                                     'value':value,'original':original})
        exclusion=re.match(r'^\s*-\s*([^：:]+)[：:]\s*(.+[；;]按原文注释，表中该合计扣除[^。]+)[。]?$',plain)
        if common_selection and exclusion:
            located=qualify(coordinate.findall(exclusion[1]))
            if located and not coordinate.sub('',exclusion[1]).strip(' 、，,'):
                result['dimensions'].append({'id':common_selection,'ranges':located,'value':exclusion[2],'original':original})
                continue
        common_entry=re.match(r'^\s*-\s*((?:'+coordinate.pattern+r'[\s、，,；;]*)+)(?:“[^”]*”)?[：:]\s*([^。]+)',plain)
        if common_entry and marker.search(common_entry[2]):common_entry=None
        if common_selection and common_entry and re.match(r'(?:\{|mode\s*=|(?:members|predicate)\s*(?:[:：，,；;]|—+))',common_entry[2]):
            line=re.sub(r'([：:])(?=\s*`?(?:\{|mode\s*=|(?:members|predicate)\b))',lambda m:m[1]+common_selection+'=',line,count=1)
            plain=line.replace('`','');common_entry=None
        if common_selection and common_entry and not coordinate.sub('',common_entry[1]).strip(' 、，,；;'):
            from .A语义阅读 import addresses
            located=qualify(coordinate.findall(common_entry[1]));description=common_entry[2]
            row_cells={}
            for area in located:
                for cell in addresses(area.rsplit('!',1)[-1]):
                    full=(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell
                    row_cells.setdefault(int(re.search(r'\d+$',cell)[0]),[]).append(full)
            ordered=re.search(r'依次为\s*\*\*([^*]+)\*\*',original) or re.search(r'依次为\s*(.+)$',description)
            if ordered:
                separator=r'[；;]' if re.search(r'[；;]',ordered[1]) else r'[、]'
                items=[value.strip() for value in re.split(separator,ordered[1])]
                if len(items)==len(row_cells):
                    for cells,value in zip(row_cells.values(),items):
                        result['dimensions'].append({'id':common_selection,'ranges':cells,'value':value,'original':original})
                    continue
            if re.search(re.escape(common_selection)+r'\s*(?:为)?未知',description):
                result['dimensions'].append({'id':common_selection,'ranges':located,'value':'未知','original':original})
                continue
            sequence=re.fullmatch(r'分别(?:是|为)([^“”]*)(“[^”]+”(?:“[^”]+”)*)',description)
            if sequence:
                items=[sequence[1]+value for value in re.findall(r'“([^”]+)”',sequence[2])]
                if len(items)==len(row_cells):
                    for (row_number,cells),value in zip(row_cells.items(),items):
                        result['dimensions'].append({'id':common_selection,'ranges':cells,'value':value,'original':original})
                        row_objects.setdefault(common_selection,{})[row_number]=value
                    ordered_values[common_selection]=items
                continue
            if description.startswith('以上') and '明细' in description:
                for row_number,cells in row_cells.items():
                    value=row_objects.get(common_selection,{}).get(row_number)
                    if value is not None:
                        result['dimensions'].append({'id':common_selection,'ranges':cells,'value':value,'original':original})
                continue
            if (located and len(re.findall(r'“[^”]+”',description))<=1 and
                not description.rstrip().endswith(('：',':')) and
                not re.search(r'分别|未知|未说明|未明确|无法确认',description) and not coordinate.search(description)):
                if re.search(r'[：:]\s*`(all|unknown)`(?=[，,。]|$)',line):
                    description=re.search(r'[：:]\s*`(all|unknown)`',line)[1]
                description=re.split(r'[；;]\s*两格财务数据',description,maxsplit=1)[0]
                result['dimensions'].append({'id':common_selection,'ranges':located,'value':description,'original':original})
                if not explicit:continue
        row_declaration=re.search(r'各行\s+(\w+)\s*[：:]',plain)
        if row_declaration and row_declaration[1] in standard['dimension']:row_field=row_declaration[1]
        sequence=re.search(r'对应[^。；;]+依次为[：:]\s*([^。]+)',plain)
        if sequence:
            from .A语义阅读 import addresses
            earlier=[area for statement in section_lines for area in coordinate.findall(statement)]
            row_numbers=sorted({int(re.search(r'\d+$',cell)[0]) for area in earlier for cell in addresses(area.rsplit('!',1)[-1])})
            values=[v.strip() for v in re.split(r'[、，,]',sequence[1])]
            if len(row_numbers)==len(values):
                for key,wording in re.findall(r'(\w+)\s*为“([^”]*该行所列区间[^”]*)”',plain):
                    if key in standard['dimension']:
                        row_texts[key]={row:wording.replace('该行所列区间',value) for row,value in zip(row_numbers,values)
                                        if not re.search(r'空白|合计',value)}
        explicit_selection=False
        selection_metric=explicit[0] if len(explicit)==1 else metric_group['id'] if metric_group else None
        selection_refs=[r['id'] for r in standard['metric'].get(selection_metric,{}).get('dimension_refs',[])
                        if standard['dimension'][r['id']]['value_type']=='selection']
        if len(selection_refs)==1:
            for sentence in plain.split('。'):
                for clause in re.split(r'[；;]',sentence):
                    restated=re.search(r'范围仍为([^。]+)',clause)
                    stated_cells=qualify(coordinate.findall(clause))
                    if restated and stated_cells:
                        result['dimensions'].append({'id':selection_refs[0],'ranges':stated_cells,'value':restated[1],'original':original})
                quoted=re.search(r'(?:[：:]\s*分别为|是|为)(“[^”]+”(?:“[^”]+”)*)',sentence)
                if not quoted:continue
                prefix=sentence[:quoted.start()].lstrip(' -')
                located=list(coordinate.finditer(prefix))
                if not located or coordinate.sub('',prefix).strip(' 、，,/'):continue
                values=re.findall(r'“([^”]+)”',quoted[1]);groups=[]
                if len(values)==1:groups=[qualify([m[0] for m in located])]
                elif len(located)==2*len(values) and all(prefix[located[i].end():located[i+1].start()]=='/' for i in range(0,len(located),2)):
                    groups=[qualify([located[i][0],located[i+1][0]]) for i in range(0,len(located),2)]
                if groups:
                    explicit_selection=True
                    result['dimensions'].extend({'id':selection_refs[0],'ranges':group,'value':value,'original':original}
                                                for group,value in zip(groups,values))
        duration=re.match(r'变动期间为\s*('+date_pattern+r')',plain)
        if duration and '至' in duration[1]:
            section_duration=re.sub(r'\s+','',duration[1])
            if not coordinate.search(plain):continue
        paired_names=re.search(r'[“"]?(期末|期初|本期|上期)[、，,](期末|期初|本期|上期)[”"]?分别(?:定义为|为)\s*('+date_pattern+r')[、，,]\s*('+date_pattern+r')',plain)
        if paired_names:
            named_periods.update(zip(paired_names.group(1,2),paired_names.group(3,4)))
        row_roles.extend((int(m[2]),int(m[3]),m[1]) for m in re.finditer(
            r'(期末|期初|本期|上期)分类为行(\d+)[—\-至](\d+)',plain))
        row_roles.extend((int(m[1]),int(m[1]),m[2]) for m in re.finditer(
            r'(?:行|[、，])\s*(\d+)(期初|期末)余额',plain))
        for m in re.finditer(r'(期末|期初|本期|上期)列([A-Z]+(?:[、，,][A-Z]+)*)(?=分别|为|[；;。])',plain):
            section_roles.update({col:m[1] for col in re.findall(r'[A-Z]+',m[2])})
        # 只消费本节已指明的字段及原词，不从数值、指标含义或原表推断。
        sentence_scope=[];ratio_positions={}
        for sentence in plain.split('。'):
            located=qualify(coordinate.findall(sentence))
            if located:sentence_scope=located
            same=re.search(r'同为\s*(?:`?((?:all|unknown))`?|“([^”]+)”)',sentence)
            same_fields={names[m[1]] for m in marker.finditer(line)
                         if standard['dimension'].get(names[m[1]],{}).get('value_type')=='selection'}
            if located and same and len(same_fields)==1:
                result['dimensions'].append({'id':next(iter(same_fields)),'ranges':located,
                    'value':next(value for value in same.groups() if value),'original':original})
            ratio_value=re.search(r'的总计分子和分母都对应(.+)$',sentence)
            if ratio_value and located and set(ratio_names)=={'分子','分母'} and any(
                set(ratio_names.values())<={r['id'] for r in standard['metric'][key].get('dimension_refs',[])}
                for key in explicit):
                result['dimensions'].extend({'id':key,'ranges':located,'value':ratio_value[1],'original':original}
                                            for key in ratio_names.values())
                from .A语义阅读 import addresses
                ratio_positions.update({key:{(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell
                                            for area in located for cell in addresses(area.rsplit('!',1)[-1])}
                                        for key in ratio_names.values()})
            bounded_values=re.findall(r'([A-Z]+(?:[、，,][A-Z]+)*)列(?:对应|为)[“"]([^”"]+)[”"]',sentence)
            enum_keys=[key for key in declared_fields|set(refs) if bounded_values and
                       all(value in standard['dimension'][key].get('allowed_values',[]) for _,value in bounded_values)]
            if len(enum_keys)==1:
                from .A语义阅读 import addresses
                for columns,value in bounded_values:
                    located=[(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell
                             for area in sentence_scope for cell in addresses(area.rsplit('!',1)[-1])
                             if re.match(r'[A-Z]+',cell)[0] in re.findall(r'[A-Z]+',columns)]
                    if located:result['dimensions'].append({'id':enum_keys[0],'ranges':located,'value':value,'original':original})
            change_refs=[key for key in refs if key.endswith('_change_selection')]
            if len(change_refs)==1:
                for clause in re.split(r'[；;]',sentence):
                    columns=re.match(r'\s*([A-Z]+(?:[、，,][A-Z]+)*)列',clause)
                    stated=re.search(r'(?:事项|'+re.escape(change_refs[0])+r')(分别)?为((?:“[^”]+”)+)',clause)
                    if not columns or not stated:continue
                    cols=re.findall(r'[A-Z]+',columns[1]);values=re.findall(r'“([^”]+)”',stated[2])
                    if not stated[1] and len(values)==1:values*=len(cols)
                    if len(cols)!=len(values):continue
                    from .A语义阅读 import addresses
                    for col,value in zip(cols,values):
                        located=[(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell
                                 for area in sentence_scope for cell in addresses(area.rsplit('!',1)[-1])
                                 if re.match(r'[A-Z]+',cell)[0]==col]
                        if located:result['dimensions'].append({'id':change_refs[0],'ranges':located,'value':value,'original':original})
        for declaration in re.finditer(r'([A-Z]+)\s*列\s*(?:period\s*[:：=＝]|(?:(?:取值)?期间)?为)\s*('+date_pattern+r')',plain):
            column_periods[declaration[1]]=re.sub(r'\s+','',declaration[2])
            following=plain[declaration.end():].split('。',1)[0]
            role=re.match(r'(?:[（(]表头|[，,]\s*(?:模板期间)?角色为)“(本年|上年)”',following)
            if role:column_roles[declaration[1]]={'本年':'本期','上年':'上期'}[role[1]]
        for declaration in re.finditer(r'列\s*([A-Z]+(?:[、，,]\s*[A-Z]+)*)\s*(?:均)?为\s*('+date_pattern+r')[、，,]\s*(期初|期末|本期|上期)',plain):
            for column in re.findall(r'[A-Z]+',declaration[1]):
                column_periods[column]=declaration[2];column_roles[column]=declaration[3]
        for declaration in re.finditer(r'([A-Z]+)\s*=\s*('+date_pattern+r')',plain):
            column_periods[declaration[1]]=re.sub(r'\s+','',declaration[2])
        for column,role in re.findall(r'列([A-Z]+)为(本年|上年)',plain):
            column_roles[column]={'本年':'本期','上年':'上期'}[role]
        if re.search(r'上述[^。]*依列相同',plain):
            locations=list(dict.fromkeys(area.rsplit('!',1)[0]+'!'+cell
                for item in result['metrics'][section_metric_start:] for area in qualify(item['ranges']) if area!='*'
                for cell in addresses(area.rsplit('!',1)[-1])))
            for column,date in re.findall(r'([A-Z]+)列为\s*('+date_pattern+r')',plain):
                selected=[loc for loc in locations if re.match(r'[A-Z]+',loc.rsplit('!',1)[-1])[0]==column]
                if selected:table_period_rows.append({'id':'period','ranges':selected,'value':date,'original':original})
        for role,date in re.findall(r'(本年|上年)期间为('+date_pattern+r')',plain):
            named_periods[{'本年':'本期','上年':'上期'}[role]]=date
        for label,boundary in re.findall(r'(?:的|[，,])([\u4e00-\u9fff]+)按相应期间(起日|末日)',plain):
            period_boundaries[label]=0 if boundary=='起日' else -1
        paired_roles=re.search(r'([A-Z]+)[、，,]([A-Z]+)(?:列)?分别(?:为|对应)(本期|上期|期末|期初)[与和、，,](本期|上期|期末|期初)',plain)
        if paired_roles and not coordinate.search(plain):
            column_roles.update({paired_roles[1]:paired_roles[3],paired_roles[2]:paired_roles[4]})
        scoped_columns=[m for clause in re.split(r'[；;。]',plain)
                        if (m:=re.fullmatch(r'\s*(.+?)的\s*([A-Z]+\s*列.+)',clause))]
        if not coordinate.search(plain) and not scoped_columns:
            column_roles.update(re.findall(r'([A-Z]+)\s*列\s*(?:均)?为\s*(本期|上期|期末|期初)(?!增加|减少)',plain))
            column_directions.update(re.findall(r'([A-Z]+)(?:列)?为本期(增加|减少)',plain))
        paired_columns=re.search(r'([A-Z]+)[/、]([A-Z]+)列分别(?:为|对应)本期[（(]('+date_pattern+r')[）)]和上期[（(]('+date_pattern+r')[）)]',plain)
        if paired_columns:
            column_periods.update({paired_columns[1]:re.sub(r'\s+','',paired_columns[3]),
                                   paired_columns[2]:re.sub(r'\s+','',paired_columns[4])})
            column_roles.update({paired_columns[1]:'本期',paired_columns[2]:'上期'})
        period_names='|'.join(re.escape(k) for k,v in names.items() if v=='period')
        for declaration in re.finditer(r'(本期|上期|期末|期初)(?:余额|变动|对应)?(?:的)?\s*(?:(?:'+period_names+r')\s*)?(?:[:：=＝]|为)\s*(?:时点|期间)?\s*('+date_pattern+r')',plain):
            named_periods[declaration[1]]=re.sub(r'\s+','',declaration[2])
        for declaration in re.finditer(r'(年初|期初|期末|本期|上期)(?:余额|变动期间|变动)?(?:对应|为)('+date_pattern+r')',plain):
            named_periods[declaration[1]]=declaration[2]
        paired_dates=re.search(r'(年初|期初|期末|本期|上期)[、，](年初|期初|期末|本期|上期)分别为('+date_pattern+r')[、，]('+date_pattern+r')',plain)
        if paired_dates:named_periods.update({paired_dates[1]:paired_dates[3],paired_dates[2]:paired_dates[4]})
        period_objects=bool(re.search(r'[“"][^”"]+[”"](?:列为|对应)\s*template_period_role\s*=',plain))
        for declaration in re.finditer(r'(本期|上期|期末|期初)[”"]?[^。；;，,]*?\bperiod\s*(?:=|为)\s*(\{[^{}]+\})',plain):
            pairs=re.findall(r'(\w+)\s*:\s*([^,{}]+)',declaration[2])
            if pairs and re.sub(r'\w+\s*:\s*[^,{}]+|[,{}\s]','',declaration[2])=='':
                named_periods[declaration[1]]={key:value.strip().strip('\"\'') for key,value in pairs}
                period_objects=True
        for declaration in scoped_columns:
            stated=re.sub(r'\s+','',declaration[1]).casefold()
            matching=[name for name in sheet_names if stated in
                      {re.sub(r'\s+','',name).casefold()+suffix for suffix in ('','表','工作表')}]
            if len(matching)!=1:continue
            for column,label in re.findall(r'([A-Z]+)\s*列\s*为\s*(本期|上期|期末|期初)',declaration[2]):
                if label in named_periods:sheet_columns[(matching[0],column)]=(label,deepcopy(named_periods[label]))
        # 日期只连接同一坐标分句明示的角色，不按指标或原表补期间。
        caption=re.match(r'^[-*]?\s*((?:'+coordinate.pattern+r'[、，,\s]*)+)[（(][^（）()]*?(期末|期初)余额[）)]',plain)
        if caption:
            result['dimensions'].append({'id':'template_period_role','ranges':qualify(coordinate.findall(caption[1])),
                'value':{'期末':'closing','期初':'opening'}[caption[2]],'original':original})
        for clause in re.split(r'[；;。]|[，,](?=[A-Z]+\d)',plain):
            clause=clause.lstrip(' -')
            bare_role=re.fullmatch(r'('+coordinate.pattern+r')\s*[:：]\s*(期初余额|期末余额|本期增加|本期减少|本期|上期)',clause)
            if bare_role:
                result['dimensions'].append({'id':'template_period_role','ranges':qualify([bare_role[1]]),
                    'value':{'期初余额':'opening','期末余额':'closing','本期增加':'current','本期减少':'current',
                             '本期':'current','上期':'prior'}[bare_role[2]],'original':original})
            if not coordinate.match(clause) or re.search(r'未找到|未确认|未知|冲突',clause) or re.search(r'[A-Z]+列',clause):continue
            located=list(coordinate.finditer(clause))
            paired_labels=re.search(r'(年初|期初|期末|本期|上期)[、，](年初|期初|期末|本期|上期)',clause)
            if paired_labels and len(located)==2 and all(label in named_periods for label in paired_labels.groups()):
                for item,label in zip(located,paired_labels.groups()):
                    result['dimensions'].extend({'id':key,'ranges':qualify([item[0]]),'value':value,'original':original}
                        for key,value in [('period',named_periods[label]),('template_period_role',
                            {'年初':'opening','期初':'opening','期末':'closing','本期':'current','上期':'prior'}[label])])
                continue
            labels=set(re.findall(r'本期|上期|期末|期初',clause))
            for i,item in enumerate(located):
                local=labels if len(labels)==1 else set(re.findall(r'本期|上期|期末|期初',
                    clause[item.end():located[i+1].start() if i+1<len(located) else len(clause)]))
                if len(local)!=1:continue
                label=next(iter(local))
                if label not in named_periods:continue
                result['dimensions'].extend({'id':key,'ranges':qualify([item[0]]),'value':value,'original':original}
                    for key,value in [('period',named_periods[label]),
                                      ('template_period_role',{'本期':'current','上期':'prior','期末':'closing','期初':'opening'}[label])])
        period_header=re.fullmatch(r'(期末|期初)列[（(]([A-Z、,，]+)[；;]期间(\d{4}-\d{2}-\d{2})[）)]',line)
        if period_header:
            ranges=[];columns=set(re.findall(r'[A-Z]+',period_header[2]))
            for following in lines[number+1:]:
                following=following.strip().replace('**','')
                if re.match(r'(期末|期初)列[（(]',following):break
                for area in coordinate.findall(following):
                    cols=re.fullmatch(r'(?:[^!]+!)?([A-Z]+)\d+(?::([A-Z]+)\d+)?',area)
                    if cols and cols[1] in columns and (not cols[2] or cols[2] in columns):
                        ranges.append(area)
            if ranges:
                result['dimensions'].extend({'id':key,'ranges':qualify(list(dict.fromkeys(ranges))),'value':value,'original':original}
                    for key,value in [('period',period_header[3]),('template_period_role',{'期末':'closing','期初':'opening'}[period_header[1]])])
            continue
        # 编号与行标签在标题中，财务坐标在冒号后的正文中；二者不能混用。
        heading=re.match(r'^[-*]\s*[^：:]*[｜|][^：:]*[：:]\s*(.*)',line)
        if heading and coordinate.search(line[:heading.start(1)]):heading=None
        body=heading[1] if heading else line
        body_cells=coordinate.findall(body) if heading else []
        if heading:
            line=body
            if body_cells:scope=qualify(body_cells)
            for item in re.finditer(r'(\$?[A-Z]+\$?\d+)\s*[（(]?(期末|期初)(?:余额)?',line):
                result['dimensions'].append({'id':'template_period_role','ranges':qualify([item[1]]),
                    'value':{'期末':'closing','期初':'opening'}[item[2]],'original':original})
        dated=[]
        if 'period' in standard['dimension']:
            for clause in re.split(r'[；;。]',plain):
                balance_date=re.match(r'('+coordinate.pattern+r')\s*[:：]\s*[^；;。()（）]*余额\s*[（(]('+date_pattern+r')[）)]',clause.lstrip(' -'))
                if balance_date:
                    dated.append({'id':'period','ranges':qualify([balance_date[1]]),'value':balance_date[2],'original':original})
            for item in re.finditer(r'(\$?[A-Z]+\$?\d+)\s*(?:期末|期初)(?:余额)?(?:为|是|[：:=＝])\s*(\d{4}年\d{1,2}月\d{1,2}日|\d{4}-\d{2}-\d{2})',line):
                dated.append({'id':'period','ranges':qualify([item[1]]),'value':item[2],'original':original})
            result['dimensions'].extend(dated)
        for declared in re.finditer('('+coordinate.pattern+r')\s*[（(]\s*('+date_pattern+r')\s*[）)]',plain):
            table_period_rows.append({'id':'period','ranges':qualify([declared[1]]),'value':declared[2],'original':original})
        for declared in re.finditer('('+coordinate.pattern+r')\s*[:：]\s*(期末|期初)[，,]\s*实际时点\s*('+date_pattern+r')',line.replace(chr(96),'')):
            table_period_rows.extend({'id':key,'ranges':qualify([declared[1]]),'value':value,'original':original}
                for key,value in [('period',declared[3]),('template_period_role',{'期末':'closing','期初':'opening'}[declared[2]])])
        for key in standard['dimension']:
            unknown=re.search(r'`'+re.escape(key)+r'`[^；;。\n]*(?:保留(?:为)?未知)',line)
            if unknown:result['dimensions'].append({'id':key,'ranges':scope[:],'value':'未知','original':original})
        common=column_scope.fullmatch(line)
        if common:
            # 共同取值只展开到答复下文列出的同列坐标，不扩成整表或整列。
            scope=[]
            for following in lines[number+1:]:
                following=following.strip().replace('**','')
                if column_scope.fullmatch(following) or following.startswith('#'):break
                cell=coordinate.match(following.lstrip(' -'))
                if cell and re.fullmatch(r'(?:[^!]+!)?\$?'+common[1]+r'\$?\d+(?::\$?'+common[1]+r'\$?\d+)?',cell[0]):
                    if cell[0] not in scope:scope.append(cell[0])
            continue
        if line.startswith('|'):
            for column,role in column_roles.items():
                if column not in column_periods and role in named_periods:column_periods[column]=named_periods[role]
            cells=[v.strip() for v in line.strip('|').split('|')]
            if all(re.fullmatch(r'[-: ]+',v) for v in cells):continue
            labels=[{'金标准维度名':'字段','维度名':'字段'}.get(v.strip('`'),v.strip('`')) for v in cells]
            if number+1<len(lines) and re.fullmatch(r'\s*\|[\s|:\-]+\|\s*',lines[number+1]):
                headers=['metric_id' if re.match(r'metric(?:[（(]|$)',v) or ('槽位' in v and '维度' not in v) or '标准编号' in v else v for v in labels]
                headers=[{'实际期间与角色':'period','期间与角色':'period'}.get(v,v) for v in headers]
                headers=[re.sub(r'\s*(?:实际)?取值$','',v.replace('`','')) if re.sub(r'\s*(?:实际)?取值$','',v.replace('`','')) in standard['dimension'] else v for v in headers]
                headers=[re.sub(r'^(本期|上期|期末|期初)\s*[:：]\s*([A-Z]+)(?=[^A-Za-z]|$)(.*)$',r'\2:\1\3',v) for v in headers]
                for key,(others,_) in shared_field_aliases.items():
                    pattern=r'\s*[、与和]\s*'.join(re.escape(k) for k in [key,*others])+r'\s*(?:的)?(?:范围和值|范围与值|实际取值|取值)'
                    headers=[key if re.fullmatch(pattern,v.replace('`','')) else v for v in headers]
                row_metric_columns={};previous_table_cells=None;column_table_scope=None;continue
            if 'metric_id' in labels:
                headers=labels;row_metric_columns={};continue
            if headers and len(cells)==len(headers):
                if any(label in headers[0] for label in ('财务格','数据格','金额格','坐标','位置')):
                    declared=cells[0].split('（',1)[0].split('(',1)[0]
                    if coordinate.search(declared):table_scope=coordinate.findall(declared)
                    column=re.search(r'(?<![A-Z0-9])([A-Z]+)列',declared.replace('`',''))
                    if column and (table_scope or '同表' in declared):
                        if table_scope:column_table_scope=qualify(table_scope)
                        if column_table_scope:
                            table_scope=[area.rsplit('!',1)[0]+'!'+cell for area in column_table_scope
                                for cell in addresses(area.rsplit('!',1)[-1]) if re.match(r'[A-Z]+',cell)[0]==column[1]]
                if previous_table_cells:
                    cells=[previous_table_cells[i] if value.strip('`')=='同上' else value for i,value in enumerate(cells)]
                previous_table_cells=cells[:]
                for header,value in zip(headers,cells):
                    dated=re.search(r'(current|prior|opening|closing)\s*[；;]\s*('+date_pattern+r')',header.replace('`',''))
                    locations=qualify(coordinate.findall(value))
                    if dated and locations:
                        result['dimensions'].extend({'id':key,'ranges':locations,'value':actual,'original':original}
                            for key,actual in [('template_period_role',dated[1]),('period',dated[2])])
                for cell_index,value in enumerate(cells[1:],1):
                    inherited=re.match(r'同\s*`?([A-Z]+\d+)`?\s*[,，]但\s*(.+)',value)
                    if inherited:
                        fields={names[m[1]] for m in marker.finditer(inherited[2])}
                        field_references.append((qualify(coordinate.findall(cells[0])),qualify([inherited[1]])[0],fields,original))
                        cells[cell_index]=inherited[2]
                if headers[0].startswith('行坐标'):table_scope=coordinate.findall(cells[0])
                row_member=re.search(r'各行的\s*`?(\w+)`?\s*分别为对应行名的\s*`?members',line)
                if table_scope and row_member and row_member[1] in standard['dimension']:
                    for entry in re.split(r'[；;]',cells[0].replace('`','')):
                        stated=re.fullmatch(r'\s*('+coordinate.pattern+r')\s+(.+?)\s*',entry)
                        if stated:result['dimensions'].append({'id':row_member[1],'ranges':qualify([stated[1]]),
                            'value':{'mode':'members','members':[stated[2]]},'original':original})
            if heading_metric and headers and len(cells)==len(headers):
                sources=[(header,qualify(coordinate.findall(value))) for header,value in zip(headers,cells)
                         if header.startswith('B') and '来源' in header]
                positions=[location for _,locations in sources for location in locations]
                if positions:
                    result['metrics'].append({'id':heading_metric,'ranges':positions,'original':original})
                    keys=[r['id'] for r in standard['metric'][heading_metric]['dimension_refs']
                          if standard['dimension'][r['id']].get('value_type')=='selection']
                    if len(keys)==1 and '范围' in headers:
                        value=cells[headers.index('范围')].replace('`','')
                        if value.startswith('all，'):value='all'
                        result['dimensions'].append({'id':keys[0],'ranges':positions,'value':value,'original':original})
                    for header,locations in sources:
                        dates=[date for label,date in header_dates.items() if label in header]
                        if len(dates)==1:result['dimensions'].append({'id':'period','ranges':locations,'value':dates[0],'original':original})
                    continue
            numbered=re.match(r'^(\d+)(?=\s|[:：「“（(])',cells[0])
            if not numbered and re.fullmatch(r'\d+',cells[0]):numbered=re.match(r'(\d+)',cells[0])
            row_span=re.fullmatch(r'(\d+)\s*[-–—至]\s*(\d+)',cells[0].strip('`'))
            if row_span:numbered=re.match(r'(\d+)',cells[0])
            row_numbers=[str(n) for n in range(int(row_span[1]),int(row_span[2])+1)] if row_span else [numbered[1]] if numbered else []
            declared_data_cells=coordinate.findall(' '.join(value for header,value in zip(headers or [],cells)
                if re.fullmatch(r'[A-Z]+(?:[、，,][A-Z]+)*格',header))) if numbered else []
            if declared_data_cells:
                table_scope=declared_data_cells
                result.setdefault('data_ranges',[]).extend(qualify(declared_data_cells))
            if numbered and requested:
                table_scope=[loc for loc in requested if re.search(r'\d+$',loc)[0]==numbered[1]] or None
            if numbered and headers:
                split_columns=[match[1] for header,value in zip(headers,cells) if re.match(r'^列(?:$|及|名|[（(])',header)
                    and (match:=re.match(r'^`?([A-Z]+)`?(?=$|\s|列|[:：])',value))]
                if len(split_columns)==1:table_scope=qualify([split_columns[0]+n for n in row_numbers])
            if numbered and len(row_column_groups)==1:
                table_scope=[column+numbered[1] for column in next(iter(row_column_groups))]
            if numbered:
                for column in re.findall(r'(?:项目|名称)\s*[（(]([A-Z]+)列[）)]',' '.join(headers or [])):
                    result.setdefault('non_business',[]).append({'range':qualify([column+numbered[1]])[0],
                        'category':'label','reason':'答复明确的行项目名称列','original':original})
            if numbered and headers and len(cells)==len(headers):
                for header,value in zip(headers,cells):
                    owner=re.match(r'^([A-Z]+)列',header)
                    if owner:
                        refs=[ref for ref in coordinate.findall(value) if re.match(r'[A-Z]+',ref.rsplit('!',1)[-1])[0]==owner[1]]
                        for field,values in [('period',column_periods),('template_period_role',column_roles)]:
                            if refs and owner[1] in values:result['dimensions'].append(
                                {'id':field,'ranges':qualify(refs),'value':values[owner[1]],'original':original})
            if numbered and headers and len(cells)==len(headers) and 'metric_id' not in headers and not any(value.strip('`') in standard['metric'] for value in cells):
                assigned_row_field=False
                for header,value in zip(headers,cells):
                    keys=[key for key in standard['dimension'] if re.search(r'(?<![\w.])'+re.escape(key)+r'(?![\w.])',header)]
                    if len(keys)!=1:continue
                    key=keys[0]
                    selected=sorted({loc for metric in result['metrics'] if key in
                        {ref['id'] for ref in standard['metric'].get(metric['id'],{}).get('dimension_refs',[])}
                        for area in qualify(metric['ranges']) if area!='*' for loc in
                        [(area.rsplit('!',1)[0]+'!' if '!' in area else '')+c for c in addresses(area.rsplit('!',1)[-1])]
                        if re.search(r'\d+$',loc)[0]==numbered[1]})
                    declared=declared_data_cells or coordinate.findall(' '.join(v for h,v in zip(headers,cells) if h.startswith('适用')))
                    if declared:
                        selected=sorted({loc for area in qualify(declared) for loc in
                            [(area.rsplit('!',1)[0]+'!' if '!' in area else '')+c for c in addresses(area.rsplit('!',1)[-1])]
                            if re.search(r'\d+$',loc)[0]==numbered[1]})
                    elif not selected:
                        selected=[loc for loc,keys in requested.items() if key in keys and re.search(r'\d+$',loc)[0]==numbered[1]]
                    columns=[match[1] for title in headers if (match:=re.match(r'^([A-Z]+)列[（(]',title))]
                    if columns:selected=qualify([col+n for col in columns for n in row_numbers])
                    if not selected and re.fullmatch(r'(?:原表)?行(?:号)?',headers[0]) and not metrics.search(' '.join(cells)) and not row_metric_columns:
                        pending_row_fields.append((sheet,row_numbers,key,value.replace(chr(96),''),original))
                        assigned_row_field=True
                        continue
                    column=re.search(r'(?<![A-Z])([A-Z]+(?:[、，,]\s*[A-Z]+)*)\s*列',header)
                    if column and '原文及对应' not in header:selected=[loc for loc in selected if re.match(r'[A-Z]+',loc.rsplit('!',1)[-1])[0] in re.findall(r'[A-Z]+',column[1])]
                    if selected:
                        paired=re.fullmatch(r'('+date_pattern+r')\s*[/／]\s*(期末|期初|本期|上期)',value)
                        if key=='period' and 'template_period_role' in header and paired:
                            assigned_row_field=True
                            result['dimensions'].extend({'id':field,'ranges':selected,'value':literal,'original':original}
                                for field,literal in [('period',paired[1]),('template_period_role',paired[2])])
                            continue
                        if '范围说明' in headers:value+='；'+cells[headers.index('范围说明')]
                        assigned_row_field=True
                        result['dimensions'].append({'id':key,'ranges':selected,'value':value.replace('`',''),'original':original})
                        for loc in selected:
                            column=re.match(r'[A-Z]+',loc.rsplit('!',1)[-1])[0]
                            for field,values in [('period',column_periods),('template_period_role',column_roles)]:
                                if column in values:result['dimensions'].append({'id':field,'ranges':[loc],'value':values[column],'original':original})
                if assigned_row_field:
                    for title,literal in zip(headers,cells):
                        column=re.match(r'^([A-Z]+)列[（(](期末|期初|本期|上期)[）)]',title)
                        dates=list(dict.fromkeys(re.findall(date_pattern,literal.replace(chr(96),''))))
                        if column and len(dates)==1:
                            locations=qualify([column[1]+n for n in row_numbers])
                            result['dimensions'].extend({'id':key,'ranges':locations,'value':value,'original':original}
                                for key,value in [('period',dates[0]),('template_period_role',column[2])])
                    continue
            if numbered and headers and len(cells)==len(headers):
                declaration=cells[headers.index('metric_id')] if 'metric_id' in headers else ''
                explicit_columns=column_metrics(declaration)
                last_column_metric=None;declared_columns=[]
                for header,value in zip(headers,cells):
                    match=re.fullmatch(r'([A-Z]+)(?:列)?(?:[:：].*)?',header)
                    if not match:continue
                    declared_columns.append(match[1])
                    literal=metrics.match(value.removeprefix('各行均为').strip(' `'))
                    key=literal[1] if literal else None
                    if not key and value.startswith('同上'):key=row_metric_columns.get(match[1])
                    if not key and re.match(r'(?:各行)?同一(?:变动)?指标',value):key=last_column_metric
                    if key:explicit_columns[match[1]]=key;last_column_metric=key
                if explicit_columns:row_metric_columns=explicit_columns
                elif not re.search(r'同上[^。]*指标',declaration):row_metric_columns={}
                if row_metric_columns:
                    selected=qualify([col+n for col in row_metric_columns for n in row_numbers])
                    result['metrics'].extend({'id':key,'ranges':qualify([col+n for n in row_numbers]),'original':original}
                        for col,key in row_metric_columns.items())
                    for column in declared_columns or row_metric_columns:
                        if column in column_periods:result['dimensions'].append({'id':'period','ranges':qualify([column+n for n in row_numbers]),'value':column_periods[column],'original':original})
                        if column in column_roles:result['dimensions'].append({'id':'template_period_role','ranges':qualify([column+n for n in row_numbers]),'value':column_roles[column],'original':original})
                    declared_values=[]
                    for header,value in zip(headers,cells):
                        keys=[key for key in standard['dimension'] if re.search(r'(?<![\w.])'+re.escape(key)+r'(?![\w.])',header)]
                        if len(keys)==1:
                            leading=re.match(r'^`(all|unknown)`(?=$|[：:，,])',value)
                            if leading:value=leading[1]
                            declared_values.append((keys[0],value,None))
                        else:
                            owner=re.fullmatch(r'([A-Z]+)(?:列)?(?:[:：].*)?',header)
                            metric=row_metric_columns.get(owner[1]) if owner else None
                            refs=[ref['id'] for ref in standard['metric'].get(metric,{}).get('dimension_refs',[])]
                            for label,suffix,values in [('方向','_direction','增加|减少'),('事项范围','_change_selection','all|unknown')]:
                                keys=[key for key in refs if key.endswith(suffix)]
                                if len(keys)==1:
                                    value=re.sub(r'(?<![\w])'+label+r'\s*`?('+values+r')`?(?=$|[，,；;。])',lambda m:keys[0]+'='+m[1],value)
                            matches=list(marker.finditer(value))
                            for i,match in enumerate(matches):
                                key=names[match[1]]
                                if key in standard['dimension']:
                                    literal=value[match.end():matches[i+1].start() if i+1<len(matches) else len(value)]
                                    owner=re.fullmatch(r'([A-Z]+)(?:列)?(?:[:：].*)?',header)
                                    declared_values.append((key,re.split(r'[；;]',literal)[0].strip(' `。，,'),owner[1] if owner else None))
                    for key,value,owner in declared_values:
                        applicable=qualify([col+n for col in declared_columns or row_metric_columns for n in row_numbers
                            if (owner is None or col==owner) and (col not in row_metric_columns or key in
                                {ref['id'] for ref in standard['metric'][row_metric_columns[col]].get('dimension_refs',[])})])
                        ordered=re.fullmatch(r'各行分别为\s*members\s*=\s*\[([^\]]+)\]',value.replace('`',''))
                        members=re.split(r'[／/]',ordered[1]) if ordered else []
                        if ordered and len(members)==len(row_numbers):
                            for n,member in zip(row_numbers,members):
                                locations=[loc for loc in applicable if re.search(r'\d+$',loc)[0]==n]
                                if locations:result['dimensions'].append({'id':key,'ranges':locations,'value':{'mode':'members','members':[member.strip()]},'original':original})
                            continue
                        if applicable:result['dimensions'].append({'id':key,'ranges':applicable,'value':value.replace('`',''),'original':original})
                    continue
            bounded_values=[(match[1],value,header) for header,value in zip(headers or [],cells)
                if (match:=re.fullmatch(r'([A-Z]+)列[:：].+',header))]
            if numbered and headers[0] in ('原表行','行号') and bounded_values:
                prior_selection=None
                for column,value,header in bounded_values:
                    locations=qualify([column+n for n in row_numbers]);literal=value.replace(chr(96),'')
                    result.setdefault('data_ranges',[]).extend(locations)
                    parsed={names[m[1]]:re.split(r'[；;]',literal[m.end():])[0].strip(' ，,。')
                            for m in marker.finditer(literal) if names[m[1]] in standard['dimension']}
                    selection=re.match(r'选择范围\s*[:：]\s*([^；;]+)',literal)
                    if common_selection and selection:parsed[common_selection]=selection[1]
                    same=re.match(r'同一([^；;]+)',literal)
                    labels={'选择范围','合计范围',re.split(r'[，,；;]',str(prior_selection))[0]}
                    if common_selection not in parsed and same and same[1] in labels and prior_selection is not None:
                        parsed[common_selection]=deepcopy(prior_selection)
                    if common_selection in parsed:prior_selection=parsed[common_selection]
                    dates=list(dict.fromkeys(re.findall(date_pattern,literal)))
                    if 'period' in standard['dimension'] and len(dates)==1:parsed['period']=dates[0]
                    role=re.search(r'(?:角色)?[“\"](?:\s*)(本期|上期|期末|期初)[”\"]',literal)
                    if not role:role=re.search(r'列[:：]\s*(本期|上期|期末|期初)',header)
                    if role:parsed['template_period_role']=role[1]
                    refs=[ref for declaration in re.findall(r'期间栏位标签位于([^。\n]+)',text.replace('**','').replace(chr(96),''))
                          for ref in qualify(coordinate.findall(declaration))
                          if re.match(r'[A-Z]+',ref.rsplit('!',1)[-1])[0]==column]
                    for key,value in parsed.items():
                        row={'id':key,'ranges':locations,'value':value,'original':original}
                        if key=='period' and refs:row.update(evidence=refs,label_references=refs)
                        (table_period_rows if key in ('period','template_period_role') else result['dimensions']).append(row)
                for header in headers:
                    label=re.fullmatch(r'([A-Z]+)列项目原文',header)
                    if label:
                        result.setdefault('non_business',[]).extend({'range':location,'category':'label',
                            'reason':'答复表头明示的项目原文列','original':original}
                            for location in qualify([label[1]+n for n in row_numbers]))
                continue
            if any(v in ('范围','位置') or v.endswith('坐标') for v in labels):headers=labels;continue
            if headers and len(cells)==len(headers):
                fields=list(zip(headers,cells))
                for header,value in fields:
                    role=re.fullmatch(r'(?:\d{4})?(本期|上期|期初|期末)来源',header)
                    refs=coordinate.findall(value)
                    if role and refs:
                        result['dimensions'].append({'id':'template_period_role','ranges':qualify(refs),
                            'value':{'本期':'current','上期':'prior','期初':'opening','期末':'closing'}[role[1]],'original':original})
                        if role[1] in named_periods:result['dimensions'].append({'id':'period','ranges':qualify(refs),
                            'value':named_periods[role[1]],'original':original})
                fields=[('坐标' if key.endswith('坐标') else key,value) for key,value in fields]
                for index,(key,value) in enumerate(fields):
                    declared=marker.match(value)
                    if key not in names and declared and len(list(marker.finditer(value)))==1 and names[declared[1]] in standard['dimension']:
                        fields[index]=(names[declared[1]],value[declared.end():])
                positions=[value for key,value in fields if re.match(r'^(?:[AB]\s*)?(?:财务格|数据格|财务位置|位置|坐标|范围)(?:$|及|[（(])',key)]
                if len(positions)==1 and table_scope is None:table_scope=coordinate.findall(positions[0])
                column_fields=[]
                for header,value in fields:
                    column=re.fullmatch(r'([A-Z]+)列',header)
                    dated=re.match(r'(本期|上期|期初|期末)[，,、]\s*('+date_pattern+r')',value.replace(chr(96),''))
                    if not (table_scope and column and dated):continue
                    locations=[area.rsplit('!',1)[0]+'!'+cell for area in qualify(table_scope)
                               for cell in addresses(area.rsplit('!',1)[-1]) if re.match(r'[A-Z]+',cell)[0]==column[1]]
                    table_period_rows.extend({'id':key,'ranges':locations,'value':actual,'original':original}
                        for key,actual in [('period',dated[2]),('template_period_role',dated[1])])
                    column_fields.append(header)
                fields=[(key,value) for key,value in fields if key not in column_fields]
                for header,value in fields:
                    role=re.match(r'^(current|prior|opening|closing|本期|上期|期初|期末)(?=$|[；;，,])',value.replace(chr(96),''))
                    if table_scope and 'template_period_role' in header and role:
                        refs=qualify(coordinate.findall(value));locations=qualify(table_scope)
                        label={'current':'本期','prior':'上期','opening':'期初','closing':'期末'}.get(role[1],role[1])
                        table_period_rows.append({'id':'template_period_role','ranges':locations,'value':role[1],'original':original})
                        if refs or label in common_periods:table_period_rows.append({'id':'period','ranges':locations,
                            **({'value':common_periods[label]} if label in common_periods else {}),
                            'label_references':refs,'evidence':refs,'original':original})
                if table_scope and requested:
                    locations={loc for area in qualify(table_scope) for loc in
                        [(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell for cell in addresses(area.rsplit('!',1)[-1])]}
                    selections={key for loc in locations for key in requested.get(loc,())
                        if standard['dimension'].get(key,{}).get('value_type')=='selection'}
                    if len(selections)==1:
                        fields=[(next(iter(selections)) if key not in names and
                            re.match(r'^(?:members|all|predicate|unknown)\s*[:：]',value.replace(chr(96),'')) else key,value) for key,value in fields]
                split_fields=[]
                for key,value in fields:
                    keys=[part.strip(' `') for part in re.split(r'[；;]',key)]
                    values=[part.strip(' `') for part in re.split(r'[；;]',value)]
                    if len(keys)>1 and len(keys)==len(values) and all(k in names for k in keys):
                        split_fields.extend(zip(keys,values))
                    else:split_fields.append((key,value))
                fields=split_fields
                combined_period=[]
                for key,value in fields:
                    pair=re.fullmatch(r'(本期|上期|期初|期末|current|prior|opening|closing)\s*[；;]\s*('+date_pattern+r')',value.replace('`',''))
                    reversed_pair=re.fullmatch(r'('+date_pattern+r')\s*[；;]\s*(本期|上期|期初|期末|current|prior|opening|closing)',value.replace('`',''))
                    if 'template_period_role' in key and re.search(r'(?<![\w.])period(?![\w.])',key) and pair:
                        combined_period.extend([('template_period_role',pair[1]),('period',pair[2])])
                    elif 'template_period_role' in key and re.search(r'(?<![\w.])period(?![\w.])',key) and reversed_pair:
                        combined_period.extend([('period',reversed_pair[1]),('template_period_role',reversed_pair[2])])
                    else:combined_period.append((key,value))
                fields=combined_period
                fields=[('坐标' if k=='坐标行' else k,v) for k,v in fields]
                if common_selection:
                    fields=[(common_selection if k in ('成员','成员或范围') else k,v) for k,v in fields]
                    fields=[(k,'未知' if re.search(r'未识别出[^。]*成员|不是[^。]*成员',v)
                             else v.split('；',1)[0]) if k==common_selection else (k,v) for k,v in fields]
                if common_selection:fields=[(common_selection if re.match(r'实际(?:取值|范围)(?=[（(]|$)',k) else k,v) for k,v in fields]
                fields=[(keys[0] if len(keys:=[key for key in standard['dimension'] if re.search(r'(?<![\w.])'+re.escape(key)+r'(?![\w.])',header)])==1 else header,value)
                        for header,value in fields]
                expanded=[]
                for key,value in fields:
                    located=re.match(r'^((?:'+coordinate.pattern+r'[\s、，,；;]*)+)(?:的分子)?[:：]\s*(?:均为)?(.+)$',value)
                    parts=[re.fullmatch(r'('+coordinate.pattern+r')\s*[:：]\s*(.+)',part.strip())
                        for part in re.split(r'[；;]',value)]
                    if (key in standard['dimension'] or key=='template_period_role') and len(parts)>1 and all(parts):
                        result['dimensions'].extend({'id':key,'ranges':qualify([part[1]]),'value':part[2],'original':original} for part in parts)
                        continue
                    if key in standard['dimension'] and located:
                        expanded.extend([('坐标',located[1]),(key,located[2])])
                    else:expanded.append((key,value))
                fields=expanded
                fields=[(key,value) for key,value in fields if not (key in standard['dimension']
                    and re.search(r'按所在行[^。]*见下方',value))]
                if '字段' in headers and '实际取值' in headers:
                    field=cells[headers.index('字段')].strip('`')
                    if field in names:
                        fields=[(field if k=='实际取值' else k,v) for k,v in fields if k!='字段']
                fields.sort(key=lambda item:not (item[0] in ('范围','位置') or item[0].endswith('坐标')))
                table_literals={key:value for key,value in fields
                    if key in standard['dimension'] and not coordinate.search(value)}
                for key,value in fields[:]:
                    if key not in standard['dimension'] or not value.startswith('依行：') or table_scope is None:continue
                    entries=re.findall(r'(\d+)\s+(.+?)(?=[；;]\s*\d+\s|$)',value[3:])
                    for row_number,literal in entries:
                        selected=[cell for area in qualify(table_scope) for cell in
                            [(area.rsplit('!',1)[0]+'!' if '!' in area else '')+c for c in addresses(area.rsplit('!',1)[-1])]
                            if re.search(r'\d+$',cell)[0]==row_number]
                        if selected:result['dimensions'].append({'id':key,'ranges':selected,'value':literal.strip(' `'),'original':original})
                    if entries:fields.remove((key,value));table_literals.pop(key,None)
                table_metric_text=dict(fields).get('metric_id')
                parts=[k+'：'+v for k,v in fields if v]
                line='；'.join(parts);offset=0
                for part in parts:
                    offset+=len(part);table_cell_ends.append(offset);offset+=1
                explicit=list(dict.fromkeys(metrics.findall(line)))
        handled_fields=set()
        row_list=re.search(r'行范围依次为([^；;。]+)',plain)
        detail_list=re.search(r'(?:各)?明细依次为([^；;。]+)',plain)
        fields={names[m[1]] for m in marker.finditer(line)
                if standard['dimension'].get(names[m[1]],{}).get('value_type')=='selection'}
        if row_list and detail_list and len(fields)==1 and re.search(r'合计[^；;。]*范围为[^；;。]*`all`',line):
            from .A语义阅读 import addresses
            key=next(iter(fields));handled_fields.add(key);groups={}
            for area in qualify(coordinate.findall(line)):
                for cell in addresses(area.rsplit('!',1)[-1]):
                    groups.setdefault(int(re.search(r'\d+$',cell)[0]),[]).append(
                        (area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell)
            labels=[v.strip() for v in row_list[1].split('、')]
            values=iter(v.strip() for v in detail_list[1].split('、'))
            if len(labels)==len(groups) and labels.count('合计')==1 and len(detail_list[1].split('、'))==len(groups)-1:
                for label,cells in zip(labels,groups.values()):
                    result['dimensions'].append({'id':key,'ranges':list(dict.fromkeys(cells)),
                        'value':'all' if label=='合计' else next(values),'original':original})
            else:result['unresolved'].append({'id':key,'ranges':qualify(coordinate.findall(line)),
                'reason':'合计及明细顺序未唯一对应坐标','original':original})
        found=[m for m in marker.finditer(line) if names[m[1]] not in handled_fields]
        found=[m for m in found if m[1]==names[m[1]] or not any(
            later[1]==names[m[1]] and later.start()>m.end()
            and not coordinate.search(line[m.end():later.start()])
            and not re.search(r'[；;。]',line[m.end():later.start()]) for later in found)]
        found=[m for m in found if not (names[m[1]]=='period' and any(names[later[1]]=='period' and later.start()>=m.end()
            and re.fullmatch(r'(?:[^；;。]*[，,]\s*实际|[ `]*)',line[m.end():later.start()])
            and not coordinate.search(line[m.end():later.start()]) for later in found))]
        head=line[:found[0].start()] if found else line
        declared_fields.update(names[m[1]] for m in found if names[m[1]] in standard['dimension'])
        if len(explicit)==1:
            for clause in re.split(r'[；;]',plain):
                if coordinate.match(clause.lstrip(' -')) and re.search(r'同一(?:比例|变动)?指标',clause):
                    result['metrics'].append({'id':explicit[0],'ranges':qualify(coordinate.findall(clause)),'original':original})
        if column_table_scope and re.search(r'(?:本区|本表)(?:各|全部)数据',line):table_scope=column_table_scope[:]
        spans=table_scope if table_scope is not None else coordinate.findall(head)
        labelled=re.match(r'^\s*[-*]?\s*`?('+coordinate.pattern+r')`?\s*[，,]\s*`?('
            +coordinate.pattern+r')`?\s*[“「"]([^”」"]+)[”」"]\s*[:：]',head)
        if labelled and table_scope is None:
            # “目标范围，标签位置‘名称’：编号”中的旁述引用不属于金额目标。
            spans=[labelled[1]]
            result.setdefault('non_business',[]).append({'range':qualify([labelled[2]])[0],
                'category':'label','reason':'答复在目标范围旁引用的名称标签','original':original})
        nested_metric=metric_group is not None and bool(re.match(r'^\s+[-*]\s',original))
        if spans and not found and not explicit and not nested_metric and not heading and '共同' not in head and not line.startswith(('|','#')) and not coordinate.match(line.lstrip(' -`').removeprefix('其中').lstrip(' `')) and not re.search(r'(?:分别)?(?:采用|选用|取自)\s*B\s*'+coordinate.pattern,head):
            continue
        if spans and not heading:scope=qualify(spans)
        for declared in re.finditer('('+coordinate.pattern+r')`?\s*[（(](期末|期初|本期|上期)(?:余额|发生额|原值)?[）)]',line):
            result['dimensions'].append({'id':'template_period_role','ranges':qualify([declared[1]]),
                                         'value':declared[2],'original':original})
        if not row_field and section_fields and '按行范围' in plain and metric_context:
            from .A语义阅读 import addresses
            located=[(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell
                     for area in scope if area!='*' for cell in addresses(area.rsplit('!',1)[-1])]
            ids=[metric_context.get(cell) for cell in located]
            candidates=[key for key in section_fields if ids and all(
                key in {ref['id'] for ref in standard['metric'].get(metric,{}).get('dimension_refs',[])} for metric in ids)]
            if len(candidates)==1:row_field=candidates[0];explicit_selection=True
        if row_field:
            numbered=re.match(r'^\s*-\s*(\d+(?:、\d+)*)[：:]\s*`?(\{[^{}]+\})`?',line)
            if numbered:
                from .A语义阅读 import addresses
                try:value=json.loads(numbered[2])
                except ValueError:value=None
                if value is not None:
                    rows=set(numbered[1].split('、'))
                    located=[(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell for area in scope if area!='*'
                             for cell in addresses(area.rsplit('!',1)[-1]) if re.search(r'\d+$',cell)[0] in rows]
                    if located:
                        result['dimensions'].append({'id':row_field,'ranges':located,'value':value,'original':original})
                        for row_number in rows:row_objects.setdefault(row_field,{})[int(row_number)]=value
            for item in re.finditer(r'第(\d+)行(?:合计)?[：:]\s*([^；;。]+)',plain):
                from .A语义阅读 import addresses
                located=[(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell
                         for area in scope if area!='*' for cell in addresses(area.rsplit('!',1)[-1])
                         if int(re.search(r'\d+$',cell)[0])==int(item[1])]
                if located:
                    value='未知' if item[2].endswith('范围未知') else item[2]
                    result['dimensions'].append({'id':row_field,'ranges':located,'value':value,'original':original})
                    row_objects.setdefault(row_field,{})[int(item[1])]=value
        metric_statement=next((line[m.end():found[i+1].start() if i+1<len(found) else len(line)]
            for i,m in enumerate(found) if names[m[1]]=='metric_id'),'').lstrip(' `')
        refused=bool(spans and (declined.search(line) or metric_statement and not metrics.match(metric_statement)
            and re.search(r'金标准(?:无|没有)|无(?:准确|相应|对应|匹配)[^。；;]*槽位|(?:标准槽位|槽位)(?:未提供|缺少|未找到)',metric_statement)))
        if refused:
            result['unresolved'].append({'id':'metric_id','value':'未知','ranges':scope[:],'reason':'原答复明确不采用此槽位','original':original})
            explicit=[]
        selected=[] if refused else metrics.findall(head)
        if spans and len(explicit)==1:selected=explicit[:]
        if found and spans and not selected and len(explicit)==1:selected=explicit[:]
        referenced_metric=refused or bool(re.match(r'^上述\s*`?'+metrics.pattern+r'`?\s*格的',source_line))
        if len(selected)==1:last_metric=selected[0]
        elif spans and last_metric and re.search(r'同(?:一)?(?:变动)?(?:槽位|指标)',line):
            referenced_metric=True
            selected=[last_metric]
            if not found:result['metrics'].append({'id':last_metric,'ranges':scope[:],'original':original})
            if last_selection and re.search(r'范围[^。]*(?:保留(?:为)?未知)',line):
                result['dimensions'].append({'id':last_selection,'ranges':scope[:],'value':'未知','original':original})
        if line.startswith('#') and not found and not selected:continue
        if spans:
            if (metric_group is not None and not selected and not referenced_metric
                    and (table_scope is not None or nested_metric or coordinate.match(line.lstrip(' -`')))
                    and not any(names[m[1]]=='metric_id' for m in found)):
                metric_group['ranges']=list(dict.fromkeys(metric_group['ranges']+qualify(coordinate.findall(line))))
            from .A语义阅读 import addresses
            local_dates={m[2]:re.sub(r'\s+','',m[1]) for m in
                         re.finditer(r'('+date_pattern+r')\s*[（(]([A-Z]+)[）)]',plain)}
            for stated in re.finditer(r'([A-Z]+)列(?:按共同说明)?为('+date_pattern+r'|\d{4}年度)',plain):
                value=stated[2]
                if value.endswith('年度'):
                    periods={re.sub(r'\s+','',m[0]) for m in re.finditer(date_pattern,text)
                             if '至' in m[0] and all(d.strip().startswith(value[:4]) for d in m[0].split('至'))}
                    if len(periods)!=1:continue
                    value=next(iter(periods))
                local_dates[stated[1]]=value
            tail_role=re.search(r'[。；;]\s*(期末|期初|本期|上期)\s*[。；;]\s*\|$',plain) if plain.startswith('|') and len(scope)==1 else None
            if tail_role and tail_role[1] in named_periods:
                result['dimensions'].extend({'id':key,'ranges':scope[:],'value':value,'original':original}
                    for key,value in [('period',named_periods[tail_role[1]]),
                        ('template_period_role',{'期末':'closing','期初':'opening','本期':'current','上期':'prior'}[tail_role[1]])])
            local_roles={column:m[2] for m in re.finditer(r'([A-Z]+(?:\s*[、,，]\s*[A-Z]+)*)(?:列)?\s*(?:是|为)?\s*[“"「]?(期末|期初|本期|上期)',plain)
                         for column in re.findall(r'[A-Z]+',m[1])}
            for m in re.finditer(r'(期末|期初|本期|上期)列为\s*([A-Z]+)(?:\s*[:：]\s*([A-Z]+))?',plain):
                local_roles.update({re.match(r'[A-Z]+',cell)[0]:m[1]
                    for cell in addresses(f'{m[2]}1:{m[3] or m[2]}1')})
            local_roles.update(re.findall(r'([A-Z]+)[“"](期末|期初|本期|上期)(?:余额|发生额|增加|减少)?[”"]',plain))
            if re.search(r'列角色及日期同上',plain):local_roles.update(last_column_roles)
            ordered_roles=re.search(r'列角色依次为([^，,；;。]+)',plain)
            if ordered_roles and len(scope)==1:
                columns=list(dict.fromkeys(re.match(r'[A-Z]+',cell)[0] for cell in addresses(scope[0].rsplit('!',1)[-1])))
                labels=ordered_roles[1].split('、')
                if len(columns)==len(labels) and all(re.fullmatch(r'(?:期初余额|期末余额|本期增加|本期减少)',label) for label in labels):
                    local_roles.update({col:label[:2] for col,label in zip(columns,labels)})
            stated_titles=re.search(r'这些格按原文列名分别是([^。；;]+)',plain)
            if stated_titles:
                titles=re.findall(r'[“"]([^”"]+)[”"]',stated_titles[1])
                roles={title[:2] for title in titles}
                if len(titles)>1 and len(roles)==1 and roles<={'本期','上期','期初','期末'}:
                    result['dimensions'].append({'id':'template_period_role','ranges':scope[:],
                        'value':{'本期':'current','上期':'prior','期初':'opening','期末':'closing'}[next(iter(roles))],'original':original})
            if local_roles:last_column_roles=local_roles.copy()
            paired_roles=re.search(r'([A-Z]+)[、，,]([A-Z]+)列分别(?:为|对应)(期末|期初|本期|上期)[、，,](期末|期初|本期|上期)',plain)
            if paired_roles:local_roles.update({paired_roles[1]:paired_roles[3],paired_roles[2]:paired_roles[4]})
            local_dates.update({column:named_periods[label] for column,label in local_roles.items() if label in named_periods})
            local_dates.update({column:named_periods[label] for column,label in section_roles.items() if label in named_periods})
            local_dates.update({column:named_periods[label] for column,label in column_roles.items() if label in named_periods and column not in local_dates})
            named_line=re.match(r'^[-*]\s*([^：:]+)[：:]',plain)
            if named_line and named_line[1] in period_boundaries:
                local_dates={column:date.split('至')[period_boundaries[named_line[1]]]
                             for column,date in local_dates.items()}
                # 答复明确的本年余额起末角色不能被本年列的发生额角色覆盖；上年列仍保留其比较年度。
                local_roles.update({column:'期初' if period_boundaries[named_line[1]]==0 else '期末'
                                    for column,label in column_roles.items() if label=='本期'})
            directions=column_directions|dict(re.findall(r'([A-Z]+)(?:列)?(?:为|[“"])本期(增加|减少)',plain))
            for area in scope:
                if area=='*':continue
                for cell in addresses(area.rsplit('!',1)[-1]):
                    location=(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell
                    value=directions.get(re.match(r'[A-Z]+',cell)[0])
                    metric=explicit[0] if len(explicit)==1 else (metric_context or {}).get(location)
                    keys=[r['id'] for r in standard['metric'].get(metric,{}).get('dimension_refs',[])
                          if value and r['id'].endswith('_direction') and value in standard['dimension'][r['id']].get('allowed_values',[])]
                    if len(keys)==1:result['dimensions'].append({'id':keys[0],'ranges':[location],'value':value,'original':original})
            for area in scope:
                for first,last,label in row_roles:
                    located=[(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell
                             for cell in addresses(area.rsplit('!',1)[-1])
                             if first<=int(re.search(r'\d+$',cell)[0])<=last]
                    if located and label in named_periods:
                        result['dimensions'].extend({'id':key,'ranges':located,'value':value,'original':original}
                            for key,value in [('period',named_periods[label]),
                                ('template_period_role',{'本期':'current','上期':'prior','期末':'closing','期初':'opening'}[label])])
            ordered_dates=re.search(r'(?:时点|期间)分别为\s*('+date_pattern+r')[、，]\s*('+date_pattern+r')',plain)
            if ordered_dates and len(scope)==2:
                local_dates.update({re.match(r'[A-Z]+',area.rsplit('!',1)[-1])[0]:re.sub(r'\s+','',date)
                                    for area,date in zip(scope,ordered_dates.groups())})
            stated_duration=re.search(r'期间按('+date_pattern+r')理解',plain)
            if stated_duration and '至' in stated_duration[1]:
                result['dimensions'].append({'id':'period','ranges':scope[:],'value':stated_duration[1],'original':original})
            bounded_values=[('period',column,date) for column,date in (column_periods|local_dates).items()]
            bounded_values.extend(('template_period_role',column,{'本期':'current','上期':'prior','期末':'closing','期初':'opening'}[label])
                                 for column,label in (column_roles|section_roles|local_roles).items())
            stated={key:{cell for row in result['dimensions'] if row['id']==key and row.get('original')==original
                         for area in row['ranges'] if area!='*'
                         for cell in [(area.rsplit('!',1)[0]+'!' if '!' in area else '')+a for a in addresses(area.rsplit('!',1)[-1])]}
                    for key in ('period','template_period_role')}
            for key,column,value in bounded_values:
                located=[(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell
                         for area in qualify(table_scope if table_scope is not None else coordinate.findall(line)) for cell in addresses(area.rsplit('!',1)[-1])
                         if re.match(r'[A-Z]+',cell)[0]==column]
                located=[cell for cell in located if cell not in stated[key]]
                if located:
                    result['dimensions'].append({'id':key,'ranges':located,'value':value,'original':original})
                    if column in column_label_refs:period_links.append((qualify([column_label_refs[column]]),located))
            for (source_sheet,column),(label,date) in sheet_columns.items():
                located=[source_sheet+'!'+cell for area in qualify(table_scope if table_scope is not None else coordinate.findall(line))
                         if area.rpartition('!')[0]==source_sheet for cell in addresses(area.rsplit('!',1)[-1])
                         if re.match(r'[A-Z]+',cell)[0]==column]
                if located:
                    result['dimensions'].extend({'id':key,'ranges':located,'value':value,'original':original}
                        for key,value in [('period',date),('template_period_role',
                            {'本期':'current','上期':'prior','期末':'closing','期初':'opening'}[label])])
            for sentence in plain.split('。'):
                paired=re.search(r'分别为(本期|上期|期末|期初)[、，,](本期|上期|期末|期初)',sentence)
                located=qualify(coordinate.findall(sentence))
                if paired and len(located)==2 and all(k in named_periods for k in paired.groups()):
                    for area,label in zip(located,paired.groups()):
                        result['dimensions'].append({'id':'period','ranges':[area],'value':named_periods[label],'original':original})
                        result['dimensions'].append({'id':'template_period_role','ranges':[area],
                            'value':{'本期':'current','上期':'prior','期末':'closing','期初':'opening'}[label],'original':original})
        row_reference=re.search(r'行范围按上述第(\d+)[—\-至](\d+)行对象对应',plain)
        if number and explicit and re.search(r'上述同一[^；;。]*类资产|同一[^；;。]*行范围',plain):
            from .A语义阅读 import addresses
            previous_values={}
            for row in result['dimensions']:
                if row.get('original')!=lines[number-1] or row['id'] not in refs or standard['dimension'].get(row['id'],{}).get('value_type')!='selection':continue
                for area in row['ranges']:
                    if area=='*':continue
                    if '!' in area and area.rsplit('!',1)[0]!=sheet:continue
                    for cell in addresses(area.rsplit('!',1)[-1]):
                        previous_values.setdefault((row['id'],int(re.search(r'\d+$',cell)[0])),[]).append(row['value'])
            # 只承接紧邻原句明确的同一行范围，不按科目或相似文字补范围。
            for area in qualify(coordinate.findall(line)):
                for cell in addresses(area.rsplit('!',1)[-1]):
                    for (key,row_number),values in previous_values.items():
                        if row_number==int(re.search(r'\d+$',cell)[0]) and all(v==values[0] for v in values):
                            result['dimensions'].append({'id':key,'ranges':qualify([cell]),'value':deepcopy(values[0]),'original':original})
                            explicit_selection=True
        row_values=row_objects if row_reference else row_texts
        named_rows={names[match[1]] for match in found}&set(row_values)
        expanded_row_fields=set()
        if row_values and (explicit or row_reference or named_rows):
            from .A语义阅读 import addresses
            mentioned_refs=({ref['id'] for key in explicit for ref in standard['metric'][key].get('dimension_refs',[])}
                            if explicit else set(row_values) if row_reference else named_rows)
            for key,values in row_values.items():
                if key not in mentioned_refs:continue
                for area in dict.fromkeys(qualify(coordinate.findall(line)) or (scope if named_rows else [])):
                    for cell in addresses(area.rsplit('!',1)[-1]):
                        row_number=int(re.search(r'\d+$',cell)[0])
                        if row_reference and not int(row_reference[1])<=row_number<=int(row_reference[2]):continue
                        value=values.get(row_number)
                        if value is not None:
                            result['dimensions'].append({'id':key,'ranges':[(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell],'value':value,'original':original})
                            expanded_row_fields.add(key)
                            if row_reference:explicit_selection=True
        rules=[(label,key) for label,key in re.findall(r'([^\s，,；;。]+)用\s*(metric\.[\w.]+)',plain) if key in standard['metric']]
        inherited=bool(re.search(r'指标及变动方向同上',plain))
        if rules:
            group_rules=rules;group_direction=None
            directions=re.search(r'方向分别为([^，,；;。]+)',plain)
            if directions:
                values=directions[1].split('、')
                linked=[(key,ref['id']) for _,key in rules for ref in standard['metric'][key].get('dimension_refs',[])
                        if ref['id'].endswith('_direction') and all(v in standard['dimension'][ref['id']].get('allowed_values',[]) for v in values)]
                if len(linked)==1:group_direction=(*linked[0],values)
        if item_groups and (rules or inherited):
            for areas,description,source in item_groups:
                choices=[key for label,key in group_rules if label in description]
                direction=group_direction
                if direction and len(areas)==len(direction[2]) and all(v in description for v in direction[2]):
                    positions=[description.index(v) for v in direction[2]]
                    if positions==sorted(set(positions)):
                        choices.append(direction[0])
                        result['dimensions'].extend({'id':direction[1],'ranges':[area],'value':value,
                            'original':{'original':source,'binding_reference':original}} for area,value in zip(areas,direction[2]))
                if len(set(choices))==1:
                    result['metrics'].append({'id':choices[0],'ranges':areas,
                        'original':{'original':source,'binding_reference':original}})
                    referenced_metric=True
        for unknown_columns in re.finditer(r'(?<![A-Z0-9])([A-Z]+(?:[、，,]\s*[A-Z]+)*)(?:列)?\s*(?:均)?无适用标准',plain):
            columns=set(re.findall(r'[A-Z]+',unknown_columns[1]))
            locations=[(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell for area in scope if area!='*'
                       for cell in addresses(area.rsplit('!',1)[-1]) if re.match(r'[A-Z]+',cell)[0] in columns]
            if locations:result['unresolved'].append({'id':'metric_id','value':'未知','ranges':locations,
                'reason':'原答复明确无适用标准','original':original})
        explicit_columns=column_metrics(line) if explicit else {}
        located_metric=None
        for clause in re.split(r'[；;]',line.replace('`','')):
            assignment=re.search(r'((?:'+coordinate.pattern+r'[\s、，,]*)+)[:：](?!\s*\$?[A-Z]+\$?\d)\s*(.*)',clause)
            if not assignment:continue
            ids=metrics.findall(assignment[2])
            if len(ids)==1:located_metric=ids[0]
            elif not (located_metric and re.match(r'同一(?:变动)?槽位',assignment[2])):continue
            result['metrics'].append({'id':located_metric,'ranges':qualify(coordinate.findall(assignment[1])),'original':original})
            direction=re.search(r'(?:\w+_direction|方向)\s*=\s*([^，,；;。]+)',assignment[2])
            keys=[r['id'] for r in standard['metric'][located_metric]['dimension_refs'] if direction and r['id'].endswith('_direction')
                  and direction[1] in standard['dimension'][r['id']].get('allowed_values',[])]
            if len(keys)==1:result['dimensions'].append({'id':keys[0],'ranges':qualify(coordinate.findall(assignment[1])),
                'value':direction[1],'original':original})
            referenced_metric=True
        directions=re.search(r'方向分别为([^，,；;。]+)',line.replace('`',''))
        if directions and explicit_columns and table_scope:
            values=directions[1].split('、')
            for key in standard['dimension']:
                columns=[col for col,metric in explicit_columns.items() if key.endswith('_direction')
                    and key in {ref['id'] for ref in standard['metric'][metric]['dimension_refs']}]
                if len(columns)==len(values) and all(value in standard['dimension'][key].get('allowed_values',[]) for value in values):
                    for col,value in zip(columns,values):
                        locations=[(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell for area in qualify(table_scope)
                            for cell in addresses(area.rsplit('!',1)[-1]) if re.match(r'[A-Z]+',cell)[0]==col]
                        result['dimensions'].append({'id':key,'ranges':locations,'value':value,'original':original})
        if re.search(r'[A-Z]+列为',line) and re.search(r'[\u4e00-\u9fff]+格适用',line):referenced_metric=True
        continued_columns=(nested_scope or number>0 and original[:1].isspace() and coordinate.search(lines[number-1])) and scope and '*' not in scope
        if explicit_columns and not spans and column_table_scope:
            scope=column_table_scope[:];continued_columns=True
        stated_roles=dict(re.findall(r'(?<![A-Z])([A-Z]+)列为(期初|期末|本期|上期)',line))
        if (explicit_columns or stated_roles) and (spans or continued_columns):
            from .A语义阅读 import addresses
            for area in scope:
                for cell in addresses(area.rsplit('!',1)[-1]):
                    column=re.match(r'[A-Z]+',cell)[0];key=explicit_columns.get(column)
                    role=stated_roles.get(column)
                    if key:
                        row={'id':key,'ranges':qualify([cell]),'original':original}
                        if role:row['template_period_role']=role
                        result['metrics'].append(row)
                        denominator=re.search(r'分母范围(?:为|是|[:：])\s*([^；;。]+)',line)
                        keys=[ref['id'] for ref in standard['metric'][key]['dimension_refs'] if ref['id'].endswith('_denominator_selection')]
                        if denominator and len(keys)==1:result['dimensions'].append({'id':keys[0],'ranges':qualify([cell]),'value':denominator[1],'original':original})
                    elif role:
                        result['dimensions'].append({'id':'template_period_role','ranges':qualify([cell]),'value':role,'original':original})
            if explicit_columns or not explicit:referenced_metric=True
        if len(selected)==len(spans)>1 and '分别' in head:
            result['metrics'].extend({'id':key,'ranges':[area],'original':original} for area,key in zip(spans,selected))
            if not found:continue
        if found and len(selected)==1 and not referenced_metric:result['metrics'].append({'id':selected[0],'ranges':scope[:],'original':original})
        if not found:
            selected=[] if refused else metrics.findall(line)
            if len(selected)==1 and not referenced_metric:
                row={'id':selected[0],'ranges':scope[:] if spans or original[:1].isspace() and scope!=['*'] else [],'original':original}
                result['metrics'].append(row)
                if not spans:metric_group=row
            elif spans and not dated and not referenced_metric and not explicit_selection and not nested_metric and not any(
                row.get('original')==original for row in result['dimensions']) and line.strip(' -#:：;；') not in spans:
                result['unresolved'].append({'ranges':scope[:],'reason':'这段答复未明确可机械读取的字段归属','original':original})
            section_lines.append(plain)
            continue
        field_ranges=scope[:];period_refs={}
        if re.match(r'以上坐标共同取值[:：]',line):
            field_ranges=list(dict.fromkeys(area for row in result['metrics'] for area in row['ranges'] if area!='*'))
        for i,match in enumerate(found):
            key=names[match[1]];end=found[i+1].start() if i+1<len(found) else len(line)
            if match[1]!=key and any(m[1]==key and m.end()<match.start() and not coordinate.search(line[m.end():match.start()]) for m in found):continue
            field_columns=[]
            if nested_scope:
                column_parts=list(re.finditer(r'([A-Z]+(?:[、，,][A-Z]+)*)列',line))
                preceding=[part for part in column_parts if part.end()<=match.start()]
                if preceding:
                    field_columns=re.findall(r'[A-Z]+',preceding[-1][1])
                    following=next((part.start() for part in column_parts if part.start()>match.end()),None)
                    if following is not None:end=min(end,following)
            if key in expanded_row_fields and '该行所列区间' in line[match.end():end]:continue
            if key in ('period','template_period_role') and period_objects and not coordinate.search(line):continue
            ordered=key!='metric_id' and ('依次' in match[0] or '分别' in match[0] or re.match(r'(?:按行)?(?:依次|分别)为',line[match.end():])
                or re.search(r'各行分别取\s*`?$',line[:match.start()]))
            multiple_metrics=key=='metric_id' and len(set(metrics.findall(line[match.end():end])))>1
            # 行内代码已有明确边界，其后的解释不能混成字段取值。
            quoted=next((m for m in re.finditer(r'`[^`]*`',line) if m.start()<match.start()<m.end()),None)
            if quoted and quoted.end()>match.end() and not multiple_metrics and not ordered:end=quoted.end()-1
            leading=re.match(r'`([^`]+)`',line[match.end():])
            components=re.match(r'`?(?:domain|mode|members|predicate)`?\s*=',line[match.end():])
            mode_prefix=leading and leading[1] in ('members','predicate') and re.match(r'\s*(?:[:：，,]|—+)',line[match.end()+leading.end():])
            if leading and not mode_prefix and not components and not multiple_metrics and not ordered and (key!='metric_id' or leading[1] in standard['metric']):end=match.end()+leading.end()
            quoted_value=re.match(r'“([^”]+)”',line[match.end():])
            quote_continues=quoted_value and standard['dimension'].get(key,{}).get('value_type')=='selection' and re.match(r'[^\s。，,；;）)]',line[match.end()+quoted_value.end():])
            if quoted_value and not ordered and not quote_continues and not re.match(r'(?:全体|总体)\s*[（(](?:mode\s*[:：]\s*)?all[）)]',line[match.end()+quoted_value.end():]):end=match.end()+quoted_value.end()
            boundary=next((point for point in table_cell_ends if point>=match.end()),None)
            if boundary is not None:end=min(end,boundary)
            segment=line[match.end():end];column_selections=[]
            if nested_scope:segment=re.split(r'。(?=第\d+行|各格)',segment,maxsplit=1)[0]
            annotation=re.match(r'\s*`?(all|unknown|members|predicate)`?\s*[）)]',segment)
            if annotation and re.search(r'[（(]\s*`?$',line[:match.start()]):
                literal=re.search(r'[A-Za-z_\u4e00-\u9fff]+?\s*(?:为|[:：=＝])\s*([^；;。（()）]+?)\s*[（(]\s*`?$',line[:match.start()])
                if annotation[1] in ('all','unknown'):segment=annotation[1]
                elif literal:segment=annotation[1]+'：'+literal[1]
                else:
                    result['unresolved'].append({'id':key,'ranges':field_ranges[:],
                        'reason':'只明确表达方式，未明确可机械读取的取值','original':original})
                    continue
            if key in ratio_names.values():
                segment=re.split(r'。\s*[^。]*的总计分子和分母都对应',segment,maxsplit=1)[0]
            if i+1<len(found):segment=re.sub(r'[,，；;]\s*(?:分子|分母)(?:维度)?\s*`?$','',segment)
            if key=='entity':segment=re.split(r'。|[；;]|[,，]\s*(?:报告主体|主体|报表范围|报表口径)',segment,maxsplit=1)[0]
            if key=='currency':segment=re.split(r'[。]|[、，,]\s*(?:金额单位|(?:金额)?倍率)',segment,maxsplit=1)[0].rstrip(' 、，,')
            segment=re.split(r'[；;]\s*分别为(?:本期|上期|期末|期初)[、，,](?:本期|上期|期末|期初)(?:余额)?(?=[。；;]|$)',segment,maxsplit=1)[0]
            segment=re.split(r'[；;]\s*(?:时点|期间)分别为\s*'+date_pattern+r'[、，]\s*'+date_pattern+r'(?=[。；;]|$)',segment,maxsplit=1)[0]
            if standard['dimension'].get(key,{}).get('value_type')=='selection':
                segment=re.split(r'。\s*`?(?:'+coordinate.pattern+r'|[A-Z]+)`?(?:列)?(?:为|[:：])\s*`?(?:期末|期初|本期|上期|'+date_pattern+r')',segment,maxsplit=1)[0]
                trailing_metric=re.search(r'[,，]\s*`([^`]+)`[。]?\s*$',segment)
                if trailing_metric and trailing_metric[1] in standard['metric']:segment=segment[:trailing_metric.start()]
                segment=re.split(r'[；;]\s*(?:其他必需维度|[A-Z]+列(?:期末|期初|本期|上期))',segment,maxsplit=1)[0]
                segment=re.split(r'[。；;]\s*[A-Z]+(?:[、，,]\s*[A-Z]+)*\s*(?:无适用标准|用\s*`?metric\.)',segment,maxsplit=1)[0]
                segment=re.split(r'[；;]\s*'+coordinate.pattern+r'\s*(?:期间|period)',segment,maxsplit=1)[0]
                column_selections=list(re.finditer(r'[，,；;]\s*([A-Z]+)列为(.*?)(?=[，,；;]\s*[A-Z]+列为|[。；;]|$)',segment.split('。',1)[0]))
                if column_selections:segment=segment[:column_selections[0].start()]
            value=segment.lstrip(' \t:：').strip(' \t；;|，、。`')
            if quoted and quoted.end()==end+1 and not components:value=segment.strip()
            if leading and not mode_prefix and not components and not multiple_metrics and not ordered:value=leading[1]
            if key in table_literals:
                literal=re.match(r'^`([^`]+)`',table_literals[key])
                value=literal[1] if literal else table_literals[key].strip(' \t。；;')
                if literal and literal[1] in ('members','predicate') and re.match(r'\s*[:：，,=]',table_literals[key][literal.end():]):
                    value+=table_literals[key][literal.end():]
            if standard['dimension'].get(key,{}).get('value_type')=='selection' and isinstance(value,str) and re.match(
                    r'^(?:上表对应行|该行|对应行|所在行)(?:所选|的)?范围(?=$|[，,。；;])',value):
                # 这是对各行取值的引用，不能另存为覆盖各行的共同取值。
                continue
            if key=='metric_id':
                stated_ids=list(dict.fromkeys(metrics.findall(table_metric_text or segment)))
                if len(stated_ids)==1:value=stated_ids[0]
                elif referenced_metric and stated_ids:continue
            if key=='currency':
                code=re.search(r'(?:代码|编码)\s*`?([A-Z]{3})`?',value)
                if code:value=code[1]
            literal=re.match(r'(all|unknown)`?(?=[。；;：:])',value)
            if literal and standard['dimension'].get(key,{}).get('value_type')=='selection':value=literal[1]
            if standard['dimension'].get(key,{}).get('value_type')=='selection':
                common_literal=re.match(r'^(?:所有|全部|上述|以下)[^；;。]*?均为\s*`([^`]+)`',value)
                if common_literal:value=common_literal[1]
                explicit_all=re.match(r'^(?:各格|全部格|上述格|各项)?\s*(?:均为|均是|均取)?\s*`?(all|unknown)`?(?=$|[。；;，,\s])',value)
                if explicit_all:value=explicit_all[1]
                if re.match(r'^(?:未知|unknown)(?=$|[。；;，,:：\s])',value):value='未知'
                literal_all=re.match(r'domain“([^”]+)”[、，,]\s*`?mode\s*=\s*all\b',value)
                if literal_all:value=json.dumps({'domain':literal_all[1],'mode':'all'},ensure_ascii=False)
            if not value and i+1<len(found) and names[found[i+1][1]]==key:
                continue
            if key in standard['dimension'] and re.match(r'^(?:未知|未确定|unknown)(?=$|[。；;，,:：\s])',value):value='未知'
            if re.fullmatch(r'原文未说明(?:是|为|是否)[^，,。；;]*[，,]留缺',value):value='未知'
            if row_span and '各格具体' in plain and value.startswith('对应'):
                continue
            if (key=='period' or standard['dimension'].get(key,{}).get('value_type')=='selection') and re.search(r'(?:期间|范围)\s*未知$',value):
                value='未知'
            stated_role=(re.fullmatch(r'(?:时点|期间)?\s*'+date_pattern+
                r'\s*[（(](?:表头角色[:：]\s*)?(本期|上期|期末|期初)(?:余额|发生额|增加|减少)?[）)]',value)
                if key=='period' else None)
            if row_declaration and key==row_field and not value:continue
            label_refs=[]
            if key=='template_period_role':
                value=re.split(r'[；;。]',value,maxsplit=1)[0].strip()
                value=re.sub(r'[（(](?:本年|上年)[）)]$','',value)
                value=re.sub(r'[,，]\s*实际$','',value)
                label=re.search(r'\s*[（(]原表期间标签引用[:：]\s*([^（）()]+)[）)]\s*$',value)
                if label and not coordinate.sub('',label[1]).strip(' `、，,'):
                    label_refs=qualify(coordinate.findall(label[1]));value=value[:label.start()].strip()
            if key=='period' and isinstance(value,str):
                role=value.split('。',1)[0].strip()
                if role in named_periods:value=named_periods[role]
            members=re.match(r'[^；;。`]*域内的\s*`?members`?\s*[:：]\s*`([^`]+)`',value)
            if members:value={'mode':'members','members':[members[1]]}
            parenthesized=re.fullmatch(r'(members|predicate)\s*[（(](.+)[）)]',value) if isinstance(value,str) else None
            if parenthesized and not (parenthesized[1]=='members' and re.search(r'[,，]',parenthesized[2])):
                value=parenthesized[1]+'：'+parenthesized[2]
            if isinstance(value,str) and re.match(r'^unknown\s*[（(]',value):value='未知'
            literal_predicate=re.fullmatch(r'predicate\[([^\[\]]+)\]',value) if isinstance(value,str) else None
            if literal_predicate:value={'mode':'predicate','predicate':literal_predicate[1]}
            quoted_mode=re.match(r'^(members|predicate)\s*[「“"]([^」”"]+)[」”"]',value) if isinstance(value,str) else None
            if quoted_mode:value={'mode':quoted_mode[1],quoted_mode[1]:[quoted_mode[2]] if quoted_mode[1]=='members' else quoted_mode[2]}
            stated_mode=re.match(r'^`?(members|predicate)`?\s*(?:[:：，,；;]|—+)\s*(.+)',value) if isinstance(value,str) and standard['dimension'].get(key,{}).get('value_type')=='selection' else None
            if stated_mode:
                content=stated_mode[2].split('。',1)[0].strip()
                wrapped=bool(re.fullmatch(r'“[^“”]*”|「[^「」]*」|"[^"]*"',content))
                if wrapped:content=content[1:-1]
                if stated_mode[1]=='members':
                    content=re.sub(r'[（(]原文列明的并列类别[）)]$','',content)
                    try:listed=json.loads(content)
                    except ValueError:listed=None
                    stated_members=re.match(r'^(?:实际)?成员为\s*((?:“[^”]+”|「[^」]+」|"[^"]+")(?:\s*(?:和|及|、)\s*(?:“[^”]+”|「[^」]+」|"[^"]+"))*)',content)
                    if stated_members:listed=re.findall(r'[“「"]([^”」"]+)[”」"]',stated_members[1])
                    parts=[content] if wrapped or re.search(r'[；;]',content) else re.split(r'[、，,]',content)
                    parts=[name.strip(' `') for name in parts]
                    value={'mode':'members','members':listed if isinstance(listed,list) else
                           [name[1:-1] if re.fullmatch(r'“[^“”]*”|「[^「」]*」|"[^"]*"',name) else name for name in parts]}
                else:value={'mode':'predicate','predicate':content}
            if not ordered and isinstance(value,str) and value.startswith('“') and value.endswith('”'):value=value[1:-1]
            value_ranges=above_ranges.get(key,field_ranges)[:]
            if field_columns:
                value_ranges=[(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell
                    for area in value_ranges if area!='*' for cell in addresses(area.rsplit('!',1)[-1])
                    if re.match(r'[A-Z]+',cell)[0] in field_columns]
            if nested_scope:
                prefix=line[:match.start()]
                sentence=prefix.rsplit('。',1)[-1]
                row_text='、'.join(re.findall(r'第([\d、，,]+)行',sentence))
                if not row_text and i==0:row_text='、'.join(re.findall(r'第([\d、，,]+)行',prefix))
                if row_text:
                    numbers=re.findall(r'\d+',row_text)
                    value_ranges=[(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell
                        for area in value_ranges if area!='*' for cell in addresses(area.rsplit('!',1)[-1])
                        if re.search(r'\d+$',cell)[0] in numbers]
                    listed=re.findall(r'“([^”]+)”',segment)
                    if ordered and len(listed)==len(numbers):
                        for row_number,literal in zip(numbers,listed):
                            result['dimensions'].append({'id':key,'ranges':[cell for cell in value_ranges if re.search(r'\d+$',cell)[0]==row_number],
                                'value':literal,'original':original})
                        continue
            if table_scope and re.search(r'行范围\s*=\s*`?$',line[:match.start()]):value_ranges=qualify(table_scope)
            same_field=re.match(r'^同\s*('+coordinate.pattern+r')(?=$|[；;。])',value.replace('`','')) if isinstance(value,str) else None
            if same_field:
                field_references.append((value_ranges,qualify([same_field[1]])[0],set(names.values())-{key},original))
                continue
            same_literal=re.fullmatch(r'与(.+?)相同',value) if isinstance(value,str) else None
            if same_literal:
                referenced=[r for r in result['dimensions'] if r['id']==key and same_literal[1] in str(r.get('value',''))]
                choices={json.dumps(r['value'],ensure_ascii=False,sort_keys=True):r for r in referenced}
                if len(choices)==1:
                    previous=next(iter(choices.values()))
                    result['dimensions'].append({'id':key,'ranges':value_ranges,'value':deepcopy(previous['value']),
                        'original':{'original':previous.get('original'),'field_reference':original}})
                    continue
            if value=='同上':
                previous=next((r for r in reversed(result['dimensions']) if r['id']==key and any(
                    area!='*' and area.rpartition('!')[0]==sheet for area in r['ranges'])),None)
                if previous:
                    result['dimensions'].append({'id':key,'ranges':value_ranges,'value':deepcopy(previous['value']),
                        'original':{'original':previous.get('original'),'field_reference':original}})
                else:result['unresolved'].append({'id':key,'ranges':value_ranges,'reason':'同上没有明确前项','original':original})
                continue
            same_row=re.match(r'^分别与同一行([A-Z]+)列[^。；;]*一致',value) if isinstance(value,str) else None
            if same_row:
                from .A语义阅读 import addresses
                for area in value_ranges:
                    for cell in addresses(area.rsplit('!',1)[-1]):
                        target=(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell
                        source=(area.rsplit('!',1)[0]+'!' if '!' in area else '')+same_row[1]+re.search(r'\d+$',cell)[0]
                        choices=[r for r in result['dimensions'] if r['id']==key and any(
                            source in [(span.rsplit('!',1)[0]+'!' if '!' in span else sheet+'!')+a
                                       for a in addresses(span.rsplit('!',1)[-1])]
                            for span in r['ranges'] if span!='*')]
                        if choices and all(r['value']==choices[0]['value'] for r in choices):
                            result['dimensions'].append({'id':key,'ranges':[target],'value':deepcopy(choices[0]['value']),
                                'original':{'original':choices[0].get('original'),'field_reference':original}})
                        else:result['unresolved'].append({'id':key,'ranges':[target],'reason':'同行引用没有唯一前项','original':original})
                continue
            if key=='period' and re.match(r'(?:时点|期间)?\s*\d{4}',value):
                value=value.split('。',1)[0]
            if isinstance(value,str) and value.startswith(('按上列格子','按所在行分别取')):
                linked=[]
                definition=standard['dimension'].get(key,{})
                for earlier in section_lines:
                    for clause in re.split(r'[；;。]',earlier):
                        located=qualify(coordinate.findall(clause))
                        if not located:continue
                        if value.startswith('按上列格子'):
                            choices=[v for v in definition.get('allowed_values',[]) if clause.strip().endswith(v)]
                            if len(choices)!=1:continue
                            actual=choices[0]
                        else:
                            item=re.search(r'\d+\s+([^：:]+)[：:]各阶段',clause)
                            if not item or item[1] not in value:continue
                            actual=item[1]
                        source_original=next((r['original'] for r in result['unresolved']
                            if isinstance(r.get('original'),str) and
                            r['original'].strip().replace('`','').replace('**','')==earlier),earlier)
                        linked.append({'id':key,'ranges':located,'value':actual,
                            'original':{'original':source_original,'field_reference':original}})
                if linked:
                    result['dimensions'].extend(linked)
                    continue
            if key=='period' and (re.search(r'[A-Z]+列',value) or re.search(r'(?:(本期|上期|期末|期初)(?:余额|变动|对应)?(?:的)?|[A-Z]+列)\s*`?$',line[:match.start()])):
                continue
            if key=='period' and re.search(date_pattern+r'\s*[（(][A-Z]+[）)]',value):continue
            if key=='period':
                date=re.match(r'(?:(?:实际)?(?:日期|时点|期间)(?:为)?\s*`?|instant\s*[,，]\s*)?('+date_pattern+r')(?=$|时点|期间|[\s，,；;。（()）])',value)
                if date and len(re.findall(date_pattern,value))==1:value=re.sub(r'\s+','',date[1])
                year=re.match(r'(\d{4})年度(?=$|[。；;])',value)
                if year and section_duration and all(date.startswith(year[1]) for date in section_duration.split('至')):
                    value=section_duration
            if components:
                if leading and re.search(r'[,，；;]\s*(?:domain|mode|members|predicate)\s*=',leading[1]):
                    value=leading[1]
                pairs=re.findall(r'`?(domain|mode|members|predicate)`?\s*=\s*(\[[^\]]*\]|.*?)(?=\s*[,，；;]\s*`?(?:domain|mode|members|predicate)`?\s*=|$)',value)
                quoted_components={k:v for k,v in re.findall(r'`(domain|mode|members|predicate)\s*=\s*([^`]+)`',segment)
                    if not re.search(r'[,，；;]\s*(?:domain|mode|members|predicate)\s*=',v)}
                parsed={}
                for name,part in pairs:
                    part=quoted_components[name] if name in quoted_components else part.strip(' `。')
                    if name=='members':part=part.split('。',1)[0]
                    if name=='mode':part=part.split('。',1)[0]
                    try:part=json.JSONDecoder().raw_decode(part)[0] if name=='members' and part.startswith('[') else json.loads(part)
                    except ValueError:
                        if part.startswith('[') and part.endswith(']'):part=[x.strip(' "\'“”「」') for x in re.split(r'[,，、]',part[1:-1])]
                    if name=='members' and isinstance(part,str):part=[part.split('。',1)[0]]
                    parsed[name]=part
                if parsed:value=parsed
            elif standard['dimension'].get(key,{}).get('value_type')=='number' and isinstance(value,str):
                number_value=re.match(r'^[+-]?\d+(?:\.\d+)?(?=$|[。；;\s（）()]|[，,]\s*原文(?:说明|为))',value)
                if not number_value and key=='unit_scale':
                    number_value=re.match(r'^[+-]?\d+(?:\.\d+)?(?=元(?:$|[。；;\s])|[，,]\s*(?:人民币|金额单位))',value)
                if number_value:value=number_value[0]
            # 每一分句中的坐标限定该分句字段；前置共同字段适用正文全部坐标。
            prefix=line[found[i-1].end() if i else 0:match.start()].rsplit('；',1)[-1].rsplit(';',1)[-1].rsplit('。',1)[-1]
            located=coordinate.findall(prefix)
            if located:
                value_ranges=qualify(located);field_ranges=value_ranges[:]
            bounded_values={m[1]:m[2].strip() for m in column_selections}
            declaration=re.search(r'([A-Z]+(?:[、，,]\s*[A-Z]+)+)列分别(?:属于|为)([^。；;]+)',line[:match.start()])
            if declaration and isinstance(value,str):
                columns=re.split(r'[、，,]\s*',declaration[1])
                sequence=re.split(r'[,，]\s*[A-Z]+列',declaration[2],maxsplit=1)[0]
                items=[part.strip() for part in re.split(r'[、，,]',sequence)]
                ordinal=r'第[一二三四五六七八九十\d]+'
                suffix=re.fullmatch('('+ordinal+r')(.+)',items[-1]) if items else None
                if (suffix and len(columns)==len(items) and len(set(columns))==len(columns)
                    and all(re.fullmatch(ordinal,item) for item in items[:-1]) and '该'+suffix[2] in value):
                    items=[item+suffix[2] for item in items[:-1]]+[items[-1]]
                    bounded_values.update((column,value.replace('该'+suffix[2],item)) for column,item in zip(columns,items))
            if bounded_values:
                from .A语义阅读 import addresses
                positions=[(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell
                           for area in value_ranges if area!='*' for cell in addresses(area.rsplit('!',1)[-1])]
                for column,bounded_value in bounded_values.items():
                    selected=[cell for cell in positions if re.match(r'[A-Z]+',cell.rsplit('!',1)[-1])[0]==column]
                    if selected:
                        result['dimensions'].append({'id':key,'ranges':selected,'value':bounded_value,'original':original})
                value_ranges=[cell for cell in positions if re.match(r'[A-Z]+',cell.rsplit('!',1)[-1])[0] not in bounded_values]
                if not value_ranges:continue
            if stated_role:
                result['dimensions'].append({'id':'template_period_role','ranges':value_ranges[:],
                    'value':{'本期':'current','上期':'prior','期末':'closing','期初':'opening'}[stated_role[1]],'original':original})
            if heading:
                predicate=re.match(r'predicate\s*[「“"]([^」”"]+)[」”"]',value) if isinstance(value,str) else None
                if predicate:value={'mode':'predicate','predicate':predicate[1]}
            if standard['dimension'].get(key,{}).get('value_type')=='selection' and isinstance(value,str):
                explicit_mode=re.match(r'([^；;。]+?)[（(]\s*`?(?:mode\s*[:：]\s*)?(all|members)`?\s*[）)](?=$|[；;。])',value)
                if explicit_mode:
                    value={'mode':explicit_mode[2]}
                    if explicit_mode[2]=='members':value['members']=[part.strip() for part in explicit_mode[1].split('、')]
            if isinstance(value,str):
                literal_object=re.fullmatch(r'(\{[^{}]*\})[，,；;。].*',value)
                if literal_object:value=literal_object[1]
            if isinstance(value,str) and (value.startswith('{') and value.endswith('}') or re.fullmatch(r'domain\s*:\s*[^,{}]+,\s*mode\s*:\s*(?:all|members|unknown|predicate)',value)):
                # 只读取答复已写明的平面键值，不生成任何省略字段。
                pair_pattern=r'(\w+)\s*:\s*(\[[^\[\]]*\]|[^,{}]+)'
                pairs=re.findall(pair_pattern,value)
                if pairs and re.sub(pair_pattern+r'|[,{}\s]','',value)=='':
                    value={k:v.strip().strip('"\'') for k,v in pairs}
                    for name,part in value.items():
                        if part.startswith('[') and part.endswith(']'):
                            try:value[name]=json.loads(part)
                            except ValueError:value[name]=[x.strip(' "\'“”「」') for x in re.split(r'[,，、]',part[1:-1])]
            placed=re.match(r'(.+?)\s*(?:均为|为)\s*(.+)',value) if isinstance(value,str) else None
            if placed and coordinate.findall(placed[1]) and not coordinate.sub('',placed[1]).strip(' 、,，'):
                value_ranges=coordinate.findall(placed[1]);value=placed[2]
                quoted_value=re.match(r'[“"]([^”"]+)[”"]',value)
                if quoted_value:value=quoted_value[1]
            try:value=json.loads(value)
            except (ValueError,TypeError):pass
            if ordered and isinstance(value,str):
                by_row=value.startswith('按行')
                sequence=re.sub(r'^(?:按行)?(?:依次|分别)为','',value).split('。',1)[0].strip(' ；;')
                reference=re.fullmatch(r'上述([一二三四五六七八九十]|\d+)(?:档|项|个)(?:[\u4e00-\u9fff]+)?',sequence)
                quotes=list(re.finditer(r'“([^”]+)”',sequence))
                items=ordered_values.get(key,[]) if reference else [item[1] for item in quotes]
                literals=re.findall(r'members\[[^\]]+\]',sequence)
                if not reference and not quotes and len(literals)>1:items=literals
                if quotes and not reference:
                    prefix=sequence[:quotes[0].start()];suffix=sequence[quotes[-1].end():]
                    separators=[sequence[left.end():right.start()] for left,right in zip(quotes,quotes[1:])]
                    # 共同措辞逐项原样带入，不把引号内简称当成完整范围。
                    if all(not part.strip(' 、，,') for part in separators):
                        items=[prefix+item+suffix for item in items]
                    else:items=[]
                if not reference and not items and not quotes:
                    sequence=re.split(r'[；;]',sequence,maxsplit=1)[0]
                    items=[v.strip(' “”,，').removeprefix('以及') for v in re.split(r'[、，,]',sequence)]
                    ordinal=r'第[一二三四五六七八九十\d]+'
                    suffix=re.fullmatch('('+ordinal+r')(.+)',items[-1]) if items else None
                    if suffix and all(re.fullmatch(ordinal,item) for item in items[:-1]):
                        items=[item+suffix[2] for item in items[:-1]]+[items[-1]]
                    elif any(re.fullmatch(ordinal,item) for item in items):items=[]
                count=(int(reference[1]) if reference[1].isdigit() else '一二三四五六七八九十'.index(reference[1])+1) if reference else len(items)
                if items and len(items)==count:
                    from .A语义阅读 import addresses
                    located=[(area.rsplit('!',1)[0]+'!' if '!' in area else '')+c
                             for area in value_ranges if area!='*' for c in addresses(area.rsplit('!',1)[-1])]
                    grouped=[];seen=set();prefix=line[:match.start()].replace('`','')
                    pairs=[]
                    for clause in re.split(r'[；;]',prefix):
                        coords=list(coordinate.finditer(clause))
                        if len(coords)==2 and clause[coords[0].end():coords[1].start()]=='/':
                            pairs.append(qualify([c[0] for c in coords]))
                    if len(pairs)==count and all(':' not in cell.rsplit('!',1)[-1] for pair in pairs for cell in pair):
                        result['dimensions'].extend({'id':key,'ranges':pair,'value':item,'original':original} for pair,item in zip(pairs,items))
                        continue
                    for clause in re.split(r'[；;]',prefix):
                        coords=list(coordinate.finditer(clause))
                        if not coords:continue
                        label=clause[coords[-1].end():].lstrip(' ，,：:')
                        label=re.sub(r'[。.]?\s*(?:维度|对应)\s*[:：]?\s*$','',metrics.sub('',label)).strip(' :：。')
                        group=[(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell
                               for area in qualify([c[0] for c in coords]) for cell in addresses(area.rsplit('!',1)[-1])]
                        choices=[item for item in items if item==label]
                        if len(choices)==1:
                            grouped.extend((cell,choices[0]) for cell in group);seen.add(choices[0])
                        elif len(group)==len(items) and re.search(r'(?:本期|分别为|依次为)'+re.escape('、'.join(items))+r'$',label):
                            grouped.extend(zip(group,items));seen.update(items)
                    row_labels=re.findall(r'行(\d+)([^、，；;。]+)',prefix)
                    if row_labels:
                        details={(int(row),label):value for row,label,value in re.findall(r'行(\d+)的“([^”]+)”按原文说明为([^。]+)',plain)
                                 if label in items}
                        for row,label in details:
                            if any(int(r)==row and name!=label for r,name in row_labels):
                                result['unresolved'].append({'id':key,'ranges':[cell for cell in located if int(re.search(r'\d+$',cell)[0])==row],
                                    'reason':'原答复同一行标签与详细说明不一致','original':original})
                        grouped.extend((cell,details.get((int(row),label),label)) for row,label in row_labels if label in items
                                       for cell in located if int(re.search(r'\d+$',cell)[0])==int(row))
                        seen.update(label for row,label in row_labels if label in items)
                    if grouped and seen==set(items) and all(cell in located for cell,item in grouped):
                        result['dimensions'].extend({'id':key,'ranges':[cell],'value':item,'original':original}
                            for cell,item in grouped if cell in located)
                        continue
                    groups=[[(area.rsplit('!',1)[0]+'!' if '!' in area else '')+c
                             for c in addresses(area.rsplit('!',1)[-1])] for area in value_ranges if area!='*']
                    if by_row:
                        groups=[[cell for cell in group if cell.rsplit('!',1)[-1].rstrip('0123456789')==col]
                                for group in groups for col in dict.fromkeys(cell.rsplit('!',1)[-1].rstrip('0123456789') for cell in group)]
                    if len(groups)==count and all(len({re.search(r'\d+$',cell)[0] for cell in group})==1 for group in groups):
                        for group,item in zip(groups,items):
                            role=re.fullmatch(r'(期末|期初|本期|上期)(?:时点|期间)?\s*('+date_pattern+r')?',item) if key=='period' else None
                            value=deepcopy(named_periods.get(role[1],item)) if role else item
                            if role and role[2]:value=re.sub(r'\s+','',role[2])
                            result['dimensions'].append({'id':key,'ranges':group,'value':value,'original':original})
                            if role:result['dimensions'].append({'id':'template_period_role','ranges':group,
                                'value':{'期末':'closing','期初':'opening','本期':'current','上期':'prior'}[role[1]],'original':original})
                        continue
                    if sum(map(len,groups))==count:groups=[sum(groups,[])]
                    if groups and all(len(group)==count for group in groups):
                        for group in groups:
                            for cell,item in zip(group,items):
                                role=re.fullmatch(r'(期末|期初|本期|上期)(?:时点|期间)?\s*('+date_pattern+r')?',item) if key=='period' else None
                                label=role[1] if role else item
                                value=deepcopy(named_periods.get(label,item)) if key=='period' else item
                                if role and role[2]:value=re.sub(r'\s+','',role[2])
                                if selection_defaults.get(key,{}).get('mode')=='members':
                                    value={**selection_defaults[key],'members':[item]}
                                declared=sheet_columns.get((cell.rsplit('!',1)[0] if '!' in cell else sheet,
                                    re.match(r'[A-Z]+',cell.rsplit('!',1)[-1])[0]))
                                if key=='period' and declared and declared[0]==label:value=deepcopy(declared[1])
                                result['dimensions'].append({'id':key,'ranges':[cell],'value':value,'original':original})
                                if role:result['dimensions'].append({'id':'template_period_role','ranges':[cell],
                                    'value':{'期末':'closing','期初':'opening','本期':'current','上期':'prior'}[role[1]],'original':original})
                        ordered_values[key]=items
                        continue
                result['unresolved'].append({'id':key,'ranges':value_ranges,'reason':'原答复顺序列表未唯一对应坐标','original':original})
                continue
            if key=='period' and isinstance(value,str):
                role_dates=dict(re.findall(r'(期末|期初|本期|上期)\s*('+date_pattern+r')',value.split('。',1)[0]))
                if role_dates:
                    wanted={loc.rsplit('!',1)[0]+'!'+cell for loc in value_ranges if '!' in loc
                            for cell in addresses(loc.rsplit('!',1)[-1])}
                    applied=set()
                    for stated in result['dimensions'][:]:
                        if stated['id']!='template_period_role' or stated.get('value') not in role_dates:continue
                        targets={loc.rsplit('!',1)[0]+'!'+cell for loc in stated['ranges'] if '!' in loc
                                 for cell in addresses(loc.rsplit('!',1)[-1])}&wanted
                        if targets:
                            result['dimensions'].append({'id':'period','ranges':sorted(targets),
                                'value':role_dates[stated['value']],'original':original})
                            applied.update(targets)
                    if wanted and applied==wanted:continue
                    if applied:value_ranges=sorted(wanted-applied)
            if key in ratio_positions:
                from .A语义阅读 import addresses
                value_ranges=[(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell
                              for area in value_ranges if area!='*' for cell in addresses(area.rsplit('!',1)[-1])
                              if (area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell not in ratio_positions[key]]
                if not value_ranges:continue
            if (isinstance(value,str) and value not in ('未知','unknown','all') and selection_defaults.get(key,{}).get('mode')=='members'
                and not re.match(r'members\s*[:：]',value)):
                value={**selection_defaults[key],'members':[value]}
            row={'id':key,'ranges':value_ranges,'value':value,'original':original}
            if label_refs:period_refs[tuple(value_ranges)]=label_refs
            if key=='period' and tuple(value_ranges) in period_refs:
                row['evidence']=period_refs[tuple(value_ranges)][:]
                row['label_references']=period_refs[tuple(value_ranges)][:]
            if key=='metric_id':
                if mixed_metric_explanation or refused or explicit_columns and (spans or continued_columns):continue
                if referenced_metric and selected and re.search(r'同(?:一)?(?:变动)?(?:槽位|指标)',line):value=selected[0]
                literal=metrics.match(value) if isinstance(value,str) else None
                if literal and re.match(r'[，,。；;\s]|$',value[literal.end():]):value=literal[1]
                if isinstance(value,str) and value in standard['metric']:
                    row.pop('value');row['id']=value;result['metrics'].append(row);last_metric=value
                    if table_scope is None and not nested_scope and not continued_columns and not coordinate.search(line) and not re.search(r'上述|上列|同上|这些格|全部坐标',line):
                        row['ranges']=[];metric_group=row
                else:
                    unknown=re.match(r'(?:metric_id\s*)?(未知|未决|unknown)(?:$|[：:，,；;。／/\s])',str(value).replace('`',''))
                    if unknown:row['value']=unknown[1]
                    result['unresolved'].append({**row,'reason':str(value) if unknown else '答复未给出确定的现有标准编号'})
            else:
                result['dimensions'].append(row)
                if standard['dimension'].get(key,{}).get('value_type')=='selection':last_selection=key
        if column_period_declared or any(re.search(r'template_period_role|实际(?:日期|期间)为|\bperiod\s*[:：=]',column_text[d.end():]) for d in declarations):
            result['dimensions']=[row for row in result['dimensions'] if not (
                row.get('original')==original and row['id'] in ('period','template_period_role') and row['ranges']==['*'])]
        common_fields=re.search(r'(?:(?:各格)?共同维度(?:取值)?|主体及其他维度实际身份|共同取值|共同适用信息)\s*[:：]',source_line.replace('**',''))
        common_value=re.match(r'^\s*(?:上述|以下|本范围)各格的\s*`?(\w+)`?\s*均为',source_line)
        if common_fields or common_value:
            for row in result['dimensions']:
                if row.get('original')==original and (common_fields or row['id']==common_value[1]):
                    row['ranges']=list(dict.fromkeys(area for item in result['metrics'][section_metric_start:]
                        for area in item['ranges'] if area!='*')) or ['*']
        scoped_metric=re.match(r'^上述\s*`?('+metrics.pattern+r')`?\s*格的',source_line)
        if scoped_metric:
            locations=list(dict.fromkeys(loc for row in result['metrics'] if row['id']==scoped_metric[1] for loc in row['ranges']))
            for row in result['dimensions']:
                if row.get('original')==original and locations:row['ranges']=locations[:]
        section_lines.append(plain)
    # 同一条行说明已明列金额位置时，范围字段随该引用走，不停留在名称格。
    for source_line in text.splitlines():
        plain=source_line.replace('`','').replace('**','')
        label=re.match(r'^\s*[-*]?\s*('+coordinate.pattern+r')[“「"][^”」"]+[”」"][:：]',plain)
        if not label:continue
        sentences=plain.split('。');first=sentences[0]
        targets=[cell for sentence in sentences[1:] if re.search(r'分别对应上述[^。；;]*列指标',sentence)
                 for cell in coordinate.findall(sentence.split('分别对应上述',1)[0])]
        if not targets:continue
        for row in result['dimensions']:
            original=row.get('original','')
            if (isinstance(original,str) and original.strip().replace('`','').replace('**','').rstrip('。')==first.strip()
                and standard['dimension'].get(row['id'],{}).get('value_type')=='selection'
                and qualify(row['ranges'])==qualify([label[1]])):
                row['ranges']=qualify(targets)
                row['original']={'original':original,'field_reference':source_line}
    # 成员名及跨列同指均取自答复原文，不从原表名称推断。
    member_rows=[]
    for row in result['dimensions']:
        original=row.get('original','');key=row['id']
        if row['value']!='members' or not isinstance(original,str) or standard['dimension'].get(key,{}).get('value_type')!='selection':
            member_rows.append(row);continue
        cells=[(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell
               for area in qualify(row['ranges']) if area!='*' for cell in addresses(area.rsplit('!',1)[-1])]
        sequence=re.search(r'依次为([^。；;|]+)',original) if '成员分别为上述各行名称' in original else None
        same=re.search(r'成员与对应([A-Z]+)列相同',original)
        values={}
        if sequence:
            labels=[value.strip(' `“”') for value in re.split(r'[、，,]',sequence[1])]
            rows=list(dict.fromkeys(re.search(r'\d+$',cell)[0] for cell in cells))
            if len(rows)==len(labels):
                values={cell:{'mode':'members','members':[dict(zip(rows,labels))[re.search(r'\d+$',cell)[0]]]} for cell in cells}
        elif same:
            known={cell:item['value'] for item in member_rows if item['id']==key for cell in item['ranges']}
            for cell in cells:
                source=cell.rsplit('!',1)[0]+'!'+same[1]+re.search(r'\d+$',cell)[0]
                if isinstance(known.get(source),dict) and known[source].get('mode')=='members':values[cell]=deepcopy(known[source])
        if values and len(values)==len(cells):
            member_rows.extend({**row,'ranges':[cell],'value':value} for cell,value in values.items())
        else:member_rows.append(row)
    result['dimensions']=member_rows
    for targets,source,overrides,original in field_references:
        for row in result['dimensions'][:]:
            if row['id'] in overrides:continue
            if any(source in [(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell
                for cell in addresses(area.rsplit('!',1)[-1])] for area in row['ranges'] if area!='*'):
                result['dimensions'].append({**deepcopy(row),'ranges':targets,
                    'original':{'original':row.get('original'),'field_reference':original}})
    for (column,key),(value,original) in column_defaults.items():
        locations=list(dict.fromkeys((area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell
            for row in result['metrics'] if key in {ref['id'] for ref in standard['metric'][row['id']]['dimension_refs']}
            for area in qualify(row['ranges']) if area!='*' for cell in addresses(area.rsplit('!',1)[-1])
            if re.match(r'[A-Z]+',cell)[0]==column))
        if locations:result['dimensions'].append({'id':key,'ranges':locations,'value':value,'original':original})
    if result.get('data_ranges'):
        owned={area.rsplit('!',1)[0]+'!'+cell for area in qualify(result['data_ranges'])
               for cell in addresses(area.rsplit('!',1)[-1])}
        for row in result['metrics']:
            cells={area.rsplit('!',1)[0]+'!'+cell for area in qualify(row['ranges']) if area!='*'
                   for cell in addresses(area.rsplit('!',1)[-1])}
            selected=owned if not row['ranges'] or '*' in row['ranges'] else cells & owned
            declaration=row.get('original','')
            columns={col for col,key in column_metrics(declaration).items() if key==row['id']} if isinstance(declaration,str) else set()
            row['ranges']=sorted(loc for loc in selected if not columns or re.match(r'[A-Z]+',loc.rsplit('!',1)[-1])[0] in columns)
        result['metrics']=[row for row in result['metrics'] if row['ranges']]
    plain_text=text.replace('`','').replace('**','')
    for declaration in plain_text.splitlines():
        label=re.search(r'标签(?:位于|在)(?:合并单元格)?\s*('+coordinate.pattern+r')',declaration)
        rows=re.search(r'适用于第?(\d+)(?:至|到|-)(\d+)行',declaration)
        if not label or not rows:continue
        refs=qualify([label[1]]);source,_,span=refs[0].rpartition('!')
        columns={re.match(r'[A-Z]+',cell)[0] for cell in addresses(span)}
        areas=sorted({source+'!'+cell for row in result['metrics'] for area in qualify(row['ranges'])
            if area!='*' and area.rpartition('!')[0]==source for cell in addresses(area.rsplit('!',1)[-1])
            if re.match(r'[A-Z]+',cell)[0] in columns and int(rows[1])<=int(re.search(r'\d+$',cell)[0])<=int(rows[2])})
        if not areas:continue
        period_links.append((refs,areas))
        date=re.search(r'\bperiod\s*(?:均为|为|[:：=])\s*(?:实际日期\s*)?('+date_pattern+r')',declaration)
        role=re.search(r'template_period_role\s*(?:均为|为|[:：=])\s*(期初|期末|本期|上期|opening|closing|current|prior)',declaration)
        for key,value in [('period',date[1] if date else None),('template_period_role',role[1] if role else None)]:
            if value is not None:result['dimensions'].append({'id':key,'ranges':areas,'value':value,'original':declaration})
    applies=re.search(r'以上标签适用于下述各组成部分的明细行和合计行|适用于(?:下列|以下)全部明细行?和合计行|这两个标签适用于以下明细行及合计行|这些标签适用于明细行和合计行|该列标签适用于本表明细行及合计行|这些期间角色适用于各明细、合计|行\s*\d+\s*至\s*\d+\s*均适用上述表头',plain_text)
    if not applies:applies=re.search(r'适用于(?:第)?(?:\d+行明细|明细行\d+)和(?:第)?(?:\d+行合计|合计行\d+)',plain_text)
    declared_headers=re.findall(r'期间(?:列)?标签在\s*([^。\n]+)',plain_text)
    if applies or declared_headers:
        refs=[ref for clause in declared_headers for ref in qualify(coordinate.findall(clause))];header_roles={}
        for declaration in re.finditer(r'(?:表头(?:原文)?位于|(?:期间|列)标签(?:位于|在)|表头位置|在)\s*((?:'+coordinate.pattern+r'[、，,\s]*)+)',plain_text):
            refs.extend(ref for area in qualify(coordinate.findall(declaration[1])) for ref in
                [(area.rsplit('!',1)[0]+'!' if '!' in area else '')+cell for cell in addresses(area.rsplit('!',1)[-1])])
        for declaration in re.finditer(r'^\|\s*([A-Z]+\d+)\s*\|\s*(期初余额|期末余额|本期增加|本期减少)\s*\|\s*(期初|期末|本期)\s*\|',plain_text,re.M):
            ref=qualify([declaration[1]])[0];refs.append(ref)
            header_roles[ref]={'期初':'opening','期末':'closing','本期':'current'}[declaration[3]]
        for declaration in re.finditer(r'('+coordinate.pattern+r')\s*原文[：:]\s*(本期|上期|期初|期末)(?:发生额|余额)',plain_text):
            ref=qualify([declaration[1]])[0];refs.append(ref)
            header_roles[ref]={'本期':'current','上期':'prior','期初':'opening','期末':'closing'}[declaration[2]]
        target_cells={cell for row in [*result['metrics'],*[r for r in result['dimensions'] if r['id']=='template_period_role']]
            for area in qualify(row['ranges']) if area!='*' for cell in
            [(area.rsplit('!',1)[0]+'!' if '!' in area else '')+c for c in addresses(area.rsplit('!',1)[-1])]}
        for ref in dict.fromkeys(refs):
            source,_,cell=ref.rpartition('!');column=re.match(r'[A-Z]+',cell)[0]
            selected=sorted(loc for loc in target_cells-set(refs) if loc.rpartition('!')[0]==source
                            and re.match(r'[A-Z]+',loc.rsplit('!',1)[-1])[0]==column)
            if selected:
                period_links.append(([ref],selected))
                for row in result['dimensions'][:]:
                    if row['id'] in ('period','template_period_role') and row.get('value') is not None and ref in qualify(row['ranges']):
                        result['dimensions'].append({**deepcopy(row),'ranges':selected,
                            'original':{'original':row.get('original'),'适用声明':applies[0] if applies else declared_headers[0]}})
                if ref in header_roles:
                    result['dimensions'].append({'id':'template_period_role','ranges':selected,
                        'value':header_roles[ref],'original':applies[0] if applies else declared_headers[0]})
                    role={'current':'本期','prior':'上期','opening':'期初','closing':'期末'}[header_roles[ref]]
                    if role in named_periods:result['dimensions'].append({'id':'period','ranges':selected,
                        'value':named_periods[role],'label_references':[ref],'original':applies[0] if applies else declared_headers[0]})
    for declaration in re.finditer(r'(?:以上|这些)([^。\n；;]+?)均为(本期|上期|期初|期末)',plain_text):
        columns=re.findall(r'(?:(%s))?([A-Z]+)列' % '|'.join(re.escape(name) for name in sheet_names),declaration[1])
        locations=list(dict.fromkeys(area.rsplit('!',1)[0]+'!'+cell for row in result['metrics']
            for area in row['ranges'] if '!' in area for cell in addresses(area.rsplit('!',1)[1])))
        selected=[location for location in locations if any(
            (not sheet or location.rsplit('!',1)[0]==sheet) and re.match(r'[A-Z]+',location.rsplit('!',1)[1])[0]==col
            for sheet,col in columns)]
        if selected:result['dimensions'].append({'id':'template_period_role','ranges':selected,
            'value':{'本期':'current','上期':'prior','期初':'opening','期末':'closing'}[declaration[2]],'original':declaration[0]})
    for declared in re.finditer(r'[“"]([^”"]+)[”"]列为\s*template_period_role\s*=\s*(\w+)\s*[、，,]\s*period\s*=\s*(\{[^{}]+\})',plain_text):
        if not re.search(r'这些坐标均在[“"]'+re.escape(declared[1])+r'[”"]列',plain_text):continue
        periods={key:value.strip().strip('\"\'') for key,value in re.findall(r'(\w+)\s*:\s*([^,{}]+)',declared[3])}
        ranges=list(dict.fromkeys(area for row in result['metrics'] for area in row['ranges'] if area!='*'))
        if ranges:
            result['dimensions'].extend({'id':key,'ranges':ranges,'value':value,'original':declared[0]}
                for key,value in [('period',periods),('template_period_role',declared[2])])
    # 只展开同节明示的标准编号、维度名及“分别/依次”配对，不读取原表猜成员。
    for section in re.split(r'(?m)^#{1,6}\s+',plain_text):
        declaration=re.search(r'金标准\s*[:：]\s*(metric\.[\w.]+)\s*[；;]\s*范围维度\s*[:：]\s*(\w+)',section)
        if not declaration or declaration[1] not in standard['metric'] or declaration[2] not in standard['dimension']:continue
        metric,key=declaration.groups()
        stated_sheets={area.rsplit('!',1)[0].strip("'") for area in coordinate.findall(section) if '!' in area}&set(sheet_names)
        default=next(iter(stated_sheets)) if len(stated_sheets)==1 else None
        def qualified(area):
            return area if '!' in area else default+'!'+area if default else None
        entries=[]
        for paragraph in re.split(r'\n\n|\n(?=\s*-\s)',section):
            if re.search(r'各自为\s*members',paragraph):
                for item in re.finditer('('+coordinate.pattern+r')[“"]([^”"]+)[”"]',paragraph):
                    area=qualified(item[1])
                    if area:entries.append((area,{'mode':'members','members':[item[2]]},paragraph))
            span=re.search('('+coordinate.pattern+r')\s*(?:为明细|明细行分别取)',paragraph)
            names=re.search(r'(?:各自取|分别取|依次对应下列)\s*members[^：。\n]*[:：]\s*([^。\n；;]+)',paragraph)
            if span and names:
                area=qualified(span[1])
                values=[v.strip(' “”"') for v in names[1].split('、')]
                quoted=re.findall(r'[“"]([^”"]+)[”"]',names[1])
                if quoted and not re.sub(r'[“"][^”"]+[”"]','',names[1]).strip(' 、，,'):values=quoted
                if area:
                    source,_,ranges=area.rpartition('!');by_row={}
                    for cell in addresses(ranges):by_row.setdefault(re.search(r'\d+$',cell)[0],[]).append(source+'!'+cell)
                    if len(by_row)==len(values):
                        entries.extend((locations,{'mode':'members','members':[value]},paragraph) for locations,value in zip(by_row.values(),values))
            if re.search(r'合计(?:取|为)\s*all|“合计”[:：]\s*all',paragraph):
                for total in re.finditer('('+coordinate.pattern+r')\s*(?:[“"]合计[”"]|为合计|合计取)',paragraph):
                    area=qualified(total[1])
                    if area:entries.append((area,{'mode':'all'},paragraph))
        roles=dict(re.findall(r'([A-Z]+)列(?:是|为)(本期|上期)',section))
        for area,value,original in entries:
            locations=area if isinstance(area,list) else [area]
            result['metrics'].append({'id':metric,'ranges':locations,'original':original})
            result['dimensions'].append({'id':key,'ranges':locations,'value':value,'original':original})
            for span in locations:
                source,_,ranges=span.rpartition('!')
                for cell in addresses(ranges):
                    role=roles.get(re.match(r'[A-Z]+',cell)[0])
                    if not role:continue
                    loc=source+'!'+cell
                    result['dimensions'].append({'id':'template_period_role','ranges':[loc],'value':role,'original':section})
                    if role in named_periods:result['dimensions'].append({'id':'period','ranges':[loc],'value':named_periods[role],'original':section})
    for refs,areas in period_links:
        from .A语义阅读 import addresses
        covered={area.rsplit('!',1)[0]+'!'+cell for area in areas if '!' in area
                 for cell in addresses(area.rsplit('!',1)[-1])}
        linked=set()
        for row in result['dimensions'][:]:
            if row['id']!='period' or '*' in row['ranges']:continue
            cells={area.rsplit('!',1)[0]+'!'+cell for area in qualify(row['ranges']) if '!' in area
                   for cell in addresses(area.rsplit('!',1)[-1])}
            selected=cells&covered
            if not selected:continue
            linked.update(selected)
            if selected!=cells:
                result['dimensions'].append({**deepcopy(row),'ranges':sorted(cells-selected)})
                row['ranges']=sorted(selected)
            row['evidence']=list(dict.fromkeys([*row.get('evidence',[]),*refs]))
            row['label_references']=list(dict.fromkeys([*row.get('label_references',[]),*refs]))
        if covered-linked:
            result['dimensions'].append({'id':'period','ranges':sorted(covered-linked),'evidence':refs,'label_references':refs})
    for declared in re.finditer(r'((?:[a-z_]+\s*[、，,]\s*)+[a-z_]+)\s*均(?:为)?未知',text.replace('`','')):
        for key in re.findall(r'[a-z_]+',declared[1]):
            if key in standard['dimension']:
                result['dimensions'].append({'id':key,'ranges':['*'],'value':'未知','original':declared[0]})
    common_dates={}
    for declaration in re.finditer(r'期间角色与日期[:：]([^。\n]+)',plain_text):
        for value in re.finditer(r'(本期|上期)\s*[=＝]\s*\d{4}年度[（(]('+date_pattern+r')[）)]',declaration[1]):
            common_dates[value[1]]=re.sub(r'\s+','',value[2])
    label_roles={qualify([m[1]])[0]:m[2] for m in re.finditer('('+coordinate.pattern+r')[“"](本期|上期)(?:发生额|余额)[”"]',plain_text)}
    for row in result['dimensions'][:]:
        if row['id']!='period':continue
        roles={label_roles[ref] for ref in row.get('label_references',[]) if ref in label_roles}
        if len(roles)==1:
            role=next(iter(roles))
            result['dimensions'].append({'id':'template_period_role','ranges':row['ranges'],'value':role,'original':row.get('original',plain_text)})
            if role in common_dates:row['value']=common_dates[role]
    result['dimensions']=[row for row in result['dimensions'] if not (isinstance(row.get('value'),str)
        and row['value'].startswith('指标均为 ') and '取值' in str(row.get('original','')))]
    # 只接入答复明说的名称位置，不根据原表文字或空白排除金额。
    for declared in re.finditer(r'([A-Z]+)列是项目名称',text):
        locations={area.rsplit('!',1)[0]+'!'+cell for row in result['metrics']
                   for area in qualify(row['ranges']) if area!='*' for cell in addresses(area.rsplit('!',1)[-1])
                   if re.match(r'[A-Z]+',cell)[0]==declared[1]}
        result.setdefault('non_business',[]).extend({'range':loc,'category':'label',
            'reason':'原答复明确此列是项目名称','original':declared[0]} for loc in sorted(locations))
    for declared in re.finditer('('+coordinate.pattern+r')为对应的行项目名称[，,]不是金额格',text.replace('`','')):
        result.setdefault('non_business',[]).append({'range':qualify([declared[1]])[0],
            'category':'label','reason':'原答复明确是行项目名称，不是金额格','original':declared[0]})
    if len({r['id'] for r in result['metrics']})==1:
        for row in result['metrics']:
            if not row['ranges']:row['ranges']=['*']
    if table_period_rows:
        explicit={key:{area.rsplit('!',1)[0]+'!'+cell for row in table_period_rows if row['id']==key
            for area in row['ranges'] for cell in addresses(area.rsplit('!',1)[-1])} for key in ('period','template_period_role')}
        retained=[]
        for row in result['dimensions']:
            if row['id'] not in explicit or '*' in row['ranges']:retained.append(row);continue
            locations={area.rsplit('!',1)[0]+'!'+cell for area in qualify(row['ranges']) for cell in addresses(area.rsplit('!',1)[-1])}
            remaining=locations-explicit[row['id']]
            if remaining:retained.append({**row,'ranges':sorted(remaining)})
        result['dimensions']=[*retained,*table_period_rows]
    for row in result['dimensions'][:]:
        if row['id'] in shared_field_aliases:
            keys,declaration=shared_field_aliases[row['id']]
            result['dimensions'].extend({**deepcopy(row),'id':key,
                'original':{'original':row.get('original'),'field_binding':declaration}} for key in keys)
    # 前文按行、后文按列声明时，只用同一答复明列的指标坐标接回行字段。
    for owner,rows,key,value,original in pending_row_fields:
        locations=sorted({(area.rsplit('!',1)[0] if '!' in area else owner)+'!'+cell
            for metric in result['metrics'] if key in {ref['id'] for ref in standard['metric'].get(metric['id'],{}).get('dimension_refs',[])}
            for area in metric['ranges'] if area!='*' and ('!' not in area or area.rsplit('!',1)[0]==owner)
            for cell in addresses(area.rsplit('!',1)[-1]) if re.search(r'\d+$',cell)[0] in rows})
        if locations:result['dimensions'].append({'id':key,'ranges':locations,'value':value,'original':original})
    _require(any(result.values()),'答复未明确可机械读取的坐标、标准编号或维度归属；原文保留')
    return result


def direct_answer(text, standard, sheet_names=(), metric_context=None, default_sheet=None, field_scope=None):
    """按答复已有的字段键和坐标展开，不从原表或文字推断取值。"""
    if isinstance(text,str):
        text=text.strip()
        fence=re.fullmatch(r'```(?:json)?\s*\n(.*)\n```',text,re.S)
        if fence:text=fence[1]
        try:answer,end=json.JSONDecoder().raw_decode(text)
        except ValueError:return _plain_answer(text,standard,sheet_names,metric_context,default_sheet,field_scope)
        _require(text[end:].strip() in ('','喵'),'答复正文后有其他内容；原文保留')
    else:answer=deepcopy(text)
    if isinstance(answer,list):answer={'results':answer}
    _require(isinstance(answer,dict),'答复不是业务字段对象')
    if 'tables' in answer:
        return {**answer,'tables':[direct_answer(t,standard) for t in answer['tables']]}
    result={**answer,'metrics':[],'dimension_bindings':[],'dimensions':[],'unresolved':[]}
    keys=set(standard['dimension'])|{'metric_id','dimension_id','template_period_role'}
    def ranges(value):
        parts=[p.strip() for x in (value if isinstance(value,list) else [value])
               if isinstance(x,str) for p in re.split(r'[,，、;；]',x)]
        return parts if parts and all(p=='*' or re.fullmatch(r"(?:[^!]+!)?\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?",p) for p in parts) else None
    def emit(key,value,scope,original,normalized=False):
        # 已解析行保留原答复出处，回读不再把整行套入新的 original。
        origin=original.get('original',original) if normalized and isinstance(original,dict) else original
        row={'id':key,'ranges':scope,'value':deepcopy(value),'original':deepcopy(origin)}
        if isinstance(original,dict):
            for name in ('evidence','label_references','label_reference_original','template_period_role'):
                if name in original:row[name]=deepcopy(original[name])
        if normalized and isinstance(original,dict) and original.get('逐格原字段') is True:
            row['逐格原字段']=True
        if value is None or value in ('','保留空白','原空白保留，不填'):
            row['reason']='原答复未提供此字段实际取值';result['unresolved'].append(row);return
        if key=='dimension_id':
            if isinstance(value,str) and value in standard['dimension']:
                row.pop('value');row['id']=value;result['dimension_bindings'].append(row)
            else:
                row['reason']='原答复未给出确定的现有维度标准编号';result['unresolved'].append(row)
        elif key=='metric_id':
            metric=value
            if isinstance(value,dict):
                metric=next((value[k] for k in ('selected_standard','standard','id','value') if k in value),None)
                for k in ('actual_dimensions','dimensions'):
                    if k in value:walk(value[k],scope)
            if isinstance(metric,str) and metric in standard['metric']:
                row.pop('value');row['id']=metric;result['metrics'].append(row)
            else:
                row['reason']='原答复未给出确定的现有标准编号';result['unresolved'].append(row)
        else:result['dimensions'].append(row)
    def field(key,value,scope):
        location_keys=('ranges','range','coordinates','coordinate','范围','坐标')
        if isinstance(value,list) and all(isinstance(x,dict) and any(k in x for k in location_keys) for x in value):
            for row in value:field(key,row,scope)
        elif isinstance(value,dict) and any(k in value for k in location_keys):
            scope=ranges(next(value[k] for k in location_keys if k in value))
            if 'values_by_row' in value:
                from .A语义阅读 import addresses
                cells=[a for area in scope for a in addresses(area.rsplit('!',1)[-1].replace('$',''))]
                for numbers,val in value['values_by_row'].items():
                    rows=set(re.split(r'[/,，、]',numbers))
                    selected=[a for a in cells if re.search(r'\d+$',a)[0] in rows]
                    if selected:emit(key,val,selected,value)
            else:
                emit(key,value.get('value',value.get('取值')),scope,value)
                if key=='metric_id':walk({k:v for k,v in value.items() if k in keys and k!=key},scope)
        elif isinstance(value,dict) and any(ranges(k) for k in value):
            for loc,val in value.items():
                if ranges(loc):emit(key,val,ranges(loc),value)
        elif key=='period' and isinstance(value,dict) and value and all(ranges(v) for v in value.values()):
            for val,loc in value.items():emit(key,val,ranges(loc),value)
        else:emit(key,value,scope,value)
    def walk(obj,scope):
        if isinstance(obj,list):
            for row in obj:walk(row,scope)
            return
        if not isinstance(obj,dict):return
        for name in ('ranges','range','coordinates','coordinate','范围','坐标'):
            local=ranges(obj.get(name,[]))
            if local:scope=local;break
        if obj.get('field') in keys and 'value' in obj:emit(obj['field'],obj['value'],scope,obj)
        if isinstance(obj.get('standard'),str) and obj['standard'] in standard['metric']:emit('metric_id',obj['standard'],scope,obj)
        for key,value in obj.items():
            if key in keys:field(key,value,scope)
            elif ranges(key):
                if isinstance(value,dict):walk(value,ranges(key))
                elif value is None:result['unresolved'].append({'ranges':ranges(key),'reason':'原答复此范围为空','original':None})
            else:
                pair=re.fullmatch(r'((?:[^!]+!)?[A-Z]+\d+)\s+(\w+)',key)
                if pair and pair[2] in keys:emit(pair[2],value,[pair[1]],obj)
                elif key in {'results','坐标识别','识别结果','fields','待纠正字段','实际维度取值'} or key=='坐标' and isinstance(value,list):walk(value,scope)
                elif key not in {'original','evidence','标准依据','说明','corrections'} and isinstance(value,list):walk(value,scope)
                elif key not in {'original','evidence','标准依据','说明'} and isinstance(value,dict) and any(
                        ranges(k) or k in {'coordinate','coordinates','range','ranges','坐标','field'} or isinstance(v,(dict,list)) for k,v in value.items()):walk(value,scope)
    for kind in ('metrics','dimension_bindings','dimensions','unresolved'):
        rows=answer.get(kind,[])
        if isinstance(rows,dict):walk(rows,['*']);continue
        for original in rows:
            row=deepcopy(original)
            row['ranges']=ranges(row.get('ranges',row.get('range',[]))) or row.get('ranges',[row['range']] if row.get('range') is not None else [])
            if kind=='dimensions' and 'value' not in row and (row.get('label_references') is not None or
                    row.get('id')=='period' and row.get('evidence')):
                result[kind].append(row);continue
            if kind=='unresolved':result[kind].append(row)
            else:emit({'metrics':'metric_id','dimension_bindings':'dimension_id'}.get(kind,row.get('id',row.get('field'))),row.get('id') if kind in ('metrics','dimension_bindings') else row.get('value'),row['ranges'],row,normalized=True)
    walk({k:v for k,v in answer.items() if k not in ('metrics','dimension_bindings','dimensions','unresolved')},['*'])
    for row in answer.get('corrections',[]):
        scope=ranges(row.get('ranges',row.get('range')))
        if 'field' in row or answer.get('field') in keys:field(row.get('field',answer.get('field')),row,scope)
        elif 'values' in row:walk(row['values'],scope)
    return result


def merge_direct_answer(old, reply, field_scope, withdrawn=None):
    """按原表已登记位置分配字段范围，复用同一字段合并逻辑。"""
    from .A语义阅读 import addresses
    if old.get('carrier')!='B' or old.get('逐格字段版本')==1:
        tables=deepcopy(old['outline']['subtables']);sheets={t['sheet'] for t in tables}
        if old.get('逐格原样保存'):
            from openpyxl.utils.cell import range_boundaries, get_column_letter
            snapshot={s['name']:s for s in old['机械记录']['sheets']}
            owned={t['sheet']+'!'+cell for t in tables for area in t['data_ranges'] for cell in addresses(area)}
            fresh={}
            for location in field_scope:
                sheet,sep,area=location.rpartition('!')
                _require(sep or len(sheets)==1,'多工作表字段范围须写明工作表')
                sheet=sheet.strip("'") if sep else next(iter(sheets));area=(area if sep else location).replace('$','')
                _require(sheet in snapshot,'字段范围引用未知工作表')
                source=snapshot[sheet];bounds=range_boundaries(area)
                _require(bounds[2]<=source['max_column'] and bounds[3]<=source['max_row'],'字段范围越出原表')
                fresh.setdefault(sheet,set()).update(cell for cell in addresses(area) if sheet+'!'+cell not in owned)
            # 新位置只登记坐标；指标和维度仍由该格自己的原答复提供。
            for sheet,cells in fresh.items():
                if cells:
                    source=snapshot[sheet]
                    tables.append({'id':'定向新增位置_'+sheet+'_'+definition_hash(sorted(cells))[:12],'sheet':sheet,
                        'range':f'A1:{get_column_letter(source["max_column"])}{source["max_row"]}',
                        'data_ranges':sorted(cells),'metrics':[],'dimensions':[],'unresolved':[]})
        def local_scope(scope,table):
            owned={cell for area in table['data_ranges'] for cell in addresses(area)}
            result={}
            for location,keys in (scope or {}).items():
                sheet,sep,area=location.rpartition('!')
                _require(sep or len(sheets)==1,'多工作表字段范围须写明工作表')
                if sep and sheet.strip("'")!=table['sheet']:continue
                for cell in addresses((area if sep else location).replace('$','')):
                    if cell in owned:result.setdefault(cell,[]).extend(keys)
            return result
        result=[];outside=[];replaced={}
        for table in tables:
            mask=local_scope(field_scope,table);removed=local_scope(withdrawn,table)
            if not mask and not removed:
                result.append(deepcopy(table));continue
            merged,unused,changed=_merge_direct_table(
                {**old,'outline':{'subtables':[table]}},reply,mask,removed)
            result.append({**deepcopy(table),**merged})
            outside.extend({'子表':table['id'],**item} for item in unused)
            for key,cells in changed.items():
                replaced.setdefault(key,[]).extend((table['sheet']+'!' if len(tables)>1 else '')+c for c in cells)
        return ({'tables':result} if len(tables)>1 else result[0]),outside,replaced
    snapshot=old['机械记录'];names=[s['name'] for s in snapshot['sheets']]
    source_sheets=list(dict.fromkeys(area.rsplit('!',1)[0] for area in field_scope if '!' in area))
    if not source_sheets:source_sheets=list(dict.fromkeys(r['sheet'] for r in old.get('b_source_ranges',[])))
    if isinstance(reply,str):reply=_own_answer_text(reply,{**old,'b_source_ranges':[{'sheet':s} for s in source_sheets]},snapshot)
    context={location:row['metric_id'] for location,row in _records(old) if row.get('metric_id')}
    delta=direct_answer(reply,_standard(old),names,context,source_sheets[0] if len(source_sheets)==1 else None,field_scope)
    if isinstance(reply,str) and (declared:=re.search(r'适用范围[:：]\s*([^\n]+)',reply)):
        pattern='('+ '|'.join(re.escape(name) for name in sorted(names,key=len,reverse=True))+r')!([A-Z]+\d+(?::[A-Z]+\d+)?)'
        scopes=[sheet+'!'+area for sheet,area in re.findall(pattern,declared[1])]
        allowed={span.rsplit('!',1)[0]+'!'+cell for span in scopes for cell in addresses(span.rsplit('!',1)[-1])}
        date=r'\d{4}-\d{2}-\d{2}(?:至\d{4}-\d{2}-\d{2})?'
        rules=dict(re.findall(r'(期初|期末|本期|上期)(?:时点|发生额或变动额|发生额|变动额)为\s*('+date+r')',reply))
        aliases={'opening':'期初','closing':'期末','current':'本期','prior':'上期','期初余额':'期初','期末余额':'期末','本期增加':'本期','本期减少':'本期'}
        for row in delta['dimensions'][:]:
            if row.get('ranges')!=['*'] or not scopes:continue
            row['ranges']=scopes[:]
            if row['id']=='period' and rules:delta['dimensions'].remove(row)
            elif row['id'] in ('currency','report_scope') and isinstance(row.get('value'),str):
                literal=re.match(r'([A-Za-z_]+)(?=[（(，,\s]|$)',row['value'])
                if literal:row['value']=literal[1]
        if rules and allowed:
            for table in old['outline']['subtables']:
                for row in table['metrics']+table['dimensions']:
                    role=row.get('template_period_role') if 'template_period_role' in row else row.get('value') if row.get('id')=='template_period_role' else None
                    if not isinstance(role,str):continue
                    value=rules.get(aliases.get(role,role))
                    if value is None:continue
                    spans=table['data_ranges'] if row['ranges']==['*'] else row['ranges']
                    selected=[table['sheet']+'!'+cell for span in spans for cell in addresses(span.rsplit('!',1)[-1])
                              if table['sheet']+'!'+cell in allowed]
                    if selected:delta['dimensions'].append({'id':'period','ranges':selected,'value':value,'original':reply})
    fresh=_answer_tables(delta,{**old,'答复指纹':definition_hash(reply),'field_scope':field_scope},snapshot)
    def scoped_rows(table,key):
        rows=deepcopy(table.get(key,[]))
        for row in rows:
            # 合并同表子表前固定星号原范围，不能把局部共同声明扩成整张表的身份。
            if '*' in row.get('ranges',[]):
                row['ranges']=[area for span in row['ranges'] for area in
                               (table['data_ranges'] if span=='*' else [span])]
        return rows
    tables={}
    for table in old['outline']['subtables']:
        sheet=table['sheet']
        if sheet not in tables:tables[sheet]=deepcopy(table)
        else:
            for key in ('metrics','dimension_bindings','dimensions','unresolved'):
                tables[sheet][key]=scoped_rows(tables[sheet],key)+scoped_rows(table,key)
            tables[sheet]['data_ranges'].extend(deepcopy(table['data_ranges']))
    allowed={area['sheet']+'!'+cell for area in old.get('b_source_ranges',[])
             for cell in addresses(area['range'])}
    allowed.update(sheet+'!'+cell for sheet,table in tables.items()
                   for area in table['data_ranges'] for cell in addresses(area))
    def qualified(scope):
        result={}
        for location,keys in (scope or {}).items():
            if '!' not in location:
                _require(len(tables)==1,'多表字段范围须写明工作表')
                location=next(iter(tables))+'!'+location
            sheet,area=location.rsplit('!',1)
            for cell in addresses(area):
                if sheet+'!'+cell in allowed:result[sheet+'!'+cell]=keys
        return result
    mask=qualified(field_scope);removed=qualified(withdrawn)
    incoming={}
    for table in fresh:
        sheet=table['sheet'];group=incoming.setdefault(sheet,{k:[] for k in ('metrics','dimension_bindings','dimensions','unresolved')})
        for key in group:group[key].extend(scoped_rows(table,key))
        permitted={cell for area in table['data_ranges'] for cell in addresses(area)
                   if sheet+'!'+cell in mask}
        if not permitted:continue
        if sheet not in tables:
            tables[sheet]={**deepcopy(table),'data_ranges':[],'metrics':[],'dimensions':[],'unresolved':[]}
        owned={cell for area in tables[sheet]['data_ranges'] for cell in addresses(area)}
        tables[sheet]['data_ranges']=sorted(owned|permitted)
        from openpyxl.utils.cell import range_boundaries, get_column_letter
        bounds=[range_boundaries(tables[sheet]['range']),*[range_boundaries(cell) for cell in permitted]]
        tables[sheet]['range']=f'{get_column_letter(min(v[0] for v in bounds))}{min(v[1] for v in bounds)}:{get_column_letter(max(v[2] for v in bounds))}{max(v[3] for v in bounds)}'
    result=[];outside=[];replaced={}
    for sheet,table in tables.items():
        local_mask={loc.rsplit('!',1)[1]:keys for loc,keys in mask.items() if loc.rsplit('!',1)[0]==sheet}
        local_removed={loc.rsplit('!',1)[1]:keys for loc,keys in removed.items() if loc.rsplit('!',1)[0]==sheet}
        if not local_mask and not local_removed:
            result.append(table);continue
        merged,unused,changed=_merge_direct_table(
            {**old,'outline':{'subtables':[table]}},
            incoming.get(sheet,{'metrics':[],'dimensions':[],'unresolved':[]}),local_mask,local_removed)
        result.append({**table,**merged})
        outside.extend({'工作表':sheet,**item} for item in unused)
        for key,cells in changed.items():replaced.setdefault(key,[]).extend(sheet+'!'+cell for cell in cells)
    return {'tables':result},outside,replaced


def _merge_direct_table(old, reply, field_scope, withdrawn=None):
    """按明确纠正范围替换声明；只依现有标准引用接通编号与维度。"""
    from collections import defaultdict
    from .A语义阅读 import addresses
    standard=_standard(old);table=deepcopy(old['outline']['subtables'][0]);sheet=table['sheet']
    owned={a for span in table['data_ranges'] for a in addresses(span)}
    mask={a:set(v) for a,v in field_scope.items()};withdrawn=withdrawn or {}
    range_cells={}
    def cells(row):
        spans=tuple(row.get('ranges',[row['range']] if 'range' in row else []))
        if spans in range_cells:return range_cells[spans]
        result=set()
        for span in spans:
            if span=='*':result.update(owned);continue
            sh,sep,area=span.rpartition('!')
            _require(not sep or sh.strip("'")==sheet,'答复引用了其他工作表')
            result.update(addresses((area if sep else span).replace('$','')))
        range_cells[spans]=frozenset(result)
        return range_cells[spans]
    def field(row,kind):
        if kind=='metrics':return 'metric_id'
        if kind=='dimension_bindings':return 'dimension_id'
        if row.get('reason')=='原答复明确不采用此槽位':return 'metric_id'
        key=next((row[k] for k in ('id','field','dimension','metric') if row.get(k)),None)
        return 'metric_id' if key in {'metric','metric_id'} or str(key).startswith('metric.') else key
    context={location:row['metric_id'] for location,row in _records(old) if row.get('metric_id')}
    delta=direct_answer(reply,standard,[sheet],context,field_scope=field_scope)
    if 'tables' in delta:
        _require(len(delta['tables'])==1,'纠正答复包含多个表区');delta=delta['tables'][0]
    for row in delta['unresolved']:
        original=row
        while isinstance(original,dict):original=original.get('original')
        if not isinstance(original,str) or not re.search(r'无法确认[^。]*期间角色[^。]*范围[^。]*保留未知',original):continue
        # 明确说范围及期间角色未知时，仅撤下本次询问的这两类字段。
        for address in cells(row)&owned:
            for key in mask.get(address,set()):
                if key=='template_period_role' or standard['dimension'].get(key,{}).get('value_type')=='selection':
                    delta['dimensions'].append({'id':key,'ranges':[address],'value':'未知','original':original})
    previous_records={loc.rsplit('!',1)[-1]:row for loc,row in _records(old) if loc.rsplit('!',1)[0]==sheet}
    previous={a:row.get('dimensions',{}) for a,row in previous_records.items()}
    # 编号未决的格仍有原维度声明，重新选编号不能覆盖这些已有取值。
    for a in owned-previous.keys():
        rows=[r for r in table['dimensions'] if 'value' in r and a in cells(r)]
        previous[a]={key:deepcopy(values[0]) for key in {r['id'] for r in rows}
            if (values:=[r['value'] for r in rows if r['id']==key]) and all(v==values[0] for v in values)}
    reference_outside=[]
    for row in delta['dimensions'][:]:
        if 'label_references' not in row:continue
        if 'value' not in row:delta['dimensions'].remove(row)
        for a in cells(row)&owned:
            if 'label_references' not in mask.get(a,set()):
                if 'value' not in row:reference_outside.append({'条目':row,'未采用位置':[a],'原因':'不在本次标签引用范围'})
                continue
            prior=[r for r in table['dimensions'] if r['id']==row['id'] and
                   r.get('逐格原字段') is True and 'value' in r and a in cells(r)]
            if not prior:
                reference_outside.append({'条目':row,'未采用位置':[a],'原因':'本格尚无直接保存的该维度；引用不能补成身份'})
                continue
            for saved in prior:
                saved['label_references']=deepcopy(row['label_references'])
                saved['label_reference_original']=deepcopy(row['original'])
    for row in delta['dimensions'][:]:
        if row.get('id')!='period' or 'value' in row or not row.get('evidence'):continue
        delta['dimensions'].remove(row)
        for a in cells(row)&owned:
            if 'period' not in mask.get(a,set()) or 'period' not in previous.get(a,{}):continue
            delta['dimensions'].append({**row,'ranges':[a],'value':deepcopy(previous[a]['period'])})
    known={a:{k for k,v in dims.items() if not isinstance(v,dict) or v.get('mode')!='unknown'}
           for a,dims in previous.items()}
    required=defaultdict(set);changed_metrics=set();incoming=defaultdict(list);replaced=defaultdict(set);outside=reference_outside
    for row in delta['metrics']:
        for a in cells(row)&owned:
            if 'metric_id' in mask.get(a,set()):
                required[a].update(r['id'] for r in standard['metric'][row['id']]['dimension_refs'])
                if previous_records.get(a,{}).get('metric_id')!=row['id']:changed_metrics.add(a)
    for kind in ('metrics','dimension_bindings','dimensions','unresolved'):
        for original in delta.get(kind,[]):
            row=deepcopy(original);key=field(row,kind)
            try:positions=cells(row)
            except (ValueError,TypeError) as error:
                outside.append({'条目':row,'原因':str(error)});continue
            allowed={a for a in positions&owned if key in mask.get(a,set()) or
                     not old.get('逐格原样保存') and a in changed_metrics and key!='report_scope' and key in required[a] and key not in known.get(a,set())}
            if kind=='metrics' and 'template_period_role' in row:
                for a in (positions&owned)-allowed:
                    if 'template_period_role' not in mask.get(a,set()):continue
                    matched=[m for m in table['metrics'] if m['id']==row['id'] and a in cells(m)]
                    if matched:
                        incoming['metrics'].append({**deepcopy(matched[0]),'ranges':[a],'template_period_role':deepcopy(row['template_period_role'])})
                        replaced['metric_id'].add(a)
            if positions-allowed:outside.append({'条目':row,'未采用位置':sorted(positions-allowed),'原因':'不在本次字段范围'})
            if not allowed:continue
            row.pop('range',None);row['ranges']=sorted(allowed)
            if kind=='dimensions' and key=='template_period_role':
                incoming['dimensions'].append(row)
                replaced[key].update(allowed)
                if old.get('逐格原样保存'):
                    for address in allowed:
                        if 'metric_id' in mask.get(address,set()):continue
                        for metric in table['metrics']:
                            if address in cells(metric):
                                incoming['metrics'].append({**deepcopy(metric),'ranges':[address],'template_period_role':deepcopy(row['value'])})
                                replaced['metric_id'].add(address)
                continue
            if kind=='dimensions' and not old.get('逐格原样保存'):
                if row.get('value') in ('未知','unknown','未决') and standard['dimension'][key]['value_type']!='selection':
                    incoming['unresolved'].append({**row,'reason':'原答复明确此字段未知'});replaced[key].update(allowed);continue
                try:
                    row=_direct_table({**table,'dimensions':[row]},old,standard)['dimensions'][0]
                except (ValueError,KeyError,TypeError,IndexError) as error:
                    incoming['unresolved'].append({'id':key,'ranges':sorted(allowed),'reason':str(error),'original':row});continue
            if kind=='dimensions':
                for a in allowed:
                    copied=deepcopy(row);copied['ranges']=[a];prior=previous.get(a,{}).get(key,{})
                    saved=previous_records.get(a,{})
                    prior_direct=any(r.get('逐格原字段') is True and r['id']==key and a in cells(r)
                                     for r in table['dimensions'])
                    if (not old.get('逐格字段版本') or prior_direct) and key in saved.get('dimensions',{}) and row['value']==saved['dimensions'][key]:
                        if not row.get('evidence'):
                            copied['evidence']=deepcopy(saved.get('dimension_evidence',{}).get(key,[]))
                        if 'label_references' not in row and key in saved.get('label_references',{}):
                            copied['label_references']=deepcopy(saved['label_references'][key])
                    if not old.get('逐格原样保存') and isinstance(row['value'],dict) and row['value'].get('mode')=='unknown' and isinstance(prior,dict) and prior.get('mode')=='unknown':
                        for part in ('known_scope','known','unknown'):
                            if part in prior:copied['value'].setdefault(part,deepcopy(prior[part]))
                    incoming[kind].append(copied)
            elif kind in ('metrics','dimension_bindings'):
                other='dimension_id' if kind=='metrics' else 'metric_id'
                replaced[other].update(a for a in allowed if other in mask.get(a,set()))
                for a in allowed:
                    copied={**row,'ranges':[a]}
                    if 'template_period_role' not in mask.get(a,set()):
                        copied.pop('template_period_role',None)
                    prior=[m for m in table.get(kind,[]) if a in cells(m)]
                    roles=[m.get('template_period_role') for m in prior]
                    if ('template_period_role' not in copied and
                        'template_period_role' not in mask.get(a,set())|set(withdrawn.get(a,[])) and
                        prior and all(m['id']==row['id'] for m in prior) and
                        roles[0] in ('current','prior','opening','closing','本期','上期','期初','期末',
                                     '本期发生额','上期发生额','期初余额','期末余额') and all(v==roles[0] for v in roles)):
                        copied['template_period_role']=deepcopy(roles[0])
                    incoming[kind].append(copied)
            else:incoming[kind].append(row)
            if key and (kind!='unresolved' or 'value' in row and (row['value'] is None or
                    key=='metric_id' and row['value'] in ('未知','未决','unknown'))):
                replaced[key].update(allowed)
    removed=deepcopy(replaced)
    for a,keys in withdrawn.items():
        for key in keys:removed[key].add(a)
    # 新答复重述编号及全部所需字段时，旧的整句解析失败不再阻断新身份。
    restated={a for a,keys in required.items() if keys and
              keys<={key for key,positions in replaced.items() if a in positions}}
    referenced={};unparsed=defaultdict(set)
    for row in incoming['unresolved']:
        if isinstance(row.get('original'),str):unparsed[row['original']].update(cells(row))
    for row in incoming['metrics']:
        source=row.get('original',{})
        if isinstance(source,dict) and source.get('binding_reference') and isinstance(source.get('original'),str):
            referenced.setdefault(source['original'],set()).update(cells(row)-unparsed[source['original']])
    def statement(row):
        value=row
        while isinstance(value,dict):value=value.get('original')
        return value if isinstance(value,str) else None
    rebound={(field(row,kind),statement(row)) for kind,rows in incoming.items() for row in rows if statement(row)}
    rebound.update(('metric_id',source) for source in delta.get('unscoped_metric_explanations',[]))
    combined={}
    for kind in ('metrics','dimension_bindings','dimensions','unresolved'):
        combined[kind]=[]
        for original in table.get(kind,[]):
            key=field(original,kind)
            try:
                positions=cells(original)
                superseded=removed.get(key,set())
                # 同一原句重新接入时，撤下它在本次字段范围内的旧落点，不动其他答复或范围外内容。
                if (key,statement(original)) in rebound:
                    rebound_positions={a for a in positions if key in mask.get(a,set())}
                    superseded=superseded|rebound_positions
                    replaced[key].update(rebound_positions)
                if kind=='unresolved' and key is None and original.get('reason')=='这段答复未明确可机械读取的字段归属':
                    superseded=superseded|restated
                    if isinstance(original.get('original'),str):superseded|=referenced.get(original['original'],set())
                if not positions&superseded:
                    combined[kind].append(deepcopy(original));continue
                remaining=positions-superseded
            except (ValueError,TypeError):combined[kind].append(deepcopy(original));continue
            if remaining:combined[kind].append({**deepcopy(original),'ranges':sorted(remaining)})
        combined[kind].extend(incoming[kind])
    combined['non_business']=deepcopy(table.get('non_business',[]))
    for row in delta.get('non_business',[]):
        for address in sorted(cells(row)&owned):
            if 'metric_id' in mask.get(address,set()):
                combined['non_business'].append({**deepcopy(row),'range':address})
    return combined,outside,{key:sorted(value) for key,value in replaced.items()}


def _save_direct_relations(mapping, reply, replaced=None):
    """保存答复明示的同范围关系；两端取值不改，不按措辞猜同义。"""
    confirmed=mapping.setdefault('已确认关系',[])
    gaps=mapping.setdefault('关系缺口',[])
    reference=mapping.get('A需要',{})
    def literal(value):
        if isinstance(value,dict) and value.get('mode')=='predicate':
            return {**value,'predicate':value['predicate'].removesuffix('。')}
        return value
    for table in mapping['outline']['subtables']:
        for row in table.get('dimensions',[]):
            key=row['id']
            scope=[table['sheet']+'!'+area for area in row['ranges']]
            if replaced is not None:
                scope=[loc for loc in replaced.get(key,[]) if loc.rsplit('!',1)[0]==table['sheet']
                       and _inside(loc.rsplit('!',1)[1],row['ranges'],table)]
                if not scope:continue
            original=row.get('original')
            while isinstance(original,dict):original=original.get('original')
            if not isinstance(original,str) or original not in reply:continue
            # 仅承接肯定句；否定、待确认和相似描述不形成相等关系。
            if not re.search(r'(?:^|[。；;|])\s*(?:这)?与\s*A(?:所需|需要)(?:的[^。；;|]*?)?范围相同(?:[。；;|]|$)',original):continue
            if key not in _standard(mapping)['dimension'] or 'value' not in row:continue
            declared=set()
            for item in table.get('dimensions',[]):
                source=item.get('original')
                while isinstance(source,dict):source=source.get('original')
                if source==original:declared.add(item['id'])
            needs=[item for item in reference.get('需要',[]) if key in item['dimensions']]
            values={definition_hash(item['dimensions'][key]) for item in needs}
            if len(values)>1:
                # 已有肯定关系中，逐字相同的条件只忽略句末句号；不改原取值或比较其他语义。
                needs=[item for item in needs if literal(item['dimensions'][key])==literal(row['value'])]
                values={definition_hash(item['dimensions'][key]) for item in needs}
            if len(values)!=1 or declared!={key}:
                gap={'dimension':key,'record_ranges':scope,'content':original,
                     'reason':'原答复已说明与A范围相同，但本次A需要、所指维度或取值未唯一确定'}
                if gap not in gaps:gaps.append(gap)
                continue
            relation={'dimension':key,'record_ranges':scope,'value':deepcopy(row['value']),
                      'A需要':{'source_hash':reference['source_hash'],'gold_hash':reference['gold_hash'],
                               'locations':[item['location'] for item in needs],
                               'value':deepcopy(needs[0]['dimensions'][key])},
                      'content':original,'answer_hash':definition_hash(reply)}
            if relation not in confirmed:confirmed.append(relation)
            gaps[:]=[gap for gap in gaps if not (gap.get('dimension')==key and gap.get('record_ranges')==scope)]


def _list_relation_conflicts(mapping, reply, source_ranges):
    """保存答复明说的引用冲突，不用公式重新裁定模型结论。"""
    book={sheet['name']:sheet['cells'] for sheet in mapping['机械记录']['sheets']}
    mentioned=[name for name in book if name+'!' in reply]
    source_sheets={row['sheet'] for row in source_ranges}
    conflicts=[];unused=[];conflict_section=False
    text=reply.replace('`','').replace('**','')
    for line in text.splitlines():
        if line.lstrip().startswith('#'):
            conflict_section=any(word in line for word in ('不一致','差异','冲突'))
        if not conflict_section and '但公式' not in line and '不一致' not in line:continue
        if '标签' not in line or not ('不一致' in line or '标签来自' in line or '标签分别来自' in line or '但公式' in line):continue
        # 工作表必须由原答复明确写出；不按业务名称猜来源表。
        direct=[name for name in mentioned if re.search(re.escape(name)+r'![A-Z]+\d+',line)]
        candidates=[name for name in source_sheets if name in book]
        sheet=direct[0] if direct and '但公式' in line else candidates[0] if len(candidates)==1 else None
        if sheet is None:continue
        if '但公式' in line:
            prefix=line.split('标签')[0]
            refs=[sheet+'!'+cell for cell in re.findall(r'\b([A-Z]+\d+)\b',prefix)]
            if refs:conflicts.append({'locations':refs,'content':line.strip()})
            continue
        for clause in re.split(r'[；;]',line):
            claim=re.search(r'((?:[A-Z]+\d+\s*[、,，]?\s*)+)\s*(?:均)?引用\s*(?:[^\s!，;；]+!)?([A-Z]+\d+)',clause)
            if not claim:continue
            refs=re.findall(r'[A-Z]+\d+',claim[1])
            for cell in refs:
                conflicts.append({'locations':[sheet+'!'+cell],'content':clause.strip()})
    return conflicts,unused


def _declared_segmented_list_tables(a, reply, dimension, b=None):
    """转存答复明确的分段行界和名单标签；不从标签列补金额身份。"""
    from openpyxl.utils.cell import range_boundaries, get_column_letter
    text=reply.replace('`','').replace('**','')
    # 来源矩阵及插行后的分段行界都已由模型列明；只展开文字中的顺序和引用。
    expansion=re.search(r'每段由原来的\s*(\d+)\s*行扩为\s*(\d+)\s*行',text)
    bands=re.findall(r'([^；：。\n]+?)合计第?(\d+)行[、，,]\s*明细(\d+)\s*[–—-]\s*(\d+)行',text)
    amount_columns=re.search(r'([A-Z]+):([A-Z]+)\s*的([^；。\n]+?)列角色不变',text)
    source_order=re.search(r'明细来源按([^；。\n]+?)排列',text)
    matrix=[list(map(str.strip,line.strip().strip('|').split('|'))) for line in text.splitlines()
            if line.strip().startswith('|') and not re.fullmatch(r'[\s|:\-]+',line)]
    source_sheet=re.search(r'B\s*坐标均为\s*([^\n]+?)\s*工作表',text)
    kept=re.search(r'另保留\s*A\s*原有的([^；。\n]+?)行',text)
    if all((expansion,bands,amount_columns,source_order,source_sheet,kept,matrix)):
        old,new=map(int,expansion.groups());delta=new-old
        bands=[(name.strip(),int(total),int(first),int(last)) for name,total,first,last in bands]
        from openpyxl.utils.cell import column_index_from_string
        cols=[get_column_letter(c) for c in range(column_index_from_string(amount_columns[1]),column_index_from_string(amount_columns[2])+1)]
        order=source_order[1].split('、');header,*members=matrix
        areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
               if range_boundaries(area['range'])[1]<bands[0][2]<=bands[-1][3]-delta*len(bands)<=range_boundaries(area['range'])[3]]
        if len(areas)==1 and len(order)==len(cols) and len(members)+1==new and delta>=0:
            name,area=areas[0];left,top,_,_=range_boundaries(area['range']);sheet=area['sheet'];label=get_column_letter(left)
            cells=next(s['cells'] for s in a['机械记录']['sheets'] if s['name']==sheet)
            prior='\n'.join(r.get('content','') for r in (b or {}).get('已确认关系',[]) if r.get('dimension')==dimension)
            aliases={new_name:old_name for old_name,new_name in re.findall(r'A\s*[“「"]([^”」"]+)[”」"]对应\s*B\s*[“「"]([^”」"]+)[”」"]',prior)}
            new_names={name.strip() for names in re.findall(r'新增([^。；\n]+?)(?:没有\s*B|无\s*B)',text) for name in names.split('、')}
            def original_row(member,first):
                stated=aliases.get(member,member)
                found=[row for row in range(first,first+old)
                       if str(cells.get(f'{label}{row}',{}).get('value','')).strip() in (stated,'其中：'+stated)]
                return found[0] if len(found)==1 else None
            source_rows=[int(re.search(r'[A-Z]+(\d+)',member[1])[1]) for member in members]
            tables=[]
            for index,(segment,total,start,end) in enumerate(bands):
                first=start-delta*index
                titles=[i for i,title in enumerate(header) if title.startswith(segment)]
                selected=list(range(len(cols)));sources={}
                if len(titles)==1:
                    title=header[titles[0]]
                    for suffix,roles in re.findall(r'([^，；。\n]+?)仅有([^，；。\n]+?)(?:两项|三项|四项|各项)',text):
                        if title.startswith(suffix.strip()) or suffix.strip() in title:
                            selected=[order.index(role) for role in roles.split('、') if role in order]
                    entries=[re.findall(r'\b[A-Z]+\d+\b',member[titles[0]]) for member in members]
                    if not selected or any(len(values)!=len(selected) for values in entries):continue
                    sources={str(row):{cols[i]:source_sheet[1].strip()+'!'+value for i,value in zip(selected,values)}
                             for row,values in zip(source_rows,entries)}
                elif not re.search(re.escape(segment)+r'明细[：:]\s*B\s*无[^。\n]*来源',text):
                    continue
                retained=original_row(kept[1].strip(),first)
                if end-start+1!=new or retained is None:continue
                blank_values={}
                for source_row,member in zip(source_rows,members):
                    original=original_row(member[0],first)
                    if original is None and member[0] not in new_names:break
                    blank_values[str(source_row)]={col:({'source':f'{sheet}!{col}{original}'} if original is not None else {'value':None}) for col in cols}
                if len(blank_values)!=len(members):continue
                mapping={col:re.match(r'[A-Z]+',ref.rsplit('!',1)[1])[0] for col,ref in next(iter(sources.values()),{}).items()}
                tables.append({'selection_dimension':dimension,'range_name':name,'source_sheet':source_sheet[1].strip(),
                    'source_rows':source_rows,'old_detail_count':old,'first_row_offset':first-top,
                    'columns':mapping,'row_sources':sources,'labels':{label:None},
                    'declared_labels':{str(row):{'value':member[0]} for row,member in zip(source_rows,members)},
                    'blank_values':blank_values,'retained_members':[{'index':len(members),
                        'cells':{col:{'source':f'{sheet}!{col}{retained}'} for col in [label,*cols]}}],
                    'operations':[{'action':'insert','axis':'row','index':first+old-top,'count':delta,'template_index':first-top}] if delta else []})
            if tables:return tables
    if re.search(r'不增行|不扩展明细',text):
        columns={target:source for source,target in re.findall(
            r'([A-Z]+)列[“"「][^”"」]+[”"」]搬至A对应行([A-Z]+)列',text)}
        for source,target in re.findall(r'B的([A-Z、]+)列[^。\n]*?对应A的([A-Z、]+)列',text):
            source,target=source.split('、'),target.split('、')
            if len(source)==len(target):columns.update(zip(target,source))
        tables=[]
        def row_numbers(value):
            return [number for part in value.replace('A','').split('、') for ends in [part.split(':')]
                    for number in range(int(ends[0]),int(ends[-1])+1)]
        for line in text.splitlines():
            band=re.search(r'[（(]A(\d+):A?(\d+)[）)]',line)
            source=re.search(r'B《([^》]+)》',line)
            if not band or not source or not columns:continue
            first,last=map(int,band.groups());pairs=[]
            for target,rows in re.findall(
                r'A(\d+(?::A?\d+)?(?:、A?\d+)*)\s*(?:分别|依次)?对应B(?:《[^》]+》|第)?(\d+(?::\d+)?(?:、\d+)*)行?',line):
                left,right=row_numbers(target),row_numbers(rows)
                if len(left)==len(right):pairs.extend(zip(left,right))
            if not pairs or [row for row,_ in pairs]!=list(range(first,first+len(pairs))):continue
            areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
                   if range_boundaries(area['range'])[1]<=first<=last<=range_boundaries(area['range'])[3]]
            if len(areas)!=1:continue
            name,area=areas[0];top=range_boundaries(area['range'])[1]
            retained=[int(number) for number in re.findall(r'A(\d+)[“"「][^”"」]+[”"」]保留原位置',line)]
            tables.append({'selection_dimension':dimension,'range_name':name,'source_sheet':source[1],
                'source_rows':[row for _,row in pairs],
                'old_detail_count':last-first+1,'first_row_offset':first-top,'columns':columns,'labels':{},
                'operations':[],'retained_rows':retained})
        if tables:return tables
    target=re.search(r'A[，,]\s*([^!\n]+)!([A-Z]+\d+:[A-Z]+\d+)[（(]([^）)]+)[）)][^\n]*?每组先列合计[^：\n]*明细[：:]([^。\n]+)',text)
    source=re.search(r'B[，,]\s*([^!\n]+)!([A-Z]+\d+:[A-Z]+\d+)[^\n]*?明细标签在\s*([A-Z]+)\s*列',text)
    names=re.search(r'其中[^。\n；;]*?(\d+)\s*[–—-]\s*(\d+)显示这些名称',text)
    arrangement=re.search(r'应保留\s*([A-Z]+\d+:[A-Z]+\d+)的列角色[：:]([^；。\n]+)',text)
    if not all((target,source,names,arrangement)) or not re.search(r'将每组[^；。\n]*明细项改为[^；。\n]*标签',text):return []
    area=a['机械记录'].get('names',{}).get(target[3],[])
    if len(area)!=1 or area[0]['sheet']!=target[1] or area[0]['range'].replace('$','')!=target[2] or arrangement[1]!=target[2]:return []
    left,top,right,bottom=range_boundaries(target[2])
    roles=arrangement[2].split('、')
    if len(roles)!=right-left+1 or roles[0]!='项目':return []
    segments=[tuple(map(int,row)) for row in re.findall(r'(\d+)\s*/\s*(\d+)\s*[–—-]\s*(\d+)',target[4])]
    first,last=map(int,names.groups());source_rows=list(range(first,last+1))
    sl,st,sr,sb=range_boundaries(source[2])
    if not source_rows or not st<=first<=last<=sb or not segments:return []
    tables=[]
    for total,start,end in segments:
        if not top<=total<start<=end<=bottom:return []
        tables.append({'selection_dimension':dimension,'range_name':target[3],
            'source_sheet':source[1],'source_rows':source_rows,'old_detail_count':end-start+1,
            'first_row_offset':start-top,'columns':{get_column_letter(col):None for col in range(left+1,right+1)},
            'labels':{get_column_letter(left):source[3]},'source_ranges':[{'sheet':source[1],'range':source[2]}],
            'source_evidence':[source[1]+'!'+source[3]+str(row) for row in source_rows],
            'coverage_reason':'原答复明确各分段保留合计并按已列标签扩展明细；金额仍用已识别身份对应'})
    return tables


def _declared_list_tables(a, reply, dimension, source_ranges, b=None):
    """只转存答复明示的连续名单位置；金额对应留给已有身份比较。"""
    from .A语义阅读 import addresses
    from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries, get_column_letter
    text=reply.replace('`','').replace('**','')
    if b:text=_source_aliases(text,{**b,'target_table':None,'target_tables':[],'b_source_ranges':source_ranges})
    sections=re.split(r'\n(?=#{1,6}\s)',text)
    inserted=[]
    source_sheets={r['sheet'] for r in source_ranges if re.search(r'金额采用\s*B\s*'+re.escape(r['sheet'])+r'\s*的',text)}
    for block in text.split('\n\n'):
        placement=re.search(r'放在原?[“"]?([^”"；，、\n]+)[”"]?之后[、，,]\s*原?[“"]?([^”"；，、\n]+)[”"]?之前',block)
        entries=re.findall(r'(?:名称取\s*)?B\s+([A-Z]+)(\d+)[^；;\n]*[；;]\s*金额(?:依次)?(?:取)?\s*([A-Z]+\d+(?:[、，,][A-Z]+\d+)*)',block)
        if not (placement and entries and len(source_sheets)==1):continue
        source_sheet=next(iter(source_sheets));candidates=[]
        for name,parts in a['机械记录'].get('names',{}).items():
            for area in parts:
                left,top,right,bottom=range_boundaries(area['range']);sheet=area['sheet']
                cells=next(s['cells'] for s in a['机械记录']['sheets'] if s['name']==sheet)
                labels={row:str(cells.get(get_column_letter(left)+str(row),{}).get('value') or '').strip() for row in range(top,bottom+1)}
                previous=[row for row,value in labels.items() if value==placement[1]]
                following=[row for row,value in labels.items() if value==placement[2]]
                if len(previous)==len(following)==1 and following[0]==previous[0]+1:
                    candidates.append((name,area,previous[0],following[0]))
        if len(candidates)!=1:continue
        name,area,previous,point=candidates[0];left,top,right,bottom=range_boundaries(area['range'])
        from openpyxl.utils.cell import get_column_letter
        columns=sorted({coordinate_to_tuple(cell)[1] for table in a.get('outline',{}).get('subtables',[])
            if table['sheet']==area['sheet'] and table['range'].replace('$','')==area['range'].replace('$','')
            for metric in table.get('metrics',[]) for span in metric['ranges'] if span!='*' for cell in addresses(span)
            if left<coordinate_to_tuple(cell)[1]<=right})
        amounts=[re.findall(r'[A-Z]+\d+',entry[2]) for entry in entries]
        if not columns or any(len(values)!=len(columns) for values in amounts) or len({entry[0] for entry in entries})!=1:continue
        rows=[int(entry[1]) for entry in entries]
        inserted.append({'selection_dimension':dimension,'range_name':name,'source_sheet':source_sheet,
            'source_rows':rows,'target_row_offsets':list(range(point-top,point-top+len(rows))),
            'first_row_offset':point-top,'old_detail_count':0,
            'columns':{get_column_letter(col):re.match(r'[A-Z]+',amounts[0][i])[0] for i,col in enumerate(columns)},
            'row_sources':{str(row):{get_column_letter(col):source_sheet+'!'+value for col,value in zip(columns,values)} for row,values in zip(rows,amounts)},
            'labels':{get_column_letter(left):entries[0][0]},
            'blank_values':{str(row):{get_column_letter(col):{'value':None} for col in columns} for row in rows},
            'operations':[{'axis':'row','action':'insert','index':point-top,'count':len(rows),'template_index':previous-top}]})
    if inserted:return inserted
    # 新排布引用先归入所属原表；不能把它当成原表中同名坐标的金额对应。
    for section in sections:
        placement=re.search(r'A\s+([^!\s]+)!([A-Z]+)(\d+):\2(\d+)\s*依次列',section)
        total=re.search(r'A\s+([^!\s]+)!([A-Z]+)(\d+)\s*为合计',section)
        names=re.search(r'B\s+([^!\s]+)!([A-Z]+)(\d+):\2(\d+)\s*的实际名单',section)
        if not (b and placement and total and names):continue
        from openpyxl.utils.cell import get_column_letter
        sheet,label,first,last=placement[1],placement[2],int(placement[3]),int(placement[4])
        source_sheet,source_label,source_first,source_last=names[1],names[2],int(names[3]),int(names[4])
        areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
            if area['sheet']==sheet and area['range'].replace('$','') in section
            and range_boundaries(area['range'])[1]<first<=range_boundaries(area['range'])[3]]
        if len(areas)!=1 or last-first!=source_last-source_first or (total[1],total[2])!=(sheet,label):continue
        name,area=areas[0];left,top,right,old_total=range_boundaries(area['range']);new_total=int(total[3])
        cells=next(s['cells'] for s in a['机械记录']['sheets'] if s['name']==sheet)
        kept=[]
        for target,member in re.findall(r'A\s+'+re.escape(sheet)+r'!'+label+r'(\d+)\s*保留原[“"]([^”"]+)[”"]及其原值',section):
            originals=[row for row in range(first,old_total) if cells.get(label+str(row),{}).get('value')==member]
            if len(originals)==1:
                kept.append({'index':int(target)-first,'cells':{get_column_letter(col):{'source':sheet+'!'+get_column_letter(col)+str(originals[0])}
                    for col in range(left,right+1)}})
        if new_total!=last+len(kept)+1 or new_total<old_total:continue
        targets=[sheet+'!'+get_column_letter(col)+str(row) for row in [*range(first,last+1),new_total] for col in range(left+1,right+1)]
        pairs=source_relation(a,b,section,targets)['bindings'];columns={};source_total=None
        for pair in pairs:
            target_row,target_col=coordinate_to_tuple(pair['target'].rsplit('!',1)[1])
            sh,cell=pair['source'].rsplit('!',1);source_row,source_col=coordinate_to_tuple(cell)
            if sh!=source_sheet:continue
            if target_row==new_total:source_total=source_row
            elif target_row-first==source_row-source_first:columns[get_column_letter(target_col)]=get_column_letter(source_col)
        if not columns or source_total is None:continue
        # 旧成员与来源成员的对应只复用先前答复明说的关系，名称仅用于定位原内容。
        originals={}
        for prior in b.get('已确认关系',[]):
            previous=prior.get('content','').replace('`','').replace('**','')
            for declared in re.finditer(r'A现有项目([^。\n]+?)对应B\s+([^。\n]+)',previous):
                members=re.findall(r'“([^”]+)”',declared[1])
                references=re.findall(r'(?:'+re.escape(source_sheet)+r'!)?'+source_label+r'(\d+)\s*“([^”]+)”',declared[2])
                if len(members)!=len(references):continue
                for member,(row,_) in zip(members,references):
                    old=[r for r in range(first,old_total) if cells.get(label+str(r),{}).get('value')==member]
                    if len(old)==1:originals[int(row)]=old[0]
        preserved=set()
        for paragraph in section.split('\n\n'):
            if '保留原内容' not in paragraph:continue
            preserved.update(col for col,start,end in re.findall(r'A\s+'+re.escape(sheet)+r'!([A-Z]+)(\d+):\1(\d+)',paragraph) if (int(start),int(end))==(first,last))
        rows=list(range(source_first,source_last+1))
        return [{'selection_dimension':dimension,'range_name':name,'source_sheet':source_sheet,
            'source_rows':rows,'old_detail_count':old_total-first,'first_row_offset':first-top,
            'columns':columns,'labels':{label:source_label},'retained_members':kept,
            'blank_values':{str(row):{col:({'source':sheet+'!'+col+str(originals[row])} if row in originals else {'value':None})
                for col in set(columns)|preserved} for row in rows},
            'totals':[{'target_row':old_total,'source_row':source_total}],
            'operations':[{'axis':'row','action':'insert','index':old_total-top,'count':new_total-old_total,'template_index':first-top}]}]
    # 数字行序表直接使用答复给出的来源段、原样式及保留成员，不重认名称含义。
    styles=re.findall(r'明细行\s*(\d+)(?:\s*[–—至-]\s*(\d+))?',text)
    columns=dict(re.findall(r'A的([A-Z]+)←B的([A-Z]+)',text))
    row_plan=[];header=None
    for line in text.splitlines():
        values=[value.strip() for value in line.strip().strip('|').split('|')]
        if 'A位置' in values and any('名称' in value for value in values):
            header=values;continue
        if header and len(values)==len(header):
            number=values[header.index('A位置')]
            if re.fullmatch(r'\d+(?:[–—至-]\d+)?',number):
                row_plan.append((list(map(int,re.split('[–—至-]',number))),values))
    groups=[]
    for index,(rows,values) in enumerate(row_plan):
        for source_sheet in {r['sheet'] for r in source_ranges}:
            match=re.search(r'B\s*'+re.escape(source_sheet)+r'\s+([A-Z]+)(\d+)[–—至-]\1(\d+)', ' '.join(values))
            if match:groups.append((index,rows,source_sheet,match[1],int(match[2]),int(match[3])))
    anchors=[int(n) for n in re.findall(r'原[^。；;|\n]*?样式行\s*(\d+)',text)]
    if groups and len(styles)==len(groups) and columns and anchors:
        from openpyxl.utils.cell import get_column_letter
        starts=[int(pair[0]) for pair in styles]
        areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
               if all(range_boundaries(area['range'])[1]<row<range_boundaries(area['range'])[3] for row in starts)
               and max(anchors)<=range_boundaries(area['range'])[3]]
        if len(areas)==1:
            name,area=areas[0];left,top,right,bottom=range_boundaries(area['range']);sheet=area['sheet']
            cells=next(s['cells'] for s in a['机械记录']['sheets'] if s['name']==sheet);tables=[]
            for start,(index,rows,source_sheet,source_col,first,last) in zip(starts,groups):
                stop=min((n for n in anchors if n>start),default=None)
                if stop is None or len(rows)!=2 or rows[1]-rows[0]!=last-first:break
                kept=[]
                for kept_rows,values in row_plan[index+1:]:
                    if not any('保留A原“' in v for v in values):break
                    names=re.findall(r'保留A原“([^”]+)”',' '.join(values))
                    old=[n for n in range(start,stop) if str(cells.get(get_column_letter(left)+str(n),{}).get('value','')).strip() in names]
                    if len(old)!=1 or len(kept_rows)!=1:break
                    kept.append({'index':kept_rows[0]-rows[0],'cells':{
                        get_column_letter(col):{'source':sheet+'!'+get_column_letter(col)+str(old[0])} for col in range(left,right+1)}})
                added=last-first+1+len(kept)-(stop-start)
                if added<0:break
                tables.append({'selection_dimension':dimension,'range_name':name,'source_sheet':source_sheet,
                    'source_rows':list(range(first,last+1)),'old_detail_count':stop-start,'first_row_offset':start-top,
                    'columns':columns,'labels':{get_column_letter(left):source_col},'retained_members':kept,
                    'blank_values':{str(n):{col:{'value':None} for col in columns} for n in range(first,last+1)},
                    'operations':[{'axis':'row','action':'insert','index':stop-top,'count':added,'template_index':start-top}] if added else []})
            if len(tables)==len(groups):return tables
    # 名称和金额分处两表时，沿用已识别列角色；仅展开答复明列的顺序与行数。
    for section in sections:
        target=re.search(r'\bA\s+([^!\s]+)!([A-Z]+\d+:[A-Z]+\d+)',section)
        old=re.search(r'A\s*原版式为\s*(\d+)\s*个明细行加\s*1\s*个合计行',section)
        names=re.search(r'名称及明细顺序取自\s*B\s+([^!\s]+)!([A-Z]+)(\d+):\2(\d+)',section)
        amounts=re.search(r'金额采用对应的\s*B\s+([^!\s]+)!([^。；\n]+)',section)
        total=re.search(r'A\s*合计行采用\s*B\s+([^!\s]+)!([A-Z]+)(\d+):([A-Z]+)\3',section)
        if not all((target,old,names,amounts,total)):continue
        areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
               if area['sheet']==target[1] and range_boundaries(area['range'])==range_boundaries(target[2])]
        spans=re.findall(r'([A-Z]+)(\d+):\1(\d+)[（(](期末|期初|本期|上期)[）)]',amounts[2])
        if len(areas)!=1 or not spans or amounts[1]!=total[1]:continue
        name,area=areas[0];left,top,right,bottom=range_boundaries(area['range']);old_count=int(old[1]);first=bottom-old_count
        source_rows=list(range(int(spans[0][1]),int(spans[0][2])+1))
        label_rows=list(range(int(names[3]),int(names[4])+1))
        if len(source_rows)!=len(label_rows) or any((start,end)!=(spans[0][1],spans[0][2]) for _,start,end,_ in spans):continue
        roles={'期末':'closing','期初':'opening','本期':'current','上期':'prior'}
        columns={}
        for column,_,_,role in spans:
            candidates={re.match(r'[A-Z]+',cell)[0] for table in a.get('outline',{}).get('subtables',[])
                if table['sheet']==area['sheet'] for metric in table.get('metrics',[])
                if metric.get('template_period_role')==roles[role] for span in metric['ranges']
                for cell in addresses(table['range'] if span=='*' else span)
                if cell in addresses(area['range']) and coordinate_to_tuple(cell)[0]==first}
            if len(candidates)==1:columns[next(iter(candidates))]=column
        from openpyxl.utils.cell import get_column_letter
        labels={get_column_letter(col) for col in range(left,right+1)}-columns.keys()
        if len(columns)!=len(spans) or len(labels)!=1:continue
        return [{'selection_dimension':dimension,'range_name':name,'source_sheet':amounts[1],
            'source_rows':source_rows,'old_detail_count':old_count,'first_row_offset':first-top,
            'columns':columns,'labels':{next(iter(labels)):None},
            'declared_labels':{str(row):{'source':names[1]+'!'+names[2]+str(label)} for row,label in zip(source_rows,label_rows)},
            'totals':[{'target_row':bottom,'source_row':int(total[3])}]}]
    scoped=[part for part in sections if re.search(r'(?<!\w)'+re.escape(dimension)+r'(?!\w)',part) and 'A沿用原样式' in part]
    if b and len(scoped)==1:
        from openpyxl.utils.cell import get_column_letter, column_index_from_string
        section=scoped[0]
        placement=re.search(r'明细[^。\n]*?([A-Z]+)(\d+):\1(\d+)[^。\n]*?合计(?:行|移至)(\d+)',section)
        target_columns=dict(re.findall(r'A的([A-Z]+)列接B(本期|上期)',text))
        declaration=re.search(r'金标准[:：]\s*(metric\.[\w.]+)',section)
        if placement and target_columns and declaration:
            label,first,last,new_total=placement.groups();first,last,new_total=map(int,(first,last,new_total))
            candidates=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
                if range_boundaries(area['range'])[0]==column_index_from_string(label)
                and range_boundaries(area['range'])[1]<first<=range_boundaries(area['range'])[3]]
            right=dict(_records(b));left=dict(_records(a));referenced=set()
            for sheet in {item['sheet'] for item in source_ranges}:
                for span in re.findall(re.escape(sheet)+r'!([A-Z]+\d+(?::[A-Z]+\d+)?)',section):
                    referenced.update(sheet+'!'+cell for cell in addresses(span))
            declared_sheets={loc.rsplit('!',1)[0] for loc in referenced}
            if len(declared_sheets)==1:
                source_sheet=next(iter(declared_sheets))
                for span in re.findall(r'(?<![\w!])([A-Z]+\d+(?::[A-Z]+\d+)?)',section.split('A沿用原样式',1)[0]):
                    referenced.update(source_sheet+'!'+cell for cell in addresses(span))
            known=[(loc,row) for loc,row in right.items() if loc in referenced and row.get('metric_id')==declaration[1]
                and isinstance(row.get('dimensions',{}).get(dimension),dict)]
            members={};totals=set();source_columns={}
            for loc,row in known:
                cell=loc.rsplit('!',1)[1];number,col=coordinate_to_tuple(cell);value=row['dimensions'][dimension]
                if value.get('mode')=='members' and len(value.get('members',[]))==1:members[number]=value['members'][0]
                if value.get('mode')=='all':totals.add(number)
                roles={item.get('template_period_role') for table in b.get('outline',{}).get('subtables',[]) if table['sheet']==loc.rsplit('!',1)[0]
                    for item in table['metrics'] if isinstance(item.get('template_period_role'),str)
                    and any(span=='*' or cell in addresses(span) for span in item['ranges'])}
                period_role=next(iter(roles)) if len(roles)==1 else None
                if period_role in ('current','prior'):source_columns.setdefault(period_role,set()).add(get_column_letter(col))
            columns={col:next(iter(source_columns[{'本期':'current','上期':'prior'}[period]])) for col,period in target_columns.items()
                if len(source_columns.get({'本期':'current','上期':'prior'}[period],()))==1}
            if len(candidates)==len(declared_sheets)==len(totals)==1 and members and len(columns)==len(target_columns):
                name,area=candidates[0];sheet=area['sheet'];_,top,_,_=range_boundaries(area['range'])
                owned={sheet+'!'+cell for cell in addresses(area['range'])}
                old_totals={coordinate_to_tuple(loc.rsplit('!',1)[1])[0] for loc,row in left.items() if loc in owned
                    and re.match(r'[A-Z]+',loc.rsplit('!',1)[1])[0] in columns and row.get('metric_id')==declaration[1] and isinstance(row.get('dimensions',{}).get(dimension),dict)
                    and row['dimensions'][dimension].get('mode')=='all'}
                if len(old_totals)==1:
                    old_total=next(iter(old_totals));old_count=old_total-first
                    kept=re.search(re.escape(sheet)+r'!([A-Z]+\d+:[A-Z]+\d+)[^。]*行及其原内容保留',section)
                    retained=[]
                    if kept and re.search(r'再保留原[“"][^”"]+[”"]行',section):
                        retained=[{'index':len(members),'cells':{re.match(r'[A-Z]+',cell)[0]:{'source':sheet+'!'+cell} for cell in addresses(kept[1])}}]
                    if len(members)+len(retained)==last-first+1 and new_total==last+1 and new_total>=old_total:
                        rows=sorted(members);total=next(iter(totals))
                        table={'selection_dimension':dimension,'range_name':name,'source_sheet':next(iter(declared_sheets)),
                            'source_rows':rows,'old_detail_count':old_count,'first_row_offset':first-top,'columns':columns,
                            'labels':{label:None},'declared_labels':{str(row):{'value':members[row]} for row in rows},
                            'retained_members':retained,'totals':[{'target_row':old_total,'source_row':total}]}
                        if new_total>old_total:table['operations']=[{'axis':'row','action':'insert','index':old_total-top,
                            'count':new_total-old_total,'template_index':first-top}]
                        preserved=set(re.findall(r'原(?:有)?([A-Z]+)列(?:内容|空白)随原成员(?:和|及)合计保留[，,]新成员\1列留空',section))
                        cells=next(item['cells'] for item in a['机械记录']['sheets'] if item['name']==sheet)
                        for col in preserved:
                            for row in rows:
                                same=[old for old in range(first,old_total) if cells.get(label+str(old),{}).get('value')==members[row]]
                                table.setdefault('blank_values',{}).setdefault(str(row),{})[col]=(
                                    {'source':sheet+'!'+col+str(same[0])} if len(same)==1 else {'value':None})
                            table.setdefault('blank_values',{}).setdefault(str(total),{})[col]={'source':sheet+'!'+col+str(old_total)}
                        return [table]
    columns_stated=re.search(r'本期[、，,]上期金额列分别取自\s*B\s*的\s*([A-Z]+)[、，,]([A-Z]+)\s*列',text)
    if b and columns_stated and re.search(r'按\s*A\s*原表样式',text) and re.search(r'更新为上述[^。；;]*明细后接合计',text):
        right=dict(_records(b));left_records=dict(_records(a));items=[]
        for line in text.splitlines():
            label=re.match(r'\s*-\s*B\s+([^!\s]+)!([A-Z]+)(\d+)\s+([^：:]+)[：:]',line)
            if not label:continue
            source=label[1]+'!'+columns_stated[1]+label[3];record=right.get(source,{})
            selection=record.get('dimensions',{}).get(dimension)
            if isinstance(selection,dict) and selection.get('mode') in ('members','all'):
                items.append((int(label[3]),label[4],label[1],label[2],record['metric_id'],selection['mode'],line))
        metrics={item[4] for item in items};totals=[item for item in items if item[5]=='all'];details=[item for item in items if item[5]=='members']
        candidates=[]
        if len(metrics)==1 and len(totals)==1 and details:
            for name,parts in a['机械记录'].get('names',{}).items():
                if len(parts)!=1:continue
                area=parts[0];owned={area['sheet']+'!'+cell for cell in addresses(area['range'])}
                records=[(loc,row) for loc,row in left_records.items() if loc in owned and row['metric_id'] in metrics and dimension in row['dimensions']]
                old_totals={coordinate_to_tuple(loc.rsplit('!',1)[1])[0] for loc,row in records if isinstance(row['dimensions'][dimension],dict) and row['dimensions'][dimension].get('mode')=='all'}
                old_details={coordinate_to_tuple(loc.rsplit('!',1)[1])[0] for loc,row in records if isinstance(row['dimensions'][dimension],dict) and row['dimensions'][dimension].get('mode')=='members'}
                if len(old_totals)==1 and old_details:candidates.append((name,area,records,old_details,next(iter(old_totals))))
        if len(candidates)==1 and len({item[2] for item in items})==1:
            from openpyxl.utils.cell import get_column_letter
            name,area,records,old_details,total=candidates[0];sheet=area['sheet'];first=min(old_details);last=max(old_details)
            left,top,_,_=range_boundaries(area['range']);columns={};roles={'current':columns_stated[1],'prior':columns_stated[2],'本期':columns_stated[1],'上期':columns_stated[2]}
            owned={loc.rsplit('!',1)[1] for loc,_ in records}
            for table in a.get('outline',{}).get('subtables',[]):
                if table['sheet']!=sheet:continue
                for metric in table['metrics']:
                    role=metric.get('template_period_role')
                    if not isinstance(role,str) or role not in roles:continue
                    for span in metric['ranges']:
                        for cell in (owned if span=='*' else owned&set(addresses(span))):columns[re.match(r'[A-Z]+',cell)[0]]=roles[role]
            retained=set()
            for kept in re.finditer(r'A\s*原明细[^。]*?'+re.escape(sheet)+r'!([A-Z]+\d+:[A-Z]+\d+)[^。]*?[。；;]\s*这组\s*A\s*原内容应保留在原行',text):
                retained.update(coordinate_to_tuple(cell)[0] for cell in addresses(kept[1]))
            retained&=old_details
            # shortcut: 复用连续明细及末尾合计的样式；交错保留行仍读取其他明确排布。
            if columns and old_details==set(range(first,last+1)) and total==last+1 and retained==set(range(first,first+len(retained))) and len(details)+len(retained)>=len(old_details):
                count=len(details)+len(retained)-len(old_details)
                start=first+len(retained)
                table={'selection_dimension':dimension,'range_name':name,'source_sheet':details[0][2],
                    'source_rows':[item[0] for item in details],'old_detail_count':len(old_details),'first_row_offset':first-top,
                    'target_row_offsets':[start-top+i for i in range(len(details))],
                    'columns':columns,'labels':{get_column_letter(left):None},
                    'declared_labels':{str(item[0]):{'value':item[1]} for item in details},
                    'retained_rows':sorted(retained),'totals':[{'target_row':total,'source_row':totals[0][0]}]}
                if count>0:table['operations']=[{'axis':'row','action':'insert','index':total-top,'count':count,'template_index':first-top}]
                for item in details:
                    for source_column in re.findall(r'(?<![A-Z])([A-Z]+)'+str(item[0])+r'\s*原文空白[，,]不填金额',item[6]):
                        table.setdefault('blank_values',{})[str(item[0])]={col:{'value':None} for col,source in columns.items() if source==source_column}
                return [table]
    relative=[]
    for line in text.splitlines():
        part=re.match(r'\s*-\s*[^：]+：\s*([^。]+?)之后列B[^：]+：([^。]+)',line)
        if not part:continue
        names=[name.strip() for name in re.split(r'[；;]' if re.search(r'[；;]',part[2]) else '、',part[2])]
        retained=[]
        kept=re.search(r'单列保留A旧成员([^。；]+)',line)
        if kept:retained.extend(re.findall(r'[“"]([^”"]+)[”"]',kept[1]))
        retained.extend(re.findall(r'A原有[“"]([^”"]+)[”"][^。]*单列保留',line))
        relative.append((part[1].strip(),names,retained))
    areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
           if area.get('range','').replace('$','') in text]
    if relative and len(areas)==1 and b:
        from .表格 import _REFERENCE
        from openpyxl.utils.cell import get_column_letter
        name,area=areas[0];left,top,right,bottom=range_boundaries(area['range']);label=get_column_letter(left)
        target_sheet=area['sheet'];target_cells=next(s['cells'] for s in a['机械记录']['sheets'] if s['name']==target_sheet)
        def literal(cell):return str(cell.get('cached_value') if cell.get('formula') else cell.get('value','')).strip()
        def named_row(member):
            declared={int(row) for col,row in re.findall(r'A\s*([A-Z]+)(\d+)\s*[“"]?'+re.escape(member),text)
                      if col==label and top<=int(row)<=bottom}
            if len(declared)==1:return next(iter(declared))
            found=[row for row in range(top,bottom+1) if literal(target_cells.get(label+str(row),{}))==member]
            return found[0] if len(found)==1 else None
        starts=[named_row(anchor) for anchor,_,_ in relative]
        targets=[target_sheet+'!'+cell for cell in addresses(area['range'])]
        pairs=source_relation(a,{**b,'target_table':None,'target_tables':[],'b_source_ranges':source_ranges},reply,targets)['bindings']
        labels=re.findall(r'\|\s*B\s+([A-Z]+)(\d+)(?=\s|\|)',text)
        source_labels={col for col,_ in labels};source_rows={int(row) for _,row in labels}
        chosen={};headers=None
        paths=b.get('答复接入来源',{}).get('本次原答复','')+'\n'+reply
        blank_columns=set(re.findall(r'([A-Z]+)列为空白[^。]*?保留为空白[^。]*?不以A旧值替代',paths.replace('`','')))
        for line in paths.replace('`','').replace('**','').splitlines():
            if not line.strip().startswith('|'):continue
            values=[v.strip() for v in line.strip().strip('|').split('|')]
            if '采用的B来源' in values:headers=values;continue
            if not headers or len(values)!=len(headers):continue
            column=re.search(r'([A-Z]+)列',values[0])
            if not column:continue
            refs=re.findall(r'([^\s!；;]+)!([A-Z]+\d+:[A-Z]+\d+)',values[headers.index('采用的B来源')])
            chosen[column[1]]={sheet+'!'+cell for sheet,span in refs for cell in addresses(span)}
        if all(row is not None for row in starts) and starts==sorted(starts) and len(source_labels)==1:
            tables=[];source_label=next(iter(source_labels))
            for index,((_,members,kept),anchor) in enumerate(zip(relative,starts)):
                source_pairs=[pair for pair in pairs if coordinate_to_tuple(pair['target'].rsplit('!',1)[1])[0]==anchor]
                sources={pair['source'].rsplit('!',1)[0] for pair in source_pairs}
                if len(sources)!=1:break
                source_sheet=next(iter(sources));source_cells=next(s['cells'] for s in b['机械记录']['sheets'] if s['name']==source_sheet)
                columns={re.match(r'[A-Z]+',pair['target'].rsplit('!',1)[1])[0]:re.match(r'[A-Z]+',pair['source'].rsplit('!',1)[1])[0] for pair in source_pairs}
                rows=[[row for row in sorted(source_rows) if literal(source_cells.get(source_label+str(row),{}))==member] for member in members]
                kept_rows=[named_row(member) for member in kept]
                if any(len(found)!=1 for found in rows) or any(row is None for row in kept_rows):break
                rows=[found[0] for found in rows];first=anchor+1;last=(starts[index+1] if index+1<len(starts) else bottom)-1
                old_count=last-first+1;added=len(rows)+len(kept_rows)-old_count
                if added<0:break
                row_sources={}
                for row in rows:
                    for target_col,source_col in columns.items():
                        formula=source_cells.get(source_col+str(row),{}).get('formula','')
                        direct=_REFERENCE.fullmatch(formula.removeprefix('='))
                        if direct and not direct['b']:
                            sheet=(direct['sheet'] or source_sheet).strip("'").replace("''","'")
                            location=sheet+'!'+direct['a'].replace('$','')
                            if location in chosen.get(source_col,set()):row_sources.setdefault(str(row),{})[target_col]=location
                tables.append({'selection_dimension':dimension,'range_name':name,'source_sheet':source_sheet,
                    'source_rows':rows,'first_row_offset':first-top,'old_detail_count':old_count,'columns':columns,
                    'labels':{label:source_label},'row_sources':row_sources,
                    'blank_values':{row:{col:{'value':None} for col,source in sources.items()
                        if re.match(r'[A-Z]+',source.rsplit('!',1)[1])[0] in blank_columns} for row,sources in row_sources.items()},
                    'retained_members':[{'index':len(rows)+i,'cells':{get_column_letter(col):{'source':target_sheet+'!'+get_column_letter(col)+str(row)}
                        for col in range(left,right+1)}} for i,row in enumerate(kept_rows)],
                    'operations':[{'action':'insert','axis':'row','index':first+old_count-top,'count':added,'template_index':first-top}] if added else []})
            if len(tables)==len(relative):return tables
    # 行对应表和后续保留说明共同引用原答复；名称只与已有成员逐字连接。
    prior='\n'.join(row.get('content','') for row in (b or {}).get('已确认关系',[])
        if row.get('kind')=='名单编排' and row.get('dimension')==dimension and row.get('source_ranges')==source_ranges)
    previous='\n'.join(part.get('识别答复','') for mapping in (a,b or {}) for part in mapping.get('合集来源',[]))
    combined=(prior+'\n'+text).replace('`','').replace('**','')
    members=[];header=None
    for line in combined.splitlines():
        values=[part.strip() for part in line.strip().strip('|').split('|')]
        titles=[re.sub(r'\s+','',part) for part in values]
        if all(name in titles for name in ('B来源行','搬至A行','成员')):
            header=titles;members=[];continue
        if header and len(values)==len(header):
            source,target=values[header.index('B来源行')],values[header.index('搬至A行')]
            if source.isdigit() and target.isdigit():members.append((int(source),int(target),values[header.index('成员')]))
    retained=re.search(r'无法确认同义的[^：:\n]+[:：]\s*([^，。\n]+)',text)
    kept_range=re.search(r'作为单列保留项[^。\n]*?依次放在\s*A(\d+):A(\d+)',text)
    kept_columns=re.search(r'名称和原\s*A\s*的\s*([A-Z、，,]+)\s*内容',text)
    prior_column=re.search(r'原\s*A\s*上期值[^。\n]*?保留在\s*([A-Z]+)\s*列',text)
    total=re.search(r'合计移至\s*A(\d+)',text)
    style=re.search(r'沿用\s*A\s*第(\d+)行样式',text)
    old_label=next((match for match in re.finditer(r'A原明细为\s*([^\s!]+)!([A-Z]+)(\d+):\2(\d+)',previous)
                    if members and int(match[3])==members[0][1]),None)
    source_label=re.search(r'名称取\s*([A-Z]+)\s*列',combined)
    source_current=re.search(r'本期金额取\s*([A-Z]+)\s*列',combined)
    if all((members,retained,kept_range,kept_columns,prior_column,total,style,old_label,source_label,source_current)):
        sheet,label,first,last=old_label[1],old_label[2],int(old_label[3]),int(old_label[4])
        areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
               if area['sheet']==sheet and range_boundaries(area['range'])[1]<first<=last<range_boundaries(area['range'])[3]]
        source_sheets={area['sheet'] for area in source_ranges if all(
            source_label[1]+str(row) in addresses(area['range']) and source_current[1]+str(row) in addresses(area['range']) for row,_,_ in members)}
        columns={re.match(r'[A-Z]+',cell)[0] for table in a.get('outline',{}).get('subtables',[]) if table['sheet']==sheet
                 for item in table.get('metrics',[]) if item.get('template_period_role') in ('current','本期')
                 for span in item['ranges'] if span!='*' for cell in addresses(span)
                 if coordinate_to_tuple(cell)[0]==first}
        old_members={}
        for loc,record in _records(a):
            sh,cell=loc.rsplit('!',1);row=coordinate_to_tuple(cell)[0]
            value=record.get('dimensions',{}).get(dimension)
            if sh==sheet and first<=row<=last and isinstance(value,dict) and value.get('mode')=='members' and len(value.get('members',[]))==1:
                old_members.setdefault(value['members'][0],set()).add(row)
        kept_names=[value.strip() for value in retained[1].split('、')]
        same_pairs={int(target):member.strip() for member,target in re.findall(r'([^：:、。\n]+)→A(\d+)',text)}
        old_names=[*kept_names,*same_pairs.values()]
        known=all(len(old_members.get(member,()))==1 for member in old_names)
        if len(areas)==len(source_sheets)==len(columns)==1 and known:
            name,area=areas[0];top=range_boundaries(area['range'])[1];kept_first,kept_last=map(int,kept_range.groups())
            if ([target for _,target,_ in members]==list(range(first,kept_first))
                and kept_last-kept_first+1==len(kept_names) and int(total[1])==kept_last+1
                and all(member==same_pairs[target] for _,target,member in members if target in same_pairs)
                and '新增成员无对应旧值的保持空白' in combined):
                retained_values=[label,*re.findall(r'[A-Z]+',kept_columns[1])]
                return [{'selection_dimension':dimension,'range_name':name,'source_sheet':next(iter(source_sheets)),
                    'source_rows':[source for source,_,_ in members],'first_row_offset':first-top,'old_detail_count':last-first+1,
                    'columns':{next(iter(columns)):source_current[1]},'labels':{label:source_label[1]},
                    'blank_values':{str(source):{prior_column[1]:({'source':sheet+'!'+prior_column[1]+str(next(iter(old_members[same_pairs[target]])))}
                        if target in same_pairs else {'value':None})} for source,target,_ in members},
                    'retained_members':[{'index':kept_first-first+index,'cells':{col:{'source':sheet+'!'+col+str(next(iter(old_members[member])))}
                        for col in retained_values}} for index,member in enumerate(kept_names)],
                    'operations':[{'axis':'row','action':'insert','index':last+1-top,
                        'count':kept_last-last,'template_index':int(style[1])-top}]}]
    old=re.search(r'A原有明细为\s*([^\s!]+)!([A-Z]+)(\d+):\2(\d+)',text)
    insertion=re.search(r'第(\d+)行用作该明细[，,]\s*将合计移至第(\d+)行',text)
    if old and insertion and int(insertion[2])==int(insertion[1])+1 and '沿用明细样式行' in text:
        sheet,label,first,last=old[1],old[2],int(old[3]),int(old[4]);point=int(insertion[1])
        areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
               if area['sheet']==sheet and range_boundaries(area['range'])[1]<first<=last<point<=range_boundaries(area['range'])[3]]
        new_members=[]
        for source_sheet in {r['sheet'] for r in source_ranges}:
            for refs,member in re.findall(re.escape(source_sheet)+r'!([A-Z]+\d+(?:、[A-Z]+\d+)*)\s*为[“「"]([^”」"]+)[”」"][^。\n]*?新成员',text):
                new_members.append((source_sheet,refs.split('、'),member))
        if len(areas)==len(new_members)==1:
            name,area=areas[0];top=range_boundaries(area['range'])[1];source_sheet,refs,member=new_members[0]
            relations=source_relation(a,{**(b or {}),'target_table':None,'target_tables':[],'b_source_ranges':source_ranges},
                reply,[sheet+'!'+cell for cell in addresses(area['range'])])
            columns={re.match(r'[A-Z]+',row['target'].rsplit('!',1)[1])[0]:re.match(r'[A-Z]+',row['source'].rsplit('!',1)[1])[0]
                     for row in relations['bindings'] if coordinate_to_tuple(row['target'].rsplit('!',1)[1])[0]==first and row['source'].rsplit('!',1)[0]==source_sheet}
            sources={re.match(r'[A-Z]+',ref)[0]:ref for ref in refs}
            if columns and all(col in sources for col in columns.values()) and len({coordinate_to_tuple(ref)[0] for ref in refs})==1:
                row=coordinate_to_tuple(refs[0])[0]
                return [{'selection_dimension':dimension,'range_name':name,'source_sheet':source_sheet,
                    'source_rows':[row],'target_row_offsets':[point-top],'old_detail_count':0,'first_row_offset':point-top,'columns':columns,
                    'labels':{label:None},'declared_labels':{str(row):{'value':member}},
                    'blank_values':{str(row):{col:{'value':None} for col in columns}},
                    'operations':[{'axis':'row','action':'insert','index':point-top,'count':1,'template_index':first-top}]}]
    if len({r['sheet'] for r in source_ranges})==1:
        source_sheet=source_ranges[0]['sheet']
        for name,areas in a['机械记录'].get('names',{}).items():
            if len(areas)!=1:continue
            area=areas[0];sheet=area['sheet'];left,top,right,bottom=range_boundaries(area['range'])
            style=re.search(r'新增的([一二两三四五六七八九十]|\d+)行沿用原\s*'+re.escape(sheet)+r'!([A-Z]+)(\d+)',text)
            if not style or not top<=int(style[3])<=bottom:continue
            added=int(style[1]) if style[1].isdigit() else '一二三四五六七八九十'.index(style[1].replace('两','二'))+1
            columns={target:source for source,target in re.findall(r'([A-Z]+)列[“「"][^”」"]+[”」"]搬至\s*'+re.escape(sheet)+r'\s*的([A-Z]+)列',text)}
            if not columns:continue
            pairs=[];expanded=None;kept=[]
            pattern=r'(?<![A-Za-z0-9])(?:'+re.escape(source_sheet)+r'!)?([A-Z]+)(\d+)(?::\1(\d+))?[^。；;\n]*?对应\s*(?:扩展后(?:的)?\s*)?'+re.escape(sheet)+r'!([A-Z]+)(\d+)(?::\4(\d+))?'
            for match in re.finditer(pattern,text):
                sf,sl,tf,tl=int(match[2]),int(match[3] or match[2]),int(match[5]),int(match[6] or match[5])
                if sl-sf!=tl-tf:continue
                pairs.extend((target,source,match[4],match[1]) for target,source in zip(range(tf,tl+1),range(sf,sl+1)))
                if tf==int(style[3]) and tl>tf:expanded=(tf,tl)
                if re.match(r'[，,]\s*A的原空白保留',text[match.end():]):kept.append((tf,sf))
            if not pairs or not expanded:continue
            pairs=sorted(set(pairs));labels={(target_col,source_col) for _,_,target_col,source_col in pairs}
            if len(labels)!=1:continue
            first=expanded[0];old_count=expanded[1]-first+1-added
            if old_count<1:continue
            return [{'selection_dimension':dimension,'range_name':name,'source_sheet':source_sheet,
                'source_rows':[source for _,source,_,_ in pairs],'target_row_offsets':[target-top for target,_,_,_ in pairs],
                'first_row_offset':pairs[0][0]-top,'old_detail_count':len(pairs)-added,
                'columns':columns,'labels':dict(labels),
                'operations':[{'axis':'row','action':'insert','index':first+old_count-top,'count':added,'template_index':int(style[3])-top}],
                'row_columns':{str(source):{} for _,source in kept},
                'blank_values':{str(source):{column:{'source':sheet+'!'+column+str(target)} for column in columns} for target,source in kept}}]
    retained_total=re.search(r'沿用A原合计行\s*([A-Z]+)(\d+):([A-Z]+)\2的值与空白',text)
    excluded_total=re.search(r'不搬入B(\d+)',text)
    if retained_total and excluded_total and len({r['sheet'] for r in source_ranges})==1:
        rows=[];header=None
        for line in text.splitlines():
            if not line.strip().startswith('|'):continue
            values=[v.strip() for v in line.strip().strip('|').split('|')]
            if 'A位置' in values and 'B名称来源' in values:header=values;continue
            if not header or len(values)!=len(header):continue
            target=re.fullmatch(r'A(\d+)',values[header.index('A位置')])
            source=re.fullmatch(r'([A-Z]+)(\d+)',values[header.index('B名称来源')].rsplit('!',1)[-1])
            if not target or not source:continue
            columns={}
            for value in values:
                if '→' not in value:continue
                left,right=value.split('→',1)
                left=left.replace(source_ranges[0]['sheet']+'!','')
                old_cells=re.findall(r'[A-Z]+\d+',left);new_cells=re.findall(r'[A-Z]+\d+',right)
                if len(old_cells)==len(new_cells):
                    columns.update((re.match(r'[A-Z]+',new)[0],re.match(r'[A-Z]+',old)[0])
                        for old,new in zip(old_cells,new_cells))
            rows.append((int(target[1]),int(source[2]),source[1],columns))
        if rows:
            first=rows[0][0];old_total=int(retained_total[2])
            areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
                if range_boundaries(area['range'])[1]<first<old_total==range_boundaries(area['range'])[3]]
            if len(areas)==1 and [r[0] for r in rows]==list(range(first,first+len(rows))):
                from openpyxl.utils.cell import get_column_letter
                name,area=areas[0];left,top,_,_=range_boundaries(area['range'])
                columns={key:value for _,_,_,cols in rows for key,value in cols.items()}
                keep=re.search(r'不搬入A表([A-Z、]+)列',text)
                kept_columns=re.findall(r'[A-Z]+',keep[1]) if keep else []
                retained={new:old for old_values,new_values in re.findall(
                    r'原([A-Z]+\d+(?:、[A-Z]+\d+)*)到新([A-Z]+\d+(?:、[A-Z]+\d+)*)',text)
                    for old,new in zip(old_values.split('、'),new_values.split('、'))}
                blank_values={str(source):{col:({'source':area['sheet']+'!'+retained[col+str(target)]}
                    if col+str(target) in retained else {'value':None})
                    for col in [*kept_columns,*[col for col in columns if col not in cols]]}
                    for target,source,_,cols in rows}
                total_columns=[re.match(r'[A-Z]+',cell)[0] for cell in addresses(retained_total[0].split('合计行',1)[1].split('的',1)[0])]
                total_source=int(excluded_total[1])
                blank_values[str(total_source)]={col:{'source':area['sheet']+'!'+col+str(old_total)} for col in total_columns}
                return [{'selection_dimension':dimension,'range_name':name,'source_sheet':source_ranges[0]['sheet'],
                    'source_rows':[r[1] for r in rows],'first_row_offset':first-top,'old_detail_count':old_total-first,
                    'columns':columns,'row_columns':{**{str(source):cols for _,source,_,cols in rows},str(total_source):{}},
                    'labels':{get_column_letter(left):rows[0][2]},'blank_values':blank_values,
                    'totals':[{'target_row':old_total,'source_row':total_source}]}]
    ordered=re.search(r'按\s*B\s*表第\s*(\d+)[–—-](\d+)\s*行顺序写入\s*A\s*表\s*([A-Z]+)(\d+):\3(\d+)',text)
    original_total=re.search(r'A原合计内容从\s*([A-Z]+\d+(?:、[A-Z]+\d+)*)\s*保留在合计位置\s*([A-Z]+\d+(?:、[A-Z]+\d+)*)',text)
    if ordered and original_total and len({r['sheet'] for r in source_ranges})==1:
        sf,sl,first,last=map(int,ordered.group(1,2,4,5));label=ordered[3]
        old_totals=[coordinate_to_tuple(cell)[0] for cell in original_total[1].split('、')]
        names={int(source):name.strip('。 ') for source,_,name in re.findall(r'B(\d+)→A(\d+)\s+([^；。\n]+)',text)}
        pairs=re.findall(r'B\s*表\s*([A-Z]+)(\d+):\1(\d+)\s*→\s*A\s*表\s*([A-Z]+)(\d+):\4(\d+)',text)
        columns={target:source for source,start,end,target,tstart,tend in pairs
                 if (int(start),int(end),int(tstart),int(tend))==(sf,sl,first,last)}
        total=re.search(r'合计[：:]\s*B(\d+)\s*名称\s*→\s*A\s*表\s*[A-Z]+(\d+)',text)
        areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
               if range_boundaries(area['range'])[1]<first<=range_boundaries(area['range'])[3]
               and len(set(old_totals))==1 and old_totals[0]==range_boundaries(area['range'])[3]]
        if total and len(areas)==1 and sl-sf==last-first and set(names)==set(range(sf,sl+1)) and columns:
            name,area=areas[0];top=range_boundaries(area['range'])[1]
            table={'selection_dimension':dimension,'range_name':name,'source_sheet':source_ranges[0]['sheet'],
                'source_rows':list(range(sf,sl+1)),'first_row_offset':first-top,'old_detail_count':old_totals[0]-first,
                'labels':{label:None},'declared_labels':{str(row):{'value':names[row]} for row in names},'columns':columns,
                'totals':[{'target_row':old_totals[0],'source_row':int(total[1])}]}
            kept_values={}
            for old,new in re.findall(r'从\s*([A-Z]+\d+(?:、[A-Z]+\d+)*)\s*保留至\s*([A-Z]+\d+(?:、[A-Z]+\d+)*)',text):
                left,right=old.split('、'),new.split('、')
                if len(left)==len(right):kept_values.update(zip(right,left))
            empty=bool(re.search(r'其他新增成员[^。\n]*金额留空',text))
            table['blank_values']={str(row):{col:({'source':area['sheet']+'!'+kept_values[col+str(first+row-sf)]}
                if col+str(first+row-sf) in kept_values else {'value':None})
                for col in columns if empty or col+str(first+row-sf) in kept_values} for row in table['source_rows']}
            return [table]
    kept=re.search(r'沿用已保存安排[：:]\s*B表第(\d+)[–—-](\d+)行[^。\n]*?搬至A表第(\d+)[–—-](\d+)行',text)
    added=re.search(r'名称从B表([A-Z]+)(\d+)搬至A表([A-Z]+)(\d+)',text)
    if kept and added and b and len({r['sheet'] for r in source_ranges})==1:
        start,end,target_start,target_end=map(int,kept.groups())
        prior=[table for relation in b.get('已确认关系',[]) if relation.get('dimension')==dimension
               for table in relation.get('tables',[]) if table['source_rows']==list(range(start,end+1))
               and table['source_sheet']==source_ranges[0]['sheet']]
        prior=list({json.dumps(table,sort_keys=True,ensure_ascii=False):table for table in prior}.values())
        amount_pairs=re.findall(r'从([A-Z]+)(\d+)搬至([A-Z]+)(\d+)',text)
        columns={target:source for source,row,target,target_row in amount_pairs
                 if row==added[2] and target_row==added[4] and target!=added[3]}
        if len(prior)==1 and end-start==target_end-target_start and int(added[4])==target_end+1 and set(columns)==set(prior[0]['columns']):
            table=deepcopy(prior[0]);source_row=int(added[2]);table['source_rows'].append(source_row)
            table.setdefault('row_columns',{})[str(source_row)]=columns
            table.setdefault('declared_labels',{})[str(source_row)]={'source':table['source_sheet']+'!'+added[1]+added[2]}
            style=re.search(r'扩展明细行沿用A表第(\d+)行样式',text)
            area=a['机械记录']['names'][table['range_name']][0];top=range_boundaries(area['range'])[1]
            if style:
                table['operations']=[{'axis':'row','action':'insert','index':table['first_row_offset']+table['old_detail_count'],
                    'count':len(table['source_rows'])-table['old_detail_count'],'template_index':int(style[1])-top}]
            return [table]
    old=re.search(r'A\s*(?:旧|原)明细(?:起止)?(?:为)?\s*(?:第)?\s*(\d+)\s*[—–-]\s*(\d+)\s*行[、，,]\s*合计(?:为)?\s*(?:第)?\s*(\d+)\s*行',text)
    if not old:old=re.search(r'A原表\s*[^\n。]+?的明细为第\s*(\d+)\s*至\s*(\d+)\s*行[，,]\s*合计为第\s*(\d+)\s*行',text)
    named_old=re.search(r'明细为\s*(?:[^!\s]+!)?([A-Z]+)(\d+):[A-Z]+(\d+)[，,]\s*合计(?:行)?(?:为)?\s*(?:第)?\s*(?:[^!\s]+!)?[A-Z]*(\d+)',text)
    list_sources={r['sheet'] for r in source_ranges if r['sheet']+'!' in text}
    adopted={r['sheet'] for r in source_ranges if re.search(r'本次采用的B来源(?:为|是)\s*'+re.escape(r['sheet']),text)}
    if adopted:list_sources=adopted
    if not list_sources:list_sources={r['sheet'] for r in source_ranges}
    if (old or named_old) and len(list_sources)==1:
        first,last,total=map(int,old.groups() if old else named_old.group(2,3,4));rows=[];header=None;columns={}
        list_source=next(iter(list_sources));new_total=re.search(r'名单更新后[^。]*?合计(?:为)?\s*第\s*(\d+)\s*行',text)
        layout=re.search(r'A第(\d+)[–—至-](\d+)行排B的(\d+)项',text)
        ordered_columns=re.search(r'A的\s*([A-Z]+(?:/[A-Z]+)+)\s*列分别',text)
        new_total_row=re.search(r'A第(\d+)行为合计[^。\n]*',text)
        if layout and ordered_columns and new_total_row:
            target_columns=ordered_columns[1].split('/');source_rows=[];row_sources={};declared_labels={};matrix=False
            for line in text.splitlines():
                values=[value.strip() for value in line.strip().strip('|').split('|')]
                if len(values)==3 and values[0]=='B行' and '名称' in values[1]:matrix=True;continue
                if not matrix or len(values)!=3 or not values[0].isdigit():continue
                refs=re.findall(r'[A-Z]+\d+',values[2])
                if len(refs)!=len(target_columns):continue
                source_row=int(values[0]);source_rows.append(source_row)
                row_sources[str(source_row)]={col:list_source+'!'+ref for col,ref in zip(target_columns,refs)}
                declared_labels[str(source_row)]={'value':values[1]}
            total_refs=re.findall(r'[A-Z]+\d+',new_total_row[0]);areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
                if range_boundaries(area['range'])[1]<first<=last<total<=range_boundaries(area['range'])[3]]
            if len(areas)==1 and len(source_rows)==int(layout[3])==int(layout[2])-int(layout[1])+1 and len(total_refs)==len(target_columns):
                name,area=areas[0];left,top,right,bottom=range_boundaries(area['range']);sheet=area['sheet']
                cells=next(s['cells'] for s in a['机械记录']['sheets'] if s['name']==sheet);kept=[]
                for target,label in re.findall(r'A第(\d+)行保留[“"]([^”"]+)[”"]',text):
                    originals=[row for row in range(first,last+1) if str(cells.get(get_column_letter(left)+str(row),{}).get('value') or '').strip()==label]
                    if len(originals)==1:kept.append({'index':int(target)-first,'cells':{get_column_letter(col):{'source':sheet+'!'+get_column_letter(col)+str(originals[0])} for col in range(left,right+1)}})
                added=len(source_rows)+len(kept)-(last-first+1)
                if added>=0 and int(new_total_row[1])==total+added:
                    source_total=coordinate_to_tuple(total_refs[0])[0]
                    row_sources[str(source_total)]={col:list_source+'!'+ref for col,ref in zip(target_columns,total_refs)}
                    return [{'selection_dimension':dimension,'range_name':name,'source_sheet':list_source,'source_rows':source_rows,
                        'old_detail_count':last-first+1,'first_row_offset':first-top,'columns':{col:re.match(r'[A-Z]+',ref)[0] for col,ref in zip(target_columns,total_refs)},
                        'row_sources':row_sources,'labels':{get_column_letter(left):None},'declared_labels':declared_labels,'retained_members':kept,
                        'totals':[{'target_row':total,'source_row':source_total}],
                        'operations':[{'axis':'row','action':'insert','index':total-top,'count':added,'template_index':first-top}] if added else []}]
        columns.update(re.findall(r'A的([A-Z]+)列[“"][^”"]+[”"]取B的([A-Z]+)列',text))
        for line in text.splitlines():
            listed=re.match(r'\s*-\s*A第(\d+)行[：:]\s*([^，,]+)[，,]\s*金额来源为B\s*'+re.escape(list_source)+r'!([A-Z]+)(\d+):([A-Z]+)(\d+)',line)
            if listed and listed[4]==listed[6]:
                rows.append((int(listed[1]),listed[2].strip(),int(listed[4])))
            if not line.startswith('|'):continue
            cells=[v.strip() for v in line.strip('|').split('|')]
            labels=[re.sub(r'\s+|的','',v) for v in cells];labels=['项目' if v=='名称' else v for v in labels]
            if '项目' in labels and ('更新后A行' in labels or any('B' in v and '来源' in v and '→' in v for v in labels)):
                header=labels;continue
            if header and len(cells)==len(header) and any('→' in value for value in cells):
                pairs=[re.search(re.escape(list_source)+r'!([A-Z]+)(\d+)\s*→\s*A\s*[^|]*?([A-Z]+)(\d+)\s*$',value) for value in cells[1:]]
                if pairs and all(pairs) and len({p[2] for p in pairs})==len({p[4] for p in pairs})==1:
                    rows.append((int(pairs[0][4]),cells[header.index('项目')],int(pairs[0][2])))
                    columns.update((p[3],p[1]) for p in pairs)
                continue
            if not header or len(cells)!=len(header) or not cells[0].isdigit():continue
            sources=[re.match(r'([A-Z]+)(\d+)(?=[:：]|$)',value) for key,value in zip(header,cells)
                     if key.startswith('B') and '来源' in key or key.endswith('金额来源')]
            if sources and all(sources) and len({m[2] for m in sources})==1:
                rows.append((int(cells[0]),cells[header.index('项目')],int(sources[0][2])))
        for match in re.finditer(r'([A-Z]+)\d+/([A-Z]+)\d+取(?:B\s*)?([A-Z]+)\d+/([A-Z]+)\d+',text):
            columns.update(zip(match.group(1,2),match.group(3,4)))
        areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
               if range_boundaries(area['range'])[1]<first<=last<total<=range_boundaries(area['range'])[3]]
        if not columns and len(areas)==1 and b:
            area=areas[0][1];targets=[area['sheet']+'!'+cell for cell in addresses(area['range'])]
            pairs=source_relation(a,{**b,'target_tables':[{'sheet':source_ranges[0]['sheet']}]},reply,targets)['bindings']
            choices={}
            for pair in pairs:
                target=re.match(r'[A-Z]+',pair['target'].rsplit('!',1)[1])[0]
                choices.setdefault(target,set()).add(re.match(r'[A-Z]+',pair['source'].rsplit('!',1)[1])[0])
            columns={target:next(iter(values)) for target,values in choices.items() if len(values)==1}
        totals=[row for row in rows if row[0]==int(new_total[1])] if new_total else [row for row in rows if row[1]=='合计']
        details=[row for row in rows if row not in totals]
        if len(areas)==1 and len(totals)==1 and details and columns:
            from openpyxl.utils.cell import get_column_letter
            name,area=areas[0];left,top,_,_=range_boundaries(area['range'])
            table={'selection_dimension':dimension,'range_name':name,'source_sheet':list_source,
                'source_rows':[row[2] for row in details],'old_detail_count':last-first+1,'first_row_offset':first-top,
                'columns':columns,'labels':{get_column_letter(left):None},
                'declared_labels':{str(row[2]):{'value':row[1]} for row in details},
                'totals':[{'target_row':total,'source_row':totals[0][2]}]}
            insertion=re.search(r'在\s*A\s*第\s*(\d+)\s*行新增',text)
            if insertion and len(details)>last-first+1:
                table['operations']=[{'axis':'row','action':'insert','index':int(insertion[1])-top,
                    'count':len(details)-(last-first+1),'template_index':first-top}]
            return [table]
    named_sources='|'.join(re.escape(s['sheet']) for s in sorted(source_ranges,key=lambda s:len(s['sheet']),reverse=True))
    adopted=re.search(r'采用来源[：:]\s*B\s+('+named_sources+r')(?=[，。\s]|$)',text) if named_sources else None
    if adopted:source_ranges=[s for s in source_ranges if s['sheet']==adopted[1]]
    old=re.search(r'A\s*原有明细为\s*A!([A-Z]+)(\d+):\1(\d+)[^\n。]*?合计为\s*A!\1(\d+)',text)
    members=re.search(r'B\s*实际[^\n。]*?名单更新为(?:上述)?第\s*(\d+)\s*[–—至-]\s*(\d+)\s*行[，,]\s*合计为第\s*(\d+)\s*行',text)
    pair=re.search(r'金额列对应\s*A!([A-Z]+\d+:[A-Z]+\d+)\s*→\s*([^\n!]+)!([A-Z]+\d+:[A-Z]+\d+)',text)
    if old and members and pair and len(source_ranges)==1:
        source_sheet=source_ranges[0]['sheet'];first,last,total=map(int,old.groups()[1:])
        source_first,source_last,source_total=map(int,members.groups())
        target_cells,source_cells=addresses(pair[1]),addresses(pair[3])
        areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
               if range_boundaries(area['range'])[1]<first<=last<total<=range_boundaries(area['range'])[3]]
        labels={int(re.search(r'\d+$',cell)[0]):label for cell,label in re.findall(
            r'\b([A-Z]+\d+):[A-Z]+\d+\s*[“「"]([^”」"]+)[”」"]',text)}
        rows=list(range(source_first,source_last+1))
        if len(areas)==1 and pair[2].strip()==source_sheet and len(target_cells)==len(source_cells) and all(row in labels for row in rows):
            name,area=areas[0]
            return [{'selection_dimension':dimension,'range_name':name,'source_sheet':source_sheet,
                'source_rows':rows,'old_detail_count':last-first+1,'first_row_offset':first-range_boundaries(area['range'])[1],
                'columns':{re.match(r'[A-Z]+',target)[0]:re.match(r'[A-Z]+',source)[0] for target,source in zip(target_cells,source_cells)},
                'labels':{old[1]:None},'declared_labels':{str(row):{'value':labels[row]} for row in rows},
                'totals':[{'target_row':total,'source_row':source_total}]}]
    text=reply.replace('`','').replace('**','')
    # 分段名单的旧行、新行和列对应均取自模型答复；只展开已写明的安排。
    segments=[]
    for line in text.splitlines():
        pairs=[(int(source),int(target),name.strip()) for source,target,name in re.findall(
            r'B(\d+)\s*→\s*(\d+)\s+([^；|]+)',line)]
        if pairs:segments.append({'pairs':pairs,'columns':{target:source for source,target in re.findall(r'\b([A-Z])→([A-Z])\b',line)}})
        span=re.search(r'B(\d+):B(\d+)[^。\n]*?依次对应\s*A(\d+):A(\d+)',line)
        total=re.search(r'B(\d+)\s*为合计[^；。\n]*?A(\d+)',line)
        if span and total:
            sf,sl,tf,tl=map(int,span.groups())
            if sl-sf==tl-tf:segments.append({'pairs':[(s,t,None) for s,t in zip(range(sf,sl+1),range(tf,tl+1))],'source_total':int(total[1])})
    if segments and len({area['sheet'] for area in source_ranges})==1:
        styles=re.search(r'分别以([^。；\n]+)行的样式',text) or re.search(r'样式行[（(]原\s*([^）)]+)',text)
        starts=[int(n) for n in re.findall(r'\d+',styles[1])] if styles else []
        areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
               if starts and range_boundaries(area['range'])[1]<min(starts)<=max(starts)<=range_boundaries(area['range'])[3]]
        if len(starts)==len(segments) and len(areas)==1:
            from openpyxl.utils.cell import get_column_letter
            name,area=areas[0];left,top,right,bottom=range_boundaries(area['range'])
            source_sheet=source_ranges[0]['sheet'];prior='\n'.join(p['识别答复'] for p in (b or {}).get('合集来源',[])
                if source_sheet in p.get('识别答复',''))
            columns={};label_column=None
            match=re.search(r'A\s*表金额列\s*([A-Z、]+)\s*分别对应\s*B\s*表金额列\s*([A-Z、]+)',prior)
            if match:
                target,source=match[1].split('、'),match[2].split('、')
                if len(target)==len(source):columns=dict(zip(target,source))
            match=re.search(r'A\s*表中项目列\s*([A-Z]+)\s*对应\s*B\s*表项目列\s*([A-Z]+)',prior)
            if match:label_column=(match[1],match[2])
            tables=[]
            for index,(start,segment) in enumerate(zip(starts,segments)):
                mapping=segment.get('columns') or columns
                if not mapping:break
                end=starts[index+1]-2 if index+1<len(starts) else bottom
                table={'selection_dimension':dimension,'range_name':name,'source_sheet':source_sheet,
                    'source_rows':[s for s,_,_ in segment['pairs']],'old_detail_count':end-start+1,
                    'first_row_offset':start-top,'columns':mapping,'labels':{}}
                if all(label is not None for _,_,label in segment['pairs']):
                    table['labels']={get_column_letter(left):None}
                    table['declared_labels']={str(s):{'value':label} for s,_,label in segment['pairs']}
                elif label_column:table['labels']=dict([label_column])
                else:break
                if segment.get('source_total'):
                    table['totals']=[{'target_row':start-1,'source_row':segment['source_total']}]
                tables.append(table)
            if len(tables)==len(segments):return tables
    # 自然表格中的明确箭头逐项转存；A/B 文件简称不当作实际列名。
    sources={area['sheet'] for area in source_ranges}
    details=re.search(r'明细扩为\s*([A-Z]+)(\d+):\1(\d+)',text)
    amounts=re.search(r'金额列分别落在\s*A\s*表\s*([A-Z]+):([A-Z]+)',text)
    total=re.search(r'合计取自\s*B\s*表第\s*(\d+)\s*行[，,]\s*放在\s*A\s*表第\s*(\d+)\s*行',text)
    rows=re.findall(r'(?:B\s*表\s*)?第\s*(\d+)\s*行[^。\n]*?[（(]([A-Z]+\d+:[A-Z]+\d+)[）)][^。\n]*?对应\s*A\s*第\s*(\d+)\s*行',text)
    if details and amounts and total and rows and len(sources)==1:
        from openpyxl.utils.cell import get_column_letter
        first,last=int(details[2]),int(details[3]);pairs=sorted((int(target),int(source),span) for source,span,target in rows)
        areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
               if area.get('range') and range_boundaries(area['range'])[1]<first<=range_boundaries(area['range'])[3]]
        prior='\n'.join(part['识别答复'] for part in a.get('合集来源',[]))
        if len(areas)==1 and int(total[2])==last+1 and [target for target,_,_ in pairs]==list(range(first,last+1)):
            name,area=areas[0];left,top,right,bottom=range_boundaries(area['range'])
            old_totals={int(value) for value in re.findall(r'第(\d+)行(?:为)?(?:该[^。\n；;]*?范围的)?合计',prior) if first<=int(value)<=bottom}
            source_columns=[list(dict.fromkeys(re.match(r'[A-Z]+',cell)[0] for cell in addresses(span))) for _,_,span in pairs]
            target_columns=[details[1],*[re.match(r'[A-Z]+',cell)[0] for cell in addresses(amounts[1]+'1:'+amounts[2]+'1')]]
            if len(old_totals)==1 and all(cols==source_columns[0] for cols in source_columns) and len(source_columns[0])==len(target_columns):
                old_total=next(iter(old_totals));columns=dict(zip(target_columns,source_columns[0]));label=details[1]
                return [{'selection_dimension':dimension,'range_name':name,'source_sheet':next(iter(sources)),
                    'source_rows':[source for _,source,_ in pairs],'old_detail_count':old_total-first,
                    'first_row_offset':first-top,'labels':{label:columns.pop(label)},'columns':columns,
                    'totals':[{'target_row':old_total,'source_row':int(total[1])}]}]
    text=re.sub(r'→\s*A\s*表\s*(?=[A-Z]+\d)', '→ ', text)
    row_pairs=[];column_pairs={};total_pair=None;name_columns=None;name_index=None
    layout_header=None;layout_sources=set();role_columns={}
    for line in text.splitlines():
        backward=re.search(r'\bA(\d+)[“"]([^”"]+)[”"]\s*←\s*B表\s*([A-Z]+)(\d+)',line)
        if backward:
            target,source=int(backward[1]),int(backward[4])
            if backward[2]=='合计':total_pair=(target,source)
            else:row_pairs.append((target,source))
            old_label=re.search(r'A\s*原明细[：:]\s*([A-Z]+)\d+',text)
            if old_label:name_columns=(old_label[1],backward[3])
        if 'B表' in line and '列→' in line:
            column_pairs.update((target,source) for source,target in re.findall(r'([A-Z]+)列→(?:A表\s*)?([A-Z]+)列',line))
        values=[value.strip() for value in line.strip().strip('|').split('|')]
        if 'A版式位置' in values and 'B来源名称' in values:
            layout_header=values;continue
        if layout_header and len(values)==len(layout_header):
            target=re.fullmatch(r'([A-Z]+)(\d+)',values[layout_header.index('A版式位置')])
            source=re.fullmatch(r'([^!]+)!([A-Z]+)(\d+)',values[layout_header.index('B来源名称')])
            if target and source:
                first=int(target[2]) if not row_pairs else row_pairs[0][0]
                if not role_columns:
                    for table in a.get('outline',{}).get('subtables',[]):
                        for metric in table.get('metrics',[]):
                            role=metric.get('template_period_role')
                            if not isinstance(role,str):continue
                            for span in metric['ranges']:
                                for cell in addresses(table['range'] if span=='*' else span):
                                    if coordinate_to_tuple(cell)[0]==first:
                                        role_columns.setdefault(role,set()).add(re.match(r'[A-Z]+',cell)[0])
                for title,value in zip(layout_header,values):
                    role=next((role for label,role in [('期末','closing'),('期初','opening'),('本期','current'),('上期','prior')]
                               if label in title and '金额来源' in title),None)
                    ref=re.fullmatch(re.escape(source[1])+r'!([A-Z]+)'+source[3],value)
                    if ref and len(role_columns.get(role,()))==1:
                        column_pairs[next(iter(role_columns[role]))]=ref[1]
                name_columns=(target[1],source[2]);layout_sources.add(source[1])
                if '合计' in values:total_pair=(int(target[2]),int(source[3]))
                else:row_pairs.append((int(target[2]),int(source[3])))
        if '名称来源' in values:name_index=values.index('名称来源')
        elif name_index is not None and len(values)>name_index:
            label_pair=re.search(r'([A-Z]+)\d+\s*→\s*([A-Z]+)\d+',values[name_index])
            if label_pair:name_columns=(label_pair[2],label_pair[1])
        row_only=re.search(r'B!(\d+)\s*→\s*A!(\d+)',line)
        if row_only:row_pairs.append((int(row_only[2]),int(row_only[1])))
        pairs=re.findall(r'(?<![A-Za-z0-9])([A-Z]+\d+(?::[A-Z]+\d+)?)(?:[「“][^」”]+[」”])?\s*(?:→|合计对应)\s*(?:[^\s!|]+!)?([A-Z]+\d+(?::[A-Z]+\d+)?)',line)
        for source,target in pairs:
            scells,tcells=addresses(source),addresses(target)
            if len(scells)!=len(tcells):continue
            for scell,tcell in zip(scells,tcells):
                sr,sc=coordinate_to_tuple(scell);tr,tc=coordinate_to_tuple(tcell)
                column_pairs[re.match(r'[A-Z]+',tcell)[0]]=re.match(r'[A-Z]+',scell)[0]
                if '合计' in line:total_pair=(tr,sr)
                else:row_pairs.append((tr,sr))
            if '名称' in line and len(scells)==len(tcells)==1:
                name_columns=(re.match(r'[A-Z]+',tcells[0])[0],re.match(r'[A-Z]+',scells[0])[0])
        columns=re.search(r'B表\s*(.+?)分别对应A表\s*([^。]+)',line)
        if columns:
            left=re.findall(r'([A-Z]+)列',columns[1]);right=re.findall(r'([A-Z]+)列',columns[2])
            if len(left)==len(right):column_pairs.update(zip(right,left))
            sn=re.search(r'([A-Z]+)列名称',columns[1]);tn=re.search(r'([A-Z]+)列名称',columns[2])
            if sn and tn:name_columns=(tn[1],sn[1])
    # 表头已写“名称：来源→位置”，其首个箭头列就是明确的名称列。
    if not name_columns and re.search(r'名称[：:]B来源\s*→\s*A位置',text):
        first=re.search(r'\|\s*明细[^|]*\|\s*[^|]*?([A-Z]+)\d+[「“][^」”]+[」”]\s*→\s*[^!|]+!([A-Z]+)\d+',text)
        if first:name_columns=(first[2],first[1])
    row_pairs=sorted(set(row_pairs))
    if layout_sources:sources=layout_sources
    if row_pairs and column_pairs and name_columns and total_pair and len(sources)==1:
        first=row_pairs[0][0]
        areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
               if area.get('range') and range_boundaries(area['range'])[1]<first<=range_boundaries(area['range'])[3]]
        old_total=(re.search(r'(?:原合计行|合计行[^。]*?沿用A原合计行)\s*A?(\d+)',text)
                   or re.search(r'原合计[：:]\s*[A-Z]+(\d+)',text)
                   or re.search(r'A原表[^。\n]*明细为[^。\n]*合计为\s*[A-Z]+(\d+)',text))
        if len(areas)==1:
            name,area=areas[0];top,bottom=range_boundaries(area['range'])[1::2]
            old_end=int(old_total[1]) if old_total else total_pair[0] if total_pair[0]==bottom else None
            if old_end is None and re.search(r'原合计行顺延',text):
                prior='\n'.join(part['识别答复'] for part in a.get('合集来源',[]))
                totals={int(number) for number in re.findall(r'^\|\s*`?(?:[^!|]+!)?[A-Z]+(\d+)`?\s*\|[^|\n]*合计',prior,re.M)
                        if first<=int(number)<=bottom}
                if len(totals)==1:old_end=next(iter(totals))
            if old_end and [r for r,_ in row_pairs]==list(range(first,total_pair[0])):
                label,source_label=name_columns;column_pairs.pop(label,None)
                table={'selection_dimension':dimension,'range_name':name,'source_sheet':next(iter(sources)),
                    'source_rows':[r for _,r in row_pairs],'old_detail_count':old_end-first,
                    'first_row_offset':first-top,'columns':column_pairs,'labels':{label:source_label},
                    'totals':[{'target_row':old_end,'source_row':total_pair[1]}]}
                style=re.search(r'沿用\s*A\s*表明细行\s*(\d+)\s*的样式',text)
                if style and len(row_pairs)>old_end-first:
                    table['operations']=[{'axis':'row','action':'insert','index':old_end-top,
                        'count':len(row_pairs)-(old_end-first),'template_index':int(style[1])-top}]
                return [table]
    # 模型明说整行对应时，展开该表行；列角色继续使用原来识别的数据列。
    grouped=[];names={};source_names={};whole_rows=False
    for line in text.splitlines():
        pair=re.search(r'\bA(\d+)(?::A?(\d+))?\s*(?:分别|依次)?对应\s*B(\d+)(?::B?(\d+))?',line)
        if pair:
            target_rows=list(range(int(pair[1]),int(pair[2] or pair[1])+1))
            source_rows=list(range(int(pair[3]),int(pair[4] or pair[3])+1))
            if len(target_rows)==len(source_rows):
                if '合计' in line:total_pair=(target_rows[0],source_rows[0])
                else:grouped.extend(zip(target_rows,source_rows));whole_rows=True
        values=[part.strip() for part in line.strip().strip('|').split('|')]
        row=re.fullmatch(r'第(\d+)行',values[0]) if values else None
        if row and len(values)>=4:
            amounts=re.findall(r'([A-Z]+)(\d+)',values[3])
            if amounts and len({r for _,r in amounts})==1:
                target=int(row[1]);grouped.append((target,int(amounts[0][1])));names[target]=values[1]
                label=re.fullmatch(r'([^!]+)!([A-Z]+\d+)',values[2])
                if label:source_names[target]=label[0]
                target_cols=re.search(r'分别搬到\s*A\s*的\s*([A-Z、，,\s]+)\s*列',text)
                if target_cols:
                    cols=re.findall(r'[A-Z]+',target_cols[1])
                    if len(cols)==len(amounts):column_pairs.update(zip(cols,[col for col,_ in amounts]))
        elif len(values)>=2 and '→' in values[1]:
            pair=re.search(r'([A-Z]+)(\d+):[A-Z]+\d+\s*→\s*([A-Z]+)(\d+):[A-Z]+\d+',values[1])
            if pair and values[0]!='合计':names[int(pair[4])]=values[0]
    grouped=sorted(set(grouped or row_pairs))
    if total_pair:grouped=[pair for pair in grouped if pair!=total_pair]
    old_rows=re.search(r'A\s*原(?:有明细行|第)\s*(\d+)\s*[:–—-]\s*(\d+)\s*行?',text)
    old_total=re.search(r'沿用\s*A(\d+)\s*(?:的)?合计',text)
    if grouped and len(source_ranges)==1 and (old_rows or old_total):
        first=grouped[0][0]
        areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
               if area.get('range') and range_boundaries(area['range'])[1]<first<=range_boundaries(area['range'])[3]]
        if len(areas)==1:
            name,area=areas[0];left,top,right,bottom=range_boundaries(area['range'])
            source=source_ranges[0];sl,_,sr,_=range_boundaries(source['range'])
            from openpyxl.utils.cell import get_column_letter
            if whole_rows and right-left==sr-sl:
                column_pairs={get_column_letter(col):get_column_letter(sl+col-left) for col in range(left,right+1)}
            data_cols={re.match(r'[A-Z]+',cell)[0] for table in a['outline']['subtables'] if table['sheet']==area['sheet']
                       for span in table['data_ranges'] for cell in addresses(span) if coordinate_to_tuple(cell)[0]==first}
            row_columns={get_column_letter(col) for col in range(left,right+1)}
            label_cols=row_columns-data_cols
            if len(label_cols)==1 and data_cols and data_cols<=column_pairs.keys():
                label=next(iter(label_cols));source_label=get_column_letter(sl+coordinate_to_tuple(label+'1')[1]-left)
                old_end=int(old_rows[2])+1 if old_rows else int(old_total[1])
                table={'selection_dimension':dimension,'range_name':name,'source_sheet':source['sheet'],
                    'source_rows':[row for _,row in grouped],'old_detail_count':old_end-first,'first_row_offset':first-top,
                    'columns':{col:column_pairs[col] for col in data_cols},'labels':{label:source_label},
                    'totals':[{'target_row':old_end,'source_row':total_pair[1]}] if total_pair else []}
                # 答复列出的名称直接保留；跨表名称引用不改成同一行的空白格。
                table['declared_labels']={str(source_row):({'source':source_names[target]} if target in source_names
                    else {'value':names[target]}) for target,source_row in grouped if target in names}
                return [table]
    rows=[(int(target),int(source)) for target,source in re.findall(r'A表(\d+)行[^\n]*?←\s*B表(\d+)行',text)]
    rows.extend((int(target),int(source)) for source,target in re.findall(r'\bB(\d+)\s*→\s*A(\d+)[:：]',text))
    columns={}
    for target,source in re.findall(r'A表([A-Z、，,]+)分别取B表([A-Z、，,]+)',text):
        left,right=re.findall('[A-Z]+',target),re.findall('[A-Z]+',source)
        if len(left)==len(right):columns.update(zip(left,right))
    columns.update((target,source) for source,target in re.findall(
        r'B!([A-Z]+)\d+:[A-Z]+\d+\s*→\s*A!([A-Z]+)\d+:[A-Z]+\d+',text))
    sources={area['sheet'] for area in source_ranges}
    if rows and columns and len(sources)==1:
        rows=sorted(set(rows));first=rows[0][0]
        areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
               if range_boundaries(area['range'])[1]<first<=range_boundaries(area['range'])[3]]
        prior='\n'.join(row.get('content','') for row in a.get('已确认关系',[])
                        if row.get('kind')=='名单编排' and row.get('dimension')==dimension)
        named_columns={left or right for left,right in re.findall(
            r'(?:名称在|明细名称在)\s*([A-Z]+)\s*列|([A-Z]+)列为项目名称',prior.replace('`',''))}
        total=re.search(r'合计移至A表(\d+)行，对应B表(\d+)行',text)
        total_pair=(int(total[1]),int(total[2])) if total else None
        total=re.search(r'合计从\s*B![A-Z]+(\d+)[^\n→]*→\s*A![A-Z]+(\d+)',text)
        if total:total_pair=(int(total[2]),int(total[1]))
        if len(areas)==len(named_columns)==1 and total_pair:
            name,area=areas[0];top,bottom=range_boundaries(area['range'])[1::2]
            original_totals={int(left or right) for left,right in re.findall(
                r'合计在第(\d+)行|第(\d+)行为[“"]合计[”"]',prior.replace('`',''))
                if top<=int(left or right)<=bottom}
            if len(original_totals)==1:
                old_total=next(iter(original_totals));old=old_total-first
                if [r for r,_ in rows]==list(range(first,first+len(rows))) and total_pair[0]==first+len(rows):
                    label=next(iter(named_columns));source=next(iter(sources))
                    table={'selection_dimension':dimension,'range_name':name,'source_sheet':source,
                        'source_rows':[r for _,r in rows],'old_detail_count':old,'first_row_offset':first-top,
                        'columns':columns,'labels':{label:label},
                        'totals':[{'target_row':old_total,'source_row':total_pair[1]}]}
                    style=re.search(r'沿用A表(\d+)行',text)
                    if len(rows)>old and style:
                        table['operations']=[{'action':'insert','axis':'row','index':old_total-top,
                            'count':len(rows)-old,'template_index':int(style[1])-top}]
                    return [table]
    segmented=_declared_segmented_list_tables(a,reply,dimension,b)
    if segmented:return segmented
    lines=reply.replace('`','').replace('**','').splitlines()
    arrangements=[line for line in lines if re.search(r'沿用\s*A\s*的\s*\S+\s*表区和列角色',line)]
    sheets={area['sheet'] for area in source_ranges}
    # shortcut: 只接已有明确行界的连续或分段名单；其他安排保留原文。
    if len(arrangements)!=1 or len(sheets)!=1:return []
    text=arrangements[0];source_sheet=next(iter(sheets))
    name=re.search(r'沿用\s*A\s*的\s*(\S+)\s*表区和列角色',text)
    label=re.search(r'([A-Z]+)\s*列放名单标签',text)
    first=re.search(r'从\s*([A-Z]+\d+)\s*起连续排布',text)
    old=re.search(r'A\s*当前有\s*(\d+)\s*个明细行',text)
    count=re.search(r'B\s*的\s*(\d+)\s*个具名位置接入为连续明细',text)
    def roles(line):
        found=re.search(r'([A-Z]+(?:[、，,][A-Z]+)+)\s*列(?:继续)?分别(?:为|承载)([^；;。]+)',line)
        if not found:return {}
        cols=re.split('[、，,]',found[1]);values=found[2].strip().split('、')
        return dict(zip(values,cols)) if len(cols)==len(values)==len(set(values)) else {}
    source_lines=[line for line in lines if source_sheet+'!' in line and re.search(r'名称在\s*[A-Z]+\s*列',line)]
    if not all((name,label,first,old,count)) or len(source_lines)!=1:return []
    target_roles=roles(text);source_roles=roles(source_lines[0])
    if not target_roles or target_roles.keys()!=source_roles.keys():return []
    source_label=re.search(r'名称在\s*([A-Z]+)\s*列',source_lines[0])[1]
    area=a['机械记录'].get('names',{}).get(name[1],[])
    if len(area)!=1:return []
    top=range_boundaries(area[0]['range'])[1];first_row,_=coordinate_to_tuple(first[1])
    if re.match(r'[A-Z]+',first[1])[0]!=label[1] or first_row<top:return []
    rows=[]
    allowed={cell for item in source_ranges for cell in addresses(item['range'])}
    for line in lines:
        item=re.match(r'\s*-\s*(?:([^!]+)!)?([A-Z]+\d+:[A-Z]+\d+)\s*[:：]\s*(\d+)\s*个名称',line)
        if not item or item[1] and item[1]!=source_sheet:continue
        cells=addresses(item[2])
        if len(cells)!=int(item[3]) or any(re.match(r'[A-Z]+',cell)[0]!=source_label or cell not in allowed for cell in cells):return []
        rows.extend(coordinate_to_tuple(cell)[0] for cell in cells)
    if len(rows)!=int(count[1]) or len(rows)!=len(set(rows)):return []
    return [{'selection_dimension':dimension,'range_name':name[1],'source_sheet':source_sheet,
             'source_rows':rows,'old_detail_count':int(old[1]),'first_row_offset':first_row-top,
             'columns':{target_roles[role]:source_roles[role] for role in target_roles},
             'labels':{label[1]:source_label},'source_ranges':deepcopy(source_ranges),
             'source_evidence':[source_sheet+'!'+source_label+str(row) for row in rows],
             'coverage_reason':'原答复明确的全部具名位置按顺序接入连续明细'}]


def _expand_declared_row_columns(text):
    """行、列关系分开写时只拼接答复明示的位置，字段原文继续沿用。"""
    lines=[];header=None;last_columns=None
    for line in text.splitlines():
        if not line.lstrip().startswith('|'):
            if line.lstrip().startswith('#'):header=None;last_columns=None
            lines.append(line);continue
        values=[v.strip().replace(chr(96),'') for v in line.strip().strip('|').split('|')]
        a=[i for i,v in enumerate(values) if re.match(r'^A\s*行',v)]
        b=[i for i,v in enumerate(values) if re.match(r'^B\s*(?:对应)?行',v)]
        cols=[i for i,v in enumerate(values) if re.fullmatch(r'A\s*列\s*←\s*B\s*列',v)]
        if len(a)==len(b)==len(cols)==1:
            header=(a[0],b[0],cols[0]);last_columns=None
            values[a[0]]='A位置';values[b[0]]='B来源位置';values[cols[0]]='metric_id'
            lines.append('| '+' | '.join(values)+' |');continue
        if header and len(values)>max(header):
            left,right,column=header
            ar=re.match(r'^([^!|]+)!(\d+)(?=\s|$)',values[left])
            br=re.match(r'^([^!|]+)!(\d+)(?=\s|$)',values[right])
            declaration=last_columns if values[column]=='同上' else values[column]
            pairs=re.findall(r'([A-Z]+)\s*←\s*([A-Z]+)\s*[:：]\s*(metric\.[A-Za-z0-9_.]+)',declaration or '')
            if ar and br and pairs:
                last_columns=declaration
                for ac,bc,metric in pairs:
                    expanded=values[:]
                    expanded[left]=ar[1]+'!'+ac+ar[2]
                    expanded[right]=br[1]+'!'+bc+br[2]
                    expanded[column]=metric
                    lines.append('| '+' | '.join(expanded)+' |')
                continue
        lines.append(line)
    return '\n'.join(lines)


def _source_aliases(text, mapping):
    """只接通已有工作表名的引号、唯一编号及单表B简称。"""
    sheets={t['sheet'] for t in mapping.get('target_tables',[])}
    sheets.update(t['sheet'] for t in mapping.get('b_source_ranges',[]))
    if mapping.get('target_table'):sheets.add(mapping['target_table']['sheet'])
    quoted_names=sheets|{s['name'] for s in mapping.get('机械记录',{}).get('sheets',[])}
    for sheet in sorted(quoted_names,key=len,reverse=True):
        for left,right in (('「','」'),('『','』'),('“','”'),('《','》'),('"','"'),("'","'")):
            text=re.sub(r'(?:B\s*)?'+re.escape(left+sheet+right)+r'\s*(?:!\s*`?|`?(?=\$?[A-Z]+\$?[1-9]\d*))',lambda _:sheet+'!',text)
    aliases={}
    for sheet in sheets:
        code=re.match(r'^[A-Za-z]+\d+(?=\s)',sheet)
        if code:aliases.setdefault(code[0],set()).add(sheet)
    if len(sheets)==1:aliases['B']=sheets
    # 答复明确把原文标题用作页名时，仅解析同文引用，不据标题选择业务来源。
    titles=re.findall(r'B来源(?:页名为)?[“「"]([^”」"]+)[”」"]',text)
    for title in titles:
        if title in quoted_names:continue
        names={sheet['name'] for sheet in mapping.get('机械记录',{}).get('sheets',[])
               if sheet['name'] in sheets and any(
                   isinstance(value:=cell.get('cached_value',cell.get('value')),str) and value.strip()==title
                   for cell in sheet.get('cells',{}).values())}
        if len(names)==1:aliases[title]=names
    for alias,names in aliases.items():
        if len(names)==1:text=re.sub(r'(?<![\w])'+re.escape(alias)+r'!',next(iter(names))+'!',text)
    return text


def _known_period_columns(mapping, sheet, row):
    """只读取A识别已保存的列角色，不根据表头文字另作判断。"""
    from .A语义阅读 import addresses
    from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
    aliases={'本期':'current','上期':'prior','期初':'opening','期末':'closing'};columns={}
    for table in mapping.get('outline',{}).get('subtables',[]):
        if table['sheet']!=sheet:continue
        for metric in table.get('metrics',[]):
            role=metric.get('template_period_role');role=aliases.get(role,role) if isinstance(role,str) else None
            if role not in aliases.values():continue
            for span in metric['ranges']:
                if span=='*':continue
                for cell in addresses(span):
                    r,c=coordinate_to_tuple(cell)
                    if r==row:columns.setdefault(role,set()).add(get_column_letter(c))
    return {role:next(iter(cols)) for role,cols in columns.items() if len(cols)==1}


def declared_named_row_relations(a, b, reply):
    """展开答复明说的原表名称引用、来源行及同序列关系；不按名称选择业务来源。"""
    from openpyxl.utils.cell import range_boundaries, get_column_letter, coordinate_to_tuple
    text=reply.replace('`','').replace('**','');sections=[]
    for name,areas in a.get('机械记录',{}).get('names',{}).items():
        for area in areas:
            ref=area['sheet']+'!'+area['range'].replace('$','')
            for match in re.finditer(r'(?m)^[^\n]*'+re.escape(ref)+r'[^\n]*$',text):
                sections.append((match.start(),match.end(),name,area))
    sections.sort();bindings=[];tables=[]
    extra=re.search(r'于原(?:小计|合计)\s*(\d+)\s*之前插入一条业务行[，,]名称保留空白',text)
    if extra:
        point=int(extra[1]);sources=re.findall(r'新增行(期末|期初|本期|上期)余额\s*→\s*([^!\s]+)!([A-Z]+\d+)',text)
        owners=[(name,area) for name,parts in a.get('机械记录',{}).get('names',{}).items() for area in parts
                if area['sheet']+'!'+area['range'].replace('$','') in text
                and range_boundaries(area['range'])[1]<point<=range_boundaries(area['range'])[3]]
        if sources and len(owners)==1:
            name,area=owners[0];left,top,_,_=range_boundaries(area['range']);roles=_known_period_columns(a,area['sheet'],point)
            aliases={'期末':'closing','期初':'opening','本期':'current','上期':'prior'}
            row_sources={roles[aliases[role]]:sheet+'!'+cell for role,sheet,cell in sources if aliases[role] in roles}
            if len(row_sources)==len(sources):
                source_sheet=sources[0][1];source_row=coordinate_to_tuple(sources[0][2])[0];label=get_column_letter(left)
                columns={col:re.match(r'[A-Z]+',ref.rsplit('!',1)[1])[0] for col,ref in row_sources.items()}
                tables.append({'selection_dimension':None,'range_name':name,'source_sheet':source_sheet,'source_rows':[source_row],
                    'old_detail_count':0,'first_row_offset':point-top,'target_row_offsets':[point-top],
                    'columns':columns,'row_sources':{str(source_row):row_sources},'labels':{label:None},
                    'declared_labels':{str(source_row):{'value':None}},'totals':[],
                    'blank_values':{str(source_row):{col:{'value':None} for col in columns}},
                    'operations':[{'axis':'row','action':'insert','index':point-top,'count':1,'template_index':point-top-1}]})
    for extra in re.finditer(r'于原A合计行\s*A\s*(\d+)\s*之前另列[^，。；;\n]+[，,]\s*取\s*B\s*([^!\s]+)!([A-Z]+\d+:[A-Z]+\d+)',text):
        point=int(extra[1]);sl,row,sr,last=range_boundaries(extra[3])
        if row!=last:continue
        candidates=[]
        for name,parts in a.get('机械记录',{}).get('names',{}).items():
            for area in parts:
                left,top,right,bottom=range_boundaries(area['range'])
                if bottom!=point:continue
                requested=[area['sheet']+'!'+get_column_letter(col)+str(point) for col in range(left,right+1)]
                pairs=source_relation(a,b,reply,requested)['bindings']
                sources={p['target']:p['source'] for p in pairs if p['source'].rsplit('!',1)[0]==extra[2]}
                columns={get_column_letter(col):re.match(r'[A-Z]+',sources[loc].rsplit('!',1)[1])[0]
                         for col,loc in zip(range(left,right+1),requested) if loc in sources}
                if len(columns)==len(requested) and list(columns.values())==[get_column_letter(col) for col in range(sl,sr+1)]:
                    candidates.append((name,area,left,top,columns))
        if len(candidates)!=1:continue
        name,area,left,top,columns=candidates[0];label=get_column_letter(left);source_label=columns.pop(label)
        tables.append({'selection_dimension':None,'range_name':name,'source_sheet':extra[2],
            'source_rows':[row],'old_detail_count':0,'first_row_offset':point-top,'target_row_offsets':[point-top],
            'columns':columns,'labels':{label:source_label},'totals':[],
            'blank_values':{str(row):{col:{'value':None} for col in columns}},
            'operations':[{'axis':'row','action':'insert','index':point-top,'count':1,'template_index':point-top-1}]})
    for block in re.split(r'\n(?=#{1,6}\s)',text):
        extra_lines=[line for line in block.splitlines() if re.search(r'原合计行之前|合计前',line) and re.search(r'新增|另列',line)]
        extra=re.search(r'(?:采用B|来源为)\s*(?:B\s*)?([^!\s]+)!([A-Z]+\d+:[A-Z]+\d+)',extra_lines[0]) if len(extra_lines)==1 else None
        total=re.search(r'A原\s*([^!\s]+)!([A-Z]+\d+:[A-Z]+\d+)\s*[“"]合计[”"]\s*↔\s*B\s*([^!\s]+)!([A-Z]+\d+:[A-Z]+\d+)',block)
        if not (extra and total):continue
        tl,point,tr,end=range_boundaries(total[2]);sl,source_total,sr,send=range_boundaries(total[4])
        xl,row,xr,last=range_boundaries(extra[2])
        areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
               if area['sheet']==total[1] and range_boundaries(area['range'])[1]<point==range_boundaries(area['range'])[3]]
        if len(areas)!=1 or point!=end or source_total!=send or row!=last or extra[1]!=total[3] or tr-tl!=sr-sl:continue
        name,area=areas[0];left,top,right,_=range_boundaries(area['range'])
        if tr!=right or xr!=sr or not (tl==left and xl==sl or tl==left+1 and xl==sl-1):continue
        # 新成员沿用同段已明说的整行列对应，原合计仍单独留在后面。
        columns={get_column_letter(c):get_column_letter(sl+c-tl) for c in range(left+1,right+1)}
        tables.append({'selection_dimension':None,'range_name':name,'source_sheet':extra[1],
            'source_rows':[row],'old_detail_count':0,'first_row_offset':point-top,'target_row_offsets':[point-top],
            'columns':columns,'labels':{get_column_letter(left):get_column_letter(xl)},'totals':[],
            'blank_values':{str(row):{col:{'value':None} for col in columns}},
            'operations':[{'axis':'row','action':'insert','index':point-top,'count':1,'template_index':point-top-1}]})
    for line in text.splitlines():
        if '已有空名称业务行' not in line or '不另新增一行' not in line:continue
        sources=re.findall(r'(期末|期初|本期|上期)金额\s*←\s*B\s*([^!\s]+)!([A-Z]+\d+)',line)
        owners={name for name,parts in a['机械记录'].get('names',{}).items() for area in parts
            if any(area['sheet']+'!'+span in text for span in re.findall(r'[A-Z]+\d+:[A-Z]+\d+',text)
                   if range_boundaries(area['range'])[1]<=range_boundaries(span)[1]<=range_boundaries(span)[3]<=range_boundaries(area['range'])[3])}
        prior=[table for rel in b.get('已确认关系',[]) for table in rel.get('tables',[])
               if table['range_name'] in owners and table['old_detail_count']==0 and len(table['source_rows'])==1]
        if not sources or len(prior)!=1:continue
        table=deepcopy(prior[0]);area=a['机械记录']['names'][table['range_name']][0]
        point=range_boundaries(area['range'])[1]+table['first_row_offset'];roles=_known_period_columns(a,area['sheet'],point)
        aliases={'期末':'closing','期初':'opening','本期':'current','上期':'prior'};row=str(table['source_rows'][0])
        explicit={roles[aliases[role]]:sheet+'!'+cell for role,sheet,cell in sources if aliases[role] in roles}
        if len(explicit)!=len(sources):continue
        table.setdefault('row_sources',{}).setdefault(row,{}).update(explicit)
        table['row_columns']={row:{col:re.match(r'[A-Z]+',ref.rsplit('!',1)[1])[0] for col,ref in explicit.items()}}
        tables.append(table)
    for block in re.split(r'\n(?=#{1,6}\s)',text):
        moved=re.search(r'插入后[，,]原独立合计移至第(\d+)行',block)
        first=re.search(r'A\s*([^!\s]+)!([A-Z]+\d+:[A-Z]+\d+)\s*↔\s*B\s*([^!\s]+)!([A-Z]+\d+:[A-Z]+\d+)',block)
        if not (moved and first and '原独立合计之前' in block):continue
        left,point,right,end=range_boundaries(first[2]);new_total=int(moved[1])
        areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
               if area['sheet']==first[1] and range_boundaries(area['range'])[1]<point==range_boundaries(area['range'])[3]]
        if len(areas)!=1 or point!=end or new_total<=point:continue
        requested=[first[1]+'!'+get_column_letter(c)+str(r) for r in range(point,new_total) for c in range(left,right+1)]
        pairs=source_relation(a,b,block,requested)['bindings'];rows=[];columns={};label=get_column_letter(left)
        for row in range(point,new_total):
            group=[p for p in pairs if coordinate_to_tuple(p['target'].rsplit('!',1)[1])[0]==row]
            source_rows={coordinate_to_tuple(p['source'].rsplit('!',1)[1])[0] for p in group}
            sources={p['source'].rsplit('!',1)[0] for p in group}
            mapping={re.match(r'[A-Z]+',p['target'].rsplit('!',1)[1])[0]:re.match(r'[A-Z]+',p['source'].rsplit('!',1)[1])[0] for p in group}
            if sources!={first[3]} or len(source_rows)!=1 or set(mapping)!={get_column_letter(c) for c in range(left,right+1)}:break
            if columns and mapping!=columns:break
            columns=mapping;rows.append(next(iter(source_rows)))
        if len(rows)!=new_total-point:continue
        name,area=areas[0];top=range_boundaries(area['range'])[1];source_label=columns.pop(label)
        tables.append({'selection_dimension':None,'range_name':name,'source_sheet':first[3],
            'source_rows':rows,'old_detail_count':0,'first_row_offset':point-top,'target_row_offsets':list(range(point-top,new_total-top)),
            'columns':columns,'labels':{label:source_label},'totals':[],
            'blank_values':{str(row):{col:{'value':None} for col in columns} for row in rows},
            'operations':[{'axis':'row','action':'insert','index':point-top,'count':len(rows),'template_index':point-top-1}]})
    for block in re.split(r'\n(?=#{1,6}\s)',text):
        old=re.search(r'A\s*([^!\s]+)!([A-Z]+\d+:[A-Z]+\d+)\s*的[^。\n]*不能按行序对应B\s*([^!\s]+)!([A-Z]+\d+:[A-Z]+\d+)',block)
        total=re.search(r'合计仍独立[^。\n]*A\s*([^!\s]+)!([A-Z]+\d+:[A-Z]+\d+)\s*对应B\s*([^!\s]+)!([A-Z]+\d+:[A-Z]+\d+)',block)
        if not (old and total and '旧名单另行保留' in block and '名称空白的业务行' in block and '列分别对应B的' in block):continue
        left,first,right,last=range_boundaries(old[2]);sl,sfirst,sr,slast=range_boundaries(old[4]);tl,point,tr,end=range_boundaries(total[2])
        areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts if area['sheet']==old[1]
               and range_boundaries(area['range'])[1]<first<=last<point<=range_boundaries(area['range'])[3]]
        if len(areas)!=1 or old[1]!=total[1] or old[3]!=total[3] or point!=end or point!=last+1 or (tl,tr)!=(left,right) or right-left!=sr-sl:continue
        name,area=areas[0];top=range_boundaries(area['range'])[1];rows=list(range(sfirst,slast+1))
        columns={get_column_letter(c):get_column_letter(sl+c-left) for c in range(left+1,right+1)}
        # 两组已被模型区分；依用户保留原位置的要求，把新增行放在旧明细之后。
        tables.append({'selection_dimension':None,'range_name':name,'source_sheet':old[3],
            'source_rows':rows,'old_detail_count':0,'first_row_offset':point-top,'target_row_offsets':list(range(point-top,point-top+len(rows))),
            'columns':columns,'labels':{get_column_letter(left):get_column_letter(sl)},'totals':[],
            'blank_values':{str(row):{col:{'value':None} for col in columns} for row in rows},
            'operations':[{'axis':'row','action':'insert','index':point-top,'count':len(rows),'template_index':first-top}]})
    for block in re.split(r'\n(?=#{1,6}\s)',text):
        extra=re.search(r'B\s*(\d+)\s*是未标明[^。\n]*的额外行',block)
        parent=re.search(r'A\s*(\d+)[“"]合计[”"]\s*↔\s*B\s*(\d+)[“"]合计[”"]',block)
        if not (extra and parent and 'A表没有对应的未命名明细行' in block):continue
        point,source_total=int(parent[1]),int(parent[2]);new_row=int(extra[1])
        prior=[pair for r in b.get('已确认关系',[]) if r.get('content')==reply for pair in r.get('bindings',[])
               if coordinate_to_tuple(pair['target'].rsplit('!',1)[1])[0]==point and coordinate_to_tuple(pair['source'].rsplit('!',1)[1])[0]==source_total]
        target_sheets={p['target'].rsplit('!',1)[0] for p in prior};source_sheets={p['source'].rsplit('!',1)[0] for p in prior}
        if len(target_sheets)!=1 or len(source_sheets)!=1:continue
        sheet=next(iter(target_sheets));areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
            if area['sheet']==sheet and range_boundaries(area['range'])[1]<point<=range_boundaries(area['range'])[3]]
        if len(areas)!=1:continue
        name,area=areas[0];left,top,_,_=range_boundaries(area['range']);label=get_column_letter(left)
        columns={re.match(r'[A-Z]+',p['target'].rsplit('!',1)[1])[0]:re.match(r'[A-Z]+',p['source'].rsplit('!',1)[1])[0] for p in prior}
        columns.pop(label,None)
        if not columns:continue
        tables.append({'selection_dimension':None,'range_name':name,'source_sheet':next(iter(source_sheets)),
            'source_rows':[new_row],'old_detail_count':0,'first_row_offset':point-top,'target_row_offsets':[point-top],
            'columns':columns,'labels':{label:None},'declared_labels':{str(new_row):{'value':None}},'totals':[],
            'blank_values':{str(new_row):{col:{'value':None} for col in columns}},
            'operations':[{'axis':'row','action':'insert','index':point-top,'count':1,'template_index':point-top-1}]})
    for block in re.split(r'\n(?=#{1,6}\s)',text):
        detail=re.search(r'A\s*([^!\s]+)!([A-Z]+\d+:[A-Z]+\d+)[^。\n]*?分别对应B\s*([^!\s]+)!([A-Z]+\d+:[A-Z]+\d+)',block)
        added=re.search(r'原合计行之前新增[^。\n]*?对应B\s*([^!\s]+)!([A-Z]+\d+:[A-Z]+\d+)',block)
        total=re.search(r'A原合计行对应B第(\d+)行',block)
        if not (detail and added and total and '名称保持空白' in block):continue
        dl,first,dr,last=range_boundaries(detail[2]);sl,sfirst,sr,slast=range_boundaries(detail[4]);al,extra,ar,end=range_boundaries(added[2])
        areas=[(n,p) for n,parts in a['机械记录'].get('names',{}).items() for p in parts if p['sheet']==detail[1]
            and range_boundaries(p['range'])[1]<=first<=last<=range_boundaries(p['range'])[3]]
        if len(areas)!=1 or detail[3]!=added[1] or dr-dl!=sr-sl or last-first!=slast-sfirst or (al,ar)!=(sl,sr) or extra!=end:continue
        name,area=areas[0];top=range_boundaries(area['range'])[1];label=get_column_letter(dl)
        cells=next(s['cells'] for s in a['机械记录']['sheets'] if s['name']==detail[1])
        totals=[r for r in range(last+1,range_boundaries(area['range'])[3]+1) if str(cells.get(label+str(r),{}).get('value') or '').strip()=='合计']
        if len(totals)!=1 or totals[0]!=last+1:continue
        columns={get_column_letter(c):get_column_letter(sl+c-dl) for c in range(dl+1,dr+1)}
        tables.append({'selection_dimension':None,'range_name':name,'source_sheet':detail[3],
            'source_rows':list(range(sfirst,slast+1))+[extra],'old_detail_count':last-first+1,'first_row_offset':first-top,
            'columns':columns,'labels':{label:get_column_letter(sl)},'totals':[{'target_row':totals[0],'source_row':int(total[1])}],
            'blank_values':{str(extra):{col:{'value':None} for col in columns}},
            'operations':[{'axis':'row','action':'insert','index':totals[0]-top,'count':1,'template_index':last-top}]})
    for index,(start,end,name,area) in enumerate(sections):
        block=text[start:sections[index+1][0] if index+1<len(sections) else len(text)]
        old=re.search(r'原有明细为\s*([^。\n]+?)[，,]\s*合计为\s*([A-Z]+)(\d+)',block)
        roles=re.search(r'期间列为\s*([A-Z]+)\s*期末余额[、，,]\s*([A-Z]+)\s*期初余额',block)
        items=re.findall(r'(?m)^\s*\d+[.]\s*(.+?)[：:]期末对应\s*([^!\s]+)!([A-Z]+)(\d+)[，,]\s*期初对应\s*\2!([A-Z]+)\4',block)
        if old and roles and len(items)>1 and items[-1][0]=='合计' and '名单随' in block:
            old_rows=[int(r) for r in re.findall(re.escape(old[2])+r'(\d+)',old[1])]
            sources={item[1] for item in items};column_sets={(item[2],item[4]) for item in items}
            if old_rows and old_rows==list(range(old_rows[0],int(old[3]))) and len(sources)==len(column_sets)==1:
                top=range_boundaries(area['range'])[1];first=old_rows[0];added=len(items)-1-len(old_rows)
                if added>=0:
                    source=next(iter(sources));columns=dict(zip(roles.groups(),next(iter(column_sets))))
                    tables.append({'selection_dimension':None,'range_name':name,'source_sheet':source,
                        'source_rows':[int(item[3]) for item in items[:-1]],'old_detail_count':len(old_rows),'first_row_offset':first-top,
                        'columns':columns,'labels':{old[2]:None},
                        'declared_labels':{str(item[3]):{'value':item[0]} for item in items[:-1]},
                        'totals':[{'target_row':int(old[3]),'source_row':int(items[-1][3])}],
                        'blank_values':{item[3]:{col:{'value':None} for col in columns} for item in items[len(old_rows):-1]},
                        'operations':[{'axis':'row','action':'insert','index':int(old[3])-top,'count':added,'template_index':old_rows[-1]-top}] if added else []})
        old=re.search(r'A原明细为\s*([A-Z]+)(\d+):\1(\d+)[，,\s]+独立合计为\s*\1(\d+)',block)
        span=re.search(r'A明细区\s*([A-Z]+)(\d+):([A-Z]+)\s*对应B\s*([^!\s]+)!([A-Z]+)(\d+):([A-Z]+)(\d+)',block)
        total=re.search(r'A合计行对应B\s*([^!\s]+)!([A-Z]+)(\d+):([A-Z]+)\3',block)
        if not (old and span and total and '逐行按行号对应' in block):continue
        if old.groups()[:2]!=span.groups()[:2] or span[4]!=total[1]:continue
        left,top,right,bottom=range_boundaries(area['range'])
        target_left,_,target_right,_=range_boundaries(span[1]+span[2]+':'+span[3]+span[2])
        source_left,first,source_right,last=range_boundaries(span[5]+span[6]+':'+span[7]+span[8])
        first_old,last_old,total_old=map(int,old.groups()[1:]);added=last-first+1-(last_old-first_old+1)
        if target_right-target_left!=source_right-source_left or added<0 or not top<=first_old<=last_old<total_old<=bottom:continue
        columns={get_column_letter(target_left+i):get_column_letter(source_left+i) for i in range(1,target_right-target_left+1)}
        tables.append({'selection_dimension':None,'range_name':name,'source_sheet':span[4],
            'source_rows':list(range(first,last+1)),'old_detail_count':last_old-first_old+1,'first_row_offset':first_old-top,
            'columns':columns,'labels':{span[1]:span[5]},'totals':[{'target_row':total_old,'source_row':int(total[3])}],
            'blank_values':{str(row):{col:{'value':None} for col in columns} for row in range(first+last_old-first_old+1,last+1)},
            'operations':[{'axis':'row','action':'insert','index':total_old-top,'count':added,'template_index':last_old-top}] if added else []})
    for index,(start,end,name,area) in enumerate(sections):
        block=text[start:sections[index+1][0] if index+1<len(sections) else len(text)]
        before=re.search(r'原([^，。；\n]+)之前插入一行',block)
        new_value=re.search(r'(?:期末|期初|本期|上期)业务格取\s*B\s*([A-Z]+)(\d+)',block)
        blank_label=re.search(r'B\s*([A-Z]+)(\d+)\s*为空',block)
        if not (before and new_value and blank_label and new_value[2]==blank_label[2] and '新增行' in block and '留空' in block):continue
        sheet=area['sheet'];left,top,right,bottom=range_boundaries(area['range']);label=get_column_letter(left)
        cells=next(s['cells'] for s in a['机械记录']['sheets'] if s['name']==sheet)
        anchors={}
        for row in range(top,bottom+1):
            literal=cells.get(label+str(row),{}).get('value')
            if isinstance(literal,str):anchors.setdefault(literal.strip(),[]).append(row)
        anchor=anchors.get(before[1],[])
        if len(anchor)!=1:continue
        from .A语义阅读 import addresses
        requested=[sheet+'!'+cell for cell in addresses(area['range'])]
        pairs=source_relation(a,b,block,requested)['bindings']
        columns={(pair['source'].rsplit('!',1)[0],re.match(r'[A-Z]+',pair['target'].rsplit('!',1)[1])[0])
            for pair in pairs if re.match(r'[A-Z]+',pair['source'].rsplit('!',1)[1])[0]==new_value[1]}
        if len(columns)!=1:continue
        source_sheet,column=next(iter(columns));row=anchor[0];source_row=int(new_value[2])
        tables.append({'selection_dimension':None,'range_name':name,'source_sheet':source_sheet,
            'source_rows':[source_row],'first_row_offset':row-top,'old_detail_count':0,'target_row_offsets':[row-top],
            'columns':{column:new_value[1]},'labels':{label:blank_label[1]},'totals':[],
            'blank_values':{str(source_row):{get_column_letter(col):{'value':None} for col in range(left+1,right+1)}},
            'operations':[{'axis':'row','action':'insert','index':row-top,'count':1,'template_index':row-top-1}]})
        # 行名在此仅定位答复明说的原行；来源列与目标列沿用同段明示对应。
        for literal,rows in anchors.items():
            if len(rows)!=1:continue
            source=re.search(r'(?:A\s*)?'+re.escape(literal)+r'行对应\s*B\s*(?:'+re.escape(source_sheet)+r'!)?([A-Z]+\d+)',block)
            if source:bindings.append({'target':sheet+'!'+column+str(rows[0]),'source':source_sheet+'!'+source[1]})
    if re.search(r'A位置[（(]插入后[）)]',text) and '新增' in text and re.search(r'原合计后移(?:一|1)行',text):
        from .A语义阅读 import addresses
        requested=[r['sheet']+'!'+cell for r in b.get('a_required_ranges',[]) for cell in addresses(r['range'])]
        pairs=source_relation(a,b,reply,requested)['bindings'];grouped={}
        for pair in pairs:
            sheet,cell=pair['target'].rsplit('!',1);row,col=coordinate_to_tuple(cell)
            grouped.setdefault((sheet,row),[]).append(pair)
        for (sheet,row),members in grouped.items():
            areas=[(name,area) for name,parts in a.get('机械记录',{}).get('names',{}).items() for area in parts
                if area['sheet']==sheet and all(pair['target'].split('!',1)[1] in addresses(area['range']) for pair in members)]
            sources={pair['source'].rsplit('!',1)[0] for pair in members}
            source_rows={coordinate_to_tuple(pair['source'].rsplit('!',1)[1])[0] for pair in members}
            if len(areas)!=1 or len(sources)!=1 or len(source_rows)!=1:continue
            name,area=areas[0];left,top,_,_=range_boundaries(area['range']);label=get_column_letter(left)
            columns={re.match(r'[A-Z]+',pair['target'].rsplit('!',1)[1])[0]:re.match(r'[A-Z]+',pair['source'].rsplit('!',1)[1])[0] for pair in members}
            if label not in columns:continue
            source_row=next(iter(source_rows));source_label=columns.pop(label)
            tables.append({'selection_dimension':None,'range_name':name,'source_sheet':next(iter(sources)),
                'source_rows':[source_row],'first_row_offset':row-top,'old_detail_count':0,
                'target_row_offsets':[row-top],'columns':columns,'labels':{label:source_label},'totals':[],
                'blank_values':{str(source_row):{col:{'value':None} for col in columns}},
                'operations':[{'axis':'row','action':'insert','index':row-top,'count':1,'template_index':row-top-1}]})
    # 名单表的行序来自答复，金额列只沿用两边已识别角色或答复明确的列简写。
    aliases={'本期':'current','上期':'prior','期初':'opening','期末':'closing'}
    known_columns=_known_period_columns
    if re.search(r'原(?:小计|合计)之前增加一行',text):
        detail=re.search(r'A\s*([A-Z]+)(\d+):\1(\d+)\s*分别对应\s*B\s*([^!\s]+)!([A-Z]+)(\d+):\5(\d+)',text)
        extra=re.search(r'新增A明细行的名称保持空白[，,]对应\s*B\s*([A-Z]+)(\d+)',text)
        if detail and extra:
            label,first,last,source=detail[1],int(detail[2]),int(detail[3]),detail[4]
            areas=[(name,area) for name,parts in a['机械记录'].get('names',{}).items() for area in parts
                   if area['sheet']+'!'+area['range'].replace('$','') in text and range_boundaries(area['range'])[1]<first<=last<range_boundaries(area['range'])[3]]
            if len(areas)==1:
                name,area=areas[0];left,top,right,_=range_boundaries(area['range']);source_row=int(extra[2]);columns={}
                roles=known_columns(a,area['sheet'],first)
                for role,col,row in re.findall(r'(期末|期初)(?:账面)?余额对应\s*B\s*([A-Z]+)(\d+)',text):
                    if int(row)==source_row and aliases[role] in roles:columns[roles[aliases[role]]]=col
                tables.append({'selection_dimension':None,'range_name':name,'source_sheet':source,'source_rows':[source_row],
                    'old_detail_count':0,'first_row_offset':last+1-top,'target_row_offsets':[last+1-top],
                    'columns':columns,'labels':{label:extra[1]},'totals':[],
                    'blank_values':{str(source_row):{get_column_letter(c):{'value':None} for c in range(left+1,right+1)}},
                    'operations':[{'axis':'row','action':'insert','index':last+1-top,'count':1,'template_index':last-top}]})
    for index,(start,end,name,area) in enumerate(sections):
        block=text[start:sections[index+1][0] if index+1<len(sections) else len(text)]
        statement=re.search(r'A\s*各[^。\n]+?余额采用\s*B\s*([^!\s]+)!([^。\n]+)',block)
        if not statement or '一条无名称行' not in block:continue
        source_sheet=statement[1]
        if source_sheet not in {s['name'] for s in b['机械记录']['sheets']}:continue
        sources={aliases[role]:(col,int(first),int(last)) for col,first,last,role in re.findall(
            r'([A-Z]+)(\d+):\1(\d+)\s*[（(](期末|期初|本期|上期)[^）)]*[）)]',statement[2])}
        targets=re.search(r'A\s*((?:[A-Z]+\d+:[A-Z]+\d+[、，,\s]*)+)\s*[:：]',block)
        if not sources or not targets:continue
        ranges=re.findall(r'([A-Z]+)(\d+):\1(\d+)',targets[1])
        row_spans={(int(first),int(last)) for _,first,last in ranges}
        source_spans={(first,last) for _,first,last in sources.values()}
        if len(row_spans)!=1 or len(source_spans)!=1:continue
        first,last=next(iter(row_spans));source_first,source_last=next(iter(source_spans))
        if source_last-source_first!=last-first+1:continue
        columns={known_columns(a,area['sheet'],first).get(role):values[0] for role,values in sources.items()}
        if None in columns or not set(columns)<={col for col,_,_ in ranges}:continue
        for target_row,source_row in zip(range(first,last+1),range(source_first,source_last)):
            bindings.extend({'target':area['sheet']+'!'+col+str(target_row),'source':source_sheet+'!'+source_col+str(source_row)} for col,source_col in columns.items())
        left,top,_,_=range_boundaries(area['range']);offset=last+1-top;label=get_column_letter(left)
        tables.append({'selection_dimension':None,'range_name':name,'source_sheet':source_sheet,
            'source_rows':[source_last],'old_detail_count':0,'first_row_offset':offset,'target_row_offsets':[offset],
            'columns':columns,'labels':{label:None},'declared_labels':{str(source_last):{'value':None}},'totals':[],
            'blank_values':{str(source_last):{col:{'value':None} for col in columns}},
            'operations':[{'axis':'row','action':'insert','index':offset,'count':1,'template_index':offset-1}]})
    for block in re.split(r'\n(?=#{1,6}\s)',text):
        old=re.search(r'A\s*原明细为\s*([^\s!]+)!([A-Z]+)(\d+):\2(\d+)[^。\n]*?合计为\s*\1!\2(\d+)',block)
        if not old:
            old=re.search(r'A(?:《[^》]+》)?\s*([^!\s]+)!([A-Z]+)(\d+):\2(\d+)原有[^。\n]*?，\s*\2(\d+)为合计',block)
        source=re.search(r'B[《「]([^》」]+)[》」]',block)
        if not (old and source):continue
        sheet,label,first,last,total=old.groups();first,last,total=map(int,(first,last,total));source=source[1]
        areas=[(name,area) for name,parts in a.get('机械记录',{}).get('names',{}).items() for area in parts
               if area['sheet']==sheet and range_boundaries(area['range'])[1]<first<=last<total<=range_boundaries(area['range'])[3]]
        if len(areas)!=1 or source not in {s['name'] for s in b['机械记录']['sheets']}:continue
        header=None;entries=[];source_roles={};source_index=None;target_index=None
        declared_count=re.search(r'实际列示\s*(\d+)\s*项[^。\n]*原明细区扩',block)
        for line in block.splitlines():
            if not line.startswith('|'):continue
            values=[v.strip() for v in line.strip('|').split('|')]
            source_headers=[i for i,v in enumerate(values) if re.sub(r'\s+','',v) in ('采用的B位置','搬运的B来源及B成员名称')]
            if any('更新后' in v and 'A' in v for v in values) and len(source_headers)==1:
                header=values;source_index=source_headers[0]
                target_index=next((i for i,v in enumerate(values) if re.sub(r'\s+','',v)=='更新后A位置'),None)
                continue
            if not header or len(values)!=len(header):continue
            ref=values[source_index]
            found=re.search(r'(?<![A-Z])([A-Z]+)(\d+)(?::([A-Z]+)(\d+))?',ref)
            if not found or found[4] and found[4]!=found[2]:continue
            pair=re.search(r'(本期|上期|期初|期末)\s*/\s*(本期|上期|期初|期末)[^；;]*?([A-Z]+)\s*/\s*([A-Z]+)\d+',ref)
            if pair:source_roles.update({aliases[pair[1]]:pair[3],aliases[pair[2]]:pair[4]})
            is_total=any('原合计' in v for v in values)
            if target_index is not None:
                planned=re.fullmatch(re.escape(label)+r'(\d+)',values[target_index])
                if not declared_count or not planned or int(planned[1])!=first+len(entries):continue
                is_total=int(planned[1])==first+int(declared_count[1])
                for amount_column,amount_row,role in re.findall(r'([A-Z]+)(\d+)\s*(期末|期初|本期|上期)',' '.join(values)):
                    if amount_row==found[2]:source_roles[aliases[role]]=amount_column
            entries.append((int(found[2]),found[1],is_total))
        details=[entry for entry in entries if not entry[2]];totals=[entry for entry in entries if entry[2]]
        if not details or len(totals)!=1:continue
        source_roles={**known_columns(b,source,details[0][0]),**source_roles}
        target_roles=known_columns(a,sheet,first)
        columns={col:source_roles[role] for role,col in target_roles.items() if role in source_roles}
        if not target_roles or len(columns)!=len(target_roles) or len({e[1] for e in entries})!=1:continue
        name,area=areas[0];top=range_boundaries(area['range'])[1];added=len(details)-(last-first+1)
        if added<0:continue
        if target_index is not None:
            bindings.append({'target':sheet+'!'+label+str(total),'source':source+'!'+totals[0][1]+str(totals[0][0]),'label':True})
        tables.append({'selection_dimension':None,'range_name':name,'source_sheet':source,
            'source_rows':[entry[0] for entry in details],'old_detail_count':last-first+1,'first_row_offset':first-top,
            'columns':columns,'labels':{label:details[0][1]},'totals':[{'target_row':total,'source_row':totals[0][0]}],
            'operations':[{'axis':'row','action':'insert','index':total-top,'count':added,'template_index':first-top}] if added else []})
    def rows(value):
        result=[]
        for part in re.split('[、，,]',value):
            numbers=[int(n) for n in re.findall(r'\d+',part)]
            result.extend(range(numbers[0],numbers[-1]+1) if len(numbers)==2 else numbers)
        return result
    for i,(start,end,name,area) in enumerate(sections):
        block=text[end:sections[i+1][0] if i+1<len(sections) else len(text)]
        detail=re.search(r'A\s*第\s*(\d+)[—–至-](\d+)\s*行分别对应\s*B\s*([^!\n]+)!([A-Z]+)(\d+):\4(\d+)',block)
        added_row=re.search(r'新增的?\s*A\s*第\s*(\d+)\s*行对应\s*B\s*第\s*(\d+)\s*行',block)
        moved_total=re.search(r'原\s*A\s*第\s*(\d+)\s*行的独立合计移至\s*A\s*第\s*(\d+)\s*行，对应\s*B\s*([^!\n]+)!([A-Z]+)(\d+):([A-Z]+)\5',block)
        if detail and added_row and moved_total:
            first,last,sheet,source_label,sfirst,slast=detail.groups()
            first,last,sfirst,slast=map(int,(first,last,sfirst,slast))
            left,top,_,_=range_boundaries(area['range']);columns={}
            for match in re.finditer(r'A\s*列?\s*([A-Z](?:[、，,]\s*[A-Z])+)[）)]?\s*分别对应\s*B\s*([^。；\n]+)',block):
                target_cols=re.findall(r'[A-Z]',match[1])
                source_cols=re.findall(r'([A-Z]+)\d+:[A-Z]+\d+',match[2])
                if len(target_cols)==len(source_cols):columns.update(zip(target_cols,source_cols))
            if (columns and last-first==slast-sfirst and int(added_row[1])==last+1==int(moved_total[1])
                    and int(moved_total[2])==int(moved_total[1])+1 and sheet==moved_total[3]
                    and sheet in {s['name'] for s in b['机械记录']['sheets']}):
                tables.append({'selection_dimension':None,'range_name':name,'source_sheet':sheet,
                    'source_rows':[*range(sfirst,slast+1),int(added_row[2])],
                    'old_detail_count':last-first+1,'first_row_offset':first-top,
                    'columns':columns,'labels':{get_column_letter(left):source_label},
                    'blank_values':{added_row[2]:{col:{'value':None} for col in columns}},
                    'totals':[{'target_row':int(moved_total[1]),'source_row':int(moved_total[5])}],
                    'operations':[{'axis':'row','action':'insert','index':last+1-top,'count':1,'template_index':first-top}]})
        if '原值' in block and '单列保留' in block:
            from .A语义阅读 import addresses
            left,top,right,bottom=range_boundaries(area['range']);sheet=area['sheet'];label=get_column_letter(left)
            cells=next(s.get('cells',{}) for s in a['机械记录']['sheets'] if s['name']==sheet)
            kept=[row for row in range(top,bottom) if (value:=str(cells.get(label+str(row),{}).get('value') or '').strip())
                  and re.search(re.escape(value)+r'[^。\n]*原值[^。\n]*单列保留',block)]
            pairs=source_relation(a,b,block,[sheet+'!'+cell for cell in addresses(area['range'])])['bindings']
            targets={item['target']:item['source'] for item in pairs};complete=[]
            for row in range(top,bottom+1):
                refs=[targets.get(sheet+'!'+get_column_letter(col)+str(row)) for col in range(left,right+1)]
                if all(refs) and len({ref.rsplit('!',1)[0] for ref in refs})==1 and len({coordinate_to_tuple(ref.rsplit('!',1)[1])[0] for ref in refs})==1:
                    complete.append((row,refs))
            if kept and complete and complete[-1][0]==bottom and complete[0][0]<=min(kept) and len(complete)>1:
                first=complete[0][0];source=complete[0][1][0].rsplit('!',1)[0]
                column_sets={tuple(re.match(r'[A-Z]+',ref.rsplit('!',1)[1])[0] for ref in refs) for _,refs in complete}
                if len(column_sets)==1 and [row for row,_ in complete]==list(range(first,bottom+1)):
                    source_cols=next(iter(column_sets));source_rows=[coordinate_to_tuple(refs[0].rsplit('!',1)[1])[0] for _,refs in complete]
                    tables.append({'selection_dimension':None,'range_name':name,'source_sheet':source,
                        'source_rows':source_rows[:-1],'old_detail_count':bottom-first,'first_row_offset':first-top,
                        'columns':dict(zip((get_column_letter(col) for col in range(left+1,right+1)),source_cols[1:])),
                        'labels':{label:source_cols[0]},'totals':[{'target_row':bottom,'source_row':source_rows[-1]}],
                        'retained_members':[{'index':len(source_rows)-1+index,'cells':{get_column_letter(col):{'source':sheet+'!'+get_column_letter(col)+str(row)}
                            for col in range(left,right+1)}} for index,row in enumerate(kept)],
                        'blank_values':{str(source_rows[row-first]):{get_column_letter(col):{'value':None} for col in range(left+1,right+1)} for row in kept},
                        'operations':[{'axis':'row','action':'insert','index':bottom-top,'count':len(kept),'template_index':first-top}]})
        target_column=re.search(r'A的[^。\n]*?列\s*([A-Z]+)\d+:[A-Z]+\d+',block)
        existing=[t for relation in b.get('已确认关系',[]) for t in relation.get('tables',[]) if t['range_name']==name]
        if target_column and len(existing)==1:
            prior_table=existing[0];col=target_column[1];chosen={}
            sheet_names=[s['name'] for s in b['机械记录']['sheets']]
            refs=re.compile(r'('+ '|'.join(re.escape(s) for s in sorted(sheet_names,key=len,reverse=True))+r')\s*!?\s*([A-Z]+)(\d+)')
            for line in block.splitlines():
                values=line.strip().strip('|').split('|')
                if len(values)<2:continue
                selected=refs.findall(values[0]);member_refs=refs.findall('|'.join(values[1:]))
                members={int(row) for sh,_,row in member_refs if sh==prior_table['source_sheet'] and int(row) in prior_table['source_rows']}
                if len(selected)==len(members)==1:
                    sh,sc,row=selected[0];chosen[next(iter(members))]=sh+'!'+sc+row
            if set(chosen)==set(prior_table['source_rows']):
                copied=deepcopy(prior_table);copied['columns'][col]=re.match(r'[A-Z]+',next(iter(chosen.values())).rsplit('!',1)[1])[0]
                for row,source_ref in chosen.items():copied.setdefault('row_sources',{}).setdefault(str(row),{})[col]=source_ref
                total=re.search(r'合计使用B[^。\n]*',block)
                sources=refs.findall(total[0]) if total else []
                if len(sources)==1 and len(copied.get('totals',[]))==1:
                    sh,sc,row=sources[0];source_ref=sh+'!'+sc+row;declared=copied['totals'][0]
                    copied.setdefault('row_sources',{}).setdefault(str(declared['source_row']),{})[col]=source_ref
                    bindings.append({'target':area['sheet']+'!'+col+str(declared['target_row']),'source':source_ref})
                tables.append(copied)
        if not re.search(r'各列依次对应|分别对应A',block):continue
        sources={s['name'] for s in b['机械记录']['sheets'] if s['name']+'!' in block}
        if len(sources)!=1:continue
        source=next(iter(sources));left,top,right,bottom=range_boundaries(area['range'])
        mentioned=[int(n) for n in re.findall(r'B\s*第\s*(\d+)',block)]
        candidates=[r for r in b.get('b_source_ranges',[]) if r['sheet']==source and mentioned
            and all(range_boundaries(r['range'])[1]<=n<=range_boundaries(r['range'])[3] for n in mentioned)]
        candidates=list({r['range']:r for r in candidates}.values())
        if len(candidates)!=1:continue
        sl,_,sr,_=range_boundaries(candidates[0]['range'])
        if right-left!=sr-sl:continue
        columns={get_column_letter(c):get_column_letter(sl+c-left) for c in range(left+1,right+1)}
        label=get_column_letter(left);source_label=get_column_letter(sl);sheet=area['sheet']
        cells=next(s['cells'] for s in a['机械记录']['sheets'] if s['name']==sheet)
        source_cells=next(s['cells'] for s in b['机械记录']['sheets'] if s['name']==source)
        labels={r:str(cells.get(label+str(r),{}).get('value') or '').strip() for r in range(top,bottom+1)}
        prior={}
        for relation in b.get('已确认关系',[]):
            for pair in relation.get('bindings',[]):
                ts,tc=pair['target'].rsplit('!',1);ss,sc=pair['source'].rsplit('!',1)
                tr=coordinate_to_tuple(tc)[0]
                if ts==sheet and ss==source and top<=tr<=bottom:
                    prior.setdefault(coordinate_to_tuple(sc)[0],set()).add(tr)
        groups=re.split(r'\n(?=-\s*[^\n：:]+段[：:])',block)
        for group in groups:
            pairs=[];parents=[]
            for sentence in re.split(r'[。；;\n]|(?<=[，,])(?=[^\d，,。；;\n]*(?:→|对应)\s*B)',group):
                found=re.search(r'(.*?)\s*(?:分别)?(?:→|对应)\s*B\s*第\s*([\d、，,—–至-]+)\s*行',sentence)
                if not found:continue
                source_rows=rows(found[2]);prefix=found[1]
                named=sorted(((prefix.find(value),row) for row,value in labels.items() if value and value in prefix),key=lambda item:item[0])
                target_rows=[row for _,row in named]
                parent=bool(re.search(r'合计|父项',prefix))
                if len(target_rows)!=len(source_rows) and parent and len(source_rows)==1:
                    named_parent=str(source_cells.get(source_label+str(source_rows[0]),{}).get('value') or '').strip()
                    exact=[row for row,value in labels.items() if named_parent and value==named_parent]
                    if len(exact)==1:target_rows=exact
                if len(target_rows)!=len(source_rows) and parent and len(source_rows)==1 and len(prior.get(source_rows[0],[]))==1:
                    target_rows=list(prior[source_rows[0]])
                if len(target_rows)!=len(source_rows):continue
                for target_row,source_row in zip(target_rows,source_rows):
                    pairs.append((target_row,source_row))
                    if parent:parents.append((target_row,source_row))
                    for c,sc in {label:source_label,**columns}.items():
                        bindings.append({'target':sheet+'!'+c+str(target_row),'source':source+'!'+sc+str(source_row),
                            **({'label':True} if c==label else {})})
            added=[]
            for sentence in re.split(r'[。；;\n]',group):
                if not re.search(r'新增|增列|排入|保留(?:\s*B)?\s*第',sentence):continue
                for stated in re.finditer(r'(?:B\s*)?第\s*([\d、，,—–至-]+)\s*行',sentence):added.extend(rows(stated[1]))
            parents=list(dict.fromkeys(parents));pairs=list(dict.fromkeys(pairs))
            detail=sorted(pair for pair in pairs if pair not in parents)
            added=[r for r in dict.fromkeys(added) if r not in {s for _,s in pairs}]
            if not added or len(parents)!=1:continue
            parent,source_parent=parents[0]
            if detail:first=detail[0][0];old=len(detail)
            elif '父项下排入' in group:first=parent+1;old=0
            else:continue
            if detail and [r for r,_ in detail]!=list(range(first,first+old)):continue
            chosen=[s for _,s in detail]+added
            style=first if old else min((top+t['first_row_offset'] for t in tables
                if t['range_name']==name and t['old_detail_count']),default=None)
            if style is None:continue
            tables.append({'selection_dimension':None,'range_name':name,'source_sheet':source,'source_rows':chosen,
                'first_row_offset':first-top,'old_detail_count':old,'columns':columns,'labels':{label:source_label},
                **({'target_row_offsets':[first-top+sum(len(t['source_rows'])-t['old_detail_count'] for t in tables
                    if t['range_name']==name and t['first_row_offset']<first-top)+i for i in range(len(chosen))]} if not old else {}),
                'totals':[{'target_row':parent,'source_row':source_parent}],
                'operations':[{'axis':'row','action':'insert','index':first+old-top,'count':len(added),'template_index':style-top}],
                'coverage_reason':'答复明确的原成员对应、同序列关系及新增来源行；未补名单维度'})
    return bindings,tables


def source_relation(a, b, reply, targets, *, arranged_targets=False):
    """只展开原答复明示的位置关系，供共用入口及已有答复批量接入复用。"""
    from .A语义阅读 import addresses
    from openpyxl.utils.cell import coordinate_to_tuple,range_boundaries,get_column_letter
    bindings=[]
    sheets={sheet['name'] for sheet in b['机械记录']['sheets']}
    target_sheets={target.rsplit('!',1)[0] for target in targets}
    text=_expand_declared_row_columns(_source_aliases(reply,b)).replace('`','').replace('**','')
    aliases=[sheet for sheet in sheets if re.search(r'[“「"\']?…[”」"\']?\s*均指\s*[“「"\']?'+re.escape(sheet),text)]
    def locations(value,names,default=None,column_major=False):
        pattern=r'(?:(?P<sheet>'+ '|'.join(re.escape(name) for name in sorted(names,key=len,reverse=True))+r'|…)!)?(?P<area>[A-Z]+\d+(?::[A-Z]+\d+)?)'
        result=[];sheet=default
        for found in re.finditer(pattern,value):
            if found['sheet']:sheet=(aliases[0] if len(aliases)==1 else sheet) if found['sheet']=='…' else found['sheet']
            if sheet in names:
                cells=addresses(found['area'])
                if column_major:cells=sorted(cells,key=lambda cell:(coordinate_to_tuple(cell)[1],coordinate_to_tuple(cell)[0]))
                result.extend(sheet+'!'+cell for cell in cells)
        return list(dict.fromkeys(result))
    def connect(left,right,source_default=None,row_reference=False):
        if re.search(r'[A-Z]+\d+:[A-Z]+(?![A-Z0-9])',left):return
        start=r"(?:合计[:：]\s*)?(?:本次(?:采用|选用)的)?(?:B(?:来源(?:位置)?|表|位置)?\s*(?:的|中)?\s*[:：]?\s*)?(?:'?(?:"+'|'.join(re.escape(name) for name in sorted(sheets,key=len,reverse=True))+r"|…)'?!|!)?[A-Z]+\d+"
        if not re.match(start,right.strip(' \"“「')):return
        right=re.sub(r'(?<=[0-9])\s*[（(]\d{4}-\d{2}-\d{2}[）)]','',right)
        right=re.sub(r'(?<=[0-9])[（(](?:期末|期初|本期|上期)(?:余额|发生额)?[）)]','',right)
        names=[match.span() for name in sheets for match in re.finditer(re.escape(name)+'!',right)]
        cuts=[match.start() for match in re.finditer(r'[，,](?!\s*[A-Z]+\d)|[；;（(]',right)
              if not any(start<=match.start()<end for start,end in names)
              and not (match[0] in '；;' and re.match(r'\s*(?:(?:'+
                  '|'.join(re.escape(name) for name in sheets)+r')!)?[A-Z]+\d+',right[match.end():]))]
        if cuts:right=right[:min(cuts)]
        selected=locations(left,target_sheets,next(iter(target_sheets)) if len(target_sheets)==1 else None)
        spans=re.findall(r'([A-Z]+)\d+:([A-Z]+)\d+',left)
        # 多个列段与一个矩形对应时，保持答复给出的列分组，不交错两列。
        column_major=len(spans)>1 and all(first==last for first,last in spans)
        right_spans=re.findall(r'([A-Z]+)\d+:([A-Z]+)\d+',right)
        if len(right_spans)>1 and all(first==last for first,last in right_spans):
            selected=locations(left,target_sheets,next(iter(target_sheets)) if len(target_sheets)==1 else None,True)
        sources=locations(right,sheets,source_default,column_major)
        if row_reference and len(selected)==1 and selected[0] not in targets:
            target_sheet,target_cell=selected[0].rsplit('!',1);row_number=coordinate_to_tuple(target_cell)[0]
            declared=sorted((location for location in targets if location.rsplit('!',1)[0]==target_sheet
                and coordinate_to_tuple(location.rsplit('!',1)[1])[0]==row_number),
                key=lambda location:coordinate_to_tuple(location.rsplit('!',1)[1])[1])
            if len(declared)==len(sources):selected=declared
        if len(sources)==1:sources=sources*len(selected)
        if len(selected)==1 and selected[0] in targets:
            bindings.extend({'target':selected[0],'source':source} for source in sources)
        elif len(selected)==len(sources):bindings.extend({'target':target,'source':source} for target,source in zip(selected,sources) if target in targets)
        elif selected and sources:
            # 行关系已明说且列编号已识别时，按编号接列，不按行宽挤配名称列。
            if len({coordinate_to_tuple(loc.rsplit('!',1)[1])[0] for loc in selected})==len({coordinate_to_tuple(loc.rsplit('!',1)[1])[0] for loc in sources})==1:
                left_records,right_records=dict(_records(a)),dict(_records(b))
                keys={left_records.get(loc,{}).get('metric_id') for loc in selected}- {None}
                for key in keys:
                    left=[loc for loc in selected if loc in targets and left_records.get(loc,{}).get('metric_id')==key]
                    right=[loc for loc in sources if right_records.get(loc,{}).get('metric_id')==key]
                    if len(left)==len(right)==1:bindings.append({'target':left[0],'source':right[0]})

    declared_sheets={t['sheet'] for t in b.get('target_tables',[])}
    if b.get('target_table'):declared_sheets.add(b['target_table']['sheet'])
    if not declared_sheets:declared_sheets={t['sheet'] for t in b.get('b_source_ranges',[])}
    source_default=next(iter(declared_sheets)) if len(declared_sheets)==1 else None
    adopted={sheet for sheet in sheets if re.search(r'本次(?:采用的\s*B来源\s*(?:为|是|[:：])|(?:统一)?采用\s*B\s*(?:的)?)\s*[“「"\']?'+re.escape(sheet),text)}
    if len(adopted)==1:source_default=next(iter(adopted))
    named_sources={sheet for sheet in sheets if re.search(r'B\s*'+re.escape(sheet)+r'!',text)}
    if source_default is None and len(named_sources)==1:source_default=next(iter(named_sources))
    choices=re.findall(r'(?m)^\s*[-*]?\s*(?:选择|采用|选用)的\s*B来源\s*[:：]\s*([^\n]+)',text)
    if len(targets)==1 and len(choices)==1:connect(next(iter(targets)),choices[0],source_default)
    replaced_targets=[]
    for cancelled in re.finditer(r'撤销上次错误落到\s*(.+?)\s*的来源对应',text):
        replaced_targets.extend(loc for loc in locations(cancelled[1],target_sheets,next(iter(target_sheets)) if len(target_sheets)==1 else None) if loc in targets)
    retained=[]
    for block in re.split(r'\n(?=#{1,6}\s)',text):
        heading=block.splitlines()[0] if block.splitlines() else ''
        if re.search(r'保留\s*A?\s*原内容',heading):
            retained.extend(loc for loc in locations(block,target_sheets,next(iter(target_sheets)) if len(target_sheets)==1 else None) if loc in targets)
    for kept in re.finditer(r'([^。\n]+?)(?<!不)(?:均)?保留\s*(?:A\s*)?原(?:空白|内容|值)',text):
        retained.extend(loc for loc in locations(kept[1],target_sheets,next(iter(target_sheets)) if len(target_sheets)==1 else None) if loc in targets)
    for kept in re.finditer(r'([^。\n]+?)原有金额(?:均)?独立保留',text):
        retained.extend(loc for loc in locations(kept[1],target_sheets,next(iter(target_sheets)) if len(target_sheets)==1 else None) if loc in targets)
    for cancelled in re.finditer(r'撤销\s*A(?:原先)?\s*(.+?)\s*从\s*[^。\n]+?取数(?:的)?对应',text):
        retained.extend(loc for loc in locations(cancelled[1],target_sheets,next(iter(target_sheets)) if len(target_sheets)==1 else None) if loc in targets)
    columns=None;table_source=source_default;name_columns=None;named_labels=[];row_reference=False
    for line in text.splitlines():
        declared=[sheet for sheet in sheets if sheet in line and re.search(r'以下\s*B\s*(?:坐标|位置)(?:均)?(?:属于|在)',line)]
        if len(declared)==1:table_source=declared[0]
        if not line.strip().startswith('|'):continue
        values=[value.strip() for value in line.strip().strip('|').split('|')]
        for value in values:
            if '←' in value:
                target,source=value.split('←',1);connect(target,source,table_source)
        if all(key in values for key in ('A 目标行','B 来源行','更新后的事项名称')):
            name_columns=(values.index('A 目标行'),values.index('B 来源行'));continue
        if name_columns and len(values)>max(name_columns) and table_source and len(target_sheets)==1:
            target,source=(values[index] for index in name_columns)
            if re.fullmatch(r'[A-Z]+\d+',target) and re.fullmatch(r'[A-Z]+\d+',source):
                named_labels.append({'action':'copy','target':next(iter(target_sheets))+'!'+target,'source':table_source+'!'+source})
        left=[i for i,value in enumerate(values) if 'A' in value and any(word in value for word in ('位置','目标','需要','坐标','来源'))]
        right=[i for i,value in enumerate(values) if 'B' in value and any(word in value for word in ('来源','取数','位置','坐标'))]
        paths=[i for i in right if '来源路径' in values[i]]
        if len(right)>1 and len(paths)==1:right=paths
        adopted_columns=[i for i in right if re.search(r'采用|选定|选用|搬运',values[i])]
        if len(right)>1 and len(adopted_columns)==1:right=adopted_columns
        if len(left)==len(right)==1 and not re.search(r'[A-Z]+\d+',line):
            stated=[sheet for sheet in sheets if re.search(r'[（(]'+re.escape(sheet)+r'[）)]',values[right[0]])]
            if len(stated)==1:table_source=stated[0]
            columns=(left[0],right[0]);row_reference='项目' in values[left[0]];continue
        if columns is None or len(values)<=max(columns):continue
        target,source=(values[index] for index in columns)
        if any(word in line for word in ('本次不扩展','A原值保留')):
            retained.extend(loc for loc in locations(target,target_sheets,next(iter(target_sheets)) if len(target_sheets)==1 else None) if loc in targets)
            continue
        stated={sheet for sheet in sheets if sheet+'!' in source}
        if len(stated)==1:table_source=next(iter(stated))
        source_cell=r'(?:'+'|'.join(re.escape(name)+'!' for name in sorted(sheets,key=len,reverse=True))+r')?[A-Z]+\d+(?::[A-Z]+\d+)?'
        target_cell=r'(?:'+'|'.join(re.escape(name)+'!' for name in sorted(target_sheets,key=len,reverse=True))+r')?[A-Z]+\d+(?::[A-Z]+\d+)?'
        source_group=source_cell+r'(?:\s*[、，,/]\s*'+source_cell+r')*'
        target_group=target_cell+r'(?:\s*[、，,/]\s*'+target_cell+r')*'
        # 期间文字在答复中充当分组标签；公式箭头后的下层位置不是本列采用的来源。
        period_targets={period:cells for cells,period in re.findall(
            '('+target_group+r')\s*[（(](\d{4})(?:年度|年)?[）)]',target)}
        grouped=False
        for period,cells in re.findall(r'(\d{4})(?:年度|年)?[:：]\s*('+source_group+')',source):
            stated=locations(cells,sheets,table_source)
            if stated:table_source=stated[0].rsplit('!',1)[0]
            if period in period_targets:
                connect(period_targets[period],cells,table_source);grouped=True
        if grouped:continue
        arrows=re.findall('('+source_cell+r')(?:(?![A-Z]+\d)[^→；;。])*→\s*(?:A\s*)?('+target_cell+')',source)
        if arrows:
            for source_location,target_location in arrows:connect(target_location,source_location,table_source)
            continue
        connect(re.split(r'→|->',target)[0],source,table_source,row_reference)
    for paragraph in text.split('\n\n'):
        declaration=re.search(r'直接同范围[^。；\n]*来源为\s*(B\s+.+)',paragraph)
        if declaration:connect(paragraph.split('。',1)[0],declaration[1],source_default)
    section_source=source_default
    for line in text.splitlines():
        if line.strip().startswith('|'):continue
        stated={sheet for sheet in sheets if sheet+'!' in line}
        declared={sheet for sheet in sheets if re.search(r'B\s*'+re.escape(sheet)+r'(?:!|\s|$)',line)}
        if not adopted:
            if len(declared)==1:section_source=next(iter(declared))
            elif len(stated)==1:section_source=next(iter(stated))
        default=source_default if adopted else next(iter(stated)) if len(stated)==1 else section_source
        target_row=re.search(r'A\s*([^!\s]+)!([A-Z]+\d+:[A-Z]+\d+)\s*[:：]',line)
        role_sources=re.findall(r'(期末|期初|本期|上期)\s*←\s*B\s*([^!\s]+)!([A-Z]+\d+)',line)
        if target_row and role_sources:
            left,row,right,last=range_boundaries(target_row[2]);roles=_known_period_columns(a,target_row[1],row)
            role_keys={'期末':'closing','期初':'opening','本期':'current','上期':'prior'}
            if row==last:
                for role,sheet,cell in role_sources:
                    col=roles.get(role_keys[role])
                    if col and col+str(row) in addresses(target_row[2]):connect(target_row[1]+'!'+col+str(row),sheet+'!'+cell)
            continue
        first_source=re.match(r'^\s*[-*]\s*B\s+((?:[^\s!]+!)?[A-Z]+\d+(?::[A-Z]+\d+)?)\s*对应B',line)
        final_target=re.search(r'。对应A\s*((?:[^\s!]+!)?[A-Z]+\d+(?::[A-Z]+\d+)?)',line)
        if first_source and final_target:
            connect(final_target[1],first_source[1],default)
        row=re.search(r'A\s+(?:[^!\s]+!)?([A-Z]+\d+:[A-Z]+\d+)',line)
        amounts=re.search(r'(?:其金额列为|金额依次取)\s*([A-Z]+\d+(?:[、，,][A-Z]+\d+)*)',line)
        name=re.search(r'(?:对应\s*B|名称采用\s*B)\s+([A-Z]+\d+)',line)
        if row and amounts and name:
            selected=locations(row[1],target_sheets,next(iter(target_sheets)) if len(target_sheets)==1 else None)
            values=locations(name[1]+'、'+amounts[1],sheets,default)
            if len(selected)==len(values) and len({coordinate_to_tuple(loc.rsplit('!',1)[1])[0] for loc in selected})==1:
                connect(row[1],name[1]+'、'+amounts[1],default)
                continue
        for sentence in re.split(r'。|[；;](?!\s*分别(?:对应|采用|取))|[，,]\s*(?=A\s*!?[A-Z]+\d+)',line):
            reverse=re.search(r'^\s*[-*]?\s*B\s+(.+?)\s*(?:对应|→|↔)\s*A\s*(.+)',sentence)
            if reverse:
                connect(reverse[2],reverse[1],default)
                continue
            claim=re.search(r'(?:(?:^|[：:，,；;\s）)])(?:分别采用|采用|选用|取自|应取|分别取|来源为|来源确定为|对应来源为|分别对应|可对应|对应)|(?<=[0-9])(?:对应|采用)|→|←|↔|->)\s*(?:B来源\s+)?(.+)',sentence)
            reverse=re.search(r'采用\s*([A-Z]+\d+(?::[A-Z]+\d+)?)的[^。；;]*?对应A(?:明细表|表)?的\s*([A-Z]+\d+(?::[A-Z]+\d+)?)',sentence)
            if reverse:connect(reverse[2],reverse[1],default)
            if not claim:continue
            prefix=re.sub(r'[（(]实际位置[^）)]*[）)]','',sentence[:claim.start()])
            prefix=re.split(r'更新后为',prefix)[0]
            local=re.split(r'：|:(?![A-Z]+\d)',prefix)[-1]
            if re.search(r'[A-Z]+\d+',local):prefix=local
            connect(prefix,claim[1],default)
    # 行号与列字母分别声明时，只展开答复明确给出的笛卡尔对应。
    for block in re.split(r'\n(?=#{1,6}\s|\*\*[^*\n]+\*\*\s*\n)',_source_aliases(reply,b)):
        block=re.sub(r'(\d+)\s*[至–—-]\s*(\d+)',r'\1:\2',block.replace('`','').replace('**',''))
        declared={sheet for sheet in sheets if sheet+'!' in block}
        if not declared and source_default:declared={source_default}
        if len(declared)!=1 or len(target_sheets)!=1:continue
        source_sheet=next(iter(declared));target_sheet=next(iter(target_sheets))
        rows=[]
        def numbers(value):
            return [number for first,last in re.findall(r'(\d+)(?::(\d+))?',value)
                    for number in range(int(first),int(last or first)+1)]
        for match in re.finditer(r'A\s*(?:第|行\s*)?([\d:、，,\s]+)(?:行)?(?:[“「"][^”」"\n]*[”」"]|[；;](?=\s*分别对应\s*B)|[^A-Za-z0-9。；;\n←→↔])*?(?:→|←|↔|分别采用|分别对应|对应|采用)\s*B\s*(?:第|行\s*)?(?:'+re.escape(source_sheet)+r'!)?([\d:、，,\s]+)',block):
            left,right=numbers(match[1]),numbers(match[2])
            if len(left)==len(right):rows.extend(zip(left,right))
        columns=re.findall(r'(?<![A-Za-z0-9:/、])([A-Z]+)(?![A-Za-z0-9:/、])[^A-Z0-9\n→]*→\s*(?:B\s+)?([A-Z]+)(?![A-Za-z0-9:/、])',block)
        columns.extend(re.findall(r'A\s*(?:的)?\s*([A-Z]+)\s*列\s*对应\s*B\s*(?:的)?\s*([A-Z]+)\s*列',block))
        column_group=r'([A-Z]+(?:[/、][A-Z]+)+)'
        for pair in re.finditer(r'A(?:表)?(?:的)?(?:期末|期初|年末|年初|本期|上期)?(?:列)?\s*'+column_group+
                r'(?:列)?\s*(?:分别)?(?:对应|采用)\s*B(?:表)?(?:的)?(?:期末|期初|年末|年初|本期|上期)?(?:列)?\s*'+column_group,block):
            left,right=re.split(r'[/、]',pair[1]),re.split(r'[/、]',pair[2])
            if len(left)==len(right):columns.extend(zip(left,right))
        for pair in re.finditer(r'A\s*[^A-Z。\n]*?([A-Z]+(?::[A-Z]+){2,})\s*→\s*B\s*[^A-Z。\n]*?([A-Z]+(?::[A-Z]+){2,})',block):
            left,right=pair[1].split(':'),pair[2].split(':')
            if len(left)==len(right):columns.extend(zip(left,right))
        table_pairs=set();area=r'([A-Z]+\d+:[A-Z]+\d+)'
        for pair in re.finditer(r'A\s*(?:'+re.escape(target_sheet)+r'!)?[^A-Z0-9。\n]*?'+area+
                r'\s*(?:→|[^。；;\n]*?对应)\s*B\s*(?:'+re.escape(source_sheet)+r'!)?'+area,block):
            if re.search(r'不能|不(?:直接|按[^，。]*?)?对应',pair[0]):continue
            left,_,right,_=range_boundaries(pair[1]);sl,_,sr,_=range_boundaries(pair[2])
            if right-left==sr-sl:
                columns.extend((get_column_letter(c),get_column_letter(sl+c-left)) for c in range(left,right+1))
                old,new=addresses(pair[1]),addresses(pair[2])
                if len(old)==len(new):table_pairs.update((target_sheet+'!'+t,source_sheet+'!'+s) for t,s in zip(old,new))
        for pair in re.finditer(r'A\s*'+re.escape(target_sheet)+r'!'+area+r'[^。\n]*?对应B第(\d+)行',block):
            left,first,right,last=range_boundaries(pair[1])
            if first==last and all(get_column_letter(c) in dict(columns) for c in range(left,right+1)):
                rows.append((first,int(pair[2])))
        cells=r'[A-Z]+\d+(?:/[A-Z]+\d+)+'
        for statement in re.finditer(r'([^。\n]+?)[；;]\s*分别按A([^。\n]+?对应)',block):
            groups=lambda part:{role:value for role,value in re.findall(r'(期末|期初|本期|上期)[^A-Z。\n]*?('+cells+r')',part)}
            sources=groups(statement[1]);target_groups=groups(statement[2])
            for role,left in target_groups.items():
                if role in sources:connect(left,sources[role],source_sheet)
        if re.search(r'上述每一行[^。\n]*列对应',block):
            for pair in re.finditer(r'A\s*'+re.escape(target_sheet)+r'!([A-Z]+)(\d+)[^\n↔→]*[↔→]\s*B\s*'+re.escape(source_sheet)+r'![A-Z]+(\d+)',block):
                rows.append((int(pair[2]),int(pair[3])))
            for pair in re.finditer(r'B\s*'+re.escape(source_sheet)+r'!([^；;。\n]+?)\s*[↔→]\s*A\s*'+re.escape(target_sheet)+r'!([^；;。\n]+)',block):
                source_cols=re.findall(r'(?<![A-Za-z0-9_])([A-Z]+)(?![A-Za-z0-9_])',pair[1])
                target_cols=re.findall(r'(?<![A-Za-z0-9_])([A-Z]+)(?![A-Za-z0-9_])',pair[2])
                if len(source_cols)==len(target_cols):columns.extend(zip(target_cols,source_cols))
        row_bindings=[]
        for old,new in rows:
            for left,right in columns:
                target=f'{target_sheet}!{left}{old}'
                if target in targets:row_bindings.append({'target':target,'source':f'{source_sheet}!{right}{new}'})
        # 表区对应只给出边界；明示的明细行关系不能被整表顺序展开覆盖。
        row_targets={row['target'] for row in row_bindings}
        bindings=[row for row in bindings if row['target'] not in row_targets or (row['target'],row['source']) not in table_pairs]
        bindings.extend(row_bindings)
    # 明说来源行及期间栏时，列号只复用原答复已指出的该表期间标签。
    if len(target_sheets)==1:
        for line in text.splitlines():
            source_row=re.search(r'→\s*B[“"]([^”"]+)[”"]第(\d+)行[^。\n]*对应期末、期初栏',line)
            if not source_row or source_row[1] not in sheets:continue
            source_sheet=source_row[1];roles={}
            for prior in b.get('合集来源',[]):
                for paragraph in prior.get('识别答复','').replace('`','').replace('**','').split('\n\n'):
                    if not re.search(r'B\s*的相应期间标签位置为\s*'+re.escape(source_sheet),paragraph):continue
                    for col,role in re.findall(r'([A-Z]+)\s*列为(期末|期初)',paragraph):roles.setdefault(role,set()).add(col)
            for cell,role in re.findall(r'([A-Z]+\d+)\s*(期末|期初)',line.split('→')[0]):
                target=next(iter(target_sheets))+'!'+cell
                if target in targets and len(roles.get(role,()))==1:
                    bindings.append({'target':target,'source':source_sheet+'!'+next(iter(roles[role]))+source_row[2]})
    bindings=list({(row['target'],row['source']):row for row in bindings if row['target'] not in retained}.values())
    label_actions=named_labels
    if len(target_sheets)==1:
        target_sheet=next(iter(target_sheets));business=dict(_records(a))
        for stated in re.finditer(r'A[^。\n]*?明细行按B[^。\n]*?排列[:：]\s*(.+?)[；;]\s*对应A行\s*(\d+)\s*[—–至-]\s*(\d+)',reply):
            names=re.findall(r'`([^`]+)`',stated[1])
            rows=list(range(int(stated[2]),int(stated[3])+1))
            if len(names)!=len(rows):continue
            for row,name in zip(rows,names):
                labels=[item for item in a.get('excluded',[]) if item.get('category')=='label'
                    and item['sheet']==target_sheet and coordinate_to_tuple(item['cell'])[0]==row]
                if len(labels)==1:
                    label_actions.append({'action':'set','target':target_sheet+'!'+labels[0]['cell'],
                        'value':name,'original':stated[0]})
        for line in text.splitlines():
            if not any(word in line for word in ('名称','标签','标题','表头')):continue
            stated=[name for name in target_sheets if name in line]
            sheet=stated[0] if len(stated)==1 else target_sheet
            for item in re.finditer(r'([A-Z]+\d+)(?:[“"]([^”"]*)[”"])?\s*(?:标签|标题|名称)?\s*(?:改为|更新为|替换为|设为)\s*[“"]([^”"]*)[”"]',line):
                label_actions.append({'action':'set','target':sheet+'!'+item[1],
                    'value':item[3],'original':item[0]})
            replaced=re.search(r'(?<![\w!])(?:'+re.escape(sheet)+r'!)?([A-Z]+\d+)[^。\n]*?原文[“"]([^”"]*)[”"][；;：:\s]*(?:替换为|更新为|改为)[“"]([^”"]*)[”"]',line)
            if replaced:
                label_actions.append({'action':'set','target':sheet+'!'+replaced[1],
                    'value':replaced[3],'original':replaced[0]})
            ordered=re.search(r'((?:[A-Z]+\d+[、，,及和\s]*)+)分别(?:明确为|改为|更新为|设为)\s*(.+)',line)
            if ordered:
                cells=re.findall(r'[A-Z]+\d+',ordered[1]);values=re.findall(r'[“"]([^”"]*)[”"]',ordered[2])
                if len(cells)==len(values):
                    label_actions.extend({'action':'set','target':sheet+'!'+cell,'value':value,'original':ordered[0]}
                                         for cell,value in zip(cells,values))
        for block in re.split(r'\n(?=#{1,6}\s)',text):
            if '局部期间表头' in block and '替换为' in block:
                areas=[area for values in a['机械记录'].get('names',{}).values() for area in values
                       if area['sheet']==target_sheet and area['range'].replace('$','') in block]
                cells=next(s.get('cells',{}) for s in a['机械记录']['sheets'] if s['name']==target_sheet)
                for line in block.splitlines():
                    change=re.search(r'([A-Z]+(?:、[A-Z]+)*)列(.+?)\s*→\s*(.+)',line)
                    if not change:continue
                    cols=change[1].split('、');old=re.findall(r'[“"]([^”"]+)[”"]',change[2]);new=change[3].strip().split('、')
                    if len(cols)!=len(old) or len(cols)!=len(new):continue
                    for area in areas:
                        for col,prior,replacement in zip(cols,old,new):
                            found=[cell for cell in addresses(area['range']) if re.match(r'[A-Z]+',cell)[0]==col and str(cells.get(cell,{}).get('value','')).strip()==prior]
                            if len(found)==1:label_actions.append({'action':'set','target':target_sheet+'!'+found[0],
                                'value':replacement.strip(' “”。'),'original':line})
            values=dict(re.findall(r'(期末|期初|本期|上期)角色(?:的表头文字)?为[“"]([^”"]+)[”"]',block))
            if not values:continue
            for line in block.splitlines():
                if not re.match(r'\s*[-*]\s*A表',line):continue
                for clause in re.split(r'[；;]',line):
                    label=re.search(re.escape(target_sheet)+r'!([A-Z]+\d+(?::[A-Z]+\d+)?)',clause)
                    if not label or not re.search(r'列头|表头',clause):continue
                    cells=addresses(label[1]);roles=re.findall(r'期末|期初|本期|上期',clause)
                    if not roles:
                        for cell in cells:
                            column=re.match(r'[A-Z]+',cell)[0]
                            found={metric.get('template_period_role') for table in a.get('outline',{}).get('subtables',[])
                                if table['sheet']==target_sheet and cell in addresses(table['range'])
                                for metric in table.get('metrics',[]) if any(re.match(r'[A-Z]+',part)[0]==column
                                    for span in metric['ranges'] for part in addresses(table['range'] if span=='*' else span))}
                            found-={None}
                            if len(found)==1:roles.append({'closing':'期末','opening':'期初','current':'本期','prior':'上期'}.get(next(iter(found)),next(iter(found))))
                    if len(cells)==len(roles) and all(role in values for role in roles):
                        label_actions.extend({'action':'set','target':target_sheet+'!'+cell,
                            'value':values[role],'original':clause} for cell,role in zip(cells,roles))
        for block in re.split(r'\n(?=\s*-\s+)',text):
            span=re.search(re.escape(target_sheet)+r'!([A-Z]+)(\d+):([A-Z]+)\2',block)
            replacement=re.search(r'替换为[:：]([^\n]+)',block)
            if span and replacement:
                for column,value in re.findall(r'([A-Z]+)[“"]([^”"]+)[”"]',replacement[1]):
                    if column in (span[1],span[3]):
                        label_actions.append({'action':'set','target':target_sheet+'!'+column+span[2],'value':value,'original':block})
        for label in re.finditer(re.escape(target_sheet)+r'!([A-Z]+\d+:[A-Z]+\d+)合并表头为[“"]([^”"]+)[”"]',text):
            label_actions.extend({'action':'set','target':target_sheet+'!'+cell,'value':label[2],'original':label[0]}
                                 for cell in addresses(label[1]))
        for label in re.finditer(r'将\s*A\s*'+re.escape(target_sheet)+r'!([A-Z]+\d+(?::[A-Z]+\d+)?)\s*(?:合并单元格)?的文字替换为[:：]\s*[“"]([^”"]+)[”"]',text):
            label_actions.extend({'action':'set','target':target_sheet+'!'+cell,'value':label[2],'original':label[0]}
                                 for cell in addresses(label[1]))
        anchors={id(mapping):{sheet['name']+'!'+cell:sheet['name']+'!'+area.split(':')[0]
            for sheet in mapping['机械记录']['sheets'] for area in sheet.get('merges',[]) for cell in addresses(area)}
            for mapping in (a,b)}
        for line in text.splitlines():
            if not any(word in line for word in ('名称','标签','标题','表头','保留原文')):continue
            names=re.search(r'原A位置\s*'+re.escape(target_sheet)+r'!([A-Z]+\d+:[A-Z]+\d+)[^\n]*?名称依B原文更新为([^。\n]+)',line)
            if names:
                cells=addresses(names[1]);values=re.findall(r'[“"]([^”"]+)[”"]',names[2])
                if len(cells)==len(values):label_actions.extend({'action':'set','target':target_sheet+'!'+cell,
                    'value':value,'original':names[0]} for cell,value in zip(cells,values))
            for adopted in re.finditer(r'\bA\s*([A-Z]+\d+(?::[A-Z]+\d+)?)[^，。；;\n]*?采用\s*B\s*([^，。；;\n]+)',line):
                left=list(dict.fromkeys(anchors[id(a)].get(loc,loc) for loc in locations(adopted[1],target_sheets,target_sheet)))
                right=list(dict.fromkeys(anchors[id(b)].get(loc,loc) for loc in locations(adopted[2],sheets,source_default)))
                if len(left)==len(right)==1:label_actions.append({'action':'copy','target':left[0],'source':right[0]})
            explicit=re.search(r'(?:表头|标签|标题)[:：]\s*A\s*(.+?)\s*(?:→|←|对应)\s*B\s*(.+)',line)
            if explicit:
                left=locations(explicit[1],target_sheets,target_sheet)
                right=locations(re.split(r'[，,。]',explicit[2])[0],sheets,source_default)
                if len(left)==len(right):label_actions.extend({'action':'copy','target':t,'source':s} for t,s in zip(left,right))
            stated_name=re.search(r'(?:名称|标签|标题)\s*(?:采用|取自|更新为)\s*(?:B\s*)?(.+)',line)
            if stated_name:
                left=locations(line[:stated_name.start()],target_sheets,target_sheet)
                right=locations(stated_name[1],sheets,source_default)
                if len(left)==len(right)==1:
                    label_actions.append({'action':'copy','target':left[0],'source':right[0]})
            stated=[target_sheet+'!'+cell for span in re.findall(r'\bA!([A-Z]+\d+(?::[A-Z]+\d+)?)',line)
                    for cell in addresses(span)]
            labels=list(dict.fromkeys(anchors[id(a)].get(loc,loc) for loc in stated if loc not in business))
            if re.search(r'(?<!不)保留原(?:标签|文)',line):
                label_actions.extend({'action':'retain','target':loc} for loc in labels)
            elif '对应采用' in line:
                named=re.findall(r'('+ '|'.join(re.escape(name) for name in sorted(sheets,key=len,reverse=True))+
                    r')!([A-Z]+\d+(?::[A-Z]+\d+)?)',line)
                sources=list(dict.fromkeys(anchors[id(b)].get(sheet+'!'+cell,sheet+'!'+cell)
                    for sheet,span in named for cell in addresses(span)))
                if len(labels)==len(sources)==1:
                    label_actions.append({'action':'copy','target':labels[0],'source':sources[0]})
        for label in re.finditer(r'('+r'(?:'+re.escape(target_sheet)+r'!)?[A-Z]+\d+(?::[A-Z]+\d+)?'+r')[（(]合并表头[“"]([^”"]+)[”"][）)]\s*对应\s*(\d{4}年\d{1,2}月\d{1,2}日)',text):
            for location in locations(label[1],target_sheets,target_sheet):
                label_actions.append({'action':'set','target':location,'value':label[2]+'（'+label[3]+'）','original':label[0]})
        label_actions=[{**action,'target':anchors[id(a)].get(action['target'],action['target']),
            **({'source':anchors[id(b)].get(action['source'],action['source'])} if 'source' in action else {})} for action in label_actions]
        label_actions=list({json.dumps(action,ensure_ascii=False,sort_keys=True):action for action in label_actions}.values())
    captions=[]
    for paragraph in text.split('\n\n'):
        header=re.search(r'A表头\s*(.+?)\s*的[“"]',paragraph)
        changes=re.findall(r'将[“"]([^”"]+)[”"]替换为[“"]([^”"]+)[”"]',paragraph)
        if not (header and changes and '仅就A' in paragraph and '不据此修改其他表' in paragraph):continue
        refs=set(locations(header[1],target_sheets,next(iter(target_sheets)) if len(target_sheets)==1 else None))
        owners={name for name,areas in a['机械记录'].get('names',{}).items()
            if refs and any(refs<={area['sheet']+'!'+c for c in addresses(area['range'])} for area in areas)}
        if len(owners)!=1:continue
        name=next(iter(owners));titles={item['title'] for item in a.get('source_context',[])
            if isinstance(item,dict) and item.get('range_name')==name and item.get('title')}
        if len(titles)==1:
            title=next(iter(titles))
            # 共用原说明不动；仅把模型明列的新文字附到它指定的表旁。
            captions.append({'range_name':name,'old_text':title,'text':title+' '+ '；'.join(new for _,new in changes)})
    for block in re.split(r'\n(?=#{1,6}\s)',text):
        note=re.search(r'本表旁[^。\n]*?以下完整文字[^\n]*\n\s*[“"]([^”"]+)[”"]',block)
        if not note:continue
        names={name for name,areas in a['机械记录'].get('names',{}).items() for area in areas
            if area['sheet']+'!'+area['range'].replace('$','') in block.splitlines()[0]}
        if len(names)!=1:continue
        name=next(iter(names));titles={item['title'] for item in a.get('source_context',[])
            if isinstance(item,dict) and item.get('range_name')==name and item.get('title')}
        if len(titles)==1:
            title=next(iter(titles));captions.append({'range_name':name,'old_text':title,'text':title+' '+note[1]})
    for block in re.split(r'\n(?=TA_\d+)',text):
        owner=re.match(r'(TA_\d+)',block.strip())
        replacement=re.search(r'将原标题[“"]([^”"]+)[”"]替换为[：:]\s*\n+\s*>\s*([^\n]+)',block)
        if owner and replacement:
            captions.append({'range_name':owner[1],'old_text':replacement[1],'text':replacement[2].strip()})
        note=re.match(r'(TA_\d+(?:[、，,]\s*TA_\d+)*)[：:][^\n]*增加以下说明[：:]\s*\n+\s*>\s*([^\n]+)',block.strip())
        if note:
            name=re.findall(r'TA_\d+',note[1])[0]
            titles={item['title'] for item in a.get('source_context',[]) if isinstance(item,dict) and item.get('range_name')==name and item.get('title')}
            if len(titles)==1:
                title=next(iter(titles))
                # 共用说明只附在明列的首张表标题后，文字照搬，不另造期间标签。
                captions.append({'range_name':name,'old_text':title,'text':title+' '+note[2].strip()})
    original=re.search(r'原(?:段落|文|说明)[:：]\s*[“"]([^”"]+)[”"]',text)
    replacement=re.search(r'(?:替换文字|更新后文字)[^\n]*?(?<![A-Za-z0-9_])(TA_\d+)(?![A-Za-z0-9_])[^\n]*?[:：]\s*[“"]([^”"]+)[”"]',text)
    if original and replacement:
        captions.append({'range_name':replacement[1],'old_text':original[1],'text':replacement[2]})
    target_area=re.search(r'目标表\s*[:：][^\n]*?([^ \t:：;；,，]+)!([A-Z]+\d+:[A-Z]+\d+)',text)
    title_change=re.search(r'(?:表头)?标题\s*[:：]\s*原文[“"]([^”"]+)[”"][；;：:\s]*(?:替换为|更新为|改为)[“"]([^”"]+)[”"]',text)
    if target_area and title_change:
        names=[name for name,areas in a['机械记录'].get('names',{}).items()
            if any(area['sheet']==target_area[1] and area['range'].replace('$','')==target_area[2] for area in areas)]
        if len(names)==1:
            captions.append({'range_name':names[0],'old_text':title_change[1],'text':title_change[2]})
    source_period_note=re.search(r'(本次纳入的B相关附注表[^。；\n]+)[；;]因此对这些目标表适用的A期间说明',text)
    if source_period_note:
        selected=set(targets)
        for name,areas in a['机械记录'].get('names',{}).items():
            if not any(selected&{area['sheet']+'!'+cell for cell in addresses(area['range'])} for area in areas):continue
            titles={item['title'] for item in a.get('source_context',[]) if isinstance(item,dict) and item.get('range_name')==name and item.get('title')}
            if len(titles)==1:
                title=next(iter(titles))
                # 保留原答复对B来源的限定，避免把旧成员原值一并改标期间。
                captions.append({'range_name':name,'old_text':title,'text':title+' '+source_period_note[1]+'。'})
    local_dates=re.search(r'标明(期末为\s*\d{4}-\d{2}-\d{2}、期初为\s*\d{4}-\d{2}-\d{2})',text)
    if not local_dates and '仅在本表局部期间说明处标明上述实际日期' in text:
        local_dates=re.search(r'(本表期末采用的实际日期为\d{4}-\d{2}-\d{2}[；;]期初采用的B实际日期为\d{4}-\d{2}-\d{2}。)',text)
    if local_dates and ('本表标题旁的局部期间文字相应更新' in text or '仅在本表局部期间说明处标明上述实际日期' in text):
        names={name for name,areas in a['机械记录'].get('names',{}).items()
            if any(set(targets)&{area['sheet']+'!'+cell for cell in addresses(area['range'])} for area in areas)}
        if len(names)==1:
            name=next(iter(names));titles={item['title'] for item in a.get('source_context',[])
                if isinstance(item,dict) and item.get('range_name')==name and item.get('title')}
            if len(titles)==1:
                title=next(iter(titles));captions.append({'range_name':name,'old_text':title,'text':title+' '+local_dates[1]})
    for block in re.split(r'\n(?=#{1,6}\s)',text):
        added_period=re.search(r'(新增(?:成员按B来源期间角色保留实际日期|匿名行继承B的期末角色)[^。\n]+。)',block)
        target_area=re.search(r'A\s*([^!\s]+)!([A-Z]+\d+:[A-Z]+\d+)',block)
        if not (added_period and target_area and '原独立合计之前' in block):continue
        refs={target_area[1]+'!'+cell for cell in addresses(target_area[2])}
        if not refs&set(targets):continue
        names=[name for name,areas in a['机械记录'].get('names',{}).items()
               if any(refs<={area['sheet']+'!'+cell for cell in addresses(area['range'])} for area in areas)]
        if len(names)!=1:continue
        name=names[0];titles={item['title'] for item in a.get('source_context',[]) if isinstance(item,dict) and item.get('range_name')==name and item.get('title')}
        if len(titles)!=1:continue
        title=next(iter(titles));previous=[p['text'] for r in b.get('已确认关系',[]) for p in r.get('caption_updates',[])
            if p['range_name']==name and p['old_text']==title]
        base=max([title,*previous],key=len)
        if any(not base.startswith(value) for value in previous):continue
        captions.append({'range_name':name,'old_text':title,'text':base if added_period[1] in base else base+' '+added_period[1]})
    return {'kind':'取数对应','bindings':bindings,'retained_targets':list(dict.fromkeys([*retained,*replaced_targets])),'label_actions':label_actions,'caption_updates':captions,
            'requested_targets':list(targets),
            'arranged_targets':arranged_targets,'content':reply,'answer_hash':definition_hash(reply),
            'source_a_hash':a['source_hash'],'source_b_hash':b['source_hash'],'gold_hash':a['gold_hash']}


def save_source_relations(a, b, reply, targets, destination, *, usage=None, arranged_targets=False):
    raise ValueError("旧位置取数安排入口已停用；只保存本格身份证。")


def append_declared_relation(relations, relation):
    """同一表区的新明确安排替换旧安排，原答复记录继续保留。"""
    relation=deepcopy(relation)
    for old in relations:
        if not relation.get('answer_hash') or (old.get('kind'),old.get('answer_hash'),old.get('dimension'))!=(relation.get('kind'),relation.get('answer_hash'),relation.get('dimension')):continue
        if relation['kind']=='取数对应':
            targets=set(relation.get('requested_targets',[]))|set(relation.get('retained_targets',[]))|{r['target'] for r in relation.get('bindings',[])}
            relation['bindings']=[r for r in old.get('bindings',[]) if r['target'] not in targets]+relation.get('bindings',[])
            relation['retained_targets']=[loc for loc in old.get('retained_targets',[]) if loc not in targets]+relation.get('retained_targets',[])
            for key,field in [('label_actions','target'),('caption_updates','range_name')]:
                changed={r[field] for r in relation.get(key,[])}
                relation[key]=[r for r in old.get(key,[]) if r[field] not in changed]+relation.get(key,[])
        elif relation['kind']=='名单编排':
            changed={(r['range_name'],r['first_row_offset']) for r in relation.get('tables',[])}
            relation['tables']=[r for r in old.get('tables',[]) if (r['range_name'],r['first_row_offset']) not in changed]+relation.get('tables',[])
    relations[:]=[row for row in relations if (row.get('kind'),row.get('answer_hash'),row.get('dimension'))!=
                  (relation.get('kind'),relation.get('answer_hash'),relation.get('dimension'))]
    replacements={(table['range_name'],table['first_row_offset']):table for table in relation.get('tables',[])}
    requested=deepcopy(replacements)
    for prior in relations:
        if prior.get('kind')=='名单编排':
            for old in prior.get('tables',[]):
                new=replacements.get((old['range_name'],old['first_row_offset']))
                if not new or 'row_columns' not in new or new['source_sheet']!=old['source_sheet']:continue
                # 局部来源列只替换该成员明列的列，其余仍引用原安排。
                for row,columns in new['row_columns'].items():
                    if int(row) not in old['source_rows']+[t['source_row'] for t in old.get('totals',[])]:continue
                    for col,source_col in old.get('row_columns',{}).get(row,old['columns']).items():
                        if col in columns:continue
                        columns[col]=source_col
                        new.setdefault('row_sources',{}).setdefault(row,{})[col]=old.get('row_sources',{}).get(row,{}).get(col) or (
                            f"{old['source_sheet']}!{source_col}{int(row)+old.get('column_row_offsets',{}).get(col,0)}")
                new['labels']={**old.get('labels',{}),**new.get('labels',{})}
            prior['tables']=[table for table in prior.get('tables',[])
                if (table['range_name'],table['first_row_offset']) not in replacements]
        elif prior.get('kind')=='取数对应':
            prior['bindings']=[item for item in prior.get('bindings',[]) if not (
                (member:=item.get('target_member')) and
                (table:=requested.get((member['range_name'],member['first_row_offset']))) and
                member['column'] in table.get('row_columns',{}).get(str(member['source_row']),table['columns']) and
                all(value is not None for value in table.get('row_columns',{}).get(str(member['source_row']),table['columns']).values()))]
    relations.append(deepcopy(relation))


def save_list_relations(a, b, reply, dimension, destination, *, usage=None, source_ranges=()):
    raise ValueError("旧名单执行安排入口已停用；只保存本格身份证。")








def _per_cell_answer(reply, standard, sheet_names, default_sheet=None):
    """仅读取各格自己写出的字段，不承接共同声明或其他格的取值。"""
    class ConflictingValues(list):
        pass
    def field_pairs(items):
        value={}
        for key,item in items:
            if isinstance(value.get(key),ConflictingValues):value[key].append(item)
            elif key in value and value[key]!=item:
                if not isinstance(value[key],ConflictingValues):value[key]=ConflictingValues([value[key]])
                value[key].append(item)
            else:value[key]=item
        return value
    def conflicting(value):
        if isinstance(value,ConflictingValues):return True
        if isinstance(value,dict):return any(conflicting(item) for item in value.values())
        return isinstance(value,list) and any(conflicting(item) for item in value)
    result={'metrics':[],'dimension_bindings':[],'dimensions':[],'unresolved':[]}
    section_sheet=None
    single=re.compile(r"(?:(?:'([^']+)'|([^!]+))!)?\$?([A-Z]+)\$?([1-9]\d*)$")
    def location(value):
        match=single.fullmatch(str(value).strip())
        if not match:return None
        sheet=match[1] or match[2] or section_sheet or default_sheet or (sheet_names[0] if len(sheet_names)==1 else None)
        return sheet+'!'+match[3]+match[4] if sheet in sheet_names else None
    def table_location(value):
        parts=re.split(r'[；;]',value,maxsplit=1)
        if len(parts)>1 and re.search(r'(?<![A-Za-z0-9_])\$?[A-Z]+\$?[1-9]\d*(?![A-Za-z0-9_])',parts[1]):return None
        return location(parts[0].strip().strip('`'))
    def keep(where,fields,original):
        own={k:v for k,v in fields.items() if k in standard['dimension'] or k in ('metric_id','dimension_id','template_period_role')}
        grouped=fields.get('dimensions',{})
        if isinstance(grouped,ConflictingValues):
            result['unresolved'].append({'id':'dimensions','ranges':[where],'reason':'该格维度分组重复且取值不同',
                'conflicting_values':deepcopy(grouped),'original':deepcopy(original)})
            grouped={}
        if isinstance(grouped,dict):
            for key,value in grouped.items():
                if key not in standard['dimension'] and key!='template_period_role':continue
                if key in own and own[key]!=value:
                    result['unresolved'].append({'id':key,'ranges':[where],'reason':'该格同一维度重复且取值不同',
                        'conflicting_values':[deepcopy(own.pop(key)),deepcopy(value)],'original':deepcopy(original)})
                else:own[key]=value
        for key in list(own):
            if conflicting(own[key]):
                result['unresolved'].append({'id':key,'ranges':[where],'reason':'该格字段内部重复且取值不同',
                    'conflicting_values':deepcopy(own.pop(key)),'original':deepcopy(original)})
        if (isinstance(own.get('metric_id'),str) and own['metric_id'] in standard['metric']
                and isinstance(own.get('dimension_id'),str) and own['dimension_id'] in standard['dimension']):
            for key in ('metric_id','dimension_id'):
                result['unresolved'].append({'id':key,'ranges':[where],'reason':'该格同时给出槽位编号和维度编号，不能代选',
                    'conflicting_values':[deepcopy(own.pop(key))],'original':deepcopy(original)})
        if own.get('metric_id') is None and isinstance(own.get('dimension_id'),str) and own['dimension_id'] in standard['dimension']:
            own.pop('metric_id',None)
        if own.get('dimension_id') is None and isinstance(own.get('metric_id'),str) and own['metric_id'] in standard['metric']:
            own.pop('dimension_id',None)
        for key,value in own.items():
            row={'id':key,'ranges':[where],'original':deepcopy(original),'evidence':[],'逐格原字段':True}
            if key in ('metric_id','dimension_id'):
                category='metric' if key=='metric_id' else 'dimension'
                if isinstance(value,str) and value in standard[category]:
                    result['metrics' if key=='metric_id' else 'dimension_bindings'].append({**row,'id':value,
                        **({'template_period_role':deepcopy(own['template_period_role'])} if 'template_period_role' in own else {})})
                else:result['unresolved'].append({**row,'reason':'该格未给出现有金标准编号','value':deepcopy(value)})
                continue
            row['value']=deepcopy(value)
            evidence=fields.get('dimension_evidence',{})
            row['evidence']=deepcopy(evidence.get(key,[]) if isinstance(evidence,dict) else [])
            labels=fields.get('label_references',{})
            refs=labels.get(key) if isinstance(labels,dict) else None
            if refs is not None:row['label_references']=deepcopy(refs)
            result['dimensions'].append(row)
        labels=fields.get('label_references',{})
        if isinstance(labels,dict):
            for key,refs in labels.items():
                if key in standard['dimension'] and key not in own:
                    result['dimensions'].append({'id':key,'ranges':[where],
                        'label_references':deepcopy(refs),'evidence':[],'original':deepcopy(original)})
    def objects(value):
        if isinstance(value,list):
            for item in value:objects(item)
        elif isinstance(value,dict):
            if any(conflicting(value[k]) for k in ('coordinate','location','坐标','位置','range') if k in value):
                result['unresolved'].append({'id':'location','reason':'本格位置字段重复且内容不同',
                    'conflicting_values':[deepcopy(value)],'original':deepcopy(value)})
                return
            own=next((location(value[k]) for k in ('coordinate','location','坐标','位置','range') if k in value and location(value[k])),None)
            if own:keep(own,value,value);return
            for key,item in value.items():
                own=location(key)
                if own and isinstance(item,ConflictingValues):
                    result['unresolved'].append({'id':'location','ranges':[own],'reason':'同一本格位置重复且内容不同',
                        'conflicting_values':deepcopy(item),'original':deepcopy(item)})
                elif own and isinstance(item,dict):keep(own,item,item)
                elif isinstance(item,(dict,list)):objects(item)
    text=reply.strip() if isinstance(reply,str) else None
    if text is not None:
        def fenced_fields(match):
            try:native=json.loads(match[1],object_pairs_hook=field_pairs)
            except ValueError:return match[0]
            if not isinstance(native,(dict,list)):return match[0]
            objects(native)
            return ''
        text=re.sub(r'^```(?:json)?[ \t]*\r?\n(.*?)\r?\n```[ \t]*$',
                    fenced_fields,text,flags=re.M|re.S)
        try:value=json.loads(text,object_pairs_hook=field_pairs)
        except ValueError:
            try:
                value=[];remaining=text.strip();decoder=json.JSONDecoder(object_pairs_hook=field_pairs)
                if remaining.startswith('[') and remaining.endswith(']'):
                    try:decoder.raw_decode(remaining)
                    except ValueError:
                        remaining=remaining[1:-1].strip()
                while remaining:
                    item,end=decoder.raw_decode(remaining)
                    value.append(item);remaining=remaining[end:].lstrip()
                    if remaining.startswith(','):
                        remaining=remaining[1:].lstrip()
                        if not remaining:raise ValueError('缺少下一个完整JSON对象')
            except ValueError:
                if value:
                    objects(value);text=remaining
                value=None
    else:value=reply
    if isinstance(value,(dict,list)):
        objects(value)
    elif text is not None:
        field_ids=set(standard['dimension'])|{'metric_id','dimension_id','dimensions','template_period_role','dimension_evidence','label_references'}
        table_fields=field_ids|{'dimensions'}
        field_line=re.compile(r'^\s*(?:[-*]\s+)?`?('+'|'.join(re.escape(k) for k in sorted(table_fields,key=len,reverse=True))+r')`?\s*[:：=＝→]\s*(?:`([^`]*)`|(\{.*\}|\[.*\]))\s*$')
        keys='|'.join(re.escape(k) for k in sorted(field_ids,key=len,reverse=True))
        dimension_keys='|'.join(re.escape(k) for k in sorted(standard['dimension'],key=len,reverse=True))
        marker=re.compile(r'(?<![\w.])`?(?:dimensions\.(?=(?:'+dimension_keys+r')`?\s*[:：=＝→]))?('+keys+r')`?\s*(?:(?:(?:[；;]\s*)?实际取值\s*)?[:：=＝→]|为\s*|(?=`))\s*')
        address=r"(?:(?:'[^']+'|[^\s!：:;；|]+)!)?\$?[A-Z]+\$?[1-9]\d*(?::\$?[A-Z]+\$?[1-9]\d*)?"
        heading=re.compile(r'^\s*(?:[-*#]+\s*|\d+[.)、]\s*)?(?:(?:`?location`?|(?:[AB]\s*)?(?:来源|原?(?:位置|坐标)))(?:\s*[:：=＝→]\s*|\s+))?`?('+address+r')`?\s*(?:[:：;；—|｜(（-]|$)')
        blocks=[];current=None;headers=None
        for line in text.replace('**','').splitlines():
            try:native,end=json.JSONDecoder(object_pairs_hook=field_pairs).raw_decode(line.lstrip())
            except ValueError:native=None
            if isinstance(native,(dict,list)):
                objects(native)
                tail=line.lstrip()[end:].strip()
                own=location(native.get('location')) if isinstance(native,dict) else None
                pairs=[]
                while own and tail.startswith(','):
                    rest=tail[1:].lstrip()
                    try:
                        key,stop=json.JSONDecoder().raw_decode(rest)
                        rest=rest[stop:].lstrip()
                        if not isinstance(key,str) or key in ('location','coordinate','坐标','位置','range') or not rest.startswith(':'):break
                        rest=rest[1:].lstrip()
                        value,stop=json.JSONDecoder(object_pairs_hook=field_pairs).raw_decode(rest)
                    except ValueError:break
                    pairs.append((key,value));tail=rest[stop:].lstrip()
                if pairs:keep(own,field_pairs(pairs),line)
                if tail.strip(',}] \t\r\n'):
                    result['unresolved'].append({'reason':'未给出本格归属的原文未接为字段','original':tail})
                continue
            section=re.fullmatch(r'\s*#{1,6}\s+(.+?)\s*',line)
            if section:
                name=section[1].strip('`').strip("'")
                if name in sheet_names:
                    section_sheet=name;current=None;headers=None;continue
                if not location(name):section_sheet=None;current=None;headers=None
            if line.lstrip().startswith('|'):
                values=[v.strip() for v in line.strip().strip('|').split('|')]
                values=[v[1:-1] if v.startswith('`') and v.endswith('`') else v for v in values]
                positions=[(i,where) for i,v in enumerate(values) if (where:=table_location(v)) and
                           not (headers and len(headers)==len(values) and re.search(r'依据|说明|证据',headers[i]))]
                if current is not None and len(values)==2 and not positions:
                    if values[0] in table_fields:
                        if len(current)==3:current.append([])
                        current[3].append((values[0],values[1]))
                        current[2]+='\n'+line
                    continue
                if not positions:
                    if not all(re.fullmatch(r':?-+:?',v) for v in values):headers=values
                    current=None;continue
                if headers:
                    if len(positions)!=1:continue
                    column,own=positions[0]
                    if len(values)!=len(headers):
                        result['unresolved'].append({'ranges':[own],'reason':
                            f'原答复表头{len(headers)}列、本行{len(values)}格；已保留明示标准编号，未移动或补齐其余字段','original':line})
                        blocks.append((own,'',line,[(k,v) for k,v in zip(headers,values)
                            if k in ('metric_id','dimension_id')]))
                        continue
                    blocks.append((own,'；'.join(v for i,(k,v) in enumerate(zip(headers,values))
                        if i!=column and k not in table_fields and not re.search(r'依据|说明|证据',k)),line,
                        [(k,v) for i,(k,v) in enumerate(zip(headers,values)) if i!=column and k in table_fields]))
                continue
            match=heading.match(line)
            if match:
                own=location(match[1]);current=None;current_indent=len(line)-len(line.lstrip())
                current_list=bool(re.match(r'^\s*(?:[-*]\s+|\d+[.)、]\s*)',line))
                if own:current=[own,line[match.end():],line];blocks.append(current)
                else:result['unresolved'].append({'reason':'原答复未逐格写明身份','original':line})
            elif current is not None and (len(line)-len(line.lstrip())>current_indent or
                    not current_list and re.match(r'^[-*]\s+',line) or
                    len(line)-len(line.lstrip())==current_indent and marker.match(re.sub(r'^\s*[-*]\s+','',line).lstrip())):
                literal=field_line.fullmatch(line)
                if literal:
                    if len(current)==3:current.append([])
                    current[3].append((literal[1],literal[2] if literal[2] is not None else literal[3]))
                    current[1]+='\n-\n'
                else:current[1]+='\n'+line
                current[2]+='\n'+line
            elif line.strip():current=None
        for block in blocks:
            own,body,original=block[:3]
            # 字段赋值外的成对括号是语法边界，引用和实际取值里的括号仍保留。
            brackets=[];boundaries=set();quote=None;escaped=False
            pairs={'（':'）','(':')'}
            for index,char in enumerate(body):
                if quote:
                    if escaped:escaped=False
                    elif char=='\\':escaped=True
                    elif char==quote:quote=None
                elif char in "\"'`“‘":quote={'“':'”','‘':'’'}.get(char,char)
                elif char in pairs:brackets.append((index,char))
                elif brackets and char==pairs[brackets[-1][1]]:
                    start,_=brackets.pop()
                    if marker.match(body[start+1:index].lstrip()):boundaries.update((start,index))
            if boundaries:body=''.join('；' if i in boundaries else char for i,char in enumerate(body))
            fields={};conflicts=set();matches=[];field_values=[]
            quoted=[(m.start(),m.end()) for m in re.finditer(r'`[^`]*`',body)]
            # 元数据里的字段名属于引用说明，不是本格第二份维度取值。
            for match in marker.finditer(body):
                if matches:
                    native_text=body[matches[-1].end():].lstrip('` \t\r\n')
                    if native_text.startswith(('{','[','"')):
                        try:_,stop=json.JSONDecoder().raw_decode(native_text)
                        except ValueError:pass
                        else:
                            if match.start()<len(body)-len(native_text)+stop:continue
                if matches and matches[-1][1] in ('metric_id','dimension_id') and not body[matches[-1].end():match.start()].strip(" \t\r\n`\"'“‘"):continue
                grouped=(matches and matches[-1][1]=='dimensions' and
                    any(start+1==match.start(1) for start,end in quoted) and
                    re.match(r'`?'+re.escape(match[1])+r'`?\s*[:：=＝→]',match[0]))
                if matches and not grouped and any(start<=match.start(1)<match.end(1)<=end and
                    not body[matches[-1].end():start].strip() for start,end in quoted):continue
                if not re.search(r'[:：=＝→]|为',match[0]):
                    following=body[match.end():].lstrip('` \t')
                    if not following or following[0] in ',，;；]':continue
                    if any(start<=match.start()<match.end()<=end for start,end in quoted):continue
                if matches and matches[-1][1] in ('label_references','dimension_evidence') and '\n' not in body[matches[-1].end():match.start()]:continue
                matches.append(match)
            for i,match in enumerate(matches):
                raw=body[match.end():matches[i+1].start() if i+1<len(matches) else len(body)].strip()
                raw=re.split(r'\n\s*[-*](?:\s|$)',raw,maxsplit=1)[0].strip()
                line_value=(bool(re.fullmatch(r'[ \t]*[-*][ \t]+',body[body.rfind('\n',0,match.start())+1:match.start()])) and
                            (i+1==len(matches) or '\n' in body[match.end():matches[i+1].start()]))
                if i+1<len(matches) and not line_value:raw=raw.rstrip(' \t\r\n,，;；')
                if match[0].count('`') % 2:
                    end=raw.find('`')
                    if end>=0:
                        raw=raw[:end].strip();line_value=True
                if line_value and (quoted:=re.match(r'^`([^`]*)`',raw)):raw=quoted[1]
                field_values.append((match[1],raw,line_value))
            # 表格已明确字段列边界，值内的分号或字段名不再拆成另一项。
            field_values.extend((key,value,True) for key,value in (block[3] if len(block)>3 else []))
            # 独立字段行与表格列均以原有边界保留文字，不按句内标点截断。
            for key,raw,bounded_value in field_values:
                literal=raw.lstrip('`')
                if key in ('metric_id','dimension_id') and (identifier:=re.fullmatch(r'([A-Za-z_][A-Za-z0-9_.]*)[ \t]*[,，\r\n][ \t\r\n]*[\u4e00-\u9fff].*',literal,re.S)):
                    literal=identifier[1]
                try:
                    parsed,end=json.JSONDecoder(object_pairs_hook=field_pairs).raw_decode(literal)
                    tail=literal[end:].strip()
                    if tail and (bounded_value or tail[0] not in '`。；;') and not (parsed is None and tail[0] in '（('):raise ValueError('非完整字面值')
                except ValueError:
                    quoted=re.match(r'^(?:`([^`]*)`|“([^”]*)”|‘([^’]*)’)',raw)
                    parsed=next(v for v in quoted.groups() if v is not None) if quoted else (literal if bounded_value else re.split(r'[。；;]|\n(?:依据|原文依据)',literal,maxsplit=1)[0].strip().rstrip('`'))
                    if literal.startswith(('{','[')):
                        import yaml
                        # 只补对象冒号后的语法空格；引号内及值中的冒号保持原文。
                        formatted=[];containers=[];quote=None;escaped=False;code=[]
                        for char in literal:
                            if quote=='`':
                                if char=='`':
                                    token=''.join(code)
                                    try:json.loads(token)
                                    except ValueError:token=json.dumps(token,ensure_ascii=False)
                                    formatted.append(token);quote=None
                                else:code.append(char)
                                continue
                            if quote is None and char=='`':quote='`';code=[];continue
                            if quote in ('”','’'):
                                if char==quote:formatted.append(json.dumps(''.join(code),ensure_ascii=False));quote=None
                                else:code.append(char)
                                continue
                            if quote is None and char in ('“','‘') and ''.join(formatted).rstrip().endswith(('{','[',':',',')):
                                quote={'“':'”','‘':'’'}[char];code=[];continue
                            formatted.append(char)
                            if quote:
                                if escaped:escaped=False
                                elif char=='\\':escaped=True
                                elif char==quote:quote=None
                            elif char in '\"\'':quote=char
                            elif char in '{[':containers.append([char,char=='{'])
                            elif char in '}]' and containers:
                                containers.pop()
                                if not containers:break
                            elif char==',' and containers:containers[-1][1]=containers[-1][0]=='{'
                            elif char==':' and containers and containers[-1][1]:
                                formatted.append(' ');containers[-1][1]=False
                        def literal_value(node):
                            if isinstance(node,yaml.MappingNode):return field_pairs((key.value,literal_value(value)) for key,value in node.value)
                            if isinstance(node,yaml.SequenceNode):return [literal_value(value) for value in node.value]
                            if node.style in ('\"',"'"):return node.value
                            try:return json.loads(node.value)
                            except ValueError:return node.value
                        try:
                            if quote is None:parsed=literal_value(yaml.compose(''.join(formatted)))
                        except yaml.YAMLError:pass
                if key in conflicts:continue
                if key in fields and fields[key]!=parsed:
                    result['unresolved'].append({'ranges':[own],'id':key,'reason':'该格字段重复且取值不同','original':original});fields.pop(key)
                    conflicts.add(key)
                else:fields[key]=parsed
            if fields:keep(own,fields,original)
    _require(result['metrics'] or result['dimension_bindings'] or result['dimensions'] or any(r.get('conflicting_values') or r.get('id') in ('metric_id','dimension_id') and 'value' in r and r['value'] is None for r in result['unresolved']),
             '原答复未提供各格自己的金标准身份；原答复保留，不展开共同声明或位置对应')
    return result


def _relocate_direct_reply(reply, mapping, standard, field_scope, source_hash, chain):
    """只按真实执行记录接续旧本格位置，原字段值及原答复不变。"""
    from .表格 import file_hash, read_workbook
    from .A语义阅读 import addresses
    _require(mapping.get('carrier')=='A' and isinstance(reply,str),'旧原答复位置接续只适用于A原文')
    _require(isinstance(source_hash,str) and re.fullmatch(r'[0-9a-f]{64}',source_hash) and chain and field_scope,
             '旧原答复须明确源指纹、实际执行链及本次字段范围')
    validate_direct_collection(mapping,for_matching=True)
    hashes={}
    def check(path,expected,message):
        key=str(Path(path).resolve())
        if key not in hashes:hashes[key]=file_hash(path)
        _require(hashes[key]==expected,message)
    positions=None;expected=source_hash;source_path=None;last_output=None
    for reference in chain:
        _require(isinstance(reference,dict) and reference.get('path') and reference.get('sha256'),
                 '位置接续须引用已保存实际执行记录的文件及指纹')
        check(reference['path'],reference['sha256'],'实际执行记录已变化')
        receipt=json.loads(Path(reference['path']).read_text(encoding='utf-8-sig'),object_pairs_hook=json_dates)
        _require(receipt.get('schema')=='附注更新执行记录-v1','位置接续不是实际执行记录')
        refs=receipt['collections']
        _require(refs['A']['source_hash']==expected,'原答复源文件与实际执行链不一致')
        if source_path is None:source_path=refs['A']['source_path']
        for item in refs.values():
            _require(item['gold_hash']==mapping['gold_hash'],'位置链金标准不一致')
            check(item['mapping_path'],item['mapping_hash'],'实际执行记录的合集已变化')
            check(item['source_path'],item['source_hash'],'实际执行记录的原文件已变化')
            check(item['gold_path'],item['gold_hash'],'实际执行记录的金标准已变化')
        plan_path=Path(receipt['plan_path'])
        seal=json.loads(plan_path.with_name('更新清单校验.json').read_text(encoding='utf-8-sig'),object_pairs_hook=json_dates)
        check(plan_path,seal['sha256'],'实际清单已变化')
        plan=json.loads(plan_path.read_text(encoding='utf-8-sig'),object_pairs_hook=json_dates)
        _require(plan['source_a_hash']==refs['A']['source_hash'] and plan['source_b_hash']==refs['B']['source_hash']
                 and plan['gold_hash']==mapping['gold_hash'] and plan['template_layout']==receipt['template_layout'],
                 '实际清单与位置记录不一致')
        check(receipt['output_path'],receipt['output_hash'],'实际副本已变化')
        step=receipt['template_layout']['positions']
        positions=dict(step) if positions is None else {old:step.get(dest) if dest is not None else None for old,dest in positions.items()}
        expected=receipt['output_hash'];last_output=receipt['output_path']
    _require(expected==mapping['source_hash'] and Path(last_output).resolve()==Path(mapping['source_path']).resolve(),
             '实际位置链没有到达当前A副本')
    snapshot=read_workbook(source_path)
    _require(snapshot['sha256']==source_hash,'原答复源文件已变化')
    requested={sheet+'!'+cell:set(keys) for area,keys in field_scope.items()
               for sheet,span in [area.rsplit('!',1)] for cell in addresses(span)}
    originals={}
    for old,current in positions.items():
        if current in requested:originals.setdefault(current,[]).append(old)
    _require(all(len(originals.get(loc,[]))==1 for loc in requested),'新增或不唯一位置不能继承旧格身份')
    selected={originals[loc][0]:loc for loc in requested}
    sheets={loc.rsplit('!',1)[0] for loc in selected}
    parsed=_per_cell_answer(reply,standard,[sheet['name'] for sheet in snapshot['sheets']],next(iter(sheets)) if len(sheets)==1 else None)
    moved={key:[] for key in parsed}
    for kind,rows in parsed.items():
        for row in rows:
            key={'metrics':'metric_id','dimension_bindings':'dimension_id'}.get(kind,row.get('id'))
            for old in row.get('ranges',[]):
                current=selected.get(old)
                if current is None or key not in requested[current]:continue
                item=deepcopy(row);item['ranges']=[current]
                for field in ('evidence','label_references'):
                    if isinstance(item.get(field),list):
                        item[field]=[positions.get(ref,ref) for ref in item[field] if positions.get(ref,ref) is not None]
                moved[kind].append(item)
    _require(any(moved.values()),'旧原答复在本次实际位置及字段范围内没有逐格原字段')
    return moved,selected


def direct_ingestion(answer, field_scope, standard, sheet_names):
    """核对本次请求与逐格原字段的接入情况；不补身份，不限制后续正确对应。"""
    from .A语义阅读 import addresses
    supplied={};selected={};problems=[]
    for kind in ('metrics','dimension_bindings','dimensions','unresolved'):
        for row in answer.get(kind,[]):
            key={'metrics':'metric_id','dimension_bindings':'dimension_id'}.get(kind,row.get('id'))
            if kind=='unresolved' and not row.get('ranges'):
                problems.append({'reason':row.get('reason')})
            for location in row.get('ranges',[]):
                supplied.setdefault(location,set()).add(key)
                if 'label_references' in row:supplied[location].add('label_references')
                if kind in ('metrics','dimension_bindings'):selected[location]=(kind,row['id'])
                if kind=='unresolved':problems.append({'location':location,'field':key,'reason':row.get('reason')})
    requested={}
    for span,fields in (field_scope or {}).items():
        sheet,sep,area=span.rpartition('!')
        if not sep:
            _require(len(sheet_names)==1,'接入记录的多工作表范围须写明工作表')
            sheet=sheet_names[0];area=span
        for cell in addresses(area.replace('$','')):
            requested.setdefault(sheet.strip("'")+'!'+cell,set()).update(fields)
    missing=[]
    for location,fields in requested.items():
        own=supplied.get(location,set());expected=set(fields)
        primary=expected&{'metric_id','dimension_id'}
        if primary:
            expected-=primary
            if not own&primary:missing.append({'location':location,'field':'metric_id / dimension_id' if len(primary)==2 else next(iter(primary))})
        # 全维度目录只是允许范围，不能要求每格交付与自身标准无关的属性。
        if set(standard['dimension'])<=expected and primary:
            kind,identifier=selected.get(location,(None,None))
            definition=standard['metric' if kind=='metrics' else 'dimension'].get(identifier,{})
            relevant={r['id'] for r in definition.get('dimension_refs',[])}
            if kind=='dimension_bindings':relevant.add(identifier)
            expected-=set(standard['dimension'])-relevant
        missing.extend({'location':location,'field':key} for key in sorted(expected-own))
    scoped=[];outside=[]
    for problem in problems:
        location=problem.get('location')
        chosen=selected.get(location)
        if chosen and problem.get('field') in {'metric_id','dimension_id'} and problem['field']!=('metric_id' if chosen[0]=='metrics' else 'dimension_id'):
            outside.append(problem);continue
        if location is None or (location in requested and problem.get('field') is None) or problem.get('field') in requested.get(location,set()):scoped.append(problem)
        else:outside.append(problem)
    return {'complete':bool(requested) and not missing and not scoped,
            '请求格数':len(requested),'未接入字段':missing,'未确定字段':scoped,'范围外未接入问题':outside}


def direct_recognition_status(mapping):
    """保存成功不能代替本次识别完整接入；历史无接入记录的合集保留待核实。"""
    receipt=mapping.get('本次逐格接入',{})
    return 'complete' if receipt.get('complete') is True else 'partial'


def save_direct_answer(mapping, reply, destination, *, field_scope=None, withdrawn=None, snapshot=None, standard=None, not_adopted_reason=None, reference=None, target_mapping=None, reply_source_hash=None, reply_position_chain=None):
    """正常识别与已有答复共用同一保存出口；不调用模型或读取原表补身份。"""
    from .统一勾稽 import load_fact_mapping
    result=dict(mapping);standard=standard or _standard(mapping);replaced={}
    original_reply=reply
    requested_sheets={area.rsplit('!',1)[0].strip("'") for area in field_scope or {} if '!' in area}
    moved_positions=None
    if reply_position_chain is not None or reply_source_hash is not None:
        reply,moved_positions=_relocate_direct_reply(reply,mapping,standard,field_scope,reply_source_hash,reply_position_chain)
    else:
        try:
            reply=_per_cell_answer(reply,standard,[s['name'] for s in (snapshot if snapshot is not None else mapping['机械记录'])['sheets']],
                next(iter(requested_sheets)) if len(requested_sheets)==1 else None)
        except ValueError as error:
            reply={'metrics':[],'dimension_bindings':[],'dimensions':[],
                   'unresolved':[{'reason':str(error),'original':deepcopy(original_reply)}]}
    scope=field_scope
    if scope is None:
        tables=mapping.get('target_tables') or ([mapping['target_table']] if mapping.get('target_table') else [])
        scope={table['sheet']+'!'+area:['metric_id','dimension_id']
               for table in tables for area in table['data_ranges']}
        if not scope:
            scope={area['sheet']+'!'+area['range']:['metric_id','dimension_id']
                   for area in mapping.get('b_source_ranges',[])}
    result['本次逐格接入']=direct_ingestion(reply,scope,standard,
        [s['name'] for s in (snapshot if snapshot is not None else mapping['机械记录'])['sheets']])
    result['逐格原样保存']=True
    result['逐格字段版本']=1
    result['本次逐格原答复']=original_reply
    if reference is not None:result['A需要']=deepcopy(reference)
    if field_scope is not None:
        combined,outside,replaced=merge_direct_answer(result,reply,field_scope,withdrawn)
        result.pop('target_table',None);result.pop('target_tables',None)
        result['答复接入来源']={'原合集':mapping.get('mapping_path'),'原答复指纹':mapping['答复指纹'],
            '本次原答复':original_reply,'字段范围':deepcopy(field_scope),'撤下旧字段':deepcopy(withdrawn or {}),'范围外未采用':outside}
        if moved_positions is not None:
            result['答复接入来源'].update(原答复源文件指纹=reply_source_hash,本格位置接续记录=deepcopy(reply_position_chain),本格原位置至现位置=moved_positions)
        if mapping.get('b_source_ranges'):result['答复接入来源']['来源范围']=deepcopy(mapping['b_source_ranges'])
        if not mapping.get('mapping_path') or not Path(mapping['mapping_path']).is_file():
            result['答复接入来源']['原识别答复']=mapping['识别答复']
        text=json.dumps(combined,ensure_ascii=False,separators=(',',':'))
    else:text=json.dumps(reply,ensure_ascii=False,separators=(',',':'))
    result.update(识别答复=text,答复指纹=definition_hash(text))
    if not_adopted_reason:
        result['不采用答复']={'答复指纹':result['答复指纹'],'原因':not_adopted_reason}
    result=direct_collection(result,snapshot if snapshot is not None else mapping['机械记录'],standard)
    if mapping.get('collection_version')==2:
        if mapping.get('source_hash')==result.get('source_hash'):
            originals={row['source_reference']['location']:row for row in mapping.get('facts',[])}
            for row in result.get('facts',[]):
                original=originals.get(row['source_reference']['location'])
                if original is not None:
                    for key in ('raw_value','raw_value_state'):
                        if key in original:row[key]=deepcopy(original[key])
        # 局部接入不重排已有来源，避免同义候选的机械取值顺序随保存而改变。
        order={location:index for index,(location,_) in enumerate(_records(mapping))}
        result['facts'].sort(key=lambda row:order.get(row['source_reference']['location'],len(order)))
        result['slot_bindings'].sort(key=lambda row:order.get(row['sheet']+'!'+row['cell'],len(order)))
    if target_mapping is not None:
        from .A语义阅读 import direct_reference
        result['A需要指纹']=definition_hash(direct_reference(target_mapping,compact=True))
    evidence=result.get('calculation_evidence')
    from .统一勾稽 import _calculation_formulas
    formulas=_calculation_formulas(result)
    if evidence and formulas!={loc:row['formula'] for loc,row in result.get('formula_values',{}).items()}:
        from .表格 import file_hash
        from .统一勾稽 import _calculated_values
        _require(file_hash(evidence['path'])==evidence['sha256'],'公式计算依据已变化')
        receipt=json.loads(Path(evidence['path']).read_text(encoding='utf-8-sig'),object_pairs_hook=json_dates)
        result['formula_values']=_calculated_values(result,receipt)
    result['封存校验']=direct_digest(result)
    path=Path(destination);path.parent.mkdir(parents=True,exist_ok=True);result['mapping_path']=str(path.resolve())
    with path.open('x',encoding='utf-8') as file:json.dump(result,file,ensure_ascii=False,indent=2,default=json_date)
    return load_fact_mapping(path,{})


def _period_rows(row, table):
    """只连接答复已有的坐标、角色与日期；不读取原表推断角色。"""
    import re
    from .A语义阅读 import addresses
    value=row['value']
    if not isinstance(value,str) or not re.search(r'[:：]',value):return [row],[]
    pairs={}
    def cells(span):
        sheet,sep,area=span.rpartition('!')
        _require(not sep or sheet.strip("'")==table['sheet'],'答复范围不属于本方工作表')
        return addresses((area if sep else span).replace('$',''))
    for part in re.split(r'[;；]',value):
        match=re.fullmatch(r'\s*([^:：]+)[:：]\s*(\d{4}-\d{2}-\d{2}(?:[至/]\d{4}-\d{2}-\d{2})?)\s*',part)
        if not match:return [row],[]
        label=match[1].strip()
        if label in pairs and pairs[label]!=match[2]:return [row],[]
        pairs[label]=match[2]
    roles={}; problems=[]; selected=set()
    for span in row['ranges']:
        try:selected.update(cells(span))
        except ValueError as error:
            problems.append({'id':row['id'],'ranges':[span],'reason':str(error),'original':deepcopy(row)})
    if all(re.fullmatch(r'[A-Z]+\d+(?::[A-Z]+\d+)?(?:[、,，][A-Z]+\d+(?::[A-Z]+\d+)?)*',label) for label in pairs):
        declared={};grouped={}
        for label,date in pairs.items():
            for span in re.split(r'[、,，]',label):
                for address in cells(span):declared.setdefault(address,set()).add(date)
        for address in selected:
            values=declared.get(address,set())
            if len(values)==1:grouped.setdefault(next(iter(values)),[]).append(address)
            else:problems.append({'id':row['id'],'ranges':[address],
                'reason':'答复日期未唯一对应此坐标','original':deepcopy(row)})
        return [{**deepcopy(row),'value':value,'ranges':ranges} for value,ranges in grouped.items()],problems
    for metric in table.get('metrics',[]):
        declared=metric.get('template_period_role')
        if not isinstance(declared,dict):continue
        covered={a for span in metric['ranges'] for a in cells(span)}
        for span,label in declared.items():
            if not isinstance(label,str):continue
            label=re.sub(r'(余额|发生额)$','',label.strip())
            if label not in pairs:continue
            for address in cells(span):
                if address in selected & covered:roles.setdefault(address,set()).add(pairs[label])
    grouped={}; missing=[]
    for address in sorted(selected):
        choices=roles.get(address,set())
        if len(choices)==1:grouped.setdefault(next(iter(choices)),[]).append(address)
        else:missing.append(address)
    if missing:problems.append({'id':row['id'],'ranges':missing,
        'reason':'答复未将期间角色唯一对应到这些坐标','original':deepcopy(row)})
    return [{**deepcopy(row),'value':value,'ranges':ranges} for value,ranges in grouped.items()],problems


def _direct_table(table, mapping, standard):
    """只统一已有答复的写法；不决定成员、条件、口径或完整程度。"""
    from .A语义阅读 import normalize_evidence
    from .统一证据 import context_catalog
    import re
    table=deepcopy(table)
    role_aliases={'期初':'opening','期末':'closing','本期':'current','上期':'prior',
                  '期初余额':'opening','期末余额':'closing','本期发生额':'current','上期发生额':'prior',
                  '本期增加':'current','本期减少':'current','本期增加额':'current','本期减少额':'current','上期增加':'prior','上期减少':'prior',
                  '本期期初余额':'opening','本期期末余额':'closing','期初所占比例':'opening','期末所占比例':'closing',
                  'period_end':'closing','period_start':'opening','current_period':'current','prior_period':'prior'}
    def period_role(value):
        if isinstance(value,dict):return {k:period_role(v) for k,v in value.items()}
        statement=re.fullmatch(r'(本期|上期)变动事项[（(]\d{4}[^（）()]+[）)](?:；[^；]*)?',value) if isinstance(value,str) else None
        if statement:return {'本期':'current','上期':'prior'}[statement[1]]
        return role_aliases.get(value,value) if isinstance(value,str) else value
    aliases={f'@C{i}':ref for i,ref in enumerate(context_catalog(mapping['source_context']))}
    def local(area):
        sheet,sep,cell=area.rpartition('!')
        if sep:
            _require(sheet.strip("'")==table['sheet'],'答复范围不属于本方工作表')
        return (cell if sep else area).replace('$','')
    for row in table.get('unresolved',[]):
        # 旧解析说明中的跨表引用保留为引用，不能再次当作本表范围导致保存退出。
        foreign=[area for area in row.get('ranges',[]) if '!' in area and area.rsplit('!',1)[0].strip("'")!=table['sheet']]
        if foreign:
            row['referenced_ranges']=list(dict.fromkeys([*row.get('referenced_ranges',[]),*foreign]))
            row['ranges']=[area for area in row['ranges'] if area not in foreign]
    for row in table.get('metrics',[])+table.get('dimensions',[])+table.get('unresolved',[]):
        if 'ranges' in row:row['ranges']=[local(area) for area in row['ranges']]
        if 'range' in row:row['range']=local(row['range'])
        role=row.get('template_period_role')
        if isinstance(role,list) and len(role)==1 and isinstance(role[0],str):
            row['template_period_role']=role[0]
        if 'template_period_role' in row:row['template_period_role']=period_role(row['template_period_role'])
    roles=[row for row in table.get('dimensions',[]) if row['id']=='template_period_role']
    for row in roles:row['value']=period_role(row['value'])
    if roles:
        from .A语义阅读 import addresses
        metrics=[]
        for row in table['metrics']:
            for area in row['ranges']:
                for cell in (a for s in table['data_ranges'] for a in addresses(s)) if area=='*' else addresses(area):
                    values=[r['value'] for r in roles if _inside(cell,r['ranges'],table)]
                    copied={**deepcopy(row),'ranges':[cell]}
                    if values and all(v==values[0] for v in values):copied['template_period_role']=values[0]
                    elif values:table.setdefault('unresolved',[]).append({'id':'template_period_role','ranges':[cell],'reason':'原答复期间角色冲突','original':values})
                    metrics.append(copied)
        table['metrics']=metrics
    # 编号尚缺时，已明示的角色仍留在原表区声明中供后续复用。
    dimensions=deepcopy(roles)
    for row in table.get('dimensions',[]):
        if row['id']=='template_period_role' or row['value'] is None:continue
        definition=standard['dimension'][row['id']]
        value=row['value']
        if isinstance(value,str):
            literal=re.fullmatch(r'`([^`]+)`',value.strip())
            if literal:value=literal[1]
        if definition['value_type']=='period' and isinstance(value,str):
            value=re.sub(r'^(?:实际日期|时点|期间)\s*(?=\d{4})','',value.strip())
            value=re.sub(r'\s*[（(](?:instant|duration|时点|期间)[）)]$','',value.strip())
            value=re.sub(r'(?<=\d)(?:时点|期间)$','',value)
            value=re.sub(r'\s*至\s*','至',value)
            value=re.sub(r'(\d{4})[年/.](\d{1,2})[月/.](\d{1,2})日?',
                         lambda m:f'{m[1]}-{int(m[2]):02d}-{int(m[3]):02d}',value)
        if definition['value_type']=='text' and isinstance(value,str) and (definition.get('allowed_values') or row['id']=='report_scope'):
            if re.search(r'(?:^|[，,；;])\s*保留(?:为)?未知$',value):
                row.setdefault('original',deepcopy(row));value='未知'
            if value.startswith('「') and value.endswith('」') and value[1:-1] in definition.get('allowed_values',[]):
                value=value[1:-1]
            value=re.sub(r'`([^`]+)`',r'\1',value)
            if row['id']=='report_scope' and re.fullmatch(r'(单体|个别|单户|合并|单体报表|合并报表)[（(][^（）()]+[）)]',value):
                row.setdefault('original',deepcopy(row));value=re.split(r'[（(]',value,maxsplit=1)[0]
            if row['id']=='report_scope' and value in ('单体','个别','单户','合并','单体报表','合并报表'):
                value={'单体':'standalone','个别':'standalone','单户':'standalone','合并':'consolidated','单体报表':'standalone','合并报表':'consolidated'}[value]
            explicit=re.fullmatch(r'[^（）()]+[（(]([^（）()]+)[）)]',value)
            if explicit and explicit[1] in definition.get('allowed_values',('standalone','consolidated','parent')):
                row.setdefault('original',deepcopy(row));value=explicit[1]
            first=re.split(r'[；;：:。（()）]',value,maxsplit=1)[0].strip()
            if first in definition.get('allowed_values',('standalone','consolidated','parent')):
                row.setdefault('original',deepcopy(row));value=first
        if row['id']=='currency' and isinstance(value,str):
            value=value.replace('`','')
            code=re.match(r'^(?:为\s*)?([A-Z]{3})(?=$|[；;。、，,）)])',value)
            if code:
                row.setdefault('original',deepcopy(row));value=code[1]
            code=re.fullmatch(r'[^，,]+[，,]\s*([A-Z]{3})',value)
            if code:value=code[1]
            code=re.fullmatch(r'([A-Z]{3})[、，,]\s*(?:倍率|金额倍率|unit_scale)\s*(?:为|[:：=])\s*[+-]?\d+(?:\.\d+)?',value)
            if code:
                row.setdefault('original',deepcopy(row));value=code[1]
            code=re.fullmatch(r'[^（）()]+[（(]([A-Z]{3})[）)]',value)
            if code:
                row.setdefault('original',deepcopy(row));value=code[1]
            code=re.fullmatch(r'([A-Z]{3})[（(][^（）()]+[）)]',value)
            if code:
                row.setdefault('original',deepcopy(row));value=code[1]
        if definition['value_type']=='period' and isinstance(value,str) and re.fullmatch(r'\d{4}-\d{2}-\d{2}',value):
            value={'kind':'instant','date':value}
        elif definition['value_type']=='period' and isinstance(value,str) and re.fullmatch(r'\d{4}-\d{2}-\d{2}[至/]\d{4}-\d{2}-\d{2}',value):
            start,end=re.split(r'[至/]',value);value={'kind':'duration','start':start,'end':end}
        elif definition['value_type']=='period' and isinstance(value,dict) and set(value)=={'start','end'}:
            value={'kind':'duration',**value}
        elif definition['value_type']=='selection':
            shared=re.fullmatch(r'(?:两列|各列)均为\s*`?(all|unknown)`?(?:[:：].*)?',value) if isinstance(value,str) else None
            if shared:value=shared[1]
            if isinstance(value,str) and re.fullmatch(re.escape(definition['domain'])+
                r'(?:全体|总体)[（(]`?all`?[，,]\s*completeness\s*=\s*complete[）)](?:。.*)?',value):
                value={'mode':'all','completeness':'complete'}
            stated_predicate=re.match(r'^predicate\s*(?:[=:：＝]|—+)\s*(.+)$',value) if isinstance(value,str) else None
            if stated_predicate:value={'mode':'predicate','predicate':stated_predicate[1]}
            predicate=re.match(r'predicate\s*[「“"]([^」”"]+)[」”"](?=$|[，,；;。])',value) if isinstance(value,str) else None
            if predicate:value={'mode':'predicate','predicate':predicate[1]}
            if isinstance(value,str) and re.search(r'(?:^|保留为\s*`?)partial[:：]',value):
                value={'mode':'unknown','known_scope':value}
            if isinstance(value,str) and re.fullmatch(re.escape(definition['domain'])+r'(?:全体|总体)[（(]?`?all`?[）)]?',value):
                value='all'
            if isinstance(value,str) and value.startswith(definition['domain']):
                rest=value[len(definition['domain']):].lstrip('；;，, ')
                if rest=='all' or rest.startswith(('members:','members：')):value=rest
            quoted_members=re.fullmatch(r'members\s*[:：]?\s*(?:「([^」]+)」|“([^”]+)”)(?:[（(]原文合列(?:[，,]保留合并范围)?[）)])?',value) if isinstance(value,str) else None
            if quoted_members:value={'mode':'members','members':[next(v for v in quoted_members.groups() if v is not None)]}
            if value=='未知':value={'mode':'unknown'}
            elif isinstance(value,str) and re.match(r'^all\s*[:：]',value):
                value={'mode':'all','condition':re.split(r'[:：]',value,maxsplit=1)[1].strip()}
            if value in ('all','unknown'):value={'mode':value}
            if isinstance(value,str) and re.fullmatch(r'members\s*[=:：]?\s*\[.*\]',value):
                names=value[value.index('[')+1:-1]
                value={'mode':'members','members':[x.strip().strip('"\'') for x in re.split(r'[,，、]',names)]}
            if isinstance(value,str) and re.fullmatch(r'members[:：]\s*[^、,，;；]+',value):
                value={'mode':'members','members':[re.split(r'[:：]',value,maxsplit=1)[1].strip()]}
            if isinstance(value,dict):
                if 'uncertainty' in value:
                    _require('unknown' not in value or value['unknown']==value['uncertainty'],'未知说明写法冲突')
                    value['unknown']=value.pop('uncertainty')
                for key in ('type','kind','op'):
                    if key in value:
                        _require('mode' not in value or value['mode']==value[key],'集合写法冲突')
                        value['mode']=value.pop(key)
                if 'details' in value:
                    row['original']=deepcopy(row);value.pop('details')
                if value.get('mode')=='partial':
                    row.setdefault('original',deepcopy(row));value['mode']='unknown'
                if value.get('mode')=='other' and set(value)=={'mode','parent','excludes'}:
                    value={'mode':'complement','base':value['parent'],'excluded':value['excludes']}
                for mode,alias,field in (('members','values','members'),('members','include','members'),
                                         ('predicate','value','predicate'),('predicate','include','predicate'),
                                         ('predicate','condition','predicate')):
                    if value.get('mode')==mode and alias in value:
                        _require(field not in value or value[field]==value[alias],'集合写法冲突')
                        value[field]=value.pop(alias)
                if value.get('mode')=='predicate' and 'conditions' in value:
                    conditions=value['conditions']
                    _require(isinstance(conditions,list) and len(conditions)==1,'多个条件未声明组合关系，原答复保留')
                    _require('predicate' not in value or value['predicate']==conditions[0],'条件写法冲突')
                    value['predicate']=value.pop('conditions')[0]
                if value.get('mode')=='predicate' and isinstance(value.get('predicate'),dict):
                    condition=value['predicate']
                    if set(condition)=={'field','value'} and all(isinstance(x,str) for x in condition.values()):
                        value['predicate']=condition['field']+'='+condition['value']
                if value.get('mode')=='predicate' and 'where' in value:
                    _require('predicate' not in value or value['predicate']==value['where'],'条件写法冲突')
                    value['predicate']=value.pop('where')
                if value.get('mode')=='predicate' and isinstance(value.get('all'),list) and value['all']:
                    _require('predicate' not in value,'条件写法冲突')
                    row.setdefault('original',deepcopy(row))
                    # 原答复已明示all关系，只序列化原条件，不替conditions决定且或。
                    value['predicate']=json.dumps({'all':value.pop('all')},ensure_ascii=False,sort_keys=True,separators=(',',':'))
                value.setdefault('domain',definition['domain'])
        row['value']=value
        refs=row.get('evidence',[])
        refs=[refs] if isinstance(refs,str) else refs
        refs=[part.strip() for ref in refs for part in re.split(r'[;；]|(?<=[0-9])、(?=(?:[^、;；!]+!)?\$?[A-Z]+\$?[0-9])',ref) if part.strip()]
        normalized=[]
        for ref in refs:
            ref=ref.removeprefix('原文.json.context.')
            ref=ref.replace('本方工作表!',table['sheet']+'!').removeprefix('本方工作表 ')
            ref=re.sub(r'(?<![@\w])C(\d+)(?=第)',r'@C\1',ref)
            ref=re.sub(r'^(共同说明(?:引用)?[：:]?\s*)C(\d+)(?=第)',r'\1@C\2',ref)
            if len(aliases)==1:
                paragraph=re.fullmatch(r'共同说明[：:]\s*(?:原文)?(第[\d、,，—段第-]+段).*',ref)
                if paragraph:ref=next(iter(aliases))+paragraph[1]
            alias=re.fullmatch(r'(?:共同说明(?:引用)?[：:]?\s*)?(?:原文\.json\s*)?(?:context\s*)?(@C\d+)\s*(第[\d、,，—段第-]+段)?(?:[（(].*[）)])?',ref)
            if alias:
                if alias[1] not in aliases:
                    row.setdefault('original',deepcopy(row));normalized.append(ref);continue
                normalized.append(aliases[alias[1]]+('#'+alias[2] if alias[2] else ''))
            else:
                if not re.match(r'^(?:[^!]+!)?\$?[A-Z]+\$?\d',ref) and not ref.startswith('@'):
                    row.setdefault('original',deepcopy(row))
                    normalized.append(ref)
                    continue
                try:
                    ref_sheet=ref.rsplit('!',1)[0].strip("'") if '!' in ref else table['sheet']
                    normalized.extend(item if item.startswith('@') or '!' in item else ref_sheet+'!'+item
                        for item in normalize_evidence([ref],ref_sheet,aliases))
                except ValueError:
                    row.setdefault('original',deepcopy(row));normalized.append(ref)
        row['evidence']=list(dict.fromkeys(normalized))
        if 'label_references' in row:
            refs=row['label_references'];refs=[refs] if isinstance(refs,str) else refs
            normalized=[]
            for ref in refs:
                try:items=normalize_evidence([ref],table['sheet'],aliases)
                except ValueError:items=[ref]
                normalized.extend(item if '!' in item else table['sheet']+'!'+item for item in items)
            row['label_references']=list(dict.fromkeys(normalized))
        dimensions.append(row)
    table['dimensions']=dimensions
    return table


def _answer_tables(answer, mapping, snapshot):
    from .A语义阅读 import addresses
    if mapping.get('carrier')=='B':
        answer=deepcopy(answer)
        for block in answer.get('tables',[answer]):
            # 共同声明属于答复明列的身份，候选材料及允许修改范围不是已采用来源。
            stated=list(dict.fromkeys(
                (block['sheet']+'!'+area if block.get('sheet') and '!' not in area else area)
                for kind in ('metrics','dimension_bindings','dimensions') for row in block.get(kind,[])
                for area in row.get('ranges',[]) if area!='*'))
            for kind in ('metrics','dimension_bindings','dimensions'):
                for row in block.get(kind,[]):
                    if '*' not in row.get('ranges',[]):continue
                    selected=stated
                    if not selected and kind=='dimensions':
                        selected=[area for area,keys in mapping.get('field_scope',{}).items() if row['id'] in keys]
                    if selected:
                        row['ranges']=[part for area in row['ranges'] for part in (selected if area=='*' else [area])]
    def with_labels(tables):
        for table in tables:
            for row in answer.get('non_business',[]):
                sheet,sep,span=row['range'].rpartition('!')
                if sep and sheet.strip("'")!=table['sheet']:continue
                span=span or row['range']
                if set(addresses(span))<=set(addresses(table['range'])):
                    table.setdefault('non_business',[]).append({**deepcopy(row),'range':span})
        return tables
    if mapping.get('carrier')=='B' and mapping.get('target_tables') and all(not t['data_ranges'] for t in mapping['target_tables']):
        return _answer_tables(answer,{**mapping,'target_tables':[]},snapshot)
    if mapping.get('target_tables'):
        tables=[]
        for declared in mapping['target_tables']:
            table=deepcopy(declared)
            owned={cell for area in table['data_ranges'] for cell in addresses(area)}
            for kind in ('metrics','dimension_bindings','dimensions','unresolved'):
                table[kind]=[]
                for block in answer.get('tables',[answer]):
                    if block.get('sheet') and block['sheet']!=table['sheet']:continue
                    for row in block.get(kind,[]):
                        spans=[]
                        for area in row.get('ranges',[]):
                            sheet,sep,span=area.rpartition('!')
                            if sep and sheet.strip("'")!=table['sheet']:continue
                            parts=block.get('data_ranges',table['data_ranges']) if area=='*' else [span or area]
                            for part in parts:
                                cells=list(addresses(part));kept=[cell for cell in cells if cell in owned]
                                spans.extend([part] if cells==kept else kept)
                        if spans:table[kind].append({**deepcopy(row),'ranges':spans})
            tables.append(table)
        return with_labels(tables)
    if mapping.get('target_table') is not None:
        table=deepcopy(mapping['target_table'])
        for kind in ('metrics','dimension_bindings','dimensions','unresolved'):
            table[kind]=[]
            for block in answer.get('tables',[answer]):
                for original in block.get(kind,[]):
                    row=deepcopy(original)
                    if '*' in row.get('ranges',[]) and block.get('data_ranges'):
                        row['ranges']=[part for span in row['ranges'] for part in
                                       (block['data_ranges'] if span=='*' else [span])]
                    table[kind].append(row)
        return with_labels([table])
    tables=answer.get('tables')
    if tables is None:
        # 未提供固定表区时，只收集答复明示的坐标，不扫描原表补认财务格。
        tables=[]
        for source in snapshot['sheets']:
            name=source['name'];local={k:[] for k in ('metrics','dimension_bindings','dimensions','unresolved')};areas=set()
            for kind in local:
                for row in answer.get(kind,[]):
                    scope=[]
                    for area in row.get('ranges',[]):
                        sh,sep,cell=area.rpartition('!')
                        if sep and sh.strip("'")==name:scope.append(cell)
                        elif not sep and (area=='*' or len(snapshot['sheets'])==1):scope.append(area)
                    if scope:local[kind].append({**row,'ranges':scope});areas.update(a for a in scope if a!='*')
            # 未解析说明中的引用不扩大共同槽位的适用范围。
            identity_areas=list(dict.fromkeys(area for kind in ('metrics','dimension_bindings','dimensions')
                for row in local[kind] for area in row['ranges'] if area!='*'))
            for row in local['metrics']:
                if '*' in row['ranges']:
                    row['ranges']=[area for span in row['ranges'] for area in (identity_areas if span=='*' else [span])]
            if areas:
                from openpyxl.utils.cell import get_column_letter
                table_id=name+'_'+mapping['答复指纹'][:12] if mapping['carrier']=='B' else name
                tables.append({'id':table_id,'sheet':name,'range':f'A1:{get_column_letter(source["max_column"])}{source["max_row"]}',
                    'data_ranges':sorted({a for area in areas for a in addresses(area)}),**local})
    _require(isinstance(tables,list),'缺少本方表区及身份记录')
    return with_labels(tables)


def _own_answer_text(text, mapping, snapshot):
    """B答复中明确写出的A位置仅是对应对象，不当作B身份的坐标。"""
    if mapping.get('carrier')!='B':return text
    names={s['name'] for s in snapshot['sheets']}
    text=_expand_declared_row_columns(_source_aliases(text,mapping))
    other={'A'}|{r['location'].rsplit('!',1)[0] for r in mapping.get('A需要',{}).get('需要',[])}
    other-=names
    if not other:return text
    coord=r'\$?[A-Z]+\$?[1-9]\d*(?::\$?[A-Z]+\$?[1-9]\d*)?'
    foreign_sheet=r"'?(?:"+'|'.join(re.escape(s) for s in sorted(other,key=len,reverse=True))+r")'?!"
    foreign=re.compile(r'(?:'+foreign_sheet+coord+r'|(?<![A-Za-z0-9_])A(?:\s+|需要中的\s*|需要的\s*|的\s*)(?:'+foreign_sheet+r')?'+coord+r')'
        +r'(?:[、，,/／]\s*(?:'+foreign_sheet+r')?'+coord+r')*')
    own=re.compile(r"(?:'?(?:"+'|'.join(re.escape(s) for s in sorted(names,key=len,reverse=True))+r")'?!)"+coord+r'|(?<![\w])!'+coord)
    lines=[];source_sheet=None;headers=None
    for line in text.replace('`','').splitlines():
        if line.lstrip().startswith('|') and own.search(line):
            line=re.sub(r'(?:→|->)\s*A(?=\s)[^|]*','',line)
        declared=[name for name in names if re.search(r'以下\s*B\s*(?:坐标|位置)(?:均)?(?:属于|在)\s*'+re.escape(name),line)]
        if len(declared)==1:
            source_sheet=declared[0];lines.append(source_sheet+'：')
        if line.lstrip().startswith('|'):
            values=[v.strip() for v in line.strip().strip('|').split('|')]
            sources=[i for i,v in enumerate(values) if 'B' in v and any(word in v for word in ('来源','位置','坐标')) and not re.search(coord,v)]
            targets=[i for i,v in enumerate(values) if 'A' in v and any(word in v for word in ('需要','位置','目标','坐标')) and not re.search(coord,v)]
            if len(sources)==len(targets)==1:
                stated=[name for name in names if re.search(r'[（(]'+re.escape(name)+r'[）)]',values[sources[0]])]
                if len(stated)==1:source_sheet=stated[0]
                headers=(sources[0],targets[0],len(values));values[sources[0]]='位置';values[targets[0]]='A对应'
                line='| '+' | '.join(values)+' |'
            elif headers and len(values)==headers[2] and not all(re.fullmatch(r'[-: ]+',v) for v in values):
                source_index,target_index,_=headers;values[target_index]='[A位置]'
                value=re.sub(r'(?:→|->)\s*(?:A\s*)?(?:'+foreign_sheet+r')?'+coord,r'[A位置]',values[source_index])
                if source_sheet:
                    qualified=r"(?:'?(?:"+'|'.join(re.escape(name) for name in sorted(names,key=len,reverse=True))+r")'?!)"+coord
                    value=re.sub(qualified+r'|(?<![A-Za-z0-9_!])'+coord,
                        lambda m:m[0] if '!' in m[0] else source_sheet+'!'+m[0],value)
                values[source_index]=value;line='| '+' | '.join(values)+' |'
        parts=[]
        for part in re.split(r'(?<=[。；;])',line):
            if re.match(r'^\s*(?:A\s*(?:的|需求|需要|列(?:期间|角色)|该格)|按\s*A\s*原文)',part) and not own.search(part):continue
            if foreign.search(part):
                if not own.search(part):continue
                part=foreign.sub('[A位置]',part)
            parts.append(part)
        lines.append(''.join(parts))
    return '\n'.join(lines)


def direct_collection(mapping, snapshot, standard):
    """单次答复直接落入既有事实字段；只展开明确范围，不重建或补认身份。"""
    from .A语义阅读 import addresses, expand_scopes
    from .统一语义 import identity_completeness
    result=deepcopy({k:v for k,v in mapping.items() if k!='机械记录'})
    result.update(collection_version=2, facts=[], owners={}, slot_bindings=[], slot_binding_version=1,
                  unresolved=[], excluded=[], source_identities=[])
    result['机械记录']=deepcopy(snapshot)
    result['封存定义原文']=mapping['金标准原文'].splitlines()
    try:
        rejected=mapping.get('不采用答复',{})
        if rejected.get('答复指纹')==mapping['答复指纹']:
            raise ValueError(rejected['原因'])
        declared={t['sheet'] for t in mapping.get('target_tables',[])}
        if mapping.get('target_table'):declared.add(mapping['target_table']['sheet'])
        text=mapping['识别答复'] if mapping.get('逐格原样保存') else _own_answer_text(mapping['识别答复'],mapping,snapshot)
        answer=direct_answer(text,standard,
            [s['name'] for s in snapshot['sheets']],default_sheet=next(iter(declared)) if len(declared)==1 else None)
        tables=_answer_tables(answer,mapping,snapshot)
    except (ValueError,KeyError,TypeError) as error:
        tables=[{**deepcopy(mapping['target_table']), 'metrics':[], 'dimensions':[]}] if mapping.get('target_table') else []
        label='原答复不采用：' if rejected.get('答复指纹')==mapping['答复指纹'] else '原答复未能读取：'
        result['unresolved'].append({'reason':label+str(error)})
    normalized=[]
    for table in tables:
        if mapping.get('逐格原样保存'):
            copied=deepcopy(table)
            for kind in ('metrics','dimension_bindings','dimensions','unresolved'):
                for row in copied.get(kind,[]):
                    local=[]
                    for span in row.get('ranges',[]):
                        sheet,sep,cell=span.rpartition('!')
                        _require(not sep or sheet.strip("'")==table['sheet'],'逐格身份引用了其他工作表')
                        local.append(cell if sep else span)
                    if 'ranges' in row:row['ranges']=local
            for row in copied.get('dimensions',[]):row.setdefault('evidence',[])
            normalized.append(copied)
            continue
        copied=_direct_table({**table,'dimensions':[r for r in table.get('dimensions',[]) if r['id']=='template_period_role']},mapping,standard)
        for row in table.get('dimensions',[]):
            if row['id']=='template_period_role':continue
            try:
                candidates,problems=_period_rows(row,table) if row.get('id')=='period' else ([row],[])
                copied.setdefault('unresolved',[]).extend(problems)
                rows=_direct_table({**table,'metrics':[],'unresolved':[],'dimensions':candidates},mapping,standard)['dimensions']
                copied['dimensions'].extend(rows)
            except (ValueError,KeyError,TypeError) as error:
                copied.setdefault('unresolved',[]).append({'id':row.get('id'),'ranges':row.get('ranges',[]),
                    'reason':str(error),'original':deepcopy(row)})
        # 同一原句的明确字段已接入后，只撤掉对应位置的旧机械解析标记。
        declared={}
        for row in copied.get('metrics',[])+copied['dimensions']:
            original=row.get('original');value=row.get('value')
            while isinstance(original,dict):original=original.get('original')
            if not isinstance(original,str) or row.get('id') not in standard['metric'] and row.get('id') not in standard['dimension']:continue
            if 'value' in row and (value in (None,'未知','unknown','未决') or isinstance(value,dict) and value.get('mode')=='unknown'):continue
            declared.setdefault(original,[]).extend(row.get('ranges',[]))
        pending=[]
        for row in copied.get('unresolved',[]):
            original=row.get('original')
            while isinstance(original,dict):original=original.get('original')
            if row.get('reason')=='这段答复未明确可机械读取的字段归属' and not any(row.get(k) for k in ('id','field','dimension','metric')) and isinstance(original,str) and original in declared:
                positions={a for span in row.get('ranges',[row['range']] if 'range' in row else [])
                           for a in ({a for area in copied['data_ranges'] for a in addresses(area)} if span=='*' else addresses(span))}
                remaining=positions-{a for a in positions if _inside(a,declared[original],copied)}
                if not remaining:continue
                row={**row,'ranges':sorted(remaining)};row.pop('range',None)
            pending.append(row)
        copied['unresolved']=pending
        normalized.append(copied)
    tables=normalized
    result['outline']={'subtables':tables}
    cells={s['name']:s['cells'] for s in snapshot['sheets']}
    merged={s['name']:{a for area in s['merges'] for a in addresses(area)} for s in snapshot['sheets']}
    provided=None
    material=mapping.get('原文材料',{})
    if mapping.get('carrier')=='B' and material.get('sheets'):
        provided=set()
        for source in material['sheets']+material.get('公式及来源',{}).get('来源行列原文',[]):
            for line in source.get('rows','').splitlines():
                match=re.match(r'^(\d+)(?:-(\d+))?:\s*(.*)',line)
                if not match:continue
                columns=re.findall(r'(?:^| \| )([A-Z]+)=',match[3])
                provided.update(source['sheet']+'!'+column+str(row)
                    for row in range(int(match[1]),int(match[2] or match[1])+1) for column in columns)
    occupied=set()
    strict={**standard,'_complete_identity':True}
    for table in tables:
        sheet=table['sheet']
        _require(sheet in cells,'答复工作表不存在')
        data={a for area in table['data_ranges'] for a in addresses(area)}
        pending=[]
        for row in table.get('unresolved',[]):
            if mapping.get('逐格字段版本')==1 and row.get('逐格原字段') is not True:continue
            key=next((row[k] for k in ('id','field','dimension') if row.get(k) in standard['dimension']),None)
            for area in row.get('ranges', [row['range']] if 'range' in row else []):
                if key:
                    try:affected=addresses(area)
                    except ValueError:continue
                    # 缺失项仍在完整性中；同字段既声明取值又未决时，两份声明均保留。
                    if not any(dim['id']==key and not (isinstance(dim['value'],dict) and dim['value'].get('mode')=='unknown')
                               and any(_inside(a,dim['ranges'],table) for a in affected) for dim in table['dimensions']):continue
                pending.append({'range':area,'reason':row.get('reason','原答复未确认'),'original':deepcopy(row)})
        annotations={}
        for row in table.get('non_business',[]):
            _require(row.get('category') in {'label','annotation'},'说明位置类别未能读取')
            for address in addresses(row['range']):
                _require(address in addresses(table['range']),'说明位置越出本方表区')
                if row['category']=='label' and mapping.get('逐格字段版本')==1 and any(
                    item.get('逐格原字段') is True and address in item['ranges']
                    for kind in ('metrics','dimension_bindings') for item in table.get(kind,[])):continue
                annotations[address]=row
        indexed={kind:{} for kind in ('metrics','dimension_bindings','dimensions','unresolved')}
        for kind in indexed:
            for row in pending if kind=='unresolved' else table.get(kind,[]):
                if kind!='unresolved' and mapping.get('逐格字段版本')==1 and row.get('逐格原字段') is not True:continue
                ranges=[row['range']] if kind=='unresolved' else row['ranges']
                locations=set().union(*(data if area=='*' else set(addresses(area)) for area in ranges))
                for address in locations&(data|set(annotations)):
                    indexed[kind].setdefault(address,[]).append(row)
        expanded={}
        for address in data|set(annotations):
            if provided is not None and sheet+'!'+address not in provided:
                result['unresolved'].append({
                    'reason':'答复坐标不在本次提供的B原文范围内；保留原声明，不取该位置原值或冒充B候选',
                    'original':{'sheet':sheet,'cell':address,
                        **{kind:deepcopy(indexed[kind].get(address,[]))
                           for kind in ('metrics','dimension_bindings','dimensions')}}})
                continue
            if address in annotations:
                occupied.add(sheet+'!'+address)
                raw=cells[sheet].get(address,{}).get('value')
                result['excluded'].append({'sheet':sheet,'cell':address,**deepcopy(annotations[address]),
                    'raw_value':raw,'raw_value_state':'blank' if raw in (None,'') else 'formula' if cells[sheet].get(address,{}).get('formula') else 'value'})
                continue
            # 仅展开本格已有声明；冲突保留原声明，不选择其中一个答案。
            local={**table,'range':address,'data_ranges':[address]}
            answer={kind:[{**row,'ranges':[address]} for row in indexed[kind].get(address,[])
                          if kind!='dimensions' or row['id']!='template_period_role']
                    for kind in ('metrics','dimension_bindings','dimensions')}
            issues=indexed['unresolved'].get(address,[])[:]
            # 不在冲突值之间择一；其他字段仍按原答复保存。
            compatible=[]
            for key in dict.fromkeys(row['id'] for row in answer['dimensions']):
                rows=[row for row in answer['dimensions'] if row['id']==key]
                if any(row['value']!=rows[0]['value'] for row in rows[1:]):
                    issues.append({'id':key,'reason':'同一数据位置的维度取值冲突','original':deepcopy(rows)})
                else:compatible.extend(rows)
            original={**answer,'unresolved':issues};answer['dimensions']=compatible
            try:expanded.update(expand_scopes(answer,local,{address:cells[sheet].get(address,{'value':None})},strict))
            except (ValueError,KeyError,TypeError) as error:
                occupied.add(sheet+'!'+address)
                raw=cells[sheet].get(address,{}).get('value')
                result['unresolved'].append({'sheet':sheet,'cell':address,'reason':str(error),'original':deepcopy(original),
                    'raw_value':raw,'raw_value_state':'blank' if raw in (None,'') else 'formula' if cells[sheet].get(address,{}).get('formula') else 'value'})
            else:
                if issues:
                    raw=cells[sheet].get(address,{}).get('value')
                    result['unresolved'].append({'sheet':sheet,'cell':address,
                        'reason':'；'.join(dict.fromkeys(row['reason'] for row in issues)),
                        'original':deepcopy(issues),'raw_value':raw,
                        'raw_value_state':'blank' if raw in (None,'') else 'formula' if cells[sheet].get(address,{}).get('formula') else 'value'})
        for address,item in expanded.items():
            location=sheet+'!'+address
            _require(location not in occupied,'同一目标重复声明：'+location)
            occupied.add(location)
            if not item.get('metric_id') and not item.get('dimension_id'):
                if any(row.get('sheet')==sheet and row.get('cell')==address for row in result['unresolved']):continue
                raw=cells[sheet].get(address,{}).get('value')
                result['unresolved'].append({'sheet':sheet,'cell':address,'reason':item.get('reason','未提供标准编号'),'original':deepcopy(item),
                    'raw_value':raw,'raw_value_state':'blank' if raw is None or isinstance(raw,str) and not raw.strip()
                    else 'formula' if cells[sheet].get(address,{}).get('formula') else 'value'})
                continue
            declared=({ref['id'] for ref in standard['metric'][item['metric_id']]['dimension_refs']}
                      if item.get('metric_id') else set(standard['dimension']))
            # 共用字段按金标准已有引用装入身份；全部原声明仍保存在表区和原答复中。
            dims={k:v['value'] for k,v in item['dimensions'].items() if k in declared}
            evidence={k:deepcopy(v['evidence']) for k,v in item['dimensions'].items() if k in declared}
            # 原值及空白状态来自原格，绝不要求模型复写金额。
            raw=cells[sheet].get(address,{}).get('value')
            fact={'id':location,'record_type':'metric_fact' if item.get('metric_id') else 'dimension_binding',
                  **{key:item[key] for key in ('metric_id','dimension_id') if item.get(key)},
                  'dimensions':dims,'dimension_evidence':evidence,'owner_reference':location,
                  'source_reference':{'file_hash':mapping['source_hash'],'location':location},
                  'recognition_provenance':{'method':'direct','record_id':mapping['答复指纹']},
                  'raw_value':raw,'raw_value_state':'blank' if raw is None or isinstance(raw,str) and not raw.strip()
                  else ('formula' if cells[sheet][address].get('formula') else 'value')}
            labels={k:deepcopy(v['label_references']) for k,v in item['dimensions'].items() if k in declared and 'label_references' in v}
            if labels:fact['label_references']=labels
            try:
                complete=identity_completeness(fact,strict)
                if not complete['complete'] and fact['record_type']=='metric_fact':
                    known={k:deepcopy(fact[k]) for k in ('metric_id','dimensions','dimension_evidence','label_references') if k in fact}
                    known.update(missing_dimensions=complete['missing_dimensions'],evidence_cells=[location])
                    result['slot_bindings'].append({'sheet':sheet,'cell':address,'partial_semantics':known,
                        'raw_value':raw,'raw_value_state':fact['raw_value_state'],
                        'recognition_provenance':fact['recognition_provenance']})
                else:
                    result['facts'].append(fact);result['owners'][location]=deepcopy(dims)
            except (ValueError,KeyError,TypeError) as error:
                result['unresolved'].append({'sheet':sheet,'cell':address,'reason':str(error),'original':deepcopy(item)})
    if mapping['carrier']=='A' and mapping.get('target_table') is None:
        covered={table['sheet']+'!'+cell for table in tables for cell in addresses(table['range'])}
        cited={ref for row in tables for dim in row.get('dimensions',[]) for ref in dim.get('evidence',[])}
        for sheet,items in cells.items():
            for cell,item in items.items():
                location=sheet+'!'+cell
                if item.get('value') is not None and str(item['value']).strip() and location not in covered|cited:
                    result['unresolved'].append({'sheet':sheet,'cell':cell,'reason':'本方表区及已引用说明未覆盖此原位置'})
    result['coverage']={'financial_cells':len(occupied),'unresolved_cells':len(result['unresolved'])}
    result.pop('complete',None)
    result['完整性']={r['source_reference']['location']:identity_completeness(r,strict) for r in result['facts']}
    result['完整性'].update({r['sheet']+'!'+r['cell']:identity_completeness(r['partial_semantics'],strict) for r in result['slot_bindings']})
    result['封存校验']=direct_digest(result)
    return result


def direct_digest(mapping):
    return definition_hash({k:v for k,v in mapping.items() if k not in {'封存校验','mapping_path','mapping_hash','usage'}})


def validate_direct_collection(mapping, *, for_matching=False):
    _require(mapping.get('collection_version')==2 and mapping.get('封存校验')==direct_digest(mapping),'合集内容发生变化')
    _require(definition_hash(mapping['识别答复'])==mapping['答复指纹'],'原答复发生变化')
    if for_matching:
        _require(mapping.get('逐格字段版本')==1,'此合集仍含旧展开身份；须从已有逐格原答复保存后再匹配，不需重答')
        declared={}
        for table in mapping['outline']['subtables']:
            for kind in ('metrics','dimension_bindings','dimensions'):
                for row in table.get(kind,[]):
                    if row.get('逐格原字段') is not True:continue
                    for cell in row['ranges']:
                        _require(re.fullmatch(r'[A-Z]+[1-9]\d*',cell) is not None,'逐格原字段必须保留单格位置')
                        key={'metrics':'metric_id','dimension_bindings':'dimension_id'}.get(kind,row['id'])
                        value=row['id'] if kind in ('metrics','dimension_bindings') else row.get('value')
                        declared.setdefault((table['sheet']+'!'+cell,key),[]).append(value)
        for location,record in _records(mapping):
            identifiers={key:record[key] for key in ('metric_id','dimension_id') if record.get(key)}
            _require(len(identifiers)==1,'单格须有唯一的本格金标准编号：'+location)
            for key,value in [*identifiers.items(),*record['dimensions'].items()]:
                values=declared.get((location,key),[])
                _require(values and all(definition_hash(v)==definition_hash(value) for v in values),
                         '保存身份没有相同的本格原字段：'+location+' '+key)
    # 只验证保存来源及完整性，不解释原答复或补齐缺失字段。
    return True


def combine_direct_collections(parts):
    """只汇总同一原件的已保存记录；不重新展开答复或补齐身份。"""
    _require(bool(parts),'合集索引为空')
    first=parts[0]
    shared=('carrier','source_hash','gold_hash','schema_version','recognition_method','collection_version')
    _require(all(p.get('recognition_method')=='direct' and
                 all(p.get(k)==first.get(k) for k in shared) for p in parts),'分组合集不是同一原件、载体或标准')
    result={k:deepcopy(first[k]) for k in (*shared,'source_path','gold_path','slot_binding_version') if k in first}
    for k in ('facts','slot_bindings','unresolved','excluded','source_identities','source_context','已确认关系','关系缺口'):
        result[k]=[]
    result.update(owners={},完整性={},outline={'subtables':[]},合集来源=[])
    result['机械记录']=deepcopy(first['机械记录'])
    if all(p.get('逐格字段版本')==1 for p in parts):
        result.update(逐格字段版本=1,逐格原样保存=True)
    occupied=set();table_ids=set();definitions={}
    for part in parts:
        _require({k:v for k,v in part['机械记录'].items() if k!='context'}==
                 {k:v for k,v in first['机械记录'].items() if k!='context'},'分组合集的原位置或原值不一致')
        locations={location for location,_ in _records(part)}
        locations.update(r['sheet']+'!'+r['cell'] for r in part.get('slot_bindings',[])
                         if r.get('sheet') and r.get('cell'))
        locations.update(r['source_reference']['location'] for r in part.get('source_identities',[]))
        _require(not occupied&locations,'分组合集位置重叠，须先明确采用哪份已有结果')
        occupied.update(locations)
        ids={t['id'] for t in part['outline']['subtables']}
        _require(not table_ids&ids,'分组合集的表区编号重复')
        table_ids.update(ids)
        for k in ('facts','slot_bindings','unresolved','excluded','source_identities'):
            result[k].extend(deepcopy(part.get(k,[])))
        for k in ('已确认关系','关系缺口'):
            for row in part.get(k,[]):
                if row not in result[k]:result[k].append(deepcopy(row))
        for k in ('owners','完整性'):
            _require(not result[k].keys()&part.get(k,{}).keys(),'分组合集记录编号重复：'+k)
            result[k].update(deepcopy(part.get(k,{})))
        result['outline']['subtables'].extend(deepcopy(part['outline']['subtables']))
        for context in part.get('source_context',[]):
            if context not in result['source_context']:result['source_context'].append(deepcopy(context))
        for line in part['封存定义原文']:
            key=json.loads(line)['id']
            _require(key not in definitions or definitions[key]==line,'分组合集的标准原文不一致：'+key)
            definitions[key]=line
        result['合集来源'].append({k:deepcopy(part[k]) for k in
            ('mapping_path','mapping_hash','封存校验','识别答复','答复指纹','答复接入来源','A纠正来源','本次用户事实接入') if k in part})
    result['机械记录']['context']=deepcopy(result['source_context'])
    result['封存定义原文']=list(definitions.values())
    result['金标准原文']='\n'.join(result['封存定义原文'])
    result['识别答复']='\n\n'.join(p['识别答复'] for p in parts)
    result['答复指纹']=definition_hash(result['识别答复'])
    result['coverage']={'financial_cells':sum(p['coverage']['financial_cells'] for p in parts),
                        'unresolved_cells':len(result['unresolved'])}
    result['recognition_scope']={'mode':'direct'}
    receipts=[p.get('本次逐格接入',{}) for p in parts]
    result['本次逐格接入']={'complete':all(r.get('complete') is True for r in receipts),
        '请求格数':sum(r.get('请求格数',0) for r in receipts),
        '未接入字段':[item for r in receipts for item in r.get('未接入字段',[])],
        '未确定字段':[item for r in receipts for item in r.get('未确定字段',[])],
        '缺接入记录的分组数':sum(not r for r in receipts)}
    result['封存校验']=direct_digest(result)
    return result


def resume_direct_output(receipt_path, snapshot, context, standard, field_scope=None):
    """按实际执行记录移动已有A身份；新增位置不继承模板或B的身份。"""
    from .表格 import file_hash
    from .统一勾稽 import load_fact_mapping
    from .A语义阅读 import addresses
    from openpyxl.utils.cell import range_boundaries, get_column_letter
    receipt_path=Path(receipt_path)
    receipt=json.loads(receipt_path.read_text(encoding='utf-8-sig'),object_pairs_hook=json_dates)
    _require(receipt.get('schema')=='附注更新执行记录-v1' and
        receipt.get('output_hash')==snapshot['sha256'] and
        Path(receipt.get('output_path','')).resolve()==Path(snapshot['path']).resolve(),
        '执行记录与所选A副本不一致')
    sources={}
    for side,ref in receipt['collections'].items():
        _require(file_hash(ref['mapping_path'])==ref['mapping_hash'],'执行记录引用的原合集已变化')
        saved=load_fact_mapping(ref['mapping_path'],standard)
        _require(saved['carrier']==side and saved['source_hash']==ref['source_hash'] and
            saved['gold_hash']==ref['gold_hash']==standard['sha256'],'执行记录的载体或金标准不一致')
        _require(file_hash(ref['source_path'])==ref['source_hash'],'执行记录引用的原件已变化')
        sources[side]=saved
    plan_path=Path(receipt['plan_path'])
    seal=json.loads(plan_path.with_name('更新清单校验.json').read_text(encoding='utf-8-sig'),object_pairs_hook=json_dates)
    _require(file_hash(plan_path)==seal['sha256'],'执行记录引用的更新清单已变化')
    plan=json.loads(plan_path.read_text(encoding='utf-8-sig'),object_pairs_hook=json_dates)
    _require(plan['source_a_hash']==sources['A']['source_hash'] and plan['source_b_hash']==sources['B']['source_hash']
        and plan['gold_hash']==standard['sha256'] and plan['template_layout']==receipt['template_layout'],
        '执行记录的原件或位置移动与更新清单不一致')
    old=sources['A'];result=deepcopy(old);positions=receipt['template_layout']['positions']
    move=lambda loc:positions.get(loc,loc)
    def area_range(sheet,area):
        bounds=[range_boundaries(move(sheet+'!'+cell).rsplit('!',1)[1]) for cell in addresses(area) if move(sheet+'!'+cell) is not None]
        if not bounds:return None
        return f'{get_column_letter(min(v[0] for v in bounds))}{min(v[1] for v in bounds)}:{get_column_letter(max(v[2] for v in bounds))}{max(v[3] for v in bounds)}'
    def context_ranges(table):
        refs=[]
        for ref in table.get('context_references',[]):
            if isinstance(ref,dict) and ref.get('sheet') and ref.get('range'):
                span=area_range(ref['sheet'],ref['range'])
                if span:refs.append({**ref,'range':span})
            elif isinstance(ref,str) and '!' in ref:
                sheet,span=ref.rsplit('!',1);span=area_range(sheet,span)
                if span:refs.append(sheet+'!'+span)
            else:refs.append(ref)
        table['context_references']=refs
    def references(row):
        for key in ('dimension_evidence','label_references'):
            if key not in row:continue
            if isinstance(row[key],dict):
                row[key]={dim:[move(ref) for ref in refs if move(ref) is not None] for dim,refs in row[key].items()}
            else:row[key]=[move(ref) for ref in row[key] if move(ref) is not None]
        for key in ('evidence','evidence_cells'):
            if key in row:row[key]=[move(ref) for ref in row[key] if move(ref) is not None]
    cells={s['name']+'!'+cell:raw for s in snapshot['sheets'] for cell,raw in s['cells'].items()}
    for key in ('facts','slot_bindings','unresolved','excluded','source_identities'):
        rows=[]
        for item in result.get(key,[]):
            loc=item.get('source_reference',{}).get('location') or (item['sheet']+'!'+item['cell'] if item.get('sheet') and item.get('cell') else None)
            if loc:
                dest=move(loc)
                if dest is None:continue
                if item.get('source_reference'):item['source_reference']={**item['source_reference'],'location':dest,'file_hash':snapshot['sha256']}
                if 'sheet' in item and 'cell' in item:item['sheet'],item['cell']=dest.rsplit('!',1)
                for field in ('id','owner_reference'):
                    if item.get(field)==loc:item[field]=dest
                raw=cells.get(dest,{})
                if 'raw_value' in item:
                    item['raw_value']=deepcopy(raw.get('value'))
                    item['raw_value_state']='blank' if raw.get('value') in (None,'') else 'formula' if raw.get('formula') else 'value'
            references(item)
            if 'partial_semantics' in item:references(item['partial_semantics'])
            rows.append(item)
        result[key]=rows
    for key in ('owners','完整性'):
        result[key]={move(loc):value for loc,value in result.get(key,{}).items() if move(loc) is not None}
    tables=[];owned=set()
    for table in result['outline']['subtables']:
        sheet=table['sheet'];data={cell for span in table['data_ranges'] for cell in addresses(span)}
        def relocated(spans):
            return [move(sheet+'!'+cell).rsplit('!',1)[1] for span in spans
                    for cell in (sorted(data) if span=='*' else addresses(span)) if move(sheet+'!'+cell) is not None]
        table['data_ranges']=list(dict.fromkeys(relocated(table['data_ranges'])))
        owned.update(sheet+'!'+cell for cell in table['data_ranges'])
        table['range']=area_range(sheet,table['range'])
        if table['range'] is None:continue
        context_ranges(table)
        for kind in ('metrics','dimension_bindings','dimensions','unresolved','non_business'):
            rows=[]
            for item in table.get(kind,[]):
                if kind=='non_business':
                    for dest in relocated([item['range']]):rows.append({**item,'range':dest})
                    continue
                if 'range' in item:item['ranges']=[item.pop('range')]
                if 'ranges' in item:
                    item['ranges']=list(dict.fromkeys(relocated(item['ranges'])))
                    if not item['ranges']:continue
                roles=item.get('template_period_role')
                if isinstance(roles,dict):
                    item['template_period_role']={dest:value for span,value in roles.items()
                        for dest in relocated([cell for cell in data if cell.rstrip('0123456789')==span] if span.isalpha() else [span])}
                references(item);rows.append(item)
            table[kind]=rows
        tables.append(table)
    requested={sheet+'!'+cell for span in field_scope or {} for sheet,area in [span.rsplit('!',1)] for cell in addresses(area)}
    for sheet in snapshot['sheets']:
        fresh=sorted(loc.rsplit('!',1)[1] for loc in requested-owned if loc.rsplit('!',1)[0]==sheet['name'])
        if fresh:tables.append({'id':'副本新位置_'+sheet['name'],'sheet':sheet['name'],
            'range':f'A1:{get_column_letter(sheet["max_column"])}{sheet["max_row"]}',
            'data_ranges':fresh,'metrics':[],'dimensions':[],'unresolved':[]})
    result['outline']={'subtables':tables}
    result.pop('target_table',None);result.pop('target_tables',None)
    for group in result.get('正常入口范围认识',{}).get('ranges',[]):
        for table in group.get('tables',[group]):
            sheet=table['sheet'];table['range']=area_range(sheet,table['range'])
            table['data_ranges']=[move(sheet+'!'+cell).rsplit('!',1)[1] for span in table['data_ranges']
                for cell in addresses(span) if move(sheet+'!'+cell) is not None]
            context_ranges(table)
    history=deepcopy(old.get('副本接续',{'执行记录':[],'已执行关系':[],'已执行安排':[],'保留原值':[],'保留标签':[],'名单维度':[]}))
    history['执行记录'].append({'path':str(receipt_path.resolve()),'sha256':file_hash(receipt_path)})
    history['已执行关系']=list(dict.fromkeys(history['已执行关系']+[definition_hash(row) for row in sources['B'].get('已确认关系',[])
        if row.get('kind') in ('名单编排','取数对应') and row.get('source_a_hash')==old['source_hash']]))
    history['已执行安排'].append({'source_b_hash':sources['B']['source_hash'],'purpose':deepcopy(plan['update_purpose'])})
    history['保留原值']=list(dict.fromkeys([move(loc) for loc in history['保留原值'] if move(loc) is not None]+
        [row['sheet']+'!'+row['cell'] for row in plan.get('retained_values',[])]))
    retained_labels=plan.get('retained_label_targets')
    if retained_labels is None:
        retained_labels=[move(row['target']) for relation in sources['B'].get('已确认关系',[]) for row in relation.get('label_actions',[])
            if row.get('action')=='retain' and relation.get('source_a_hash')==old['source_hash'] and move(row['target']) is not None]
    retained_labels=list(retained_labels)+[row['sheet']+'!'+row['cell'] for row in plan.get('labels',[])
        if 'source_sheet' not in row]
    history['保留标签']=list(dict.fromkeys([move(loc) for loc in history['保留标签'] if move(loc) is not None]+retained_labels))
    history['名单维度']=list(dict.fromkeys(history['名单维度']+plan.get('list_dimensions',[])+
        [row['selection_dimension'] for row in plan['update_purpose'].get('tables',[]) if row.get('selection_dimension')]))
    result.update(source_path=snapshot['path'],source_hash=snapshot['sha256'],source_context=deepcopy(context),
        机械记录=deepcopy(snapshot),副本接续=history)
    for key,value in (('original_path',snapshot['path']),('original_hash',snapshot['sha256'])):
        if key in result:result[key]=value
    if receipt.get('a_link'):result['a_link']=receipt['a_link']
    result['封存校验']=direct_digest(result)
    return result,receipt,sources


def direct_output_record(a,b,path,plan,plan_path):
    """只记录副本的实际搬运；两份合集原身份分别引用，不生成新的语义身份。"""
    from .表格 import read_workbook, file_hash
    snapshot=read_workbook(path)
    cells={sheet['name']+'!'+cell:raw for sheet in snapshot['sheets'] for cell,raw in sheet['cells'].items()}
    result={'schema':'附注更新执行记录-v1','output_path':str(Path(path).resolve()),
        'output_hash':file_hash(path),'plan_path':str(plan_path),
        'collections':{}}
    for carrier,record in (('A',a),('B',b)):
        reference={key:record.get(key) for key in
            ('mapping_path','mapping_hash','source_path','source_hash','gold_path','gold_hash')}
        reference['mapping_path']=reference['mapping_path'] or plan.get(carrier.lower()+'_mapping')
        reference['mapping_hash']=reference['mapping_hash'] or plan.get('inputs',{}).get(reference['mapping_path'])
        result['collections'][carrier]=reference
    for key in ('update_purpose','comparison_requirements','template_layout','identity_matches',
                 'updates','labels','operations','word_operations','retained_blanks','retained_values','unprovided_fields','issues',
                 'retained_label_targets','list_dimensions'):
        result[key]=deepcopy(plan.get(key,[]))
    for item in result['updates']+result['labels']+result['retained_blanks']+result['retained_values']:
        raw=cells.get(item['sheet']+'!'+item['cell'],{})
        item['actual_value']=deepcopy(raw.get('value'))
        item['actual_formula']=raw.get('formula')
    return result


def seal_collection(mapping, snapshot, standard):
    """只从本次识别原任务/原回答定位字段，旧重建输出不能充当出处。"""
    from .A语义阅读 import standard_material
    from .表格编排 import _stage_receipt
    result=deepcopy(mapping);result['collection_version']=1
    documents={};ordered=[]
    for ref in mapping.get('request_receipts',[]):
        path=Path(ref['path']);task_path=path.parent/'任务.json'
        task=json.loads(task_path.read_text(encoding='utf-8-sig'),object_pairs_hook=json_dates)
        _stage_receipt(ref,mapping,task['payload']['task'])
        key=task['task_id']
        if key in documents:continue
        documents[key]={'task_text':task_path.read_bytes().decode('utf-8'),
                        'receipt_text':path.read_bytes().decode('utf-8'),
                        'task_sha256':hashlib.sha256(task_path.read_bytes()).hexdigest(),
                        'receipt_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
                        'task_path':str(task_path),'receipt_path':str(path)}
        ordered.append(key)
    result['原判断与核查']=documents;parsed=_documents(result)
    result['未通过答复']=[]
    for name in mapping.get('failed_request_receipts',[]):
        path=Path(name);task_path=path.parent/'任务.json'
        text=path.read_bytes().decode('utf-8');task_text=task_path.read_bytes().decode('utf-8')
        failure={'path':name,'text':text,'sha256':_digest(text),'task_text':task_text,'task_sha256':_digest(task_text)}
        response_path=json.loads(text.lstrip('\ufeff')).get('response_evidence')
        if response_path:
            response_text=Path(response_path).read_bytes().decode('utf-8')
            failure['response_evidence']={'path':str(response_path),'text':response_text,'sha256':_digest(response_text)}
        result['未通过答复'].append(failure)
    ids={r['metric_id'] for _,r in _records(result)}
    for task,_ in parsed.values():ids.update(m['id'] for m in task['payload'].get('standards',{}).get('metrics',[]))
    definitions=standard_material(sorted(ids),standard)
    selected={r['id'] for key in ('metrics','dimensions') for r in definitions[key]}
    _require(hashlib.sha256(Path(mapping['gold_path']).read_bytes()).hexdigest()==mapping['gold_hash'],'封存标准原文件指纹变化')
    result['封存定义原文']=[line for line in Path(mapping['gold_path']).read_text(encoding='utf-8-sig').splitlines()
                         if line.strip() and json.loads(line)['id'] in selected]
    result['字段来源']={};result['来源缺口']=[];result['完整性']={};result['已确认关系']=[];result['关系缺口']=[]
    for location,record in _records(result):
        sheet,address=location.rsplit('!',1);sources={}
        for task_id in ordered:
            task,answer=parsed[task_id]
            if task['payload']['task'] not in {'a_semantic_bind','a_semantic_review'}:continue
            tables={t['id']:t for t in task['payload'].get('reading',[])}
            for index,row in enumerate(answer.get('tables',[])):
                if not isinstance(row,dict) or row.get('id') not in tables:continue
                table=tables[row['id']]
                if table['sheet']!=sheet:continue
                review=task['payload']['task']=='a_semantic_review'
                content=row.get('corrections',{}) if review else row
                root=['tables',index]+(['corrections'] if review else [])
                for name,key_field,value_field in [('metrics','metric_id','id'),('dimensions',None,'value'),('patches',None,'value')]:
                    for i,item in enumerate(content.get(name,[])):
                        if not _inside(address,item['ranges'],table):continue
                        key=key_field or 'dimensions/'+item['id']
                        ref={'task_id':task_id,'path':root+[name,i,value_field],'table_id':row['id'],'ranges':deepcopy(item['ranges'])}
                        if review:sources[key]=[ref]
                        else:sources.setdefault(key,[]).append(ref)
        values={'metric_id':record['metric_id'],**{'dimensions/'+k:v for k,v in record['dimensions'].items()}}
        result['字段来源'][location]={}
        for key,value in values.items():
            refs=sources.get(key,[])
            if refs and all(_at(parsed[r['task_id']][1],r['path'])==value for r in refs):
                result['字段来源'][location][key]=refs
            else:result['来源缺口'].append({'location':location,'field':key,'reason':'来源尚未查清或与原回答不一致；不得以method或更新要求补证'})
        result['完整性'][location]=identity_completeness(record,standard)
    for task_id in ordered:
        task,answer=parsed[task_id]
        if task['payload']['task']!='a_semantic_review':continue
        recognized={r['id']:r for r in task['payload']['recognized']['tables']}
        for n,row in enumerate(answer.get('tables',[])):
            if not isinstance(row,dict) or row.get('id') not in recognized:continue
            original=recognized[row['id']].get('relations',[])
            corrected={r['index']:r['value'] for r in row.get('corrections',{}).get('relations',[])}
            confirmed=set(row.get('checked_relations',[]))|set(corrected)
            if not original:result['关系缺口'].append({'table_id':row['id'],'reason':'原答复没有结构化关系；不从reading文字提取'})
            for i,relation in enumerate(original):
                if i not in confirmed:
                    result['关系缺口'].append({'table_id':row['id'],'index':i,'reason':'关系未确认'});continue
                value=corrected.get(i,relation)
                if _relation_blocked(value,answer,task['payload']['reading']):
                    result['关系缺口'].append({'table_id':row['id'],'index':i,'reason':'关系所指范围仍有明确核查问题'});continue
                original_refs=[{'task_id':key,'path':['tables',j,'relations',k]} for key,(raw_task,raw_answer) in parsed.items()
                    if raw_task['payload']['task']=='a_semantic_bind'
                    for j,raw_row in enumerate(raw_answer.get('tables',[])) if raw_row.get('id')==row['id']
                    for k,raw_relation in enumerate(raw_row.get('relations',[])) if raw_relation==relation]
                if not original_refs:
                    result['关系缺口'].append({'table_id':row['id'],'index':i,'reason':'核查材料中的关系尚无原绑定答复出处'});continue
                result['已确认关系'].append({'task_id':task_id,'row':n,'index':i,'原有关系内容':deepcopy(value),
                    '本方原任务引用':original_refs,'核查修正引用':{'task_id':task_id,'path':['tables',n]},
                    '原答复位置':original_refs[0]['path']})
    result['机械记录']={'sheets':deepcopy(snapshot['sheets'])}
    if mapping.get('a_link'):
        path=Path(mapping['a_link']);result['机械记录']['Excel—Word关联']={'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'content':json.loads(path.read_text(encoding='utf-8-sig'),object_pairs_hook=json_dates)}
    result['原文证据']={'context':deepcopy(mapping['source_context']), 'cells':deepcopy(snapshot.get('_all_source_evidence',{}))}
    result['封存校验']=definition_hash({k:result.get(k) for k in FIELDS})
    validate_collection(result)
    return result


def validate_collection(mapping):
    """仅核对本文件内已存值、原判断及引用；不读原表、外部标准或旧映射。"""
    _require(mapping.get('collection_version')==1,'不是已封存的新A合集')
    _require(mapping.get('封存校验')==definition_hash({k:mapping.get(k) for k in FIELDS}),'封存内容发生变化')
    parsed=_documents(mapping);standard=_standard(mapping)
    for metric in standard['metric'].values():
        _require(all(ref['id'] in standard['dimension'] for ref in metric['dimension_refs']),'封存定义缺少依赖维度原文')
        ratio=metric['definition'].get('ratio',{})
        _require(all(row['metric_id'] in standard['metric'] for row in ratio.values()),'封存定义缺少比例依赖指标原文')
    contexts={'@context:'+definition_hash(item) for item in mapping['原文证据']['context']}
    cells=mapping['原文证据']['cells']
    mechanical={s['name']+'!'+a:c for s in mapping['机械记录']['sheets'] for a,c in s['cells'].items()}
    _require(all(ref in mechanical and mechanical[ref]==cell for ref,cell in cells.items()),'原文证据与封存机械原格不符')
    for failed in mapping['未通过答复']:
        _require(_digest(failed['text'])==failed['sha256'] and _digest(failed['task_text'])==failed['task_sha256'],'未通过原答复封存变化')
        task=json.loads(failed['task_text'].lstrip('\ufeff'));answer=json.loads(failed['text'].lstrip('\ufeff'))
        _require(task['task_id']==answer['task_id'] and task['source']==mapping['source_hash'],'未通过答复缺少本次任务出处')
        if failed.get('response_evidence'):
            proof=failed['response_evidence']
            _require(_digest(proof['text'])==proof['sha256'],'解析失败原响应封存变化')
    for relation in mapping['已确认关系']:
        task,answer=parsed[relation['task_id']];row=answer['tables'][relation['row']];i=relation['index']
        declared=next(r for r in task['payload']['recognized']['tables'] if r['id']==row['id'])['relations'][i]
        _require(bool(relation['本方原任务引用']),'关系缺少原答复出处')
        for ref in relation['本方原任务引用']:
            _require(_at(parsed[ref['task_id']][1],ref['path'])==declared,'关系缺少一致的原答复出处')
        fixes={r['index']:r['value'] for r in row.get('corrections',{}).get('relations',[])}
        _require(i in row.get('checked_relations',[]) or i in fixes,'关系没有明确核查')
        _require(relation['原有关系内容']==fixes.get(i,declared),'关系与原判断及明确修正不一致')
        _require(not _relation_blocked(relation['原有关系内容'],answer,task['payload']['reading']),'关系所指范围仍有核查问题')
    for location,record in _records(mapping):
        for refs in record.get('dimension_evidence',{}).values():
            for ref in refs:
                _require(ref in cells or ref in contexts,'合集缺少原文证据：'+location+' → '+ref)
        _require(mapping['完整性'][location]==identity_completeness(record,standard),'封存完整性说明不符')
        values={'metric_id':record['metric_id'],**{'dimensions/'+k:v for k,v in record['dimensions'].items()}}
        for field,value in values.items():
            refs=mapping['字段来源'][location].get(field,[])
            if not refs:
                _require(any(r['location']==location and r['field']==field for r in mapping['来源缺口']),'封存字段没有来源或缺口记录');continue
            for ref in refs:
                _require(ref['task_id'] in parsed,'封存引用缺少原任务')
                task,answer=parsed[ref['task_id']];table=next(t for t in task['payload']['reading'] if t['id']==ref['table_id'])
                sheet,address=location.rsplit('!',1)
                _require(table['sheet']==sheet and _inside(address,ref['ranges'],table),'封存字段引用越出原范围')
                item=_at(answer,ref['path'][:-1])
                _require(item['ranges']==ref['ranges'] and (item['id']==record['metric_id'] if field=='metric_id' else item['id']==field.split('/',1)[1]),'封存字段引用不是该指标或维度的原声明')
                _require(_at(answer,ref['path'])==value,'封存值与原回答字段不一致')
    return True
