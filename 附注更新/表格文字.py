# coding: utf-8
"""只呈现文件明确记录的网格和合并，不判断表头角色或业务身份。"""
import json
from openpyxl.utils.cell import range_boundaries, coordinate_to_tuple, get_column_letter


# 与已验证的文字识别要求逐字一致，日常不读取外部技能文件。
TABLE_STRUCTURE_READING = ' 二维网格保留原行列和明确合并覆盖；先读清本格所在行列及受哪些原表头约束，再选择金标准并给出本格身份。并入位置不是独立空白，不凭空白补表头，不把原文顺序当作已经确定的表头关系。表头角色及作用范围由你依据完整原文判断，程序没有预选。'


def grid_text(rows, columns, values, merges, coordinate):
    covered = {}
    anchors = {}
    for top, left, bottom, right in merges:
        anchor = coordinate(top, left)
        anchors.setdefault((top, left), []).append(coordinate(top, left) + ':' + coordinate(bottom, right))
        for row in range(top, bottom + 1):
            for column in range(left, right + 1):
                if (row, column) != (top, left):
                    covered.setdefault((row, column), []).append(anchor)
    lines = ['| 行号 | ' + ' | '.join(columns.values()) + ' |',
             '| --- | ' + ' | '.join('---' for _ in columns) + ' |']
    for row in rows:
        fields = []
        for column in columns:
            at = (row, column)
            if at in covered:
                text = '并入' + '、'.join(covered[at])
                if at in values and values[at] not in (None, ''):
                    text += '；本位置另有原记录：' + json.dumps(values[at], ensure_ascii=False, default=str)
            elif at not in values:
                text = '未记录此位置的内容'
            else:
                text = coordinate(row, column) + ' ' + json.dumps(values[at], ensure_ascii=False, default=str)
                if values[at] in (None, ''):
                    text += '〔原格空白〕'
                if at in anchors:
                    text += '〔合并范围：' + '、'.join(anchors[at]) + '〕'
            fields.append(text.replace('|', '\\|').replace('\r', '\\r').replace('\n', '\\n'))
        lines.append('| ' + str(row) + ' | ' + ' | '.join(fields) + ' |')
    return '\n'.join(lines)


def workbook_table_ranges(snapshot, areas):
    """读取已有物理表区，不判断内容含义。"""
    tables = []
    for area in areas:
        candidates = []
        left, top, right, bottom = range_boundaries(area['range'])
        for name, named in snapshot.get('names', {}).items():
            for part in named:
                if part['sheet'] != area['sheet']:
                    continue
                l, t, r, b = range_boundaries(part['range'])
                if l <= left and r >= right and t <= top and b >= bottom:
                    candidates.append((name, part['sheet'], part['range']))
        tables.extend(candidates or [('', area['sheet'], area['range'])])
    return list(dict.fromkeys(tables))


def workbook_table_text(snapshot, areas):
    """优先呈现既有物理表区；没有表区记录时使用明确阅读范围。"""
    tables = workbook_table_ranges(snapshot, areas)
    bounds = [range_boundaries(span) for _, _, span in tables]
    lines = ['以下为原文件二维网格。并入表示明确的合并覆盖，不是独立空白；原格空白与未记录分开列示。']
    for index, (name, sheet_name, span) in enumerate(tables):
        left, top, right, bottom = bounds[index]
        containing = [j for j, (_, sheet, _) in enumerate(tables) if sheet == sheet_name
                      and bounds[j][0] <= left and bounds[j][1] <= top
                      and bounds[j][2] >= right and bounds[j][3] >= bottom]
        parent = max(containing, key=lambda j: ((bounds[j][2]-bounds[j][0]+1)*(bounds[j][3]-bounds[j][1]+1), -j))
        if parent != index:
            lines.append('原表区：' + name + '；工作表：' + sheet_name + '；范围：' + span +
                         '；原内容已列在同工作表网格 ' + tables[parent][2] + ' 中。')
            continue
        sheet = next(s for s in snapshot['sheets'] if s['name'] == sheet_name)
        values = {coordinate_to_tuple(address): (cell.get('cached_value') if cell.get('formula') else cell.get('value'))
                  for address, cell in sheet['cells'].items()
                  if top <= cell['row'] <= bottom and left <= cell['column'] <= right}
        merges = []
        for merged in sheet['merges']:
            l, t, r, b = range_boundaries(merged)
            if l <= right and r >= left and t <= bottom and b >= top:
                merges.append((t, l, b, r))
                if (t, l) not in values and (cell := sheet['cells'].get(get_column_letter(l) + str(t))):
                    lines.append('范围外合并起点原文：' + sheet_name + '!' + get_column_letter(l) + str(t) +
                                 ' ' + json.dumps(cell.get('value'), ensure_ascii=False, default=str))
        lines.append('\n工作表：' + sheet_name + '；原表区：' + name + '；范围：' + span)
        linked = [{key: item[key] for key in ('table_id', 'range_name', 'title', 'link_error') if item.get(key) is not None}
                  for item in snapshot.get('context', [])
                  if name and isinstance(item, dict) and item.get('range_name') == name and item.get('sheet') == sheet_name]
        if linked:
            lines.append('原表号及标题：' + json.dumps(linked, ensure_ascii=False))
        lines.append(grid_text(range(top, bottom + 1),
                               {c: get_column_letter(c) for c in range(left, right + 1)}, values, merges,
                               lambda r, c: get_column_letter(c) + str(r)))
    return '\n'.join(lines)
