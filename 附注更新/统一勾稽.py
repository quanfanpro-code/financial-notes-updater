# coding: utf-8
"""按统一业务身份生成更新计划；不以坐标、旧编号或金额相等寻找来源。"""
from collections import defaultdict
from copy import deepcopy
import re
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

from .表格 import _plain, file_hash, recalculate_copy, same_excel_value
from .金标准 import _pairs, _invalid_number
from .数值副本 import derive_value_copy
from .统一证据 import context_catalog, resolve_context
from .统一语义 import (_hash, _require, semantic_key, validate_facts, value_for_target,
                     ensure_standard_compatible, read_standard_v2, validate_partial_semantics, validate_fact)


def _location(fact):
    sheet, separator, cell = fact['source_reference']['location'].rpartition('!')
    _require(separator and sheet and re.fullmatch(r'[A-Z]{1,3}[1-9]\d*', cell), '事实缺少可回写的原格位置')
    return sheet, cell


def _check_mapping(mapping, standard):
    _require(mapping.get('schema_version') == 2 and
             (mapping.get('gold_hash') == standard.get('sha256') or mapping.get('_validated_standard_hash') == standard.get('sha256')),
             '映射与统一标准版本不符')
    _hash(mapping['source_hash'])
    validate_facts(mapping['facts'], standard, mapping['owners'])
    if mapping.get('slot_bindings'):
        _require(mapping.get('slot_binding_version')==1 and (mapping.get('a_semantic_reading') or mapping.get('semantic_reading') or mapping.get('external_review_v2') or mapping.get('derivation')),
                 '槽位绑定缺少共用识别及核查回执')
        for row in mapping['slot_bindings']:validate_partial_semantics(row['partial_semantics'],standard)
    occupied = set()
    for fact in mapping['facts']:
        _require(fact['source_reference']['file_hash'] == mapping['source_hash'], '事实与载体源文件版本不符')
        if fact['record_type'] == 'metric_fact':
            position = _location(fact)
            _require(position not in occupied, '同一财务数据格有重复事实')
            occupied.add(position)


def compare_semantic_slots(target, source, standard):
    """已发证格按同一标准槽位对应；位置和实际维度独立保存，不读取金额。"""
    from .统一语义 import slot_key
    for mapping in (target, source):_check_mapping(mapping, standard)
    def cells(mapping):
        result=[]
        for fact in mapping['facts']:
            if fact['record_type']!='metric_fact':continue
            sheet,cell=_location(fact)
            result.append({'sheet':sheet,'cell':cell,**{k:deepcopy(fact[k]) for k in ('metric_id','dimensions','dimension_evidence')}})
        for row in mapping.get('slot_bindings',[]):
            known=row['partial_semantics']
            result.append({'sheet':row['sheet'],'cell':row['cell'],
                           **{k:deepcopy(known[k]) for k in ('metric_id','dimensions','dimension_evidence')},
                           'missing_dimensions':deepcopy(known['missing_dimensions'])})
        _require(len(result)==len({(r['sheet'],r['cell']) for r in result}),'同一位置存在重复槽位')
        return result
    targets=cells(target);sources=cells(source);index=defaultdict(list)
    for row in sources:index[slot_key(row['metric_id'],standard)].append(row)
    groups={};issues=[];matched=0
    for row in targets:
        key=slot_key(row['metric_id'],standard);available=index.get(key,[])
        row['source_group']=key[0] if available else None
        if available:
            matched+=1
            groups.setdefault(key[0],{'slot_key':list(key),'cells':available})
        else:issues.append({'kind':'missing_standard_slot','sheet':row['sheet'],'cell':row['cell'],
                            'reason':'B已确认记录中没有此金标准槽位：'+row['metric_id']})
    for carrier,mapping in (('A',target),('B',source)):
        issues.extend({'kind':'unresolved_identity','carrier':carrier,**deepcopy(row)} for row in mapping.get('unresolved',[]))
        issues.extend({'kind':'unbound_identity','carrier':carrier,'record':deepcopy(row),
                       'reason':'此原文记录尚未绑定金标准槽位，须继续识别'} for row in mapping.get('source_identities',[]))
    identity_coverage=target.get('scope_coverage',target.get('a_identity_coverage',target.get('coverage',{})))
    if identity_coverage.get('complete') is False:
        issues.append({'kind':'target_recognition_incomplete','reason':'A逐格识别尚未完成，全部已确认目标仍保留在对应清单'})
    return {'schema':'附注语义对应-v2','targets':targets,'source_groups':groups,'issues':issues,
            'target_count':len(targets),'matched_targets':matched,
            'source_count':sum(len(g['cells']) for g in groups.values()),
            'complete':bool(targets) and not issues,
            'comparison_rule':'仅比较相同金标准槽位及维度定义；具体维度取值分别保存，同槽位位置组不表示同一个具体成员或可直接取数。'}


def _calculated_values(mapping, receipt):
    """重读Excel原生重算副本；从未把来源原有缓存当作已核实结果。"""
    _require(receipt.get('schema')=='统一语义公式计算-v1' and receipt.get('method')=='Excel.CalculateFullRebuild',
             '公式计算依据类型不符')
    _require(receipt['source_hash']==mapping['source_hash'] and file_hash(receipt['source_path'])==mapping['source_hash'],
             '公式计算依据与原件版本不符')
    _require(file_hash(receipt['workbook_path'])==receipt['workbook_hash'],'公式计算副本已变化')
    original=load_workbook(mapping['source_path'],data_only=False)
    calculated=load_workbook(receipt['workbook_path'],data_only=False)
    values=load_workbook(receipt['workbook_path'],data_only=True)
    result={}
    try:
        _require(original.sheetnames==calculated.sheetnames,'公式计算前后工作表不一致')
        for name in original.sheetnames:
            def contents(sheet):
                return {cell.coordinate:cell.value for cell in sheet._cells.values() if cell.value is not None}
            _require(contents(original[name])==contents(calculated[name]),'公式计算前后原值或公式发生变化')
        for fact in mapping['facts']:
            if fact['record_type']!='metric_fact' or fact['raw_value_state']!='formula':continue
            sheet,address=_location(fact);location=sheet+'!'+address
            _require(original[sheet][address].data_type=='f' and original[sheet][address].value==fact['raw_value'],
                     '公式记录与原件不符')
            cell=values[sheet][address];value=cell.value
            if cell.data_type=='e':state,reason='error','Excel计算错误：'+str(value)
            elif value is None and cell.data_type not in ('str','s'):
                state,reason='error','重算后仍未取得计算结果'
            elif value is None or value=='':state,reason='blank','公式结果为空，不解释成零'
            elif not isinstance(value,(str,int,float,bool)):
                state,reason='error','计算结果不是可搬运的数值或文字';value=str(value)
            else:state,reason='value','已由Excel原生重算并回读'
            result[location]={'formula':fact['raw_value'],'state':state,'value':value,'reason':reason}
    finally:
        original.close();calculated.close();values.close()
    _require(file_hash(mapping['source_path'])==mapping['source_hash'] and file_hash(receipt['workbook_path'])==receipt['workbook_hash'],
             '读取计算结果期间文件已变化')
    return result


def calculate_mapping_formulas(mapping, folder):
    """语义识别仍绑定原文件，计算副本与原公式分开保存。"""
    if not any(row['record_type']=='metric_fact' and row['raw_value_state']=='formula' for row in mapping['facts']):return
    from .流程 import save_json
    from .表格 import read_workbook
    folder=Path(folder);folder.mkdir(parents=True,exist_ok=True)
    source=Path(mapping['source_path']);calculated=folder/('公式重算副本'+source.suffix)
    _require(file_hash(source)==mapping['source_hash'],'公式计算前原件已变化')
    recalculate_copy(source,calculated)
    snapshot=read_workbook(calculated)
    _require(snapshot.get('recalculated') and snapshot['recalculation']['source_sha256']==mapping['source_hash'],
             '未取得本次Excel原生重算依据')
    receipt={'schema':'统一语义公式计算-v1','method':snapshot['recalculation']['method'],
             'source_path':str(source.resolve()),'source_hash':mapping['source_hash'],
             'workbook_path':str(calculated.resolve()),'workbook_hash':snapshot['sha256']}
    checked=_calculated_values(mapping,receipt)
    path=save_json(folder/'公式重算依据.json',receipt)
    mapping['calculation_evidence']={'path':path,'sha256':file_hash(path)}
    mapping['formula_values']=checked


def calculation_files(mapping):
    evidence=mapping.get('calculation_evidence')
    if not evidence:
        _require(not mapping.get('formula_values'),'公式计算结果缺少计算依据')
        return []
    _require(file_hash(evidence['path'])==evidence['sha256'],'公式计算依据已变化')
    receipt=json.loads(Path(evidence['path']).read_text(encoding='utf-8-sig'),object_pairs_hook=_pairs,parse_constant=_invalid_number)
    _require(_calculated_values(mapping,receipt)==mapping.get('formula_values'),'保存的公式计算结果与实际副本不符')
    return [evidence['path'],receipt['workbook_path'],receipt['source_path']]


def _effective_source(fact, mapping, standard):
    if fact['raw_value_state']!='formula':return fact
    calculated=mapping.get('formula_values',{}).get(fact['source_reference']['location'])
    if not calculated or not mapping.get('calculation_evidence') or calculated['state']=='error':return fact
    _require(calculated['formula']==fact['raw_value'],'计算结果与原公式不符')
    result=deepcopy(fact);result['raw_value_state']=calculated['state'];result['raw_value']=calculated['value']
    from .统一语义 import validate_fact
    validate_fact(result,standard)
    return result


def load_fact_mapping(path, standard, _ancestors=frozenset()):
    """独立映射读回时核对当前原件、逐格原值、证据位置及已保存回执版本。"""
    path=Path(path).resolve(); raw=path.read_bytes()
    _require(path not in _ancestors,'派生映射存在循环引用')
    ancestors=_ancestors|{path}
    mapping=json.loads(raw.decode('utf-8-sig'),object_pairs_hook=_pairs,parse_constant=_invalid_number)
    mapping.pop('_validated_standard_hash',None)
    if mapping.get('value_derivation'):
        from .数值副本 import validate_value_copy
        return validate_value_copy(path,mapping,standard,ancestors)
    ensure_standard_compatible(mapping,standard)
    mapping['_validated_standard_hash']=standard['sha256']
    _check_mapping(mapping,standard)
    def verify_classification(content):
        if content.get('a_semantic_reading') or content.get('semantic_reading'):
            from .A语义阅读 import verify_mapping
            verify_mapping(content,standard)
            return
        if content.get('identity_review'):
            from .统一识别 import verify_identity_review_mapping
            verify_identity_review_mapping(content,standard)
            return
        _require('packet_mappings' in content or content.get('reused_mapping') or content.get('derivation')
                 or not any(row.get('recognition_provenance',{}).get('method')=='model_v2_recognition_checked'
                            for row in content['facts']+content.get('excluded',[])), '识别后核查结果缺少实际核查依据')
        from .统一识别 import agree_rounds
        for row in content.get('unresolved',[]):
            if not row.get('partial_semantics'):continue
            validate_partial_semantics(row['partial_semantics'],standard)
            rounds=row.get('rounds',[])
            _require(len(rounds)==2 and all(item.get('kind') in {'partial_metric_fact','metric_fact'} for item in rounds),
                     '部分语义缺少两轮独立记录')
            for item in rounds:
                if item['kind']=='partial_metric_fact':validate_partial_semantics(item['partial_semantics'],standard)
                else:validate_fact(item['fact'],standard)
            if not content.get('equivalence_proofs') and 'packet_mappings' not in content:
                expected=agree_rounds({row['cell']:rounds[0]},{row['cell']:rounds[1]},standard)
                _require(expected['unresolved'][0].get('partial_semantics')==row['partial_semantics'],'保存的部分语义与两轮识别不符')
        if not content.get('equivalence_proofs'):
            if not content.get('derivation') and 'packet_mappings' not in content:
                _require(not any(row['recognition_provenance']['method']=='model_v2_two_rounds_equivalence' for row in content['facts']),
                         '同义识别事实缺少保存的复核依据')
            return
        historical=standard if content['gold_hash']==standard['sha256'] else read_standard_v2(content['gold_path'])
        _require(historical['sha256']==content['gold_hash'],'同义复核所用历史标准变化')
        _require(len(content.get('classification_rounds',[]))==2,'同义识别缺少原始两轮分类')
        for proof in content['equivalence_proofs']:
            for pair in proof['pairs']:
                for side in ('left','right'):
                    _require(pair.get(side+'_context',[])==content.get('source_context',[]),'同义识别复核说明与实际分类不符')
                    if pair[side].get('record_type')=='partial_metric_fact':
                        reference=pair[side]['source_reference'];address=reference['location'].rpartition('!')[2]
                        _require(reference=={'file_hash':content['source_hash'],'location':content['sheet']+'!'+address} and
                                 Path(pair[side+'_path']).resolve()==Path(content['source_path']).resolve(),'部分语义复核未绑定实际原格与原件')
        expected=agree_rounds(*content['classification_rounds'],historical,equivalence_proofs=content['equivalence_proofs'],
                              agreement_version=content.get('agreement_version',1))
        for key in ('facts','owners','unresolved','excluded'):
            _require(content[key]==expected[key],'同义识别结果与原分类和复核不符：'+key)
    if mapping.get('external_review_v2'):
        from .统一外部复核 import validate_external_mapping
        validate_external_mapping(mapping,standard,ancestors)
    else:
        _require(not any(row['recognition_provenance']['method']=='external_semantic_review_v2' for row in mapping['facts'])
                 or mapping.get('derivation') or mapping.get('reused_mapping'),'外部语义缺少实际采纳依据')
        verify_classification(mapping)
    calculation_files(mapping)
    if mapping.get('derivation'):
        derivation=mapping['derivation']; plan_path=Path(derivation['plan_path'])
        _require(hashlib.sha256(plan_path.read_bytes()).hexdigest()==derivation['plan_hash'],'派生映射的更新清单已变化')
        plan=json.loads(plan_path.read_text(encoding='utf-8-sig'),object_pairs_hook=_pairs,parse_constant=_invalid_number)
        for original,digest in plan['inputs'].items():
            _require(hashlib.sha256(Path(original).read_bytes()).hexdigest()==digest,'派生映射的依据文件已变化')
        historical=standard if plan['gold_hash']==standard['sha256'] else read_standard_v2(plan['gold_path'])
        _require(historical['sha256']==plan['gold_hash'],'派生依据所用历史标准已变化')
        parent=load_fact_mapping(plan['a_mapping'],historical,ancestors)
        source_mapping=load_fact_mapping(plan['b_mapping'],historical,ancestors)
        fresh=verify_fact_update_plan(plan,parent,source_mapping,historical)
        expected=derive_fact_mapping(fresh.pop('_working_target'),mapping['source_path'],fresh,plan_path,mapping.get('a_link'))
        for key in ('facts','owners','excluded','unresolved','slot_bindings','coverage','complete','source_context','derivation'):
            _require(mapping.get(key)==expected.get(key),'派生映射与原语义及更新清单不符：'+key)
    source=Path(mapping['source_path']).resolve()
    _require(hashlib.sha256(source.read_bytes()).hexdigest()==mapping['source_hash'],'映射绑定的原件已变化')
    book=load_workbook(source,data_only=False,keep_vba=source.suffix.lower()=='.xlsm')
    try:
        catalog=context_catalog(mapping.get('source_context',[]),source)
        for fact in mapping['facts']:
            sheet,cell=_location(fact)
            _require(sheet in book,'事实引用的工作表不存在')
            original=book[sheet][cell].value
            if fact['record_type']=='metric_fact':
                _require(original==fact['raw_value'],'映射原值与实际工作簿不符：'+sheet+'!'+cell)
                references=[ref for values in fact['dimension_evidence'].values() for ref in values]
                for reference in references:
                    if reference.startswith('@context:'):
                        resolve_context(reference,catalog,fact['source_reference']['location'])
                        continue
                    evidence_sheet,separator,address=reference.rpartition('!')
                    _require(separator and evidence_sheet in book and re.fullmatch(r'[A-Z]{1,3}[1-9]\d*',address),'维度证据不是真实原格位置')
                    value=book[evidence_sheet][address].value
                    _require(value is not None and str(value).strip(),'维度证据原格为空')
        for row in mapping.get('unresolved',[])+mapping.get('slot_bindings',[]):
            known=row.get('partial_semantics')
            if not known:continue
            sheet_name=row.get('sheet',mapping.get('sheet'))
            _require(sheet_name in book and re.fullmatch(r'[A-Z]{1,3}[1-9]\d*',row['cell']),'部分语义原格不存在')
            references=known['evidence_cells']+[ref for values in known['dimension_evidence'].values() for ref in values]
            for reference in references:
                if reference.startswith('@context:'):
                    resolve_context(reference,catalog,sheet_name+'!'+row['cell']);continue
                sh,separator,address=reference.rpartition('!')
                _require(separator and sh in book and re.fullmatch(r'[A-Z]{1,3}[1-9]\d*',address),'部分语义证据位置无效')
                value=book[sh][address].value
                _require(value is not None and str(value).strip(),'部分语义证据原格为空')
    finally:book.close()
    for receipt in mapping.get('request_receipts',[]):
        _require(hashlib.sha256(Path(receipt['path']).read_bytes()).hexdigest()==receipt['sha256'],'识别回执已变化')
    packet_facts=[];packet_excluded=[];packet_owners={};packet_partial=[]
    if mapping.get('reused_mapping'):
        reused=mapping['reused_mapping']
        _require(file_hash(reused['path'])==reused['sha256'],'复用的原语义映射已变化')
        if reused.get('selection') == 'unchanged_definitions':
            from .统一语义 import compatible_classification
            original=json.loads(Path(reused['path']).read_text(encoding='utf-8-sig'),object_pairs_hook=_pairs,parse_constant=_invalid_number)
            previous=load_fact_mapping(reused['path'],read_standard_v2(original['gold_path']),ancestors)
            selected=compatible_classification(previous,standard)
        else:
            previous=load_fact_mapping(reused['path'],standard,ancestors)
            selected=previous
        _require(previous['source_hash']==mapping['source_hash'] and previous.get('source_context',[])==mapping.get('source_context',[]),
                 '复用映射的原件或上下文不一致')
        packet_facts.extend(selected['facts']);packet_excluded.extend(selected.get('excluded',[]));packet_owners.update(selected['owners'])
    for packet in mapping.get('packet_mappings',[]):
        _require(hashlib.sha256(Path(packet['path']).read_bytes()).hexdigest()==packet['sha256'],'识别任务包映射已变化')
        content=json.loads(Path(packet['path']).read_text(encoding='utf-8-sig'),object_pairs_hook=_pairs,parse_constant=_invalid_number)
        _require(content['source_hash']==mapping['source_hash'],'识别任务包的原件不符')
        # 历史识别包保留真实版本；只有引用的指标及维度未改义时才能接续。
        if content['gold_hash'] != mapping['gold_hash']:
            ensure_standard_compatible(content, standard)
        if 'source_context' in content:
            _require(content['source_context']==mapping.get('source_context',[]),'识别任务包与整份映射的语义说明不符')
        verify_classification(content)
        packet_facts.extend(content['facts']);packet_owners.update(content['owners'])
        packet_excluded.extend({'sheet':content['sheet'],**row} for row in content.get('excluded',[]))
        packet_partial.extend({'sheet':content['sheet'],**row} for row in content.get('unresolved',[]) if row.get('partial_semantics'))
    if 'packet_mappings' in mapping:
        _require(mapping['facts']==packet_facts and mapping['owners']==packet_owners and mapping.get('excluded',[])==packet_excluded,
                 '整份映射中的语义或覆盖分类与已保存识别任务包不符')
        _require([row for row in mapping.get('unresolved',[]) if row.get('partial_semantics')]==packet_partial,
                 '保存的部分语义与两轮识别不符：整份映射未保持已验证任务包的结果')
    if 'coverage' in mapping:
        from .表格 import read_workbook,materialize_candidate_blanks
        from .统一识别 import workbook_coverage
        snapshot=read_workbook(source,context=mapping.get('source_context',[]));materialize_candidate_blanks(snapshot)
        coverage=workbook_coverage(mapping,snapshot)
        _require(mapping['coverage']==coverage and mapping.get('complete')==coverage['complete'],'整表覆盖声明与原件不符')
    if mapping.get('recognition_scope') and mapping['recognition_scope'].get('mode')=='a_semantics':
        _require(bool(mapping.get('a_semantic_reading')),'A语义识别缺少实际阅读及确认依据')
    elif mapping.get('recognition_scope') and not mapping.get('external_review_v2'):
        # 外部复核已递归核验历史范围及本次采纳，当前身份不冒充历史识别原结果。
        from .表格 import read_workbook, materialize_candidate_blanks
        from .表格编排 import validate_saved_scope
        snapshot=read_workbook(source,context=mapping.get('source_context',[]));materialize_candidate_blanks(snapshot)
        validate_saved_scope(mapping,snapshot,standard)
    _require(hashlib.sha256(source.read_bytes()).hexdigest()==mapping['source_hash'],'核验期间原件已变化')
    _require(path.read_bytes()==raw,'核验期间映射已变化')
    mapping['mapping_path']=str(path); mapping['mapping_hash']=hashlib.sha256(raw).hexdigest()
    return mapping


def verify_fact_update_plan(plan, target, source, standard):
    """写入前重新按语义取数，并核对全部可执行内容。"""
    _require(plan.get('schema')=='附注更新清单-v2' and plan.get('complete') and not plan.get('issues'),
             '更新清单尚未确认完整')
    _require(plan.get('gold_hash')==standard['sha256'],'更新清单的标准版本不符')
    for mapping in (target,source):
        _require(plan['inputs'].get(mapping['mapping_path'])==mapping['mapping_hash'],'更新清单未绑定当前映射版本')
    original_target=target
    scope=plan.get('validation_scope')
    if scope:
        _require(isinstance(scope.get('label'),str) and scope['label'].strip(),'局部验证须明确名称')
        _require(not plan.get('structure'),'局部验证不隐式执行结构增减')
        target=select_mapping_scope(target,scope['target']);source=select_mapping_scope(source,scope['source'])
    from .表格编排 import source_ready
    _require(target.get('scope_coverage' if scope else 'coverage',{}).get('complete'),'A范围语义尚未完整')
    _require(source.get('scope_coverage',{}).get('complete') if scope else source_ready(source,target,standard),'B完整初识或当前A所需内容尚未完成')
    if plan.get('structure'):
        from .统一结构 import load_structural_target
        target=load_structural_target(plan['structure'],target,source,standard)
        _require(plan['inputs'].get(target['mapping_path'])==target['mapping_hash'],'清单未绑定调整后独立语义映射')
    expected_operations=plan.get('structure',{}).get('layout',{}).get('operations',[]) or plan.get('update_purpose',{}).get('operations',[])
    _require(plan.get('word_operations',[])==expected_operations,'Word结构变更与已核验语义布局不符')
    fresh=build_fact_update_plan(target,source,standard,equivalence_proofs=plan.get('equivalence_proofs',[]),update_purpose=plan.get('update_purpose'))
    from .表格 import complete_numeric_formats
    complete_numeric_formats(target,fresh)
    from .统一未采用来源 import validate_unused_source_reviews
    unused=validate_unused_source_reviews(plan.get('unused_source_reviews',[]),target,source,standard,fresh)
    _require(plan.get('unused_sources',[])==unused,'未采用来源与原文及两轮复核依据不一致')
    fresh['unused_sources']=unused
    unused_ids={row['source_id'] for row in unused}
    fresh['unplaced_sources']=[row for row in fresh['unplaced_sources'] if row['id'] not in unused_ids]
    _require(not fresh['issues'] and not fresh['unplaced_sources'],'重新勾稽仍有未处理事项')
    for row in fresh['updates']:
        definition=standard['metric'][row['metric_id']]['definition']
        row['semantic_field']=' / '.join(definition[key] for key in ('economic_object','measure','measurement_basis'))
    for key in ('updates','retained_blanks','unprovided_fields','operations','labels','unplaced_sources','source_a_hash','source_b_hash','template_layout'):
        _require(plan.get(key)==fresh.get(key),'清单与重新语义取数不一致：'+key)
    _require(plan.get('caption_updates',[])==fresh.get('caption_updates',[]),'标题安排与重新语义取数不一致')
    fresh['_working_target']=original_target if scope else target
    return fresh


def derive_fact_mapping(parent, source_path, plan, plan_path, word_link=None):
    """原格语义不变时沿用已识别含义，原值必须与实际更新结果逐格一致。"""
    from .表格 import read_workbook,materialize_candidate_blanks
    from .统一识别 import workbook_coverage
    _require(not plan.get('operations') or plan.get('template_layout'),'布局变化须先迁移逐格语义，不能沿用原位置')
    _require(not plan.get('labels') or plan.get('update_purpose',{}).get('bindings'), '文字更新缺少本次模板维度绑定')
    if plan.get('operations'):
        from .模板布局 import relocate_mapping,verify_output
        verify_output(parent,source_path,plan)
        parent=relocate_mapping(parent,plan['template_layout'])
    path=Path(source_path).resolve(); digest=hashlib.sha256(path.read_bytes()).hexdigest()
    result=deepcopy(parent)
    result['facts']=[row for row in result['facts'] if row['record_type']!='metric_fact']+list(_template_cell_records(parent).values())
    result['slot_bindings']=[]
    for key in ('packet_mappings','request_receipts','mapping_hash','mapping_path','a_link','review_path',
                'calculation_evidence','formula_values','calculation_issues','_validated_standard_hash','reused_mapping',
                'classification_rounds','equivalence_proofs','external_review_v2','outline','recognition_scope',
                'a_semantic_reading','semantic_reading','identity_review','_template_layout'):
        result.pop(key,None)
    saved_plan=json.loads(Path(plan_path).read_text(encoding='utf-8-sig'),object_pairs_hook=_pairs,parse_constant=_invalid_number)
    if saved_plan.get('update_purpose'):
        if saved_plan['update_purpose'].get('bindings'):
            dimensions={row['sheet']+'!'+row['cell']:row['dimensions']
                        for row in [*plan['updates'],*plan.get('unprovided_fields',[])]}
        else:
            selected=select_mapping_scope(parent,saved_plan['validation_scope']['target']) if saved_plan.get('validation_scope') else parent
            projected=target_for_purpose(selected,saved_plan['update_purpose'])
            dimensions={row['id']:row['dimensions'] for row in projected['facts'] if row['record_type']=='metric_fact'}
        from .统一语义 import definition_hash
        evidence='本次更新用途：'+json.dumps(saved_plan['update_purpose'],ensure_ascii=False,sort_keys=True)
        result.setdefault('source_context',[]).append(evidence)
        for row in result['facts']:
            if row['id'] in dimensions:
                for key,value in dimensions[row['id']].items():
                    if row['dimensions'].get(key)!=value:row['dimension_evidence'][key]=['@context:'+definition_hash(evidence)]
                row['dimension_evidence']={key:row['dimension_evidence'][key] for key in dimensions[row['id']]}
                row['dimensions']=deepcopy(dimensions[row['id']])
                row['owner_reference']=row['id']
                result['owners'][row['owner_reference']]=deepcopy(row['dimensions'])
    result.update(source_path=str(path),source_hash=digest,original_path=str(path),original_hash=digest,carrier='A',usage={},
                  gold_path=saved_plan['gold_path'],gold_hash=plan['gold_hash'],
                  derivation={'plan_path':str(Path(plan_path).resolve()),'plan_hash':hashlib.sha256(Path(plan_path).read_bytes()).hexdigest()})
    if word_link:result['a_link']=str(Path(word_link).resolve())
    if saved_plan.get('update_purpose',{}).get('bindings'):
        from .统一语义 import required_cell_dimensions,partial_pending_dimensions
        standard=read_standard_v2(saved_plan['gold_path'])
        full=[]
        for row in result['facts']:
            if row['record_type']!='metric_fact':full.append(row);continue
            known={k:deepcopy(row[k]) for k in ('metric_id','dimensions','dimension_evidence')}
            known['missing_dimensions']=sorted(required_cell_dimensions(standard['metric'][row['metric_id']],standard)-set(row['dimensions']))
            known['evidence_cells']=list(dict.fromkeys(ref for refs in row['dimension_evidence'].values() for ref in refs))
            if partial_pending_dimensions(known,standard):
                sheet,cell=_location(row)
                result.setdefault('slot_bindings',[]).append({'sheet':sheet,'cell':cell,'partial_semantics':known,
                    'recognition_provenance':deepcopy(row['recognition_provenance'])})
                result['slot_binding_version']=1
            else:full.append(row)
        result['facts']=full
        used_owners={row.get('owner_reference') for row in full}
        result['owners']={key:value for key,value in result['owners'].items() if key in used_owners}
    updates={(row['sheet'],row['cell']):row for row in [*plan['updates'],*plan.get('labels',[])]}
    # openpyxl 保存会把空字符串格归一化为 None；两种表示同为空白，与写入回读校验一致。
    def _same_value(actual,wanted):return same_excel_value(actual,wanted)
    book=load_workbook(path,data_only=False)
    try:
        for row in result['facts']:
            if row['record_type']!='metric_fact':
                row['source_reference']['file_hash']=digest
                continue
            sheet,cell=_location(row); original=row['raw_value']
            expected=updates.get((sheet,cell),{}).get('value',original)
            value=book[sheet][cell].value
            _require(_same_value(value,expected),'派生映射对应的更新值与实际工作簿不符：'+sheet+'!'+cell)
            row['source_reference']['file_hash']=digest
            row['raw_value']=value
            # 空白语义与 validate_fact 一致：None 或纯空白字符串都算空白，不因全角空格误判为实际值。
            row['raw_value_state']='blank' if value is None or isinstance(value,str) and not value.strip() else ('formula' if book[sheet][cell].data_type=='f' else 'value')
    finally:book.close()
    snapshot=read_workbook(path,context=result.get('source_context',[]));materialize_candidate_blanks(snapshot)
    previous=read_workbook(parent['source_path'],context=parent.get('source_context',[])) if not plan.get('operations') else snapshot
    materialize_candidate_blanks(previous)
    _require([s['name'] for s in snapshot['sheets']]==[s['name'] for s in previous['sheets']],'派生表格的工作表结构已变化')
    for before,after in zip(previous['sheets'],snapshot['sheets']):
        for key in ('merges','hidden','hidden_rows','hidden_columns'):
            _require(before[key]==after[key],'派生表格的布局已变化：'+key)
        _require(set(before['cells'])==set(after['cells']),'派生表格的单元格范围已变化')
        for address,cell in before['cells'].items():
            expected=updates.get((before['name'],address),{}).get('value',cell['value'])
            _require(_same_value(after['cells'][address]['value'],expected),'派生表格存在清单外的改动：'+before['name']+'!'+address)
    if plan.get('template_layout'):
        occupied={r['source_reference']['location'] for r in result['facts']}
        occupied.update(r['sheet']+'!'+r['cell'] for kind in ('excluded','unresolved','slot_bindings') for r in result.get(kind,[]))
        actual_cells={sheet['name']+'!'+cell for sheet in snapshot['sheets'] for cell in sheet['cells']}
        for location,item in plan['template_layout']['added'].items():
            if location not in occupied and location in actual_cells:
                sheet,cell=location.rsplit('!',1)
                result.setdefault('excluded',[]).append({'sheet':sheet,'cell':cell,'category':'label','reason':'模板增行的名称或留白，财务格已另行绑定'})
    result['coverage']=workbook_coverage(result,snapshot);result['complete']=result['coverage']['complete']
    return result


def select_mapping_scope(mapping, ranges):
    """仅生成局部执行视图；原映射及其整份覆盖声明始终保留。"""
    from openpyxl.utils.cell import range_boundaries, coordinate_to_tuple
    _require(isinstance(ranges,list) and ranges, '局部验证须明确表格范围')
    bounds=[]
    for row in ranges:
        _require(set(row)=={'sheet','range'} and isinstance(row['sheet'],str), '局部验证范围格式不符')
        left,top,right,bottom=range_boundaries(row['range'])
        _require(all(type(v) is int and v>0 for v in (left,top,right,bottom)) and left<=right and top<=bottom, '局部范围无效')
        bounds.append((row['sheet'],left,top,right,bottom))
    def selected(location):
        sheet,_,address=location.rpartition('!');r,c=coordinate_to_tuple(address)
        return any(sheet==s and l<=c<=rr and t<=r<=b for s,l,t,rr,b in bounds)
    expected=mapping.get('coverage',{}).get('expected_cells')
    if expected is None:
        from .表格 import read_workbook,materialize_candidate_blanks
        snapshot=read_workbook(mapping['source_path'],context=mapping.get('source_context',[]));materialize_candidate_blanks(snapshot)
        expected=[sheet['name']+'!'+cell for sheet in snapshot['sheets'] for cell in sheet['cells']]
    expected={location for location in expected if selected(location)}
    _require(expected,'所选范围没有候选格')
    view=deepcopy(mapping)
    view['facts']=[row for row in view['facts'] if selected(row['source_reference']['location'])]
    assigned=[row['source_reference']['location'] for row in view['facts'] if row['record_type']=='metric_fact']
    if 'slot_bindings' in view:
        view['slot_bindings']=[row for row in view['slot_bindings'] if selected(row['sheet']+'!'+row['cell'])]
        assigned.extend(row['sheet']+'!'+row['cell'] for row in view['slot_bindings'])
    for kind in ('excluded','unresolved'):
        view[kind]=[row for row in view.get(kind,[]) if selected(row.get('sheet',mapping.get('sheet',''))+'!'+row['cell'])]
        assigned.extend(row.get('sheet',mapping.get('sheet',''))+'!'+row['cell'] for row in view[kind])
    _require(len(assigned)==len(set(assigned)) and set(assigned)==expected,'局部验证范围存在遗漏或重复分类')
    view['scope_coverage']={'scope':'local_validation','ranges':deepcopy(ranges),'expected_count':len(expected),
        'financial_cells':len(view['facts'])+len(view.get('slot_bindings',[])),'unresolved_cells':len(view['unresolved']),
        'outside_unresolved_count':len(mapping.get('unresolved',[]))-len(view['unresolved']),
        'complete':not view['unresolved']}
    return view


def _template_cell_records(mapping):
    """模板搬运按已识别槽位取原格；未填姓名仍保留原状态。"""
    rows={r['source_reference']['location']:r for r in mapping['facts'] if r['record_type']=='metric_fact'}
    if not mapping.get('slot_bindings'):return rows
    book=load_workbook(mapping['source_path'],data_only=False)
    try:
        for item in mapping['slot_bindings']:
            location=item['sheet']+'!'+item['cell'];known=item['partial_semantics']
            layout=mapping.get('_template_layout',{});added=layout.get('added',{}).get(location)
            original=added['template'] if added else next((old for old,new in layout.get('positions',{}).items() if new==location),location)
            sh,address=original.rsplit('!',1);cell=book[sh][address]
            rows[location]={'id':location,'record_type':'metric_fact','owner_reference':location,
                **{k:deepcopy(known[k]) for k in ('metric_id','dimensions','dimension_evidence')},
                'source_reference':{'file_hash':mapping['source_hash'],'location':location},
                'recognition_provenance':deepcopy(item['recognition_provenance']),
                'raw_value':None if added else _plain(cell.value),'raw_value_state':'blank' if added else ('formula' if cell.data_type=='f' else
                ('blank' if cell.value is None or isinstance(cell.value,str) and not cell.value.strip() else 'value'))}
    finally:book.close()
    return rows


def target_for_purpose(mapping, purpose, source=None, standard=None):
    """本次模板的维度取自已确定B来源，原始识别不改写。"""
    if not purpose:return mapping
    if 'bindings' in purpose:
        purpose=deepcopy(purpose)
        _require({'bindings','labels','reason'}<=set(purpose)<={'bindings','labels','reason','operations','unprovided'} and isinstance(purpose['reason'],str) and purpose['reason'].strip(), '模板更新须有明确依据')
        _require(source is not None and standard is not None,'模板更新缺少B及金标准')
        view=deepcopy(mapping)
        if purpose.get('operations'):
            from .模板布局 import project_layout
            view=project_layout(mapping,source,standard,purpose['operations'],
                                source_locations={row['source'] for row in purpose['bindings']})
        targets=_template_cell_records(view)
        sources=_template_cell_records(source)
        _require(isinstance(purpose['bindings'],list) and purpose['bindings'],'模板更新缺少逐格来源')
        _require(isinstance(purpose['labels'],list),'模板更新缺少文字安排')
        labels=[];bound={};copied_text=[];copied_periods={};copied_period_targets={};unprovided={}
        for item in purpose['bindings']:
            _require(set(item)=={'target','source'} and item['target'] in targets and item['source'] in sources,'模板绑定须引用已识别的实际数据格')
            _require(item['target'] not in bound,'模板数据格重复绑定')
            old,new=targets[item['target']],sources[item['source']]
            _require(old['metric_id']==new['metric_id'],'模板更新不能改变金标准槽位')
            bound[item['target']]=item['source']
        for current in (mapping,source):_require(file_hash(current['source_path'])==current['source_hash'],'模板原件已变化')
        a=load_workbook(mapping['source_path'],data_only=False);b=load_workbook(source['source_path'],data_only=False)
        try:
            from .表格 import _apply_operation
            for operation in purpose.get('operations',[]):_apply_operation(a,operation)
            from .统一证据 import context_catalog
            caption_contexts=[]
            for record in context_catalog(view.get('source_context',[]),workbook=a).values():
                item,scope=record['content'],record['scope']
                if not (isinstance(item,dict) and item.get('range_name') and scope):continue
                top,left,bottom,right=scope['bounds']
                caption_contexts.append({**item,'first_row':top,'first_column':left,
                    'row_count':bottom-top+1,'column_count':right-left+1})
            period_headers=defaultdict(list)
            from openpyxl.utils.cell import coordinate_to_tuple,range_boundaries
            from .表格 import _definitions,_destinations
            named_ranges=[(sh,*range_boundaries(area)) for definition in _definitions(a) for sh,area in _destinations(definition)]
            for destination,origin in bound.items():
                old,new=targets[destination],sources[origin]
                if old['dimensions'].get('period')==new['dimensions'].get('period'):continue
                before_refs=old['dimension_evidence'].get('period',[]);after_refs=new['dimension_evidence'].get('period',[])
                if any((destination in label.get('period_targets',[]) or label['target'] in before_refs)
                       and label['source'] in after_refs and (label.get('period_source') or label.get('period_sources'))
                       for label in purpose['labels']):continue
                sheet,address=destination.rsplit('!',1);row,column=coordinate_to_tuple(address);headers=[]
                for ref in before_refs:
                    if ref.startswith('@context:') or ref in targets:continue
                    sh,cell=ref.rsplit('!',1);r,c=coordinate_to_tuple(cell)
                    if (sh==sheet and c==column and r<row and isinstance(a[sh][cell].value,str)
                            and any(s==sheet and left<=column<=right and top<=r<row<=bottom for s,left,top,right,bottom in named_ranges)):
                        headers.append((r,ref))
                if headers:period_headers[max(headers)[1]].append((destination,origin,after_refs))
            for header,items in period_headers.items():
                common=set.intersection(*(set(refs) for _,_,refs in items))
                choices=[]
                for ref in common:
                    if ref.startswith('@context:') or ref in sources:continue
                    sheet,cell=ref.rsplit('!',1);row,_=coordinate_to_tuple(cell)
                    if isinstance(b[sheet][cell].value,str):choices.append((row,ref))
                if not choices:continue
                source_header=max(choices)[1]
                existing=next((label for label in purpose['labels'] if label['target']==header),None)
                if existing is not None and existing['source']!=source_header:continue
                label=existing if existing is not None else {'target':header,'source':source_header,'preserve_text':True}
                label.update(period_sources=list(dict.fromkeys(origin for _,origin,_ in items)),
                             period_targets=list(dict.fromkeys(destination for destination,_,_ in items)))
                label.pop('period_source',None)
                if existing is None:purpose['labels'].append(label)
            caption_sources=defaultdict(list)
            from openpyxl.utils.cell import coordinate_to_tuple,get_column_letter
            for destination,origin in bound.items():
                old,new=targets[destination],sources[origin]
                target_sheet,target_address=destination.rsplit('!',1);target_row,target_col=coordinate_to_tuple(target_address)
                source_sheet,source_address=origin.rsplit('!',1);source_row,_=coordinate_to_tuple(source_address)
                for key,value in new['dimensions'].items():
                    if old['dimensions'].get(key)==value or standard['dimension'][key]['value_type']!='selection':continue
                    before_refs=old['dimension_evidence'].get(key,[]);after_refs=new['dimension_evidence'][key]
                    if any(item['target'] in before_refs and any(ref in after_refs for ref in [item['source'],*item.get('source_aliases',[])])
                           for item in purpose['labels']):continue
                    identical=[]
                    for left in before_refs:
                        if left.startswith('@context:') or left in targets:continue
                        sh,cell=left.rsplit('!',1);text=a[sh][cell].value
                        if not isinstance(text,str) or not text.strip() or a[sh][cell].data_type=='f':continue
                        for right in after_refs:
                            if right.startswith('@context:') or right in sources:continue
                            source_sh,source_cell=right.rsplit('!',1)
                            if b[source_sh][source_cell].data_type!='f' and b[source_sh][source_cell].value==text:
                                identical.append((left,right))
                    if len(identical)==1 and not any(item['target']==identical[0][0] for item in purpose['labels']):
                        left,right=identical[0];purpose['labels'].append({'target':left,'source':right})
                    elif len(identical)==1:
                        left,right=identical[0]
                        existing=next(item for item in purpose['labels'] if item['target']==left)
                        if existing.get('preserve_text') and right!=existing['source']:
                            aliases=existing.setdefault('source_aliases',[])
                            if right not in aliases:aliases.append(right)
                    else:
                        before_labels=[ref for ref in before_refs if not ref.startswith('@context:') and ref not in targets]
                        after_labels=[ref for ref in after_refs if not ref.startswith('@context:') and ref not in sources]
                        if len(before_labels)==len(after_labels)==1 and not any(label['target']==before_labels[0] for label in purpose['labels']):
                            left,right=before_labels[0],after_labels[0]
                            sh,cell=left.rsplit('!',1);source_sh,source_cell=right.rsplit('!',1)
                            if isinstance(a[sh][cell].value,str) and isinstance(b[source_sh][source_cell].value,str):
                                purpose['labels'].append({'target':left,'source':right})
                    for context in caption_contexts:
                        if not (isinstance(context,dict) and context.get('range_name') and context.get('sheet')==target_sheet
                                and context['first_row']<=target_row<context['first_row']+context['row_count']
                                and context['first_column']<=target_col<context['first_column']+context['column_count']):continue
                        caption=f"{target_sheet}!{get_column_letter(context['first_column'])}{context['first_row']-1}"
                        if caption not in before_refs or a[target_sheet][caption.rsplit('!',1)[1]].value!=context.get('title'):continue
                        candidates=[]
                        for ref in after_refs:
                            if ref.startswith('@context:') or ref in sources:continue
                            sh,cell=ref.rsplit('!',1);row,_=coordinate_to_tuple(cell)
                            if sh==source_sheet and row<source_row and isinstance(b[sh][cell].value,str) and b[sh][cell].data_type!='f':
                                candidates.append((row,ref))
                        if candidates:
                            caption_sources[caption].append({ref for _,ref in candidates})
            for caption,groups in caption_sources.items():
                if any(item['target']==caption for item in purpose['labels']):continue
                origins=set.intersection(*groups)
                if origins:
                    nearest=max(coordinate_to_tuple(ref.rsplit('!',1)[1])[0] for ref in origins)
                    origins={ref for ref in origins if coordinate_to_tuple(ref.rsplit('!',1)[1])[0]==nearest}
                else:origins=set.union(*groups)
                origins=sorted(origins)
                texts={b[ref.rsplit('!',1)[0]][ref.rsplit('!',1)[1]].value for ref in origins}
                _require(len(texts)==1,'同一表外名称的B原文不一致，须核准来源：'+caption)
                purpose['labels'].append({'target':caption,'source':origins[0],
                    **({'source_aliases':origins[1:]} if len(origins)>1 else {})})
            occupied=set()
            label_targets={ref for row in targets.values() for refs in row['dimension_evidence'].values() for ref in refs if not ref.startswith('@context:')}
            label_targets.update(f"{item['sheet']}!{get_column_letter(item['first_column'])}{item['first_row']-1}"
                                 for item in caption_contexts if item.get('first_row'))
            for item in purpose['labels']:
                _require({'target','source'}<=set(item)<={'target','source','source_aliases','period_source','period_sources','period_targets','preserve_text','selection_source','selection_dimension'},'模板文字须绑定实际来源位置')
                _require(not ('period_source' in item and 'period_sources' in item),'期间来源不能重复声明')
                period_locations=item.get('period_sources',([item['period_source']] if item.get('period_source') else []))
                _require(isinstance(period_locations,list) and all(isinstance(ref,str) for ref in period_locations),'期间来源须为数据格清单')
                _require(type(item.get('preserve_text',False)) is bool and (not item.get('preserve_text') or period_locations),
                         '仅更新期间时才能保留原表头文字')
                destination,origin=item['target'],item['source']
                _require(destination not in targets and origin not in sources and destination not in occupied,'维度文字不能覆盖数据格或重复安排')
                sheet,cell=destination.rsplit('!',1);source_sheet,source_cell=origin.rsplit('!',1)
                _require(sheet in a and source_sheet in b,'维度文字工作表不存在')
                period_targets=item.get('period_targets',[])
                _require(destination in label_targets or period_targets,'维度文字没有对应的财务格或合法表头：'+destination)
                _require(isinstance(period_targets,list) and all(isinstance(ref,str) for ref in period_targets)
                         and len(period_targets)==len(set(period_targets)) and (not period_targets or period_locations),
                         '新期间表头须明确所应用的数据格及B实际期间')
                if period_targets:
                    from openpyxl.utils.cell import coordinate_to_tuple,range_boundaries
                    from .表格 import _definitions,_destinations
                    header_row,header_col=coordinate_to_tuple(cell)
                    for reference in period_targets:
                        _require(reference in targets,'期间表头引用了未识别的数据格')
                        data_sheet,address=reference.rsplit('!',1);data_row,data_col=coordinate_to_tuple(address)
                        _require(data_sheet==sheet and data_col==header_col and data_row>header_row,
                                 '新期间须写在对应数据列的上方表头')
                        _require(any(sh==sheet and left<=header_col<=right and top<=header_row<data_row<=bottom
                            for definition in _definitions(a) for sh,area in _destinations(definition)
                            for left,top,right,bottom in [range_boundaries(area)]),'期间表头与数据须位于同一实际表格')
                    copied_period_targets[destination]=period_targets
                incoming=b[source_sheet][source_cell]
                text_value=_plain(incoming.value)
                reference_chain=[]
                if incoming.data_type=='f':
                    from .同义复核 import _referenced_text
                    referenced=_referenced_text(b,source_sheet,source_cell)
                    text_value=referenced['text'];reference_chain=referenced['reference_chain']
                aliases=item.get('source_aliases',[])
                _require(isinstance(aliases,list) and all(isinstance(ref,str) for ref in aliases)
                         and len(aliases)==len(set(aliases)) and origin not in aliases,'重复文字依据须为不同原文位置')
                for alias in aliases:
                    _require(alias not in sources and any(alias in refs for ref in bound.values()
                        for refs in sources[ref]['dimension_evidence'].values()),'重复文字依据须来自已绑定B格的识别原文')
                    alias_sheet,alias_cell=alias.rsplit('!',1)
                    _require(alias_sheet in b,'重复文字依据工作表不存在')
                    aliased=b[alias_sheet][alias_cell];alias_value=_plain(aliased.value)
                    if aliased.data_type=='f':
                        from .同义复核 import _referenced_text
                        alias_value=_referenced_text(b,alias_sheet,alias_cell)['text']
                    _require(alias_value==(_plain(a[sheet][cell].value) if item.get('preserve_text') else text_value),
                             '同一A名称的各处B原文不一致')
                _require(('selection_source' in item)==('selection_dimension' in item),'范围说明须同时指定B来源及选择维度')
                if 'selection_source' in item:
                    reference,key=item['selection_source'],item['selection_dimension']
                    _require(reference in bound.values() and key in sources[reference]['dimensions'], '范围说明须来自本次已绑定的B格')
                    selected=sources[reference]['dimensions'][key]
                    previous=[targets[loc]['dimensions'][key] for loc in bound
                              if destination in targets[loc]['dimension_evidence'].get(key,[])]
                    _require(standard['dimension'][key]['value_type']=='selection' and selected['mode'] in {'complement','predicate'}
                             and selected['completeness']=='complete' and origin in sources[reference]['dimension_evidence'][key],
                             '范围说明须有B原文及明确排除范围')
                    _require(previous and all(value==previous[0] and value['completeness']=='complete'
                             and value['domain']==selected['domain'] for value in previous),
                             '同一名称的原范围不一致，须分别说明')
                    notes=[]
                    if selected['mode']=='predicate':
                        if selected!=previous[0]:notes.append('范围：'+selected['predicate'])
                    else:
                        _require(previous[0]['mode']=='complement' and previous[0]['base']==selected['base'],
                                 '同一名称的原范围不一致，须分别说明')
                        included=[value for value in previous[0]['excluded'] if value not in selected['excluded']]
                        excluded=[value for value in selected['excluded'] if value not in previous[0]['excluded']]
                        if included:notes.append('范围包含：'+'、'.join(included))
                        if excluded:notes.append('范围不含：'+'、'.join(excluded))
                    if notes:text_value=str(text_value)+'（'+'；'.join(notes)+'）'
                if period_locations:
                    periods=[]
                    for reference in period_locations:
                        period_fact=sources.get(reference)
                        _require(period_fact is not None and origin in period_fact['dimension_evidence'].get('period',[]),
                                 '期间文字须引用本格已识别的实际期间及表头依据')
                        period=period_fact['dimensions']['period']
                        if period not in periods:periods.append(period)
                    copied_periods[destination]=periods
                    from datetime import date
                    def date_text(value):
                        d=date.fromisoformat(value);return f'{d.year}年{d.month}月{d.day}日'
                    spans=[p for p in periods if p['kind']=='duration']
                    # 同一年度列可同时列期初、变动和期末；只合并显示，逐格期间仍分别校验。
                    shown=spans if len(spans)==1 and all(p['kind']=='duration' or spans[0]['start']<=p['date']<=spans[0]['end'] for p in periods) else periods
                    dates='；'.join(date_text(p['date']) if p['kind']=='instant' else date_text(p['start'])+'至'+date_text(p['end']) for p in shown)
                    if item.get('preserve_text'):text_value=_plain(a[sheet][cell].value)
                    text_value=dates+('\n'+str(text_value) if text_value is not None else '')
                labels.append({'sheet':sheet,'cell':cell,'value':text_value,'old_value':_plain(a[sheet][cell].value),
                               'source_sheet':source_sheet,'source_cell':source_cell,
                               **({'source_text_reference_chain':reference_chain} if reference_chain else {}),
                               **({'source_aliases':aliases} if aliases else {})})
                copied_text.extend((destination,ref) for ref in [origin,*aliases]);occupied.add(destination)
            missing=purpose.get('unprovided',[])
            _require(isinstance(missing,list),'未提供字段须逐格列明')
            for item in labels:a[item['sheet']][item['cell']]=item['value']
            for item in missing:
                _require({'target','reason','source_evidence'}<=set(item)<={'target','reason','source_evidence','dimensions','dimension_evidence','action'},
                         '未提供字段须有目标、原文范围及核查原因')
                action=item.get('action','retain')
                _require(action in ('retain','mark_missing'),'未提供字段只能保留或明确标注B未提供')
                destination=item['target']
                _require(destination in targets and destination not in bound and destination not in unprovided,'未提供字段与取数绑定重复或未识别')
                _require(isinstance(item['reason'],str) and item['reason'].strip(),'未提供字段缺少原文核查结论')
                from .统一结构 import _text_evidence
                _text_evidence(b,item['source_evidence'])
                old=targets[destination]
                changed_text=any(label['sheet']+'!'+label['cell'] in refs and not same_excel_value(label['old_value'],label['value'])
                                 for refs in old['dimension_evidence'].values() for label in labels)
                supplied='dimensions' in item and 'dimension_evidence' in item
                _require(('dimensions' in item)==('dimension_evidence' in item) and (not changed_text or supplied),
                         '对应文字改变后须重新确认未提供字段的语义，不能沿用旧对象')
                dimensions=deepcopy(item.get('dimensions',old['dimensions']))
                evidence=deepcopy(item.get('dimension_evidence',old['dimension_evidence']))
                sh,address=destination.rsplit('!',1)
                from .统一识别 import validate_response
                packet={'source_path':mapping['source_path'],'source_hash':mapping['source_hash'],'sheet':sh,
                        'candidate_cells':[address],'document_context':view.get('source_context',[]),
                        'cells':{cell.coordinate:{'value':_plain(cell.value)} for row in a[sh] for cell in row}}
                response={'kind':'metric_fact','metric_id':old['metric_id'],'dimensions':dimensions,
                          'dimension_evidence':{key:[ref.removeprefix(sh+'!') for ref in refs] for key,refs in evidence.items()},'reason':item['reason']}
                checked=validate_response({'cells':{address:response}},packet,standard,context_workbook=a)[address]['fact']
                _require(all(dimensions.get('period') in copied_periods[header] for header,refs in copied_period_targets.items() if destination in refs),
                         '未提供字段的期间与本次表头不一致')
                key=semantic_key(checked,standard)
                for candidate in sources.values():
                    try:same=semantic_key(candidate,standard)==key
                    except ValueError:continue
                    _require(not same,'B已有相同语义来源，不能登记为未提供：'+destination)
                # 没有B值不能把旧值改属新对象，也不能以维度变化为由清空。
                _require(action=='mark_missing' or dimensions==old['dimensions'] or old['raw_value_state']=='blank',
                         'B未提供本字段，不能清空旧值或将旧值改属新维度；须补齐来源或保留原行及其语义：'+destination)
                unprovided[destination]={'action':action,'reason':item['reason'],'source_evidence':deepcopy(item['source_evidence'])}
                old['dimensions']=dimensions;old['dimension_evidence']=checked['dimension_evidence']
                old['owner_reference']=old['id'];view['owners'][old['owner_reference']]=deepcopy(dimensions)
            unhandled_labels=[]
            for destination,row in targets.items():
                if destination in bound or destination in unprovided:continue
                if (any(label['sheet']+'!'+label['cell'] in refs and not same_excel_value(label['old_value'],label['value'])
                        for refs in row['dimension_evidence'].values() for label in labels)
                        or any(destination in refs for refs in copied_period_targets.values())):
                    unhandled_labels.append(destination)
            _require(not unhandled_labels,'名称或期间改变后，所管辖的数据格须重新绑定或核实未提供：'+'；'.join(unhandled_labels))
            unlinked_dimensions=[]
            for destination,origin in bound.items():
                old,new=targets[destination],sources[origin]
                dimensions=deepcopy(new['dimensions'])
                for key in ('entity','report_scope'):
                    if key not in dimensions and key in old['dimensions']:
                        dimensions[key]=deepcopy(old['dimensions'][key])
                period_labels=[(left,right) for left,right in copied_text
                    if left in copied_periods and (left in old['dimension_evidence'].get('period',[]) or destination in copied_period_targets.get(left,[]))]
                linked_period_labels=[(left,right) for left,right in period_labels
                    if right in new['dimension_evidence'].get('period',[])]
                _require(not period_labels or linked_period_labels,'已更新期间文字未连接本格B期间依据：'+destination)
                # 证据可能同时引用期初、期末来解释变动期间，只核验双方实际对应的期间标签。
                for left,right in linked_period_labels:
                    _require(new['dimensions'].get('period') in copied_periods[left],'同一表头的期间与所绑定B数据格不一致：'+destination)
                for key,value in new['dimensions'].items():
                    if old['dimensions'].get(key)!=value:
                        linked=any(left in old['dimension_evidence'].get(key,[]) and right in new['dimension_evidence'][key]
                                   for left,right in copied_text)
                        if key=='period' and not linked:
                            linked=any(destination in copied_period_targets.get(left,[]) and right in new['dimension_evidence'][key] for left,right in copied_text)
                        if not linked and standard['dimension'][key]['value_type']=='selection':
                            # B分块列示同一类别时，A共用名称可引用任一已绑定且维度值相同的来源。
                            linked=any(left in old['dimension_evidence'].get(key,[]) and
                                any(sources[ref]['dimensions'].get(key)==value and right in sources[ref]['dimension_evidence'].get(key,[])
                                    for ref in bound.values()) for left,right in copied_text)
                        if key in {'entity','report_scope'} and not linked:
                            if key in old['dimensions']:dimensions[key]=deepcopy(old['dimensions'][key])
                            else:dimensions.pop(key,None)
                            continue
                        empty_member=standard['dimension'][key]['value_type']=='selection' and value.get('completeness')!='complete'
                        if empty_member:
                            from openpyxl.utils.cell import coordinate_to_tuple
                            sh,addr=origin.rsplit('!',1)
                            linked=linked or any(left in old['dimension_evidence'].get(key,[]) and
                                right.rsplit('!',1)[0]==sh and coordinate_to_tuple(right.rsplit('!',1)[1])[0]==coordinate_to_tuple(addr)[0] and
                                b[sh][right.rsplit('!',1)[1]].value is None for left,right in copied_text)
                        if not linked:unlinked_dimensions.append(destination+' / '+key)
                old['dimensions']=dimensions
                old['owner_reference']=old['id'];view['owners'][old['owner_reference']]=deepcopy(old['dimensions'])
            _require(not unlinked_dimensions,'维度变化须同步更新对应的文字：'+'；'.join(unlinked_dimensions))
            if view.get('_template_layout'):
                original={r['source_reference']['location'] for r in mapping['facts']}
                for destination,item in view['_template_layout']['added'].items():
                    _require(item['template'] not in original or destination in bound or destination in unprovided,'新增财务格须绑定本次B来源或核实未提供')
                    if item['template'] not in original:
                        sheet,cell=destination.rsplit('!',1)
                        _require(destination in occupied or a[sheet][cell].value is None,'新增文字须来自B')
        finally:a.close();b.close()
        existing={r['id'] for r in view['facts']}
        view['facts'].extend(targets[location] for location in bound if location not in existing)
        view['slot_bindings']=[r for r in view.get('slot_bindings',[]) if r['sheet']+'!'+r['cell'] not in bound]
        view['_purpose_sources']=bound;view['_purpose_labels']=labels
        view['_purpose_caption_updates']=[{'range_name':item['range_name'],'old_text':label['old_value'],'text':label['value']}
            for label in labels if not same_excel_value(label['value'],label['old_value']) for item in caption_contexts
            if item['sheet']==label['sheet'] and item.get('title')==label['old_value']
            and label['cell']==f"{get_column_letter(item['first_column'])}{item['first_row']-1}"]
        view['_effective_update_purpose']=purpose
        if unprovided:view['_purpose_unprovided']=unprovided
        view['_purpose_source_facts']={location:sources[location] for location in set(bound.values())}
        return view
    _require(set(purpose)=={'entity','reason'} and isinstance(purpose['reason'],str) and purpose['reason'].strip(), '本次更新用途须有明确依据')
    entity=purpose['entity']
    _require(isinstance(entity,dict) and set(entity)=={'from','to'} and all(isinstance(v,str) and v.strip() for v in entity.values()), '本次用途须明确原主体及目标主体')
    view=deepcopy(mapping)
    for row in view['facts']:
        if row['dimensions'].get('entity')==entity['from']:
            row['dimensions']['entity']=entity['to']
            if row.get('owner_reference'):view['owners'][row['owner_reference']]['entity']=entity['to']
    return view


def build_fact_update_plan(target_mapping, source_mapping, standard, equivalence_proofs=(), update_purpose=None):
    """输入为已识别事实。文件级执行者仍须校验原件和映射版本后再写入。"""
    for mapping in (target_mapping, source_mapping):
        _check_mapping(mapping, standard)
        _require(not mapping.get('source_identities'), '识别结果只有文字或最近标准，须先准确对应金标准再比较')
    original_dimensions={row['id']:deepcopy(row['dimensions']) for row in _template_cell_records(target_mapping).values()}
    recognition_target=target_mapping
    target_mapping=target_for_purpose(target_mapping,update_purpose,source_mapping,standard)
    if target_mapping.get('_template_layout'):
        from .模板布局 import relocate_mapping
        before=relocate_mapping(recognition_target,target_mapping['_template_layout'])
        original_dimensions={row['id']:deepcopy(row['dimensions']) for row in _template_cell_records(before).values()}
    from .同义复核 import approved_mapping_pairs
    approved=approved_mapping_pairs(equivalence_proofs,target_mapping,source_mapping,standard)
    plan = {'schema_version': 2, 'gold_hash': standard['sha256'],
            'source_a_hash': target_mapping['source_hash'], 'source_b_hash': source_mapping['source_hash'],
            'updates': [], 'issues': [], 'unplaced_sources': [], 'retained_blanks': [], 'operations': [],
            'labels':deepcopy(target_mapping.get('_purpose_labels',[]))}
    plan['caption_updates']=deepcopy(target_mapping.get('_purpose_caption_updates',[]))
    if update_purpose:plan['update_purpose']=deepcopy(target_mapping.get('_effective_update_purpose',update_purpose))
    if target_mapping.get('_template_layout'):
        plan['template_layout']=deepcopy(target_mapping['_template_layout'])
        plan['operations']=deepcopy(plan['template_layout']['operations'])
    candidates = defaultdict(list); consumed = set()
    directed=defaultdict(set)
    positions=target_mapping.get('_template_layout',{}).get('positions',{})
    for detail in source_mapping.get('recognition_scope',{}).get('details',[]):
        for need in detail.get('targets',[]):
            destination=positions.get(need['id'],need['id'])
            if destination is not None:directed[destination].update(need['candidates'])
    source_facts = [row for row in source_mapping['facts'] if row['record_type'] == 'metric_fact']
    for row in source_facts:
        try:
            candidates[semantic_key(row, standard)].append(_effective_source(row,source_mapping,standard))
        except ValueError as error:
            if row['source_reference']['location'] in target_mapping.get('_purpose_source_facts',{}):continue
            sheet, cell = _location(row)
            plan['issues'].append({'kind': 'unresolved_source', 'sheet': sheet, 'cell': cell, 'reason': str(error)})
    known_source_ids={r['id'] for r in source_facts}
    source_facts.extend(r for r in target_mapping.get('_purpose_source_facts',{}).values() if r['id'] not in known_source_ids)
    for mapping, carrier in ((target_mapping, '附注表格'), (source_mapping, '更新数据')):
        for row in mapping.get('unresolved', []):
            plan['issues'].append({'kind': 'unresolved_mapping', 'carrier': carrier, 'record': deepcopy(row)})
    for row in target_mapping.get('slot_bindings',[]):
        plan['issues'].append({'kind':'target_dimensions_unresolved','sheet':row['sheet'],'cell':row['cell'],
                               'reason':'槽位已确认，本格取数所需的具体维度尚未明确'})
    for target in target_mapping['facts']:
        if target['record_type'] != 'metric_fact':
            continue
        sheet, cell = _location(target)
        issue = {'sheet': sheet, 'cell': cell, 'metric_id': target['metric_id']}
        missing=target_mapping.get('_purpose_unprovided',{}).get(sheet+'!'+cell)
        if missing:
            item={**issue,**deepcopy(missing),'old_value':target['raw_value'],'dimensions':deepcopy(target['dimensions'])}
            plan.setdefault('unprovided_fields',[]).append(item)
            if missing['action']=='mark_missing':
                # 经原文确认的缺项只写明状态，不能冒充B取数或沿用另一对象的旧金额。
                plan['updates'].append({**issue,'value':'B未提供','old_value':target['raw_value'],
                    'dimensions':deepcopy(target['dimensions']),'sources':[],'blank_sources':[],
                    'unprovided_reason':missing['reason']})
            continue
        selected=target_mapping.get('_purpose_sources',{}).get(sheet+'!'+cell)
        try:
            matches = ([_effective_source(target_mapping['_purpose_source_facts'][selected],source_mapping,standard)]
                       if selected else candidates.get(semantic_key(target, standard), []))
            matches=list(matches)
            matches.extend(_effective_source(row,source_mapping,standard) for row in source_facts
                           if (target['id'],row['id']) in approved and row['id'] not in {item['id'] for item in matches})
            if selected:matches=[row for row in matches if row['source_reference']['location']==selected]
            preferred=[row for row in matches if row['source_reference']['location'] in directed[sheet+'!'+cell]]
            if preferred:matches=preferred
        except ValueError as error:
            plan['issues'].append({**issue, 'kind': 'unresolved_target', 'reason': str(error)})
            continue
        # 先保存身份证对应；是否能取数是后续处理，不能隐藏已有对应关系。
        plan.setdefault('identity_matches',[]).append({**issue,'dimensions':deepcopy(target['dimensions']),
            'source_locations':[row['source_reference']['location'] for row in matches]})
        explicit = [row for row in matches if row['raw_value_state'] == 'value']
        if any(row['raw_value_state'] == 'formula' for row in matches):
            plan['issues'].append({**issue, 'kind': 'unevaluated_source_formula', 'reason': '来源含尚未取得实际计算值的公式，须先核验计算值'})
            continue
        if not explicit:
            partial=[row for row in source_mapping.get('slot_bindings',[]) if row['partial_semantics']['metric_id']==target['metric_id']]
            if not matches and partial:
                plan['issues'].append({**issue,'kind':'source_dimensions_unresolved',
                    'reason':'同槽位来源已保留，具体成员等维度尚未明确，未作为该目标的确定取数来源',
                    'source_locations':[row['sheet']+'!'+row['cell'] for row in partial]})
                continue
            if target['raw_value_state'] == 'blank':
                plan['retained_blanks'].append({**issue, 'reason': '来源没有明确新值，保留原空白，不计作已更新'})
                consumed.update(row['id'] for row in matches)
            else:
                plan['issues'].append({**issue, 'kind': 'source_unknown_blank' if matches else 'missing_source',
                                       'reason': '来源空白不能清空原金额' if matches else '没有相同指标和实际维度的来源'})
            continue
        def converted(row):
            if selected:return deepcopy(row['raw_value'])
            if (target['id'],row['id']) not in approved:return value_for_target(row,target,standard)
            projected=deepcopy(row);projected['dimensions']=deepcopy(target['dimensions'])
            if 'unit_scale' in row['dimensions']:projected['dimensions']['unit_scale']=row['dimensions']['unit_scale']
            return value_for_target(projected,target,standard)
        try:
            values = [converted(row) for row in explicit]
        except ValueError as error:
            plan['issues'].append({**issue, 'kind': 'source_value_needs_review', 'reason': str(error),
                                   'source_locations': [row['source_reference']['location'] for row in explicit]})
            consumed.update(row['id'] for row in matches)
            continue
        if any(value != values[0] for value in values[1:]):
            plan['issues'].append({**issue, 'kind': 'conflicting_sources', 'reason': '相同业务身份的多个来源金额冲突，未选取或求和',
                                   'source_locations': [row['source_reference']['location'] for row in explicit]})
            continue
        sources = []
        for row in explicit:
            source_sheet, address = _location(row)
            source={'sheet': source_sheet, 'cell': address, 'value': row['raw_value'], 'fact_id': row['id']}
            original=next(item for item in source_facts if item['id']==row['id'])
            if original['raw_value_state']=='formula':source['formula']=original['raw_value']
            if (target['id'],row['id']) in approved:source['equivalence_pair']=approved[(target['id'],row['id'])]
            sources.append(source)
        plan['updates'].append({**issue, 'value': _plain(values[0]), 'old_value': target['raw_value'],
                                'dimensions': deepcopy(target['dimensions']), 'sources': sources,
                                **({'original_dimensions':original_dimensions[target['id']]} if update_purpose else {}),
                                'source_blank_locations': [row['source_reference']['location'] for row in matches if row not in explicit]})
        consumed.update(row['id'] for row in matches)
    if source_mapping.get('recognition_scope',{}).get('mode')=='b_required':
        # 按需B已有用途及候选去向；未采用候选保留事实，不启动整B来源安置循环。
        from .表格编排 import source_ready
        local_ready=all(m.get('scope_coverage',{}).get('complete') for m in (source_mapping,target_mapping))
        if not local_ready and not source_ready(source_mapping,recognition_target,standard):
            plan['issues'].append({'kind':'required_scope_incomplete','reason':'B初识或当前A所需内容尚未确认完整'})
        for target in target_mapping['facts']:
            if target['record_type']!='metric_fact':continue
            if target['id'] in target_mapping.get('_purpose_sources',{}):continue
            if target['id'] in target_mapping.get('_purpose_unprovided',{}):continue
            matches=candidates.get(semantic_key(target,standard),[])
            if not matches and not any(left==target['id'] for left,right in approved):
                sheet,cell=_location(target)
                if not any(x.get('sheet')==sheet and x.get('cell')==cell for x in plan['issues']):
                    plan['issues'].append({'kind':'missing_source','sheet':sheet,'cell':cell,'reason':'本次A需要没有已确认来源，空白目标也不能漏记'})
        plan['unplaced_sources']=[]
    else:
        plan['unplaced_sources'] = [deepcopy(row) for row in source_facts if row['id'] not in consumed]
    return plan
