# coding: utf-8
"""只改数据值的副本沿用原语义；原映射和非数据内容仍逐项验证。"""
from copy import deepcopy
from pathlib import Path
from openpyxl import load_workbook
from .表格 import file_hash, same_excel_value
from .统一语义 import _require


def derive_value_copy(parent, source_path):
    path=Path(source_path).resolve()
    _require(file_hash(parent['mapping_path'])==parent['mapping_hash'], '数值副本的原映射变化')
    _require(file_hash(parent['source_path'])==parent['source_hash'], '数值副本的原工作簿变化')
    _require(path!=Path(parent['source_path']).resolve(), '数值副本不能覆盖原件')
    positions={tuple(f['source_reference']['location'].rsplit('!',1)) for f in parent['facts'] if f['record_type']=='metric_fact'}
    positions.update((r['sheet'],r['cell']) for r in parent.get('slot_bindings',[]))
    before=load_workbook(parent['source_path'],data_only=False)
    after=load_workbook(path,data_only=False)
    result=deepcopy(parent);digest=file_hash(path)
    try:
        _require(before.sheetnames==after.sheetnames,'数值副本的工作表结构变化')
        for left,right in zip(before.worksheets,after.worksheets):
            _require(str(left.merged_cells)==str(right.merged_cells) and left.sheet_state==right.sheet_state,
                     '数值副本的合并或可见性变化')
            for point in set(left._cells)|set(right._cells):
                old=left.cell(*point);new=right.cell(*point)
                _require(old._style==new._style,'数值副本的格式变化：'+left.title+'!'+old.coordinate)
                if same_excel_value(new.value,old.value):continue
                _require((left.title,old.coordinate) in positions,'数据格以外发生变化：'+left.title+'!'+old.coordinate)
                _require(new.value is None or isinstance(new.value,(int,float)) and not isinstance(new.value,bool),
                         '数值副本只接受明确数值或空白：'+left.title+'!'+old.coordinate)
        for fact in result['facts']:
            fact['source_reference']['file_hash']=digest
            if fact['record_type']!='metric_fact':continue
            sheet,cell=fact['source_reference']['location'].rsplit('!',1)
            item=after[sheet][cell];value=item.value
            fact['raw_value']=value
            fact['raw_value_state']='blank' if value is None or isinstance(value,str) and not value.strip() else ('formula' if item.data_type=='f' else 'value')
    finally:
        before.close();after.close()
    for key in ('mapping_path','mapping_hash','_validated_standard_hash','calculation_evidence','formula_values','calculation_issues','review_path','a_link'):
        result.pop(key,None)
    result.update(source_path=str(path),source_hash=digest,original_path=str(path),original_hash=digest,
                  value_derivation={'parent_path':parent['mapping_path'],'parent_hash':parent['mapping_hash']})
    _require(file_hash(path)==digest,'核验期间数值副本发生变化')
    return result


def validate_value_copy(path, mapping, standard, ancestors):
    from .统一勾稽 import load_fact_mapping, _check_mapping
    proof=mapping['value_derivation']
    _require(file_hash(proof['parent_path'])==proof['parent_hash'],'数值副本的原映射变化')
    parent=load_fact_mapping(proof['parent_path'],standard,ancestors)
    expected=derive_value_copy(parent,mapping['source_path'])
    _require(mapping==expected,'数值副本身份或原值与已核验的原映射及副本不符')
    expected['_validated_standard_hash']=standard['sha256'];_check_mapping(expected,standard)
    expected.update(mapping_path=str(path),mapping_hash=file_hash(path))
    return expected
