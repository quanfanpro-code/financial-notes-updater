# coding: utf-8
"""合成原文和模拟原模型验证本格维度标准，不发网络请求。"""
from copy import deepcopy
import json
from pathlib import Path
import unittest
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import range_boundaries, get_column_letter

from 附注更新.A合集 import _records, save_direct_answer, validate_direct_collection, direct_digest
from 附注更新.A语义阅读 import direct_reference
from 附注更新.分步流程 import run_step
from 附注更新.统一识别 import UnifiedSemanticEngine
from 附注更新.统一勾稽 import load_fact_mapping
from 附注更新.表格 import file_hash
from 测试 import 测试直接识别 as fixture


class 维度格身份证入口测试(unittest.TestCase):
    def setUp(self):
        self.case=fixture.直接识别测试();self.case.setUp()

    def test_从正常范围入口按自身身份证搬运名称期间和数据到两个副本(self):
        word=self.case.folder/'维度格原附注.docx';doc=Document()
        doc.add_heading('六、财务报表重要项目的说明',level=1)
        table=doc.add_table(rows=2,cols=3)
        for row,values in zip(table.rows,[['项目','原期末','原期初'],['原名称','10','9']]):
            for cell,value in zip(row.cells,values):cell.text=value
        doc.save(word)
        config={**self.case.config,'word_path':str(word)}
        for side in ('a','b'):config.pop(side+'_target_table',None)
        extracted=run_step('extract_a',config)
        self.assertEqual(extracted['status'],'complete',extracted);config.update(extracted['files'])
        book=load_workbook(config['a_path']);sheet,area=next(next(iter(book.defined_names.values())).destinations)
        l,t,r,b=range_boundaries(area);cols=[get_column_letter(l+i) for i in range(3)];book.close()
        acells=[cols[1]+str(t),cols[2]+str(t),cols[0]+str(t+1),cols[1]+str(t+1),cols[2]+str(t+1)]
        bcells=['F2','G9','H3','K5','L8'];bpath=self.case.folder/'维度格B原表.xlsx'
        book=Workbook();ws=book.active;ws.title='来源'
        for cell,value in zip(bcells,['新期末','新期初','新名称',20,11]):ws[cell]=value
        book.save(bpath);book.close();config['b_path']=str(bpath)
        before={path:file_hash(path) for path in (word,Path(config['a_path']),bpath,self.case.gold)}
        selection={'mode':'members','domain':'固定资产','members':['同一设备'],'completeness':'complete'}
        def answer(name,cells,year):
            result=[]
            for index,cell in enumerate(cells):
                dims={'entity':'主体'+str(year),'report_scope':'standalone','unit_scale':1,
                      'asset_selection':selection,'currency':'CNY'}
                value={'location':name+'!'+cell,'dimensions':dims}
                if index in (0,1,3,4):
                    role='closing' if index in (0,3) else 'opening'
                    dims['period']={'kind':'instant','date':str(year if role=='closing' else year-1)+'-12-31'}
                    value['template_period_role']=role
                value['dimension_id' if index<3 else 'metric_id']='period' if index<2 else 'asset_selection' if index==2 else '固定资产原值'
                result.append(value)
            return json.dumps(result,ensure_ascii=False)
        ascope=f'|工作表|完整范围|数据范围|相关标准|\n|---|---|---|---|\n|{sheet}|{area}|{",".join(acells)}|固定资产原值, period, asset_selection|'
        bscope=f'|工作表|完整范围|数据范围|相关标准|A范围|\n|---|---|---|---|---|\n|来源|F2:L9|{",".join(bcells)}|固定资产原值, period, asset_selection|{sheet}!{area}|'
        replies=[ascope,answer(sheet,acells,2024),bscope,answer('来源',bcells,2025)]
        calls=[]
        def respond(engine,prompt,materials):
            calls.append((prompt,materials))
            return replies.pop(0)
        config['comparison_requirements']=[{'维度ID':'entity','要求含义':'不参与比较','出处':'本合成验证'}]
        with patch.object(UnifiedSemanticEngine,'request_text',respond):
            for step in ('recognize_a','recognize_b','match','write'):
                result=run_step(step,config)
                self.assertEqual(result['status'],'complete',result);config.update(result['files'])
        self.assertEqual(len(calls),4);self.assertFalse(replies)
        self.assertIn('维度标准目录.json',calls[0][1])
        for call in (calls[1],calls[3]):
            definitions=[json.loads(line) for line in call[1]['金标准.jsonl'].splitlines()]
            self.assertTrue(any(row['id']=='period' and 'meaning' in row for row in definitions))
            self.assertIn('完整工作表原文',call[1].get('原文.json',call[1].get('B原文.json',{})))
        mapping=load_fact_mapping(config['a_mapping'],{})
        validate_direct_collection(mapping,for_matching=True)
        rows=dict(_records(mapping))
        self.assertEqual(rows[sheet+'!'+acells[0]]['dimension_id'],'period')
        self.assertEqual(len(direct_reference(mapping,compact=True)['需要']),5)
        plan=json.loads(Path(config['update_plan']).read_text(encoding='utf-8-sig'))
        self.assertEqual({row['cell']:row['source_locations'][0] for row in plan['identity_matches']},dict(zip(acells,['来源!'+c for c in bcells])))
        self.assertEqual(plan['labels'],[])
        actual=Document(config['word'])
        self.assertEqual([actual.tables[0].cell(0,c).text for c in (1,2)],['新期末','新期初'])
        self.assertEqual(actual.tables[0].cell(1,0).text,'新名称')
        book=load_workbook(config['a_prime']);ws=book[sheet]
        self.assertEqual([ws[cell].value for cell in acells],['新期末','新期初','新名称',20,11]);book.close()
        self.assertEqual({path:file_hash(path) for path in before},before)
        # 原格身份证只随真实位置链接续，不能由写入的文字或B身份重建。
        from 附注更新.A合集 import resume_direct_output, _standard
        from 附注更新.表格 import read_workbook
        resumed,_,_=resume_direct_output(config['a_prime_mapping'],read_workbook(config['a_prime'],context=mapping['source_context']),mapping['source_context'],_standard(mapping))
        validate_direct_collection(resumed,for_matching=True)
        self.assertEqual(dict(_records(resumed))[sheet+'!'+acells[0]]['dimensions'],rows[sheet+'!'+acells[0]]['dimensions'])

    def test_定向字段不外溢共同归属不补身份重复不选来源(self):
        from 附注更新.模板更新编排 import direct_update_plan
        config={**self.case.config}
        answer=[{'location':'合成表!B3','dimension_id':'period','dimensions':{'entity':'甲','period':None}},
                {'location':'合成表!C3','dimension_id':'period','dimensions':{'entity':'乙'}}]
        with patch.object(UnifiedSemanticEngine,'request_text',return_value=json.dumps(answer,ensure_ascii=False)):
            result=run_step('recognize_a',config)
        self.assertEqual(result['status'],'complete',result)
        old=load_fact_mapping(result['files']['a_mapping'],{})
        reply=[{'location':'合成表!B3','dimension_id':'asset_selection','dimensions':{'entity':'越界'}},
               {'location':'合成表!C3','dimension_id':'asset_selection'}]
        saved=save_direct_answer(old,json.dumps(reply,ensure_ascii=False),self.case.folder/'定向维度编号.json',field_scope={'合成表!B3':['dimension_id']})
        before=dict(_records(old));after=dict(_records(saved))
        self.assertEqual(after['合成表!B3']['dimension_id'],'asset_selection')
        self.assertEqual(after['合成表!B3']['dimensions'],before['合成表!B3']['dimensions'])
        self.assertEqual({k:v for k,v in after['合成表!C3'].items() if k!='recognition_provenance'},
                         {k:v for k,v in before['合成表!C3'].items() if k!='recognition_provenance'})
        b=deepcopy(old);b['carrier']='B';b['封存校验']=direct_digest(b)
        a=deepcopy(old);a['facts']=[r for r in a['facts'] if r['source_reference']['location']=='合成表!B3'];a['封存校验']=direct_digest(a)
        # 排除主体后两个来源身份相同，不按值、格址或声明顺序择一。
        plan=direct_update_plan(a,b,comparison_requirements=[{'维度ID':'entity'}])
        self.assertFalse(plan['updates'])
        self.assertTrue(any(r.get('kind')=='ambiguous_source' for r in plan['issues']))
        from 附注更新.A合集 import _per_cell_answer,_standard
        for raw in ({'dimension_id':'period','cells':[{'location':'合成表!B3','owner_reference':'合成表!C3'}]},
                    '合成表!B3:C3：dimension_id=period；entity=甲。'):
            with self.assertRaises(ValueError):_per_cell_answer(raw,_standard(old),['合成表'])

    def test_副本接续和已定位来源的默认字段均保存本格维度编号(self):
        config={**self.case.config}
        reply=json.dumps([{'location':'合成表!B3','dimension_id':'period',
            'dimensions':{'period':{'kind':'instant','date':'2025-12-31'}},
            'template_period_role':'closing'}],ensure_ascii=False)
        with patch.object(UnifiedSemanticEngine,'request_text',return_value=reply):
            for step in ('recognize_a','recognize_b','match','write'):
                result=run_step(step,config)
                self.assertEqual(result['status'],'complete',result)
                config.update(result['files'])
        changed=json.dumps([{'location':'合成表!B3','dimension_id':'asset_selection',
            'dimensions':{'asset_selection':{'mode':'members','members':['原文设备']}}}],ensure_ascii=False)
        scopes=[('recognize_a',{'a_path':config['a_prime'],'a_mapping':config['a_prime_mapping']}),
                ('recognize_b',{'b_source_targets':['既有目标'],
                    'b_source_ranges':[{'sheet':'合成表','range':'B3'}]})]
        for step,extra in scopes:
            with self.subTest(step=step),patch.object(UnifiedSemanticEngine,'request_text',return_value=changed) as request:
                result=run_step(step,{**config,**extra})
                self.assertEqual(result['status'],'complete',result)
                request.assert_called_once()
                mapping=load_fact_mapping(result['files']['a_mapping' if step=='recognize_a' else 'b_mapping'],{})
                self.assertEqual(dict(_records(mapping))['合成表!B3']['dimension_id'],'asset_selection')


    def test_旧原答复经真实执行位置链接回不改身份和原文(self):
        from openpyxl.workbook.defined_name import DefinedName
        from 附注更新.A合集 import resume_direct_output, _standard
        from 附注更新.表格 import read_workbook
        source=self.case.folder/'位移原表.xlsx'
        book=load_workbook(self.case.source)
        book.defined_names.add(DefinedName('合成位置区',attr_text="'合成表'!$A$2:$C$8"))
        book.save(source);book.close()
        config={**self.case.config,'a_path':str(source)}
        reply=json.dumps({'合成表!B3':{'metric_id':'固定资产原值'}},ensure_ascii=False)
        with patch.object(UnifiedSemanticEngine,'request_text',return_value=reply):
            for step in ('recognize_a','recognize_b'):
                result=run_step(step,config)
                self.assertEqual(result['status'],'complete',result);config.update(result['files'])
        old=load_fact_mapping(config['a_mapping'],{})
        config['update_purpose']={'reason':'合成真实位置变化','operations':[{'range_name':'合成位置区','axis':'row','action':'insert','index':1,'count':1,'template_index':1}]}
        for step in ('match','write'):
            result=run_step(step,config)
            self.assertEqual(result['status'],'complete',result);config.update(result['files'])
        resumed,receipt,_=resume_direct_output(config['a_prime_mapping'],read_workbook(config['a_prime'],context=old['source_context']),old['source_context'],_standard(old))
        resumed['mapping_path']=str(self.case.folder/'当前位置合集.json');resumed['封存校验']=direct_digest(resumed)
        Path(resumed['mapping_path']).write_text(json.dumps(resumed,ensure_ascii=False),encoding='utf-8')
        self.assertEqual(receipt['template_layout']['positions']['合成表!B3'],'合成表!B4')
        late=json.dumps({'合成表!B3':{'currency':'CNY','asset_selection':{'mode':'members','members':['原文含合成表!B3，不改写']},'label_references':{'asset_selection':['合成表!A3']}}},ensure_ascii=False)
        chain=[{'path':config['a_prime_mapping'],'sha256':file_hash(config['a_prime_mapping'])}]
        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('位置接续不能调用模型')):
            saved=save_direct_answer(resumed,late,self.case.folder/'晚接字段合集.json',field_scope={'合成表!B4':['currency','asset_selection']},reply_source_hash=old['source_hash'],reply_position_chain=chain)
        validate_direct_collection(saved,for_matching=True)
        row=dict(_records(saved))['合成表!B4']
        self.assertEqual(row['metric_id'],'固定资产原值')
        self.assertEqual(row['dimensions'],{'currency':'CNY','asset_selection':{'mode':'members','members':['原文含合成表!B3，不改写']}})
        self.assertNotIn('合成表!B3',dict(_records(saved)))
        self.assertEqual(saved['本次逐格原答复'],late)
        self.assertEqual(saved['机械记录'],resumed['机械记录'])
        self.assertEqual(saved['答复接入来源']['本格原位置至现位置'],{'合成表!B3':'合成表!B4'})
        self.assertEqual(row['label_references']['asset_selection'],['合成表!A4'])
        wrong=deepcopy(receipt);wrong['template_layout']['positions']['合成表!B3']='合成表!B5'
        wrong_path=self.case.folder/'伪造位置记录.json';wrong_path.write_text(json.dumps(wrong,ensure_ascii=False),encoding='utf-8')
        checks=[(old['source_hash'],chain,{'合成表!B3':['currency']}),('0'*64,chain,{'合成表!B4':['currency']}),(old['source_hash'],[{'path':str(wrong_path),'sha256':file_hash(wrong_path)}],{'合成表!B4':['currency']}),(old['source_hash'],[{'path':chain[0]['path'],'sha256':'f'*64}],{'合成表!B4':['currency']})]
        for n,(sha,links,scope) in enumerate(checks):
            with self.subTest(n=n),self.assertRaises(ValueError):
                save_direct_answer(resumed,late,self.case.folder/('拒绝位置链'+str(n)+'.json'),field_scope=scope,reply_source_hash=sha,reply_position_chain=links)

if __name__=='__main__':unittest.main()
