# coding: utf-8
"""分层识别行为验证：只替代外部模型，实际保存、回读及原件校验均保留。"""
import copy
import json
import unittest
import uuid
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from 附注更新.统一识别 import UnifiedSemanticEngine
from 附注更新.表格 import file_hash
from 测试.测试统一语义 import records, fact


class 分层识别测试(unittest.TestCase):
    def setUp(self):
        self.folder = Path(__file__).resolve().parents[1] / '测试输出' / ('分层识别_' + uuid.uuid4().hex)
        self.folder.mkdir(parents=True)
        self.gold = self.folder / '测试标准.jsonl'
        self.gold.write_text('\n'.join(json.dumps(x, ensure_ascii=False) for x in records()), encoding='utf-8-sig')
        self.source = self.folder / '多个科目及子表.xlsx'
        book = Workbook(); sheet = book.active; sheet.title = '综合'
        for address, value in {'A1': '合成企业甲 单户 人民币元 2025年12月31日',
                               'A2': '固定资产原值合计', 'B2': '期末原值', 'C2': '累计折旧',
                               'A3': '全部固定资产', 'B3': 125, 'C3': 25,
                               'A6': '固定资产前五名', 'A7': '供应商甲', 'B7': 13,
                               'A10': '存货', 'A11': '材料', 'B11': 80}.items(): sheet[address] = value
        continuation = book.create_sheet('续页'); continuation['A1'] = '存货续表'; continuation['A2'] = '商品'; continuation['B2'] = 40
        book.save(self.source); book.close()
        self.requests = []
        self.maxDiff = 500

    def response(self, engine, system, payload):
        self.requests.append(copy.deepcopy(payload))
        if payload['task'] == 'workbook_outline':
            if len(payload['sheets']) == 1 and payload['sheets'][0]['max_row'] == 3:
                return {'subjects': ['固定资产'], 'subtables': [{'id': '原值', 'subject': '固定资产', 'purpose': '固定资产原值合计',
                        'ranges': [{'sheet': '综合', 'range': 'A1:B3'}], 'evidence': ['综合!A2']}],
                        'unlocated_ranges': [], 'context_ranges': []}
            return {'subjects': ['固定资产', '存货'], 'subtables': [
                {'id': '原值', 'subject': '固定资产', 'purpose': '固定资产原值合计',
                 'ranges': [{'sheet': '综合', 'range': 'A1:C3'}], 'evidence': ['综合!A2']},
                {'id': '前五名', 'subject': '固定资产', 'purpose': '固定资产前五名',
                 'ranges': [{'sheet': '综合', 'range': 'A6:B7'}], 'evidence': ['综合!A6']},
                {'id': '存货', 'subject': '存货', 'purpose': '存货及续表',
                 'ranges': [{'sheet': '综合', 'range': 'A10:B11'}, {'sheet': '续页', 'range': 'A1:B2'}],
                 'evidence': ['综合!A10', '续页!A1']}], 'unlocated_ranges': [], 'context_ranges': []}
        if payload['task'] == 'unified_classify':
            cells = {}
            for a in payload['candidate_cells']:
                if payload['sheet'] == '综合' and a == 'B3':
                    cells[a] = {'kind': 'metric_fact', 'metric_id': '固定资产原值', 'dimensions': fact()['dimensions'],
                                'dimension_evidence': {k: ['A1', 'A2', 'A3'] for k in fact()['dimensions']}, 'reason': '原文原值合计'}
                else:
                    evidence = next(row['cells'][0][0] for row in payload['source_text_rows'])
                    cells[a] = {'kind': 'non_business', 'category': 'label', 'reason': '合成例标签', 'evidence_cells': [evidence]}
            return {'cells': cells}
        if payload['task'] == 'select_subtables':
            return {'subtables': [{'id': t['id'], 'decision': 'needed' if t['id'] == '原值' else 'unneeded',
                                  'target_ids': [f['id'] for f in payload['targets']] if t['id'] == '原值' else [],
                                  'reason': '原文用途已明确', 'evidence': t['evidence']} for t in payload['outline']['subtables']],
                    'unresolved_targets': []}
        if payload['task'] == 'subtable_detail':
            return {'allowed_metrics': ['固定资产原值'] if payload['subtable']['id'] == '原值' else [],
                    'structure': '期末原值列与累计折旧列分别列示', 'evidence': payload['subtable']['evidence'],
                    'targets': [{'id': f['id'], 'candidates': ['综合!B3'], 'reason': '仅原值列满足需要'}
                                for f in payload.get('targets', [])]}
        raise AssertionError('未预设外部任务：' + payload['task'])

    def engine(self):
        return UnifiedSemanticEngine({'model': '本地替身', 'base_url': 'http://localhost', 'timeout': 1}, self.gold, self.folder / '回执')

    def target(self, value=100):
        path = self.folder / ('A仅需原值_' + str(value) + '.xlsx')
        book = Workbook(); book.active.title = '综合'
        for a, v in {'A1': '合成企业甲 单户 人民币元 2025年12月31日', 'A2': '固定资产原值合计',
                     'B2': '期末原值', 'A3': '全部固定资产', 'B3': value}.items(): book.active[a] = v
        book.save(path); book.close()
        return self.engine().recognize_workbook(path)

    def test_b_keeps_source_identities_through_selection_and_saved_scope(self):
        from 附注更新.表格编排 import validate_saved_scope, target_requirements
        from 附注更新.表格 import read_workbook, materialize_candidate_blanks
        with patch.object(UnifiedSemanticEngine, '_http', lambda e, s, p: self.response(e, s, p)):
            target = self.target()
            target['source_identities'] = [
                {'source_reference': {'file_hash': target['source_hash'], 'location': '综合!C3'},
                 'table_id': '说明', 'meaning': '原值期末金额，报表范围未说明', 'evidence_cells': ['A2'],
                 'closest_metric_id': '固定资产原值', 'standard_difference': '原文未说明报表范围'},
                {'source_reference': {'file_hash': target['source_hash'], 'location': '综合!D3'},
                 'table_id': '说明', 'meaning': '固定资产原值备注', 'evidence_cells': ['A2'],
                 'closest_metric_id': None, 'standard_difference': '无对应文字指标'}]
            target['unresolved'] = [{'sheet': '综合', 'cell': a, 'reason': '缺少对应标准'} for a in ('C3', 'D3')]
            self.requests.clear()
            result = self.engine().recognize_workbook(self.source, target_mapping=target)
        scope = result['recognition_scope']
        self.assertEqual({t['location'] for t in scope['targets']}, {'综合!B3', '综合!C3', '综合!D3'})
        self.assertEqual(len({t['id'] for t in scope['targets']}), 3)
        for source, need in zip(target['source_identities'], scope['targets'][1:]):
            for key in ('meaning', 'closest_metric_id', 'standard_difference', 'evidence_cells'):
                self.assertEqual(need[key], source[key])
            self.assertIsNone(need.get('semantic_key'))
        for payload in self.requests:
            if payload['task'] in ('select_subtables', 'subtable_detail'):
                self.assertEqual(payload['targets'], scope['targets'])
        self.assertEqual(scope['target_unresolved'], [])
        self.assertTrue(scope['coverage']['complete'])
        self.assertFalse(scope['coverage']['satisfied'], '最近指标相同不能直接认作精确来源')
        self.assertEqual([r['exact_sources'] for r in scope['coverage']['targets']], [['综合!B3'], [], []])
        reread = json.loads(Path(result['mapping_path']).read_text(encoding='utf-8-sig'))
        snapshot = read_workbook(self.source); materialize_candidate_blanks(snapshot)
        validate_saved_scope(reread, snapshot, self.engine().standard, target)
        changed = copy.deepcopy(target); changed['source_identities'][0]['meaning'] = '原值期初金额'
        self.assertEqual([r['id'] for r in target_requirements(changed, self.engine().standard)],
                         [r['id'] for r in scope['targets']])
        with self.assertRaisesRegex(ValueError, 'A需要已变化'):
            validate_saved_scope(reread, snapshot, self.engine().standard, changed)

    def test_b_source_only_a_preserves_real_unresolved(self):
        with patch.object(UnifiedSemanticEngine, '_http', lambda e, s, p: self.response(e, s, p)):
            target = self.target(); target['facts'] = []
            target['source_identities'] = [{'source_reference': {'file_hash': target['source_hash'], 'location': '综合!B3'},
                'meaning': '固定资产原值期末合计', 'closest_metric_id': '固定资产原值',
                'standard_difference': '原文缺少口径说明', 'evidence_cells': ['A2'], 'table_id': '原值'}]
            target['unresolved'] = [{'sheet': '综合', 'cell': 'B3', 'reason': '缺少对应标准'},
                                    {'sheet': '综合', 'cell': 'C3', 'reason': '此处原义尚未读懂'}]
            result = self.engine().recognize_workbook(self.source, target_mapping=target)
        scope = result['recognition_scope']
        self.assertEqual([r['location'] for r in scope['targets']], ['综合!B3'])
        self.assertEqual(scope['target_unresolved'], [{'sheet': '综合', 'cell': 'C3', 'reason': '此处原义尚未读懂'}])
        self.assertFalse(scope['coverage']['complete'])

    def test_outline_stops_at_subject_and_subtable(self):
        engine = self.engine()
        metadata = {'carrier': 'B', 'original_path': str(self.source), 'original_hash': file_hash(self.source)}
        with patch.object(UnifiedSemanticEngine, '_http', lambda e, s, p: self.response(e, s, p)):
            result = engine.recognize_workbook(self.source, artifact_metadata=metadata)
        fine = [p for p in self.requests if p['task'] == 'unified_classify']
        self.assertEqual(fine, [], '没有 A 时只能完整初识 B，不得逐格识别')
        self.assertEqual({x['id'] for x in result['outline']['subtables']}, {'原值', '前五名', '存货'})
        self.assertEqual(result['outline']['subtables'][2]['ranges'][1], {'sheet': '续页', 'range': 'A1:B2'})
        self.assertTrue(result['outline']['complete'])
        self.assertFalse(result['complete'])
        self.assertEqual(result['facts'], [])
        from 附注更新.统一勾稽 import load_fact_mapping
        self.assertEqual(load_fact_mapping(result['mapping_path'],engine.standard)['recognition_scope']['mode'],'b_outline')

    def test_a_uses_related_standard_and_scoped_context(self):
        with patch.object(UnifiedSemanticEngine, '_http', lambda e, s, p: self.response(e, s, p)):
            result = self.engine().recognize_workbook(self.source)
        fine = [p for p in self.requests if p['task'] == 'unified_classify']
        self.assertNotIn('A10', fine[0]['cells'], '原值子表不应重复携带存货的整页文字')
        self.assertEqual([x['source_reference']['location'] for x in result['facts']], ['综合!B3'])

    def test_b_recognizes_only_required_subtables(self):
        with patch.object(UnifiedSemanticEngine, '_http', lambda e, s, p: self.response(e, s, p)):
            target = self.target()
            self.requests.clear()
            metadata = {'carrier': 'B', 'original_path': str(self.source), 'original_hash': file_hash(self.source)}
            result = self.engine().recognize_workbook(self.source, target_mapping=target, artifact_metadata=metadata)
        self.assertEqual({r['source_reference']['location'] for r in result['facts']}, {'综合!B3'})
        self.assertNotIn('综合!C3', {r['source_reference']['location'] for r in result['facts']})
        self.assertEqual({p['subtable']['id'] for p in self.requests if p['task'] == 'subtable_detail'}, {'原值'})
        self.assertEqual({a for p in self.requests if p['task'] == 'unified_classify' for a in p['candidate_cells']}, {'B3'})

    def test_required_scope_survives_real_save_reload_and_match(self):
        from 附注更新.统一勾稽 import load_fact_mapping
        from 附注更新.分步流程 import run_step
        with patch.object(UnifiedSemanticEngine, '_http', lambda e, s, p: self.response(e, s, p)):
            target = self.target()
            metadata = {'carrier': 'B', 'original_path': str(self.source), 'original_hash': file_hash(self.source)}
            result = self.engine().recognize_workbook(self.source, target_mapping=target, artifact_metadata=metadata)
            reread = load_fact_mapping(result['mapping_path'], self.engine().standard)
            config = {'gold_path': str(self.gold), 'a_mapping': target['mapping_path'], 'b_mapping': reread['mapping_path'],
                      'output_dir': str(self.folder / '流程'), 'base_url': 'http://localhost', 'model': '本地替身', 'timeout': 1}
            matched = run_step('match', config)
        self.assertEqual(matched['status'], 'complete', matched['message'])
        plan = json.loads(Path(matched['files']['update_plan']).read_text(encoding='utf-8-sig'))
        self.assertEqual([(u['sheet'], u['cell'], u['value']) for u in plan['updates']], [('综合', 'B3', 125)])
        self.assertFalse(reread['coverage']['complete'], '不能谎称整B逐格识别完成')
        self.assertTrue(reread['recognition_scope']['coverage']['complete'])

    def test_normal_entry_passes_a_needs_to_b(self):
        from 附注更新.分步流程 import run_step
        with patch.object(UnifiedSemanticEngine, '_http', lambda e, s, p: self.response(e, s, p)):
            target = self.target()
            config = {'gold_path': str(self.gold), 'a_mapping': target['mapping_path'], 'b_path': str(self.source),
                      'output_dir': str(self.folder / '正式入口'), 'base_url': 'http://localhost', 'model': '本地替身', 'timeout': 1}
            output = run_step('recognize_b', config)
        self.assertEqual(output['status'], 'complete', output['message'])
        result = json.loads(Path(output['files']['b_mapping']).read_text(encoding='utf-8-sig'))
        self.assertEqual(result['recognition_scope']['mode'], 'b_required')
        self.assertEqual([f['source_reference']['location'] for f in result['facts']], ['综合!B3'])

    def test_changed_a_reuses_b_facts_but_invalidates_selection(self):
        from 附注更新.表格编排 import source_ready
        from 附注更新.统一勾稽 import load_fact_mapping
        metadata = {'carrier': 'B', 'original_path': str(self.source), 'original_hash': file_hash(self.source)}
        with patch.object(UnifiedSemanticEngine, '_http', lambda e, s, p: self.response(e, s, p)):
            first = self.target()
            previous = self.engine().recognize_workbook(self.source, target_mapping=first, artifact_metadata=metadata)
            changed = self.target(101)
            with self.assertRaisesRegex(ValueError, 'A需要已变化'):
                source_ready(previous, changed, self.engine().standard)
            self.requests.clear()
            result = self.engine().recognize_workbook(self.source, target_mapping=changed, artifact_metadata=metadata,
                                                      previous_result=previous['mapping_path'])
        self.assertEqual(result['facts'], previous['facts'])
        self.assertEqual([p for p in self.requests if p['task'] == 'unified_classify'], [], '已有合法B身份不能因A变动全部重跑')
        self.assertTrue(load_fact_mapping(result['mapping_path'], self.engine().standard)['recognition_scope']['coverage']['complete'])

    def test_outline_rejects_forged_evidence_and_keeps_unlocated(self):
        from 附注更新.表格编排 import validate_outline
        from 附注更新.表格 import read_workbook
        snapshot = read_workbook(self.source)
        raw = self.response(None, '', {'task':'workbook_outline', 'sheets':snapshot['sheets']})
        bad = copy.deepcopy(raw); bad['subtables'][0]['evidence'] = ['综合!Z99']
        with self.assertRaisesRegex(ValueError, '目录证据不是原文'): validate_outline(bad, snapshot)
        bad = copy.deepcopy(raw); bad['subtables'][0]['ranges'][0]['range'] = 'A1:Z99'
        with self.assertRaisesRegex(ValueError, '越界'): validate_outline(bad, snapshot)
        raw['subtables'].pop()
        checked = validate_outline(raw, snapshot)
        self.assertFalse(checked['complete'])
        self.assertTrue(any(x['sheet'] == '续页' for x in checked['unlocated_ranges']))

    def test_saved_scope_rejects_missing_candidates_and_false_completion(self):
        from 附注更新.统一勾稽 import load_fact_mapping
        with patch.object(UnifiedSemanticEngine, '_http', lambda e, s, p: self.response(e, s, p)):
            target = self.target()
            metadata = {'carrier': 'B', 'original_path': str(self.source), 'original_hash': file_hash(self.source)}
            result = self.engine().recognize_workbook(self.source, target_mapping=target, artifact_metadata=metadata)
        forged = copy.deepcopy(result); forged['recognition_scope']['candidate_locations'] = []
        path = self.folder/'篡改候选.json'; path.write_text(json.dumps(forged,ensure_ascii=False),encoding='utf-8-sig')
        with self.assertRaisesRegex(ValueError, '候选遗漏'): load_fact_mapping(path, self.engine().standard)
        forged = copy.deepcopy(result); forged['complete'] = True
        path = self.folder/'伪造完整.json'; path.write_text(json.dumps(forged,ensure_ascii=False),encoding='utf-8-sig')
        with self.assertRaisesRegex(ValueError, '覆盖声明'): load_fact_mapping(path, self.engine().standard)

    def test_required_mapping_writes_and_reloads_actual_values(self):
        from 附注更新.分步流程 import run_step
        from 附注更新.统一勾稽 import load_fact_mapping
        with patch.object(UnifiedSemanticEngine, '_http', lambda e, s, p: self.response(e, s, p)):
            target = self.target()
            metadata = {'carrier': 'B', 'original_path': str(self.source), 'original_hash': file_hash(self.source)}
            source = self.engine().recognize_workbook(self.source, target_mapping=target, artifact_metadata=metadata)
            config = {'gold_path':str(self.gold),'a_mapping':target['mapping_path'],'b_mapping':source['mapping_path'],
                      'output_dir':str(self.folder/'实际更新'),'base_url':'http://localhost','model':'本地替身','timeout':1}
            matched = run_step('match', config)
            self.assertEqual(matched['status'], 'complete', matched['message'])
            config['update_plan'] = matched['files']['update_plan']
            written = run_step('write', config)
        self.assertEqual(written['status'], 'complete', written['message'])
        actual = load_fact_mapping(written['files']['a_prime_mapping'], self.engine().standard)
        self.assertEqual(actual['facts'][0]['raw_value'], 125)

    def test_b_candidates_use_own_evidence_and_keep_alternatives(self):
        from openpyxl import load_workbook
        from 附注更新.统一勾稽 import build_fact_update_plan
        book = load_workbook(self.source); book['综合']['A4'] = '固定资产原值合计补充来源'; book['综合']['B4'] = 126
        self.source = self.folder/'多个合理来源.xlsx'; book.save(self.source); book.close()
        def response(engine, system, payload):
            result = self.response(engine, system, payload)
            if payload['task']=='workbook_outline' and len(payload['sheets'])>1:
                result['subtables'][0]['ranges'][0]['range']='A1:C4'
            if payload['task']=='subtable_detail' and payload.get('targets'):
                for row in result['targets']: row['candidates']=['综合!B3','综合!B4']
            if payload['task']=='unified_classify' and 'B4' in payload['candidate_cells']:
                result['cells']['B4']=copy.deepcopy(result['cells']['B3'])
            return result
        with patch.object(UnifiedSemanticEngine, '_http', response):
            target=self.target()
            metadata={'carrier':'B','original_path':str(self.source),'original_hash':file_hash(self.source)}
            source=self.engine().recognize_workbook(self.source,target_mapping=target,artifact_metadata=metadata)
        self.assertEqual({x['source_reference']['location'] for x in source['facts']},{'综合!B3','综合!B4'})
        plan=build_fact_update_plan(target,source,self.engine().standard)
        self.assertEqual(plan['updates'],[])
        self.assertTrue(any(x['kind']=='conflicting_sources' for x in plan['issues']))

    def test_uncertain_subtable_is_not_redundant(self):
        from 附注更新.表格编排 import validate_selection
        from 附注更新.表格 import read_workbook
        snapshot=read_workbook(self.source)
        outline=self.response(None,'',{'task':'workbook_outline','sheets':snapshot['sheets']})
        target={'id':'A需要'}
        answer=self.response(None,'',{'task':'select_subtables','outline':outline,'targets':[target]})
        answer['subtables'][2].update(decision='uncertain',target_ids=['A需要'],reason='用途仍可能包含A来源')
        checked=validate_selection(answer,outline,[target],snapshot)
        self.assertEqual(checked['subtables'][2]['decision'],'uncertain')
        answer['subtables'][0]['target_ids']=[];answer['subtables'][2]['target_ids']=[]
        with self.assertRaisesRegex(ValueError,'遗漏了A需要'): validate_selection(answer,outline,[target],snapshot)

    def test_empty_metric_selection_never_falls_back_to_full_standard(self):
        from 附注更新.统一识别 import build_payload
        from 测试.测试统一识别 import packet
        with self.assertRaisesRegex(ValueError,'零个指标'):
            build_payload(packet(),self.engine().standard,allowed_metrics=[])

    def test_outline_http_does_not_receive_global_cell_classification_task(self):
        import io
        from unittest.mock import Mock
        from 附注更新.表格编排 import OUTLINE_PROMPT
        body=[]
        def receive(request, **kwargs):
            body.append(json.loads(request.data.decode('utf-8')))
            return io.BytesIO(json.dumps({'choices':[{'message':{'content':'{}'},'finish_reason':'stop'}],
                                          'usage':{'completion_tokens':1}}).encode())
        opener=Mock();opener.open.side_effect=receive
        with patch('附注更新.语义.urllib.request.build_opener',return_value=opener):
            self.engine()._http(OUTLINE_PROMPT,{'task':'workbook_outline','sheets':[]})
        self.assertNotIn('目标是识别新布局中每格财务数据',body[0]['messages'][0]['content'], '初识HTTP不能仍混入逐格总任务')
        self.assertEqual(body[0]['response_format']['type'],'json_schema', '初识目录必须有专用结构约束')

    def test_outline_keeps_contiguous_unmerged_title_clues(self):
        from 附注更新.表格编排 import outline_payload
        from 附注更新.表格 import read_workbook
        snapshot=read_workbook(self.source)
        # 同一连续区块内、没有合并样式的标题不能因“只取前两行”而漏掉。
        snapshot['sheets'][0]['cells']['A3']['value']='其中：另一个披露子表'
        payload=outline_payload(snapshot,[])
        self.assertIn('其中：另一个披露子表',json.dumps(payload,ensure_ascii=False))

    def test_a_independent_context_still_receives_verified_classification(self):
        from 附注更新.统一勾稽 import load_fact_mapping
        def response(engine, system, payload):
            result=self.response(engine,system,payload)
            if payload['task']=='workbook_outline':
                result['subtables'][0]['ranges'][0]['range']='A2:B3'
                result['context_ranges']=[{'sheet':'综合','range':'A1:B1'}]
            return result
        with patch.object(UnifiedSemanticEngine,'_http',response): result=self.target()
        self.assertTrue(result['complete'], '公共说明不能因在子表之外而永远未决')
        self.assertEqual({r['cell'] for r in result['excluded']},{'A1','B1','A2','B2','A3'})
        self.assertTrue(load_fact_mapping(result['mapping_path'],self.engine().standard)['complete'])

    def test_target_purpose_does_not_invalidate_original_b_selection(self):
        from 附注更新.统一勾稽 import build_fact_update_plan
        with patch.object(UnifiedSemanticEngine,'_http',lambda e,s,p:self.response(e,s,p)):
            target=self.target()
            source=self.engine().recognize_workbook(self.source,target_mapping=target)
        purpose={'entity':{'from':fact()['dimensions']['entity'],'to':'本次目标企业'},'reason':'明确的本次更新用途'}
        plan=build_fact_update_plan(target,source,self.engine().standard,update_purpose=purpose)
        self.assertEqual(len(plan['updates']),1)

    def test_incomplete_outline_cannot_claim_complete_a(self):
        def response(engine,system,payload):
            result=self.response(engine,system,payload)
            if payload['task']=='workbook_outline': result['unlocated_ranges']=[{'sheet':'综合','range':'A3:B3'}]
            return result
        with patch.object(UnifiedSemanticEngine,'_http',response): result=self.target()
        self.assertFalse(result['outline']['complete'])
        self.assertFalse(result['complete'],'逐格结果不能掩盖目录未定位问题')

    def test_disabled_style_reference_does_not_require_historical_assets(self):
        from 附注更新.统一识别 import build_payload
        from 测试.测试统一识别 import packet
        with patch('附注更新.统一识别.recognition_expressions',side_effect=AssertionError('关闭五样式仍读取历史资产')):
            payload=build_payload(packet(),self.engine().standard,allowed_metrics=['固定资产原值'],style_references=False)
        self.assertEqual(payload['metrics'][0]['definition'],self.engine().standard['metric']['固定资产原值']['definition'])
        self.assertEqual({d['id'] for d in payload['dimensions']},{r['id'] for r in self.engine().standard['metric']['固定资产原值']['dimension_refs']})
        self.assertNotIn('standard_expressions',payload)

    def test_subtable_reading_reaches_both_rounds_without_becoming_identity(self):
        from 附注更新.统一勾稽 import load_fact_mapping
        with patch.object(UnifiedSemanticEngine,'_http',lambda e,s,p:self.response(e,s,p)):
            target=self.target()
        fine=[p for p in self.requests if p['task']=='unified_classify']
        self.assertEqual([p['round'] for p in fine],[1,2])
        for request in fine:
            reading=request.get('subtable_reading',{})
            self.assertEqual(reading.get('structure'),'期末原值列与累计折旧列分别列示')
            self.assertEqual(reading.get('ranges'),[{'sheet':'综合','range':'A1:B3'}])
            self.assertEqual(reading.get('evidence'),['综合!A2'])
            self.assertNotIn('dimensions',reading)
        loaded=load_fact_mapping(target['mapping_path'],self.engine().standard)
        self.assertEqual(loaded['facts'][0]['dimensions'],fact()['dimensions'])

    def test_failed_extra_review_keeps_agreed_a_facts_and_real_disagreement(self):
        from 附注更新.统一勾稽 import load_fact_mapping
        from 附注更新.语义 import SemanticCancelled
        source = self.folder / 'A一致与分歧.xlsx'
        book = Workbook(); sheet = book.active; sheet.title = '综合'
        for address, value in {'A1':'合成企业甲 单户 人民币元 2025年12月31日',
                               'A2':'固定资产原值', 'B2':'期末', 'A3':'全部固定资产',
                               'B3':125, 'A4':'其中：境外固定资产', 'B4':None}.items(): sheet[address]=value
        book.save(source); book.close()
        def response(engine, system, payload):
            if payload['task']=='unified_equivalence': raise OSError('合成的追加复核连接失败')
            answer=self.response(engine,system,payload)
            if payload['task']=='workbook_outline':
                answer={'subjects':['固定资产'],'subtables':[{'id':'原值','subject':'固定资产',
                        'purpose':'固定资产原值','ranges':[{'sheet':'综合','range':'A1:B4'}],
                        'evidence':['综合!A2']}],'unlocated_ranges':[],'context_ranges':[]}
            elif payload['task']=='unified_classify':
                row=copy.deepcopy(answer['cells']['B3'])
                row['dimensions']['asset_selection']={'domain':'固定资产','mode':'predicate',
                    'predicate':'位于境外的固定资产' if payload['round']==1 else '所在地为境外的固定资产',
                    'completeness':'complete'}
                row['dimension_evidence']['asset_selection']=['A4']; answer['cells']['B4']=row
            return answer
        engine=self.engine()
        with patch.object(UnifiedSemanticEngine,'_http',response): result=engine.recognize_workbook(source)
        self.assertEqual([f['source_reference']['location'] for f in result['facts']],['综合!B3'])
        self.assertEqual({r['cell'] for r in result['unresolved']},{'B4'})
        self.assertFalse(result['complete'])
        loaded=load_fact_mapping(result['mapping_path'],engine.standard)
        self.assertEqual(loaded['facts'][0]['raw_value'],125)
        self.assertEqual({r['cell'] for r in loaded['unresolved']},{'B4'})
        with patch('附注更新.同义复核.review_pairs',side_effect=SemanticCancelled('取消仍须停止')):
            with self.assertRaises(SemanticCancelled): engine.recognize_workbook(source)

    def test_a_checks_first_identity_without_regenerating_or_extra_equivalence(self):
        from 附注更新.统一勾稽 import load_fact_mapping
        def response(engine, system, payload):
            if payload['task']=='unified_identity_review':
                self.requests.append(copy.deepcopy(payload))
                self.assertIn('first_reading',payload)
                return {'checked_cells':payload['candidate_cells'],'issues':[]}
            return self.response(engine,system,payload)
        # 普通入口提供carrier=A；无载体的旧调用方式继续兼容。
        with patch.object(UnifiedSemanticEngine,'_http',response):
            initial=self.target()
            self.requests.clear()
            result=self.engine().recognize_workbook(initial['source_path'],artifact_metadata={
                'carrier':'A','original_path':initial['source_path'],'original_hash':initial['source_hash']})
        self.assertEqual([p['task'] for p in self.requests],['unified_identity_review'])
        loaded=load_fact_mapping(result['mapping_path'],self.engine().standard)
        self.assertTrue(loaded['complete'])
        self.assertEqual(loaded['facts'][0]['recognition_provenance']['method'],'model_v2_recognition_checked')
        self.assertEqual(loaded['facts'][0]['dimensions'],fact()['dimensions'])
        # 从空缓存进入也只能一次识别、一次结果核查；不能暗中另跑第二套身份。
        engine=self.engine();engine.work_dir=self.folder/'新的A回执';engine.work_dir.mkdir()
        self.requests.clear()
        with patch.object(UnifiedSemanticEngine,'_http',response):
            fresh=engine.recognize_workbook(initial['source_path'],artifact_metadata={
                'carrier':'A','original_path':initial['source_path'],'original_hash':initial['source_hash']})
        self.assertEqual([p['task'] for p in self.requests],
                         ['workbook_outline','subtable_detail','unified_classify','unified_identity_review'])
        self.assertEqual([p['round'] for p in self.requests if p['task']=='unified_classify'],[1])
        packet=json.loads(Path(fresh['packet_mappings'][0]['path']).read_text(encoding='utf-8-sig'))
        from 附注更新.统一识别 import verify_identity_review_mapping
        changed=copy.deepcopy(packet);changed['facts'][0]['dimensions']['period']['date']='2024-12-31'
        with self.assertRaisesRegex(ValueError,'核查结果与实际回执不符'):
            verify_identity_review_mapping(changed,engine.standard)
        changed=copy.deepcopy(packet);changed['identity_review']['review_receipts'][0]['sha256']='0'*64
        with self.assertRaises(ValueError):verify_identity_review_mapping(changed,engine.standard)

    def test_a_review_errors_and_unchecked_cells_stay_pending_without_retry(self):
        from 附注更新.统一勾稽 import load_fact_mapping
        with patch.object(UnifiedSemanticEngine,'_http',lambda e,s,p:self.response(e,s,p)): initial=self.target()
        def response(engine, system, payload):
            self.assertEqual(payload['task'],'unified_identity_review')
            self.requests.append(copy.deepcopy(payload))
            return {'checked_cells':['B3'],'issues':[{'cells':['B3'],'reason':'项目范围没有充分依据','evidence_cells':['A3']}]}
        self.requests.clear()
        with patch.object(UnifiedSemanticEngine,'_http',response):
            result=self.engine().recognize_workbook(initial['source_path'],artifact_metadata={
                'carrier':'A','original_path':initial['source_path'],'original_hash':initial['source_hash']})
        self.assertEqual(len(self.requests),1)
        self.assertFalse(result['complete']);self.assertEqual(result['facts'],[])
        self.assertEqual({r['cell'] for r in result['unresolved']},{'A1','B1','A2','B2','A3','B3'})
        self.assertFalse(load_fact_mapping(result['mapping_path'],self.engine().standard)['complete'])

    def test_a_review_bad_coverage_does_not_start_another_model_round(self):
        from 附注更新.统一勾稽 import load_fact_mapping
        from 附注更新.语义 import SemanticCancelled
        with patch.object(UnifiedSemanticEngine,'_http',lambda e,s,p:self.response(e,s,p)): initial=self.target()
        calls=[]
        def response(engine,system,payload):
            calls.append(payload['task'])
            return {'checked_cells':['Z99'],'issues':[]}
        metadata={'carrier':'A','original_path':initial['source_path'],'original_hash':initial['source_hash']}
        with patch.object(UnifiedSemanticEngine,'_http',response):
            result=self.engine().recognize_workbook(initial['source_path'],artifact_metadata=metadata)
        self.assertEqual(calls,['unified_identity_review'])
        self.assertFalse(result['complete']);self.assertEqual(result['facts'],[])
        self.assertEqual(len(result['unresolved']),6)
        self.assertFalse(load_fact_mapping(result['mapping_path'],self.engine().standard)['complete'])
        with patch.object(UnifiedSemanticEngine,'_http',side_effect=SemanticCancelled('用户取消')):
            with self.assertRaises(SemanticCancelled):
                self.engine().recognize_workbook(initial['source_path'],artifact_metadata=metadata)

    def test_a_review_accepts_complete_json_with_plain_explanation_but_not_two_answers(self):
        import io
        from 附注更新.语义 import SemanticResponseError
        engine=self.engine();expected={'checked_cells':['B3'],'issues':[]}
        payload={'task':'unified_identity_review','candidate_cells':['B3'],'cells':{'A3':{'value':'全部固定资产'}}}
        def reply(content):
            stream=io.BytesIO(json.dumps({'choices':[{'finish_reason':'stop','message':{'content':content}}],
                            'usage':{'prompt_tokens':10,'completion_tokens':20}},ensure_ascii=False).encode())
            stream.headers={'Content-Type':'application/json'}
            return stream
        text=json.dumps(expected,ensure_ascii=False)
        with patch('附注更新.语义.urllib.request.build_opener') as opener:
            opener.return_value.open.return_value.__enter__.return_value=reply(text+'\n核查了表头和行项目，范围有原文依据。')
            self.assertEqual(engine._http('只核查原结果',payload),expected)
            self.assertEqual(engine._last_transport_repair['repairs'][0]['kind'],'trailing_review_commentary')
        with patch('附注更新.语义.urllib.request.build_opener') as opener:
            opener.return_value.open.return_value.__enter__.return_value=reply('```json\n'+text+'\n```\n核查说明：\n- 采用 `metric.example`，维度引用 [0,1,2] 有依据。')
            self.assertEqual(engine._http('只核查原结果',payload),expected)
        for content in (text+'\n'+text,text+'\n[]',text+'\n```json\n'+text+'\n```',text[:-1]):
            with patch('附注更新.语义.urllib.request.build_opener') as opener:
                opener.return_value.open.return_value.__enter__.return_value=reply(content)
                # 原来允许补唯一根结束符，故截断改为缺内部数组，不能靠补根闭合恢复。
                if content==text[:-1]:opener.return_value.open.return_value.__enter__.return_value=reply('{"checked_cells":["B3"')
                with self.assertRaises(SemanticResponseError):engine._http('只核查原结果',payload)

    def test_missing_candidate_rechecks_other_relevant_subtable_without_full_b(self):
        from openpyxl import load_workbook
        from 附注更新.统一勾稽 import load_fact_mapping, build_fact_update_plan
        book=load_workbook(self.source)
        book['综合']['A3']='生产设备'
        book['综合']['A6']='合成企业甲 单户 人民币元 2025年12月31日 固定资产原值总额'
        book['综合']['A7']='全部固定资产原值'
        self.source=self.folder/'其他子表含完整来源.xlsx';book.save(self.source);book.close()
        def response(engine,system,payload):
            result=self.response(engine,system,payload)
            if payload['task']=='workbook_outline' and len(payload['sheets'])>1:
                result['subtables'][1]['purpose']='固定资产原值总额，含完整统计范围'
            if payload['task']=='select_subtables' and payload.get('candidate_followup'):
                row=result['subtables'][1]
                row.update(decision='needed',target_ids=[t['id'] for t in payload['targets']],reason='总额子表可能提供完整范围的来源')
            if payload['task']=='subtable_detail' and payload.get('targets'):
                for row in result['targets']:
                    row['candidates']=['综合!B7'] if payload['subtable']['id']=='前五名' else []
                result['allowed_metrics']=['固定资产原值']
            if payload['task']=='unified_classify' and 'B7' in payload['candidate_cells']:
                result['cells']['B7']={'kind':'metric_fact','metric_id':'固定资产原值','dimensions':fact()['dimensions'],
                    'dimension_evidence':{k:['A6','A7'] for k in fact()['dimensions']},'reason':'原文总额及完整范围'}
            return result
        with patch.object(UnifiedSemanticEngine,'_http',response):
            target=self.target();self.requests.clear()
            source=self.engine().recognize_workbook(self.source,target_mapping=target)
        self.assertEqual({f['source_reference']['location'] for f in source['facts']},{'综合!B7'})
        self.assertEqual([p['subtable']['id'] for p in self.requests if p['task']=='subtable_detail'],['原值','前五名'])
        self.assertEqual({a for p in self.requests if p['task']=='unified_classify' for a in p['candidate_cells']},{'B7'})
        loaded=load_fact_mapping(source['mapping_path'],self.engine().standard)
        plan=build_fact_update_plan(target,loaded,self.engine().standard)
        self.assertEqual([(u['cell'],u['value']) for u in plan['updates']],[('B3',13)])
        self.assertEqual(plan['issues'],[])

    def test_unavailable_optional_expressions_keep_current_standard_and_source(self):
        from 附注更新.统一识别 import build_payload
        from 测试.测试统一识别 import packet
        for failure in (FileNotFoundError('参考资产缺失'),ValueError('参考资产与标准不兼容')):
            with self.subTest(failure=str(failure)), patch('附注更新.统一识别.recognition_expressions',side_effect=failure) as reference:
                payload=build_payload(packet(),self.engine().standard,allowed_metrics=['固定资产原值'])
            reference.assert_called_once()
            self.assertNotIn('standard_expressions',payload)
            self.assertEqual([m['id'] for m in payload['metrics']],['固定资产原值'])
            self.assertEqual(payload['candidate_cells'],packet()['candidate_cells'])
            self.assertIn('未采用',payload['standard_expression_status'])

    def test_added_a_period_only_recognizes_new_b_candidates(self):
        from openpyxl import load_workbook
        from 附注更新.统一勾稽 import load_fact_mapping
        book=load_workbook(self.source);book['综合']['C2']='2024年12月31日期初原值'
        source=self.folder/'两个期间的B.xlsx';book.save(source);book.close()
        def response(engine,system,payload):
            result=self.response(engine,system,payload)
            if payload['task']=='workbook_outline' and len(payload['sheets'])==1:
                if payload['sheets'][0]['max_column']==3:result['subtables'][0]['ranges'][0]['range']='A1:C3'
            if payload['task']=='subtable_detail':
                for need, row in zip(payload['targets'],result['targets']):
                    row['candidates']=['综合!'+('C3' if need['dimensions']['period']['date']=='2024-12-31' else 'B3')]
            if payload['task']=='unified_classify' and 'C3' in payload['candidate_cells']:
                dims=copy.deepcopy(fact()['dimensions']);dims['period']['date']='2024-12-31'
                result['cells']['C3']={'kind':'metric_fact','metric_id':'固定资产原值','dimensions':dims,
                    'dimension_evidence':{k:['A1','A2','A3','C2'] for k in dims},'reason':'期初原值合计的完整原文'}
            return result
        with patch.object(UnifiedSemanticEngine,'_http',response):
            first=self.target()
            previous=self.engine().recognize_workbook(source,target_mapping=first)
            book=load_workbook(first['source_path']);book['综合']['C2']='2024年12月31日期初原值';book['综合']['C3']=20
            expanded=self.folder/'A增加期初需要.xlsx';book.save(expanded);book.close()
            target=self.engine().recognize_workbook(expanded)
            self.requests.clear()
            updated=self.engine().recognize_workbook(source,target_mapping=target,previous_result=previous['mapping_path'])
        self.assertEqual({f['source_reference']['location'] for f in updated['facts']},{'综合!B3','综合!C3'})
        self.assertEqual({a for p in self.requests if p['task']=='unified_classify' for a in p['candidate_cells']},{'C3'})
        self.assertEqual(len([p for p in self.requests if p['task']=='subtable_detail']),1,'两个需要共用子表，不重复详看')
        coverage=load_fact_mapping(updated['mapping_path'],self.engine().standard)['recognition_scope']['coverage']
        self.assertEqual(len(coverage['targets']),2)
        self.assertTrue(coverage['satisfied'])

    def test_incomplete_outline_retries_without_changing_saved_evidence(self):
        from 附注更新.统一勾稽 import load_fact_mapping
        metadata={'carrier':'B','original_path':str(self.source),'original_hash':file_hash(self.source)}
        def uncertain(engine,system,payload):
            answer=self.response(engine,system,payload)
            answer['unlocated_ranges']=[{'sheet':'综合','range':'A3:C3'}]
            return answer
        with patch.object(UnifiedSemanticEngine,'_http',uncertain): old=self.engine().recognize_workbook(self.source,artifact_metadata=metadata)
        self.assertFalse(old['outline']['complete'])
        with patch.object(UnifiedSemanticEngine,'_http',lambda e,s,p:self.response(e,s,p)):
            fresh=self.engine().recognize_workbook(self.source,artifact_metadata=metadata)
        self.assertTrue(fresh['outline']['complete'],'未定位目录不可永远命中旧缓存')
        self.assertFalse(load_fact_mapping(old['mapping_path'],self.engine().standard)['outline']['complete'],'旧进度应仍可核验')

    def test_ui_does_not_describe_required_b_as_unfinished_whole_book(self):
        from 测试.测试界面进度 import 界面进度测试
        metadata={'carrier':'B','original_path':str(self.source),'original_hash':file_hash(self.source)}
        with patch.object(UnifiedSemanticEngine,'_http',lambda e,s,p:self.response(e,s,p)):
            target=self.target();source=self.engine().recognize_workbook(self.source,target_mapping=target,artifact_metadata=metadata)
        ui=界面进度测试();ui.setUp()
        try:
            app=ui.make_app({**ui.base,'gold_path':str(self.gold),'b_mapping':source['mapping_path']})
            app.show_step('recognize_b')
            text=app.mapping_summary_labels['b_mapping'].cget('text')
            self.assertIn('按需识别',text)
            self.assertNotIn('部分成果',text)
        finally:ui.tearDown()

    def test_unrelated_standard_change_keeps_prior_scope_readable(self):
        from 附注更新.统一勾稽 import load_fact_mapping
        from 附注更新.统一语义 import read_standard_v2
        rows=records();extra=copy.deepcopy(rows[-1]);extra['id']='另一个未使用指标';extra['definition']['measure']='未使用口径'
        rows.append(extra)
        self.gold=self.folder/'含未使用指标的标准.jsonl'
        self.gold.write_text('\n'.join(json.dumps(row,ensure_ascii=False) for row in rows),encoding='utf-8-sig')
        with patch.object(UnifiedSemanticEngine,'_http',lambda e,s,p:self.response(e,s,p)):
            target=self.target();source=self.engine().recognize_workbook(self.source,target_mapping=target)
        changed=copy.deepcopy(rows);changed[-1]['definition']['measure']='修订未使用口径'
        new_gold=self.folder/'修订未使用定义的标准.jsonl'
        new_gold.write_text('\n'.join(json.dumps(row,ensure_ascii=False) for row in changed),encoding='utf-8-sig')
        loaded=load_fact_mapping(source['mapping_path'],read_standard_v2(new_gold))
        self.assertEqual(loaded['facts'],source['facts'])


if __name__ == '__main__':
    unittest.main()
