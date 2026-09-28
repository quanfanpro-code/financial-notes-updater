# coding: utf-8
"""正式入口取消后保留载体关联；请求去重必须逐字可还原定义。"""
import copy
import json
from pathlib import Path
import threading
import unittest
from unittest.mock import patch
import uuid
from docx import Document
from 附注更新.分步流程 import run_step
from 附注更新.统一识别 import build_payload
from 附注更新.统一语义 import validate_records,read_standard_v2
from 附注更新.统一勾稽 import load_fact_mapping
from 附注更新.界面 import file_selection_updates
from 附注更新.外部复核 import export_review_template
from 测试.测试统一语义 import records
from 测试.测试统一识别 import packet


class ProgressPayloadTests(unittest.TestCase):
    def test_重复单元格属性无损共用且保留原包(self):
        source=packet();source['cells']={}
        for number in range(1,41):
            source['cells']['A'+str(number)]={'row':number,'column':1,'value':'原文项目'+str(number),
                'cached_value':'原文项目'+str(number),'formula':None,'number_format':'#,##0.00',
                'style_id':15,'hidden_row':False,'hidden_column':False,'hidden_sheet':False,'hidden':False}
        source['candidate_cells']=['A1'];source['cells']['A40']['hidden_row']=True
        source['cells']['A38'].update(value='=1+2',cached_value=3,formula='=1+2')
        source['cells']['A39'].update(value=0,cached_value=False)
        before=copy.deepcopy(source);standard=validate_records(records())
        result=build_payload(source,standard)
        self.assertTrue(result.get('shared_cell_properties'))
        self.assertTrue(any(row.get('cached_value')=={'from_cell':'value'}
                            for row in result['shared_cell_properties'].values()))
        restored={}
        for address,row in result['cells'].items():
            current=copy.deepcopy(row);key=current.pop('shared_properties',None)
            restored[address]={**result['shared_cell_properties'].get(key,{}),**current}
            if restored[address].get('cached_value')=={'from_cell':'value'}:
                restored[address]['cached_value']=copy.deepcopy(restored[address]['value'])
        self.assertEqual(restored,source['cells']);self.assertEqual(source,before)
        self.assertIs(restored['A39']['cached_value'],False)
        plain=build_payload(source,standard,compact_cells=False)
        self.assertEqual(plain['cells'],source['cells'])
        self.assertEqual({k:v for k,v in result.items() if k not in {'cells','shared_cell_properties','cell_property_instructions'}},
                         {k:v for k,v in plain.items() if k!='cells'})
        self.assertLess(len(json.dumps(result,ensure_ascii=False)),len(json.dumps(plain,ensure_ascii=False)))

    def test_共用定义无损传递且不改原标准(self):
        rows=records();common='本段是多个维度共有的原文约束，必须完整保留每一句限定。'*6
        for row in rows:
            if row['record_type']=='dimension':row['meaning']=common+'  '+row['meaning']+' '+common
        standard=validate_records(rows);before=copy.deepcopy(standard)
        payload=build_payload(packet(),standard)
        self.assertTrue(payload.get('shared_dimension_rules'))
        restored=[]
        for row in payload['dimensions']:
            current=copy.deepcopy(row)
            if 'meaning_parts' in current:
                current['meaning']=' '.join(part if isinstance(part,str) else payload['shared_dimension_rules'][part['rule']] for part in current.pop('meaning_parts'))
            restored.append(current)
        self.assertEqual(restored,[standard['dimension'][row['id']] for row in restored])
        self.assertEqual(standard,before)

    def test_首个进度可从界面恢复原件说明和Word关联(self):
        root=Path(__file__).resolve().parents[1]/'测试结果'/('语义进度关联_'+uuid.uuid4().hex);root.mkdir()
        gold=root/'标准.jsonl';gold.write_text('\n'.join(json.dumps(r,ensure_ascii=False) for r in records()),encoding='utf-8')
        word=root/'附注.docx';doc=Document();doc.add_paragraph('六、财务报表重要项目的说明');doc.add_paragraph('固定资产')
        table=doc.add_table(rows=2,cols=2)
        for row,values in zip(table.rows,[['项目','期末原值'],['固定资产','125']]):
            for cell,value in zip(row.cells,values):cell.text=value
        doc.add_paragraph('七、其他事项');doc.save(word)
        config={'word_path':str(word),'gold_path':str(gold),'output_dir':str(root/'正式流程'),
                'base_url':'http://127.0.0.1:1/v1','model':'受控本地测试','api_key':'','a_context_text':'合成企业甲；单户；人民币元；2025年12月31日'}
        extracted=run_step('extract_a',config);self.assertEqual(extracted['status'],'complete')
        config.update(extracted['files']);cancel=threading.Event()
        with patch('附注更新.统一识别.UnifiedSemanticEngine._request',side_effect=AssertionError('取消前不应调用模型')):
            stopped=run_step('recognize_a',config,cancel=cancel,progress=lambda event:cancel.set())
        self.assertEqual(stopped['status'],'cancelled')
        mapping=stopped['files']['a_mapping']
        restored=file_selection_updates('a_mapping',mapping,{})
        self.assertEqual(restored['a_path'],config['a_path'])
        self.assertEqual(restored['a_context_text'],config['a_context_text'])
        self.assertEqual(restored['a_link'],config['a_link'])
        saved=load_fact_mapping(mapping,read_standard_v2(gold))
        self.assertEqual(saved['carrier'],'A');self.assertFalse(saved['complete'])
        self.assertTrue(saved['source_context'])
        opinion=root/'进度复核格式.json';export_review_template(mapping,opinion,gold_path=gold)
        self.assertTrue(json.loads(opinion.read_text(encoding='utf-8-sig'))['items'])


if __name__=='__main__':unittest.main()
