# coding: utf-8
"""实际来源比例在xlsx保存时只损失浮点尾位，不放行金额差错。"""
import unittest,uuid
from pathlib import Path
from openpyxl import Workbook,load_workbook
from 附注更新.表格 import write_updated_workbook


class Excel数值回读测试(unittest.TestCase):
    def test_原封保存的小数仍与自身相等(self):
        from 附注更新.表格 import same_excel_value
        for value in (23417817.779999997, -4083105.6300000027,
                      -1793066.6599999927, 2771089.1099999994):
            self.assertTrue(same_excel_value(value, value))
            self.assertFalse(same_excel_value(value + 0.01, value))

    def test_按实际保存精度回读且拒绝改值(self):
        out=Path(__file__).resolve().parents[1]/'测试输出'/('数值回读_'+uuid.uuid4().hex);out.mkdir()
        source=out/'原表.xlsx';result=out/'更新副本.xlsx'
        book=Workbook();book.active['A1']=10.56;book.save(source);book.close()
        write_updated_workbook(source,result,{'updates':[{'sheet':'Sheet','cell':'A1','value':10.562338425840736}]})
        book=load_workbook(result);actual=book.active['A1'].value;book.close()
        self.assertEqual(actual,10.56233842584074)
        from 附注更新.表格 import same_excel_value
        self.assertTrue(same_excel_value(actual,10.562338425840736))
        self.assertFalse(same_excel_value(10.56233842594074,10.562338425840736))
        self.assertFalse(same_excel_value(1.01,1.00))
        self.assertTrue(same_excel_value(97.18,97.18))
        self.assertFalse(same_excel_value(97.1800000001,97.18))


if __name__=='__main__':unittest.main()
