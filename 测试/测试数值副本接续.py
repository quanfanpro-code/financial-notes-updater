# coding: utf-8
"""数值副本复用语义必须核验实际副本，不能连同表头或身份一起改掉。"""
import json
import unittest
from openpyxl import load_workbook
from 测试 import 测试统一勾稽 as 夹具
from 附注更新 import 统一勾稽 as api
from 附注更新.统一语义 import read_standard_v2


class 数值副本接续测试(unittest.TestCase):
    def test_数值改变可正式回读但表头及身份变化被拒绝(self):
        self.assertTrue(hasattr(api, 'derive_value_copy'), '缺少仅数值变化的正式接续入口')
        case=夹具.MatchTests();case.setUp();root,gold,paths=case.file_fixture()
        parent=api.load_fact_mapping(paths[1],read_standard_v2(gold))
        book=load_workbook(parent['source_path']);book['来源']['J3']=900.5
        copy_path=root/'测试来源.xlsx';book.save(copy_path);book.close()
        derived=api.derive_value_copy(parent,copy_path)
        path=root/'测试来源映射.json';path.write_text(json.dumps(derived,ensure_ascii=False),encoding='utf-8-sig')
        actual=api.load_fact_mapping(path,read_standard_v2(gold))
        self.assertEqual(actual['facts'][0]['raw_value'],900.5)
        self.assertEqual(actual['facts'][0]['dimensions'],parent['facts'][0]['dimensions'])
        derived['facts'][0]['raw_value']=901
        wrong=root/'伪改映射.json';wrong.write_text(json.dumps(derived,ensure_ascii=False),encoding='utf-8-sig')
        with self.assertRaisesRegex(ValueError,'数值副本'):
            api.load_fact_mapping(wrong,read_standard_v2(gold))
        book=load_workbook(copy_path);book['来源']['A1']='替换了含义';changed=root/'改了表头.xlsx';book.save(changed);book.close()
        with self.assertRaisesRegex(ValueError,'数据格以外'):
            api.derive_value_copy(parent,changed)


if __name__=='__main__':unittest.main()
