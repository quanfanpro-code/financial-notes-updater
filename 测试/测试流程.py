"""真实文件流程测试；模型边界使用已声明测试回执，不代表语义准确率。"""
import importlib
import threading
import unittest
import uuid
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/"测试结果"/("流程_"+uuid.uuid4().hex[:8])
WORK.mkdir(parents=True,exist_ok=True)

class FakeEngine:
    def __init__(self,*args,**kwargs): self.gold_hash="test-only";self.usage={}
    def recognize(self,snapshot,reference=None):
        mappings=[];excluded=[]
        for sheet in snapshot["sheets"]:
            for address,cell in sheet["cells"].items():
                if isinstance(cell.get("value"),(int,float)):
                    mappings.append({"sheet":sheet["name"],"cell":address,"slot_id":"C-N087-T001-S001",
                    "dimensions":{"currency":"CNY","unit":"元","scale":1},"value_type":"monetary",
                    "semantic_field":"库存现金期末余额","scope":"consolidated","reason":"测试回执"})
                else: excluded.append({"sheet":sheet["name"],"cell":address,"reason":"测试标题"})
        return {"source_hash":snapshot["sha256"],"gold_hash":"test-only","mappings":mappings,
            "excluded":excluded,"unresolved":[],"complete":True,"usage":{}}

class FlowTests(unittest.TestCase):
    def api(self):
        try:return importlib.import_module("附注更新.流程")
        except ImportError:self.fail("尚未实现完整流程")
    def config(self,word=False):
        folder=WORK/uuid.uuid4().hex[:8];folder.mkdir()
        b=Workbook();s=b.active;s.append(["项目","库存现金"]);s.append(["期末",125])
        bp=folder/"新数据.xlsx";b.save(bp)
        ap=folder/("文档一.docx" if word else "表格A.xlsx")
        if word:
            d=Document();d.add_paragraph("合并财务报表附注 2025年12月31日 单位：人民币元")
            t=d.add_table(rows=2,cols=2);t.cell(0,0).text="项目";t.cell(0,1).text="期末"
            t.cell(1,0).text="库存现金";t.cell(1,1).text="7";d.save(ap)
        else:
            a=Workbook();a.active.append(["项目","期末"]);a.active.append(["库存现金",7]);a.save(ap)
        gold=folder/"金标准.jsonl";gold.write_text("{}",encoding="utf-8-sig")
        return {"source_a":str(ap),"source_b":str(bp),"output_dir":str(folder/"结果"),"gold_path":str(gold),
                "base_url":"http://127.0.0.1:1/v1","model":"test","api_key":"","mode":"update","review":True}
    def test_excel_pipeline_reads_source_and_writes_new_copy(self):
        api=self.api();cfg=self.config()
        with patch.object(api,"SemanticEngine",FakeEngine): result=api.run_job(cfg,lambda _:None,threading.Event())
        self.assertEqual(result["status"],"complete")
        self.assertEqual(load_workbook(result["excel"]).active["B2"].value,125)
        self.assertEqual(load_workbook(cfg["source_a"]).active["B2"].value,7)
    def test_word_pipeline_updates_original_structure(self):
        api=self.api();cfg=self.config(word=True)
        with patch.object(api,"SemanticEngine",FakeEngine): result=api.run_job(cfg,lambda _:None,threading.Event())
        self.assertEqual(result["status"],"complete")
        self.assertEqual(Document(result["word"]).tables[0].cell(1,1).text,"125")
        self.assertEqual(Document(cfg["source_a"]).tables[0].cell(1,1).text,"7")

    def test_failure_is_recorded_without_exposing_key(self):
        api=self.api();cfg=self.config();cfg["api_key"]="test-secret-never-save"
        class Failed(FakeEngine):
            def recognize(self,snapshot,reference=None):
                raise RuntimeError("测试失败 "+cfg["api_key"])
        with patch.object(api,"SemanticEngine",Failed):
            result=api.run_job(cfg,lambda _:None,threading.Event())
        self.assertEqual(result["status"],"failed")
        record=Path(result["output_dir"])/"运行结果.json"
        self.assertTrue(record.exists())
        self.assertNotIn(cfg["api_key"],record.read_text(encoding="utf-8-sig"))


    def test_usage_includes_both_a_and_b(self):
        api=self.api();cfg=self.config()
        class Counted(FakeEngine):
            def recognize(self,snapshot,reference=None):
                result=super().recognize(snapshot,reference)
                self.usage={"requests":3 if reference is None else 4,"prompt_tokens":20}
                result["usage"]=dict(self.usage)
                return result
        with patch.object(api,"SemanticEngine",Counted):
            result=api.run_job(cfg,lambda _:None,threading.Event())
        self.assertEqual(result["status"],"complete")
        self.assertEqual(result["usage"]["requests"],7)
        self.assertEqual(result["usage"]["prompt_tokens"],40)


    def test_word_new_customer_row_flows_through_a_prime(self):
        api=self.api();cfg=self.config(word=True)
        # 新用例使用独立文件，避免改写前一份测试原件。
        folder=Path(cfg["source_a"]).parent
        doc=Document();doc.add_paragraph("合并应收账款明细，2025年12月31日，人民币元")
        table=doc.add_table(rows=4,cols=2);table.style="Table Grid"
        for row,values in zip(table.rows,[["客户","期末余额"],["甲","10"],["乙","20"],["合计","30"]]):
            for cell,value in zip(row.cells,values):cell.text=value
        word=folder/"客户原报告.docx";doc.save(word)
        b=Workbook();b.active.append(["客户","甲","乙","丙","合计"]);b.active.append(["期末余额",10,20,30,60])
        bp=folder/"客户横排B.xlsx";b.save(bp)
        cfg.update(source_a=str(word),source_b=str(bp))
        class StructureEngine(FakeEngine):
            log=staticmethod(lambda message:None)
            def _check_cancel(self):pass
            def recognize(self,snapshot,reference=None):
                s=snapshot["sheets"][0];maps=[]
                is_a=bool(snapshot.get("names"))
                for address,c in s["cells"].items():
                    if not isinstance(c.get("value"),(int,float)):continue
                    from openpyxl.utils import get_column_letter
                    label=(get_column_letter(c["column"]-1)+str(c["row"])) if is_a else get_column_letter(c["column"])+"1"
                    obj=s["cells"][label]["value"]
                    maps.append({"sheet":s["name"],"cell":address,"slot_id":"C-N001-T001-S002" if obj=="合计" else "C-N001-T001-S001",
                        "dimensions":{"counterparty":obj,"currency":"CNY","unit":"元","scale":1},"scope":"consolidated",
                        "value_type":"monetary","semantic_field":"期末余额","table_range":"B2:C5" if is_a else "A1:E2",
                        "aggregation":"total" if obj=="合计" else "component","reviewed":True})
                return {"complete":True,"source_hash":snapshot["sha256"],"mappings":maps,"excluded":[],"unresolved":[],"usage":{}}
            def _request(self,instruction,payload,validator):
                return validator({"operations":[{"range_name":"TA_001","axis":"row","action":"insert","index":3,"count":1,"template_index":2,"reason":"B新增客户丙"}],
                    "bindings":[{"source_sheet":"Sheet","source_cell":"D2","target_sheet":"sheet1","target_cell":"C5","template_sheet":"sheet1","template_cell":"C4","reason":"丙期末余额"}],
                    "labels":[{"sheet":"sheet1","cell":"B5","value":"丙","source_sheet":"Sheet","source_cell":"D1","reason":"B客户名称"}],"unresolved":[]})
        with patch.object(api,"SemanticEngine",StructureEngine):
            result=api.run_job(cfg,lambda _:None,threading.Event())
        self.assertEqual(result["status"],"complete",result)
        updated=Document(result["word"]).tables[0]
        self.assertEqual(len(updated.rows),5)
        self.assertEqual(updated.cell(3,0).text,"丙")
        self.assertEqual(updated.cell(3,1).text,"30")
        self.assertEqual(updated.cell(4,1).text,"60")
        self.assertEqual(len(Document(word).tables[0].rows),4)


    def test_word_new_period_column_preserves_amount_format(self):
        api=self.api();cfg=self.config(word=True);folder=Path(cfg["source_a"]).parent
        doc=Document();doc.add_paragraph("合并应收账款，人民币元")
        table=doc.add_table(rows=4,cols=2);table.style="Table Grid"
        for row,values in zip(table.rows,[["客户","期末余额"],["甲","10.00"],["乙","20.00"],["合计","30.00"]]):
            for cell,value in zip(row.cells,values):cell.text=value
        word=folder/"单期原报告.docx";doc.save(word)
        b=Workbook();b.active.append(["客户","甲","乙","合计"])
        b.active.append(["期末余额",10,20,30]);b.active.append(["期初余额",7,8,15])
        bp=folder/"双期横排B.xlsx";b.save(bp);cfg.update(source_a=str(word),source_b=str(bp))
        class PeriodEngine(FakeEngine):
            log=staticmethod(lambda message:None)
            def _check_cancel(self):pass
            def recognize(self,snapshot,reference=None):
                from openpyxl.utils import get_column_letter
                s=snapshot["sheets"][0];is_a=bool(snapshot.get("names"));maps=[]
                for address,c in s["cells"].items():
                    if not isinstance(c.get("value"),(int,float)):continue
                    obj=s["cells"][get_column_letter(c["column"]-1)+str(c["row"]) if is_a else get_column_letter(c["column"])+"1"]["value"]
                    opening=not is_a and c["row"]==3
                    slot=("S004" if opening else "S003") if obj=="合计" else ("S002" if opening else "S001")
                    maps.append({"sheet":s["name"],"cell":address,"slot_id":"C-N001-T001-"+slot,
                        "dimensions":{"counterparty":obj,"currency":"CNY","unit":"元","scale":1,"metric":"期初余额" if opening else "期末余额"},
                        "scope":"consolidated","value_type":"monetary","semantic_field":"期初余额" if opening else "期末余额",
                        "table_range":"B2:C5" if is_a else "A1:D3","aggregation":"total" if obj=="合计" else "component","reviewed":True})
                return {"complete":True,"source_hash":snapshot["sha256"],"mappings":maps,"excluded":[],"unresolved":[],"usage":{}}
            def _request(self,instruction,payload,validator):
                return validator({"operations":[{"range_name":"TA_001","axis":"column","action":"insert","index":2,"count":1,"template_index":1,"reason":"新增期初指标"}],
                    "bindings":[{"source_sheet":"Sheet","source_cell":src,"target_sheet":"sheet1","target_cell":"D"+str(row),"template_sheet":"sheet1","template_cell":"C"+str(row),"reason":"同客户期初余额"} for row,src in [(3,"B3"),(4,"C3"),(5,"D3")]],
                    "labels":[{"sheet":"sheet1","cell":"D2","value":"期初余额","source_sheet":"Sheet","source_cell":"A3","reason":"B期初表头"}],"unresolved":[]})
        with patch.object(api,"SemanticEngine",PeriodEngine):result=api.run_job(cfg,lambda _:None,threading.Event())
        self.assertEqual(result["status"],"complete",result)
        updated=Document(result["word"]).tables[0]
        self.assertEqual(len(updated.columns),3)
        self.assertEqual(updated.cell(0,2).text,"期初余额")
        self.assertEqual(updated.cell(1,2).text,"7.00")
        self.assertEqual(updated.cell(3,2).text,"15.00")
        self.assertEqual(updated.cell(1,1).text,"10.00")
        self.assertEqual(len(Document(word).tables[0].columns),2)

    def test_cancel_before_start_creates_no_final(self):
        api=self.api();cfg=self.config();event=threading.Event();event.set()
        result=api.run_job(cfg,lambda _:None,event)
        self.assertEqual(result["status"],"cancelled")
    def test_unresolved_is_not_completed_or_word_published(self):
        api=self.api();cfg=self.config()
        class Unresolved(FakeEngine):
            def recognize(self,snapshot,reference=None):
                result=super().recognize(snapshot,reference)
                result["complete"]=False;result["unresolved"]=[{"reason":"测试缺上下文"}]
                return result
        with patch.object(api,"SemanticEngine",Unresolved): result=api.run_job(cfg,lambda _:None,threading.Event())
        self.assertEqual(result["status"],"partial");self.assertNotIn("word",result)
