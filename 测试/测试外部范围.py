"""外部范围确认的本地虚构数据验收，不调用模型。"""
import copy
import importlib
import json
from pathlib import Path
import unittest
import uuid
from unittest.mock import patch
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from 附注更新 import 分步流程 as flow
from 附注更新.表格 import read_workbook,file_hash,materialize_candidate_blanks
from 附注更新.识别范围 import build_target_meanings

ROOT=Path(__file__).resolve().parents[1]

class ExternalScopeTests(unittest.TestCase):
    def setUp(self):
        self.work=ROOT/"测试结果"/("外部范围_"+uuid.uuid4().hex);self.work.mkdir(parents=True)
        self.source=self.work/"本地虚构更新数据.xlsx";book=Workbook();sheet=book.active;sheet.title="混合资料"
        for row in [["货币资金","期末余额"],["新银行账户",0],[None,None],["人数统计","人数"],["新职工",9]]:sheet.append(row)
        sheet["B2"].number_format="0.00";book.save(self.source);book.close()
        self.target=self.work/"本地虚构附注.xlsx";book=Workbook();sheet=book.active;sheet.title="附注"
        for row in [["货币资金","期末余额"],["银行存款",None]]:sheet.append(row)
        book.defined_names.add(DefinedName("TA1",attr_text="'附注'!$A$1:$B$2"));book.save(self.target);book.close()
        word=self.work/"虚构Word原件.docx";word.write_bytes(b"local fixture only")
        linked=self.work/"虚构Word连接.docx";linked.write_bytes(b"local fixture linked only")
        ctx={"scope":{"chapter_title":flow.CHAPTER_TITLE},"tables":[{"sheet":"附注","table_id":"t1","range_name":"TA1","title":"货币资金","chapter_context":[flow.CHAPTER_TITLE,"人民币元"]}]}
        self.link=self.save("位置关联",{"schema":flow.LINK_SCHEMA,"scope":ctx["scope"],"context":ctx,"a_path":str(self.target),"a_hash":file_hash(self.target),"word_path":str(word),"word_hash":file_hash(word),"linked_word":str(linked),"linked_word_hash":file_hash(linked)})
        self.gold=self.work/"本地标准.jsonl";self.gold.write_text(json.dumps({"id":"C-N087-T001-S001","status":"active","scope":"consolidated","note":{"id":"C-N087","name":"货币资金"},"table":{"id":"C-N087-T001","name":"资金"},"slot":{"name":"余额"},"value_type":"monetary","dimensions":[]},ensure_ascii=False),encoding="utf-8-sig")
        self.original={str(p):file_hash(p) for p in [self.source,self.target,self.link,self.gold,word,linked]}
    def api(self):return importlib.import_module("附注更新.外部范围")
    def save(self,name,value):
        p=self.work/(name+uuid.uuid4().hex+".json")
        with p.open("x",encoding="utf-8-sig") as f:json.dump(value,f,ensure_ascii=False,indent=2)
        return p
    def opinion(self):
        p=self.work/("意见模板"+uuid.uuid4().hex+".json");self.api().export_scope_review_template(self.source,self.link,p,gold_path=self.gold)
        doc=json.loads(p.read_text(encoding="utf-8-sig"));doc["reviewer"]={"name":"本地模拟高级AI","method":"advanced_ai","reviewed_at":"2026-09-20T11:00:00+08:00"}
        doc["areas"]=[{"sheet":"混合资料","range":"A1:B2","decision":"include","reason":"同一货币资金业务，新账户仍保留","evidence_cells":["A1","A2"]},
          {"sheet":"混合资料","range":"A3:B3","decision":"uncertain","reason":"原空位置尚未确定业务，保留","evidence_cells":[]},
          {"sheet":"混合资料","range":"A4:B5","decision":"out_of_scope","reason":"人员数量与资金余额为不同指标","evidence_cells":["A4","B4"],"semantic_comparison":{"source_business":"职工人数","target_business":"货币资金余额","difference_kind":"value_type","reason":"人数与货币金额不同，不由新职工身份排除","target_evidence":["货币资金","期末余额"]}}]
        return doc
    def preview(self,op=None):return self.api().preview_external_scope_review(self.source,self.link,self.save("模拟意见",op or self.opinion()),gold_path=self.gold)
    def test_template_has_complete_candidates_and_does_not_preapprove(self):
        p=self.work/"模板.json";self.api().export_scope_review_template(self.source,self.link,p,gold_path=self.gold)
        d=json.loads(p.read_text(encoding="utf-8-sig"));self.assertEqual(d["candidate_basis"],"source_grid_v1");self.assertEqual(d["candidate_count"],10)
        self.assertFalse(d["reviewer"]["name"]);self.assertTrue(all(a["decision"]=="uncertain" for a in d["areas"]))
        with self.assertRaises(FileExistsError):self.api().export_scope_review_template(self.source,self.link,p,gold_path=self.gold)
    def test_confirmed_partition_roundtrips_without_api_or_source_mutation(self):
        api=self.api()
        with patch("urllib.request.build_opener",side_effect=AssertionError("禁止网络")):
            preview=self.preview();self.assertEqual(preview["coverage"],{"candidate_count":10,"selected_count":6,"out_of_scope_count":4})
            with self.assertRaises(ValueError):api.apply_external_scope_review(preview,self.work/"无确认")
            result=api.apply_external_scope_review(preview,self.work/"成果",confirmed=True,confirm_method="本地模拟确认")
            selection=json.loads(Path(result["files"]["b_scope"]).read_text(encoding="utf-8-sig"))
            self.assertEqual(selection["review_method"],"external_scope_review");self.assertTrue(all("round" not in a for a in selection["areas"]))
            snap=read_workbook(self.source);before=copy.deepcopy(snap)
            target=build_target_meanings(json.loads(self.link.read_text(encoding="utf-8-sig")),read_workbook(self.target,context=json.loads(self.link.read_text(encoding="utf-8-sig"))["context"]["tables"]))
            checked=api.read_verified_scope(selection,snap,target);self.assertEqual(checked,selection);self.assertEqual(snap,before)
            self.assertIn("B2",checked["selected_cells"]["混合资料"]);self.assertIn("B3",checked["selected_cells"]["混合资料"])
        self.assertEqual(self.original,{p:file_hash(p) for p in self.original})
    def test_rejects_omission_overlap_wrong_evidence_and_dimension_filter(self):
        changes=[lambda d:d["areas"].pop(1),lambda d:d["areas"].append(copy.deepcopy(d["areas"][0])),lambda d:d.update(source_hash="0"*64),lambda d:d.update(target_hash="0"*64),lambda d:d.update(candidate_hash="0"*64),lambda d:d["areas"][2].update(evidence_cells=["B5"]),lambda d:d["areas"][2]["semantic_comparison"].update(difference_kind="counterparty"),lambda d:d["areas"][2]["semantic_comparison"].update(target_evidence=["不存在的指标"]),lambda d:d["areas"][2].update(round=1)]
        for change in changes:
            d=self.opinion();change(d)
            with self.subTest(d=d),self.assertRaises(ValueError):self.preview(d)
    def test_preview_and_saved_proof_tampering_are_rejected(self):
        api=self.api();p=self.preview();bad=copy.deepcopy(p);bad["coverage"]["selected_count"]=1
        with self.assertRaises(ValueError):api.apply_external_scope_review(bad,self.work/"拒绝",confirmed=True,confirm_method="本地模拟确认")
        result=api.apply_external_scope_review(p,self.work/"成果",confirmed=True,confirm_method="本地模拟确认")
        selected=json.loads(Path(result["files"]["b_scope"]).read_text(encoding="utf-8-sig"));snap=read_workbook(self.source);target=selected["target_meanings"]
        for change in [lambda s:s["selected_cells"]["混合资料"].remove("B3"),lambda s:s["external_scope_review"].update(confirm_method="假确认"),lambda s:s["areas"][2].update(reason="替换审核意见")]:
            bad=copy.deepcopy(selected);change(bad)
            with self.subTest(bad=bad),self.assertRaises(ValueError):api.read_verified_scope(bad,snap,target)
    def test_derived_blank_can_be_excluded_but_forged_blank_is_rejected(self):
        source=self.work/"本地虚构人数空白.xlsx";book=Workbook();sheet=book.active;sheet.title="混合资料"
        for row in [["货币资金","期末余额"],["新银行账户",0],[None,None],["人数统计","人数"],["新职工",None]]:sheet.append(row)
        book.save(source);book.close();self.source=source
        api=self.api();p=self.preview();result=api.apply_external_scope_review(p,self.work/"含派生格成果",confirmed=True,confirm_method="本地模拟确认")
        selection=json.loads(Path(result["files"]["b_scope"]).read_text(encoding="utf-8-sig"));snapshot=read_workbook(source)
        self.assertNotIn("B5",snapshot["sheets"][0]["cells"]);materialize_candidate_blanks(snapshot)
        self.assertTrue(snapshot["sheets"][0]["cells"]["B5"]["derived_blank"])
        self.assertNotIn("B5",api.read_verified_scope(selection,snapshot,selection["target_meanings"])["selected_cells"]["混合资料"])
        snapshot["sheets"][0]["cells"]["B5"]["value"]=123
        with self.assertRaises(ValueError):api.read_verified_scope(selection,snapshot,selection["target_meanings"])

    def test_legacy_scope_is_not_reinterpreted_as_external_review(self):
        self.assertIsNone(self.api().read_verified_scope({"areas":[{"round":1}]},read_workbook(self.source),{"disclosures":["货币资金"]}))

if __name__=="__main__":unittest.main()
