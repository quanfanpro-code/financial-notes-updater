import importlib
import unittest
import uuid
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.workbook.defined_name import DefinedName
ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/"测试结果"/("表格_"+uuid.uuid4().hex[:8])
WORK.mkdir(parents=True,exist_ok=True)
class WorkbookTests(unittest.TestCase):
    def api(self):
        try: return importlib.import_module("附注更新.表格")
        except ImportError: self.fail("尚未实现真实Excel读取与语义勾稽")
    def make(self,name,value):
        w=Workbook();s=w.active;s.title="附注"
        s.append(["项目","期末"]);s.append(["库存现金",value])
        s["B2"].font=Font(name="宋体",bold=True)
        w.defined_names.add(DefinedName("TA_001",attr_text="'附注'!$A$1:$B$2"))
        p=WORK/name;w.save(p);return p
    def mapping(self,cell="B2",scale=1,**dims):
        return {"sheet":"附注","cell":cell,"slot_id":"C-N087-T001-S001","dimensions":{"currency":"CNY","unit":"元","scale":scale,**dims},"value_type":"monetary","semantic_field":"库存现金期末余额"}
    def cash_format_fixture(self):
        from openpyxl.styles import Alignment, Border, Side, PatternFill, Protection
        api=self.api();token=uuid.uuid4().hex[:8];a=WORK/("金额显示A_"+token+".xlsx");b=WORK/("金额显示B_"+token+".xlsx")
        w=Workbook();s=w.active;s.title="附注"
        for address,value in {"B2":"项目","D2":"期初余额","B4":"银行存款","D4":1200,"B5":"其他货币资金","B7":"合计","D7":1200,"B10":"另一表","D10":2}.items():s[address]=value
        for address in ("D4","D7"):s[address].number_format="#,##0.00"
        s["D10"].number_format="0.0000"
        c=s["D5"];c.number_format="@";c.font=Font(name="宋体",bold=True,size=11)
        c.alignment=Alignment(horizontal="right",vertical="center",wrap_text=True)
        c.border=Border(bottom=Side(style="thin",color="FF112233"));c.fill=PatternFill("solid",fgColor="FFEEDDCC")
        c.protection=Protection(locked=False)
        w.save(a);w.close()
        w=Workbook();s=w.active;s.title="附注"
        for address,value in {"B4":"银行存款","C4":1200,"B6":"其他货币资金","C6":992997,"B7":"合计","C7":994197}.items():s[address]=value
        s["C6"].number_format="0.0000"  # 不得从来源格式照抄。
        w.save(b);w.close()
        mappings=[]
        for address,slot in (("D4","C-N087-T001-S004"),("D5","C-N087-T001-S006"),("D7","C-N087-T001-S010")):
            mappings.append({**self.mapping(address,period="2024-12-31"),"slot_id":slot,"scope":"standalone","reviewed":True,"table_range":"B2:D8"})
        sources=[{**m,"cell":address,"table_range":"B2:C7"} for m,address in zip(mappings,("C4","C6","C7"))]
        return a,b,api.read_workbook(a),api.read_workbook(b),mappings,sources

    def test_blank_monetary_cell_borrows_only_unique_confirmed_target_column_format(self):
        api=self.api();a,b,sa,sb,am,bm=self.cash_format_fixture();ah=api.file_hash(a);bh=api.file_hash(b)
        plan=api.build_update_plan(sa,sb,am,bm);entry=next(e for e in plan["updates"] if e["cell"]=="D5")
        self.assertFalse(plan["issues"]);self.assertEqual(entry["value"],992997)
        proof=entry["number_format_update"]
        self.assertEqual(proof["old_format"],"@");self.assertEqual(proof["new_format"],"#,##0.00")
        self.assertEqual(proof["source_a_hash"],sa["sha256"]);self.assertEqual(proof["table_range"],"B2:D8")
        self.assertEqual({s["cell"] for s in proof["sources"]},{"D4","D7"})
        self.assertTrue(all(s["number_format"]=="#,##0.00" for s in proof["sources"]))
        before=load_workbook(a);old_style=before["附注"]["D5"]._style.__copy__();before.close()
        output=WORK/("金额格式结果_"+uuid.uuid4().hex[:8]+".xlsx");api.write_updated_workbook(a,output,plan)
        check=load_workbook(output);cell=check["附注"]["D5"]
        self.assertEqual(cell.value,992997);self.assertEqual(cell.number_format,"#,##0.00")
        old_style.numFmtId=cell._style.numFmtId;self.assertEqual(cell._style,old_style)
        self.assertEqual(check["附注"]["D10"].number_format,"0.0000");check.close()
        self.assertEqual(api.file_hash(a),ah);self.assertEqual(api.file_hash(b),bh)

    def test_monetary_format_does_not_cross_business_or_borrow_unconfirmed_empty_wrong_units(self):
        import copy
        api=self.api();a,b,sa,sb,am,bm=self.cash_format_fixture()
        variants=[]
        for key,value in (("value_type","number"),("reviewed",False),("table_range","B2:D10")):
            altered=copy.deepcopy(am)
            for m in (altered[0],altered[2]):m[key]=value
            variants.append((copy.deepcopy(sa),altered))
        for key,value in (("currency","USD"),("unit","万元"),("scale",10000),("period","2025-12-31")):
            altered=copy.deepcopy(am)
            for m in (altered[0],altered[2]):m["dimensions"][key]=value
            variants.append((copy.deepcopy(sa),altered))
        empty=copy.deepcopy(sa)
        for cell in ("D4","D7"):empty["sheets"][0]["cells"][cell]["value"]=None
        variants.append((empty,copy.deepcopy(am)))
        for index,(snapshot,mappings) in enumerate(variants):
            with self.subTest(index=index):
                entry=next(e for e in api.build_update_plan(snapshot,sb,mappings,bm)["updates"] if e["cell"]=="D5")
                self.assertNotIn("number_format_update",entry)

    def test_monetary_format_conflicts_percent_dates_and_existing_format_are_not_changed(self):
        import copy
        api=self.api();a,b,sa,sb,am,bm=self.cash_format_fixture()
        for formats in (("0.00","#,##0.00"),("0.00%","0.00%"),("yyyy-mm-dd","yyyy-mm-dd"),("General","@")):
            snapshot=copy.deepcopy(sa)
            for address,fmt in zip(("D4","D7"),formats):snapshot["sheets"][0]["cells"][address]["number_format"]=fmt
            entry=next(e for e in api.build_update_plan(snapshot,sb,am,bm)["updates"] if e["cell"]=="D5")
            self.assertNotIn("number_format_update",entry)
        for current in ("0.000", "#,##0", "yyyy-mm-dd"):
            snapshot=copy.deepcopy(sa);snapshot["sheets"][0]["cells"]["D5"]["number_format"]=current
            entry=next(e for e in api.build_update_plan(snapshot,sb,am,bm)["updates"] if e["cell"]=="D5")
            self.assertNotIn("number_format_update",entry)
        snapshot=copy.deepcopy(sa);snapshot["sheets"][0]["cells"]["D5"]["value"]=0
        entry=next(e for e in api.build_update_plan(snapshot,sb,am,bm)["updates"] if e["cell"]=="D5")
        self.assertNotIn("number_format_update",entry)

    def test_writer_rejects_format_proof_that_disagrees_with_original_target_styles(self):
        import copy
        api=self.api();a,b,sa,sb,am,bm=self.cash_format_fixture();plan=api.build_update_plan(sa,sb,am,bm)
        entry=next(e for e in plan["updates"] if e["cell"]=="D5");self.assertIn("number_format_update",entry)
        for fmt in ("0.00%","0.0000"):
            altered=copy.deepcopy(plan)
            next(e for e in altered["updates"] if e["cell"]=="D5")["number_format_update"]["new_format"]=fmt
            with self.assertRaisesRegex(ValueError,"格式"):
                api.write_updated_workbook(a,WORK/("拒绝格式_"+uuid.uuid4().hex[:8]+".xlsx"),altered)

    def test_source_grid_adds_unformatted_blank_without_claiming_business(self):
        import copy
        api=self.api();w=Workbook();s=w.active;s.title="附注"
        s.append(["项目","期末"]);s.append(["甲",10]);s.append(["乙"])
        path=WORK/("候选无格式空白_"+uuid.uuid4().hex[:8]+".xlsx");w.save(path);w.close()
        before=api.file_hash(path);snapshot=api.read_workbook(path);original=copy.deepcopy(snapshot)
        self.assertNotIn("B3",snapshot["sheets"][0]["cells"])
        added=api.materialize_candidate_blanks(snapshot)
        self.assertEqual(added,{"附注":["B3"]})
        cell=api.cell_at(snapshot,"附注","B3");proof=cell["blank_evidence"]
        self.assertTrue(cell["derived_blank"]);self.assertIsNone(cell["value"])
        self.assertIsNone(cell["formula"]);self.assertIsNone(cell["cached_value"])
        self.assertEqual((proof["kind"],proof["source_hash"],proof["sheet"],proof["cell"],proof["bounds"]),
                         ("source_grid",before,"附注","B3","A1:B3"))
        self.assertFalse({"note_ids","scope","table_semantic"}&set(proof))
        self.assertNotIn("confirmed_business_ranges",snapshot["sheets"][0])
        self.assertEqual(api.materialize_candidate_blanks(snapshot),{})
        check=copy.deepcopy(snapshot);del check["sheets"][0]["cells"]["B3"]
        self.assertEqual(check,original);self.assertEqual(api.file_hash(path),before)
        # 进入候选并不把来源空白变成可清除旧金额的零。
        target=copy.deepcopy(snapshot);target["sheets"][0]["cells"]["B3"]["value"]=7
        plan=api.build_update_plan(target,snapshot,[self.mapping("B3")],[self.mapping("B3")])
        self.assertEqual([x["kind"] for x in plan["issues"]],["source_unknown_blank"])
        self.assertFalse(plan["updates"])

    def test_source_grid_respects_word_omissions_physical_range_merges_and_hidden(self):
        import copy
        api=self.api();w=Workbook();s=w.active;s.title="附注"
        s["B1"]="本表标题";s["B2"]="项目";s["C2"]="期末";s["D2"]="期初"
        s["B3"]="甲";s["B4"]="合计";s["D4"]=10;s["F4"]="表外说明";s.merge_cells("B4:C4")
        s.row_dimensions[3].hidden=True;s.column_dimensions["C"].hidden=True
        w.create_sheet("空页");s.sheet_state="hidden"
        path=WORK/("候选Word边界_"+uuid.uuid4().hex[:8]+".xlsx");w.save(path);w.close()
        context=[{"table_id":"/word/document.xml#body/1","range_name":"TA_001","sheet":"附注",
            "first_row":2,"first_column":2,"row_count":3,"column_count":3,
            "structural_omissions":[{"sheet":"附注","range":"D3","reason":"Word 原表该位置没有单元格"}]}]
        snapshot=api.read_workbook(path,context);original=copy.deepcopy(snapshot);before=api.file_hash(path)
        self.assertEqual(api.materialize_candidate_blanks(snapshot),{"附注":["C3"]})
        cell=api.cell_at(snapshot,"附注","C3")
        for flag in ("hidden_row","hidden_column","hidden_sheet","hidden"):self.assertTrue(cell[flag])
        for address in ("A3","D3","C4","E3"):self.assertNotIn(address,snapshot["sheets"][0]["cells"])
        self.assertEqual(snapshot["sheets"][1]["cells"],{})
        del snapshot["sheets"][0]["cells"]["C3"]
        self.assertEqual(snapshot,original);self.assertEqual(api.file_hash(path),before)

    def test_source_grid_keeps_empty_merge_anchor_and_existing_formula(self):
        import copy
        api=self.api();w=Workbook();s=w.active;s.title="附注"
        s["A1"]="说明";s["A3"]="金额";s["D3"]="=1+2";s.merge_cells("B2:C2")
        path=WORK/("候选合并_"+uuid.uuid4().hex[:8]+".xlsx");w.save(path);w.close()
        snapshot=api.read_workbook(path);original=copy.deepcopy(snapshot["sheets"][0]["cells"])
        added=api.materialize_candidate_blanks(snapshot)
        self.assertIn("B2",added["附注"]);self.assertNotIn("C2",snapshot["sheets"][0]["cells"])
        for address,cell in original.items():self.assertEqual(snapshot["sheets"][0]["cells"][address],cell)
        self.assertEqual(snapshot["sheets"][0]["cells"]["D3"]["formula"],"=1+2")

    def test_source_grid_rejects_abnormal_boundaries_and_missing_source_identity(self):
        import copy
        api=self.api();snapshot=api.read_workbook(self.make("候选边界检查.xlsx",1))
        for key,value,message in (("max_row",500001,"50万"),("max_column",True,"边界"),("max_row",0,"边界")):
            with self.subTest(key=key,value=value):
                changed=copy.deepcopy(snapshot);changed["sheets"][0][key]=value;before=copy.deepcopy(changed)
                with self.assertRaisesRegex(api.WorkbookError,message):api.materialize_candidate_blanks(changed)
                self.assertEqual(changed,before)
        changed=copy.deepcopy(snapshot);changed["sha256"]=""
        with self.assertRaisesRegex(api.WorkbookError,"来源"):api.materialize_candidate_blanks(changed)

    def test_visible_grid_keeps_business_blank_and_formula(self):
        api=self.api();p=self.make("读取.xlsx",None);w=load_workbook(p);s=w.active
        s["B2"].number_format="#,##0.00";s["B3"]="=SUM(B2:B2)"
        s["D1"]="隐藏";s.column_dimensions["D"].hidden=True
        s.merge_cells("A4:B4");s["A4"]="说明";w.save(WORK/"读取2.xlsx")
        got=api.read_workbook(WORK/"读取2.xlsx");cells=got["sheets"][0]["cells"]
        self.assertIn("B2",cells);self.assertIn("D1",cells);self.assertTrue(cells["D1"]["hidden_column"]);self.assertNotIn("B4",cells)
        self.assertEqual(cells["B3"]["formula"],"=SUM(B2:B2)")
    def test_word_omissions_stay_out_of_candidates_and_business_blanks(self):
        import json
        api=self.api();w=Workbook();s=w.active;s.title="附注"
        s["B2"]="项目";s["C2"]="期末";s["D2"]="期初";s["C3"]=8;s["D4"]=12
        s["B3"].font=Font(name="宋体")
        w.defined_names.add(DefinedName("TA_001",attr_text="'附注'!$B$2:$D$3"))
        path=WORK/"Word行首省略.xlsx";w.save(path);w.close()
        omission={"sheet":"附注","range":"B3","reason":"Word 原表该位置没有单元格"}
        context=[{"sheet":"附注","range_name":"TA_001","structural_omissions":[omission]},"金额单位：元"]
        before=api.file_hash(path);snapshot=api.read_workbook(path,context=context)
        sheet=snapshot["sheets"][0]
        self.assertEqual(sheet["structural_omissions"],[omission])
        self.assertNotIn("B3",sheet["cells"])
        self.assertIn("D3",sheet["cells"]);self.assertIsNone(sheet["cells"]["D3"]["value"])
        table={"range":"B2:D4","table_semantic":"余额明细","note_ids":["C-N087"],
               "scope":"consolidated","header_cells":["B2","C2"]}
        added=api.materialize_business_blanks(snapshot,"附注",table)
        self.assertIn("C4",added);self.assertNotIn("B3",added)
        restored=api.read_workbook(path,context=json.loads(json.dumps(snapshot["context"])))
        api.materialize_business_blanks(restored,"附注",table)
        self.assertEqual(restored["sheets"][0]["cells"],sheet["cells"])
        self.assertEqual(restored["sheets"][0]["structural_omissions"],[omission])
        self.assertIn("B3",api.read_workbook(path)["sheets"][0]["cells"])
        self.assertEqual(api.file_hash(path),before)

    def test_word_physical_range_limits_only_new_blanks(self):
        api=self.api();w=Workbook();s=w.active;s.title="附注"
        s["B1"]="本表标题";s["B2"]="项目";s["C2"]="期末";s["D2"]="期初"
        s["B3"]="客户甲";s["B4"]="合计";s["D4"]=10;s["F4"]="真实表外说明"
        s.merge_cells("B4:C4")
        path=WORK/"Word真实边界无名称.xlsx";w.save(path);w.close()
        omission={"sheet":"附注","range":"D3","reason":"Word 原表该位置没有单元格"}
        context={"tables":[{"table_id":"/word/document.xml#body/1","range_name":"TA_001","sheet":"附注",
            "first_row":2,"first_column":2,"row_count":3,"column_count":3,"structural_omissions":[omission]}]}
        snapshot=api.read_workbook(path,context);sheet=snapshot["sheets"][0]
        self.assertNotIn("C3",sheet["cells"])
        original=dict(sheet["cells"])
        table={"range":"B1:F4","table_semantic":"款项余额","note_ids":["C-N020"],"scope":"standalone","header_cells":["B2","C2"]}
        added=api.materialize_business_blanks(snapshot,"附注",table)
        self.assertEqual(added,["C3"])
        self.assertEqual(sheet["word_table_ranges"],[{"table_id":"/word/document.xml#body/1","range_name":"TA_001","range":"B2:D4"}])
        for address,cell in original.items():self.assertEqual(sheet["cells"][address],cell)
        self.assertNotIn("D3",sheet["cells"]);self.assertNotIn("C4",sheet["cells"])
        self.assertEqual(sheet["cells"]["B1"]["value"],"本表标题")
        self.assertEqual(sheet["cells"]["F4"]["value"],"真实表外说明")

    def test_word_physical_range_uses_current_name_after_move(self):
        api=self.api();w=Workbook();s=w.active;s.title="附注"
        s["B1"]="原段落";s["B5"]="项目";s["C5"]="期末";s["B6"]="客户甲";s["F8"]="后文"
        w.defined_names.add(DefinedName("TA_001",attr_text="'附注'!$B$5:$D$7"))
        path=WORK/"Word命名区域已移动.xlsx";w.save(path);w.close()
        context=[{"table_id":"/word/document.xml#body/1","range_name":"TA_001","sheet":"附注",
                  "first_row":2,"first_column":2,"row_count":2,"column_count":2}]
        snapshot=api.read_workbook(path,context);sheet=snapshot["sheets"][0]
        table={"range":"B1:F8","table_semantic":"款项余额","note_ids":["C-N020"],"scope":"standalone","header_cells":["B5","C5"]}
        self.assertEqual(api.materialize_business_blanks(snapshot,"附注",table),[])
        self.assertEqual(sheet["word_table_ranges"][0]["range"],"B5:D7")
        self.assertIn("D7",sheet["cells"]);self.assertIsNone(sheet["cells"]["D7"]["value"])
        self.assertNotIn("C2",sheet["cells"]);self.assertNotIn("E6",sheet["cells"])

    def test_generic_excel_without_word_identity_can_still_materialize(self):
        api=self.api();w=Workbook();s=w.active;s.title="附注"
        s["A1"]="项目";s["B1"]="期末";s["A2"]="客户甲";s["C2"]=1
        path=WORK/"普通更新数据补格.xlsx";w.save(path);w.close()
        snapshot=api.read_workbook(path,context=[{"sheet":"附注","range_name":"TA_001"}])
        table={"range":"A1:C2","table_semantic":"款项余额","note_ids":["C-N020"],"scope":"standalone","header_cells":["A1","B1"]}
        self.assertIn("B2",api.materialize_business_blanks(snapshot,"附注",table))
        self.assertNotIn("word_table_ranges",snapshot["sheets"][0])

    def test_word_omissions_reject_values_formulas_merges_and_bad_metadata(self):
        api=self.api()
        omission={"sheet":"附注","range":"B3","reason":"Word 原表该位置没有单元格"}
        for name,value,merge in [("数值",0,None),("公式","=1+2",None),
                                 ("合并主格",None,"B3:C3"),("合并从格",None,"A3:B3")]:
            with self.subTest(name=name):
                w=Workbook();s=w.active;s.title="附注";s["B2"]="项目";s["C4"]=1
                if value is not None:s["B3"]=value
                if merge:s.merge_cells(merge)
                path=WORK/("省略冲突_"+name+".xlsx");w.save(path);w.close()
                with self.assertRaisesRegex(api.WorkbookError,"省略"):
                    api.read_workbook(path,context={"tables":[{"structural_omissions":[omission]}]})
        path=self.make("省略范围核验.xlsx",1)
        for invalid in [dict(omission,sheet="不存在"),dict(omission,range="B"),
                        dict(omission,range="XFE1"),dict(omission,range="B0"),
                        dict(omission,range="B3:A2")]:
            with self.subTest(invalid=invalid):
                with self.assertRaisesRegex(api.WorkbookError,"省略"):
                    api.read_workbook(path,context=[{"structural_omissions":[invalid]}])

    def test_standalone_scope_stays_distinct_and_preserves_blank_evidence(self):
        api=self.api();a=self.make("单户附注.xlsx",1);b=self.make("单户更新.xlsx",2)
        am={**self.mapping(),"scope":"standalone","definition_scope":"consolidated"}
        bm={**am,"scope":"consolidated"}
        self.assertNotEqual(api.semantic_key(am),api.semantic_key(bm))
        self.assertTrue(api.build_update_plan(api.read_workbook(a),api.read_workbook(b),[am],[bm])["issues"])
        plan=api.build_update_plan(api.read_workbook(a),api.read_workbook(b),[am],[am])
        self.assertEqual(plan["issues"],[]);self.assertEqual(plan["updates"][0]["value"],2)
        w=Workbook();s=w.active;s.title="附注";s["A1"]="项目";s["B1"]="期末余额";s["A2"]="库存现金";s["C2"]=1
        source=WORK/"单户业务空白.xlsx";w.save(source);w.close()
        snapshot=api.read_workbook(source)
        table={"range":"A1:C2","table_semantic":"资金余额","note_ids":["C-N087"],
               "scope":"standalone","header_cells":["A1","B1"]}
        self.assertIn("B2",api.materialize_business_blanks(snapshot,"附注",table))
        self.assertEqual(snapshot["sheets"][0]["confirmed_business_ranges"][0]["scope"],"standalone")

    def test_same_semantics_transposed_source_keeps_given_value_without_scale_conversion(self):
        api=self.api();a=self.make("目标.xlsx",7)
        w=Workbook();s=w.active;s.title="新底稿";s.append(["指标","库存现金"]);s.append(["年末数",2.5])
        b=WORK/"来源.xlsx";w.save(b);am=self.mapping();bm={**self.mapping(scale=10000),"sheet":"新底稿"}
        plan=api.build_update_plan(api.read_workbook(a),api.read_workbook(b),[am],[bm])
        self.assertEqual(plan["issues"],[]);self.assertEqual(plan["updates"][0]["value"],2.5)
        output=WORK/"更新.xlsx";api.write_updated_workbook(a,output,plan)
        out=load_workbook(output);self.assertEqual(out.active["B2"].value,2.5)
        self.assertEqual(out.active["B2"].font.name,"宋体")
        self.assertEqual(load_workbook(a).active["B2"].value,7)
    def test_conflicting_sources_do_not_write(self):
        api=self.api();a=self.make("冲突A.xlsx",8);b=self.make("冲突B.xlsx",2)
        w=load_workbook(b);w.active["B3"]=3;w.save(WORK/"冲突B2.xlsx")
        plan=api.build_update_plan(api.read_workbook(a),api.read_workbook(WORK/"冲突B2.xlsx"),[self.mapping()],[self.mapping(),self.mapping(cell="B3")])
        self.assertEqual(plan["updates"],[]);self.assertTrue(plan["issues"])
    def test_blank_is_not_zero(self):
        api=self.api();a=self.make("空白A.xlsx",4);b=self.make("空白B.xlsx",None)
        plan=api.build_update_plan(api.read_workbook(a),api.read_workbook(b),[self.mapping()],[self.mapping()])
        self.assertFalse(plan["updates"])
        self.assertEqual(plan["issues"][0]["kind"],"source_unknown_blank")
    def test_unknown_source_and_original_blank_are_preserved_without_update(self):
        api=self.api();a=self.make("双方原空白A.xlsx",None);b=self.make("双方原空白B.xlsx",None)
        plan=api.build_update_plan(api.read_workbook(a),api.read_workbook(b),[self.mapping()],[self.mapping()])
        self.assertFalse(plan["updates"]);self.assertFalse(plan["issues"]);self.assertFalse(plan["unplaced_sources"])
        self.assertEqual(len(plan["retained_blanks"]),1)
        self.assertIsNone(plan["retained_blanks"][0]["blank_sources"][0]["value"])

    def test_known_same_identity_source_is_used_but_blank_is_not_zero(self):
        api=self.api();a=self.make("明确加未知A.xlsx",7);b=self.make("明确加未知B.xlsx",0)
        book=load_workbook(b);book.active["A3"]="库存现金";book.active["B3"].number_format="0.00"
        source=WORK/"明确加未知B完整.xlsx";book.save(source);book.close()
        plan=api.build_update_plan(api.read_workbook(a),api.read_workbook(source),[self.mapping()],[self.mapping(),self.mapping(cell="B3")])
        self.assertFalse(plan["issues"]);self.assertFalse(plan["unplaced_sources"])
        self.assertEqual(plan["updates"][0]["value"],0)
        self.assertEqual(plan["updates"][0]["sources"],[{"sheet":"附注","cell":"B2","value":0}])
        self.assertEqual(plan["updates"][0]["blank_sources"],[{"sheet":"附注","cell":"B3","value":None}])

    def test_original_blank_without_source_is_preserved_and_not_counted_as_update(self):
        api=self.api();a=self.make("原空白无来源.xlsx",None);b=self.make("另一业务来源.xlsx",5)
        plan=api.build_update_plan(api.read_workbook(a),api.read_workbook(b),[self.mapping()],[self.mapping(counterparty="另一对象")])
        self.assertFalse(plan["issues"]);self.assertFalse(plan["updates"])
        self.assertEqual(len(plan["retained_blanks"]),1)
        self.assertEqual(plan["retained_blanks"][0]["semantic_key"],api.semantic_key(self.mapping()))
        self.assertIsNone(plan["retained_blanks"][0]["old_value"])
        self.assertEqual(len(plan["unplaced_sources"]),1)

    def test_nonblank_or_formula_target_cannot_be_retained_as_blank_without_source(self):
        api=self.api();b=self.make("缺来源保护B.xlsx",9)
        for index,value in enumerate((0,12,"——"," ","=SUM(1,2)")):
            a=self.make("缺来源保护A"+str(index)+".xlsx",value)
            plan=api.build_update_plan(api.read_workbook(a),api.read_workbook(b),[self.mapping()],[])
            self.assertFalse(plan["retained_blanks"]);self.assertTrue(plan["issues"])

    def test_blank_target_does_not_hide_existing_source_read_failure(self):
        api=self.api();a=self.make("原空白来源失败A.xlsx",None);b=self.make("原空白来源失败B.xlsx","=NOW()")
        plan=api.build_update_plan(api.read_workbook(a),api.read_workbook(b),[self.mapping()],[self.mapping()])
        self.assertFalse(plan["retained_blanks"]);self.assertTrue(plan["issues"]);self.assertFalse(plan["updates"])

    def test_distinct_open_objects_never_merge(self):
        api=self.api();a=self.make("对象A.xlsx",4);b=self.make("对象B.xlsx",9)
        p=api.build_update_plan(api.read_workbook(a),api.read_workbook(b),[self.mapping(counterparty="甲公司")],[self.mapping(counterparty="乙公司")])
        self.assertFalse(p["updates"]);self.assertTrue(p["issues"])
    def test_source_formula_without_result_is_not_number(self):
        api=self.api();a=self.make("公式A.xlsx",4);b=self.make("公式B.xlsx","=NOW()")
        p=api.build_update_plan(api.read_workbook(a),api.read_workbook(b),[self.mapping()],[self.mapping()])
        self.assertFalse(p["updates"]);self.assertTrue(p["issues"])
    def test_insert_row_clones_style_and_expands_name(self):
        api=self.api();a=self.make("增行A.xlsx",4);output=WORK/"增行结果.xlsx"
        plan={"updates":[{"sheet":"附注","cell":"B3","value":12}],"issues":[],"operations":[{"range_name":"TA_001","axis":"row","action":"insert","index":2,"count":1,"template_index":1}]}
        api.write_updated_workbook(a,output,plan)
        w=load_workbook(output);self.assertEqual(w.active["B3"].value,12)
        self.assertEqual(w.active["B3"].font.name,"宋体")
        self.assertEqual(list(w.defined_names["TA_001"].destinations),[("附注","$A$1:$B$3")])
    def test_reject_source_overwrite(self):
        api=self.api();a=self.make("不可覆盖.xlsx",4)
        with self.assertRaises(ValueError):api.write_updated_workbook(a,a,{"updates":[],"operations":[]})

    def test_safe_sum_ignores_untrusted_cached_result(self):
        api=self.api();p=self.make("安全公式.xlsx",2);w=load_workbook(p);w.active["B3"]=4;w.active["B4"]="=SUM(B2:B3)*2";out=WORK/"安全公式输入.xlsx";w.save(out)
        snap=api.read_workbook(out);snap["sheets"][0]["cells"]["B4"]["cached_value"]=999
        self.assertEqual(api.actual_value(snap,"附注","B4"),12)

    def test_format_only_last_row_is_ignored(self):
        api=self.api();p=self.make("污染前.xlsx",None);w=load_workbook(p);w.active["B1048576"].font=Font(bold=True);q=WORK/"百万行格式污染.xlsx";w.save(q)
        s=api.read_workbook(q)["sheets"][0];self.assertEqual(s["max_row"],2);self.assertNotIn("B1048576",s["cells"]);self.assertIn("B2",s["cells"])

    def test_local_column_insert_keeps_other_table(self):
        api=self.api();p=self.make("插列A.xlsx",4);w=load_workbook(p);w.active["A5"]="另一张表";w.active["B5"]=99;w.defined_names.add(DefinedName("TA_002",attr_text="'附注'!$A$5:$B$5"));q=WORK/"插列输入.xlsx";w.save(q)
        output=WORK/"局部插列.xlsx";plan={"operations":[{"range_name":"TA_001","axis":"column","action":"insert","index":1,"template_index":1}],"labels":[{"sheet":"附注","cell":"B1","value":"新增列"}],"updates":[{"sheet":"附注","cell":"B2","value":8}]}
        api.write_updated_workbook(q,output,plan);got=load_workbook(output)
        self.assertEqual(got.active["B2"].value,8);self.assertEqual(got.active["C2"].value,4);self.assertEqual(got.active["B5"].value,99)
        self.assertEqual(list(got.defined_names["TA_001"].destinations),[("附注","$A$1:$C$2")]);self.assertEqual(list(got.defined_names["TA_002"].destinations),[("附注","$A$5:$B$5")])

    def test_row_insert_moves_other_names_and_formula_references(self):
        api=self.api();p=self.make("引用A.xlsx",4);w=load_workbook(p);w.active["B5"]="=B2*2";w.defined_names.add(DefinedName("TA_002",attr_text="'附注'!$A$5:$B$5"));q=WORK/"引用输入.xlsx";w.save(q)
        output=WORK/"引用结果.xlsx";api.write_updated_workbook(q,output,{"operations":[{"range_name":"TA_001","axis":"row","action":"insert","index":1,"template_index":1}]})
        got=load_workbook(output);self.assertEqual(got.active["B6"].value,"=B3*2");self.assertEqual(list(got.defined_names["TA_002"].destinations),[("附注","$A$6:$B$6")])

    def test_percent_display_keeps_given_value_and_currency_alias(self):
        api=self.api();a=self.make("比例A.xlsx",0.1);w=load_workbook(a);w.active["B2"].number_format="0.00%";a2=WORK/"比例A2.xlsx";w.save(a2);b=self.make("比例B.xlsx",25)
        am={**self.mapping(),"value_type":"percentage","dimensions":{"unit":"%","scale":1}};bm={**am,"semantic_field":"自由说明不同"}
        plan=api.build_update_plan(api.read_workbook(a2),api.read_workbook(b),[am],[bm]);self.assertEqual(plan["updates"][0]["value"],25)
        self.assertEqual(api.semantic_key(self.mapping(currency="RMB")),api.semantic_key(self.mapping(currency="CNY")))
        self.assertNotEqual(api.semantic_key(self.mapping(metric="土地面积")),api.semantic_key(self.mapping(metric="建筑面积")))

    def test_text_starting_equal_never_becomes_formula_and_hash_guard(self):
        api=self.api();a=self.make("文本A.xlsx",4);h=api.file_hash(a);output=WORK/"文本输出.xlsx"
        api.write_updated_workbook(a,output,{"source_a_hash":h,"updates":[{"sheet":"附注","cell":"B2","value":"=HYPERLINK(\"https://invalid\")"}]})
        cell=load_workbook(output).active["B2"];self.assertEqual(cell.data_type,"s")
        with self.assertRaises(ValueError):api.write_updated_workbook(a,WORK/"过期.xlsx",{"source_a_hash":"错误hash","updates":[]})

    def test_delete_row_shrinks_name_and_rejects_cross_merge(self):
        api=self.api();a=self.make("删除A.xlsx",4);w=load_workbook(a);w.active.append(["其他",8]);w.defined_names["TA_001"].attr_text="'附注'!$A$1:$B$3";p=WORK/"删除输入.xlsx";w.save(p);out=WORK/"删除结果.xlsx"
        api.write_updated_workbook(p,out,{"operations":[{"range_name":"TA_001","axis":"row","action":"delete","index":1,"count":1,"template_index":2}]})
        got=load_workbook(out);self.assertEqual(got.active["B2"].value,8);self.assertEqual(list(got.defined_names["TA_001"].destinations),[("附注","$A$1:$B$2")])
        w.active.merge_cells("A2:A3");merged=WORK/"合并输入.xlsx";w.save(merged)
        with self.assertRaises(ValueError):api.write_updated_workbook(merged,WORK/"拒绝合并.xlsx",{"operations":[{"range_name":"TA_001","axis":"row","action":"insert","index":2,"template_index":0}]})
    def test_missing_source_has_distinct_issue_kind(self):
        api=self.api();a=self.make("缺来源A.xlsx",4);b=self.make("缺来源B.xlsx",8)
        plan=api.build_update_plan(api.read_workbook(a),api.read_workbook(b),[self.mapping(counterparty="甲")],[self.mapping(counterparty="乙")])
        self.assertEqual(plan["issues"][0]["kind"],"missing_source")

    def test_monetary_blank_display_is_literal_and_records_field(self):
        api=self.api();a=self.make("横线A.xlsx",4);b=self.make("横线B.xlsx","—")
        plan=api.build_update_plan(api.read_workbook(a),api.read_workbook(b),[self.mapping()],[self.mapping()])
        self.assertEqual(plan["issues"],[]);self.assertEqual(plan["updates"][0]["value"],"—");self.assertEqual(plan["updates"][0]["semantic_field"],"库存现金期末余额")

    def test_hidden_sheet_and_rows_remain_available_for_semantics(self):
        api=self.api();a=self.make("隐藏前.xlsx",4);w=load_workbook(a);w.active.row_dimensions[2].hidden=True;s=w.create_sheet("隐藏业务表");s["A1"]="财务数据";s.sheet_state="hidden";p=WORK/"含隐藏业务.xlsx";w.save(p)
        snap=api.read_workbook(p);self.assertTrue(api.cell_at(snap,"附注","B2")["hidden_row"]);hidden=next(s for s in snap["sheets"] if s["name"]=="隐藏业务表");self.assertTrue(hidden["hidden"]);self.assertTrue(hidden["cells"]["A1"]["hidden"])
    def test_workbook_without_calculation_properties_can_update(self):
        api=self.api();a=self.make("无计算属性A.xlsx",4);w=load_workbook(a);w.calculation=None;p=WORK/"无计算属性输入.xlsx";w.save(p);out=WORK/"无计算属性结果.xlsx"
        api.write_updated_workbook(p,out,{"updates":[{"sheet":"附注","cell":"B2","value":6}]})
        self.assertEqual(load_workbook(out).active["B2"].value,6)
    def test_macro_archives_are_closed_before_garbage_collection(self):
        import gc
        import sys
        import zipfile
        api=self.api()
        plain=self.make("宏归档源.xlsx",4)
        source=WORK/"宏归档源.xlsm"
        payload=b"non-executable-macro-fixture"
        with zipfile.ZipFile(plain) as original,zipfile.ZipFile(source,"w") as target:
            for info in original.infolist():target.writestr(info,original.read(info.filename))
            target.writestr("xl/vbaProject.bin",payload)
        gc.collect()
        errors=[]
        previous=sys.unraisablehook
        sys.unraisablehook=lambda event:errors.append((event.exc_type.__name__,str(event.exc_value)))
        try:
            api.read_workbook(source)
            output=WORK/"宏归档结果.xlsm"
            api.write_updated_workbook(source,output,{"updates":[{"sheet":"附注","cell":"B2","value":9}]})
            self.assertEqual(api.cell_at(api.read_workbook(output),"附注","B2")["value"],9)
            with zipfile.ZipFile(output) as archive:self.assertEqual(archive.read("xl/vbaProject.bin"),payload)
            with self.assertRaises(ValueError):
                api.write_updated_workbook(source,WORK/"宏错误计划.xlsm",{"updates":[{"sheet":"不存在","cell":"B2","value":9}]})
            for _ in range(3):gc.collect()
        finally:
            sys.unraisablehook=previous
        self.assertEqual(errors,[],"宏归档析构阶段不应发生隐藏异常")

    def test_internal_structured_formula_recalculates_with_real_excel(self):
        from openpyxl.worksheet.table import Table
        api=self.api();w=Workbook();s=w.active;s.title="附注"
        s.append(["项目","余额"]);s.append(["甲",3]);s.append(["乙",7]);s.add_table(Table(displayName="Table1",ref="A1:B3"));s["D1"]="=SUM(Table1[余额])"
        source=WORK/"内部结构化公式.xlsx";output=WORK/"内部结构化公式_Excel重算.xlsx";w.save(source);w.close()
        before=api.file_hash(source);unverified=api.read_workbook(source)
        self.assertEqual(unverified["external_references"],[])
        with self.assertRaises(ValueError):api.actual_value(unverified,"附注","D1")
        api.recalculate_copy(source,output);verified=api.read_workbook(output)
        self.assertTrue(verified["recalculated"]);self.assertEqual(api.actual_value(verified,"附注","D1"),10)
        self.assertEqual(api.file_hash(source),before)
