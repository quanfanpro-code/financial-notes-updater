# coding: utf-8
"""用模拟原答复和合成文件检查输入、保存及机械搬运，不请求模型。"""
from datetime import date, datetime
import json
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from 附注更新.A合集 import _per_cell_answer, _records, direct_digest, save_direct_answer
from 附注更新.A语义阅读 import recognize_a
from 附注更新.分步流程 import run_step, _read
from 附注更新.流程 import save_json
from 附注更新.统一勾稽 import load_fact_mapping, _calculated_values
from 附注更新.表格 import file_hash
from 附注更新.统一语义 import json_date, json_dates
from 附注更新.统一识别 import UnifiedSemanticEngine
from 附注更新.模板更新编排 import direct_update_plan
from 测试 import 测试直接识别 as fixture


class 识别输入与保存测试(unittest.TestCase):
    def setUp(self):
        self.case=fixture.直接识别测试();self.case.setUp()
        self.standard={'metric':{'m1':{},'m2':{}},'dimension':{'currency':{},'asset_selection':{}}}

    def answer(self):
        return {f'合成表!{col}3':{'metric_id':'固定资产原值',
            'dimensions':{'entity':'合成企业','report_scope':'standalone','currency':'CNY','unit_scale':1,
                'period':{'kind':'instant','date':'2025-12-31'},
                'asset_selection':{'mode':'members','members':['设备'+col],'domain':'固定资产','completeness':'complete'}},
            'template_period_role':'closing'} for col in ('B','C')}

    def recognize(self):
        config={**self.case.config,'a_target_table':{'id':'本表','sheet':'合成表','range':'A1:C8','data_ranges':['B3:C3']}}
        with patch.object(UnifiedSemanticEngine,'request_text',return_value=json.dumps(self.answer(),ensure_ascii=False)):
            result=run_step('recognize_a',config)
        self.assertEqual(result['status'],'complete',result)
        return load_fact_mapping(result['files']['a_mapping'],{}),config,result

    def test_JSON代码块前后说明不丢逐格原字段(self):
        standard={'metric':{'m1':{}},'dimension':{'currency':{},'entity':{}}}
        row={'location':'S!A1','metric_id':'m1','dimensions':{'currency':'CNY'}}
        reply='说明文字\n```json\n'+json.dumps(row)+'\n```\n共同属性entity=某主体'
        parsed=_per_cell_answer(reply,standard,['S'])
        self.assertEqual(len(parsed['metrics']),1)
        self.assertEqual(parsed['metrics'][0]['original'],row)
        self.assertEqual([(x['id'],x['value']) for x in parsed['dimensions']],[('currency','CNY')])
        parsed=_per_cell_answer(json.dumps(row)+'\n说明文字',standard,['S'])
        self.assertEqual(parsed['metrics'][0]['original'],row)
        second={'location':'S!A2','metric_id':'m1','dimensions':{'currency':'USD'}}
        parsed=_per_cell_answer(reply+'\n```json\n'+json.dumps(second)+'\n```',standard,['S'])
        self.assertEqual({x['ranges'][0] for x in parsed['metrics']},{'S!A1','S!A2'})

    def test_编号后的说明不是编号且不据此推造维度值(self):
        raw='- location:S!A1；metric_id:null；dimensions:dimension_id=asset_selection，实际取值={mode:members,members:[“设备甲”]}；currency=CNY'
        parsed=_per_cell_answer(raw,self.standard,['S'])
        self.assertEqual(parsed['dimension_bindings'][0]['id'],'asset_selection')
        self.assertEqual(parsed['dimension_bindings'][0]['original'],raw)
        self.assertEqual([(v['id'],v['value']) for v in parsed['dimensions']],[('currency','CNY')])
        for identity in ['asset_selection.extra','asset_selection，currency']:
            bad=_per_cell_answer('- S!A1：dimension_id='+identity+'；currency=CNY',self.standard,['S'])
            self.assertFalse(bad['dimension_bindings'])
        structured=_per_cell_answer({'location':'S!A1','dimension_id':'asset_selection，说明','dimensions':{'currency':'CNY'}},self.standard,['S'])
        self.assertFalse(structured['dimension_bindings'])
        missing=_per_cell_answer('- S!A1：metric_id=null  \n依据：未提供标准。；currency=CNY',self.standard,['S'])
        self.assertIsNone(missing['unresolved'][0]['value'])

    def test_箭头只是原字段赋值分隔符且不吞入实际值(self):
        standard={'metric':{'m1':{}},'dimension':{'currency':{},'entity':{}}}
        reply='| location | metric_id | dimensions（维度ID → 实际值） |\n|---|---|---|\n| S!A1 | m1 | `entity`→原主体；`currency`→CNY |'
        parsed=_per_cell_answer(reply,standard,['S'])
        self.assertEqual(parsed['metrics'][0]['id'],'m1')
        self.assertEqual({x['id']:x['value'] for x in parsed['dimensions']},{'entity':'原主体','currency':'CNY'})
        parsed=_per_cell_answer(reply.replace('CNY','CNY；`currency`→USD'),standard,['S'])
        self.assertTrue(any(x.get('id')=='currency' for x in parsed['unresolved']))

    def test_同格维度名称清单不冒充已给出的维度实际值(self):
        standard={'metric':{'m1':{}},'dimension':{'currency':{},'entity':{}}}
        reply='| location | metric_id | dimensions（原始维度ID） | currency |\n|---|---|---|---|\n| S!A1 | m1 | `entity`, `currency` | CNY |'
        parsed=_per_cell_answer(reply,standard,['S'])
        self.assertEqual([(x['id'],x['value']) for x in parsed['dimensions']],[('currency','CNY')])
        self.assertFalse(parsed['unresolved'])

    def test_答复列数不齐仍保留明示标准编号且不猜其他列(self):
        from 附注更新.A合集 import direct_ingestion
        raw='| location | metric_id | currency | unit_scale |\n|---|---|---|---|\n| S!A1 | m1 | USD |\n| S!A2 | m2 | CNY | 1 |'
        parsed=_per_cell_answer(raw,self.standard,['S'])
        self.assertEqual([(x['ranges'][0],x['id']) for x in parsed['metrics']],[('S!A1','m1'),('S!A2','m2')])
        self.assertEqual([(x['ranges'][0],x['id'],x['value']) for x in parsed['dimensions']],[('S!A2','currency','CNY')])
        self.assertEqual(parsed['unresolved'][0]['original'],raw.splitlines()[2])
        receipt=direct_ingestion(parsed,{'S!A1:A2':['metric_id']},self.standard,['S'])
        self.assertFalse(receipt['complete'])
        self.assertEqual(receipt['未接入字段'],[])
        self.assertEqual(receipt['未确定字段'][0]['location'],'S!A1')
        self.assertIn('表头4列',receipt['未确定字段'][0]['reason'])

    def test_逐格对象及同一本格行原字段读取但不展开共同声明(self):
        raw='\n'.join(['{"answers":[',
            '{"note":"共同说明"},"template_period_role":"opening"},',
            '{"location":"S!A1","metric_id":"m1","dimensions":{"currency":"CNY"}},',
            '{"location":"S!A2","dimension_id":"asset_selection","dimensions":{"asset_selection":{"mode":"members","members":["甲"]}}},"template_period_role":"closing"},',
            '{"location":"S!A3","metric_id":"m1","dimensions":{"currency":"CNY"}}',
            ']}'])
        parsed=_per_cell_answer(raw,self.standard,['S'])
        self.assertEqual([r['ranges'] for r in parsed['metrics']],[['S!A1'],['S!A3']])
        self.assertEqual(parsed['dimension_bindings'][0]['id'],'asset_selection')
        self.assertEqual(parsed['dimension_bindings'][0]['ranges'],['S!A2'])
        roles=[(r['ranges'],r['value']) for r in parsed['dimensions'] if r['id']=='template_period_role']
        self.assertEqual(roles,[(['S!A2'],'closing')])
        self.assertEqual(next(r['value'] for r in parsed['dimensions'] if r['id']=='asset_selection'),{'mode':'members','members':['甲']})

    def test_未采用的另一主键不否定有效编号且原文保留(self):
        from 附注更新.A合集 import direct_ingestion
        for row,kind in [({'metric_id':'无','dimension_id':'asset_selection','dimensions':{'asset_selection':{'mode':'members','members':['甲']}}},'dimension_bindings'),
                         ({'metric_id':'m1','dimension_id':'不存在的编号','dimensions':{'currency':'CNY'}},'metrics')]:
            own={'location':'S!A1',**row}
            parsed=_per_cell_answer(own,self.standard,['S'])
            self.assertEqual(len(parsed[kind]),1)
            self.assertEqual(parsed[kind][0]['original'],own)
            self.assertTrue(parsed['unresolved'])
            receipt=direct_ingestion(parsed,{'S!A1':['metric_id','dimension_id']},self.standard,['S'])
            self.assertTrue(receipt['complete'],receipt)
            self.assertTrue(receipt['范围外未接入问题'])
        parsed=_per_cell_answer({'location':'S!A1','metric_id':'m1','dimension_id':'asset_selection','dimensions':{'currency':'CNY'}},self.standard,['S'])
        self.assertFalse(parsed['metrics']);self.assertFalse(parsed['dimension_bindings'])
        receipt=direct_ingestion(parsed,{'S!A1':['metric_id','dimension_id']},self.standard,['S'])
        self.assertFalse(receipt['complete'])

    def test_已选有效指标或维度时另一空主键不是漏交(self):
        from 附注更新.A合集 import direct_ingestion
        for row,kind in [({'location':'S!A1','metric_id':'m1','dimension_id':None},'metrics'),
                         ({'location':'S!A1','metric_id':None,'dimension_id':'asset_selection',
                           'dimensions':{'asset_selection':{'mode':'members','members':['甲']}}},'dimension_bindings')]:
            with self.subTest(kind=kind):
                parsed=_per_cell_answer(json.dumps(row,ensure_ascii=False),self.standard,['S'])
                self.assertEqual(len(parsed[kind]),1)
                self.assertFalse(parsed['unresolved'])
                self.assertEqual(parsed[kind][0]['original'],row)
                receipt=direct_ingestion(parsed,{'S!A1':['metric_id','dimension_id']},self.standard,['S'])
                self.assertTrue(receipt['complete'])
        parsed=_per_cell_answer({'location':'S!A1','metric_id':None,'dimension_id':None},self.standard,['S'])
        self.assertFalse(parsed['metrics']);self.assertFalse(parsed['dimension_bindings'])
        self.assertTrue(parsed['unresolved'])

    def test_重复字段四种写法保留其他格和正确字段(self):
        replies=[
            '{"S!A1":{"metric_id":"m1","metric_id":"m2","currency":"CNY"},"S!A2":{"metric_id":"m1"}}',
            '{"S!A1":{"metric_id":"m1","currency":"CNY","currency":"USD"}},{"S!A2":{"metric_id":"m1"}}',
            '- S!A1：metric_id=m1；dimensions={"currency":"CNY","currency":"USD"}\n- S!A2：metric_id=m1',
            '- S!A1：\n  - metric_id: `m1`\n  - dimensions: {currency: CNY, currency: USD}\n- S!A2：metric_id=m1']
        for reply in replies:
            with self.subTest(reply=reply):
                result=_per_cell_answer(reply,self.standard,['S'])
                self.assertTrue(any(r['ranges']==['S!A2'] for r in result['metrics']))
                self.assertTrue(result['unresolved'])
                self.assertFalse(any(r.get('value')=='USD' for r in result['dimensions']))
                self.assertFalse(any(r['id']=='m2' for r in result['metrics']))

    def test_JSON外层名称不影响本格原字段且不展开共同声明(self):
        cell={'location':'S!A1','metric_id':'m1','dimensions':{'currency':'CNY'}}
        expected=_per_cell_answer([cell],self.standard,['S'])
        for value in [{'answers':[cell]},{'任意容器':{'多层资料':[cell]}},{'cells':[cell]}]:
            self.assertEqual(_per_cell_answer(value,self.standard,['S']),expected)
        for value in [{'answers':{'共同声明':{'metric_id':'m1','currency':'CNY'},
                                  '位置列表':['S!A1','S!A2']}},
                      {'answers':[{'location':'S!A1','金标准编号':'m1','币种':'CNY'}]}]:
            with self.assertRaises(ValueError):_per_cell_answer(value,self.standard,['S'])
        nested={'location':'S!A1','metric_id':'m1','dimensions':{
            'asset_selection':{'location':'S!A2','metric_id':'m2'}}}
        result=_per_cell_answer({'任意容器':[nested]},self.standard,['S'])
        self.assertEqual([r['ranges'] for r in result['metrics']],[['S!A1']])

    def test_重复嵌套值只排除冲突维度(self):
        raw='{"S!A1":{"metric_id":"m1","currency":"CNY","asset_selection":{"mode":"members","members":["甲"],"members":["乙"]}}}'
        result=_per_cell_answer(raw,self.standard,['S'])
        self.assertEqual([r['id'] for r in result['metrics']],['m1'])
        self.assertEqual([(r['id'],r['value']) for r in result['dimensions']],[('currency','CNY')])
        self.assertEqual(result['unresolved'][0]['id'],'asset_selection')

    def test_同值重复不增门槛且第三次重复不能抹掉冲突(self):
        same=_per_cell_answer('{"S!A1":{"metric_id":"m1","metric_id":"m1"}}',self.standard,['S'])
        self.assertFalse(same['unresolved'])
        bad=_per_cell_answer('{"S!A1":{"currency":"CNY","currency":"USD","currency":["CNY","USD"]}}',self.standard,['S'])
        self.assertFalse(bad['dimensions']);self.assertTrue(bad['unresolved'])

    def test_错误补入不覆盖已保存的正确身份(self):
        old,_,_=self.recognize()
        raw='{"合成表!B3":{"metric_id":"固定资产原值","metric_id":"不存在的标准"}}'
        saved=save_direct_answer(old,raw,self.case.folder/'冲突原答复保存.json',field_scope={'合成表!B3':['metric_id']})
        self.assertEqual(dict(_records(saved))['合成表!B3']['metric_id'],'固定资产原值')
        self.assertEqual(saved['本次逐格原答复'],raw)
        self.assertFalse(saved['本次逐格接入']['complete'])

    def test_新槽位不能夹带未采用期间角色覆盖旧值(self):
        old,_,_=self.recognize()
        raw=json.dumps([{'location':'合成表!B3','metric_id':'固定资产原值','template_period_role':'不采用的新角色'},
                        {'location':'合成表!C3','metric_id':'固定资产原值','template_period_role':'prior'}],ensure_ascii=False)
        saved=save_direct_answer(old,raw,self.case.folder/'期间角色按范围保存.json',
            field_scope={'合成表!B3':['metric_id'],'合成表!C3':['metric_id','template_period_role']})
        roles={cell:r.get('template_period_role') for t in saved['outline']['subtables'] for r in t['metrics'] for cell in r['ranges']}
        self.assertEqual(roles['B3'],'closing')
        self.assertEqual(roles['C3'],'prior')
        self.assertEqual(saved['本次逐格原答复'],raw)

    def test_新维度编号也不能夹带未采用期间角色(self):
        old,_,_=self.recognize()
        seed={'location':'合成表!A3','dimension_id':'asset_selection',
              'dimensions':self.answer()['合成表!B3']['dimensions'],'template_period_role':'closing'}
        old=save_direct_answer(old,json.dumps(seed,ensure_ascii=False),self.case.folder/'维度角色原值.json',
            field_scope={'合成表!A3':['dimension_id','template_period_role',*seed['dimensions']]})
        raw=json.dumps({**seed,'template_period_role':'不采用的新角色'},ensure_ascii=False)
        saved=save_direct_answer(old,raw,self.case.folder/'维度角色按范围保存.json',field_scope={'合成表!A3':['dimension_id']})
        roles=[r.get('template_period_role') for t in saved['outline']['subtables'] for r in t['dimension_bindings'] if r['ranges']==['A3']]
        self.assertEqual(roles,['closing'])
        self.assertEqual(saved['本次逐格原答复'],raw)

    def test_日期封存往返保留类型且普通日期文字不改(self):
        original={'date':date(2026,10,29),'datetime':datetime(2026,10,29,12,34,56),'text':'2026-10-29'}
        text=json.dumps(original,default=json_date)
        restored=json.loads(text,object_pairs_hook=json_dates)
        self.assertEqual(original,restored)
        self.assertEqual(direct_digest(original),direct_digest(restored))
        path=save_json(self.case.folder/'日期原值.json',original)
        self.assertEqual(_read(path),original)
        self.assertIs(type(_read(path)['date']),date)
        self.assertIs(type(_read(path)['datetime']),datetime)

    def test_只交一格不能报告两格识别完成(self):
        reply=self.answer();reply.pop('合成表!C3')
        config={**self.case.config,'a_target_table':{'id':'本表','sheet':'合成表','range':'A1:C8','data_ranges':['B3:C3']}}
        with patch.object(UnifiedSemanticEngine,'request_text',return_value=json.dumps(reply,ensure_ascii=False)):
            result=run_step('recognize_a',config)
        self.assertEqual(result['status'],'partial')
        mapping=load_fact_mapping(result['files']['a_mapping'],{})
        self.assertEqual(mapping['本次逐格接入']['未接入字段'],[{'location':'合成表!C3','field':'metric_id / dimension_id'}])

    def test_原表日期在识别合集保存和回读后仍是日期(self):
        source=self.case.folder/'日期上下文原表.xlsx'
        book=Workbook();ws=book.active;ws.title='合成表'
        ws['A1']=datetime(2026,10,29,12,34,56);ws['B3']=1;ws['C3']=2;ws['A8']='表后说明'
        book.save(source);book.close();self.case.config['a_path']=str(source)
        mapping,_,_=self.recognize()
        value=mapping['机械记录']['sheets'][0]['cells']['A1']['value']
        self.assertEqual(value,datetime(2026,10,29,12,34,56))
        self.assertIs(type(value),datetime)

    def test_AB最终提示词简短且完整材料仍在(self):
        a,config,run=self.recognize()
        a['facts'].append({'record_type':'metric_fact','metric_id':'无关槽位','dimensions':{'currency':'USD'},
                          'source_reference':{'location':'合成表!B4'}})
        for carrier,target in [('A',None),('B',a)]:
            with self.subTest(carrier=carrier):
                engine=self.case.engine()
                with patch.object(engine,'request_text',side_effect=RuntimeError('发送前截获')) as request:
                    with self.assertRaisesRegex(RuntimeError,'发送前截获'):
                        recognize_a(engine,self.case.source,[],artifact_metadata={'carrier':carrier,'normal_entry':True,
                            'comparison_requirements':[{'维度ID':'entity','要求含义':'不参与比较'}],
                            '既有另一方身份参照':{'载体':'另一方','原样身份':{'需要':[
                                {'metric_id':'固定资产原值','dimensions':{'currency':'CNY'}},
                                {'metric_id':'无关槽位','dimensions':{'currency':'USD'}}]}},
                            'additional_context':'按本次明确要求换期。'},
                            target_mapping=target,target_table=config['a_target_table'],
                            allowed_metrics=['固定资产原值'])
                prompt,materials=request.call_args.args
                self.assertLessEqual(len(prompt),2000)
                self.assertIn('各格分别是什么意思',prompt)
                self.assertIn('共同属性也各格写明，未知照实写',prompt)
                self.assertIn('本次原表网格.txt',materials)
                self.assertIn('测试项目3',materials['本次原表网格.txt'])
                self.assertTrue(materials['金标准.jsonl'])
                self.assertNotIn('A原表.json',materials)
                self.assertNotIn('A原表网格.txt',materials)
                if carrier=='B':
                    self.assertTrue(materials['A需要.json']['需要'])
                    self.assertIn('outline',materials['A需要.json'])
                    self.assertTrue(all(row['id']=='template_period_role'
                        for table in materials['A需要.json']['outline']['subtables'] for row in table['dimensions']))
                    self.assertNotIn('unresolved',materials['A需要.json'])
                    self.assertTrue(all(row.get('metric_id')!='无关槽位' for row in materials['A需要.json']['需要']))
                self.assertIn('用户明确的比较排除.json',materials)
                self.assertNotIn('既有另一方身份证表达.json',materials)
                source=materials.get('原文.json',materials.get('B原文.json'))
                self.assertNotIn('完整工作表原文',source)

    def test_A只在明确B原表范围时传递既有原生表达(self):
        a,config,run=self.recognize()
        with patch.object(UnifiedSemanticEngine,'request_text',return_value=json.dumps(self.answer(),ensure_ascii=False)):
            b=run_step('recognize_b',{**config,'a_mapping':run['files']['a_mapping'],
                'b_target_table':config['a_target_table']})
        self.assertEqual(b['status'],'complete',b)
        for ranges in (None,[{'sheet':'合成表','range':'A2:B4'}]):
            with patch.object(UnifiedSemanticEngine,'request_text',side_effect=RuntimeError('只离线截获')) as request:
                run_step('recognize_a',{**config,'b_mapping':b['files']['b_mapping'],'b_source_ranges':ranges})
            materials=request.call_args.args[1]
            if ranges:
                reference=materials['既有另一方身份证表达.json']
                self.assertEqual(reference['source_ranges'],ranges)
                self.assertEqual([row['location'] for row in reference['原样身份']['需要']],['合成表!B3'])
                self.assertEqual(reference['原样身份']['需要'][0]['dimensions'],self.answer()['合成表!B3']['dimensions'])
            else:self.assertNotIn('既有另一方身份证表达.json',materials)

    def test_长附加指令发送前停止且不截断原文(self):
        engine=self.case.engine()
        with patch.object(engine,'request_text') as request:
            with self.assertRaisesRegex(ValueError,'超过2000'):
                recognize_a(engine,self.case.source,[],artifact_metadata={'carrier':'A','additional_context':'要求'*2000},
                    target_table=self.case.config['a_target_table'],allowed_metrics=['固定资产原值'])
            request.assert_not_called()

    def test_已有阅读范围保留本表及共同说明但不扩整工作表(self):
        a,_,_=self.recognize()
        table={'id':'本表','sheet':'合成表','range':'A2:C4','data_ranges':['B3:C3'],
               'context_references':[{'sheet':'合成表','range':'A1:C1'},
                                     {'sheet':'合成表','range':'A8:C8'}]}
        for carrier,target in [('A',None),('B',a)]:
            engine=self.case.engine()
            with patch.object(engine,'request_text',side_effect=RuntimeError('离线截获')) as request:
                with self.assertRaisesRegex(RuntimeError,'离线截获'):
                    recognize_a(engine,self.case.source,[],artifact_metadata={'carrier':carrier},
                        target_mapping=target,target_table=table,allowed_metrics=['固定资产原值'])
            materials=request.call_args.args[1];source=materials['原文.json']
            self.assertNotIn('完整工作表原文',source)
            self.assertNotIn('测试项目3',json.dumps(source.get('sheets',[]),ensure_ascii=False))
            text=json.dumps(source,ensure_ascii=False)+materials['本次原表网格.txt']
            self.assertIn('合成企业；单体；人民币；元',text)
            self.assertIn('测试项目8',text)
            self.assertNotIn('测试项目6',text)

    def test_B保留本表及公式来源行列但不附来源整张工作表(self):
        a,_,_=self.recognize()
        source=self.case.folder/'B定向公式原表.xlsx'
        book=load_workbook(self.case.source)
        book['合成表']['B3']="='原始来源'!B3"
        original=book.create_sheet('原始来源')
        original['A1']='来源共同标题原文';original.merge_cells('A1:B1')
        original['A3']='实际来源项目';original['B3']=10
        original['A80']='来源表无关另一项目';original['B80']=999
        book.save(source);book.close()
        engine=self.case.engine()
        with patch.object(engine,'request_text',side_effect=RuntimeError('离线截获')) as request:
            with self.assertRaisesRegex(RuntimeError,'离线截获'):
                recognize_a(engine,source,[],artifact_metadata={'carrier':'B'},target_mapping=a,
                    target_table={'id':'本表','sheet':'合成表','range':'A1:C4','data_ranges':['B3:C3']},
                    allowed_metrics=['固定资产原值'])
        materials=request.call_args.args[1]
        sent=json.dumps(materials['原文.json'],ensure_ascii=False)+materials['本次原表网格.txt']
        self.assertIn('合成企业；单体；人民币；元',sent)
        self.assertIn('来源共同标题原文',sent)
        self.assertIn('实际来源项目',sent)
        self.assertEqual(materials['原文.json']['公式及来源']['公式原文']['合成表!B3'],"='原始来源'!B3")
        self.assertNotIn('来源表无关另一项目',sent)
        self.assertNotIn('测试项目8',materials['本次原表网格.txt'])

    def test_A完整Word正文不重复附整张Excel且普通Excel保留表外原文(self):
        table={'id':'本表','sheet':'合成表','range':'A2:C4','data_ranges':['B3:C4']}
        title='原Word正文顺序（表格引用指向实际表区，各段适用范围由模型判断）'
        text='完整Word正文：合成企业，人民币元；本期为2025年度；本表之后的限制原文。'
        for complete in (True,False):
            context=[{'title':title if complete else '普通说明','text':text}]
            engine=self.case.engine()
            with patch.object(engine,'request_text',side_effect=RuntimeError('离线截获')) as request:
                with self.assertRaisesRegex(RuntimeError,'离线截获'):
                    recognize_a(engine,self.case.source,context,artifact_metadata={'carrier':'A'},
                                target_table=table,allowed_metrics=['固定资产原值'])
            materials=request.call_args.args[1];source=materials['原文.json']
            self.assertIn(text,json.dumps(source,ensure_ascii=False))
            self.assertIn('本期',materials['本次原表网格.txt'])
            self.assertIn('测试项目3',materials['本次原表网格.txt'])
            self.assertNotIn('sheets',source)
            self.assertNotIn('完整工作表原文',source)
            if complete:self.assertNotIn('测试项目6',materials['本次原表网格.txt'])
            else:self.assertIn('测试项目6',materials['本次原表网格.txt'])

    def test_B公式失败仍发布逐格成果并搬运其他原值(self):
        a,config,run=self.recognize()
        source=self.case.folder/'B公式原表.xlsx'
        book=Workbook();ws=book.active;ws.title='合成表';ws['B3']='=2+3';ws['C3']=20
        book.save(source);book.close()
        config.update(a_mapping=run['files']['a_mapping'],b_path=str(source));config.pop('b_target_table')
        replies=['工作表：合成表；完整范围：A1:C3；数据范围：B3:C3；相关标准：固定资产原值；对应A：合成表!A1:C8',
                 json.dumps(self.answer(),ensure_ascii=False)]
        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=replies) as request, \
             patch('附注更新.统一勾稽.calculate_mapping_formulas',side_effect=RuntimeError('模拟重算不可用')):
            result=run_step('recognize_b',config)
        self.assertEqual(request.call_count,2)
        self.assertIn('b_mapping',result['files'],result)
        self.assertIn('公式取值未完成',result['message'])
        b=load_fact_mapping(result['files']['b_mapping'],{})
        self.assertIn('模拟重算不可用',b['公式取值问题'])
        self.assertEqual(len(list(_records(b))),2)
        plan=direct_update_plan(a,b)
        self.assertEqual([(r['cell'],r['value']) for r in plan['updates']],[('C3',20)])

    def test_原生日期经AB保存对应和Excel副本写回仍是日期(self):
        config={**self.case.config,'update_purpose':{'reason':'日期随B，原值搬运'}}
        replies=[]
        for side,year in [('a',2025),('b',2026)]:
            path=self.case.folder/(side+'日期原表.xlsx')
            book=Workbook();ws=book.active;ws.title='合成表';ws['A1']='期末日期'
            ws['B3']=datetime(year,12,31,12,34,56);book.save(path);book.close()
            config[side+'_path']=str(path)
            config[side+'_target_table']={'id':'日期表','sheet':'合成表','range':'A1:B3','data_ranges':['B3']}
            replies.append(json.dumps({'合成表!B3':{'dimension_id':'period','template_period_role':'closing',
                'dimensions':{'period':{'kind':'instant','date':str(year)+'-12-31'}}}},ensure_ascii=False))
        hashes={config[key]:file_hash(config[key]) for key in ('a_path','b_path')}
        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=replies) as request:
            for step in ('recognize_a','recognize_b','match','write'):
                result=run_step(step,config)
                self.assertEqual(result['status'],'complete',result);config.update(result['files'])
        self.assertEqual(request.call_count,2)
        plan=_read(config['update_plan'])
        self.assertIs(type(plan['updates'][0]['value']),datetime)
        book=load_workbook(config['a_prime']);value=book['合成表']['B3'].value;book.close()
        self.assertEqual(value,datetime(2026,12,31,12,34,56));self.assertIs(type(value),datetime)
        self.assertEqual({path:file_hash(path) for path in hashes},hashes)

    def test_日期公式的机械回读不降成文字(self):
        source=self.case.folder/'公式原件.xlsx';calculated=self.case.folder/'模拟公式结果.xlsx'
        for path,value in [(source,'=DATE(2026,10,29)'),(calculated,datetime(2026,10,29))]:
            book=Workbook();ws=book.active;ws.title='合成表';ws['B3']=value
            ws['B3'].number_format='yyyy-mm-dd';book.save(path);book.close()
        mapping={'source_path':str(source),'source_hash':file_hash(source)}
        receipt={'schema':'统一语义公式计算-v1','method':'Excel.CalculateFullRebuild',
            'source_path':str(source),'source_hash':file_hash(source),
            'workbook_path':str(calculated),'workbook_hash':file_hash(calculated)}
        with patch('附注更新.统一勾稽._calculation_formulas',return_value={'合成表!B3':'=DATE(2026,10,29)'}):
            result=_calculated_values(mapping,receipt)
        self.assertEqual(result['合成表!B3']['state'],'value')
        self.assertEqual(result['合成表!B3']['value'],datetime(2026,10,29))


if __name__=='__main__':
    unittest.main()
