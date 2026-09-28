"""独立评测的行为测试：统计错误、遗漏、重复及维度和取数差异。"""
from pathlib import Path
import importlib.util
import json
import sys
import tempfile
import unittest
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
MODULE = ROOT / "附注更新" / "评测.py"


class BenchmarkTests(unittest.TestCase):
    def setUp(self):
        self.assertTrue(MODULE.exists(), "独立评测模块尚未实现")
        spec = importlib.util.spec_from_file_location("独立评测", MODULE)
        self.module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(self.module)
        base = Path(r"C:\Users\27651\Documents\Codex\2026-09-19\du-q\work\评测测试")
        base.mkdir(parents=True, exist_ok=True)
        self.source = Path(tempfile.mkdtemp(prefix="用例_", dir=base))
        mappings = self.source / "work" / "excel-mapping-v1" / "style1-manual-review-v2"
        mappings.mkdir(parents=True)
        (self.source / "语义金标准").mkdir()
        (self.source / "附注样式示例").mkdir()
        book = Workbook()
        sheet = book.active
        sheet.title = "财务"
        for address, value in {"A1":"表头", "B2":100, "C2":200, "D2":300, "E2":400, "F2":500}.items():
            sheet[address] = value
        book.save(self.source / "附注样式示例" / "附注样式1.xlsx")
        records = [
            {"worksheet":"财务", "cell":cell, "target_slot_id":f"S{i}",
             "dimensions":{"期间角色":"期末", "币种":"人民币", "单位":"元"},
             "target_slot":{"status":"active", "value_type":"monetary"}}
            for i, cell in enumerate(["B2","C2","D2","E2","F2"],1)]
        self.write_jsonl(mappings / "样式1-合并映射.jsonl", records)
        self.write_jsonl(self.source / "语义金标准" / "财务报表附注语义金标准-v1.jsonl",
                         [{"id":f"S{i}","status":"active","scope":"consolidated","value_type":"monetary"} for i in range(1,7)])
        self.predictions = {"mappings":[
            {"sheet":"财务","cell":"$B$2","slot_id":"S1","scope":"consolidated","dimensions":{"period_role":"closing","currency":"CNY","unit":"元"},"value":"100.00"},
            {"sheet":"财务","cell":"C2","slot_id":"S6","scope":"consolidated","dimensions":{}},
            {"sheet":"财务","cell":"E2","slot_id":"S4","scope":"consolidated","dimensions":{}},
            {"sheet":"财务","cell":"E2","slot_id":"S4","scope":"consolidated","dimensions":{}},
            {"sheet":"财务","cell":"F2","slot_id":"S5","scope":"consolidated","dimensions":{"period_role":"opening","currency":"CNY","unit":"元"},"value":501},
            {"sheet":"财务","cell":"G2","slot_id":"S6","scope":"consolidated","dimensions":{}}],
            "excluded":[], "unresolved":[]}

    @staticmethod
    def write_jsonl(path, values):
        path.write_text("\n".join(json.dumps(v,ensure_ascii=False) for v in values)+"\n",encoding="utf-8-sig")

    def test_counts_do_not_credit_duplicates_or_hide_dimension_errors(self):
        result = self.module.compare_predictions(self.predictions, 1, self.source)
        counts = result["counts"]
        self.assertEqual(counts["expected"],5)
        for key, expected in {"correct":2,"wrong":1,"missing":1,"extra":1,"duplicate":1,
                              "dimension_review":1,"value_mismatch":1,"values_checked":2}.items():
            self.assertEqual(counts[key], expected, (key,result))
        self.assertEqual(result["details"]["dimension_review"][0]["cell"],"F2")
        self.assertFalse(result["full_semantic_verified"])

    def test_benchmark_reads_actual_workbook_values_and_checks_gold_status(self):
        benchmark = self.module.read_benchmark(1,self.source)
        self.assertEqual(len(benchmark),5)
        self.assertEqual(benchmark[0]["value"],100)
        self.assertEqual(benchmark[0]["cell"],"B2")
        self.assertEqual(benchmark[0]["slot_id"],"S1")
        self.assertTrue(benchmark[0]["source_hash"])
        gold = self.source / "语义金标准" / "财务报表附注语义金标准-v1.jsonl"
        changed = gold.read_text(encoding="utf-8-sig").replace('"active"','"deprecated"',1)
        # 用例目录属于测试自身；保留另一个源目录以验证停用ID，不覆盖正式资料。
        second = self.source / "停用来源"
        import shutil
        shutil.copytree(self.source / "work",second / "work")
        shutil.copytree(self.source / "附注样式示例",second / "附注样式示例")
        (second / "语义金标准").mkdir()
        (second / "语义金标准" / gold.name).write_text(changed,encoding="utf-8-sig")
        with self.assertRaises(ValueError):
            self.module.read_benchmark(1,second)

    def test_source_hash_mismatch_and_conflicting_classifications_are_visible(self):
        self.predictions["source_hash"] = "not-the-source-hash"
        with self.assertRaises(ValueError):
            self.module.compare_predictions(self.predictions,1,self.source)
        self.predictions.pop("source_hash")
        self.predictions["excluded"] = [{"sheet":"财务","cell":"B2","category":"header"}]
        result = self.module.compare_predictions(self.predictions,1,self.source)
        self.assertEqual(result["counts"]["duplicate"],2)
        self.assertEqual(result["counts"]["correct"],1)

    def test_cli_writes_new_result_and_refuses_overwrite(self):
        import subprocess
        prediction = self.source / '预测.json'
        prediction.write_text(json.dumps(self.predictions,ensure_ascii=False),encoding='utf-8-sig')
        output = self.source / '独立结果.json'
        command = [sys.executable,'-X','utf8',str(MODULE),'--source-root',str(self.source),
                   '--prediction',str(prediction),'--style','1','--output',str(output)]
        result = subprocess.run(command,capture_output=True,text=True,encoding='utf-8')
        self.assertEqual(result.returncode,0,result.stderr)
        self.assertEqual(json.loads(output.read_text(encoding='utf-8-sig'))['counts']['correct'],2)
        before = output.read_bytes()
        again = subprocess.run(command,capture_output=True,text=True,encoding='utf-8')
        self.assertNotEqual(again.returncode,0)
        self.assertEqual(output.read_bytes(),before)
    def test_style2_uses_all_mapping_files_and_excludes_structure_records(self):
        folder = self.source / "work" / "excel-mapping-v1" / "style2-manual-review-v2"
        folder.mkdir()
        self.write_jsonl(folder / "001-样式2-人工映射.jsonl",[
            {"worksheet":"财务","cell":"A1","target_slot_id":None,"status":"结构格排除"},
            {"worksheet":"财务","cell":"B2","target_slot_id":"S1","dimensions":{}}])
        import shutil
        shutil.copy2(self.source / "附注样式示例" / "附注样式1.xlsx", self.source / "附注样式示例" / "附注样式2.xlsx")
        benchmark = self.module.read_benchmark(2,self.source)
        self.assertEqual(len(benchmark),1)
        self.assertEqual(benchmark[0]["cell"],"B2")
        result = self.module.compare_predictions({"mappings":[{"sheet":"财务","cell":"A1","slot_id":"S1"}]},2,self.source)
        self.assertEqual(result["counts"]["extra"],1)
        self.assertEqual(result["counts"]["missing"],1)

    def test_recalculated_copy_is_bound_to_original_and_actual_values(self):
        from openpyxl import load_workbook
        original=self.source/"附注样式示例"/"附注样式1.xlsx"
        copy=self.source/"重算副本.xlsx"
        book=load_workbook(original);book.save(copy);book.close()
        self.predictions.update(source_hash=self.module._hash(copy),original_hash=self.module._hash(original),source_path=str(copy))
        result=self.module.compare_predictions(self.predictions,1,self.source)
        self.assertTrue(result["source_hash_verified"])
        self.assertEqual(result["value_source"],str(copy))
        # 仅声明original_hash不能授权替换业务资料。
        changed=self.source/"错误副本.xlsx"
        book=load_workbook(original);book.active["B2"]=999;book.save(changed);book.close()
        self.predictions.update(source_hash=self.module._hash(changed),source_path=str(changed))
        with self.assertRaises(ValueError):
            self.module.compare_predictions(self.predictions,1,self.source)


if __name__ == "__main__":
    unittest.main(verbosity=2)
