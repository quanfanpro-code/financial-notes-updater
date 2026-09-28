# coding: utf-8
"""新版语义结构变更的真实Excel验证；不将模拟语义当作模型识别验收。"""
import copy
import unittest
import uuid
import json
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,Border,Side
from openpyxl.workbook.defined_name import DefinedName
from 测试.测试统一语义 import records,fact
from 附注更新.统一语义 import validate_records
from 附注更新.表格 import file_hash


class UnifiedStructureTests(unittest.TestCase):
    def setUp(self):
        self.folder=Path(__file__).resolve().parents[1]/'测试结果'/('统一结构_'+uuid.uuid4().hex);self.folder.mkdir()
        self.standard=validate_records(records());self.standard['sha256']='a'*64
        self.target=self.make('原附注',[['资产类别','年末余额'],['甲类设备',11],['乙类设备',22],['合计',33]],
                              [('B2','甲类设备'),('B3','乙类设备'),('B4',None)])
        self.source=self.make('更新数据',[['项目','期末金额'],['甲类设备',77],['乙类设备',77],['丙类设备',77],['合计',231]],
                              [('B2','甲类设备'),('B3','乙类设备'),('B4','丙类设备'),('B5',None)])
        self.proposal={'operations':[{'range_name':'附注表','axis':'row','action':'insert','index':3,'count':1,'template_index':2,'reason':'新增资产类别'}],
             'bindings':[{'source_sheet':'明细','source_cell':'B4','target_sheet':'明细','target_cell':'B4',
                          'template_sheet':'明细','template_cell':'B3','reason':'丙类设备的原值余额'}],
             'labels':[{'sheet':'明细','cell':'A4','value':'丙类设备','source_sheet':'明细','source_cell':'A4','reason':'来源的实际资产类别名称'}],
             'unresolved':[]}

    def make(self,name,rows,meanings):
        path=self.folder/(name+'.xlsx');book=Workbook();sheet=book.active;sheet.title='明细'
        for row in rows:sheet.append(row)
        for row in sheet:
            for cell in row:cell.font=Font(name='宋体',size=10);cell.border=Border(bottom=Side(style='thin'))
        sheet['A8']='后方说明';sheet['B8']='不得丢失'
        book.defined_names.add(DefinedName('附注表',attr_text="'明细'!$A$1:$"+chr(64+len(rows[0]))+'$'+str(len(rows))))
        book.save(path);book.close();digest=file_hash(path);facts=[];owners={}
        for address,category in meanings:
            row=fact();row['id']='明细!'+address;row['source_reference']={'file_hash':digest,'location':row['id']}
            row['dimensions']['asset_selection']={'domain':'固定资产','mode':'all','completeness':'complete'} if category is None else {
                'domain':'固定资产','mode':'members','members':[category],'completeness':'complete'}
            if address.startswith('C'):row['dimensions']['period']={'kind':'instant','date':'2024-12-31'}
            row['dimension_evidence']={key:['明细!A'+address[1:]] for key in row['dimensions']}
            row['owner_reference']=row['id'];owners[row['id']]=copy.deepcopy(row['dimensions'])
            book=load_workbook(path);row['raw_value']=book['明细'][address].value;book.close();facts.append(row)
        return {'schema_version':2,'gold_hash':self.standard['sha256'],'source_path':str(path),'source_hash':digest,
                'facts':facts,'owners':owners,'source_context':[]}

    def apply(self,proposal=None):
        from 附注更新.统一结构 import apply_structural_layout
        return apply_structural_layout(self.target,self.source,self.standard,proposal or self.proposal,self.folder/'调整后.xlsx')

    def test_insert_row_keeps_old_values_and_creates_blank_semantically_bound_cell(self):
        result=self.apply();book=load_workbook(result['source_path']);sheet=book['明细']
        self.assertEqual([sheet['A'+str(r)].value for r in range(1,6)],['资产类别','甲类设备','乙类设备','丙类设备','合计'])
        self.assertEqual([sheet['B'+str(r)].value for r in range(2,6)],[11,22,None,33])
        self.assertEqual(sheet['B4'].font.name,'宋体');self.assertEqual(sheet['A9'].value,'后方说明')
        book.close();self.assertEqual(result['positions']['明细!B4'],'明细!B5')
        self.assertEqual(result['new_bindings'][0]['source_id'],'明细!B4')
        self.assertEqual(file_hash(self.target['source_path']),self.target['source_hash'])

    def test_inserting_column_keeps_periods_distinct_and_new_values_blank(self):
        self.source=self.make('含上年数据',[['类别','年末','年初'],['甲类设备',77,77],['乙类设备',77,77],['合计',154,154]],
                              [('B2','甲类设备'),('C2','甲类设备'),('B3','乙类设备'),('C3','乙类设备'),('B4',None),('C4',None)])
        proposal={'operations':[{'range_name':'附注表','axis':'column','action':'insert','index':2,'count':1,'template_index':1,'reason':'新增独立年初期间'}],
                  'bindings':[{'source_sheet':'明细','source_cell':'C'+str(r),'target_sheet':'明细','target_cell':'C'+str(r),
                               'template_sheet':'明细','template_cell':'B'+str(r),'reason':'对应资产类别的年初余额'} for r in range(2,5)],
                  'labels':[{'sheet':'明细','cell':'C1','value':'年初','source_sheet':'明细','source_cell':'C1','reason':'原始期间标题'}],'unresolved':[]}
        result=self.apply(proposal);book=load_workbook(result['source_path']);self.assertEqual(book['明细']['C1'].value,'年初')
        self.assertTrue(all(book['明细']['C'+str(r)].value is None for r in range(2,5)));book.close()
        self.assertEqual(len(result['new_bindings']),3)

    def test_deletion_requires_explicit_source_evidence_and_cannot_remove_still_present_semantics(self):
        proposal={'operations':[{'range_name':'附注表','axis':'row','action':'delete','index':2,'count':1,'template_index':1,
                                'reason':'来源未见乙不能直接解释为删除'}],'bindings':[],'labels':[],'unresolved':[]}
        with self.assertRaises(ValueError):self.apply(proposal)
        proposal['operations'][0]['source_evidence']=['明细!A3']
        with self.assertRaises(ValueError):self.apply(proposal)

    def test_foreign_entity_or_fabricated_labels_cannot_be_inserted(self):
        self.source['facts'][2]['dimensions']['entity']='另一个报告主体'
        self.source['owners'][self.source['facts'][2]['owner_reference']]['entity']='另一个报告主体'
        with self.assertRaises(ValueError):self.apply()
        self.source['facts'][2]['dimensions']['entity']=self.target['facts'][0]['dimensions']['entity']
        self.source['owners'][self.source['facts'][2]['owner_reference']]['entity']=self.target['facts'][0]['dimensions']['entity']
        bad=copy.deepcopy(self.proposal);bad['labels'][0]['value']='虚构的新类别'
        with self.assertRaises(ValueError):self.apply(bad)

    def test_every_inserted_financial_position_needs_its_own_semantic_binding(self):
        bad=copy.deepcopy(self.proposal);bad['bindings']=[]
        with self.assertRaises(ValueError):self.apply(bad)

    def test_explicit_complete_replacement_can_remove_a_missing_category_without_touching_other_values(self):
        self.source=self.make('完整列示的数据',[['本期完整类别明细，仅甲类设备','余额'],['甲类设备',77],['合计',77]],
                              [('B2','甲类设备'),('B3',None)])
        proposal={'operations':[{'range_name':'附注表','axis':'row','action':'delete','index':2,'count':1,'template_index':1,
                  'reason':'完整明细中乙类设备已不属于本期披露范围','coverage_reason':'来源标题明确这是完整类别明细且仅列甲类设备',
                  'source_evidence':['明细!A1']}],'bindings':[],'labels':[],'unresolved':[]}
        result=self.apply(proposal);book=load_workbook(result['source_path'])
        self.assertEqual([book['明细']['A'+str(r)].value for r in range(1,4)],['资产类别','甲类设备','合计'])
        self.assertEqual(book['明细']['B3'].value,33);self.assertEqual(book['明细']['A7'].value,'后方说明');book.close()
        self.assertEqual(result['removed_fact_ids'],['明细!B3'])

    def test_explicit_period_scope_can_remove_a_column_while_keeping_current_totals(self):
        self.target=self.make('双期原附注',[['类别','期末','期初'],['甲类设备',11,9],['乙类设备',22,8],['合计',33,17]],
                              [('B2','甲类设备'),('C2','甲类设备'),('B3','乙类设备'),('C3','乙类设备'),('B4',None),('C4',None)])
        self.source=self.make('仅本期完整披露',[['本次完整披露仅2025年末，不再列示比较期','余额'],['甲类设备',77],['乙类设备',77],['合计',154]],
                              [('B2','甲类设备'),('B3','乙类设备'),('B4',None)])
        proposal={'operations':[{'range_name':'附注表','axis':'column','action':'delete','index':2,'count':1,'template_index':1,
                  'reason':'本次明确不再披露比较期','coverage_reason':'来源说明完整披露范围仅为2025年末','source_evidence':['明细!A1']}],
                  'bindings':[],'labels':[],'unresolved':[]}
        result=self.apply(proposal);book=load_workbook(result['source_path'])
        self.assertEqual(book['明细']['B4'].value,33);self.assertIsNone(book['明细']['C4'].value)
        self.assertEqual(book['明细']['A8'].value,'后方说明');book.close()
        self.assertEqual(set(result['removed_fact_ids']),{'明细!C2','明细!C3','明细!C4'})

    def test_semantic_row_insertion_then_value_matching_reaches_original_word_bridge(self):
        from docx import Document
        from docx.shared import Pt
        from 测试.测试Word桥接 import WordBridgeTests
        from 附注更新.统一结构 import apply_structural_layout
        from 附注更新.统一勾稽 import build_fact_update_plan
        from 附注更新.表格 import write_updated_workbook
        bridge=WordBridgeTests();bridge.folder=self.folder;bridge.word=self.folder/'原始资产附注.docx'
        doc=Document();doc.add_heading('五、税项',1);doc.add_paragraph('范围外文字不得改动')
        doc.add_heading('六、财务报表重要项目的说明',1);doc.add_paragraph('固定资产原值；单户人民币元；2025年12月31日')
        table=doc.add_table(rows=4,cols=2);table.style='Table Grid'
        for row,values in zip(table.rows,[['资产类别','期末余额'],['甲类设备','11.00'],['乙类设备','22.00'],['合计','33.00']]):
            for cell,value in zip(row.cells,values):
                cell.text=value;cell.paragraphs[0].runs[0].font.name='宋体';cell.paragraphs[0].runs[0].font.size=Pt(10)
        doc.add_heading('七、其他说明',1);doc.add_paragraph('范围外后文也保留');doc.save(bridge.word)
        original=bridge.word.read_bytes();excel,linked,context=bridge.extract()
        target=copy.deepcopy(self.target);target.update(source_path=str(excel),source_hash=file_hash(excel));target['facts']=[];target['owners']={}
        book=load_workbook(excel)
        for old,address in zip(self.target['facts'],['C3','C4','C5']):
            row=copy.deepcopy(old);row.update(id='sheet1!'+address,owner_reference='sheet1!'+address)
            row['source_reference']={'file_hash':target['source_hash'],'location':row['id']};row['raw_value']=book['sheet1'][address].value
            row['dimension_evidence']={key:['sheet1!B'+address[1:]] for key in row['dimensions']}
            target['facts'].append(row);target['owners'][row['id']]=copy.deepcopy(row['dimensions'])
        book.close()
        proposal=copy.deepcopy(self.proposal);proposal['operations'][0]['range_name']='TA_001'
        proposal['bindings'][0].update(target_sheet='sheet1',target_cell='C5',template_sheet='sheet1',template_cell='C4')
        proposal['labels'][0].update(sheet='sheet1',cell='B5')
        layout=apply_structural_layout(target,self.source,self.standard,proposal,self.folder/'结构调整.xlsx')
        adjusted=copy.deepcopy(target);adjusted.update(source_path=layout['source_path'],source_hash=layout['source_hash'],facts=[],owners={})
        for row in target['facts']:
            moved=copy.deepcopy(row);location=layout['positions'][row['id']];moved.update(id=location,owner_reference=location)
            moved['source_reference']={'file_hash':layout['source_hash'],'location':location}
            moved['dimension_evidence']={key:[layout['positions'][ref] for ref in refs] for key,refs in row['dimension_evidence'].items()}
            adjusted['facts'].append(moved);adjusted['owners'][location]=copy.deepcopy(moved['dimensions'])
        added=copy.deepcopy(self.source['facts'][2]);added.update(id='sheet1!C5',owner_reference='sheet1!C5',raw_value=None,raw_value_state='blank')
        added['source_reference']={'file_hash':layout['source_hash'],'location':'sheet1!C5'}
        added['dimension_evidence']={key:['sheet1!B5'] for key in added['dimensions']}
        adjusted['facts'].append(added);adjusted['owners'][added['id']]=copy.deepcopy(added['dimensions'])
        # 此处语义由受控夹具给定，只测试结构与写回，不冒充真实模型识别。
        plan=build_fact_update_plan(adjusted,self.source,self.standard)
        self.assertFalse(plan['issues']);self.assertFalse(plan['unplaced_sources'])
        updated=self.folder/'完成语义取数.xlsx';write_updated_workbook(layout['source_path'],updated,plan)
        output=self.folder/'更新后资产附注.docx'
        result=bridge.call('update',word=linked,excel=updated,output=output,operations=layout['operations'],chapter_title='财务报表重要项目的说明')
        self.assertTrue(result['output_verified']);after=Document(output)
        self.assertEqual([row.cells[0].text for row in after.tables[0].rows],['资产类别','甲类设备','乙类设备','丙类设备','合计'])
        self.assertEqual([float(row.cells[1].text.replace(',','')) for row in after.tables[0].rows[1:]],[77,77,77,231])
        self.assertEqual([p.text for p in after.paragraphs],[p.text for p in doc.paragraphs])
        self.assertEqual(after.tables[0].rows[3].cells[1].paragraphs[0].runs[0].font.name,'宋体')
        self.assertEqual(bridge.word.read_bytes(),original)

    def test_model_numbers_and_unreviewed_plans_do_not_create_outputs(self):
        from 附注更新.统一结构 import plan_structural_layout
        case=self
        class Engine:
            standard=case.standard;gold_hash=case.standard['sha256'];usage={};_review_receipts=[]
            def _request(self,system,payload,validator):
                if payload['task']=='unified_structure':return validator(copy.deepcopy(case.proposal))
                return validator({'accepted':False,'reason':'语义范围尚不充分','source_evidence':['明细!A4'],'target_evidence':['明细!A3']})
            def _check_cancel(self):pass
        result=plan_structural_layout(Engine(),self.target,self.source,self.folder/'规划')
        self.assertFalse(result['accepted']);self.assertNotIn('layout',result)
        self.assertFalse(list((self.folder/'规划').glob('*.xlsx')))
        bad=copy.deepcopy(self.proposal);bad['bindings'][0]['value']=999
        with self.assertRaises(ValueError):self.apply(bad)

    def test_two_reviews_keep_metadata_but_never_copy_financial_values_into_new_cells(self):
        from 附注更新.统一结构 import plan_structural_layout
        case=self;calls=[]
        class Engine:
            standard=case.standard;gold_hash=case.standard['sha256'];usage={};_review_receipts=[]
            def _request(self,system,payload,validator):
                calls.append(payload)
                if payload['task']=='unified_structure':return validator(copy.deepcopy(case.proposal))
                return validator({'accepted':True,'reason':'丙类设备为既有指标的新类别，来源有独立语义，原合计保留',
                                  'source_evidence':['明细!A4'],'target_evidence':['明细!A3']})
            def _check_cancel(self):pass
        result=plan_structural_layout(Engine(),self.target,self.source,self.folder/'规划')
        self.assertTrue(result['accepted']);self.assertEqual(len(calls),3)
        self.assertTrue(all('raw_value' not in str(payload) for payload in calls))
        book=load_workbook(result['layout']['source_path']);self.assertIsNone(book['明细']['B4'].value);book.close()

    def complete_mapping(self,mapping):
        from 附注更新.表格 import read_workbook,materialize_candidate_blanks
        from 附注更新.统一识别 import workbook_coverage
        from 附注更新.统一勾稽 import load_fact_mapping
        snapshot=read_workbook(mapping['source_path'],context=mapping.get('source_context',[]));materialize_candidate_blanks(snapshot)
        positions={row['source_reference']['location'] for row in mapping['facts']}
        mapping['excluded']=[{'sheet':sheet['name'],'cell':address,'category':'annotation','reason':'受控分类，仅验证程序行为'}
                             for sheet in snapshot['sheets'] for address in sheet['cells'] if sheet['name']+'!'+address not in positions]
        mapping['unresolved']=[];mapping['coverage']=workbook_coverage(mapping,snapshot);mapping['complete']=True
        path=self.folder/(Path(mapping['source_path']).stem+'映射.json')
        path.write_text(json.dumps(mapping,ensure_ascii=False),encoding='utf-8-sig')
        return load_fact_mapping(path,self.standard)

    def test_adjusted_semantics_must_be_independently_recognized_and_match_every_survivor(self):
        from 附注更新.统一结构 import verify_adjusted_semantics
        layout=self.apply()
        adjusted=self.make('独立识别示意',[['资产类别','年末余额'],['甲类设备',11],['乙类设备',22],['丙类设备',None],['合计',33]],
                           [('B2','甲类设备'),('B3','乙类设备'),('B4','丙类设备'),('B5',None)])
        adjusted['source_path']=layout['source_path'];adjusted['source_hash']=layout['source_hash']
        for row in adjusted['facts']:
            row['source_reference']['file_hash']=layout['source_hash']
            row['raw_value_state']='blank' if row['raw_value'] is None else 'value'
        verify_adjusted_semantics(self.target,self.source,adjusted,self.standard,layout)
        bad=copy.deepcopy(adjusted)
        bad['facts'][0]['dimensions'],bad['facts'][1]['dimensions']=bad['facts'][1]['dimensions'],bad['facts'][0]['dimensions']
        bad['owners']={row['id']:copy.deepcopy(row['dimensions']) for row in bad['facts']}
        with self.assertRaisesRegex(ValueError,'含义'):verify_adjusted_semantics(self.target,self.source,bad,self.standard,layout)
        bad=copy.deepcopy(adjusted);bad['facts'].pop(2)
        with self.assertRaisesRegex(ValueError,'业务格'):verify_adjusted_semantics(self.target,self.source,bad,self.standard,layout)

    def test_formal_match_and_write_recheck_adjusted_mapping_and_reject_structure_tampering(self):
        self.formal_case()

    def test_formal_word_structure_write_and_second_update_preserve_outside_chapter(self):
        self.formal_case(word=True)

    def test_rejected_structure_keeps_adjusted_independent_mapping_visible(self):
        self.formal_case(reject=True)

    def test_context_scope_follows_inserted_rows_and_named_word_table(self):
        from 附注更新.统一结构 import structural_context
        from 附注更新.统一证据 import context_catalog
        self.target['source_context']=[{'text':'仅后方说明适用','sheet':'明细','first_row':8,'first_column':1,'row_count':1,'column_count':2},
             {'title':'表内说明','sheet':'明细','table_id':'1','range_name':'附注表','first_row':1,'first_column':1,'row_count':4,'column_count':2},
             '适用于全表的原始业务说明']
        layout=self.apply();context=structural_context(self.target,layout)
        self.assertEqual(context[0]['first_row'],9)
        self.assertEqual(context[1],self.target['source_context'][1])
        scopes=[row['scope'] for row in context_catalog(context,layout['source_path']).values() if row['scope']]
        self.assertEqual(scopes,[{'sheet':'明细','bounds':[9,1,9,2]},{'sheet':'明细','bounds':[1,1,5,2]}])

    def formal_case(self,word=False,reject=False):
        from unittest.mock import patch
        from 附注更新.统一识别 import UnifiedSemanticEngine
        from 附注更新.统一语义 import read_standard_v2
        from 附注更新.统一勾稽 import load_fact_mapping
        from 附注更新.分步流程 import run_step
        word_files={};target_sheet='明细';value_column='B';label_column='A';value_start=2
        if word:
            from docx import Document
            from docx.shared import Pt
            original_word=self.folder/'正式入口原附注.docx';doc=Document()
            doc.add_heading('五、税项',1);outside=doc.add_table(rows=1,cols=1);outside.cell(0,0).text='章外888不改'
            outside_xml=outside._tbl.xml;doc.add_heading('六、财务报表重要项目的说明',1)
            doc.add_paragraph('合成企业甲，单户，人民币元，2025年12月31日固定资产原值')
            table=doc.add_table(rows=4,cols=2);table.style='Table Grid'
            for row,values in zip(table.rows,[['资产类别','年末余额'],['甲类设备','11.00'],['乙类设备','22.00'],['合计','33.00']]):
                for cell,value in zip(row.cells,values):
                    cell.text=value;cell.paragraphs[0].runs[0].font.name='宋体';cell.paragraphs[0].runs[0].font.size=Pt(10)
            doc.add_heading('七、其他说明',1);doc.add_paragraph('章外后文999不改');doc.save(original_word)
            original_bytes=original_word.read_bytes()
            extracted=run_step('extract_a',{'word_path':str(original_word),'output_dir':str(self.folder/'正式提取')})
            self.assertEqual(extracted['status'],'complete',extracted['message']);word_files=extracted['files']
            self.target.update(source_path=word_files['a_path'],source_hash=file_hash(word_files['a_path']))
            link=json.loads(Path(word_files['a_link']).read_text(encoding='utf-8-sig'))
            self.target['source_context']=link['context']['tables'];self.target['owners']={}
            book=load_workbook(self.target['source_path'])
            for row,address in zip(self.target['facts'],['C3','C4','C5']):
                row.update(id='sheet1!'+address,owner_reference='sheet1!'+address,raw_value=book['sheet1'][address].value)
                row['source_reference']={'file_hash':self.target['source_hash'],'location':row['id']}
                row['dimension_evidence']={key:['sheet1!B'+address[1:]] for key in row['dimensions']}
                self.target['owners'][row['id']]=copy.deepcopy(row['dimensions'])
            self.proposal['operations'][0]['range_name']=next(iter(book.defined_names));book.close()
            self.proposal['bindings'][0].update(target_sheet='sheet1',target_cell='C5',template_sheet='sheet1',template_cell='C4')
            self.proposal['labels'][0].update(sheet='sheet1',cell='B5')
            target_sheet='sheet1';value_column='C';label_column='B';value_start=3
        gold=self.folder/'受控标准.jsonl';gold.write_text('\n'.join(json.dumps(row,ensure_ascii=False) for row in records()),encoding='utf-8')
        self.standard=read_standard_v2(gold)
        for mapping in (self.target,self.source):mapping.update(gold_path=str(gold),gold_hash=self.standard['sha256'])
        target=self.complete_mapping(self.target);source=self.complete_mapping(self.source);calls=[]
        def response(engine,system,payload):
            calls.append(payload['task'])
            if payload['task']=='unified_structure':return copy.deepcopy(self.proposal)
            if payload['task']=='unified_structure_review':return {'accepted':True,'reason':'受控试验新增明确类别，原合计保留',
                                                                   'source_evidence':['明细!A4'],'target_evidence':[target_sheet+'!'+label_column+str(value_start+1)]}
            if payload['task']=='unified_equivalence':
                return {'pairs':{key:{'decision':'different','reason':'受控试验不同资产类别',
                                     'left_evidence':list(row['left']['evidence']),'right_evidence':list(row['right']['evidence'])}
                                 for key,row in payload['pairs'].items()}}
            self.assertEqual(payload['task'],'unified_classify')
            # 受控回执按此测试原文分类；没有把这一规则加入产品。
            result={}
            for address in payload['candidate_cells']:
                label_address=label_column+address[1:];label=payload['cells'].get(label_address,{}).get('value')
                if address.startswith(value_column) and label in ('甲类设备','乙类设备','丙类设备','合计'):
                    dims=copy.deepcopy(fact()['dimensions'])
                    if label!='合计':dims['asset_selection']={'domain':'固定资产','mode':'members','members':[label],'completeness':'complete'}
                    if reject and label=='丙类设备':dims['asset_selection']['members']=['乙类设备']
                    result[address]={'kind':'metric_fact','metric_id':'固定资产原值','dimensions':dims,
                                     'dimension_evidence':{key:[label_address] for key in dims},'reason':'受控独立分类'}
                else:result[address]={'kind':'non_business','category':'annotation','reason':'受控标题或排版','evidence_cells':[label_column+str(value_start-1)]}
            return {'cells':result}
        config={'gold_path':str(gold),'a_path':target['source_path'],'b_path':source['source_path'],
                'a_mapping':target['mapping_path'],'b_mapping':source['mapping_path'],'output_dir':str(self.folder/'正式入口'),
                'model':'受控模型','base_url':'http://localhost','api_key':'synthetic-structure-key-0d8'}
        if word:config['a_link']=word_files['a_link']
        with patch.object(UnifiedSemanticEngine,'_http',response):result=run_step('match',config)
        if reject:
            self.assertEqual(result['status'],'partial',result['message'])
            self.assertTrue(Path(result['files']['adjusted_a_mapping']).is_file())
            config['update_plan']=result['files']['update_plan'];refused=run_step('write',config)
            self.assertEqual(refused['status'],'failed');self.assertNotIn('a_prime',refused.get('files',{}))
            book=load_workbook(result['files']['adjusted_a']);self.assertIsNone(book['明细']['B4'].value);book.close()
            return
        self.assertEqual(result['status'],'complete',result['message'])
        plan_path=Path(result['files']['update_plan']);plan=json.loads(plan_path.read_text(encoding='utf-8-sig'))
        self.assertIn('structure',plan);self.assertIn('unified_classify',calls)
        config['update_plan']=str(plan_path)
        with patch.object(UnifiedSemanticEngine,'_http',side_effect=AssertionError('写入不可重新询问模型')):
            written=run_step('write',config)
        self.assertEqual(written['status'],'complete',written['message'])
        book=load_workbook(written['files']['a_prime']);self.assertEqual([book[target_sheet][value_column+str(r)].value for r in range(value_start,value_start+4)],[77,77,77,231]);book.close()
        derived=load_fact_mapping(written['files']['a_prime_mapping'],self.standard)
        self.assertTrue(derived['complete']);self.assertEqual(len(derived['facts']),4)
        if word:
            after=Document(written['files']['word'])
            self.assertEqual(after.tables[0]._tbl.xml,outside_xml)
            self.assertEqual([p.text for p in after.paragraphs],[p.text for p in doc.paragraphs])
            self.assertEqual([row.cells[0].text for row in after.tables[1].rows],['资产类别','甲类设备','乙类设备','丙类设备','合计'])
            self.assertEqual([float(row.cells[1].text.replace(',','')) for row in after.tables[1].rows[1:]],[77,77,77,231])
            self.assertEqual(after.tables[1].rows[3].cells[1].paragraphs[0].runs[0].font.name,'宋体')
            self.assertEqual(original_word.read_bytes(),original_bytes)
            again={**config,'a_path':written['files']['a_prime'],'a_mapping':written['files']['a_prime_mapping'],'a_link':written['files']['word_link']}
            with patch.object(UnifiedSemanticEngine,'_http',response):matched=run_step('match',again)
            self.assertEqual(matched['status'],'complete',matched['message']);again['update_plan']=matched['files']['update_plan']
            with patch.object(UnifiedSemanticEngine,'_http',side_effect=AssertionError('写入不重新识别')):second=run_step('write',again)
            self.assertEqual(second['status'],'complete',second['message'])
            self.assertEqual(len(Document(second['files']['word']).tables[1].rows),5)
        tampered=copy.deepcopy(plan);tampered['structure']['layout']['operations'][0]['index']=1
        bad=self.folder/'篡改结构清单.json';bad.write_text(json.dumps(tampered,ensure_ascii=False),encoding='utf-8-sig')
        bad.with_name('更新清单校验.json').write_text(json.dumps({'sha256':file_hash(bad)}),encoding='utf-8')
        config['update_plan']=str(bad)
        refused=run_step('write',config)
        self.assertEqual(refused['status'],'failed');self.assertNotIn('a_prime',refused.get('files',{}))


if __name__=='__main__':unittest.main()
