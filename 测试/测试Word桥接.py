"""真实 Word / Excel 的桥接验收：不覆盖输入，回写数值并保留原表格式。"""
from __future__ import annotations
import json
from pathlib import Path
import subprocess
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET
from copy import copy
from docx import Document
from docx.shared import Pt
from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName

ROOT = Path(__file__).resolve().parents[1]
DLL = ROOT / "Word桥接" / "发布" / "WordBridge.dll"
DOTNET = Path(r"C:\Program Files\dotnet\dotnet.exe")
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class WordBridgeTests(unittest.TestCase):
    def setUp(self):
        base = ROOT / "Word桥接" / "测试结果"
        base.mkdir(exist_ok=True)
        self.folder = Path(tempfile.mkdtemp(prefix="验收_", dir=base))
        self.word = self.folder / "原始报告.docx"
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = "必须保留的报告页眉"
        doc.add_heading("六、合并财务报表项目注释", level=1)
        doc.add_paragraph("（一）货币资金")
        doc.add_paragraph("单位：人民币元")
        table = doc.add_table(rows=4, cols=3)
        table.style = "Table Grid"
        for row, values in zip(table.rows, [
            ["项目", "期末余额", "期初余额"],
            ["库存现金", "100.00", "90.00"],
            ["银行存款", "200.00", "180.00"],
            ["合计", "300.00", "270.00"],
        ]):
            for cell, value in zip(row.cells, values):
                cell.text = value
                run = cell.paragraphs[0].runs[0]
                run.font.name = "宋体"
                run.font.size = Pt(10)
        for cell in table.rows[-1].cells:
            cell.paragraphs[0].runs[0].bold = True
        doc.add_paragraph("表后说明必须保留。")
        doc.save(self.word)
        self.original = self.word.read_bytes()

    def call(self, action, success=True, **kwargs):
        self.assertTrue(DLL.exists(), "缺少 Word 桥接可执行程序，行为尚未实现")
        request = self.folder / ("请求_" + str(len(list(self.folder.glob("请求_*")))) + ".json")
        request.write_text(json.dumps(dict(action=action, **{k:str(v) if isinstance(v,Path) else v for k,v in kwargs.items()}), ensure_ascii=False), encoding="utf-8-sig")
        process = subprocess.run([str(DOTNET), str(DLL), str(request)], capture_output=True, text=True, encoding="utf-8")
        try:
            result = json.loads(process.stdout.lstrip("\ufeff"))
        except ValueError:
            self.fail(f"桥接没有返回 JSON：{process.stdout}\n{process.stderr}")
        self.assertEqual(process.returncode == 0, success, result)
        self.assertEqual(result["ok"], success, result)
        return result

    def extract(self):
        excel, linked, context = (self.folder / n for n in ["表格A.xlsx", "带连接报告.docx", "上下文.json"])
        result = self.call("extract", word=self.word, excel=excel, linked_word=linked, context=context)
        self.assertEqual(result["table_count"], 1)
        return excel, linked, context

    def save_changed(self, excel, rows):
        book = load_workbook(excel)
        sheet = book["sheet1"]
        for row in sheet.iter_rows(min_row=2, max_row=12, min_col=2, max_col=8):
            for cell in row:
                cell.value = None
        for r, values in enumerate(rows, 2):
            for c, value in enumerate(values, 2):
                sheet.cell(r, c, value)
        end_col = chr(65 + len(rows[0]))
        book.defined_names["TA_001"] = DefinedName("TA_001", attr_text=f"'sheet1'!$B$2:${end_col}${len(rows)+1}".replace("\\$", "$"))
        changed = self.folder / ("新数据_" + str(len(list(self.folder.glob("新数据_*")))) + ".xlsx")
        book.save(changed)
        return changed

    def test_chapter_scope_limits_read_extract_update_and_preserves_other_parts(self):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        chapter = "财务报表重要项目的说明"
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = "范围外页眉保留"
        doc.add_heading("五、税项", level=1)
        outside = doc.add_table(rows=2, cols=2)
        outside.cell(0, 0).text = "税项不得处理"
        outside.cell(1, 1).text = "666"
        record = OxmlElement("w:tblDescription")
        record.set(qn("w:val"), '<#'+json.dumps({"SavePath": str(self.folder / "无关底稿.xlsx"), "RelativePath": "无关底稿.xlsx", "LinkName": "OUTSIDE"}, ensure_ascii=False)+'#>')
        outside._tbl.tblPr.append(record)
        doc.add_heading("六、" + chapter, level=1)
        doc.add_paragraph("本章单位：人民币元")
        doc.add_heading("（一）货币资金", level=3)
        table = doc.add_table(rows=2, cols=2)
        for cell, value in zip([table.cell(0,0),table.cell(0,1),table.cell(1,0),table.cell(1,1)], ["项目","年末","库存现金","7"]):
            cell.text = value
        doc.add_heading("七、或有事项", level=1)
        doc.add_table(rows=2, cols=2).cell(0,0).text = "其他章节不可提取"
        source = self.folder / "含范围外链接.docx"
        doc.save(source)
        original = source.read_bytes()
        read = self.call("read", word=source, chapter_title=chapter)
        self.assertEqual(len(read["tables"]), 1)
        self.assertEqual(read["scope"]["table_count"], 1)
        out = self.folder / "不同输出目录"
        out.mkdir()
        excel, linked, context = (out / name for name in ("附注表格.xlsx","关联.docx","上下文.json"))
        self.call("extract", word=source, excel=excel, linked_word=linked, context=context, chapter_title=chapter)
        data = json.loads(context.read_text(encoding="utf-8-sig"))
        self.assertEqual(len(data["tables"]), 1)
        self.assertNotIn("五、税项", data["tables"][0]["chapter_context"])
        self.assertEqual(data["scope"]["table_count"], 1)
        book=load_workbook(excel)
        self.assertEqual(list(book.defined_names), ["TA_002"])
        book["sheet1"]["C3"]=125
        changed=out / "新附注表格.xlsx"
        book.save(changed)
        book.close()
        final=out / "更新后附注.docx"
        update=self.call("update", word=linked, excel=changed, output=final, chapter_title=chapter)
        self.assertEqual(update["successful_tables"], 1)
        self.assertEqual(update["scope"]["table_count"], 1)
        self.assertEqual(Document(final).tables[1].cell(1,1).text, "125")
        for candidate in (linked, final):
            with zipfile.ZipFile(source) as before, zipfile.ZipFile(candidate) as after:
                self.assertEqual(set(before.namelist()),set(after.namelist()))
                for name in before.namelist():
                    if name != "word/document.xml": self.assertEqual(before.read(name),after.read(name),name)
                a=ET.fromstring(before.read("word/document.xml")); b=ET.fromstring(after.read("word/document.xml"))
                abody=a.find(W+"body"); bbody=b.find(W+"body")
                selected=[i for i,n in enumerate(abody) if n.tag==W+"tbl"][1]
                for i in range(len(abody)):
                    if i != selected: self.assertEqual(self.xml_signature(abody[i]),self.xml_signature(bbody[i]))
        rejected=out / "不得生成.docx"
        bad=self.call("update", success=False, word=linked, excel=changed, output=rejected, chapter_title=chapter,
            operations=[{"range_name":"OUTSIDE","axis":"row","action":"delete","index":0}])
        self.assertIn("不存在",bad["error"])
        self.assertFalse(rejected.exists())
        self.assertEqual(source.read_bytes(),original)

    def test_chapter_scope_ignores_invalid_path_in_unrelated_link(self):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        doc=Document()
        doc.add_heading("五、税项",level=1)
        outside=doc.add_table(rows=1,cols=1)
        record=OxmlElement("w:tblDescription")
        record.set(qn("w:val"), '<#'+json.dumps({"SavePath":"irrelevant.xlsx", "RelativePath":"bad\u0000.xlsx", "LinkName":"OUTSIDE"})+'#>')
        outside._tbl.tblPr.append(record)
        doc.add_heading("六、财务报表重要项目的说明",level=1)
        doc.add_table(rows=1,cols=1).cell(0,0).text="7"
        source=self.folder / "范围外无效路径.docx"
        doc.save(source)
        out=self.folder / "另存"
        out.mkdir()
        excel,linked,context=(out / n for n in ("附注表格.xlsx","关联.docx","上下文.json"))
        self.call("extract",word=source,excel=excel,linked_word=linked,context=context,chapter_title="财务报表重要项目的说明")
        final=out / "原样回写.docx"
        self.call("update",word=linked,excel=excel,output=final,chapter_title="财务报表重要项目的说明")
        for candidate in (linked,final):
            self.assertEqual(Document(source).tables[0]._tbl.tblPr.xml,Document(candidate).tables[0]._tbl.tblPr.xml)

    def test_chapter_scope_missing_or_ambiguous_never_falls_back(self):
        chapter="财务报表重要项目的说明"
        missing=self.call("read", success=False, word=self.word, chapter_title=chapter)
        self.assertIn("未找到",missing["error"])
        duplicate=self.folder / "重复章节.docx"
        doc=Document()
        doc.add_heading(chapter,level=1)
        doc.add_table(rows=1,cols=1)
        doc.add_heading(chapter,level=1)
        doc.add_table(rows=1,cols=1)
        doc.save(duplicate)
        failed=self.call("read",success=False,word=duplicate,chapter_title=chapter)
        self.assertIn("多个",failed["error"])

    def test_real_chapter_scope_extracts_43_tables_and_roundtrips(self):
        source=Path(r"C:\Users\27651\Desktop\附注数据示例\示例附注.docx")
        if not source.exists(): self.skipTest("本机真实验收原件不存在")
        original=source.read_bytes()
        chapter="财务报表重要项目的说明"
        read=self.call("read",word=source,chapter_title=chapter)
        self.assertEqual(len(read["tables"]),43)
        self.assertEqual(read["scope"]["end_heading"],"或有事项")
        excel,linked,context=(self.folder / n for n in ("真实附注表格.xlsx","真实关联.docx","真实上下文.json"))
        self.call("extract",word=source,excel=excel,linked_word=linked,context=context,chapter_title=chapter)
        result=self.call("update",word=linked,excel=excel,output=self.folder / "真实原样回写.docx",chapter_title=chapter)
        self.assertEqual(result["successful_tables"],43)
        for candidate in (linked,self.folder / "真实原样回写.docx"):
            self.assertTrue(self.xml_signature(self.document_without_links(source)) == self.xml_signature(self.document_without_links(candidate)), "真实原样回写改变了正文结构或文字：" + str(candidate))
            with zipfile.ZipFile(source) as before, zipfile.ZipFile(candidate) as after:
                for name in before.namelist():
                    if name != "word/document.xml": self.assertEqual(before.read(name),after.read(name),name)
        self.assertEqual(source.read_bytes(),original)

    def test_default_table_roundtrip_preserves_entire_document(self):
        plain = self.folder / "默认表格报告.docx"
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = "页眉不得变化"
        doc.sections[0].footer.paragraphs[0].text = "页脚不得变化"
        doc.add_paragraph("合并财务报表附注 2025年12月31日 单位：人民币元")
        table = doc.add_table(rows=2, cols=2)
        for cell, text in zip([table.cell(0, 0), table.cell(0, 1), table.cell(1, 0), table.cell(1, 1)], ["项目", "期末", "库存现金", "7"]):
            cell.text = text
        doc.add_paragraph("表后正文与分节属性必须完整保留")
        doc.save(plain)
        original = plain.read_bytes()
        excel, linked, context = (self.folder / name for name in ("默认A.xlsx", "默认连接.docx", "默认上下文.json"))
        self.call("extract", word=plain, excel=excel, linked_word=linked, context=context)
        def document_tree(path):
            with zipfile.ZipFile(path) as package:
                root = ET.fromstring(package.read("word/document.xml"))
            for parent in root.iter():
                for child in list(parent):
                    if child.tag == W + "tblDescription":
                        parent.remove(child)
            return root
        def signature(node):
            return node.tag, tuple(sorted(node.attrib.items())), node.text or "", tuple(signature(child) for child in node), node.tail or ""
        self.assertEqual(signature(document_tree(plain)), signature(document_tree(linked)))
        book = load_workbook(excel)
        self.assertEqual(book["sheet1"]["C3"].value, 7)
        book["sheet1"]["C3"] = 125
        revised_excel = self.folder / "默认A更新.xlsx"
        book.save(revised_excel)
        book.close()
        output = self.folder / "默认报告更新.docx"
        self.call("update", word=linked, excel=revised_excel, output=output)
        expected = document_tree(plain)
        expected.find(".//" + W + "tbl").findall(W + "tr")[1].findall(W + "tc")[1].find(".//" + W + "t").text = "125"
        self.assertEqual(signature(expected), signature(document_tree(output)))
        with zipfile.ZipFile(plain) as before, zipfile.ZipFile(output) as after:
            for name in before.namelist():
                if name.startswith("word/") and name != "word/document.xml":
                    if name.endswith((".xml", ".rels")):
                        self.assertEqual(signature(ET.fromstring(before.read(name))), signature(ET.fromstring(after.read(name))), name)
                    else:
                        self.assertEqual(before.read(name), after.read(name), name)
        self.assertEqual(plain.read_bytes(), original)
    def test_extract_captures_context_and_links(self):
        excel, linked, context = self.extract()
        content = json.loads(context.read_text(encoding="utf-8-sig"))
        table = content["tables"][0]
        self.assertEqual(table["range_name"], "TA_001")
        self.assertIn("合并财务报表", "\n".join(table["chapter_context"]))
        self.assertIn("单位：人民币元", table["unit_text_candidates"])
        read = self.call("read", word=linked)
        self.assertEqual(read["tables"][0]["range_name"], "TA_001")
        self.assertEqual(self.word.read_bytes(), self.original)

    def test_context_preserves_table_break_without_assigning_previous_unit(self):
        doc = Document()
        doc.add_heading("六、财务报表重要项目的说明", level=1)
        doc.add_paragraph("本章除另有说明外，金额单位为人民币元。")
        for name in ("项目甲", "项目乙"):
            doc.add_heading(name, level=2)
            if name == "项目甲":
                doc.add_paragraph("以下这张表的金额单位为人民币万元，仅适用于本表。")
            table = doc.add_table(rows=2, cols=2)
            for row, values in zip(table.rows, [("项目", "期末余额"), ("库存现金", "10")]):
                for cell, value in zip(row.cells, values):
                    cell.text = value
        doc.add_heading("七、其他事项", level=1)
        source = self.folder / "单位说明有局部限定.docx"
        doc.save(source)
        before = source.read_bytes()
        excel, linked, context = (self.folder / name for name in ("分界表格.xlsx", "分界关联.docx", "分界原文.json"))
        self.call("extract", word=source, excel=excel, linked_word=linked, context=context,
                  chapter_title="财务报表重要项目的说明")
        tables = json.loads(context.read_text(encoding="utf-8-sig"))["tables"]
        self.assertEqual(len(tables), 2)
        second = tables[1]
        self.assertNotIn("unit_text", second, "不能以最近出现的说明替第二张表指定单位")
        self.assertEqual(second["unit_text_candidates"], [
            "本章除另有说明外，金额单位为人民币元。",
            "以下这张表的金额单位为人民币万元，仅适用于本表。"])
        blocks = second["preceding_blocks"]
        self.assertEqual([block["kind"] for block in blocks], ["paragraph"] * 4 + ["table", "paragraph"])
        self.assertEqual(blocks[3]["text"], "以下这张表的金额单位为人民币万元，仅适用于本表。")
        self.assertEqual(blocks[4]["table_id"], tables[0]["table_id"])
        self.assertEqual(blocks[5]["text"], "项目乙")
        self.assertEqual(source.read_bytes(), before)

    def test_update_preserves_format_headers_and_source(self):
        excel, linked, _ = self.extract()
        book = load_workbook(excel)
        book["sheet1"]["C3"] = 1234.5
        changed = self.folder / "表格A更新.xlsx"
        book.save(changed)
        output = self.folder / "更新报告.docx"
        result = self.call("update", word=linked, excel=changed, output=output, operations=[])
        self.assertTrue(result["output_verified"])
        doc = Document(output)
        self.assertEqual(doc.tables[0].cell(1, 1).text, "1234.50")
        self.assertEqual(doc.sections[0].header.paragraphs[0].text, "必须保留的报告页眉")
        self.assertEqual(doc.tables[0].cell(1, 1).paragraphs[0].runs[0].font.size, Pt(10))
        self.assertTrue(doc.tables[0].cell(3, 1).paragraphs[0].runs[0].bold)
        self.assertIn("表后说明必须保留。", [p.text for p in doc.paragraphs])
        self.assertEqual(self.word.read_bytes(), self.original)

    def test_empty_cell_inherits_own_paragraph_mark_font_and_keeps_existing_run_font(self):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        doc = Document(self.word)
        blank = doc.tables[0].cell(1, 1).paragraphs[0]
        for run in list(blank._p.findall(qn("w:r"))):
            blank._p.remove(run)
        properties = OxmlElement("w:rPr")
        for name, value in [("spacing", "-20"), ("sz", "24"), ("szCs", "24"), ("color", "123456")]:
            node = OxmlElement("w:" + name)
            node.set(qn("w:val"), value)
            properties.append(node)
        blank._p.get_or_add_pPr().append(properties)
        blank_signature = self.xml_signature(ET.fromstring(blank._p.xml))
        # 已有空文字片段具有自己的9磅字体，不能被段落标记12磅覆盖。
        existing = doc.tables[0].cell(1, 2).paragraphs[0]
        existing.runs[0].text = ""
        existing.runs[0].font.size = Pt(9)
        own_run_font = ET.fromstring(existing.runs[0]._r.rPr.xml)
        from copy import deepcopy
        existing._p.get_or_add_pPr().append(deepcopy(properties))
        source = self.folder / "含空白金额段落字体.docx"
        doc.save(source)
        source_bytes = source.read_bytes()
        excel, linked, context = (self.folder / n for n in ["空白格提取.xlsx", "空白格连接.docx", "空白格上下文.json"])
        self.call("extract", word=source, excel=excel, linked_word=linked, context=context)
        book = load_workbook(excel)
        for address, value in [("C3", 992997), ("D3", 17), ("C4", 456)]:
            book["sheet1"][address] = value
            book["sheet1"][address].number_format = "#,##0.00"
        changed = self.folder / "空白格更新.xlsx"
        book.save(changed)
        book.close()
        output = self.folder / "空白格回写.docx"
        self.call("update", word=linked, excel=changed, output=output)
        final = Document(output)
        new_blank = final.tables[0].cell(1, 1).paragraphs[0]
        self.assertEqual(new_blank.text, "992,997.00")
        self.assertEqual(new_blank.runs[0].font.size, Pt(12))
        self.assertEqual(self.xml_signature(ET.fromstring(new_blank.runs[0]._r.rPr.xml)),
                         self.xml_signature(ET.fromstring(properties.xml)))
        restored = deepcopy(new_blank._p)
        for run in list(restored.findall(qn("w:r"))):
            restored.remove(run)
        self.assertEqual(self.xml_signature(ET.fromstring(restored.xml)), blank_signature)
        self.assertEqual(self.xml_signature(ET.fromstring(final.tables[0].cell(1, 2).paragraphs[0].runs[0]._r.rPr.xml)),
                         self.xml_signature(own_run_font))
        self.assertEqual(final.tables[0].cell(2, 1).paragraphs[0].runs[0].font.size, Pt(10))
        self.assertEqual(source.read_bytes(), source_bytes)

    def test_insert_row_uses_detail_style_and_delete_row(self):
        excel, linked, _ = self.extract()
        changed = self.save_changed(excel, [
            ["项目","期末余额","期初余额"],["库存现金",100,90],
            ["银行存款",200,180],["其他货币资金",50,40],["合计",350,310]])
        output = self.folder / "增加明细.docx"
        self.call("update", word=linked, excel=changed, output=output, operations=[
            dict(range_name="TA_001", axis="row", action="insert", index=3, count=1, template_index=2)])
        doc = Document(output)
        self.assertEqual(len(doc.tables[0].rows),5)
        self.assertEqual(doc.tables[0].cell(3,0).text,"其他货币资金")
        self.assertFalse(bool(doc.tables[0].cell(3,1).paragraphs[0].runs[0].bold))
        self.assertTrue(doc.tables[0].cell(4,1).paragraphs[0].runs[0].bold)
        deleted = self.save_changed(excel,[["项目","期末余额","期初余额"],["库存现金",100,90],["合计",100,90]])
        output2 = self.folder / "减少明细.docx"
        self.call("update", word=linked, excel=deleted, output=output2, operations=[
            dict(range_name="TA_001",axis="row",action="delete",index=2,count=1)])
        self.assertEqual(len(Document(output2).tables[0].rows),3)
        self.assertEqual(Document(output2).tables[0].cell(2,0).text,"合计")

    def test_insert_and_delete_columns_updates_grid(self):
        excel, linked, _ = self.extract()
        changed = self.save_changed(excel,[["项目","期末余额","变动","期初余额"],["库存现金",100,10,90],["银行存款",200,20,180],["合计",300,30,270]])
        output = self.folder / "增加列.docx"
        self.call("update", word=linked,excel=changed,output=output,operations=[
            dict(range_name="TA_001",axis="column",action="insert",index=2,count=1,template_index=1)])
        doc = Document(output)
        self.assertEqual(len(doc.tables[0].columns),4)
        self.assertEqual(doc.tables[0].cell(0,2).text,"变动")
        reduced = self.save_changed(excel,[["项目","期末余额"],["库存现金",100],["银行存款",200],["合计",300]])
        output2 = self.folder / "减少列.docx"
        self.call("update",word=linked,excel=reduced,output=output2,operations=[
            dict(range_name="TA_001",axis="column",action="delete",index=2,count=1)])
        self.assertEqual(len(Document(output2).tables[0].columns),2)

    def test_refuses_input_overwrite_and_invalid_operations(self):
        excel, linked, _ = self.extract()
        before = linked.read_bytes()
        self.call("update",success=False,word=linked,excel=excel,output=linked,operations=[])
        self.assertEqual(linked.read_bytes(),before)
        output = self.folder / "不能生成.docx"
        self.call("update",success=False,word=linked,excel=excel,output=output,operations=[
            dict(range_name="TA_001",axis="row",action="delete",index=0,count=10)])
        self.assertFalse(output.exists())

    def test_merged_header_update_preserves_merge(self):
        doc = Document(self.word)
        doc.tables[0].cell(0, 1).merge(doc.tables[0].cell(0, 2))
        merged = self.folder / '合并表.docx'
        doc.save(merged)
        excel, linked, context = (self.folder / n for n in ['合并源.xlsx', '合并源连接.docx', '合并源上下文.json'])
        self.call('extract', word=merged, excel=excel, linked_word=linked, context=context)
        book = load_workbook(excel)
        book['sheet1']['C3'] = 888.5
        changed = self.folder / '合并新数.xlsx'
        book.save(changed)
        output = self.folder / '合并更新.docx'
        self.call('update', word=linked, excel=changed, output=output, operations=[])
        result = Document(output)
        self.assertEqual(result.tables[0].cell(1,1).text, '888.50')
        self.assertEqual(result.tables[0].cell(0,1)._tc, result.tables[0].cell(0,2)._tc)

    def test_rejects_vertical_merge_cut_and_dimension_mismatch(self):
        doc = Document(self.word)
        doc.tables[0].cell(1,0).merge(doc.tables[0].cell(2,0))
        merged = self.folder / '纵向合并.docx'
        doc.save(merged)
        excel, linked, context = (self.folder / n for n in ['纵向A.xlsx','纵向连接.docx','纵向上下文.json'])
        self.call('extract',word=merged,excel=excel,linked_word=linked,context=context)
        output = self.folder / '未生成.docx'
        result = self.call('update',success=False,word=linked,excel=excel,output=output,operations=[
            dict(range_name='TA_001',axis='row',action='delete',index=1,count=1)])
        self.assertIn('合并',result['error'])
        self.assertFalse(output.exists())
        book = load_workbook(excel)
        book.defined_names['TA_001'] = DefinedName('TA_001',attr_text="'sheet1'!$B$2:$D$4")
        wrong = self.folder / '区域错误.xlsx'
        book.save(wrong)
        self.call('update',success=False,word=linked,excel=wrong,output=output,operations=[])
        self.assertFalse(output.exists())
    def test_rejects_malformed_operation_protocol(self):
        excel, linked, _ = self.extract()
        for number, operations in enumerate([
            {'range_name':'TA_001','axis':'row','action':'delete','index':1},
            [dict(range_name='TA_001',axis='row',action='delete',index=1.2,count=1)],
            [dict(range_name='TA_001',axis='row',action='insert',index=2,count=True,template_index=1)]
        ]):
            output = self.folder / f'无效协议{number}.docx'
            self.call('update',success=False,word=linked,excel=excel,output=output,operations=operations)
            self.assertFalse(output.exists())
    @staticmethod
    def document_without_links(path):
        with zipfile.ZipFile(path) as package:
            root=ET.fromstring(package.read("word/document.xml"))
        for parent in root.iter():
            for child in list(parent):
                if child.tag==W+"tblDescription" and (child.get(W+"val") or "").startswith("<#"):
                    parent.remove(child)
        return root

    @staticmethod
    def xml_signature(node):
        return node.tag,tuple(sorted(node.attrib.items())),node.text or "",tuple(WordBridgeTests.xml_signature(c) for c in node),node.tail or ""

    def test_grid_before_roundtrip_and_real_cell_update_preserve_original_structure(self):
        source=ROOT/"依赖"/"audit-notes-tool"/"核心测试"/"测试资料"/"Word编号"/"07_行首省略一列.docx"
        before=source.read_bytes()
        excel,linked,context=(self.folder/n for n in ("行首省略.xlsx","行首省略连接.docx","行首省略上下文.json"))
        self.call("extract",word=source,excel=excel,linked_word=linked,context=context)
        book=load_workbook(excel);sheet=book["sheet1"]
        self.assertIsNone(sheet["B3"].value)
        self.assertEqual(sheet["C3"].value,"丁：从第二网格列开始")
        self.assertEqual(sheet["D3"].value,"戊")
        self.assertEqual([sheet.cell(2,c).value for c in range(2,5)],["甲","乙","丙"])
        self.assertEqual([sheet.cell(4,c).value for c in range(2,5)],["己","庚","辛"])
        with zipfile.ZipFile(linked) as package:
            tree=ET.fromstring(package.read("word/document.xml"))
        description=tree.find(".//"+W+"tblDescription").get(W+"val")
        self.assertTrue(json.loads(description[2:-2])["UseGridColumns"])
        unchanged=self.folder/"行首省略原样回写.docx"
        self.call("update",word=linked,excel=excel,output=unchanged,operations=[])
        original=self.document_without_links(source)
        self.assertEqual(self.xml_signature(original),self.xml_signature(self.document_without_links(unchanged)))
        sheet["C3"]="仅更新原本存在的第一格"
        changed=self.folder/"行首省略新数据.xlsx";book.save(changed);book.close()
        output=self.folder/"行首省略单格更新.docx"
        self.call("update",word=linked,excel=changed,output=output,operations=[])
        rows=original.find(".//"+W+"tbl").findall(W+"tr")
        self.assertEqual(len(rows[1].findall(W+"tc")),2)
        rows[1].findall(W+"tc")[0].find(".//"+W+"t").text="仅更新原本存在的第一格"
        self.assertEqual(self.xml_signature(original),self.xml_signature(self.document_without_links(output)))
        with zipfile.ZipFile(source) as old,zipfile.ZipFile(output) as new:
            self.assertEqual(set(old.namelist()),set(new.namelist()))
            for name in old.namelist():
                if name=="word/document.xml":continue
                if name.endswith((".xml",".rels")):
                    self.assertEqual(self.xml_signature(ET.fromstring(old.read(name))),self.xml_signature(ET.fromstring(new.read(name))),name)
                else:self.assertEqual(old.read(name),new.read(name),name)
        self.assertEqual(source.read_bytes(),before)

    def test_grid_before_omitted_position_cannot_receive_data(self):
        source=ROOT/"依赖"/"audit-notes-tool"/"核心测试"/"测试资料"/"Word编号"/"07_行首省略一列.docx"
        before=source.read_bytes()
        excel,linked,context=(self.folder/n for n in ("省略位置.xlsx","省略位置连接.docx","省略位置上下文.json"))
        self.call("extract",word=source,excel=excel,linked_word=linked,context=context)
        book=load_workbook(excel);book["sheet1"]["B3"]="原Word不存在此格，不可忽略这个新值"
        changed=self.folder/"误写省略位置.xlsx";book.save(changed);book.close()
        output=self.folder/"省略位置不得生成.docx"
        result=self.call("update",success=False,word=linked,excel=changed,output=output,operations=[])
        self.assertIn("省略",result["error"])
        self.assertFalse(output.exists());self.assertEqual(source.read_bytes(),before)
        for index,area in enumerate(("B3:C3","B2:B3","B2:C3")):
            book=load_workbook(excel);book["sheet1"].merge_cells(area)
            wrong=self.folder/("跨省略合并"+str(index)+".xlsx");book.save(wrong);book.close()
            rejected=self.folder/("跨省略合并不得生成"+str(index)+".docx")
            result=self.call("update",success=False,word=linked,excel=wrong,output=rejected,operations=[])
            self.assertIn("省略",result["error"]);self.assertFalse(rejected.exists())
        self.assertEqual(source.read_bytes(),before)

    def test_grid_before_with_vertical_merge_and_compressed_column_boundaries(self):
        source=ROOT/"依赖"/"audit-notes-tool"/"核心测试"/"测试资料"/"Word编号"/"07_行首省略一列.docx"
        for kind in ("纵向合并", "压缩边界"):
            with self.subTest(kind=kind):
                with zipfile.ZipFile(source) as original:
                    entries={name:original.read(name) for name in original.namelist()}
                document=ET.fromstring(entries["word/document.xml"])
                rows=document.find(".//"+W+"tbl").findall(W+"tr")
                if kind=="纵向合并":
                    first=rows[0].findall(W+"tc")[1];second=rows[1].findall(W+"tc")[0]
                    ET.SubElement(first.find(W+"tcPr"),W+"vMerge",{W+"val":"restart"})
                    ET.SubElement(second.find(W+"tcPr"),W+"vMerge")
                    for text in second.iter(W+"t"):text.text=""
                else:
                    for row in (rows[0],rows[2]):
                        cells=row.findall(W+"tc");row.remove(cells[1])
                        ET.SubElement(cells[0].find(W+"tcPr"),W+"gridSpan",{W+"val":"2"})
                        cells[0].find(W+"tcPr").find(W+"tcW").set(W+"w","4800")
                    cells=rows[1].findall(W+"tc");rows[1].remove(cells[1])
                    rows[1].find(W+"trPr").find(W+"gridBefore").set(W+"val","2")
                    rows[1].find(W+"trPr").find(W+"wBefore").set(W+"w","4800")
                derived=self.folder/(kind+"省略行首.docx")
                with zipfile.ZipFile(derived,"x",zipfile.ZIP_DEFLATED) as package:
                    for name,raw in entries.items():package.writestr(name,ET.tostring(document,encoding="utf-8",xml_declaration=True) if name=="word/document.xml" else raw)
                original_bytes=derived.read_bytes()
                excel,linked,context=(self.folder/(kind+n) for n in ("提取.xlsx","连接.docx","上下文.json"))
                self.call("extract",word=derived,excel=excel,linked_word=linked,context=context)
                book=load_workbook(excel);sheet=book["sheet1"]
                self.assertIsNone(sheet["B3"].value)
                if kind=="纵向合并":
                    self.assertIn("C2:C3",str(sheet.merged_cells));self.assertEqual(sheet["D3"].value,"戊")
                else:
                    self.assertEqual(sheet["C3"].value,"丁：从第二网格列开始")
                    self.assertEqual(list(book.defined_names["TA_001"].destinations)[0][1],"$B$2:$C$4")
                book.close()
                output=self.folder/(kind+"原样回写.docx")
                self.call("update",word=linked,excel=excel,output=output,operations=[])
                self.assertEqual(self.xml_signature(self.document_without_links(derived)),self.xml_signature(self.document_without_links(output)))
                self.assertEqual(derived.read_bytes(),original_bytes)

    def test_rejects_merge_crossing_without_publishing(self):
        doc = Document(self.word)
        doc.tables[0].cell(0,1).merge(doc.tables[0].cell(0,2))
        merged = self.folder / "合并表头.docx"
        doc.save(merged)
        excel, linked, context = (self.folder / n for n in ["合并A.xlsx","合并连接.docx","合并上下文.json"])
        self.call("extract",word=merged,excel=excel,linked_word=linked,context=context)
        output = self.folder / "不能切开合并.docx"
        result = self.call("update",success=False,word=linked,excel=excel,output=output,operations=[
            dict(range_name="TA_001",axis="column",action="delete",index=2,count=1)])
        self.assertIn("合并",result["error"])
        self.assertFalse(output.exists())


if __name__ == "__main__":
    unittest.main(verbosity=2)
