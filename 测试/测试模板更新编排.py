# coding: utf-8
"""已识别内容的编排只展开位置，不重写身份或计算原值。"""
import unittest
import json
from unittest.mock import patch
from 测试 import 测试统一结构 as 夹具
from 附注更新.统一勾稽 import build_fact_update_plan


class 模板更新编排测试(unittest.TestCase):
    def test_科目编排不能修改其他科目财务格(self):
        from 附注更新.模板更新编排 import compile_template_purpose
        case=夹具.UnifiedStructureTests();case.setUp()
        spec={'reason':'此回复只负责首个目标','tables':[],
              'bindings':[{'target':'明细!B3','source':'明细!B3'}],'labels':[]}
        with self.assertRaisesRegex(ValueError,'科目之外'):
            compile_template_purpose(case.target,case.source,case.standard,spec,target_locations={'明细!B2'})



    def test_换表名及移动行列后沿用同一编排机制(self):
        import copy
        from openpyxl import load_workbook
        from openpyxl.utils.cell import coordinate_to_tuple,get_column_letter
        from openpyxl.workbook.defined_name import DefinedName
        from 附注更新.表格 import file_hash,write_updated_workbook
        from 附注更新.模板更新编排 import compile_template_purpose
        case=夹具.UnifiedStructureTests();case.setUp()
        for mapping,name in [(case.target,'旧年明细副本'),(case.source,'新年来源表')]:
            book=load_workbook(mapping['source_path']);sheet=book['明细'];sheet.title=name
            sheet.insert_rows(1);sheet.insert_cols(1)
            end=5 if mapping is case.target else 6
            book.defined_names.clear();book.defined_names.add(DefinedName('移动后的附注',attr_text=f"'{name}'!$B$2:$C${end}"))
            book.save(mapping['source_path']);book.close();mapping['source_hash']=file_hash(mapping['source_path'])
            def moved(ref):
                _,address=ref.rsplit('!',1);row,col=coordinate_to_tuple(address)
                return name+'!'+get_column_letter(col+1)+str(row+1)
            mapping['owners']={}
            for fact in mapping['facts']:
                fact['id']=moved(fact['id']);fact['owner_reference']=fact['id']
                fact['source_reference']={'location':fact['id'],'file_hash':mapping['source_hash']}
                fact['dimension_evidence']={key:[moved(ref) for ref in refs] for key,refs in fact['dimension_evidence'].items()}
                mapping['owners'][fact['id']]=copy.deepcopy(fact['dimensions'])
        spec={'reason':'仅位置和表名改变，已核准身份相同','tables':[{'range_name':'移动后的附注',
            'source_sheet':'新年来源表','source_rows':[3,4,5],'old_detail_count':2,'first_row_offset':1,
            'columns':{'C':'C'},'labels':{'B':'B'},'totals':[{'target_row':5,'source_row':6}]}]}
        plan=build_fact_update_plan(case.target,case.source,case.standard,
            update_purpose=compile_template_purpose(case.target,case.source,case.standard,spec))
        self.assertEqual(plan['issues'],[])
        output=case.folder/'变式结果.xlsx';write_updated_workbook(case.target['source_path'],output,plan)
        book=load_workbook(output)
        self.assertEqual([[book['旧年明细副本'].cell(r,c).value for c in (2,3)] for r in range(3,7)],
            [['甲类设备',77],['乙类设备',77],['丙类设备',77],['合计',231]])
        book.close()

    def test_普通入口在需要增行时自动生成安排并写入(self):
        from 测试.测试统一语义 import records
        from 附注更新.统一语义 import read_standard_v2
        from 附注更新.统一识别 import UnifiedSemanticEngine
        from 附注更新.分步流程 import run_step
        from openpyxl import load_workbook
        case=夹具.UnifiedStructureTests();case.setUp()
        gold=case.folder/'本次标准.jsonl';gold.write_text('\n'.join(json.dumps(x,ensure_ascii=False) for x in records()),encoding='utf-8')
        case.standard=read_standard_v2(gold)
        for mapping in (case.target,case.source):mapping.update(gold_hash=case.standard['sha256'],gold_path=str(gold))
        a=case.complete_mapping(case.target);b=case.complete_mapping(case.source)
        spec={'reason':'三类资产的已确认名单随B更新','tables':[{'range_name':'附注表','source_sheet':'明细',
            'source_rows':[2,3,4],'old_detail_count':2,'first_row_offset':1,'columns':{'B':'B'},'labels':{'A':'A'},
            'totals':[{'target_row':4,'source_row':5}]}]}
        def response(engine,system,payload):
            if payload['task']!='template_update_layout':raise AssertionError('不应重新识别或比较A/B语义')
            return spec
        config={'gold_path':str(gold),'a_mapping':a['mapping_path'],'b_mapping':b['mapping_path'],
                'output_dir':str(case.folder/'普通入口'),'base_url':'http://localhost','model':'受控结构回复','api_key':'synthetic-test'}
        with patch.object(UnifiedSemanticEngine,'_http',response):
            matched=run_step('match',config)
            result=run_step('write',{**config,'update_plan':matched['files']['update_plan']})
        self.assertEqual(result['status'],'complete',result['message'])
        book=load_workbook(result['files']['a_prime'])
        self.assertEqual([[book['明细'].cell(r,c).value for c in (1,2)] for r in range(2,6)],
                         [['甲类设备',77],['乙类设备',77],['丙类设备',77],['合计',231]])
        book.close()

    def test_共用编排按来源行展开名单数据及合计(self):
        try:from 附注更新.模板更新编排 import compile_template_purpose
        except ImportError:self.fail('缺少正式入口可复用的模板更新编排')
        case=夹具.UnifiedStructureTests();case.setUp()
        spec={'reason':'已识别的三类资产随B更新','tables':[{
            'range_name':'附注表','source_sheet':'明细','source_rows':[2,3,4],
            'old_detail_count':2,'first_row_offset':1,'columns':{'B':'B'},'labels':{'A':'A'},
            'totals':[{'target_row':4,'source_row':5}]}]}
        purpose=compile_template_purpose(case.target,case.source,case.standard,spec)
        plan=build_fact_update_plan(case.target,case.source,case.standard,update_purpose=purpose)
        self.assertEqual({row['cell']:row['value'] for row in plan['updates']},{'B2':77,'B3':77,'B4':77,'B5':231})
        self.assertEqual([row['value'] for row in plan['labels']],['甲类设备','乙类设备','丙类设备'])
        self.assertEqual(len(plan['operations']),1)
        self.assertEqual(plan['operations'][0]['index'],3)


if __name__=='__main__':unittest.main()
