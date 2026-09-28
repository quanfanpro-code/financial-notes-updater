# coding: utf-8
"""语义身份独立于人工原值的格式，原值不被猜测或修正。"""
import copy
import unittest

from 附注更新.统一语义 import validate_records, validate_fact, semantic_key
from 附注更新.统一识别 import validate_response, agree_rounds
from 附注更新.统一勾稽 import build_fact_update_plan
from 测试.测试统一语义 import records, fact
from 测试.测试统一识别 import packet, answer
from 测试.测试统一勾稽 import mapping


class MeaningAndValueTests(unittest.TestCase):
    def setUp(self):
        self.standard = validate_records(records())
        self.standard['sha256'] = 'a'*64

    def test_original_value_format_does_not_discard_recognized_meaning(self):
        for value in ('—', '人工保留内容', '1,250.00', -125):
            with self.subTest(value=value):
                source = packet(); source['cells']['D7']['value'] = value
                checked = validate_response(answer(), source, self.standard)
                result = agree_rounds(checked, checked, self.standard)
                actual = next(row for row in result['facts'] if row['id'].endswith('!D7'))
                self.assertEqual(actual['raw_value'], value)
                self.assertEqual(semantic_key(actual, self.standard), semantic_key(fact(), self.standard))
                self.assertEqual(result['unresolved'], [])

    def test_same_unit_values_are_copied_without_interpretation(self):
        for value in ('—', '人工保留内容', '1,250.00', -125):
            with self.subTest(value=value):
                source = mapping('新布局!Z19', value)
                before = copy.deepcopy(source)
                plan = build_fact_update_plan(mapping('附注!D8'), source, self.standard)
                self.assertEqual(plan['updates'][0]['value'], value)
                self.assertEqual(plan['issues'], [])
                self.assertEqual(source, before)

    def test_unconvertible_value_is_not_a_semantic_failure(self):
        source = mapping('新布局!Z19', '人工保留内容', 10000)
        self.assertEqual(semantic_key(source['facts'][0], self.standard), semantic_key(fact(), self.standard))
        plan = build_fact_update_plan(mapping('附注!D8'), source, self.standard)
        self.assertEqual(plan['updates'], [])
        self.assertEqual([r['kind'] for r in plan['issues']], ['source_value_needs_review'])
        self.assertEqual(plan['issues'][0]['source_locations'], ['新布局!Z19'])
        self.assertEqual(plan['unplaced_sources'], [])
        self.assertEqual(source['facts'][0]['raw_value'], '人工保留内容')

    def test_raw_container_and_formula_state_are_still_validated(self):
        for value in ([], {}, float('nan'), float('inf'), '=1+1', ''):
            with self.subTest(value=value):
                row = fact(); row['raw_value'] = value
                with self.assertRaises(ValueError):
                    validate_fact(row, self.standard)

    def test_formal_excel_and_word_write_preserves_original_literal(self):
        # 合成夹具预设业务含义，验证真实文件入口；不冒充真实模型识别验收。
        import json
        import uuid
        from pathlib import Path
        from docx import Document
        from openpyxl import Workbook, load_workbook
        from 附注更新.表格 import file_hash, read_workbook, materialize_candidate_blanks
        from 附注更新.统一识别 import workbook_coverage
        from 附注更新.统一语义 import read_standard_v2
        from 附注更新.统一勾稽 import load_fact_mapping
        from 附注更新.分步流程 import run_step
        root = Path(__file__).resolve().parents[1]/'测试结果'/('语义与人工原值回写_'+uuid.uuid4().hex)
        root.mkdir()
        gold = root/'标准.jsonl'
        gold.write_text('\n'.join(json.dumps(r, ensure_ascii=False) for r in records()), encoding='utf-8')
        word = root/'合成附注.docx'; doc = Document()
        doc.add_paragraph('五、其他说明'); doc.add_paragraph('前章保持原样')
        doc.add_paragraph('六、财务报表重要项目的说明')
        doc.add_paragraph('合成企业甲；单户；人民币元；2025年12月31日')
        table = doc.add_table(rows=2, cols=2); table.style = 'Table Grid'
        for row, values in zip(table.rows, [['项目', '期末原值'], ['固定资产原值', '125.00']]):
            for cell, value in zip(row.cells, values): cell.text = value
        doc.add_paragraph('七、其他事项'); doc.add_paragraph('后章保持原样'); doc.save(word)
        config = {'output_dir': str(root/'正式入口'), 'gold_path': str(gold), 'word_path': str(word)}
        extracted = run_step('extract_a', config)
        self.assertEqual(extracted['status'], 'complete')
        target = Path(extracted['files']['a_path']); book = load_workbook(target)
        name = book.worksheets[0].title; old = book.worksheets[0]['C3'].value; book.close()
        source = root/'人工来源.xlsx'; book = Workbook(); sheet = book.active; sheet.title = '来源'
        sheet['A1'] = '合成企业甲；单户；人民币元；2025年12月31日固定资产原值合计'
        literal = '人工保留内容（原样）'; sheet['J3'] = literal; book.save(source); book.close()
        protected = {str(p): file_hash(p) for p in (word, target, source, gold)}
        paths = []
        for role, carrier, location, value, label in [('目标', target, name+'!C3', old, name+'!B1'), ('来源', source, '来源!J3', literal, '来源!A1')]:
            item = mapping(location, value)
            item.update(source_path=str(carrier), source_hash=file_hash(carrier), gold_path=str(gold), gold_hash=file_hash(gold))
            row = item['facts'][0]; row['source_reference']['file_hash'] = item['source_hash']
            row['dimension_evidence'] = {k: [label] for k in row['dimensions']}
            if role == '目标': item['a_link'] = extracted['files']['a_link']
            snapshot = read_workbook(carrier); materialize_candidate_blanks(snapshot)
            item['excluded'] = [{'sheet': s['name'], 'cell': address, 'kind': 'non_business', 'category': 'label', 'reason': '合成夹具的表头或留白'}
                                for s in snapshot['sheets'] for address in s['cells'] if s['name']+'!'+address != location]
            item['coverage'] = workbook_coverage(item, snapshot); item['complete'] = True
            path = root/(role+'语义映射.json'); path.write_text(json.dumps(item, ensure_ascii=False), encoding='utf-8-sig'); paths.append(path)
        config.update(a_mapping=str(paths[0]), b_mapping=str(paths[1]), a_link=extracted['files']['a_link'])
        matched = run_step('match', config); self.assertEqual(matched['status'], 'complete', matched['message'])
        written = run_step('write', {**config, **matched['files']})
        self.assertEqual(written['status'], 'complete', written['message'])
        result = Document(written['files']['word'])
        self.assertEqual(result.tables[0].cell(1, 1).text, literal)
        self.assertEqual(result.tables[0].style.style_id, table.style.style_id)
        self.assertEqual([p.text for p in result.paragraphs], [p.text for p in doc.paragraphs])
        derived = load_fact_mapping(written['files']['a_prime_mapping'], read_standard_v2(gold))
        self.assertEqual(derived['facts'][0]['raw_value'], literal)
        self.assertTrue(all(file_hash(p) == h for p, h in protected.items()))


if __name__ == '__main__':
    unittest.main()
