# coding: utf-8
"""开放维度同义复核行为；模拟回执只验证程序，不代表模型准确率。"""
import copy
import json
import unittest
import uuid
from pathlib import Path
from openpyxl import Workbook
from 测试.测试统一语义 import records, fact
from 附注更新.统一语义 import validate_records
from 附注更新.表格 import file_hash


class EquivalenceTests(unittest.TestCase):
    def setUp(self):
        from 附注更新 import 同义复核
        self.api=同义复核
        self.standard=validate_records(records());self.standard['sha256']='a'*64
        self.left=fact();self.right=fact()
        self.left['dimensions']['asset_selection']={'domain':'固定资产','mode':'members','members':['办公用电脑'],'completeness':'complete'}
        self.right['dimensions']['asset_selection']={'domain':'固定资产','mode':'predicate','predicate':'办公用电脑（用于办公的计算机）','completeness':'complete'}

    def test_only_open_dimensions_can_enter_review(self):
        self.assertTrue(self.api.reviewable(self.left,self.right,self.standard,same_cell=True))
        for key,value in [('entity','另一主体'),('unit_scale',10000),('period',{'kind':'instant','date':'2024-12-31'})]:
            changed=copy.deepcopy(self.right);changed['dimensions'][key]=value
            self.assertFalse(self.api.reviewable(self.left,changed,self.standard,same_cell=True))
        self.right['dimensions']['unit_scale']=10000
        self.assertTrue(self.api.reviewable(self.left,self.right,self.standard,same_cell=False))

    def test_partial_or_unknown_never_enters_review(self):
        for state in ['partial','unknown']:
            self.right['dimensions']['asset_selection']['completeness']=state
            self.assertFalse(self.api.reviewable(self.left,self.right,self.standard,same_cell=False))

    def fixture(self):
        folder=Path(__file__).resolve().parents[1]/'测试结果'/('同义复核行为_'+uuid.uuid4().hex);folder.mkdir()
        book=Workbook();sheet=book.active;sheet.title='当前表'
        sheet['A1']='合成企业甲，单户，人民币元，2025-12-31；办公用电脑即用于办公的计算机。'
        sheet['D8']=125;sheet['F9']=250
        path=folder/'合成资料.xlsx';book.save(path);book.close()
        for item,cell,value in [(self.left,'D8',125),(self.right,'F9',250)]:
            item['source_reference']={'file_hash':file_hash(path),'location':'当前表!'+cell}
            item['raw_value']=value
            item['dimension_evidence']={key:['当前表!A1'] for key in item['dimensions']}
        pair={'left':self.left,'right':self.right,'left_path':str(path),'right_path':str(path)}
        return folder,pair

    def response(self,payload,decision='equivalent'):
        return {'pairs':{key:{'decision':decision,'reason':'两侧原文完整说明同一集合',
                  'left_evidence':list(item['left']['evidence']), 'right_evidence':list(item['right']['evidence'])}
                        for key,item in payload['pairs'].items()}}

    def test_payload_reads_original_text_without_amounts(self):
        _,pair=self.fixture()
        payload=self.api.build_payload([pair],self.standard,same_cell=False)
        self.assertNotIn('raw_value',json.dumps(payload))
        self.assertIn('办公用电脑',json.dumps(payload,ensure_ascii=False))
        self.assertNotIn('125',json.dumps(list(payload['pairs'].values())))
        self.assertNotIn('250',json.dumps(list(payload['pairs'].values())))
        pair['right']['dimension_evidence']['asset_selection']=['当前表!Q99']
        with self.assertRaises(ValueError):self.api.build_payload([pair],self.standard,same_cell=False)

    def test_repeated_evidence_read_reuses_parse_but_changed_source_is_rechecked(self):
        from unittest.mock import patch
        from shutil import copy2
        from datetime import datetime
        folder,pair=self.fixture();path=Path(pair['left_path'])
        with patch.object(self.api,'load_workbook',wraps=self.api.load_workbook) as reader:
            first=self.api.build_payload([pair],self.standard,same_cell=False)
            again=self.api.build_payload([pair],self.standard,same_cell=False)
            self.assertEqual(first,again)
            self.assertEqual(reader.call_count,1,'同一版本原件反复核验不应重复解析整份工作簿')
            backup=Path(r'C:\Users\27651\BackUp')/('附注自动更新系统_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'_合成原文读取_'+uuid.uuid4().hex)/path.name
            backup.parent.mkdir(parents=True);copy2(path,backup)
            self.assertEqual(file_hash(path),file_hash(backup))
            book=Workbook();sheet=book.active;sheet.title='当前表'
            sheet['A1']='修改后的明确业务说明：办公用电脑即用于办公的计算机。'
            sheet['D8']=125;sheet['F9']=250;book.save(path);book.close()
            with self.assertRaisesRegex(ValueError,'原件版本已变化'):
                self.api.build_payload([pair],self.standard,same_cell=False)
            for side in ('left','right'):pair[side]['source_reference']['file_hash']=file_hash(path)
            fresh=self.api.build_payload([pair],self.standard,same_cell=False)
            self.assertEqual(reader.call_count,2)
            self.assertIn('修改后的明确业务说明',json.dumps(fresh,ensure_ascii=False))
            self.assertNotEqual(first,fresh)

    def test_response_requires_all_pairs_and_both_sides_evidence(self):
        _,pair=self.fixture();payload=self.api.build_payload([pair],self.standard,same_cell=False)
        response=self.response(payload)
        self.api.validate_response(response,payload)
        key=next(iter(response['pairs']))
        response['pairs'][key]['right_evidence']=['别的表!A1']
        with self.assertRaises(ValueError):self.api.validate_response(response,payload)
        with self.assertRaises(ValueError):self.api.validate_response({'pairs':{}},payload)
        isolated=copy.deepcopy(payload)
        isolated['pairs'][key]['left']['evidence']['当前表!A2']='只说明报告日期'
        row=self.response(isolated);row['pairs'][key]['left_evidence']=['当前表!A2']
        with self.assertRaises(ValueError):self.api.validate_response(row,isolated)

    def test_label_reference_evidence_reads_text_without_calculating_amounts(self):
        folder,pair=self.fixture()
        def make_source(name,formula):
            book=Workbook();sheet=book.active;sheet.title='当前表';other=book.create_sheet('名称 表')
            sheet['A1']="='名称 表'!$B$1";other['B1']=formula;other['C1']='办公用电脑即用于办公的计算机'
            sheet['D8']=125;sheet['F9']=250
            path=folder/name;book.save(path);book.close();current=copy.deepcopy(pair)
            for side in ('left','right'):
                current[side+'_path']=str(path);current[side]['source_reference']['file_hash']=file_hash(path)
            return current
        current=make_source('文字引用.xlsx','=C1');payload=self.api.build_payload([current],self.standard,same_cell=False)
        evidence=next(iter(payload['pairs'].values()))['left']['evidence']['当前表!A1']
        self.assertEqual(evidence['text'],'办公用电脑即用于办公的计算机')
        self.assertEqual(len(evidence['reference_chain']),3)
        for name,formula in [('数值公式.xlsx','=1+2'),('循环引用.xlsx','=B1'),('外部引用.xlsx',"='[外部.xlsx]名称'!A1"),('数字原文.xlsx','12345')]:
            invalid=make_source(name,formula)
            with self.assertRaises(ValueError):self.api.build_payload([invalid],self.standard,same_cell=False)

    def engine(self,folder):
        from 附注更新.统一识别 import UnifiedSemanticEngine
        gold=folder/'合成金标准.jsonl'
        gold.write_text('\n'.join(json.dumps(row,ensure_ascii=False) for row in records()),encoding='utf-8-sig')
        engine=UnifiedSemanticEngine({'model':'本地测试','base_url':'https://example.invalid/v1'},gold,folder/'回执')
        self.standard=engine.standard
        engine._http=lambda system,payload:self.response(payload)
        return engine

    def test_two_reviews_are_bound_and_single_disagreement_blocks(self):
        folder,pair=self.fixture();engine=self.engine(folder)
        proof=self.api.review_pairs(engine,[pair],same_cell=False)
        approved=self.api.verify_proof(proof,[pair],self.standard,same_cell=False)
        self.assertEqual(len(approved),1)
        changed=copy.deepcopy(pair);changed['right']['dimensions']['entity']='其他企业'
        with self.assertRaises(ValueError):self.api.verify_proof(proof,[changed],self.standard,same_cell=False)
        folder2=folder/'不同意见';folder2.mkdir();engine2=self.engine(folder2)
        engine2._http=lambda system,payload:self.response(payload,'equivalent' if payload['round']==1 else 'different')
        proof2=self.api.review_pairs(engine2,[pair],same_cell=False)
        self.assertEqual(self.api.verify_proof(proof2,[pair],self.standard,same_cell=False),set())

    def test_proof_rejects_changed_receipt_or_source(self):
        folder,pair=self.fixture();engine=self.engine(folder)
        proof=self.api.review_pairs(engine,[pair],same_cell=False)
        changed=copy.deepcopy(proof);changed['rounds'][0]['receipt']['sha256']='0'*64
        with self.assertRaises(ValueError):self.api.verify_proof(changed,[pair],self.standard,same_cell=False)

    def test_review_protocol_upgrade_keeps_old_proofs_and_rejects_relabeling(self):
        from unittest.mock import patch
        self.assertEqual(getattr(self.api,'PROOF_SCHEMA',None),'开放维度同义复核-v2')
        folder,pair=self.fixture();engine=self.engine(folder)
        # 创建旧协议下的真实本地任务与回执，再按升级后的读取器回读。
        with patch.object(self.api,'PROOF_SCHEMA','开放维度同义复核-v1'),patch.object(self.api,'PROMPT',self.api.LEGACY_PROMPT):
            legacy=self.api.review_pairs(engine,[pair],same_cell=False,shared=False)
        self.assertEqual(len(self.api.verify_proof(legacy,[pair],self.standard,same_cell=False)),1)
        current=self.api.review_pairs(engine,[pair],same_cell=False,shared=False)
        self.assertEqual(current['schema'],'开放维度同义复核-v2')
        self.assertEqual(len(self.api.verify_proof(current,[pair],self.standard,same_cell=False)),1)
        self.assertNotEqual(legacy['rounds'][0]['task']['sha256'],current['rounds'][0]['task']['sha256'])
        for source,version in [(legacy,'开放维度同义复核-v2'),(current,'开放维度同义复核-v1'),(current,'开放维度同义复核-v999')]:
            altered=copy.deepcopy(source);altered['schema']=version
            with self.assertRaises(ValueError):self.api.verify_proof(altered,[pair],self.standard,same_cell=False)

    def test_classification_uses_review_and_mapping_reloads(self):
        from 附注更新.统一勾稽 import load_fact_mapping
        folder,pair=self.fixture();engine=self.engine(folder)
        def respond(system,payload):
            if payload['task']=='unified_equivalence':return self.response(payload)
            selected=self.left if payload['round']==1 else self.right
            return {'cells':{'D8':{'kind':'metric_fact','metric_id':selected['metric_id'],
                'dimensions':selected['dimensions'],'dimension_evidence':{key:['A1'] for key in selected['dimensions']},'reason':'合成原文依据'}}}
        engine._http=respond
        packet={'source_path':pair['left_path'],'source_hash':self.left['source_reference']['file_hash'],'sheet':'当前表',
                'candidate_cells':['D8'],'cells':{'A1':{'value':'合成企业甲，单户，人民币元，2025-12-31；办公用电脑即用于办公的计算机。'},'D8':{'value':125}}}
        result=engine.classify_packet(packet)
        self.assertEqual(len(result['facts']),1)
        self.assertFalse(result['unresolved'])
        path=next(engine.work_dir.glob('逐格语义映射_*.json'))
        loaded=load_fact_mapping(path,engine.standard)
        self.assertEqual(loaded['facts'],result['facts'])
        # 篡改输出维度而保留原复核证明，应在加载时被拒绝。
        tampered=copy.deepcopy(result);tampered['facts'][0]['dimensions']['entity']='别的主体'
        tampered['owners'][tampered['facts'][0]['owner_reference']]['entity']='别的主体'
        other=folder/'篡改映射.json';other.write_text(json.dumps(tampered,ensure_ascii=False),encoding='utf-8-sig')
        with self.assertRaises(ValueError):load_fact_mapping(other,engine.standard)
        removed=copy.deepcopy(result);removed.pop('equivalence_proofs');removed.pop('classification_rounds')
        other=folder/'缺少复核依据.json';other.write_text(json.dumps(removed,ensure_ascii=False),encoding='utf-8-sig')
        with self.assertRaises(ValueError):load_fact_mapping(other,engine.standard)

    def test_frozen_reviews_survive_cache_retry_and_require_source_hash(self):
        folder,pair=self.fixture();engine=self.engine(folder)
        proof=self.api.review_pairs(engine,[pair],same_cell=False)
        engine._retry_receipts(0)
        self.assertEqual(len(self.api.verify_proof(proof,[pair],self.standard,same_cell=False)),1)
        changed=copy.deepcopy(pair);changed['left']['source_reference']['file_hash']='0'*64
        with self.assertRaises(ValueError):self.api.build_payload([changed],self.standard,same_cell=False)

    def test_partial_review_preserves_known_scope_without_inventing_missing_fields(self):
        from 附注更新.统一勾稽 import load_fact_mapping
        folder,pair=self.fixture();engine=self.engine(folder)
        def respond(system,payload):
            if payload['task']=='unified_equivalence':return self.response(payload)
            selected=self.left if payload['round']==1 else self.right
            dims={'asset_selection':selected['dimensions']['asset_selection']}
            return {'cells':{cell:({'kind':'partial_metric_fact','metric_id':selected['metric_id'],'dimensions':dims,
                'dimension_evidence':{'asset_selection':['A1']},'missing_dimensions':['entity','period','report_scope','currency','unit_scale'],
                'evidence_cells':['A1'],'reason':'仅保存模型有据识别的范围'} if cell in {'D8','F9'} else
                {'kind':'non_business','category':'annotation','reason':'合成测试说明或留白','evidence_cells':['A1']}) for cell in payload['candidate_cells']}}
        engine._http=respond
        result=engine.recognize_workbook(pair['left_path'])
        self.assertFalse(result['facts']);self.assertFalse(result['complete'])
        self.assertEqual(len(result['unresolved']),2)
        for row in result['unresolved']:
            known=row['partial_semantics']
            self.assertEqual(known['dimensions'],{'asset_selection':self.left['dimensions']['asset_selection']})
            self.assertEqual(set(known['missing_dimensions']),{'entity','period','report_scope','currency','unit_scale'})
        loaded=load_fact_mapping(result['mapping_path'],engine.standard)
        self.assertEqual(loaded['unresolved'],result['unresolved'])
        proof=self.api.mapping_review_proofs(loaded)[0]
        self.assertFalse(self.api.reviewable(proof['pairs'][0]['left'],proof['pairs'][0]['right'],engine.standard,same_cell=False))
        self.assertNotIn('raw_value',json.dumps(self.api.build_payload(proof['pairs'],engine.standard,same_cell=True)))
        altered_proof=copy.deepcopy(proof);altered_proof['schema']='开放维度同义复核-v2'
        with self.assertRaises(ValueError):self.api.verify_proof(altered_proof,proof['pairs'],engine.standard,same_cell=True)
        uncertain=copy.deepcopy(proof['pairs'][0]);uncertain['right']['dimensions']['asset_selection']['completeness']='partial'
        self.assertFalse(self.api.reviewable(uncertain['left'],uncertain['right'],engine.standard,same_cell=True))
        # 整表中的部分语义必须和已验证任务包相符，不能丢掉未决格或篡改范围。
        changed=copy.deepcopy(result);changed['unresolved'][0]['partial_semantics']['dimensions']['asset_selection']['members']=['厂房']
        altered=folder/'篡改部分语义.json';altered.write_text(json.dumps(changed,ensure_ascii=False),encoding='utf-8-sig')
        with self.assertRaises(ValueError):load_fact_mapping(altered,engine.standard)
        packets=[json.loads(Path(item['path']).read_text(encoding='utf-8-sig')) for item in result['packet_mappings']]
        packet=next(item for item in packets if item.get('equivalence_proofs'))
        packet.pop('equivalence_proofs');packet.pop('classification_rounds')
        altered=folder/'删除部分复核依据.json';altered.write_text(json.dumps(packet,ensure_ascii=False),encoding='utf-8-sig')
        with self.assertRaises(ValueError):load_fact_mapping(altered,engine.standard)

    def test_review_workbook_shows_both_meanings_and_decisions(self):
        from openpyxl import load_workbook
        from 附注更新.流程 import trace_workbook
        folder,pair=self.fixture();engine=self.engine(folder)
        proof=self.api.review_pairs(engine,[pair],same_cell=False)
        path=folder/'核对表.xlsx';trace_workbook(path,{'equivalence_proofs':[proof]})
        book=load_workbook(path)
        try:
            self.assertIn('维度同义复核',book.sheetnames)
            values=list(book['维度同义复核'].values)
            self.assertEqual(len(values),2)
            text=json.dumps(values,ensure_ascii=False)
            for expected in ('办公用电脑','用于办公的计算机','两轮同意','当前表!D8','当前表!F9'):
                self.assertIn(expected,text)
        finally:book.close()

    def mappings(self,pair):
        result=[]
        for side in ('left','right'):
            row=copy.deepcopy(pair[side]);row['id']=row['source_reference']['location'];row['owner_reference']='所属对象'
            result.append({'schema_version':2,'gold_hash':self.standard['sha256'],'source_hash':row['source_reference']['file_hash'],
                'source_path':pair[side+'_path'],'facts':[row],'owners':{'所属对象':copy.deepcopy(row['dimensions'])}})
        return result

    def test_cross_table_equivalence_moves_original_value_and_keeps_meaning(self):
        from 附注更新.统一勾稽 import build_fact_update_plan
        folder,pair=self.fixture();engine=self.engine(folder);left,right=self.mappings(pair)
        pairs=self.api.mapping_pairs(left,right,self.standard)
        proof=self.api.review_pairs(engine,pairs,same_cell=False)
        plan=build_fact_update_plan(left,right,self.standard,equivalence_proofs=[proof])
        self.assertEqual(plan['updates'][0]['value'],250)
        self.assertFalse(plan['issues']);self.assertFalse(plan['unplaced_sources'])
        self.assertEqual(right['facts'][0]['dimensions'],self.right['dimensions'])
        self.assertTrue(plan['updates'][0]['sources'][0]['equivalence_pair'])

    def test_equivalence_does_not_hide_conflicting_source_values(self):
        from 附注更新.统一勾稽 import build_fact_update_plan
        folder,pair=self.fixture();engine=self.engine(folder);left,right=self.mappings(pair)
        # 另一个已精确识别的来源拥有同样语义但不同金额，须和同义来源一起检查冲突。
        exact=copy.deepcopy(left['facts'][0]);exact['id']='另一个来源';exact['owner_reference']='另一个所属对象'
        right['owners'][exact['owner_reference']]=copy.deepcopy(exact['dimensions']);right['facts'].append(exact)
        pairs=self.api.mapping_pairs(left,right,self.standard)
        proof=self.api.review_pairs(engine,pairs,same_cell=False)
        plan=build_fact_update_plan(left,right,self.standard,equivalence_proofs=[proof])
        self.assertFalse(plan['updates'])
        self.assertIn('conflicting_sources',[row['kind'] for row in plan['issues']])

    def test_formal_match_write_and_derived_mapping_with_equivalence(self):
        from unittest.mock import patch
        from 附注更新.统一识别 import UnifiedSemanticEngine,workbook_coverage
        from 附注更新.统一勾稽 import load_fact_mapping
        from 附注更新.表格 import read_workbook,materialize_candidate_blanks
        from 附注更新.分步流程 import run_step
        folder,pair=self.fixture();engine=self.engine(folder)
        paths=[]
        for side in ('left','right'):
            item=copy.deepcopy(pair[side]);book=Workbook();sheet=book.active;sheet.title='当前表'
            sheet['A1']='合成企业甲，单户，人民币元，2025-12-31；办公用电脑即用于办公的计算机。'
            sheet['B2']=item['raw_value'];path=folder/(side+'资料.xlsx');book.save(path);book.close()
            item['source_reference']={'file_hash':file_hash(path),'location':'当前表!B2'}
            item.update(id='当前表!B2',owner_reference='本表对象')
            mapping={'schema_version':2,'gold_hash':engine.gold_hash,'gold_path':engine.gold_path,
                     'source_path':str(path),'source_hash':file_hash(path),'facts':[item],
                     'owners':{'本表对象':copy.deepcopy(item['dimensions'])},'unresolved':[]}
            snapshot=read_workbook(path);materialize_candidate_blanks(snapshot)
            mapping['excluded']=[{'sheet':s['name'],'cell':address,'category':'label','reason':'合成夹具明确的表头或空白'}
                for s in snapshot['sheets'] for address in s['cells'] if address!='B2']
            mapping['coverage']=workbook_coverage(mapping,snapshot);mapping['complete']=True
            destination=folder/(side+'映射.json');destination.write_text(json.dumps(mapping,ensure_ascii=False),encoding='utf-8-sig');paths.append(destination)
        config={'gold_path':engine.gold_path,'a_mapping':str(paths[0]),'b_mapping':str(paths[1]),
                'output_dir':str(folder/'正式入口'),'model':'本地测试','base_url':'https://example.invalid/v1','api_key':'','timeout':600}
        with patch.object(UnifiedSemanticEngine,'_http',lambda instance,system,payload:self.response(payload)):
            matched=run_step('match',config)
        self.assertEqual(matched['status'],'complete',matched['message'])
        with patch.object(UnifiedSemanticEngine,'_http',side_effect=AssertionError('写入不能重调模型')):
            written=run_step('write',{**config,**matched['files']})
        self.assertEqual(written['status'],'complete',written['message'])
        loaded=load_fact_mapping(written['files']['a_prime_mapping'],engine.standard)
        self.assertEqual(loaded['facts'][0]['raw_value'],250)
        self.assertEqual(loaded['facts'][0]['dimensions'],self.left['dimensions'])

    def test_routing_restricts_candidates_to_matched_b_tables(self):
        """A端导向表级路由：A端格只与路由到的B端逻辑表内事实组合，路由外同指标格不进复核。"""
        folder,pair=self.fixture()
        left,right=self.mappings(pair)
        routed=copy.deepcopy(right['facts'][0]);routed['id']='当前表!F9';routed['owner_reference']='所属对象'
        routed['source_reference']={'file_hash':right['facts'][0]['source_reference']['file_hash'],
                                    'location':'QCF42 按账龄披露应收账款!C6'}
        outside=copy.deepcopy(right['facts'][0]);outside['id']='QCF24!C6';outside['owner_reference']='路由外所属'
        outside['source_reference']={'file_hash':right['facts'][0]['source_reference']['file_hash'],
                                     'location':'QCF24 单项计提坏账准备的应收账款!C6'}
        right['facts']=[routed,outside]
        right['owners']={'所属对象':copy.deepcopy(routed['dimensions']),'路由外所属':copy.deepcopy(outside['dimensions'])}
        routing={'a_tables':[{'sheet':'当前表','rows':(1,20),'title':'按账龄披露应收账款'}],
                 'b_tables':[{'sheet':'QCF42 按账龄披露应收账款','rows':None,'title':'按账龄披露应收账款'},
                             {'sheet':'QCF24 单项计提坏账准备的应收账款','rows':None,'title':'单项计提坏账准备的应收账款'}],
                 'links':[(0,0)]}
        pairs=self.api.mapping_pairs(left,right,self.standard,routing=routing)
        involved={item['right']['source_reference']['location'] for item in pairs}
        self.assertTrue(any(loc.startswith('QCF42') for loc in involved),'路由到的B端逻辑表候选必须保留')
        self.assertFalse(any(loc.startswith('QCF24') for loc in involved),'路由外的同指标格不得进入复核')

    def test_routing_matches_titles_when_links_absent(self):
        """路由表未给出显式links时，按A端表标题与B端逻辑表标题的字面包含自动匹配。"""
        folder,pair=self.fixture()
        left,right=self.mappings(pair)
        routed=copy.deepcopy(right['facts'][0]);routed['id']='当前表!F9';routed['owner_reference']='所属对象'
        routed['source_reference']={'file_hash':right['facts'][0]['source_reference']['file_hash'],
                                    'location':'货币资金余额情况!C7'}
        outside=copy.deepcopy(right['facts'][0]);outside['id']='坏账!C6';outside['owner_reference']='路由外所属'
        outside['source_reference']={'file_hash':right['facts'][0]['source_reference']['file_hash'],
                                     'location':'坏账准备情况!C6'}
        right['facts']=[routed,outside]
        right['owners']={'所属对象':copy.deepcopy(routed['dimensions']),'路由外所属':copy.deepcopy(outside['dimensions'])}
        routing={'a_tables':[{'sheet':'当前表','rows':(1,20),'title':'货币资金'}],
                 'b_tables':[{'sheet':'货币资金余额情况','rows':None,'title':'货币资金余额情况'},
                             {'sheet':'坏账准备情况','rows':None,'title':'坏账准备情况'}]}
        pairs=self.api.mapping_pairs(left,right,self.standard,routing=routing)
        involved={item['right']['source_reference']['location'] for item in pairs}
        self.assertTrue(any(loc.startswith('货币资金余额情况') for loc in involved),'标题包含匹配的B端表候选保留')
        self.assertFalse(any(loc.startswith('坏账准备情况') for loc in involved),'标题不相关的B端表不进复核')

    def test_mapping_pairs_prefilters_unrelated_candidates(self):
        """A端导向：每格只与确定性兼容的候选复核；成员完全无关的候选被预筛排除。"""
        folder,pair=self.fixture()
        left,right=self.mappings(pair)
        unrelated=copy.deepcopy(right['facts'][0])
        unrelated['id']='当前表!B3';unrelated['owner_reference']='无关所属对象'
        unrelated['source_reference']={'file_hash':right['facts'][0]['source_reference']['file_hash'],
                                       'location':'当前表!B3'}
        unrelated['dimensions']['asset_selection']={'domain':'固定资产','mode':'members','members':['厂房'],'completeness':'complete'}
        right['owners']['无关所属对象']=copy.deepcopy(unrelated['dimensions'])
        right['facts'].append(unrelated)
        pairs=self.api.mapping_pairs(left,right,self.standard)
        involved={(item['left']['id'],item['right']['id']) for item in pairs}
        self.assertTrue(any(item[1]=='当前表!F9' for item in involved),'同义候选必须保留送复核')
        self.assertFalse(any(item[1]=='当前表!B3' for item in involved),'成员完全无关的候选必须被预筛排除')


if __name__=='__main__':unittest.main()
