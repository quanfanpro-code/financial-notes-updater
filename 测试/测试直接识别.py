# coding: utf-8
"""只用合成文件和模拟传输，验证正式入口没有隐藏的第二次识别。"""
import json
from pathlib import Path
import unittest
import uuid
from unittest.mock import patch

from openpyxl import Workbook
from 附注更新.分步流程 import run_step
from 附注更新.统一识别 import UnifiedSemanticEngine
from 附注更新.统一勾稽 import load_fact_mapping, compare_semantic_slots
from 附注更新.表格 import file_hash
from 附注更新.语义 import SemanticError
from 测试.测试统一语义 import records


WORK = Path(__file__).resolve().parents[1] / '测试结果'


class 直接识别测试(unittest.TestCase):
    def setUp(self):
        self.folder = WORK / ('直接识别离线_' + uuid.uuid4().hex)
        self.folder.mkdir(parents=True)
        self.source = self.folder / '合成原表.xlsx'
        self.gold = self.folder / '合成金标准.jsonl'
        self.gold.write_text('\n'.join(json.dumps(r, ensure_ascii=False) for r in records()), encoding='utf-8')
        book = Workbook()
        sheet = book.active
        sheet.title = '合成表'
        sheet['A1']='合成企业；单体；人民币；元'
        sheet['B2'] = '本期'
        sheet['C2'] = '上期'
        for row in range(3, 9):
            sheet.cell(row, 1, '测试项目' + str(row))
            for column in (2, 3):
                cell = sheet.cell(row, column)
                cell.number_format = '0.00'
                if row < 5:
                    cell.value = row * column
        book.save(self.source)
        book.close()
        self.config = {'a_path': str(self.source), 'b_path': str(self.source), 'gold_path': str(self.gold),
                       'output_dir': str(self.folder / '运行'), 'base_url': 'codex://local', 'model': 'offline-test'}
        # 本组测试的是已给定范围的一次识别；整份自动范围由正常入口贯通测试覆盖。
        for side in ('a','b'):
            self.config[side+'_target_table']={'id':'本表','sheet':'合成表','range':'A2:C8','data_ranges':['B3:C8']}
        self.hashes = (file_hash(self.source), file_hash(self.gold))

    def engine(self, url='codex://local'):
        return UnifiedSemanticEngine({'base_url': url, 'model': 'offline-test', 'timeout': 5},
                                     self.gold, self.folder / '缓存')

    def answer(self,year=2025):
        dims=[{'id':key,'value':value,'ranges':['*'],'evidence':['合成表!A1']} for key,value in
              [('entity','合成企业'),('report_scope','standalone'),('currency','CNY'),('unit_scale',1)]]
        for col,date in [('B',f'{year}-12-31'),('C',f'{year-1}-12-31')]:
            dims.append({'id':'period','value':{'kind':'instant','date':date},'ranges':[f'{col}3:{col}8'],
                         'evidence':[f'合成表!{col}2'],'label_references':[f'合成表!{col}2']})
        for row in range(3,9):
            dims.append({'id':'asset_selection','value':{'domain':'固定资产','mode':'members','members':['测试项目'+str(row)],'completeness':'complete'},
                         'ranges':[f'B{row}:C{row}'],'evidence':[f'合成表!A{row}'],'label_references':[f'合成表!A{row}']})
        return json.dumps({'tables':[{'id':'本表','sheet':'合成表','range':'A2:C8','data_ranges':['B3:C8'],
            'metrics':[{'id':'固定资产原值','ranges':['B3:B8'],'template_period_role':'closing'},
                       {'id':'固定资产原值','ranges':['C3:C8'],'template_period_role':'opening'}],
            'dimensions':dims}]},ensure_ascii=False)

    def test_裸期间角色保存及同编号重绑保留旧角色(self):
        from copy import deepcopy
        from 附注更新.A合集 import save_direct_answer, _records, _merge_direct_table, _standard
        from 附注更新.A语义阅读 import addresses
        with patch.object(UnifiedSemanticEngine,'_http',return_value=self.answer()):
            run=run_step('recognize_a',self.config)
        old=load_fact_mapping(run['files']['a_mapping'],{})
        reply='- B3:B4：期初余额\n- B5:B6：本期增加\n- C3:C4：本期减少\n- C5:C6：期末余额'
        def role_and_label(engine,prompt,payload):
            self.assertIn('同时明确本期、上期、期初或期末角色',prompt)
            self.assertIn('已有期间只作只读参照',prompt)
            return reply
        with patch.object(UnifiedSemanticEngine,'request_text',role_and_label):
            preview=run_step('recognize_a',{**self.config,'a_mapping':run['files']['a_mapping'],
                'a_context_text':'period只补明确的原表标签引用，不改日期',
                'a_field_scope':{'合成表!B3:C6':['period','template_period_role']}})
        self.assertEqual(preview['status'],'complete',preview)
        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('不得重答')):
            saved=save_direct_answer(old,reply,self.folder/'裸期间角色.json',field_scope={
                area:['template_period_role'] for area in ('B3:B6','C3:C6')})
            rebound=save_direct_answer(saved,'合成表!B3:C6：metric_id=固定资产原值。',
                self.folder/'同编号保留角色.json',field_scope={'B3:C6':['metric_id']})
        self.assertEqual(rebound['机械记录'],old['机械记录'])
        self.assertEqual({k:v['dimensions'] for k,v in _records(rebound)},
                         {k:v['dimensions'] for k,v in _records(old)})
        roles={a:m.get('template_period_role') for t in rebound['outline']['subtables']
               for m in t['metrics'] for span in m['ranges'] for a in addresses(span)}
        for area,expected in [('B3:B4','opening'),('B5:B6','current'),('C3:C4','current'),('C5:C6','closing')]:
            for cell in addresses(area):self.assertEqual(roles[cell],expected)
        standard=deepcopy(_standard(old))
        standard['metric']['另一指标']={**standard['metric']['固定资产原值'],'id':'另一指标'}
        with patch('附注更新.A合集._standard',return_value=standard):
            for identity,fields in [('另一指标',['metric_id']),('固定资产原值',['metric_id','template_period_role'])]:
                merged,_,_=_merge_direct_table(old,{'metrics':[{'id':identity,'ranges':['B3']}],
                    'dimensions':[],'unresolved':[]},{'B3':fields})
                self.assertNotIn('template_period_role',next(m for m in merged['metrics'] if m['ranges']==['B3']))

    def test_期间标签双向写法及分别对应不交叉(self):
        from 附注更新.A合集 import direct_answer
        engine=self.engine()
        reply=('合成表!B3:B4对应原表标签坐标 合成表!B2，标签原文为“本期”；期间角色为本期。\n'
               '合成表!C3:C4对应 合成表!C2“上期”，合成表!D3:D4对应 合成表!D2“期初余额”。\n'
               'E2、F2 的“期末余额”分别适用于 E3:E4、F3:F4；这些列的角色为期末。')
        parsed=direct_answer(reply,engine.standard,['合成表'])
        periods={tuple(r['ranges']):r['evidence'] for r in parsed['dimensions'] if r['id']=='period'}
        for col in 'BCDEF':self.assertEqual(periods[(f'合成表!{col}3',f'合成表!{col}4')],[f'合成表!{col}2'])
        roles={tuple(r['ranges']):r['value'] for r in parsed['dimensions'] if r['id']=='template_period_role'}
        for col,role in [('B','current'),('C','prior'),('D','opening'),('E','closing'),('F','closing')]:
            self.assertEqual(roles[(f'合成表!{col}3:{col}4',)],role)
        parsed=direct_answer('合成表!B3、B4对应合成表!B2“本年”，C3、C4对应C2“上年”。',engine.standard,['合成表'])
        periods={tuple(r['ranges']):r['evidence'] for r in parsed['dimensions'] if r['id']=='period'}
        for col in 'BC':self.assertEqual(periods[(f'合成表!{col}3',f'合成表!{col}4')],[f'合成表!{col}2'])
        parsed=direct_answer('- **合成表!B3:C4**：template_period_role 为**期末**，包括空白数据格 B4:C4。对应原表合并表头 **合成表!B2:C2「期末余额」**。',engine.standard,['合成表'])
        labels=[r for r in parsed['dimensions'] if r['id']=='period']
        self.assertEqual(len(labels),1)
        self.assertEqual(set(labels[0]['ranges']),{'合成表!B3','合成表!C3','合成表!B4','合成表!C4'})
        self.assertEqual(labels[0]['evidence'],['合成表!B2:C2'])
        parsed=direct_answer('合成表!B3:B4、C3:C4：均对应合成表!B2:C2合并标签“期初余额”，适用于所列数据格。',engine.standard,['合成表'])
        labels=[r for r in parsed['dimensions'] if r['id']=='period']
        self.assertEqual(len(labels),1)
        self.assertEqual(set(labels[0]['ranges']),{'合成表!B3','合成表!C3','合成表!B4','合成表!C4'})
        self.assertEqual(labels[0]['evidence'],['合成表!B2:C2'])
        parsed=direct_answer('metric_id=固定资产原值。B/C列分别为本期（2025-01-01至2025-12-31）和上期（2024-01-01至2024-12-31）。\n'
            '| 坐标 | asset_selection |\n|---|---|\n| 合成表!B3、合成表!C3 | all |',engine.standard,['合成表'])
        self.assertTrue(any(r['id']=='template_period_role' and r['ranges']==['合成表!B3'] and r['value']=='current' for r in parsed['dimensions']))
        self.assertTrue(any(r['id']=='template_period_role' and r['ranges']==['合成表!C3'] and r['value']=='prior' for r in parsed['dimensions']))

    def test_范围尾部独立标准编号不混入业务取值(self):
        from 附注更新.A合集 import direct_answer
        standard=self.engine().standard
        for tail,expected in [('固定资产原值','测试项目'),('原文字词','测试项目，`原文字词')]:
            parsed=direct_answer(f'合成表!B3：asset_selection=测试项目，`{tail}`。',standard,['合成表'])
            row=next(r for r in parsed['dimensions'] if r['id']=='asset_selection')
            self.assertEqual(row['value'],expected)

    def test_已明示顺序的同行字段不依赖重复指标编号(self):
        from 附注更新.A合集 import direct_answer
        reply=('### 明细\n合成表!B3:B4、C3:C4为明细范围。\n'
               '对应账龄依次为：一年以内、一年以上。明细的 asset_selection 为“固定资产中，区间为该行所列区间的项目”。')
        parsed=direct_answer(reply,self.engine().standard,['合成表'])
        rows=[r for r in parsed['dimensions'] if r['id']=='asset_selection']
        self.assertEqual(len(rows),4)
        for row in rows:
            expected='一年以内' if row['ranges'][0].endswith('3') else '一年以上'
            self.assertEqual(row['value'],f'固定资产中，区间为{expected}的项目')

    def test_同句明确列顺序接通指代而不跨句撤销字段(self):
        from 附注更新.A合集 import direct_answer
        raw=('- **合成表!B3:E4**：明细表；B、C、D列分别属于第一、第二、第三阶段，E列为这些阶段的合计。'
             '各格选 `固定资产原值`；`asset_selection` 为固定资产中该阶段的项目，E列为各阶段相应事项的合计范围。'
             '原文未明确的方向或变动内容保留未知。')
        def rows(text):
            return [r for r in direct_answer(text,self.engine().standard,['合成表'])['dimensions'] if r['id']=='asset_selection']
        result=rows(raw)
        self.assertFalse(any(r['value']=='未知' for r in result))
        for col,name in [('B','第一'),('C','第二'),('D','第三')]:
            row=next(r for r in result if f'合成表!{col}3' in r['ranges'])
            self.assertEqual(row['value'],f'固定资产中{name}阶段的项目')
        self.assertEqual(next(r['value'] for r in result if '合成表!E3' in r['ranges']),'各阶段相应事项的合计范围')
        for text in (raw.replace('分别属于','属于'),raw.replace('B、C、D列','B、C列')):
            self.assertFalse(any(r['value']=='固定资产中第一阶段的项目' for r in rows(text)))
        self.assertTrue(any(r['value']=='未知' for r in rows('合成表!B3：`asset_selection` 保留未知。')))

    def test_原句分类编号与方向仅按明示分组接入(self):
        from copy import deepcopy
        from 附注更新.A合集 import direct_answer
        from 附注更新.A语义阅读 import addresses
        standard=deepcopy(self.engine().standard)
        standard['dimension']['test_direction']={'id':'test_direction','value_type':'enum','allowed_values':['甲向','乙向']}
        for key in ('metric.test.balance','metric.test.change'):
            standard['metric'][key]={**deepcopy(standard['metric']['固定资产原值']),'id':key}
        standard['metric']['metric.test.change']['dimension_refs'].append({'id':'test_direction'})
        raw=('C列为期初；期初为2024-01-01。\n### 本表\n'
             '期初、期末分别为2024-01-01、2024-12-31，本期变动期间为2024-01-01至2024-12-31。\n'
             '- 合成表!B3、E3：期初、期末余额；C3、D3：本期甲向、乙向。\n'
             '  余额用 `metric.test.balance`，变动用 `metric.test.change`，变动方向分别为甲向、乙向。\n'
             '- 合成表!B4、E4：期初、期末余额；C4、D4：本期甲向、乙向。指标及变动方向同上。')
        parsed=direct_answer(raw,standard,['合成表'])
        def values(kind,key=None):
            result={}
            for row in parsed[kind]:
                if key and row['id']!=key:continue
                for span in row['ranges']:
                    if span=='*':continue
                    for cell in addresses(span.split('!')[-1]):result.setdefault(cell,set()).add(str(row.get('value',row['id'])))
            return result
        ids=values('metrics');roles=values('dimensions','template_period_role');dates=values('dimensions','period')
        for n in (3,4):
            for col in 'BE':self.assertEqual(ids[col+str(n)],{'metric.test.balance'})
            for col in 'CD':
                self.assertEqual(ids[col+str(n)],{'metric.test.change'})
                self.assertEqual(roles[col+str(n)],{'current'})
                self.assertEqual(dates[col+str(n)],{'2024-01-01至2024-12-31'})
            self.assertEqual(dates['B'+str(n)],{'2024-01-01'})
            self.assertEqual(dates['E'+str(n)],{'2024-12-31'})
        self.assertEqual(values('dimensions','test_direction')['C4'],{'甲向'})
        missing=direct_answer(raw.replace('，变动方向分别为甲向、乙向',''),standard,['合成表'])
        self.assertFalse(any(r['id']=='metric.test.change' for r in missing['metrics']))
        separated=direct_answer(raw.replace('- 合成表!B4','### 另一段\n- 合成表!B4'),standard,['合成表'])
        self.assertFalse(any('合成表!C4' in r['ranges'] for r in separated['metrics']))

    def test_原句本期列名及明示顺序角色不遗漏(self):
        from 附注更新.A合集 import direct_answer
        from 附注更新.A语义阅读 import addresses
        for declaration in ('B“期初余额”、C“本期增加”、D“本期减少”、E“期末余额”',
                            '列角色依次为期初余额、本期增加、本期减少、期末余额',
                            'C为本期增加、D为本期减少',
                            '这些格按原文列名分别是“本期应交”和“本期已交”'):
            scope='C3:D4' if declaration.startswith(('C为','这些')) else 'B3:E4'
            parsed=direct_answer(f'合成表!{scope}：固定资产原值；{declaration}。',self.engine().standard,['合成表'])
            roles={cell:r['value'] for r in parsed['dimensions'] if r['id']=='template_period_role'
                   for span in r['ranges'] for cell in addresses(span.rsplit('!',1)[-1])}
            for col in 'CD':
                for n in (3,4):self.assertEqual(roles[col+str(n)],'current')
        parsed=direct_answer('合成表!C3:D4：这些格按原文列名分别是“本期应交”和“上期已交”。',self.engine().standard,['合成表'])
        self.assertFalse(any(r['id']=='template_period_role' for r in parsed['dimensions']))

    def test_标签对应列简写只承接本句明确行范围(self):
        from 附注更新.A合集 import direct_answer
        raw=('在“合成表”中，B2“本期增加”适用于 B3:B4、B6；C2“本期减少”适用于对应的 C 列各范围。\n'
             '在“另一表”中，B2“本期应交”适用于 B5；C2“本期已交”适用于对应的 C 列各范围。')
        rows=direct_answer(raw,self.engine().standard,['合成表','另一表'])['dimensions']
        refs={cell:row['evidence'] for row in rows if row['id']=='period' for cell in row['ranges']}
        for cell in ('C3','C4','C6'):self.assertEqual(refs['合成表!'+cell],['合成表!C2'])
        self.assertEqual(refs['另一表!C5'],['另一表!C2'])
        self.assertNotIn('另一表!C6',refs)
        with self.assertRaises(ValueError):
            direct_answer('在“合成表”中，C2“本期减少”适用于对应的 C 列各范围。',self.engine().standard,['合成表','另一表'])

    def test_期间普通叙述不生成字段且币种止于原句(self):
        from 附注更新.A合集 import direct_answer
        raw=('金额币种为人民币（CNY）。原文另有说明。合成表的B列为本期、C列为上期；'
             '依报表年度2025年度识别期间：本期为2025-01-01至2025-12-31，上期为2024-01-01至2024-12-31。\n'
             '合成表!B3:C3：固定资产原值；asset_selection=金融工具持有期间取得的收益；期间分别为本期、上期。')
        parsed=direct_answer(raw,self.engine().standard,['合成表'])
        self.assertEqual([r['value'] for r in parsed['dimensions'] if r['id']=='currency'],['人民币（CNY）'])
        self.assertTrue(all('得的收益'!=r['value'] and not str(r['value']).startswith('本期为')
                            for r in parsed['dimensions'] if r['id']=='period'))
        self.assertFalse(any(r['id']=='period' and r['value'] in ('本期','上期') for r in parsed['dimensions']))
        self.assertTrue(any(r['id']=='asset_selection' and r['value']=='金融工具持有期间取得的收益' for r in parsed['dimensions']))
        changed=direct_answer(raw.replace('\n合成表!','\n别表!A1：普通说明。\n合成表!'),self.engine().standard,['合成表','别表'])
        self.assertFalse(any(r['id']=='period' and r['value'] in ('本期','上期') for r in changed['dimensions']))

    def test_共同选取方式标题不当作实际取值(self):
        from 附注更新.A合集 import direct_answer
        raw=('本次列明的4格，asset_selection均为 **predicate**。共同说明。\n'
             '- **合成表!B3:C3**：完整范围甲。\n- **合成表!B4:C4**：完整范围乙。')
        rows=[r for r in direct_answer(raw,self.engine().standard,['合成表'])['dimensions'] if r['id']=='asset_selection']
        self.assertEqual([(r['ranges'],r['value']) for r in rows],[(['合成表!B3:C3'],'完整范围甲'),(['合成表!B4:C4'],'完整范围乙')])

    def test_加粗字段组与成员数组仅读取字面列表(self):
        from copy import deepcopy
        from 附注更新.A合集 import direct_answer
        standard=deepcopy(self.engine().standard)
        standard['dimension']['第二范围']={**standard['dimension']['asset_selection'],'id':'第二范围'}
        reply='- **合成表!B3:C3 — `asset_selection`：** `mode=members`，`members=["甲", "乙"]`；**`第二范围`：** `mode=members`，`members=["事项"]`。解释文字。'
        parsed=direct_answer(reply,standard,['合成表'])
        self.assertEqual([(r['id'],r['ranges'],r['value']) for r in parsed['dimensions']],
            [('asset_selection',['合成表!B3:C3'],{'mode':'members','members':['甲','乙']}),
             ('第二范围',['合成表!B3:C3'],{'mode':'members','members':['事项']})])
        reply='**asset_selection（domain：固定资产）**：以下各坐标均为 `mode=members`，且只选该行类别。\n| B坐标 | asset_selection 实际取值 |\n|---|---|\n| 合成表!B3:C3 | members：房屋 |'
        parsed=direct_answer(reply,standard,['合成表'])
        self.assertEqual(len(parsed['dimensions']),1)
        self.assertEqual(parsed['dimensions'][0]['ranges'],['合成表!B3:C3'])
        self.assertEqual(parsed['dimensions'][0]['value'],'members：房屋')
        parsed=direct_answer('合成表!B3：asset_selection={domain: 固定资产, mode: members, members: [房屋, 设备]}。\n'
            '**期间实际取值**（期间为 2024-01-01 至 2024-12-31）：\n'
            '合成表!B3：period={start: 2024-01-01, end: 2024-12-31, kind: duration}。',standard,['合成表'])
        self.assertEqual(next(r['value'] for r in parsed['dimensions'] if r['id']=='asset_selection')['members'],['房屋','设备'])
        self.assertTrue(any(r['id']=='period' and r['value']=='2024-01-01至2024-12-31' for r in parsed['dimensions']))

    def test_同一行范围仅承接紧邻原句明示取值(self):
        from 附注更新.A合集 import direct_answer
        standard=self.engine().standard
        first='合成表!B3:B4、C3:C4：相应固定资产范围依次为房屋、设备；对应`固定资产原值`。\n'
        for reference in ('上述同一两类资产','这些资产'):
            parsed=direct_answer(first+f'合成表!D3:E4：{reference}的变动格；指标为`固定资产原值`。',standard,['合成表'])
            found={(r['ranges'][0],r['value']) for r in parsed['dimensions'] if r['id']=='asset_selection' and r['ranges'][0].startswith(('合成表!D','合成表!E'))}
            self.assertEqual(found,{(f'合成表!{col}{row}',value) for col in 'DE' for row,value in [(3,'房屋'),(4,'设备')]} if reference.startswith('上述同一') else set())
        parsed=direct_answer('合成表!B5：asset_selection=all；`固定资产原值`。\n合成表!D5：同一行范围；指标为`固定资产原值`。',standard,['合成表'])
        self.assertTrue(any(r['id']=='asset_selection' and r['ranges']==['合成表!D5'] and r['value']=='all' for r in parsed['dimensions']))
        parsed=direct_answer('合成表!B3:B4、E3:E4：相应固定资产范围依次为房屋、设备。B列为期初余额，E列为期末余额。对应`固定资产原值`。\n合成表!C3:D4：上述同一两类资产的变动格；指标为`固定资产原值`。',standard,['合成表'])
        values=[r for r in parsed['dimensions'] if r['id']=='asset_selection']
        for col in 'BCDE':
            for row,value in [(3,'房屋'),(4,'设备')]:
                self.assertEqual([r['value'] for r in values if f'合成表!{col}{row}' in r['ranges']],[value])

    def test_期间标签不被无关维度引用阻断(self):
        from copy import deepcopy
        from 附注更新.A合集 import save_direct_answer
        from 附注更新.模板更新编排 import direct_update_plan
        with patch.object(UnifiedSemanticEngine,'_http',return_value=self.answer()):
            result=run_step('recognize_a',self.config)
        base=load_fact_mapping(result['files']['a_mapping'],{})
        snapshot=deepcopy(base['机械记录']);snapshot['sheets'][0]['cells']['B2']['value']='新的本期'
        source=save_direct_answer({**base,'carrier':'B'},self.answer(2026),self.folder/'新期间B.json',snapshot=snapshot)
        for shared in (False,True):
            answer=json.loads(self.answer());dims=answer['tables'][0]['dimensions'];period=dims[4]
            period['ranges']=['B3']
            dims.append({**period,'ranges':['B5:B8'],'evidence':['合成表!C2']})
            dims.append({**period,'ranges':['B4'],'evidence':['合成表!B2' if shared else '合成表!C2']})
            dims[7]['evidence']=['合成表!B2']
            target=save_direct_answer(base,json.dumps(answer,ensure_ascii=False),self.folder/f'同标签其他维度{shared}.json')
            plan=direct_update_plan(target,source,{'reason':'换期'},
                {'target':[{'sheet':'合成表','range':'B3:B4'}],'source':[{'sheet':'合成表','range':'B3'}]})
            self.assertEqual([r['cell'] for r in plan['updates']],[] if shared else ['B3'])
        unchanged=save_direct_answer({**base,'carrier':'B'},self.answer(2026),self.folder/'标签原字不变B.json')
        plan=direct_update_plan(target,unchanged,{'reason':'换期'},
            {'target':[{'sheet':'合成表','range':'B3:B4'}],'source':[{'sheet':'合成表','range':'B3'}]})
        self.assertEqual([r['cell'] for r in plan['updates']],['B3'])
        self.assertEqual(plan['labels'],[])
        snapshot['sheets'][0]['cells']['D2']={**snapshot['sheets'][0]['cells']['B2'],'column':4}
        split=json.loads(self.answer(2026));dims=split['tables'][0]['dimensions'];dims[4]['ranges']=['B3','B5:B8']
        dims.append({**dims[4],'ranges':['B4'],'evidence':['合成表!D2']})
        source=save_direct_answer({**base,'carrier':'B'},json.dumps(split,ensure_ascii=False),self.folder/'同字不同出处B.json',snapshot=snapshot)
        plan=direct_update_plan(target,source,{'reason':'换期'},
            {'target':[{'sheet':'合成表','range':'B3:B4'}],'source':[{'sheet':'合成表','range':'B3:B4'}]})
        self.assertEqual({r['cell'] for r in plan['updates']},{'B3','B4'})
        self.assertEqual(len(plan['labels']),1)

    def test_单表字段纠正保留其他工作表的完整标签坐标(self):
        from copy import deepcopy
        from 附注更新.A合集 import save_direct_answer,_records
        with patch.object(UnifiedSemanticEngine,'_http',return_value=self.answer()):
            result=run_step('recognize_a',self.config)
        base=load_fact_mapping(result['files']['a_mapping'],{});base['carrier']='B'
        snapshot=deepcopy(base['机械记录']);extra=deepcopy(snapshot['sheets'][0]);extra['name']='附注汇总（单体）';snapshot['sheets'].append(extra)
        base=save_direct_answer(base,self.answer(),self.folder/'带其他工作表B.json',snapshot=snapshot)
        saved=save_direct_answer(base,'合成表!B3:B4对应`附注汇总（单体）!B2`“本年”；C3:C4对应`附注汇总（单体）!C2`“上年”。',
            self.folder/'跨表标签引用.json',field_scope={'合成表!B3:C4':['period']})
        rows=dict(_records(saved))
        for col in 'BC':
            for row in (3,4):self.assertEqual(rows[f'合成表!{col}{row}']['dimension_evidence']['period'],[f'附注汇总（单体）!{col}2'])

    def test_正式A一次答复及八个空白保存回读并复用(self):
        answer = self.answer()
        sent = []
        def fake(engine, prompt, payload):
            sent.append((prompt, payload))
            self.assertEqual(payload['task'], 'direct_recognition')
            material = payload['materials']['原文.json']
            target_rows = [row for row in material['sheets'][0]['rows'].splitlines() if 3 <= int(row.split(':', 1)[0]) <= 8]
            self.assertEqual(len(target_rows), 6)
            self.assertEqual('\n'.join(target_rows).count('[空白]'), 8)
            self.assertNotIn('返回JSON', prompt)
            self.assertNotIn('checked_relations', prompt)
            return answer
        with patch.object(UnifiedSemanticEngine, '_http', fake):
            first = run_step('recognize_a', self.config)
            self.assertEqual(first['status'], 'complete', first)
            again = run_step('recognize_a', {**self.config, 'a_mapping': first['files']['a_mapping']})
        self.assertEqual(len(sent), 1)
        self.assertEqual(again['usage']['cached_requests'], 1)
        path = Path(first['files']['a_mapping'])
        saved_paths = {p.resolve() for p in self.folder.rglob('A识别结果.json')}
        self.assertEqual(saved_paths, {path.resolve(), Path(again['files']['a_mapping']).resolve()},
                         '共用出口已保存的合集不应在外层再保存一份')
        original = Path.read_bytes
        def only_result(p):
            self.assertEqual(p.resolve(), path.resolve(), '回读调用了原件或旧核查链')
            return original(p)
        with patch.object(Path, 'read_bytes', only_result):
            saved = load_fact_mapping(path, {})
        self.assertEqual(saved['识别答复'], answer)
        self.assertEqual(len(saved['facts']),12)
        self.assertEqual(sum(r['raw_value_state']=='blank' for r in saved['facts']),8)
        from 附注更新.界面 import _read_record, _mapping_sources
        self.assertEqual(_read_record(path, '附注语义映射-v1', 'A')['识别答复'], answer)
        self.assertEqual(_mapping_sources(path, 'A')['a_path'], str(self.source))
        from 附注更新.A合集 import _standard
        self.assertEqual(compare_semantic_slots(saved,saved,_standard(saved))['matched_targets'],12)
        self.assertEqual(self.hashes, (file_hash(self.source), file_hash(self.gold)))

    def test_B只发一次任务并复用A原答复不复制身份(self):
        answers = iter([self.answer(), self.answer(2026)])
        sent = []
        def fake(engine, prompt, payload):
            sent.append(payload)
            return next(answers)
        with patch.object(UnifiedSemanticEngine, '_http', fake):
            a = run_step('recognize_a', self.config)
            b = run_step('recognize_b', {**self.config, 'a_mapping': a['files']['a_mapping']})
        self.assertEqual(b['status'], 'complete', b)
        self.assertEqual(len(sent), 2)
        needs=sent[1]['materials']['A需要.json']
        self.assertEqual(len(needs['需要']),12)
        self.assertNotIn('识别答复',needs)
        result = load_fact_mapping(b['files']['b_mapping'], {})
        self.assertEqual(result['识别答复'], self.answer(2026))
        self.assertNotIn('semantic_reading', result)

    def test_正常入口换期匹配写回及合集接续(self):
        from openpyxl import load_workbook
        bpath=self.folder/'B.xlsx'
        book=load_workbook(self.source);ws=book.active
        ws['B2']='2026年末';ws['C2']='2025年末'
        for row in range(3,9):
            for col in (2,3):ws.cell(row,col,1000+row*col)
        book.save(bpath);book.close()
        config={**self.config,'b_path':str(bpath)}
        a_reply=self.answer().replace('"closing"','"期末余额"').replace('"opening"','"期初余额"')
        with patch.object(UnifiedSemanticEngine,'_http',side_effect=[a_reply,self.answer(2026)]) as send:
            a=run_step('recognize_a',config);config.update(a['files'])
            b=run_step('recognize_b',config);config.update(b['files'])
            match=run_step('match',config);self.assertEqual(match['status'],'complete',match);config.update(match['files'])
            written=run_step('write',config)
            self.assertEqual(send.call_count,2)
        self.assertEqual(written['status'],'complete',written)
        self.assertEqual(written['updated_cells'],12)
        book=load_workbook(written['files']['a_prime']);ws=book.active
        self.assertEqual(ws['B3'].value,1006);self.assertEqual(ws['C8'].value,1024)
        self.assertEqual(ws['B2'].value,'2026年末');book.close()
        record=json.loads(Path(written['files']['a_prime_mapping']).read_text(encoding='utf-8-sig'))
        self.assertEqual(record['schema'],'附注更新执行记录-v1')
        self.assertEqual(record['collections']['A']['mapping_path'],a['files']['a_mapping'])
        self.assertEqual(record['collections']['B']['mapping_path'],b['files']['b_mapping'])
        self.assertNotIn('facts',record)
        def reuse(engine,prompt,materials):
            self.assertEqual(materials['上次搬运记录.json'],record)
            self.assertIn('上次A原判断.json',materials)
            self.assertIn('上次B原判断.json',materials)
            self.assertIn('仅对变动后含义或上下文尚未确定的部分作识别',prompt)
            return self.answer(2026)
        with patch.object(UnifiedSemanticEngine,'request_text',reuse):
            continued=run_step('recognize_a',{**config,**written['files'],'a_path':written['files']['a_prime'],'a_mapping':''})
        self.assertEqual(continued['status'],'complete',continued)
        self.assertEqual(self.hashes,(file_hash(self.source),file_hash(self.gold)))

    def test_公式副本取值经合集读取及正常安排入口(self):
        from copy import deepcopy
        from openpyxl import load_workbook
        from 附注更新.A合集 import direct_digest
        from 附注更新.模板更新编排 import direct_update_plan
        from 附注更新.分步流程 import _prepare_update_v2
        from 附注更新.流程 import save_json
        bpath=self.folder/'公式B.xlsx'
        book=load_workbook(self.source);ws=book.active
        for cell in ('B3','C3','B4'):ws[cell]='=1+2'
        book.save(bpath);book.close()
        config={**self.config,'b_path':str(bpath)}
        with patch.object(UnifiedSemanticEngine,'_http',return_value=self.answer()):
            a=run_step('recognize_a',config);config.update(a['files'])
            b=run_step('recognize_b',config)
        ar=load_fact_mapping(a['files']['a_mapping'],{});br=load_fact_mapping(b['files']['b_mapping'],{})
        self.assertEqual(len(direct_update_plan(ar,br)['updates']),1)
        evidence=save_json(self.folder/'合成计算依据.json',{'workbook_path':str(bpath)})
        br['calculation_evidence']={'path':evidence,'sha256':file_hash(evidence)}
        br['formula_values']={
            '合成表!B3':{'formula':'=1+2','state':'value','value':3,'reason':'离线合成计算结果'},
            '合成表!C3':{'formula':'=1+2','state':'blank','value':'','reason':'离线合成空白'},
            '合成表!B4':{'formula':'=1+2','state':'error','value':'#VALUE!','reason':'离线合成计算错误'}}
        br['封存校验']=direct_digest(br)
        before=deepcopy(br)
        with patch('openpyxl.load_workbook',side_effect=AssertionError('匹配不能读表补身份')):
            plan=direct_update_plan(ar,br)
        self.assertEqual(br,before)
        self.assertEqual({r['cell']:r['value'] for r in plan['updates']},{'B3':3,'C4':12})
        self.assertTrue(any(r['cell']=='C3' for r in plan['retained_blanks']))
        self.assertTrue(any(r.get('reason')=='离线合成计算错误' for r in plan['issues']))
        missing=deepcopy(br);missing.pop('calculation_evidence');missing['封存校验']=direct_digest(missing)
        with self.assertRaisesRegex(ValueError,'缺少计算依据'):direct_update_plan(ar,missing)
        wrong=deepcopy(br);wrong['formula_values']['合成表!B3']['formula']='=9';wrong['封存校验']=direct_digest(wrong)
        with self.assertRaisesRegex(ValueError,'与原公式不符'):direct_update_plan(ar,wrong)
        saved=save_json(self.folder/'带计算依据B合集.json',br)
        # Excel原生计算不在离线测试中启动；这里只验证已有证据的消费路径。
        with patch('附注更新.统一勾稽.calculation_files',return_value=[evidence,str(bpath)]) as checked:
            loaded=load_fact_mapping(saved,{})
            checked.assert_called_once_with(loaded)
            prepared=_prepare_update_v2(config,self.folder,lambda _:None,None,
                                        Path(a['files']['a_mapping']),Path(saved))
        from 附注更新.A合集 import save_direct_answer
        def existing_values(mapping,receipt):
            return {row['source_reference']['location']:br['formula_values'][row['source_reference']['location']]
                    for row in mapping['facts'] if row['raw_value_state']=='formula'}
        with patch('附注更新.统一勾稽._calculated_values',side_effect=existing_values) as reused, \
             patch('附注更新.统一勾稽.calculation_files',return_value=[evidence,str(bpath)]), \
             patch('附注更新.统一勾稽.recalculate_copy',side_effect=AssertionError('不重新计算')):
            partial=save_direct_answer(br,'合成表!B3：asset_selection=未知。',self.folder/'部分身份沿用公式副本.json',
                                      field_scope={'合成表!B3':['asset_selection']})
            reused.assert_called_once()
            self.assertEqual(set(partial['formula_values']),{'合成表!C3','合成表!B4'})
            self.assertEqual(br,before)
        arranged=json.loads(Path(prepared['files']['update_plan']).read_text(encoding='utf-8-sig'))
        self.assertEqual(arranged['updates'],plan['updates'])
        self.assertIn(str(Path(evidence).resolve()),arranged['inputs'])
        self.assertIn(str(bpath.resolve()),arranged['inputs'])
        self.assertEqual(self.hashes,(file_hash(self.source),file_hash(self.gold)))

    def test_名单改值缺口不影响无共享标签的其他动作(self):
        from openpyxl import load_workbook
        from 附注更新.模板更新编排 import direct_update_plan
        bpath=self.folder/'名单B.xlsx';book=load_workbook(self.source);ws=book.active
        ws['A3']='新项目';ws['B3']=300;ws['C3']=301
        book.save(bpath);book.close()
        answer=json.loads(self.answer())
        selection=next(r for r in answer['tables'][0]['dimensions'] if r['id']=='asset_selection')
        selection['value']['members']=['新项目']
        # 第4行具体范围未知；明确指定来源也不能放行。
        next(r for r in answer['tables'][0]['dimensions'] if r['id']=='asset_selection' and r['ranges']==['B4:C4'])['value']={'domain':'固定资产','mode':'unknown','completeness':'unknown'}
        config={**self.config,'b_path':str(bpath)}
        with patch.object(UnifiedSemanticEngine,'_http',side_effect=[self.answer(),json.dumps(answer,ensure_ascii=False)]):
            a=run_step('recognize_a',config);config.update(a['files']);b=run_step('recognize_b',config)
        ar=load_fact_mapping(a['files']['a_mapping'],{});br=load_fact_mapping(b['files']['b_mapping'],{})
        purpose={'reason':'合成离线要求：名单随B更新','bindings':[{'target':f'合成表!{c}{r}','source':f'合成表!{c}{r}'} for r in (3,4) for c in 'BC'],
                 'labels':[{'target':'合成表!A3','source':'合成表!A3'}]}
        # 编排只读合集，即使禁止读取任何工作簿也应产生相同安排。
        with patch('openpyxl.load_workbook',side_effect=AssertionError('编排不得重新读表判断')):
            plan=direct_update_plan(ar,br,purpose)
        self.assertEqual({r['cell'] for r in plan['updates']},{'B3','C3'})
        self.assertEqual(plan['labels'][0]['value'],'新项目')
        self.assertTrue(any('身份仍缺失' in r['reason'] for r in plan['issues']))
        config.update(b['files']);config['update_purpose']=purpose
        with patch.object(UnifiedSemanticEngine,'_http',side_effect=AssertionError('后续不能调用模型')):
            match=run_step('match',config);config.update(match['files']);written=run_step('write',config)
        self.assertEqual(written['status'],'partial',written)
        book=load_workbook(written['files']['a_prime']);ws=book.active
        self.assertEqual(ws['A3'].value,'新项目');self.assertEqual(ws['B3'].value,300)
        self.assertEqual(ws['B4'].value,8);book.close()
        from copy import deepcopy
        from 附注更新.A合集 import direct_digest
        empty=deepcopy(br);empty['机械记录']['sheets'][0]['cells']['A3']['value']=None
        empty['封存校验']=direct_digest(empty)
        blocked=direct_update_plan(ar,empty,purpose)
        self.assertEqual(blocked['updates'],[])
        self.assertTrue(any('标签没有可直接搬运' in r['reason'] for r in blocked['issues']))

    def test_名单新增行只搬运B身份并经正常入口写回(self):
        from openpyxl import load_workbook
        from openpyxl.workbook.defined_name import DefinedName
        book=load_workbook(self.source);book.defined_names.add(DefinedName('名单表',attr_text="'合成表'!$A$2:$C$8"))
        apath=self.folder/'带范围A.xlsx';book.save(apath)
        ws=book.active;ws['A9']='新增项目'
        for row in range(3,10):
            for col in (2,3):ws.cell(row,col,row*100+col)
        bpath=self.folder/'新增名单B.xlsx';book.save(bpath);book.close()
        answer=json.loads(self.answer());t=answer['tables'][0];t['range']='A2:C9';t['data_ranges']=['B3:C9']
        for m in t['metrics']:m['ranges']=[m['ranges'][0].replace('8','9')]
        for dim in t['dimensions']:
            if dim['id']=='period':dim['ranges']=[dim['ranges'][0].replace('8','9')]
        t['dimensions'].append({'id':'asset_selection','ranges':['B9:C9'],'value':{'domain':'固定资产','mode':'members','members':['新增项目'],'completeness':'complete'},'evidence':['合成表!A9']})
        purpose={'reason':'合成明确要求：采用B七个成员，新增一行，已有名单行作样式',
            'tables':[{'range_name':'名单表','source_sheet':'合成表','source_rows':list(range(3,10)),
                       'old_detail_count':6,'first_row_offset':1,'columns':{'B':'B','C':'C'},'labels':{'A':'A'}}]}
        config={**self.config,'a_path':str(apath),'b_path':str(bpath),'update_purpose':purpose}
        hashes=(file_hash(apath),file_hash(bpath))
        with patch.object(UnifiedSemanticEngine,'_http',side_effect=[self.answer(),json.dumps(answer,ensure_ascii=False)]) as send:
            a=run_step('recognize_a',config);self.assertEqual(a['status'],'complete',a);config.update(a['files'])
            b=run_step('recognize_b',config);self.assertEqual(b['status'],'complete',b);config.update(b['files'])
            match=run_step('match',config);config.update(match['files'])
            written=run_step('write',config)
        self.assertEqual(written['status'],'complete',written)
        self.assertEqual(send.call_count,2)
        book=load_workbook(written['files']['a_prime']);ws=book.active
        self.assertEqual(ws['A9'].value,'新增项目');self.assertEqual(ws['C9'].value,903);book.close()
        saved=json.loads(Path(written['files']['a_prime_mapping']).read_text(encoding='utf-8-sig'))
        row=next(r for r in saved['updates'] if r['cell']=='B9')
        self.assertEqual(row['source_identity']['dimensions']['asset_selection']['members'],['新增项目'])
        self.assertEqual(saved['collections']['B']['source_hash'],file_hash(bpath))
        self.assertNotIn('facts',saved)
        self.assertEqual(hashes,(file_hash(apath),file_hash(bpath)))

    def test_合成Word经提取识别更新到Word副本(self):
        from docx import Document
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries,get_column_letter
        word=self.folder/'合成附注.docx';doc=Document()
        doc.add_heading('六、财务报表重要项目的说明',level=1)
        table=doc.add_table(rows=2,cols=3)
        for row,values in zip(table.rows,[['项目','2025年末','2024年末'],['测试设备','10','9']]):
            for cell,value in zip(row.cells,values):cell.text=value
        doc.save(word);before=file_hash(word)
        config={**self.config,'word_path':str(word)}
        extracted=run_step('extract_a',config);self.assertEqual(extracted['status'],'complete',extracted)
        config.update(extracted['files'])
        book=load_workbook(config['a_path']);name=next(iter(book.defined_names.values()))
        sheet,area=next(name.destinations);l,t,r,b=range_boundaries(area);ws=book[sheet]
        label=f'{get_column_letter(l)}{t+1}'
        cols=[get_column_letter(l+1),get_column_letter(l+2)]
        ws[f'{cols[0]}{t}']='2026年末';ws[f'{cols[1]}{t}']='2025年末'
        ws[f'{cols[0]}{t+1}']=20;ws[f'{cols[1]}{t+1}']=10
        bpath=self.folder/'Word更新B.xlsx';book.save(bpath);book.close();config['b_path']=str(bpath)
        def answer(year):
            dims=[{'id':k,'value':v,'ranges':['*'],'evidence':[sheet+'!'+label]} for k,v in
                  [('entity','合成企业'),('report_scope','standalone'),('currency','CNY'),('unit_scale',1),
                   ('asset_selection',{'domain':'固定资产','mode':'members','members':['测试设备'],'completeness':'complete'})]]
            metrics=[]
            for col,role,date in zip(cols,['closing','opening'],[f'{year}-12-31',f'{year-1}-12-31']):
                metrics.append({'id':'固定资产原值','ranges':[f'{col}{t+1}'],'template_period_role':role})
                dims.append({'id':'period','value':{'kind':'instant','date':date},'ranges':[f'{col}{t+1}'],
                             'evidence':[f'{sheet}!{col}{t}'],'label_references':[f'{sheet}!{col}{t}']})
            return json.dumps({'tables':[{'id':'表一','sheet':sheet,'range':f'{get_column_letter(l)}1:{get_column_letter(r)}{b}','data_ranges':[f'{cols[0]}{t+1}:{cols[1]}{t+1}'],
                                           'metrics':metrics,'dimensions':dims}]},ensure_ascii=False)
        with patch.object(UnifiedSemanticEngine,'_http',side_effect=[answer(2025),answer(2026)]) as send:
            a=run_step('recognize_a',config);config.update(a['files'])
            b=run_step('recognize_b',config);config.update(b['files'])
            match=run_step('match',config);config.update(match['files'])
            written=run_step('write',config)
        self.assertEqual(written['status'],'complete',written)
        self.assertEqual(send.call_count,2)
        actual=Document(written['files']['word'])
        self.assertEqual(actual.tables[0].cell(0,1).text,'2026年末')
        self.assertIn('20',actual.tables[0].cell(1,1).text)
        saved=json.loads(Path(written['files']['a_prime_mapping']).read_text(encoding='utf-8-sig'))
        self.assertEqual(saved['a_link'],written['files']['word_link'])
        self.assertEqual(saved['output_hash'],file_hash(written['files']['a_prime']))
        self.assertNotIn('facts',saved)
        self.assertEqual(file_hash(word),before)

    def test_明确名单缩减沿用机械删行器且不修改原件(self):
        from openpyxl import load_workbook
        from openpyxl.workbook.defined_name import DefinedName
        book=load_workbook(self.source);book.defined_names.add(DefinedName('名单表',attr_text="'合成表'!$A$2:$C$8"))
        apath=self.folder/'缩减A.xlsx';book.save(apath)
        book.active.delete_rows(8)
        for r in range(3,8):
            for c in (2,3):book.active.cell(r,c,100+r*c)
        bpath=self.folder/'缩减B.xlsx';book.save(bpath);book.close()
        answer=json.loads(self.answer());t=answer['tables'][0];t['range']='A2:C7';t['data_ranges']=['B3:C7']
        for m in t['metrics']:m['ranges']=[m['ranges'][0].replace('8','7')]
        t['dimensions']=[d for d in t['dimensions'] if d['ranges']!=['B8:C8']]
        for d in t['dimensions']:
            if d['id']=='period':d['ranges']=[d['ranges'][0].replace('8','7')]
        purpose={'reason':'合成要求：名单采用B明确列出的五项，移除最后一项模板行','bindings':[],
                 'operations':[{'range_name':'名单表','axis':'row','action':'delete','index':6,'count':1,'template_index':1,
                                'source_evidence':['合成表!A7'],'coverage_reason':'合成要求明确采用已识别五项名单'}]}
        config={**self.config,'a_path':str(apath),'b_path':str(bpath),'update_purpose':purpose}
        before=file_hash(apath)
        with patch.object(UnifiedSemanticEngine,'_http',side_effect=[self.answer(),json.dumps(answer,ensure_ascii=False)]):
            a=run_step('recognize_a',config);config.update(a['files']);b=run_step('recognize_b',config);config.update(b['files'])
            matched=run_step('match',config);config.update(matched['files']);written=run_step('write',config)
        self.assertEqual(written['status'],'complete',written)
        saved=json.loads(Path(written['files']['a_prime_mapping']).read_text(encoding='utf-8-sig'))
        self.assertEqual(len(saved['updates']),10)
        self.assertFalse(any(r['cell']=='B8' for r in saved['updates']))
        self.assertNotIn('facts',saved)
        self.assertEqual(file_hash(apath),before)

    def test_没有A时只读取B目录且不传金标准(self):
        calls = []
        def fake(engine, prompt, payload):
            calls.append(prompt)
            self.assertNotIn('金标准.jsonl', payload['materials'])
            return '仅目录认识'
        with patch.object(UnifiedSemanticEngine, '_http', fake):
            result = run_step('recognize_b', self.config)
        self.assertEqual(len(calls), 1)
        self.assertIn('只保存目录认识', calls[0])
        self.assertEqual(result['status'], 'saved', result)

    def test_失败不自动重答旧入口也不能恢复分轮(self):
        engine = self.engine()
        with patch.object(engine, '_http', side_effect=TimeoutError('离线模拟超时')) as send:
            for _ in range(2):
                with self.assertRaises((TimeoutError, SemanticError)):
                    engine.request_text('测试', {'原文.json': {'空白': None}})
            self.assertEqual(send.call_count, 1)
            with self.assertRaisesRegex(SemanticError, '已停用'):
                engine._request('旧绑定', {'task': 'a_semantic_bind'}, lambda x: x)
            self.assertEqual(send.call_count, 1)

    def test_Codex材料一次送入不带结构模板(self):
        engine = self.engine()
        data = {'choices': [{'finish_reason': 'stop', 'message': {'content': '  原样自由答复\n'}}],
                'usage': {'prompt_tokens': 2, 'completion_tokens': 3, 'total_tokens': 5}}
        def fake(settings, work_dir, body, cancel):
            self.assertNotIn('response_format', body)
            self.assertEqual(json.loads(body['messages'][1]['content']), {'原文.json': {'text': '本组完整材料'}})
            self.assertNotIn('材料目录', body['messages'][1]['content'])
            return data, b'raw'
        with patch('附注更新.Codex语义调用.call_codex', side_effect=fake) as send:
            result = engine.request_text('直接回答', {'原文.json': {'text': '本组完整材料'}})
        self.assertEqual(send.call_count, 1)
        self.assertEqual(result, '  原样自由答复\n')
        self.assertEqual(engine.usage['requests'], 1)

    def test_HTTP一次传输不解析JSON且不重试(self):
        engine = self.engine('https://example.invalid/v1')
        captured = []
        def fake(request, simple, payload):
            body = json.loads(request.data)
            captured.append(body)
            self.assertNotIn('response_format', body)
            return {'choices': [{'finish_reason': 'stop', 'message': {'content': '{不完整JSON也应原样保存'}}]}, b'raw'
        with patch.object(engine, '_read_http_response', side_effect=fake):
            answer = engine.request_text('直接回答', {'原文.json': {'x': None}})
        self.assertEqual(answer, '{不完整JSON也应原样保存')
        self.assertEqual(len(captured), 1)
        with patch.object(engine, '_read_http_response', side_effect=TimeoutError('离线超时')) as send:
            with self.assertRaises(TimeoutError):
                engine.request_text('另一次明确任务', {'原文.json': {'x': None}})
            self.assertEqual(send.call_count, 1)

    def test_原生只启动一个进程并保存非JSON答案(self):
        from 附注更新.Codex语义调用 import call_codex
        commands = []
        class FakeProcess:
            returncode = 0
            def __init__(self, command, **kwargs):
                commands.append(command)
                Path(command[command.index('--output-last-message') + 1]).write_text('直接回答', encoding='utf-8')
            def communicate(self, **kwargs):
                return b'{"type":"turn.completed","usage":{"input_tokens":1,"output_tokens":1}}\n', b''
        with patch('附注更新.Codex语义调用.shutil.which', return_value=str(self.folder / 'codex.exe')), \
             patch('附注更新.Codex语义调用.subprocess.Popen', FakeProcess):
            data, _ = call_codex({'model': 'gpt-6-luna', 'model_reasoning_effort': 'high'}, self.folder,
                                {'messages': [{'content': '只读合成材料，直接回答'}]}, lambda: None)
        self.assertEqual(len(commands), 1)
        self.assertNotIn('--output-schema', commands[0])
        self.assertIn('web_search="disabled"', commands[0])
        self.assertIn('read-only', commands[0])
        self.assertIn('gpt-6-luna', commands[0])
        self.assertIn('model_reasoning_effort="high"', commands[0])
        self.assertIn(('--disable','shell_tool'), list(zip(commands[0],commands[0][1:])))
        self.assertIn(('--enable','skip_host_skill_discovery'), list(zip(commands[0],commands[0][1:])))
        self.assertIn('skills.bundled.enabled=false',commands[0])
        self.assertIn('skills.include_instructions=false',commands[0])
        self.assertNotIn('skills.max_context_tokens=1',commands[0])
        self.assertTrue(any(x.startswith('developer_instructions=') for x in commands[0]))
        self.assertEqual(data['choices'][0]['message']['content'], '直接回答')

    def test_已有答复机械写法及维度未知不抹掉已知身份(self):
        from 附注更新.A合集 import direct_collection, _standard
        from 附注更新.统一语义 import _dimension_value
        from 附注更新.表格 import read_workbook, materialize_candidate_blanks
        engine=self.engine();snapshot=read_workbook(self.source,context=['第4段：合成企业'])
        materialize_candidate_blanks(snapshot)
        table=json.loads(self.answer())['tables'][0]
        for row in table['metrics']+table['dimensions']:
            row['ranges']=['合成表!'+r if r!='*' else r for r in row['ranges']]
        for row in table['dimensions']:
            if row['id']=='period':row['value']=row['value']['date']
            if row['id']=='report_scope':row['value']=None
            if row['id']=='entity':row['evidence']=['@C0 第4段']
            if row['id']=='asset_selection':
                row['value']={'type':'members','members':row['value']['members']}
        table['dimensions'][-1]['value']={'type':'predicate','conditions':['原答复的单项条件']}
        table['unresolved']=[{'id':'report_scope','ranges':['合成表!B3:C8'],'value':None}]
        answer=json.dumps({k:table[k] for k in ('metrics','dimensions','unresolved')},ensure_ascii=False)+'\n喵'
        from 附注更新.统一语义 import definition_hash
        mapping={'source_hash':snapshot['sha256'],'source_context':['第4段：合成企业'],
            'gold_hash':engine.gold_hash,'金标准原文':self.gold.read_text(encoding='utf-8'),
            '识别答复':answer,'答复指纹':definition_hash(answer),'carrier':'A',
            'target_table':{k:table[k] for k in ('id','sheet','range','data_ranges')}}
        with patch.object(UnifiedSemanticEngine,'_http',side_effect=AssertionError('重读不能调用模型')):
            saved=direct_collection(mapping,snapshot,engine.standard)
        self.assertEqual(saved['识别答复'],answer)
        self.assertEqual(saved['coverage']['financial_cells'],12)
        self.assertEqual(saved['unresolved'],[])
        self.assertEqual(len(saved['slot_bindings']),12)
        self.assertEqual(sum(r['raw_value_state']=='blank' for r in saved['slot_bindings']),8)
        for row in saved['slot_bindings']:
            known=row['partial_semantics'];dims=known['dimensions']
            self.assertEqual(known['missing_dimensions'],['report_scope'])
            self.assertEqual(known['metric_id'],'固定资产原值')
            self.assertNotIn('completeness',dims['asset_selection'])
            self.assertEqual(dims['period']['kind'],'instant')
            self.assertTrue(dims['entity']=='合成企业')
            self.assertTrue(known['dimension_evidence']['entity'][0].endswith('#第4段'))
        definition=engine.standard['dimension']['asset_selection']
        v={'domain':definition['domain'],'mode':'all'}
        self.assertEqual(_dimension_value(v,definition,matching=True),
                         _dimension_value({**v,'completeness':'complete'},definition,matching=True))
        with self.assertRaises(ValueError):_dimension_value({**v,'completeness':'unknown'},definition,matching=True)
        self.assertEqual(self.hashes,(file_hash(self.source),file_hash(self.gold)))

    def test_未决和维度冲突不撤销其他已知内容(self):
        from 附注更新.A合集 import direct_collection
        from 附注更新.表格 import read_workbook, materialize_candidate_blanks
        from 附注更新.统一语义 import definition_hash
        from 附注更新.模板更新编排 import direct_update_plan
        engine=self.engine();snapshot=read_workbook(self.source,context=[])
        materialize_candidate_blanks(snapshot)
        table=json.loads(self.answer())['tables'][0]
        table['dimensions'].append({'id':'period','value':{'kind':'instant','date':'2024-12-31'},
                                    'ranges':['B4'],'evidence':['合成表!B2']})
        table['unresolved']=[{'metric':'固定资产原值','ranges':['B3'],'reason':'原答复同时保留的未决说明'}]
        target={k:table[k] for k in ('id','sheet','range','data_ranges')}
        target['data_ranges']=['B3:C7','B8']
        target['non_business']=[{'range':'C8','category':'annotation','reason':'测试输入明确的说明位置','evidence':['合成表!A8']}]
        answer=json.dumps({k:table[k] for k in ('metrics','dimensions','unresolved')},ensure_ascii=False)
        mapping={'recognition_method':'direct','source_hash':snapshot['sha256'],'source_context':[],
                 'gold_hash':engine.gold_hash,'金标准原文':'\n'.join(line for line in self.gold.read_text(encoding='utf-8').splitlines()
                     if json.loads(line)['record_type'] in {'metric','dimension'}),
                 '识别答复':answer,'答复指纹':definition_hash(answer),'carrier':'A','target_table':target}
        with patch.object(UnifiedSemanticEngine,'_http',side_effect=AssertionError('离线不能调用模型')):
            saved=direct_collection(mapping,snapshot,engine.standard)
        facts={r['source_reference']['location']:r for r in saved['facts']}
        slots={r['cell']:r['partial_semantics'] for r in saved['slot_bindings']}
        self.assertEqual(facts['合成表!B3']['metric_id'],'固定资产原值')
        self.assertEqual(slots['B4']['metric_id'],'固定资产原值')
        self.assertEqual(slots['B4']['dimensions']['entity'],'合成企业')
        self.assertNotIn('period',slots['B4']['dimensions'])
        self.assertEqual({r['cell'] for r in saved['unresolved']},{'B3','B4'})
        self.assertEqual([r['cell'] for r in saved['excluded']],['C8'])
        self.assertEqual(saved['coverage']['financial_cells'],12)
        path=self.folder/'未决并存合集.json';path.write_text(json.dumps(saved,ensure_ascii=False),encoding='utf-8')
        restored=load_fact_mapping(path,{})
        self.assertEqual(restored['识别答复'],answer)
        plan=direct_update_plan(restored,restored,{'reason':'合成测试指定来源也不能绕过未决',
                                 'bindings':[{'target':'合成表!C3','source':'合成表!B3'}]})
        matched={r['sheet']+'!'+r['cell'] for r in plan['identity_matches']}
        self.assertFalse(matched & {'合成表!B3','合成表!B4','合成表!C3'})
        self.assertIn('合成表!C4',matched)
        self.assertEqual(self.hashes,(file_hash(self.source),file_hash(self.gold)))

    def test_机械表示接通但不猜条件和期间角色(self):
        from 附注更新.A合集 import direct_collection
        from 附注更新.表格 import read_workbook, materialize_candidate_blanks
        from 附注更新.统一语义 import definition_hash
        engine=self.engine();snapshot=read_workbook(self.source,context=[]);materialize_candidate_blanks(snapshot)
        table=json.loads(self.answer())['tables'][0]
        table['metrics'][0]['template_period_role']={'B3:B8':'本期余额'}
        table['metrics'][1]['template_period_role']={'C3:C6':'上期余额','C7':'本期期初','C8':'上期余额'}
        table['dimensions']=[r for r in table['dimensions'] if r['id']!='period']
        table['dimensions'].append({'id':'period','value':'本期：2025-12-31；上期：2024-12-31',
            'ranges':['B3:C8','C8:B8'],'evidence':['合成表!B2:C2']})
        selection=[r for r in table['dimensions'] if r['id']=='asset_selection']
        selection[0]['value']={'type':'members','members':['测试项目3'],'details':'原有附加说明'}
        selection[0]['evidence']=['合成表!A3；原答复附带说明']
        selection[1]['value']={'type':'predicate','conditions':[{'field':'项目','value':'甲'}]}
        selection[2]['value']={'type':'predicate','conditions':['甲','乙']}
        selection[3]['value']={'type':'partial','known':'仅已知原有范围文字','unknown':'具体条件未定'}
        answer=json.dumps({'tables':[table]},ensure_ascii=False)
        mapping={'recognition_method':'direct','source_hash':snapshot['sha256'],'source_context':[],
                 'gold_hash':engine.gold_hash,'金标准原文':'\n'.join(line for line in self.gold.read_text(encoding='utf-8').splitlines()
                     if json.loads(line)['record_type'] in {'metric','dimension'}),
                 '识别答复':answer,'答复指纹':definition_hash(answer),'carrier':'A'}
        saved=direct_collection(mapping,snapshot,engine.standard)
        all_rows={r['source_reference']['location']:r for r in saved['facts']}
        all_rows.update({r['sheet']+'!'+r['cell']:r['partial_semantics'] for r in saved['slot_bindings']})
        self.assertEqual(all_rows['合成表!B3']['dimensions']['asset_selection']['members'],['测试项目3'])
        self.assertEqual(all_rows['合成表!B4']['dimensions']['asset_selection']['predicate'],'项目=甲')
        self.assertNotIn('asset_selection',all_rows['合成表!B5']['dimensions'])
        self.assertEqual(all_rows['合成表!B6']['dimensions']['asset_selection']['mode'],'unknown')
        self.assertNotIn('period',all_rows['合成表!C7']['dimensions'])
        self.assertEqual(all_rows['合成表!C8']['dimensions']['period']['date'],'2024-12-31')
        self.assertEqual(saved['识别答复'],answer)
        self.assertTrue(any('C8:B8' in r.get('ranges',[]) for r in saved['outline']['subtables'][0]['unresolved']))
        path=self.folder/'机械表示合集.json';path.write_text(json.dumps(saved,ensure_ascii=False),encoding='utf-8')
        self.assertEqual(load_fact_mapping(path,{})['slot_bindings'],saved['slot_bindings'])
        from 附注更新.模板更新编排 import direct_update_plan
        plan=direct_update_plan(saved,saved)
        self.assertIn('B3',{r['cell'] for r in plan['identity_matches']})
        self.assertNotIn('C7',{r['cell'] for r in plan['identity_matches']})

    def test_显式同时满足及坐标日期只转换表示(self):
        from 附注更新.A合集 import _direct_table, _period_rows
        engine=self.engine();table=json.loads(self.answer())['tables'][0]
        table['metrics'][0]['template_period_role']=['本期']
        self.assertEqual(_direct_table({**table,'dimensions':[]},{'source_context':[]},engine.standard)
                         ['metrics'][0]['template_period_role'],'本期')
        conditions=[{'type':'members','members':['事项甲']},{'type':'condition','value':'条件乙'}]
        row={'id':'asset_selection','value':{'type':'predicate','all':conditions},
             'ranges':['B3'],'evidence':['合成表!A3']}
        saved=_direct_table({**table,'dimensions':[row]},{'source_context':[]},engine.standard)['dimensions'][0]
        self.assertEqual(json.loads(saved['value']['predicate']),{'all':conditions})
        period={'id':'period','value':'B3、B4：2025-12-31；C3、C4：2024-01-01',
                'ranges':['B3:C3'],'evidence':['合成表!B2:C2']}
        expanded,problems=_period_rows(period,table)
        self.assertFalse(problems)
        self.assertEqual({a:r['value'] for r in expanded for a in r['ranges']},
                         {'B3':'2025-12-31','C3':'2024-01-01'})
        self.assertEqual(row['value'],{'type':'predicate','all':conditions})

    def test_固定目标不靠答复圈范围且漏空白保留(self):
        table={'id':'本表','sheet':'合成表','range':'A1:C8','data_ranges':['B3:C8']}
        config={**self.config,'a_target_table':table,'a_allowed_metrics':['固定资产原值']}
        answer=json.loads(self.answer())['tables'][0]
        answer={k:answer[k] for k in ('metrics','dimensions')}
        # 省略一个空白格的身份，程序仍应保留该位置。
        answer['metrics'][1]['ranges']=['C3:C7']
        sent=[]
        def fake(engine,prompt,payload):
            sent.append(payload)
            self.assertEqual(payload['materials']['原文.json']['固定目标'],table)
            actual=payload['materials']['金标准.jsonl']
            lines=self.gold.read_text(encoding='utf-8').splitlines()
            metric=next(json.loads(l) for l in lines if json.loads(l).get('id')=='固定资产原值')
            ids={metric['id']}|{r['id'] for r in metric['dimension_refs']}
            expected='\n'.join(l for l in lines if json.loads(l).get('id') in ids)
            self.assertEqual(actual,expected)
            return json.dumps(answer,ensure_ascii=False)
        with patch.object(UnifiedSemanticEngine,'_http',fake):
            run=run_step('recognize_a',config)
        self.assertEqual(len(sent),1)
        self.assertEqual(run['status'],'complete')
        saved=load_fact_mapping(run['files']['a_mapping'],{})
        self.assertEqual(saved['coverage']['financial_cells'],12)
        self.assertEqual(saved['outline']['subtables'][0]['data_ranges'],['B3:C8'])
        self.assertEqual({r['id'] for r in saved['facts']}|{r['sheet']+'!'+r['cell'] for r in saved['unresolved']},
                         {f'合成表!{c}{r}' for c in 'BC' for r in range(3,9)})
        self.assertEqual(sum(r['raw_value_state']=='blank' for r in saved['facts']+saved['unresolved']),8)
        self.assertEqual(saved['unresolved'][0]['cell'],'C8')
        self.assertEqual(saved['unresolved'][0]['original']['dimensions']['entity']['value'],'合成企业')
        self.assertEqual(self.hashes,(file_hash(self.source),file_hash(self.gold)))

    def test_固定目标完整答复直接进入合集及匹配(self):
        table={'id':'本表','sheet':'合成表','range':'A1:C8','data_ranges':['B3:C8']}
        answer=json.loads(self.answer())['tables'][0]
        answer={k:answer[k] for k in ('metrics','dimensions')}
        with patch.object(UnifiedSemanticEngine,'_http',return_value=json.dumps(answer,ensure_ascii=False)) as send:
            run=run_step('recognize_a',{**self.config,'a_target_table':table,'a_allowed_metrics':['固定资产原值']})
        self.assertEqual(send.call_count,1)
        self.assertEqual(run['status'],'complete',run)
        saved=load_fact_mapping(run['files']['a_mapping'],{})
        self.assertEqual(len(saved['facts']),12)
        self.assertEqual(sum(r['raw_value_state']=='blank' for r in saved['facts']),8)
        from 附注更新.A合集 import _standard
        self.assertEqual(compare_semantic_slots(saved,saved,_standard(saved))['matched_targets'],12)

    def test_固定目标答复不可用仍保存十二格不重答(self):
        table={'id':'本表','sheet':'合成表','range':'A1:C8','data_ranges':['B3:C8']}
        for answer in ('[]','{未结束'):
            with self.subTest(answer=answer), patch.object(UnifiedSemanticEngine,'_http',return_value=answer) as send:
                run=run_step('recognize_a',{**self.config,'a_target_table':table,
                    'output_dir':str(self.folder/uuid.uuid4().hex)})
                self.assertEqual(send.call_count,1)
                self.assertEqual(run['status'],'complete')
                saved=load_fact_mapping(run['files']['a_mapping'],{})
                self.assertEqual(saved['识别答复'],answer)
                self.assertEqual(saved['coverage']['financial_cells'],12)
                missing=[r for r in saved['unresolved'] if r.get('cell')]
                self.assertEqual(len(missing),12)
                self.assertEqual(sum(r['raw_value_state']=='blank' for r in missing),8)

    def test_分组合集经普通入口汇总保留缺项且不阻断其他格(self):
        from copy import deepcopy
        from 附注更新.A合集 import save_direct_answer, direct_digest, _records
        from 附注更新.A语义阅读 import direct_reference
        from 附注更新.模板更新编排 import direct_update_plan
        from 附注更新.表格 import read_workbook, materialize_candidate_blanks
        engine=self.engine();snapshot=read_workbook(self.source,context=[]);materialize_candidate_blanks(snapshot)
        index=[];parts=[]
        for start,end in ((3,5),(6,8)):
            table=json.loads(self.answer())['tables'][0]
            table.update(id=str(start),data_ranges=[f'B{start}:C{end}'])
            table['metrics']=[{**r,'ranges':[f'{col}{start}:{col}{end}']} for col,r in zip(('B','C'),table['metrics'])]
            table['dimensions']=table['dimensions'][:4]+[
                {**r,'ranges':[f'{col}{start}:{col}{end-1 if start==6 and col=="B" else end}']}
                for col,r in zip(('B','C'),table['dimensions'][4:6])]+table['dimensions'][6+start-3:6+end-2]
            text=json.dumps(table,ensure_ascii=False)
            mapping={'recognition_method':'direct','source_hash':snapshot['sha256'],'source_context':[],
                'gold_hash':engine.gold_hash,'金标准原文':'\n'.join(line for line in self.gold.read_text(encoding='utf-8').splitlines()
                    if json.loads(line)['record_type'] in {'metric','dimension'}),
                'carrier':'A','target_table':{k:v for k,v in table.items() if k not in ('metrics','dimensions')}}
            destination=self.folder/f'{start}合集.json'
            with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('不能调用模型')):
                part=save_direct_answer(mapping,text,destination,snapshot=snapshot,standard=engine.standard)
            parts.append(part);index.append({'最终合集':str(destination),'sha256':file_hash(destination)})
        path=self.folder/'合集索引.json'
        path.write_text(json.dumps(index,ensure_ascii=False),encoding='utf-8')
        combined=load_fact_mapping(path,{})
        self.assertEqual(dict(_records(combined)),{loc:r for p in parts for loc,r in _records(p)})
        self.assertEqual((len(combined['facts']),len(combined['slot_bindings'])),(11,1))
        self.assertEqual(sum(r['raw_value_state']=='blank' for r in combined['facts']),7)
        reference=direct_reference(combined)
        self.assertEqual(reference['slot_bindings'][0]['cell'],'B8')
        from 附注更新.统一语义 import identity_completeness
        with patch('附注更新.统一语义.identity_completeness',wraps=identity_completeness) as completeness:
            plan=direct_update_plan(combined,combined)
        self.assertLessEqual(completeness.call_count,24,'同一B记录完整性只计算一次，不对每个A格重新计算')
        self.assertEqual(len(plan['updates']),4)
        self.assertTrue(any(r.get('cell')=='B8' for r in plan['issues']))
        standalone=self.folder/'整份合集.json'
        standalone.write_text(json.dumps(combined,ensure_ascii=False),encoding='utf-8')
        original=Path.read_bytes
        def only_saved(p):
            self.assertEqual(p.resolve(),standalone.resolve(),'正式合集回读不再依赖分组、原表或金标准')
            return original(p)
        with patch.object(Path,'read_bytes',only_saved):
            self.assertEqual(load_fact_mapping(standalone,{})['facts'],combined['facts'])
        # 同一工作表内矩形范围相同，按各子表已登记财务格接入第二表。
        original_hash=file_hash(standalone)
        correction=json.dumps({'dimensions':[{'id':'period','value':'2025-12-31','ranges':['合成表!B7:B8'],'evidence':['@C0']},
            {'id':'template_period_role','value':'closing','ranges':['合成表!B7:B8']}]},ensure_ascii=False)
        sent=[]
        def correct(engine,prompt,payload):
            sent.append(payload)
            return correction
        with patch.object(UnifiedSemanticEngine,'_http',correct):
            replay=run_step('recognize_a',{**self.config,'a_mapping':str(standalone),'a_context_text':'本次期间关系原文',
                'a_field_scope':{'合成表!B7:B8':['period','template_period_role']}})
        self.assertEqual(replay['status'],'complete',replay)
        self.assertEqual(len(sent),1)
        corrected=load_fact_mapping(replay['files']['a_mapping'],{})
        self.assertEqual([t['id'] for t in corrected['outline']['subtables']],['3','6'])
        for key in ('id','sheet','range','data_ranges','non_business'):
            self.assertEqual(corrected['outline']['subtables'][0].get(key),combined['outline']['subtables'][0].get(key))
        self.assertEqual([{k:v for k,v in m.items() if k!='original'} for m in corrected['outline']['subtables'][0]['metrics']],
                         [{k:v for k,v in m.items() if k!='original'} for m in combined['outline']['subtables'][0]['metrics']])
        self.assertEqual([t['data_ranges'] for t in corrected['outline']['subtables']],
                         [t['data_ranges'] for t in combined['outline']['subtables']])
        before_rows=dict(_records(combined));after_rows=dict(_records(corrected))
        self.assertEqual(set(before_rows),set(after_rows))
        self.assertEqual(corrected['机械记录']['sheets'],combined['机械记录']['sheets'])
        for loc,row in before_rows.items():
            if loc not in {'合成表!B7','合成表!B8'}:
                self.assertEqual(row['dimensions'],after_rows[loc]['dimensions'])
                self.assertEqual(row['dimension_evidence'],after_rows[loc]['dimension_evidence'])
        self.assertEqual((len(corrected['facts']),len(corrected['slot_bindings'])),(12,0))
        self.assertEqual(sum(r['raw_value_state']=='blank' for r in corrected['facts']),8)
        self.assertEqual(after_rows['合成表!B8']['dimensions']['period'],
                         {'kind':'instant','date':'2025-12-31'})
        from 附注更新.统一证据 import context_catalog
        self.assertEqual(after_rows['合成表!B8']['dimension_evidence']['period'],list(context_catalog(['本次期间关系原文'])))
        self.assertEqual(file_hash(standalone),original_hash)
        for bad in (index+index[:1],):
            path.write_text(json.dumps(bad,ensure_ascii=False),encoding='utf-8')
            with self.assertRaisesRegex(ValueError,'位置重叠'):load_fact_mapping(path,{})
        other=deepcopy(parts[1]);other['source_hash']='0'*64;other['封存校验']=direct_digest(other)
        different=self.folder/'异源合集.json';different.write_text(json.dumps(other,ensure_ascii=False),encoding='utf-8')
        path.write_text(json.dumps([index[0],{'最终合集':str(different),'sha256':file_hash(different)}]),encoding='utf-8')
        with self.assertRaisesRegex(ValueError,'不是同一原件'):load_fact_mapping(path,{})

    def test_B范围认识重复行无损且普通入口只调用一次(self):
        from 附注更新.A语义阅读 import scope_material, source_material
        from 附注更新.表格 import read_workbook, materialize_candidate_blanks
        snapshot=read_workbook(self.source,context=[]);materialize_candidate_blanks(snapshot)
        for cell in snapshot['sheets'][0]['cells'].values():
            if cell['row']>=5:cell['value']=None;cell.pop('formula',None)
        snapshot['semantic_formula_sources']=5
        compact=scope_material(snapshot,[])
        original,_=source_material({**snapshot,'semantic_formula_sources':0},[])
        for actual,expected in zip(compact['sheets'],original['sheets']):
            expanded=[]
            for line in actual['rows'].splitlines():
                span,body=line.split(': ',1);ends=span.split('-')
                expanded.extend(f'{n}: {body}' for n in range(int(ends[0]),int(ends[-1])+1))
            self.assertEqual('\n'.join(expanded),expected['rows'])
            self.assertEqual(actual['merges'],expected['merges'])
        def answer(engine,prompt,payload):
            self.assertNotIn('金标准.jsonl',payload['materials'])
            self.assertNotIn('公式引用依据',payload['materials']['原文.json'])
            self.assertIn('不逐格识别',prompt)
            return '合成表 A2:C8：固定资产原值，本期与上期；包括待填空白。'
        with patch.object(UnifiedSemanticEngine,'_http',answer):
            run=run_step('recognize_b',self.config)
        self.assertEqual(run['status'],'saved')
        saved=json.loads(Path(run['files']['b_mapping']).read_text(encoding='utf-8-sig'))
        self.assertEqual(saved['recognition_scope']['mode'],'b_outline')
        self.assertNotIn('facts',saved)

    def test_B实际公式逐格给出且A需要按范围保留(self):
        from 附注更新.A语义阅读 import source_material, compact_formula_material, direct_reference, addresses
        from 附注更新.表格 import read_workbook, materialize_candidate_blanks
        book=Workbook();sheet=book.active;sheet.title='公式'
        for row in range(1,7):
            sheet.cell(row,1,row);sheet.cell(row,2,f'=A{row}+$A$1')
        for row in (1,2,4):sheet.cell(row,3,'=未定义名称')
        book.create_sheet('另一公式表')['C1']='=未定义名称'
        path=self.folder/'公式.xlsx';book.save(path);book.close()
        snapshot=read_workbook(path,context=[]);materialize_candidate_blanks(snapshot)
        snapshot.update(semantic_formula_sources=5,semantic_identity_version=2)
        raw,_=source_material(snapshot,[])
        compact=compact_formula_material(snapshot,raw)
        restored=compact['公式及来源']['公式原文']
        self.assertEqual(restored,raw['公式原文'])
        self.assertNotIn('公式分组',compact['公式及来源'])
        self.assertEqual(restored['公式!B1'],'=A1+$A$1')
        self.assertEqual(restored['公式!B6'],'=A6+$A$1')
        unresolved=[]
        for row in compact['公式及来源']['未展开']:
            for location in row['位置']:
                sheet,_,span=location.rpartition('!')
                unresolved.extend({**row,'位置':sheet+'!'+cell} for cell in addresses(span))
        self.assertCountEqual(unresolved,raw['公式引用依据']['未展开'])
        self.assertEqual(next(r['位置'] for r in compact['公式及来源']['未展开']
                              if r['位置'][0].startswith('公式!')),['公式!C1:C2','公式!C4'])
        with patch.object(UnifiedSemanticEngine,'_http',return_value=self.answer()):
            run=run_step('recognize_a',self.config)
        saved=load_fact_mapping(run['files']['a_mapping'],{})
        needs=direct_reference(saved,compact=True,ranges=[{'sheet':'合成表','range':'B3:C3'}])['需要']
        self.assertEqual({r['location'] for r in needs},{'合成表!B3','合成表!C3'})
        originals={r['source_reference']['location']:r for r in saved['facts']}
        for row in needs:self.assertEqual(row['dimensions'],originals[row['location']]['dimensions'])

    def test_B普通入口字段纠正保留多表并接新增候选(self):
        from copy import deepcopy
        from openpyxl import load_workbook
        from 附注更新.A合集 import save_direct_answer, _records
        from 附注更新.表格 import read_workbook, materialize_candidate_blanks
        bpath=self.folder/'两表B.xlsx'
        book=load_workbook(self.source);book.copy_worksheet(book.active).title='另一表'
        book.save(bpath);book.close()
        snapshot=read_workbook(bpath,context=[]);materialize_candidate_blanks(snapshot)
        engine=self.engine()
        base={'recognition_method':'direct','source_path':str(bpath),'source_hash':snapshot['sha256'],
              'source_context':[],'gold_path':str(self.gold),'gold_hash':engine.gold_hash,
              '金标准原文':'\n'.join(line for line in self.gold.read_text(encoding='utf-8').splitlines()
                  if json.loads(line)['record_type'] in {'metric','dimension'}),'carrier':'A'}
        a=save_direct_answer(base,self.answer(),self.folder/'所需A.json',snapshot=snapshot,standard=engine.standard)
        tables=[]
        for name in ('合成表','另一表'):
            table=json.loads(self.answer())['tables'][0]
            table.update(id=name,sheet=name,data_ranges=['B3','B5'])
            table['metrics']=[{**table['metrics'][0],'ranges':['B3','B5']}]
            table['dimensions']=table['dimensions'][:4]+[
                {**table['dimensions'][4],'ranges':['B3','B5']},
                table['dimensions'][6],table['dimensions'][8]]
            tables.append(table)
        old=save_direct_answer({**base,'carrier':'B'},json.dumps({'tables':tables},ensure_ascii=False),
                               self.folder/'原两表B.json',snapshot=snapshot,standard=engine.standard)
        before=deepcopy(dict(_records(old)));old_hash=file_hash(old['mapping_path'])
        source_other={**old,'b_source_ranges':[{'sheet':'另一表','range':'A1:C5'}]}
        bare=save_direct_answer(source_other,'B3 的 asset_selection 为“裸坐标明确范围”。',
            self.folder/'范围表名B.json',field_scope={'合成表!B3':['asset_selection']})
        self.assertEqual(dict(_records(bare))['合成表!B3']['dimensions']['asset_selection']['predicate'],'裸坐标明确范围')
        for key in ('metric_id','dimensions','raw_value','raw_value_state'):
            self.assertEqual(dict(_records(bare))['另一表!B3'].get(key),dict(_records(old))['另一表!B3'].get(key))
        blank_source={**old,'b_source_ranges':[{'sheet':'合成表','range':'A1:C8'}]}
        unknown=save_direct_answer(blank_source,'合成表!B7：本年期末余额（2025-12-31）；C7：上年期末余额（2024-12-31）。',
            self.folder/'未给编号空白B.json',field_scope={'合成表!B7:C7':['metric_id','period'],'合成表!B8':['metric_id']})
        table=next(t for t in unknown['outline']['subtables'] if t['sheet']=='合成表')
        self.assertEqual(set(table['data_ranges']),{'B3','B5','B7','C7'})
        self.assertNotIn('B8',table['data_ranges'])
        self.assertFalse({'合成表!B7','合成表!C7','合成表!B8'}&dict(_records(unknown)).keys())
        periods={cell:r['value'] for r in table['dimensions'] if r['id']=='period' for cell in r['ranges']}
        self.assertEqual(periods['B7'],{'kind':'instant','date':'2025-12-31'})
        self.assertEqual(periods['C7'],{'kind':'instant','date':'2024-12-31'})
        fields=['metric_id','entity','report_scope','currency','unit_scale','period','asset_selection','template_period_role']
        scope={'合成表!B3':['asset_selection'],'合成表!B4:B6':fields,'合成表!C4':fields}
        config={**self.config,'a_path':str(bpath),'b_path':str(bpath),'a_mapping':a['mapping_path'],
                'b_mapping':old['mapping_path'],'b_field_scope':scope,
                'b_source_ranges':[{'sheet':'合成表','range':'A1:C5'}]}
        reply='合成表!B3 的 asset_selection 为“本次明确更正范围”。\n'
        reply+='合成表!B4、B6：固定资产原值；entity=合成企业；report_scope=standalone；currency=CNY；unit_scale=1；period=2025-12-31；asset_selection=新增明确项目；template_period_role=closing。'
        with patch.object(UnifiedSemanticEngine,'request_text',return_value=reply) as send:
            run=run_step('recognize_b',config)
        send.assert_called_once()
        self.assertEqual(Path(run['files']['recognition_answer']).read_text(encoding='utf-8-sig'),reply)
        saved=load_fact_mapping(run['files']['b_mapping'],{})
        rows=dict(_records(saved))
        table=next(t for t in saved['outline']['subtables'] if t['sheet']=='合成表')
        self.assertEqual(set(table['data_ranges']),{'B3','B4','B5'})
        self.assertNotIn('C4',table['data_ranges'])
        self.assertEqual(set(rows),set(before)|{'合成表!B4'})
        self.assertEqual(rows['合成表!B3']['dimensions']['asset_selection']['predicate'],'本次明确更正范围')
        for loc,row in before.items():
            self.assertEqual(rows[loc]['raw_value_state'],row['raw_value_state'])
            self.assertEqual(rows[loc].get('raw_value'),row.get('raw_value'))
            self.assertEqual(rows[loc]['metric_id'],row['metric_id'])
            for key,value in row['dimensions'].items():
                if (loc,key)!=('合成表!B3','asset_selection'):
                    self.assertEqual(rows[loc]['dimensions'][key],value,(loc,key))
        self.assertEqual(rows['合成表!B4']['dimensions']['asset_selection']['predicate'],'新增明确项目')
        self.assertEqual([row['template_period_role'] for table in saved['outline']['subtables']
                          if table['sheet']=='合成表' for row in table['metrics'] if row['ranges']==['B4']],['closing'])
        self.assertEqual({loc for loc,row in rows.items() if row['raw_value_state']=='blank'},
                         {'合成表!B5','另一表!B5'})
        self.assertEqual({t['sheet'] for t in saved['outline']['subtables']},{'合成表','另一表'})
        self.assertEqual(saved['答复接入来源']['本次原答复'],reply)
        self.assertEqual(file_hash(old['mapping_path']),old_hash)
        rejected=save_direct_answer(old,old['识别答复'],self.folder/'空outlineB.json',
                                    not_adopted_reason='合成旧答复坐标错误，尚无采用的位置')
        self.assertFalse(rejected['outline']['subtables'])
        config.update(b_mapping=rejected['mapping_path'],b_field_scope={'合成表!B4':fields,'合成表!B6':fields})
        with patch.object(UnifiedSemanticEngine,'request_text',return_value=reply.splitlines()[1]) as send:
            run=run_step('recognize_b',config)
        send.assert_called_once()
        self.assertEqual(Path(run['files']['recognition_answer']).read_text(encoding='utf-8-sig'),reply.splitlines()[1])
        restored=load_fact_mapping(run['files']['b_mapping'],{})
        self.assertEqual(set(dict(_records(restored))),{'合成表!B4'})
        self.assertEqual(dict(_records(restored))['合成表!B4']['dimensions']['asset_selection']['predicate'],'新增明确项目')
        self.assertEqual(self.hashes,(file_hash(self.source),file_hash(self.gold)))
        self.assertEqual(file_hash(bpath),snapshot['sha256'])

    def test_已有字段分组坐标围栏及空引用直接保存(self):
        from 附注更新.A合集 import direct_collection, direct_answer
        from 附注更新.表格 import read_workbook, materialize_candidate_blanks
        from 附注更新.统一语义 import definition_hash, semantic_key
        engine=self.engine();snapshot=read_workbook(self.source,context=[]);materialize_candidate_blanks(snapshot)
        table=json.loads(self.answer())['tables'][0]
        target={k:v for k,v in table.items() if k not in ('metrics','dimensions')}
        answer={'standard':'固定资产原值','dimensions':{'entity':'合成企业','report_scope':'standalone',
            'currency':'CNY','unit_scale':1,'period':{'B3:B8':'2025-12-31','C3:C8':'2024-12-31'}},
            'asset_selection':[{'range':'B3:B8,C3:C8','value':{'type':'all','condition':'仅原答复已经写明的范围条件'}}],
            'template_period_role':{'B3:B8':'closing','C3:C8':'opening'}}
        text='```json\n'+json.dumps(answer,ensure_ascii=False)+'\n```'
        mapping={'recognition_method':'direct','source_hash':snapshot['sha256'],'source_context':[],
                 'gold_hash':engine.gold_hash,'金标准原文':'\n'.join(line for line in self.gold.read_text(encoding='utf-8').splitlines()
                     if json.loads(line)['record_type'] in {'metric','dimension'}),
                 '识别答复':text,'答复指纹':definition_hash(text),'carrier':'A','target_table':target}
        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('不调用模型')):
            saved=direct_collection(mapping,snapshot,engine.standard)
        self.assertEqual(len(saved['facts']),12)
        self.assertEqual(sum(r['raw_value_state']=='blank' for r in saved['facts']),8)
        for row in saved['facts']:
            self.assertTrue(all(refs==[] for refs in row['dimension_evidence'].values()))
            self.assertEqual(row['dimensions']['asset_selection']['condition'],'仅原答复已经写明的范围条件')
        altered=json.loads(json.dumps(saved['facts'][0]));altered['dimensions']['asset_selection']['condition']='另一条件'
        self.assertNotEqual(semantic_key(saved['facts'][0],engine.standard),semantic_key(altered,engine.standard))
        path=self.folder/'无引用已有答复合集.json';path.write_text(json.dumps(saved,ensure_ascii=False),encoding='utf-8')
        restored=load_fact_mapping(path,{})
        self.assertEqual(restored['识别答复'],text)
        self.assertEqual(restored['facts'],saved['facts'])
        shapes=[{'B3':{'currency':'CNY'}},{'currency':{'B3':'CNY'}},
                {'corrections':[{'ranges':['B3'],'field':'currency','value':'CNY'}]}]
        for shape in shapes:
            row=direct_answer(shape,engine.standard)['dimensions'][0]
            self.assertEqual((row['id'],row['ranges'],row['value']),('currency',['B3'],'CNY'))
        inverse=direct_answer({'period':{'2024-12-31':['B3、B4']}},engine.standard)['dimensions'][0]
        self.assertEqual(inverse['ranges'],['B3','B4'])
        for shape in ({'currency':{'coordinates':['B3'],'value':'CNY'}},
                      {'currency':[{'范围':'B3','取值':'CNY'}]},
                      {'field':'currency','corrections':[{'range':'B3','value':'CNY'}]}):
            row=direct_answer(shape,engine.standard)['dimensions'][0]
            self.assertEqual((row['ranges'],row['value']),(['B3'],'CNY'))
        grouped={'corrections':[{'field':'currency','ranges':['B3:C5'],
                                'values_by_row':{'3/5':'CNY','4':'USD'}}]}
        grouped=direct_answer(grouped,engine.standard)['dimensions']
        self.assertEqual(set(grouped[0]['ranges']),{'B3','C3','B5','C5'})
        self.assertEqual(set(grouped[1]['ranges']),{'B4','C4'})
        blank=direct_answer({'asset_selection':{'B3':'保留空白'}},engine.standard)
        self.assertFalse(blank['dimensions']);self.assertEqual(blank['unresolved'][0]['ranges'],['B3'])
        unknown={'id':'asset_selection','value':{'type':'partial','known_scope':'已有范围','unknown':'未给成员'},'ranges':['B3']}
        from 附注更新.A合集 import _direct_table
        value=_direct_table({**target,'metrics':[], 'dimensions':[unknown]},mapping,engine.standard)['dimensions'][0]['value']
        self.assertEqual(value['known_scope'],'已有范围');self.assertEqual(value['mode'],'unknown')
        literal={**engine.standard,'dimension':{**engine.standard['dimension'],
            '测试方向':{'id':'测试方向','value_type':'text','allowed_values':['增加','减少']}}}
        enum={'id':'测试方向','ranges':['B3'],'value':'减少；原答复附带的说明'}
        row=_direct_table({**target,'metrics':[], 'dimensions':[enum]},mapping,literal)['dimensions'][0]
        self.assertEqual(row['value'],'减少');self.assertEqual(row['original']['value'],enum['value'])

    def test_文字和表格经正式A入口保存十二格八空白(self):
        text='B3:C8 标准编号：固定资产原值；entity：合成企业；report_scope：standalone；currency：CNY；unit_scale：1\n'
        text+='| 坐标 | period |\n|---|---|\n| B3:B8 | 2025年12月31日 |\n| C3:C8 | 2024-12-31 |\n'
        text+='\n'.join(f'B{r}:C{r} asset_selection：测试项目{r}' for r in range(3,9))
        def fake(engine,prompt,payload):
            self.assertIn('不要求JSON或固定模板',prompt)
            self.assertNotIn('只返回metrics',prompt)
            self.assertNotIn('空白只是不填值',prompt)
            return text
        with patch.object(UnifiedSemanticEngine,'_http',fake):
            result=run_step('recognize_a',self.config)
        self.assertEqual(result['status'],'complete',result)
        saved=load_fact_mapping(result['files']['a_mapping'],{})
        self.assertEqual(saved['识别答复'],text)
        self.assertEqual(len(saved['facts']),12)
        self.assertEqual(sum(r['raw_value_state']=='blank' for r in saved['facts']),8)
        self.assertEqual(self.hashes,(file_hash(self.source),file_hash(self.gold)))

    def test_新外层字段及未知字面接入共用出口(self):
        from 附注更新.A合集 import save_direct_answer
        with patch.object(UnifiedSemanticEngine,'_http',return_value=self.answer()):
            run=run_step('recognize_a',self.config)
        old=load_fact_mapping(run['files']['a_mapping'],{})
        old['target_table']={k:old['outline']['subtables'][0][k] for k in ('id','sheet','range','data_ranges')}
        answer={'results':[{'coordinate':'B3','fields':{'asset_selection':{'type':'unknown','known_scope':'测试项目3','uncertainty':'具体条件未定'}}}],
                '识别结果':[{'坐标':['B4'],'实际维度取值':{'B4 asset_selection':'all：仅保留原答复条件'}}],
                '坐标':[{'coordinate':'B5','asset_selection':'未知'}]}
        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('不调用模型')):
            saved=save_direct_answer(old,json.dumps(answer,ensure_ascii=False),self.folder/'共用出口.json',
                field_scope={f'B{r}':['asset_selection'] for r in (3,4,5)})
        slots={r['cell']:r['partial_semantics'] for r in saved['slot_bindings']}
        self.assertEqual(slots['B3']['dimensions']['asset_selection']['known_scope'],'测试项目3')
        self.assertEqual(slots['B3']['dimensions']['asset_selection']['unknown'],'具体条件未定')
        self.assertEqual(slots['B5']['dimensions']['asset_selection']['mode'],'unknown')
        facts={r['id']:r for r in saved['facts']}
        self.assertEqual(facts['合成表!B4']['dimensions']['asset_selection']['mode'],'all')
        self.assertEqual(facts['合成表!B4']['dimensions']['asset_selection']['condition'],'仅保留原答复条件')
        before=next(r for r in old['facts'] if r['id']=='合成表!C3')
        self.assertEqual(facts['合成表!C3']['dimensions'],before['dimensions'])
        self.assertEqual(saved['coverage']['financial_cells'],12)

        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('不调用模型')):
            again=save_direct_answer(saved,'B3 asset_selection：未知',self.folder/'保留已知部分.json',field_scope={'B3':['asset_selection']})
        kept=next(r['partial_semantics'] for r in again['slot_bindings'] if r['cell']=='B3')
        self.assertEqual(kept['dimensions']['asset_selection'],slots['B3']['dimensions']['asset_selection'])

    def test_编号切换按所选标准接收维度且撤下明确旧错(self):
        from copy import deepcopy
        from 附注更新.A合集 import merge_direct_answer
        with patch.object(UnifiedSemanticEngine,'_http',return_value=self.answer()):
            run=run_step('recognize_a',self.config)
        old=load_fact_mapping(run['files']['a_mapping'],{})
        std=self.engine().standard
        dim=deepcopy(std['dimension']['asset_selection']);dim['id']='另一范围'
        metric=deepcopy(std['metric']['固定资产原值']);metric['id']='另一标准'
        metric['dimension_refs']=[{**r,'id':'另一范围'} if r['id']=='asset_selection' else r for r in metric['dimension_refs']]
        old['封存定义原文']+=list(map(lambda r:json.dumps(r,ensure_ascii=False),[dim,metric]))
        answer='B3 标准编号：另一标准；另一范围：原答复明确的范围\nB4 asset_selection：未知'
        merged,_,_=merge_direct_answer(old,answer,{'B3':['metric_id'],'B4':['asset_selection']},{'B5':['asset_selection']})
        self.assertTrue(any(r['id']=='另一范围' and r['ranges']==['B3'] for r in merged['dimensions']))
        self.assertTrue(any(r['id']=='asset_selection' and r['ranges']==['B4'] and r['value']['mode']=='unknown' for r in merged['dimensions']))
        self.assertFalse(any(r['id']=='asset_selection' and 'B5' in r['ranges'] for r in merged['dimensions']))

    def test_行内取值边界和编号明确未决经共用出口保存(self):
        from 附注更新.A合集 import save_direct_answer
        with patch.object(UnifiedSemanticEngine,'_http',return_value=self.answer()):
            run=run_step('recognize_a',self.config)
        old=load_fact_mapping(run['files']['a_mapping'],{})
        old['target_table']={k:old['outline']['subtables'][0][k] for k in ('id','sheet','range','data_ranges')}
        answer='B3 `asset_selection=all`，此处是原答复的解释。\n'
        answer+='| 坐标 | 标准槽位编号 | 其余维度实际取值 |\n|---|---|---|\n'
        answer+='| B4 | `固定资产原值` | `asset_selection=all`（原答复解释） |\n'
        answer+='| B5:C5 | `metric_id` 未决：原答复尚未选定编号。 | 原有其他维度保留 |\n'
        answer+='| B6 | 无法读取的编号文字 | 原有其他维度保留 |'
        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('不得重答')):
            saved=save_direct_answer(old,answer,self.folder/'行内取值与未决.json',
                field_scope={'B3':['asset_selection'],'B4':['metric_id','asset_selection'],'B5':['metric_id'],'C5':['metric_id'],'B6':['metric_id']})
        facts={r['id']:r for r in saved['facts']}
        self.assertEqual(facts['合成表!B3']['dimensions']['asset_selection']['mode'],'all')
        self.assertNotIn('合成表!B5',facts);self.assertNotIn('合成表!C5',facts)
        self.assertIn('合成表!B6',facts)
        pending={r['cell']:r for r in saved['unresolved']}
        self.assertNotIn('B4',pending)
        self.assertEqual(facts['合成表!B4']['dimensions']['asset_selection']['mode'],'all')
        for cell in ('B5','C5'):
            self.assertIn('未决',pending[cell]['reason']);self.assertEqual(pending[cell]['raw_value_state'],'blank')
        from 附注更新.A合集 import _inside
        before=old['outline']['subtables'][0];after=saved['outline']['subtables'][0]
        for cell in ('B5','C5'):
            values=lambda table:{r['id']:r['value'] for r in table['dimensions'] if _inside(cell,r['ranges'],table)}
            self.assertEqual(values(after),values(before))
        self.assertEqual(saved['coverage']['financial_cells'],12)
        self.assertEqual(saved['机械记录'],old['机械记录'])
        self.assertEqual(saved['答复接入来源']['本次原答复'],answer)

        reply='B5 标准编号：固定资产原值；period：2030-12-31；asset_selection：all'
        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('不得重答')):
            rebound=save_direct_answer(saved,reply,self.folder/'编号恢复后保留原维度.json',field_scope={'B5':['metric_id']})
        restored=next(r for r in rebound['facts'] if r['id']=='合成表!B5')
        original=next(r for r in old['facts'] if r['id']=='合成表!B5')
        self.assertEqual(restored['dimensions'],original['dimensions'])

    def test_分别对应的编号和紧邻中文的坐标(self):
        from 附注更新.A合集 import direct_answer,_direct_table
        standard={'metric':{'metric.a':{},'metric.b':{}},'dimension':{'selection':{'domain':'范围'}}}
        text='C3、D3分别为 `metric.a` 和 `metric.b`；`selection=all`。\nH8均为同一范围，selection：all'
        parsed=direct_answer(text,standard)
        self.assertEqual([(r['id'],r['ranges']) for r in parsed['metrics']],[('metric.a',['C3']),('metric.b',['D3'])])
        self.assertEqual(parsed['dimensions'][-1]['ranges'],['H8'])
        self.assertEqual(parsed['dimensions'][0]['value'],'all')
        standard['dimension']['selection']['value_type']='selection'
        table={'sheet':'合成表','metrics':[],'dimensions':[{'id':'selection','ranges':['C3'],'value':'members[原材料]'}]}
        value=_direct_table(table,{'source_context':[]},standard)['dimensions'][0]['value']
        self.assertEqual(value['mode'],'members');self.assertEqual(value['members'],['原材料'])
        two=direct_answer('| 坐标 | 槽位 |\n|---|---|\n| B3:C3 | `metric.a` / `metric.b` |',standard)
        self.assertFalse(two['metrics'],'同一范围列出两个槽位不能截取第一个冒充唯一槽位')
        self.assertEqual(two['unresolved'][0]['ranges'],['B3:C3'])
        self.assertIn('metric.a',two['unresolved'][0]['value'])
        self.assertIn('metric.b',two['unresolved'][0]['value'])

    def test_文字明示期间引号及同槽位引用不丢失(self):
        from 附注更新.A合集 import direct_answer
        text='''主体为**合成企业**，后面是解释文字。
共同取值：主体“合成企业”、报告口径 `standalone`。
### 《合成（单体）》
本表C=2025年度、D=2024年度；下列期间均为对应自然年度。
- `合成（单体）!C3`、`D3`：固定资产原值；asset_selection=“测试项目3”。
- `合成（单体）!C4`、`D4`：同一槽位；合计范围未说明，保留未知。
- `C5、D5`：同一指标；选择维度=“测试项目5”；period=2025-12-31（C）、2024-12-31（D）。
- `C6、D6`：不据此标作固定资产原值候选。
- `C7、D7`：同一指标；该选择维度取值=“测试项目7”；时点分别为2025-01-01、2024-01-01。
期末 period=2025-12-31（时点），期初 period=2024-12-31（时点）。
- `B8`、`E8`：同一指标；asset_selection=“测试项目8”；B、E列分别为期初、期末余额。
'''
        parsed=direct_answer(text,self.engine().standard,['合成（单体）'])
        self.assertEqual(next(r['value'] for r in parsed['dimensions'] if r['id']=='entity'),'合成企业')
        self.assertEqual({r['value'] for r in parsed['dimensions'] if r['id']=='entity'},{'合成企业'})
        targets={a for r in parsed['metrics'] for a in r['ranges']}
        self.assertIn('合成（单体）!C4',targets)
        self.assertNotIn('合成（单体）!C6',targets)
        self.assertTrue(any(r['reason']=='原答复明确不采用此槽位' for r in parsed['unresolved']))
        dims={(r['id'],a):r['value'] for r in parsed['dimensions'] for a in r['ranges']}
        self.assertEqual(dims['asset_selection','合成（单体）!C4'],'未知')
        self.assertEqual(dims['asset_selection','合成（单体）!C7'],'测试项目7')
        self.assertEqual(dims['period','合成（单体）!C3'],'2025-01-01至2025-12-31')
        self.assertEqual(dims['period','合成（单体）!C5'],'2025-12-31')
        self.assertEqual(dims['period','合成（单体）!D7'],'2024-01-01')
        self.assertEqual(dims['period','合成（单体）!B8'],'2024-12-31')
        self.assertEqual(dims['period','合成（单体）!E8'],'2025-12-31')
        from 附注更新.A合集 import _direct_table
        table=json.loads(self.answer())['tables'][0]
        table['dimensions']=[{'id':key,'ranges':['B3'],'value':value} for key,value in
                             [('currency','人民币（`CNY`）'),('report_scope','单体（`standalone`）')]]
        normalized=_direct_table(table,{'source_context':[]},self.engine().standard)
        self.assertEqual({r['id']:r['value'] for r in normalized['dimensions']},
                         {'currency':'CNY','report_scope':'standalone'})

    def test_独立句子编号按行取值及文字边界(self):
        from copy import deepcopy
        from 附注更新.A合集 import direct_answer
        standard=deepcopy(self.engine().standard)
        standard['metric']['另一编号']={**standard['metric']['固定资产原值'],'id':'另一编号'}
        text='''B原文主体为合成企业，主体不匹配，原文说明保持。
- B3:B4：`固定资产原值`。其中C3:C4选`另一编号`。
- B5:C8：槽位 `固定资产原值`；asset_selection：按行分别为项目甲、项目乙、项目丙、项目丁。
- B9：槽位 `固定资产原值`；asset_selection：mode=members；members=[项目甲]。其余说明原样保留。
- B10：槽位 `固定资产原值`；asset_selection：包含“其他项目”。
- B11、C11：槽位 `固定资产原值`；asset_selection：明确范围；时点分别为2025-12-31、2024-12-31。
A所需固定资产“其他”在B未找到对应明细。
- B12：槽位 `固定资产原值` / `另一编号`。
'''
        parsed=direct_answer(text,standard)
        self.assertTrue(any(r['id']=='固定资产原值' and r['ranges']==['B3:B4'] for r in parsed['metrics']))
        self.assertTrue(any(r['id']=='另一编号' and r['ranges']==['C3:C4'] for r in parsed['metrics']))
        self.assertFalse(any('B12' in r['ranges'] for r in parsed['metrics']))
        dims={(r['id'],a):r['value'] for r in parsed['dimensions'] for a in r['ranges']}
        self.assertEqual(dims['entity','*'],'合成企业')
        self.assertEqual(dims['asset_selection','B5'],'项目甲')
        self.assertEqual(dims['asset_selection','C8'],'项目丁')
        self.assertEqual(dims['asset_selection','B9'],{'mode':'members','members':['项目甲']})
        self.assertEqual(dims['asset_selection','B10'],'包含“其他项目”')
        self.assertEqual(dims['asset_selection','B11'],'明确范围')
        self.assertEqual(dims['period','C11'],'2024-12-31')

    def test_期间引用只连接明确坐标且不跨表(self):
        from 附注更新.A合集 import direct_answer, _direct_table
        text='''### 《合成表》
共同说明：期末余额的报告期间为时点2025-12-31；期初余额为时点2025-01-01；本期变动 period=2025-01-01至2025-12-31。
- B3（期末）、C3（期初）：固定资产原值；asset_selection=测试项目。
- B4:B5 为期末原值，C4:C5 为期初原值；槽位固定资产原值。
- B6:C6：本期增加和本期减少；asset_selection=测试项目。
- B7：固定资产原值，未确认本期期间。
- B8：固定资产原值；报告期间为时点2025-12-31；后续说明原样保留。
- B9：固定资产原值；报告期间为2025-12-31或2024-12-31。
- B10、C10：asset_selection=甲；分别为期末、期初。B11、C11：asset_selection=乙；分别为期末、期初。
### 《第二张表》
- B3（期末）：固定资产原值。
'''
        parsed=direct_answer(text,self.engine().standard,['合成表','第二张表'])
        periods=[r for r in parsed['dimensions'] if r['id']=='period']
        self.assertFalse(any('*' in r['ranges'] for r in periods))
        dims={a:r['value'] for r in periods for a in r['ranges']}
        self.assertEqual(dims['合成表!B3'],'2025-12-31')
        self.assertEqual(dims['合成表!C3'],'2025-01-01')
        self.assertEqual(dims['合成表!B4:B5'],'2025-12-31')
        self.assertEqual(dims['合成表!C4:C5'],'2025-01-01')
        self.assertEqual(dims['合成表!B6:C6'],'2025-01-01至2025-12-31')
        self.assertNotIn('合成表!B7',dims)
        self.assertNotIn('第二张表!B3',dims)
        self.assertEqual(dims['合成表!B8'],'2025-12-31')
        self.assertEqual(dims['合成表!B9'],'2025-12-31或2024-12-31')
        self.assertEqual(dims['合成表!B11'],'2025-12-31')
        self.assertEqual(dims['合成表!C11'],'2025-01-01')
        selections={a:r['value'] for r in parsed['dimensions'] if r['id']=='asset_selection' for a in r['ranges']}
        self.assertEqual(selections['合成表!B10'],'甲')
        self.assertEqual(selections['合成表!B11'],'乙')
        table=json.loads(self.answer())['tables'][0]
        table['dimensions']=[{'id':'period','ranges':['B3'],'value':'时点2025-01-01'}]
        normalized=_direct_table(table,{'source_context':[]},self.engine().standard)
        self.assertEqual(normalized['dimensions'][0]['value'],{'kind':'instant','date':'2025-01-01'})
        columns=direct_answer('共同期间：B列为2025-01-01至2025-12-31；C列为2024-01-01至2024-12-31。B、C分别对应本期与上期发生额。\n- B3、C3：固定资产原值。',self.engine().standard)
        roles={a:r['value'] for r in columns['dimensions'] if r['id']=='template_period_role' for a in r['ranges']}
        self.assertEqual(roles,{'B3':'current','C3':'prior'})

    def test_明确字段别称行列引用及日期声明只接本范围(self):
        from copy import deepcopy
        from 附注更新.A合集 import direct_answer
        from 附注更新.A语义阅读 import addresses
        standard=deepcopy(self.engine().standard)
        definitions=[
            ('cash_flow_activity','text','现金流量活动类别'),
            ('cash_flow_direction','text','现金收付方向'),
            ('cash_flow_presentation','text','现金流量列报口径'),
            ('cash_flow_selection','selection','现金流量项目'),
            ('cash_and_cash_equivalents_selection','selection','现金及现金等价物'),
            ('liability_selection','selection','负债'),
            ('liability_change_direction','text','负债变动方向'),
            ('liability_change_form','text','负债变动性质'),
            ('receivable_selection','selection','应收款项'),
            ('receivable_loss_change_direction','text','应收款项信用损失准备变动方向'),
            ('receivable_loss_change_selection','selection','应收款项信用损失准备变动事项'),
        ]
        for key,kind,domain in definitions:
            standard['dimension'][key]={'id':key,'value_type':kind,'domain':domain}
        standard['dimension']['liability_change_direction']['allowed_values']=['增加','减少','净增加','净减少']
        standard['dimension']['receivable_loss_change_direction']['allowed_values']=['增加','减少','净变动']
        common=['entity','report_scope','period','currency','unit_scale']
        fields={
            'metric.cash_flow.amount':['cash_flow_activity','cash_flow_direction','cash_flow_presentation','cash_flow_selection'],
            'metric.cash_and_cash_equivalents.net_change':['cash_and_cash_equivalents_selection'],
            'metric.liability.movement':['liability_selection','liability_change_direction','liability_change_form'],
            'metric.receivable.loss_allowance.change':['receivable_selection','receivable_loss_change_direction','receivable_loss_change_selection'],
            'metric.receivable.loss_allowance.balance':['receivable_selection'],
            'metric.two_selections':['cash_flow_selection','cash_and_cash_equivalents_selection'],
        }
        for key,refs in fields.items():
            standard['metric'][key]={**standard['metric']['固定资产原值'],'id':key,
                'dimension_refs':[{'id':ref,'required':True} for ref in common+refs]}
        def values(parsed,key,cell):
            return {row['value'] for row in parsed['dimensions'] if row['id']==key
                    and any(area!='*' and cell in addresses(area.rsplit('!',1)[-1]) for area in row['ranges'])}
        parsed=direct_answer('''- C15、D15：metric.cash_flow.amount；活动＝经营活动，方向＝净流出，列报口径＝净额列报，项目范围＝经营活动现金流量总体。
- C41、D41：metric.cash_and_cash_equivalents.net_change；选择维度＝现金及现金等价物。
- C50：metric.two_selections；选择维度＝未指定归属。
''',standard)
        for key,value in [('cash_flow_activity','经营活动'),('cash_flow_direction','净流出'),
                          ('cash_flow_presentation','净额列报'),('cash_flow_selection','经营活动现金流量总体')]:
            for cell in ['C15','D15']:self.assertEqual(values(parsed,key,cell),{value})
        for cell in ['C41','D41']:
            self.assertEqual(values(parsed,'cash_and_cash_equivalents_selection',cell),{'现金及现金等价物'})
        for key in fields['metric.two_selections']:self.assertEqual(values(parsed,key,'C50'),set())
        parsed=direct_answer('''### metric.liability.movement
- C621:D621：liability_change_direction依次为“增加”“减少”，liability_change_form＝“全部变动”。
- C629:D630、C642:D642：本期增加、本期减少；C列对应“增加”、D列对应“减少”，liability_change_form＝“全部变动”。
- C640:D640：其他项目，方向未知。
''',standard)
        for row in [629,630,642]:
            self.assertEqual(values(parsed,'liability_change_direction','C'+str(row)),{'增加'})
            self.assertEqual(values(parsed,'liability_change_direction','D'+str(row)),{'减少'})
        for cell in ['C640','D640']:
            self.assertTrue(values(parsed,'liability_change_direction',cell)<={'未知'})
        parsed=direct_answer('''### 信用损失准备变动表
- B269:D269 本期计提：各阶段增加；B270:D270 本期转回：各阶段减少；B271:D271 本期转销：各阶段减少；B272:D272 本期核销：各阶段减少。
- 对上述变动格，槽位均为 metric.receivable.loss_allowance.change，报告期间为2025-01-01至2025-12-31；receivable_loss_change_direction 按上列格子的增加或减少取值。receivable_loss_change_selection 按所在行分别取：本期计提、本期转回、本期转销、本期核销。
- B273:D273 为“其他变动”，但原文没有说明其变动方向；方向维度未知，不从金额或符号推断。
''',standard)
        for row,name in [(269,'本期计提'),(270,'本期转回'),(271,'本期转销'),(272,'本期核销')]:
            for col in 'BCD':
                self.assertEqual(values(parsed,'receivable_loss_change_direction',col+str(row)),{'增加' if row==269 else '减少'})
                self.assertEqual(values(parsed,'receivable_loss_change_selection',col+str(row)),{name})
        for col in 'BCD':
            self.assertTrue(values(parsed,'receivable_loss_change_direction',col+'273')<={'未知'})
            self.assertEqual(values(parsed,'receivable_loss_change_selection',col+'273'),set())
        parsed=direct_answer('''附注共同说明把“期末、期初”分别定义为2024-12-31、2024-01-01；其他报表注明2025年度，不能替代附注日期。
### 分类表
期末分类为行94—98，期初分类为行103—107。
- B94:B98：固定资产原值。
- B103:B107：固定资产原值。
### 账龄表
期末列B、C、D分别为账面余额、余额比例、坏账准备；期初列E、F、G分别为账面余额、余额比例、坏账准备。
- B114:B115、C114:C115、D114:D115、E114:E115、F114:F115、G114:G115：固定资产原值。
### 变动表
- B130:B134：metric.receivable.loss_allowance.balance，期间时点为2024-01-01；G130:G134：同一指标，时点为2025-12-31。
- B140：固定资产原值，期间未定。
- B141：固定资产原值，期间为2025年度。
''',standard)
        for row in range(94,99):self.assertEqual(values(parsed,'period','B'+str(row)),{'2024-12-31'})
        for row in range(103,108):self.assertEqual(values(parsed,'period','B'+str(row)),{'2024-01-01'})
        for col in 'BCDEFG':
            for row in [114,115]:
                self.assertEqual(values(parsed,'period',col+str(row)),{'2024-12-31' if col in 'BCD' else '2024-01-01'})
        for row in range(130,135):
            self.assertEqual(values(parsed,'period','B'+str(row)),{'2024-01-01'})
            self.assertEqual(values(parsed,'period','G'+str(row)),{'2025-12-31'})
            for col in 'BG':
                bound={item['id'] for item in parsed['metrics']
                       if any(area!='*' and col+str(row) in addresses(area.rsplit('!',1)[-1]) for area in item['ranges'])}
                self.assertEqual(bound,{'metric.receivable.loss_allowance.balance'})
        self.assertFalse(values(parsed,'period','B140')-{'未定','未知'})
        self.assertFalse(values(parsed,'period','B141')-{'2025年度'})

    def test_明确列事项成对坐标及共同后缀不补未知(self):
        from copy import deepcopy
        from 附注更新.A合集 import direct_answer
        from 附注更新.A语义阅读 import addresses
        standard=deepcopy(self.engine().standard)
        for key,kind,domain in [
            ('receivable_selection','selection','应收款项'),
            ('receivable_denominator_selection','selection','应收款项比例分母'),
            ('receivable_loss_change_direction','text','应收款项信用损失准备变动方向'),
            ('receivable_loss_change_selection','selection','应收款项信用损失准备变动事项'),
            ('liability_selection','selection','负债'),
        ]:
            standard['dimension'][key]={'id':key,'value_type':kind,'domain':domain}
        standard['dimension']['receivable_loss_change_direction']['allowed_values']=['增加','减少','净变动']
        common=['entity','report_scope','period','currency','unit_scale']
        for key,refs in [
            ('metric.receivable.loss_allowance.change',['receivable_selection','receivable_loss_change_direction','receivable_loss_change_selection']),
            ('metric.receivable.loss_allowance.balance',['receivable_selection']),
            ('metric.receivable.gross_balance.share_ratio',['receivable_selection','receivable_denominator_selection']),
            ('metric.liability.balance',['liability_selection']),
        ]:
            standard['metric'][key]={**standard['metric']['固定资产原值'],'id':key,
                'dimension_refs':[{'id':ref,'required':True} for ref in common+refs]}
        def values(parsed,key,cell):
            return {row['value'] for row in parsed['dimensions'] if row['id']==key
                    and any(area!='*' and cell in addresses(area.rsplit('!',1)[-1]) for area in row['ranges'])}
        parsed=direct_answer('''- C130:F134：metric.receivable.loss_allowance.change；receivable_selection=应收款项总体。C列为“增加”且事项为“计提”；D、E列为“减少”，事项分别为“收回或转回”“核销”；F列为“净变动”，限于其他但范围未说明。
''',standard)
        known_items={}
        for row in range(130,135):
            for col,direction in [('C','增加'),('D','减少'),('E','减少'),('F','净变动')]:
                self.assertEqual(values(parsed,'receivable_loss_change_direction',col+str(row)),{direction})
            for col,item in [('C','计提'),('D','收回或转回'),('E','核销')]:
                cell=col+str(row)
                known_items[cell]=values(parsed,'receivable_loss_change_selection',cell)
                self.assertEqual(known_items[cell],{item})
            self.assertTrue(values(parsed,'receivable_loss_change_selection','F'+str(row))<={'未知'})
        self.assertEqual(len(known_items),15)
        for key in ['receivable_loss_change_direction','receivable_loss_change_selection']:
            self.assertEqual(values(parsed,key,'C135'),set())
        parsed=direct_answer('''### metric.liability.balance
- B645/E645、B646/E646、B647/E647：分别为“应付职工薪酬；离职后福利；设定提存计划；基本养老保险”“应付职工薪酬；离职后福利；设定提存计划；失业保险费”“应付职工薪酬；离职后福利；设定提存计划；其他年金”。
- B648、E648是“应付职工薪酬；离职后福利；设定提存计划合计”，不是职工薪酬总合计。
- B625、E625为“应付职工薪酬合计”的期初、期末余额。
''',standard)
        for row,item in [(645,'基本养老保险'),(646,'失业保险费'),(647,'其他年金')]:
            for col in 'BE':
                self.assertEqual(values(parsed,'liability_selection',col+str(row)),
                                 {'应付职工薪酬；离职后福利；设定提存计划；'+item})
        for col in 'BE':
            self.assertEqual(values(parsed,'liability_selection',col+'648'),{'应付职工薪酬；离职后福利；设定提存计划合计'})
            self.assertEqual(values(parsed,'liability_selection',col+'625'),{'应付职工薪酬合计'})
            self.assertEqual(values(parsed,'liability_selection',col+'649'),set())
        parsed=direct_answer('- B264:D264：metric.receivable.loss_allowance.balance；receivable_selection分别为第一、第二、第三阶段的其他应收款项。',standard)
        for col,stage in [('B','第一'),('C','第二'),('D','第三')]:
            self.assertEqual(values(parsed,'receivable_selection',col+'264'),{stage+'阶段的其他应收款项'})
        parsed=direct_answer('''分子维度 `receivable_selection` 分别为单项计提、按组合计提、低风险组合、信用风险组合、其他应收款项分类总计；分母维度 `receivable_denominator_selection` 均为对应时点的其他应收款项总体。
- C249:C255、F249:F255：metric.receivable.gross_balance.share_ratio；分母维度 `receivable_denominator_selection` 为同期间信用风险组合的全部款项。C255、F255的总计分子和分母都对应信用风险组合全部款项。
''',standard)
        for col in 'CF':
            self.assertEqual(values(parsed,'receivable_selection',col+'255'),{'信用风险组合全部款项'})
            self.assertEqual(values(parsed,'receivable_selection',col+'249'),set())
            self.assertEqual(values(parsed,'receivable_denominator_selection',col+'249'),{'同期间信用风险组合的全部款项'})
        parsed=direct_answer('''期末余额的报告期间为时点2025-12-31；期初余额为时点2025-01-01。
- B249:B255 为期末各账龄原值，D249:D255 为期末坏账准备；E249:E255 为期初各账龄原值，G249:G255 为期初坏账准备。
- 对应账龄依次为：3个月以内、4—12个月、1—2年、2—3年、3年以上、空白账龄行、合计。账龄明细格的 receivable_selection 为“其他应收款项信用风险组合中，账龄为该行所列区间的款项”。
- C249:C255 和 F249:F255 是对应期间各行原值占比，槽位 metric.receivable.gross_balance.share_ratio。
''',standard)
        for row,age in enumerate(['3个月以内','4—12个月','1—2年','2—3年','3年以上'],249):
            for col in 'CF':
                self.assertEqual(values(parsed,'receivable_selection',col+str(row)),
                                 {'其他应收款项信用风险组合中，账龄为'+age+'的款项'})
        for col in 'CF':
            self.assertTrue(values(parsed,'receivable_selection',col+'254')<={'未知'})
            for row in range(249,256):self.assertEqual(values(parsed,'period',col+str(row)),set())
        parsed=direct_answer('''### metric.liability.balance
- `B687:C687`、`B689:C689`、`B692:C692`分别为其他应付款中的保证金及质保金、代扣代缴、其他，B列期末、C列期初；`B696`、`C696`是按性质列示的其他应付款合计，范围仍为其他应付款总体。
''',standard)
        for col in 'BC':
            self.assertEqual(values(parsed,'liability_selection',col+'696'),
                             {'其他应付款总体'})
            for row in [687,689,692]:
                self.assertNotIn('其他应付款总体',values(parsed,'liability_selection',col+str(row)))

    def test_答复依次列表和已声明的分子分母直接引用(self):
        from copy import deepcopy
        from 附注更新.A合集 import direct_answer
        standard=deepcopy(self.engine().standard)
        standard['dimension']['denominator_selection']={**standard['dimension']['asset_selection'],'id':'denominator_selection','domain':'分母范围'}
        text='''类别范围中，`asset_selection` 的 mode=predicate。
- B3:B6：固定资产原值；类别依次为项目甲、项目乙、项目丙、项目丁。
- C3:C6：固定资产原值；分子 `asset_selection` 依次为上述四项类别，分母 `denominator_selection` 均为 domain=分母范围, mode=all。
- D3:D6：固定资产原值；类别依次为上述四项。
- B7：固定资产原值；分子和分母的选择均为 domain=固定资产, mode=all。
- C8：固定资产原值；asset_selection分别为“项目甲”“项目乙”。
- B9：槽位 `固定资产原值`。C9:C10：槽位 `固定资产原值`；分子维度asset_selection分别为项目甲、项目乙；分母维度denominator_selection为这两个项目总体。
- B11:D11：槽位 `固定资产原值`；asset_selection分别为第一、第二、第三类项目。
### 新段落
- B12：原文项目。
- 对上述格，槽位 `固定资产原值`；asset_selection为已答项目。
'''
        parsed=direct_answer(text,standard)
        dims={(r['id'],a):r['value'] for r in parsed['dimensions'] for a in r['ranges']}
        self.assertEqual(dims['asset_selection','B3'],'项目甲')
        self.assertEqual(dims['asset_selection','C6'],'项目丁')
        self.assertEqual(dims['asset_selection','D4'],'项目乙')
        self.assertEqual(dims['asset_selection','B7'],{'domain':'固定资产','mode':'all'})
        self.assertEqual(dims['denominator_selection','B7'],{'domain':'固定资产','mode':'all'})
        self.assertNotIn(('asset_selection','C8'),dims)
        self.assertEqual(dims['asset_selection','C10'],'项目乙')
        self.assertEqual(dims['denominator_selection','C9:C10'],'这两个项目总体')
        self.assertNotIn(('denominator_selection','B9'),dims)
        for col,name in [('B','第一类项目'),('C','第二类项目'),('D','第三类项目')]:
            self.assertEqual(dims['asset_selection',col+'11'],name)
        grouped=[r for r in parsed['metrics'] if 'C9:C10' in r['ranges']]
        self.assertEqual(len(grouped),1)
        self.assertNotIn('B12',grouped[0]['ranges'])
        self.assertTrue(any(r['ranges']==['B12'] for r in parsed['metrics']))
        grouped=direct_answer('### 固定资产原值\n- B3：asset_selection=甲。C3：asset_selection=乙。\n### 下一段\n- B4：原文项目。',standard)
        self.assertEqual(grouped['metrics'][0]['ranges'],['B3','C3'])
        self.assertTrue(any(r['ranges']==['C8'] and '顺序列表' in r['reason'] for r in parsed['unresolved']))

    def test_B文字工作表共同字段及行内坐标直接保存(self):
        from copy import deepcopy
        from 附注更新.A合集 import save_direct_answer, _records
        with patch.object(UnifiedSemanticEngine,'_http',return_value=self.answer()):
            run=run_step('recognize_a',self.config)
        old=load_fact_mapping(run['files']['a_mapping'],{})
        old.pop('target_table',None);old['carrier']='B'
        snapshot=deepcopy(old['机械记录'])
        snapshot['sheets'].append({**deepcopy(snapshot['sheets'][0]),'name':'第二张表'})
        text='共同取值（以下均为 `固定资产原值`）：`entity=合成企业`；`report_scope=standalone`；`currency=CNY`、`unit_scale=1`。\n'
        text+='`合成表` 的候选如下：\n'
        text+='- **B1｜A3 测试项目**：`asset_selection`＝predicate「测试项目范围」；B3对应 `period={date:2025-12-31, kind:instant}`，C3对应 `period={date:2024-12-31, kind:instant}`。\n'
        text+='- **B2｜A5 测试项目**：B5期末为2025年12月31日，C5期初为2024年12月31日；`asset_selection` 的范围保留为未知。喵'
        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('不得重答')):
            saved=save_direct_answer(old,text,self.folder/'B自然文字接入.json',snapshot=snapshot)
        rows=dict(_records(saved))
        self.assertEqual(set(rows),{'合成表!'+c for c in ('B3','C3','B5','C5')})
        self.assertEqual(saved['识别答复'],text)
        self.assertFalse(saved['unresolved'],saved['unresolved'])
        for cell,row in rows.items():
            self.assertEqual(row['dimensions']['period']['date'],'2025-12-31' if '!B' in cell else '2024-12-31')
            self.assertEqual(row['dimensions']['entity'],'合成企业')
        self.assertEqual(rows['合成表!B3']['dimensions']['asset_selection']['predicate'],'测试项目范围')
        self.assertEqual(rows['合成表!B5']['dimensions']['asset_selection']['mode'],'unknown')
        self.assertEqual(len(saved['facts']),2)
        self.assertEqual(len(saved['slot_bindings']),2)
        self.assertTrue(all(r['raw_value_state']=='blank' for r in saved['slot_bindings']))
        from 附注更新.A合集 import direct_answer
        text='### 明细：`合成表!B3:C5`\n期末列（B；期间2025-12-31）\n- `B3:B5`：`固定资产原值`。\n期初列（C；期间2024-12-31）\n- `C3:C5`：`固定资产原值`。'
        parsed=direct_answer(text,self.engine().standard,['合成表','第二张表'])
        self.assertEqual([r['ranges'] for r in parsed['metrics']],[['合成表!B3:B5'],['合成表!C3:C5']])
        self.assertEqual([(r['ranges'],r['value']) for r in parsed['dimensions'] if r['id']=='period'],
                         [(['合成表!B3:B5'],'2025-12-31'),(['合成表!C3:C5'],'2024-12-31')])
        text='共同取值：`entity` 均为合成企业；`report_scope` 为 `standalone`；`currency` 为 CNY；`unit_scale` 为 1。\n'
        text+='| 坐标 | 槽位 | asset_selection |\n|---|---|---|\n| 合成表!B3、C3 | 固定资产原值 | `domain`=固定资产；`mode`=members；`members`=[项目甲] |\n'
        text+='| 坐标 | 槽位 | 范围实际取值 |\n| 合成表!B5、C5 | 固定资产原值 | `asset_selection`：`domain`=固定资产, mode=all |'
        parsed=direct_answer(text,self.engine().standard,['合成表','第二张表'])
        selections=[r for r in parsed['dimensions'] if r['id']=='asset_selection']
        self.assertEqual([r['ranges'] for r in selections],[['合成表!B3','合成表!C3'],['合成表!B5','合成表!C5']])
        self.assertEqual(selections[0]['value'],{'domain':'固定资产','mode':'members','members':['项目甲']})
        self.assertEqual(selections[1]['value'],{'domain':'固定资产','mode':'all'})
        old['原文材料']={'sheets':[{'sheet':'合成表','rows':'3: B=6 | C=9'}]}
        saved=save_direct_answer(old,'`合成表!B5`：固定资产原值',self.folder/'B答复串入未提供坐标.json',snapshot=snapshot)
        self.assertEqual([(r['original']['sheet'],r['original']['cell']) for r in saved['unresolved']],[('合成表','B5')])
        self.assertEqual(len(saved['slot_bindings']),0,'越出已提供原文的位置保留为原声明，不污染其他组的真实B候选')
        text='以下候选均为 `固定资产原值`。\n`entity`：合成企业；`report_scope`：standalone（单体）；`currency`：人民币（CNY）；`unit_scale`：1\n'
        text+='`period`：B列为2025-12-31；C列为2024-12-31。\n'
        text+='| B、C坐标 | asset_selection |\n|---|---|\n| 合成表!B3、C3 | 固定资产；members：项目甲 |\n'
        text+='本段说明提到C3，但没有撤销任何字段。'
        old.pop('原文材料',None)
        saved=save_direct_answer(old,text,self.folder/'B分组编号及按列期间.json',snapshot=snapshot)
        self.assertEqual(len(saved['facts']),2);self.assertFalse(saved['unresolved'])
        for loc,row in _records(saved):
            self.assertEqual(row['dimensions']['report_scope'],'standalone')
            self.assertEqual(row['dimensions']['currency'],'CNY')
            self.assertEqual(row['dimensions']['period']['date'],'2025-12-31' if '!B' in loc else '2024-12-31')
        rejected=save_direct_answer(saved,text,self.folder/'B已确认坐标错误保留原答复.json',
                                   not_adopted_reason='本次答复把A坐标写成B坐标，尚无正确B位置')
        self.assertEqual(rejected['识别答复'],text)
        self.assertFalse(rejected['facts'] or rejected['slot_bindings'])
        self.assertIn('不采用',rejected['unresolved'][0]['reason'])
        reread=save_direct_answer(rejected,text,self.folder/'同一错误答复重读仍不采用.json')
        self.assertFalse(reread['facts'] or reread['slot_bindings'])

    def test_列共同取值和反引号字段经共用出口保存(self):
        from 附注更新.A合集 import save_direct_answer, direct_answer, _direct_table
        text='B列共同适用：\n- `asset_selection`：共同范围。\n各格沿用上述共同取值：\n- **B5**：原答复说明。\n- **B6**：原答复说明。\n'
        text+='C列共同适用：\n- `asset_selection`：另一范围。\n- **C5**：原答复说明。'
        parsed=direct_answer(text,self.engine().standard)
        self.assertEqual([(r['ranges'],r['value']) for r in parsed['dimensions']],
                         [(['B5','B6'],'共同范围'),(['C5'],'另一范围')])
        with patch.object(UnifiedSemanticEngine,'_http',return_value=self.answer()):
            run=run_step('recognize_a',self.config)
        old=load_fact_mapping(run['files']['a_mapping'],{})
        old['target_table']={k:old['outline']['subtables'][0][k] for k in ('id','sheet','range','data_ranges')}
        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('不得重答')):
            saved=save_direct_answer(old,text,self.folder/'列共同取值.json',
                                    field_scope={c:['asset_selection'] for c in ('B5','B6','C5')})
        before={r['id']:r for r in old['facts']}
        for row in saved['facts']:
            cell=row['source_reference']['location'].split('!')[-1]
            if cell in ('B5','B6','C5'):
                self.assertEqual(row['dimensions']['asset_selection']['predicate'],'另一范围' if cell=='C5' else '共同范围')
            else:self.assertEqual(row['dimensions'],before[row['id']]['dimensions'])
        self.assertEqual(len(saved['facts']),12)
        self.assertEqual(sum(r['raw_value_state']=='blank' for r in saved['facts']),8)
        self.assertEqual(saved['机械记录'],old['机械记录'])
        self.assertEqual(saved['答复接入来源']['本次原答复'],text)
        standard={'metric':{},'dimension':{'测试方向':{'value_type':'text','allowed_values':['增加','减少','净变动']}}}
        row=direct_answer('B5 `测试方向`：**净变动**。原答复说明。',standard)
        table=_direct_table({'sheet':'合成表',**row},{'source_context':[]},standard)
        self.assertEqual(table['dimensions'][0]['value'],'净变动')

    def test_日期括号角色和顺序列表共同措辞保存回读(self):
        from 附注更新.A合集 import save_direct_answer, _records
        from 附注更新.A语义阅读 import addresses
        with patch.object(UnifiedSemanticEngine,'_http',return_value=self.answer()):
            run=run_step('recognize_a',self.config)
        old=load_fact_mapping(run['files']['a_mapping'],{})
        old['carrier']='B'
        reply=('合成表!B5:B6：period＝2024-01-01（期初余额）。\n'
               '合成表!C5:C6：period＝2024-12-31（表头角色：期末余额）。\n'
               '合成表!B5:B6：asset_selection分别为类别“甲”“乙”的固定资产。')
        fields={'B5:B6':['period','template_period_role','asset_selection'],
                'C5:C6':['period','template_period_role']}
        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('不得重答')):
            saved=save_direct_answer(old,reply,self.folder/'括号角色与共同措辞.json',field_scope=fields)
            reread=load_fact_mapping(saved['mapping_path'],{})
        original=dict(_records(old));actual=dict(_records(reread))
        self.assertEqual(set(actual),set(original))
        self.assertEqual(reread['机械记录'],old['机械记录'])
        self.assertEqual(sum(r['raw_value_state']=='blank' for r in actual.values()),8)
        pending=save_direct_answer(old,'合成表!B7：period＝2024-01-01（期初余额）。',
                                   self.folder/'尚缺编号的期间角色.json')
        self.assertFalse(pending['facts'])
        declared=[r for t in pending['outline']['subtables'] for r in t['dimensions']
                  if r['id']=='template_period_role']
        self.assertEqual([(r['ranges'],r['value']) for r in declared],[(['B7'],'opening')])
        rebound=save_direct_answer(pending,'合成表!B7：metric_id＝固定资产原值。',
                                  self.folder/'补编号沿用已有角色.json',field_scope={'合成表!B7':['metric_id']})
        self.assertEqual(rebound['outline']['subtables'][0]['metrics'][0]['template_period_role'],'opening')
        for cell,item in [('B5','甲'),('B6','乙')]:
            self.assertEqual(actual['合成表!'+cell]['dimensions']['asset_selection']['predicate'],
                             '类别'+item+'的固定资产')
        for cell,role in [('B5','opening'),('B6','opening'),('C5','closing'),('C6','closing')]:
            values={m.get('template_period_role') for t in reread['outline']['subtables']
                    for m in t['metrics'] if any(cell in addresses(s) for s in m['ranges'])}
            self.assertEqual(values,{role})
        for loc in original:
            if loc not in {'合成表!'+c for c in ('B5','B6','C5','C6')}:
                self.assertEqual(actual[loc]['dimensions'],original[loc]['dimensions'])

    def test_字段后的坐标及括号内坐标明确未知(self):
        from 附注更新.A合集 import direct_answer, save_direct_answer
        text='report_scope：B5、C5 均为“consolidated”。原答复说明。'
        parsed=direct_answer(text,self.engine().standard)['dimensions'][0]
        self.assertEqual((parsed['ranges'],parsed['value']),(['B5','C5'],'consolidated'))
        description=direct_answer('共同维度（B5:C5）\n`asset_selection`（实际范围）：其他项目；范围未确定。',self.engine().standard)['dimensions'][0]
        self.assertEqual((description['ranges'],description['value']),(['B5:C5'],'其他项目；范围未确定'))
        with patch.object(UnifiedSemanticEngine,'_http',return_value=self.answer()):
            run=run_step('recognize_a',self.config)
        old=load_fact_mapping(run['files']['a_mapping'],{})
        old['target_table']={k:old['outline']['subtables'][0][k] for k in ('id','sheet','range','data_ranges')}
        reply='共同取值（合成表!B5、C5）：`report_scope=未知`。'
        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('不得重答')):
            saved=save_direct_answer(old,reply,self.folder/'明确未知不留旧值.json',
                                    field_scope={c:['report_scope'] for c in ('B5','C5')})
        self.assertEqual(len(saved['slot_bindings']),2)
        for row in saved['slot_bindings']:
            self.assertNotIn('report_scope',row['partial_semantics']['dimensions'])
        self.assertEqual(saved['机械记录'],old['机械记录'])
        self.assertEqual(saved['答复接入来源']['本次原答复'],reply)
        parsed=direct_answer("共同取值（'合成表 (1)'!B5）：`report_scope=未知`。",self.engine().standard)
        self.assertEqual(parsed['dimensions'][0]['ranges'],["'合成表 (1)'!B5"])

    def test_表头字段分母同行引用及扣除限定原样接入(self):
        from copy import deepcopy
        from 附注更新.A合集 import direct_answer, _direct_table
        from 附注更新.A语义阅读 import addresses
        standard=deepcopy(self.engine().standard)
        for key,domain in [('fixed_asset_selection','固定资产'),('receivable_selection','应收款项'),
                           ('receivable_denominator_selection','应收款项分母范围'),
                           ('receivable_loss_change_selection','信用损失准备变动事项'),('liability_selection','负债')]:
            standard['dimension'][key]={**standard['dimension']['asset_selection'],'id':key,'domain':domain}
        for key,fields in [('metric.fixed_asset.cost.balance',['fixed_asset_selection']),
                           ('metric.fixed_asset.cost.change',['fixed_asset_selection']),
                           ('metric.receivable.loss_allowance.balance',['receivable_selection'])]:
            standard['metric'][key]={**standard['metric']['固定资产原值'],'id':key,
                                    'dimension_refs':[{'id':field,'required':True} for field in fields]}
        def values(parsed,key,cell):
            return {row['value'] for row in parsed['dimensions'] if row['id']==key
                    and any(area!='*' and cell in addresses(area.rsplit('!',1)[-1]) for area in row['ranges'])}
        fixed=direct_answer('''以下表格坐标均在`word附注`工作表。
| 行项目 | `fixed_asset_selection`实际范围 | `metric_id` |
|---|---|---|
| 400 账面原值合计 | `all`，完整固定资产总体 | `metric.fixed_asset.cost.balance`（B、E列余额）；`metric.fixed_asset.cost.change`（C、D列变动） |
| 401 房屋建筑物 | 房屋及建筑物 | 同上原值指标 |
''',standard,['word附注'])
        for col in 'BCDE':
            self.assertEqual(values(fixed,'fixed_asset_selection',col+'400'),{'all'})
            self.assertEqual(values(fixed,'fixed_asset_selection',col+'401'),{'房屋及建筑物'})
            self.assertFalse(values(fixed,'fixed_asset_selection',col+'402'))
        normalized=_direct_table({'sheet':'word附注',**fixed},{'source_context':[]},standard)
        self.assertEqual(normalized['dimensions'][0]['value'],{'mode':'all','domain':'固定资产'})
        reply='''#### 比例字段：receivable_selection、receivable_denominator_selection
- word附注!C94:C98：分子为同行应收账款范围；分母为应收账款总体。
- word附注!C103:C107：分子为同行应收账款范围；分母为应收账款总体。
- word附注!C114:C120、F114:F120：分子为同行的信用风险组合账龄范围；分母为信用风险组合应收账款总体。
#### 变动表：receivable_selection、receivable_loss_change_selection
`word附注!B130:B134`和`G130:G134`分别标为“期初余额”“期末余额”，对应2024-01-01和2024-12-31。按行范围：
- 第130行：单项计提预期信用损失的应收账款。
- 第131行：按组合计提预期信用损失的应收账款总体，包括其下类别。
- 第132行：组合计提范围内的账龄组合应收账款。
- 第133行：组合计提范围内的无风险组合应收账款。
- 第134行合计：应收账款总体，包括单项计提及组合计提。
`word附注!C130:F134`标为本期变动，对应2024-01-01至2024-12-31。行范围按上述第130—134行对象对应；列范围分别为：`C130:C134`计提，`D130:D134`收回或转回，`E130:E134`核销，`F130:F134`其他变动。其他变动的具体事项范围未披露。
- word附注!H130：metric.receivable.loss_allowance.balance。
#### 下一节
- word附注!B135:G135：范围未知。
'''
        context={'word附注!'+col+str(row):'metric.receivable.loss_allowance.balance' for col in 'BG' for row in range(130,136)}
        parsed=direct_answer(reply,standard,['word附注'],context)
        for area in ['C94:C98','C103:C107','C114:C120','F114:F120']:
            expected='信用风险组合应收账款总体' if area in ['C114:C120','F114:F120'] else '应收账款总体'
            for cell in addresses(area):self.assertEqual(values(parsed,'receivable_denominator_selection',cell),{expected})
        row_values=['单项计提预期信用损失的应收账款','按组合计提预期信用损失的应收账款总体，包括其下类别',
                    '组合计提范围内的账龄组合应收账款','组合计提范围内的无风险组合应收账款','应收账款总体，包括单项计提及组合计提']
        for row,value in zip(range(130,135),row_values):
            for col in 'BCDEFG':self.assertEqual(values(parsed,'receivable_selection',col+str(row)),{value})
            self.assertFalse(values(parsed,'receivable_loss_change_selection','F'+str(row)))
        for col in 'BCDEFG':self.assertFalse(values(parsed,'receivable_selection',col+'135'))
        self.assertFalse(values(parsed,'receivable_denominator_selection','C121'))
        self.assertFalse(values(parsed,'receivable_selection','B103'))
        self.assertFalse(values(parsed,'receivable_selection','H130'),'同一行其他格没有引用，不复制范围')
        self.assertEqual([row['ranges'] for row in parsed['metrics']],[['word附注!H130']],
                         '已有编号只用于明确字段归属，不生成新的编号声明')
        liability=direct_answer('''本次列示的负债选取范围均使用金标准维度 liability_selection，总体为负债。
- `word附注!B674`、`C674`：其他应付款合计；按原文注释，表中该合计扣除应付利息、应付股利。
- `word附注!B675`、`C675`：原文未说明。
''',standard,['word附注'])
        for col in 'BC':
            self.assertEqual(values(liability,'liability_selection',col+'674'),
                             {'其他应付款合计；按原文注释，表中该合计扣除应付利息、应付股利'})
            self.assertFalse(values(liability,'liability_selection',col+'675'))
        normalized=_direct_table({'sheet':'word附注',**liability},{'source_context':[]},standard)
        self.assertEqual(normalized['dimensions'][0]['value']['mode'],'predicate')

    def test_共同字段只接明确分别及以上明细的同行范围(self):
        from copy import deepcopy
        from 附注更新.A合集 import direct_answer
        from 附注更新.A语义阅读 import addresses
        standard=deepcopy(self.engine().standard)
        standard['dimension']['liability_selection']={**standard['dimension']['asset_selection'],'id':'liability_selection','domain':'负债'}
        text='''本次列示的负债选取范围均使用金标准维度 `liability_selection`，总体为负债。
- `word附注!C621:D621`：应付职工薪酬中的短期薪酬。
- `word附注!B629:B630`、`B632`：分别是短期薪酬中的“工资”“职工福利费”“社会保险费”。父项与其子项分别保留。
- `word附注!C629:D632`：以上短期薪酬明细及 `C634:D635` 的“生育保险费”“大病医疗保险费”。C列实际为本期增加，D列实际为本期减少；这里仅列 `liability_selection`。
- `word附注!C645:D647`：设定提存计划中的“基本养老保险”“失业保险费”“其他年金”。
- `word附注!B651:E652`、`B664:E664`：分别为应交税费中的“增值税”“企业所得税”“其他”。
- `word附注!B667:E667`：应交税费合计；当前列出的税费明细不替代合计总体。
- `word附注!B674`、`C674`：其他应付款合计；按原文注释，表中该合计扣除应付利息、应付股利。
- `word附注!B692`、`C692`：其他应付款中的“其他”；此明细不等同整个其他应付款合计。
'''
        parsed=direct_answer(text,standard,['word附注','另一表'])
        values={}
        for item in parsed['dimensions']:
            if item['id']!='liability_selection':continue
            for area in item['ranges']:
                for cell in addresses(area.rsplit('!',1)[-1]):values.setdefault(cell,set()).add(item['value'])
        for cell in ['C621','D621']:self.assertEqual(values[cell],{'应付职工薪酬中的短期薪酬'})
        for row,value in [(629,'工资'),(630,'职工福利费'),(632,'社会保险费')]:
            for col in 'BCD':self.assertEqual(values[col+str(row)],{'短期薪酬中的'+value})
            self.assertNotIn('E'+str(row),values)
        for row,value in [(651,'增值税'),(652,'企业所得税'),(664,'其他')]:
            for col in 'BCDE':self.assertEqual(values[col+str(row)],{'应交税费中的'+value})
        for area in ['C631:D631','C634:D635','C645:D647']:
            for cell in addresses(area):self.assertNotIn(cell,values)
        for col in 'BCDE':self.assertEqual(values[col+'667'],{'应交税费合计；当前列出的税费明细不替代合计总体'})
        for col in 'BC':
            self.assertEqual(values[col+'674'],{'其他应付款合计；按原文注释，表中该合计扣除应付利息、应付股利'})
            self.assertEqual(values[col+'692'],{'其他应付款中的“其他”；此明细不等同整个其他应付款合计'})
        self.assertFalse(any(row['id']=='liability_change_direction' for row in parsed['dimensions']))
        other=direct_answer(text+'\n- `另一表!C629:D629`：以上短期薪酬明细。',standard,['word附注','另一表'])
        self.assertFalse(any(area.startswith('另一表!') for row in other['dimensions'] for area in row['ranges']))
        heading='以下仅列本次范围内的 `liability_selection`，不填写金额。'
        alternate=direct_answer(heading+'\n'+text.split('\n',1)[1],standard,['word附注','另一表'])
        self.assertEqual([(row['id'],row['ranges'],row['value']) for row in alternate['dimensions']],
                         [(row['id'],row['ranges'],row['value']) for row in parsed['dimensions']])

    def test_原标准定义未裁剪且后台不进入B或后续(self):
        from 附注更新.无界面 import run_background
        with patch.object(UnifiedSemanticEngine, '_http', return_value=self.answer()):
            a = run_step('recognize_a', self.config)
        saved = load_fact_mapping(a['files']['a_mapping'], {})
        expected = [line for line in self.gold.read_text(encoding='utf-8').splitlines()
                    if (lambda row: row.get('record_type') == 'dimension' or row.get('record_type') == 'metric' and row.get('status') == 'active')(json.loads(line))]
        self.assertEqual(saved['金标准原文'], '\n'.join(expected))
        seen = []
        def fake(step, config, *args):
            seen.append(step)
            return {'status': 'complete' if step == 'extract_a' else 'saved', 'files': {}, 'message': '离线'}
        with patch('附注更新.无界面.run_step', side_effect=fake):
            result = run_background({**self.config, 'word_path': str(self.source)})
        self.assertEqual(seen, ['extract_a', 'recognize_a'])
        self.assertEqual(result['status'], 'saved')


    def test_字段标题合计明细独立说明及角引号不混入取值(self):
        from copy import deepcopy
        from 附注更新.A合集 import direct_answer, _direct_table
        from 附注更新.A语义阅读 import addresses
        standard=deepcopy(self.engine().standard)
        standard['dimension']['change_direction']={'id':'change_direction','domain':'变动方向',
            'value_type':'text','allowed_values':['增加','减少']}
        reply='''- **asset_selection**（固定资产原值）
  - B3、C3“合计”：`all`，为固定资产总体。
合成表!B4:B6、C4:C6：行范围依次为合计、甲、乙；合计的固定资产范围为 `all`，明细依次为甲、乙。
合成表!B7、C7：固定资产“合计”，asset_selection为固定资产总体`all`；B列期末余额（2025-12-31），C列期初余额（2024-12-31）。
合成表!B8：change_direction=「减少」。
'''
        parsed=direct_answer(reply,standard,['合成表'])
        table=_direct_table({**parsed,'sheet':'合成表','data_ranges':['B3:C8']},
                            {'source_context':[]},standard)
        def values(cell,key):
            return [row['value'] for row in table['dimensions'] if row['id']==key
                    and any(cell in addresses(span) for span in row['ranges'])]
        self.assertEqual(values('B3','asset_selection'),[{'mode':'all','domain':'固定资产'}])
        self.assertEqual(values('B4','asset_selection'),[{'mode':'all','domain':'固定资产'}])
        self.assertEqual(values('C5','asset_selection'),[{'mode':'predicate','predicate':'甲','domain':'固定资产'}])
        self.assertEqual(values('B6','asset_selection'),[{'mode':'predicate','predicate':'乙','domain':'固定资产'}])
        self.assertEqual(values('B7','asset_selection'),[{'mode':'all','domain':'固定资产'}])
        self.assertEqual(values('B8','change_direction'),['减少'])


    def test_同值未给新引用逐格保留旧引用(self):
        from 附注更新.A合集 import save_direct_answer, _records
        from 附注更新.表格 import read_workbook, materialize_candidate_blanks
        engine=self.engine();snapshot=read_workbook(self.source,context=[]);materialize_candidate_blanks(snapshot)
        answer=json.loads(self.answer());dimensions=answer['tables'][0]['dimensions']
        dimensions[4]['ranges']=['B5:B8']
        dimensions.extend({'id':'period','value':{'kind':'instant','date':'2025-12-31'},
                           'ranges':[cell],'evidence':[ref]} for cell,ref in [('B3','合成表!B2'),('B4','合成表!A2')])
        base={'recognition_method':'direct','source_path':str(self.source),'source_hash':snapshot['sha256'],
              'source_context':[],'gold_path':str(self.gold),'gold_hash':engine.gold_hash,'carrier':'A',
              '金标准原文':'\n'.join(line for line in self.gold.read_text(encoding='utf-8').splitlines()
                  if json.loads(line)['record_type'] in {'metric','dimension'})}
        old=save_direct_answer(base,json.dumps(answer,ensure_ascii=False),self.folder/'保留引用原合集.json',
                              snapshot=snapshot,standard=engine.standard)
        reply=json.dumps({'dimensions':[
            {'id':'period','value':'2025-12-31','ranges':['B3:B4']},
            {'id':'period','value':'2025-12-31','ranges':['B5'],'evidence':['合成表!C2']},
            {'id':'period','value':'2026-12-31','ranges':['B6']}]},ensure_ascii=False)
        saved=save_direct_answer(old,reply,self.folder/'保留引用新合集.json',
                                field_scope={f'合成表!B{n}':['period'] for n in range(3,7)})
        rows=dict(_records(saved));before=dict(_records(old))
        self.assertEqual(set(rows),set(before))
        for cell in ('B3','B4'):
            self.assertEqual(rows['合成表!'+cell]['dimension_evidence']['period'],
                             before['合成表!'+cell]['dimension_evidence']['period'])
        self.assertEqual(rows['合成表!B5']['dimension_evidence']['period'],['合成表!C2'])
        self.assertEqual(rows['合成表!B6']['dimension_evidence']['period'],[])
        self.assertEqual(rows['合成表!B6']['dimensions']['period'],{'kind':'instant','date':'2026-12-31'})

    def test_角色括号标签引用及instant日期保存回读(self):
        from 附注更新.A合集 import save_direct_answer, _records
        from 附注更新.表格 import read_workbook, materialize_candidate_blanks
        engine=self.engine();snapshot=read_workbook(self.source,context=[]);materialize_candidate_blanks(snapshot)
        base={'recognition_method':'direct','source_path':str(self.source),'source_hash':snapshot['sha256'],
              'source_context':[],'gold_path':str(self.gold),'gold_hash':engine.gold_hash,'carrier':'A',
              '金标准原文':'\n'.join(line for line in self.gold.read_text(encoding='utf-8').splitlines()
                  if json.loads(line)['record_type'] in {'metric','dimension'})}
        answer=json.loads(self.answer())
        for row in answer['tables'][0]['dimensions'][4:6]:row['evidence']=['合成表!A2']
        old=save_direct_answer(base,json.dumps(answer,ensure_ascii=False),self.folder/'角色引用原合集.json',snapshot=snapshot,standard=engine.standard)
        reply='\n'.join(f'- `合成表!{cell}`（测试项目）：`template_period_role` = 期末余额（原表期间标签引用：`合成表!{ref}`）；`period` = instant，{date}。'
            for cell,ref,date in [('B3','B2','2025-12-31'),('B4','B2','2025-12-31'),('C3','C2','2024-12-31')])
        reply+='\n合成表!B5：`template_period_role`＝“期末余额”（引用原表期间标签）；金标准维度 `period`＝时点 `2025-12-31`。'
        saved=save_direct_answer(old,reply,self.folder/'角色引用新合集.json',
            field_scope={f'合成表!{cell}':['period','template_period_role'] for cell in ('B3','B4','C3','B5')})
        rows=dict(_records(saved))
        self.assertEqual(set(rows),set(dict(_records(old))))
        self.assertEqual(rows['合成表!B5']['dimensions']['period'],{'kind':'instant','date':'2025-12-31'})
        self.assertFalse(any(row.get('cell')=='B5' for row in saved['unresolved']))
        for cell,ref,date in [('B3','B2','2025-12-31'),('B4','B2','2025-12-31'),('C3','C2','2024-12-31')]:
            self.assertEqual(rows['合成表!'+cell]['dimensions']['period'],{'kind':'instant','date':date})
            self.assertEqual(rows['合成表!'+cell]['dimension_evidence']['period'],['合成表!'+ref])
            roles=[row['template_period_role'] for table in saved['outline']['subtables']
                   for row in table['metrics'] if row['ranges']==[cell]]
            self.assertEqual(roles,['期末余额'])

    def test_独立期间标签声明及合并引用接通匹配(self):
        from copy import deepcopy
        from openpyxl import load_workbook
        from 附注更新.A合集 import save_direct_answer, direct_answer, _records
        from 附注更新.表格 import read_workbook
        from 附注更新.模板更新编排 import direct_update_plan
        engine=self.engine()
        reply='合成表!B3:B4：period = 2026-12-31。\n期间标签坐标为 `合成表!B2:C2`“期末余额”，适用于 `B3`；本次只用已列明坐标。'
        parsed=direct_answer(reply,engine.standard,['合成表'])
        periods={tuple(row['ranges']):row for row in parsed['dimensions'] if row['id']=='period'}
        self.assertEqual(periods[('合成表!B3',)]['evidence'],['合成表!B2:C2'])
        self.assertFalse(periods[('合成表!B4',)].get('evidence'))
        mappings=[]
        for year,carrier in ((2025,'A'),(2026,'B')):
            source=self.folder/(carrier+'合并标签.xlsx')
            book=load_workbook(self.source);ws=book.active;ws.merge_cells('B2:C2');ws['B2']=str(year)+'年末'
            book.save(source);book.close()
            snapshot=read_workbook(source,context=[])
            base={'recognition_method':'direct','source_path':str(source),'source_hash':snapshot['sha256'],
                  'source_context':[],'gold_path':str(self.gold),'gold_hash':engine.gold_hash,'carrier':carrier,
                  '金标准原文':'\n'.join(line for line in self.gold.read_text(encoding='utf-8').splitlines()
                      if json.loads(line)['record_type'] in {'metric','dimension'})}
            answer=json.loads(self.answer(year));table=answer['tables'][0]
            table['data_ranges']=['B3'];table['metrics']=[{**table['metrics'][0],'ranges':['B3']}]
            table['dimensions']=[*table['dimensions'][:4],table['dimensions'][4],table['dimensions'][6]]
            for row in table['dimensions']:row['ranges']=['B3']
            table['dimensions'][4]['evidence']=['合成表!B2','合成表!C2'] if carrier=='A' else []
            saved=save_direct_answer(base,json.dumps(answer,ensure_ascii=False),self.folder/(carrier+'合并合集.json'),snapshot=snapshot,standard=engine.standard)
            if carrier=='B':
                saved=save_direct_answer(saved,reply,self.folder/'B接入标签合集.json',field_scope={'合成表!B3':['period']})
            mappings.append(saved)
        before=deepcopy(mappings)
        plan=direct_update_plan(*mappings)
        self.assertEqual(len(plan['updates']),1)
        self.assertEqual([(row['cell'],row['value']) for row in plan['labels']],[('B2','2026年末')])
        self.assertEqual(mappings,before)

    def test_列范围分句和币种尾文仅接入明示值(self):
        from 附注更新.A合集 import save_direct_answer, _records
        from 附注更新.表格 import read_workbook, materialize_candidate_blanks
        engine=self.engine();snapshot=read_workbook(self.source,context=[]);materialize_candidate_blanks(snapshot)
        base={'recognition_method':'direct','source_path':str(self.source),'source_hash':snapshot['sha256'],
              'source_context':[],'gold_path':str(self.gold),'gold_hash':engine.gold_hash,'carrier':'A',
              '金标准原文':'\n'.join(line for line in self.gold.read_text(encoding='utf-8').splitlines()
                  if json.loads(line)['record_type'] in {'metric','dimension'})}
        old=save_direct_answer(base,self.answer(),self.folder/'分句原合集.json',snapshot=snapshot,standard=engine.standard)
        reply='''- 合成表!B3:C4：固定资产原值；asset_selection 为甲范围，C列为乙范围。行3期初余额、4期末余额。其他说明保留。
- 合成表!B5：asset_selection 为丙范围。行5变动。没有说明的内容未知。
- 合成表!B3:C4：currency 为CNY、倍率为1。'''
        scope={f'合成表!{c}{r}':['asset_selection','currency'] for c in 'BC' for r in (3,4)}
        scope['合成表!B5']=['asset_selection']
        saved=save_direct_answer(old,reply,self.folder/'分句新合集.json',field_scope=scope)
        rows=dict(_records(saved))
        for c,value in [('B','甲范围'),('C','乙范围')]:
            for r in (3,4):
                dims=rows[f'合成表!{c}{r}']['dimensions']
                self.assertEqual(dims['asset_selection']['predicate'],value)
                self.assertEqual(dims['currency'],'CNY')
        self.assertEqual(rows['合成表!B5']['dimensions']['asset_selection']['predicate'],'丙范围。行5变动。没有说明的内容未知')
        self.assertEqual(saved['机械记录'],old['机械记录'])
        self.assertEqual(set(rows),set(dict(_records(old))))
        self.assertEqual(saved['答复接入来源']['本次原答复'],reply)
        self.assertFalse(saved['unresolved'])
        self.assertEqual(self.hashes,(file_hash(self.source),file_hash(self.gold)))
        from 附注更新.A合集 import direct_answer, _direct_table
        parsed=direct_answer('合成表!B3:D3：asset_selection 为甲范围，C列为乙范围，D列为丙范围。',engine.standard,['合成表'])
        self.assertEqual({r['ranges'][0]:r['value'] for r in parsed['dimensions']},
                         {'合成表!B3':'甲范围','合成表!C3':'乙范围','合成表!D3':'丙范围'})
        uncertain='CNY、倍率为1，但本格币种尚未确认'
        table=_direct_table({'sheet':'合成表','dimensions':[{'id':'currency','ranges':['B3'],'value':uncertain}]},
                            {'source_context':[]},engine.standard)
        self.assertEqual(table['dimensions'][0]['value'],uncertain)

    def test_明确范围不重复展开整表且引用只撤对应机械未决(self):
        from 附注更新.A合集 import _inside, save_direct_answer
        from 附注更新.A语义阅读 import addresses
        from 附注更新.表格 import read_workbook, materialize_candidate_blanks
        with patch('附注更新.A语义阅读.addresses',wraps=addresses) as expand:
            self.assertTrue(_inside('B3',['B3'],{'data_ranges':['B3:C8']}))
            self.assertEqual([c.args[0] for c in expand.call_args_list],['B3'])
        self.assertTrue(_inside('C8',['*'],{'data_ranges':['B3:C8']}))
        engine=self.engine();snapshot=read_workbook(self.source,context=[]);materialize_candidate_blanks(snapshot)
        base={'recognition_method':'direct','source_path':str(self.source),'source_hash':snapshot['sha256'],
              'source_context':[],'gold_path':str(self.gold),'gold_hash':engine.gold_hash,'carrier':'A',
              '金标准原文':'\n'.join(line for line in self.gold.read_text(encoding='utf-8').splitlines()
                  if json.loads(line)['record_type'] in {'metric','dimension'})}
        statement='- `B3:C3` 本期计提：各阶段增加；`B4:C4` 本期转回：各阶段减少。'
        answer=json.loads(self.answer())
        answer['tables'][0]['unresolved']=[{'ranges':['B3:C4'],'reason':'这段答复未明确可机械读取的字段归属','original':statement},
                                           {'ranges':['B8'],'reason':'原文仍有未说明事项','original':'保留此项'}]
        old=save_direct_answer(base,json.dumps(answer,ensure_ascii=False),self.folder/'引用旧标记.json',snapshot=snapshot,standard=engine.standard)
        reply=statement+'\n- asset_selection 按所在行分别取：本期计提、本期转回。'
        saved=save_direct_answer(old,reply,self.folder/'引用已接通.json',field_scope={'合成表!B3:C4':['asset_selection']})
        self.assertEqual(saved['机械记录'],old['机械记录'])
        self.assertTrue(saved['unresolved'])
        self.assertTrue(all(r['cell']=='B8' and r['reason']=='原文仍有未说明事项' for r in saved['unresolved']))

    def test_顺序值沿原句小组及明确行号保存限定(self):
        from 附注更新.A合集 import direct_answer
        from 附注更新.A语义阅读 import addresses
        standard=self.engine().standard
        def values(reply):
            parsed=direct_answer(reply,standard,['合成表'])
            return {cell:r['value'] for r in parsed['dimensions'] if r['id']=='asset_selection'
                    for area in r['ranges'] if area!='*' for cell in addresses(area.rsplit('!',1)[-1])}
        self.assertEqual(values('B3/C3，甲；B4/C4，乙：固定资产原值。asset_selection 分别为“甲”“乙”。'),
                         {'B3':'甲','C3':'甲','B4':'乙','C4':'乙'})
        self.assertEqual(values('B3/C3，项目本期增加、减少；B4/C4，另一项目本期增加、减少：固定资产原值。asset_selection 分别为“增加”“减少”。'),
                         {'B3':'增加','C3':'减少','B4':'增加','C4':'减少'})
        self.assertEqual(values('B3:C6：行3甲、行4乙、行5丙、行6合计。固定资产原值；asset_selection 分别为甲、乙、丙；合计另列。行5的“丙”按原文说明为排除甲、乙后的丙。'),
                         {'B3':'甲','C3':'甲','B4':'乙','C4':'乙','B5':'排除甲、乙后的丙','C5':'排除甲、乙后的丙'})
        conflict=direct_answer('B3:C4：行3甲、行4乙。固定资产原值；asset_selection 分别为“甲”“乙”。行3的“乙”按原文说明为乙的明确限定。',standard,['合成表'])
        self.assertTrue(any(r['reason']=='原答复同一行标签与详细说明不一致' for r in conflict['unresolved']))
        limited=direct_answer('B3/C3，甲：仅限设备；B4/C4，乙：仅限车辆。固定资产原值。asset_selection 分别为“甲”“乙”。',standard,['合成表'])
        self.assertFalse(any(r['id']=='asset_selection' for r in limited['dimensions']))
        self.assertTrue(limited['unresolved'])

    def test_表格末尾角色及行内列日期不扩到未知列(self):
        from 附注更新.A合集 import direct_answer
        answer='''本表“期末”取 period={date:2024-12-31, kind:instant}，“期初”取 period={date:2024-01-01, kind:instant}。
| B原文坐标 | 金标准编号及实际范围 |
|---|---|
| 合成表!B3 | 固定资产原值；asset_selection=all。期末。 |
| 合成表!C3 | 固定资产原值；asset_selection=all。期初。 |
本期变动期间为2024-01-01至2024-12-31。
### 另一表
- 合成表!B4、C4：“期初余额”。B列按共同说明为2024-01-01；C列实际日期无法确定。
- 合成表!B5、C5：“本期增加”。B列为2024年度；C列期间日期未知。
- 合成表!B6、C6：“期末余额”。B列按共同说明为2024-12-31；C列实际日期无法确定。'''
        parsed=direct_answer(answer,self.engine().standard,['合成表'])
        values={(row['id'],tuple(row['ranges'])):row['value'] for row in parsed['dimensions']}
        for cell,role in [('B3','closing'),('C3','opening')]:
            self.assertEqual(values[('template_period_role',('合成表!'+cell,))],role)
        self.assertEqual(values[('period',('合成表!B4',))],'2024-01-01')
        self.assertEqual(values[('period',('合成表!B5',))],'2024-01-01至2024-12-31')
        self.assertEqual(values[('period',('合成表!B6',))],'2024-12-31')
        self.assertFalse(any(key[0]=='period' and any(cell in key[1] for cell in ('合成表!C4','合成表!C5','合成表!C6')) for key in values))

    def test_表头后缀冒号对象重复期间及明确未知保存回读(self):
        from 附注更新.A合集 import save_direct_answer, _records
        from 附注更新.表格 import read_workbook, materialize_candidate_blanks
        engine=self.engine();snapshot=read_workbook(self.source,context=[]);materialize_candidate_blanks(snapshot)
        base={'recognition_method':'direct','source_path':str(self.source),'source_hash':snapshot['sha256'],
              'source_context':[],'gold_path':str(self.gold),'gold_hash':engine.gold_hash,'carrier':'A',
              '金标准原文':'\n'.join(line for line in self.gold.read_text(encoding='utf-8').splitlines()
                  if json.loads(line)['record_type'] in {'metric','dimension'})}
        old=save_direct_answer(base,self.answer(),self.folder/'机械句式原合集.json',snapshot=snapshot,standard=engine.standard)
        reply='''| B坐标 | `asset_selection` 实际取值 |
|---|---|
| B3 | `domain: 固定资产, mode: all` |
- **合成表!B4 — period**：本格为期末余额，实际期间为 **2025-12-31（时点）**。
- **合成表!C4 — period**：原文没有对应期间标题，期间**未知**。
- **合成表!B5 — asset_selection**：原文没有对应项目，范围**未知**。
- 合成表!B6 的 asset_selection 均为：`{"domain":"固定资产","mode":"all"}`。
- 合成表!C6 的 asset_selection 为“固定资产”全体（all）。
- 合成表!B7 的 asset_selection 为固定资产总体（`mode: all`）。后面的说明不进入取值。
- 合成表!C7 的 asset_selection 为甲、乙、丙（`mode: members`）。
- 合成表!B8 的 currency 为原文未说明是CNY还是USD，留缺。
'''
        saved=save_direct_answer(old,reply,self.folder/'机械句式新合集.json',field_scope={
            '合成表!B3':['asset_selection'],'合成表!B4':['period'],'合成表!C4':['period'],'合成表!B5':['asset_selection'],
            **{f'合成表!{cell}':['asset_selection'] for cell in ('B6','C6','B7','C7')},'合成表!B8':['currency']})
        rows=dict(_records(saved))
        self.assertEqual(rows['合成表!B3']['dimensions']['asset_selection']['mode'],'all')
        self.assertEqual(rows['合成表!B4']['dimensions']['period'],{'kind':'instant','date':'2025-12-31'})
        self.assertNotIn('period',rows['合成表!C4']['dimensions'])
        self.assertEqual(rows['合成表!B5']['dimensions']['asset_selection']['mode'],'unknown')
        self.assertFalse(any(row.get('cell')=='B4' for row in saved['unresolved']))
        for cell in ('B6','C6','B7'):self.assertEqual(rows['合成表!'+cell]['dimensions']['asset_selection']['mode'],'all')
        self.assertEqual(rows['合成表!C7']['dimensions']['asset_selection']['members'],['甲','乙','丙'])
        self.assertNotIn('currency',rows['合成表!B8']['dimensions'])
        self.assertIn('currency',rows['合成表!B8']['missing_dimensions'])

    def test_明确members标题和配对行号仅按明示范围展开(self):
        from copy import deepcopy
        from 附注更新.A合集 import direct_answer
        standard=deepcopy(self.engine().standard)
        standard['dimension']['change_form']={'value_type':'text','allowed_values':['全部变动']}
        standard['dimension']['change_direction']={'value_type':'text','allowed_values':['增加','减少']}
        reply='''共同适用的选取维度均为 `asset_selection`，业务域为“固定资产”，模式为 `members`。
### 余额
以下各格的 `asset_selection.members` 取值为：
- `合成表!B3`、`C3`：项目甲；原文限定
### 变动
以下各格除 `asset_selection` 外，`change_form` 均为“全部变动”；增加列的 `change_direction` 为“增加”，减少列为“减少”。
- `合成表!B4`：增加；项目乙
- `合成表!C4`：减少；项目乙
'''
        parsed=direct_answer(reply,standard,['合成表'])
        selected=[row for row in parsed['dimensions'] if row['id']=='asset_selection']
        self.assertEqual([(row['ranges'],row['value']) for row in selected],[
            (['合成表!B3','合成表!C3'],{'domain':'固定资产','mode':'members','members':['项目甲；原文限定']}),
            (['合成表!B4'],{'domain':'固定资产','mode':'members','members':['项目乙']}),
            (['合成表!C4'],{'domain':'固定资产','mode':'members','members':['项目乙']})])
        self.assertEqual([(row['ranges'],row['value']) for row in parsed['dimensions'] if row['id']=='change_direction'],
                         [(['合成表!B4'],'增加'),(['合成表!C4'],'减少')])
        standard['metric']['metric.test.balance']=deepcopy(standard['metric']['固定资产原值'])
        text='分类表（合成表第3–4、5–6行）：B列为 `metric.test.balance`，C列为 `metric.test.balance`。3/5行对象为项目甲；4/6行是本表合计范围。'
        paired=direct_answer(text,standard,['合成表'])
        self.assertEqual([(row['ranges'],row['value']) for row in paired['dimensions'] if row['id']=='asset_selection'],
            [(['合成表!B3','合成表!C3','合成表!B5','合成表!C5'],'项目甲'),
             (['合成表!B4','合成表!C4','合成表!B6','合成表!C6'],'本表合计范围')])
        scoped=direct_answer('账龄表（合成表第7–8行）：B、C列为 `metric.test.balance`。'
                             '各格具体 `asset_selection` 为对应行的对象；合计行范围按本表。\n'
                             +text+'\n合成表!B8：period=2025-12-31。',standard,['合成表'])
        self.assertEqual([row for row in scoped['dimensions'] if row['id']=='asset_selection'],
                         [row for row in paired['dimensions'] if row['id']=='asset_selection'])
        self.assertEqual([row['ranges'] for row in scoped['dimensions'] if row['id']=='period'],[['合成表!B8']])
        with self.assertRaisesRegex(ValueError,'未明确可机械读取'):
            direct_answer(text.replace('B列为','为').replace('C列为','为'),standard,['合成表'])
        standard['metric']['metric.test.change']=deepcopy(standard['metric']['固定资产原值'])
        tail=direct_answer('合成表!B7、C7：asset_selection=甲，`metric.test.balance`。\n'
                           '合成表!B8：asset_selection=乙，`metric.test.balance`或`metric.test.change`。',standard,['合成表'])
        self.assertEqual([(row['ranges'],row['id']) for row in tail['metrics']],
                         [(['合成表!B7','合成表!C7'],'metric.test.balance')])

    def test_表格字段列直接连接实际取值不补期间标签坐标(self):
        from 附注更新.A合集 import direct_answer, _direct_table
        standard=self.engine().standard
        reply='''| 坐标 | 字段 | 原表期间标签 | 实际取值 |
|---|---|---|---|
| 合成表!B3 | period | 期末 | 2025-12-31 |
| 合成表!C3 | `period` | 期初 | 2024-12-31 |

依据：第370段对“期末”“期初”的日期定义；表格第2行的列标题对应关系。
'''
        parsed=direct_answer(reply,standard,['合成表'])
        table=_direct_table({**parsed,'sheet':'合成表','data_ranges':['B3:C3']},
                            {'source_context':[]},standard)
        self.assertEqual([(row['id'],row['ranges'],row['value'],row['evidence'])
                          for row in table['dimensions']], [
            ('period',['B3'],{'kind':'instant','date':'2025-12-31'},[]),
            ('period',['C3'],{'kind':'instant','date':'2024-12-31'},[])])


    def test_共同变动事项同句业务域不生成空字段并保留明确未知(self):
        from copy import deepcopy
        from 附注更新.A合集 import direct_answer, _direct_table
        standard=deepcopy(self.engine().standard)
        key='fixed_asset_change_selection'
        standard['dimension'][key]={**standard['dimension']['asset_selection'],
                                    'id':key,'domain':'固定资产变动事项'}
        reply='所列固定资产变动格中的变动事项范围均为该格相应方向下的全部变动事项（`fixed_asset_change_selection`：固定资产变动事项 `all`）。'
        parsed=direct_answer(reply,standard,['合成表'])
        selected=[row for row in parsed['dimensions'] if row['id']==key]
        self.assertTrue(selected)
        self.assertEqual({row['value'] for row in selected},{'all'})
        self.assertFalse(any(row.get('id')==key for row in parsed['unresolved']))
        partial='partial：已知属于固定资产，具体选取范围未确定'
        table=_direct_table({'sheet':'合成表','dimensions':[
            {'id':'report_scope','ranges':['B3'],'value':'单体'},
            {'id':'asset_selection','ranges':['B3'],'value':partial}]},
            {'source_context':[]},standard)
        values={row['id']:row['value'] for row in table['dimensions']}
        self.assertEqual(values['report_scope'],'standalone')
        self.assertEqual(values['asset_selection'],
                         {'mode':'unknown','known_scope':partial,'domain':'固定资产'})

    def test_正文列方向仅展开到同段已声明坐标(self):
        from copy import deepcopy
        from 附注更新.A合集 import direct_answer
        from 附注更新.A语义阅读 import addresses
        standard=deepcopy(self.engine().standard)
        key='fixed_asset_change_direction'
        standard['dimension'][key]={'id':key,'domain':'固定资产余额变动方向',
                                    'value_type':'text','allowed_values':['增加','减少','净变动']}
        standard['metric']['metric.fixed_asset.cost.change']={
            **standard['metric']['固定资产原值'],
            'dimension_refs':[{'id':key,'required':True}]}
        reply='''以下坐标的变动格C列为本期增加、D列为本期减少。
word附注!C401:D402：metric.fixed_asset.cost.change。
'''
        parsed=direct_answer(reply,standard,['word附注','另一表'])
        values={(area.rsplit('!',1)[0]+'!'+cell,row['value'])
                for row in parsed['dimensions'] if row['id']==key
                for area in row['ranges'] for cell in addresses(area.rsplit('!',1)[-1])}
        self.assertEqual(values,{('word附注!C401','增加'),('word附注!C402','增加'),
                                 ('word附注!D401','减少'),('word附注!D402','减少')})

    def test_共同期间与列角色保留明确工作表(self):
        from 附注更新.A合集 import direct_answer
        from 附注更新.A语义阅读 import addresses
        reply='''期间说明：word附注!A1写明期末为2024年12月31日、期初为2024年1月1日；这些日期适用于下列相关余额。
以下表格坐标均在`word附注`工作表。
对于`word附注!B400:E400`，列角色为B“期初余额”、C“本期增加”、D“本期减少”、E“期末余额”；各余额的期间按上述共同期间说明。
'''
        parsed=direct_answer(reply,self.engine().standard,['word附注','另一表'])
        values={(area.rsplit('!',1)[0]+'!'+cell,row['value'])
                for row in parsed['dimensions'] if row['id']=='period'
                for area in row['ranges'] for cell in addresses(area.rsplit('!',1)[-1])}
        self.assertEqual(values,{('word附注!B400','2024-01-01'),('word附注!E400','2024-12-31')})


    def test_顺序角色尾句不吞取值且明确期间对象去除引号(self):
        from 附注更新.A合集 import direct_answer, _direct_table
        from 附注更新.A语义阅读 import addresses
        standard=self.engine().standard
        reply='''期间说明：word附注!A1写明期末为2024年12月31日、期初为2024年1月1日；这些日期适用于下列相关余额。
word附注!B415:E419：列角色依次为期初余额、本期增加、本期减少、期末余额，期间按共同期间说明。
'''
        for text,complete in [(reply,True),(reply.replace('本期减少、',''),False)]:
            with self.subTest(角色完整=complete):
                parsed=direct_answer(text,standard,['word附注','另一表'])
                values={(row['id'],area.rsplit('!',1)[0]+'!'+cell,row['value'])
                        for row in parsed['dimensions'] if row['id'] in ('period','template_period_role')
                        for area in row['ranges'] for cell in addresses(area.rsplit('!',1)[-1])}
                expected={(key,f'word附注!{column}{row}',value)
                          for row in range(415,420)
                          for column,date,role in [('B','2024-01-01','opening'),('E','2024-12-31','closing')]
                          for key,value in [('period',date),('template_period_role',role)]} if complete else set()
                self.assertEqual(values,expected)
                self.assertFalse(any(location.rsplit('!',1)[-1].startswith(('C','D'))
                                     for key,location,value in values if key=='period'))
        parsed=direct_answer('word附注!C415：period={kind:"duration", start:"2024-01-01", end:"2024-12-31"}。',
                             standard,['word附注','另一表'])
        table=_direct_table({**parsed,'sheet':'word附注','data_ranges':['C415']},
                            {'source_context':[]},standard)
        self.assertEqual([(row['ranges'],row['value']) for row in table['dimensions'] if row['id']=='period'],
                         [(['C415'],{'kind':'duration','start':'2024-01-01','end':'2024-12-31'})])


    def test_纯期间标签关系仅给已有期间添加引用(self):
        from 附注更新.A合集 import save_direct_answer, _records
        from 附注更新.表格 import read_workbook, materialize_candidate_blanks
        engine=self.engine();snapshot=read_workbook(self.source,context=[]);materialize_candidate_blanks(snapshot)
        reply='期间标签坐标为 合成表!B2:C2，适用于 合成表!B3:B4。'
        for carrier in ('A','B'):
            for has_period in (True,False):
                with self.subTest(合集=carrier,原有期间=has_period):
                    base={'recognition_method':'direct','source_path':str(self.source),'source_hash':snapshot['sha256'],
                          'source_context':[],'gold_path':str(self.gold),'gold_hash':engine.gold_hash,'carrier':carrier,
                          '金标准原文':'\n'.join(line for line in self.gold.read_text(encoding='utf-8').splitlines()
                              if json.loads(line)['record_type'] in {'metric','dimension'})}
                    answer=json.loads(self.answer());table=answer['tables'][0]
                    if has_period:
                        table['dimensions'][4]['evidence']=['合成表!A2']
                    else:
                        table['dimensions']=[row for row in table['dimensions'] if row['id']!='period']
                    suffix=carrier+str(has_period)
                    old=save_direct_answer(base,json.dumps(answer,ensure_ascii=False),self.folder/(suffix+'标签原合集.json'),
                                          snapshot=snapshot,standard=engine.standard)
                    saved=save_direct_answer(old,reply,self.folder/(suffix+'标签新合集.json'),
                                            field_scope={f'合成表!B{row}':['period'] for row in (3,4)})
                    before=dict(_records(old));rows=dict(_records(saved))
                    self.assertEqual(set(rows),set(before))
                    for location,row in rows.items():
                        if location not in ('合成表!B3','合成表!B4'):
                            self.assertEqual({k:v for k,v in row.items() if k!='recognition_provenance'},
                                             {k:v for k,v in before[location].items() if k!='recognition_provenance'})
                        elif has_period:
                            self.assertEqual(row['dimensions'],before[location]['dimensions'])
                            self.assertEqual(row['dimensions']['period'],{'kind':'instant','date':'2025-12-31'})
                            self.assertEqual(row['dimension_evidence']['period'],['合成表!B2','合成表!C2'])
                        else:
                            self.assertNotIn('period',row['dimensions'])
        self.assertEqual(self.hashes,(file_hash(self.source),file_hash(self.gold)))

    def test_共用期间对象的列角色严格限定工作表(self):
        from 附注更新.A合集 import direct_answer
        from 附注更新.A语义阅读 import addresses
        reply='''本期 period={kind:"duration", start:"2024-01-01", end:"2024-12-31"}；上期 period={kind:"duration", start:"2023-01-01", end:"2023-12-31"}。
Excel 附注汇总表的 C 列为本期、D 列为上期；Word 附注表的 B 列为本期、C 列为上期。
word附注!B980:C980：固定资产原值。
EXCEL附注汇总（单体）!C1438:D1438：固定资产原值。
'''
        parsed=direct_answer(reply,self.engine().standard,['EXCEL附注汇总（单体）','word附注'])
        values={}
        for row in parsed['dimensions']:
            if row['id'] not in ('period','template_period_role'):continue
            for area in row['ranges']:
                for cell in addresses(area.rsplit('!',1)[-1]):
                    values.setdefault((area.rsplit('!',1)[0]+'!'+cell,row['id']),[]).append(row['value'])
        self.assertEqual(values,{
            ('word附注!B980','period'):[{'kind':'duration','start':'2024-01-01','end':'2024-12-31'}],
            ('word附注!B980','template_period_role'):['current'],
            ('word附注!C980','period'):[{'kind':'duration','start':'2023-01-01','end':'2023-12-31'}],
            ('word附注!C980','template_period_role'):['prior']})


    def test_所在附注定位引用不覆盖原范围且同为只接声明字段(self):
        from 附注更新.A合集 import direct_answer
        from 附注更新.A语义阅读 import addresses
        reply='''合成表!B3：固定资产原值；asset_selection=其他；currency=CNY。
合成表!B3 所在附注的固定资产合计为 合成表!B4；asset_selection=all。Word对应为 word附注!B5，同为all。
'''
        parsed=direct_answer(reply,self.engine().standard,['合成表','word附注'])
        values={}
        for row in parsed['dimensions']:
            for area in row['ranges']:
                for cell in addresses(area.rsplit('!',1)[-1]):
                    values.setdefault((area.rsplit('!',1)[0]+'!'+cell,row['id']),[]).append(row['value'])
        self.assertEqual(values,{('合成表!B3','asset_selection'):['其他'],
                                 ('合成表!B4','asset_selection'):['all'],
                                 ('word附注!B5','asset_selection'):['all'],
                                 ('合成表!B3','currency'):['CNY']})


    def test_斜杠多角色仅停止对应格(self):
        from 附注更新.A合集 import save_direct_answer
        from 附注更新.表格 import read_workbook
        from 附注更新.模板更新编排 import direct_update_plan
        engine=self.engine();snapshot=read_workbook(self.source,context=[])
        base={'recognition_method':'direct','source_path':str(self.source),'source_hash':snapshot['sha256'],
              'source_context':[],'gold_path':str(self.gold),'gold_hash':engine.gold_hash,
              '金标准原文':'\n'.join(line for line in self.gold.read_text(encoding='utf-8').splitlines() if json.loads(line)['record_type'] in {'metric','dimension'})}
        for index,mixed in enumerate(('期末/期初',{'B3':'期末／期初'})):
            mappings=[]
            for year,carrier in ((2025,'A'),(2026,'B')):
                answer=json.loads(self.answer(year));table=answer['tables'][0]
                table['data_ranges']=['B3:B4']
                table['metrics']=[{'id':'固定资产原值','ranges':['B3'],'template_period_role':mixed},
                                  {'id':'固定资产原值','ranges':['B4'],'template_period_role':'期末余额'}]
                table['dimensions']=[row for row in table['dimensions'] if row['id']!='period']
                for cell,label in (('B3','B2'),('B4','C2')):
                    table['dimensions'].append({'id':'period','ranges':[cell],'value':{'kind':'instant','date':f'{year}-12-31'},
                                                'evidence':['合成表!'+label]})
                mappings.append(save_direct_answer({**base,'carrier':carrier},json.dumps(answer,ensure_ascii=False),
                    self.folder/f'斜杠角色{index}{carrier}.json',snapshot=snapshot,standard=engine.standard))
            plan=direct_update_plan(*mappings)
            self.assertEqual([row['cell'] for row in plan['updates']],['B4'])
            self.assertTrue(any(row.get('cell')=='B3' for row in plan['issues']))

    def test_共同成员标题接入且不按缺失字段猜成员(self):
        from 附注更新.A合集 import save_direct_answer, _records
        from 附注更新.表格 import read_workbook
        engine=self.engine();snapshot=read_workbook(self.source,context=[])
        base={'recognition_method':'direct','source_path':str(self.source),'source_hash':snapshot['sha256'],
              'source_context':[],'gold_path':str(self.gold),'gold_hash':engine.gold_hash,'carrier':'A',
              '金标准原文':'\n'.join(line for line in self.gold.read_text(encoding='utf-8').splitlines() if json.loads(line)['record_type'] in {'metric','dimension'})}
        old=save_direct_answer(base,self.answer(),self.folder/'共同成员原合集.json',snapshot=snapshot,standard=engine.standard)
        for index,mode in enumerate(('', '，选取方式均为 `members`')):
            reply=f'以下均为 `asset_selection`{mode}。\n| 坐标 | 实际取值 |\n| --- | --- |\n| B3 | 父项；成员甲 |\n| B4 | 父项；成员乙 |'
            saved=save_direct_answer(old,reply,self.folder/f'共同成员{index}.json',field_scope={'合成表!B3:B4':['asset_selection']})
            rows=dict(_records(saved))
            for cell,value in (('B3','父项；成员甲'),('B4','父项；成员乙')):
                selection=rows['合成表!'+cell]['dimensions']['asset_selection']
                if mode:
                    self.assertEqual(selection['mode'],'members')
                    self.assertEqual(selection['members'],[value])
                else:
                    self.assertNotEqual(selection['mode'],'members')
                    self.assertEqual(selection.get('predicate'),value)


    def test_单个开放文本字段依请求保存原话不猜归属(self):
        from 附注更新.A合集 import direct_answer
        standard=self.engine().standard
        standard['dimension']['test_basis']={'id':'test_basis','value_type':'text','domain':'测试口径'}
        reply='- **合成表!B3:C3（原表注解）**：按本表列示口径；已实施事项甲，未实施事项乙。'
        scope={'合成表!B3:C3':['test_basis']}
        result=direct_answer(reply,standard,['合成表'],field_scope=scope)
        rows=[r for r in result['dimensions'] if r['id']=='test_basis']
        self.assertEqual(len(rows),1)
        self.assertEqual(rows[0]['ranges'],['合成表!B3','合成表!C3'])
        self.assertEqual(rows[0]['value'],'按本表列示口径；已实施事项甲，未实施事项乙。')
        self.assertEqual(rows[0]['original'],reply)
        for mask in (None,{'合成表!B3:C3':['test_basis','entity']},{'合成表!B3':['test_basis']}):
            parsed=direct_answer(reply,standard,['合成表'],field_scope=mask)
            self.assertFalse(any(r['id']=='test_basis' for r in parsed['dimensions']))

    def test_集合方式保留其显式取值及未知(self):
        from 附注更新.A合集 import save_direct_answer,_records
        from 附注更新.表格 import read_workbook
        engine=self.engine();snapshot=read_workbook(self.source,context=[])
        base={'recognition_method':'direct','source_path':str(self.source),'source_hash':snapshot['sha256'],
              'source_context':[],'gold_path':str(self.gold),'gold_hash':engine.gold_hash,'carrier':'A',
              '金标准原文':'\n'.join(line for line in self.gold.read_text(encoding='utf-8').splitlines() if json.loads(line)['record_type'] in {'metric','dimension'})}
        old=save_direct_answer(base,self.answer(),self.folder/'显式方式原合集.json',snapshot=snapshot,standard=engine.standard)
        cases=[('`members`：甲、乙（原文列明的并列类别）。',{'mode':'members','members':['甲','乙']}),
               ('members「甲及乙」（原文合列，保留合并范围）。',{'mode':'members','members':['甲及乙']}),
               ('`predicate`：仅所列项目；不含其他项目。',{'mode':'predicate','predicate':'仅所列项目；不含其他项目'}),
               ('未知。原文没有明确范围。',None)]
        for i,(text,expected) in enumerate(cases):
            saved=save_direct_answer(old,'合成表!B3：asset_selection＝'+text,
                self.folder/f'显式方式{i}.json',field_scope={'合成表!B3':['asset_selection']})
            value=dict(_records(saved))['合成表!B3']['dimensions'].get('asset_selection')
            if expected is None:self.assertEqual(value.get('mode'),'unknown')
            else:self.assertEqual({k:v for k,v in value.items() if k!='domain'},expected)

    def test_列名及编号两段分别只按唯一原文字尾连接(self):
        from 附注更新.A合集 import direct_answer
        standard=self.engine().standard
        for name in ('metric.x','metric.y','metric.z'):standard['metric'][name]={'id':name,'dimension_refs':[]}
        raw='合成表!B3:F4：B、D、F列分别为甲余额、乙准备、丙价值。余额、准备、价值分别选 `metric.x`、`metric.y`、`metric.z`。'
        result=direct_answer(raw,standard,['合成表'])
        got={cell:row['id'] for row in result['metrics'] for cell in row['ranges']}
        self.assertEqual(got,{f'合成表!{col}{row}':metric for col,metric in zip('BDF',('metric.x','metric.y','metric.z')) for row in (3,4)})
        ambiguous=direct_answer(raw.replace('乙准备','乙余额'),standard,['合成表'])
        self.assertFalse(any(row['id']=='metric.x' for row in ambiguous['metrics']))

    def test_同段列名说明不遮住另句明确期间(self):
        from 附注更新.A合集 import direct_answer
        raw='期末余额的 period 为 {date: 2024-12-31, kind: instant}；期初余额的 period 为 {date: 2024-01-01, kind: instant}。\n合成表!B3:F4、B5:F6：B列为金额。B3:F4为期末，B5:F6为期初。'
        result=direct_answer(raw,self.engine().standard,['合成表'])
        dated={tuple(row['ranges']):row['value'] for row in result['dimensions'] if row['id']=='period'}
        self.assertEqual(dated[('合成表!B3:F4',)],{'date':'2024-12-31','kind':'instant'})
        self.assertEqual(dated[('合成表!B5:F6',)],{'date':'2024-01-01','kind':'instant'})

    def test_字段标题承接显式组成和限定行号(self):
        from 附注更新.A合集 import direct_answer
        for heading in ('**asset_selection**','### asset_selection','- **asset_selection**'):
            raw=heading+'\n- `合成表!B3:C3`：`mode=predicate`，`predicate=范围甲`。\n- **合成表!B4:C5**：按行范围为：\n  - 4、5：`{"mode":"members","members":["甲","乙"]}`'
            rows=direct_answer(raw,self.engine().standard,['合成表'])['dimensions']
            self.assertTrue(any(row['id']=='asset_selection' and row['value']=={'mode':'predicate','predicate':'范围甲'} for row in rows))
            selected=[row for row in rows if row['id']=='asset_selection' and row['value']=={'mode':'members','members':['甲','乙']}]
            self.assertEqual(len(selected),1)
            self.assertEqual(sum(row['id']=='asset_selection' for row in rows),2)
            self.assertEqual(set(selected[0]['ranges']),{'合成表!B4','合成表!C4','合成表!B5','合成表!C5'})

    def test_分子分母依金标准关系承接同行原值(self):
        from copy import deepcopy
        from 附注更新.A合集 import direct_answer
        standard=deepcopy(self.engine().standard)
        standard['dimension']['other_selection']={**standard['dimension']['asset_selection'],'id':'other_selection'}
        standard['metric']['比例']={'definition':{'ratio':{part:{'dimension_bindings':{'x':{'op':'dimension','key':key}}}
            for part,key in [('numerator','asset_selection'),('denominator','other_selection')]}},'dimension_refs':[]}
        raw='### asset_selection\n- 合成表!B3:B4：按行范围为：\n  - 3、4：{"mode":"predicate","predicate":"原范围甲"}\n- 合成表!C3:C4：分子分别同上3–4；分母均为 {"mode":"predicate","predicate":"原总体"}。\n- 合成表!C5：分子、分母均为原范围乙。'
        rows=direct_answer(raw,standard,['合成表'],{f'合成表!C{r}':'比例' for r in (3,4,5)})['dimensions']
        self.assertEqual([(r['ranges'],r['value']) for r in rows if r['id']=='asset_selection' and r['ranges']==['合成表!C3']],
                         [(['合成表!C3'],{'mode':'predicate','predicate':'原范围甲'})])
        self.assertTrue(any(r['id']=='other_selection' and r['value']=={'mode':'predicate','predicate':'原总体'} for r in rows))
        self.assertEqual({r['id'] for r in rows if r['value']=='原范围乙'},{'asset_selection','other_selection'})

    def test_回读保留原文及证据而不再套历史行(self):
        from 附注更新.A合集 import direct_answer
        standard=self.engine().standard
        raw={'metrics':[{'id':'固定资产原值','ranges':['B3'],'template_period_role':'closing','original':'最初槽位答复'}],
             'dimensions':[{'id':'asset_selection','ranges':['B3'],'value':{'mode':'members','members':['甲']},
                            'evidence':['合成表!A3'],'original':{'original':'最初范围答复','field_reference':'本句引用'}}]}
        once=direct_answer(raw,standard);twice=direct_answer(once,standard)
        self.assertEqual(once,twice)
        self.assertEqual(twice['dimensions'][0]['original'],raw['dimensions'][0]['original'])
        self.assertEqual(twice['dimensions'][0]['evidence'],['合成表!A3'])
        self.assertEqual(twice['metrics'][0]['template_period_role'],'closing')

    def test_多指标唯一未命名范围按本句行号分配(self):
        from copy import deepcopy
        from 附注更新.A合集 import direct_answer
        standard=deepcopy(self.engine().standard)
        standard['dimension']['asset_change_selection']={**standard['dimension']['asset_selection'],'id':'asset_change_selection'}
        for key in ('metric.one','metric.two'):
            standard['metric'][key]={'dimension_refs':[{'id':'asset_selection'},{'id':'asset_change_selection'}]}
        raw='合成表!B3、B4：metric.one；B5：metric.two。asset_change_selection=all；资产范围分别为：行3、5选all；行4选members「甲及乙」。'
        rows=direct_answer(raw,standard,['合成表'])['dimensions']
        selected=[r for r in rows if r['id']=='asset_selection']
        self.assertTrue(any(set(r['ranges'])=={'合成表!B3','合成表!B5'} and r['value']=='all' for r in selected))
        self.assertTrue(any(r['ranges']==['合成表!B4'] and r['value']=='members「甲及乙」' for r in selected))
        raw='合成表!B3/C3，第一组；B4/C4，第二组；B5/C5，第三组：metric.one。asset_change_selection=all；对象范围依次为范围甲、范围乙、以及甲乙并集。'
        rows=direct_answer(raw,standard,['合成表'])['dimensions']
        selected=[r for r in rows if r['id']=='asset_selection']
        self.assertEqual([(r['ranges'],r['value']) for r in selected],
                         [(['合成表!B3','合成表!C3'],'范围甲'),(['合成表!B4','合成表!C4'],'范围乙'),(['合成表!B5','合成表!C5'],'甲乙并集')])

    def test_标签句保留明确期间角色不截标题前缀(self):
        from 附注更新.A合集 import direct_answer
        raw=('原表 B2“本期期末余额”适用于 B3；角色为期末。\n'
             '原表 C2 标签“上年年末账面余额”，适用于 C3；角色为上期（上年年末）。')
        rows=direct_answer(raw,self.engine().standard,['合成表'])['dimensions']
        roles={loc:r['value'] for r in rows if r['id']=='template_period_role' for loc in r['ranges']}
        refs={loc:r['evidence'] for r in rows if r['id']=='period' for loc in r['ranges']}
        self.assertEqual(roles,{'合成表!B3':'closing','合成表!C3':'prior'})
        self.assertEqual(refs,{'合成表!B3':['合成表!B2'],'合成表!C3':['合成表!C2']})

    def test_数据格适用标签及分别引用不混合(self):
        from 附注更新.A合集 import direct_answer
        raw=('合成表!B3、合成表!B4（均为空白数据格）适用标签 合成表!B2「本期发生额」。\n'
             '合成表!B5、合成表!C5（均为空白数据格）分别适用标签 合成表!B2「本期增加」和 合成表!C2「本期减少」。')
        rows=direct_answer(raw,self.engine().standard,['合成表'])['dimensions']
        refs={loc:r['evidence'] for r in rows if r['id']=='period' for loc in r['ranges']}
        self.assertEqual(refs,{'合成表!B3':['合成表!B2'],'合成表!B4':['合成表!B2'],
                               '合成表!B5':['合成表!B2'],'合成表!C5':['合成表!C2']})

    def test_内联范围组件只取声明值且保留值内限定(self):
        from 附注更新.A合集 import direct_answer
        for body,expected in [
                ('`mode=predicate`，`predicate=信用风险组合款项合计`；不扩为全部款项。','信用风险组合款项合计'),
                ('`mode=predicate`，`predicate=甲类款项；不含乙类`。','甲类款项；不含乙类')]:
            raw='合成表!B3：asset_selection='+body
            rows=direct_answer(raw,self.engine().standard,['合成表'])['dimensions']
            value=next(r['value'] for r in rows if r['id']=='asset_selection')
            self.assertEqual(value,{'mode':'predicate','predicate':expected})

    def test_字段标题与末列坐标承接声明原值(self):
        from 附注更新.A合集 import direct_answer,_direct_table
        raw=('**合成表** 的 `asset_selection`（均为 predicate 模式）：\n'
             '| 坐标 | 实际取值 |\n|---|---|\n| B3:C3 | `甲。`（原话保留） |\n'
             '**asset_selection**\n| 原文行 | 实际范围（mode=predicate） | 适用坐标 |\n'
             '|---|---|---|\n| 4行 | 乙。 | B4；C4 |\n'
             '- `asset_selection`：\n  - `B5；C5`：predicate「丙」。\n'
             '- `asset_selection`：`B6:C6` 均为 predicate「丁」，说明仍留原文。')
        answer=direct_answer(raw,self.engine().standard,['合成表'])
        table=_direct_table({**answer,'sheet':'合成表','data_ranges':['B3:C6']},
                            {'source_context':[]},self.engine().standard)
        from 附注更新.A语义阅读 import addresses
        got={cell:r['value']['predicate'] for r in table['dimensions'] if r['id']=='asset_selection'
             for area in r['ranges'] for cell in addresses(area)}
        self.assertEqual(got,{f'{col}{row}':value for row,value in [(3,'甲。'),(4,'乙。'),(5,'丙'),(6,'丁')]
                              for col in 'BC'})

    def test_范围模式与其后原值不互相吞并(self):
        from 附注更新.A合集 import direct_answer,_direct_table
        raw=('以下均为 `asset_selection`。\n'
             '- 合成表!B3：predicate＝甲类 且 条件乙。\n'
             '合成表!C3 — `asset_selection`：`predicate`，类别为“甲”且阶段为“乙”。原文没有其他条件。')
        answer=direct_answer(raw,self.engine().standard,['合成表'])
        table=_direct_table({**answer,'sheet':'合成表','data_ranges':['B3:C3']},
                            {'source_context':[]},self.engine().standard)
        got={loc:r['value'] for r in table['dimensions'] if r['id']=='asset_selection' for loc in r['ranges']}
        self.assertEqual(got,{'B3':{'mode':'predicate','predicate':'甲类 且 条件乙','domain':'固定资产'},
                              'C3':{'mode':'predicate','predicate':'类别为“甲”且阶段为“乙”','domain':'固定资产'}})

    def test_自然表格坐标在取值列仍保留声明内容(self):
        from 附注更新.A合集 import direct_answer,_direct_table
        raw=('**`asset_selection`**\n| 坐标 | 实际取值 |\n|---|---|\n'
             '| 合成表!B3、C3 | predicate：甲类 且 条件乙 |\n'
             '| 行 | `asset_selection` 实际取值 |\n|---|---|\n'
             '| 4 | B4；C4的分子：乙类 且 条件丙。 |\n'
             '| `asset_selection` 实际取值 |\n|---|\n| B5:C5：均为丙类总体；不含丁类。 |')
        answer=direct_answer(raw,self.engine().standard,['合成表'])
        table=_direct_table({**answer,'sheet':'合成表','data_ranges':['B3:C5']},
                            {'source_context':[]},self.engine().standard)
        from 附注更新.A语义阅读 import addresses
        got={cell:r['value']['predicate'] for r in table['dimensions'] if r['id']=='asset_selection'
             for area in r['ranges'] for cell in addresses(area)}
        self.assertEqual(got,{f'{col}{row}':value for row,value in [(3,'甲类 且 条件乙'),(4,'乙类 且 条件丙。'),(5,'丙类总体；不含丁类。')]
                              for col in 'BC'})

    def test_同上及指定同行引用只搬运明确原值(self):
        from 附注更新.A合集 import direct_answer,_direct_table
        raw=('**B候选：合成表!B3:C3。**\n'
             '- 两格的 `asset_selection` 均为 **predicate**：甲类且条件乙。\n'
             '| 坐标 | 字段及实际取值 |\n|---|---|\n'
             '| B4 | **asset_selection**：predicate——乙类，条件丙。 |\n'
             '| C4 | **asset_selection**：同上。 |\n'
             '- B5：asset_selection为predicate——丙类总体。\n'
             '- C5：同一总体范围；period为2025-12-31。\n'
             '- D4:D5：asset_selection分别与同一行B列的范围一致。')
        answer=direct_answer(raw,self.engine().standard,['合成表'])
        table=_direct_table({**answer,'sheet':'合成表','data_ranges':['B3:D5']},
                            {'source_context':[]},self.engine().standard)
        from 附注更新.A语义阅读 import addresses
        got={cell:r['value']['predicate'] for r in table['dimensions'] if r['id']=='asset_selection'
             for area in r['ranges'] for cell in addresses(area)}
        self.assertEqual(got,{f'{col}{row}':value for row,cols,value in
                             [(3,'BC','甲类且条件乙'),(4,'BCD','乙类，条件丙'),(5,'BCD','丙类总体')] for col in cols})

    def test_比较说明不继承末行且并列字段共享原值(self):
        from copy import deepcopy
        from 附注更新.A合集 import direct_answer
        standard=deepcopy(self.engine().standard)
        standard['dimension']['other_selection']={**standard['dimension']['asset_selection'],'id':'other_selection'}
        raw=('合成表!B3:C3：template_period_role=opening。\n'
             'B的template_period_role为closing，而非A对应格的opening。\n'
             '合成表!B4:C4：分子 asset_selection 与分母 other_selection 均为范围甲；B4期间为2025-12-31，C4期间为2024-12-31。')
        rows=direct_answer(raw,standard,['合成表'])['dimensions']
        self.assertEqual([r['value'] for r in rows if r['id']=='template_period_role'],['opening'])
        self.assertEqual({r['id']:r['value'] for r in rows if r['id'].endswith('_selection')},
                         {'asset_selection':'范围甲','other_selection':'范围甲'})

    def test_原答复行号余额角色接已有日期(self):
        from 附注更新.A合集 import direct_answer, _direct_table
        table={'sheet':'合成表','data_ranges':['B3:C3'],'dimensions':[],
               'metrics':[{'id':'固定资产原值','ranges':['B3:C3'],'template_period_role':{'B':'本期增加','C':'本期减少'}}]}
        normalized=_direct_table(table,{'source_context':[]},self.engine().standard)
        self.assertEqual(normalized['metrics'][0]['template_period_role'],{'B':'current','C':'current'})
        self.assertEqual(table['metrics'][0]['template_period_role'],{'B':'本期增加','C':'本期减少'})
        table['metrics'][0]['template_period_role']='本期变动事项（2025年1月1日至12月31日）；方向未知仍保留未知'
        self.assertEqual(_direct_table(table,{'source_context':[]},self.engine().standard)['metrics'][0]['template_period_role'],'current')
        rows=direct_answer('- B3、C3（现金的期末余额）：metric_id=固定资产原值。',self.engine().standard,['合成表'])['dimensions']
        self.assertTrue(any(r['id']=='template_period_role' and r['value']=='closing' and r['ranges']==['合成表!B3','合成表!C3'] for r in rows))
        rows=direct_answer('合成表!B3：template_period_role=上期发生额，实际期间为 2024-01-01 至 2024-12-31。',self.engine().standard,['合成表'])['dimensions']
        self.assertEqual({r['id']:r['value'] for r in rows},{'template_period_role':'上期发生额','period':'2024-01-01至2024-12-31'})
        rows=direct_answer('合成表!B3：period=2024-01-01至2024-12-31（本期减少）。',self.engine().standard,['合成表'])['dimensions']
        self.assertTrue(any(r['id']=='template_period_role' and r['value']=='current' for r in rows))
        raw='期末余额的 period 为 {date: 2024-12-31, kind: instant}；期初余额的 period 为 {date: 2024-01-01, kind: instant}。\n合成表!B3:C5：行3期初余额、4变动、5期末余额。'
        rows=direct_answer(raw,self.engine().standard,['合成表'])['dimensions']
        dates={loc:row['value'] for row in rows if row['id']=='period' for loc in row['ranges']}
        self.assertEqual(dates['合成表!C3'],{'date':'2024-01-01','kind':'instant'})
        self.assertEqual(dates['合成表!C5'],{'date':'2024-12-31','kind':'instant'})
        self.assertNotIn('合成表!C4',dates)


if __name__ == '__main__':
    unittest.main()
