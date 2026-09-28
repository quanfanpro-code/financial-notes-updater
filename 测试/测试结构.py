"""结构规划边界使用受控回执，Excel插入、保存、命名范围及重读使用真实工作簿。"""
import copy
import pathlib
import unittest
import uuid

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Font, Side
from openpyxl.workbook.defined_name import DefinedName
from 附注更新.结构 import resolve_structure
import 附注更新.结构 as structure
from 附注更新.表格 import read_workbook, build_update_plan, write_updated_workbook

ROOT = pathlib.Path(__file__).resolve().parents[1]


def mapped(cell, obj, slot="C-N001-T001-S001", label="期末余额"):
    return {"sheet": "附注", "cell": cell, "slot_id": slot, "scope": "consolidated", "dimensions": {"counterparty": obj, "currency": "CNY", "unit": "元", "scale": 1, "period": "2025-12-31"}, "semantic_field": label, "value_type": "monetary", "row_label": obj, "column_label": label, "table_range": "A1:B4", "table_semantic": "应收账款客户明细", "aggregation": "total" if obj == "合计" else "component", "reviewed": True}


class ReplyEngine:
    def __init__(self, first, second=None):
        self.first, self.second = first, second or first
        self.calls = 0
        self.log = lambda text: None

    def _check_cancel(self):
        pass

    def _request(self, instruction, payload, validator):
        self.calls += 1
        return validator(copy.deepcopy(self.first if self.calls == 1 else self.second))


class StructureTests(unittest.TestCase):
    def setUp(self):
        self.work = ROOT / "测试输出" / ("结构_" + uuid.uuid4().hex)
        self.work.mkdir(parents=True)
        self.a_path = self.work / "A.xlsx"
        self.b_path = self.work / "B.xlsx"
        self.write_book(self.a_path, [["客户", "期末余额"], ["甲", 10], ["乙", 20], ["合计", 30]], "A1:B4")
        self.write_book(self.b_path, [["客户", "期末余额"], ["甲", 10], ["乙", 20], ["丙", 30], ["合计", 60]], "A1:B5")
        self.a, self.b = read_workbook(str(self.a_path)), read_workbook(str(self.b_path))
        self.ar = {"complete": True, "mappings": [mapped("B2", "甲"), mapped("B3", "乙"), mapped("B4", "合计", "C-N001-T001-S002")]}
        self.br = {"complete": True, "mappings": [mapped("B2", "甲"), mapped("B3", "乙"), mapped("B4", "丙"), mapped("B5", "合计", "C-N001-T001-S002")]}
        self.proposal = {"operations": [{"range_name": "表格一", "axis": "row", "action": "insert", "index": 3, "count": 1, "template_index": 2, "reason": "新增客户丙，合计仍保留"}], "bindings": [{"source_sheet": "附注", "source_cell": "B4", "target_sheet": "附注", "target_cell": "B4", "template_sheet": "附注", "template_cell": "B3", "reason": "丙的期末余额"}], "labels": [{"sheet": "附注", "cell": "A4", "value": "丙", "source_sheet": "附注", "source_cell": "A4", "reason": "B新增客户名称"}], "unresolved": []}

    @staticmethod
    def write_book(path, rows, area):
        wb = Workbook()
        ws = wb.active
        ws.title = "附注"
        for row in rows:
            ws.append(row)
        for row in ws:
            for cell in row:
                cell.font = Font(name="宋体", size=10)
                cell.border = Border(bottom=Side(style="thin"))
        ws.column_dimensions["B"].width = 19
        wb.defined_names.add(DefinedName("表格一", attr_text="'附注'!" + area))
        wb.save(path)
        wb.close()

    def run_plan(self, proposal=None, second=None):
        return resolve_structure(ReplyEngine(proposal or self.proposal, second), self.a, self.b, self.ar, self.br, {}, str(self.work))

    def test_insert_customer_preserves_original_and_moves_total(self):
        original = self.a_path.read_bytes()
        result = self.run_plan()
        self.assertEqual(result["issues"], [])
        self.assertEqual(self.a_path.read_bytes(), original)
        self.assertNotEqual(pathlib.Path(result["source_a"]), self.a_path)
        wb = load_workbook(result["source_a"])
        self.assertEqual(wb["附注"]["A4"].value, "丙")
        self.assertEqual(wb["附注"]["A5"].value, "合计")
        self.assertEqual(wb["附注"]["B4"].font.name, "宋体")
        wb.close()
        positions = {m["dimensions"]["counterparty"]: m["cell"] for m in result["a_mappings"]}
        self.assertEqual(positions, {"甲": "B2", "乙": "B3", "丙": "B4", "合计": "B5"})

    def test_layout_then_real_semantic_value_update(self):
        result = self.run_plan()
        self.assertEqual(result["issues"], [])
        plan = build_update_plan(result["a_snapshot"], self.b, result["a_mappings"], self.br["mappings"])
        self.assertEqual(plan["issues"], [])
        self.assertEqual(plan["unplaced_sources"], [])
        updated = self.work / "A更新完成.xlsx"
        write_updated_workbook(result["source_a"], str(updated), plan)
        wb = load_workbook(updated)
        self.assertEqual([wb["附注"]["B"+str(row)].value for row in (2, 3, 4, 5)], [10, 20, 30, 60])
        self.assertEqual(wb["附注"]["B4"].font.name, "宋体")
        wb.close()

    def test_structure_preserves_macro_suffix_and_payload(self):
        import zipfile
        macro=self.work/"宏A.xlsm"
        with zipfile.ZipFile(self.a_path) as source,zipfile.ZipFile(macro,"w") as target:
            for info in source.infolist():target.writestr(info,source.read(info.filename))
            target.writestr("xl/vbaProject.bin",b"non-executable-test-macro-payload")
        self.a=read_workbook(macro)
        result=self.run_plan()
        self.assertEqual(result["issues"],[])
        self.assertEqual(pathlib.Path(result["source_a"]).suffix,".xlsm")
        with zipfile.ZipFile(result["source_a"]) as archive:
            self.assertEqual(archive.read("xl/vbaProject.bin"),b"non-executable-test-macro-payload")

    def test_prior_table_row_insert_moves_later_word_omissions(self):
        source=self.work/"附注含后方省略表.xlsx"
        wb=load_workbook(self.a_path);sheet=wb["附注"]
        sheet["A7"]="期末";sheet["B7"]="期初";sheet["B8"]=50
        wb.defined_names.add(DefinedName("表格二",attr_text="'附注'!A7:B8"))
        wb.save(source);wb.close()
        context=[{"range_name":"表格一","sheet":"附注","first_row":1,"first_column":1,
                  "row_count":4,"column_count":2,"structural_omissions":[]},
                 {"range_name":"表格二","sheet":"附注","first_row":7,"first_column":1,
                  "row_count":2,"column_count":2,"structural_omissions":[
                     {"sheet":"附注","range":"A8","reason":"Word 原表该位置没有单元格"}]},
                 "金额单位：元"]
        self.a=read_workbook(source,context=context)
        result=self.run_plan()
        self.assertEqual(result["issues"],[])
        updated=result["a_snapshot"];second=updated["context"][1]
        self.assertEqual(updated["context"][0]["row_count"],5)
        self.assertEqual(second["first_row"],8)
        self.assertEqual(second["structural_omissions"][0]["range"],"A9")
        self.assertEqual(updated["context"][2],"金额单位：元")
        cells=updated["sheets"][0]["cells"]
        self.assertNotIn("A9",cells);self.assertEqual(cells["B9"]["value"],50)
        restored=read_workbook(result["source_a"],context=updated["context"])
        self.assertNotIn("A9",restored["sheets"][0]["cells"])
        self.assertEqual(context[1]["structural_omissions"][0]["range"],"A8")
        self.assertEqual(read_workbook(source,context=context)["sha256"],self.a["sha256"])

    def test_refuse_structure_edits_inside_word_omission_table(self):
        source=self.work/"附注自身含省略.xlsx"
        wb=load_workbook(self.a_path);sheet=wb["附注"];sheet["A3"]=None
        wb.save(source);wb.close()
        self.a=read_workbook(source,context=[{"range_name":"表格一","sheet":"附注",
            "structural_omissions":[{"sheet":"附注","range":"A3","reason":"Word 原表该位置没有单元格"}]}])
        result=self.run_plan()
        self.assertTrue(any("省略" in issue for issue in result["issues"]),result)
        self.assertEqual(result["source_a"],self.a["path"])
        self.assertEqual(list(self.work.glob("*结构调整*")),[])

    def test_insert_metric_column(self):
        self.write_book(self.work / "A列.xlsx", [["客户", "期末余额"], ["甲", 10], ["乙", 20]], "A1:B3")
        self.write_book(self.work / "B列.xlsx", [["客户", "期末余额", "期初余额"], ["甲", 10, 7], ["乙", 20, 8]], "A1:C3")
        self.a = read_workbook(str(self.work / "A列.xlsx"))
        self.b = read_workbook(str(self.work / "B列.xlsx"))
        self.ar["mappings"] = [mapped("B2", "甲"), mapped("B3", "乙")]
        self.br["mappings"] = self.ar["mappings"] + [mapped("C2", "甲", "C-N001-T001-S003", "期初余额"), mapped("C3", "乙", "C-N001-T001-S003", "期初余额")]
        proposal = {"operations": [{"range_name": "表格一", "axis": "column", "action": "insert", "index": 2, "count": 1, "template_index": 1, "reason": "新增期初指标列"}], "bindings": [{"source_sheet": "附注", "source_cell": "C"+str(row), "target_sheet": "附注", "target_cell": "C"+str(row), "template_sheet": "附注", "template_cell": "B"+str(row), "reason": "新增期初余额"} for row in (2, 3)], "labels": [{"sheet": "附注", "cell": "C1", "value": "期初余额", "source_sheet": "附注", "source_cell": "C1", "reason": "指标表头"}], "unresolved": []}
        result = self.run_plan(proposal)
        self.assertEqual(result["issues"], [])
        self.assertEqual({m["cell"] for m in result["a_mappings"]}, {"B2", "B3", "C2", "C3"})
        wb = load_workbook(result["source_a"])
        self.assertEqual(wb["附注"]["C1"].value, "期初余额")
        wb.close()
        # 真实引擎提供规范metric，即使未返回行列标签，也足以证明该表头。
        for mapping in self.br["mappings"]:
            mapping["dimensions"]["metric"] = mapping["semantic_field"]
            mapping.pop("row_label", None)
            mapping.pop("column_label", None)
        with self.subTest("规范指标证据"):
            self.assertEqual(self.run_plan(proposal)["issues"], [])
        # 自由说明没有确认指标的效力，不能据此绕过标签与业务语义关联检查。
        for mapping in self.br["mappings"]:
            mapping["dimensions"].pop("metric", None)
        with self.subTest("仅说明文字不能作证据"):
            rejected = self.run_plan(proposal)
            self.assertTrue(any("业务对象或指标无关" in issue for issue in rejected["issues"]), rejected)
            self.assertEqual(rejected["source_a"], self.a["path"])

    def test_refuse_deleting_still_disclosed_customer(self):
        proposal = {"operations": [{"range_name": "表格一", "axis": "row", "action": "delete", "index": 1, "count": 1, "reason": "错误删除甲"}], "bindings": [], "labels": [], "unresolved": []}
        self.assertTrue(self.run_plan(proposal)["issues"])

    def test_refuse_fabricated_source_binding_and_label(self):
        proposal = copy.deepcopy(self.proposal)
        proposal["bindings"][0]["source_cell"] = "Z999"
        self.assertTrue(self.run_plan(proposal)["issues"])
        proposal = copy.deepcopy(self.proposal)
        proposal["labels"][0]["value"] = "伪造客户"
        self.assertTrue(self.run_plan(proposal)["issues"])

    def test_delete_customer_absent_from_b_keeps_total(self):
        self.write_book(self.work / "B删行.xlsx", [["客户", "期末余额"], ["甲", 10], ["合计", 10]], "A1:B3")
        self.b = read_workbook(str(self.work / "B删行.xlsx"))
        self.br["mappings"] = [mapped("B2", "甲"), mapped("B3", "合计", "C-N001-T001-S002")]
        proposal = {"operations": [{"range_name": "表格一", "axis": "row", "action": "delete", "index": 2, "count": 1, "reason": "客户乙在B确认已不再披露"}], "bindings": [], "labels": [], "unresolved": []}
        result = self.run_plan(proposal)
        self.assertEqual(result["issues"], [])
        positions = {m["dimensions"]["counterparty"]: m["cell"] for m in result["a_mappings"]}
        self.assertEqual(positions, {"甲": "B2", "合计": "B3"})

    def test_missing_new_customer_label_is_rejected(self):
        proposal = copy.deepcopy(self.proposal)
        proposal["labels"] = []
        result = self.run_plan(proposal)
        self.assertTrue(result["issues"])
        self.assertEqual(result["source_a"], self.a["path"])
    def test_disagreement_never_writes_layout(self):
        other = copy.deepcopy(self.proposal)
        other["operations"][0]["index"] = 2
        result = self.run_plan(second=other)
        self.assertTrue(result["issues"])
        self.assertEqual(result["source_a"], self.a["path"])


class UnusedSourceTests(unittest.TestCase):
    def setUp(self):
        self.work=ROOT/"测试输出"/("未采用来源_"+uuid.uuid4().hex)
        self.work.mkdir(parents=True)
        self.a=self.book("原附注.xlsx",[["项目","期末余额"],["库存现金",10],["本表仅披露货币资金余额",None]])
        self.b=self.book("更新来源.xlsx",[["项目","期末余额"],["库存现金",20],["银行账户数量",3]])
        self.am=[mapped("B2","库存现金")]
        self.bm=[mapped("B2","库存现金"),mapped("B3","银行账户数量","C-N001-T002-S001")]
        for mapping in self.am+self.bm:mapping["table_range"]="A1:B3"
        self.plan=build_update_plan(self.a,self.b,self.am,self.bm)
        self.item={"sheet":"附注","cell":"B3","reason":"账户数量不是原附注本次列示的货币资金余额指标",
            "evidence":[{"carrier":"A","sheet":"附注","cell":"A3","text":"本表仅披露货币资金余额"},
                        {"carrier":"B","sheet":"附注","cell":"A3","text":"银行账户数量"}]}
        self.reply={"unused_sources":[self.item],"unresolved":[]}

    def book(self,name,rows):
        path=self.work/name;wb=Workbook();ws=wb.active;ws.title="附注"
        for row in rows:ws.append(row)
        wb.save(path);wb.close()
        return read_workbook(path)

    def review(self,first=None,second=None):
        function=getattr(structure,"review_unused_sources",None)
        self.assertTrue(callable(function),"缺少独立的未采用来源双轮复核接口")
        engine=ReplyEngine(first or self.reply,second)
        result=function(engine,self.a,self.b,self.am,self.bm,self.plan)
        return result,engine

    def test_extra_indicator_review_works_without_named_ranges_and_keeps_evidence(self):
        before=(copy.deepcopy(self.a),copy.deepcopy(self.b),copy.deepcopy(self.plan))
        second=copy.deepcopy(self.reply);second["unused_sources"][0]["reason"]="原目标只列资金余额，不列账户个数"
        result,engine=self.review(second=second)
        self.assertEqual(engine.calls,2);self.assertEqual(result["issues"],[])
        record=result["unused_sources"][0]
        self.assertEqual((record["sheet"],record["cell"]),("附注","B3"))
        self.assertEqual(record["source_a_hash"],self.a["sha256"])
        self.assertEqual(record["source_b_hash"],self.b["sha256"])
        self.assertEqual(len(record["evidence"]),2)
        self.assertEqual(structure.validate_unused_sources(self.a,self.b,self.am,self.bm,self.plan,result["unused_sources"]),result["unused_sources"])
        self.assertEqual((self.a,self.b,self.plan),before)

    def test_no_extras_needs_no_model(self):
        self.plan["unplaced_sources"]=[]
        result,engine=self.review()
        self.assertEqual(result,{"unused_sources":[],"issues":[]});self.assertEqual(engine.calls,0)

    def test_disagreement_or_unproved_evidence_remains_blocking(self):
        for variant in ("disagree","fabricated","number","one_side","missing_reason"):
            with self.subTest(variant=variant):
                first=copy.deepcopy(self.reply);second=copy.deepcopy(self.reply)
                if variant=="disagree":second["unused_sources"]=[]
                elif variant=="fabricated":first["unused_sources"][0]["evidence"][0]["text"]="原表没有的限制"
                elif variant=="number":first["unused_sources"][0]["evidence"][1].update(cell="B3",text="3")
                elif variant=="one_side":first["unused_sources"][0]["evidence"]=first["unused_sources"][0]["evidence"][1:]
                else:first["unused_sources"][0]["reason"]=""
                result,_=self.review(first,second)
                self.assertTrue(result["issues"]);self.assertEqual(result["unused_sources"],[])

    def test_new_customer_age_or_period_cannot_be_ignored_without_explicit_limit(self):
        for dimension,value in (("counterparty","乙"),("age_bucket",{"lower":1,"upper":2,"lower_inclusive":False,"upper_inclusive":True}),("period","2024-12-31")):
            with self.subTest(dimension=dimension):
                self.bm[1]=copy.deepcopy(self.am[0]);self.bm[1]["cell"]="B3";self.bm[1]["dimensions"][dimension]=value
                self.plan=build_update_plan(self.a,self.b,self.am,self.bm)
                result,_=self.review()
                self.assertTrue(result["issues"]);self.assertEqual(result["unused_sources"],[])

    def test_same_slot_requires_full_scope_statement_and_consistent_allowed_values(self):
        self.a=self.book("仅甲附注.xlsx",[["客户","期末余额"],["甲",10],["本表仅披露甲客户余额",None]])
        self.b=self.book("甲乙来源.xlsx",[["客户","期末余额"],["甲",20],["乙",3]])
        self.am=[mapped("B2","甲")];self.bm=[mapped("B2","甲"),mapped("B3","乙")]
        for mapping in self.am+self.bm:mapping["table_range"]="A1:B3"
        self.plan=build_update_plan(self.a,self.b,self.am,self.bm)
        self.item["evidence"][0].update(text="本表仅披露甲客户余额")
        self.item["evidence"][1].update(text="乙")
        self.item["scope_limit"]={"dimension":"counterparty","allowed_values":["甲"],"evidence":copy.deepcopy(self.item["evidence"][0])}
        result,_=self.review();self.assertEqual(result["issues"],[])
        saved=result["unused_sources"]
        with self.subTest("只列甲不是范围限制"):
            self.item["scope_limit"]["evidence"].update(cell="A2",text="甲")
            result,_=self.review();self.assertTrue(result["issues"])
        with self.subTest("允许值不能包括拟排除来源"):
            self.item["scope_limit"]["evidence"]=copy.deepcopy(self.item["evidence"][0])
            self.item["scope_limit"]["allowed_values"].append("乙")
            result,_=self.review();self.assertTrue(result["issues"])
        self.assertEqual(structure.validate_unused_sources(self.a,self.b,self.am,self.bm,self.plan,saved),saved)

    def test_negative_exclusion_text_is_not_a_positive_allowed_population(self):
        self.a=self.book("排除甲附注.xlsx",[["客户","期末余额"],["甲",10],["本表不包括甲客户余额",None]])
        self.b=self.book("负面范围来源.xlsx",[["客户","期末余额"],["甲",20],["乙",3]])
        self.am=[mapped("B2","甲")];self.bm=[mapped("B2","甲"),mapped("B3","乙")]
        for mapping in self.am+self.bm:mapping["table_range"]="A1:B3"
        self.plan=build_update_plan(self.a,self.b,self.am,self.bm)
        self.item["evidence"][0].update(text="本表不包括甲客户余额")
        self.item["evidence"][1].update(text="乙")
        self.item["scope_limit"]={"dimension":"counterparty","allowed_values":["甲"],"evidence":copy.deepcopy(self.item["evidence"][0])}
        result,_=self.review()
        self.assertTrue(result["issues"]);self.assertEqual(result["unused_sources"],[])

    def test_existing_identity_used_source_and_duplicate_cannot_be_unused(self):
        for variant in ("identity","used","binding","duplicate"):
            with self.subTest(variant=variant):
                plan=copy.deepcopy(self.plan);reply=copy.deepcopy(self.reply)
                if variant=="identity":
                    reply["unused_sources"][0]["cell"]="B2";plan["unplaced_sources"].append(self.bm[0])
                elif variant=="used":plan["updates"][0]["sources"].append({"sheet":"附注","cell":"B3"})
                elif variant=="binding":plan["bindings"]=[{"source_sheet":"附注","source_cell":"B3"}]
                else:reply["unused_sources"].append(copy.deepcopy(reply["unused_sources"][0]))
                function=getattr(structure,"review_unused_sources",None)
                self.assertTrue(callable(function),"缺少独立的未采用来源双轮复核接口")
                result=function(ReplyEngine(reply),self.a,self.b,self.am,self.bm,plan)
                self.assertTrue(result["issues"]);self.assertEqual(result["unused_sources"],[])

    def test_many_sources_are_batched_with_compact_target_context_and_no_amounts(self):
        import json
        rows=[["项目","期末余额"],["库存现金",20]]+[["账户项目"+str(i),987654321] for i in range(70)]
        self.b=self.book("大量不同指标.xlsx",rows)
        self.bm=[mapped("B2","库存现金")]+[mapped("B"+str(i+3),"账户项目"+str(i),"C-N001-T002-S001") for i in range(70)]
        for item in self.bm:item["table_range"]="A1:B72"
        arows=[["项目","期末余额"],["库存现金",10],["本表仅披露货币资金余额",None]]+[["原客户"+str(i),987654321] for i in range(200)]
        self.a=self.book("大量目标对象.xlsx",arows)
        self.am+=[mapped("B"+str(i+4),"原客户"+str(i)) for i in range(200)]
        for item in self.am:item["table_range"]="A1:B203"
        self.plan=build_update_plan(self.a,self.b,self.am,self.bm)
        owner=self
        class BatchedEngine:
            def __init__(self):self.payloads=[]
            def _check_cancel(self):pass
            def _request(self,instruction,payload,validator):
                self.payloads.append(copy.deepcopy(payload))
                items=[]
                for source in payload["unplaced_sources"]:
                    item=copy.deepcopy(owner.item);item["cell"]=source["cell"]
                    address="A"+source["cell"][1:]
                    item["evidence"][1].update(cell=address,text=owner.b["sheets"][0]["cells"][address]["value"])
                    items.append(item)
                return validator({"unused_sources":items,"unresolved":[]})
        engine=BatchedEngine();result=structure.review_unused_sources(engine,self.a,self.b,self.am,self.bm,self.plan)
        self.assertEqual(result["issues"],[],result)
        self.assertEqual(len(result["unused_sources"]),70)
        self.assertGreater(len(engine.payloads),2)
        for payload in engine.payloads:
            self.assertLessEqual(len(payload["unplaced_sources"]),24)
            self.assertLessEqual(len(json.dumps(payload,ensure_ascii=False)),24000)
            self.assertNotIn("987654321",json.dumps(payload,ensure_ascii=False))
            self.assertLess(len(payload["a_mappings"]),len(self.am))

    def test_uncovered_remaining_source_is_not_silently_accepted(self):
        result,_=self.review({"unused_sources":[],"unresolved":[]})
        self.assertTrue(result["issues"]);self.assertEqual(result["unused_sources"],[])

    def test_scope_limit_must_restrict_the_claimed_dimension_not_another_subject(self):
        self.a=self.book("仅期末附注.xlsx",[["客户","期末余额"],["甲",10],["本表仅披露期末余额",None]])
        self.b=self.book("新增乙来源.xlsx",[["客户","期末余额"],["甲",20],["乙",3]])
        self.am=[mapped("B2","甲")];self.bm=[mapped("B2","甲"),mapped("B3","乙")]
        for mapping in self.am+self.bm:mapping["table_range"]="A1:B3"
        self.plan=build_update_plan(self.a,self.b,self.am,self.bm)
        self.item["evidence"][0].update(text="本表仅披露期末余额")
        self.item["evidence"][1].update(text="乙")
        self.item["scope_limit"]={"dimension":"counterparty","allowed_values":["甲"],"evidence":copy.deepcopy(self.item["evidence"][0])}
        result,_=self.review();self.assertTrue(result["issues"]);self.assertEqual(result["unused_sources"],[])

    def test_conflicting_cached_receipts_are_marked_for_retry(self):
        class ReceiptEngine(ReplyEngine):
            def __init__(self,first,second):super().__init__(first,second);self._review_receipts=[];self.retries=[]
            def _request(self,*args):
                result=super()._request(*args);self._review_receipts.append(str(self.calls));return result
            def _retry_receipts(self,start):self.retries.append(start)
        second=copy.deepcopy(self.reply)
        second["unused_sources"][0]["evidence"][0].update(cell="B1",text="期末余额")
        engine=ReceiptEngine(self.reply,second)
        result=structure.review_unused_sources(engine,self.a,self.b,self.am,self.bm,self.plan)
        self.assertTrue(result["issues"]);self.assertEqual(engine.retries,[0])

    def test_local_revalidation_rejects_hash_identity_coordinate_and_evidence_changes(self):
        result,_=self.review();self.assertEqual(result["issues"],[])
        for variant in ("hash","identity","coordinate","evidence","mapping","used"):
            with self.subTest(variant=variant):
                saved=copy.deepcopy(result["unused_sources"]);bmaps=copy.deepcopy(self.bm);plan=copy.deepcopy(self.plan)
                if variant=="hash":saved[0]["source_b_hash"]="changed"
                elif variant=="identity":saved[0]["semantic_key"]="wrong"
                elif variant=="coordinate":saved[0]["cell"]="Z999"
                elif variant=="evidence":saved[0]["evidence"][0]["text"]="改变了的文字"
                elif variant=="mapping":bmaps[1]["dimensions"]["period"]="2023-12-31"
                else:plan["updates"][0]["sources"].append({"sheet":"附注","cell":"B3"})
                with self.assertRaises(ValueError):structure.validate_unused_sources(self.a,self.b,self.am,bmaps,plan,saved)


if __name__ == "__main__":
    unittest.main()