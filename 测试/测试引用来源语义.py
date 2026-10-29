# coding: utf-8
"""引用的空白业务格和来源标签进入共同阅读，不用数值判断身份。"""
import json
import unittest
from 附注更新.A语义阅读 import source_material


def 原格(row, column, value=None, formula=None):
    return {'row':row,'column':column,'value':value,'formula':formula,'cached_value':771188 if formula else None}


class 引用来源语义测试(unittest.TestCase):
    def test_跨表空白组成项与合并标题均可读且金额隐藏(self):
        snapshot={'path':'小例.xlsx','semantic_formula_sources':2,'sheets':[
            {'name':'展示','merges':[],'cells':{'A2':原格(2,1,'设备'),
                'B2':原格(2,2,formula="='分类'!C3+'分类'!C4+'分类'!C5")}},
            {'name':'分类','merges':['A2:A5'],'cells':{'A2':原格(2,1,'生产设施原值'),
                'B2':原格(2,2,'年初数'),'C1':原格(1,3,'本年'),
                'B3':原格(3,2,'甲类'),'B4':原格(4,2,'乙类'),
                'B5':原格(5,2,formula='=B6'),'B6':原格(6,2,'丙类'),
                'C3':原格(3,3),'C4':原格(4,3),'C5':原格(5,3,829977)}}]}
        tables=[{'sheet':'展示','range':'A2:B2','data_ranges':['B2']}]
        material,_=source_material(snapshot,[],tables,semantic_only=True)
        evidence=json.dumps(material.get('公式引用依据',{}),ensure_ascii=False)
        for label in ['甲类','乙类','丙类','生产设施原值','年初数','本年','分类!C3','分类!C4']:
            self.assertIn(label,evidence)
        self.assertNotIn('771188',evidence);self.assertNotIn('829977',evidence)
        snapshot['sheets'][1]['cells']['C3']['value']=558866
        changed,_=source_material(snapshot,[],tables,semantic_only=True)
        self.assertEqual(material['公式引用依据'],changed['公式引用依据'])
        snapshot['semantic_formula_sources']=1
        legacy,_=source_material(snapshot,[],tables,semantic_only=True)
        self.assertNotIn('公式引用依据',legacy)

    def test_循环引用终止且未知名称保留原引用(self):
        snapshot={'path':'小例.xlsx','semantic_formula_sources':2,'sheets':[
            {'name':'展示','merges':[],'cells':{'A2':原格(2,1,'设备'),
                'B2':原格(2,2,formula='=C2+未定义名称'),'C2':原格(2,3,formula='=B2')}}]}
        material,_=source_material(snapshot,[],[{'sheet':'展示','range':'A2:B2','data_ranges':['B2']}],semantic_only=True)
        evidence=material.get('公式引用依据',{})
        self.assertIn('展示!C2',json.dumps(evidence,ensure_ascii=False))
        self.assertEqual(evidence['未展开'],[{'位置':'展示!B2','引用':'未定义名称'}])
        snapshot['sheets'][0]['cells']['B2']['formula']='="未闭合'
        material,_=source_material(snapshot,[],[{'sheet':'展示','range':'A2:B2','data_ranges':['B2']}],semantic_only=True)
        self.assertEqual(material['公式引用依据']['未展开'],[{'位置':'展示!B2','引用':'="未闭合'}])

    def test_普通入口识别及正式回读包含引用语义(self):
        from unittest.mock import patch
        from openpyxl import load_workbook
        from 测试.测试B范围定位 import B范围定位测试
        from 附注更新.统一识别 import UnifiedSemanticEngine
        from 附注更新.统一勾稽 import load_fact_mapping
        case=B范围定位测试();case.setUp()
        source=case.case.folder/'带空白来源的业务表.xlsx'
        book=load_workbook(case.case.source)
        book['综合']['B3']="='组成明细'!C3+'组成明细'!C4"
        detail=book.create_sheet('组成明细')
        detail.append(['固定资产','','期末原值'])
        detail['B3']='甲类';detail['B4']='乙类'
        book.save(source);book.close()
        requests=[]
        def reply(engine,system,payload):
            if payload['task'] in {'a_semantic_bind','a_semantic_review'}:
                requests.append(payload)
                text=json.dumps(payload['source']['公式引用依据'],ensure_ascii=False)
                self.assertIn('甲类',text);self.assertIn('乙类',text)
                self.assertIn('组成明细!C3',text);self.assertIn('组成明细!C4',text)
            return case.response(engine,system,payload)
        with patch.object(UnifiedSemanticEngine,'_http',reply):
            engine=case.case.engine()
            result=engine.recognize_workbook(source,target_mapping=case.case.target())
        self.assertEqual({p['task'] for p in requests},{'a_semantic_bind','a_semantic_review'})
        self.assertEqual(result['formula_source_version'],5)
        saved=load_fact_mapping(result['mapping_path'],engine.standard)
        self.assertEqual(saved['recognition_scope']['coverage']['targets'][0]['exact_sources'],['综合!B3'])


if __name__=='__main__':unittest.main()
