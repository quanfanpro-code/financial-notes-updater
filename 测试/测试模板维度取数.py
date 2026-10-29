# coding: utf-8
"""同槽位的类别、期间和金额一起更新，原识别及原件保持。"""
import copy
import json
import unittest
from pathlib import Path
from openpyxl import load_workbook
from 测试 import 测试统一勾稽 as 夹具
from 附注更新.表格 import file_hash, read_workbook, materialize_candidate_blanks
from 附注更新.统一识别 import workbook_coverage
from 附注更新.统一语义 import read_standard_v2
from 附注更新.统一勾稽 import load_fact_mapping, build_fact_update_plan
from 附注更新.分步流程 import run_step


class 模板维度取数测试(unittest.TestCase):
    def test_槽位增行后保留原维度及复制来源(self):
        from 测试 import 测试统一结构 as 结构夹具
        from 测试.测试槽位记录调整行数 import 槽位记录
        case=结构夹具.UnifiedStructureTests();case.setUp();target=槽位记录(case.target)
        target.update(slot_binding_version=1,semantic_reading={'synthetic':True})
        purpose={'reason':'已核准槽位增行后仍须保留更新前已知维度','operations':case.proposal['operations'],
            'bindings':[{'target':'明细!B'+str(row),'source':'明细!B'+str(row)} for row in range(2,6)],
            'labels':[{'target':'明细!A4','source':'明细!A4'}]}
        plan=build_fact_update_plan(target,case.source,case.standard,update_purpose=purpose)
        self.assertEqual(plan['issues'],[])
        added=next(row for row in plan['updates'] if row['cell']=='B4')
        self.assertEqual(added['original_dimensions']['asset_selection']['members'],['乙类设备'])
        self.assertNotIn('entity',added['original_dimensions'])
        self.assertEqual(added['dimensions']['asset_selection']['members'],['丙类设备'])

    def test_原槽位记录保存时同步最新类别及已知维度(self):
        from 测试 import 测试统一结构 as 结构夹具
        from 测试.测试槽位记录调整行数 import 槽位记录
        from 附注更新.表格 import write_updated_workbook
        from 附注更新.统一勾稽 import derive_fact_mapping,_template_cell_records,_check_mapping
        case=结构夹具.UnifiedStructureTests();case.setUp();target=槽位记录(case.target)
        target.update(slot_binding_version=1,semantic_reading={'synthetic':True})
        snapshot=read_workbook(target['source_path']);materialize_candidate_blanks(snapshot)
        financial={row['sheet']+'!'+row['cell'] for row in target['slot_bindings']}
        target['excluded']=[{'sheet':s['name'],'cell':cell,'category':'label','reason':'受控测试的标题或留白'}
            for s in snapshot['sheets'] for cell in s['cells'] if s['name']+'!'+cell not in financial]
        purpose={'reason':'槽位与数据同时更新，不保留原模板行的旧类别','operations':case.proposal['operations'],
            'bindings':[{'target':f'明细!B{row}','source':f'明细!B{row}'} for row in range(2,6)],
            'labels':[{'target':'明细!A4','source':'明细!A4'}]}
        plan=build_fact_update_plan(target,case.source,case.standard,update_purpose=purpose)
        gold=case.folder/'标准.jsonl';gold.write_text('\n'.join(json.dumps(row,ensure_ascii=False) for row in 结构夹具.records()),encoding='utf-8')
        plan.update(gold_path=str(gold),gold_hash=file_hash(gold))
        pp=case.folder/'更新清单.json';pp.write_text(json.dumps(plan,ensure_ascii=False),encoding='utf-8')
        output=case.folder/'最新槽位更新.xlsx';write_updated_workbook(target['source_path'],output,plan)
        saved=derive_fact_mapping(target,output,plan,pp);_check_mapping(saved,read_standard_v2(gold))
        actual=_template_cell_records(saved)
        self.assertEqual(len(actual),4)
        for item in plan['updates']:
            self.assertEqual(actual[item['sheet']+'!'+item['cell']]['dimensions'],item['dimensions'])
        self.assertEqual(actual['明细!B4']['dimensions']['asset_selection']['members'],['丙类设备'])

    def test_表头改变时一次列出所有漏安排的格(self):
        from 测试 import 测试统一结构 as 结构夹具
        case=结构夹具.UnifiedStructureTests();case.setUp();maps=[]
        for year in ('2024','2025'):
            item=case.make(year,[['项目','年末'],['甲设备',77],['乙设备',88],['丙设备',99]],
                           [('B2','甲设备'),('B3','乙设备'),('B4','丙设备')])
            for fact in item['facts']:
                fact['dimensions']['period']={'kind':'instant','date':year+'-12-31'}
                fact['dimension_evidence']['period']=['明细!B1'];item['owners'][fact['owner_reference']]=copy.deepcopy(fact['dimensions'])
            maps.append(item)
        purpose={'reason':'表头更新必须覆盖所管辖的全部数据格','bindings':[{'target':'明细!B2','source':'明细!B2'}],
                 'labels':[{'target':'明细!B1','source':'明细!B1','period_source':'明细!B2'}]}
        with self.assertRaisesRegex(ValueError,'B3.*B4'):
            build_fact_update_plan(*maps,case.standard,update_purpose=purpose)

    def test_已识别的单一名称引用随来源同步(self):
        from 测试 import 测试统一结构 as 结构夹具
        case=结构夹具.UnifiedStructureTests();case.setUp()
        a=case.make('旧名称',[['项目','期末'],['甲设备',77]],[('B2','甲设备')])
        b=case.make('公式名称',[['项目','期末'],['乙设备',88]],[('B2','乙设备')])
        book=load_workbook(b['source_path']);book.create_sheet('原文')['A1']='乙设备';book['明细']['A2']="='原文'!A1"
        book.save(b['source_path']);book.close();b['source_hash']=file_hash(b['source_path']);b['facts'][0]['source_reference']['file_hash']=b['source_hash']
        purpose={'reason':'身份已核准，按同一名称字段同步实际引用文字','bindings':[{'target':'明细!B2','source':'明细!B2'}],'labels':[]}
        try:plan=build_fact_update_plan(a,b,case.standard,update_purpose=purpose)
        except ValueError as error:self.fail('名称的原文引用没有同步：'+str(error))
        self.assertEqual(plan['labels'][0]['value'],'乙设备')
        self.assertTrue(plan['labels'][0]['source_text_reference_chain'])

    def test_按已识别期间补齐同列日期表头(self):
        from 测试 import 测试统一结构 as 结构夹具
        case=结构夹具.UnifiedStructureTests();case.setUp()
        a=case.make('旧期初',[['项目','期初余额'],['甲设备',77]],[('B2','甲设备')])
        b=case.make('新期初',[['2025年度','金额'],['甲设备',88]],[('B2','甲设备')])
        for item,date,ref in [(a,'2024-12-31','明细!B1'),(b,'2025-01-01','明细!A1')]:
            fact=item['facts'][0];fact['dimensions']['period']={'kind':'instant','date':date}
            fact['dimension_evidence']['period']=[ref];item['owners'][fact['owner_reference']]=copy.deepcopy(fact['dimensions'])
        purpose={'reason':'按原文已核准的新期初日期更新对应列','bindings':[{'target':'明细!B2','source':'明细!B2'}],'labels':[]}
        try:plan=build_fact_update_plan(a,b,case.standard,update_purpose=purpose)
        except ValueError as error:self.fail('实际期间已有依据，应同步列表头：'+str(error))
        self.assertEqual(plan['labels'][0]['cell'],'B1')
        self.assertEqual(plan['labels'][0]['value'],'2025年1月1日\n期初余额')

    def test_一次列出各格未安排同步的维度(self):
        from 测试 import 测试统一结构 as 结构夹具
        case=结构夹具.UnifiedStructureTests();case.setUp();maps=[]
        for year in ('2024','2025'):
            item=case.make(year,[['项目','年末'],['甲设备',77],['乙设备',88]],[('B2','甲设备'),('B3','乙设备')])
            for fact in item['facts']:
                fact['dimensions']['period']={'kind':'instant','date':year+'-12-31'}
                item['owners'][fact['owner_reference']]=copy.deepcopy(fact['dimensions'])
            maps.append(item)
        purpose={'reason':'两格的期间都缺对应表头安排','bindings':[{'target':'明细!'+cell,'source':'明细!'+cell} for cell in ('B2','B3')],'labels':[]}
        with self.assertRaisesRegex(ValueError,'B2 / period.*B3 / period'):
            build_fact_update_plan(*maps,case.standard,update_purpose=purpose)

    def test_表外标题从合法文字清单生成Word安排(self):
        from 测试 import 测试统一结构 as 结构夹具
        case=结构夹具.UnifiedStructureTests();case.setUp()
        a=case.make('旧表外标题',[['旧类别','期末'],['甲设备',11]],[('B2','甲设备')])
        b=case.make('新表外标题',[['新类别','期末'],['甲设备',77]],[('B2','甲设备')])
        from openpyxl.workbook.defined_name import DefinedName
        book=load_workbook(a['source_path']);book.defined_names['附注表']=DefinedName('附注表',attr_text='明细!$A$2:$B$2')
        book.save(a['source_path']);book.close()
        a['source_hash']=file_hash(a['source_path']);a['facts'][0]['source_reference']['file_hash']=a['source_hash']
        a['source_context']=[{'table_id':'测试表路径','range_name':'附注表','title':'旧类别',
                             'sheet':'明细','first_row':2,'first_column':1,'row_count':1,'column_count':2}]
        purpose={'reason':'来源已确认，标题与数据一起更新',
                 'bindings':[{'target':'明细!B2','source':'明细!B2'}],
                 'labels':[{'target':'明细!A1','source':'明细!A1'}]}
        plan=build_fact_update_plan(a,b,case.standard,update_purpose=purpose)
        self.assertEqual(plan.get('caption_updates'),[{'range_name':'附注表','old_text':'旧类别','text':'新类别'}])

    def test_范围变化时补齐已识别的表外名称并拒绝不同来源混用(self):
        from 测试 import 测试统一结构 as 结构夹具
        from openpyxl.workbook.defined_name import DefinedName
        case=结构夹具.UnifiedStructureTests();case.setUp()
        maps=[]
        for title in ('旧范围','新范围'):
            item=case.make(title,[[title,'年末'],['甲设备',77]],[('B2','甲设备')])
            book=load_workbook(item['source_path']);book.defined_names['附注表']=DefinedName('附注表',attr_text='明细!$A$2:$B$2')
            book.save(item['source_path']);book.close()
            item['source_hash']=file_hash(item['source_path']);fact=item['facts'][0];fact['source_reference']['file_hash']=item['source_hash']
            fact['dimensions']['asset_selection']={'domain':'固定资产','mode':'predicate','predicate':title,'completeness':'complete'}
            fact['dimension_evidence']['asset_selection']=['明细!A1']
            item['owners'][fact['owner_reference']]=copy.deepcopy(fact['dimensions'])
            item['source_context']=[{'table_id':'测试表路径','range_name':'附注表','title':title,
                                     'sheet':'明细','first_row':2,'first_column':1,'row_count':1,'column_count':2}]
            maps.append(item)
        purpose={'reason':'已经发证的范围名称须随来源更新','bindings':[{'target':'明细!B2','source':'明细!B2'}],'labels':[]}
        original=copy.deepcopy(purpose)
        try:plan=build_fact_update_plan(*maps,case.standard,update_purpose=purpose)
        except ValueError as error:self.fail('已识别的表外范围名称应同步更新：'+str(error))
        self.assertEqual(plan['labels'][0]['value'],'新范围')
        self.assertEqual(plan['caption_updates'],[{'range_name':'附注表','old_text':'旧范围','text':'新范围'}])
        self.assertEqual(purpose,original)
        unsupported=copy.deepcopy(maps[1]);unsupported['facts'][0]['dimension_evidence']['asset_selection']=['明细!A9']
        with self.assertRaises(ValueError):build_fact_update_plan(maps[0],unsupported,case.standard,update_purpose=purpose)

    def test_前表增行后后表标题按实际命名区关联(self):
        from 测试 import 测试统一结构 as 结构夹具
        from openpyxl.workbook.defined_name import DefinedName
        for start,old_title,new_title in [(4,'范围甲','范围乙'),(11,'范围丙','范围丁')]:
            with self.subTest(start=start):
                case=结构夹具.UnifiedStructureTests();case.setUp();maps=[]
                for title in (old_title,new_title):
                    rows=[['项类','期末'],['甲设备',77]]+[[None,None] for _ in range(start-3)]+[[title,None],['乙设备',88]]
                    item=case.make(title,rows,[('B2','甲设备'),(f'B{start+1}','乙设备')])
                    book=load_workbook(item['source_path'])
                    book.defined_names['附注表']=DefinedName('附注表',attr_text='明细!$A$1:$B$2')
                    book.defined_names['后表']=DefinedName('后表',attr_text=f'明细!$A${start+1}:$B${start+1}')
                    book.active[f'A{start}']=title;book.save(item['source_path']);book.close()
                    item['source_hash']=file_hash(item['source_path'])
                    for fact in item['facts']:fact['source_reference']['file_hash']=item['source_hash']
                    fact=item['facts'][1]
                    fact['dimensions']['asset_selection']={'domain':'固定资产','mode':'predicate','predicate':title,'completeness':'complete'}
                    fact['dimension_evidence']['asset_selection']=[f'明细!A{start}','明细!A1']
                    item['owners'][fact['id']]=copy.deepcopy(fact['dimensions'])
                    item['source_context']=[{'table_id':'后表路径','range_name':'后表','title':title,'sheet':'明细',
                        'first_row':start+1,'first_column':1,'row_count':1,'column_count':2}]
                    maps.append(item)
                purpose={'reason':'前表扩行后，后表已识别范围仍须随来源更新',
                    'operations':[{'range_name':'附注表','axis':'row','action':'insert','index':2,'count':1,'template_index':1,
                        'reason':'已确认完整类别名单新增一行'}],
                    'bindings':[{'target':f'明细!B{r}','source':'明细!B2'} for r in (2,3)]+
                        [{'target':f'明细!B{start+2}','source':f'明细!B{start+1}'}],
                    'labels':[{'target':'明细!A3','source':'明细!A2'}]}
                original=copy.deepcopy(maps[0])
                plan=build_fact_update_plan(*maps,case.standard,update_purpose=purpose)
                self.assertEqual(plan['issues'],[])
                self.assertEqual(plan['caption_updates'],[{'range_name':'后表','old_text':old_title,'text':new_title}])
                self.assertEqual(next(x for x in plan['labels'] if x['cell']==f'A{start+1}')['value'],new_title)
                self.assertEqual(maps[0],original)

    def test_公共表外名称不取各列的局部表头(self):
        from 测试 import 测试统一结构 as 结构夹具
        from openpyxl.workbook.defined_name import DefinedName
        case=结构夹具.UnifiedStructureTests();case.setUp();maps=[]
        for title in ('旧总范围','新总范围'):
            item=case.make(title,[[title,None,None],['项目','阶段一','阶段二'],['甲设备',77,88]],
                           [('B3','甲设备'),('C3','甲设备')])
            book=load_workbook(item['source_path']);book.defined_names['附注表']=DefinedName('附注表',attr_text='明细!$A$2:$C$3')
            book.save(item['source_path']);book.close();item['source_hash']=file_hash(item['source_path'])
            for index,fact in enumerate(item['facts']):
                fact['source_reference']['file_hash']=item['source_hash']
                fact['dimensions']['asset_selection']={'domain':'固定资产','mode':'predicate','predicate':title+str(index),'completeness':'complete'}
                fact['dimension_evidence']['asset_selection']=['明细!A1','明细!'+('B2' if index==0 else 'C2')]
                item['owners'][fact['owner_reference']]=copy.deepcopy(fact['dimensions'])
            item['source_context']=[{'table_id':'测试表路径','range_name':'附注表','title':title,
                                     'sheet':'明细','first_row':2,'first_column':1,'row_count':2,'column_count':3}]
            maps.append(item)
        purpose={'reason':'公共标题与各列表头分层保留','bindings':[{'target':'明细!'+cell,'source':'明细!'+cell} for cell in ('B3','C3')],'labels':[]}
        try:plan=build_fact_update_plan(*maps,case.standard,update_purpose=purpose)
        except ValueError as error:self.fail('共同标题不能取各列局部表头：'+str(error))
        self.assertEqual(plan['caption_updates'],[{'range_name':'附注表','old_text':'旧总范围','text':'新总范围'}])

    def test_相同原文名称接续不同集合表达的来源依据(self):
        from 测试 import 测试统一结构 as 结构夹具
        case=结构夹具.UnifiedStructureTests();case.setUp();maps=[]
        for members in (['甲设备','乙设备'],['甲设备或乙设备']):
            item=case.make('共同原文'+str(len(maps)),[['项目','年末'],['甲设备或乙设备',77]],[('B2','甲设备或乙设备')])
            fact=item['facts'][0];fact['dimensions']['asset_selection']['members']=members
            item['owners'][fact['owner_reference']]=copy.deepcopy(fact['dimensions']);maps.append(item)
        purpose={'reason':'以B已确认身份更新，原文名称相同仍接续来源依据',
                 'bindings':[{'target':'明细!B2','source':'明细!B2'}],'labels':[]}
        try:plan=build_fact_update_plan(*maps,case.standard,update_purpose=purpose)
        except ValueError as error:self.fail('相同原文依据没有接续：'+str(error))
        self.assertEqual(plan['labels'][0]['value'],'甲设备或乙设备')
        self.assertEqual(plan['updates'][0]['dimensions']['asset_selection']['members'],['甲设备或乙设备'])

    def test_先删行再增行时新格名称依据不能退回模板行(self):
        from 测试 import 测试统一结构 as 结构夹具
        case=结构夹具.UnifiedStructureTests();case.setUp()
        b=case.make('替换部分名单',[['项目','年末'],['乙类设备',55],['丙类设备',66],['合计',121]],
                    [('B2','乙类设备'),('B3','丙类设备'),('B4',None)])
        for f in case.target['facts']:f['dimension_evidence']['period']=['明细!B1']
        purpose={'reason':'保留乙，删除甲后新增丙，新增格数字地址恰好等于原模板地址',
            'operations':[{'range_name':'附注表','axis':'row','action':'delete','index':1,'count':1,'template_index':0,
                           'source_evidence':['明细!A1','明细!A2','明细!A3','明细!A4'],'coverage_reason':'来源完整名单乙、丙及合计'},
                          {'range_name':'附注表','axis':'row','action':'insert','index':2,'count':1,'template_index':1}],
            'bindings':[{'target':'明细!B'+str(r),'source':'明细!B'+str(r)} for r in (2,3,4)],
            'labels':[{'target':'明细!A'+str(r),'source':'明细!A'+str(r)} for r in (2,3)]}
        plan=build_fact_update_plan(case.target,b,case.standard,update_purpose=purpose)
        self.assertEqual(plan['issues'],[])
        row=next(x for x in plan['updates'] if x['cell']=='B3')
        self.assertEqual(row['dimensions']['asset_selection']['members'],['丙类设备'])
        from 附注更新.模板布局 import project_layout
        view=project_layout(case.target,b,case.standard,purpose['operations'])
        fact=next(x for x in view['facts'] if x['id']=='明细!B3')
        self.assertEqual(fact['dimension_evidence']['asset_selection'],['明细!A3'])
        self.assertEqual(fact['dimension_evidence']['period'],['明细!B1'])

    def test_其他范围随B改变时将范围差异写入名称(self):
        from 测试 import 测试统一结构 as 结构夹具
        case=结构夹具.UnifiedStructureTests();case.setUp();maps=[]
        for name,excluded in [('旧其他',['甲设备','乙设备']),('新其他',['甲设备'])]:
            item=case.make(name,[['项目','余额'],['其他',12]],[('B2','其他')])
            item['facts'][0]['dimensions']['asset_selection']={'domain':'固定资产','mode':'complement','base':'固定资产',
                'excluded':excluded,'completeness':'complete'}
            item['owners']['明细!B2']=copy.deepcopy(item['facts'][0]['dimensions']);maps.append(item)
        purpose={'reason':'其他的实际范围改变，名称中说明纳入的原排除项目',
            'bindings':[{'target':'明细!B2','source':'明细!B2'}],
            'labels':[{'target':'明细!A2','source':'明细!A2','selection_source':'明细!B2','selection_dimension':'asset_selection'}]}
        plan=build_fact_update_plan(*maps,case.standard,update_purpose=purpose)
        self.assertEqual(plan['issues'],[]);self.assertEqual(plan['labels'][0]['value'],'其他（范围包含：乙设备）')
        self.assertEqual(plan['updates'][0]['dimensions']['asset_selection']['excluded'],['甲设备'])

    def test_其他范围已用完整条件表达时随B更新说明(self):
        from 测试 import 测试统一结构 as 结构夹具
        case=结构夹具.UnifiedStructureTests();case.setUp();maps=[]
        for name,predicate in [('旧其他条件','固定资产中除甲、乙设备以外的其他'),('新其他条件','固定资产中除甲设备以外的其他')]:
            item=case.make(name,[['项目','余额'],['其他',12]],[('B2','其他')])
            item['facts'][0]['dimensions']['asset_selection']={'domain':'固定资产','mode':'predicate',
                'predicate':predicate,'completeness':'complete'}
            item['owners']['明细!B2']=copy.deepcopy(item['facts'][0]['dimensions']);maps.append(item)
        purpose={'reason':'按已识别的完整条件说明其他的实际范围，不解析或改写语义',
            'bindings':[{'target':'明细!B2','source':'明细!B2'}],
            'labels':[{'target':'明细!A2','source':'明细!A2','selection_source':'明细!B2','selection_dimension':'asset_selection'}]}
        original=copy.deepcopy(maps)
        plan=build_fact_update_plan(*maps,case.standard,update_purpose=purpose)
        self.assertEqual(plan['issues'],[])
        self.assertEqual(plan['labels'][0]['value'],'其他（范围：固定资产中除甲设备以外的其他）')
        self.assertEqual(plan['updates'][0]['dimensions']['asset_selection'],maps[1]['facts'][0]['dimensions']['asset_selection'])
        same=build_fact_update_plan(maps[1],maps[1],case.standard,update_purpose=purpose)
        self.assertEqual(same['labels'][0]['value'],'其他')
        missing=copy.deepcopy(maps[1]);missing['facts'][0]['dimension_evidence']['asset_selection']=['明细!A1']
        with self.assertRaisesRegex(ValueError,'范围说明须有B原文'):
            build_fact_update_plan(maps[0],missing,case.standard,update_purpose=purpose)
        self.assertEqual(maps,original)

    def test_原期间来自表外时可明确写入对应列表头(self):
        from 测试 import 测试统一结构 as 结构夹具
        case=结构夹具.UnifiedStructureTests();case.setUp()
        a=case.make('表外原期间',[['项目','余额'],['甲设备',11]],[('B2','甲设备')])
        b=case.make('新期间表头',[['项目','2024年12月31日余额'],['甲设备',77]],[('B2','甲设备')])
        book=load_workbook(a['source_path']);book['明细']['A8']='2025年度附注';book.save(a['source_path']);book.close()
        a['source_hash']=file_hash(a['source_path']);a['facts'][0]['source_reference']['file_hash']=a['source_hash']
        a['facts'][0]['dimension_evidence']['period']=['明细!A8']
        b['facts'][0]['dimensions']['period']={'kind':'instant','date':'2024-12-31'}
        b['facts'][0]['dimension_evidence']['period']=['明细!B1'];b['owners']['明细!B2']=copy.deepcopy(b['facts'][0]['dimensions'])
        purpose={'reason':'实际新期间明确落到同表同列的表头，随Word表格回写',
            'bindings':[{'target':'明细!B2','source':'明细!B2'}],
            'labels':[{'target':'明细!B1','source':'明细!B1','period_source':'明细!B2','period_targets':['明细!B2'],'preserve_text':True}]}
        plan=build_fact_update_plan(a,b,case.standard,update_purpose=purpose)
        self.assertEqual(plan['issues'],[]);self.assertEqual(plan['labels'][0]['value'],'2024年12月31日\n余额')
        invalid=copy.deepcopy(purpose);invalid['labels'][0]['target']='明细!A1'
        with self.assertRaises(ValueError):build_fact_update_plan(a,b,case.standard,update_purpose=invalid)

    def test_名单变更有来源随B更新无来源不能清空旧值(self):
        from 测试 import 测试统一结构 as 结构夹具
        case=结构夹具.UnifiedStructureTests();case.setUp()
        a=case.make('旧项目附属值',[['项目','总体','指定范围'],['甲设备',11,12]],[('B2','甲设备'),('C2','甲设备')])
        b=case.make('新项目仅有总体',[['项目','总体'],['乙设备',77]],[('B2','乙设备')])
        old=a['facts'][1];old['dimensions']['asset_selection']={'domain':'固定资产','mode':'predicate','predicate':'甲设备；指定范围','completeness':'complete'}
        a['owners'][old['owner_reference']]=copy.deepcopy(old['dimensions'])
        dims=copy.deepcopy(old['dimensions']);dims['asset_selection']['predicate']='乙设备；指定范围'
        missing={'target':'明细!C2','reason':'B完整表只列总体，没有指定范围的数值','source_evidence':['明细!A1','明细!B1'],
                 'dimensions':dims,'dimension_evidence':copy.deepcopy(old['dimension_evidence'])}
        purpose={'reason':'名单及其数值随B更新，缺少附属值不能沿用旧项目数字',
                 'bindings':[{'target':'明细!B2','source':'明细!B2'}],
                 'labels':[{'target':'明细!A2','source':'明细!A2'}],'unprovided':[missing]}
        with self.assertRaisesRegex(ValueError,'不能清空'):
            build_fact_update_plan(a,b,case.standard,update_purpose=purpose)
        self.assertEqual(a['facts'][1]['raw_value'],12)
        marked=copy.deepcopy(purpose);marked['unprovided'][0]['action']='mark_missing'
        marked_plan=build_fact_update_plan(a,b,case.standard,update_purpose=marked)
        self.assertEqual(marked_plan['unprovided_fields'][0]['action'],'mark_missing')
        self.assertEqual([(x['cell'],x['value']) for x in marked_plan['updates']],[('B2',77),('C2','B未提供')])
        from 附注更新.表格 import write_updated_workbook
        marked_path=case.folder/'明确注明缺字段.xlsx'
        write_updated_workbook(a['source_path'],marked_path,marked_plan)
        actual=load_workbook(marked_path)
        self.assertEqual([actual['明细'][c].value for c in ('A2','B2','C2')],['乙设备',77,'B未提供']);actual.close()
        supplied=case.make('新项目也有指定范围',[['项目','总体','指定范围'],['乙设备',77,88]],
                           [('B2','乙设备'),('C2','乙设备')])
        supplied['facts'][1]['dimensions']=copy.deepcopy(dims)
        supplied['owners'][supplied['facts'][1]['owner_reference']]=copy.deepcopy(dims)
        full=copy.deepcopy(purpose);full.pop('unprovided');full['bindings'].append({'target':'明细!C2','source':'明细!C2'})
        with self.assertRaisesRegex(ValueError,'B已有相同语义来源'):
            build_fact_update_plan(a,supplied,case.standard,update_purpose=marked)
        updated=build_fact_update_plan(a,supplied,case.standard,update_purpose=full)
        self.assertEqual([(x['cell'],x['value']) for x in updated['updates']],[('B2',77),('C2',88)])
        self.assertEqual(updated['labels'][0]['value'],'乙设备')
        self.assertEqual(updated['updates'][1]['dimensions'],dims)
        keep_b=case.make('旧项目仅有总体',[['项目','总体'],['甲设备',77]],[('B2','甲设备')])
        keep=copy.deepcopy(purpose);keep['unprovided'][0].pop('dimensions');keep['unprovided'][0].pop('dimension_evidence')
        retained=build_fact_update_plan(a,keep_b,case.standard,update_purpose=keep)
        self.assertEqual(retained['unprovided_fields'][0]['action'],'retain')
        self.assertEqual([x['cell'] for x in retained['updates']],['B2'])
        available=case.make('已提供指定范围',[['项目','总体','指定范围'],['甲设备',77,88]],[('B2','甲设备'),('C2','甲设备')])
        available['facts'][1]['dimensions']=copy.deepcopy(old['dimensions'])
        available['owners'][available['facts'][1]['owner_reference']]=copy.deepcopy(old['dimensions'])
        with self.assertRaises(ValueError):build_fact_update_plan(a,available,case.standard,update_purpose=keep)
        for fact in available['facts']:
            fact['dimensions']['entity']='另一企业'
            available['owners'][fact['owner_reference']]=copy.deepcopy(fact['dimensions'])
        with self.assertRaisesRegex(ValueError,'B已有相同语义来源'):
            build_fact_update_plan(a,available,case.standard,update_purpose=keep)
        omitted=copy.deepcopy(purpose);omitted.pop('unprovided')
        with self.assertRaises(ValueError):build_fact_update_plan(a,b,case.standard,update_purpose=omitted)
        invalid=copy.deepcopy(purpose);invalid['unprovided'][0].pop('dimensions');invalid['unprovided'][0].pop('dimension_evidence')
        with self.assertRaises(ValueError):build_fact_update_plan(a,b,case.standard,update_purpose=invalid)
        book=load_workbook(a['source_path']);book['明细']['C2']=None;book.save(a['source_path']);book.close()
        a['source_hash']=file_hash(a['source_path'])
        for fact in a['facts']:fact['source_reference']['file_hash']=a['source_hash']
        a['facts'][1].update(raw_value=None,raw_value_state='blank')
        plan=build_fact_update_plan(a,b,case.standard,update_purpose=purpose)
        self.assertEqual([(x['cell'],x['value']) for x in plan['updates']],[('B2',77)])
        self.assertEqual(plan['unprovided_fields'][0]['action'],'retain')
        gold=case.folder/'标准.jsonl';gold.write_text('\n'.join(json.dumps(x,ensure_ascii=False) for x in 结构夹具.records()),encoding='utf-8')
        paths=[]
        for side,item in [('A',a),('B',b)]:
            item.update(gold_path=str(gold),gold_hash=file_hash(gold),unresolved=[])
            snapshot=read_workbook(item['source_path']);materialize_candidate_blanks(snapshot)
            financial={r['source_reference']['location'] for r in item['facts']}
            item['excluded']=[{'sheet':s['name'],'cell':address,'category':'label','reason':'标题或留白'}
                              for s in snapshot['sheets'] for address in s['cells'] if s['name']+'!'+address not in financial]
            item['coverage']=workbook_coverage(item,snapshot);item['complete']=item['coverage']['complete']
            path=case.folder/(side+'映射.json');path.write_text(json.dumps(item,ensure_ascii=False),encoding='utf-8-sig');paths.append(path)
        config={'output_dir':str(case.folder/'正式入口'),'gold_path':str(gold),'a_mapping':str(paths[0]),'b_mapping':str(paths[1]),'update_purpose':purpose}
        matched=run_step('match',config);self.assertEqual(matched['status'],'complete',matched)
        written=run_step('write',{**config,'update_plan':matched['files']['update_plan']});self.assertEqual(written['status'],'complete',written)
        saved=json.loads(Path(written['files']['a_prime_mapping']).read_text(encoding='utf-8-sig'))
        self.assertNotIn('facts',saved)
        self.assertEqual(saved['collections']['A']['mapping_path'],str(paths[0]))
        book=load_workbook(written['files']['a_prime'])
        self.assertIsNone(book['明细']['C2'].value);book.close()
        trace=load_workbook(written['files']['trace']);self.assertIn('B未提供字段',trace.sheetnames);trace.close()
        marked_config={**config,'output_dir':str(case.folder/'缺字段标注正式入口'),'update_purpose':marked}
        matched=run_step('match',marked_config);self.assertEqual(matched['status'],'complete',matched)
        written=run_step('write',{**marked_config,'update_plan':matched['files']['update_plan']})
        self.assertEqual(written['status'],'complete',written);self.assertEqual(written['updated_cells'],1)
        saved=json.loads(Path(written['files']['a_prime_mapping']).read_text(encoding='utf-8-sig'))
        self.assertNotIn('facts',saved)
        book=load_workbook(written['files']['a_prime'])
        self.assertEqual(book['明细']['C2'].value,'B未提供');book.close()
        trace=load_workbook(written['files']['trace'])
        self.assertEqual(trace['更新记录'].max_row,2)
        self.assertEqual(trace['B未提供字段']['D2'].value,'标注B未提供，未计为取数更新');trace.close()

    def test_同一名称可明确连接各区块的原文依据(self):
        from 测试 import 测试统一结构 as 结构夹具
        case=结构夹具.UnifiedStructureTests();case.setUp()
        a=case.make('并列不同范围',[['项目','总体','指定范围'],['甲设备',11,12]],[('B2','甲设备'),('C2','甲设备')])
        b=case.make('同名不同区块',[['项目','余额'],['乙设备',77],['乙设备',88]],[('B2','乙设备'),('B3','乙设备')])
        for mapping,member in [(a,'甲设备'),(b,'乙设备')]:
            fact=mapping['facts'][1]
            fact['dimensions']['asset_selection']={'domain':'固定资产','mode':'predicate','predicate':member+'；指定范围','completeness':'complete'}
            fact['dimensions']['period']={'kind':'instant','date':'2024-12-31'}
            mapping['owners'][fact['owner_reference']]=copy.deepcopy(fact['dimensions'])
        purpose={'reason':'同名原文在各区块各自证明对应范围，明确登记多个来源',
            'bindings':[{'target':'明细!B2','source':'明细!B2'},{'target':'明细!C2','source':'明细!B3'}],
            'labels':[{'target':'明细!A2','source':'明细!A2','source_aliases':['明细!A3']}]}
        plan=build_fact_update_plan(a,b,case.standard,update_purpose=purpose)
        self.assertEqual(plan['issues'],[])
        self.assertEqual([row['value'] for row in plan['updates']],[77,88])
        self.assertEqual(plan['labels'][0]['source_aliases'],['明细!A3'])
        for alias in ['明细!B3','明细!A1']:
            invalid=copy.deepcopy(purpose);invalid['labels'][0]['source_aliases']=[alias]
            with self.assertRaises(ValueError):build_fact_update_plan(a,b,case.standard,update_purpose=invalid)

    def test_模板名称追溯原工作簿的单格文字引用(self):
        from 测试 import 测试统一结构 as 结构夹具
        case=结构夹具.UnifiedStructureTests();case.setUp()
        a=case.make('原名称',[['项目','余额'],['甲设备',11]],[('B2','甲设备')])
        purpose={'reason':'B名称引用其原工作簿中的实际文字',
            'bindings':[{'target':'明细!B2','source':'明细!B2'}],
            'labels':[{'target':'明细!A2','source':'明细!A2'}]}
        b=case.make('引用名称',[['项目','余额'],['=A3',77],['乙设备',None]],[('B2','乙设备')])
        from 附注更新.统一结构 import _text_evidence
        book=load_workbook(b['source_path'])
        try:_text_evidence(book,['明细!A2'])
        finally:book.close()
        plan=build_fact_update_plan(a,b,case.standard,update_purpose=purpose)
        self.assertEqual(plan['issues'],[])
        self.assertEqual(plan['labels'][0]['value'],'乙设备')
        self.assertEqual(plan['labels'][0]['source_text_reference_chain'],
            [{'location':'明细!A2','formula':'=A3'},{'location':'明细!A3','text':'乙设备'}])
        for name,formula,value in [('循环','=A2','乙设备'),('计算','=SUM(A3)',33),('数字','=A3',33)]:
            bad=case.make(name,[['项目','余额'],[formula,77],[value,None]],[('B2','乙设备')])
            with self.assertRaises(ValueError):
                build_fact_update_plan(a,bad,case.standard,update_purpose=purpose)

    def test_本次完整名单替换不受B其他表旧成员阻止(self):
        from 测试 import 测试统一结构 as 结构夹具
        case=结构夹具.UnifiedStructureTests();case.setUp()
        a=case.target
        a['facts'][0]['dimension_evidence']['asset_selection'].append('明细!A3')
        b=case.make('新名单与另表旧成员',[['本次明细仅甲，合计以下是另一张表','余额'],['甲类设备',77],['合计',77],
            ['另一张表',None],['乙类设备',88]],[('B2','甲类设备'),('B3',None),('B5','乙类设备')])
        purpose={'reason':'用本次B完整名单替换A，另一张表仍独立保留',
            'operations':[{'range_name':'附注表','axis':'row','action':'delete','index':2,'count':1,'template_index':1,
                'reason':'本次表仅列甲及合计','source_evidence':['明细!A1','明细!A3'],'coverage_reason':'原文明确合计以下是另一张表'}],
            'bindings':[{'target':'明细!B2','source':'明细!B2'},{'target':'明细!B3','source':'明细!B3'}],'labels':[]}
        plan=build_fact_update_plan(a,b,case.standard,update_purpose=purpose)
        self.assertEqual(plan['issues'],[])
        self.assertEqual([r['value'] for r in plan['updates']],[77,77])
        self.assertEqual([r['id'] for r in plan['unplaced_sources']],['明细!B5'])
        self.assertEqual(plan['template_layout']['removed'],['明细!B3'])
        self.assertIn('明细!A3',a['facts'][0]['dimension_evidence']['asset_selection'])

    def test_B分块列示的同一类别可共用A一处名称(self):
        from 测试 import 测试统一结构 as 结构夹具
        case=结构夹具.UnifiedStructureTests();case.setUp()
        a=case.make('横向两期',[['项目','本年','上年'],['甲设备',11,12]],[('B2','甲设备'),('C2','甲设备')])
        b=case.make('分别列示两期',[['项目','余额'],['乙设备',77],['乙设备',88]],[('B2','乙设备'),('B3','乙设备')])
        b['facts'][1]['dimensions']['period']={'kind':'instant','date':'2024-12-31'}
        b['owners'][b['facts'][1]['owner_reference']]=copy.deepcopy(b['facts'][1]['dimensions'])
        purpose={'reason':'A一个名称对应两期，B同一类别在两个区块分别列示',
            'bindings':[{'target':'明细!B2','source':'明细!B2'},{'target':'明细!C2','source':'明细!B3'}],
            'labels':[{'target':'明细!A2','source':'明细!A2'}]}
        plan=build_fact_update_plan(a,b,case.standard,update_purpose=purpose)
        self.assertEqual(plan['issues'],[])
        self.assertEqual([r['value'] for r in plan['updates']],[77,88])
        self.assertEqual(plan['labels'][0]['value'],'乙设备')
        b['facts'][1]['dimensions']['asset_selection']['members']=['丙设备']
        b['owners'][b['facts'][1]['owner_reference']]=copy.deepcopy(b['facts'][1]['dimensions'])
        with self.assertRaisesRegex(ValueError,'维度变化须同步'):
            build_fact_update_plan(a,b,case.standard,update_purpose=purpose)

    def test_同一年度表头覆盖期初变动及期末且保留原文字(self):
        from 测试 import 测试统一结构 as 结构夹具
        case=结构夹具.UnifiedStructureTests();case.setUp()
        metric=copy.deepcopy(case.standard['metric']['固定资产原值'])
        metric['id']='固定资产增加';metric['definition'].update(measure='增加额',time_kind='duration')
        case.standard['metric'][metric['id']]=metric
        maps=[]
        for name,year in [('原年度',2025),('来源年度',2024)]:
            item=case.make(name,[['项目','本期数据'],['期初余额',10],['本期增加',20],['期末余额',30]],
                           [('B2',None),('B3',None),('B4',None)])
            for index,fact in enumerate(item['facts']):
                if index==1:
                    fact['metric_id']='固定资产增加'
                    fact['dimensions']['period']={'kind':'duration','start':f'{year}-01-01','end':f'{year}-12-31'}
                else:fact['dimensions']['period']={'kind':'instant','date':f'{year}-'+('01-01' if index==0 else '12-31')}
                fact['dimension_evidence']['period']=['明细!B1']
                item['owners'][fact['owner_reference']]=copy.deepcopy(fact['dimensions'])
            maps.append(item)
        purpose={'reason':'同一年度列内含期初、发生额和期末，三种实际期间分别来自B',
            'bindings':[{'target':f'明细!B{r}','source':f'明细!B{r}'} for r in (2,3,4)],
            'labels':[{'target':'明细!B1','source':'明细!B1','period_sources':['明细!B2','明细!B3','明细!B4'],
                       'preserve_text':True}]}
        plan=build_fact_update_plan(*maps,case.standard,update_purpose=purpose)
        self.assertEqual(plan['issues'],[])
        self.assertEqual(plan['labels'][0]['value'],'2024年1月1日至2024年12月31日\n本期数据')
        self.assertEqual([x['dimensions']['period'] for x in plan['updates']],
                         [x['dimensions']['period'] for x in maps[1]['facts']])
        for item in maps:item['facts'][0]['dimension_evidence']['period'].append('明细!A2')
        maps[0]['facts'][1]['dimension_evidence']['period'].append('明细!A2')
        broader=copy.deepcopy(purpose)
        broader['labels'].append({'target':'明细!A2','source':'明细!A2','period_source':'明细!B2'})
        self.assertEqual(build_fact_update_plan(*maps,case.standard,update_purpose=broader)['issues'],[])
        bad=copy.deepcopy(purpose);bad['labels'][0]['period_sources'].pop()
        with self.assertRaisesRegex(ValueError,'同一表头的期间'):
            build_fact_update_plan(*maps,case.standard,update_purpose=bad)

    def test_行数变化后原逐格来源仍跟随目标位置(self):
        from 测试 import 测试统一结构 as 结构夹具
        case=结构夹具.UnifiedStructureTests();case.setUp();a,b=case.target,case.source
        book=load_workbook(b['source_path']);book.active['A6']='另一处同槽位合计';book.active['B6']=999
        book.save(b['source_path']);book.close();b['source_hash']=file_hash(b['source_path'])
        duplicate=copy.deepcopy(b['facts'][-1]);duplicate.update(id='明细!B6',owner_reference='明细!B6',raw_value=999)
        duplicate['source_reference']['location']='明细!B6'
        duplicate['dimension_evidence']={k:['明细!A6'] for k in duplicate['dimensions']}
        b['facts'].append(duplicate);b['owners'][duplicate['id']]=copy.deepcopy(duplicate['dimensions'])
        for fact in b['facts']:fact['source_reference']['file_hash']=b['source_hash']
        b['recognition_scope']={'details':[{'targets':[{'id':'明细!B4','candidates':['明细!B5']}]}]}
        purpose={'reason':'新增类别后原合计对应继续沿用','operations':case.proposal['operations'],
                 'bindings':[{'target':'明细!B'+str(r),'source':'明细!B'+str(r)} for r in range(2,5)],
                 'labels':[{'target':'明细!A4','source':'明细!A4'}]}
        plan=build_fact_update_plan(a,b,case.standard,update_purpose=purpose)
        self.assertEqual(plan['issues'],[])
        total=next(r for r in plan['updates'] if r['cell']=='B5')
        self.assertEqual(total['value'],231)
        self.assertEqual(total['sources'][0]['cell'],'B5')

    def test_空白名单按已绑定行搬值并保留未填姓名(self):
        self.核验空白名单()

    def test_缺主体槽位保存时不冒充完整业务对象(self):
        self.核验空白名单(missing_entity=True)

    def 核验空白名单(self,missing_entity=False):
        from 测试 import 测试统一结构 as 结构夹具
        case=结构夹具.UnifiedStructureTests();case.setUp()
        a=case.target
        snapshot=read_workbook(a['source_path']);materialize_candidate_blanks(snapshot)
        financial={r['id'] for r in a['facts']}
        a['excluded']=[{'sheet':s['name'],'cell':address,'category':'label','reason':'测试表标题及留白'}
                       for s in snapshot['sheets'] for address in s['cells'] if s['name']+'!'+address not in financial]
        a['unresolved']=[]
        b=case.make('空白名单来源',[['本期完整类别明细','年末余额'],[None,0.5],['合计',0.5]],
                    [('B2','未知类别'),('B3',None)])
        item=b['facts'].pop(0);known={k:copy.deepcopy(item[k]) for k in ('metric_id','dimensions','dimension_evidence')}
        known['dimensions']['asset_selection']={'domain':'固定资产','mode':'unknown','completeness':'unknown'}
        known['dimension_evidence']={k:['明细!A1'] for k in known['dimensions']}
        known.update(missing_dimensions=[],evidence_cells=['明细!A1'])
        if missing_entity:
            known['dimensions'].pop('entity');known['dimension_evidence'].pop('entity')
            known['missing_dimensions']=['entity'];b['owners'].pop(item['owner_reference'])
        b.update(slot_binding_version=1,semantic_reading={'synthetic':True},slot_bindings=[
            {'sheet':'明细','cell':'B2','partial_semantics':known,'recognition_provenance':item['recognition_provenance']}])
        purpose={'reason':'以B一行空白姓名的明细替换A两行旧名单',
                 'operations':[{'range_name':'附注表','axis':'row','action':'delete','index':2,'count':1,'template_index':1,
                    'reason':'按B完整明细行数更新模板','source_evidence':['明细!A1'],
                    'coverage_reason':'用户要求名单及数据一起更新，来源明细为一行，随后为合计'}],
                 'bindings':[{'target':'明细!B2','source':'明细!B2'},{'target':'明细!B3','source':'明细!B3'}],
                 'labels':[{'target':'明细!A2','source':'明细!A2'}]}
        from 附注更新.统一勾稽 import select_mapping_scope
        local=select_mapping_scope(a,[{'sheet':'明细','range':'A1:B4'}])
        plan=build_fact_update_plan(local,b,case.standard,update_purpose=purpose)
        self.assertEqual(plan['issues'],[])
        self.assertEqual([r['value'] for r in plan['updates']],[0.5,0.5])
        self.assertIsNone(plan['labels'][0]['value'])
        self.assertEqual(plan['updates'][0]['dimensions']['asset_selection'],known['dimensions']['asset_selection'])
        self.assertEqual(plan['template_layout']['removed'],['明细!B3'])
        self.assertEqual(a['facts'][0]['dimensions']['asset_selection']['members'],['甲类设备'])
        from 附注更新.表格 import write_updated_workbook
        from 附注更新.统一勾稽 import derive_fact_mapping,_check_mapping
        gold=case.folder/'标准.jsonl'
        gold.write_text('\n'.join(json.dumps(r,ensure_ascii=False) for r in 结构夹具.records()),encoding='utf-8-sig')
        plan['gold_path']=str(gold);plan['gold_hash']=file_hash(gold)
        pp=case.folder/'更新清单.json';pp.write_text(json.dumps(plan,ensure_ascii=False),encoding='utf-8-sig')
        output=case.folder/'更新A.xlsx';write_updated_workbook(a['source_path'],output,plan)
        saved=derive_fact_mapping(a,output,plan,pp)
        _check_mapping(saved,read_standard_v2(gold))
        self.assertEqual(len(saved['facts']),1)
        if missing_entity:
            self.assertEqual(saved['slot_bindings'][0]['partial_semantics']['dimensions']['entity'],a['facts'][0]['dimensions']['entity'])
            self.assertNotIn('entity',b['slot_bindings'][0]['partial_semantics']['dimensions'])
            self.assertEqual(set(saved['owners']),{row['owner_reference'] for row in saved['facts']})
        self.assertEqual(saved['slot_bindings'][0]['cell'],'B2')
        self.assertEqual(saved['slot_bindings'][0]['partial_semantics']['dimensions']['asset_selection'],known['dimensions']['asset_selection'])
        book=load_workbook(output)
        self.assertEqual([[book.active.cell(r,c).value for c in (1,2)] for r in (2,3)],[[None,0.5],['合计',0.5]])
        book.close()

    def test_名单增行沿用模板并保存所有新旧格的实际维度(self):
        self.核验名单增行()

    def test_Word表区增行不登记旁侧不存在的空位(self):
        self.核验名单增行(word_table=True)

    def 核验名单增行(self,word_table=False):
        from 测试 import 测试统一结构 as 结构夹具
        case=结构夹具.UnifiedStructureTests();case.setUp()
        book=load_workbook(case.target['source_path']);book.active['D8']='旁侧说明';book.save(case.target['source_path']);book.close()
        case.target['source_hash']=file_hash(case.target['source_path'])
        for fact in case.target['facts']:fact['source_reference']['file_hash']=case.target['source_hash']
        from 附注更新.统一语义 import definition_hash
        from 附注更新.统一证据 import context_catalog, resolve_context
        context={'text':'本表按资产类别列示','sheet':'明细','first_row':1,'first_column':1,'row_count':4,'column_count':2}
        if word_table:context.update(table_id='原Word表路径',range_name='附注表')
        case.target['source_context']=[context]
        old_reference='@context:'+definition_hash(context)
        case.target['facts'][-1]['dimension_evidence']['asset_selection'].append(old_reference)
        gold=case.folder/'标准.jsonl'
        gold.write_text('\n'.join(json.dumps(row,ensure_ascii=False) for row in 结构夹具.records()),encoding='utf-8')
        paths=[]
        for side,item in [('A',case.target),('B',case.source)]:
            item.update(gold_path=str(gold),gold_hash=file_hash(gold),unresolved=[])
            snapshot=read_workbook(item['source_path'],context=item.get('source_context',[]));materialize_candidate_blanks(snapshot)
            financial={r['source_reference']['location'] for r in item['facts']}
            item['excluded']=[{'sheet':s['name'],'cell':address,'category':'label','reason':'测试表的标题或留白'}
                              for s in snapshot['sheets'] for address in s['cells'] if s['name']+'!'+address not in financial]
            item['coverage']=workbook_coverage(item,snapshot);item['complete']=item['coverage']['complete']
            path=case.folder/(side+'映射.json');path.write_text(json.dumps(item,ensure_ascii=False),encoding='utf-8-sig');paths.append(path)
        purpose={'reason':'按B的完整类别名单更新A模板','operations':case.proposal['operations'],
                 'bindings':[{'target':'明细!B'+str(r),'source':'明细!B'+str(r)} for r in range(2,6)],
                 'labels':[{'target':'明细!A4','source':'明细!A4'}]}
        config={'output_dir':str(case.folder/'正式入口'),'gold_path':str(gold),'a_mapping':str(paths[0]),
                'b_mapping':str(paths[1]),'update_purpose':purpose}
        matched=run_step('match',config)
        written=run_step('write',{**config,'update_plan':matched['files']['update_plan']})
        self.assertEqual(written['status'],'complete',written)
        book=load_workbook(written['files']['a_prime']);sheet=book['明细']
        self.assertEqual([[sheet.cell(r,c).value for c in (1,2)] for r in range(2,6)],
                         [['甲类设备',77],['乙类设备',77],['丙类设备',77],['合计',231]])
        self.assertEqual(sheet['A9'].value,'后方说明');self.assertEqual(sheet['D9'].value,'旁侧说明');book.close()
        saved=json.loads(Path(written['files']['a_prime_mapping']).read_text(encoding='utf-8-sig'))
        self.assertNotIn('facts',saved)
        self.assertEqual(saved['collections']['A']['mapping_path'],str(paths[0]))
        self.assertEqual(saved['collections']['B']['mapping_path'],str(paths[1]))
        self.assertEqual(next(r for r in saved['updates'] if r['cell']=='B4')['actual_value'],77)
        self.assertEqual(case.target['source_context'][0]['row_count'],4)

    def test_类别及年度更新只记执行动作而不拼出新身份(self):
        case=夹具.MatchTests();case.setUp()
        root,gold,old_paths=case.file_fixture();paths=[]
        for side,old_path,name,year,period_cell,label_cell in (
                ('A',old_paths[0],'甲设备',2024,'D7','A8'),
                ('B',old_paths[1],'乙运输工具',2025,'J2','A3')):
            item=json.loads(old_path.read_text(encoding='utf-8-sig'))
            book=load_workbook(item['source_path']);sheet=book.active
            sheet['A1']='合成企业甲；单户；固定资产原值；人民币元'
            sheet[period_cell]='年末余额';sheet[label_cell]=name
            path=root/(side+'不同维度.xlsx');book.save(path);book.close()
            item.update(source_path=str(path),source_hash=file_hash(path))
            fact=item['facts'][0];fact['source_reference']['file_hash']=item['source_hash']
            fact['dimensions']['period']={'kind':'instant','date':f'{year}-12-31'}
            if side=='B':fact['dimensions']['entity']='来源主体乙'
            fact['dimensions']['asset_selection']={'domain':'固定资产','mode':'members','members':[name],'completeness':'complete'}
            fact['dimension_evidence']['period']=[sheet.title+'!'+period_cell]
            fact['dimension_evidence']['asset_selection']=[sheet.title+'!'+label_cell]
            item['owners'][fact['owner_reference']]=copy.deepcopy(fact['dimensions'])
            snapshot=read_workbook(path);materialize_candidate_blanks(snapshot)
            item['excluded']=[{'sheet':s['name'],'cell':address,'category':'label','reason':'测试表的标题或留白'}
                              for s in snapshot['sheets'] for address in s['cells']
                              if s['name']+'!'+address!=fact['source_reference']['location']]
            item['coverage']=workbook_coverage(item,snapshot);item['complete']=item['coverage']['complete']
            mapping=root/(side+'原识别.json');mapping.write_text(json.dumps(item,ensure_ascii=False),encoding='utf-8-sig');paths.append(mapping)
        standard=read_standard_v2(gold)
        a,b=[load_fact_mapping(path,standard) for path in paths]
        before=copy.deepcopy(a);hashes={p:file_hash(p) for p in [*paths,Path(a['source_path']),Path(b['source_path'])]}
        purpose={'reason':'用户已确认以A为模板，具体类别、日期和金额随B更新',
                 'bindings':[{'target':'附注!D8','source':'来源!J3'}],
                 'labels':[{'target':'附注!A8','source':'来源!A3'}, {'target':'附注!D7','source':'来源!J2','period_source':'来源!J3'}]}
        config={'output_dir':str(root/'正式入口'),'gold_path':str(gold),'a_mapping':str(paths[0]),
                'b_mapping':str(paths[1]),'update_purpose':purpose}
        matched=run_step('match',config)
        self.assertEqual(matched['status'],'complete',matched)
        written=run_step('write',{**config,'update_plan':matched['files']['update_plan']})
        self.assertEqual(written['status'],'complete',written)
        result=json.loads(Path(written['files']['a_prime_mapping']).read_text(encoding='utf-8-sig'))
        book=load_workbook(written['files']['a_prime']);sheet=book['附注']
        self.assertEqual([sheet[a].value for a in ['A8','D7','D8']],['乙运输工具','2025年12月31日\n年末余额',900]);book.close()
        self.assertNotIn('facts',result)
        self.assertNotIn('recognition_provenance',result)
        self.assertEqual(result['collections']['A']['mapping_hash'],hashes[paths[0]])
        self.assertEqual(result['collections']['B']['mapping_hash'],hashes[paths[1]])
        self.assertEqual(a,before)
        self.assertEqual({p:file_hash(p) for p in hashes},hashes)
        missing_labels=copy.deepcopy(purpose);missing_labels['labels']=[]
        with self.assertRaisesRegex(ValueError,'维度.*文字'):
            build_fact_update_plan(a,b,standard,update_purpose=missing_labels)
        wrong_place=copy.deepcopy(purpose);wrong_place['labels'][0]['target']='附注!A9'
        with self.assertRaisesRegex(ValueError,'维度.*文字'):
            build_fact_update_plan(a,b,standard,update_purpose=wrong_place)
        altered=copy.deepcopy(b);altered['facts'][0]['metric_id']='另一个计量槽位'
        different=copy.deepcopy(standard);different['metric']['另一个计量槽位']=copy.deepcopy(different['metric']['固定资产原值'])
        different['metric']['另一个计量槽位']['id']='另一个计量槽位'
        different['metric']['另一个计量槽位']['definition']['measurement_basis']='累计折旧'
        with self.assertRaisesRegex(ValueError,'不能改变金标准槽位'):
            build_fact_update_plan(a,altered,different,update_purpose=purpose)


if __name__=='__main__':unittest.main()
