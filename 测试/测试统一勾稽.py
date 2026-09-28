# coding: utf-8
"""统一事实生成更新计划；位置变化、业务不符、冲突和空白的行为核验。"""
import copy
import importlib
import unittest
import json
import hashlib
import uuid
from pathlib import Path
from openpyxl import Workbook

from 附注更新.统一语义 import validate_records
from 测试.测试统一语义 import records, fact


def mapping(location, value=125, scale=1):
    row = fact(); row['raw_value'] = value
    row['raw_value_state'] = 'blank' if value is None else 'value'
    row['dimensions']['unit_scale'] = scale
    row['source_reference']['location'] = location
    row.update(id=location, owner_reference='对象')
    return {'schema_version': 2, 'gold_hash': 'a'*64, 'source_hash': 'b'*64,
            'facts': [row], 'owners': {'对象': copy.deepcopy(row['dimensions'])}, 'unresolved': []}


class MatchTests(unittest.TestCase):
    def setUp(self):
        self.api = importlib.import_module('附注更新.统一勾稽')
        self.standard = validate_records(records()); self.standard['sha256'] = 'a'*64

    def test_different_units_preserve_value_without_conversion_or_blocking(self):
        plan = self.api.build_fact_update_plan(mapping('附注!D8'), mapping('新布局!Z19', 2, 10000), self.standard)
        self.assertEqual(plan['updates'][0]['value'], 2)
        self.assertEqual(plan['issues'], [])

    def test_semantics_find_moved_source_and_preserve_value(self):
        plan = self.api.build_fact_update_plan(mapping('附注!D8'), mapping('新布局!Z19', -2.5), self.standard)
        self.assertEqual(plan['updates'][0]['value'], -2.5)
        self.assertEqual(plan['updates'][0]['cell'], 'D8')
        self.assertEqual(plan['updates'][0]['sources'][0]['cell'], 'Z19')
        self.assertEqual(plan['issues'], [])

    def test_same_position_does_not_override_different_entity(self):
        source = mapping('附注!D8')
        source['facts'][0]['dimensions']['entity'] = '另一企业'
        source['owners']['对象']['entity'] = '另一企业'
        plan = self.api.build_fact_update_plan(mapping('附注!D8'), source, self.standard)
        self.assertEqual(plan['updates'], [])
        self.assertEqual(plan['issues'][0]['kind'], 'missing_source')

    def test_conflicting_same_meaning_sources_are_not_chosen_or_summed(self):
        source = mapping('来源!B2', 10)
        other = copy.deepcopy(source['facts'][0]); other['id'] = '来源!B3'
        other['source_reference']['location'] = '来源!B3'; other['raw_value'] = 20
        source['facts'].append(other)
        plan = self.api.build_fact_update_plan(mapping('附注!D8'), source, self.standard)
        self.assertEqual(plan['updates'], [])
        self.assertEqual(plan['issues'][0]['kind'], 'conflicting_sources')

    def test_blank_does_not_clear_existing_amount(self):
        plan = self.api.build_fact_update_plan(mapping('附注!D8'), mapping('来源!B2', None), self.standard)
        self.assertEqual(plan['updates'], [])
        self.assertEqual(plan['issues'][0]['kind'], 'source_unknown_blank')

    def test_standard_and_source_identity_must_match(self):
        for field in ('gold_hash', 'source_hash'):
            source = mapping('来源!B2'); source[field] = 'c'*64
            with self.subTest(field=field), self.assertRaises(ValueError):
                self.api.build_fact_update_plan(mapping('附注!D8'), source, self.standard)

    def test_duplicate_target_location_is_rejected(self):
        target = mapping('附注!D8')
        other = copy.deepcopy(target['facts'][0]); other['id'] = '另一编号同格'
        target['facts'].append(other)
        with self.assertRaises(ValueError):
            self.api.build_fact_update_plan(target, mapping('来源!B2'), self.standard)

    def file_fixture(self):
        root=Path(__file__).resolve().parents[1]/'测试结果'/('统一勾稽文件测试_'+uuid.uuid4().hex)
        root.mkdir()
        gold=root/'标准.jsonl'
        gold.write_text('\n'.join(json.dumps(r,ensure_ascii=False) for r in records()),encoding='utf-8')
        gold_hash=hashlib.sha256(gold.read_bytes()).hexdigest()
        paths=[]
        for name,address,value in [('附注','D8',125),('来源','J3',900)]:
            workbook=Workbook(); sheet=workbook.active; sheet.title=name
            sheet[address]=value; sheet['A1']='合成企业甲；单户；固定资产原值；2025年12月31日；人民币元'
            source=root/(name+'.xlsx'); workbook.save(source); workbook.close()
            item=mapping(name+'!'+address,value); item.update(source_path=str(source),source_hash=hashlib.sha256(source.read_bytes()).hexdigest(),gold_path=str(gold),gold_hash=gold_hash)
            item['facts'][0]['source_reference']['file_hash']=item['source_hash']
            item['facts'][0]['dimension_evidence']={k:[name+'!A1'] for k in item['facts'][0]['dimensions']}
            path=root/(name+'映射.json'); path.write_text(json.dumps(item,ensure_ascii=False),encoding='utf-8-sig'); paths.append(path)
        return root,gold,paths

    def test_saved_mapping_values_are_checked_against_source_file(self):
        root,gold,paths=self.file_fixture()
        from 附注更新.统一语义 import read_standard_v2
        standard=read_standard_v2(gold)
        self.api.load_fact_mapping(paths[0],standard)
        forged=json.loads(paths[0].read_text(encoding='utf-8-sig'))
        forged['facts'][0]['raw_value']=999
        bad=root/'错误原值映射.json'; bad.write_text(json.dumps(forged,ensure_ascii=False),encoding='utf-8-sig')
        with self.assertRaisesRegex(ValueError,'原值'):
            self.api.load_fact_mapping(bad,standard)

    def test_official_match_entry_reads_two_saved_v2_maps(self):
        root,gold,paths=self.file_fixture()
        from 附注更新.分步流程 import run_step
        result=run_step('match',{'output_dir':str(root/'正式入口'),'gold_path':str(gold),'a_mapping':str(paths[0]),'b_mapping':str(paths[1])})
        self.assertEqual(result['status'],'partial',result['message'])
        plan=json.loads(Path(result['files']['update_plan']).read_text(encoding='utf-8-sig'))
        self.assertEqual(plan['updates'][0]['value'],900)
        self.assertEqual(plan['updates'][0]['sources'][0]['cell'],'J3')
        self.assertFalse(plan['complete'])
        self.assertTrue(any(i['kind']=='coverage_not_verified' for i in plan['issues']))

    def complete_fixture_mapping(self,path):
        from 附注更新.表格 import read_workbook,materialize_candidate_blanks
        from 附注更新.统一识别 import workbook_coverage
        item=json.loads(path.read_text(encoding='utf-8-sig'))
        snapshot=read_workbook(item['source_path'],context=item.get('source_context',[]))
        materialize_candidate_blanks(snapshot)
        positions={row['source_reference']['location'] for row in item['facts']}
        item['excluded']=[{'sheet':sheet['name'],'cell':address,'kind':'non_business','category':'label',
                           'reason':'合成夹具中明确的标题或留白'}
                          for sheet in snapshot['sheets'] for address in sheet['cells']
                          if sheet['name']+'!'+address not in positions]
        item['coverage']=workbook_coverage(item,snapshot);item['complete']=True
        path.write_text(json.dumps(item,ensure_ascii=False),encoding='utf-8-sig')
        return item

    def test_v2_write_saves_reusable_mapping_and_rejects_changed_meaning(self):
        from 附注更新.分步流程 import run_step
        from 附注更新.统一语义 import read_standard_v2
        from openpyxl import load_workbook
        root,gold,paths=self.file_fixture()
        originals=[self.complete_fixture_mapping(path) for path in paths]
        config={'output_dir':str(root/'正式入口'),'gold_path':str(gold),
                'a_mapping':str(paths[0]),'b_mapping':str(paths[1])}
        matched=run_step('match',config)
        self.assertEqual(matched['status'],'complete',matched['message'])
        written=run_step('write',{**config,**matched['files']})
        self.assertEqual(written['status'],'complete',written['message'])
        standard=read_standard_v2(gold)
        derived_path=Path(written['files']['a_prime_mapping'])
        derived=self.api.load_fact_mapping(derived_path,standard)
        self.assertEqual(derived['facts'][0]['raw_value'],900)
        self.assertEqual(derived['facts'][0]['dimensions'],originals[0]['facts'][0]['dimensions'])
        self.assertTrue(derived['coverage']['complete'])
        workbook=load_workbook(written['files']['a_prime']);self.assertEqual(workbook['附注']['D8'].value,900);workbook.close()
        for original in originals:
            self.assertEqual(hashlib.sha256(Path(original['source_path']).read_bytes()).hexdigest(),original['source_hash'])
        again=run_step('match',{**config,'a_mapping':str(derived_path)})
        self.assertEqual(again['status'],'complete',again['message'])
        second=run_step('write',{**config,'a_mapping':str(derived_path),**again['files']})
        self.assertEqual(second['status'],'complete',second['message'])
        self.api.load_fact_mapping(second['files']['a_prime_mapping'],standard)
        forged=json.loads(derived_path.read_text(encoding='utf-8-sig'))
        forged['facts'][0]['dimensions']['entity']='错误企业'
        forged['owners']['对象']['entity']='错误企业'
        changed=root/'更改了语义的派生映射.json';changed.write_text(json.dumps(forged,ensure_ascii=False),encoding='utf-8-sig')
        with self.assertRaisesRegex(ValueError,'派生'):
            self.api.load_fact_mapping(changed,standard)
        changed_book=root/'改了主体表头.xlsx'
        workbook=load_workbook(written['files']['a_prime']);workbook['附注']['A1']='另一企业的资料'
        workbook.save(changed_book);workbook.close()
        forged=json.loads(derived_path.read_text(encoding='utf-8-sig'))
        digest=hashlib.sha256(changed_book.read_bytes()).hexdigest()
        forged.update(source_path=str(changed_book),source_hash=digest)
        forged['facts'][0]['source_reference']['file_hash']=digest
        changed=root/'表头变化的派生映射.json';changed.write_text(json.dumps(forged,ensure_ascii=False),encoding='utf-8-sig')
        with self.assertRaisesRegex(ValueError,'清单外的改动'):
            self.api.load_fact_mapping(changed,standard)

    def test_v2_word_write_preserves_other_chapters_and_reuses_new_link(self):
        from docx import Document
        from openpyxl import load_workbook
        from 附注更新.分步流程 import run_step
        from 附注更新.统一语义 import read_standard_v2
        root,gold,paths=self.file_fixture()
        word=root/'合成附注.docx';doc=Document()
        doc.add_paragraph('五、其他说明')
        outside=doc.add_table(rows=1,cols=1);outside.cell(0,0).text='章前资料原样保留123'
        doc.add_paragraph('六、财务报表重要项目的说明')
        doc.add_paragraph('合成企业甲；单户；人民币元；2025年12月31日')
        table=doc.add_table(rows=2,cols=2);table.style='Table Grid'
        for row,values in zip(table.rows,[['项目','期末原值'],['固定资产原值','125.00']]):
            for cell,value in zip(row.cells,values):cell.text=value
        doc.add_paragraph('七、其他事项')
        after=doc.add_table(rows=1,cols=1);after.cell(0,0).text='章后资料原样保留456'
        doc.save(word);original_hash=hashlib.sha256(word.read_bytes()).hexdigest()
        config={'output_dir':str(root/'正式入口'),'gold_path':str(gold),'word_path':str(word)}
        extracted=run_step('extract_a',config)
        self.assertEqual(extracted['status'],'complete',extracted['message'])
        original=json.loads(paths[0].read_text(encoding='utf-8-sig'))
        a_path=Path(extracted['files']['a_path']);book=load_workbook(a_path)
        sheet=book.worksheets[0];name=sheet.title;value=sheet['C3'].value;book.close()
        original.update(source_path=str(a_path),source_hash=hashlib.sha256(a_path.read_bytes()).hexdigest(),a_link=extracted['files']['a_link'])
        row=original['facts'][0];row.update(id=name+'!C3',raw_value=value)
        row['source_reference']={'file_hash':original['source_hash'],'location':name+'!C3'}
        row['dimension_evidence']={k:[name+'!B1'] for k in row['dimensions']}
        paths[0].write_text(json.dumps(original,ensure_ascii=False),encoding='utf-8-sig')
        for path in paths:self.complete_fixture_mapping(path)
        config.update(a_mapping=str(paths[0]),b_mapping=str(paths[1]))
        matched=run_step('match',config)
        self.assertEqual(matched['status'],'complete',matched['message'])
        written=run_step('write',{**config,**matched['files']})
        self.assertEqual(written['status'],'complete',written['message'])
        updated=Document(written['files']['word'])
        self.assertEqual(updated.tables[1].cell(1,1).text,'900.00')
        for index in (0,2):self.assertEqual(updated.tables[index]._tbl.xml,doc.tables[index]._tbl.xml)
        def format_properties(node):
            # 表格描述承载Excel关联；属性顺序不影响Word格式。
            return (node.tag,dict(node.attrib),[format_properties(child) for child in node
                    if not child.tag.endswith('}tblDescription')])
        self.assertEqual(format_properties(updated.tables[1]._tbl.tblPr),format_properties(doc.tables[1]._tbl.tblPr))
        self.assertEqual(format_properties(updated.tables[1].cell(1,1)._tc.tcPr),format_properties(doc.tables[1].cell(1,1)._tc.tcPr))
        self.assertEqual([p.text for p in updated.paragraphs],[p.text for p in doc.paragraphs])
        self.assertEqual(hashlib.sha256(word.read_bytes()).hexdigest(),original_hash)
        derived=self.api.load_fact_mapping(written['files']['a_prime_mapping'],read_standard_v2(gold))
        self.assertEqual(derived['a_link'],written['files']['word_link'])
        second_config={**config,'a_mapping':written['files']['a_prime_mapping']}
        matched_again=run_step('match',second_config)
        self.assertEqual(matched_again['status'],'complete',matched_again['message'])
        second=run_step('write',{**second_config,**matched_again['files']})
        self.assertEqual(second['status'],'complete',second['message'])
        self.assertEqual(Document(second['files']['word']).tables[1].cell(1,1).text,'900.00')

    def test_v2_write_recomputes_even_when_edited_plan_has_new_seal(self):
        from 附注更新.分步流程 import run_step
        root,gold,paths=self.file_fixture()
        for path in paths:self.complete_fixture_mapping(path)
        config={'output_dir':str(root/'正式入口'),'gold_path':str(gold),'a_mapping':str(paths[0]),'b_mapping':str(paths[1])}
        result=run_step('match',config);self.assertEqual(result['status'],'complete',result['message'])
        plan_path=Path(result['files']['update_plan']);plan=json.loads(plan_path.read_text(encoding='utf-8-sig'))
        plan['updates'][0]['value']=123456
        plan_path.write_text(json.dumps(plan,ensure_ascii=False),encoding='utf-8-sig')
        plan_path.with_name('更新清单校验.json').write_text(json.dumps({'sha256':hashlib.sha256(plan_path.read_bytes()).hexdigest()}),encoding='utf-8-sig')
        written=run_step('write',{**config,**result['files']})
        self.assertEqual(written['status'],'failed')
        self.assertIn('重新语义取数不一致',written['message'])
        self.assertFalse(list(Path(written['output_dir']).glob('*.xlsx')))

    def formula_fixture(self,formula):
        from openpyxl import load_workbook
        from zipfile import ZipFile
        from xml.etree import ElementTree as ET
        root,gold,paths=self.file_fixture()
        item=json.loads(paths[1].read_text(encoding='utf-8-sig'))
        source=Path(item['source_path']);book=load_workbook(source)
        book['来源']['J3']=formula;book.save(source);book.close()
        # 故意放入过期缓存，验证程序采用本次Excel重算结果。
        with ZipFile(source) as archive:parts={name:archive.read(name) for name in archive.namelist()}
        ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        xml=ET.fromstring(parts['xl/worksheets/sheet1.xml'])
        xml.find(".//s:c[@r='J3']/s:v",ns).text='123456'
        parts['xl/worksheets/sheet1.xml']=ET.tostring(xml,encoding='utf-8',xml_declaration=True)
        with ZipFile(source,'w') as archive:
            for name,data in parts.items():archive.writestr(name,data)
        item['source_hash']=hashlib.sha256(source.read_bytes()).hexdigest()
        item['facts'][0].update(raw_value_state='formula',raw_value=formula)
        item['facts'][0]['source_reference']['file_hash']=item['source_hash']
        paths[1].write_text(json.dumps(item,ensure_ascii=False),encoding='utf-8-sig')
        for path in paths:self.complete_fixture_mapping(path)
        return root,gold,paths

    def test_real_excel_formula_evidence_is_saved_and_used_by_official_write(self):
        from openpyxl import load_workbook
        from 附注更新.分步流程 import run_step
        from 附注更新.统一语义 import read_standard_v2
        root,gold,paths=self.formula_fixture('=IF(1=1,900,0)')
        item=json.loads(paths[1].read_text(encoding='utf-8-sig'))
        original_hash=item['source_hash']
        cached=load_workbook(item['source_path'],data_only=True)
        self.assertEqual(cached['来源']['J3'].value,123456);cached.close()
        self.assertTrue(hasattr(self.api,'calculate_mapping_formulas'),'统一流程缺少公式结果依据')
        self.api.calculate_mapping_formulas(item,root/'公式计算')
        self.assertEqual(item['facts'][0]['raw_value'],'=IF(1=1,900,0)')
        self.assertEqual(item['formula_values']['来源!J3']['value'],900)
        self.assertEqual(hashlib.sha256(Path(item['source_path']).read_bytes()).hexdigest(),original_hash)
        paths[1].write_text(json.dumps(item,ensure_ascii=False),encoding='utf-8-sig')
        loaded=self.api.load_fact_mapping(paths[1],read_standard_v2(gold))
        self.assertEqual(loaded['formula_values']['来源!J3']['value'],900)
        config={'output_dir':str(root/'正式入口'),'gold_path':str(gold),'a_mapping':str(paths[0]),'b_mapping':str(paths[1])}
        matched=run_step('match',config);self.assertEqual(matched['status'],'complete',matched['message'])
        plan=json.loads(Path(matched['files']['update_plan']).read_text(encoding='utf-8-sig'))
        self.assertEqual(plan['updates'][0]['value'],900)
        self.assertEqual(plan['updates'][0]['sources'][0]['formula'],'=IF(1=1,900,0)')
        written=run_step('write',{**config,**matched['files']})
        self.assertEqual(written['status'],'complete',written['message'])
        self.api.load_fact_mapping(written['files']['a_prime_mapping'],read_standard_v2(gold))
        bad=copy.deepcopy(item);bad['formula_values']['来源!J3']['value']=123456
        changed=root/'更改计算结果.json';changed.write_text(json.dumps(bad,ensure_ascii=False),encoding='utf-8-sig')
        with self.assertRaisesRegex(ValueError,'计算结果'):
            self.api.load_fact_mapping(changed,read_standard_v2(gold))

    def test_formula_error_remains_unresolved_instead_of_zero(self):
        from 附注更新.统一语义 import read_standard_v2
        root,gold,paths=self.formula_fixture('=1/0')
        item=json.loads(paths[1].read_text(encoding='utf-8-sig'))
        self.assertTrue(hasattr(self.api,'calculate_mapping_formulas'),'统一流程缺少公式结果依据')
        self.api.calculate_mapping_formulas(item,root/'公式计算')
        self.assertEqual(item['formula_values']['来源!J3']['state'],'error')
        target=json.loads(paths[0].read_text(encoding='utf-8-sig'))
        plan=self.api.build_fact_update_plan(target,item,read_standard_v2(gold))
        self.assertEqual(plan['updates'],[])
        self.assertTrue(any(row['kind']=='unevaluated_source_formula' for row in plan['issues']))

    def test_empty_formula_result_cannot_erase_existing_amount(self):
        from 附注更新.统一语义 import read_standard_v2
        root,gold,paths=self.formula_fixture('=IF(1=0,900,"")')
        item=json.loads(paths[1].read_text(encoding='utf-8-sig'))
        self.api.calculate_mapping_formulas(item,root/'公式计算')
        self.assertEqual(item['formula_values']['来源!J3']['state'],'blank')
        target=json.loads(paths[0].read_text(encoding='utf-8-sig'))
        plan=self.api.build_fact_update_plan(target,item,read_standard_v2(gold))
        self.assertEqual(plan['updates'],[])
        self.assertTrue(any(row['kind']=='source_unknown_blank' for row in plan['issues']))

    def test_calculation_failure_preserves_recognized_mapping(self):
        from unittest.mock import patch
        from 附注更新.统一识别 import UnifiedSemanticEngine
        from 附注更新.分步流程 import run_step
        root,gold,paths=self.formula_fixture('=SUM(800,100)')
        item=json.loads(paths[1].read_text(encoding='utf-8-sig'));item['usage']={}
        config={'output_dir':str(root/'正式入口'),'gold_path':str(gold),'b_path':item['source_path'],
                'base_url':'http://unused.example/v1','model':'受控测试模型','api_key':'','timeout':600}
        with patch.object(UnifiedSemanticEngine,'recognize_workbook',return_value=copy.deepcopy(item)), \
             patch.object(self.api,'calculate_mapping_formulas',side_effect=OSError('Excel暂不可用')):
            result=run_step('recognize_b',config)
        self.assertEqual(result['status'],'partial',result['message'])
        saved=json.loads(Path(result['files']['b_mapping']).read_text(encoding='utf-8-sig'))
        self.assertEqual(saved['facts'],item['facts'])
        self.assertEqual(len(saved['calculation_issues']),1)
        self.assertIn('Excel暂不可用',saved['calculation_issues'][0]['reason'])
        self.assertTrue(Path(result['files']['b_review']).is_file())

    def test_additive_standard_reuses_existing_maps_and_derived_history(self):
        from 附注更新.金标准 import publish_revision,preview_revision,ensure_compatible
        from 附注更新.统一语义 import read_standard_v2
        from 附注更新.分步流程 import run_step
        root,gold,paths=self.file_fixture()
        for path in paths:self.complete_fixture_mapping(path)
        candidate=records();added=copy.deepcopy(candidate[-1]);added['id']='固定资产账面净值'
        added['definition']['measurement_basis']='账面净值';candidate.append(added)
        new=root/'扩充标准.jsonl';new.write_text('\n'.join(json.dumps(row,ensure_ascii=False) for row in candidate),encoding='utf-8-sig')
        published=publish_revision(gold,new,root=root,expected_hashes=preview_revision(gold,new)['source_hashes'])
        original=json.loads(paths[0].read_text(encoding='utf-8-sig'))
        self.assertTrue(ensure_compatible(original,published['path']))
        config={'output_dir':str(root/'正式入口'),'gold_path':published['path'],'a_mapping':str(paths[0]),'b_mapping':str(paths[1])}
        matched=run_step('match',config);self.assertEqual(matched['status'],'complete',matched['message'])
        written=run_step('write',{**config,**matched['files']});self.assertEqual(written['status'],'complete',written['message'])
        derived=self.api.load_fact_mapping(written['files']['a_prime_mapping'],read_standard_v2(published['path']))
        self.assertEqual(derived['gold_hash'],published['sha256'])
        self.assertEqual(json.loads(paths[0].read_text(encoding='utf-8-sig'))['gold_hash'],original['gold_hash'])
        candidate.append({'record_type':'dimension','id':'asset_location','meaning':'实际资产所在地','value_type':'text','domain':'资产所在地'})
        next_gold=root/'再次扩充标准.jsonl';next_gold.write_text('\n'.join(json.dumps(row,ensure_ascii=False) for row in candidate),encoding='utf-8-sig')
        self.api.load_fact_mapping(written['files']['a_prime_mapping'],read_standard_v2(next_gold))
        changed=copy.deepcopy(candidate);changed_metric=next(row for row in changed if row['id']=='固定资产原值')
        changed_metric['definition']['measurement_basis']='另一计量基础'
        bad=root/'改变既有含义.jsonl';bad.write_text('\n'.join(json.dumps(row,ensure_ascii=False) for row in changed),encoding='utf-8-sig')
        with self.assertRaisesRegex(ValueError,'含义|定义|识别'):
            self.api.load_fact_mapping(paths[0],read_standard_v2(bad))


if __name__ == '__main__':
    unittest.main()
