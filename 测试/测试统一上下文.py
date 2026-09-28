# coding: utf-8
"""章节说明与人工业务说明只证明语义，不产生财务数值。"""
import copy
import json
import unittest
import uuid
from pathlib import Path
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from 测试.测试统一语义 import records, fact
from 附注更新.表格 import file_hash
from 附注更新.统一识别 import UnifiedSemanticEngine, build_payload, validate_response
from 附注更新.统一勾稽 import load_fact_mapping


class ContextTests(unittest.TestCase):
    def setUp(self):
        self.folder=Path(__file__).resolve().parents[1]/'测试结果'/('统一上下文_'+uuid.uuid4().hex)
        self.folder.mkdir()
        self.source=self.folder/'章节表格.xlsx'
        book=Workbook();sheet=book.active;sheet.title='附注'
        sheet.append(['项目','期末余额']);sheet.append(['固定资产原值',-125])
        sheet['A6']='其他表格';sheet['B6']=10
        book.defined_names.add(DefinedName('原表一',attr_text="'附注'!$A$1:$B$2"))
        book.save(self.source);book.close()
        self.context=[{'table_id':'word#1','range_name':'原表一','sheet':'附注',
                      'source':str(self.source),
                      'first_row':1,'first_column':1,'row_count':2,'column_count':2,
                      'chapter_context':['2025年12月31日，单户报表，人民币元，全部固定资产。']},
                      '用户确认：模板中的企业名称只作格式来源，目标主体为合成企业甲。',
                      {'sheet':'附注','first_row':6,'first_column':1,'row_count':1,'column_count':2,
                       'chapter_context':['本表日期为2024年12月31日。']}]
        gold=self.folder/'金标准.jsonl'
        gold.write_text('\n'.join(json.dumps(row,ensure_ascii=False) for row in records()),encoding='utf-8-sig')
        self.engine=UnifiedSemanticEngine({'model':'本地行为测试','base_url':'https://example.invalid/v1'},gold,self.folder/'回执')
        self.packet={'source_path':str(self.source),'source_hash':file_hash(self.source),'sheet':'附注',
                     'cells':{'A1':{'value':'项目'},'B1':{'value':'期末余额'},'A2':{'value':'固定资产原值'},'B2':{'value':-125}},
                     'candidate_cells':['B2'],'physical_table':[1,1,2,2],'document_context':self.context}

    def response(self,payload):
        refs=payload['context_evidence']
        chapter=next(key for key,row in refs.items() if isinstance(row['content'],dict))
        user=next(key for key,row in refs.items() if isinstance(row['content'],str))
        dims=fact()['dimensions']
        return {'cells':{'B2':{'kind':'metric_fact','metric_id':'固定资产原值','dimensions':dims,
                'dimension_evidence':{key:[user if key=='entity' else chapter] for key in dims},'reason':'原表标题及章节说明明确'}}}

    def test_actual_http_shared_dimensions_keep_original_reply_and_full_mapping(self):
        from io import BytesIO
        from unittest.mock import patch
        from 附注更新.回执格式 import normalize_response
        from 附注更新.统一语义 import definition_hash
        requests=[]
        def respond(request,timeout):
            body=json.loads(request.data);payload=json.loads(body['messages'][1]['content']);requests.append(payload)
            self.assertEqual(payload['response_contract'],'unified_dimension_refs_v1')
            self.assertNotIn('chat_template_kwargs',body)
            full=self.response(payload);row=full['cells']['B2'];pool=[]
            for key,value in row.pop('dimensions').items():
                pool.append({'dimension_id':key,'value':value,'evidence':row['dimension_evidence'][key]})
            row.pop('dimension_evidence');row['dimension_refs']=list(range(len(pool)))
            wire={**full,'dimension_values':pool}
            return BytesIO(json.dumps({'choices':[{'finish_reason':'stop','message':{'content':json.dumps(wire)}}],
                                      'usage':{'prompt_tokens':10,'completion_tokens':20,'total_tokens':30}}).encode())
        with patch('附注更新.语义.urllib.request.build_opener') as opener:
            opener.return_value.open.side_effect=respond
            result=self.engine.classify_packet(self.packet)
        self.assertEqual(result['facts'][0]['raw_value'],-125)
        self.assertEqual(len(requests),2)
        for ref in result['request_receipts']:
            receipt=json.loads(Path(ref['path']).read_text(encoding='utf-8-sig'))
            task=json.loads(Path(ref['path']).with_name('任务.json').read_text(encoding='utf-8-sig'))
            representation=receipt['transport_representations'][0]
            self.assertEqual(normalize_response(task['payload'],representation['raw_result']),receipt['result'])
            self.assertEqual(representation['normalized_hash'],definition_hash(receipt['result']))
        mapped=next(self.engine.work_dir.glob('逐格语义映射_*.json'))
        self.assertEqual(load_fact_mapping(mapped,self.engine.standard)['facts'],result['facts'])
        again=self.engine.classify_packet(self.packet)
        self.assertEqual(again['facts'],result['facts'])
        self.assertEqual(len(requests),2,'再次识别应复用已核验回执，不重复调用模型')
        self.assertEqual(self.engine.usage['cached_requests'],2)

    def test_chapter_and_user_context_are_saved_separately_from_cell_values(self):
        payload=build_payload(self.packet,self.engine.standard)
        self.assertEqual(len(payload['context_evidence']),2)
        raw=self.response(payload)
        result=validate_response(raw,self.packet,self.engine.standard)['B2']['fact']
        self.assertEqual(result['raw_value'],-125)
        self.assertTrue(all(ref.startswith('@context:') for refs in result['dimension_evidence'].values() for ref in refs))
        self.assertNotIn('B6',payload['cells'])
        from 附注更新.回执格式 import _unified_classify
        self.assertIn('@context:',json.dumps(_unified_classify(payload)))

    def test_other_table_context_cannot_prove_this_cells_dimensions(self):
        from 附注更新.统一证据 import context_catalog
        payload=build_payload(self.packet,self.engine.standard);raw=self.response(payload)
        catalog=context_catalog(self.context,self.source)
        other=next(key for key,row in catalog.items() if row['content']==self.context[-1])
        raw['cells']['B2']['dimension_evidence']['period']=[other]
        with self.assertRaises(ValueError):validate_response(raw,self.packet,self.engine.standard)

    def test_actual_request_sends_context_once_and_preserves_scoped_evidence(self):
        original=copy.deepcopy(self.packet);before=file_hash(self.source)
        def respond(system,payload):
            self.assertNotIn('document_context',payload,'正文已完整存在证据目录，不应再次发送')
            contexts=list(payload['context_evidence'].values())
            self.assertEqual(len(contexts),2)
            chapter=next(row for row in contexts if row['kind']=='document_context')
            user=next(row for row in contexts if row['kind']=='user_context')
            expected=copy.deepcopy(self.context[0]);expected.pop('source')
            self.assertEqual(chapter['content'],expected)
            self.assertEqual(chapter['scope'],{'sheet':'附注','bounds':[1,1,2,2]})
            self.assertEqual(user['content'],'用户确认：模板中的企业名称只作格式来源，目标主体为合成企业甲。')
            self.assertIsNone(user['scope'])
            self.assertEqual(payload['cells']['B2']['value'],-125)
            return self.response(payload)
        self.engine._http=respond
        self.engine.classify_packet(self.packet)
        saved=next(self.engine.work_dir.glob('逐格语义映射_*.json'))
        result=load_fact_mapping(saved,self.engine.standard)
        self.assertEqual(result['facts'][0]['raw_value'],-125)
        self.assertEqual(result['source_context'],self.context)
        self.assertEqual(result['facts'][0]['dimensions']['entity'],'合成企业甲')
        self.assertEqual(len(result['request_receipts']),2)
        self.assertEqual(self.packet,original)
        self.assertEqual(file_hash(self.source),before)

    def test_header_and_unresolved_can_cite_the_same_scoped_original_context(self):
        payload=build_payload(self.packet,self.engine.standard)
        reference=next(key for key,row in payload['context_evidence'].items() if row['kind']=='document_context')
        packet=copy.deepcopy(self.packet);packet['candidate_cells']=['B1','B2']
        raw={'cells':{'B1':{'kind':'non_business','category':'header','reason':'章节说明解释期末所指日期','evidence_cells':[reference]},
                      'B2':{'kind':'unresolved','reason':'保留未决语义及相关说明供进一步复核','evidence_cells':[reference]}}}
        checked=validate_response(raw,packet,self.engine.standard)
        self.assertEqual(checked['B1']['kind'],'non_business');self.assertEqual(checked['B2']['kind'],'unresolved')
        from 附注更新.回执格式 import _unified_classify
        schema=_unified_classify(payload)
        for branch in schema['$defs']['unified_record']['anyOf'][-2:]:
            evidence=branch['properties']['evidence_cells']
            if '$ref' in evidence:evidence=schema['$defs'][evidence['$ref'].rsplit('/',1)[-1]]
            self.assertIn(reference,evidence['items']['enum'])

    def test_independent_rounds_overlap_and_keep_their_own_receipts_and_usage(self):
        from threading import Barrier
        from unittest.mock import patch
        barrier=Barrier(2)
        def respond(instance,system,payload):
            instance.usage['requests']+=1
            instance.usage['prompt_tokens']+=7 if payload['round']==1 else 11
            barrier.wait(timeout=2)
            return self.response(payload)
        with patch.object(UnifiedSemanticEngine,'_http',respond):
            result=self.engine.classify_packet(self.packet)
        self.assertEqual(len(result['facts']),1)
        self.assertEqual(result['facts'][0]['raw_value'],-125)
        self.assertEqual(self.engine.usage['requests'],2)
        self.assertEqual(self.engine.usage['prompt_tokens'],18)
        self.assertEqual(len(result['request_receipts']),2)
        rounds=[];identities=[]
        for record in result['request_receipts']:
            path=Path(record['path']);task=json.loads(path.with_name('任务.json').read_text(encoding='utf-8-sig'))
            receipt=json.loads(path.read_text(encoding='utf-8-sig'))
            rounds.append(task['payload']['round']);identities.append(task['task_id'])
            self.assertEqual(receipt['task_id'],task['task_id'])
        self.assertEqual(rounds,[1,2]);self.assertEqual(len(set(identities)),2)
        saved=next(self.engine.work_dir.glob('逐格语义映射_*.json'))
        self.assertEqual(load_fact_mapping(saved,self.engine.standard)['facts'],result['facts'])

    def test_failed_parallel_round_keeps_other_receipt_without_confirming_a_fact(self):
        from threading import Barrier
        from unittest.mock import patch
        from 附注更新.语义 import SemanticError
        barrier=Barrier(2)
        def respond(instance,system,payload):
            instance.usage['requests']+=1
            barrier.wait(timeout=2)
            if payload['round']==2:raise SemanticError('受控连接失败')
            return self.response(payload)
        with patch.object(UnifiedSemanticEngine,'_http',respond):
            with self.assertRaisesRegex(SemanticError,'受控连接失败'):
                self.engine.classify_packet(self.packet)
        self.assertEqual(self.engine.usage['requests'],2)
        self.assertEqual(len(self.engine._review_receipts),1)
        task=json.loads(self.engine._review_receipts[0].with_name('任务.json').read_text(encoding='utf-8-sig'))
        self.assertEqual(task['payload']['round'],1)
        self.assertEqual(list(self.engine.work_dir.glob('逐格语义映射_*.json')),[])

    def test_named_table_scope_follows_actual_workbook_range(self):
        from 附注更新.统一证据 import context_catalog, applicable_context
        from openpyxl import load_workbook
        book=load_workbook(self.source)
        book.defined_names['原表一'].attr_text="'附注'!$D$10:$E$11"
        shifted=self.folder/'移动后的表格.xlsx';book.save(shifted);book.close()
        catalog=context_catalog(self.context,shifted)
        chapter=next(key for key,row in catalog.items() if row['content']==self.context[0])
        self.assertIn(chapter,applicable_context(catalog,'附注',[10,4,11,5]))
        self.assertNotIn(chapter,applicable_context(catalog,'附注',[1,1,2,2]))

    def test_saved_mapping_reloads_and_rejects_replaced_context(self):
        self.engine._http=lambda system,payload:self.response(payload)
        result=self.engine.classify_packet(self.packet)
        path=next(self.engine.work_dir.glob('逐格语义映射_*.json'))
        self.assertEqual(load_fact_mapping(path,self.engine.standard)['facts'],result['facts'])
        changed=copy.deepcopy(result);changed['source_context'][1]='目标主体为另一个企业'
        other=self.folder/'变更说明.json';other.write_text(json.dumps(changed,ensure_ascii=False),encoding='utf-8-sig')
        with self.assertRaises(ValueError):load_fact_mapping(other,self.engine.standard)

    def test_same_cell_equivalence_can_review_context_evidence(self):
        def respond(system,payload):
            if payload['task']=='unified_equivalence':
                return {'pairs':{key:{'decision':'equivalent','reason':'双方都明确全部固定资产',
                        'left_evidence':list(row['left']['evidence']),'right_evidence':list(row['right']['evidence'])}
                        for key,row in payload['pairs'].items()}}
            result=self.response(payload)
            if payload['round']==2:
                result['cells']['B2']['dimensions']['asset_selection']={'domain':'固定资产','mode':'predicate',
                        'predicate':'全部固定资产','completeness':'complete'}
            return result
        self.engine._http=respond
        result=self.engine.classify_packet(self.packet)
        self.assertEqual(len(result['facts']),1);self.assertTrue(result['equivalence_proofs'])
        path=next(self.engine.work_dir.glob('逐格语义映射_*.json'))
        self.assertEqual(load_fact_mapping(path,self.engine.standard)['facts'][0]['raw_value'],-125)
        changed=copy.deepcopy(result);changed['equivalence_proofs'][0]['pairs'][0]['left_context']=[]
        other=self.folder/'缺少复核上下文.json';other.write_text(json.dumps(changed,ensure_ascii=False),encoding='utf-8-sig')
        with self.assertRaises(ValueError):load_fact_mapping(other,self.engine.standard)

    def test_standard_proposal_can_cite_scoped_chapter_text(self):
        from 附注更新.统一补充 import _proposal
        payload=build_payload(self.packet,self.engine.standard)
        reference=next(key for key,row in payload['context_evidence'].items() if isinstance(row['content'],dict))
        metric=copy.deepcopy(records()[-1]);metric.pop('source_refs')
        metric['id']='固定资产累计折旧';metric['definition']['measurement_basis']='累计已计提折旧'
        proposal={'decisions':[{'cell':'B2','decision':'revision','reason':'独立计量基础',
                               'record_ids':[metric['id']],'evidence_cells':[reference]}],'records':[metric]}
        payload.update(source_hash=file_hash(self.source),mapping_hash='a'*64,text_evidence_cells=['A1','A2','B1'])
        result=_proposal(proposal,payload,records(),self.folder/'复核记录.json')
        sources=[row for row in result['rows'] if row['record_type']=='source_example' and row['id']!='来源一']
        self.assertEqual(len(sources),1)
        self.assertEqual(sources[0]['origin']['location'],reference)
        self.assertIn('2025年12月31日',sources[0]['text_evidence'][0])
        changed=copy.deepcopy(payload);changed['candidate_cells']=['B6']
        changed_proposal=copy.deepcopy(proposal);changed_proposal['decisions'][0]['cell']='B6'
        with self.assertRaises(ValueError):_proposal(changed_proposal,changed,records(),self.folder/'错误复核.json')

    def test_review_workbook_shows_original_context(self):
        from 附注更新.统一证据 import append_context_sheet
        book=Workbook();append_context_sheet(book,self.context,self.source)
        self.assertIn('语义说明依据',book.sheetnames)
        text=json.dumps(list(book['语义说明依据'].values),ensure_ascii=False)
        self.assertIn('合成企业甲',text);self.assertIn('@context:',text)
        book.close()

    def test_context_survives_independent_mappings_match_write_and_derived_mapping(self):
        from unittest.mock import patch
        from 附注更新.分步流程 import run_step
        originals={};mappings=[]
        for side,value in [('附注表格',118),('更新数据',-125)]:
            source=self.folder/(side+'.xlsx')
            book=Workbook();sheet=book.active;sheet.title='附注'
            sheet.append(['项目','期末余额']);sheet.append(['固定资产原值',value]);book.save(source);book.close()
            originals[str(source)]=file_hash(source)
            engine=UnifiedSemanticEngine(self.engine.settings,self.engine.gold_path,self.folder/(side+'识别'))
            def respond(system,payload,side=side):
                raw=self.response(payload)
                if side=='更新数据':
                    raw['cells']['B2']['dimensions']['asset_selection']={'domain':'固定资产','mode':'predicate',
                            'predicate':'全部固定资产','completeness':'complete'}
                for address in payload['candidate_cells']:
                    if address!='B2':raw['cells'][address]={'kind':'non_business','category':'header','reason':'原表标题',
                                                         'evidence_cells':[address]}
                return raw
            engine._http=respond
            mapping=engine.recognize_workbook(source,context=self.context)
            self.assertTrue(mapping['complete']);load_fact_mapping(mapping['mapping_path'],engine.standard)
            mappings.append(mapping)
        def review(instance,system,payload):
            return {'pairs':{key:{'decision':'equivalent','reason':'双方说明都为全部固定资产',
                    'left_evidence':list(row['left']['evidence']),'right_evidence':list(row['right']['evidence'])}
                    for key,row in payload['pairs'].items()}}
        config={**self.engine.settings,'mode':'online','gold_path':self.engine.gold_path,
                'a_mapping':mappings[0]['mapping_path'],'b_mapping':mappings[1]['mapping_path'],
                'output_dir':str(self.folder/'正式入口'),'api_key':''}
        with patch.object(UnifiedSemanticEngine,'_http',review):matched=run_step('match',config)
        self.assertEqual(matched['status'],'complete',matched)
        with patch.object(UnifiedSemanticEngine,'_http',side_effect=AssertionError('写入只搬运人工提供的值')):
            written=run_step('write',{**config,**matched['files']})
        self.assertEqual(written['status'],'complete',written)
        updated=load_fact_mapping(written['files']['a_prime_mapping'],self.engine.standard)
        self.assertEqual(updated['facts'][0]['raw_value'],-125)
        self.assertEqual(updated['facts'][0]['dimensions'],mappings[0]['facts'][0]['dimensions'])
        self.assertEqual(updated['source_context'],self.context)
        self.assertTrue(all(file_hash(path)==digest for path,digest in originals.items()))


if __name__=='__main__':unittest.main()
