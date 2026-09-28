# coding: utf-8
"""新标准真实分类入口的本地行为验证，不冒充模型识别准确率。"""
import copy
import importlib
import importlib.util
import unittest
import json
import uuid
from pathlib import Path
from openpyxl import Workbook
from 附注更新.统一识别 import UnifiedSemanticEngine

from 附注更新.统一语义 import validate_records
from 测试.测试统一语义 import records, fact


def packet():
    return {"source_hash": "b" * 64, "sheet": "任意布局", "cells": {
        "A1": {"value": "合成企业甲，单户人民币元，2025年12月31日固定资产原值合计"},
        "D7": {"value": 125.0}, "F9": {"value": None}}, "candidate_cells": ["D7", "F9"]}


def answer():
    result = {"cells": {}}
    for address in packet()["candidate_cells"]:
        result["cells"][address] = {"kind": "metric_fact", "metric_id": "固定资产原值",
                                     "dimensions": fact()["dimensions"],
                                     "dimension_evidence": {key: ["A1"] for key in fact()["dimensions"]},
                                     "reason": "合成例原表明确说明范围"}
    return result


class RecognitionTests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(importlib.util.find_spec("附注更新.统一识别"), "缺少统一标准识别入口")
        self.api = importlib.import_module("附注更新.统一识别")
        self.standard = validate_records(records())

    def partial_answer(self):
        raw=answer()
        for row in raw['cells'].values():
            row.update(kind='partial_metric_fact',missing_dimensions=['entity'],evidence_cells=['A1'])
            row['dimensions']=copy.deepcopy(row['dimensions']);row['dimensions'].pop('entity')
            row['dimension_evidence'].pop('entity')
        return raw

    def test_partial_identity_retains_known_meaning_without_matching_or_values(self):
        checked=self.api.validate_response(self.partial_answer(),packet(),self.standard)
        output=self.api.agree_rounds(checked,checked,self.standard)
        self.assertEqual(output['facts'],[])
        self.assertEqual(len(output['unresolved']),2)
        for row in output['unresolved']:
            known=row['partial_semantics']
            self.assertEqual(known['metric_id'],'固定资产原值')
            self.assertEqual(known['missing_dimensions'],['entity'])
            self.assertEqual(known['dimensions']['asset_selection'],fact()['dimensions']['asset_selection'])
            self.assertNotIn('raw_value',known)

    def test_partial_rounds_keep_only_agreed_actual_dimensions(self):
        first=self.api.validate_response(self.partial_answer(),packet(),self.standard)
        other=self.partial_answer()
        for row in other['cells'].values():row['dimensions']['period']['date']='2024-12-31'
        second=self.api.validate_response(other,packet(),self.standard)
        output=self.api.agree_rounds(first,second,self.standard)
        for row in output['unresolved']:
            self.assertEqual(row['partial_semantics']['missing_dimensions'],['entity','period'])
            self.assertNotIn('period',row['partial_semantics']['dimensions'])

    def test_partial_and_full_round_keep_shared_meaning_without_upgrading_missing_information(self):
        partial=self.api.validate_response(self.partial_answer(),packet(),self.standard)
        full=self.api.validate_response(answer(),packet(),self.standard)
        for first,second in ((partial,full),(full,partial)):
            result=self.api.agree_rounds(first,second,self.standard)
            self.assertEqual(result['facts'],[])
            for row in result['unresolved']:
                self.assertEqual(row['partial_semantics']['metric_id'],'固定资产原值')
                self.assertEqual(row['partial_semantics']['missing_dimensions'],['entity'])
                self.assertNotIn('entity',row['partial_semantics']['dimensions'])

    def test_partial_and_complete_range_keep_confirmed_members_with_partial_status(self):
        partial=answer();full=answer()
        for row in partial['cells'].values():
            row.update(kind='partial_metric_fact',missing_dimensions=[],evidence_cells=['A1'])
            row['dimensions']=copy.deepcopy(row['dimensions'])
            row['dimensions']['asset_selection']={'domain':'固定资产','mode':'members','members':['机器设备'],'completeness':'partial'}
        for row in full['cells'].values():
            row['dimensions']=copy.deepcopy(row['dimensions'])
            row['dimensions']['asset_selection']={'domain':'固定资产','mode':'members','members':['机器设备'],'completeness':'complete'}
        a=self.api.validate_response(partial,packet(),self.standard);b=self.api.validate_response(full,packet(),self.standard)
        for first,second in ((a,b),(b,a)):
            result=self.api.agree_rounds(first,second,self.standard);self.assertEqual(result['facts'],[])
            selection=result['unresolved'][0]['partial_semantics']['dimensions']['asset_selection']
            self.assertEqual(selection,partial['cells']['D7']['dimensions']['asset_selection'])

    def test_partial_identity_rejects_false_missing_or_invalid_dimension_and_model_value(self):
        variants=[]
        raw=self.partial_answer();raw['cells']['D7']['missing_dimensions']=[];variants.append(raw)
        raw=self.partial_answer();raw['cells']['D7']['dimensions']['unit_scale']=0;variants.append(raw)
        raw=self.partial_answer();raw['cells']['D7']['raw_value']=123;variants.append(raw)
        raw=self.partial_answer();raw['cells']['D7']['evidence_cells']=['Z99'];variants.append(raw)
        raw=self.partial_answer();raw['cells']['D7']['dimensions']['period']={'kind':'duration','start':'2025-01-01','end':'2025-12-31'};variants.append(raw)
        for raw in variants:
            with self.assertRaises(ValueError):self.api.validate_response(raw,packet(),self.standard)

    def test_full_rounds_with_disputed_range_keep_shared_slot_and_other_dimensions(self):
        left=answer();right=answer()
        for row in right['cells'].values():
            row['dimensions']=copy.deepcopy(row['dimensions'])
            row['dimensions']['asset_selection']={'domain':'固定资产','mode':'members','members':['机器设备'],'completeness':'complete'}
        first=self.api.validate_response(left,packet(),self.standard);second=self.api.validate_response(right,packet(),self.standard)
        result=self.api.agree_rounds(first,second,self.standard)
        self.assertEqual(result['facts'],[]);self.assertEqual(result['owners'],{})
        for row in result['unresolved']:
            known=row['partial_semantics']
            self.assertEqual(known['metric_id'],'固定资产原值')
            self.assertEqual(known['missing_dimensions'],['asset_selection'])
            self.assertEqual(known['dimensions']['entity'],'合成企业甲')
            self.assertNotIn('asset_selection',known['dimensions'])
            self.assertNotIn('raw_value',known)
        legacy=self.api.agree_rounds(first,second,self.standard,agreement_version=1)
        self.assertNotIn('partial_semantics',legacy['unresolved'][0])

    def test_full_rounds_disagreeing_on_slot_do_not_invent_common_slot(self):
        rows=records();other=copy.deepcopy(rows[-1]);other['id']='其他计量';other['definition']['measurement_basis']='另一种计量';rows.append(other)
        standard=validate_records(rows);left=answer();right=answer()
        for row in right['cells'].values():row['metric_id']='其他计量'
        result=self.api.agree_rounds(self.api.validate_response(left,packet(),standard),self.api.validate_response(right,packet(),standard),standard)
        self.assertEqual(result['facts'],[])
        self.assertTrue(all('partial_semantics' not in row for row in result['unresolved']))

    def test_partial_transport_contract_accepts_only_declared_known_dimensions(self):
        from 附注更新.回执格式 import _unified_classify
        import jsonschema
        schema=_unified_classify(self.api.build_payload(packet(),self.standard))
        jsonschema.validate(self.partial_answer(),schema)
        raw=self.partial_answer();raw['cells']['D7']['dimensions']['invented']='猜测'
        with self.assertRaises(jsonschema.ValidationError):jsonschema.validate(raw,schema)

    def test_partial_range_is_preserved_but_never_becomes_a_complete_fact(self):
        raw=answer()
        for row in raw['cells'].values():
            row.update(kind='partial_metric_fact',missing_dimensions=[],evidence_cells=['A1'])
            row['dimensions']=copy.deepcopy(row['dimensions'])
            row['dimensions']['asset_selection']={'domain':'固定资产','mode':'members','members':['机器设备'],'completeness':'partial'}
        checked=self.api.validate_response(raw,packet(),self.standard)
        output=self.api.agree_rounds(checked,checked,self.standard)
        self.assertEqual(output['facts'],[])
        self.assertEqual(output['unresolved'][0]['partial_semantics']['dimensions']['asset_selection']['completeness'],'partial')
        self.assertIn('asset_selection',output['unresolved'][0]['reason'])
        from 附注更新.回执格式 import _unified_classify
        import jsonschema
        jsonschema.validate(raw,_unified_classify(self.api.build_payload(packet(),self.standard)))

    def test_blank_target_reference_is_redundant_not_a_semantic_proof(self):
        raw=self.partial_answer();raw['cells']['F9']['evidence_cells']=['F9','A1']
        raw['cells']['F9']['dimension_evidence']['asset_selection']=['F9','A1']
        checked=self.api.validate_response(raw,packet(),self.standard)
        self.assertEqual(checked['F9']['partial_semantics']['evidence_cells'],['任意布局!A1'])
        raw['cells']['F9']['evidence_cells']=['F9']
        with self.assertRaises(ValueError):self.api.validate_response(raw,packet(),self.standard)

    def test_values_are_reloaded_and_financial_blank_keeps_identity(self):
        checked = self.api.validate_response(answer(), packet(), self.standard)
        self.assertEqual(checked["D7"]["fact"]["raw_value"], 125.0)
        self.assertEqual(checked["F9"]["fact"]["raw_value_state"], "blank")
        self.assertEqual(checked["F9"]["fact"]["metric_id"], "固定资产原值")

    def test_omissions_foreign_addresses_and_model_values_are_rejected(self):
        variants = []
        one = answer(); del one["cells"]["F9"]; variants.append(one)
        one = answer(); one["cells"]["Z99"] = one["cells"].pop("F9"); variants.append(one)
        one = answer(); one["cells"]["D7"]["raw_value"] = 999; variants.append(one)
        for other in variants:
            with self.subTest(other=other), self.assertRaises(ValueError):
                self.api.validate_response(other, packet(), self.standard)

    def test_evidence_must_exist_in_actual_packet(self):
        one = answer(); one["cells"]["D7"]["dimension_evidence"]["period"] = ["Q99"]
        with self.assertRaises(ValueError):
            self.api.validate_response(one, packet(), self.standard)

    def test_matching_rounds_require_same_actual_dimensions(self):
        first = self.api.validate_response(answer(), packet(), self.standard)
        raw = answer(); raw["cells"]["D7"]["dimensions"] = copy.deepcopy(raw["cells"]["D7"]["dimensions"])
        raw["cells"]["D7"]["dimensions"]["entity"] = "合成企业乙"
        second = self.api.validate_response(raw, packet(), self.standard)
        output = self.api.agree_rounds(first, second, self.standard)
        self.assertEqual(len(output["facts"]), 1)
        self.assertEqual(output["unresolved"][0]["cell"], "D7")

    def test_shared_unknown_population_is_not_approved_by_agreement(self):
        raw = answer()
        for row in raw["cells"].values():
            row["dimensions"] = copy.deepcopy(row["dimensions"])
            row["dimensions"]["asset_selection"] = {"domain": "固定资产", "mode": "unknown", "completeness": "unknown"}
        checked = self.api.validate_response(raw, packet(), self.standard)
        output = self.api.agree_rounds(checked, checked, self.standard)
        self.assertEqual(output["facts"], [])
        self.assertEqual(len(output["unresolved"]), 2)

    def test_same_cell_rounds_must_agree_on_representation_scale(self):
        first = self.api.validate_response(answer(), packet(), self.standard)
        raw = answer()
        for row in raw["cells"].values():
            row["dimensions"] = copy.deepcopy(row["dimensions"])
            row["dimensions"]["unit_scale"] = "10000"
        second = self.api.validate_response(raw, packet(), self.standard)
        output = self.api.agree_rounds(first, second, self.standard)
        self.assertEqual(output["facts"], [])
        self.assertEqual(len(output["unresolved"]), 2)

    def test_payload_retains_physical_order_without_classifying_headers_in_code(self):
        self.assertTrue(hasattr(self.api, "build_payload"), "原始布局须通过统一构造入口传给模型")
        data = packet(); data["cells"].update({"A10": {"value": "后面的原文"}, "B2": {"value": "前面的原文"}})
        payload = self.api.build_payload(data, self.standard)
        self.assertEqual([row["row"] for row in payload["source_text_rows"]], [1, 2, 10])
        self.assertEqual(payload["source_text_rows"][1]["cells"], [["B2", "前面的原文"]])
        self.assertEqual(payload["candidate_cells"], ["D7", "F9"])
        self.assertEqual(payload["cells"], data["cells"])

    def test_missing_selection_domain_reports_missing_field_for_correction(self):
        raw = answer()
        for row in raw['cells'].values():
            row['dimensions'] = copy.deepcopy(row['dimensions'])
            del row['dimensions']['asset_selection']['domain']
        with self.assertRaisesRegex(ValueError, '缺少字段.*domain'):
            self.api.validate_response(raw, packet(), self.standard)

    def test_non_business_needs_agreement_and_blank_is_not_automatically_excluded(self):
        data=packet(); data['candidate_cells']=['A1','D7','F9']
        raw=answer(); raw['cells']['A1']={'kind':'non_business','category':'title','reason':'合成资料总标题','evidence_cells':['A1']}
        first=self.api.validate_response(raw,data,self.standard)
        result=self.api.agree_rounds(first,first,self.standard)
        self.assertEqual(len(result['facts']),2)
        self.assertEqual(result['excluded'][0]['cell'],'A1')
        self.assertEqual(result['facts'][1]['raw_value_state'],'blank')
        bad=copy.deepcopy(raw); bad['cells']['D7']={'kind':'non_business','category':'empty_padding','reason':'不能排除实际数值','evidence_cells':['A1']}
        with self.assertRaises(ValueError):self.api.validate_response(bad,data,self.standard)
        other=copy.deepcopy(first); other['A1']={'kind':'unresolved','reason':'是否属于业务待填格证据不足','evidence_cells':['A1']}
        result=self.api.agree_rounds(first,other,self.standard)
        self.assertEqual(result['excluded'],[])
        self.assertEqual(result['unresolved'][0]['cell'],'A1')

    def test_non_business_display_labels_are_not_business_identity(self):
        data=packet();data['candidate_cells']=['A1']
        for category in ('header','label','unit','annotation'):
            with self.subTest(category=category):
                first={'cells':{'A1':{'kind':'non_business','category':'title','reason':'非业务说明','evidence_cells':['A1']}}}
                second=copy.deepcopy(first);second['cells']['A1']['category']=category
                result=self.api.agree_rounds(self.api.validate_response(first,data,self.standard),
                                             self.api.validate_response(second,data,self.standard),self.standard)
                self.assertEqual(result['unresolved'],[])
                self.assertEqual(set(result['excluded'][0]['category'].split(' / ')),{'title',category})
                self.assertEqual(len(result['excluded'][0]['rounds']),2)

    def test_evidence_addresses_do_not_define_the_non_business_meaning(self):
        data=packet(); data['cells']['B2']={'value':'资料说明'};data['candidate_cells']=['A1']
        first={'cells':{'A1':{'kind':'non_business','category':'title','reason':'总标题','evidence_cells':['A1']}}}
        second=copy.deepcopy(first);second['cells']['A1']['evidence_cells']=['A1','B2']
        result=self.api.agree_rounds(self.api.validate_response(first,data,self.standard),self.api.validate_response(second,data,self.standard),self.standard)
        self.assertEqual(result['unresolved'],[])
        self.assertEqual(set(result['excluded'][0]['evidence_cells']),{'A1','B2'})


class WholeWorkbookTests(unittest.TestCase):
    def test_partial_meaning_saved_displayed_and_not_treated_as_new_gold_or_complete(self):
        self._check_partial_flow(both_full=False)

    def test_full_round_disagreement_is_saved_and_displayed_without_gold_or_value_changes(self):
        self._check_partial_flow(both_full=True)

    def _check_partial_flow(self, *, both_full):
        from unittest.mock import patch
        from openpyxl import load_workbook
        from 附注更新.分步流程 import run_step
        from 附注更新.统一语义 import read_standard_v2
        from 附注更新.统一勾稽 import load_fact_mapping,build_fact_update_plan
        root=Path(__file__).resolve().parents[1]/'测试结果'/('部分语义正式入口_'+uuid.uuid4().hex);root.mkdir()
        gold=root/'标准.jsonl';gold.write_text('\n'.join(json.dumps(r,ensure_ascii=False) for r in records()),encoding='utf-8')
        source=root/'原表.xlsx';book=Workbook();sheet=book.active;sheet.title='任意布局'
        sheet['A1']='单户，人民币元，2025年12月31日，全部固定资产原值。未提供报告主体。';sheet['B2']=-125;book.save(source);book.close()
        calls=[]
        class ControlledPartial(UnifiedSemanticEngine):
            def _request(self,system,payload,validator):
                calls.append(payload['task']);self.assert_task=payload['task']
                if payload['task']!='unified_classify':raise AssertionError('实际主体缺失不得触发金标准新增')
                response={'cells':{}}
                for address in payload['candidate_cells']:
                    if address=='B2':
                        row=RecognitionTests().partial_answer()['cells']['D7']
                        if payload['round']==2:row=answer()['cells']['D7']
                        if both_full:
                            row=answer()['cells']['D7'];row['dimensions']=copy.deepcopy(row['dimensions'])
                            if payload['round']==2:row['dimensions']['entity']='另一合成主体'
                    else:row={'kind':'non_business','category':'title' if address=='A1' else 'empty_padding','reason':'受控试验原文','evidence_cells':['A1']}
                    response['cells'][address]=row
                return validator(response)
        config={'mode':'online','model':'受控模型','base_url':'http://localhost','api_key':'受控测试',
                'gold_path':str(gold),'b_path':str(source),'output_dir':str(root/'正式入口')}
        with patch('附注更新.统一识别.UnifiedSemanticEngine',ControlledPartial):result=run_step('recognize_b',config)
        self.assertEqual(result['status'],'partial');self.assertNotIn('gold_path',result['files'])
        standard=read_standard_v2(gold);mapping=load_fact_mapping(result['files']['b_mapping'],standard)
        self.assertFalse(mapping['complete']);self.assertEqual(mapping['facts'],[])
        self.assertEqual(mapping['unresolved'][0]['partial_semantics']['missing_dimensions'],['entity'])
        self.assertEqual(build_fact_update_plan(mapping,mapping,standard)['updates'],[])
        review=load_workbook(result['files']['b_review']);rows=list(review['逐格语义识别'].values);review.close()
        shown=next(row for row in rows if row[0]=='任意布局!B2')
        self.assertEqual(shown[2],'固定资产原值');self.assertIn('asset_selection',shown[3]);self.assertIn('entity',shown[5])
        self.assertEqual(shown[4],-125)
        self.assertTrue(all(task=='unified_classify' for task in calls))
        tampered=copy.deepcopy(mapping);tampered['unresolved'][0]['partial_semantics']['dimensions']['period']['date']='2024-12-31'
        bad=root/'被改语义.json';bad.write_text(json.dumps(tampered,ensure_ascii=False),encoding='utf-8')
        with self.assertRaisesRegex(ValueError,'部分语义与两轮识别不符'):load_fact_mapping(bad,standard)

    def test_extended_standard_reuses_reviewed_cells_and_only_retries_unresolved(self):
        root=Path(__file__).resolve().parents[1]/'测试结果'/('统一标准扩充接续_'+uuid.uuid4().hex);root.mkdir()
        gold=root/'原标准.jsonl';gold.write_text('\n'.join(json.dumps(row,ensure_ascii=False) for row in records()),encoding='utf-8')
        source=root/'原表.xlsx';book=Workbook();sheet=book.active;sheet.title='任意布局'
        sheet['A1']='合成企业甲；单户人民币元；2025年12月31日；固定资产原值及累计折旧'
        sheet['B2']=125;sheet['C2']=25;book.save(source);book.close()
        calls=[]
        class ControlledEngine(UnifiedSemanticEngine):
            def _request(self,system,payload,validator):
                calls.append(payload['candidate_cells'])
                available={row['id'] for row in payload['metrics']};response={'cells':{}}
                for address in payload['candidate_cells']:
                    if address=='C2' and '固定资产累计折旧' not in available:
                        row={'kind':'unresolved','reason':'缺少累计折旧指标','evidence_cells':['A1']}
                    elif address in ('B2','C2'):
                        row=copy.deepcopy(answer()['cells']['D7'])
                        if address=='C2':row['metric_id']='固定资产累计折旧'
                    else:row={'kind':'non_business','category':'title' if address=='A1' else 'empty_padding',
                              'reason':'受控测试，不计为模型准确率','evidence_cells':['A1']}
                    response['cells'][address]=row
                return validator(response)
        first=ControlledEngine({},gold,root/'第一轮').recognize_workbook(source)
        self.assertEqual(len(first['facts']),1);self.assertEqual(len(first['unresolved']),1)
        from 附注更新.统一识别 import standard_review_material
        from 附注更新.统一语义 import read_standard_v2
        material=standard_review_material(first,read_standard_v2(gold))
        self.assertEqual(material['sheets'][0]['candidate_cells'],['C2'])
        self.assertIn('A1',material['sheets'][0]['cells'])
        self.assertEqual(len(material['metrics']),1)
        self.assertTrue(material['dimension_definitions'])
        rows=records();new=copy.deepcopy(rows[-1]);new['id']='固定资产累计折旧';new['definition']['measurement_basis']='累计折旧';rows.append(new)
        extended=root/'扩充标准.jsonl';extended.write_text('\n'.join(json.dumps(row,ensure_ascii=False) for row in rows),encoding='utf-8')
        calls.clear()
        second=ControlledEngine({},extended,root/'第二轮').recognize_workbook(source,previous_result=first['mapping_path'])
        self.assertEqual(calls,[['C2'],['C2']])
        self.assertTrue(second['complete']);self.assertEqual(len(second['facts']),2)
        self.assertEqual(second['facts'][0],first['facts'][0])
        from 附注更新.统一勾稽 import load_fact_mapping
        load_fact_mapping(second['mapping_path'],read_standard_v2(extended))
        self.assertEqual(second['reused_mapping']['path'],first['mapping_path'])

    def test_whole_grid_and_hidden_financial_cell_are_covered_and_saved(self):
        root=Path(__file__).resolve().parents[1]/'测试结果'/('统一整表规则验证_'+uuid.uuid4().hex); root.mkdir()
        gold=root/'测试标准.jsonl'; gold.write_text('\n'.join(json.dumps(row,ensure_ascii=False) for row in records()),encoding='utf-8')
        source=root/'合成原表.xlsx'; book=Workbook(); sheet=book.active; sheet.title='任意布局'
        sheet['A1']='合成企业甲，单户人民币元，2025年12月31日固定资产原值合计'
        sheet.merge_cells('A1:D1'); sheet['D7']=125; sheet.row_dimensions[7].hidden=True
        book.save(source);book.close()
        class ControlledEngine(UnifiedSemanticEngine):
            def _request(self, system, payload, validator):
                response={'cells':{}}
                for address in payload['candidate_cells']:
                    if address=='D7':response['cells'][address]=answer()['cells']['D7']
                    else:response['cells'][address]={'kind':'non_business','category':'title' if address=='A1' else 'empty_padding',
                        'reason':'受控测试回执，不是模型准确率','evidence_cells':['A1']}
                return validator(response)
        engine=ControlledEngine({},gold,root/'识别结果')
        result=engine.recognize_workbook(source)
        self.assertTrue(result['complete'])
        self.assertEqual(len(result['facts']),1)
        self.assertEqual(len(result['excluded']),24)
        self.assertEqual(result['coverage']['expected_count'],25)
        self.assertNotIn('任意布局!B1',result['coverage']['expected_cells'])
        self.assertIn('任意布局!D7',result['coverage']['expected_cells'])
        self.assertTrue(Path(result['mapping_path']).is_file())
        from 附注更新.统一勾稽 import load_fact_mapping
        loaded=load_fact_mapping(result['mapping_path'],engine.standard)
        self.assertTrue(loaded['coverage']['complete'])
        forged=copy.deepcopy(result); forged['excluded'].pop()
        bad=root/'遗漏却声称完整.json'; bad.write_text(json.dumps(forged,ensure_ascii=False),encoding='utf-8-sig')
        with self.assertRaisesRegex(ValueError,'覆盖'):
            load_fact_mapping(bad,engine.standard)
        changed=copy.deepcopy(result)
        changed['facts'][0]['dimensions']['entity']='其他企业'
        changed['owners'][changed['facts'][0]['owner_reference']]['entity']='其他企业'
        bad_meaning=root/'与原回执不同的语义.json';bad_meaning.write_text(json.dumps(changed,ensure_ascii=False),encoding='utf-8-sig')
        with self.assertRaisesRegex(ValueError,'任务包'):
            load_fact_mapping(bad_meaning,engine.standard)
        from unittest.mock import patch
        from 附注更新.分步流程 import run_step
        progress=[]
        with patch.object(UnifiedSemanticEngine,'_request',ControlledEngine._request):
            formal=run_step('recognize_a',{'gold_path':str(gold),'a_path':str(source),'output_dir':str(root/'正式入口'),
                           'base_url':'http://unused.example/v1','model':'受控测试模型','api_key':'','timeout':600},progress=progress.append)
        self.assertEqual(formal['status'],'complete',formal['message'])
        self.assertTrue(Path(formal['files']['a_mapping']).is_file())
        self.assertTrue(Path(formal['files']['a_review']).is_file())
        self.assertTrue(progress)
        import threading
        from 附注更新.语义 import SemanticCancelled
        cancelled=threading.Event(); checkpoints=[]
        cancel_source=root/'取消合成原表.xlsx';cancel_book=Workbook();cancel_sheet=cancel_book.active
        for row in range(1,482):cancel_sheet.cell(row,1,'合成说明文本')
        cancel_book.save(cancel_source);cancel_book.close()
        def stop_after_one_packet(record):
            checkpoints.append(record)
            if record['excluded']:cancelled.set()
        interrupted=ControlledEngine({},gold,root/'取消后保留',cancel=cancelled)
        with self.assertRaises(SemanticCancelled):
            interrupted.recognize_workbook(cancel_source,progress=stop_after_one_packet)
        latest=load_fact_mapping(checkpoints[-1]['mapping_path'],interrupted.standard)
        self.assertFalse(latest['complete'])
        self.assertGreater(len(latest['excluded']),0)
        self.assertGreater(len(latest['unresolved']),0)
        self.assertEqual(len(latest['excluded'])+len(latest['unresolved']),481)


if __name__ == "__main__":
    unittest.main()
