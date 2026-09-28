# coding: utf-8
"""实际文件及正式入口验证未采用来源；模型响应为本地测试，不代表真实模型准确率。"""
import copy
import json
from pathlib import Path
import unittest
import uuid
from unittest.mock import patch
from docx import Document
from openpyxl import Workbook, load_workbook
from 附注更新.分步流程 import run_step
from 附注更新.统一识别 import UnifiedSemanticEngine
from 附注更新.统一勾稽 import load_fact_mapping, verify_fact_update_plan
from 附注更新.统一语义 import read_standard_v2
from 附注更新.表格 import file_hash
from 测试.测试统一语义 import records, fact


class UnusedSourceTests(unittest.TestCase):
    def setUp(self):
        self.root=Path(__file__).resolve().parents[1]/'测试结果'/('统一未采用来源_'+uuid.uuid4().hex);self.root.mkdir()
        self.gold=self.root/'标准.jsonl';self.gold.write_text('\n'.join(json.dumps(r,ensure_ascii=False) for r in records()),encoding='utf-8-sig')
        self.word=self.root/'原附注.docx';doc=Document();doc.add_heading('六、财务报表重要项目的说明',level=1)
        doc.add_paragraph('合成企业甲，单户人民币元，2025年12月31日。本表仅披露办公设备原值，不包括生产机器。')
        table=doc.add_table(rows=2,cols=2);table.style='Table Grid'
        for row,values in zip(table.rows,[['项目','期末原值'],['办公设备','10']]):
            for cell,value in zip(row.cells,values):cell.text=value
        doc.add_heading('七、其他事项',level=1);doc.add_paragraph('范围外文字保持。');doc.save(self.word)
        self.source=self.root/'更新来源.xlsx';book=Workbook();sheet=book.active;sheet.title='来源'
        sheet['A1']='合成企业甲，单户人民币元，2025年12月31日固定资产原值';sheet['A2']='办公设备';sheet['B2']=50
        sheet['A4']='生产机器';sheet['B4']=70;book.save(self.source);book.close()
        self.config={'gold_path':str(self.gold),'word_path':str(self.word),'b_path':str(self.source),
                     'base_url':'https://example.invalid/v1','model':'本地行为测试','output_dir':str(self.root/'正式流程')}
        self.mode='accept'
        extracted=run_step('extract_a',self.config);self.assertEqual(extracted['status'],'complete',extracted)
        self.config.update(extracted['files'])
        # 真实识别、覆盖、回执保存；只替换远程模型响应。
        with patch.object(UnifiedSemanticEngine,'_http',lambda engine,system,payload:self.respond(payload)):
            for step in ['recognize_a','recognize_b']:
                result=run_step(step,self.config);self.assertEqual(result['status'],'complete',result);self.config.update(result['files'])
        self.standard=read_standard_v2(self.gold)
        self.maps=[load_fact_mapping(self.config[key],self.standard) for key in ('a_mapping','b_mapping')]
        self.frozen={p:file_hash(p) for p in [str(self.word),str(self.source),self.config['a_mapping'],self.config['b_mapping']]}

    def respond(self,payload):
        if payload['task']=='unified_classify':
            result={};sheet=payload['sheet']
            for address in payload['candidate_cells']:
                business=address=='C3' if sheet=='sheet1' else address in {'B2','B4'}
                evidence='B3' if sheet=='sheet1' else ('A4' if address=='B4' else 'A2')
                if business:
                    dims=copy.deepcopy(fact()['dimensions']);dims['asset_selection']={'domain':'固定资产','mode':'members',
                        'members':['生产机器' if address=='B4' else '办公设备'],'completeness':'complete'}
                    result[address]={'kind':'metric_fact','metric_id':'固定资产原值','dimensions':dims,
                        'dimension_evidence':{key:[evidence] for key in dims},'reason':'合成资料原文明确'}
                else:result[address]={'kind':'non_business','category':'label','reason':'合成文字或布局格','evidence_cells':[evidence]}
            return {'cells':result}
        if payload['task']=='unified_equivalence':
            return {'pairs':{key:{'decision':'different','reason':'办公设备与生产机器为不同对象',
                'left_evidence':list(row['left']['evidence']),'right_evidence':list(row['right']['evidence'])} for key,row in payload['pairs'].items()}}
        if payload['task']=='unified_structure':
            if self.mode=='rejected_structure':
                return {'operations':[{'range_name':'TA_001','axis':'row','action':'insert','index':2,'count':1,'template_index':1,'reason':'待复核增行'}],
                    'bindings':[{'source_sheet':'来源','source_cell':'B4','target_sheet':'sheet1','target_cell':'C4',
                                 'template_sheet':'sheet1','template_cell':'C3','reason':'待复核新增成员'}],
                    'labels':[{'sheet':'sheet1','cell':'B4','value':'生产机器','source_sheet':'来源','source_cell':'A4','reason':'来源原文'}], 'unresolved':[]}
            return {'operations':[],'bindings':[],'labels':[],'unresolved':[]}
        if payload['task']=='unified_structure_review':
            return {'accepted':False,'reason':'原说明明确不披露生产机器，不应增行','source_evidence':['来源!A4'],'target_evidence':['sheet1!B1']}
        if payload['task']=='unified_unused_sources':
            target_ref=next(key for key,value in payload['target']['context_evidence'].items() if '仅披露' in json.dumps(value,ensure_ascii=False))
            rows={}
            for source in payload['source']['facts']:
                rows[source['id']]={'decision':'unresolved' if self.mode=='disagree' and payload['round']==2 else 'unused',
                    'reason':'附注明确只披露办公设备，生产机器不在本表披露范围。',
                    'target_evidence':[target_ref],'source_evidence':['来源!A4'],
                    'scope_limit':None if self.mode=='no_limit' else {'dimension_id':'asset_selection',
                        'reason':'原说明明确仅披露办公设备，不包括生产机器。','evidence':[target_ref]}}
            return {'sources':rows}
        raise AssertionError('意外模型任务：'+payload['task'])

    def match(self):
        with patch.object(UnifiedSemanticEngine,'_http',lambda engine,system,payload:self.respond(payload)):
            return run_step('match',self.config)

    def test_full_entry_saves_unused_proof_writes_word_and_rejects_forgery(self):
        result=self.match();self.assertEqual(result['status'],'complete',result)
        plan=json.loads(Path(result['files']['update_plan']).read_text(encoding='utf-8-sig'))
        from 附注更新.界面 import file_selection_updates
        restored=file_selection_updates('update_plan',result['files']['update_plan'],{})
        self.assertEqual(restored.get('unused_source_review'),result['files']['unused_source_review'])
        self.assertEqual(len(plan['updates']),1);self.assertEqual(len(plan['unused_sources']),1)
        self.assertEqual(plan['unused_sources'][0]['cell'],'B4');self.assertFalse(plan['unplaced_sources'])
        altered=copy.deepcopy(plan);altered['unused_source_reviews']=[]
        with self.assertRaises(ValueError):verify_fact_update_plan(altered,*self.maps,self.standard)
        altered=copy.deepcopy(plan);altered['unused_sources'][0]['cell']='B2'
        with self.assertRaises(ValueError):verify_fact_update_plan(altered,*self.maps,self.standard)
        written=run_step('write',{**self.config,'update_plan':result['files']['update_plan']})
        self.assertEqual(written['status'],'complete',written)
        book=load_workbook(written['files']['a_prime']);self.assertEqual(book['sheet1']['C3'].value,50);book.close()
        doc=Document(written['files']['word']);self.assertEqual(doc.tables[0].cell(1,1).text,'50')
        self.assertEqual(len(doc.tables[0].rows),2);self.assertEqual(doc.paragraphs[-1].text,'范围外文字保持。')
        self.assertEqual(len(load_fact_mapping(written['files']['a_prime_mapping'],self.standard)['facts']),1)
        self.assertTrue(all(file_hash(p)==digest for p,digest in self.frozen.items()))

    def test_same_metric_without_scope_limit_or_two_round_agreement_stays_pending(self):
        for mode in ['no_limit','disagree']:
            self.mode=mode;result=self.match();self.assertEqual(result['status'],'partial',result)
            plan=json.loads(Path(result['files']['update_plan']).read_text(encoding='utf-8-sig'))
            self.assertFalse(plan.get('unused_sources'));self.assertEqual(len(plan['unplaced_sources']),1)

    def test_rejected_unnecessary_expansion_does_not_block_proven_unused_source(self):
        self.mode='rejected_structure';result=self.match()
        self.assertEqual(result['status'],'complete',result)
        plan=json.loads(Path(result['files']['update_plan']).read_text(encoding='utf-8-sig'))
        self.assertFalse(plan.get('word_operations'));self.assertEqual(len(plan['unused_sources']),1)
        report=json.loads(Path(result['files']['structure_review']).read_text(encoding='utf-8-sig'))
        self.assertFalse(report['accepted']);self.assertEqual(len(report['reviews']),2)


if __name__=='__main__':unittest.main()
