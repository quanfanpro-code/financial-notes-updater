# coding: utf-8
"""正式入口逐格保存身份；模拟模型返回，不发网络请求。"""
import json
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch

from 测试 import 测试直接识别 as fixture
from 附注更新.A合集 import _records
from 附注更新.A语义阅读 import CELL_IDENTITY_OUTPUT
from 附注更新.分步流程 import run_step
from 附注更新.统一识别 import UnifiedSemanticEngine
from 附注更新.统一勾稽 import load_fact_mapping


class 逐格识别入口测试(unittest.TestCase):
    def setUp(self):
        self.case=fixture.直接识别测试();self.case.setUp()
        self.config={**self.case.config}
        for side in ('a','b'):
            self.config[side+'_target_table']={'id':'本表','sheet':'合成表','range':'A2:C5','data_ranges':['B3:C5']}

    def answer(self):
        return json.dumps({f'合成表!{column}{row}':{
            'metric_id':'固定资产原值','entity':'合成企业','report_scope':'standalone',
            'currency':'CNY','unit_scale':1,'period':{'kind':'instant','date':date},
            'template_period_role':role,
            'asset_selection':{'domain':'固定资产','mode':'members','members':['测试项目'+str(row)],'completeness':'complete'}}
            for row in (3,4,5) for column,date,role in
            [('B','2025-12-31','closing'),('C','2024-12-31','opening')]},ensure_ascii=False)

    def recognize(self,side,reply,extra=None):
        with ExitStack() as stack:
            call=stack.enter_context(patch.object(UnifiedSemanticEngine,'request_text',return_value=reply))
            for name in ('_attach_declared_name_refs','_attach_declared_row_labels','_attach_declared_source_periods','_save_direct_relations',
                         '_plain_answer','_own_answer_text','_period_rows','_direct_table'):
                stack.enter_context(patch('附注更新.A合集.'+name,side_effect=AssertionError('逐格身份不得经额外关系加工')))
            stack.enter_context(patch('附注更新.正常入口识别.answer_relations',side_effect=AssertionError('识别不生成取数安排')))
            result=run_step('recognize_'+side,{**self.config,**(extra or {})})
        self.assertEqual(result['status'],'complete',result)
        call.assert_called_once()
        prompt,materials=call.call_args.args
        self.assertTrue(prompt.endswith(CELL_IDENTITY_OUTPUT))
        self.assertIn('金标准对应的是本格自身承载的业务含义和计量值',prompt)
        self.assertIn('维度使用金标准原始ID',prompt)
        self.assertIn('原位置字段沿用location，维度分组字段沿用dimensions',prompt)
        self.assertIn('不另加自命名的外层包装',prompt)
        self.assertIn('不用字段别名或解释文字代替',prompt)
        self.assertIn('财务数据格的原值为空或具体成员未知则仍按已明确含义选择指标',prompt)
        self.assertIn('selection使用含mode的取值对象',prompt)
        self.assertIn('三位大写代码',prompt)
        self.assertIn('closing为期末、opening为期初、current为本期、prior为上期',prompt)
        for text in ('共同取值可以合并','共用取值可以合并','更新后的A排布由你','给出本次确定的安排',
                     '本次采用的必要来源及其他未采用候选'):
            self.assertNotIn(text,prompt)
        if side=='b' and 'A需要.json' in materials:
            self.assertIn('使用A已保存且符合金标准的同一实际取值',prompt)
            self.assertIn('双方原文不同或无法确定时照实保留',prompt)
        self.assertNotIn('本次需要确定来源的A位置',materials)
        self.assertNotIn('已保存的当前候选对应',materials)
        return load_fact_mapping(result['files'][side+'_mapping'],{}),materials

    def test_AB请求具名维度字段且泛称分组不转换(self):
        selection={'mode':'members','members':['测试项目3']}
        forms=[('具名',{'asset_selection':selection}),('未具名对象',selection),
               ('通用值列表',[{'dimension_id':'asset_selection','value':selection}])]
        a_mapping=None
        for side in ('a','b'):
            for name,dimensions in forms:
                with self.subTest(side=side,form=name):
                    reply=json.dumps({'location':'合成表!A3','dimension_id':'asset_selection',
                                      'dimensions':dimensions},ensure_ascii=False)
                    local={**self.config,side+'_target_table':{'id':'本表','sheet':'合成表','range':'A2:C5','data_ranges':['A3']}}
                    if side=='b':
                        self.assertIsNotNone(a_mapping)
                        local['a_mapping']=a_mapping
                    with patch.object(UnifiedSemanticEngine,'request_text',return_value=reply) as call:
                        result=run_step('recognize_'+side,local)
                    self.assertEqual(result['status'],'complete',result)
                    if side=='a' and name=='具名':
                        a_mapping=result['files']['a_mapping']
                    call.assert_called_once()
                    prompt,_=call.call_args.args
                    self.assertIn('逐项以金标准原始维度ID为字段名',prompt)
                    self.assertIn('不要把未具名的取值对象直接放进dimensions',prompt)
                    self.assertIn('不能由维度分组里的属性ID代替',prompt)
                    self.assertIn('文字、列表或表格均可，不限定JSON或固定模板',prompt)
                    mapping=load_fact_mapping(result['files'][side+'_mapping'],{})
                    record=dict(_records(mapping))['合成表!A3']
                    self.assertEqual(record['dimension_id'],'asset_selection')
                    if name=='具名':
                        self.assertEqual(record['dimensions']['asset_selection'],selection)
                    else:
                        self.assertNotIn('asset_selection',record['dimensions'])
                    self.assertEqual(mapping['本次逐格原答复'],reply)

    def test_逐格独立字段行直接JSON对象与代码引用原样一致(self):
        from 附注更新.A合集 import _per_cell_answer
        standard=self.case.engine().standard
        dims={'asset_selection':{'mode':'members','members':['设备甲；设备乙']},'currency':'CNY'}
        raw=json.dumps(dims,ensure_ascii=False,separators=(',',':'))
        for separator in (':','：'):
            for value in (raw,'`'+raw+'`'):
                reply=('dimensions: {"currency":"USD"}\n'
                       '- location'+separator+'合成表!B3\n'
                       '  - dimension_id'+separator+'asset_selection\n'
                       '  - dimensions'+separator+value+'\n'
                       '  - 依据：仅本格原文。')
                with self.subTest(separator=separator,quoted=value.startswith('`')):
                    result=_per_cell_answer(reply,standard,['合成表'])
                    self.assertEqual({row['id']:row['value'] for row in result['dimensions']},dims)
                    self.assertTrue(all(row['ranges']==['合成表!B3'] for row in result['dimensions']))
                    self.assertEqual([row['id'] for row in result['dimension_bindings']],['asset_selection'])

    def test_裸对象字段行不让后文依据串入前一标量字段(self):
        from 附注更新.A合集 import _per_cell_answer
        reply=('合成表!B3：\n'
               '  dimension_id：asset_selection\n'
               '  unit_scale：1\n'
               '  asset_selection：{"mode":"members","members":["设备甲"]}\n'
               '  依据：来源为本格原文。')
        result=_per_cell_answer(reply,self.case.engine().standard,['合成表'])
        dims={row['id']:row['value'] for row in result['dimensions']}
        self.assertEqual(dims['unit_scale'],1)
        self.assertEqual(dims['asset_selection'],{'mode':'members','members':['设备甲']})

    def test_编号值与字段同名时引号内仍是本格字面值(self):
        from 附注更新.A合集 import _per_cell_answer
        standard=self.case.engine().standard
        reply=('- `location`: `合成表!B3`；`dimension_id`: `asset_selection`；'
               '`asset_selection`: `{"mode":"members","members":["原文甲"]}`。')
        parsed=_per_cell_answer(reply,standard,['合成表'])
        self.assertEqual([row['id'] for row in parsed['dimension_bindings']],['asset_selection'])
        self.assertEqual({row['id']:row['value'] for row in parsed['dimensions']},
                         {'asset_selection':{'mode':'members','members':['原文甲']}})
        conflict=reply.replace('；`asset_selection`:', '；`dimension_id`: `period`；`asset_selection`:')
        parsed=_per_cell_answer(conflict,standard,['合成表'])
        self.assertFalse(parsed['dimension_bindings'])
        self.assertTrue(any(row.get('id')=='dimension_id' and '重复' in row['reason'] for row in parsed['unresolved']))

    def test_位置字段空格分隔只读取明确单格不扩展范围(self):
        from 附注更新.A合集 import _per_cell_answer
        standard=self.case.engine().standard
        for label in ('位置','原位置','location'):
            reply='- '+label+' `合成表!B3`：dimension_id=currency；currency=CNY。'
            parsed=_per_cell_answer(reply,standard,['合成表'])
            self.assertEqual([(r['id'],r['ranges']) for r in parsed['dimension_bindings']],[('currency',['合成表!B3'])])
            self.assertEqual([(r['id'],r['value']) for r in parsed['dimensions']],[('currency','CNY')])
            for bad in (reply.replace('B3','B3:B4'),reply.replace(label+' `',label+'合成表 `')):
                with self.assertRaises(ValueError):_per_cell_answer(bad,standard,['合成表'])

    def test_dimensions组内原始字段赋值保留而不从编号或泛称补值(self):
        from 附注更新.A合集 import _per_cell_answer
        standard=self.case.engine().standard
        for reply in (
            '- 合成表!B3：dimension_id=currency；dimensions：`currency`=`CNY`。',
            '- 合成表!B3：\n  dimension_id：`currency`\n  dimensions：`currency` = `CNY`',
        ):
            result=_per_cell_answer(reply,standard,['合成表'])
            self.assertEqual([(r['id'],r['value']) for r in result['dimensions']],[('currency','CNY')])
        for value in ('`currency`；实际取值=`CNY`','`CNY`','`{"mode":"members","members":["甲"]}`'):
            result=_per_cell_answer('- 合成表!B3：dimension_id=currency；dimensions：'+value,standard,['合成表'])
            self.assertEqual(result['dimensions'],[])
        reply='- 合成表!B3：dimension_id=currency；dimensions=`{"entity":"currency=CNY","currency":"EUR"}`'
        result=_per_cell_answer(reply,standard,['合成表'])
        self.assertEqual({r['id']:r['value'] for r in result['dimensions']},{'entity':'currency=CNY','currency':'EUR'})

    def test_行内dimensions对象只存本格字面值不补标准编号(self):
        from 附注更新.A合集 import _per_cell_answer
        standard=self.case.engine().standard
        dims={'asset_selection':{'mode':'members','members':['原文甲']}}
        value=json.dumps(dims,ensure_ascii=False,separators=(',',':'))
        reply='- `B3`：metric_id=null；dimensions='+value+'。依据：B3原文。'
        parsed=_per_cell_answer(reply,standard,['合成表'])
        self.assertEqual({row['id']:row['value'] for row in parsed['dimensions']},dims)
        self.assertFalse(parsed['metrics'])
        self.assertFalse(parsed['dimension_bindings'])

    def test_AB旧展示标签不屏蔽本格直接编号且说明仍保留(self):
        from copy import deepcopy
        from 附注更新.A合集 import direct_collection, save_direct_answer, _standard, validate_direct_collection
        from 附注更新.统一语义 import definition_hash
        a,_=self.recognize('a',self.answer())
        for side in ('a','b'):
            current=a if side=='a' else self.recognize('b',self.answer(),{'a_mapping':a['mapping_path']})[0]
            old=deepcopy(current);old.pop('target_table',None);old.pop('target_tables',None)
            tables=deepcopy(old['outline']['subtables'])
            labels=[{'range':'A3:B3','category':'label','reason':'旧展示标签'},
                    {'range':'B4','category':'label','reason':'旧展示标签'},
                    {'range':'C3','category':'annotation','reason':'原说明'}]
            tables[0]['non_business']=deepcopy(labels)
            old['识别答复']=json.dumps({'tables':tables},ensure_ascii=False)
            old['答复指纹']=definition_hash(old['识别答复'])
            old=direct_collection(old,old['机械记录'],_standard(old))
            reply={'合成表!B3':{'metric_id':None,'dimension_id':'asset_selection',
                   'dimensions':{'asset_selection':{'mode':'members','members':['设备甲']}}}}
            with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('离线保存不许模型调用')):
                saved=save_direct_answer(old,reply,self.case.folder/('展示标签接续'+side+'.json'),
                    field_scope={'合成表!B3':['metric_id','dimension_id','asset_selection']})
            rows=dict(_records(saved))
            self.assertIn('合成表!B3',rows)
            self.assertEqual(rows['合成表!B3']['dimension_id'],'asset_selection')
            self.assertEqual(rows['合成表!B3']['dimensions']['asset_selection'],{'mode':'members','members':['设备甲']})
            self.assertIn('合成表!B4',rows)
            excluded={row['cell'] for row in saved['excluded']}
            self.assertIn('A3',excluded);self.assertIn('C3',excluded)
            self.assertNotIn('B3',excluded);self.assertNotIn('B4',excluded)
            self.assertEqual(saved['outline']['subtables'][0]['non_business'],labels)
            self.assertEqual(saved['机械记录'],current['机械记录'])
            validate_direct_collection(saved,for_matching=True)

    def test_普通B接续不因A需要变化重答或向A传入额外关系(self):
        from copy import deepcopy
        from 附注更新.A合集 import direct_digest
        a,_=self.recognize('a',self.answer())
        b,_=self.recognize('b',self.answer(),{'a_mapping':a['mapping_path']})
        stored=deepcopy(b)
        stored['正常入口范围认识']={'原答复':'已有范围认识，未决保留','ranges':[]}
        stored['A需要指纹']='旧A需要指纹'
        stored['已确认关系']=[{'kind':'取数对应','answer_hash':'历史答复','bindings':[
            {'target':'合成表!B3','source':'合成表!C5'}]}]
        path=self.case.folder/'已有B合集普通接续.json'
        stored['mapping_path']=str(path.resolve())
        stored['封存校验']=direct_digest(stored)
        path.write_text(json.dumps(stored,ensure_ascii=False),encoding='utf-8')
        a_raw=Path(a['mapping_path']).read_bytes();b_raw=path.read_bytes()
        config={**self.config,'a_mapping':a['mapping_path'],'b_mapping':str(path),'b_target_table':None}
        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('不得重答')) as request:
            result=run_step('recognize_b',config)
        self.assertEqual(result['status'],'complete',result)
        self.assertEqual(result['files'],{'b_mapping':str(path.resolve())})
        request.assert_not_called()
        self.assertEqual(Path(a['mapping_path']).read_bytes(),a_raw)
        self.assertEqual(path.read_bytes(),b_raw)
        damaged=self.case.folder/'内容改变的B合集.json'
        stored['facts'][0]['dimensions']['currency']='错误改写'
        stored['mapping_path']=str(damaged.resolve())
        damaged.write_text(json.dumps(stored,ensure_ascii=False),encoding='utf-8')
        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('不能用重答掩盖文件变化')) as request:
            result=run_step('recognize_b',{**config,'b_mapping':str(damaged)})
        self.assertEqual(result['status'],'failed',result)
        self.assertIn('合集内容发生变化',result['message'])
        request.assert_not_called()

    def test_局部接入不能把旧共同展开字段变成逐格身份(self):
        from copy import deepcopy
        from 附注更新.A合集 import direct_digest, validate_direct_collection, combine_direct_collections
        for side in ('a','b'):
            extra={'a_mapping':current['mapping_path']} if side=='b' else {}
            current,_=self.recognize(side,self.answer(),extra)
            old=deepcopy(current);old.pop('逐格字段版本')
            for table in old['outline']['subtables']:
                for kind in ('metrics','dimensions'):
                    for row in table[kind]:
                        if ('B3' in row['ranges'] and kind=='metrics' or
                            'B4' in row['ranges'] and row['id']=='currency'):
                            row.pop('逐格原字段',None)
                            row['original']='共同声明；程序曾展开到各格'
            old['封存校验']=direct_digest(old)
            with self.assertRaisesRegex(ValueError,'旧展开身份'):
                validate_direct_collection(old,for_matching=True)
            old_path=self.case.folder/('旧展开'+side+'.json')
            old_path.write_text(json.dumps(old,ensure_ascii=False),encoding='utf-8')
            raw=json.dumps({loc:{'template_period_role':'closing'} for loc in ('合成表!B3','合成表!B4')})
            saved,_=self.recognize(side,raw,{**extra,side+'_mapping':str(old_path),side+'_target_table':None,
                side+'_field_scope':{loc:['template_period_role'] for loc in ('合成表!B3','合成表!B4')}})
            records=dict(_records(saved))
            self.assertNotIn('合成表!B3',records)
            self.assertNotIn('currency',records['合成表!B4']['dimensions'])
            self.assertEqual(records['合成表!C3']['dimensions'],dict(_records(current))['合成表!C3']['dimensions'])
            self.assertEqual(saved['机械记录'],current['机械记录'])
            self.assertTrue(any(r.get('original')=='共同声明；程序曾展开到各格'
                for t in saved['outline']['subtables'] for k in ('metrics','dimensions') for r in t[k]))
            validate_direct_collection(saved,for_matching=True)
            validate_direct_collection(combine_direct_collections([saved]),for_matching=True)

    def test_本格未知维度不阻断身份且旧标签不被同值继承(self):
        from copy import deepcopy
        from 附注更新.A合集 import direct_digest, save_direct_answer, validate_direct_collection
        current,_=self.recognize('a',self.answer())
        old=deepcopy(current);old.pop('逐格字段版本')
        for table in old['outline']['subtables']:
            for row in table['dimensions']:
                if row['id']=='currency' and 'B3' in row['ranges']:
                    row.pop('逐格原字段',None)
        for location,row in _records(old):
            if location=='合成表!B3':row['label_references']={'currency':['合成表!A1']}
        old['封存校验']=direct_digest(old)
        saved=save_direct_answer(old,json.dumps({'合成表!B3':{'currency':'CNY','asset_selection':None}}),
            self.case.folder/'只保存直接字段.json',field_scope={'合成表!B3':['currency','asset_selection']})
        row=dict(_records(saved))['合成表!B3']
        self.assertEqual(row['dimensions']['currency'],'CNY')
        self.assertNotIn('asset_selection',row['dimensions'])
        self.assertNotIn('currency',row.get('label_references',{}))
        validate_direct_collection(saved,for_matching=True)

    def test_只接本格标签引用不重答或补成身份(self):
        from copy import deepcopy
        from 附注更新.A合集 import direct_digest, save_direct_answer, validate_direct_collection
        for side in ('a','b'):
            extra={'a_mapping':current['mapping_path']} if side=='b' else {}
            current,_=self.recognize(side,self.answer(),extra)
            refs={'period':['合成表!B2'],'asset_selection':['合成表!A3']}
            raw=json.dumps({'合成表!B3':{'label_references':refs,'currency':'USD'},
                            '合成表!B4':{'label_references':{'period':['合成表!C2']}}},ensure_ascii=False)
            saved,_=self.recognize(side,raw,{**extra,side+'_mapping':current['mapping_path'],
                side+'_target_table':None,side+'_field_scope':{'合成表!B3':['label_references']}})
            before,after=dict(_records(current)),dict(_records(saved))
            self.assertEqual({k:v['dimensions'] for k,v in before.items()},
                             {k:v['dimensions'] for k,v in after.items()})
            self.assertEqual(after['合成表!B3']['label_references'],refs)
            self.assertNotIn('label_references',after['合成表!B4'])
            self.assertEqual(saved['机械记录'],current['机械记录'])
            validate_direct_collection(saved,for_matching=True)
            old=deepcopy(current)
            for table in old['outline']['subtables']:
                for row in table['dimensions']:
                    if row['id']=='asset_selection' and 'B3' in row['ranges']:row.pop('逐格原字段',None)
            old['封存校验']=direct_digest(old)
            saved=save_direct_answer(old,raw,self.case.folder/('引用不能补身份'+side+'.json'),
                field_scope={'合成表!B3':['label_references']})
            own=dict(_records(saved))['合成表!B3']
            self.assertNotIn('asset_selection',own['dimensions'])
            self.assertNotIn('asset_selection',own.get('label_references',{}))
            self.assertEqual(own['label_references']['period'],['合成表!B2'])
            validate_direct_collection(saved,for_matching=True)

    def test_未定位说明及下一组声明不成为前格字段(self):
        from 附注更新.A合集 import _per_cell_answer, direct_answer
        raw=('- 合成表!B3:B4：currency=USD。\n'
             '  - 合成表!B3：metric_id=固定资产原值；period=2027-03-31。\n'
             '    - entity=`本格自己的主体`。\n'
             '- 另组声明：entity=不属于前格。')
        parsed=_per_cell_answer(raw,self.case.engine().standard,['合成表'])
        self.assertEqual({r['id']:r['value'] for r in parsed['dimensions']},
                         {'period':'2027-03-31','entity':'本格自己的主体'})
        normalized=direct_answer(parsed,self.case.engine().standard,['合成表'])
        self.assertEqual(normalized['unresolved'][0]['ranges'],[])
        a,_=self.recognize('a',self.answer())
        b,_=self.recognize('b',self.answer(),{'a_mapping':a['mapping_path']})
        saved,_=self.recognize('b',raw,{'a_mapping':a['mapping_path'],'b_mapping':b['mapping_path'],
            'b_target_table':None,'b_field_scope':{'合成表!B3':['period','entity','currency']}})
        row=dict(_records(saved))['合成表!B3']
        self.assertEqual(row['dimensions']['entity'],'本格自己的主体')
        self.assertEqual(row['dimensions']['period'],'2027-03-31')
        self.assertEqual(row['dimensions']['currency'],'CNY')

    def test_中文引号内的分号及说明原样保留(self):
        for opening,closing in [('“','”'),('‘','’')]:
            reply=(f'- 合成表!B3：metric_id=固定资产原值；entity={opening}甲类；附加说明，未明确{closing}。')
            a,_=self.recognize('a',reply)
            b,_=self.recognize('b',reply,{'a_mapping':a['mapping_path']})
            for mapping in (a,b):
                self.assertEqual(dict(_records(mapping))['合成表!B3']['dimensions']['entity'],'甲类；附加说明，未明确')

    def test_格子标题下的字段段落不依赖段落名称(self):
        reply=('### 合成表!B3\n- metric_id=固定资产原值\n'
               '- 本格身份取值：entity=本格主体；currency=CNY。')
        a,_=self.recognize('a',reply)
        b,_=self.recognize('b',reply,{'a_mapping':a['mapping_path']})
        for mapping in (a,b):
            self.assertEqual(dict(_records(mapping))['合成表!B3']['dimensions'],{'entity':'本格主体','currency':'CNY'})

    def test_AB逐格结果原样保存并保留空白身份(self):
        a,_=self.recognize('a',self.answer())
        reply=json.loads(self.answer())
        for row in reply.values():row['dimensions']={'template_period_role':row.pop('template_period_role')}
        b,_=self.recognize('b',json.dumps(reply,ensure_ascii=False),{'a_mapping':a['mapping_path']})
        for mapping in (a,b):
            records=dict(_records(mapping))
            self.assertEqual(set(records),{'合成表!'+cell for cell in ('B3','C3','B4','C4','B5','C5')})
            for loc,row in records.items():
                self.assertEqual(row['metric_id'],'固定资产原值')
                self.assertEqual(row['dimensions']['asset_selection']['members'],['测试项目'+loc[-1]])
            roles={cell:row.get('template_period_role') for table in mapping['outline']['subtables'] for row in table['metrics'] for cell in row['ranges']}
            self.assertEqual(roles,{col+str(row):role for row in (3,4,5) for col,role in [('B','closing'),('C','opening')]})
        self.assertEqual({loc:row['dimensions'] for loc,row in _records(a)},
                         {loc:row['dimensions'] for loc,row in _records(b)})

    def test_字段和值同在代码引用内时只读引号内字面值(self):
        from 附注更新.A合集 import _per_cell_answer
        raw=('- 合成表!B3 — metric_id：`固定资产原值`；'
             '`entity=甲公司，乙单位`，`currency=CNY`，`template_period_role=closing`，'
             '`period={"kind":"instant","date":"2025-12-31"}`，'
             '`asset_selection={"mode":"members","members":["带，逗号的项目"]}`，`unit_scale=1`。依据：本格原文。')
        parsed=_per_cell_answer(raw,self.case.engine().standard,['合成表'])
        values={row['id']:row['value'] for row in parsed['dimensions']}
        self.assertEqual(values,{'entity':'甲公司，乙单位','currency':'CNY','template_period_role':'closing',
            'period':{'kind':'instant','date':'2025-12-31'},'asset_selection':{'mode':'members','members':['带，逗号的项目']},'unit_scale':1})
        self.assertEqual(parsed['metrics'][0]['template_period_role'],'closing')

    def test_来源标题和对象中文引号只读取本格字面值(self):
        from 附注更新.A合集 import _per_cell_answer
        standard=self.case.engine().standard
        for carrier in ('A','B',''):
            raw=(f'- **{carrier}来源：合成表!B3（本格项目）**\n'
                 '  `metric_id=固定资产原值`；`template_period_role=closing`；维度：entity=甲公司，乙部门，'
                 'currency=CNY，unit_scale=1，period={date:2027-03-31, kind:instant}，'
                 'asset_selection={mode:members, members:[“普通成员”,"值含“中文引号”、}和;”"]}。\n'
                 '  依据：本格原文。')
            parsed=_per_cell_answer(raw,standard,['合成表'])
            self.assertEqual(parsed['metrics'][0]['ranges'],['合成表!B3'])
            self.assertEqual({row['id']:row['value'] for row in parsed['dimensions']},{
                'template_period_role':'closing','entity':'甲公司，乙部门','currency':'CNY','unit_scale':1,
                'period':{'date':'2027-03-31','kind':'instant'},
                'asset_selection':{'mode':'members','members':['普通成员','值含“中文引号”、}和;”']}})
            with self.assertRaises(ValueError):
                _per_cell_answer(raw.replace('合成表!B3','合成表!B3:B4'),standard,['合成表'])

    def test_AB括号内字段赋值不把外层括号存入值(self):
        for opening,closing in [('（','）'),('(',')')]:
            reply=('- 合成表!B3：metric_id：`固定资产原值`；`entity`＝`甲公司（分部）`；'
                   f'`period`＝2027年3月31日{opening}`template_period_role`＝closing{closing}。')
            a,_=self.recognize('a',reply)
            b,_=self.recognize('b',reply,{'a_mapping':a['mapping_path']})
            for mapping in (a,b):
                row=dict(_records(mapping))['合成表!B3']
                self.assertEqual(row['dimensions']['period'],'2027年3月31日')
                self.assertEqual(row['dimensions']['entity'],'甲公司（分部）')
                roles=[x['template_period_role'] for t in mapping['outline']['subtables'] for x in t['metrics']]
                self.assertEqual(roles,['closing'])
        from 附注更新.A合集 import _per_cell_answer
        parsed=_per_cell_answer('- 合成表!B3：metric_id：`固定资产原值`；entity=`原文字面值（`。',
                                self.case.engine().standard,['合成表'])
        self.assertEqual(parsed['dimensions'][0]['value'],'原文字面值（')

    def test_A定向接续保存旧登记外已回答的位置且不继承身份(self):
        original=json.loads(self.answer())
        a,_=self.recognize('a',json.dumps({k:v for k,v in original.items() if k.endswith('3')},ensure_ascii=False),{
            'a_target_table':{'id':'本表','sheet':'合成表','range':'A2:C5','data_ranges':['B3:C3']}})
        reply={'合成表!B4':{'metric_id':'固定资产原值','asset_selection':{'mode':'members','members':['新位置自己的成员']}},
               '合成表!B5':original['合成表!B5']}
        saved,_=self.recognize('a',json.dumps(reply,ensure_ascii=False),{
            'a_mapping':a['mapping_path'],'a_field_scope':{'合成表!B4':['metric_id','asset_selection','currency']},
            'a_target_table':{'id':'本次位置','sheet':'合成表','range':'A2:C5','data_ranges':['B4']}})
        before=dict(_records(a));after=dict(_records(saved))
        self.assertEqual(set(after),set(before)|{'合成表!B4'})
        self.assertEqual(after['合成表!B4']['dimensions'],{'asset_selection':{'mode':'members','members':['新位置自己的成员']}})
        for loc,row in before.items():
            self.assertEqual(after[loc]['metric_id'],row['metric_id'])
            self.assertEqual(after[loc]['dimensions'],row['dimensions'])

    def test_副本新增格接续独立识别并按身份更新且不重复插行(self):
        from openpyxl import load_workbook
        from openpyxl.workbook.defined_name import DefinedName
        from 附注更新.表格 import file_hash
        book=load_workbook(self.case.source)
        book.defined_names.add(DefinedName('名单表',attr_text="'合成表'!$A$2:$C$5"))
        book['合成表']['B8']=123;book['合成表']['C8']=456
        book.save(self.case.source);book.close()
        original_hash=file_hash(self.case.source)
        reply=json.loads(self.answer())
        for loc,row in reply.items():
            row['label_references']={'asset_selection':['合成表!A'+loc[-1]],'period':[loc[:-1]+'2']}
        a,_=self.recognize('a',json.dumps(reply,ensure_ascii=False))
        for loc,row in reply.items():row['period']['date']='2026-12-31' if '!B' in loc else '2025-12-31'
        for col in ('B','C'):
            reply['合成表!'+col+'8']={**reply['合成表!'+col+'3'],
                'asset_selection':{'mode':'members','members':['新增成员']}}
        b,_=self.recognize('b',json.dumps(reply,ensure_ascii=False),{
            'a_mapping':a['mapping_path'],
            'b_target_table':{'id':'本表','sheet':'合成表','range':'A2:C8','data_ranges':['B3:C5','B8:C8']}})
        from 附注更新.A合集 import direct_digest
        b['已确认关系']=[{'kind':'取数对应','source_a_hash':a['source_hash'],'source_b_hash':b['source_hash'],
            'gold_hash':a['gold_hash'],'label_actions':[{'target':'合成表!B2','action':'set','value':'已明确的期末标签'}]}]
        b['封存校验']=direct_digest(b)
        b_path=self.case.folder/'已有明确标签安排B.json'
        b_path.write_text(json.dumps(b,ensure_ascii=False),encoding='utf-8')
        b=load_fact_mapping(b_path,{})
        purpose={'reason':'已明确的名单版式增加一个空行',
            'operations':[{'range_name':'名单表','axis':'row','action':'insert','index':2,'count':1,'template_index':1}]}
        config={**self.config,'a_mapping':a['mapping_path'],'b_mapping':b['mapping_path'],'update_purpose':purpose}
        matched=run_step('match',config);self.assertEqual(matched['status'],'complete',matched)
        written=run_step('write',{**config,**matched['files']});self.assertEqual(written['status'],'complete',written)
        receipt=Path(written['files']['a_prime_mapping'])
        fields=['metric_id','template_period_role','entity','report_scope','currency','unit_scale','period','asset_selection']
        selected={**config,**written['files'],'a_path':written['files']['a_prime'],
            'a_field_scope':{'合成表!B4':fields,'合成表!C4':fields},
            'a_target_table':{'id':'本表','sheet':'合成表','range':'A2:C6','data_ranges':['B4:C4']}}
        new_answer={f'合成表!{col}4':reply[f'合成表!{col}8'] for col in ('B','C')}
        saved,materials=self.recognize('a',json.dumps(new_answer,ensure_ascii=False),selected)
        self.assertEqual(materials['已有身份.json']['需要'],[])
        self.assertIn('上次搬运记录.json',materials)
        old=dict(_records(a));new=dict(_records(saved))
        for loc,row in old.items():
            col=loc.rsplit('!',1)[1][0];number=int(loc.rsplit('!',1)[1][1:])
            moved=f'合成表!{col}{number+(number>=4)}'
            self.assertEqual(new[moved]['metric_id'],row['metric_id'])
            self.assertEqual(new[moved]['dimensions'],row['dimensions'])
            self.assertEqual(new[moved]['label_references']['asset_selection'],
                             ['合成表!A'+str(number+(number>=4))])
        for col in ('B','C'):
            self.assertEqual(new[f'合成表!{col}4']['dimensions']['asset_selection'],{'mode':'members','members':['新增成员']})
        second=run_step('match',{**selected,'a_mapping':saved['mapping_path']})
        self.assertEqual(second['status'],'complete',second)
        plan=json.loads(Path(second['files']['update_plan']).read_text(encoding='utf-8-sig'))
        self.assertEqual(plan['operations'],[])
        updates={row['cell']:row for row in plan['updates']}
        self.assertEqual(updates['B4']['sources'][0]['cell'],'B8')
        self.assertEqual(updates['C4']['sources'][0]['cell'],'C8')
        final=run_step('write',{**selected,'a_mapping':saved['mapping_path'],**second['files']})
        self.assertEqual(final['status'],'complete',final)
        book=load_workbook(final['files']['a_prime'])
        self.assertEqual((book['合成表']['B4'].value,book['合成表']['C4'].value),(123,456))
        self.assertEqual(book['合成表']['B2'].value,'本期')
        self.assertEqual(book.defined_names['名单表'].attr_text,"'合成表'!$A$2:$C$6")
        book.close()
        self.assertEqual(file_hash(self.case.source),original_hash)
        automatic,automatic_materials=self.recognize('a',json.dumps(new_answer,ensure_ascii=False),{
            **selected,'a_mapping':'','a_field_scope':None,'a_target_table':None})
        self.assertEqual(set(automatic_materials['本次字段范围.json']['字段范围']),
                         {'合成表!A4','合成表!B4','合成表!C4'})
        self.assertEqual({loc:row['dimensions'] for loc,row in _records(automatic)},
                         {loc:row['dimensions'] for loc,row in _records(saved)})
        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('没有新位置时不重答')):
            reused=run_step('recognize_a',{**selected,**final['files'],'a_path':final['files']['a_prime'],
                'a_mapping':'','a_field_scope':None,'a_target_table':None})
        self.assertEqual(reused['status'],'complete',reused)
        reused_mapping=load_fact_mapping(reused['files']['a_mapping'],{})
        self.assertEqual({loc:row['dimensions'] for loc,row in _records(reused_mapping)},
                         {loc:row['dimensions'] for loc,row in _records(saved)})
        # 执行记录与当前副本不对应时，在调用模型之前停止。
        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('不应调用模型')):
            failed=run_step('recognize_a',{**selected,'a_mapping':str(receipt),'a_path':str(self.case.source)})
        self.assertEqual(failed['status'],'failed')
        self.assertIn('执行记录与所选A副本不一致',failed['message'])

    def test_名单入口仅改指定格字段并保留完整相关标准(self):
        a,_=self.recognize('a',self.answer())
        b,_=self.recognize('b',self.answer(),{'a_mapping':a['mapping_path']})
        reply='- **合成表!B3** — metric_id：`固定资产原值`；字段：`asset_selection`；实际取值：`{"domain":"固定资产","mode":"members","members":["另一成员"],"completeness":"complete"}`。原文依据：该行项目原文。'
        saved,materials=self.recognize('b',reply,{'a_mapping':a['mapping_path'],'b_mapping':b['mapping_path'],
            'list_dimension':'asset_selection','b_field_scope':{'合成表!B3':['asset_selection']},
            'a_required_ranges':[{'sheet':'合成表','range':'A2:C4'}],
            'b_source_ranges':[{'sheet':'合成表','range':'A2:C4'}],
            'b_allowed_metrics':['固定资产原值']})
        before=dict(_records(b));after=dict(_records(saved))
        self.assertEqual(after['合成表!B3']['dimensions']['asset_selection']['members'],['另一成员'])
        for loc,row in before.items():
            for key,value in row['dimensions'].items():
                if (loc,key)!=('合成表!B3','asset_selection'):
                    self.assertEqual(after[loc]['dimensions'][key],value)
        definitions=[json.loads(line) for line in materials['金标准.jsonl'].splitlines()]
        metric=next(row for row in definitions if row['id']=='固定资产原值')
        self.assertTrue({row['id'] for row in metric['dimension_refs']}<={row['id'] for row in definitions})
        self.assertNotIn('已有名单搬运安排',materials)
        self.assertIn('已有身份.json',materials)
        self.assertEqual([row['location'] for row in materials['已有身份.json']['需要']],['合成表!B3'])
        source=materials.get('B原文.json',materials.get('原文.json'))
        self.assertIn('合成表!A2:C4',[row['sheet']+'!'+row['range'] for row in source['本次读取范围']])
        full=next(row for row in source['完整工作表原文']['sheets'] if row['sheet']=='合成表')
        self.assertIn('5:',full['rows'])
        self.assertEqual(materials['本次字段范围.json']['字段范围'],{'合成表!B3':['asset_selection']})
        self.assertEqual(saved['机械记录'],b['机械记录'])

    def test_旧B的保留动作不替当前或另一份B决定取值(self):
        from 测试.测试唯一身份证匹配 import 唯一身份证匹配测试
        helper=唯一身份证匹配测试()
        a=helper.mapping('A',[('C3','余额',{'类别':'成员'},10)])
        b=helper.mapping('B',[('D7','余额',{'类别':'成员'},99)])
        a['副本接续']={'保留原值':['A!C3'],'已执行安排':[{'source_b_hash':'B','purpose':{}}]}
        self.assertEqual(helper.plan(a,b)['updates'][0]['value'],99)
        b['source_hash']='另一份B'
        self.assertEqual(helper.plan(a,b)['updates'][0]['value'],99)

    def test_共同取值和行列声明不生成逐格身份(self):
        from 附注更新.A合集 import _per_cell_answer
        engine=self.case.engine()
        with self.assertRaisesRegex(ValueError,'未提供各格'):
            _per_cell_answer('B列均为固定资产原值，currency=CNY；范围合成表!B3:B5。',engine.standard,['合成表'])
        with self.assertRaisesRegex(ValueError,'未提供各格'):
            _per_cell_answer('| A位置 | B位置 | 金标准 |\n|---|---|---|\n|合成表!B3|合成表!C3|固定资产原值|',engine.standard,['合成表'])
        reply='共同维度：currency=CNY。\n合成表!B3：metric_id=固定资产原值；asset_selection={"mode":"members","members":["原值"]}。\n合成表!B4：metric_id=固定资产原值。'
        parsed=_per_cell_answer(reply,engine.standard,['合成表'])
        self.assertFalse(any(r['id']=='currency' for r in parsed['dimensions']))
        self.assertEqual([(r['ranges'],r['value']) for r in parsed['dimensions']],
                         [(['合成表!B3'],{'mode':'members','members':['原值']})])

    def test_单工作表定向请求仅补全各格地址不共享身份(self):
        from 附注更新.A合集 import _per_cell_answer
        engine=self.case.engine()
        reply='共同维度：currency=CNY。\n- B3｜metric_id：固定资产原值；asset_selection={"mode":"all"}。\n- B4｜metric_id：固定资产原值。'
        parsed=_per_cell_answer(reply,engine.standard,['合成表','其他表'],default_sheet='合成表')
        self.assertEqual([r['ranges'] for r in parsed['metrics']],[['合成表!B3'],['合成表!B4']])
        self.assertEqual([(r['id'],r['ranges']) for r in parsed['dimensions']],[('asset_selection',['合成表!B3'])])
        with self.assertRaisesRegex(ValueError,'未提供各格'):
            _per_cell_answer(reply,engine.standard,['合成表','其他表'])
        parsed=_per_cell_answer('- B3｜metric_id：`固定资产原值`｜`template_period_role=current`；`currency=CNY`。依据：本格原文。',
                                engine.standard,['合成表','其他表'],default_sheet='合成表')
        self.assertEqual([(r['id'],r['template_period_role']) for r in parsed['metrics']],[('固定资产原值','current')])

    def test_明确工作表标题只补位置且指标只读本格(self):
        from 附注更新.A合集 import _per_cell_answer
        reply='''### 合成表
共同字段：currency=USD。
- **B3**：metric_id=`固定资产原值`；`currency=CNY`；`template_period_role=opening`。
### 其他表
- **B3**：metric_id=`固定资产原值`；`period=2027-12-31`。
- **C4**：依据：`固定资产原值`；`currency=EUR`。
### 未登记工作表
- **D5**：metric_id=`固定资产原值`；`currency=JPY`。
'''
        parsed=_per_cell_answer(reply,self.case.engine().standard,['合成表','其他表'])
        self.assertEqual([(r['ranges'],r['id']) for r in parsed['metrics']],
                         [(['合成表!B3'],'固定资产原值'),(['其他表!B3'],'固定资产原值')])
        self.assertEqual(parsed['metrics'][0]['template_period_role'],'opening')
        self.assertEqual([(r['ranges'],r['value']) for r in parsed['dimensions'] if r['id']=='currency'],
                         [(['合成表!B3'],'CNY'),(['其他表!C4'],'EUR')])

    def test_表格坐标附行名不继承身份且拒绝多位置(self):
        from 附注更新.A合集 import _per_cell_answer
        engine=self.case.engine()
        header='共同维度：currency=USD。\n| 位置及原文行名 | metric_id | 本格实际取值 | 原文依据 |\n|---|---|---|---|\n'
        reply=header+'| 合成表!B3；原文项目 | 固定资产原值 | currency=CNY；asset_selection={"mode":"all"} | 合成表!A3 |\n'
        reply+='| C3; 另一原文项目 | 固定资产原值 | period=2024-12-31 | 合成表!A3 |'
        parsed=_per_cell_answer(reply,engine.standard,['合成表'])
        self.assertEqual([r['ranges'] for r in parsed['metrics']],[['合成表!B3'],['合成表!C3']])
        self.assertEqual([(r['ranges'],r['value']) for r in parsed['dimensions'] if r['id']=='currency'],[(['合成表!B3'],'CNY')])
        for address in ['合成表!B3；合成表!C3','合成表!B3:C3；两个位置']:
            with self.assertRaisesRegex(ValueError,'未提供各格'):
                _per_cell_answer(header+f'| {address} | 固定资产原值 | currency=CNY | 原文 |',engine.standard,['合成表'])

    def test_逐格文字与表格取值不经程序同义改写(self):
        from 附注更新.A合集 import _per_cell_answer
        engine=self.case.engine()
        parsed=_per_cell_answer('| 位置 | metric_id | currency | period |\n|---|---|---|---|\n|合成表!B3|固定资产原值|人民币|2025-12-31|',engine.standard,['合成表'])
        self.assertEqual({r['id']:r['value'] for r in parsed['dimensions']},
                         {'currency':'人民币','period':'2025-12-31'})
        parsed=_per_cell_answer('- 合成表!B3｜metric_id：`固定资产原值`；`currency`＝`CNY`。依据：本格原文。',engine.standard,['合成表'])
        self.assertEqual(parsed['metrics'][0]['ranges'],['合成表!B3'])
        self.assertEqual(parsed['dimensions'][0]['value'],'CNY')
        parsed=_per_cell_answer('- 原位置：`合成表!B3`；metric_id：`固定资产原值`。`currency`：`CNY`。依据：本格原文。',engine.standard,['合成表'])
        self.assertEqual(parsed['metrics'][0]['ranges'],['合成表!B3'])
        self.assertEqual(parsed['dimensions'][0]['value'],'CNY')
        parsed=_per_cell_answer('| 原位置 | metric_id | 本次字段实际取值 | 原文依据 |\n|---|---|---|---|\n'
            '| 合成表!B3 | `固定资产原值` | `currency="CNY"`；`asset_selection=null` | 另一引用有currency=USD，不是本格取值 |',engine.standard,['合成表'])
        self.assertEqual(parsed['metrics'][0]['ranges'],['合成表!B3'])
        self.assertEqual({r['id']:r['value'] for r in parsed['dimensions']},{'currency':'CNY','asset_selection':None})
        parsed=_per_cell_answer('| B位置 | 身份字段实际取值 | 依据 |\n|---|---|---|\n'
            '| 合成表!B3 | `metric_id=固定资产原值`；`currency=CNY` | 合成表!A3 |',engine.standard,['合成表'])
        self.assertEqual(parsed['metrics'][0]['ranges'],['合成表!B3'])
        self.assertEqual(parsed['dimensions'][0]['value'],'CNY')

    def test_每格金标准编号可从文字和表格保存而不补共用维度(self):
        reply='''共同说明：currency=USD。
- **合成表!B3**
  - metric_id：`固定资产原值`
  - `currency`：`CNY`
  - `period`：2025-12-31

| 位置 | metric_id | period |
|---|---|---|
| 合成表!C3 | 固定资产原值 | 2024-12-31 |
'''
        saved,_=self.recognize('a',reply)
        rows=dict(_records(saved))
        self.assertEqual(rows['合成表!B3']['metric_id'],'固定资产原值')
        self.assertEqual(rows['合成表!C3']['metric_id'],'固定资产原值')
        self.assertEqual(rows['合成表!B3']['dimensions'],{'currency':'CNY','period':'2025-12-31'})
        self.assertEqual(rows['合成表!C3']['dimensions'],{'period':'2024-12-31'})

    def test_模型明确无指标时只撤下该格旧指标并保留原文(self):
        before,_=self.recognize('a',self.answer())
        after,_=self.recognize('a','合成表!B3：metric_id=null；template_period_role=closing。依据：本格原文未对应现有计量指标。',
            {'a_mapping':before['mapping_path'],'a_field_scope':{'合成表!B3':['metric_id','template_period_role']}})
        old=dict(_records(before));new=dict(_records(after))
        self.assertEqual(set(new),set(old)-{'合成表!B3'})
        for loc in new:
            self.assertEqual(new[loc]['metric_id'],old[loc]['metric_id'])
            self.assertEqual(new[loc]['dimensions'],old[loc]['dimensions'])
        self.assertEqual(after['机械记录'],before['机械记录'])
        self.assertTrue(any(r.get('cell')=='B3' for r in after['unresolved']))

    def test_AB明确空值撤下所问旧维度而不影响其他字段(self):
        a,_=self.recognize('a',self.answer())
        b,_=self.recognize('b',self.answer(),{'a_mapping':a['mapping_path']})
        for side,before in [('a',a),('b',b)]:
            with self.subTest(side=side):
                after,_=self.recognize(side,'- 原位置：合成表!B3；`asset_selection` = `null`。依据：本格范围未能确定。',
                    {'a_mapping':a['mapping_path'],side+'_mapping':before['mapping_path'],
                     side+'_field_scope':{'合成表!B3':['asset_selection']}})
                old=dict(_records(before));new=dict(_records(after))
                self.assertNotIn('asset_selection',new['合成表!B3']['dimensions'])
                for loc,row in old.items():
                    self.assertEqual(new[loc]['metric_id'],row['metric_id'])
                    self.assertEqual({k:v for k,v in new[loc]['dimensions'].items() if (loc,k)!=('合成表!B3','asset_selection')},
                                     {k:v for k,v in row['dimensions'].items() if (loc,k)!=('合成表!B3','asset_selection')})
                self.assertEqual(after['机械记录'],before['机械记录'])
                self.assertEqual(new['合成表!B3']['missing_dimensions'],['asset_selection'])
                self.assertTrue(any(r.get('id')=='asset_selection' and r.get('value') is None and r.get('ranges')==['B3']
                                    for table in after['outline']['subtables'] for r in table['unresolved']))

    def test_逐格标题下的字段保留且引用说明不改维度(self):
        before,_=self.recognize('a',self.answer())
        reply='''### 原位置：合成表!B3（本格原文名称）
- `asset_selection`：`{mode: "predicate", conditions: ["属于测试资产"]}`
- `currency`：`CNY`
- `period`：2025年12月31日
- `label_references`：`asset_selection`：A3“测试资产”；`period`：B2“期末”
- 依据：本格标签和表头。

### 合成表!C3
- `metric_id`：`null`
- 依据：本格原文未对应现有计量指标。

### 共同说明
- currency=USD
'''
        after,_=self.recognize('a',reply,{'a_mapping':before['mapping_path'],
            'a_field_scope':{'合成表!B3':['asset_selection','currency','period'],'合成表!C3':['metric_id']}})
        old=dict(_records(before));new=dict(_records(after))
        self.assertEqual(set(new),set(old)-{'合成表!C3'})
        self.assertEqual(new['合成表!B3']['dimensions']['asset_selection'],{'mode':'predicate','conditions':['属于测试资产']})
        self.assertEqual(new['合成表!B3']['dimensions']['currency'],'CNY')
        self.assertEqual(new['合成表!B3']['dimensions']['period'],'2025年12月31日')
        for loc in set(new)-{'合成表!B3'}:
            self.assertEqual(new[loc]['dimensions'],old[loc]['dimensions'])
        self.assertEqual(after['机械记录'],before['机械记录'])

    def test_对象冒号无空格不改键值和日期文字(self):
        before,_=self.recognize('a',self.answer())
        reply='''### 合成表!B3
- 实际取值：`period={date:2025-12-31, kind:instant}`；`asset_selection={mode:predicate, conditions:["文字mode:unknown保持原文", 含冒号:的文字], count:1, flag:false}`
'''
        after,_=self.recognize('a',reply,{'a_mapping':before['mapping_path'],
            'a_field_scope':{'合成表!B3':['period','asset_selection']}})
        row=dict(_records(after))['合成表!B3']
        self.assertEqual(row['dimensions']['period'],{'date':'2025-12-31','kind':'instant'})
        self.assertEqual(row['dimensions']['asset_selection'],{'mode':'predicate',
            'conditions':['文字mode:unknown保持原文','含冒号:的文字'],'count':1,'flag':False})

    def test_无代码引号的逐格对象按自身括号读取(self):
        from 附注更新.A合集 import _per_cell_answer
        reply='''- 合成表!C3
  - metric_id：固定资产原值
  - 实际取值：`period`={kind: instant, date: 2027-03-31}；`asset_selection`={mode: members, members:["含;号、}和`的文字", 普通成员]}；`currency`=CNY。
'''
        parsed=_per_cell_answer(reply,self.case.engine().standard,['合成表'])
        self.assertEqual({row['id']:row['value'] for row in parsed['dimensions']},{
            'period':{'kind':'instant','date':'2027-03-31'},
            'asset_selection':{'mode':'members','members':['含;号、}和`的文字','普通成员']},'currency':'CNY'})


    def test_字段之间的逗号不是未引用取值而引号内逗号保留(self):
        from 附注更新.A合集 import _per_cell_answer
        reply='- 合成表!B3：metric_id `固定资产原值`；`entity`=甲公司，乙单位，`unit_scale`=1，`period`=2027年度（2027-01-01至2027-12-31），`currency`=`CNY`。'
        parsed=_per_cell_answer(reply,self.case.engine().standard,['合成表'])
        self.assertEqual({row['id']:row['value'] for row in parsed['dimensions']},{
            'entity':'甲公司，乙单位','unit_scale':1,'period':'2027年度（2027-01-01至2027-12-31）','currency':'CNY'})
        reply='- 合成表!B3：metric_id `固定资产原值`；`entity=甲公司，乙单位，`，`unit_scale=1`，`currency=CNY`。'
        parsed=_per_cell_answer(reply,self.case.engine().standard,['合成表'])
        self.assertEqual({row['id']:row['value'] for row in parsed['dimensions']},{
            'entity':'甲公司，乙单位，','unit_scale':1,'currency':'CNY'})


    def test_对象内代码标记只表示自身字面值(self):
        before,_=self.recognize('a',self.answer())
        reply='- 合成表!B3：`period`={`kind`:`instant`,`date`:`2028-02-29`}；`asset_selection`={`mode`:`members`,`members`:[`带"引号的成员`,"原字符串里的`符号"],`count`:`1`,`flag`:`false`}。'
        after,_=self.recognize('a',reply,{'a_mapping':before['mapping_path'],
            'a_field_scope':{'合成表!B3':['period','asset_selection']}})
        row=dict(_records(after))['合成表!B3']
        self.assertEqual(row['dimensions']['period'],{'kind':'instant','date':'2028-02-29'})
        self.assertEqual(row['dimensions']['asset_selection'],{'mode':'members',
            'members':['带"引号的成员','原字符串里的`符号'],'count':1,'flag':False})
        self.assertEqual(after['机械记录'],before['机械记录'])

    def test_AB明确null后的括号说明不成为取值(self):
        a,_=self.recognize('a',self.answer())
        b,_=self.recognize('b',self.answer(),{'a_mapping':a['mapping_path']})
        for side,before in [('a',a),('b',b)]:
            with self.subTest(side=side):
                reply='- 合成表!B3：`asset_selection`=null（原文未给范围）；`entity`="null（这是原文字串）"。'
                after,_=self.recognize(side,reply,{'a_mapping':a['mapping_path'],
                    side+'_mapping':before['mapping_path'],side+'_field_scope':{'合成表!B3':['asset_selection','entity']}})
                rows=dict(_records(after));row=rows['合成表!B3']
                self.assertNotIn('asset_selection',row['dimensions'])
                self.assertEqual(row['dimensions']['entity'],'null（这是原文字串）')
                self.assertEqual(row['missing_dimensions'],['asset_selection'])
                for loc,old in _records(before):
                    if loc!='合成表!B3':self.assertEqual(rows[loc]['dimensions'],old['dimensions'])
                self.assertEqual(after['机械记录'],before['机械记录'])

    def test_表格明确期间角色列只读取各行原值(self):
        from 附注更新.A合集 import _per_cell_answer
        reply='| 原位置 | metric_id | 维度实际取值 | template_period_role | 来源及原文依据 |\n|---|---|---|---|---|\n|合成表!B3|固定资产原值|currency=CNY|current|另一格为prior|\n|合成表!C3|固定资产原值|currency=CNY|prior|另一格为current|'
        parsed=_per_cell_answer(reply,self.case.engine().standard,['合成表'])
        self.assertEqual([(x['ranges'],x.get('template_period_role')) for x in parsed['metrics']],
                         [(['合成表!B3'],'current'),(['合成表!C3'],'prior')])


    def test_AB表格独立字段列保留完整原文不再按标点拆分(self):
        a,_=self.recognize('a',self.answer())
        b,_=self.recognize('b',self.answer(),{'a_mapping':a['mapping_path']})
        reply=('| 原位置 | entity | period | unit_scale | asset_selection | 原文依据 |\n'
               '|---|---|---|---|---|---|\n'
               '|合成表!B3|甲；currency=原文字。乙|{"date":"2028-02-29","kind":"instant"}|1|'
               'null（本格未提供成员）|currency=不得接入|\n'
               '|合成表!C3|1；后半句保留。|2028年度；原表注明期间|1|'
               '{mode:members,members:["含;分号。的原文"]}|来源原文|')
        for side,before in [('a',a),('b',b)]:
            with self.subTest(side=side):
                after,_=self.recognize(side,reply,{'a_mapping':a['mapping_path'],
                    side+'_mapping':before['mapping_path'],side+'_field_scope':{
                        '合成表!B3':['entity','period','unit_scale','asset_selection'],
                        '合成表!C3':['entity','period','unit_scale','asset_selection']}})
                rows=dict(_records(after))
                self.assertEqual(rows['合成表!B3']['dimensions']['entity'],'甲；currency=原文字。乙')
                self.assertEqual(rows['合成表!B3']['dimensions']['period'],{'date':'2028-02-29','kind':'instant'})
                self.assertEqual(rows['合成表!B3']['dimensions']['unit_scale'],1)
                self.assertNotIn('asset_selection',rows['合成表!B3']['dimensions'])
                self.assertEqual(rows['合成表!C3']['dimensions']['entity'],'1；后半句保留。')
                self.assertEqual(rows['合成表!C3']['dimensions']['period'],'2028年度；原表注明期间')
                self.assertEqual(rows['合成表!C3']['dimensions']['asset_selection'],
                    {'mode':'members','members':['含;分号。的原文']})
                for loc in ('合成表!B3','合成表!C3'):
                    self.assertEqual(rows[loc]['dimensions']['currency'],'CNY')
                self.assertEqual(after['机械记录'],before['机械记录'])


    def test_AB列表独立字段行保留完整文字及字面类型(self):
        a,_=self.recognize('a',self.answer())
        reply=('**合成表!B3**\n'
               '- entity：本格主体；保留完整限定。第二句仍是身份文字。\n'
               '- currency：`CNY`\n'
               '- unit_scale：`1`\n'
               '- asset_selection：null（本格成员未提供）\n'
               '- 依据：另一格不属于本次身份。')
        for side in ('a','b'):
            with self.subTest(side=side):
                prior=a if side=='a' else self.recognize('b',self.answer(),{'a_mapping':a['mapping_path']})[0]
                saved,_=self.recognize(side,reply,{'a_mapping':a['mapping_path'],side+'_mapping':prior['mapping_path'],
                    side+'_field_scope':{'合成表!B3':['entity','currency','unit_scale','asset_selection']}})
                row=dict(_records(saved))['合成表!B3']
                self.assertEqual(row['dimensions']['entity'],'本格主体；保留完整限定。第二句仍是身份文字。')
                self.assertEqual(row['dimensions']['currency'],'CNY')
                self.assertEqual(row['dimensions']['unit_scale'],1)
                self.assertIsNone(row['dimensions'].get('asset_selection'))

    def test_表格混合字段列末尾代码引号保留(self):
        from 附注更新.A合集 import _per_cell_answer
        reply=('| 原位置 | 身份字段实际取值 | 依据 |\n|---|---|---|\n'
               '|合成表!B3|metric_id=`固定资产原值`；currency=`CNY`；entity=`甲；保留后半句。`|本格|')
        parsed=_per_cell_answer(reply,self.case.engine().standard,['合成表'])
        self.assertEqual({x['id']:x['value'] for x in parsed['dimensions']},
            {'currency':'CNY','entity':'甲；保留后半句。'})


    def test_AB同层列表原字段保留且另组声明不接入(self):
        reply=('- **原位置：** `合成表!B3`\n'
               '- **metric_id：** `固定资产原值`\n'
               '- **currency：** `CNY`\n'
               '- **entity：** `本格主体`\n'
               '- 另组声明：currency=USD。\n'
               '- **entity：** `另一组主体`\n'
               '- **原位置：** `合成表!C3`\n'
               '- **metric_id：** `固定资产原值`\n'
               '- **currency：** `EUR`')
        a,_=self.recognize('a',reply)
        b,_=self.recognize('b',reply,{'a_mapping':a['mapping_path']})
        for mapping in (a,b):
            rows=dict(_records(mapping))
            self.assertEqual(rows['合成表!B3']['dimensions'],{'currency':'CNY','entity':'本格主体'})
            self.assertEqual(rows['合成表!C3']['dimensions'],{'currency':'EUR'})
            self.assertEqual(mapping['本次逐格原答复'],reply)

    def test_AB单格字段表只保存原始字段和值(self):
        from 附注更新.A合集 import _per_cell_answer
        body=('| 字段 | 实际取值 |\n|---|---|\n'
              '| metric_id | `固定资产原值` |\n'
              '| currency | `CNY` |\n'
              '| 币种 | `USD` |\n'
              '| unit_scale | `1` |\n'
              '| asset_selection | `{"mode":"unknown","known":"原文；保留。"}` |')
        reply='**原位置：合成表!B3**\n\n'+body
        a,_=self.recognize('a',reply)
        b,_=self.recognize('b',reply,{'a_mapping':a['mapping_path']})
        for mapping in (a,b):
            row=dict(_records(mapping))['合成表!B3']
            self.assertEqual(row['metric_id'],'固定资产原值')
            self.assertEqual(row['dimensions'],{'currency':'CNY','unit_scale':1,
                'asset_selection':{'mode':'unknown','known':'原文；保留。'}})
            self.assertEqual(mapping['本次逐格原答复'],reply)
        for no_cell in (body,reply.replace('合成表!B3','合成表!B3:B4')):
            with self.assertRaises(ValueError):
                _per_cell_answer(no_cell,self.case.engine().standard,['合成表'])

    def test_AB表格dimensions对象与原JSON身份证同值(self):
        from 附注更新.A合集 import _per_cell_answer
        dims={'entity':'甲公司（分部）','currency':'CNY','unit_scale':1,
              'period':{'kind':'instant','date':'2027-06-30'},
              'asset_selection':{'domain':'固定资产','mode':'members','members':['含分号；和句号。的项目']}}
        row={'metric_id':'固定资产原值','template_period_role':'closing','dimensions':dims}
        reply=('| 位置 | metric_id | template_period_role | dimensions | 依据 |\n|---|---|---|---|---|\n'
               '| 合成表!B3 | 固定资产原值 | closing | `'+json.dumps(dims,ensure_ascii=False)+'` | 本格原文 |')
        standard=self.case.engine().standard
        expected=_per_cell_answer({'合成表!B3':row},standard,['合成表'])
        actual=_per_cell_answer(reply,standard,['合成表'])
        self.assertEqual({d['id']:d['value'] for d in actual['dimensions']},
                         {d['id']:d['value'] for d in expected['dimensions']})
        a,_=self.recognize('a',reply)
        b,_=self.recognize('b',reply,{'a_mapping':a['mapping_path']})
        for mapping in (a,b):
            self.assertEqual(dict(_records(mapping))['合成表!B3']['dimensions'],dims)
            self.assertEqual(mapping['本次逐格原答复'],reply)
        for invalid in (reply.replace('合成表!B3','合成表!B3:B4'),
                        '| dimensions |\n|---|\n| `'+json.dumps(dims,ensure_ascii=False)+'` |'):
            with self.assertRaises(ValueError):_per_cell_answer(invalid,standard,['合成表'])

    def test_AB逐行JSON与列表保存同一原字段且不接半份对象(self):
        from 附注更新.A合集 import _per_cell_answer
        records=[{'location':loc,**row} for loc,row in json.loads(self.answer()).items()]
        records[0]['entity']='原文字；第二句。\n仍属于本格'
        standard=self.case.engine().standard
        reply='\n\n'.join(json.dumps(row,ensure_ascii=False) for row in records)
        self.assertEqual(_per_cell_answer(reply,standard,['合成表']),
                         _per_cell_answer(records,standard,['合成表']))
        a,_=self.recognize('a',reply)
        b,_=self.recognize('b',reply,{'a_mapping':a['mapping_path']})
        for mapping in (a,b):
            actual=dict(_records(mapping))
            for row in records:
                self.assertEqual(actual[row['location']]['metric_id'],row['metric_id'])
                self.assertEqual(actual[row['location']]['dimensions'],
                                 {key:value for key,value in row.items() if key in standard['dimension']})
        with self.assertRaises(ValueError):
            _per_cell_answer(reply+'\n{"location":',standard,['合成表'])

    def test_AB独立原字段行与JSON对象保存同一身份(self):
        from 附注更新.A合集 import _per_cell_answer
        records=[{'location':loc,**row} for loc,row in json.loads(self.answer()).items()][:2]
        for row in records:
            row['dimensions']={k:row.pop(k) for k in list(row) if k in self.case.engine().standard['dimension']}
        records[0]['dimensions']['entity']='本格；currency=USD。仍是主体原文'
        records[1]['dimensions']['asset_selection']=None
        reply='\n\n'.join('- `location`: `'+row['location']+'`\n'+'\n'.join(
            '  `'+key+'`: `'+(value if isinstance(value,str) else json.dumps(value,ensure_ascii=False))+'`'
            for key,value in row.items() if key!='location')+'\n  依据：本格原文。' for row in records)
        standard=self.case.engine().standard
        parsed=_per_cell_answer(reply,standard,['合成表'])
        expected=_per_cell_answer(records,standard,['合成表'])
        for kind in ('metrics','dimensions','unresolved'):
            self.assertEqual([{k:v for k,v in row.items() if k!='original'} for row in parsed[kind]],
                             [{k:v for k,v in row.items() if k!='original'} for row in expected[kind]])
        a,_=self.recognize('a',reply)
        b,_=self.recognize('b',reply,{'a_mapping':a['mapping_path']})
        for mapping in (a,b):
            actual=dict(_records(mapping))
            self.assertEqual(actual['合成表!B3']['dimensions']['entity'],records[0]['dimensions']['entity'])
            self.assertEqual(actual['合成表!B3']['dimensions']['currency'],'CNY')
            self.assertNotIn('asset_selection',actual['合成表!C3']['dimensions'])
            self.assertEqual(mapping['本次逐格原答复'],reply)
        with self.assertRaises(ValueError):
            _per_cell_answer(reply.replace('合成表!B3','合成表!B3:B4').replace('合成表!C3','合成表!C3:C4'),standard,['合成表'])

    def test_编号代码引用后实际取值不成为第二个字段(self):
        from 附注更新.A合集 import _per_cell_answer
        standard=self.case.engine().standard
        reply='- `合成表!B3`：`dimension_id`=`currency`；实际取值=`CNY`。依据：本格原文。'
        parsed=_per_cell_answer(reply,standard,['合成表'])
        self.assertEqual([row['id'] for row in parsed['dimension_bindings']],['currency'])
        self.assertEqual(parsed['dimensions'],[])
        self.assertEqual(parsed['unresolved'],[])
        nested='- 合成表!B3：dimension_id: template_period_role\n  dimensions: {"template_period_role":"closing"}'
        parsed=_per_cell_answer(nested,standard,['合成表'])
        self.assertEqual(parsed['dimension_bindings'],[])
        self.assertEqual([(row['id'],row['value']) for row in parsed['unresolved']],[('dimension_id','template_period_role')])
        self.assertEqual([(row['id'],row['value']) for row in parsed['dimensions']],[('template_period_role','closing')])

    def test_程序不把字段别名或裸编号补成身份证(self):
        from 附注更新.A合集 import _per_cell_answer
        standard=self.case.engine().standard
        standard['dimension']['currency']['name']='币种'
        for reply in (
            '- 合成表!B3：金标准：`固定资产原值`。',
            '- 合成表!B3：匹配标准：`固定资产原值`。',
            '- 合成表!B3：币种：CNY。',
            '- 合成表!B3：`固定资产原值`。',
            {'合成表!B3':{'standard':'固定资产原值'}},
            {'合成表!B3':{'actual_dimensions':{'currency':'CNY'}}},
        ):
            with self.subTest(reply=reply),self.assertRaises(ValueError):
                _per_cell_answer(reply,standard,['合成表'])

    def test_AB只保存模型直接给出的原始字段和取值(self):
        alias='币种'
        reply=f'- 合成表!B3：metric_id=固定资产原值；{alias}=USD；template_period_role=`期末`；entity=`保留原说法`。'
        a,_=self.recognize('a',reply)
        b,_=self.recognize('b',reply,{'a_mapping':a['mapping_path']})
        for mapping in (a,b):
            row=dict(_records(mapping))['合成表!B3']
            self.assertEqual(row['dimensions'],{'entity':'保留原说法'})
            roles=[x['value'] for t in mapping['outline']['subtables'] for x in t['dimensions'] if x['id']=='template_period_role']
            self.assertEqual(roles,['期末'])
            self.assertIn(reply,mapping['本次逐格原答复'])

    def test_AB外层列表连续原生对象原样读取且拒绝残缺(self):
        from 附注更新.A合集 import _per_cell_answer
        records=[{'currency':'USD'},
            {'location':'合成表!B3','metric_id':'固定资产原值',
             'template_period_role':'closing','dimensions':{'period':'保留原日期表达',
             'asset_selection':{'mode':'members','members':['含},[括号]与逗号的原类别']}}},
            {'location':'合成表!C3','metric_id':'固定资产原值',
             'template_period_role':'opening','dimensions':{'currency':'CNY'}}]
        chunks=[json.dumps(x,ensure_ascii=False) for x in records]
        reply='[\n'+'\n'.join(chunks)+'\n]'
        standard=self.case.engine().standard
        expected=_per_cell_answer(records,standard,['合成表'])
        self.assertEqual(_per_cell_answer(reply,standard,['合成表']),expected)
        arrays=json.dumps(records[:2],ensure_ascii=False)+'\n'+json.dumps(records[2:],ensure_ascii=False)
        self.assertEqual(_per_cell_answer(arrays,standard,['合成表']),expected)
        self.assertNotIn('currency',{x['id'] for x in expected['dimensions'] if x['ranges']==['合成表!B3']})
        for bad in (reply[:-1],reply[:-1]+'{"location":\n]',reply[:-1]+',]'):
            with self.subTest(bad=bad),self.assertRaises(ValueError):
                _per_cell_answer(bad,standard,['合成表'])

    def test_AB连续JSON对象保留逐格字段并拒绝半份对象(self):
        from 附注更新.A合集 import _per_cell_answer
        from 附注更新.表格 import file_hash
        records=[
            {'location':'合成表!B3','dimension_id':'template_period_role',
             'value':'closing','period':'2031-12-31'},
            {'location':'合成表!C3','dimension_id':None,'currency':'CNY'},
            {'location':'合成表!B4','dimension_id':'asset_selection',
             'dimensions':{'asset_selection':{'mode':'members','members':['包含},逗号和"文字的类别']}}}]
        chunks=[json.dumps(row,ensure_ascii=False) for row in records]
        reply=chunks[0]+', '+chunks[1]+'\n\n'+chunks[2]
        standard=self.case.engine().standard
        self.assertEqual(_per_cell_answer(reply,standard,['合成表']),
                         _per_cell_answer(records,standard,['合成表']))
        partial=_per_cell_answer(chunks[0]+','+chunks[1],standard,['合成表'])
        self.assertEqual([(row['id'],row['value']) for row in partial['dimensions']],
                         [('period','2031-12-31'),('currency','CNY')])
        self.assertEqual(partial['dimension_bindings'],[])
        for malformed in (reply+',',reply+'\n{"location":',reply+' 未完成'):
            with self.subTest(malformed=malformed),self.assertRaises(ValueError):
                _per_cell_answer(malformed,standard,['合成表'])
        unowned=_per_cell_answer('{"currency":"USD"},'+chunks[2],standard,['合成表'])
        self.assertNotIn('currency',{row['id'] for row in unowned['dimensions']})
        a,_=self.recognize('a',reply)
        b,_=self.recognize('b',reply,{'a_mapping':a['mapping_path']})
        for mapping in (a,b):
            actual=dict(_records(mapping))
            fields={(table['sheet']+'!'+cell,row['id']):row.get('value')
                    for table in mapping['outline']['subtables'] for row in table['dimensions']
                    if row.get('逐格原字段') is True for cell in row['ranges']}
            self.assertEqual(fields[('合成表!B3','period')],'2031-12-31')
            self.assertNotIn(('合成表!B3','template_period_role'),fields)
            self.assertEqual(fields[('合成表!C3','currency')],'CNY')
            self.assertEqual(actual['合成表!B4']['dimensions'],records[2]['dimensions'])
            self.assertEqual(mapping['本次逐格原答复'],reply)
        self.assertEqual(self.case.hashes,(file_hash(self.case.source),file_hash(self.case.gold)))


    def test_AB裸编号值不是实际维度字段且不接受取值别名(self):
        from 附注更新.A合集 import _per_cell_answer
        standard=self.case.engine().standard
        value=json.dumps({'mode':'members','members':['仅原文保存，不得推字段']},ensure_ascii=False)
        replies=[
            '- 合成表!B3：dimension_id=asset_selection；实际取值='+value,
            '合成表!B3\nmetric_id=null\ndimensions：dimension_id=asset_selection；实际取值='+value,
            '| 位置 | 本格身份证与维度实际取值 | 来源和原文依据 |\n|---|---|---|\n'
            '| `合成表!B3` | `metric_id=null`；`dimension_id=asset_selection`，`value='+value+'` | 原文 |',
        ]
        for reply in replies:
            with self.subTest(reply=reply):
                parsed=_per_cell_answer(reply,standard,['合成表'])
                self.assertEqual([r['id'] for r in parsed['dimension_bindings']],['asset_selection'])
                self.assertEqual(parsed['dimensions'],[])
        a,_=self.recognize('a',replies[0])
        b,_=self.recognize('b',replies[0],{'a_mapping':a['mapping_path']})
        for mapping in (a,b):
            row=dict(_records(mapping))['合成表!B3']
            self.assertEqual(row['dimension_id'],'asset_selection')
            self.assertEqual(row['dimensions'],{})
            self.assertEqual(mapping['本次逐格原答复'],replies[0])
        explicit=_per_cell_answer('- 合成表!B3：dimension_id=asset_selection；asset_selection='+value,standard,['合成表'])
        self.assertEqual([r['value'] for r in explicit['dimensions']],[json.loads(value)])

    def test_AB对象成员正文弯引号原样保留(self):
        from 附注更新.A合集 import _per_cell_answer
        expected={'mode':'members','members':['设备甲（俗称“设备”）','设备乙（标为‘备选’）']}
        reply='- 合成表!B3：dimension_id=asset_selection；asset_selection={mode:members,members:[设备甲（俗称“设备”）,设备乙（标为‘备选’）]}'
        standard=self.case.engine().standard
        parsed=_per_cell_answer(reply,standard,['合成表'])
        self.assertEqual(parsed['dimensions'][0]['value'],expected)
        quoted='- 合成表!B3：dimension_id=asset_selection；asset_selection='+json.dumps(expected,ensure_ascii=False)
        self.assertEqual(_per_cell_answer(quoted,standard,['合成表'])['dimensions'][0]['value'],expected)
        delimiters='- 合成表!B3：dimension_id=asset_selection；asset_selection={mode:members,members:[“设备甲”,‘设备乙’]}'
        self.assertEqual(_per_cell_answer(delimiters,standard,['合成表'])['dimensions'][0]['value'],{'mode':'members','members':['设备甲','设备乙']})
        a,_=self.recognize('a',reply)
        b,_=self.recognize('b',reply,{'a_mapping':a['mapping_path']})
        for mapping in (a,b):
            row=dict(_records(mapping))['合成表!B3']
            self.assertEqual(row['dimensions']['asset_selection'],expected)
            self.assertEqual(mapping['本次逐格原答复'],reply)

    def test_AB原生location等号位置逐格读取(self):
        from 附注更新.A合集 import _per_cell_answer
        standard=self.case.engine().standard
        own={'asset_selection':{'mode':'members','members':['设备甲（保留原文）']},'currency':'CNY'}
        for separator in ('=','＝'):
            reply=('dimensions={"currency":"EUR"}\n'
                   '- location'+separator+'`合成表!B3`；dimension_id=`asset_selection`；dimensions=`'+json.dumps(own,ensure_ascii=False)+'`。依据：仅本格原文。\n'
                   '- location'+separator+'`合成表!C3`；dimension_id=`currency`；dimensions=`{"currency":"USD"}`。依据：另一格原文。')
            with self.subTest(separator=separator):
                parsed=_per_cell_answer(reply,standard,['合成表'])
                self.assertEqual([(r['ranges'],r['id']) for r in parsed['dimension_bindings']],
                                 [(['合成表!B3'],'asset_selection'),(['合成表!C3'],'currency')])
                a,_=self.recognize('a',reply)
                b,_=self.recognize('b',reply,{'a_mapping':a['mapping_path']})
                for mapping in (a,b):
                    records=dict(_records(mapping))
                    self.assertEqual(records['合成表!B3']['dimensions'],own)
                    self.assertEqual(records['合成表!C3']['dimensions'],{'currency':'USD'})
                    self.assertEqual(mapping['本次逐格原答复'],reply)
                with self.assertRaises(ValueError):
                    _per_cell_answer('- location'+separator+'`合成表!B3:B4`；dimension_id=`currency`；currency=`USD`',standard,['合成表'])

class 逐格代码赋值读取测试(unittest.TestCase):
    def standard(self):
        return {'metric': {'固定资产原值': {}}, 'dimension':
            {key: {} for key in ('entity', 'currency', 'unit_scale', 'period', 'asset_selection')}}

    def test_dimensions代码赋值首字段不丢失且原值不截断(self):
        from 附注更新.A合集 import _per_cell_answer
        values = {'entity': '原主体；currency=USD。仍属主体原文', 'currency': 'CNY',
            'unit_scale': 1, 'period': '2027年度',
            'asset_selection': {'mode': 'members', 'members': ['类别甲；乙']}}
        for first in values:
            with self.subTest(first=first):
                keys = [first] + [key for key in values if key != first]
                parts = ['`'+key+'='+(value if isinstance(value, str) else json.dumps(value, ensure_ascii=False))+'`'
                    for key in keys for value in [values[key]]]
                reply = ('- **合成表!B3**\n  - `metric_id`: `固定资产原值`\n'
                    '  - `dimensions`: '+'；'.join(parts))
                parsed = _per_cell_answer(reply, self.standard(), ['合成表'])
                self.assertEqual({row['id']: row['value'] for row in parsed['dimensions']}, values)
                self.assertEqual([row['id'] for row in parsed['metrics']], ['固定资产原值'])

    def test_dimensions引用说明和共同声明不补本格字段(self):
        from 附注更新.A合集 import _per_cell_answer
        reply = ('共同声明：currency=CNY。\n- 合成表!B3\n'
            '  - metric_id: `固定资产原值`\n  - dimensions: `引用 currency=USD 的说明`\n'
            '- 合成表!B4\n  - metric_id: `固定资产原值`')
        parsed = _per_cell_answer(reply, self.standard(), ['合成表'])
        self.assertEqual(parsed['dimensions'], [])
        self.assertEqual([row['ranges'] for row in parsed['metrics']], [['合成表!B3'], ['合成表!B4']])
        for value in ({'币种': 'CNY'}, {'mode': 'members', 'members': ['甲']}):
            parsed = _per_cell_answer({'合成表!B3': {'metric_id': '固定资产原值', 'dimensions': value}},
                self.standard(), ['合成表'])
            self.assertEqual(parsed['dimensions'], [])

    def test_dimensions代码只包字段名时仍保存原字段(self):
        from 附注更新.A合集 import _per_cell_answer
        reply = ('- 合成表!B3\n  - metric_id: `固定资产原值`\n'
            '  - dimensions: `currency`: `CNY`；`unit_scale`: `1`')
        parsed = _per_cell_answer(reply, self.standard(), ['合成表'])
        self.assertEqual({row['id']: row['value'] for row in parsed['dimensions']},
            {'currency': 'CNY', 'unit_scale': 1})


class 逐格交付接入状态测试(unittest.TestCase):
    def check_reply(self, reply, fields):
        from 附注更新.A合集 import _per_cell_answer, direct_ingestion
        standard={'metric':{'计量槽':{'dimension_refs':[{'id':'currency'}]}},
                  'dimension':{'currency':{},'period':{},'其他维度':{}}}
        answer=_per_cell_answer(reply,standard,['合成表'])
        return answer,direct_ingestion(answer,{'合成表!B3':fields},standard,['合成表'])

    def test_原字段原值直接接入且AB共用(self):
        fields=['metric_id','currency']
        for reply in ('合成表!B3：metric_id=计量槽；currency=CNY',
                      {'合成表!B3':{'metric_id':'计量槽','dimensions':{'currency':'CNY'}}}):
            answer,receipt=self.check_reply(reply,fields)
            self.assertTrue(receipt['complete'])
            self.assertEqual(answer['metrics'][0]['id'],'计量槽')
            self.assertEqual(answer['dimensions'][0]['value'],'CNY')

    def test_别名未交付不能借已有维度报完成(self):
        answer,receipt=self.check_reply('合成表!B3：金标准编号=计量槽；currency=CNY',
                                        ['metric_id','currency'])
        self.assertEqual(answer['metrics'],[])
        self.assertFalse(receipt['complete'])
        self.assertEqual(receipt['未接入字段'],[{'location':'合成表!B3','field':'metric_id'}])

    def test_未知维度不撤销槽位也不增加无关属性门槛(self):
        answer,receipt=self.check_reply({'合成表!B3':{'metric_id':'计量槽','currency':None,
                                                         'template_period_role':None}},
            ['metric_id','dimension_id','currency','period','其他维度','template_period_role'])
        self.assertEqual(answer['metrics'][0]['id'],'计量槽')
        self.assertTrue(receipt['complete'])
        self.assertIsNone(answer['dimensions'][0]['value'])

    def test_缺交字段和错误维度包装不补写(self):
        answer,receipt=self.check_reply({'合成表!B3':{'metric_id':'计量槽',
            'dimensions':[{'dimension_id':'currency','value':'CNY'}]}},['metric_id','currency'])
        self.assertFalse(receipt['complete'])
        self.assertEqual(answer['dimensions'],[])
        self.assertEqual(receipt['未接入字段'],[{'location':'合成表!B3','field':'currency'}])

    def test_只核对本次范围且保留范围外编号问题(self):
        answer,receipt=self.check_reply({'合成表!B3':{'dimension_id':'template_period_role',
            'template_period_role':'closing'}},['template_period_role'])
        self.assertTrue(receipt['complete'])
        self.assertEqual(receipt['未确定字段'],[])
        self.assertEqual(receipt['范围外未接入问题'][0]['field'],'dimension_id')
        self.assertEqual(answer['dimension_bindings'],[])
        self.assertEqual(answer['dimensions'][0]['value'],'closing')
        answer,receipt=self.check_reply({'合成表!B3':{'metric_id':None,'currency':'CNY'}},['metric_id'])
        self.assertFalse(receipt['complete'])
        self.assertEqual(receipt['未确定字段'][0]['field'],'metric_id')

    def test_范围外格子的未决不伪装成本次缺交(self):
        from 附注更新.A合集 import _per_cell_answer,direct_ingestion
        standard={'metric':{'计量槽':{}},'dimension':{'currency':{}}}
        parsed=_per_cell_answer({'合成表!B3':{'currency':'CNY'},'合成表!B4':{'metric_id':None}},standard,['合成表'])
        receipt=direct_ingestion(parsed,{'合成表!B3':['currency']},standard,['合成表'])
        self.assertTrue(receipt['complete'])
        self.assertEqual(receipt['未接入字段'],[])
        self.assertEqual(receipt['范围外未接入问题'][0]['location'],'合成表!B4')

    def test_无接入记录或编号未知不能报完成(self):
        from 附注更新.A合集 import direct_recognition_status
        self.assertEqual(direct_recognition_status({'collection_version':2}),'partial')
        answer,receipt=self.check_reply({'合成表!B3':{'metric_id':None}},['metric_id'])
        self.assertFalse(receipt['complete'])
        self.assertEqual(receipt['未接入字段'],[])
        self.assertEqual(direct_recognition_status({'本次逐格接入':receipt}),'partial')


if __name__=='__main__':
    unittest.main()
