# coding: utf-8
"""已有声明的保存、比较和更新依赖；全部使用合成材料，不调用模型或写回。"""
import json
import unittest
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

from 测试 import 测试直接识别 as fixture
from 附注更新.A合集 import save_direct_answer, _records, direct_digest, combine_direct_collections, direct_answer
from 附注更新.A语义阅读 import direct_reference
from 附注更新.模板更新编排 import direct_update_plan
from 附注更新.表格 import read_workbook
from 附注更新.分步流程 import run_step
from 附注更新.统一识别 import UnifiedSemanticEngine


class 合集关系贯通测试(unittest.TestCase):
    def setUp(self):
        self.case=fixture.直接识别测试();self.case.setUp()
        self.engine=self.case.engine();self.index=0
        self.snapshot=read_workbook(self.case.source,context=[])
        self.base={'schema_version':2,'recognition_method':'direct','source_path':str(self.case.source),
            'source_hash':self.snapshot['sha256'],'source_context':[],'gold_path':str(self.case.gold),
            'gold_hash':self.engine.gold_hash,'carrier':'A','金标准原文':'\n'.join(
                line for line in self.case.gold.read_text(encoding='utf-8').splitlines()
                if json.loads(line)['record_type'] in {'metric','dimension'})}

    def save(self,carrier,answer=None,snapshot=None):
        self.index+=1
        return save_direct_answer({**self.base,'carrier':carrier},answer or self.case.answer(),
            self.case.folder/f'{carrier}{self.index}.json',snapshot=snapshot or self.snapshot,standard=self.engine.standard)

    def corrected(self,a,b,text):
        self.index+=1
        return save_direct_answer(b,text,self.case.folder/f'引用接入{self.index}.json',
            field_scope={'合成表!B3:C3':['asset_selection']},reference=direct_reference(
                a,compact=True,ranges=[{'sheet':'合成表','range':'B3:C3'}]))

    def test_主入口来源答复随合集接入匹配且不改身份(self):
        from 附注更新.统一勾稽 import load_fact_mapping
        a=self.save('A');b=self.save('B')
        reply='A 合成表!B3：采用 B 合成表!C3，依据模型已经明确的来源关系。\nA 合成表!B4不能采用 B 合成表!C4，原文不足。'
        config={**self.case.config,'a_mapping':a['mapping_path'],'b_mapping':b['mapping_path'],
                'update_purpose':{'reason':'A作模板，日期随B更新'},
                'comparison_requirements':[{'维度ID':'entity','要求含义':'本次主体名称不参与比较，实际身份照实保留。'}],
                'b_source_targets':['合成表!B3','合成表!B4'],
                'b_context_text':'本次只明确标签动作，不重做金额来源选择。',
                'b_source_candidates':{'合成表!B3':['合成表!C3']},
                'a_required_ranges':[{'sheet':'合成表','range':'B3:B4'}],
                'b_source_ranges':[{'sheet':'合成表','range':'A2:C8'}]}
        with patch.object(UnifiedSemanticEngine,'request_text',return_value=reply) as request:
            run=run_step('recognize_b',config)
        self.assertEqual(run['status'],'complete',run)
        request.assert_called_once()
        self.assertIn(config['b_context_text'],request.call_args.args[0])
        materials=request.call_args.args[1]
        self.assertIn(config['update_purpose']['reason'],request.call_args.args[0])
        self.assertIn(config['comparison_requirements'][0]['要求含义'],request.call_args.args[0])
        self.assertNotIn('本次明确更新要求.json',materials)
        self.assertNotIn('A需要.json',materials)
        self.assertEqual(materials['已保存的当前候选对应']['合成表!B3'][0]['dimensions'],dict(_records(b))['合成表!C3']['dimensions'])
        saved=load_fact_mapping(run['files']['b_mapping'],{})
        self.assertEqual(saved['facts'],b['facts'])
        self.assertEqual(saved['slot_bindings'],b['slot_bindings'])
        self.assertEqual(saved['已确认关系'][-1]['bindings'],[{'target':'合成表!B3','source':'合成表!C3'}])
        self.assertEqual(Path(run['files']['recognition_answer']).read_text(encoding='utf-8-sig'),reply)
        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('匹配不能再识别')):
            matched=run_step('match',{**config,'b_mapping':saved['mapping_path']})
        self.assertEqual(matched['status'],'complete',matched)
        plan=json.loads(Path(matched['files']['update_plan']).read_text(encoding='utf-8-sig'))
        update=next(row for row in plan['updates'] if row['cell']=='B3')
        self.assertEqual([row['fact_id'] for row in update['sources']],['合成表!C3'])
        from 附注更新.A合集 import save_source_relations
        table='| A原位置 | 采用的B来源 | 依据 |\n|---|---|---|\n| 合成表!B3 | 合成表!C3 | 已明确来源 |'
        tabular=save_source_relations(a,b,table,['合成表!B3'],self.case.folder/'自然表格对应.json')
        self.assertEqual(tabular['已确认关系'][-1]['bindings'],saved['已确认关系'][-1]['bindings'])
        paired=('| A 原位置 → 更新后位置 | B 来源：本期 / 上期 |\n|---|---|\n'
                '| B3/C3 → B4/C4 | 合成表!B5 / C5 |\n| B4/C4 → B5/C5 | …!B6 / C6 |\n'
                '表中“…”均指 合成表。\n原 B8/C8 更新后为 B9/C9，应取合计：合成表!B8 / C8。')
        pairs=save_source_relations(a,b,paired,['合成表!'+x for x in ['B3','C3','B4','C4','B8','C8']],self.case.folder/'成对位置.json')
        self.assertEqual(pairs['已确认关系'][-1]['bindings'],[
            {'target':'合成表!'+target,'source':'合成表!'+source}
            for target,source in [('B3','B5'),('C3','C5'),('B4','B6'),('C4','C6'),('B8','B8'),('C8','C8')]])

    def test_来源识别只携带本次表区的已有名单说明(self):
        snapshot=deepcopy(self.snapshot)
        snapshot['names']={'当前名单':[{'sheet':'合成表','range':'A2:C8'}]}
        a=self.save('A',snapshot=snapshot);b=self.save('B')
        b['已确认关系']=[{'kind':'名单编排','content':text,'tables':[{'range_name':name}]}
                         for name,text in [('当前名单','本次需要的名单说明'),('别张名单','无关历史名单说明')]]
        b['封存校验']=direct_digest(b)
        path=self.case.folder/'含两张名单的已有B.json';path.write_text(json.dumps(b,ensure_ascii=False),encoding='utf-8')
        config={**self.case.config,'a_mapping':a['mapping_path'],'b_mapping':str(path),
                'b_source_targets':['合成表!B3'],'a_required_ranges':[{'sheet':'合成表','range':'B3'}],
                'b_source_ranges':[{'sheet':'合成表','range':'A2:C8'}],
                'b_source_candidates':{'合成表!B3':['合成表!B3']}}
        with patch.object(UnifiedSemanticEngine,'request_text',return_value='A 合成表!B3：采用 B 合成表!B3。') as request:
            result=run_step('recognize_b',config)
        self.assertEqual(result['status'],'complete',result);request.assert_called_once()
        self.assertEqual(request.call_args.args[1]['已有名单安排'],['本次需要的名单说明'])

    def test_直接值目标不展开旁支计算而公式目标保留引用链(self):
        from openpyxl import load_workbook
        book=load_workbook(self.case.source)
        book['合成表']['C3']="='引用来源'!B2"
        refs=book.create_sheet('引用来源');refs['A2']='实际来源项目';refs['B2']="=SUM('计算明细'!B2:B9)"
        detail=book.create_sheet('计算明细');detail['A1']='原始计算明细'
        for row in range(2,10):detail.cell(row,2,row)
        source=self.case.folder/'含公式链的合成原表.xlsx';book.save(source);book.close()
        self.snapshot=read_workbook(source,context=[])
        self.base.update(source_path=str(source),source_hash=self.snapshot['sha256'])
        a=self.save('A');b=self.save('B')
        for cell,recursive in [('B3',False),('C3',True)]:
            config={**self.case.config,'a_path':str(source),'b_path':str(source),
                    'a_mapping':a['mapping_path'],'b_mapping':b['mapping_path'],
                    'a_required_ranges':[{'sheet':'合成表','range':cell}],
                    'b_source_ranges':[{'sheet':'合成表','range':'A2:C4'}],
                    'b_field_scope':{'合成表!'+cell:['metric_id']}}
            with patch.object(UnifiedSemanticEngine,'request_text',return_value=f'合成表!{cell}：metric_id=固定资产原值。') as request:
                result=run_step('recognize_b',config)
            self.assertEqual(result['status'],'complete',result);request.assert_called_once()
            material=request.call_args.args[1]['原文.json'];proof=material['公式及来源']
            self.assertEqual(proof['公式原文']['引用来源!B2'],"=SUM('计算明细'!B2:B9)")
            self.assertEqual(any(r['sheet']=='计算明细' for r in proof['来源行列原文']),recursive)
            self.assertIn('实际来源项目',json.dumps(proof,ensure_ascii=False))
            self.assertEqual(material['本次读取范围'],[{'sheet':'合成表','range':'A2:C4','data_ranges':[]}])

    def test_明确搬运关系不要求来源另存身份且不补身份(self):
        from 附注更新.A合集 import source_relation,direct_output_record,direct_digest
        a=self.save('A');b=self.save('B')
        b['facts']=[row for row in b['facts'] if row['source_reference']['location']!='合成表!C3']
        b['slot_bindings']=[row for row in b['slot_bindings'] if row['sheet']+'!'+row['cell']!='合成表!C3']
        b.setdefault('已确认关系',[]).append(source_relation(a,b,
            'A 合成表!B3：采用 B 合成表!C3。',['合成表!B3']))
        b['封存校验']=direct_digest(b)
        plan=direct_update_plan(a,b)
        update=next(row for row in plan['updates'] if row['cell']=='B3')
        self.assertEqual(update['value'],b['机械记录']['sheets'][0]['cells']['C3']['value'])
        self.assertEqual(update['metric_id'],dict(_records(a))['合成表!B3']['metric_id'])
        self.assertEqual(update['source_identity']['dimensions'],{})
        with patch('附注更新.表格.read_workbook',return_value=deepcopy(self.snapshot)),\
             patch('附注更新.表格.file_hash',return_value='c'*64):
            derived=direct_output_record(a,b,self.case.folder/'模拟副本.xlsx',plan,'已保存安排')
        self.assertNotIn('facts',derived)
        self.assertTrue(any(row.get('cell')=='B3' and row.get('source_locations')==['合成表!C3']
                            for row in derived['identity_matches']))
        self.assertIn('合成表!B3',dict(_records(a)))

    def test_跨槽位搬运保留双方原身份而不拼接(self):
        from 附注更新.A合集 import source_relation,direct_output_record,_standard
        a=self.save('A');b=self.save('B',self.case.answer(2026))
        standard=_standard(a)
        source_metric=deepcopy(standard['metric']['固定资产原值']);source_metric['id']='另一个来源槽位'
        source_metric['dimension_refs']=[r for r in source_metric['dimension_refs'] if r['id']!='asset_selection']
        source_metric['dimension_refs'].append({'id':'来源专属范围','required':True})
        standard['metric'][source_metric['id']]=source_metric
        for mapping in (a,b):
            mapping['封存定义原文'].append(json.dumps(source_metric,ensure_ascii=False))
        record=dict(_records(b))['合成表!C3'];record['metric_id']=source_metric['id']
        record['dimensions'].pop('asset_selection');record['dimension_evidence'].pop('asset_selection')
        record['dimensions']['来源专属范围']={'mode':'members','members':['来源项']}
        b.setdefault('已确认关系',[]).append(source_relation(a,b,'A 合成表!B3：采用 B 合成表!C3。',['合成表!B3']))
        for mapping in (a,b):mapping['封存校验']=direct_digest(mapping)
        before=deepcopy((a,b));plan=direct_update_plan(a,b)
        update=next(row for row in plan['updates'] if row['cell']=='B3')
        original=dict(_records(a))['合成表!B3']
        self.assertEqual(update['metric_id'],original['metric_id'])
        self.assertEqual(update['original_dimensions'],original['dimensions'])
        self.assertEqual(update['source_identity'],{'metric_id':source_metric['id'],'dimensions':record['dimensions']})
        self.assertNotIn('dimensions',update)
        self.assertEqual(update['value'],b['机械记录']['sheets'][0]['cells']['C3']['value'])
        with patch('附注更新.表格.read_workbook',return_value=deepcopy(self.snapshot)),\
             patch('附注更新.表格.file_hash',return_value='c'*64):
            derived=direct_output_record(a,b,self.case.folder/'模拟副本.xlsx',plan,'已保存安排')
        self.assertNotIn('facts',derived)
        self.assertNotIn('recognition_method',derived)
        self.assertNotIn('识别答复',derived)
        self.assertEqual(derived['collections']['A']['mapping_hash'],a['mapping_hash'])
        self.assertEqual(derived['collections']['B']['mapping_hash'],b['mapping_hash'])
        saved=next(row for row in derived['updates'] if row['cell']=='B3')
        self.assertEqual(saved['source_identity'],update['source_identity'])
        self.assertEqual(saved['original_dimensions'],original['dimensions'])
        self.assertEqual((a,b),before)

    def test_明确保留原值也进入执行记录(self):
        from 附注更新.A合集 import direct_output_record
        a=self.save('A');b=self.save('B');before=deepcopy((a,b))
        plan={'retained_values':[{'sheet':'合成表','cell':'B3','reason':'已有答复明确保留原值'}]}
        with patch('附注更新.表格.read_workbook',return_value=deepcopy(self.snapshot)),\
             patch('附注更新.表格.file_hash',return_value=self.snapshot['sha256']):
            record=direct_output_record(a,b,self.case.source,plan,'已保存安排')
        self.assertEqual(record['retained_values'][0]['actual_value'],self.snapshot['sheets'][0]['cells']['B3']['value'])
        self.assertEqual(record['retained_values'][0]['reason'],plan['retained_values'][0]['reason'])
        self.assertNotIn('facts',record)
        self.assertEqual((a,b),before)

    def test_来源认识保留A原表上下文及空白(self):
        a=self.save('A');b=self.save('B')
        config={**self.case.config,'a_mapping':a['mapping_path'],'b_mapping':b['mapping_path'],
                'b_source_targets':['合成表!B3'],
                'a_required_ranges':[{'sheet':'合成表','range':'A2:C8'}],
                'b_source_ranges':[{'sheet':'合成表','range':'A2:C8'}]}
        with patch.object(UnifiedSemanticEngine,'request_text',return_value='A 合成表!B3：采用 B 合成表!B3。') as request:
            run=run_step('recognize_b',config)
        self.assertEqual(run['status'],'complete',run)
        prompt,materials=request.call_args.args
        self.assertIn('空白也是业务位置',prompt)
        self.assertNotIn('A原表.json',materials)
        self.assertNotIn('A原表网格.txt',materials)
        self.assertIn('A需要.json',materials)
        request.assert_called_once()

    def test_期间标签调用只携带完整期间定义(self):
        a=self.save('A');b=self.save('B')
        config={**self.case.config,'a_mapping':a['mapping_path'],'b_mapping':b['mapping_path'],
            'b_field_scope':{'合成表!B3':['period']},'b_additional_context':'period只补明确的原表标签引用，不改日期',
            'b_source_ranges':[{'sheet':'合成表','range':'A2:C8'}]}
        with patch.object(UnifiedSemanticEngine,'request_text',return_value='合成表!B2“期初”适用于合成表!B3。') as request:
            run=run_step('recognize_b',config)
        self.assertEqual(run['status'],'complete',run)
        material=request.call_args.args[1]
        self.assertEqual([json.loads(line)['id'] for line in material['金标准.jsonl'].splitlines()],['period'])
        self.assertNotIn('A需要.json',material)
        self.assertIn('只补字段范围内数据格与原表期间标签的对应关系',request.call_args.args[0])
        saved=json.loads(Path(run['files']['b_mapping']).read_text(encoding='utf-8-sig'))
        self.assertEqual(saved['source_context'],b['source_context'])

    def test_材料保留原有子表边界(self):
        from 附注更新.A语义阅读 import source_material
        snapshot=deepcopy(self.snapshot)
        snapshot['names']={'表一':[{'sheet':'合成表','range':'A2:C4'}],
                           '表二':[{'sheet':'合成表','range':'A6:C8'}],
                           '动态名称':[{'sheet':'合成表','range':''}]}
        material,_=source_material(snapshot,[],[{'sheet':'合成表','range':'A2:C4'}])
        self.assertEqual(material['原有表区'],{'表一':snapshot['names']['表一']})

    def test_期间标签依次对应不广播到所有数据(self):
        answer='合成表!B3、C3依次对应B2“期初余额”、C2“期末余额”。'
        result=direct_answer(answer,self.engine.standard,['合成表'],field_scope={'合成表!B3:C3':['period']})
        periods={tuple(r['ranges']):r['label_references'] for r in result['dimensions'] if r['id']=='period'}
        self.assertEqual(periods,{('合成表!B3',):['合成表!B2'],('合成表!C3',):['合成表!C2']})
        result=direct_answer('合成表!B3 对应同一行的标签 合成表!A3「期初余额」，适用于合成表!B3。',
            self.engine.standard,['合成表'],field_scope={'合成表!B3':['period']})
        self.assertEqual(next(row for row in result['dimensions'] if row['id']=='period')['label_references'],['合成表!A3'])
        result=direct_answer('原表 B2“期初”适用于合成表!B3；C2“期末”适用于C3。',
            self.engine.standard,['合成表','另一表'],field_scope={'合成表!B3:C3':['period'],'另一表!B3':['period']})
        periods={tuple(r['ranges']):r['label_references'] for r in result['dimensions'] if r['id']=='period'}
        self.assertEqual(periods,{('合成表!B3',):['合成表!B2'],('合成表!C3',):['合成表!C2']})

    def test_明确来源不因范围描述不同自行要求改名(self):
        from 附注更新.A合集 import source_relation,direct_digest
        a=self.save('A');b=self.save('B')
        b['已确认关系']=[source_relation(a,b,'A 合成表!B3：采用 B 合成表!B4。',['合成表!B3'])]
        b['封存校验']=direct_digest(b)
        plan=direct_update_plan(a,b)
        self.assertTrue(any(row['cell']=='B3' for row in plan['updates']))
        self.assertFalse(any(row.get('cell')=='B3' and row.get('dimension')=='asset_selection' for row in plan['issues']))
        relation=source_relation(a,b,'A B3对应B 合成表!C3；B4对应C4。',['合成表!B3','合成表!B4'])
        self.assertEqual(relation['bindings'],[{'target':'合成表!B3','source':'合成表!C3'},
                                             {'target':'合成表!B4','source':'合成表!C4'}])

    def test_名单逐行搬运复用已有名称列及明确样式行(self):
        from 附注更新.A合集 import _declared_list_tables
        a=self.save('A')
        a['机械记录']['names']={'名单表':[{'sheet':'合成表','range':'B2:D6'}]}
        a['已确认关系']=[{'kind':'名单编排','dimension':'asset_selection',
            'content':'A名称在B列，第6行为“合计”。B各明细名称在B列。'}]
        text=('A表3行成员甲 ← B表20行；A表C、D分别取B表C、D。\n'
              'A表4行成员乙 ← B表22行；A表C、D分别取B表C、D。\n'
              'A表5行成员丙 ← B表24行；A表C、D分别取B表C、D。\n'
              '扩展新增A表6行成员丁 ← B表21行；A表C、D分别取B表C、D。\n'
              '扩展新增A表7行成员戊 ← B表23行；A表C、D分别取B表C、D。\n'
              '合计移至A表8行，对应B表28行；沿用A表4行的明细样式。')
        tables=_declared_list_tables(a,text,'asset_selection',[{'sheet':'来源表','range':'B20:D28'}])
        self.assertEqual(tables[0]['source_rows'],[20,22,24,21,23])
        self.assertEqual(tables[0]['columns'],{'C':'C','D':'D'})
        self.assertEqual(tables[0]['totals'],[{'target_row':6,'source_row':28}])
        self.assertEqual(tables[0]['operations'],[{'action':'insert','axis':'row','index':4,'count':2,'template_index':2}])

    def test_明确新名单搬运不被旧成员来源覆盖(self):
        from 附注更新.A合集 import source_relation,direct_digest
        a=self.save('A');b=self.save('B')
        a['机械记录']['names']={'名单表':[{'sheet':'合成表','range':'A2:C8'}]}
        a['封存校验']=direct_digest(a)
        b.setdefault('已确认关系',[]).append(source_relation(a,b,
            'A 合成表!B3：采用 B 合成表!C3。',['合成表!B3']))
        b['封存校验']=direct_digest(b)
        purpose={'reason':'采用已明确新名单','tables':[{'range_name':'名单表','source_sheet':'合成表',
            'source_rows':[4],'old_detail_count':1,'first_row_offset':1,
            'columns':{'B':'B','C':'C'},'labels':{'A':'A'}}]}
        plan=direct_update_plan(a,b,purpose)
        self.assertEqual(next(row for row in plan['updates'] if row['cell']=='B3')['sources'][0]['fact_id'],'合成表!B4')
        purpose['bindings']=[{'target':'合成表!B3','source':'合成表!C3'}]
        specified=direct_update_plan(a,b,purpose)
        self.assertEqual(next(row for row in specified['updates'] if row['cell']=='B3')['sources'][0]['fact_id'],'合成表!C3')

    def test_新明确来源复用已有公式副本取值(self):
        from 附注更新.A合集 import source_relation,direct_digest
        from 附注更新.统一勾稽 import _calculation_formulas
        a=self.save('A');snapshot=deepcopy(self.snapshot)
        snapshot['sheets'][0]['cells']['C3'].update(value='=1+2',formula='=1+2')
        b=self.save('B',snapshot=snapshot)
        b['facts']=[row for row in b['facts'] if row['source_reference']['location']!='合成表!C3']
        b['slot_bindings']=[row for row in b['slot_bindings'] if row['sheet']+'!'+row['cell']!='合成表!C3']
        b.setdefault('已确认关系',[]).append(source_relation(a,b,
            'A 合成表!B3：采用 B 合成表!C3。',['合成表!B3']))
        receipt=self.case.folder/'已有公式副本依据.json';receipt.write_text('{}',encoding='utf-8')
        b['calculation_evidence']={'path':str(receipt)};b['封存校验']=direct_digest(b)
        self.assertEqual(_calculation_formulas(b)['合成表!C3'],'=1+2')
        table_only=deepcopy(b);table_only['已确认关系']=[{'kind':'名单编排','tables':[{
            'source_sheet':'合成表','source_rows':[4],'columns':{'C':'C'},'labels':{},
            'totals':[{'source_row':3}]}]}]
        self.assertEqual(_calculation_formulas(table_only)['合成表!C3'],'=1+2')
        with patch('附注更新.统一勾稽._calculated_values',return_value={
            '合成表!C3':{'formula':'=1+2','state':'value','value':3}}) as read:
            plan=direct_update_plan(a,b)
        read.assert_called_once()
        self.assertEqual(next(row for row in plan['updates'] if row['cell']=='B3')['value'],3)
        self.assertNotIn('formula_values',b)

    def test_来源矩形按答复列组展开再裁切请求(self):
        from 附注更新.A合集 import save_source_relations
        a=self.save('A');b=self.save('B')
        reply=('| A位置 | B来源 |\n|---|---|\n'
               '| 合成表!B3:B4、C3:C4 | 合成表!B5:C6 |\n'
               'A B8 → B 合成表!B7；C8 → C7。')
        saved=save_source_relations(a,b,reply,['合成表!B4','合成表!C3','合成表!B8','合成表!C8'],self.case.folder/'列组来源.json')
        self.assertEqual(saved['已确认关系'][-1]['bindings'],[
            {'target':'合成表!'+target,'source':'合成表!'+source}
            for target,source in [('B4','B6'),('C3','C5'),('B8','B7'),('C8','C7')]])
        reply=('以下 B 位置均在 合成表。\n| A位置 | 对应B位置 |\n|---|---|\n'
               '| B3 | !C3 |\n| B4 | !C4 |\n'
               'A B8 对应 B 合成表!C8（对应明细来源：合成表!C7）。')
        saved=save_source_relations(a,b,reply,['合成表!B3','合成表!B4','合成表!B8'],self.case.folder/'明确表名及说明引用.json')
        self.assertEqual(saved['已确认关系'][-1]['bindings'],[
            {'target':'合成表!B'+str(row),'source':'合成表!C'+str(row)} for row in [3,4,8]])
        from 附注更新.A合集 import source_relation
        relation=source_relation(a,b,'A B3:B4按项目对应：B3→合成表!C3；B4→C4。',
                                 ['合成表!B3','合成表!B4'])
        self.assertEqual(relation['bindings'],[
            {'target':'合成表!B3','source':'合成表!C3'},
            {'target':'合成表!B4','source':'合成表!C4'}])

    def test_已明确来源的Excel错误原值直接进入搬运(self):
        a=self.save('A');snapshot=deepcopy(self.snapshot)
        snapshot['sheets'][0]['cells']['B3'].update(value='=0/0',formula='=0/0')
        b=self.save('B',snapshot=snapshot)
        b['calculation_evidence']={'path':'本例已有计算记录'}
        b['formula_values']={'合成表!B3':{'formula':'=0/0','state':'error','value':'#DIV/0!','reason':'Excel计算错误：#DIV/0!'}}
        b['封存校验']=direct_digest(b)
        plan=direct_update_plan(a,b,{'bindings':[{'target':'合成表!B3','source':'合成表!B3'}]})
        item=next(row for row in plan['updates'] if row['cell']=='B3')
        self.assertEqual((item['value'],item['excel_data_type']),('#DIV/0!','e'))
        self.assertFalse(any(row.get('cell')=='B3' and row.get('kind')=='unevaluated_source_formula' for row in plan['issues']))
        self.assertEqual(b['机械记录']['sheets'][0]['cells']['B3']['formula'],'=0/0')

    def test_来源表名括号保留而说明中的引用不混入(self):
        from 附注更新.A合集 import source_relation
        a=self.save('A');b=self.save('B')
        b['机械记录']['sheets'][0]['name']='汇总（单体）'
        reply=('| A 位置及事项 | 对应的 B 位置 | 说明 |\n|---|---|---|\n'
               '| 合成表!B3、C3 | 汇总（单体）!B4、C4（说明：汇总（单体）!B8） | A位置对应B位置 |\n'
               '| 合成表!B4 | 汇总（单体）!B5 | 同上 |')
        relation=source_relation(a,b,reply,['合成表!B3','合成表!C3','合成表!B4'])
        self.assertEqual(relation['bindings'],[
            {'target':'合成表!B3','source':'汇总（单体）!B4'},
            {'target':'合成表!C3','source':'汇总（单体）!C4'},
            {'target':'合成表!B4','source':'汇总（单体）!B5'}])
        many=source_relation(a,b,'| A位置 | B位置 |\n|---|---|\n'
            '| 合成表!B3 | 汇总（单体）!B4；汇总（单体）!C4 |',['合成表!B3'])
        self.assertEqual([row['source'] for row in many['bindings']],['汇总（单体）!B4','汇总（单体）!C4'])
        backward=source_relation(a,b,'A位置 合成表!B3 ← 本次采用的B位置 汇总（单体）!C4。',['合成表!B3'])
        self.assertEqual(backward['bindings'],[{'target':'合成表!B3','source':'汇总（单体）!C4'}])

    def test_指定期间字段的标签原文不受词表限制(self):
        role=direct_answer('| B坐标 | metric_id | template_period_role | 其他实际维度取值 |\n'
            '|---|---|---|---|\n| 合成表!B3 | 固定资产原值 | 期初余额 | period=2025-01-01 |',
            self.engine.standard,['合成表'])
        self.assertEqual([row['value'] for row in role['dimensions'] if row['id']=='template_period_role'],['期初余额'])
        roles=direct_answer('甲!D3：固定资产原值。\n甲!C3：固定资产原值。\n乙!C4：固定资产原值。\n'
            '以上D列及乙C列均为上期，template_period_role=prior。未填写金额。',self.engine.standard,['甲','乙'])
        declared=roles['dimensions'][-1]
        self.assertEqual(declared['value'],'prior')
        self.assertEqual(set(declared['ranges']),{'甲!D3','乙!C4'})
        common=direct_answer('甲!B3：固定资产原值。\n乙!C4：固定资产原值。\n'
            '以上坐标共同取值：currency=CNY；unit_scale=1。',self.engine.standard,['甲','乙'])
        self.assertEqual(set(next(row for row in common['dimensions'] if row['id']=='currency')['ranges']),
                         {'甲!B3','乙!C4'})
        parsed=direct_answer('B2“年初数”适用于 B3。\nB12“年末数”适用于 B13。',self.engine.standard,
            ['甲','乙'],field_scope={'甲!B3':['period'],'乙!B13':['period']})
        self.assertEqual([(row['ranges'],row['evidence']) for row in parsed['dimensions'] if row['id']=='period'],
                         [(['甲!B3'],['甲!B2']),(['乙!B13'],['乙!B12'])])
        a=self.save('A');before=dict(_records(a))
        saved=save_direct_answer(a,'B2“年初数”适用于 B3:B4；C2“年末数”适用于 C3:C4。',
            self.case.folder/'年初年末标签.json',field_scope={'合成表!B3:C4':['period']})
        after=dict(_records(saved))
        for col in 'BC':
            for row in [3,4]:
                loc=f'合成表!{col}{row}'
                self.assertEqual(after[loc]['label_references']['period'],[f'合成表!{col}2'])
                self.assertEqual(after[loc]['dimensions'],before[loc]['dimensions'])

    def test_共用期间只用于答复明说的坐标且成员引号是表示(self):
        metric=next(iter(self.engine.standard['metric']))
        reply=('“本年”列为 template_period_role=current、period={kind:duration,start:2025-01-01,end:2025-12-31}。\n'
               f'合成表!B3：{metric}；asset_selection：domain=资产，mode=members，members=[“设备”]。\n'
               '这些坐标均在“本年”列。')
        parsed=direct_answer(reply,self.engine.standard,['合成表'],field_scope={'合成表!B3:C3':['period','asset_selection','template_period_role']})
        row=next(x for x in parsed['dimensions'] if x['id']=='period' and x.get('value',{}).get('kind')=='duration')
        self.assertEqual(row['ranges'],['合成表!B3'])
        selection=next(x for x in parsed['dimensions'] if x['id']=='asset_selection')
        self.assertEqual(selection['value']['members'],['设备'])

    def test_来源答复的行列分组和章节表名只机械展开(self):
        from 附注更新.A合集 import save_source_relations
        a=self.save('A');b=self.save('B')
        reply=('### 明细 A → B 合成表\nA B3,C3 → B B5,C5。\n'
               '**变动表**\n行对应：A 4 → B 合成表!6；A 行 7:8 → B 8:9。\n'
               '列按阶段对应：A B第一列→B B，A C第二列→B C。')
        targets=['合成表!'+c+str(r) for r in [3,4,7,8] for c in 'BC']
        saved=save_source_relations(a,b,reply,targets,self.case.folder/'分组行列来源.json')
        self.assertEqual(saved['已确认关系'][-1]['bindings'],[
            {'target':f'合成表!{c}{r}','source':f'合成表!{c}{s}'}
            for r,s in [(3,5),(4,6),(7,8),(8,9)] for c in 'BC'])

    def test_模型明确的新增行来源不再次移位(self):
        from 附注更新.A合集 import save_source_relations
        a,b,table=self.two_lists()
        saved=save_source_relations(a,b,'A 合成表!B9：采用 B 合成表!C9。',['合成表!B9'],
            self.case.folder/'编排后来源.json',arranged_targets=True)
        plan=direct_update_plan(a,saved,{'tables':[table]})
        update=next(row for row in plan['updates'] if row['cell']=='B9')
        self.assertEqual([row['fact_id'] for row in update['sources']],['合成表!C9'])

    def test_后补前方名单增行仍沿用已安排的同一目标(self):
        from 附注更新.A合集 import source_relation,save_list_relations,direct_digest
        a,b,table=self.two_lists()
        relation={'kind':'名单编排','dimension':'asset_selection','answer_hash':'先前名单',
            'source_a_hash':a['source_hash'],'source_b_hash':b['source_hash'],'gold_hash':a['gold_hash'],
            'tables':[table],'conflicts':[]}
        b['已确认关系']=[relation,source_relation(a,b,'合成表!B9采用合成表!C9。',['合成表!B9'],arranged_targets=True)]
        b['封存校验']=direct_digest(b)
        earlier={**table,'range_name':'第一表','source_rows':[3,4],'old_detail_count':1}
        with patch('附注更新.A合集._declared_list_tables',return_value=[earlier]):
            saved=save_list_relations(a,b,'新增前方名单','asset_selection',self.case.folder/'新名单')
        mapping=json.loads(Path(saved['files']['b_mapping']).read_text(encoding='utf-8'))
        plan=direct_update_plan(a,mapping)
        update=next(row for row in plan['updates'] if row['cell']=='B10')
        self.assertEqual([r['fact_id'] for r in update['sources']],['合成表!C9'])
        binding=mapping['已确认关系'][1]['bindings'][0]
        self.assertNotIn('target_anchor',binding)
        self.assertEqual(binding['target_member']['source_row'],table['source_rows'][-1])
        reordered={**table,'source_rows':list(reversed(table['source_rows']))}
        plan=direct_update_plan(a,mapping,{'tables':[earlier,reordered]})
        update=next(row for row in plan['updates'] if row['cell']=='B7')
        self.assertEqual([r['fact_id'] for r in update['sources']],['合成表!C9'])

    def test_名单自然箭头表直接接入名称金额及合计(self):
        from 附注更新.A合集 import _declared_list_tables
        a=self.save('A');a['机械记录']['names']={'明细':[{'sheet':'合成表','range':'A2:C5'}]}
        reply=('| 名单 | 名称：B来源 → A位置 | 本期：B → A | 上期：B → A |\n'
            '| 明细第1行 | 来源表!A7「甲」→ 合成表!A3 | B7 → B3 | C7 → C3 |\n'
            '| 明细第2行 | 来源表!A8「乙」→ 合成表!A4 | B8 → B4 | C8 → C4 |\n'
            '| 合计 | 来源表!A9「合计」→ 合成表!A5 | B9 → B5 | C9 → C5 |')
        tables=_declared_list_tables(a,reply,'asset_selection',[{'sheet':'来源表','range':'A6:C9'}])
        self.assertEqual(tables[0]['columns'],{'B':'B','C':'C'})
        self.assertEqual(tables[0]['labels'],{'A':'A'})
        self.assertEqual(tables[0]['source_rows'],[7,8])
        self.assertEqual(tables[0]['totals'],[{'target_row':5,'source_row':9}])

    def test_名单整行声明和独立名称按原答复保存(self):
        from 附注更新.A合集 import _declared_list_tables
        a=self.save('A');a['机械记录']['names']={'明细':[{'sheet':'合成表','range':'A2:C5'}]}
        sources=[{'sheet':'来源表','range':'A6:C9'}]
        reply='A3:A4 分别对应 B7:B8：甲、乙。\nA5 对应 B9 合计，沿用 A5 合计样式。'
        table=_declared_list_tables(a,reply,'asset_selection',sources)[0]
        self.assertEqual(table['columns'],{'B':'B','C':'C'})
        self.assertEqual(table['labels'],{'A':'A'})
        self.assertEqual(table['source_rows'],[7,8])
        reply=('A 原有明细行 3–4 保留；金额分别搬到 A 的 B、C 列。\n'
               '| A位置 | 名称 | B名称来源 | 金额来源 |\n'
               '| 第3行 | 甲 | 来源表!A7 | 来源表!B7、C7 |\n'
               '| 第4行 | 乙 | 原答复已明确的名称 | 来源表!B8、C8 |')
        table=_declared_list_tables(a,reply,'asset_selection',sources)[0]
        self.assertEqual(table['declared_labels'],{'7':{'source':'来源表!A7'},'8':{'value':'乙'}})
        b=self.save('B');table.update(source_sheet='合成表',source_rows=[3,4],declared_labels={'4':{'value':'模型已列明的新名称'}})
        a['封存校验']=direct_digest(a)
        plan=direct_update_plan(a,b,{'tables':[table]})
        self.assertEqual(next(row['value'] for row in plan['labels'] if row['cell']=='A4'),'模型已列明的新名称')

    def test_模型明确标签保留与搬运不重做金额判断(self):
        from 附注更新.A合集 import save_source_relations
        a,b,table=self.two_lists(year=2024)
        reply=('A!B3 所属期间标题：A!B2 标题对应采用 B 合成表!B2 的标题。B3不另造标题。\n'
               'A!C2：保留原文。')
        saved=save_source_relations(a,b,reply,['合成表!B3','合成表!C3'],self.case.folder/'标签动作.json')
        self.assertEqual(saved['已确认关系'][-1]['label_actions'],[
            {'action':'copy','target':'合成表!B2','source':'合成表!B2'},
            {'action':'retain','target':'合成表!C2'}])
        plan=direct_update_plan(a,saved)
        self.assertTrue(any(row['cell']=='B2' and row['value']=='新期末' for row in plan['labels']))
        self.assertFalse(any(row['cell']=='C2' for row in plan['labels']))

    def test_执行记录引用原合集并保留空白不生成身份(self):
        from 附注更新.A合集 import direct_output_record
        a=self.save('A');b=self.save('B')
        b=save_direct_answer(b,'合成表!B5：asset_selection=unknown。',self.case.folder/'已存未知.json',
                             field_scope={'合成表!B5':['asset_selection']})
        purpose={'reason':'执行已明确来源','bindings':[{'target':'合成表!B5','source':'合成表!B5'}]}
        plan=direct_update_plan(a,b,purpose);before=deepcopy((a,b))
        with patch('附注更新.表格.read_workbook',return_value=deepcopy(self.snapshot)),\
             patch('附注更新.表格.file_hash',return_value='c'*64),\
             patch('附注更新.A合集.identity_completeness',side_effect=AssertionError('派生时不能重判身份')):
            derived=direct_output_record(a,b,self.case.folder/'模拟副本.xlsx',plan,'已保存安排')
        self.assertEqual((a,b),before)
        self.assertEqual(derived['schema'],'附注更新执行记录-v1')
        self.assertEqual(derived['collections']['A']['mapping_path'],a['mapping_path'])
        self.assertEqual(derived['collections']['B']['mapping_path'],b['mapping_path'])
        retained=next(row for row in derived['retained_blanks'] if row['cell']=='B5' and not row.get('action'))
        self.assertIsNone(retained['actual_value'])
        self.assertNotIn('完整性',derived)
        self.assertNotIn('facts',derived)

    def test_普通证据不被推定为标签位置(self):
        both=[]
        for carrier,year in [('A',2025),('B',2026)]:
            answer=json.loads(self.case.answer(year))
            for row in answer['tables'][0]['dimensions']:row.pop('label_references',None)
            both.append(self.save(carrier,json.dumps(answer,ensure_ascii=False)))
        plan=direct_update_plan(*both)
        self.assertTrue(plan['identity_matches'])
        self.assertFalse(plan['labels'])
        self.assertEqual(len(plan['updates'])+len(plan['retained_blanks']),len(plan['identity_matches']))

    def test_说明声明无需附加理由且保存合并不裁定整体语义完成(self):
        from 附注更新.统一勾稽 import load_fact_mapping
        answer=json.loads(self.case.answer());table=answer['tables'][0]
        table['data_ranges']=['B3:C7','B8']
        table['non_business']=[{'range':'C8','category':'annotation'}]
        saved=self.save('A',json.dumps(answer,ensure_ascii=False))
        loaded=load_fact_mapping(saved['mapping_path'],{})
        combined=combine_direct_collections([loaded])
        for result in (saved,loaded,combined):
            self.assertNotIn('complete',result)
            self.assertNotIn('complete',result['outline'])
            self.assertNotIn('complete',result['coverage'])
            self.assertEqual([row['cell'] for row in result['excluded']],['C8'])
            self.assertNotIn('reason',result['excluded'][0])
            self.assertNotIn('evidence',result['excluded'][0])
            self.assertEqual(dict(_records(result))['合成表!B8']['raw_value_state'],'blank')
        self.assertEqual(dict(_records(combined)),dict(_records(saved)))

    def test_明确标签引用不被财务记录身份复判且空白按既定要求保留(self):
        for value in (73,0,None,''):
            with self.subTest(value=value):
                both=[]
                for carrier,year in [('A',2025),('B',2026)]:
                    answer=json.loads(self.case.answer(year))
                    for row in answer['tables'][0]['dimensions']:
                        if row['id']=='period' and row['ranges']==['B3:B8']:
                            row['label_references']=['合成表!B8']
                    snapshot=deepcopy(self.snapshot)
                    snapshot['sheets'][0]['cells']['B8']['value']='旧标签' if carrier=='A' else value
                    both.append(self.save(carrier,json.dumps(answer,ensure_ascii=False),snapshot))
                before=deepcopy(both)
                plan=direct_update_plan(*both,scope={'target':[{'sheet':'合成表','range':'B3'}],
                    'source':[{'sheet':'合成表','range':'B3:B8'}]})
                matched=next(row for row in plan['identity_matches'] if row['cell']=='B3')
                self.assertEqual(matched['label_references']['period'],['合成表!B8'])
                labels=[row for row in plan['labels'] if row['cell']=='B8']
                if value in (None,''):
                    self.assertFalse(labels)
                    self.assertTrue(any(row['cell']=='B8' and row.get('action')=='label' for row in plan['retained_blanks']))
                else:
                    self.assertEqual([row['value'] for row in labels],[value])
                self.assertFalse(any(row.get('action')=='label' for row in plan['issues']))
                self.assertEqual(both,before)

    def test_明确搬运不要求额外填写理由放行(self):
        a=self.save('A');b=self.save('B')
        plan=direct_update_plan(a,b,{'bindings':[{'target':'合成表!B3','source':'合成表!C3'}]})
        row=next(row for row in plan['identity_matches'] if row['cell']=='B3')
        self.assertEqual(row['source_locations'],['合成表!C3'])

    def test_自然答复明确标签接入局部保存及回读(self):
        from 附注更新.统一勾稽 import load_fact_mapping
        answer=json.loads(self.case.answer())
        for row in answer['tables'][0]['dimensions']:row.pop('label_references',None)
        original=self.save('A',json.dumps(answer,ensure_ascii=False));before=dict(_records(original))
        text='合成表!B2（标签位于合成表!B2）为“期末数”，适用于合成表!B3:B8，角色为期末。'
        saved=save_direct_answer(original,text,self.case.folder/'明确标签.json',field_scope={'合成表!B3:B8':['period','template_period_role']})
        loaded=load_fact_mapping(saved['mapping_path'],{});after=dict(_records(loaded))
        self.assertEqual(set(after),set(before))
        for location,record in before.items():
            self.assertEqual(after[location]['dimensions'],record['dimensions'])
            self.assertEqual(after[location]['raw_value_state'],record['raw_value_state'])
        self.assertEqual(after['合成表!B8']['label_references']['period'],['合成表!B2'])
        with patch('附注更新.统一语义._dimension_value',side_effect=AssertionError('保存不能复判业务值')):
            again=save_direct_answer(loaded,'合成表!B3:B8：period=instant,2025-12-31。',
                self.case.folder/'同值保留标签.json',field_scope={'合成表!B3:B8':['period']})
        self.assertEqual(dict(_records(again))['合成表!B8']['label_references']['period'],['合成表!B2'])

    def test_B局部纠正合并同表子表时共同声明不越界(self):
        from 附注更新.A语义阅读 import addresses
        parts=[]
        for start,end,currency in [(3,4,'CNY'),(5,8,'USD')]:
            answer=json.loads(self.case.answer());table=answer['tables'][0]
            table.update(id=f'子表{start}',data_ranges=[f'B{start}:C{end}'])
            owned=set(addresses(table['data_ranges'][0]))
            for kind in ('metrics','dimensions'):
                kept=[]
                for row in table[kind]:
                    if row['ranges']!=['*']:
                        row['ranges']=sorted({cell for area in row['ranges'] for cell in addresses(area)}&owned)
                    if row['id']=='currency':row['value']=currency
                    if row['ranges']:kept.append(row)
                table[kind]=kept
            parts.append(self.save('B',json.dumps(answer,ensure_ascii=False)))
        original=combine_direct_collections(parts);before=dict(_records(original))
        saved=save_direct_answer(original,'合成表!B3：asset_selection=members：设备甲。',
            self.case.folder/'同表局部纠正.json',field_scope={'合成表!B3':['asset_selection']})
        after=dict(_records(saved))
        self.assertEqual(set(after),set(before))
        for loc,row in before.items():
            if loc!='合成表!B3':self.assertEqual(after[loc]['dimensions'],row['dimensions'])
        self.assertEqual(after['合成表!B3']['dimensions']['currency'],'CNY')
        self.assertEqual(after['合成表!B5']['dimensions']['currency'],'USD')
        self.assertEqual(saved['机械记录'],original['机械记录'])
        self.assertFalse(saved['unresolved'])

    def test_名单关系一次识别原身份原值均不改(self):
        from 附注更新.A语义阅读 import recognize_list_relations
        from 附注更新.统一勾稽 import load_fact_mapping
        a=self.save('A');b=self.save('B');reply='名单关系原文不足，保留原结果。'
        config={'list_dimension':'asset_selection','a_required_ranges':[{'sheet':'合成表','range':'A2:C4'}],
            'b_source_ranges':[{'sheet':'合成表','range':'A2:C4'}],
            'b_allowed_metrics':['固定资产原值'],'update_purpose':{'reason':'用户要求名单随B'},
            'b_additional_context':'本次仅补充同一名单续列关系。'}
        with patch.object(self.engine,'request_text',return_value=reply) as call:
            result=recognize_list_relations(self.engine,a,b,config,self.case.folder/'名单关系')
        self.assertEqual(call.call_count,1)
        self.assertIn(config['b_additional_context'],call.call_args.args[0])
        self.assertEqual(call.call_args.args[1]['已有名单搬运安排'],[])
        for carrier,original in [('A',a),('B',b)]:
            self.assertEqual(call.call_args.args[1][carrier+'原表.json']['载体'],carrier)
            stated=call.call_args.args[1][carrier+'已有识别.json']['已有身份']
            for fact in original['facts']:
                if fact['source_reference']['location'].rsplit('!',1)[-1] in ('B3','C3','B4','C4'):
                    self.assertTrue(any(row['dimensions']==fact['dimensions'] for row in stated))
        both={}
        for carrier,original in [('a',a),('b',b)]:
            saved=load_fact_mapping(result['files'][carrier+'_mapping'],{})
            both[carrier]=saved
            for key in ('facts','slot_bindings','outline','机械记录','识别答复','unresolved'):
                self.assertEqual(saved[key],original[key])
            self.assertEqual(saved['已确认关系'][-1]['content'],reply)
        before=direct_update_plan(a,b);after=direct_update_plan(both['a'],both['b'])
        for key in ('updates','labels','retained_blanks','identity_matches','operations','issues'):
            self.assertEqual(after[key],before[key])
        self.assertEqual(after['list_relation_gaps'],[])
        self.assertEqual(after['list_relations'][0]['content'],reply)
        with patch.object(self.engine,'request_text',return_value=reply) as call:
            recognize_list_relations(self.engine,both['a'],both['b'],config,self.case.folder/'名单原答复接续')
        self.assertEqual(call.call_args.args[1]['同范围已有名单原答复'],[reply])

    def test_名单补充材料携带已保存的正确来源(self):
        from 附注更新.A语义阅读 import recognize_list_relations
        a=self.save('A');b=self.save('B');table={'source_sheet':'合成表','source_rows':[3],'columns':{'B':'B'}}
        b['已确认关系'].append({'kind':'名单编排','dimension':'asset_selection','tables':[table]})
        config={'list_dimension':'asset_selection','a_required_ranges':[{'sheet':'合成表','range':'A2:C4'}],
            'b_source_ranges':[{'sheet':'合成表','range':'A2:C4'}],'b_allowed_metrics':['固定资产原值'],
            'update_purpose':{'reason':'名单随B'}}
        with patch.object(self.engine,'request_text',return_value='已有来源复用。') as call,patch('附注更新.A合集.save_list_relations',return_value={}):
            recognize_list_relations(self.engine,a,b,config,self.case.folder/'来源复用')
        self.assertEqual(call.call_args.args[1]['已有名单搬运安排'],[table])

    def test_名单答复明确的冲突原样保存不按公式裁定(self):
        from 附注更新.A合集 import _list_relation_conflicts
        mapping={'机械记录':{'sheets':[{'name':'合成表','cells':{
            'B3':{'formula':"='来源'!C8"},'B4':{'formula':"='来源'!C9"}}}]}}
        reply='### 名称与金额来源不一致\nB3、B4 均引用 来源!C8，而标签分别来自该表 A1、A2。'
        found,unused=_list_relation_conflicts(mapping,reply,[{'sheet':'合成表','range':'A1:C4'}])
        self.assertEqual([loc for row in found for loc in row['locations']],['合成表!B3','合成表!B4'])
        self.assertEqual(unused,[])
        ordinary='B3 引用 来源!C8，标签来自 A1。'
        self.assertEqual(_list_relation_conflicts(mapping,ordinary,[{'sheet':'合成表','range':'A1:C4'}]),([],[]))

    def test_文字成员列表及分列期间不丢失(self):
        reply='''**合成表**（B自身坐标）
共同取值：`metric_id=固定资产原值`；C列期间为`2025-01-01至2025-12-31`（表头“本年”）；D列期间为`2024-01-01至2024-12-31`（表头“上年”）。下列明细的`asset_selection`均为`domain=固定资产, mode=members`，成员取值按原文列出。
- `C3:D3`（设备甲）：members=`[设备甲]`
### 合成表：指定坐标的 `asset_selection`
- `合成表!B4:C4`：`mode: members`，并列成员为“设备乙”“设备丙”。
- `合成表!B5:C5`：同为“设备乙”“设备丙”。'''
        parsed=direct_answer(reply,self.engine.standard,['合成表','其他表'])
        members=[r for r in parsed['dimensions'] if r['id']=='asset_selection']
        self.assertEqual([r['value']['members'] for r in members],[['设备甲'],['设备乙','设备丙'],['设备乙','设备丙']])
        periods={r['ranges'][0]:r['value'] for r in parsed['dimensions'] if r['id']=='period'}
        self.assertEqual(periods['合成表!C3'],'2025-01-01至2025-12-31')
        self.assertEqual(periods['合成表!D3'],'2024-01-01至2024-12-31')
        self.assertFalse(parsed['unresolved'])
        total=direct_answer('合成表!B3:C3：asset_selection 均为`domain=固定资产, mode=all`。该合计代表总体。',
                            self.engine.standard,['合成表'])
        self.assertEqual(total['dimensions'][0]['value'],{'domain':'固定资产','mode':'all'})

    def test_成员表引用来源不改变目标所属表并保留未知(self):
        reply='''**合成表**
共同取值：计量指标为`固定资产原值`；B列取值期间为2025-01-01至2025-12-31，模板期间角色为“本年”；C列为2024-01-01至2024-12-31，角色为“上年”。
`asset_selection`的domain为“固定资产”。各坐标均为`mode=members`，成员分别为：
| 坐标行 | 成员或范围 |
|---|---|
| B3 | 设备甲 |
| B4 | 引用来源为其他表!A9，未识别出固定资产成员 |
| B5 | 设备乙；其公式引用其他表!A10 |
C3:C4按名称取范围：依次为“设备甲、设备乙”。
C4不是固定资产成员。
B6:C6为`mode=members`，成员“设备丙”。
B8、C8的`asset_selection`均为domain“固定资产”、`mode=all`。'''
        parsed=direct_answer(reply,self.engine.standard,['合成表','其他表'])
        selections={area:row['value'] for row in parsed['dimensions'] if row['id']=='asset_selection' for area in row['ranges']}
        self.assertEqual(selections['合成表!B5']['members'],['设备乙'])
        self.assertEqual(selections['合成表!C3']['members'],['设备甲'])
        self.assertEqual(selections['合成表!C4'],'未知')
        self.assertEqual(selections['合成表!B4'],'未知')
        self.assertEqual(selections['合成表!B8'],{'domain':'固定资产','mode':'all'})
        self.assertFalse(any(area.startswith('其他表!') for area in selections))
        saved=save_direct_answer(self.save('B'),reply,self.case.folder/'成员原答复重读.json',
            field_scope={'合成表!B3:C8':['asset_selection']})
        from 附注更新.统一勾稽 import load_fact_mapping
        self.assertEqual(dict(_records(load_fact_mapping(self.case.folder/'成员原答复重读.json',{}))),dict(_records(saved)))
        self.assertEqual(dict(_records(saved))['合成表!B6']['dimensions']['asset_selection']['members'],['设备丙'])

    def test_名单明示说明行例外不同时保存成成员(self):
        reply=('asset_selection的domain为“固定资产”。各坐标均为mode=members。\n'
               '合成表!C3:C4按名称取范围：依次为“设备甲、说明文字”；除说明行外均为domain“固定资产”、mode=members。'
               '合成表!C4是该说明行对应的空白金额格，不是固定资产成员。')
        parsed=direct_answer(reply,self.engine.standard,['合成表'])
        values=[r['value'] for r in parsed['dimensions'] if r['id']=='asset_selection' and '合成表!C4' in r['ranges']]
        self.assertEqual(values,['未知'])
        b=self.save('B')
        saved=save_direct_answer(b,reply,self.case.folder/'说明行排除.json',
            field_scope={'合成表!C3:C4':['asset_selection']})
        rows=dict(_records(saved))
        self.assertEqual(rows['合成表!C3']['dimensions']['asset_selection']['members'],['设备甲'])
        self.assertEqual(rows['合成表!C4']['dimensions']['asset_selection']['mode'],'unknown')
        self.assertEqual(saved['机械记录'],b['机械记录'])
        ordinary=direct_answer(reply.replace('；除说明行外均为domain“固定资产”、mode=members',''),self.engine.standard,['合成表'])
        values=[r['value'] for r in ordinary['dimensions'] if r['id']=='asset_selection' and '合成表!C4' in r['ranges']]
        self.assertEqual(len(values),2)

    def test_缩进编号同一槽位及明示期间边界(self):
        quoted=direct_answer('合成表!B3:C3：period=`2025-01-01` 至 `2025-12-31`。',self.engine.standard,['合成表'])
        self.assertEqual(quoted['dimensions'][0]['value'],'2025-01-01至2025-12-31')
        standard=deepcopy(self.engine.standard)
        standard['metric']['变动编号']=deepcopy(standard['metric']['固定资产原值'])
        reply='''共同取值：列B为本年，列C为上年；本年期间为2025-01-01至2025-12-31，上年期间为2024-01-01至2024-12-31。余额行的年初数按相应期间起日取时点，年末数按相应期间末日取时点。
以下“类别成员”均为asset_selection，mode=members、domain=固定资产；“合计”行为全体，mode=all、domain=固定资产。
- 项目：`固定资产原值`
  - 年初数：B3:C3为各类别成员；B4:C4为全体合计。
  - 增加：`变动编号`，B5:C5为各类别成员。
  - 减少：同一变动槽位，B6:C6为各类别成员。
  - 年末数：`固定资产原值`，B7:C7为各类别成员。'''
        parsed=direct_answer(reply,standard,['合成表'])
        metrics={area:row['id'] for row in parsed['metrics'] for area in row['ranges']}
        self.assertEqual(metrics['合成表!B3:C3'],'固定资产原值')
        self.assertEqual(metrics['合成表!B6:C6'],'变动编号')
        periods={row['ranges'][0]:row['value'] for row in parsed['dimensions'] if row['id']=='period'}
        self.assertEqual(periods['合成表!B3'],'2025-01-01')
        self.assertEqual(periods['合成表!C7'],'2024-12-31')
        selections={r['ranges'][0]:r['value'] for r in parsed['dimensions'] if r['id']=='asset_selection'}
        self.assertEqual(selections['合成表!B4:C4'],{'mode':'all','domain':'固定资产'})
        self.assertEqual(selections['合成表!B3:C3'],{'mode':'unknown','domain':'固定资产','known_scope':{'mode':'members'}})
        base=self.save('B')
        saved=save_direct_answer(base,reply,self.case.folder/'分段成员未明确.json',field_scope={'合成表!B3:C8':['asset_selection']})
        self.assertEqual(dict(_records(saved))['合成表!B3']['dimensions']['asset_selection'],selections['合成表!B3:C3'])
        self.assertFalse(parsed['unresolved'])

    def test_名单答复整列期间说明直接接为引用不改日期(self):
        reply='''- A 原表：`合成表!A2:C8`。`A2`为项目列标题，`B2`、`C2`分别为本期、上期。金额列为`B3:B8`、`C3:C8`。
- B 原表：`来源表!A7:C12`；`A8:C8`为列标题，`A9:A10`是明细及占位行，`A11:C11`为合计，`A12:C12`为检验行。名称在A列，B、C列分别为本期、上期。名称由公式引用`汇总表!B100:B102`。
来源表!B9：其余内容原样保留。'''
        scope={'合成表!B3:C8':['period'],'来源表!B9:C11':['period']}
        parsed=direct_answer(reply,self.engine.standard,['合成表','来源表','汇总表'],field_scope=scope)
        refs={cell:row['evidence'] for row in parsed['dimensions'] if row['id']=='period' for cell in row['ranges']}
        self.assertEqual(refs['合成表!B3'],['合成表!B2'])
        self.assertEqual(refs['来源表!C11'],['来源表!C8'])
        self.assertNotIn('来源表!C12',refs)
        self.assertFalse(any(cell.startswith('汇总表!') for cell in refs))
        original=self.save('A')
        saved=save_direct_answer(original,reply,self.case.folder/'名单列关系.json',field_scope={'合成表!B3:C8':['period']})
        for loc,row in _records(saved):
            self.assertEqual(row['dimensions'],dict(_records(original))[loc]['dimensions'])
            self.assertEqual(row['dimension_evidence']['period'],['合成表!'+loc.split('!')[1][0]+'2'])

    def test_明确纠正范围后不再沿用旧名单冲突(self):
        a=self.save('A');b=self.save('B')
        relation={'kind':'名单编排','dimension':'asset_selection','answer_hash':'原名单答复',
            'source_a_hash':a['source_hash'],'source_b_hash':b['source_hash'],'gold_hash':a['gold_hash'],
            'conflicts':[{'locations':['合成表!B3'],'content':'原范围与金额来源冲突'}]}
        for mapping in (a,b):
            mapping['已确认关系']=[deepcopy(relation)];mapping['封存校验']=direct_digest(mapping)
        current=direct_update_plan(a,b)
        self.assertTrue(any(r.get('answer_hash')=='原名单答复' for r in current['issues']))
        self.assertTrue(any(r['cell']=='B3' for r in current['identity_matches']))
        changed=self.corrected(a,b,'合成表!B3：asset_selection 为 members：设备甲。')
        self.assertEqual(changed['已确认关系'][0],relation)
        self.assertFalse(any(r.get('answer_hash')=='原名单答复' for r in direct_update_plan(a,changed)['issues']))
        unknown=self.corrected(a,b,'合成表!B3：asset_selection 为 unknown。')
        self.assertTrue(any(r.get('answer_hash')=='原名单答复' for r in direct_update_plan(a,unknown)['issues']))

    def test_局部重读保留已保存来源顺序(self):
        b=self.save('B');b['facts'].reverse();b['slot_bindings'].reverse();b['封存校验']=direct_digest(b)
        saved=save_direct_answer(b,'合成表!B3：template_period_role=closing。',
            self.case.folder/'保留来源顺序.json',field_scope={'合成表!B3':['template_period_role']})
        self.assertEqual([loc for loc,_ in _records(b)],[loc for loc,_ in _records(saved)])
        a=self.save('A');purpose={'reason':'沿已保存顺序机械匹配'}
        before=direct_update_plan(a,b,purpose);after=direct_update_plan(a,saved,purpose)
        for field in ['updates','labels','identity_matches','retained_blanks']:
            self.assertEqual(before[field],after[field])

    def test_本年余额起末角色不被年度列覆盖且不改原日期(self):
        b=self.save('B')
        reply=('合成表：\n列B为本年，列C为上年；本年期间为2025-01-01至2025-12-31，'
               '上年期间为2024-01-01至2024-12-31。余额行的年初数按相应期间起日取时点，'
               '年末数按相应期间末日取时点。\n'
               '- 计量：`固定资产原值`\n'
               '  - 年初数：B3:C3。\n  - 年末数：B4:C4。\n')
        parsed=direct_answer(reply,self.engine.standard,['合成表'])
        rows={(row['ranges'][0],row['id']):row['value'] for row in parsed['dimensions']}
        for cell,role in [('B3','opening'),('B4','closing'),('C3','prior'),('C4','prior')]:
            self.assertEqual(rows[('合成表!'+cell,'template_period_role')],role)
        self.assertEqual(rows[('合成表!B3','period')],'2025-01-01')
        self.assertEqual(rows[('合成表!C3','period')],'2024-01-01')
        saved=save_direct_answer(b,reply,self.case.folder/'只接余额角色.json',
            field_scope={'合成表!B3:B4':['template_period_role']})
        before=dict(_records(b));after=dict(_records(saved))
        self.assertEqual({k:v['dimensions'] for k,v in before.items()},{k:v['dimensions'] for k,v in after.items()})
        roles={cell:row.get('template_period_role') for table in saved['outline']['subtables']
               for row in table['metrics'] for cell in row['ranges']}
        self.assertEqual(roles['B3'],'opening');self.assertEqual(roles['B4'],'closing')
        self.assertEqual(saved['机械记录'],b['机械记录'])

    def test_共同币种与倍率的中文分隔不混入取值(self):
        parsed=direct_answer('共同取值：金额币种为 CNY、金额倍率为 1。',self.engine.standard,['合成表'])
        values={r['id']:r['value'] for r in parsed['dimensions']}
        self.assertEqual(values,{'currency':'CNY','unit_scale':1})

    def test_只接币种不把含标准编号的说明当币种(self):
        b=self.save('B')
        reply='共同取值：币种为 CNY。\n合成表!B3:C3：余额选 固定资产原值；对象按原答复保留。'
        saved=save_direct_answer(b,reply,self.case.folder/'共同币种.json',
            field_scope={'合成表!B3:C3':['currency']})
        self.assertEqual(saved['unresolved'],b['unresolved'])
        self.assertEqual(saved['机械记录'],b['机械记录'])
        records=lambda mapping:{loc:{k:v for k,v in row.items() if k!='recognition_provenance'}
                                for loc,row in _records(mapping)}
        self.assertEqual(records(saved),records(b))

    def test_原文名称的明确编号引用接入且只清同句解析错误(self):
        line='- 合成表!B3 甲值、C3 乙值：原范围。'
        answer=json.loads(self.case.answer())
        answer['tables'][0]['unresolved']=[
            {'ranges':['B3','C3'],'reason':'这段答复未明确可机械读取的字段归属','original':line},
            {'ranges':['B4'],'reason':'原文对象未知','original':'此项原文对象未知'}]
        b=self.save('B',json.dumps(answer,ensure_ascii=False))
        reply=line+'\n上述甲值格用 固定资产原值；乙值格用 固定资产原值。'
        saved=save_direct_answer(b,reply,self.case.folder/'名称编号引用.json',
            field_scope={'合成表!B3':['metric_id']})
        self.assertFalse(any(r.get('cell')=='B3' for r in saved['unresolved']))
        self.assertTrue(any(r.get('cell')=='C3' for r in saved['unresolved']))
        self.assertTrue(any(r.get('reason')=='原文对象未知' for r in saved['unresolved']))
        self.assertEqual(dict(_records(saved))['合成表!B3']['dimensions'],dict(_records(b))['合成表!B3']['dimensions'])
        other=direct_answer(line+'\n### 另一节\n上述甲值格用 固定资产原值。',self.engine.standard,['合成表'])
        self.assertFalse(any('合成表!B3' in r['ranges'] for r in other['metrics']))
        unknown=direct_answer(reply.replace('原范围。','原范围未知。'),self.engine.standard,['合成表'])
        self.assertTrue(any('合成表!B3' in r['ranges'] for r in unknown['unresolved']))

    def test_只重接同一编号不接收未列出的未知维度取值(self):
        b=self.save('B')
        b=save_direct_answer(b,'合成表!B3：asset_selection 为 unknown。',self.case.folder/'原范围未知.json',
            field_scope={'合成表!B3':['asset_selection']})
        reply='合成表!B3：metric_id=固定资产原值；asset_selection 为 members：旧答复中的成员。'
        saved=save_direct_answer(b,reply,self.case.folder/'只重接编号.json',
            field_scope={'合成表!B3':['metric_id']})
        self.assertEqual(dict(_records(saved))['合成表!B3']['dimensions'],dict(_records(b))['合成表!B3']['dimensions'])
        self.assertEqual(saved['机械记录'],b['机械记录'])

    def test_明确同范围进入合集并在正常匹配入口使用(self):
        a=self.save('A');b=self.save('B')
        reply='合成表!B3:C3：asset_selection 为 members：设备甲。这与A所需范围相同。'
        b=self.corrected(a,b,reply)
        self.assertTrue(b['已确认关系'])
        self.assertEqual(dict(_records(b))['合成表!B3']['dimensions']['asset_selection']['members'],['设备甲'])
        self.assertEqual(dict(_records(a))['合成表!B3']['dimensions']['asset_selection']['members'],['测试项目3'])
        combined=combine_direct_collections([b]);self.assertEqual(combined['已确认关系'],b['已确认关系'])
        with patch.object(UnifiedSemanticEngine,'_http',side_effect=AssertionError('不调用模型')):
            result=run_step('match',{**self.case.config,'a_mapping':a['mapping_path'],'b_mapping':b['mapping_path']})
        plan=json.loads(__import__('pathlib').Path(result['files']['update_plan']).read_text(encoding='utf-8-sig'))
        selected=next(r for r in plan['identity_matches'] if r['cell']=='B3')
        self.assertEqual(selected['source_locations'],['合成表!B3'])
        self.assertTrue(selected['已确认关系'])
        changed=deepcopy(a)
        next(r for r in changed['facts'] if r['id']=='合成表!B3')['dimensions']['asset_selection']['members']=['后续改变']
        changed['封存校验']=direct_digest(changed)
        stale=direct_update_plan(changed,b)
        self.assertFalse(any(r['cell']=='B3' for r in stale['identity_matches']))
        self.assertTrue(any(r['cell']=='C3' for r in stale['identity_matches']))

    def test_模型明示的多项同范围关系全部使用不按顺序择一(self):
        a=self.save('A');b=self.corrected(a,self.save('B'),
            '合成表!B3:C3：asset_selection 为 members：设备甲。这与A所需范围相同。')
        second=deepcopy(b['已确认关系'][0])
        second['record_ranges']=['合成表!B3']
        second['A需要']['locations']=['合成表!B4']
        second['A需要']['value']=deepcopy(dict(_records(a))['合成表!B4']['dimensions']['asset_selection'])
        second['content']='合成答复明确同时对应另一项已给定范围'
        b['已确认关系'].append(second)
        b['facts']=[row for row in b['facts'] if row['source_reference']['location']=='合成表!B3']
        b['封存校验']=direct_digest(b)
        before=deepcopy((a,b));plan=direct_update_plan(a,b)
        for cell in ('B3','B4'):
            row=next(row for row in plan['updates'] if row['cell']==cell)
            self.assertEqual([source['fact_id'] for source in row['sources']],['合成表!B3'])
        self.assertEqual((a,b),before)

    def test_局部重读不按新上下文重解释未修改字段关系(self):
        a=self.save('A');b=self.save('B')
        reply='合成表!B3:C3：asset_selection 为 members：设备甲。这与A所需范围相同。'
        b=self.corrected(a,b,reply)
        references=[{},direct_reference(a,compact=True,ranges=[{'sheet':'合成表','range':'B4:C4'}])]
        for index,reference in enumerate(references):
            with self.subTest(reference=index):
                saved=save_direct_answer(b,reply+'\n合成表!B3:C3：template_period_role=closing。',
                    self.case.folder/f'仅重读期间{index}.json',
                    field_scope={'合成表!B3:C3':['template_period_role']},reference=reference)
                self.assertEqual(saved['已确认关系'],b['已确认关系'])
                self.assertEqual(saved['关系缺口'],b['关系缺口'])
                self.assertEqual(saved['机械记录'],b['机械记录'])
                for loc,row in _records(saved):
                    self.assertEqual(row['dimensions'],dict(_records(b))[loc]['dimensions'])

    def test_否定和歧义引用不能形成关系(self):
        a=self.save('A');b=self.save('B')
        negative=self.corrected(a,b,'合成表!B3:C3：asset_selection 为 members：设备甲。不能认定这与A所需范围相同。')
        self.assertFalse(negative['已确认关系'])
        ambiguous=save_direct_answer(b,'合成表!B3:C3：asset_selection 为 members：设备甲。这与A所需范围相同。',
            self.case.folder/'歧义引用.json',field_scope={'合成表!B3:C3':['asset_selection']},reference=direct_reference(a,compact=True))
        self.assertFalse(ambiguous['已确认关系']);self.assertTrue(ambiguous['关系缺口'])

    def test_多种需要中的明确同范围声明按原字句定位(self):
        answer=json.loads(self.case.answer())
        row=next(r for r in answer['tables'][0]['dimensions'] if r['id']=='asset_selection' and r['ranges']==['B3:C3'])
        row['value']={'domain':'固定资产','mode':'predicate','predicate':'设备甲且用于生产。'}
        a=self.save('A',json.dumps(answer,ensure_ascii=False))
        row['value']['predicate']='设备甲且用于生产'
        b=self.save('B',json.dumps(answer,ensure_ascii=False))
        saved=save_direct_answer(b,'合成表!B3:C3：asset_selection 为 predicate：设备甲且用于生产。与A需要的设备范围相同。',
            self.case.folder/'多范围引用.json',field_scope={'合成表!B3:C3':['asset_selection']},reference=direct_reference(a,compact=True))
        self.assertTrue(saved['已确认关系'])
        self.assertTrue(all(set(r['A需要']['locations'])=={'合成表!B3','合成表!C3'} for r in saved['已确认关系']))
        self.assertEqual(dict(_records(a))['合成表!B3']['dimensions']['asset_selection']['predicate'],'设备甲且用于生产。')
        self.assertEqual(dict(_records(saved))['合成表!B3']['dimensions']['asset_selection']['predicate'],'设备甲且用于生产')
        self.assertTrue(any(r['cell']=='B3' for r in direct_update_plan(a,saved)['identity_matches']))
        unstated=save_direct_answer(b,'合成表!B3:C3：asset_selection 为 predicate：设备甲且用于生产。',
            self.case.folder/'没有关系声明.json',field_scope={'合成表!B3:C3':['asset_selection']},reference=direct_reference(a,compact=True))
        self.assertFalse(unstated['已确认关系'])

    def test_共用B纠正出口保留本次引用材料(self):
        a=self.save('A');b=self.save('B')
        config={**self.case.config,'a_mapping':a['mapping_path'],'b_mapping':b['mapping_path'],
                'a_required_ranges':[{'sheet':'合成表','range':'B3:C3'}],
                'b_source_ranges':[{'sheet':'合成表','range':'A1:C4'}],
                'b_field_scope':{'合成表!B3:C3':['asset_selection']}}
        with patch.object(UnifiedSemanticEngine,'_http',return_value='合成表!B3:C3：asset_selection 为 members：设备甲。这与A所需范围相同。'):
            result=run_step('recognize_b',config)
        self.assertIn('b_mapping',result['files'],result)
        saved=json.loads(__import__('pathlib').Path(result['files']['b_mapping']).read_text(encoding='utf-8-sig'))
        self.assertTrue(saved['已确认关系']);self.assertIn('A需要',saved)

    def test_无候选多身份及语义缺失分别报告而金额不裁决(self):
        a=self.save('A');b=self.save('B')
        empty=deepcopy(b);empty['facts']=[];empty['slot_bindings']=[];empty['封存校验']=direct_digest(empty)
        self.assertEqual(direct_update_plan(a,empty)['issues'][0]['kind'],'missing_source')
        answer=json.loads(self.case.answer())
        dims=answer['tables'][0]['dimensions']
        source=next(r for r in dims if r['id']=='asset_selection' and r['ranges']==['B3:C3'])
        next(r for r in dims if r['id']=='asset_selection' and r['ranges']==['B4:C4'])['value']=deepcopy(source['value'])
        multiple=self.save('B',json.dumps(answer,ensure_ascii=False))
        plan=direct_update_plan(a,multiple)
        matched=next(r for r in plan['identity_matches'] if r['cell']=='B3')
        self.assertEqual(set(matched['source_locations']),{'合成表!B3','合成表!B4'})
        self.assertFalse(any(r['cell']=='B3' for r in plan['updates']))
        self.assertTrue(any(r.get('cell')=='B3' and r['kind']=='ambiguous_source' for r in plan['issues']))
        period=next(r for r in dims if r['id']=='period' and r['ranges']==['B3:B8'])
        period['ranges']=['B3','B5:B8']
        dims.append({**deepcopy(period),'ranges':['B4'],'value':{'kind':'instant','date':'2026-12-31'}})
        multiple=self.save('B',json.dumps(answer,ensure_ascii=False))
        continued=direct_update_plan(a,multiple)
        matched=next(r for r in continued['identity_matches'] if r['cell']=='B3')
        self.assertEqual(set(matched['source_locations']),{'合成表!B3','合成表!B4'})
        self.assertFalse(any(r['cell']=='B3' for r in continued['updates']))
        self.assertTrue(any(r.get('cell')=='B3' and r['kind']=='ambiguous_source' for r in continued['issues']))
        missing=self.corrected(a,b,'合成表!B3:C3：asset_selection 为 unknown。')
        partial=direct_update_plan(a,missing)
        issue=next(r for r in partial['issues'] if r.get('cell')=='B3')
        self.assertEqual(issue['kind'],'semantic_difference')
        self.assertNotIn('合成表!B3',{r['location'] for r in issue['candidates']})
        source=next(r for r in partial['unplaced_sources'] if r['id']=='合成表!B3')
        self.assertEqual(source['missing_dimensions'],['asset_selection'])
        limited=direct_update_plan(a,missing,scope={'target':[{'sheet':'合成表','range':'B3'}],
            'source':[{'sheet':'合成表','range':'B3'}]})
        self.assertEqual(limited['issues'][0]['kind'],'missing_source')
        self.assertEqual(limited['issues'][0]['candidates'],[])

    def test_历史解析说明不撤销当前身份且真实未知仍单列(self):
        a=self.save('A');b=self.save('B')
        b=save_direct_answer(b,'合成表!C3：asset_selection=未知。',self.case.folder/'真实未知.json',
            field_scope={'合成表!C3':['asset_selection']})
        for mapping in (a,b):
            mapping['unresolved'].append({'sheet':'合成表','cell':'B3','reason':'这段答复未明确可机械读取的字段归属'})
            mapping['封存校验']=direct_digest(mapping)
        plan=direct_update_plan(a,b,{'reason':'沿当前合集机械对应'})
        self.assertTrue(any(row['cell']=='B3' for row in plan['identity_matches']))
        self.assertFalse(any(row['cell']=='C3' for row in plan['identity_matches']))
        self.assertTrue(any(row['id']=='合成表!C3' and row.get('missing_dimensions')==['asset_selection']
                            for row in plan['unplaced_sources']))
        self.assertFalse(any(row.get('cell')=='C3' and row.get('kind')=='incomplete_identity' for row in plan['issues']))
        self.assertTrue(any(row.get('cell')=='B3' and row.get('carrier')=='B' for row in plan['issues']))

    def test_多来源保留对应但不选第一项且明确来源直接执行(self):
        from 附注更新.A合集 import direct_output_record
        a=self.save('A');answer=json.loads(self.case.answer());dims=answer['tables'][0]['dimensions']
        shared=next(r for r in dims if r['id']=='asset_selection' and r['ranges']==['B3:C3'])
        repeated=next(r for r in dims if r['id']=='asset_selection' and r['ranges']==['B4:C4'])
        repeated.update(value=deepcopy(shared['value']),evidence=deepcopy(shared['evidence']))
        snapshot=deepcopy(self.snapshot)
        for amount in (6,999,None):
            with self.subTest(amount=amount):
                snapshot['sheets'][0]['cells']['B4']['value']=amount
                b=self.save('B',json.dumps(answer,ensure_ascii=False),snapshot)
                plan=direct_update_plan(a,b)
                self.assertFalse(any(r['cell']=='B3' for r in plan['updates']))
                matched=next(r for r in plan['identity_matches'] if r['cell']=='B3')
                self.assertEqual(set(matched['source_locations']),{'合成表!B3','合成表!B4'})
                self.assertTrue(any(r['cell']=='B3' and r['kind']=='ambiguous_source' for r in plan['issues']))
                self.assertTrue(any(r['cell']=='B5' for r in plan['retained_blanks']))
                before=deepcopy((a,b))
                with patch('附注更新.表格.read_workbook',return_value=deepcopy(self.snapshot)), patch('附注更新.表格.file_hash',return_value=self.snapshot['sha256']):
                    derived=direct_output_record(a,b,self.case.source,plan,'已保存安排')
                self.assertEqual(derived['identity_matches'],plan['identity_matches'])
                self.assertNotIn('facts',derived)
                self.assertEqual((a,b),before)
        purpose={'reason':'已明确搬运来源','bindings':[{'target':'合成表!B3','source':'合成表!B4'}]}
        snapshot['sheets'][0]['cells']['B4']['value']=999
        b=self.save('B',json.dumps(answer,ensure_ascii=False),snapshot)
        chosen=direct_update_plan(a,b,purpose)
        update=next(r for r in chosen['updates'] if r['cell']=='B3')
        self.assertEqual(update['value'],999)
        self.assertEqual([r['fact_id'] for r in update['sources']],['合成表!B4'])

    def test_保存不重新判断期间性质允许值或引用存在(self):
        answer=json.loads(self.case.answer());dims=answer['tables'][0]['dimensions']
        next(r for r in dims if r['id']=='currency')['value']='模型原答复币种'
        period=next(r for r in dims if r['id']=='period')
        period.update(value={'kind':'duration','start':'2025-01-01','end':'2025-12-31'},evidence=['合成表!Z999','原答复保留的引用'])
        with patch('附注更新.统一语义.validate_fact',side_effect=AssertionError('保存不作业务判断')), \
             patch('附注更新.统一语义.validate_partial_semantics',side_effect=AssertionError('保存不作业务判断')):
            saved=self.save('A',json.dumps(answer,ensure_ascii=False))
            partial=deepcopy(answer)
            partial['tables'][0]['dimensions']=[r for r in dims if r['id']!='entity']
            saved_partial=self.save('A',json.dumps(partial,ensure_ascii=False))
        for mapping in (saved,saved_partial):
            row=dict(_records(mapping))['合成表!B3']
            self.assertEqual(row['dimensions']['currency'],'模型原答复币种')
            self.assertEqual(row['dimensions']['period'],period['value'])
            self.assertEqual(row['dimension_evidence']['period'],period['evidence'])
            self.assertEqual(set(dict(_records(mapping))),set(dict(_records(self.save('A')))))

    def test_未声明范围类型的文字不由程序变成条件(self):
        answer=json.loads(self.case.answer())
        row=next(r for r in answer['tables'][0]['dimensions'] if r['id']=='asset_selection')
        row['value']='原答复的范围说明'
        saved=self.save('A',json.dumps(answer,ensure_ascii=False))
        self.assertEqual(dict(_records(saved))['合成表!B3']['dimensions']['asset_selection'],row['value'])

    def test_读取已有合集不重新核验公式副本(self):
        from 附注更新.统一勾稽 import load_fact_mapping
        saved=self.save('B')
        saved['formula_values']={'合成表!B3':{'state':'value','value':6,'formula':'=1+5'}}
        saved['calculation_evidence']={'path':'原记录中的公式依据','sha256':'原记录指纹'}
        saved['封存校验']=direct_digest(saved)
        path=self.case.folder/'已有公式合集.json'
        path.write_text(json.dumps(saved,ensure_ascii=False),encoding='utf-8')
        with patch('附注更新.统一勾稽.calculation_files',side_effect=AssertionError('读取不能核验公式')), \
             patch('附注更新.统一勾稽._calculated_values',side_effect=AssertionError('读取不能重算公式')):
            loaded=load_fact_mapping(path,{})
        self.assertEqual(loaded['formula_values'],saved['formula_values'])
        self.assertEqual(loaded['calculation_evidence'],saved['calculation_evidence'])

    def test_空白及公式不撤销已确认换期和名单关系(self):
        a=self.save('A');answer=json.loads(self.case.answer(2026))
        row=next(r for r in answer['tables'][0]['dimensions'] if r['id']=='asset_selection')
        row['value']['members']=['新名单成员']
        snapshot=deepcopy(self.snapshot);cells=snapshot['sheets'][0]['cells']
        for cell,value in [('A3','新名单成员'),('B2','新期末'),('C2','新期初'),('B3',None),('C3',None)]:
            cells[cell].update(value=value,formula=None)
        purpose={'reason':'名单和期间按B更新，来源空白保留A原值',
            'bindings':[{'target':f'合成表!{col}3','source':f'合成表!{col}3'} for col in 'BC'],
            'labels':[{'target':'合成表!A3','source':'合成表!A3'}]}
        for formula in (None,'=1+5'):
            with self.subTest(formula=formula):
                cells['B3'].update(value=formula,formula=formula)
                b=self.save('B',json.dumps(answer,ensure_ascii=False),snapshot)
                original=deepcopy((a,b))
                with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('不调用模型')):
                    result=run_step('match',{**self.case.config,'a_mapping':a['mapping_path'],
                        'b_mapping':b['mapping_path'],'update_purpose':purpose})
                plan=json.loads(Path(result['files']['update_plan']).read_text(encoding='utf-8-sig'))
                self.assertTrue({'B3','C3'}<={r['cell'] for r in plan['identity_matches']})
                self.assertTrue({'A3','B2','C2'}<={r['cell'] for r in plan['labels']})
                self.assertFalse(any(r['cell'] in ('B3','C3') for r in plan['updates']))
                self.assertEqual({r['cell'] for r in plan['retained_blanks']} & {'B3','C3'},
                                 {'C3'} if formula else {'B3','C3'})
                self.assertTrue(all(r.get('kind')=='unevaluated_source_formula' for r in plan['issues']))
                self.assertEqual((a,b),original)

    def test_明确搬运来源不再复判身份字段和期间角色(self):
        a=self.save('A');answer=json.loads(self.case.answer(2026))
        dims=answer['tables'][0]['dimensions']
        dims[:]=[row for row in dims if row['id']!='entity']
        next(row for row in dims if row['id']=='currency')['value']='模型已写明的币种取值'
        next(row for row in dims if row['id']=='asset_selection' and row['ranges']==['B4:C4'])['value']={'domain':'固定资产','mode':'unknown'}
        b=self.save('B',json.dumps(answer,ensure_ascii=False));before=deepcopy((a,b))
        purpose={'reason':'原答复已经明确这项搬运对应','bindings':[{'target':'合成表!B3','source':'合成表!C4'}]}
        with patch('附注更新.统一语义.identity_completeness',side_effect=AssertionError('匹配不能重判完整性')), \
             patch('附注更新.统一语义._dimension_value',side_effect=AssertionError('匹配不能验证业务取值')), \
             patch('附注更新.统一语义.slot_key',side_effect=AssertionError('匹配不能重判槽位资格')):
            plan=direct_update_plan(a,b,purpose)
        update=next(row for row in plan['updates'] if row['cell']=='B3')
        self.assertEqual([row['fact_id'] for row in update['sources']],['合成表!C4'])
        self.assertEqual(update['value'],12)
        self.assertEqual(update['source_identity']['dimensions'],dict(_records(b))['合成表!C4']['dimensions'])
        self.assertNotIn('entity',update['source_identity']['dimensions'])
        self.assertEqual(update['source_identity']['dimensions']['asset_selection']['mode'],'unknown')
        self.assertEqual((a,b),before)

    def test_机械比较只使用已有取值不审核业务允许范围(self):
        answer=json.loads(self.case.answer())
        next(row for row in answer['tables'][0]['dimensions'] if row['id']=='currency')['value']='答复中的币种文字'
        a=self.save('A',json.dumps(answer,ensure_ascii=False));b=self.save('B',json.dumps(answer,ensure_ascii=False))
        with patch('附注更新.统一语义.identity_completeness',side_effect=AssertionError('不重判完整性')), \
             patch('附注更新.统一语义._dimension_value',side_effect=AssertionError('不验证业务值')):
            plan=direct_update_plan(a,b)
        self.assertEqual(next(row for row in plan['updates'] if row['cell']=='B3')['source_identity']['dimensions']['currency'],'答复中的币种文字')

    def test_比较排除只用明确要求且原身份原值不变(self):
        a=self.save('A')
        requirements=[{'维度ID':key,'要求来源':'合成用户要求：本次不比较该维度'}
                      for key in ('entity','report_scope','unit_scale')]
        for key,value in [('entity','来源公司'),('report_scope','consolidated'),('unit_scale',10000),('currency','USD')]:
            with self.subTest(dimension=key):
                answer=json.loads(self.case.answer())
                next(r for r in answer['tables'][0]['dimensions'] if r['id']==key)['value']=value
                b=self.save('B',json.dumps(answer,ensure_ascii=False));before=deepcopy(b)
                self.assertFalse(direct_update_plan(a,b)['identity_matches'])
                plan=direct_update_plan(a,b,comparison_requirements=requirements)
                self.assertEqual(bool(plan['identity_matches']),key!='currency')
                if key!='currency':
                    row=next(r for r in plan['updates'] if r['cell']=='B3')
                    self.assertEqual(row['source_identity']['dimensions'][key],value)
                    self.assertEqual(row['value'],6)
                self.assertEqual(b,before)
        for invalid in [[{'维度ID':'不存在','要求来源':'合成说明'}]]:
            with self.assertRaises(ValueError):direct_update_plan(a,a,comparison_requirements=invalid)

    def test_明确比较排除不依赖附加出处说明(self):
        a=self.save('A');answer=json.loads(self.case.answer())
        next(row for row in answer['tables'][0]['dimensions'] if row['id']=='entity')['value']='来源公司'
        b=self.save('B',json.dumps(answer,ensure_ascii=False));before=deepcopy((a,b))
        documented=direct_update_plan(a,b,comparison_requirements=[{'维度ID':'entity','要求来源':'合成用户要求'}])
        requirements=[{'维度ID':'entity'}]
        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('不调用模型')):
            result=run_step('match',{**self.case.config,'a_mapping':a['mapping_path'],
                                    'b_mapping':b['mapping_path'],'comparison_requirements':requirements})
        self.assertEqual(result['status'],'complete',result)
        plan=json.loads(Path(result['files']['update_plan']).read_text(encoding='utf-8-sig'))
        for key in ('updates','labels','identity_matches','retained_blanks','operations'):
            self.assertEqual(plan[key],documented[key])
        self.assertEqual(plan['comparison_requirements'],requirements)
        self.assertEqual(next(row for row in plan['updates'] if row['cell']=='B3')['source_identity']['dimensions']['entity'],'来源公司')
        self.assertEqual((a,b),before)

    def test_比较排除不把身份缺失改为完整(self):
        from 附注更新.统一语义 import identity_completeness
        a=self.save('A');answer=json.loads(self.case.answer())
        answer['tables'][0]['dimensions']=[r for r in answer['tables'][0]['dimensions'] if r['id']!='entity']
        b=self.save('B',json.dumps(answer,ensure_ascii=False));before=deepcopy(b)
        plan=direct_update_plan(a,b,comparison_requirements=[{'维度ID':'entity','要求来源':'合成用户要求'}])
        self.assertTrue(plan['identity_matches'],plan['issues'])
        self.assertEqual(b,before)
        status=identity_completeness(dict(_records(b))['合成表!B3'],self.engine.standard)
        self.assertIn('entity',status['missing_dimensions']);self.assertFalse(status['complete'])

    def test_安排保存成功不裁定语义完成也不丢弃未明确内容(self):
        a=self.save('A');b=self.save('B')
        b['facts']=[];b['slot_bindings']=[];b['封存校验']=direct_digest(b)
        path=self.case.folder/'未提供来源合集.json'
        path.write_text(json.dumps(b,ensure_ascii=False),encoding='utf-8')
        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('不调用模型')):
            result=run_step('match',{**self.case.config,'a_mapping':a['mapping_path'],'b_mapping':str(path)})
        self.assertEqual(result['status'],'complete')
        self.assertIn('安排已保存',result['message'])
        plan=json.loads(Path(result['files']['update_plan']).read_text(encoding='utf-8-sig'))
        self.assertTrue(plan['issues']);self.assertFalse(plan['updates'])
        self.assertNotIn('complete',plan)

    def test_比较要求贯通正常入口且写入不重新匹配(self):
        from 附注更新.统一勾稽 import build_fact_update_plan
        from 附注更新.分步流程 import _write_direct
        a=self.save('A');answer=json.loads(self.case.answer())
        next(r for r in answer['tables'][0]['dimensions'] if r['id']=='entity')['value']='来源公司'
        b=self.save('B',json.dumps(answer,ensure_ascii=False))
        requirements=[{'维度ID':'entity','要求来源':'合成用户要求'}]
        config={**self.case.config,'a_mapping':a['mapping_path'],'b_mapping':b['mapping_path'],
                'comparison_requirements':requirements}
        with patch.object(UnifiedSemanticEngine,'_http',side_effect=AssertionError('不调用模型')):
            result=run_step('match',config)
        path=Path(result['files']['update_plan']);plan=json.loads(path.read_text(encoding='utf-8-sig'))
        self.assertEqual(plan['comparison_requirements'],requirements)
        self.assertTrue(plan['updates'])
        generic=build_fact_update_plan(a,b,self.engine.standard,comparison_requirements=requirements)
        self.assertEqual(generic['updates'],plan['updates'])
        with self.assertRaisesRegex(ValueError,'旧映射入口不能忽略该要求'):
            build_fact_update_plan({}, {}, self.engine.standard,comparison_requirements=requirements)
        with patch('附注更新.模板更新编排.direct_update_plan',side_effect=AssertionError('写入不能重新匹配')), \
             patch('附注更新.分步流程.write_updated_workbook',side_effect=RuntimeError('已到副本入口')) as write:
            with self.assertRaisesRegex(RuntimeError,'已到副本入口'):
                _write_direct(config,self.case.folder,lambda text:None,None,path,plan,a,b)
            write.assert_called_once()
            self.assertEqual(write.call_args.args[2]['comparison_requirements'],requirements)
            write.reset_mock()
            with self.assertRaisesRegex(RuntimeError,'已到副本入口'):
                _write_direct({**config,'comparison_requirements':[]},self.case.folder,lambda text:None,None,path,plan,a,b)
            self.assertIs(write.call_args.args[2],plan)

    def test_结构动作局部停止时比较要求继续传递(self):
        snapshot=deepcopy(self.snapshot);snapshot['names']={'局部表':[{'sheet':'合成表','range':'A2:C3'}]}
        a=self.save('A',snapshot=snapshot);answer=json.loads(self.case.answer())
        next(r for r in answer['tables'][0]['dimensions'] if r['id']=='entity')['value']='来源公司'
        b=self.save('B',json.dumps(answer,ensure_ascii=False));requirements=[{'维度ID':'entity','要求来源':'合成用户要求'}]
        purpose={'reason':'合成名单删除缺少完整来源时局部停止','operations':[
            {'range_name':'局部表','axis':'row','action':'delete','index':99,'count':1,'template_index':1}]}
        plan=direct_update_plan(a,b,purpose,comparison_requirements=requirements)
        self.assertFalse(plan['operations']);self.assertEqual(plan['comparison_requirements'],requirements)
        self.assertTrue(any(row['cell']=='B4' for row in plan['identity_matches']))
        self.assertTrue(any(row['cell']=='B5' for row in plan['retained_blanks']))

    def two_lists(self,year=2025):
        snapshot=deepcopy(self.snapshot)
        snapshot['names']={'第一表':[{'sheet':'合成表','range':'A2:C3'}],
                           '第二表':[{'sheet':'合成表','range':'A5:C8'}]}
        a=self.save('A',snapshot=snapshot);snapshot['sheets'][0]['max_row']=9
        for col,value in [('A','新增项目'),('B',90),('C',91)]:
            snapshot['sheets'][0]['cells'][col+'9']={**snapshot['sheets'][0]['cells'][col+'3'],'value':value}
        answer=json.loads(self.case.answer(year));table=answer['tables'][0]
        table.update(range='A2:C9',data_ranges=['B3:C9'])
        for row in table['metrics']:row['ranges']=[span.replace('8','9') for span in row['ranges']]
        for row in table['dimensions']:
            if row['id']=='period':row['ranges']=[span.replace('8','9') for span in row['ranges']]
        table['dimensions'].append({'id':'asset_selection','ranges':['B9:C9'],
            'value':{'domain':'固定资产','mode':'members','members':['新增项目']},'evidence':['合成表!A9']})
        if year!=2025:
            snapshot['sheets'][0]['cells']['B2']['value']='新期末'
            snapshot['sheets'][0]['cells']['C2']['value']='新期初'
        b=self.save('B',json.dumps(answer,ensure_ascii=False),snapshot)
        table={'range_name':'第二表','source_sheet':'合成表','source_rows':[6,7,8,9],
               'old_detail_count':3,'first_row_offset':1,'columns':{'B':'B','C':'C'},'labels':{'A':'A'}}
        return a,b,table

    def test_明确名单删除直接展开不审核范围依据(self):
        a,b,table=self.two_lists()
        operation={'range_name':'第一表','axis':'row','action':'delete','index':1,'count':1,'template_index':1}
        purpose={'reason':'已经明确的名单删行及新增安排','tables':[table],'operations':[operation]}
        before=deepcopy((a,b));plan=direct_update_plan(a,b,purpose)
        self.assertIn(operation,plan['operations'])
        self.assertIsNone(plan['template_layout']['positions']['合成表!B3'])
        self.assertTrue(any(row['range_name']=='第二表' for row in plan['operations']))
        self.assertFalse(any('范围依据' in row.get('reason','') for row in plan['issues']))
        self.assertEqual((a,b),before)

    def test_一张名单物理位置越界不撤掉另一名单新增(self):
        a,b,table=self.two_lists()
        purpose={'reason':'合成明确名单安排','tables':[table],'operations':[
            {'range_name':'第一表','axis':'row','action':'delete','index':99,'count':1,'template_index':1}]}
        with patch.object(UnifiedSemanticEngine,'_http',side_effect=AssertionError('不调用模型')):
            result=run_step('match',{**self.case.config,'a_mapping':a['mapping_path'],
                'b_mapping':b['mapping_path'],'update_purpose':purpose})
        plan=json.loads(Path(result['files']['update_plan']).read_text(encoding='utf-8-sig'))
        self.assertEqual([r['range_name'] for r in plan['operations']],['第二表'])
        self.assertTrue(any(r['cell']=='B9' and r['value']==90 for r in plan['updates']))
        self.assertTrue(any(r['cell']=='A9' and r['value']=='新增项目' for r in plan['labels']))
        self.assertFalse(any(r['cell']=='B3' for r in plan['updates']))
        self.assertTrue(any(r['cell']=='B4' for r in plan['identity_matches']))
        self.assertEqual(plan['update_purpose'],purpose)
        from 附注更新.分步流程 import _write_direct
        with patch('附注更新.分步流程.write_updated_workbook',side_effect=RuntimeError('到达副本入口')) as write:
            with self.assertRaisesRegex(RuntimeError,'到达副本入口'):
                _write_direct(self.case.config,self.case.folder,lambda _:None,None,
                              Path(result['files']['update_plan']),plan,a,b)
            self.assertEqual(write.call_args.args[2]['operations'],plan['operations'])

    def test_名单身份缺失保留该项而不撤销明确结构及其他对应(self):
        a,b,table=self.two_lists()
        b=save_direct_answer(b,'合成表!B4:C4：asset_selection 为 unknown。',self.case.folder/'第一名单未知.json',
                             field_scope={'合成表!B4:C4':['asset_selection']})
        first={**table,'range_name':'第一表','source_rows':[3,4],'old_detail_count':1}
        purpose={'reason':'合成明确名单安排','tables':[first,table],
                 'bindings':[{'target':'合成表!B8','source':'合成表!B7'}]}
        plan=direct_update_plan(a,b,purpose)
        self.assertEqual([r['range_name'] for r in plan['operations']],['第一表','第二表'])
        self.assertTrue(any(r['cell']=='B10' and r['value']==90 for r in plan['updates']))
        self.assertTrue(any(r['cell']=='B8' for r in plan['retained_blanks']))
        self.assertEqual(plan['template_layout']['positions']['合成表!B7'],'合成表!B8')
        self.assertFalse(any(r.get('cell')=='B4' and r.get('kind')=='incomplete_identity' for r in plan['issues']))
        matched=next(r for r in plan['identity_matches'] if r['cell']=='B4')
        self.assertEqual(matched['source_locations'],['合成表!B4'])
        self.assertEqual(next(r for r in plan['updates'] if r['cell']=='B4')['source_identity']['dimensions']['asset_selection']['mode'],'unknown')
        self.assertEqual(plan['template_layout']['positions']['合成表!B4'],'合成表!B5')

    def test_共用期间标签不连带撤销已确认名单及对应(self):
        a,b,table=self.two_lists(2026)
        purpose={'reason':'合成明确名单及换期安排','tables':[table],'operations':[
            {'range_name':'第一表','axis':'row','action':'delete','index':99,'count':1,'template_index':1}]}
        plan=direct_update_plan(a,b,purpose)
        self.assertEqual([r['range_name'] for r in plan['operations']],['第二表'])
        self.assertTrue(any(r['cell']=='B2' and r['value']=='新期末' for r in plan['labels']))
        self.assertTrue(any(r['cell']=='B9' and r['value']==90 for r in plan['updates']))
        self.assertFalse(any('共享标签依赖' in str(r) for r in plan['issues']))

    def test_关系比较接续换期且不阻断无关未决(self):
        a=self.save('A');snapshot=deepcopy(self.snapshot)
        snapshot['sheets'][0]['cells']['B2']['value']='新期末'
        snapshot['sheets'][0]['cells']['C2']['value']='新期初'
        b=self.save('B',self.case.answer(2026),snapshot)
        b=self.corrected(a,b,'合成表!B3:C3：asset_selection 为 members：设备甲。这与A所需范围相同。')
        a['unresolved'].append({'sheet':'合成表','cell':'A8','reason':'不引用期间标签的独立说明尚未明确'})
        a['封存校验']=direct_digest(a)
        with patch('openpyxl.load_workbook',side_effect=AssertionError('只读合集')):
            plan=direct_update_plan(a,b,{'reason':'按已明确要求换期'})
        self.assertEqual({r['value'] for r in plan['labels']},{'新期末','新期初'})
        self.assertTrue(any(r['cell']=='B3' for r in plan['identity_matches']))

    def test_期间标签缺引用不撤销金额对应及已知名单名称(self):
        a,b,table=self.two_lists(2026)
        for row in b['facts']:
            row['dimension_evidence']['period']=[]
            row.get('label_references',{}).pop('period',None)
        b['封存校验']=direct_digest(b)
        before=deepcopy(b)
        plan=direct_update_plan(a,b,{'reason':'按已明确角色换期及更新名单','tables':[table]})
        self.assertTrue(any(r['cell']=='B9' and r['value']==90 for r in plan['updates']))
        self.assertTrue(any(r['cell']=='A9' and r['value']=='新增项目' for r in plan['labels']))
        self.assertFalse(any(r['cell'] in ('B2','C2') for r in plan['labels']))
        matched=next(r for r in plan['identity_matches'] if r['cell']=='B9')
        self.assertNotIn('period',matched.get('dimension_evidence',{}))
        self.assertTrue(any(r.get('cell')=='B9' and r.get('action')=='label' for r in plan['issues']))
        self.assertEqual(b,before)
        self.assertEqual(len(plan['operations']),1)

    def test_已明确金额来源不自行补出名单动作(self):
        a=self.save('A');answer=json.loads(self.case.answer())
        row=next(r for r in answer['tables'][0]['dimensions'] if r['id']=='asset_selection')
        row['value']['members']=['新名单成员'];row['evidence']=[];row.pop('label_references',None)
        b=self.save('B',json.dumps(answer,ensure_ascii=False))
        purpose={'reason':'用户已明确的名单来源安排',
                 'bindings':[{'target':'合成表!B3','source':'合成表!B3'}]}
        plan=direct_update_plan(a,b,purpose)
        self.assertTrue(any(r['cell']=='B3' for r in plan['identity_matches']))
        update=next(r for r in plan['updates'] if r['cell']=='B3')
        self.assertEqual(update['source_identity']['dimensions']['asset_selection']['members'],['新名单成员'])
        self.assertFalse(any(r['cell']=='A3' for r in plan['labels']))
        self.assertFalse(any(r.get('cell')=='B3' and r.get('action')=='label' for r in plan['issues']))

    def test_明确名单标签搬运不受金额身份缺项撤销(self):
        a=self.save('A');answer=json.loads(self.case.answer())
        row=next(r for r in answer['tables'][0]['dimensions'] if r['id']=='asset_selection')
        row['value']['members']=['新名单成员']
        snapshot=deepcopy(self.snapshot);snapshot['sheets'][0]['cells']['A3']['value']='新名单成员'
        b=self.save('B',json.dumps(answer,ensure_ascii=False),snapshot)
        purpose={'reason':'合成已明确名单随B更新',
                 'bindings':[{'target':f'合成表!{c}3','source':f'合成表!{c}3'} for c in 'BC'],
                 'labels':[{'target':'合成表!A3','source':'合成表!A3'}]}
        with patch('openpyxl.load_workbook',side_effect=AssertionError('不读原表补身份')):
            plan=direct_update_plan(a,b,purpose)
        self.assertTrue(any(r['cell']=='A3' and r['value']=='新名单成员' for r in plan['labels']))
        unknown=self.corrected(a,b,'合成表!B3:C3：asset_selection 为 unknown。')
        plan=direct_update_plan(a,unknown,purpose)
        self.assertTrue(any(r['cell']=='A3' and r['value']=='新名单成员' for r in plan['labels']))
        self.assertTrue(any(r['cell']=='B4' for r in plan['identity_matches']))

    def test_正常匹配入口承接已有名单的新增安排(self):
        snapshot=deepcopy(self.snapshot)
        snapshot['names']={'名单表':[{'sheet':'合成表','range':'A2:C8'}]}
        a=self.save('A',snapshot=snapshot)
        snapshot['sheets'][0]['max_row']=9
        for col,value in [('A','新增项目'),('B',90),('C',91)]:
            snapshot['sheets'][0]['cells'][col+'9']={**snapshot['sheets'][0]['cells'][col+'3'],'value':value}
        answer=json.loads(self.case.answer());table=answer['tables'][0]
        table.update(range='A2:C9',data_ranges=['B3:C9'])
        for row in table['metrics']:row['ranges']=[span.replace('8','9') for span in row['ranges']]
        for row in table['dimensions']:
            if row['id']=='period':row['ranges']=[span.replace('8','9') for span in row['ranges']]
        table['dimensions'].append({'id':'asset_selection','ranges':['B9:C9'],
            'value':{'domain':'固定资产','mode':'members','members':['新增项目']},'evidence':['合成表!A9']})
        b=self.save('B',json.dumps(answer,ensure_ascii=False),snapshot)
        purpose={'reason':'合成已明确完整名单随B更新','tables':[{'range_name':'名单表','source_sheet':'合成表',
            'source_rows':list(range(3,10)),'old_detail_count':6,'first_row_offset':1,
            'columns':{'B':'B','C':'C'},'labels':{'A':'A'}}]}
        with patch.object(UnifiedSemanticEngine,'_http',side_effect=AssertionError('不调用模型')):
            result=run_step('match',{**self.case.config,'a_mapping':a['mapping_path'],
                'b_mapping':b['mapping_path'],'update_purpose':purpose})
        plan=json.loads(__import__('pathlib').Path(result['files']['update_plan']).read_text(encoding='utf-8-sig'))
        self.assertTrue(plan['operations'],plan['issues'])
        self.assertTrue(any(row['cell']=='B9' and row['value']==90 for row in plan['updates']))
        self.assertTrue(any(row['cell']=='A9' and row['value']=='新增项目' for row in plan['labels']))

    def test_自然名单安排进入主入口按成员取异行金额并保留新增空白(self):
        from 附注更新.A合集 import save_list_relations
        from 附注更新.统一勾稽 import load_fact_mapping
        snapshot=deepcopy(self.snapshot);snapshot['names']={'名单表':[{'sheet':'合成表','range':'A2:C8'}]}
        a=self.save('A',snapshot=snapshot)
        snapshot['sheets'][0]['max_row']=9
        for col,value in [('A','新增项目'),('B',None),('C',None)]:
            snapshot['sheets'][0]['cells'][col+'9']={**snapshot['sheets'][0]['cells'][col+'3'],'value':value}
        answer=json.loads(self.case.answer());table=answer['tables'][0]
        table.update(range='A2:C9',data_ranges=['B3:C9'])
        for row in table['metrics']:row['ranges']=[span.replace('8','9') for span in row['ranges']]
        for row in table['dimensions'][:]:
            if row['id']=='period':row['ranges']=[span.replace('8','9') for span in row['ranges']]
            if row['id']=='asset_selection' and row['ranges'] in (['B3:C3'],['B4:C4']):
                number=3 if row['ranges']==['B3:C3'] else 4
                row['ranges']=[f'C{number}']
                table['dimensions'].append({**deepcopy(row),'ranges':[f'B{7-number}'],'evidence':[]})
        table['dimensions'].append({'id':'asset_selection','ranges':['B9:C9'],
            'value':{'domain':'固定资产','mode':'members','members':['新增项目']},'evidence':['合成表!A9']})
        b=self.save('B',json.dumps(answer,ensure_ascii=False),snapshot)
        reply=('B原表：合成表!A2:C9，名称在A列，B、C列分别为期末、期初。\n'
               '- 合成表!A3:A9：7 个名称。\n'
               '沿用 A 的 名单表 表区和列角色：A列放名单标签，B、C列继续分别承载期末、期初；'
               'B的7个具名位置接入为连续明细。A当前有6个明细行，需增加1行；从A3起连续排布。')
        saved=save_list_relations(a,b,reply,'asset_selection',self.case.folder/'自然名单',
                                  source_ranges=[{'sheet':'合成表','range':'A2:C9'}])
        both={key:load_fact_mapping(saved['files'][key+'_mapping'],{}) for key in ('a','b')}
        self.assertEqual(both['b']['facts'],b['facts'])
        purpose={'reason':'用户明确要求名单随B'}
        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('不调用模型')):
            result=run_step('match',{**self.case.config,**{k+'_mapping':v['mapping_path'] for k,v in both.items()},'update_purpose':purpose})
        plan=json.loads(Path(result['files']['update_plan']).read_text(encoding='utf-8-sig'))
        self.assertEqual(len(plan['operations']),1,plan['issues'])
        self.assertEqual(next(r for r in plan['updates'] if r['cell']=='B3')['sources'][0]['cell'],'B4')
        self.assertTrue(any(r['cell']=='A9' and r['value']=='新增项目' for r in plan['labels']))
        self.assertTrue({'B9','C9'}<={r['cell'] for r in plan['retained_blanks']})
        self.assertTrue({'B9','C9'}<={r['cell'] for r in plan['identity_matches']})
        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('不调用模型')):
            single=run_step('match',{**self.case.config,'a_mapping':a['mapping_path'],
                                    'b_mapping':both['b']['mapping_path'],'update_purpose':purpose})
        self.assertEqual(single['status'],'complete',single)
        single_plan=json.loads(Path(single['files']['update_plan']).read_text(encoding='utf-8-sig'))
        for key in ('updates','labels','retained_blanks','identity_matches','operations'):
            self.assertEqual(single_plan[key],plan[key])
        from 附注更新.分步流程 import _write_direct
        with patch('附注更新.分步流程.write_updated_workbook',side_effect=RuntimeError('只到写回入口')) as write:
            with self.assertRaisesRegex(RuntimeError,'只到写回入口'):
                _write_direct(self.case.config,self.case.folder,lambda _:None,None,
                              Path(result['files']['update_plan']),plan,both['a'],both['b'])
            self.assertEqual(write.call_args.args[2]['operations'],plan['operations'])
        unknown=save_direct_answer(both['b'],'合成表!B4：asset_selection 为 unknown。',self.case.folder/'名单来源未知.json',
                                  field_scope={'合成表!B4':['asset_selection']})
        stopped=direct_update_plan(both['a'],unknown,purpose)
        self.assertEqual(stopped['operations'],plan['operations'])
        self.assertTrue(any('名单标签尚无对应槽位及期间的已识别来源' in str(r) for r in stopped['issues']))
        self.assertTrue(any(r['cell']=='A9' and r['value']=='新增项目' for r in stopped['labels']))

    def test_名单按已识别身份跨表区列位取数并保留多来源(self):
        from 附注更新.A合集 import save_list_relations
        from 附注更新.统一勾稽 import load_fact_mapping
        snapshot=deepcopy(self.snapshot);snapshot['names']={'名单表':[{'sheet':'合成表','range':'A2:C8'}]}
        a=self.save('A',snapshot=snapshot)
        sheet=snapshot['sheets'][0];sheet['max_column']=6
        sheet['cells']['A3']['value']='新名单成员'
        for old,new,column in [('B','E',5),('C','F',6)]:
            sheet['cells'][new+'2']={**sheet['cells'][old+'2'],'column':column}
        for row in range(3,9):
            for old,new,column in [('B','E',5),('C','F',6)]:
                sheet['cells'][new+str(row)]={**sheet['cells'][old+str(row)],'column':column}
        original=json.loads(self.case.answer())['tables'][0];original['range']='A2:F8'
        for row in original['dimensions']:
            if row['id']=='asset_selection' and row['ranges']==['B3:C3']:
                row['value']['members']=['新名单成员']
        shifted=deepcopy(original);shifted.update(id='补充来源',range='A2:F8',data_ranges=['E3:F8'])
        for row in shifted['metrics']+shifted['dimensions']:
            row['ranges']=[span.translate(str.maketrans({'B':'E','C':'F'})) for span in row['ranges']]
            if row['id']=='period':
                row['evidence']=[ref.translate(str.maketrans({'B':'E','C':'F'})) for ref in row['evidence']]
        reply=('B原表：合成表!A2:C8，名称在A列，B、C列分别为期末、期初。\n'
               '- 合成表!A3:A8：6 个名称。\n'
               '沿用 A 的 名单表 表区和列角色：A列放名单标签，B、C列继续分别承载期末、期初；'
               'B的6个具名位置接入为连续明细。A当前有6个明细行；从A3起连续排布。')
        for duplicate in (False,True):
            with self.subTest(duplicate=duplicate):
                answer={'tables':([original] if duplicate else [])+[deepcopy(shifted)]}
                if duplicate:
                    for row in answer['tables'][-1]['dimensions']:
                        if row['id']=='asset_selection' and row['ranges']==['E3:F3']:row['evidence']=[]
                b=self.save('B',json.dumps(answer,ensure_ascii=False),snapshot)
                saved=save_list_relations(a,b,reply,'asset_selection',self.case.folder/f'跨列名单{duplicate}',
                                          source_ranges=[{'sheet':'合成表','range':'A2:C8'}])
                both={key:load_fact_mapping(saved['files'][key+'_mapping'],{}) for key in ('a','b')}
                purpose={'reason':'用户明确要求名单随B，复用本次已识别来源'}
                plan=direct_update_plan(both['a'],both['b'],purpose)
                expected={'B3','E3'} if duplicate else {'E3'}
                match=next(row for row in plan['identity_matches'] if row['cell']=='B3')
                self.assertEqual(set(match['source_locations']),{'合成表!'+cell for cell in expected})
                self.assertEqual(any(row['cell']=='B3' for row in plan['updates']),not duplicate)
                self.assertTrue(any(row['cell']=='A3' and row['value']=='新名单成员' for row in plan['labels']))
                self.assertEqual(both['b']['facts'],b['facts'])
                if not duplicate:continue
                scope={'target':[{'sheet':'合成表','range':'B3:C8'}],
                       'source':[{'sheet':'合成表','range':'B3:C8'}]}
                limited=direct_update_plan(both['a'],both['b'],purpose,scope)
                update=next(row for row in limited['updates'] if row['cell']=='B3')
                self.assertEqual({row['cell'] for row in update['sources']},{'B3'})
                conflict=deepcopy(both['b']);conflict['机械记录']['sheets'][0]['cells']['E3']['value']=999
                conflict['封存校验']=direct_digest(conflict)
                continued=direct_update_plan(both['a'],conflict,purpose)
                self.assertFalse(any(row['cell']=='B3' for row in continued['updates']))
                self.assertTrue(any(row.get('cell')=='B3' and row['kind']=='ambiguous_source' for row in continued['issues']))
                specified={**purpose,'bindings':[{'target':'合成表!B3','source':'合成表!E3'}]}
                chosen=direct_update_plan(both['a'],conflict,specified)
                update=next(row for row in chosen['updates'] if row['cell']=='B3')
                self.assertEqual(update['value'],999)
                self.assertEqual({row['cell'] for row in update['sources']},{'E3'})
                for source in ['合成表!E4','合成表!F3','合成表!E99']:
                    request={**purpose,'bindings':[{'target':'合成表!B3','source':source}]}
                    directed=direct_update_plan(both['a'],both['b'],request)
                    if source=='合成表!E99':
                        self.assertFalse(any(row['cell']=='B3' for row in directed['updates']))
                    else:
                        update=next(row for row in directed['updates'] if row['cell']=='B3')
                        self.assertEqual([row['fact_id'] for row in update['sources']],[source])
                unknown=save_direct_answer(both['b'],'合成表!E3：asset_selection 为 unknown。',
                    self.case.folder/'指定来源身份未知.json',field_scope={'合成表!E3':['asset_selection']})
                directed=direct_update_plan(both['a'],unknown,specified)
                update=next(row for row in directed['updates'] if row['cell']=='B3')
                self.assertEqual(update['source_identity']['dimensions']['asset_selection']['mode'],'unknown')
                totals=deepcopy(both['b']);totals['机械记录']['sheets'][0]['cells']['E8']['value']=888
                totals['封存校验']=direct_digest(totals)
                request={'reason':'已有明确表区合计来源选择',
                    'bindings':[{'target':'合成表!B8','source':'合成表!E8'}],
                    'tables':[{'range_name':'名单表','source_sheet':'合成表','source_rows':list(range(3,8)),
                        'old_detail_count':5,'first_row_offset':1,'columns':{'B':'B','C':'C'},'labels':{'A':'A'},
                        'totals':[{'target_row':8,'source_row':8}]}]}
                chosen=direct_update_plan(both['a'],totals,request)
                self.assertEqual(next(row for row in chosen['updates'] if row['cell']=='B8')['value'],888)

    def test_名单行已换成员时缺项不回填旧成员金额(self):
        snapshot=deepcopy(self.snapshot);snapshot['names']={'名单表':[{'sheet':'合成表','range':'A2:C8'}]}
        a=self.save('A',snapshot=snapshot)
        snapshot['sheets'][0]['cells']['A3']['value']='新成员'
        answer=json.loads(self.case.answer());dims=answer['tables'][0]['dimensions']
        old=next(row for row in dims if row['id']=='asset_selection' and row['ranges']==['B3:C3'])
        old['ranges']=['C3'];dims.append({**deepcopy(old),'ranges':['B3'],
            'value':{'mode':'members','domain':'固定资产','members':['新成员']}})
        b=self.save('B',json.dumps(answer,ensure_ascii=False),snapshot)
        purpose={'reason':'沿已明确新成员名单安排','tables':[{
            'selection_dimension':'asset_selection','range_name':'名单表','source_sheet':'合成表',
            'source_rows':[3,4,5,6,7,8],'old_detail_count':6,'first_row_offset':1,
            'columns':{'B':None,'C':None},'labels':{'A':'A'},
            'source_ranges':[{'sheet':'合成表','range':'A2:C8'}]}]}
        # 原答复只将新成员的本期金额指向新名称，旧成员的上期金额保留其自己的引用。
        for row in b['facts']:
            if row['id']=='合成表!C3':row['dimension_evidence']['asset_selection']=[]
        b['封存校验']=direct_digest(b)
        plan=direct_update_plan(a,b,purpose)
        matched={row['cell'] for row in plan['identity_matches']}
        self.assertTrue({'B3','B4','C4'}<=matched);self.assertNotIn('C3',matched)
        self.assertTrue(any(row['cell']=='A3' and row['value']=='新成员' for row in plan['labels']))

    def test_已知名单名称不因其他字段缺失撤销也不搬回旧成员金额(self):
        a,b,_=self.two_lists()
        b=save_direct_answer(b,'合成表!B9:C9：unit_scale=unknown。',self.case.folder/'新成员其他字段未知.json',
                             field_scope={'合成表!B9:C9':['unit_scale']})
        purpose={'reason':'按已明确名单替换名称','tables':[{
            'range_name':'第一表','source_sheet':'合成表','source_rows':[9],
            'old_detail_count':1,'first_row_offset':1,'columns':{'B':None,'C':None},'labels':{'A':'A'},
            'selection_dimension':'asset_selection','source_ranges':[{'sheet':'合成表','range':'A2:C9'}]}]}
        before=deepcopy((a,b));plan=direct_update_plan(a,b,purpose)
        self.assertEqual((a,b),before)
        self.assertTrue(any(row['cell']=='A3' and row['value']=='新增项目' for row in plan['labels']))
        self.assertFalse(any(row['cell'] in ('B3','C3') for row in plan['identity_matches']))
        self.assertTrue(any(row['cell']=='B4' for row in plan['identity_matches']))
        self.assertEqual(dict(_records(b))['合成表!B9']['dimensions']['asset_selection']['members'],['新增项目'])

    def test_同一名单成员的多个计量项目按模板槽位分别对应(self):
        rows=fixture.records();metric=deepcopy(next(r for r in rows if r['record_type']=='metric'))
        metric['id']='累计折旧';metric['definition']['measurement_basis']='累计折旧';rows.append(metric)
        self.case.gold=self.case.folder/'分段合成金标准.jsonl'
        self.case.gold.write_text('\n'.join(json.dumps(r,ensure_ascii=False) for r in rows),encoding='utf-8')
        self.engine=self.case.engine()
        self.base.update(gold_path=str(self.case.gold),gold_hash=self.engine.gold_hash,
                         金标准原文='\n'.join(json.dumps(r,ensure_ascii=False) for r in rows
                            if r['record_type'] in {'metric','dimension'}))
        snapshot=deepcopy(self.snapshot);snapshot['names']={'名单表':[{'sheet':'合成表','range':'A2:C8'}]}
        a=self.save('A',snapshot=snapshot);snapshot['sheets'][0]['max_column']=6
        original=json.loads(self.case.answer())['tables'][0]
        second=deepcopy(original);second.update(id='累计折旧表',range='A2:F8',data_ranges=['E3:F8'])
        for row in second['metrics']+second['dimensions']:
            row['ranges']=[span.translate(str.maketrans({'B':'E','C':'F'})) for span in row['ranges']]
            if row['id']=='固定资产原值':row['id']='累计折旧'
        for old,new in [('B','E'),('C','F')]:
            for row in range(3,9):
                snapshot['sheets'][0]['cells'][new+str(row)]={**snapshot['sheets'][0]['cells'][old+str(row)],'value':999}
        b=self.save('B',json.dumps({'tables':[original,second]},ensure_ascii=False),snapshot)
        purpose={'reason':'沿用已识别名单和模板计量项目','tables':[{
            'selection_dimension':'asset_selection','range_name':'名单表','source_sheet':'合成表',
            'source_rows':list(range(3,9)),'old_detail_count':6,'first_row_offset':1,
            'columns':{'B':'B','C':'C'},'labels':{'A':'A'},
            'source_ranges':[{'sheet':'合成表','range':'A2:F8'}]}]}
        plan=direct_update_plan(a,b,purpose)
        self.assertFalse(plan['issues'])
        for update in plan['identity_matches']:
            self.assertEqual(update['metric_id'],'固定资产原值')
            self.assertEqual(update['source_locations'],['合成表!'+update['cell']])
        self.assertTrue(plan['identity_matches'])
        self.assertTrue(any(row['metric_id']=='累计折旧' for row in plan['unplaced_sources']))

    def test_分段名单保存回读并按原位置顺延且不补金额身份(self):
        from 附注更新.A合集 import save_list_relations
        from 附注更新.统一勾稽 import load_fact_mapping
        snapshot=deepcopy(self.snapshot);snapshot['names']={'分段表':[{'sheet':'合成表','range':'A2:C8'}]}
        answer=json.loads(self.case.answer())
        for row in answer['tables'][0]['dimensions']:
            if row['id']!='asset_selection':continue
            if row['ranges'] in (['B3:C3'],['B6:C6']):row['value']={'domain':'固定资产','mode':'all'}
            for earlier,later in [(4,7),(5,8)]:
                if row['ranges'] in ([f'B{earlier}:C{earlier}'],[f'B{later}:C{later}']):
                    row['value']['members']=['测试项目'+str(earlier)]
                    row['evidence']=[f'合成表!A{earlier}',f'合成表!A{later}']
        a=self.save('A',json.dumps(answer,ensure_ascii=False),snapshot);b=self.save('B')
        reply=('A，合成表!A2:C8（分段表）：每组先列合计、再列“其中”明细：第一计量3/4–5，第二计量6/7–8。\n'
               'B，合成表!A2:C8：每段标题在A列，明细标签在A列。\n'
               '其中原值3–5显示这些名称。\n'
               '应保留A2:C8的列角色：项目、期末余额、期初余额；将每组两个明细项改为B有据的三个标签。')
        saved=save_list_relations(a,b,reply,'asset_selection',self.case.folder/'分段名单',
                                  source_ranges=[{'sheet':'合成表','range':'A2:C8'}])
        both={key:load_fact_mapping(saved['files'][key+'_mapping'],{}) for key in ('a','b')}
        for key,old in [('a',a),('b',b)]:
            for field in ('facts','slot_bindings','机械记录','unresolved'):
                self.assertEqual(both[key][field],old[field])
        tables=both['a']['已确认关系'][-1]['tables']
        self.assertEqual([t['first_row_offset'] for t in tables],[2,5])
        automatic=direct_update_plan(both['a'],both['b'])
        self.assertEqual([op['index'] for op in automatic['operations']],[7,4])
        self.assertTrue(any(row['cell']=='A6' for row in automatic['labels']))
        with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('不调用模型')):
            result=run_step('match',{**self.case.config,'a_mapping':saved['files']['a_mapping'],
                'b_mapping':saved['files']['b_mapping'],'update_purpose':{'reason':'按已有答复保留分段并更新名单'}})
        plan=json.loads(Path(result['files']['update_plan']).read_text(encoding='utf-8-sig'))
        self.assertEqual([op['index'] for op in plan['operations']],[7,4])
        self.assertEqual(plan['template_layout']['positions']['合成表!B6'],'合成表!B7')
        self.assertEqual(plan['template_layout']['positions']['合成表!B8'],'合成表!B9')
        for target,source in [('B4','B3'),('B5','B4'),('B6','B5'),('B8','B3'),('B9','B4'),('B10','B5')]:
            row=next(r for r in plan['identity_matches'] if r['cell']==target)
            self.assertEqual(row['source_locations'],['合成表!'+source])
        for target,source in [('A4','A3'),('A6','A5'),('A8','A3'),('A10','A5')]:
            self.assertTrue(any(r['cell']==target and r['source_cell']==source for r in plan['labels']))
        self.assertTrue(all(r['old_value'] is None for r in plan['labels'] if r['cell'] in ('A6','A10')))
        self.assertTrue({'B6','C6','B10','C10'}<={r['cell'] for r in plan['retained_blanks']})
        from 附注更新.分步流程 import _write_direct
        with patch('附注更新.分步流程.write_updated_workbook',side_effect=RuntimeError('已到副本入口')) as write:
            with self.assertRaisesRegex(RuntimeError,'已到副本入口'):
                _write_direct(self.case.config,self.case.folder,lambda _:None,None,
                    Path(result['files']['update_plan']),plan,both['a'],both['b'])
            self.assertEqual(write.call_args.args[2]['operations'],plan['operations'])
        unknown=save_direct_answer(both['b'],'合成表!B5:C5：asset_selection 为 unknown。',
            self.case.folder/'分段成员仍未知.json',field_scope={'合成表!B5:C5':['asset_selection']})
        partial=direct_update_plan(both['a'],unknown,{'reason':'按已有答复保留分段并更新名单'})
        self.assertEqual(partial['operations'],plan['operations'])
        self.assertFalse(any(r['cell'] in ('B6','C6','B10','C10') for r in partial['identity_matches']))
        self.assertTrue({'A6','A10'}<={r['cell'] for r in partial['labels']})
        self.assertTrue(any(r['cell']=='B4' for r in partial['identity_matches']))

    def test_名单局部未知不撤销已确认明细合计及期间标签(self):
        snapshot=deepcopy(self.snapshot)
        snapshot['names']={'名单表':[{'sheet':'合成表','range':'A2:C8'}]}
        for column in 'BC':snapshot['sheets'][0]['cells'][column+'8']['value']=80
        answer=json.loads(self.case.answer())
        total=next(r for r in answer['tables'][0]['dimensions'] if r['id']=='asset_selection' and r['ranges']==['B8:C8'])
        total['value']={'domain':'固定资产','mode':'all'}
        a=self.save('A',json.dumps(answer,ensure_ascii=False),snapshot)
        purpose={'reason':'名单随B，独立合计按已有身份取值','tables':[{
            'range_name':'名单表','source_sheet':'合成表','source_rows':[3,4,5,6,7],
            'old_detail_count':4,'first_row_offset':1,'columns':{'B':'B','C':'C'},'labels':{'A':'A'}}]}
        for year in (2025,2026):
            fresh=json.loads(self.case.answer(year))
            next(r for r in fresh['tables'][0]['dimensions'] if r['id']=='asset_selection' and r['ranges']==['B8:C8'])['value']=deepcopy(total['value'])
            source=deepcopy(snapshot)
            if year==2026:
                source['sheets'][0]['cells']['B2']['value']='2026年末'
                source['sheets'][0]['cells']['C2']['value']='2025年末'
            b=self.save('B',json.dumps(fresh,ensure_ascii=False),source)
            b=save_direct_answer(b,'合成表!B4：asset_selection 为 unknown。',self.case.folder/f'名单待定{year}.json',
                                field_scope={'合成表!B4':['asset_selection']})
            with patch.object(UnifiedSemanticEngine,'request_text',side_effect=AssertionError('不调用模型')):
                result=run_step('match',{**self.case.config,'a_mapping':a['mapping_path'],
                    'b_mapping':b['mapping_path'],'update_purpose':purpose})
            plan=json.loads(Path(result['files']['update_plan']).read_text(encoding='utf-8-sig'))
            self.assertEqual(len(plan['operations']),1)
            self.assertTrue({'B3','C3','C4'}<={r['cell'] for r in plan['updates']})
            self.assertEqual(next(r for r in plan['identity_matches'] if r['cell']=='B4')['source_locations'],['合成表!B4'])
            totals={r['cell'] for r in plan['updates']} & {'B9','C9'}
            self.assertEqual(totals,{'B9','C9'})
            label=next(r for r in plan['labels'] if r['cell']=='A7')
            self.assertEqual(label['value'],'测试项目7');self.assertIsNone(label['old_value'])
            if year==2026:self.assertEqual({r['cell'] for r in plan['labels']},{'B2','C2','A3','A4','A5','A6','A7'})

    def test_编排执行记录保留安排及实际值而不生成身份(self):
        from 附注更新.A合集 import direct_output_record
        from 附注更新.统一勾稽 import load_fact_mapping
        from openpyxl.utils.cell import coordinate_to_tuple
        a,b,_=self.two_lists(2026)
        a['unresolved'].append({'sheet':'合成表','cell':'B4','reason':'原答复保留的未决'})
        a['excluded'].append({'sheet':'合成表','cell':'A4','range':'A4:A5','category':'annotation',
                              'reason':'已有说明位置','evidence':['合成表!A2']})
        a['outline']['subtables'][0]['metrics'][0]['template_period_role']={'B3:B8':'closing'}
        a['封存校验']=direct_digest(a)
        b['unresolved'].append({'sheet':'合成表','cell':'B6','reason':'B来源待定'})
        b['封存校验']=direct_digest(b)
        purpose={'reason':'明确插入样式行，未给新行任何金额身份','operations':[
            {'range_name':'第一表','axis':'row','action':'insert','index':0,'count':1,'template_index':1}]}
        plan=direct_update_plan(a,b,purpose)
        self.assertTrue(any(r.get('carrier')=='A' and r.get('cell')=='B5' and r.get('reason')=='原答复保留的未决' for r in plan['issues']))
        self.assertTrue(any(r.get('carrier')=='B' and r.get('cell')=='B6' and r.get('reason')=='B来源待定' for r in plan['issues']))
        self.assertFalse(any(r['cell'] in ('B2','C2') for r in plan['identity_matches']))
        snapshot=deepcopy(a['机械记录']);snapshot['sha256']='d'*64;cells={}
        for cell,value in snapshot['sheets'][0]['cells'].items():
            dest=plan['template_layout']['positions']['合成表!'+cell].rsplit('!',1)[1]
            row,column=coordinate_to_tuple(dest);cells[dest]={**value,'row':row,'column':column}
        for item in plan['updates']+plan['labels']:
            row,column=coordinate_to_tuple(item['cell'])
            cells[item['cell']]={**cells.get(item['cell'],{}),'value':item['value'],'formula':None,'row':row,'column':column}
        snapshot['sheets'][0]['cells']=cells;snapshot['sheets'][0]['max_row']=10
        with patch('附注更新.表格.read_workbook',return_value=snapshot),patch('附注更新.表格.file_hash',return_value=snapshot['sha256']):
            derived=direct_output_record(a,b,self.case.folder/'未实际写出的副本.xlsx',plan,'离线安排')
        path=self.case.folder/'副本合集离线回读.json';path.write_text(json.dumps(derived,ensure_ascii=False),encoding='utf-8')
        loaded=json.loads(path.read_text(encoding='utf-8'))
        self.assertEqual(loaded['template_layout'],plan['template_layout'])
        self.assertEqual(loaded['identity_matches'],plan['identity_matches'])
        self.assertEqual(loaded['issues'],plan['issues'])
        saved=next(row for row in loaded['updates'] if row['cell']=='B4')
        self.assertEqual(saved['actual_value'],snapshot['sheets'][0]['cells']['B4']['value'])
        self.assertEqual(saved['source_identity']['dimensions'],dict(_records(b))['合成表!B3']['dimensions'])
        self.assertNotIn('facts',loaded)
        self.assertNotIn('识别答复',loaded)

    def test_标签位于声明使用精确位置而不展开整个表头(self):
        a=self.save('A');before=dict(_records(a))
        reply='B2:C2（标签位于 B2）为“期末余额”，适用于 B3:B8，角色为期末。'
        saved=save_direct_answer(a,reply,self.case.folder/'精确标签位置.json',
                                field_scope={'合成表!B3:B8':['period']})
        after=dict(_records(saved))
        self.assertEqual(set(after),set(before));self.assertEqual(saved['机械记录'],a['机械记录'])
        for loc,row in before.items():
            self.assertEqual(after[loc]['dimensions'],row['dimensions'])
            if loc.rsplit('!',1)[1].startswith('B'):
                self.assertEqual(after[loc]['dimension_evidence']['period'],['合成表!B2'])
            else:self.assertEqual({k:v for k,v in after[loc].items() if k!='recognition_provenance'},
                                  {k:v for k,v in row.items() if k!='recognition_provenance'})

    def test_并列列及列区间沿用明示期间不漏列(self):
        from 附注更新.A语义阅读 import addresses
        reply=('共同取值：期末余额的 period 为 {date: 2024-12-31, kind: instant}；'
               '期初余额的 period 为 {date: 2024-01-01, kind: instant}。\n'
               '合成表!B3:E4：B、C为期末原值、准备，D、E为期初原值、准备。\n'
               '合成表!B5:G5：期末列为B:D，期初列为E:G。\n'
               '合成表!B6:C6：B、C列分别为期末、期初。')
        parsed=direct_answer(reply,self.engine.standard,['合成表'])
        actual={}
        for row in parsed['dimensions']:
            if row['id'] not in ('period','template_period_role'):continue
            for area in row['ranges']:
                if area=='*':continue
                for cell in addresses(area.rsplit('!',1)[-1]):
                    actual.setdefault(cell,{})[row['id']]=row['value']
        closing=set(addresses('B3:C4')+addresses('B5:D5')+['B6'])
        opening=set(addresses('D3:E4')+addresses('E5:G5')+['C6'])
        self.assertEqual(set(actual),closing|opening)
        for cells,role,date in [(closing,'closing','2024-12-31'),(opening,'opening','2024-01-01')]:
            for cell in cells:
                self.assertEqual(actual[cell],{'period':{'kind':'instant','date':date},'template_period_role':role})
        before=self.save('B')
        saved=save_direct_answer(before,reply,self.case.folder/'并列期间保存.json',
            field_scope={'合成表!B3:C3':['period','template_period_role']})
        self.assertEqual(saved['机械记录'],before['机械记录'])
        for loc,row in _records(saved):
            if loc in ('合成表!B3','合成表!C3'):
                self.assertEqual(row['dimensions']['period'],{'kind':'instant','date':'2024-12-31'})
            else:
                for key in ('metric_id','dimensions','dimension_evidence','raw_value','raw_value_state'):
                    self.assertEqual(row.get(key),dict(_records(before))[loc].get(key))
        roles={cell:row.get('template_period_role') for table in saved['outline']['subtables']
               for row in table['metrics'] for area in row['ranges'] for cell in addresses(area)}
        self.assertEqual([roles['B3'],roles['C3']],['closing','closing'])

    def test_段落行范围内按明示顺序保存列期间角色(self):
        from 附注更新.A语义阅读 import addresses
        reply=('合成表第3—4行：B、C列为固定资产原值，期间分别为期初、期末。\n'
               '合成表第6—7行：B、C列为固定资产原值，期间分别为期末、期初。')
        original=self.save('B')
        saved=save_direct_answer(original,reply,self.case.folder/'段落期间角色.json',
            field_scope={area:['template_period_role'] for area in ('合成表!B3:C4','合成表!B6:C7')})
        self.assertEqual(saved['机械记录'],original['机械记录'])
        for loc,row in _records(saved):
            before=dict(_records(original))[loc]
            for key in ('metric_id','dimensions','dimension_evidence','raw_value','raw_value_state'):
                self.assertEqual(row.get(key),before.get(key))
        roles=lambda mapping:{cell:row.get('template_period_role') for table in mapping['outline']['subtables']
            for row in table['metrics'] for area in row['ranges'] for cell in addresses(area)}
        expected=roles(original)
        expected.update({cell:role for span,role in [('B3:B4','opening'),('C3:C4','closing'),
            ('B6:B7','closing'),('C6:C7','opening')] for cell in addresses(span)})
        self.assertEqual(roles(saved),expected)
        grouped=direct_answer('合成表第3—4行：B/D列期间为期末，C、E列为期初。',
            self.engine.standard,['合成表'])
        parsed={cell:row['value'] for row in grouped['dimensions'] if row['id']=='template_period_role'
                for cell in row['ranges']}
        self.assertEqual(parsed,{'合成表!'+col+str(row):role for cols,role in [('BD','closing'),('CE','opening')]
                                 for col in cols for row in (3,4)})
        ambiguous=direct_answer('合成表第3—4行：B、C列为固定资产原值，期间分别为期初、期末、本期。',
            self.engine.standard,['合成表'])
        self.assertFalse(any(r['id']=='template_period_role' for r in ambiguous['dimensions']))

    def test_顺序期间引用同段明确日期并保存角色(self):
        reply=('共同说明的期末为2026-12-31、期初为2026-01-01。\n'
               '合成表!B3/C3，合计：期间分别为期末、期初时点。')
        original=self.save('B')
        saved=save_direct_answer(original,reply,self.case.folder/'顺序期间引用.json',
            field_scope={'合成表!B3:C3':['period','template_period_role']})
        self.assertEqual(saved['机械记录'],original['机械记录'])
        rows=dict(_records(saved))
        for cell,date in [('B3','2026-12-31'),('C3','2026-01-01')]:
            self.assertEqual(rows['合成表!'+cell]['dimensions']['period'],{'kind':'instant','date':date})
        roles={cell:row.get('template_period_role') for t in saved['outline']['subtables']
               for row in t['metrics'] for cell in row['ranges']}
        self.assertEqual([roles['B3'],roles['C3']],['closing','opening'])
        dated=direct_answer('合成表!B3/C3：期间分别为期末时点2026-12-31、期初时点2026-01-01。',
            self.engine.standard,['合成表'])
        self.assertEqual([r['value'] for r in dated['dimensions'] if r['id']=='period'],['2026-12-31','2026-01-01'])
        missing=direct_answer(reply.split('\n')[1],self.engine.standard,['合成表'])
        self.assertEqual([r['value'] for r in missing['dimensions'] if r['id']=='period'],['期末','期初时点'])

    def test_原表期间标签位于只引用明示目标(self):
        answer='合成表!B3、B4：原表期间标签位于 合成表!B2，为“本期增加”；期间同为 2026-01-01 至 2026-12-31。'
        parsed=direct_answer(answer,self.engine.standard,['合成表'])
        periods=[r for r in parsed['dimensions'] if r['id']=='period']
        self.assertEqual({c for r in periods for c in r['ranges']},{'合成表!B3','合成表!B4'})
        self.assertTrue(all(r.get('evidence')==['合成表!B2'] for r in periods))
        self.assertFalse(parsed['unresolved'])
        b=self.save('B')
        balance='合成表!B3、B4：原表期间标签位于 合成表!B2，为“期末余额”；期间为 2026-12-31。'
        saved=save_direct_answer(b,balance,self.case.folder/'期间引用.json',
            field_scope={'合成表!B3:B4':['period','template_period_role']})
        for cell in ['B3','B4']:
            row=dict(_records(saved))['合成表!'+cell]
            self.assertEqual(row['dimension_evidence']['period'],['合成表!B2'])
            self.assertEqual(row['dimensions']['period'],{'kind':'instant','date':'2026-12-31'})

    def test_原表期间标签下一句保留前句明示目标(self):
        b=self.save('B')
        answer='合成表!B3、B4：原有明细说明。原表期间标签坐标：合成表!B2“期末余额”。'
        saved=save_direct_answer(b,answer,self.case.folder/'句间标签.json',
                                 field_scope={'合成表!B3:B4':['period']})
        before=dict(_records(b))
        self.assertEqual(saved['机械记录'],b['机械记录'])
        for location,row in _records(saved):
            self.assertEqual(row['dimensions'],before[location]['dimensions'])
            if location in ('合成表!B3','合成表!B4'):
                self.assertEqual(row['label_references']['period'],['合成表!B2'])

    def test_数据格适用标签与标签适用于数据格均只保存明示引用(self):
        b=self.save('B')
        answer='合成表!B3、B4适用原表标签 合成表!B2“期末余额”；C3、C4适用 合成表!C2“期初余额”。'
        saved=save_direct_answer(b,answer,self.case.folder/'正反向标签.json',
            field_scope={'合成表!B3:C4':['period']})
        for col in 'BC':
            for number in (3,4):
                location='合成表!'+col+str(number)
                row=dict(_records(saved))[location]
                self.assertEqual(row['dimension_evidence']['period'],['合成表!'+col+'2'])
                self.assertEqual(row['dimensions'],dict(_records(b))[location]['dimensions'])
        negative=direct_answer('合成表!B3不适用原表标签 合成表!C2“期初余额”。',self.engine.standard,['合成表'])
        self.assertFalse(any(r.get('evidence') for r in negative['dimensions']))

    def test_对应列数据格接入明确坐标及沿用行范围(self):
        answer=('合成表!B3:B4对应B2“期初余额”；'
                '对应的C列数据格C3:C4对应C2“本期应交”；'
                '对应的D列数据格对应D2“本期已交”；'
                '对应的E列数据格对应E2“期末余额”。')
        parsed=direct_answer(answer,self.engine.standard,['合成表'])
        labels={cell:row['label_references'] for row in parsed['dimensions']
                if row['id']=='period' and row.get('label_references') for cell in row['ranges']}
        self.assertEqual(labels,{'合成表!'+col+str(number):['合成表!'+col+'2']
                                 for col in 'BCDE' for number in (3,4)})
        original=self.save('B');before=dict(_records(original))
        for number,text in enumerate((answer,answer.replace('C列数据格C3:C4','C列数据格'))):
            saved=save_direct_answer(original,text,self.case.folder/f'对应列标签{number}.json',
                                     field_scope={'合成表!B3:C4':['period']})
            self.assertEqual(saved['机械记录'],original['机械记录'])
            for location,row in _records(saved):
                self.assertEqual(row['dimensions'],before[location]['dimensions'])
                if location in ('合成表!B3','合成表!B4','合成表!C3','合成表!C4'):
                    self.assertEqual(row['label_references']['period'],[location[:-1]+'2'])
        with self.assertRaisesRegex(ValueError,'未明确可机械读取'):
            direct_answer('对应的D列数据格对应D2“本期已交”。',self.engine.standard,['合成表'],default_sheet='合成表')

    def test_A字段纠正保留字典上下文且只使用一次已有答复(self):
        a=self.save('A');context={'text':'本期表头说明','source':'原文'}
        with patch.object(UnifiedSemanticEngine,'request_text',return_value=
                '原表 合成表!B2 标签原文为“期末余额”，适用于数据格 B3。') as call:
            saved=self.engine.recognize_workbook(self.case.source,context=[context],
                previous_result=a['mapping_path'],artifact_metadata={'carrier':'A',
                'field_scope':{'合成表!B3':['period']}})
        self.assertEqual(call.call_count,1)
        self.assertIn(context,saved['source_context'])
        self.assertEqual(saved['机械记录'],a['机械记录'])
        self.assertEqual(dict(_records(saved))['合成表!B3']['dimension_evidence']['period'],['合成表!B2'])

    def test_同行标签接入不丢其他已识别字段(self):
        b=self.save('B')
        answer='合成表!B3：asset_selection=predicate，“设备甲”。原表期间标签坐标：合成表!B2“期末余额”。'
        saved=save_direct_answer(b,answer,self.case.folder/'同行字段.json',
            field_scope={'合成表!B3':['asset_selection','period']})
        row=dict(_records(saved))['合成表!B3']
        self.assertEqual(row['dimensions']['asset_selection']['predicate'],'设备甲')
        self.assertEqual(row['dimension_evidence']['period'],['合成表!B2'])
        original=dict(_records(b))['合成表!B3']
        self.assertEqual(row['dimensions']['period'],original['dimensions']['period'])

    def test_各自名称声明只连接已明确名称列且不套用其他金额(self):
        b=self.save('B')
        for row in b['outline']['subtables'][0]['dimensions']:
            if row['id']=='asset_selection':row['evidence']=[]
        for _,row in _records(b):row['dimension_evidence']['asset_selection']=[]
        b['已确认关系']=[{'kind':'名单编排','dimension':'asset_selection',
            'source_b_hash':b['source_hash'],'gold_hash':b['gold_hash'],
            'source_ranges':[{'sheet':'合成表','range':'A2:C8'}],
            'content':'B原表为合成表!A2:C8，名称在 A 列。金额不总随同行名称。'}]
        b['封存校验']=direct_digest(b)
        reply=('合成表：\n`asset_selection`的domain为“固定资产”。各坐标均为`mode=members`，成员分别为：\n'
               '| 坐标 | 成员 |\n|---|---|\n| B3 | 设备甲 |\n'
               'C3:C4的空白上期发生额格按各自表格项目名称取范围：依次为“设备甲、设备乙”；除说明行外均为domain“固定资产”、`mode=members`。')
        saved=save_direct_answer(b,reply,self.case.folder/'各自名称.json',
            field_scope={'合成表!C3:C4':['asset_selection']})
        rows=dict(_records(saved))
        self.assertEqual(rows['合成表!C3']['dimension_evidence']['asset_selection'],['合成表!A3'])
        self.assertEqual(rows['合成表!C4']['dimension_evidence']['asset_selection'],['合成表!A4'])
        self.assertEqual(rows['合成表!B3']['dimension_evidence']['asset_selection'],[])
        b['已确认关系'][0]['content']='合成表!A2:C8，名称在 A 列。合成表!A2:C8，名称在 B 列。'
        b['封存校验']=direct_digest(b)
        ambiguous=save_direct_answer(b,reply,self.case.folder/'名称列不唯一.json',
            field_scope={'合成表!C3:C4':['asset_selection']})
        self.assertEqual(dict(_records(ambiguous))['合成表!C3']['dimension_evidence']['asset_selection'],[])

    def test_按原文行标签声明直接接成员且不按名单顺序猜范围(self):
        snapshot=deepcopy(self.snapshot);cells=snapshot['sheets'][0]['cells']
        for cell,text in {'A3':'  设备乙','A4':'其中：设备甲','A5':'设备甲别称','A6':'0','A7':'设备乙'}.items():
            cells.setdefault(cell,{}).update(value=text,formula=None)
        cells['A7'].update(value='=A3',formula='=A3',cached_value='  设备乙')
        b=self.save('B',snapshot=snapshot);before=dict(_records(b))
        b['已确认关系']=[{'kind':'名单编排','dimension':'asset_selection',
            'source_b_hash':b['source_hash'],'gold_hash':b['gold_hash'],
            'source_ranges':[],
            'tables':[{'source_ranges':[{'sheet':'合成表','range':'A3:C8'}],'labels':{'A':'A'}}],
            'content':'合成表!A4标为“设备甲”，来源标签也是“其中：设备甲”。'}]
        b['封存校验']=direct_digest(b)
        reply=('合成表：\n以下“类别成员”均为asset_selection，mode=members、domain=固定资产，成员按原文标签列出；'
               '“合计”行为全体，mode=all、domain=固定资产。类别成员包括：设备甲、设备乙。\n'
               'B3:C7为各类别成员；B8:C8为全体合计。')
        saved=save_direct_answer(b,reply,self.case.folder/'明示行标签.json',
            field_scope={'合成表!C3:C8':['asset_selection']})
        rows=dict(_records(saved))
        for cell,member in [('C3','设备乙'),('C4','设备甲'),('C7','设备乙')]:
            self.assertEqual(rows['合成表!'+cell]['dimensions']['asset_selection']['members'],[member])
            self.assertEqual(rows['合成表!'+cell]['dimension_evidence']['asset_selection'],['合成表!A'+cell[1:]])
        for cell in ['C5','C6']:
            self.assertEqual(rows['合成表!'+cell]['dimensions']['asset_selection']['mode'],'unknown')
        self.assertEqual(rows['合成表!C8']['dimensions']['asset_selection']['mode'],'all')
        self.assertEqual(rows['合成表!B3']['dimensions'],before['合成表!B3']['dimensions'])
        for loc,row in rows.items():
            self.assertEqual(row.get('raw_value',row.get('value')),before[loc].get('raw_value',before[loc].get('value')))
            self.assertEqual(row['metric_id'],before[loc]['metric_id'])
            for dim,value in row['dimensions'].items():
                if dim!='asset_selection':self.assertEqual(value,before[loc]['dimensions'][dim])
        self.assertEqual(saved['答复接入来源']['本次原答复'],reply)
        self.assertEqual(saved['机械记录'],b['机械记录'])
        b['已确认关系']=[];b['封存校验']=direct_digest(b)
        without=save_direct_answer(b,reply,self.case.folder/'未声明名称列.json',
            field_scope={'合成表!C3':['asset_selection']})
        self.assertEqual(dict(_records(without))['合成表!C3']['dimensions']['asset_selection']['mode'],'unknown')

    def test_已明确名称公式复用副本取值后进入名单安排(self):
        from openpyxl import Workbook
        from 附注更新.统一勾稽 import _calculated_values,calculation_files
        a=self.save('A');answer=json.loads(self.case.answer())
        row=next(r for r in answer['tables'][0]['dimensions'] if r['id']=='asset_selection')
        row['value']['members']=['新名单成员']
        snapshot=deepcopy(self.snapshot)
        snapshot['sheets'][0]['cells']['A3'].update(value='="新名单成员"',formula='="新名单成员"')
        b=self.save('B',json.dumps(answer,ensure_ascii=False),snapshot)
        receipt={'schema':'统一语义公式计算-v1','method':'Excel.CalculateFullRebuild',
            'source_hash':b['source_hash'],'source_path':b['source_path'],'workbook_path':'合成重算副本','workbook_hash':'副本指纹'}
        def book(path,data_only=False):
            wb=Workbook();ws=wb.active;ws.title='合成表'
            for cell,item in snapshot['sheets'][0]['cells'].items():ws[cell]=item['value']
            if data_only:ws['A3']='新名单成员'
            return wb
        with patch('附注更新.统一勾稽.load_workbook',side_effect=book) as read, \
             patch('附注更新.统一勾稽.file_hash',side_effect=lambda p:'副本指纹' if p=='合成重算副本' else b['source_hash']):
            calculated=_calculated_values(b,receipt)
        read.assert_called_once_with('合成重算副本',data_only=True)
        self.assertEqual(calculated['合成表!A3']['value'],'新名单成员')
        b['calculation_evidence']={'path':str(self.case.folder/'计算依据.json'),'sha256':'依据指纹'}
        Path(b['calculation_evidence']['path']).write_text(json.dumps(receipt),encoding='utf-8')
        with patch('附注更新.统一勾稽.file_hash',return_value='依据指纹'), \
             patch('附注更新.统一勾稽._calculated_values',side_effect=AssertionError('列出文件不重核计算结果')):
            calculation_files(b)
            b['formula_values']=deepcopy(calculated)
            b['formula_values']['合成表!A3']['value']='错误名称'
            self.assertEqual(calculation_files(b),[b['calculation_evidence']['path'],receipt['workbook_path'],receipt['source_path']])
        b['formula_values']=calculated;b['封存校验']=direct_digest(b)
        purpose={'reason':'沿已明确名单安排',
            'bindings':[{'target':f'合成表!{col}3','source':f'合成表!{col}3'} for col in 'BC'],
            'labels':[{'target':'合成表!A3','source':'合成表!A3'}]}
        plan=direct_update_plan(a,b,purpose)
        self.assertTrue(any(r['cell']=='A3' and r['value']=='新名单成员' for r in plan['labels']))
        self.assertEqual(b['机械记录']['sheets'][0]['cells']['A3']['formula'],'="新名单成员"')

    def test_明确全体模式不混入后续说明(self):
        tabular=direct_answer('| 坐标 | 金标准指标编号 | 金标准维度名 | 实际取值 |\n'
            '|---|---|---|---|\n| 合成表!B3 | 固定资产原值 | asset_selection | all：固定资产总体。 |',
            self.engine.standard,['合成表'])
        self.assertTrue(any(row['id']=='asset_selection' and row.get('value')=='all' for row in tabular['dimensions']))
        b=self.save('B')
        answer='合成表!B3：asset_selection=固定资产全体（all，completeness=complete）。期间依据另有原文说明。'
        saved=save_direct_answer(b,answer,self.case.folder/'全体取值.json',field_scope={'合成表!B3':['asset_selection']})
        value=dict(_records(saved))['合成表!B3']['dimensions']['asset_selection']
        self.assertEqual(value,{'domain':'固定资产','mode':'all','completeness':'complete'})

    def test_上述各项只引用紧邻的已声明项目(self):
        reply=('合成表!B7：固定资产原值。\n\n'
               '- 合成表!B3：固定资产原值。\n'
               '- 合成表!B4：固定资产原值。\n'
               '- 上述各项的 asset_selection 均为固定资产全体（all）。B5名称空白。')
        parsed=direct_answer(reply,self.engine.standard,['合成表'])
        selections=[r for r in parsed['dimensions'] if r['id']=='asset_selection']
        self.assertEqual(len(selections),1)
        self.assertEqual(set(selections[0]['ranges']),{'合成表!B3','合成表!B4'})
        self.assertEqual(selections[0]['value'],{'mode':'all'})

    def test_坐标分句及括号方式不吞掉下句或混用位置(self):
        reply=('合成表!B3:B4：名称未知。'
               '合成表!C3:C4：占比；分母为全部固定资产（asset_selection：predicate）。'
               '合成表!B5:B6：变动事项未限定（asset_selection：all），不扩大其他类别。')
        parsed=direct_answer(reply,self.engine.standard,['合成表'])
        dims=[r for r in parsed['dimensions'] if r['id']=='asset_selection']
        self.assertEqual(len(dims),2)
        self.assertEqual(dims[0]['ranges'],['合成表!C3:C4'])
        self.assertEqual(dims[0]['value'],{'mode':'predicate','predicate':'全部固定资产'})
        self.assertEqual(dims[1]['ranges'],['合成表!B5:B6'])
        self.assertEqual(dims[1]['value'],'all')
        missing=direct_answer('合成表!B3：事项（asset_selection：predicate）。',self.engine.standard,['合成表'])
        self.assertFalse(any(r['id']=='asset_selection' for r in missing['dimensions']))
        self.assertTrue(any(r.get('id')=='asset_selection' for r in missing['unresolved']))

    def test_列名称引用只接原答复明示的唯一编号(self):
        standard=deepcopy(self.engine.standard)
        for name in ('合成余额','合成变动'):
            standard['metric'][name]={**deepcopy(standard['metric']['固定资产原值']),'id':name}
        columns='B列为期初金额；C列为期初比例。金额格适用 合成余额、合成变动，比例格适用 固定资产原值。'
        reply='合成表!B3:C4：范围已给出。\n  '+columns
        parsed=direct_answer(reply,standard,['合成表'])
        metrics=[r for r in parsed['metrics'] if r['id']=='固定资产原值']
        self.assertEqual({loc for r in metrics for loc in r['ranges']},{'合成表!C3','合成表!C4'})
        self.assertEqual({r.get('template_period_role') for r in metrics},{'期初'})
        self.assertFalse(any(r['id'] in ('合成余额','合成变动') for r in parsed['metrics']))
        with self.assertRaisesRegex(ValueError,'未明确可机械读取的坐标'):
            direct_answer('  '+columns,standard,['合成表'])
        different=direct_answer(reply.replace('C列为期初比例','C列为期初比重'),standard,['合成表'])
        self.assertFalse(different['metrics'])
        ambiguous=direct_answer(reply.replace('比例格适用 固定资产原值','比例格适用 固定资产原值、合成变动'),standard,['合成表'])
        self.assertFalse(ambiguous['metrics'])
        single=reply.replace('金额格适用 合成余额、合成变动，','')
        parsed=direct_answer(single,standard,['合成表'])
        self.assertEqual({loc for r in parsed['metrics'] for loc in r['ranges']},{'合成表!C3','合成表!C4'})
        self.assertFalse(direct_answer(single.replace('C列为期初比例','C列为期初比重'),standard,['合成表'])['metrics'])
        with self.assertRaisesRegex(ValueError,'未明确可机械读取的坐标'):
            direct_answer(single.split('\n')[1],standard,['合成表'])

    def test_编号未定仍保存明示列角色(self):
        answer=json.loads(self.case.answer());answer['tables'][0]['metrics']=[]
        before=self.save('B',json.dumps(answer,ensure_ascii=False))
        reply='合成表!B3:C4：编号未定。\n  B列为期初金额；C列为本期增加。'
        saved=save_direct_answer(before,reply,self.case.folder/'未绑定期间角色.json',
            field_scope={'合成表!B3:C4':['template_period_role']})
        self.assertFalse(dict(_records(saved)))
        roles=[r for t in saved['outline']['subtables'] for r in t['dimensions'] if r['id']=='template_period_role']
        self.assertEqual({cell:r['value'] for r in roles for cell in r['ranges']},
                         {'B3':'期初','B4':'期初','C3':'本期','C4':'本期'})
        self.assertEqual(saved['机械记录'],before['机械记录'])
        self.assertEqual({(r['sheet'],r['cell']):r.get('raw_value') for r in saved['unresolved']},
                         {(r['sheet'],r['cell']):r.get('raw_value') for r in before['unresolved']})

    def test_未绑定位置重读只修已明示字段不补编号(self):
        answer=json.loads(self.case.answer());answer['tables'][0]['metrics']=[]
        snapshot=deepcopy(self.snapshot);snapshot['sheets'][0]['cells']['C3'].update(value='=B3',formula='=B3')
        b=self.save('B',json.dumps(answer,ensure_ascii=False),snapshot);reply=(
            '合成表!B3：名称未知。合成表!C3：分母为全部固定资产（asset_selection：predicate）。')
        saved=save_direct_answer(b,reply,self.case.folder/'未绑定字段重读.json',
            field_scope={'合成表!C3':['asset_selection']},withdrawn={'合成表!B3':['asset_selection']})
        self.assertFalse(dict(_records(saved)))
        unresolved={r['cell']:r for r in saved['unresolved'] if r.get('cell')}
        self.assertNotIn('asset_selection',unresolved['B3']['original']['dimensions'])
        value=unresolved['C3']['original']['dimensions']['asset_selection']['value']
        self.assertEqual(value,{'domain':'固定资产','mode':'predicate','predicate':'全部固定资产'})
        self.assertEqual(unresolved['C3']['reason'],'未提供标准编号')
        self.assertEqual(unresolved['C3']['raw_value'],'=B3')
        self.assertEqual(unresolved['C3']['raw_value_state'],'formula')
        self.assertEqual(saved['机械记录'],b['机械记录'])

    def test_明说同行标签只机械定位原字并保留期间(self):
        snapshot=deepcopy(self.snapshot)
        snapshot['sheets'][0]['cells']['A3']['value']='本期计提'
        snapshot['sheets'][0]['cells']['A4']['value']='本期转回'
        b=self.save('B',snapshot=snapshot)
        reply='合成表：\nB、C列第3:4行的指定数据格对应各自行标签“本期计提”“本期转回”。'
        saved=save_direct_answer(b,reply,self.case.folder/'同行标签.json',
            field_scope={'合成表!B3:C4':['period']})
        before=dict(_records(b));after=dict(_records(saved))
        for cell in ('B3','C3','B4','C4'):
            loc='合成表!'+cell
            self.assertEqual(after[loc]['dimensions']['period'],before[loc]['dimensions']['period'])
            self.assertEqual(after[loc]['label_references']['period'],['合成表!A'+cell[1:]])
        self.assertEqual(saved['机械记录'],b['机械记录'])
        missing=save_direct_answer(b,reply.replace('本期转回','另一个标签'),self.case.folder/'同行无原字.json',
            field_scope={'合成表!B3:C4':['period']})
        self.assertEqual(dict(_records(missing))['合成表!B4']['label_references'],before['合成表!B4']['label_references'])

    def test_自然分段及分列标签保存不混用且不改期间(self):
        b=self.save('B');before=dict(_records(b))
        replies=[
            ('合成表!B3:C3：原表标签位于 B2:C2，原文“期末余额”，适用于该行指定范围。',
             {'B3':['B2','C2'],'C3':['B2','C2']}),
            ('合成表!B3:C3 与 B4:C4：两段数据均对应原表第2行及第2行的列标签：B列“期末余额”、C列“期初余额”。',
             {col+str(row):[col+'2'] for col in 'BC' for row in (3,4)}),
            ('合成表!B3:C4：对应原表 B2“期末余额”和 C2“期初余额”，分别适用于 B 列和 C 列指定范围。',
             {col+str(row):[col+'2'] for col in 'BC' for row in (3,4)})]
        for number,(reply,expected) in enumerate(replies):
            saved=save_direct_answer(b,reply,self.case.folder/f'分组标签{number}.json',
                                     field_scope={'合成表!B3:C4':['period']})
            after=dict(_records(saved))
            for cell,refs in expected.items():
                self.assertEqual(after['合成表!'+cell]['label_references']['period'],['合成表!'+ref for ref in refs])
            self.assertEqual({loc:row['dimensions'] for loc,row in after.items()},
                             {loc:row['dimensions'] for loc,row in before.items()})
            self.assertEqual(saved['机械记录'],b['机械记录'])

    def test_分段名单替换旧安排且明确保留的行不自动匹配(self):
        from 附注更新.A合集 import save_list_relations
        from 附注更新.统一勾稽 import load_fact_mapping
        snapshot=deepcopy(self.snapshot)
        snapshot['names']={'名单表':[{'sheet':'合成表','range':'A2:C8'}]}
        a=self.save('A',snapshot=snapshot);b=self.save('B',snapshot=snapshot)
        table={'selection_dimension':'asset_selection','range_name':'名单表','source_sheet':'合成表',
               'source_rows':[3,4,5],'old_detail_count':2,'first_row_offset':1,'columns':{'B':'B','C':'C'},'labels':{'A':'A'}}
        b['已确认关系']=[{'kind':'名单编排','dimension':'asset_selection','answer_hash':'旧名单',
            'source_a_hash':a['source_hash'],'source_b_hash':b['source_hash'],'gold_hash':a['gold_hash'],
            'tables':[table,{**table,'first_row_offset':4,'source_rows':[6,7]}]},
            {'kind':'取数对应','source_a_hash':a['source_hash'],'source_b_hash':b['source_hash'],'gold_hash':a['gold_hash'],
             'bindings':[{'target':'合成表!B4','source':'合成表!C4','target_member':{
                 'range_name':'名单表','first_row_offset':1,'source_sheet':'合成表','source_row':4,'column':'B'}}]}]
        b['封存校验']=direct_digest(b)
        reply=('本次只调整第一段，不增行。\n'
               '- 第一段（A3:A4）：A3对应B《合成表》3行。A4“其他”保留原位置；B没有独立来源。\n'
               'B列“期末余额”搬至A对应行B列，C列“期初余额”搬至A对应行C列。')
        result=save_list_relations(a,b,reply,'asset_selection',self.case.folder/'分段替换',
                                   source_ranges=[{'sheet':'合成表','range':'A2:C8'}])
        saved=load_fact_mapping(result['files']['b_mapping'],{})
        self.assertEqual(saved['facts'],b['facts'])
        self.assertEqual([t['first_row_offset'] for t in saved['已确认关系'][0]['tables']],[4])
        self.assertEqual(saved['已确认关系'][1]['bindings'],[])
        latest=saved['已确认关系'][-1]['tables'][0]
        self.assertEqual(latest['operations'],[])
        self.assertEqual(latest['retained_rows'],[4])
        from 附注更新.A合集 import _declared_list_tables
        grouped=reply+'\nB的D、E列为“本期增加”“本期减少”；这两列均为空，对应A的D、E列保持原状。'
        self.assertEqual(_declared_list_tables(a,grouped,'asset_selection',[])[0]['columns'],
                         {'B':'B','C':'C','D':'D','E':'E'})
        plan=direct_update_plan(a,saved)
        self.assertFalse(plan['operations'])
        self.assertEqual({x['cell'] for x in plan['retained_values']},{'B4','C4'})
        self.assertFalse(any(x['cell'] in ('B4','C4') for x in plan['updates']))
        self.assertEqual(next(x for x in plan['updates'] if x['cell']=='B3')['sources'][0]['fact_id'],'合成表!B3')

    def test_共用表头已有搬运关系不重复要求各金额提供同一引用(self):
        a=self.save('A');answer=json.loads(self.case.answer())
        for row in answer['tables'][0]['dimensions']:
            if row['id']=='period' and row['value'].get('kind')=='instant':row['value']['date']='2026-12-31'
        b=self.save('B',json.dumps(answer,ensure_ascii=False))
        for _,row in _records(b):
            row.get('label_references',{}).pop('period',None)
        records=dict(_records(b));records['合成表!B3'].setdefault('label_references',{})['period']=['合成表!B2']
        b['封存校验']=direct_digest(b)
        plan=direct_update_plan(a,b,{'bindings':[{'target':'合成表!B3','source':'合成表!B3'},
                                                {'target':'合成表!B4','source':'合成表!B4'}]})
        self.assertFalse(any(item.get('action')=='label' and item.get('cell')=='B4' for item in plan['issues']))
        self.assertNotIn('period',records['合成表!B4'].get('label_references',{}))

    def test_完整重述撤下旧解析失败而局部修正不撤下(self):
        answer=json.loads(self.case.answer())
        answer['tables'][0]['unresolved']=[
            {'ranges':['B3'],'reason':'这段答复未明确可机械读取的字段归属','original':'旧的未解析描述'},
            {'ranges':['B3'],'reason':'原文仍有未解释的业务冲突','original':'真实冲突'}]
        b=self.save('B',json.dumps(answer,ensure_ascii=False))
        partial=save_direct_answer(b,'合成表!B3：asset_selection=all。',self.case.folder/'局部重述.json',
                                  field_scope={'合成表!B3':['asset_selection']})
        self.assertTrue(any('这段答复未明确可机械读取的字段归属' in r['reason'] for r in partial['unresolved']))
        keys=['metric_id',*[r['id'] for r in self.engine.standard['metric']['固定资产原值']['dimension_refs']]]
        full=save_direct_answer(b,self.case.answer(),self.case.folder/'完整重述.json',field_scope={'合成表!B3':keys})
        self.assertFalse(any('这段答复未明确可机械读取的字段归属' in r['reason'] for r in full['unresolved']))
        self.assertTrue(any('原文仍有未解释的业务冲突' in r['reason'] for r in full['unresolved']))


if __name__=='__main__':unittest.main()
