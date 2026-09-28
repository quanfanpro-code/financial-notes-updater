"""用本地 HTTP 服务验证协议和拒绝错误结果；这些测试不代表模型语义准确率。"""
import copy
import http.server
import json
import pathlib
import threading
import unittest
import uuid

from 附注更新.语义 import SemanticEngine, SemanticCancelled, SemanticError, SameOriginRedirect

ROOT = pathlib.Path(__file__).resolve().parents[1]


def snapshot(source_hash="source-one"):
    values = {"A1": "货币资金（合并，2025年，人民币元）", "A2": "项目", "B2": "期末余额", "A3": "银行存款", "B3": 123.45}
    return {"path": "虚构测试.xlsx", "sha256": source_hash, "context": [], "sheets": [{"name": "表一", "max_row": 3, "max_column": 2, "merges": [], "cells": {a: {"row": int(a[1:]), "column": ord(a[0]) - 64, "value": v, "formula": None, "cached_value": v, "number_format": "General", "style_id": 0} for a, v in values.items()}}]}


def mapping():
    return {"cell": "B3", "slot_id": "C-N001-T001-S001", "scope": "consolidated", "dimensions": {"currency": "人民币", "period": "2025-12-31", "unit": "元", "scale": 1}, "semantic_field": "银行存款期末余额", "value_type": "monetary", "reason": "A1合并口径与人民币元；A3银行存款；B2期末余额", "row_label": "银行存款", "column_label": "期末余额"}


class LocalAPI(http.server.BaseHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def do_POST(self):
        outer = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        self.server.bodies.append(outer)
        payload = json.loads(outer["messages"][-1]["content"])
        self.server.received.append(payload)
        mode = self.server.mode
        if mode == "redirect":
            self.send_response(302)
            self.send_header("Location", self.server.redirect_url)
            self.end_headers()
            return
        if payload["task"] == "layout_cells":
            result={"cells":{a:({"kind":"business","reason":"原表货币资金业务","note_id":"C-N001","scope":"consolidated",
                "table_semantic":"银行存款余额","header_cells":["A1","A3","B2"]} if a=="B3" else
                {"kind":"non_business","reason":"原始文字表头或独立标题同行留白","category":"empty_padding" if a=="B1" else "header","evidence_cells":["A1"] if a=="B1" else [a]}) for a in payload["candidate_cells"]}}
        elif payload["task"] == "layout":
            result = {"tables": [{"range": "A1:B3", "note_ids": ["C-N001"], "scope": "consolidated", "table_semantic": "货币资金", "header_cells": ["A1", "A2", "B2"]}], "carry_context": "合并货币资金，人民币元"}
            if mode == "layout_omission":
                result["tables"][0].update(range="A1:A2",header_cells=["A1","A2"])
            if mode in {"title_only", "title_conflict", "title_number"}:
                category = "annotation" if mode == "title_conflict" and payload.get("round") == 2 else "title"
                result = {"tables": [], "non_business": [{"range": "A1:A1", "category": category,
                          "reason": "A1是报告标题", "evidence_cells": ["A1"]}]}
        elif payload["task"] == "tables":
            result = {"table_ids": ["C-N001-T001"]}
        else:
            cells = payload["candidate_cells"]
            item = mapping()
            if mode == "bad_id":
                item["slot_id"] = "C-N001-T001-S999"
            if mode == "wrong_scope":
                item["scope"] = "parent"
            if mode == "wrong_type":
                item["value_type"] = "percentage"
            if mode == "no_dimension":
                item["dimensions"].pop("period")
            if mode == "conflict" and payload["round"] == 2:
                item["dimensions"]["period"] = "2024-12-31"
            if mode == "field_words":
                item["semantic_field"] = "银行存款期末余额" if payload["round"] == 1 else "银行存款年末金额"
            if mode == "metric_conflict":
                item["dimensions"]["metric"] = "土地面积" if payload["round"] == 1 else "拟开发面积"
                item["metric_evidence"] = ["A3"]
            if mode == "value_injected":
                item["value"] = 999999
            if mode == "synonyms" and payload["round"] == 2:
                item["dimensions"]["currency"] = "CNY"
                item["dimensions"]["scale"] = "1"
            result = {"mappings": [item] if "B3" in cells else [], "excluded": [{"cell": a, "reason": "标题或行列表头", "category": "header"} for a in cells if a != "B3"], "unresolved": []}
            if mode == "numeric_exclusion":
                result["mappings"] = []
                result["excluded"].append({"cell": "B3", "category": "header", "reason": "不是业务格"})
            if mode == "unresolved":
                result["mappings"] = []
                result["unresolved"] = [{"cell": "B3", "reason": "目前不能确定"}]
            if mode == "omission":
                result["mappings"] = []
            if mode == "duplicate":
                result["mappings"].append(copy.deepcopy(item))
            if mode == "outside":
                result["mappings"].append({**item, "cell": "Z999"})
            if mode == "server_error":
                self.send_response(503)
                self.end_headers()
                return
        if mode=="per_cell" and "mappings" in result:
            names={"mappings":"mapping","excluded":"excluded","unresolved":"unresolved"}
            result={"cells":{item["cell"]:{"kind":names[category],**{k:v for k,v in item.items() if k!="cell"}} for category in names for item in result[category]}}
        content=json.dumps(result,ensure_ascii=False)
        if mode=="inline_reasoning":
            content='先核对结构。{"tables":[]}只是中间草稿。\n</think>\n'+content
        elif mode=="tagged_reasoning":
            content="<think>先核对结构。</think>\n\n```json\n"+content+"\n```"
        elif mode=="unfinished_reasoning":
            content="<think>"+content
        elif mode=="reasoning_no_answer":
            content="尚未形成结果。</think>"
        elif mode=="literal_think_marker":
            result["carry_context"]="说明中含有字面标记 </think>，仍然属于JSON字符串"
            content=json.dumps(result,ensure_ascii=False)
        data = json.dumps({"choices": [{"finish_reason": "stop", "message": {"content": content}}], "usage": {"prompt_tokens": 10, "completion_tokens": 20}}).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)


class SemanticTests(unittest.TestCase):
    def test_progress_output_keeps_full_coverage_without_accepting_single_round(self):
        from 附注更新.语义 import recognition_progress
        data=snapshot()
        raw={"mappings":[dict(mapping(),sheet="表一",reviewed=False)],
             "excluded":[{"sheet":"表一","cell":"A1","reason":"只看了一轮","reviewed":False}],
             "unresolved":[{"sheet":"表一","cell":"A2","reason":"原有待核实原因"}],"complete":True}
        before=copy.deepcopy(raw)
        result=recognition_progress(data,raw)
        self.assertFalse(result["complete"])
        self.assertEqual(result["mappings"],[])
        self.assertEqual(result["excluded"],[])
        self.assertEqual({i["cell"] for i in result["unresolved"]},set(data["sheets"][0]["cells"]))
        self.assertEqual(len(result["unresolved"]),5)
        self.assertEqual(next(i for i in result["unresolved"] if i["cell"]=="A2")["reason"],"原有待核实原因")
        self.assertEqual(raw,before)
        with self.assertRaisesRegex(SemanticError,"重复|包外"):
            recognition_progress(data,{**raw,"unresolved":[{"sheet":"表一","cell":"B3","reason":"重复"}]})

    def test_word_packet_publishes_small_new_progress_after_950_reused_cells(self):
        from unittest.mock import patch
        engine=self.engine();data=snapshot();sheet=data["sheets"][0]
        sheet.update(max_row=956,max_column=1,word_table_ranges=[{"range":"A1:A952"},{"range":"A953:A956"}],
            cells={f"A{r}":{"row":r,"column":1,"value":r,"cached_value":r,"formula":None} for r in range(1,957)})
        kept={"mappings":[{"sheet":sheet["name"],"cell":f"A{r}","reviewed":True} for r in range(1,951)],"excluded":[],"unresolved":[]}
        saved=[]
        def layout(system,payload,current,owned):
            if "A953" in owned:
                self.assertGreaterEqual(len(saved),2,"第一张Word表的新2格须在下一张表开始前保存")
                pending={i["cell"]:i for i in saved[-1]["unresolved"]}
                self.assertEqual(pending["A951"]["reason"],"布局待核实：本表新判断")
            return {"tables":[],"unresolved":[{"range":a,"reason":"本表新判断","evidence_cells":[a]} for a in owned]}
        with patch.object(engine,"_reuse_previous",return_value=(kept,{})),patch.object(engine,"_request_layout",side_effect=layout):
            result=engine.recognize(data,previous_result={},progress=lambda snap,out:saved.append(copy.deepcopy(out)))
        self.assertEqual(len(result["mappings"]),950)
        self.assertEqual(len(result["unresolved"]),6)
        self.assertGreaterEqual(len(saved),3)
        for event in saved:
            self.assertEqual(sum(len(event[k]) for k in ("mappings","excluded","unresolved")),956)

    def test_plain_excel_packet_saves_before_next_packet_after_existing_business(self):
        from unittest.mock import patch
        engine=self.engine();data=snapshot();sheet=data["sheets"][0]
        sheet.update(max_row=1196,max_column=1,
            cells={f"A{r}":{"row":r,"column":1,"value":r,"cached_value":r,"formula":None} for r in range(1,1197)})
        kept={"mappings":[{"sheet":sheet["name"],"cell":f"A{r}","reviewed":True} for r in range(1,951)],"excluded":[],"unresolved":[]}
        saved=[];seen_before_next=[]
        def layout(system,payload,current,owned):
            if "A1191" in owned:
                pending={i["cell"]:i for i in saved[-1]["unresolved"]}
                seen_before_next.append(pending["A951"]["reason"])
            return {"tables":[],"unresolved":[{"range":a,"reason":"已读取本包原文，仍缺业务依据","evidence_cells":[a]} for a in owned]}
        with patch.object(engine,"_reuse_previous",return_value=(kept,{})),patch.object(engine,"_request_layout",side_effect=layout):
            result=engine.recognize(data,previous_result={},progress=lambda snap,out:saved.append(copy.deepcopy(out)))
        self.assertEqual(seen_before_next,["布局待核实：已读取本包原文，仍缺业务依据"])
        self.assertEqual(len(result["mappings"]),950)
        self.assertEqual(len(result["unresolved"]),246)
        for event in saved:
            self.assertEqual(event["mappings"],kept["mappings"])
            self.assertEqual(sum(len(event[k]) for k in ("mappings","excluded","unresolved")),1196)

    def test_model_receives_distinct_parent_populations_without_duplicate_source_examples(self):
        from 附注更新.语义 import _hash
        row=json.loads(self.gold.read_text(encoding="utf-8-sig"))
        parent={"parent_slot_id":"C-N001-T001-S002","dimension_constraints":{"category":"短期薪酬"},
                "original_period_role":"closing","original_column_path":["期末余额"],"interpretation":"按实际范围理解"}
        other={**copy.deepcopy(parent),"dimension_constraints":{"category":"离职后福利"}}
        row["aggregation"].update(parent_variants=[{**copy.deepcopy(parent),"source_table_ids":["示例甲"]},
            {**copy.deepcopy(parent),"source_table_ids":["示例乙"]},{**copy.deepcopy(other),"source_table_ids":["示例丙"]}],
            component_variants=[{"component_slot_ids":["C-N001-T001-S003"],"source_table_ids":["示例甲"]},
                {"component_slot_ids":["C-N001-T001-S003"],"source_table_ids":["示例乙"]}])
        gold=self.work/("多来源相同业务_"+uuid.uuid4().hex+".jsonl")
        gold.write_text(json.dumps(row,ensure_ascii=False),encoding="utf-8-sig");original=gold.read_bytes()
        engine=SemanticEngine(self.engine().settings,str(gold),str(self.work/uuid.uuid4().hex))
        data=snapshot();sheet=data["sheets"][0]
        table={"range":"A1:B3","note_ids":["C-N001"],"scope":"consolidated","table_semantic":"货币资金","header_cells":["A1","A2","B2"]}
        result=engine._round(sheet,table,["B3"],list(engine.slots.values()),[],{},1)
        self.assertEqual(result["B3"][0],"mappings")
        sent=self.server.received[-1]["gold_slots"][0]["aggregation"]
        self.assertEqual(sent["parent_variants"],[parent,other])
        self.assertEqual(sent["component_variants"],[{"component_slot_ids":["C-N001-T001-S003"]}])
        self.assertEqual(engine.definition_hashes[row["id"]],_hash(row))
        self.assertEqual(gold.read_bytes(),original,"发送投影不能改变完整金标准及其来源依据")

    def test_progress_save_failure_is_not_swallowed_as_model_failure(self):
        from unittest.mock import patch
        from 附注更新.语义 import SemanticProgressError
        data=snapshot();engine=self.engine();logs=[];engine.log=logs.append
        table={"range":"A1:B3","note_ids":["C-N001"],"scope":"consolidated","table_semantic":"测试余额","header_cells":["A1","A2","B2"]}
        def classify(sheet,table,candidates,*args):
            return {"mappings":[{"sheet":sheet["name"],"cell":a,"reviewed":True} for a in candidates],"excluded":[],"unresolved":[]}
        def failed(*args):raise OSError("磁盘写入失败")
        with patch.object(engine,"_request_layout",return_value={"tables":[table]}),patch.object(engine,"_classify",side_effect=classify) as classify_call:
            with self.assertRaisesRegex(SemanticProgressError,"磁盘写入失败"):
                engine.recognize(data,progress=failed)
        self.assertEqual(classify_call.call_count,1)
        self.assertFalse(any("本包未完成" in log or "布局识别未完成" in log for log in logs))

    def test_progress_cancellation_flushes_below_threshold_and_preserves_pending_cells(self):
        from unittest.mock import patch
        event=threading.Event();engine=self.engine(cancel=event);data=snapshot();sheet=data["sheets"][0]
        sheet["cells"]={f"A{i}":{"row":i,"column":1,"value":i,"formula":None,"cached_value":i} for i in range(1,1201)}
        sheet.update(max_row=1200,max_column=1)
        table={"range":"A1:A1200","note_ids":["C-N001"],"scope":"consolidated","table_semantic":"测试余额","header_cells":[]}
        calls=[];saved=[]
        def classify(sheet,table,candidates,*args):
            engine._check_cancel();calls.append(list(candidates))
            result={"mappings":[{"sheet":sheet["name"],"cell":a,"reviewed":True} for a in candidates],"excluded":[],"unresolved":[]}
            if len(calls)==2:event.set()
            return result
        with patch.object(engine,"_request_layout",return_value={"tables":[table]}),patch.object(engine,"_classify",side_effect=classify):
            with self.assertRaises(SemanticCancelled):engine.recognize(data,progress=lambda snap,result:saved.append(copy.deepcopy(result)))
        self.assertGreaterEqual(len(saved),2)
        self.assertEqual(len(calls),2)
        self.assertGreater(len(saved[0]["mappings"]),0)
        self.assertLess(len(saved[-1]["mappings"]),1200)
        self.assertEqual({m['cell'] for m in saved[0]['mappings']},set(calls[0]))
        self.assertEqual({m['cell'] for m in saved[-1]['mappings']},set(calls[0]+calls[1]))
        self.assertEqual({m['cell'] for m in saved[-1]['unresolved']},set(sheet['cells'])-set(calls[0]+calls[1]))
        self.assertFalse(saved[-1]["complete"])
        self.assertEqual(len(saved[0]["mappings"]),len(calls[0]),"此前已发布的进度不能被后续内存修改")
        for item in saved:
            all_cells=[i["cell"] for key in ("mappings","excluded","unresolved") for i in item[key]]
            self.assertEqual(len(all_cells),1200);self.assertEqual(len(set(all_cells)),1200)

    def test_范围实际发送前拒绝超限初次和纠正请求(self):
        from unittest.mock import patch
        from 附注更新.识别范围 import ScopeRequestTooLarge
        for task in ("source_scope", "source_scope_batch"):
            for field in ("target_meanings", "previous_result"):
                with self.subTest(task=task,field=field):
                    payload={"task":task,"cells":{},"context_cells":{},"candidate_cells":[],"tasks":[],field:"原"*120001}
                    with patch("urllib.request.build_opener",side_effect=AssertionError("不应进入网络")) as network:
                        with self.assertRaises(ScopeRequestTooLarge):self.engine()._http("范围请求",payload)
                    network.assert_not_called()

    def test_范围预算允许正常请求且不影响其他任务(self):
        from unittest.mock import patch
        for payload in ({"task":"source_scope","cells":{},"context_cells":{},"candidate_cells":[]},
                        {"task":"tables","catalogue":[{"id":"C-N001-T001"}],"other":"原"*120001}):
            with self.subTest(task=payload["task"]):
                with patch("urllib.request.build_opener",side_effect=AssertionError("已经进入发送边界")) as network:
                    with self.assertRaisesRegex(AssertionError,"已经进入发送边界"):self.engine()._http("测试边界",payload)
                self.assertEqual(network.call_count,1)

    @classmethod
    def setUpClass(cls):
        cls.work = ROOT / "测试输出" / ("语义_" + uuid.uuid4().hex)
        cls.work.mkdir(parents=True)
        cls.gold = cls.work / "测试金标准.jsonl"
        row = {"id": "C-N001-T001-S001", "status": "active", "scope": "consolidated", "note": {"id": "C-N001", "name": "货币资金", "meaning": "货币资金披露"}, "table": {"id": "C-N001-T001", "name": "余额构成", "meaning": "货币资金构成"}, "slot": {"name": "银行存款 / 期末余额"}, "row_path": ["银行存款"], "column_path": ["期末余额"], "value_type": "monetary", "dimensions": [{"name": n, "required": True} for n in ("currency", "period", "scale", "unit")], "aggregation": {"role": "component"}}
        cls.gold.write_text(json.dumps(row, ensure_ascii=False), encoding="utf-8-sig")
        cls.server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), LocalAPI)
        cls.server.received = []
        cls.server.mode = "normal"
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join(2)

    def engine(self, name=None, cancel=None):
        return SemanticEngine({"base_url": f"http://127.0.0.1:{self.server.server_port}/v1", "model": "local-test", "api_key": "local-test-secret", "review": True, "timeout": 2}, str(self.gold), str(self.work / (name or uuid.uuid4().hex)), cancel=cancel)


    def cell_layout_fixture(self):
        data=snapshot("a"*64);sheet=data["sheets"][0]
        sheet.update(max_row=5,max_column=4)
        sheet["cells"]["C3"]={**sheet["cells"]["B3"],"column":3,"value":0,"cached_value":0}
        sheet["cells"]["D5"]={**sheet["cells"]["B3"],"row":5,"column":4,"value":None,"cached_value":None}
        tables={a:{"range":f"{a}:{a}","note_ids":["C-N001"],"scope":"consolidated",
                  "table_semantic":"银行存款余额","header_cells":["A1","A3","B2"]} for a in ("B3","D5")}
        sheet["structural_omissions"]=[{"range":f"{c}{r}:{c}{r}"} for r in range(1,6) for c in "ABCD" if f"{c}{r}" not in sheet["cells"]]
        return data,tables

    def cell_layout_reply(self, tables, calls, second_change=None):
        def reply(system,payload):
            calls.append(copy.deepcopy(payload))
            if payload["task"]=="layout_cells":
                plan={"tables":[],"non_business":[],"unresolved":[],"carry_context":"原表合并货币资金"}
                for address in payload["candidate_cells"]:
                    if address in tables:
                        table=copy.deepcopy(tables[address])
                        if second_change and payload["round"]==2:second_change(address,table)
                        plan["tables"].append(table)
                    elif address=="C3":
                        plan["unresolved"].append({"range":"C3:C3","reason":"另一业务尚不能确定","evidence_cells":["C3"]})
                    else:
                        plan["non_business"].append({"range":f"{address}:{address}","category":"header",
                            "reason":"原始文字表头","evidence_cells":[address]})
                return plan
            if payload["task"]=="tables":return {"table_ids":["C-N001-T001"]}
            if payload["task"]=="classify":
                return {"mappings":[{**mapping(),"cell":a} for a in payload["candidate_cells"]],"excluded":[],"unresolved":[]}
            raise AssertionError("逐格布局不能退回旧矩形请求")
        return reply

    def test_cells_layout_sends_original_kind_and_text_views_in_both_rounds(self):
        from unittest.mock import patch
        from 附注更新.回执格式 import response_format
        engine=self.engine();data=snapshot();sheet=data["sheets"][0]
        values={"G2":3,"D2":0,"A2":None,"C2":"——","B2":"","E2":"1,234.50","F2":"=SUM(D2:E2)",
                "G1":"行次","A1":"项目","H3":"真实披露文字"}
        sheet.update(max_row=3,max_column=8,cells={a:{"row":int(a[1:]),"column":ord(a[0])-64,"value":v,
            "formula":v if a=="F2" else None,"cached_value":1234.5 if a=="F2" else v,
            "number_format":"General","style_id":0} for a,v in values.items()})
        owned=["A2","B2","C2","D2","E2","F2","G2"]
        payload={"task":"layout_cells","sheet":sheet["name"],"candidate_cells":owned,"cells":sheet["cells"],
                 "notes":list(engine.notes.values()),"context":{"document":["原文上下文"]},"merges":["A1:F1"]}
        before=copy.deepcopy(payload);source_before=copy.deepcopy(data);calls=[]
        def reply(system,call):
            calls.append(copy.deepcopy(call))
            return {"tables":[],"non_business":[],"unresolved":[{"range":a+":"+a,"reason":"本地协议测试未判定业务",
                "evidence_cells":[a]} for a in owned]}
        with patch.object(engine,"_http",side_effect=reply):
            plan=engine._request_layout_cells(payload,sheet,owned,"原值视图测试")
        self.assertEqual([call["round"] for call in calls],[1,2])
        expected={"A2":{"kind":"blank"},"B2":{"kind":"blank"},"C2":{"kind":"text","text":"——"},
                  "D2":{"kind":"number"},"E2":{"kind":"number_text"},"F2":{"kind":"formula"},"G2":{"kind":"number"}}
        for number,call in enumerate(calls,1):
            self.assertEqual(call.get("candidate_kind_view"),expected)
            self.assertEqual(call.get("text_evidence_view"),[["A1","项目"],["G1","行次"],["C2","——"],["H3","真实披露文字"]])
            original={**before,"round":number}
            self.assertEqual({k:v for k,v in call.items() if k not in {"candidate_kind_view","text_evidence_view"}},original)
            self.assertEqual(response_format(call),response_format(original))
        self.assertEqual(payload,before)
        self.assertEqual(data,source_before)
        self.assertEqual({item["range"] for item in plan["unresolved"]},{a+":"+a for a in owned})
        self.assertEqual(plan["tables"],[])
        self.assertEqual(plan["non_business"],[])

    def test_cells_layout_views_do_not_reuse_legacy_cache_but_resume_new_requests(self):
        from unittest.mock import patch
        from 附注更新.语义 import LAYOUT_CELLS
        engine=self.engine();data=snapshot();sheet=data["sheets"][0]
        sheet.update(max_row=1,max_column=2,cells={"A1":{"row":1,"column":1,"value":"项目"},
                                                "B1":{"row":1,"column":2,"value":"行次"}})
        owned=["A1","B1"]
        payload={"task":"layout_cells","sheet":sheet["name"],"cells":sheet["cells"],
                 "candidate_cells":owned,"notes":list(engine.notes.values())}
        answer={"tables":[],"non_business":[{"range":a+":"+a,"category":"header","reason":"原始文字表头",
            "evidence_cells":[a]} for a in owned],"unresolved":[]}
        def validate(response):
            checked=engine._layout_cell_partition(response,owned)
            engine._validate_layout(checked,sheet,owned,single_note=True)
            return checked
        with patch.object(engine,"_http",return_value=answer):
            for number in (1,2):engine._request(LAYOUT_CELLS,{**payload,"round":number},validate)
        old_tasks={path:path.read_bytes() for path in (engine.work_dir/"语义任务").glob("*/*.json")}
        calls=[]
        def reply(system,call):
            calls.append(copy.deepcopy(call))
            return copy.deepcopy(answer)
        with patch.object(engine,"_http",side_effect=reply):
            first=engine._request_layout_cells(payload,sheet,owned,"缓存原值视图测试")
        self.assertEqual([call["round"] for call in calls],[1,2],"旧无视图请求必须与带原值视图请求分开保存")
        self.assertEqual(len(first["non_business"]),2)
        with patch.object(engine,"_http",side_effect=AssertionError("新视图双轮成果应从其自身缓存续用")):
            resumed=engine._request_layout_cells(payload,sheet,owned,"缓存原值视图测试")
        self.assertEqual(resumed,first)
        self.assertEqual(engine.usage["cached_requests"],2)
        self.assertEqual({path:path.read_bytes() for path in old_tasks},old_tasks)

    def test_cells_layout_batches_sparse_same_note_without_materializing_envelope(self):
        from unittest.mock import patch
        data,tables=self.cell_layout_fixture();before=copy.deepcopy(data);calls=[];engine=self.engine()
        engine.settings["layout_mode"]="cells_v1";events=[]
        with patch.object(engine,"_http",side_effect=self.cell_layout_reply(tables,calls)):
            result=engine.recognize(data,progress=lambda snap,part:events.append(copy.deepcopy(part)))
        self.assertEqual({i["cell"] for i in result["mappings"]},{"B3","D5"})
        self.assertEqual([p["candidate_cells"] for p in calls if p["task"]=="classify"],[["B3","D5"],["B3","D5"]])
        self.assertEqual([p["round"] for p in calls if p["task"]=="layout_cells"],[1,2])
        self.assertEqual(data["sheets"][0]["cells"],before["sheets"][0]["cells"])
        self.assertEqual({i["cell"] for i in result["unresolved"]},{"C3"})
        proofs=result["confirmed_business_ranges"]["表一"]
        self.assertEqual({p["table_range"] for p in proofs},{"B3:B3","D5:D5"})
        self.assertEqual(len({p["layout_batch_id"] for p in proofs}),1)
        for item in result["mappings"]:
            self.assertEqual(item["table_range"],item["cell"]+":"+item["cell"])
            self.assertEqual(len(item["layout_review"]["rounds"]),2)
            self.assertTrue(item["reviewed"])
        self.assertEqual(result["recognition_basis"]["layout_mode"],"cells_v1")
        self.assertEqual(len(events[-1]["mappings"])+len(events[-1]["excluded"])+len(events[-1]["unresolved"]),7)

    def test_cells_layout_conflict_and_failed_round_never_confirm_business(self):
        from unittest.mock import patch
        for failure in (False,True):
            with self.subTest(failure=failure):
                data,tables=self.cell_layout_fixture();engine=self.engine();engine.settings["layout_mode"]="cells_v1";calls=[]
                def change(address,table):
                    if failure:raise SemanticError("第二轮请求失败")
                    if address=="D5":table["scope"]="parent"
                with patch.object(engine,"_http",side_effect=self.cell_layout_reply(tables,calls,change)):
                    result=engine.recognize(data)
                self.assertEqual({i["cell"] for i in result["mappings"]},set() if failure else {"B3"})
                unresolved={i["cell"]:i for i in result["unresolved"]}
                self.assertIn("D5",unresolved)
                self.assertEqual(len(unresolved["D5"]["layout_review"]["rounds"]),2)
                self.assertFalse(unresolved["D5"].get("reviewed",False))

    def test_cells_layout_grouping_respects_physical_tables_and_per_cell_range(self):
        data,tables=self.cell_layout_fixture();sheet=data["sheets"][0];engine=self.engine()
        self.assertTrue(hasattr(engine,"_cell_table_groups"),"逐格布局须有明确格址合批入口")
        groups=engine._cell_table_groups(sheet,tables,["B3","D5"],"原包身份")
        self.assertEqual(len(groups),1)
        self.assertEqual(groups[0]["candidate_cells"],["B3","D5"])
        raw={"mappings":[{**mapping(),"cell":a} for a in ("B3","D5")],"excluded":[],"unresolved":[]}
        checked=engine._validate_classification(raw,["B3","D5"],set(engine.slots),sheet,groups[0])
        self.assertEqual(checked["B3"][1]["table_range"],"B3:B3")
        self.assertEqual(checked["D5"][1]["table_range"],"D5:D5")
        sheet["cells"]["A5"]={"row":5,"column":1,"value":"本表期末余额","formula":None,"cached_value":"本表期末余额"}
        tables["D5"]["header_cells"]=["A5"]
        sheet["word_table_ranges"]=[{"range":"A1:D3"},{"range":"A4:D5"}]
        self.assertEqual(len(engine._cell_table_groups(sheet,tables,["B3","D5"],"原包身份")),2)
        sheet.pop("word_table_ranges");tables["D5"]["scope"]="parent"
        self.assertEqual(len(engine._cell_table_groups(sheet,tables,["B3","D5"],"原包身份")),2)

    def test_cells_layout_no_applicable_scope_does_not_request_table_selection(self):
        from unittest.mock import patch
        data,tables=self.cell_layout_fixture();engine=self.engine();table=tables["B3"];table["scope"]="standalone"
        with patch.object(engine,"_http",side_effect=AssertionError("已知无适用定义不得请求模型")) as request:
            result=engine._classify(data["sheets"][0],table,["B3"],[],{})
        self.assertEqual(request.call_count,0)
        self.assertEqual([i["cell"] for i in result["unresolved"]],["B3"])

    def test_cells_layout_previous_legacy_mode_is_not_reused(self):
        from unittest.mock import patch
        data,first,engine,reply=self.resume_fixture("a"*64)
        self.assertNotIn("layout_mode",first["recognition_basis"])
        engine.settings["layout_mode"]="cells_v1"
        with patch.object(engine,"_http",side_effect=AssertionError("模式不同须先明确拒绝旧成果")):
            with self.assertRaisesRegex(SemanticError,"布局模式|识别.*变化"):
                engine.recognize(data,previous_result=first)


    def test_cells_layout_real_local_protocol_and_same_mode_resume_keep_cell_proofs(self):
        from unittest.mock import patch
        engine=self.engine();engine.settings["layout_mode"]="cells_v1"
        result=engine.recognize(snapshot("a"*64))
        self.assertTrue(result["complete"])
        self.assertEqual([i["cell"] for i in result["mappings"]],["B3"])
        self.assertEqual([p["task"] for p in self.server.received],["layout_cells","layout_cells","tables","classify","classify"])
        retry=self.engine();retry.settings["layout_mode"]="cells_v1"
        with patch.object(retry,"_http",side_effect=AssertionError("同模式已确认原格不得再请求")):
            resumed=retry.recognize(snapshot("a"*64),previous_result=result)
        self.assertTrue(resumed["complete"])
        self.assertEqual(resumed["mappings"],result["mappings"])
        self.assertEqual(resumed["confirmed_business_ranges"],result["confirmed_business_ranges"])

    def test_cells_layout_nonbusiness_conflict_and_cancel_keep_full_pending_coverage(self):
        from unittest.mock import patch
        for action in ("conflict","cancel"):
            data,tables=self.cell_layout_fixture();engine=self.engine();engine.settings["layout_mode"]="cells_v1"
            calls=[];base=self.cell_layout_reply(tables,calls);progress=[]
            def reply(system,payload):
                result=base(system,payload)
                if payload["task"]=="layout_cells" and payload["round"]==2:
                    if action=="cancel":raise SemanticCancelled("本地取消")
                    for region in result["non_business"]:
                        if region["range"]=="A1:A1":region["category"]="annotation"
                return result
            with patch.object(engine,"_http",side_effect=reply):
                if action=="cancel":
                    with self.assertRaises(SemanticCancelled):
                        engine.recognize(data,progress=lambda snap,part:progress.append(copy.deepcopy(part)))
                    self.assertEqual(len(progress[-1]["unresolved"]),7)
                    self.assertEqual(progress[-1]["mappings"],[])
                else:
                    result=engine.recognize(data)
                    self.assertIn("A1",{i["cell"] for i in result["unresolved"]})
                    self.assertNotIn("A1",{i["cell"] for i in result["excluded"]})

    def test_cells_layout_limit_and_grouping_do_not_mix_notes_or_lose_shared_headers(self):
        from unittest.mock import patch
        data,tables=self.cell_layout_fixture();engine=self.engine();engine.settings["layout_mode"]="cells_v1"
        sheet=data["sheets"][0];sheet["max_row"]=330
        sheet["cells"]={"A1":sheet["cells"]["A1"],"B3":sheet["cells"]["B3"],**{
            f"C{r}":{"row":r,"column":3,"value":"原行标签","formula":None,"cached_value":"原行标签"} for r in range(4,243)},
            "B330":{**sheet["cells"]["B3"],"row":330}}
        sheet["structural_omissions"]=[{"range":f"{c}{r}:{c}{r}"} for r in range(1,331) for c in "ABCD" if f"{c}{r}" not in sheet["cells"]]
        cell_tables={a:{"range":f"{a}:{a}","note_ids":["C-N001"],"scope":"consolidated",
                       "table_semantic":"银行存款余额","header_cells":["A1"]} for a in ("B3","B330")}
        calls=[]
        with patch.object(engine,"_http",side_effect=self.cell_layout_reply(cell_tables,calls)):
            result=engine.recognize(data)
        self.assertEqual({i["cell"] for i in result["mappings"]},{"B3","B330"})
        layout=[p for p in calls if p["task"]=="layout_cells"]
        self.assertTrue(all(len(p["candidate_cells"])<=40 for p in layout))
        last=next(p for p in layout if "B330" in p["candidate_cells"])
        self.assertIn("A1",last["cells"])
        self.assertEqual(last["cells"]["A1"],sheet["cells"]["A1"])
        other=copy.deepcopy(cell_tables["B330"]);other["note_ids"]=["C-N002"]
        engine.notes["C-N002"]={**engine.notes["C-N001"],"id":"C-N002"}
        groups=engine._cell_table_groups(sheet,{"B3":cell_tables["B3"],"B330":other},["B3","B330"],"同一原包")
        self.assertEqual(len(groups),2)


    def test_cells_layout_rejects_header_from_other_word_table_but_keeps_external_title(self):
        from unittest.mock import patch
        data,tables=self.cell_layout_fixture();sheet=data["sheets"][0];engine=self.engine()
        sheet["word_table_ranges"]=[{"range":"A2:D3"},{"range":"A4:D5"}]
        # A1不属于另一张物理表，可作为原表外标题；A3明确属于第一张表。
        tables["D5"]["header_cells"]=["A1"]
        groups=engine._cell_table_groups(sheet,{"D5":tables["D5"]},["D5"],"原包")
        self.assertEqual(groups[0]["cell_tables"]["D5"]["header_cells"],["A1"])
        tables["D5"]["header_cells"]=["A3"]
        with self.assertRaisesRegex(SemanticError,"另一张|跨Word|跨.*物理表"):
            engine._cell_table_groups(sheet,{"D5":tables["D5"]},["D5"],"原包")
        payload={"task":"layout_cells","sheet":sheet["name"],"candidate_cells":["D5"],"cells":sheet["cells"],"notes":list(engine.notes.values())}
        with patch.object(engine,"_http",return_value={"tables":[tables["D5"]],"non_business":[],"unresolved":[]}):
            plan=engine._request_layout_cells(payload,sheet,["D5"],"原包")
        self.assertEqual(plan["tables"],[])
        self.assertEqual([x["range"] for x in plan["unresolved"]],["D5:D5"])
        self.assertEqual(len(plan["unresolved"][0]["layout_review"]["rounds"]),2)

    def test_cells_layout_recovers_only_verified_cells_with_independent_round_proofs(self):
        from unittest.mock import patch
        import hashlib
        data,tables=self.cell_layout_fixture();sheet=data["sheets"][0];engine=self.engine();calls=[]
        base=self.cell_layout_reply(tables,calls)
        def reply(system,payload):
            result=base(system,payload)
            result["unresolved"]=[]
            result["non_business"].append({"range":"C3:C3","category":"empty_padding","reason":"误排零金额","evidence_cells":["C3"]})
            if payload["round"]==2:
                next(item for item in result["non_business"] if item["range"]=="A1:A1")["category"]="empty_padding"
            return result
        payload={"task":"layout_cells","cells":sheet["cells"],"candidate_cells":list(sheet["cells"]),"notes":list(engine.notes.values())}
        with patch.object(engine,"_http",side_effect=reply):
            result=engine._request_layout_cells(payload,sheet,list(sheet["cells"]),"原包")
        self.assertEqual(len(calls),4)
        self.assertEqual({item["range"] for item in result["tables"]},{"B3:B3","D5:D5"})
        self.assertEqual({item["range"] for item in result["non_business"]},{"A2:A2","B2:B2","A3:A3"})
        self.assertEqual({item["range"] for item in result["unresolved"]},{"A1:A1","C3:C3"})
        for item in result["non_business"]:
            self.assertTrue(item["reviewed"])
            self.assertEqual([r["round"] for r in item["layout_review"]["rounds"]],[1,2])
            for round_proof in item["layout_review"]["rounds"]:
                self.assertEqual(round_proof["record_kind"],"non_business")
                self.assertEqual(len(round_proof["receipts"]),1)
                path=pathlib.Path(round_proof["receipts"][0]["path"])
                saved=json.loads(path.read_text(encoding="utf-8-sig"))
                self.assertTrue(saved["retry_required"])
                self.assertEqual(saved["source_evidence"]["validation_failures"],2)
                source=saved["source_evidence"]
                for key in ("response","correction_request"):
                    self.assertEqual(hashlib.sha256(pathlib.Path(source[key+"_path"]).read_bytes()).hexdigest(),source[key+"_sha256"])
                self.assertFalse(pathlib.Path(source["response_path"]).with_name("回执.json").exists())
        self.assertTrue(all(not item.get("reviewed") for item in result["unresolved"]))

    def test_cells_layout_recovery_keeps_word_header_guard_and_global_partition(self):
        from unittest.mock import patch
        for fault in ("other_word_header","missing","duplicate","outside","rectangle","forged_evidence"):
            with self.subTest(fault=fault):
                data,tables=self.cell_layout_fixture();sheet=data["sheets"][0];engine=self.engine()
                sheet["word_table_ranges"]=[{"range":"A2:D3"},{"range":"A4:D5"}]
                tables["D5"]["header_cells"]=["A3"]
                owned=["A1","D5"]
                response={"tables":[tables["D5"]],"non_business":[{"range":"A1:A1","category":"title","reason":"原表外标题","evidence_cells":["A1"]}],"unresolved":[]}
                if fault=="missing":response["tables"]=[];response["non_business"][0]["category"]="empty_padding"
                if fault=="duplicate":response["non_business"].append(copy.deepcopy(response["non_business"][0]))
                if fault=="outside":response["non_business"][0]["range"]="A2:A2"
                if fault=="rectangle":response["non_business"][0]["range"]="A1:A2"
                if fault=="forged_evidence":response["non_business"][0]["evidence_cells"]=["Z999"]
                payload={"task":"layout_cells","cells":sheet["cells"],"candidate_cells":owned,"notes":list(engine.notes.values())}
                with patch.object(engine,"_http",return_value=response):result=engine._request_layout_cells(payload,sheet,owned,"原包")
                self.assertEqual(result["tables"],[])
                if fault=="other_word_header":
                    self.assertEqual([item["range"] for item in result["non_business"]],["A1:A1"])
                    self.assertEqual([item["range"] for item in result["unresolved"]],["D5:D5"])
                else:
                    self.assertEqual(result["non_business"],[])
                    self.assertEqual({item["range"] for item in result["unresolved"]},{"A1:A1","D5:D5"})
                    self.assertTrue(all(r["request_evidence"].get("files") for r in result["unresolved"][0]["layout_review"]["rounds"]))

    def test_cells_layout_recovery_rejects_changed_correction_chain(self):
        from unittest.mock import patch
        data,tables=self.cell_layout_fixture();sheet=data["sheets"][0];engine=self.engine();owned=list(sheet["cells"])
        response=self.cell_layout_reply(tables,[])("",{"task":"layout_cells","candidate_cells":owned,"round":1})
        response["unresolved"]=[]
        response["non_business"].append({"range":"C3:C3","category":"empty_padding","reason":"误排零金额","evidence_cells":["C3"]})
        original=copy.deepcopy(response)
        payload={"task":"layout_cells","cells":sheet["cells"],"candidate_cells":owned,"round":1,"notes":list(engine.notes.values())}
        def validate(raw):return engine._validate_layout(engine._layout_cell_partition(raw,owned),sheet,owned,single_note=True)
        with patch.object(engine,"_http",return_value=response):
            with self.assertRaises(SemanticError) as caught:engine._request("本地固定格布局",payload,validate)
        failure=caught.exception.layout_failure
        self.assertEqual(failure["result"],original)
        engine._recover_layout(failure["result"],sheet,owned,failure["evidence"])
        self.assertEqual(response,original,"不能剥掉错误字段后把原结果冒充通过")
        for field in ("response_sha256","correction_request_sha256"):
            evidence=copy.deepcopy(failure["evidence"]);evidence[field]="0"*64
            with self.subTest(field=field),self.assertRaises(SemanticError):
                engine._recover_layout(failure["result"],sheet,owned,evidence)

    def test_cells_layout_parse_failure_cannot_supply_two_business_failures(self):
        from unittest.mock import patch
        from 附注更新.语义 import SemanticResponseError
        data,tables=self.cell_layout_fixture();sheet=data["sheets"][0];engine=self.engine()
        base=self.cell_layout_reply(tables,[])
        def reply(system,payload):
            if "correction" not in payload:raise SemanticResponseError(self.work/"合成格式错误.json",{},"未通过原协议")
            result=base(system,payload);result["unresolved"]=[]
            result["non_business"].append({"range":"C3:C3","category":"empty_padding","reason":"误排零金额","evidence_cells":["C3"]})
            return result
        payload={"task":"layout_cells","cells":sheet["cells"],"candidate_cells":list(sheet["cells"]),"notes":list(engine.notes.values())}
        with patch.object(engine,"_http",side_effect=reply),patch.object(engine,"_recover_layout",side_effect=AssertionError("仅一次业务校验失败不得恢复")) as recovery:
            result=engine._request_layout_cells(payload,sheet,list(sheet["cells"]),"原包")
        recovery.assert_not_called()
        self.assertEqual(result["tables"],[]);self.assertEqual(result["non_business"],[])
        self.assertEqual(len(result["unresolved"]),len(sheet["cells"]))

    def test_cells_layout_replaces_same_cell_proof_without_erasing_old_saved_evidence(self):
        data,tables=self.cell_layout_fixture();sheet=data["sheets"][0];engine=self.engine()
        first={**tables["B3"],"layout_batch_id":"原包","layout_review":{"rounds":[{"round":1},{"round":2}]}}
        engine._remember_cell_table(data,sheet,first);old=copy.deepcopy(sheet["confirmed_business_ranges"])
        revised={**first,"note_ids":["C-N002"],"scope":"parent","layout_review":{"rounds":[{"round":1,"reason":"重新核验"},{"round":2}]}}
        engine.notes["C-N002"]={**engine.notes["C-N001"],"id":"C-N002"}
        engine._remember_cell_table(data,sheet,revised)
        engine._remember_cell_table(data,sheet,revised)
        proofs=sheet["confirmed_business_ranges"]
        self.assertEqual(len(proofs),1)
        self.assertEqual(proofs[0]["note_ids"],["C-N002"])
        self.assertEqual(proofs[0]["scope"],"parent")
        self.assertEqual(old[0]["note_ids"],["C-N001"])
        self.assertEqual(old[0]["layout_review"]["rounds"][0],{"round":1})

    def test_cells_layout_classification_cannot_borrow_other_cell_evidence(self):
        data,tables=self.cell_layout_fixture();sheet=data["sheets"][0];engine=self.engine()
        tables["B3"]["header_cells"]=["A3"];tables["D5"]["header_cells"]=["A1","B2"]
        group=engine._cell_table_groups(sheet,tables,["B3","D5"],"原包")[0]
        for field in ("evidence_cells","metric_evidence"):
            item={**mapping(),field:["A1"]}
            raw={"mappings":[item],"excluded":[],"unresolved":[]}
            with self.subTest(field=field):
                with self.assertRaisesRegex(SemanticError,"原格|独立.*证据|业务.*证据"):
                    engine._validate_classification(raw,["B3"],set(engine.slots),sheet,group)
            item[field]=["A3"]
            checked=engine._validate_classification(raw,["B3"],set(engine.slots),sheet,group)
            self.assertEqual(checked["B3"][1]["table_range"],"B3:B3")

    def test_cells_layout_pending_resume_retains_one_current_proof_per_cell(self):
        from unittest.mock import patch
        data,tables=self.cell_layout_fixture();engine=self.engine();engine.settings["layout_mode"]="cells_v1"
        calls=[];base=self.cell_layout_reply(tables,calls)
        def first_reply(system,payload):
            reply=base(system,payload)
            if payload["task"]=="classify":
                reply["mappings"]=[m for m in reply["mappings"] if m["cell"]!="D5"]
                reply["unresolved"]=[{"cell":"D5","reason":"首次缺少业务依据"}] if "D5" in payload["candidate_cells"] else []
            return reply
        with patch.object(engine,"_http",side_effect=first_reply):first=engine.recognize(data)
        retry=self.engine();retry.settings["layout_mode"]="cells_v1";fresh,_=self.cell_layout_fixture()
        with patch.object(retry,"_http",side_effect=self.cell_layout_reply(tables,[])):
            result=retry.recognize(fresh,previous_result=first)
        self.assertEqual({i["cell"] for i in result["mappings"]},{"B3","D5"})
        self.assertEqual(len(result["confirmed_business_ranges"]["表一"]),2)
        self.assertEqual({p["table_range"] for p in result["confirmed_business_ranges"]["表一"]},{"B3:B3","D5:D5"})


    def test_cells_layout_source_grid_covers_real_unformatted_blank_after_scope_check(self):
        from unittest.mock import patch
        from 附注更新.语义 import _hash
        from 附注更新.表格 import read_workbook,actual_value
        path,area=self.blank_fixture("逐格布局真实无样式空白")
        data=read_workbook(path);self.assertNotIn("B3",data["sheets"][0]["cells"])
        engine=self.engine();engine.settings["layout_mode"]="cells_v1";calls=[]
        tables={a:{"range":f"{a}:{a}","note_ids":["C-N001"],"scope":"consolidated",
                   "table_semantic":"业务金额明细","header_cells":["A1","B1","A"+a[1:]]} for a in ("B2","B3")}
        target={"chapter":"本地测试目标"}
        selection={"source_hash":data["sha256"],"target_hash":_hash(target),
                   "selected_cells":{"表一":list(data["sheets"][0]["cells"])},"out_of_scope":[]}
        with patch("附注更新.识别范围.select_source_scope",return_value=selection),patch.object(engine,"_http",side_effect=self.cell_layout_reply(tables,calls)):
            result=engine.recognize(data,target_scope=target)
        self.assertEqual({i["cell"] for i in result["mappings"]},{"B2","B3"})
        self.assertTrue(result["complete"])
        self.assertIsNone(actual_value(data,"表一","B3"))
        proof=data["sheets"][0]["cells"]["B3"]["blank_evidence"]
        self.assertEqual(proof["kind"],"source_grid")
        self.assertFalse(any(k in proof for k in ("note_ids","scope","table_semantic")))
        self.assertEqual(result["recognition_basis"]["candidate_blanks"],"source_grid_v1")
        self.assertEqual({p["table_range"] for p in result["confirmed_business_ranges"]["表一"]},{"B2:B2","B3:B3"})
        self.assertIn("B3",{a for p in calls if p["task"]=="layout_cells" for a in p["candidate_cells"]})

    def test_cells_layout_previous_without_source_grid_strategy_requires_reidentification(self):
        import hashlib
        from unittest.mock import patch
        data,tables=self.cell_layout_fixture();engine=self.engine();engine.settings["layout_mode"]="cells_v1"
        with patch.object(engine,"_http",side_effect=self.cell_layout_reply(tables,[])):
            first=engine.recognize(data)
        first["recognition_basis"].pop("candidate_blanks",None)
        old=self.work/("旧候选覆盖策略_"+uuid.uuid4().hex+".json")
        raw={k:v for k,v in first.items() if k not in ("evidence_path","evidence_hash")}
        old.write_text(json.dumps(raw,ensure_ascii=False),encoding="utf-8-sig")
        prior={**raw,"evidence_path":str(old),"evidence_hash":hashlib.sha256(old.read_bytes()).hexdigest()}
        fresh,_=self.cell_layout_fixture();retry=self.engine();retry.settings["layout_mode"]="cells_v1"
        with patch.object(retry,"_http",side_effect=AssertionError("候选策略改变须先拒绝旧接续")):
            with self.assertRaisesRegex(SemanticError,"候选.*变化|覆盖.*变化"):
                retry.recognize(fresh,previous_result=prior)

    def setUp(self):
        self.server.mode = "normal"
        self.server.received.clear()
        self.server.bodies=[]


    def resume_fixture(self, source_hash="source-one"):
        from unittest.mock import patch
        data=snapshot(source_hash);sheet=data["sheets"][0];sheet["max_row"]=4
        for address,value in (("A4","新的业务指标"),("B4",67)):
            sheet["cells"][address]={**sheet["cells"]["B3"],"row":4,"column":1 if address=="A4" else 2,"value":value,"cached_value":value}
        first_engine=self.engine();calls=[]
        def reply(system,payload):
            calls.append(copy.deepcopy(payload))
            if payload["task"]=="layout":return {"tables":[{"range":"A1:B4","note_ids":["C-N001"],"scope":"consolidated","table_semantic":"货币资金","header_cells":["A1","A2","B2"]}]}
            if payload["task"]=="tables":return {"table_ids":["C-N001-T001"]}
            owned=payload["candidate_cells"]
            return {"mappings":[mapping()] if "B3" in owned else [],"excluded":[{"cell":a,"category":"label","reason":"原始文字标签"} for a in owned if a not in {"B3","B4"}],"unresolved":[{"cell":"B4","reason":"需补充真正的新指标"}] if "B4" in owned else []}
        with patch.object(first_engine,"_http",side_effect=reply):first=first_engine.recognize(data)
        new_gold=self.work/("新增定义_"+uuid.uuid4().hex+".jsonl")
        old=json.loads(self.gold.read_text(encoding="utf-8-sig"));added=copy.deepcopy(old)
        added["id"]="C-N001-T001-S002";added["slot"]["name"]="新的业务指标";added["row_path"]=["新的业务指标"]
        new_gold.write_text("\n".join(json.dumps(row,ensure_ascii=False) for row in [old,added]),encoding="utf-8-sig")
        engine=SemanticEngine(first_engine.settings,str(new_gold),str(self.work/uuid.uuid4().hex))
        return data,first,engine,reply

    def test_左右科目子区段共享表头且原格不重复分类(self):
        from unittest.mock import patch
        from collections import Counter
        from 附注更新.语义 import _inside, _range
        engine=self.engine();first=next(iter(engine.slots.values()));second=copy.deepcopy(first)
        second.update(id="C-N002-T001-S001",note={"id":"C-N002","name":"短期借款"},table={"id":"C-N002-T001","name":"余额构成"},slot={"name":"短期借款余额"},row_path=["短期借款"])
        engine.slots[second["id"]]=second
        engine.notes["C-N002"]={"id":"C-N002","name":"短期借款"}
        engine.tables["C-N002-T001"]={"id":"C-N002-T001","note_id":"C-N002","name":"余额构成"}
        values={"A1":"2025年，合并报表，人民币元","A2":"项目","B2":"行次","C2":"期末余额","D2":"期初余额",
                "E2":"项目","F2":"行次","G2":"期末余额","H2":"期初余额",
                "A3":"银行存款","B3":1,"C3":100,"D3":90,"E3":"短期借款","F3":2,"G3":20,"H3":10}
        sheet={"name":"左右列组","max_row":3,"max_column":8,"merges":["A1:H1"],
               "cells":{a:{"row":int(a[1:]),"column":ord(a[0])-64,"value":v,"cached_value":v,"formula":None} for a,v in values.items()}}
        tables=[{"range":"A3:D3","note_ids":["C-N001"],"scope":"consolidated","table_semantic":"货币资金","header_cells":["A1","A2","B2","C2","D2"]},
                {"range":"E3:H3","note_ids":["C-N002"],"scope":"consolidated","table_semantic":"短期借款","header_cells":["A1","E2","F2","G2","H2"]}]
        plan={"tables":tables,"non_business":[{"range":"A1:H2","category":"header","reason":"原始共享文字表头","evidence_cells":["A1","A2","B2","C2","D2","E2","F2","G2","H2"]}]}
        engine._validate_layout(plan,sheet,list(values));shared=engine._nonbusiness_cells(plan,sheet,list(values))
        calls=[];accepted=[]
        def reply(system,payload):
            calls.append(copy.deepcopy(payload))
            if payload["task"]=="tables":return {"table_ids":[payload["catalogue"][0]["id"]]}
            table=payload["table"];slot=payload["gold_slots"][0];left=table["range"]=="A3:D3";line="B3" if left else "F3";label="A3" if left else "E3"
            mappings=[];excluded=[]
            for address in payload["candidate_cells"]:
                if address in {line,label}:
                    item={"cell":address,"category":"label","reason":"原始项目或有明确行次列头的编号"}
                    if address==line:item["evidence_cells"]=["B2" if left else "F2"]
                    excluded.append(item)
                else:
                    item=mapping();item.update(cell=address,slot_id=slot["id"],row_label=values[label],column_label="期末余额" if address in {"C3","G3"} else "期初余额")
                    item["dimensions"]["period"]="2025-12-31" if address in {"C3","G3"} else "2024-12-31"
                    mappings.append(item)
            return {"mappings":mappings,"excluded":excluded,"unresolved":[]}
        with patch.object(engine,"_http",side_effect=reply):
            for table in tables:
                owned=[a for a in values if _inside(a,_range(table["range"]))]
                result=engine._classify(sheet,table,owned,[],{"document":[]})
                self.assertFalse(result["unresolved"])
                accepted.extend(x["cell"] for k in ("mappings","excluded") for x in result[k])
        self.assertEqual(set(shared)|set(accepted),set(values));self.assertFalse(set(shared)&set(accepted))
        self.assertEqual(len(accepted),len(set(accepted)))
        for round_number in (1,2):
            counts=Counter(a for p in calls if p["task"]=="classify" and p["round"]==round_number for a in p["candidate_cells"])
            self.assertEqual(counts,Counter(accepted))
        missing=[(p["task"],p["table"]["range"],a) for p in calls for a in p["table"]["header_cells"] if a not in p["cells"]]
        self.assertEqual(missing,[],"选表与逐格识别都必须看见子区段之外明确引用的共享原表头")
        for payload in calls:
            expected=payload["table"]["note_ids"][0]
            if payload["task"]=="tables":self.assertEqual({x["note_id"] for x in payload["catalogue"]},{expected})
            else:
                self.assertEqual({x["note"]["id"] for x in payload["gold_slots"]},{expected})
                self.assertTrue(set(payload["candidate_cells"]).isdisjoint(shared))

    def test_appended_gold_only_requests_pending_cells_and_keeps_verified_results(self):
        from unittest.mock import patch
        data,first,engine,reply=self.resume_fixture();calls=[]
        before=copy.deepcopy(first)
        def retry(system,payload):
            calls.append(copy.deepcopy(payload))
            result=reply(system,payload)
            if payload["task"]=="classify":
                self.assertEqual(payload["candidate_cells"],["B4"])
                item={**mapping(),"cell":"B4","slot_id":"C-N001-T001-S002","semantic_field":"新的业务指标"}
                result={"mappings":[item],"excluded":[],"unresolved":[]}
            return result
        with patch.object(engine,"_http",side_effect=retry):result=engine.recognize(data,previous_result=first)
        self.assertTrue(result["complete"])
        self.assertEqual({m["cell"] for m in result["mappings"]},{"B3","B4"})
        self.assertEqual(next(m for m in result["mappings"] if m["cell"]=="B3"),first["mappings"][0])
        self.assertEqual(first,before)
        self.assertEqual(result["reuse"]["counts"]["mappings"],1)
        self.assertEqual(result["reuse"]["previous_gold_hash"],first["gold_hash"])
        self.assertTrue(all("B3" not in p.get("candidate_cells",[]) for p in calls))

    def test_failed_pending_retry_does_not_replace_verified_mapping(self):
        from unittest.mock import patch
        data,first,engine,reply=self.resume_fixture()
        with patch.object(engine,"_http",side_effect=ValueError("新包仍失败")):
            result=engine.recognize(data,previous_result=first)
        self.assertFalse(result["complete"])
        self.assertEqual(result["mappings"],first["mappings"])
        self.assertEqual({m["cell"] for m in result["unresolved"]},{"B4"})

    def test_previous_result_requires_unchanged_source_context_and_evidence(self):
        from unittest.mock import patch
        data,first,engine,reply=self.resume_fixture()
        for change in ("context","cell","result","duplicate"):
            with self.subTest(change=change):
                candidate=copy.deepcopy(data);prior=copy.deepcopy(first)
                if change=="context":candidate["context"]=["已改为另一年度或另一企业"]
                if change=="cell":candidate["sheets"][0]["cells"]["B3"]["value"]=999
                if change=="result":prior["mappings"][0]["dimensions"]["period"]="2024-12-31"
                if change=="duplicate":prior["excluded"].append(copy.deepcopy(prior["excluded"][0]))
                with patch.object(engine,"_http",side_effect=AssertionError("绑定失败不得请求")):
                    with self.assertRaises((ValueError,SemanticError)):
                        engine.recognize(candidate,previous_result=prior)

    def test_partial_validated_response_is_saved_but_retried(self):
        engine=self.engine();payload={"task":"classify","candidate_cells":["B3"]}
        from unittest.mock import patch
        result={"mappings":[],"excluded":[],"unresolved":[{"cell":"B3","reason":"尚未确定"}]}
        with patch.object(engine,"_http",return_value=result) as request:
            self.assertEqual(engine._request("测试部分成果",payload,lambda r:r),result)
            receipts=list(engine.work_dir.rglob("回执.json"));self.assertEqual(len(receipts),1)
            saved=json.loads(receipts[0].read_text(encoding="utf-8-sig"))
            self.assertTrue(saved["retry_required"]);self.assertEqual(saved["result"],result)
            engine._request("测试部分成果",payload,lambda r:r)
            self.assertEqual(request.call_count,2)

    def test_real_text_header_label_disagreement_is_not_unresolved(self):
        from unittest.mock import patch
        for label in ("账面余额","减：坏账准备","账面价值","合计"):
            data=snapshot();data["sheets"][0]["cells"]["A3"]["value"]=label
            engine=self.engine()
            def reply(system,payload):
                if payload["task"]=="layout":return {"tables":[{"range":"A1:B3","note_ids":["C-N001"],"scope":"consolidated","table_semantic":"货币资金","header_cells":["A1","A2","B2"]}]}
                if payload["task"]=="tables":return {"table_ids":["C-N001-T001"]}
                return {"mappings":[mapping()],"excluded":[{"cell":a,"category":"label" if a=="A3" and payload["round"]==2 else "header","reason":"原表明确文字标签"} for a in payload["candidate_cells"] if a!="B3"],"unresolved":[]}
            with patch.object(engine,"_http",side_effect=reply):result=engine.recognize(data)
            self.assertTrue(result["complete"],label)

    def test_used_definition_change_rechecks_only_affected_mapping(self):
        from unittest.mock import patch
        data,first,engine,reply=self.resume_fixture()
        rows=[json.loads(line) for line in pathlib.Path(engine.gold_path).read_text(encoding="utf-8-sig").splitlines()]
        rows[0]["slot"]["name"]="银行存款 / 修订后的期末余额"
        revised=self.work/("修订定义_"+uuid.uuid4().hex+".jsonl")
        revised.write_text("\n".join(json.dumps(r,ensure_ascii=False) for r in rows),encoding="utf-8-sig")
        engine=SemanticEngine(engine.settings,str(revised),str(self.work/uuid.uuid4().hex));calls=[]
        def respond(system,payload):
            calls.append(copy.deepcopy(payload));return reply(system,payload)
        with patch.object(engine,"_http",side_effect=respond):result=engine.recognize(data,previous_result=first)
        self.assertFalse(result["complete"])
        self.assertEqual(result["reuse"]["counts"]["mappings"],0)
        self.assertTrue(any(x["cell"]=="B3" for x in result["reuse"]["recheck"]))
        self.assertEqual({a for p in calls if p["task"]=="classify" for a in p["candidate_cells"]},{"B3","B4"})
        self.assertEqual(result["mappings"][0]["semantic_field"],"银行存款 / 修订后的期末余额")

    def test_old_formal_mapping_reference_reuses_blank_evidence_without_rewriting(self):
        import hashlib
        from unittest.mock import patch
        from 附注更新.表格 import read_workbook
        from 附注更新.分步流程 import _record
        path,area=self.blank_fixture("旧正式成果复用")
        snapshot_before=read_workbook(path);first_engine=self.engine()
        with patch.object(first_engine,"_http",side_effect=self.blank_reply(area,[])):
            first=first_engine.recognize(snapshot_before)
        record=_record(snapshot_before,first,"A",self.gold,path,[])
        for key in ("recognition_basis","evidence_path","evidence_hash"):
            record.pop(key,None)
        record["mappings"][0]["dimensions"]["metric"]="银行存款 / 期末余额"
        previous=self.work/("旧版正式映射_"+uuid.uuid4().hex+".json")
        previous.write_text(json.dumps(record,ensure_ascii=False),encoding="utf-8-sig")
        original=previous.read_bytes();ref={"mapping_path":str(previous),"mapping_sha256":hashlib.sha256(original).hexdigest()}
        engine=self.engine();fresh=read_workbook(path)
        self.assertNotIn("B3",fresh["sheets"][0]["cells"])
        with patch.object(engine,"_http",side_effect=AssertionError("已确认成果不得再请求")):
            result=engine.recognize(fresh,previous_result=ref)
        self.assertTrue(result["complete"])
        self.assertEqual(result["reuse"]["counts"]["mappings"],2)
        self.assertEqual(len(result["mappings"])+len(result["excluded"]),len(fresh["sheets"][0]["cells"]))
        self.assertIn("B3",fresh["sheets"][0]["cells"])
        self.assertEqual(previous.read_bytes(),original)
        self.assertNotIn("metric",result["mappings"][0]["dimensions"])
        with patch.object(engine,"_http",side_effect=AssertionError("错误哈希不得请求")):
            with self.assertRaisesRegex(SemanticError,"哈希"):
                engine.recognize(read_workbook(path),previous_result={**ref,"mapping_sha256":"错误"})

    def test_legacy_numeric_header_without_evidence_is_reidentified(self):
        import hashlib
        from unittest.mock import patch
        from openpyxl import load_workbook
        from 附注更新.表格 import read_workbook
        from 附注更新.分步流程 import _record
        path,area=self.blank_fixture("旧年份表头缺证据")
        book=load_workbook(path);book.active["A1"]="2025年度";book.active["B1"]=2025
        book.save(path);book.close()
        engine=self.engine();calls=[];base=self.blank_reply(area,calls)
        def reply(system,payload):
            result=base(system,payload)
            if payload["task"]=="layout":result["tables"][0]["header_cells"].append("B1")
            if payload["task"]=="classify":
                for item in result["excluded"]:
                    if item["cell"]=="B1":item.update(category="header",evidence_cells=["A1"])
            return result
        data=read_workbook(path)
        with patch.object(engine,"_http",side_effect=reply):first=engine.recognize(data)
        self.assertTrue(first["complete"])
        record=_record(data,first,"A",self.gold,path,[])
        for key in ("recognition_basis","evidence_path","evidence_hash"):record.pop(key,None)
        for item in record["excluded"]:
            if item["cell"]=="B1":item.pop("evidence_cells",None)
        prior=self.work/("旧年头映射_"+uuid.uuid4().hex+".json")
        prior.write_text(json.dumps(record,ensure_ascii=False),encoding="utf-8-sig")
        engine=self.engine();calls.clear()
        with patch.object(engine,"_http",side_effect=reply):
            result=engine.recognize(read_workbook(path),previous_result={"mapping_path":str(prior),"mapping_sha256":hashlib.sha256(prior.read_bytes()).hexdigest()})
        self.assertTrue(result["complete"])
        self.assertEqual({a for p in calls if p["task"]=="classify" for a in p["candidate_cells"]},{"B1"})
        self.assertEqual([i["cell"] for i in result["reuse"]["recheck"]],["B1"])
        self.assertEqual(result["reuse"]["counts"]["mappings"],2)

    def test_formal_single_round_exclusion_is_not_reused_under_double_review(self):
        import hashlib
        from unittest.mock import patch
        from 附注更新.表格 import read_workbook
        from 附注更新.分步流程 import _record
        path,area=self.blank_fixture("旧排除尚未双审")
        engine=self.engine();data=read_workbook(path);calls=[];reply=self.blank_reply(area,calls)
        with patch.object(engine,"_http",side_effect=reply):first=engine.recognize(data)
        record=_record(data,first,"A",self.gold,path,[])
        for item in record["excluded"]:
            if item["cell"]=="A2":item["reviewed"]=False
        prior=self.work/("未双审排除_"+uuid.uuid4().hex+".json")
        prior.write_text(json.dumps(record,ensure_ascii=False),encoding="utf-8-sig")
        engine=self.engine();calls.clear()
        with patch.object(engine,"_http",side_effect=reply):
            result=engine.recognize(read_workbook(path),previous_result={"mapping_path":str(prior),"mapping_sha256":hashlib.sha256(prior.read_bytes()).hexdigest()})
        self.assertTrue(result["complete"])
        self.assertEqual({a for p in calls if p["task"]=="classify" for a in p["candidate_cells"]},{"A2"})
        self.assertEqual([i["cell"] for i in result["reuse"]["recheck"]],["A2"])

    def test_scope_resume_metadata_does_not_invalidate_verified_cells(self):
        from unittest.mock import patch
        from 附注更新.语义 import _hash
        data=snapshot();target={"disclosures":["货币资金"]}
        selection={"source_hash":"source-one","target_hash":_hash(target),"gold_hash":"旧版本",
                   "selected_cells":{"表一":list(data["sheets"][0]["cells"])},"out_of_scope":[]}
        first_engine=self.engine()
        with patch("附注更新.识别范围.select_source_scope",return_value=selection):
            first=first_engine.recognize(data,target_scope=target)
        changed={**copy.deepcopy(selection),"gold_hash":"新版本","resume":{"reused_round_count":2}}
        engine=self.engine()
        with patch("附注更新.识别范围.select_source_scope",return_value=changed),patch.object(engine,"_http",side_effect=AssertionError("范围结论相同不得重跑成功格")):
            result=engine.recognize(data,target_scope=target,scope_selection=selection,previous_result=first)
        self.assertTrue(result["complete"]);self.assertEqual(result["mappings"],first["mappings"])

    def test_http_format_constraint_is_only_transport_and_success_cache_is_reused(self):
        engine=self.engine();data=snapshot()
        self.assertTrue(engine.recognize(data)["complete"])
        for body in self.server.bodies:
            payload=json.loads(body["messages"][1]["content"])
            self.assertNotIn("response_format",payload)
            mode=body["response_format"]["type"]
            self.assertEqual(mode,"json_schema" if payload["task"] in {"layout","classify","tables"} else "json_object")
        count=len(self.server.bodies)
        self.assertTrue(engine.recognize(snapshot())["complete"])
        self.assertEqual(len(self.server.bodies),count)

    def test_per_cell_http_transport_preserves_semantic_validation(self):
        self.server.mode="per_cell"
        result=self.engine().recognize(snapshot())
        self.assertTrue(result["complete"]);self.assertEqual(len(result["mappings"]),1)
        self.assertEqual(result["mappings"][0]["cell"],"B3")
        bodies=[b for b in self.server.bodies if json.loads(b["messages"][1]["content"])["task"]=="classify"]
        self.assertEqual(len(bodies),2)
        for body in bodies:
            self.assertIn('"cells"',body["messages"][0]["content"])
            schema=body["response_format"]["json_schema"]["schema"]
            self.assertIn("cells",schema["properties"])
            for branch in schema["$defs"]["cell_record"]["oneOf"]:
                self.assertEqual(next(iter(branch["properties"])),"kind",
                                 "HTTP不能重排属性，让模型先选状态时只剩待核实分支")

    def test_classification_failure_preserves_verified_layout_for_resume(self):
        engine=self.engine();self.server.mode="no_dimension"
        first=engine.recognize(snapshot());self.assertFalse(first["complete"])
        layouts=sum(item["task"]=="layout" for item in self.server.received)
        self.server.mode="normal"
        self.assertTrue(engine.recognize(snapshot())["complete"])
        self.assertEqual(sum(item["task"]=="layout" for item in self.server.received),layouts,
                         "分类维度错误不应让已通过的表范围重新识别")

    def test_classification_error_identifies_the_source_cell(self):
        engine=self.engine();self.server.mode="no_dimension"
        result=engine.recognize(snapshot())
        self.assertFalse(result["complete"])
        self.assertTrue(any("表一!B3" in item["reason"] and "period" in item["reason"] for item in result["unresolved"]))

    def test_known_word_boundaries_reject_invented_columns_or_cross_table_range(self):
        engine=self.engine();sheet=snapshot()["sheets"][0]
        sheet.update(max_column=4,max_row=8,word_table_ranges=[
            {"range_name":"TA_001","range":"A2:B3","table_id":"原Word表一"},
            {"range_name":"TA_002","range":"A5:D8","table_id":"原Word表二"}])
        table={"range":"A2:B3","note_ids":["C-N001"],"scope":"consolidated",
               "table_semantic":"货币资金","header_cells":["A2","B2"]}
        owned=["A2","B2","A3","B3"]
        self.assertEqual(engine._validate_layout({"tables":[table]},sheet,owned)["tables"],[table])
        for area in ("A1:B3","A2:D3","A2:D8"):
            with self.subTest(area=area),self.assertRaisesRegex(SemanticError,"Word实际表格边界"):
                engine._validate_layout({"tables":[{**table,"range":area}]},sheet,owned)

    def test_duplicate_json_keys_are_rejected_before_normalization(self):
        import io
        from unittest.mock import patch,Mock
        for content in ('{"tables":[],"tables":[]}', '{"cells":{"B3":{"kind":"unresolved"},"B3":{"kind":"excluded"}}}'):
            with self.subTest(content=content):
                engine=self.engine();opener=Mock()
                raw={"choices":[{"finish_reason":"stop","message":{"content":content}}]}
                opener.open.return_value=io.BytesIO(json.dumps(raw).encode())
                with patch("urllib.request.build_opener",return_value=opener),self.assertRaises(SemanticError):
                    engine._http("同一字段只能出现一次",{"task":"layout"})
                files=list(engine.work_dir.rglob("响应失败_*.json"));self.assertEqual(len(files),1)
                evidence=json.loads(files[0].read_text(encoding="utf-8-sig"))
                self.assertIn("重复字段",evidence["parse_error"]["message"])

    def test_missing_only_root_closer_preserves_response_and_runs_validator(self):
        import io
        from unittest.mock import patch,Mock
        engine=self.engine();content='{"items":[{"text":"原文 } 与转义引号 \\\" 不改变"}]'
        raw={"choices":[{"finish_reason":"stop","message":{"content":content}}]}
        opener=Mock();opener.open.side_effect=lambda *a,**k:io.BytesIO(json.dumps(raw).encode())
        validator=Mock(side_effect=lambda result:result)
        with patch("urllib.request.build_opener",return_value=opener):
            result=engine._request("原文结构校验",{"task":"transport_probe"},validator)
        self.assertEqual(result,json.loads(content+'}'));validator.assert_called_once()
        receipt=json.loads(next(engine.work_dir.rglob("回执.json")).read_text(encoding="utf-8-sig"))
        proof=receipt['transport_repairs'][0]
        self.assertEqual(proof['original_content'],content);self.assertEqual(proof['appended'],'}')
        self.assertEqual(proof['raw_response']['choices'][0]['message']['content'],content)
        self.assertEqual(opener.open.call_count,1)

    def test_root_closer_does_not_repair_incomplete_values_or_wrong_endings(self):
        import io
        from unittest.mock import patch,Mock
        cases=[('{"n":12','stop'),('{"text":"abc','stop'),('{"items":[{"a":1}','stop'),
               ('{"x":{},"x":{}','stop'),('{"items":[],','stop'),('{"items":[]','length'),('{"items":[]',None)]
        for content,finish in cases:
            with self.subTest(content=content,finish=finish):
                engine=self.engine();opener=Mock()
                raw={"choices":[{"finish_reason":finish,"message":{"content":content}}]}
                opener.open.return_value=io.BytesIO(json.dumps(raw).encode())
                with patch("urllib.request.build_opener",return_value=opener),self.assertRaises(SemanticError):
                    engine._http("必须保留失败",{"task":"transport_probe"})

    def test_root_closer_cannot_bypass_missing_candidate_validation(self):
        import io
        from unittest.mock import patch,Mock
        engine=self.engine();opener=Mock()
        raw={"choices":[{"finish_reason":"stop","message":{"content":'{"cells":{"B3":{}}'}}]}
        opener.open.side_effect=lambda *a,**k:io.BytesIO(json.dumps(raw).encode())
        def validator(result):
            if set(result['cells'])!={'B3','B4'}:raise ValueError('缺少业务格B4')
        with patch("urllib.request.build_opener",return_value=opener),self.assertRaisesRegex(SemanticError,'缺少业务格'):
            engine._request("全部格须覆盖",{"task":"transport_probe"},validator)
        self.assertEqual(opener.open.call_count,2)
        self.assertFalse(list(engine.work_dir.rglob("回执.json")))
        failure=json.loads(sorted(engine.work_dir.rglob("失败_*.json"))[-1].read_text(encoding="utf-8-sig"))
        self.assertTrue(failure['transport_repairs'])

    def test_malformed_http_response_keeps_redacted_evidence_and_short_error(self):
        import io
        from unittest.mock import patch
        engine=self.engine()
        content='{"mappings":[},"note":"local-test-secret 原始失败证据"}'
        raw={"choices":[{"finish_reason":"stop","message":{"content":content}}],"usage":{"prompt_tokens":7,"completion_tokens":9}}
        opener=__import__("unittest.mock",fromlist=["Mock"]).Mock()
        opener.open.return_value=io.BytesIO(json.dumps(raw,ensure_ascii=False).encode())
        with patch("urllib.request.build_opener",return_value=opener),self.assertRaises(SemanticError) as caught:
            engine._http("原任务",{"task":"layout"})
        self.assertNotIn("local-test-secret",str(caught.exception));self.assertLess(len(str(caught.exception)),160)
        files=list(engine.work_dir.rglob("响应失败_*.json"))
        self.assertEqual(len(files),1,"格式错误不能丢失原始HTTP证据")
        saved=json.loads(files[0].read_text(encoding="utf-8-sig"))
        self.assertEqual(saved["finish_reason"],"stop");self.assertEqual(saved["usage"],raw["usage"])
        self.assertIn("原始失败证据",saved["raw_response"])
        self.assertNotIn("local-test-secret",files[0].read_text(encoding="utf-8-sig"))
        self.assertIn("parse_error",saved);self.assertNotIn("Authorization",saved)

    def test_validator_failure_gets_one_strict_correction_then_original_cache(self):
        from unittest.mock import patch
        for kind in ("missing_dimension","duplicate","invalid_category"):
            with self.subTest(kind=kind):
                engine=self.engine();item=mapping();item["dimensions"]["counterparty"]="银行存款"
                engine.slots[item["slot_id"]]["dimensions"].append({"name":"counterparty","required":True,"open":True})
                good={"mappings":[item],"excluded":[],"unresolved":[]};bad=copy.deepcopy(good)
                if kind=="missing_dimension":bad["mappings"][0]["dimensions"].pop("counterparty")
                elif kind=="duplicate":bad["mappings"].append(copy.deepcopy(item))
                else:bad={"mappings":[],"excluded":[{"cell":"B3","category":"business","reason":"原错误类别"}],"unresolved":[]}
                payload={"task":"classify","round":1,"candidate_cells":["B3"]}
                table={"range":"A1:B3","note_ids":["C-N001"],"scope":"consolidated","table_semantic":"货币资金"}
                validator=lambda raw:engine._validate_classification(raw,["B3"],{item["slot_id"]},snapshot()["sheets"][0],table)
                calls=[]
                def http(system,request):calls.append(copy.deepcopy(request));return copy.deepcopy(good if request.get("correction") else bad)
                with patch.object(engine,"_http",side_effect=http):
                    try:result=engine._request("逐格原任务",payload,validator)
                    except ValueError as error:self.fail("可纠正业务结构错误仍立即失败："+str(error))
                    self.assertEqual(len(calls),2);self.assertEqual(result["B3"][1]["slot_id"],item["slot_id"])
                    self.assertEqual(calls[1]["correction"]["previous_result"],bad)
                    self.assertTrue(calls[1]["correction"]["validation_error"])
                    self.assertEqual(engine._request("逐格原任务",payload,validator),result)
                    self.assertEqual(len(calls),2,"成功纠正应保存为原任务回执，恢复时不能重跑")
                self.assertEqual(len(list(engine.work_dir.rglob("纠正请求_*.json"))),1)
                self.assertEqual(len(list(engine.work_dir.rglob("纠正回执_*.json"))),1)

    def test_failed_correction_is_not_accepted_and_never_loops(self):
        from unittest.mock import patch
        engine=self.engine()
        def invalid(raw):raise SemanticError("仍缺必要维度 local-test-secret")
        with patch.object(engine,"_http",return_value={"mappings":[]}) as http,self.assertRaises(SemanticError) as caught:
            engine._request("原任务",{"task":"classify","round":1},invalid)
        self.assertEqual(http.call_count,2);self.assertNotIn("local-test-secret",str(caught.exception))
        self.assertFalse(list(engine.work_dir.rglob("回执.json")))
        self.assertGreaterEqual(len(list(engine.work_dir.rglob("失败_*.json"))),2)
        for path in engine.work_dir.rglob("*.json"):self.assertNotIn("local-test-secret",path.read_text(encoding="utf-8-sig"))

    def test_invalid_json_can_be_corrected_once_without_sending_secret_back(self):
        import io
        from unittest.mock import patch
        engine=self.engine();bodies=[]
        class Opener:
            def open(self,request,**kwargs):
                body=json.loads(request.data);bodies.append(body)
                content='{"ok": broken, "note":"local-test-secret"}' if len(bodies)==1 else '{"ok":true}'
                return io.BytesIO(json.dumps({"choices":[{"finish_reason":"stop","message":{"content":content}}],"usage":{"prompt_tokens":1,"completion_tokens":2}}).encode())
        with patch("urllib.request.build_opener",return_value=Opener()):
            try:result=engine._request("只返回JSON",{"task":"format_test"},lambda raw:raw)
            except ValueError as error:self.fail("无效JSON未尝试一次完整纠正："+str(error))
        self.assertEqual(result,{"ok":True});self.assertEqual(len(bodies),2)
        correction=json.loads(bodies[1]["messages"][1]["content"])["correction"]
        self.assertTrue(correction["validation_error"]);self.assertNotIn("local-test-secret",json.dumps(bodies[1]))
        self.assertEqual(len(list(engine.work_dir.rglob("响应失败_*.json"))),1)

    def test_correction_stays_with_its_round_and_does_not_feed_other_round(self):
        from unittest.mock import patch
        engine=self.engine();calls=[]
        def validator(raw):
            if not raw.get("valid"):raise SemanticError("缺少业务结构")
            return raw
        def http(system,payload):
            calls.append(copy.deepcopy(payload))
            return {"valid":True,"round":payload["round"]} if payload.get("correction") else {"marker":"原轮专属"+str(payload["round"])}
        with patch.object(engine,"_http",side_effect=http):
            for number in (1,2):
                try:engine._request("独立轮次",{"task":"classify","round":number},validator)
                except ValueError as error:self.fail("没有执行一次纠正："+str(error))
        self.assertEqual(len(calls),4)
        self.assertNotIn("correction",calls[2]);self.assertNotIn("原轮专属1",json.dumps(calls[2:],ensure_ascii=False))
        self.assertEqual(calls[3]["correction"]["previous_result"],{"marker":"原轮专属2"})

    def test_transport_cancellation_and_disk_errors_never_become_correction(self):
        from unittest.mock import patch
        for error in (SemanticError("模型服务返回 HTTP 503"),SemanticError("模型服务连接失败或超时"),SemanticCancelled("取消")):
            with self.subTest(error=str(error)):
                engine=self.engine()
                with patch.object(engine,"_http",side_effect=error) as http,self.assertRaises(type(error)):
                    engine._request("原任务",{"task":"classify"},lambda raw:raw)
                self.assertEqual(http.call_count,1);self.assertFalse(list(engine.work_dir.rglob("纠正请求_*.json")))
        engine=self.engine()
        with patch.object(engine,"_http",return_value={"ok":True}) as http:
            def disk_failure(raw):raise OSError("磁盘错误")
            with self.assertRaises((OSError,SemanticError)):engine._request("原任务",{"task":"classify"},disk_failure)
            self.assertEqual(http.call_count,1)

    def test_original_http_backoff_remains_without_semantic_correction(self):
        import io,urllib.error
        from unittest.mock import patch,Mock
        engine=self.engine();opener=Mock()
        opener.open.side_effect=[urllib.error.HTTPError("http://local.test",503,"busy",{},io.BytesIO(b"busy")) for _ in range(3)]
        with patch("urllib.request.build_opener",return_value=opener),patch("time.sleep"),self.assertRaises(SemanticError):
            engine._request("原任务",{"task":"classify"},lambda raw:raw)
        self.assertEqual(opener.open.call_count,3)
        self.assertFalse(list(engine.work_dir.rglob("纠正请求_*.json")))

    def test_cancel_after_invalid_first_reply_stops_before_correction(self):
        from unittest.mock import patch
        event=threading.Event();engine=self.engine(cancel=event)
        def validator(raw):event.set();raise SemanticError("缺少必填维度")
        with patch.object(engine,"_http",return_value={}) as http,self.assertRaises(SemanticCancelled):
            engine._request("原任务",{"task":"classify"},validator)
        self.assertEqual(http.call_count,1)

    def test_disk_failure_short_message_is_redacted_without_model_correction(self):
        from unittest.mock import patch
        engine=self.engine()
        def validator(raw):raise OSError("磁盘错误 local-test-secret")
        with patch.object(engine,"_http",return_value={}) as http,self.assertRaises((OSError,SemanticError)) as caught:
            engine._request("原任务",{"task":"classify"},validator)
        self.assertEqual(http.call_count,1);self.assertNotIn("local-test-secret",str(caught.exception))

    def test_thinking_mode_only_adds_explicit_template_kwarg(self):
        engine=self.engine()
        for mode, expected in (("default",None),("enabled",True),("disabled",False)):
            with self.subTest(mode=mode):
                engine.settings["thinking_mode"]=mode
                engine._http("测试请求选项",{"task":"layout"})
                body=self.server.bodies[-1]
                self.assertNotIn("reasoning_effort",body)
                if expected is None:self.assertNotIn("chat_template_kwargs",body)
                else:self.assertEqual(body.get("chat_template_kwargs"),{"enable_thinking":expected})
        engine.settings["thinking_mode"]="invalid"
        count=len(self.server.bodies)
        with self.assertRaises(ValueError):engine._http("无效设置",{"task":"layout"})
        self.assertEqual(len(self.server.bodies),count)

    def test_thinking_mode_switch_cannot_reuse_old_mode_receipts(self):
        engine=self.engine("thinking_cache")
        self.assertTrue(engine.recognize(snapshot())["complete"])
        before=len(self.server.received)
        engine.settings["thinking_mode"]="disabled"
        self.assertTrue(engine.recognize(snapshot())["complete"])
        self.assertGreater(len(self.server.received),before)
        before=len(self.server.received)
        self.assertTrue(engine.recognize(snapshot())["complete"])
        self.assertEqual(len(self.server.received),before)
        engine.settings["thinking_mode"]="enabled"
        self.assertTrue(engine.recognize(snapshot())["complete"])
        self.assertGreater(len(self.server.received),before)
        before=len(self.server.received)
        engine.settings["thinking_mode"]="default"
        self.assertTrue(engine.recognize(snapshot())["complete"])
        self.assertEqual(len(self.server.received),before)

    def test_period_role_is_description_only_when_declared_by_slot(self):
        from 附注更新.分步流程 import _normalize_mapping
        from 附注更新.表格 import semantic_key
        engine=self.engine();sheet=snapshot()["sheets"][0];item=mapping()
        slot=engine.slots[item["slot_id"]]
        period=next(d for d in slot["dimensions"] if d["name"]=="period")
        period.update(role="closing",type="instant")
        table={"range":"A1:B3","scope":"consolidated","table_semantic":"余额","header_cells":["A1","B2"]}
        def checked(date):
            source=copy.deepcopy(item)
            source["dimensions"]["period"]={"date":date,"role":"closing","type":"instant"}
            before=copy.deepcopy(source)
            result=engine._validate_classification({"mappings":[source],"excluded":[],"unresolved":[]},
                ["B3"],{item["slot_id"]},sheet,table)["B3"][1]
            self.assertEqual(source,before)
            self.assertEqual(result["dimensions"]["period"],date)
            self.assertEqual(_normalize_mapping(source,slot)["dimensions"],result["dimensions"])
            self.assertEqual(source,before)
            return result
        first=checked("2025-12-31");second=checked("2024-12-31")
        plain=engine._validate_classification({"mappings":[item],"excluded":[],"unresolved":[]},
            ["B3"],{item["slot_id"]},sheet,table)["B3"][1]
        self.assertEqual(semantic_key(first),semantic_key(plain))
        self.assertNotEqual(semantic_key(first),semantic_key(second))

    def test_period_role_conflicts_unknown_fields_and_absent_definition_are_rejected(self):
        from 附注更新.语义 import _normal_dimensions
        base={"date":"2025-12-31","role":"closing","type":"instant"}
        for declared in (None,{}, {"dimensions":[{"name":"period","role":None}]},
                         {"dimensions":[{"name":"period","role":"opening"}]}):
            with self.subTest(declared=declared),self.assertRaises(SemanticError):
                _normal_dimensions({"period":base},declared)
        slot={"dimensions":[{"name":"period","role":"closing","type":"instant"}]}
        for bad in ({**base,"role":"opening"},{**base,"role":True},{**base,"extra":"x"},
                    {**base,"type":"duration"},{**base,"date":"不是日期"},
                    {"start":"2025-01-01","end":"2025-12-31","role":"closing","type":"duration"},
                    {"start":"2025-01-01","end":"2025-12-31","role":"closing"}):
            with self.subTest(period=bad),self.assertRaises(SemanticError):
                _normal_dimensions({"period":bad},slot)
        duration={"start":"2025-01-01","end":"2025-12-31","type":"duration","role":"current"}
        slot={"dimensions":[{"name":"period","role":"current","type":"duration"}]}
        self.assertEqual(_normal_dimensions({"period":duration},slot)["period"],
                         {"start":"2025-01-01","end":"2025-12-31"})

    def test_period_type_metadata_does_not_create_another_actual_dimension(self):
        from 附注更新.语义 import _normal_dimensions
        engine=self.engine();sheet=snapshot()["sheets"][0]
        table={"range":"A1:B3","scope":"consolidated","table_semantic":"银行存款","header_cells":["A1","B2"]}
        meanings=[]
        for period in ("2025-12-31",{"date":"2025-12-31"},{"type":"instant","date":"2025-12-31"}):
            item=mapping();item["dimensions"]["period"]=period
            accepted=engine._validate_classification({"mappings":[item],"excluded":[],"unresolved":[]},
                ["B3"],{item["slot_id"]},sheet,table)["B3"][1]
            meanings.append(engine._meaning(accepted))
        self.assertTrue(all(item==meanings[0] for item in meanings))
        duration={"start":"2025-01-01","end":"2025-12-31"}
        self.assertEqual(_normal_dimensions({"period":dict(duration,type="duration")})["period"],duration)
        invalid=[{"type":"duration","date":"2025-12-31"},dict(duration,type="instant"),
                 {"date":"2025-12-31","role":"closing"},{"date":"2025-12-31","extra":"x"},
                 dict(duration,date="2025-12-31"),{"type":"unknown","date":"2025-12-31"}]
        for period in invalid:
            with self.subTest(period=period),self.assertRaises(SemanticError):
                _normal_dimensions({"period":period})

    def test_date_object_matches_plain_date_and_reversed_duration_is_rejected(self):
        from 附注更新.语义 import _normal_dimensions
        from 附注更新.表格 import semantic_key
        plain=mapping();wrapped=copy.deepcopy(plain)
        wrapped["dimensions"]["period"]={"date":"2025-12-31"}
        plain["dimensions"]=_normal_dimensions(plain["dimensions"])
        wrapped["dimensions"]=_normal_dimensions(wrapped["dimensions"])
        self.assertEqual(semantic_key(plain),semantic_key(wrapped))
        duration={"start":"2025-01-01","end":"2025-12-31"}
        self.assertEqual(_normal_dimensions({"period":duration})["period"],duration)
        with self.assertRaises(SemanticError):
            _normal_dimensions({"period":{"start":"2025-12-31","end":"2025-01-01"}})

    def test_standalone_requires_explicit_gold_applicability(self):
        engine=self.engine()
        sheet=snapshot()["sheets"][0]
        table={"range":"A1:B3","note_ids":["C-N001"],"scope":"standalone","table_semantic":"单户货币资金","header_cells":["A1","A2","B2"]}
        item={**mapping(),"scope":"standalone"}
        response={"mappings":[item],"excluded":[],"unresolved":[]}
        with self.assertRaises(ValueError):
            engine._validate_classification(response,["B3"],{item["slot_id"]},sheet,table)
        row=json.loads(self.gold.read_text(encoding="utf-8-sig"))
        row["applicability"]={"report_scopes":[{"scope":"standalone","reason":"银行存款期末余额不依赖集团合并关系","reviewed_by":"测试审核","reviewed_at":"2026-09-19"}]}
        path=self.work/("单户标准_"+uuid.uuid4().hex+".jsonl")
        path.write_text(json.dumps(row,ensure_ascii=False),encoding="utf-8-sig")
        engine=SemanticEngine(engine.settings,str(path),str(self.work/uuid.uuid4().hex))
        engine._validate_layout({"tables":[table]},sheet,["B3"])
        result=engine._validate_classification(response,["B3"],{item["slot_id"]},sheet,table)["B3"][1]
        self.assertEqual(result["scope"],"standalone")
        self.assertEqual(result["definition_scope"],"consolidated")
        self.assertEqual(result["slot_id"],row["id"])

    def test_service_without_key_is_usable_and_errors_are_readable(self):
        engine=self.engine()
        engine.settings["api_key"]=""
        result=engine.recognize(snapshot())
        self.assertTrue(result["complete"])
        self.server.mode="bad_id"
        failed=self.engine();failed.settings["api_key"]=""
        self.assertFalse(failed.recognize(snapshot())["complete"])
        records=list(failed.work_dir.rglob("失败_*.json"))
        self.assertTrue(records)
        for file in records:
            self.assertNotIn("[密钥]",file.read_text(encoding="utf-8-sig"))

    def test_inline_reasoning_reads_only_final_json(self):
        for mode in ("inline_reasoning","tagged_reasoning","literal_think_marker"):
            with self.subTest(mode=mode):
                self.server.mode=mode
                reply=self.engine()._http("测试最终结果",{"task":"layout"})
                self.assertEqual(reply["tables"][0]["range"],"A1:B3")
                if mode=="literal_think_marker":
                    self.assertIn("</think>",reply["carry_context"])

    def test_unfinished_reasoning_is_not_a_successful_answer(self):
        for mode in ("unfinished_reasoning","reasoning_no_answer"):
            with self.subTest(mode=mode):
                self.server.mode=mode
                engine=self.engine()
                with self.assertRaises(SemanticError):
                    engine._http("测试最终结果",{"task":"layout"})
                self.assertEqual(engine.usage["requests"],0)

    def test_confirmed_scope_rejects_model_guess_and_retains_failure_reply(self):
        engine=self.engine()
        engine.report_scope="standalone"
        result=engine.recognize(snapshot())
        self.assertFalse(result["complete"])
        failures=list(engine.work_dir.rglob("失败_*.json"))
        self.assertTrue(failures)
        receipt=json.loads(failures[0].read_text(encoding="utf-8-sig"))
        self.assertEqual(receipt["result"]["tables"][0]["scope"],"consolidated")
        self.assertNotIn("local-test-secret",json.dumps(receipt))

    def test_scope_selection_accounts_for_excluded_financial_cells(self):
        from unittest.mock import patch
        source=snapshot(); sheet=source["sheets"][0]; sheet["max_row"]=5
        sheet["cells"]["A5"]={"row":5,"column":1,"value":"人力资源统计"}
        sheet["cells"]["B5"]={"row":5,"column":2,"value":10}
        selection={"source_hash":"source-one","target_hash":__import__("附注更新.语义",fromlist=["_hash"])._hash({"disclosures":["货币资金"]}),"selected_cells":{"表一":list(snapshot()["sheets"][0]["cells"])},"out_of_scope":[{"sheet":"表一","cells":["A5","B5"],"reason":"本次不更新人力资源统计","evidence_cells":["A5"]}]}
        with patch("附注更新.识别范围.select_source_scope",return_value=selection):
            result=self.engine().recognize(source,target_scope={"disclosures":["货币资金"]})
        self.assertTrue(result["complete"])
        self.assertEqual({x["cell"] for x in result["out_of_scope"]},{"A5","B5"})
        self.assertEqual([x["cell"] for x in result["mappings"]],["B3"])

    def test_previous_scope_selection_is_validated_by_scope_resume_entry(self):
        from unittest.mock import patch
        from 附注更新.语义 import _hash
        source=snapshot();target={"disclosures":["货币资金"]}
        previous={"source_hash":"source-one","target_hash":_hash(target),
                  "selected_cells":{"表一":list(source["sheets"][0]["cells"])},"out_of_scope":[]}
        engine=self.engine()
        with patch("附注更新.识别范围.select_source_scope",return_value=previous) as select:
            result=engine.recognize(source,target_scope=target,scope_selection=previous)
        self.assertTrue(result["complete"])
        select.assert_called_once_with(engine,source,target,previous_selection=previous)

    def test_recognize_validate_and_resume_without_new_requests(self):
        engine = self.engine("resume")
        first = engine.recognize(snapshot())
        self.assertTrue(first["complete"])
        self.assertEqual([m["cell"] for m in first["mappings"]], ["B3"])
        self.assertEqual(first["mappings"][0]["dimensions"]["currency"], "CNY")
        self.assertEqual(first["source_hash"], "source-one")
        count = len(self.server.received)
        second = engine.recognize(snapshot())
        self.assertEqual(first["mappings"], second["mappings"])
        self.assertEqual(count, len(self.server.received))
        for file in (self.work / "resume").rglob("*.json"):
            self.assertNotIn("local-test-secret", file.read_text(encoding="utf-8-sig"))

    def test_bad_slot_omission_duplicate_and_outside_are_incomplete(self):
        for mode in ("bad_id", "omission", "duplicate", "outside", "wrong_scope", "wrong_type", "no_dimension"):
            with self.subTest(mode=mode):
                self.server.mode = mode
                result = self.engine().recognize(snapshot())
                self.assertFalse(result["complete"])
                self.assertTrue(result["unresolved"])
                self.assertFalse(result["mappings"])

    def test_two_independent_rounds_disagree(self):
        self.server.mode = "conflict"
        result = self.engine().recognize(snapshot())
        self.assertFalse(result["complete"])
        self.assertFalse(result["mappings"])
        self.assertIn("B3", [x["cell"] for x in result["unresolved"]])

    def test_synonymous_dimensions_are_normalized(self):
        self.server.mode = "synonyms"
        self.assertTrue(self.engine().recognize(snapshot())["complete"])

    def test_ai_values_never_become_mapping_source(self):
        self.server.mode = "value_injected"
        result = self.engine().recognize(snapshot())
        self.assertTrue(result["complete"])
        self.assertNotIn("value", result["mappings"][0])

    def test_cancel_before_request(self):
        event = threading.Event()
        event.set()
        with self.assertRaises(SemanticCancelled):
            self.engine(cancel=event).recognize(snapshot())
        self.assertEqual(self.server.received, [])

    def test_failed_package_is_retried_on_resume(self):
        engine = self.engine("retry")
        self.server.mode = "server_error"
        result = engine.recognize(snapshot())
        self.assertFalse(result["complete"])
        self.server.mode = "normal"
        result = engine.recognize(snapshot())
        self.assertTrue(result["complete"])

    def test_non_numeric_value_cannot_be_mapped_as_amount(self):
        data = snapshot()
        data["sheets"][0]["cells"]["B3"]["value"] = "不能当金额"
        result = self.engine().recognize(data)
        self.assertFalse(result["complete"])
        self.assertFalse(result["mappings"])

    def test_formula_semantics_do_not_require_cached_amount(self):
        data = snapshot()
        data["sheets"][0]["cells"]["B3"].update({"formula": "=SUM(B4:B5)", "value": "=SUM(B4:B5)", "cached_value": None})
        result = self.engine().recognize(data)
        self.assertTrue(result["complete"])
        self.assertTrue(result["mappings"])
        self.assertNotIn("value", result["mappings"][0])

    def test_changed_context_makes_new_requests(self):
        engine = self.engine("context")
        engine.recognize(snapshot())
        before = len(self.server.received)
        data = snapshot()
        data["context"] = ["新增：2025年度合并财务报表"]
        engine.recognize(data)
        self.assertGreater(len(self.server.received), before)
    def test_numeric_cell_cannot_be_excluded_by_two_agreeing_rounds(self):
        self.server.mode = "numeric_exclusion"
        result = self.engine().recognize(snapshot())
        self.assertFalse(result["complete"])
        self.assertIn("B3", [x["cell"] for x in result["unresolved"]])

    def test_nonbusiness_title_is_independently_checked(self):
        data = snapshot()
        sheet = data["sheets"][0]
        sheet.update(max_row=1, max_column=1, cells={"A1": sheet["cells"]["A1"]})
        self.server.mode = "title_only"
        result = self.engine().recognize(data)
        self.assertTrue(result["complete"])
        self.assertEqual([x["cell"] for x in result["excluded"]], ["A1"])
        self.assertGreaterEqual(sum(p["task"] == "layout" for p in self.server.received), 2)
        self.server.mode = "title_conflict"
        self.assertFalse(self.engine().recognize(data)["complete"])
        self.server.mode = "title_number"
        sheet["cells"]["A1"]["value"] = 123.45
        self.assertFalse(self.engine().recognize(data)["complete"])

    def test_default_metric_comes_from_gold_not_free_wording(self):
        self.server.mode = "field_words"
        result = self.engine().recognize(snapshot())
        self.assertTrue(result["complete"])
        mapped = result["mappings"][0]
        self.assertNotIn("metric", mapped["dimensions"])
        self.assertEqual(mapped["semantic_field"], "银行存款 / 期末余额")

    def test_actual_metric_conflict_cannot_pass_double_review(self):
        self.server.mode = "metric_conflict"
        result = self.engine().recognize(snapshot())
        self.assertFalse(result["complete"])

    def test_unresolved_and_conflicting_receipts_retry_on_resume(self):
        for mode in ("unresolved", "conflict"):
            with self.subTest(mode=mode):
                engine = self.engine("retry_" + mode)
                self.server.mode = mode
                self.assertFalse(engine.recognize(snapshot())["complete"])
                before = len(self.server.received)
                self.server.mode = "normal"
                self.assertTrue(engine.recognize(snapshot())["complete"])
                self.assertGreater(len(self.server.received), before)

    def test_cache_ignores_temporary_paths_but_binds_content(self):
        engine = self.engine("stable_paths")
        first = snapshot()
        first.update(original_hash="original-document", sha256="temporary-zip-one",
                     context=[{"source": "C:\\run-one\\表格A.xlsx", "title": "货币资金"}])
        engine.recognize(first)
        before = len(self.server.received)
        second = copy.deepcopy(first)
        second.update(sha256="temporary-zip-two", context=[{"source": "C:\\run-two\\表格A.xlsx", "title": "货币资金"}])
        engine.recognize(second)
        self.assertEqual(len(self.server.received), before)
        second["sheets"][0]["cells"]["B3"]["value"] = 999.00
        engine.recognize(second)
        self.assertGreater(len(self.server.received), before)

    def test_cross_origin_redirect_does_not_forward_key(self):
        received = []
        class Sink(http.server.BaseHTTPRequestHandler):
            def log_message(self, *args): pass
            def do_GET(self):
                received.append(self.headers.get("Authorization"))
                self.send_response(200); self.end_headers(); self.wfile.write(b"{}")
        sink = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Sink)
        thread = threading.Thread(target=sink.serve_forever, daemon=True)
        thread.start()
        try:
            self.server.mode = "redirect"
            self.server.redirect_url = f"http://127.0.0.1:{sink.server_port}/other"
            with self.assertRaises(SemanticError):
                self.engine()._http("测试", {"task": "layout"})
            self.assertEqual(received, [])
        finally:
            sink.shutdown(); sink.server_close(); thread.join(2)

    def test_redirect_keeps_same_origin_and_blocks_https_downgrade(self):
        import urllib.request
        handler = SameOriginRedirect()
        request = urllib.request.Request("https://example.test/v1", headers={"Authorization": "Bearer dummy"})
        same = handler.redirect_request(request, None, 302, "Found", {}, "https://example.test/v2")
        self.assertEqual(same.full_url, "https://example.test/v2")
        with self.assertRaises(SemanticError):
            handler.redirect_request(request, None, 302, "Found", {}, "http://example.test/v2")

    def test_year_header_needs_layout_and_text_evidence(self):
        data = snapshot()["sheets"][0]
        data["cells"]["B2"]["value"] = 2025
        raw = {"category": "header", "evidence_cells": ["A1"]}
        table = {"header_cells": ["B2"]}
        try:
            self.engine()._check_exclusion("B2", raw, data, table)
        except SemanticError as error:
            self.fail("明确2025年上下文的年份表头不应拒绝：" + str(error))
        with self.assertRaises(SemanticError):
            self.engine()._check_exclusion("B2", {"category": "header"}, data, table)
        with self.assertRaises(SemanticError):
            self.engine()._check_exclusion("B2", raw, data, {"header_cells": []})

    def test_undeclared_metric_cannot_repurpose_an_existing_slot(self):
        engine = self.engine()
        item = mapping()
        item["dimensions"]["metric"] = "土地面积"
        item["metric_evidence"] = ["A3"]
        sheet = snapshot()["sheets"][0]
        table = {"range":"A1:B3","scope":"consolidated","table_semantic":"指标","header_cells":["A1","B2"]}
        response = {"mappings":[item],"excluded":[],"unresolved":[]}
        with self.assertRaises(SemanticError):
            engine._validate_classification(response,["B3"],{item["slot_id"]},sheet,table)

    def test_value_type_cannot_change_even_for_an_open_metric(self):
        engine = self.engine()
        item = mapping()
        engine.slots[item["slot_id"]]["dimensions"].append({"name":"metric","required":False,"open":True,"description":"银行存款金额指标"})
        item["metric_evidence"] = ["A3"]
        sheet = snapshot()["sheets"][0]
        table = {"range":"A1:B3","scope":"consolidated","table_semantic":"指标","header_cells":["A1","B2"]}
        for kind, unit, metric in (("number","平方米","土地面积"),("percentage","%","银行存款比例"),("number","个","银行存款 / 期末余额")):
            with self.subTest(kind=kind,unit=unit):
                item.update(value_type=kind)
                item["dimensions"].update(unit=unit,metric=metric)
                with self.assertRaises(SemanticError):
                    engine._validate_classification({"mappings":[item],"excluded":[],"unresolved":[]},["B3"],{item["slot_id"]},sheet,table)

    def test_open_metric_requires_explicit_declaration_and_original_evidence(self):
        engine = self.engine()
        item = mapping()
        item["dimensions"]["metric"] = "受限银行存款期末余额"
        sheet = snapshot()["sheets"][0]
        sheet["cells"]["A3"]["value"] = "受限银行存款期末余额"
        table = {"range":"A1:B3","scope":"consolidated","table_semantic":"银行存款金额指标","header_cells":["A1","B2"]}
        response = {"mappings":[item],"excluded":[],"unresolved":[]}
        item["metric_evidence"] = ["A3"]
        declaration={"name":"metric","required":False,"description":"银行存款内不同受限类别的期末金额指标"}
        engine.slots[item["slot_id"]]["dimensions"].append(declaration)
        for opened in (None,False,"true"):
            declaration["open"]=opened
            with self.subTest(opened=opened),self.assertRaises(SemanticError):
                engine._validate_classification(response,["B3"],{item["slot_id"]},sheet,table)
        declaration["open"]=True
        item.pop("metric_evidence")
        with self.assertRaises(SemanticError):
            engine._validate_classification(response,["B3"],{item["slot_id"]},sheet,table)
        item["metric_evidence"]=["B3"]
        with self.assertRaises(SemanticError):
            engine._validate_classification(response,["B3"],{item["slot_id"]},sheet,table)
        item["metric_evidence"]=["A3"]
        accepted = engine._validate_classification(response,["B3"],{item["slot_id"]},sheet,table)
        self.assertEqual(accepted["B3"][1]["dimensions"]["metric"], "受限银行存款期末余额")
        self.assertEqual(accepted["B3"][1]["value_type"], "monetary")

    def test_standard_metric_alias_and_existing_open_object_remain_usable(self):
        engine=self.engine();item=mapping()
        engine.slots[item["slot_id"]]["dimensions"].append({"name":"counterparty","required":True,"open":True})
        item["dimensions"].update(metric="期末余额",counterparty="新开户银行")
        sheet=snapshot()["sheets"][0]
        table={"range":"A1:B3","scope":"consolidated","table_semantic":"银行存款","header_cells":["A1","B2"]}
        accepted=engine._validate_classification({"mappings":[item],"excluded":[],"unresolved":[]},["B3"],{item["slot_id"]},sheet,table)["B3"][1]
        self.assertNotIn("metric",accepted["dimensions"])
        self.assertEqual(accepted["semantic_field"],"银行存款 / 期末余额")
        self.assertEqual(accepted["dimensions"]["counterparty"],"新开户银行")

    def test_fixed_age_bucket_cannot_be_rebound_to_other_intervals(self):
        engine=self.engine();item=mapping();sheet=snapshot()["sheets"][0]
        table={"range":"A1:B3","scope":"consolidated","table_semantic":"账龄","header_cells":["A1","B2"]}
        declaration={"name":"age_bucket","required":True,"lower":0,"upper":1,"lower_inclusive":True,"upper_inclusive":True}
        engine.slots[item["slot_id"]]["dimensions"].append(declaration)
        exact={"lower":0,"upper":1,"lower_inclusive":True,"upper_inclusive":True}
        for actual in ({**exact,"lower":1,"upper":2},{**exact,"lower":2,"upper":3},{**exact,"lower_inclusive":False},"1年以内",{"lower":0,"upper":1}):
            with self.subTest(actual=actual),self.assertRaises(SemanticError):
                item["dimensions"]["age_bucket"]=actual
                engine._validate_classification({"mappings":[item],"excluded":[],"unresolved":[]},["B3"],{item["slot_id"]},sheet,table)
        item["dimensions"]["age_bucket"]=exact
        accepted=engine._validate_classification({"mappings":[item],"excluded":[],"unresolved":[]},["B3"],{item["slot_id"]},sheet,table)["B3"][1]
        self.assertEqual(accepted["dimensions"]["age_bucket"],exact)
        self.assertEqual(accepted["dimensions"]["period"],"2025-12-31")
        declaration["open"]=True
        item["dimensions"]["age_bucket"]={**exact,"lower":1,"upper":2}
        engine._validate_classification({"mappings":[item],"excluded":[],"unresolved":[]},["B3"],{item["slot_id"]},sheet,table)

    def test_open_age_bucket_uses_one_slot_for_new_intervals_and_rejects_invalid_values(self):
        engine=self.engine();item=mapping();sheet=snapshot()["sheets"][0]
        table={"range":"A1:B3","scope":"consolidated","table_semantic":"账龄","header_cells":["A1","B2"]}
        engine.slots[item["slot_id"]]["dimensions"].append({"name":"age_bucket","required":True,"open":True})
        expected={"lower":2,"upper":3,"lower_inclusive":False,"upper_inclusive":True}
        item["dimensions"]["age_bucket"]={**expected,"lower":2.0}
        response={"mappings":[item],"excluded":[],"unresolved":[]}
        accepted=engine._validate_classification(response,["B3"],{item["slot_id"]},sheet,table)["B3"][1]
        self.assertEqual(accepted["slot_id"],item["slot_id"])
        self.assertIs(type(accepted["dimensions"]["age_bucket"]["lower"]),int)
        for invalid in ("2至3年",{**expected,"lower":4},{**expected,"lower":-1},{**expected,"upper":None},{**expected,"lower":True},{**expected,"upper":float("inf")},{"lower":2,"upper":3}):
            with self.subTest(invalid=invalid),self.assertRaises(SemanticError):
                item["dimensions"]["age_bucket"]=invalid
                engine._validate_classification(response,["B3"],{item["slot_id"]},sheet,table)
        item["dimensions"]["age_bucket"]={"lower":3,"upper":None,"lower_inclusive":False,"upper_inclusive":False}
        accepted=engine._validate_classification(response,["B3"],{item["slot_id"]},sheet,table)["B3"][1]
        self.assertEqual(accepted["slot_id"],item["slot_id"])

    def test_period_template_role_allows_actual_dates_but_preserves_period_type(self):
        engine=self.engine();item=mapping();sheet=snapshot()["sheets"][0]
        table={"range":"A1:B3","scope":"consolidated","table_semantic":"期间","header_cells":["A1","B2"]}
        period=next(d for d in engine.slots[item["slot_id"]]["dimensions"] if d["name"]=="period")
        period.update(role="closing",type="instant")
        response={"mappings":[item],"excluded":[],"unresolved":[]}
        engine._validate_classification(response,["B3"],{item["slot_id"]},sheet,table)
        item["dimensions"]["period"]={"start":"2025-01-01","end":"2025-12-31"}
        with self.assertRaises(SemanticError):
            engine._validate_classification(response,["B3"],{item["slot_id"]},sheet,table)
        period.update(role="current",type="duration")
        engine._validate_classification(response,["B3"],{item["slot_id"]},sheet,table)
        item["dimensions"]["period"]="2025-12-31"
        with self.assertRaises(SemanticError):
            engine._validate_classification(response,["B3"],{item["slot_id"]},sheet,table)

    def classification_failure_fixture(self):
        engine=self.engine();sheet=snapshot()["sheets"][0];sheet["max_row"]=4
        sheet["cells"]["B4"]={**sheet["cells"]["B3"],"row":4,"value":67,"cached_value":67}
        table={"range":"A1:B4","note_ids":["C-N001"],"scope":"consolidated","table_semantic":"银行存款余额","header_cells":["A1","B2"]}
        broken={**mapping(),"cell":"B4","dimensions":{k:v for k,v in mapping()["dimensions"].items() if k!="period"}}
        return engine,sheet,table,broken

    def test_classification_keeps_valid_cells_only_after_two_failed_validations(self):
        from unittest.mock import patch
        import hashlib
        from pathlib import Path
        engine,sheet,table,broken=self.classification_failure_fixture();calls=[]
        def reply(system,payload):
            calls.append(copy.deepcopy(payload))
            if payload["task"]=="tables":return {"table_ids":["C-N001-T001"]}
            return {"mappings":[mapping(),copy.deepcopy(broken)],"excluded":[],"unresolved":[]}
        with patch.object(engine,"_http",side_effect=reply):
            result=engine._classify(sheet,table,["B3","B4"],[],{})
        self.assertEqual(len([p for p in calls if p["task"]=="classify"]),4)
        self.assertEqual([m["cell"] for m in result["mappings"]],["B3"])
        self.assertTrue(result["mappings"][0]["reviewed"])
        self.assertEqual([m["cell"] for m in result["unresolved"]],["B4"])
        for item in result["mappings"]+result["unresolved"]:
            proofs=item["classification_recovery"];self.assertEqual(len(proofs),2)
            self.assertEqual({proof["round"] for proof in proofs},{1,2})
            for proof in proofs:
                raw=Path(proof["response_path"]).read_bytes()
                self.assertEqual(hashlib.sha256(raw).hexdigest(),proof["response_sha256"])
                stored=json.loads(raw.decode("utf-8-sig"));self.assertFalse(stored["validated"])
                self.assertEqual(proof["raw_record"]["cell"],item["cell"])
                self.assertEqual(proof["task_id"],stored["task_id"])
                self.assertFalse(Path(proof["response_path"]).with_name("回执.json").exists())
                if item["cell"]=="B4":self.assertIn("period",proof["validation_error"])
        with patch.object(engine,"_http",side_effect=reply):engine._classify(sheet,table,["B3","B4"],[],{})
        self.assertEqual(len([p for p in calls if p["task"]=="classify"]),8,"失败包不能冒充成功缓存")

    def test_recovered_cells_still_require_two_round_semantic_agreement(self):
        from unittest.mock import patch
        engine,sheet,table,broken=self.classification_failure_fixture()
        def reply(system,payload):
            if payload["task"]=="tables":return {"table_ids":["C-N001-T001"]}
            good=mapping()
            if payload["round"]==2:good["dimensions"]["period"]="2024-12-31"
            return {"mappings":[good,copy.deepcopy(broken)],"excluded":[],"unresolved":[]}
        with patch.object(engine,"_http",side_effect=reply):result=engine._classify(sheet,table,["B3","B4"],[],{})
        self.assertEqual(result["mappings"],[])
        item=next(x for x in result["unresolved"] if x["cell"]=="B3")
        self.assertIn("两轮",item["reason"]);self.assertNotEqual(item["first"]["dimensions"],item["second"]["dimensions"])
        self.assertEqual(len(item["classification_recovery"]),2)

    def test_final_classification_identity_errors_never_allow_partial_recovery(self):
        from unittest.mock import patch
        for fault in ("duplicate","missing","outside","not_object","missing_array"):
            with self.subTest(fault=fault):
                engine,sheet,table,broken=self.classification_failure_fixture()
                def reply(system,payload):
                    if payload["task"]=="tables":return {"table_ids":["C-N001-T001"]}
                    result={"mappings":[mapping(),copy.deepcopy(broken)],"excluded":[],"unresolved":[]}
                    if fault=="duplicate":result["excluded"].append({"cell":"B3","category":"label","reason":"重复"})
                    if fault=="missing":result["mappings"].pop()
                    if fault=="outside":result["mappings"][1]["cell"]="B99"
                    if fault=="not_object":result["excluded"].append("错误记录")
                    if fault=="missing_array":del result["unresolved"]
                    return result
                with patch.object(engine,"_http",side_effect=reply):
                    with self.assertRaises(SemanticError):engine._classify(sheet,table,["B3","B4"],[],{})

    def test_parse_failure_does_not_count_as_a_failed_business_validation(self):
        from unittest.mock import patch
        from 附注更新.语义 import SemanticResponseError
        engine,sheet,table,broken=self.classification_failure_fixture();calls=[]
        def reply(system,payload):
            calls.append(payload)
            if payload["task"]=="tables":return {"table_ids":["C-N001-T001"]}
            if "correction" not in payload:raise SemanticResponseError(self.work/"本地解析失败.json",{},"模拟无法解析的回执")
            return {"mappings":[mapping(),copy.deepcopy(broken)],"excluded":[],"unresolved":[]}
        with patch.object(engine,"_http",side_effect=reply):
            with self.assertRaises(SemanticError):engine._classify(sheet,table,["B3","B4"],[],{})
        self.assertEqual(len([p for p in calls if p["task"]=="classify"]),2)

    def layout_failure_fixture(self):
        engine=self.engine();snap=snapshot();sheet=snap["sheets"][0];sheet["max_row"]=4
        for address,value in (("A4","金额合计"),("B4",0)):
            sheet["cells"][address]={"row":4,"column":ord(address[0])-64,"value":value,"cached_value":value,"formula":None}
        response={"tables":[{"range":"A3:B3","note_ids":["C-N001"],"scope":"consolidated",
                            "table_semantic":"银行存款余额","header_cells":["A1","A2","B2"]}],
                  "non_business":[{"range":"A1:A2","category":"header","reason":"原标题及项目列头","evidence_cells":["A1","A2"]},
                                  {"range":"B2:B2","category":"header","reason":"原期间列头","evidence_cells":["B2"]},
                                  {"range":"A4:B4","category":"empty_padding","reason":"误把真实合计及零金额视为空白","evidence_cells":["A4"]}],
                  "unresolved":[]}
        return engine,snap,response

    def test_layout_recovery_preserves_valid_regions_and_requires_real_dual_review(self):
        from unittest.mock import patch
        import hashlib
        engine,snap,response=self.layout_failure_fixture();calls=[]
        def reply(system,payload):
            calls.append(copy.deepcopy(payload))
            if payload["task"]=="layout":return copy.deepcopy(response)
            if payload["task"]=="tables":return {"table_ids":["C-N001-T001"]}
            return {"mappings":[mapping()],"excluded":[{"cell":"A3","category":"label","reason":"银行存款行标签"}],"unresolved":[]}
        with patch.object(engine,"_http",side_effect=reply):result=engine.recognize(snap)
        self.assertEqual(len([x for x in calls if x["task"]=="layout"]),4)
        self.assertEqual([x["cell"] for x in result["mappings"]],["B3"])
        self.assertEqual({x["cell"] for x in result["excluded"]},{"A1","A2","B2","A3"})
        self.assertEqual({x["cell"] for x in result["unresolved"]},{"A4","B4"})
        self.assertFalse(result["complete"])
        for item in result["mappings"]+result["excluded"]:
            self.assertTrue(item["reviewed"])
            for evidence in item["layout_recovery"]:
                raw=pathlib.Path(evidence["recovery_path"]).read_bytes()
                self.assertEqual(hashlib.sha256(raw).hexdigest(),evidence["recovery_sha256"])
                proof=json.loads(raw.decode("utf-8-sig"));self.assertTrue(proof["retry_required"])
                self.assertFalse(pathlib.Path(proof["source_evidence"]["response_path"]).with_name("回执.json").exists())
        item=next(x for x in result["excluded"] if x["cell"]=="A1")
        self.assertEqual({x["round"] for x in item["layout_recovery"]},{1,2})
        self.assertNotIn("B1",snap["sheets"][0]["cells"],"恢复布局不能补造候选以外的空格")
        with patch.object(engine,"_http",side_effect=reply):engine.recognize(snap)
        self.assertEqual(len([x for x in calls if x["task"]=="layout"]),8)

    def test_layout_recovery_does_not_accept_second_round_rejected_exclusion(self):
        from unittest.mock import patch
        engine,snap,response=self.layout_failure_fixture()
        def reply(system,payload):
            if payload["task"]=="layout":
                item=copy.deepcopy(response)
                if payload["round"]==2:item["non_business"][0]["category"]="empty_padding"
                return item
            if payload["task"]=="tables":return {"table_ids":["C-N001-T001"]}
            good=mapping()
            if payload["round"]==2:good["dimensions"]["period"]="2024-12-31"
            return {"mappings":[good],"excluded":[{"cell":"A3","category":"label","reason":"行标签"}],"unresolved":[]}
        with patch.object(engine,"_http",side_effect=reply):result=engine.recognize(snap)
        self.assertEqual(result["mappings"],[])
        self.assertEqual({x["cell"] for x in result["excluded"]},{"B2","A3"})
        self.assertEqual({x["cell"] for x in result["unresolved"]},{"A1","A2","B3","A4","B4"})

    def test_layout_recovery_rejects_unreliable_partition_and_forged_evidence(self):
        from unittest.mock import patch
        for fault in ("missing","overlap","outside","missing_range","bad_evidence","bad_header","duplicate_candidate"):
            with self.subTest(fault=fault):
                engine,snap,response=self.layout_failure_fixture();sheet=snap["sheets"][0];owned=list(sheet["cells"])
                if fault=="missing":response["non_business"].pop()
                if fault=="overlap":response["non_business"][-1]["range"]="A3:B4"
                if fault=="outside":response["non_business"][-1]["range"]="A4:C4"
                if fault=="missing_range":del response["non_business"][-1]["range"]
                if fault=="bad_evidence":response["non_business"][-1]["evidence_cells"]=["Z999"]
                if fault=="bad_header":response["tables"][0]["header_cells"]=["Z999"]
                if fault=="duplicate_candidate":owned.append(owned[0])
                with patch.object(engine,"_http",return_value=response):
                    with self.assertRaises(SemanticError):engine._request_layout("布局",{"task":"layout","candidate_cells":owned,"round":1},sheet,owned)

    def test_layout_parse_failure_and_tampered_final_proof_cannot_be_recovered(self):
        from unittest.mock import patch
        from 附注更新.语义 import SemanticResponseError
        engine,snap,response=self.layout_failure_fixture();sheet=snap["sheets"][0];owned=list(sheet["cells"])
        def reply(system,payload):
            if "correction" not in payload:raise SemanticResponseError(self.work/"解析失败.json",{},"模拟解析失败")
            return response
        with patch.object(engine,"_http",side_effect=reply):
            with self.assertRaises(SemanticError):engine._request_layout("布局",{"task":"layout","candidate_cells":owned},sheet,owned)
        with patch.object(engine,"_http",return_value=response):
            with self.assertRaises(SemanticError) as failure:
                engine._request("布局",{"task":"layout","candidate_cells":owned},lambda r:engine._validate_layout(r,sheet,owned,single_note=True))
        proof=failure.exception.layout_failure;proof["evidence"]["response_sha256"]="0"*64
        with self.assertRaises(SemanticError):engine._recover_layout(proof["result"],sheet,owned,proof["evidence"])

    def test_layout_recovery_ignores_only_real_presented_context_regions(self):
        from unittest.mock import patch
        engine,snap,response=self.layout_failure_fixture();sheet=snap["sheets"][0];owned=list(sheet["cells"])
        sheet.update(max_row=33,max_column=8)
        for a,value in (("E33","其他流动负债"),("F33",29),("G33",0),("H33",0)):
            sheet["cells"][a]={"row":33,"column":ord(a[0])-64,"value":value,"cached_value":value,"formula":None}
        context_table={"range":"E33:H33","note_ids":["C-N001"],"scope":"consolidated",
                       "table_semantic":"包外上下文，不在本次候选","header_cells":["B2"]}
        response["tables"].append(context_table)
        payload={"task":"layout","candidate_cells":owned,"cells":copy.deepcopy(sheet["cells"]),"round":1}
        with patch.object(engine,"_http",return_value=response):result=engine._request_layout("布局",payload,sheet,owned)
        self.assertEqual(len(result["tables"]),1)
        self.assertEqual(result["ignored_context_regions"][0]["raw_record"],context_table)
        self.assertFalse(result["ignored_context_regions"][0]["confirmed"])
        self.assertEqual(result["ignored_context_regions"][0]["candidate_cells"],[])
        self.assertEqual(set(engine._layout_unresolved_cells(result,sheet,owned)),{"A4","B4"})
        payload["cells"].pop("H33")
        with patch.object(engine,"_http",return_value=response):
            with self.assertRaisesRegex(SemanticError,"展示上下文"):
                engine._request_layout("布局",payload,sheet,owned)

    def test_layout_recovery_never_expands_a_table_beyond_current_owned_candidates(self):
        from unittest.mock import patch
        engine,snap,response=self.layout_failure_fixture();sheet=snap["sheets"][0]
        sheet["max_column"]=3;response["tables"][0]["range"]="A3:C3"
        sheet["cells"]["C3"]={**sheet["cells"]["B3"],"column":3}
        owned=[a for a in sheet["cells"] if a!="C3"]
        with patch.object(engine,"_http",return_value=response):
            result=engine._request_layout("布局",{"task":"layout","candidate_cells":owned,"round":1},sheet,owned)
        self.assertEqual(result["tables"][0]["range"],"A3:C3")
        proof=json.loads(pathlib.Path(result["layout_recovery"][0]["recovery_path"]).read_text(encoding="utf-8-sig"))
        self.assertEqual(set(proof["candidate_cells"]),set(owned));self.assertNotIn("C3",proof["candidate_cells"])

    def test_classification_retains_projected_layout_text_outside_business_bbox(self):
        from unittest.mock import patch
        from 附注更新.语义 import _source_text_rows
        data=snapshot();sheet=data["sheets"][0];sheet["max_row"]=6
        sheet["cells"]["A6"]={"row":6,"column":1,"value":"权益内部结转","formula":None}
        sheet["cells"]["B6"]={"row":6,"column":2,"value":"=B3","formula":"=B3"}
        table={"range":"B3:B3","note_ids":["C-N001"],"scope":"consolidated","table_semantic":"余额","header_cells":["A1","A3","B2"]}
        projection=_source_text_rows(dict(reversed(list(sheet["cells"].items()))))
        self.assertEqual([r["row"] for r in projection],[1,2,3,6])
        self.assertEqual(projection[-1],{"row":6,"cells":[["A6","权益内部结转"]]})
        context={"document":[],"preceding_cells":copy.deepcopy(sheet["cells"]),"source_text_rows":projection};before=copy.deepcopy(context)
        engine=self.engine();calls=[]
        def reply(system,payload):
            calls.append(copy.deepcopy(payload))
            return {"table_ids":["C-N001-T001"]} if payload["task"]=="tables" else {"mappings":[mapping()],"excluded":[],"unresolved":[]}
        with patch.object(engine,"_http",side_effect=reply):result=engine._classify(sheet,table,["B3"],[],context)
        self.assertEqual(len(result["mappings"]),1)
        self.assertEqual([p["task"] for p in calls],["tables","classify","classify"])
        for payload in calls:
            self.assertEqual(payload["context"]["source_text_rows"],projection)
            self.assertNotIn("preceding_cells",payload["context"])
            self.assertNotIn("A6",payload["cells"])
            if payload["task"]=="classify":self.assertEqual(payload["candidate_cells"],["B3"])
        self.assertEqual(context,before)

    def test_word_projected_layout_text_excludes_other_physical_tables(self):
        from unittest.mock import patch
        from 附注更新.语义 import _source_text_rows
        values={"A1":"前表标题","A2":"另一表内容","B3":100,"A7":"当前表标题","A8":"项目","B8":"期末余额","A9":"银行存款","B9":200}
        sheet={"name":"表一","max_row":9,"max_column":2,"merges":[],"word_table_ranges":[{"range":"A2:B3"},{"range":"A8:B9"}],
               "cells":{a:{"row":int(a[1:]),"column":ord(a[0])-64,"value":v,"cached_value":v,"formula":None} for a,v in values.items()}}
        table={"range":"B9:B9","note_ids":["C-N001"],"scope":"consolidated","table_semantic":"余额","header_cells":["A7","A8","B8","A9"]}
        engine=self.engine();calls=[];context={"document":[],"source_text_rows":_source_text_rows(sheet["cells"])}
        def reply(system,payload):
            calls.append(copy.deepcopy(payload))
            return {"table_ids":["C-N001-T001"]} if payload["task"]=="tables" else {"mappings":[{**mapping(),"cell":"B9"}],"excluded":[],"unresolved":[]}
        with patch.object(engine,"_http",side_effect=reply):engine._classify(sheet,table,["B9"],[],context)
        for payload in calls:
            shown={a for row in payload["context"]["source_text_rows"] for a,_ in row["cells"]}
            self.assertEqual(shown,{"A7","A8","B8","A9"})
        bad=copy.deepcopy(context);bad["source_text_rows"][0]["cells"][0][1]="未见原文"
        with self.assertRaisesRegex(SemanticError,"原文.*不一致"):
            engine._classify(sheet,table,["B9"],[],bad)

    def test_layout_packets_keep_physical_word_tables_and_outside_titles_separate(self):
        from 附注更新.语义 import _layout_packets
        sheet={"word_table_ranges":[{"range":"B4:C5"},{"range":"B9:C10"}]}
        ordered=["B3","B4","C4","B5","C5","B8","B9","C9","B10","C10"]
        packets=list(_layout_packets(sheet,ordered))
        self.assertEqual([p[0] for p in packets],[["B3"],["B4","C4","B5","C5"],["B8"],["B9","C9","B10","C10"]])
        self.assertEqual([p[1] for p in packets],[None,(4,2,5,3),None,(9,2,10,3)])
        self.assertEqual([a for p,_ in packets for a in p],ordered)
        plain=["A"+str(r) for r in range(1,482)]
        plain_packets=list(_layout_packets({},plain))
        self.assertEqual([a for p,_ in plain_packets for a in p],plain)
        self.assertTrue(all(0<len(p)<=240 for p,_ in plain_packets))
        with self.assertRaisesRegex(SemanticError,"物理表.*重叠"):
            list(_layout_packets({"word_table_ranges":[{"range":"A1:B3"},{"range":"A2:B4"}]},["A2"]))

    def test_recognize_word_packet_keeps_own_title_and_headers_after_240_cells(self):
        from unittest.mock import patch
        engine=self.engine();data=snapshot();sheet=data["sheets"][0]
        values={"A1":"第一张表标题","A2":"项目","B2":"期末余额","A165":"第二张表标题","A166":"第二张表项目","B166":"第二张表期间","A167":"原行项","B167":20}
        for r in range(3,163):values.update({f"A{r}":"实际业务项目",f"B{r}":r})
        sheet.update(max_row=167,max_column=2,word_table_ranges=[{"range":"A2:B162"},{"range":"A166:B167"}],
            cells={a:{"row":int(a[1:]),"column":ord(a[0])-64,"value":v,"cached_value":v,"formula":None} for a,v in values.items()})
        docs=[{"sheet":sheet["name"],"first_row":2,"row_count":161,"title":"原文第一表标题"},
              {"sheet":sheet["name"],"first_row":166,"row_count":2,"title":"原文第二表标题"}]
        data["context"]=docs+["用户确认人民币元"]
        calls=[]
        def layout(system,payload,current,owned):
            calls.append(copy.deepcopy(payload))
            result={"tables":[],"non_business":[],"unresolved":[{"range":a+":"+a,"reason":"测试保留待核实","evidence_cells":[a]} for a in owned],"carry_context":"不得继承其他表的自由摘要"}
            return engine._validate_layout(result,current,owned)
        with patch.object(engine,"_request_layout",side_effect=layout):result=engine.recognize(data)
        self.assertEqual(len(result["unresolved"]),len(values))
        self.assertEqual(len({i["cell"] for i in result["unresolved"]}),len(values))
        first=[p for p in calls if "B3" in p["candidate_cells"] or "B162" in p["candidate_cells"]]
        self.assertEqual(len(first),2)
        for p in first:
            self.assertLessEqual(len(p["candidate_cells"]),240)
            self.assertTrue({"A1","A2","B2"}<=set(p["cells"]))
            self.assertNotIn("A166",p["cells"])
            self.assertEqual(p["context"]["document"],[docs[0],"用户确认人民币元"])
            self.assertFalse(p["context"].get("previous_context"))
        second=next(p for p in calls if "B167" in p["candidate_cells"])
        self.assertTrue({"A165","A166","B166"}<=set(second["cells"]))
        self.assertNotIn("A162",second["cells"])
        self.assertEqual(second["context"]["document"],[docs[1],"用户确认人民币元"])
        self.assertEqual(next(p for p in calls if "A165" in p["candidate_cells"])["candidate_cells"],["A165"])

    def test_word_tables_in_one_packet_use_their_own_classification_context(self):
        from unittest.mock import patch
        engine=self.engine();values={"A1":"前包的其他应收款", "A10":"项目", "B10":"期末余额", "A11":"银行存款", "B11":100,
                                    "A20":"项目", "B20":"期末余额", "A21":"银行存款", "B21":200}
        sheet={"name":"表一","max_row":21,"max_column":2,"merges":[],
               "word_table_ranges":[{"range":"A10:B11"},{"range":"A20:B21"}],
               "cells":{a:{"row":int(a[1:]),"column":ord(a[0])-64,"value":v,"cached_value":v,"formula":None} for a,v in values.items()}}
        docs=[{"sheet":"表一","first_row":row,"first_column":1,"row_count":2,"column_count":2,
               "title":"本表"+str(row),"unit_text":"2025年末，人民币元"} for row in (10,20)]
        facts=["用户确认2025年末人民币元",{"report_scope":"consolidated","basis":"用户确认"}]
        context={"document":docs+facts,"previous_context":"前包其他应收款坏账准备", "preceding_cells":{"A1":sheet["cells"]["A1"]}}
        before=copy.deepcopy(context);calls=[]
        def reply(system,payload):
            calls.append(copy.deepcopy(payload))
            if payload["task"]=="tables":return {"table_ids":["C-N001-T001"]}
            return {"mappings":[{**mapping(),"cell":payload["candidate_cells"][0]}],"excluded":[],"unresolved":[]}
        with patch.object(engine,"_http",side_effect=reply):
            for row in (10,20):
                table={"range":f"A{row}:B{row+1}","note_ids":["C-N001"],"scope":"consolidated","table_semantic":"本表余额","header_cells":[f"A{row}",f"B{row}"]}
                result=engine._classify(sheet,table,[f"B{row+1}"],[],context)
                self.assertEqual(len(result["mappings"]),1)
        self.assertEqual(len(calls),6)
        for payload in calls:
            actual=payload["context"];row=10 if payload["table"]["range"].startswith("A10") else 20
            self.assertNotIn("previous_context",actual)
            self.assertEqual(set(actual["preceding_cells"]),{f"A{row}",f"B{row}"})
            self.assertEqual(actual["document"],[docs[0 if row==10 else 1]]+facts)
        self.assertEqual(context,before)

    def test_word_subregion_keeps_shared_headers_and_same_physical_table_ancestors(self):
        from unittest.mock import patch
        engine=self.engine();values={"A1":"用户明示2025年末人民币元","A2":"另一张表项目","A10":"项目","B10":2025,
                                    "A11":"账面原值合计","A12":"银行存款","B12":100}
        sheet={"name":"表一","max_row":12,"max_column":2,"merges":[],
               "word_table_ranges":[{"range":"A2:B2"},{"range":"A10:B12"}],
               "cells":{a:{"row":int(a[1:]),"column":ord(a[0])-64,"value":v,"cached_value":v,"formula":None} for a,v in values.items()}}
        table={"range":"A12:B12","note_ids":["C-N001"],"scope":"consolidated","table_semantic":"子区段余额","header_cells":["A1","A10","B10"]}
        context={"document":["用户确认2025年末人民币元"],"previous_context":"另一张表","preceding_cells":{"A2":sheet["cells"]["A2"]}}
        calls=[]
        def reply(system,payload):
            calls.append(copy.deepcopy(payload))
            if payload["task"]=="tables":return {"table_ids":["C-N001-T001"]}
            return {"mappings":[{**mapping(),"cell":"B12"}],"excluded":[],"unresolved":[]}
        with patch.object(engine,"_http",side_effect=reply):engine._classify(sheet,table,["B12"],[],context)
        for payload in calls:
            self.assertEqual(set(payload["context"]["preceding_cells"]),{"A1","A10","B10","A11"})
            self.assertEqual(payload["context"]["preceding_cells"]["B10"]["value"],2025)
            self.assertTrue(set(table["header_cells"])<=set(payload["cells"]))
            self.assertNotIn("previous_context",payload["context"])
            if payload["task"]=="classify":self.assertEqual(payload["candidate_cells"],["B12"])

    def test_classification_context_is_unchanged_without_unique_word_table(self):
        from unittest.mock import patch
        for ranges in (None,[],[{"range":"A20:B23"}],[{"range":"A1:B3"},{"range":"A1:B4"}]):
            with self.subTest(ranges=ranges):
                engine=self.engine();sheet=snapshot()["sheets"][0]
                if ranges is not None:sheet["word_table_ranges"]=ranges
                table={"range":"A1:B3","note_ids":["C-N001"],"scope":"consolidated","table_semantic":"货币资金","header_cells":["A1","B2"]}
                context={"document":["已确认期间单位"],"previous_context":"普通Excel前文","preceding_cells":{"Z9":{"value":"跨区说明"}}};before=copy.deepcopy(context);calls=[]
                def reply(system,payload):
                    calls.append(copy.deepcopy(payload))
                    if payload["task"]=="tables":return {"table_ids":["C-N001-T001"]}
                    return {"mappings":[mapping()],"excluded":[],"unresolved":[]}
                with patch.object(engine,"_http",side_effect=reply):engine._classify(sheet,table,["B3"],[],context)
                for payload in calls:self.assertEqual(payload["context"],before)
                self.assertEqual(context,before)

    def test_word_context_is_limited_to_current_ranges_and_preserves_global_facts(self):
        data=snapshot()
        current={"sheet":"表一","first_row":1,"first_column":1,"row_count":3,"column_count":2,"title":"本表说明","chapter_context":["公司主体","整段继承说明"]}
        outside={**current,"first_row":20,"title":"后面的另一张表"}
        other_sheet={**current,"sheet":"表二","title":"其他工作表"}
        beside={**current,"first_column":20,"title":"同排另一张表"}
        global_fact={"report_scope":"consolidated","basis":"用户确认事实"}
        data["context"]=[current,outside,other_sheet,beside,"全局币种和期间补充",global_fact]
        before=copy.deepcopy(data["context"])
        result=self.engine().recognize(data)
        self.assertTrue(result["complete"])
        for payload in self.server.received:
            document=payload["context"]["document"]
            self.assertIn(current,document);self.assertIn(global_fact,document)
            self.assertIn("全局币种和期间补充",document)
            self.assertNotIn(outside,document);self.assertNotIn(other_sheet,document)
            if payload["task"] in {"classify","tables"}:self.assertNotIn(beside,document)
        self.assertEqual(data["context"],before)

    def test_classification_preserves_complete_aggregation_and_calculation(self):
        from unittest.mock import patch
        base=json.loads(self.gold.read_text(encoding="utf-8-sig"));rows=[]
        # 真实余额构成中，其他应收款组件与股利均属于合计，不能只传同名行标签。
        for suffix,parent in (("S001","S005"),("S003","S006"),("S005",None),("S006",None)):
            row=copy.deepcopy(base);row["id"]="C-N020-T002-"+suffix
            row["note"]={"id":"C-N020","name":"其他应收款"};row["table"]={"id":"C-N020-T002","name":"余额构成"}
            component=parent is not None
            row["slot"]={"name":("其他应收款" if component else "合计")+" / "+("期末余额" if suffix in {"S001","S005"} else "期初余额")}
            row["aggregation"]={"role":"component" if component else "total","parent_slot_id":"C-N020-T002-"+parent if parent else None,
                                "parent_variants":[],"component_slot_ids":[],"component_variants":[]}
            if not component:
                members=["S001","S002","S007"] if suffix=="S005" else ["S003","S004","S008"]
                row["aggregation"]["component_variants"]=[{"source_table_ids":["000593-C-0023","600137-C-0012","600712-C-0011","603077-C-0018"],"component_slot_ids":["C-N020-T002-"+member for member in members]}]
            row["calculation"]={"expression":None,"component_slot_ids":[]};rows.append(row)
        # 合成非空公式只用于检查原样传递，不把它当成真实标准修订。
        rows[-1]["calculation"]={"expression":"S003 + S004 + S008","component_slot_ids":["C-N020-T002-S003","C-N020-T002-S004","C-N020-T002-S008"]}
        path=self.work/("完整业务关系_"+uuid.uuid4().hex+".jsonl")
        path.write_text("\n".join(json.dumps(row,ensure_ascii=False) for row in rows),encoding="utf-8-sig")
        engine=SemanticEngine(self.engine().settings,str(path),str(self.work/uuid.uuid4().hex))
        sheet=snapshot()["sheets"][0];table={"range":"A1:B3","scope":"consolidated","note_ids":["C-N020"],"table_semantic":"其他应收款余额构成","header_cells":["A1","B2"]};calls=[]
        def reply(system,payload):
            calls.append(copy.deepcopy(payload))
            if payload["task"]=="tables":return {"table_ids":["C-N020-T002"]}
            return {"mappings":[],"excluded":[],"unresolved":[{"cell":"B3","reason":"本地只验证定义传输，不模拟业务结论"}]}
        with patch.object(engine,"_http",side_effect=reply),patch("urllib.request.build_opener",side_effect=AssertionError("禁止网络")):
            result=engine._classify(sheet,table,["B3"],[],{})
        self.assertEqual(len(result["unresolved"]),1)
        actual=[payload for payload in calls if payload["task"]=="classify"]
        self.assertEqual([payload["round"] for payload in actual],[1,2])
        for payload in actual:
            sent={slot["id"]:slot for slot in payload["gold_slots"]}
            for row in rows:
                for key in ("aggregation","calculation"):
                    with self.subTest(round=payload["round"],slot=row["id"],field=key):
                        expected=copy.deepcopy(row[key])
                        if key=="aggregation":
                            # 来源表编号留在完整标准中；汇总成员及全部业务限定仍必须完整发送。
                            for variant in expected["component_variants"]:
                                variant.pop("source_table_ids",None)
                        self.assertEqual(sent[row["id"]].get(key),expected)

    def test_standalone_table_catalogue_explains_definition_and_allowed_scope(self):
        from unittest.mock import patch
        engine=self.engine();item=mapping();sheet=snapshot()["sheets"][0]
        row=json.loads(self.gold.read_text(encoding="utf-8-sig"))
        row["applicability"]={"report_scopes":[{"scope":"standalone","reason":"单户可复用","reviewed_by":"测试审核","reviewed_at":"2026-09-19"}]}
        path=self.work/("单户表目录_"+uuid.uuid4().hex+".jsonl")
        path.write_text(json.dumps(row,ensure_ascii=False),encoding="utf-8-sig")
        engine=SemanticEngine(engine.settings,str(path),str(self.work/uuid.uuid4().hex))
        table={"range":"A1:B3","scope":"standalone","note_ids":["C-N001"],"table_semantic":"货币资金","header_cells":["A1","B2"]}
        def reply(system,payload):
            if payload["task"]=="tables":
                self.assertEqual(payload.get("actual_report_scope"),"standalone")
                self.assertEqual(payload["catalogue"][0].get("definition_scope"),"consolidated")
                self.assertEqual(payload["catalogue"][0].get("allowed_scopes"),["consolidated","standalone"])
                return {"table_ids":[row["table"]["id"]]}
            return {"mappings":[{**item,"scope":"standalone"}],"excluded":[],"unresolved":[]}
        with patch.object(engine,"_http",side_effect=reply):
            result=engine._classify(sheet,table,["B3"],[],{})
        self.assertEqual(result["unresolved"],[])
        self.assertEqual(result["mappings"][0]["definition_scope"],"consolidated")

    def test_header_and_title_agreement_requires_real_nonblank_text(self):
        from unittest.mock import patch
        for value,expected in (("货币资金",True),(123.45,False),(None,False),("",False),("—",False)):
            with self.subTest(value=value):
                data=snapshot();sheet=data["sheets"][0]
                sheet["cells"]["A1"].update(value=value,cached_value=value)
                engine=self.engine()
                def reply(system,payload):
                    if payload["task"]=="layout":
                        return {"tables":[{"range":"A1:B3","note_ids":["C-N001"],"scope":"consolidated","table_semantic":"货币资金","header_cells":["A1","A2","B2"]}]}
                    if payload["task"]=="tables":return {"table_ids":["C-N001-T001"]}
                    excluded=[{"cell":a,"category":"header","reason":"真实原表文字标签"} for a in payload["candidate_cells"] if a!="B3"]
                    for item in excluded:
                        if item["cell"]=="A1":item["category"]="header" if payload["round"]==1 else "title"
                    return {"mappings":[mapping()] if "B3" in payload["candidate_cells"] else [],"excluded":excluded,"unresolved":[]}
                with patch.object(engine,"_http",side_effect=reply):result=engine.recognize(data)
                self.assertEqual(result["complete"],expected)
                if expected:
                    self.assertEqual([m["cell"] for m in result["mappings"]],["B3"])
                    self.assertEqual(result["unresolved"],[])
                else:self.assertIn("A1",[m["cell"] for m in result["unresolved"]])

    def test_failed_nonbusiness_review_retries(self):
        engine = self.engine("retry_title")
        data = snapshot()
        sheet = data["sheets"][0]
        sheet.update(max_row=1,max_column=1,cells={"A1":sheet["cells"]["A1"]})
        self.server.mode = "title_conflict"
        self.assertFalse(engine.recognize(data)["complete"])
        before = len(self.server.received)
        self.server.mode = "title_only"
        self.assertTrue(engine.recognize(data)["complete"])
        self.assertGreater(len(self.server.received),before)

    def test_unresolved_classifications_reuse_verified_layout_on_resume(self):
        engine = self.engine("retry_layout")
        self.server.mode = "unresolved"
        self.assertFalse(engine.recognize(snapshot())["complete"])
        count = sum(p["task"] == "layout" for p in self.server.received)
        self.server.mode = "normal"
        self.assertTrue(engine.recognize(snapshot())["complete"])
        self.assertEqual(sum(p["task"] == "layout" for p in self.server.received), count)

    def test_空布局应进入一次有界纠正并保留原回执证据(self):
        from unittest.mock import patch
        engine=self.engine();calls=[]
        def reply(system,payload):
            calls.append(copy.deepcopy(payload))
            if payload["task"]=="layout":
                if "correction" not in payload:return {"tables":[],"carry_context":"知道是货币资金但未返回业务范围"}
                return {"tables":[{"range":"A1:B3","note_ids":["C-N001"],"scope":"consolidated","table_semantic":"货币资金","header_cells":["A1","A2","B2"]}]}
            if payload["task"]=="tables":return {"table_ids":["C-N001-T001"]}
            return {"mappings":[mapping()],"excluded":[{"cell":a,"category":"header","reason":"真实原表文字标题"} for a in payload["candidate_cells"] if a!="B3"],"unresolved":[]}
        with patch.object(engine,"_http",side_effect=reply):result=engine.recognize(snapshot())
        self.assertTrue(result["complete"])
        layout_calls=[p for p in calls if p["task"]=="layout"]
        self.assertEqual(len(layout_calls),2)
        self.assertEqual(layout_calls[1]["correction"]["previous_result"]["tables"],[])
        self.assertIn("B3",layout_calls[1]["correction"]["validation_error"])
        failures=[json.loads(p.read_text(encoding="utf-8-sig")) for p in engine.work_dir.rglob("失败_*.json")]
        self.assertEqual(len(failures),1);self.assertEqual(failures[0]["result"]["tables"],[])
        corrected=[json.loads(p.read_text(encoding="utf-8-sig")) for p in engine.work_dir.rglob("纠正回执_*.json")]
        self.assertEqual(len(corrected),1);self.assertTrue(corrected[0]["validated"])

    def test_部分布局遗漏候选应标明原地址而不是默认通过(self):
        engine=self.engine();sheet=snapshot()["sheets"][0]
        plan={"tables":[{"range":"A1:B2","note_ids":["C-N001"],"scope":"consolidated","table_semantic":"货币资金","header_cells":["A1","A2","B2"]}]}
        with self.assertRaisesRegex(SemanticError,"遗漏.*2.*A3.*B3"):
            engine._validate_layout(plan,sheet,list(sheet["cells"]))

    def test_布局明确待核实区域保留原财务格并继续已知科目(self):
        from unittest.mock import patch
        data=snapshot();sheet=data["sheets"][0];sheet["max_row"]=4
        for address,value in (("A4","无对应科目的实际业务"),("B4",99)):
            sheet["cells"][address]={**sheet["cells"]["B3"],"row":4,"column":1 if address=="A4" else 2,"value":value,"cached_value":value}
        original=copy.deepcopy(data);engine=self.engine();calls=[]
        def reply(system,payload):
            calls.append(copy.deepcopy(payload))
            if payload["task"]=="layout":return {"tables":[{"range":"A1:B3","note_ids":["C-N001"],"scope":"consolidated","table_semantic":"货币资金","header_cells":["A1","A2","B2"]}],
                "unresolved":[{"range":"A4:B4","reason":"实际业务指标尚无法确定所属金标准科目","evidence_cells":["A4"]}]}
            if payload["task"]=="tables":return {"table_ids":["C-N001-T001"]}
            return {"mappings":[mapping()],"excluded":[{"cell":a,"category":"header","reason":"真实原表文字标签"} for a in payload["candidate_cells"] if a!="B3"],"unresolved":[]}
        with patch.object(engine,"_http",side_effect=reply):
            result=engine.recognize(data)
            self.assertEqual([m["cell"] for m in result["mappings"]],["B3"])
            self.assertEqual({m["cell"] for m in result["unresolved"]},{"A4","B4"})
            self.assertTrue(all(m["evidence_cells"]==["A4"] for m in result["unresolved"]))
            self.assertFalse(result["complete"]);self.assertEqual(sheet["cells"]["B4"]["value"],99)
            final=[m["cell"] for key in ("mappings","excluded","unresolved") for m in result[key]]
            self.assertEqual(set(final),set(sheet["cells"]));self.assertTrue(set(original["sheets"][0]["cells"])<=set(final));self.assertEqual(len(final),len(set(final)))
            before=sum(p["task"]=="layout" for p in calls);engine.recognize(copy.deepcopy(original))
            self.assertGreater(sum(p["task"]=="layout" for p in calls),before)
        task_dir=next(p.parent for p in engine.work_dir.rglob("任务.json") if json.loads(p.read_text(encoding="utf-8-sig"))["payload"]["task"]=="layout")
        self.assertTrue(json.loads((task_dir/"回执.json").read_text(encoding="utf-8-sig"))["retry_required"])
        self.assertTrue(all(not ({"A4","B4"}&set(p["candidate_cells"])) for p in calls if p["task"]=="classify"))

    def test_布局待核实区域不得重叠或缺少真实依据(self):
        engine=self.engine();sheet=snapshot()["sheets"][0]
        for region in ({"range":"A1:B3","reason":"待核实","evidence_cells":["A1"]},
                       {"range":"A1:B3","reason":"","evidence_cells":["A1"]},
                       {"range":"A1:B3","reason":"待核实","evidence_cells":["Z99"]}):
            plan={"tables":[{"range":"A1:B3","note_ids":["C-N001"],"scope":"consolidated","table_semantic":"货币资金","header_cells":["A1"]}],"unresolved":[region]}
            with self.subTest(region=region),self.assertRaises(SemanticError):engine._validate_layout(plan,sheet,list(sheet["cells"]))
        pure={"tables":[],"unresolved":[{"range":"A1:B3","reason":"需核实实际科目","evidence_cells":["A1","A3"]}]}
        engine._validate_layout(pure,sheet,list(sheet["cells"]))

    def test_布局待核实理由不能只有空白即使区域没有重叠(self):
        engine=self.engine();sheet=snapshot()["sheets"][0]
        for reason in ("", " ", "\t\n"):
            plan={"tables":[],"unresolved":[{"range":"A1:B3","reason":reason,"evidence_cells":["A1"]}]}
            with self.subTest(reason=reason),self.assertRaisesRegex(SemanticError,"具体原因"):
                engine._validate_layout(plan,sheet,list(sheet["cells"]))

    def test_新布局按单科目校验但旧回执多科目证据仍可读取(self):
        engine=self.engine();engine.notes["C-N002"]={"id":"C-N002","name":"其他业务"};sheet=snapshot()["sheets"][0]
        plan={"tables":[{"range":"A1:B3","note_ids":["C-N001","C-N002"],"scope":"consolidated","table_semantic":"旧复合业务范围","header_cells":["A1"]}]}
        engine._validate_layout(plan,sheet,list(sheet["cells"]))
        with self.assertRaisesRegex(SemanticError,"一个.*科目"):
            engine._validate_layout(plan,sheet,list(sheet["cells"]),single_note=True)
        plan["tables"][0]["note_ids"]=["C-N001"]
        engine._validate_layout(plan,sheet,list(sheet["cells"]),single_note=True)

    def test_布局编号凭原表头双轮排除而金额仍待核实(self):
        from unittest.mock import patch
        engine=self.engine();values={"A1":"项目","B1":"行次","C1":"金额","A2":"流动资产","B2":1,"C2":0}
        sheet={"name":"布局编号","max_row":2,"max_column":3,"merges":[],"cells":{a:{"row":int(a[1:]),"column":ord(a[0])-64,"value":v,"cached_value":v,"formula":None} for a,v in values.items()}}
        data={"path":"测试.xlsx","sha256":"layout-number","context":[],"sheets":[sheet]};calls=[]
        plan={"tables":[],"non_business":[{"range":"A1:C1","category":"header","reason":"真实列头","evidence_cells":["A1","B1","C1"]},
              {"range":"A2","category":"header","reason":"分区名称","evidence_cells":["A2"]},
              {"range":"B2","category":"header","reason":"B1明确为行次","evidence_cells":["B1"]}],
              "unresolved":[{"range":"C2","reason":"合计金额没有对应科目定义","evidence_cells":["A2","C1"]}]}
        def reply(system,payload):
            self.assertEqual(payload["task"],"layout");calls.append(payload);return copy.deepcopy(plan)
        with patch.object(engine,"_http",side_effect=reply):result=engine.recognize(data)
        self.assertEqual(len(calls),2);self.assertFalse(result["complete"])
        self.assertEqual([m["cell"] for m in result["unresolved"]],["C2"])
        self.assertTrue(next(m for m in result["excluded"] if m["cell"]=="B2")["reviewed"])
        self.assertEqual(sheet["cells"]["C2"]["value"],0)
        bad=copy.deepcopy(plan);bad["non_business"][-1]["evidence_cells"]=["B2"]
        with self.assertRaisesRegex(SemanticError,"B2"):
            engine._validate_layout(bad,sheet,list(values))
        sheet["cells"]["C1"]["value"]="期末余额";sheet["cells"]["C2"].update(value=2025,cached_value=2025)
        fake={"non_business":[{"range":"C2","category":"header","reason":"冒认年份表头","evidence_cells":["C1","C2"]}]}
        with self.assertRaisesRegex(SemanticError,"C2"):
            engine._nonbusiness_cells(fake,sheet,["C2"])

    def test_布局唯一纠正可同时看到多个误排金额和范围重叠(self):
        engine=self.engine();sheet=snapshot()["sheets"][0]
        sheet["max_row"]=4
        sheet["cells"]["B4"]={**sheet["cells"]["B3"],"row":4,"value":0,"cached_value":0}
        plan={"tables":[{"range":"A1:B3","note_ids":["C-N001"],"scope":"consolidated","table_semantic":"货币资金","header_cells":["A1","A2","B2"]}],
              "non_business":[{"range":"B3:B4","category":"header","reason":"错误地排除金额","evidence_cells":["B2"]}]}
        with self.assertRaises(SemanticError) as raised:engine._validate_layout(plan,sheet,list(sheet["cells"]))
        message=str(raised.exception)
        self.assertIn("B3",message);self.assertIn("B4",message);self.assertIn("重叠",message)

    def test_layout_missing_cells_is_retried_on_resume(self):
        engine=self.engine();self.server.mode="layout_omission"
        self.assertFalse(engine.recognize(snapshot())["complete"])
        before=sum(p["task"]=="layout" for p in self.server.received)
        self.server.mode="normal"
        self.assertTrue(engine.recognize(snapshot())["complete"])
        self.assertGreater(sum(p["task"]=="layout" for p in self.server.received),before)

    def test_reference_does_not_leak_positions_or_values(self):
        self.engine().recognize(snapshot(), [{**mapping(), "sheet": "禁止泄露目标表名", "cell": "XFD1048576", "value": 918273645}])
        content = json.dumps(self.server.received, ensure_ascii=False)
        self.assertNotIn("禁止泄露目标表名", content)
        self.assertNotIn("XFD1048576", content)
        self.assertNotIn("918273645", content)



    def blank_fixture(self, name, merged=False, columns=2, old_value=None):
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
        path = self.work / (name + "_" + uuid.uuid4().hex[:8] + ".xlsx")
        book = Workbook(); sheet = book.active; sheet.title = "表一"
        if columns > 2:
            for column in range(1, columns + 1):
                sheet.cell(1, column, "项目" if column == 1 else "期末指标" + str(column))
            sheet.cell(2, 1, "乙"); sheet.cell(2, 2, 10)
            area = "A1:" + get_column_letter(columns) + "2"
        elif merged:
            sheet.merge_cells("A1:B1"); sheet["A1"] = "合并货币资金，2025年，人民币元"
            sheet.append(["项目", "期末余额"]); sheet.append(["乙", old_value])
            area = "A1:B3"
        else:
            sheet.append(["项目", "期末余额"]); sheet.append(["甲", 10]); sheet.append(["乙", old_value])
            area = "A1:B3"
        book.save(path); book.close()
        return path, area

    def blank_reply(self, area, calls, omit=None):
        def reply(system, payload):
            calls.append(copy.deepcopy(payload))
            if payload["task"] == "layout":
                return {"tables": [{"range": area, "note_ids": ["C-N001"], "scope": "consolidated",
                                    "table_semantic": "业务金额明细", "header_cells": ["A1"]}]}
            if payload["task"] == "tables":
                return {"table_ids": ["C-N001-T001"]}
            result = {"mappings": [], "excluded": [], "unresolved": []}
            for address in payload["candidate_cells"]:
                if address == omit:
                    continue
                cell = payload["cells"][address]
                if cell["row"] >= 2 and cell["column"] >= 2 and not isinstance(cell.get("value"), str):
                    item = mapping(); item["cell"] = address
                    row_label = payload["cells"].get("A" + str(cell["row"]), {}).get("value", "项目")
                    item["dimensions"]["counterparty"] = row_label
                    item["dimensions"]["business_category"] = "指标" + str(cell["column"])
                    item["row_label"] = row_label
                    result["mappings"].append(item)
                else:
                    result["excluded"].append({"cell": address, "category": "label", "reason": "该格是原表文字标签"})
            return result
        return reply

    def test_unformatted_missing_amount_is_classified_but_cannot_clear_target(self):
        from unittest.mock import patch
        from openpyxl import load_workbook
        from 附注更新.表格 import read_workbook, build_update_plan, write_updated_workbook, actual_value
        b_path, area = self.blank_fixture("业务空白B")
        a_path, _ = self.blank_fixture("已有金额A", old_value=7)
        b = read_workbook(b_path); a = read_workbook(a_path)
        self.assertNotIn("B3", b["sheets"][0]["cells"])
        calls = []; engine = self.engine()
        with patch.object(engine, "_http", side_effect=self.blank_reply(area, calls)):
            recognized = engine.recognize(b)
        self.assertTrue(recognized["complete"])
        self.assertIn("B3", [m["cell"] for m in recognized["mappings"]])
        self.assertIsNone(actual_value(b, "表一", "B3"))
        proof = b["sheets"][0]["cells"]["B3"]["blank_evidence"]
        self.assertEqual(proof["source_hash"], b["sha256"])
        self.assertEqual(proof["table_range"], area)
        plan = build_update_plan(a, b, copy.deepcopy(recognized["mappings"]), recognized["mappings"])
        self.assertEqual([issue["kind"] for issue in plan["issues"]],["source_unknown_blank"])
        self.assertFalse(any(update["cell"]=="B3" for update in plan["updates"]))
        book=load_workbook(a_path);self.assertEqual(book["表一"]["B3"].value,7);book.close()

    def test_materialized_blank_skips_merged_placeholders_and_reuses_fresh_cache(self):
        from unittest.mock import patch
        from 附注更新.表格 import read_workbook
        path, area = self.blank_fixture("合并与缓存", merged=True)
        calls = []; engine = self.engine("补格缓存")
        with patch.object(engine, "_http", side_effect=self.blank_reply(area, calls)):
            first = read_workbook(path); result = engine.recognize(first)
            self.assertTrue(result["complete"])
            self.assertNotIn("B1", first["sheets"][0]["cells"])
            self.assertIn("B3", first["sheets"][0]["cells"])
            count = len(calls)
            fresh = read_workbook(path)
            self.assertEqual(engine.recognize(fresh)["mappings"], result["mappings"])
            self.assertEqual(len(calls), count)
            self.assertEqual(engine.recognize(fresh)["mappings"], result["mappings"])
            self.assertEqual(len(calls), count)

    def test_materialized_tail_columns_are_reviewed_once_across_layout_packages(self):
        from collections import Counter
        from unittest.mock import patch
        from 附注更新.表格 import read_workbook
        path, area = self.blank_fixture("宽表跨包", columns=260)
        data = read_workbook(path); calls = []; engine = self.engine()
        with patch.object(engine, "_http", side_effect=self.blank_reply(area, calls)):
            result = engine.recognize(data)
        self.assertTrue(result["complete"])
        addresses = set(data["sheets"][0]["cells"])
        self.assertEqual(len(addresses), 520)
        self.assertGreaterEqual(sum(p["task"] == "layout" for p in calls), 2)
        for round_number in (1, 2):
            counts = Counter(a for p in calls if p["task"] == "classify" and p["round"] == round_number for a in p["candidate_cells"])
            self.assertEqual(set(counts), addresses)
            self.assertEqual(set(counts.values()), {1})
        final = [x["cell"] for key in ("mappings", "excluded", "unresolved") for x in result[key]]
        self.assertEqual(len(final), len(set(final)))

    def test_failed_materialized_blank_classification_blocks_completion(self):
        from unittest.mock import patch
        from 附注更新.表格 import read_workbook
        path, area = self.blank_fixture("补格遗漏")
        data = read_workbook(path); engine = self.engine()
        with patch.object(engine, "_http", side_effect=self.blank_reply(area, [], omit="B3")):
            result = engine.recognize(data)
        self.assertFalse(result["complete"])
        self.assertIn("B3", [x["cell"] for x in result["unresolved"]])


class 行政数字排除测试(unittest.TestCase):
    def sheet(self,label,value="610000",address="B1",merges=None):
        from 附注更新.语义 import _address
        return {"name":"封面代码","max_row":3,"max_column":4,"merges":merges or [],
                "cells":{a:{"row":_address(a)[0],"column":_address(a)[1],"value":v,"cached_value":v,"formula":None}
                         for a,v in {"A1":label,address:value}.items()}}

    def test_administrative_integer_values_require_exact_adjacent_original_labels(self):
        for label,value in (("邮政编码","610000"),("邮编",610000),("电话号码（区号）","028"),("电话区号","028"),
                            ("电话号码","68174861"),("联系电话","68174861"),("手机号码","13800000000"),("设立年份","2020"),("成立年份",2020)):
            for category in ("label","annotation"):
                with self.subTest(label=label,category=category):
                    sheet=self.sheet(label,value);before=copy.deepcopy(sheet)
                    SemanticEngine._check_exclusion("B1",{"category":category,"evidence_cells":["B1","A1"]},sheet)
                    self.assertEqual(sheet,before)
        merged=self.sheet("电话 号码(区号)：","028",address="C1",merges=["A1:B1"])
        merged["word_table_ranges"]=[{"range":"A1:C1"}]
        SemanticEngine._check_exclusion("C1",{"category":"annotation","evidence_cells":["A1"]},merged)

    def test_financial_values_or_unproven_administrative_labels_remain_protected(self):
        raw={"category":"label","evidence_cells":["B1","A1"]}
        for label in ("期末余额","数量","比例","电话费金额","设立费用","年度","代码","编号","联系电话费用","企业代码（邮政编码）","备用码","手机费","手机号码费用","手机号码数量"):
            with self.subTest(label=label),self.assertRaises(SemanticError):
                SemanticEngine._check_exclusion("B1",raw,self.sheet(label))
        for value in ("12.3","-28","2020.0"):
            with self.subTest(value=value),self.assertRaises(SemanticError):
                SemanticEngine._check_exclusion("B1",raw,self.sheet("设立年份",value))
        for value in ("1600",str(__import__("datetime").date.today().year+1)):
            with self.subTest(year=value),self.assertRaises(SemanticError):
                SemanticEngine._check_exclusion("B1",raw,self.sheet("成立年份",value))
        for address in ("B2","C1"):
            with self.subTest(address=address),self.assertRaises(SemanticError):
                SemanticEngine._check_exclusion(address,raw,self.sheet("电话号码","68174861",address))
        for evidence in ([],["B1"]):
            with self.subTest(evidence=evidence),self.assertRaises(SemanticError):
                SemanticEngine._check_exclusion("B1",{**raw,"evidence_cells":evidence},self.sheet("邮编"))
        for category in ("header","title","unit"):
            with self.subTest(category=category),self.assertRaises(SemanticError):
                SemanticEngine._check_exclusion("B1",{**raw,"category":category},self.sheet("邮编"))
        formula=self.sheet("邮编");formula["cells"]["B1"]["formula"]="=C1"
        with self.assertRaises(SemanticError):SemanticEngine._check_exclusion("B1",raw,formula)
        other_table=self.sheet("邮编");other_table["word_table_ranges"]=[{"range":"A1:A1"},{"range":"B1:B1"}]
        with self.assertRaises(SemanticError):SemanticEngine._check_exclusion("B1",raw,other_table)
        wordless=self.sheet("邮编");wordless["word_table_ranges"]=[{"range":"A2:B2"}]
        with self.assertRaises(SemanticError):SemanticEngine._check_exclusion("B1",raw,wordless)


class 编号格排除测试(unittest.TestCase):
    def fixture(self):
        cells={"A1":"项目","B1":"行次","C1":"金额","B2":"1","B3":2,"C2":1}
        sheet={"name":"编号测试","max_row":5,"max_column":4,"merges":[],
               "cells":{a:{"row":int(a[1:]),"column":ord(a[0])-64,"value":v,"cached_value":v,"formula":None} for a,v in cells.items()}}
        table={"range":"A1:D5","header_cells":["A1","B1","C1"]}
        return sheet,table

    def test_原Z01两轮行次号名称不同仍能按同一原证据确认(self):
        # 来自同40格真实回执：第一轮label、第二轮annotation；只保留本用例所需原列。
        from 附注更新.语义 import _address
        values={"B4":"行次","F4":"行次",**{f"B{r}":str(r-4) for r in range(5,12)},
                **{f"F{r}":str(r+73) for r in range(5,12)}}
        sheet={"name":"Z01 资产负债表(企财01表)","max_row":11,"max_column":8,"merges":[],
               "cells":{a:{"row":_address(a)[0],"column":_address(a)[1],"value":v,"cached_value":v,"formula":None} for a,v in values.items()}}
        before=copy.deepcopy(sheet)
        for row in range(7,12):
            for column in ("B","F"):
                address=f"{column}{row}";header=column+"4"
                first={"cell":address,"range":address+":"+address,"category":"label",
                       "reason":"该格为行次编号，依据列头行次判定为行政排序号，非财务数值","evidence_cells":[header,address]}
                second={**first,"category":"annotation","reason":"该格位于行次表头下方，属于报表格式模板行次编号，非财务数值"}
                for raw in (first,second):
                    SemanticEngine._check_exclusion(address,raw,sheet,{"range":raw["range"],"header_cells":raw["evidence_cells"]},allow_period_header=False)
                with self.subTest(cell=address):
                    self.assertEqual(SemanticEngine._exclusion_kind(first,sheet),SemanticEngine._exclusion_kind(second,sheet))
                    self.assertEqual(second["category"],"annotation")
        self.assertEqual(sheet,before)

    def test_归一化不能扩展普通说明行政号码或无编号证据的财务数字(self):
        for change in ("amount","wrong_column","year","formula","blank_gap","other_word_table","annotation_text","administrative"):
            with self.subTest(change=change):
                sheet,table=self.fixture();address="B3";evidence=["B1"]
                if change=="amount":sheet["cells"]["B1"]["value"]="金额"
                if change=="wrong_column":address="C2"
                if change=="year":sheet["cells"]["B1"]["value"]="年度";sheet["cells"][address]["value"]=2025
                if change=="formula":sheet["cells"][address]["formula"]="=2"
                if change=="blank_gap":sheet["cells"]["B2"]["value"]=None
                if change=="other_word_table":sheet["word_table_ranges"]=[{"range":"A1:D2"},{"range":"A3:D5"}]
                if change=="annotation_text":sheet["cells"][address]["value"]="报表填报说明"
                if change=="administrative":
                    sheet["cells"]["A3"]={"row":3,"column":1,"value":"电话号码","formula":None};evidence=["A3"]
                    sheet["cells"][address]["value"]="68174861"
                raw={"cell":address,"category":"annotation","evidence_cells":evidence}
                self.assertEqual(SemanticEngine._exclusion_kind(raw,sheet),"annotation")

    def test_真实编号表头下的字符串和数值整数可排除(self):
        for category in ("header","label","annotation"):
            for value in ("1",2,3.0,"2025",2025,2025.0,"+2025"):
                with self.subTest(category=category,value=value):
                    sheet,table=self.fixture();sheet["cells"]["B2"]["value"]=value
                    SemanticEngine._check_exclusion("B2",{"category":category,"evidence_cells":["B1"]},sheet,table)
        sheet,table=self.fixture();sheet["cells"]["B1"]["value"]="序 号"
        sheet["merges"]=["B1:C1"];sheet["cells"].pop("C1")
        table["header_cells"]=["A1","B1"]
        SemanticEngine._check_exclusion("C2",{"category":"label","evidence_cells":["B1"]},sheet,table)

    def test_没有完整编号结构依据时不得排除数字(self):
        for change in ("amount","year","year_float","year_plus","decimal","percent","wrong_column","outside_table",
                       "missing_evidence","missing_header","wrong_header","formula","no_table"):
            with self.subTest(change=change):
                sheet,table=self.fixture();raw={"category":"label","evidence_cells":["B1"]};address="B2"
                if change=="amount":sheet["cells"]["B1"]["value"]="金额"
                if change=="year":
                    sheet["cells"]["B2"]["value"]="2025";sheet["cells"]["B1"]["value"]="年度"
                if change=="year_float":
                    sheet["cells"]["B2"]["value"]=2025.0;sheet["cells"]["B1"]["value"]="期末金额"
                if change=="year_plus":
                    sheet["cells"]["B2"]["value"]="+2025";sheet["cells"]["B1"]["value"]="年度"
                if change=="decimal":sheet["cells"]["B2"]["value"]=1.25
                if change=="percent":sheet["cells"]["B2"]["value"]="1%"
                if change=="wrong_column":address="C2"
                if change=="outside_table":table["range"]="A1:D1"
                if change=="missing_evidence":raw["evidence_cells"]=[]
                if change=="missing_header":table["header_cells"]=["A1","C1"]
                if change=="wrong_header":raw["evidence_cells"]=["C1"]
                if change=="formula":sheet["cells"]["B2"].update(formula="=1",cached_value=1)
                if change=="no_table":table=None
                with self.assertRaisesRegex(SemanticError,"数值格不能"):
                    SemanticEngine._check_exclusion(address,raw,sheet,table)

    def test_共享编号表头可位于子区段上方但不能跨Word表或空行(self):
        for physical in (None,[{"range":"A1:D5"}]):
            with self.subTest(physical=physical):
                sheet,table=self.fixture();table["range"]="A3:D5"
                if physical is not None:sheet["word_table_ranges"]=physical
                SemanticEngine._check_exclusion("B3",{"category":"label","evidence_cells":["B1"]},sheet,table)
        for boundary in ("other_word_table","blank_gap"):
            with self.subTest(boundary=boundary):
                sheet,table=self.fixture();table["range"]="A3:D5"
                if boundary=="other_word_table":sheet["word_table_ranges"]=[{"range":"A1:D2"},{"range":"A3:D5"}]
                else:sheet["cells"]["B2"]["value"]=None
                with self.assertRaisesRegex(SemanticError,"数值格不能"):
                    SemanticEngine._check_exclusion("B3",{"category":"label","evidence_cells":["B1"]},sheet,table)

    def test_更近冲突表头或越界合并不能把金额当编号(self):
        for merge,declared in ((False,True),(False,False),(True,True),(True,False)):
            with self.subTest(merge=merge,declared=declared):
                sheet,table=self.fixture();sheet["cells"]["B2"]["value"]="期末余额"
                if merge:
                    sheet["merges"]=["B2:C2"];sheet["cells"].pop("C2")
                    sheet["cells"]["C3"]={"row":3,"column":3,"value":1,"formula":None}
                    sheet["merges"].append("B1:C1");sheet["cells"].pop("C1")
                    table["header_cells"]=["B1"];address="C3"
                else:address="B3"
                if declared:table["header_cells"].append("B2")
                with self.assertRaisesRegex(SemanticError,"数值格不能"):
                    SemanticEngine._check_exclusion(address,{"category":"annotation","evidence_cells":["B1"]},sheet,table)
        sheet,table=self.fixture();sheet["merges"]=["B1:E1"];sheet["cells"].pop("C1")
        table["header_cells"]=["B1"]
        with self.assertRaisesRegex(SemanticError,"数值格不能"):
            SemanticEngine._check_exclusion("C2",{"category":"label","evidence_cells":["B1"]},sheet,table)


class SameMeaningNoteTests(unittest.TestCase):
    """同名同含义科目仅扩大合法候选；所有模型回执均在本地模拟。"""
    def fixture(self, change=None):
        work=ROOT/"测试输出"/("同名科目_"+uuid.uuid4().hex);work.mkdir(parents=True)
        base={"id":"C-N001-T001-S001","status":"active","scope":"consolidated",
              "note":{"id":"C-N001","name":"货币资金","meaning":"货币资金余额构成"},
              "table":{"id":"C-N001-T001","name":"余额构成","meaning":"资金余额"},
              "slot":{"name":"银行存款 / 期末余额"},"row_path":["银行存款"],"column_path":["期末余额"],
              "value_type":"monetary","dimensions":[{"name":n,"required":True} for n in ("currency","period","scale","unit")],
              "applicability":{"report_scopes":[{"scope":"standalone","reason":"测试：余额适用单户","reviewed_by":"本地模拟审核","reviewed_at":"2026-09-20"}]}}
        peer=copy.deepcopy(base);peer.update(id="P-N009-T001-S001",scope="parent")
        peer["note"]["id"]="P-N009";peer["table"]["id"]="P-N009-T001"
        if change:change(base,peer)
        gold=work/"测试标准.jsonl";gold.write_text("\n".join(json.dumps(r,ensure_ascii=False) for r in (base,peer)),encoding="utf-8-sig")
        engine=SemanticEngine({},str(gold),str(work))
        table={"range":"A1:B3","note_ids":["C-N001"],"scope":"standalone","table_semantic":"资金余额","header_cells":["A1","A3","B2"]}
        return engine,table,peer

    def test_classify_offers_same_meaning_approved_definition_in_catalogue_and_rounds(self):
        from unittest.mock import patch
        engine,table,peer=self.fixture();original=copy.deepcopy(table);calls=[]
        def reply(system,payload):
            calls.append(copy.deepcopy(payload))
            if payload["task"]=="tables":
                self.assertIn(peer["table"]["id"],{t["id"] for t in payload["catalogue"]})
                return {"table_ids":[peer["table"]["id"]]}
            self.assertIn(peer["id"],{s["id"] for s in payload["gold_slots"]})
            return {"mappings":[{**mapping(),"slot_id":peer["id"],"scope":"standalone"}],"excluded":[],"unresolved":[]}
        with patch.object(engine,"_http",side_effect=reply):
            result=engine._classify(snapshot()["sheets"][0],table,["B3"],[],{})
        self.assertEqual(table,original)
        self.assertEqual(len([p for p in calls if p["task"]=="classify"]),2)
        self.assertEqual(result["mappings"][0]["scope"],"standalone")
        self.assertEqual(result["mappings"][0]["definition_scope"],"parent")

    def test_candidates_require_full_note_identity_and_actual_scope(self):
        from unittest.mock import patch
        cases=[("允许单户",None,"standalone",{"C-N001","P-N009"}),
               ("未批准单户",lambda b,p:p.pop("applicability"),"standalone",{"C-N001"}),
               ("不同名称",lambda b,p:p["note"].update(name="其他资金"),"standalone",{"C-N001"}),
               ("不同含义",lambda b,p:p["note"].update(meaning="受限资金集合"),"standalone",{"C-N001"}),
               ("空含义不推定",lambda b,p:(b["note"].update(meaning=""),p["note"].update(meaning="")),"standalone",{"C-N001"}),
               ("停用定义",lambda b,p:p.update(status="deprecated"),"standalone",{"C-N001"}),
               ("合并不串母公司",None,"consolidated",{"C-N001"}),
               ("母公司不串合并",None,"parent",{"P-N009"})]
        for label,change,scope,expected in cases:
            with self.subTest(label=label):
                engine,table,_=self.fixture(change);table["scope"]=scope;seen=[]
                def request(system,payload,validate):return validate({"table_ids":[]})
                def classify(sheet,table,cells,slots,*rest):
                    seen.append({s["note"]["id"] for s in slots})
                    return {a:("unresolved",{"cell":a,"reason":"本地仅检查候选"}) for a in cells}
                with patch.object(engine,"_request",side_effect=request),patch.object(engine,"_round",side_effect=classify):
                    engine._classify(snapshot()["sheets"][0],table,["B3"],[],{})
                self.assertEqual(seen,[expected,expected])


if __name__ == "__main__":
    unittest.main()
