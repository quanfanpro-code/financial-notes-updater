# coding: utf-8
"""真实入口接收漏识别格的槽位，保留原件，并在界面与核对表中展示。"""
import copy
import json
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from openpyxl import load_workbook
from 测试 import 测试B范围定位 as 范围定位
from 测试.测试统一语义 import fact
from 附注更新.统一识别 import UnifiedSemanticEngine
from 附注更新.统一外部复核 import preview_external_review, apply_external_review, export_review_template
from 附注更新.统一勾稽 import load_fact_mapping
from 附注更新.统一语义 import definition_hash
from 附注更新.表格 import file_hash
from 附注更新.界面 import Application, format_external_review_preview


class 槽位复核接续测试(unittest.TestCase):
    def test_漏识别格可正式补入槽位且导出不丢失不计算公式(self):
        case=范围定位.B范围定位测试();case.setUp()
        book=load_workbook(case.case.source);sheet=book['综合']
        sheet['A4']='固定资产名称';sheet['C4']='期末原值';sheet['B3']='=SUM(1,2)'
        source=case.case.folder/'待补充原文.xlsx';book.save(source);book.close()
        with patch.object(UnifiedSemanticEngine,'_http',lambda e,s,p:case.response(e,s,p)):
            target=case.case.target();engine=case.case.engine()
            parent=engine.recognize_workbook(source,target_mapping=target,artifact_metadata={
                'carrier':'B','original_path':str(source),'original_hash':file_hash(source)})
        raw=fact();dimensions=copy.deepcopy(raw['dimensions'])
        dimensions['asset_selection']={'mode':'unknown','domain':'固定资产','completeness':'unknown'}
        item={'sheet':'综合','cell':'C5','kind':'partial_metric_fact','metric_id':'固定资产原值',
              'dimensions':dimensions,'dimension_evidence':{k:['A1','A4','C4'] for k in dimensions},
              'missing_dimensions':[],'evidence_cells':['C4'],'reason':'列名确认期末原值，名称尚未填写'}
        opinion={'schema':'附注外部语义复核-v2','carrier':'B','mapping_sha256':file_hash(parent['mapping_path']),
                 'source_hash':file_hash(source),'gold_hash':engine.standard['sha256'],
                 'reviewer':{'name':'本地行为检查','method':'advanced_ai','reviewed_at':'2026-10-11T19:00:00+08:00'},'items':[item]}
        path=case.case.folder/'槽位意见.json';path.write_text(json.dumps(opinion,ensure_ascii=False),encoding='utf-8-sig')
        before=file_hash(source)
        preview=preview_external_review(parent['mapping_path'],path)
        self.assertIn('固定资产原值',format_external_review_preview(preview))
        with patch('附注更新.统一勾稽.calculate_mapping_formulas',side_effect=AssertionError('语义导入不能触发计算')) as calculate:
            result=apply_external_review(preview,case.case.folder,confirmed=True,confirm_method='采纳合成语义')
            calculate.assert_not_called()
        saved=load_fact_mapping(result['files']['b_mapping'],engine.standard)
        self.assertEqual(saved['facts'],parent['facts'])
        self.assertEqual(len(saved['slot_bindings']),1)
        self.assertEqual(saved['slot_bindings'][0]['cell'],'C5')
        self.assertEqual(saved['slot_bindings'][0]['partial_semantics']['dimensions'],dimensions)
        self.assertEqual(file_hash(source),before)
        book=load_workbook(result['files']['b_review']);rows=list(book.active.values);book.close()
        self.assertEqual(sum(r[0]=='综合!C5' for r in rows),1)
        self.assertEqual(next(r for r in rows if r[0]=='综合!C5')[1],'槽位已确认')
        shown={};label=SimpleNamespace(winfo_exists=lambda:True,configure=lambda **kw:shown.update(kw))
        app=SimpleNamespace(mapping_summary_labels={'b_mapping':label},values={'b_mapping':SimpleNamespace(get=lambda:result['files']['b_mapping'])})
        Application.refresh_mapping_summary(app,'b_mapping')
        self.assertIn('准确对应金标准 2 格',shown['text'])
        exported=case.case.folder/'确认后意见格式.json'
        export_review_template(result['files']['b_mapping'],exported,include_confirmed=True)
        options=json.loads(exported.read_text(encoding='utf-8-sig'))['confirmed_mapping_options']
        slot=next(r for r in options if r['cell']=='C5')
        self.assertEqual(slot['replace_slot_hash'],definition_hash(saved['slot_bindings'][0]))
        self.assertEqual(slot['kind'],'partial_metric_fact')
        correction=copy.deepcopy(opinion);correction['mapping_sha256']=file_hash(result['files']['b_mapping'])
        correction_path=case.case.folder/'缺少原槽位绑定.json'
        correction_path.write_text(json.dumps(correction,ensure_ascii=False),encoding='utf-8-sig')
        with self.assertRaisesRegex(ValueError,'更正已确认槽位须绑定原记录'):
            preview_external_review(result['files']['b_mapping'],correction_path)
        from 附注更新 import 分步流程 as flow
        destination=case.case.folder/'正式分步识别导出';destination.mkdir()
        config={'gold_path':str(case.case.gold),'b_path':str(source),'a_mapping':target['mapping_path'],'output_dir':str(destination)}
        with patch.object(flow,'validate_settings',return_value={}), \
             patch.object(UnifiedSemanticEngine,'recognize_workbook',return_value=copy.deepcopy(saved)), \
             patch('附注更新.统一勾稽.calculate_mapping_formulas') as calculate:
            exported=flow._recognize_v2(config,'B',destination,lambda s:None,None)
            calculate.assert_not_called()
        book=load_workbook(exported['files']['b_review']);rows=list(book.worksheets[0].values);book.close()
        self.assertEqual(sum(r[0]=='综合!C5' for r in rows),1)
        self.assertNotIn('计算结果',rows[0])
        legacy=copy.deepcopy(saved);legacy['recognition_scope']['identity_coverage_version']=1
        legacy_path=case.case.folder/'旧对应界面记录.json';legacy_path.write_text(json.dumps(legacy,ensure_ascii=False),encoding='utf-8-sig')
        app.values['b_mapping']=SimpleNamespace(get=lambda:str(legacy_path))
        Application.refresh_mapping_summary(app,'b_mapping')
        self.assertIn('槽位对应清单待更新',shown['text'])


if __name__=='__main__':unittest.main()
