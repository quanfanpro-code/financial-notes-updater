# coding: utf-8
"""仅验证实际分包入口与原文范围，不把受控分类当作模型准确率。"""
import copy
import json
from pathlib import Path
import unittest
import uuid

from openpyxl import Workbook
from 附注更新.统一识别 import UnifiedSemanticEngine, build_payload
from 附注更新.统一证据 import resolve_context
from 附注更新.语义 import SemanticError
from 测试.测试统一语义 import records


class SparsePacketTests(unittest.TestCase):
    def run_source(self, rows_per_table, overlap=False):
        folder=Path(__file__).resolve().parents[1]/'测试结果'/('稀疏分包_'+uuid.uuid4().hex)
        folder.mkdir()
        gold=folder/'标准.jsonl'
        with gold.open('x',encoding='utf-8') as handle:
            handle.write('\n'.join(json.dumps(row,ensure_ascii=False) for row in records()))
        book=Workbook();sheet=book.active;sheet.title='附注'
        context=[];start=1
        for index,count in enumerate(rows_per_table):
            context.append({'table_id':'表'+str(index),'range_name':'TA'+str(index),'sheet':'附注',
                            'first_row':start,'first_column':1,'row_count':count,'column_count':2,
                            'title':'第'+str(index)+'表的专属说明'})
            for row in range(start,start+count):
                sheet.cell(row,1,'表'+str(index)+'项目'+str(row));sheet.cell(row,2,'待核实的原文')
            start+=count-1 if overlap else count+3
        source=folder/'原表.xlsx';book.save(source);book.close()
        before=copy.deepcopy(context)
        engine=UnifiedSemanticEngine({},gold,folder/'结果');packets=[]
        def classify(packet):
            packets.append(copy.deepcopy(packet))
            return {'facts':[],'owners':{},'excluded':[],
                    'unresolved':[{'cell':cell,'reason':'受控分包检查，不作业务判断'} for cell in packet['candidate_cells']],
                    'request_receipts':[]}
        engine.classify_packet=classify
        result=engine.recognize_workbook(source,context=context)
        self.assertEqual(context,before)
        return packets,result,engine.standard

    def test_同页小表合批但说明不跨格适用(self):
        packets,result,standard=self.run_source([2,2])
        self.assertEqual(len(packets),1)
        self.assertEqual(len(packets[0]['candidate_cells']),8)
        self.assertEqual(len(result['unresolved']),8)
        self.assertIsNone(packets[0]['physical_table'])
        payload=build_payload(packets[0],standard);catalog=payload['context_evidence']
        self.assertEqual(len(catalog),2)
        reference=next(key for key,row in catalog.items() if row['scope']['bounds'][0]==1)
        resolve_context(reference,catalog,'附注!A1')
        with self.assertRaises(ValueError):resolve_context(reference,catalog,'附注!A6')

    def test_上限与顺序及单表边界保持(self):
        packets,result,_=self.run_source([15,15])
        self.assertEqual([len(p['candidate_cells']) for p in packets],[24,24,12])
        addresses=[cell for p in packets for cell in p['candidate_cells']]
        self.assertEqual(len(addresses),len(set(addresses)))
        self.assertEqual(len(addresses),result['coverage']['expected_count'])
        self.assertIsNotNone(packets[0]['physical_table'])
        self.assertIsNone(packets[1]['physical_table'])
        self.assertIsNotNone(packets[2]['physical_table'])

    def test_只发送候选格适用说明并保留全局说明(self):
        packets,_,standard=self.run_source([2,2])
        packet=copy.deepcopy(packets[0])
        packet['candidate_cells']=['B1','B2']
        packet['document_context'].append('用户说明：采用单户人民币口径')
        before=copy.deepcopy(packet)
        payload=build_payload(packet,standard)
        catalog=payload['context_evidence']
        self.assertEqual(len(catalog),2)
        self.assertTrue(any(item['scope'] is None for item in catalog.values()))
        self.assertTrue(all(item['scope'] is None or item['scope']['bounds'][0]==1 for item in catalog.values()))
        self.assertEqual(packet,before)
        self.assertEqual(set(payload['cells']),set(packet['cells']))

    def test_原表范围重叠仍拒绝(self):
        with self.assertRaises(SemanticError):self.run_source([2,2],overlap=True)


if __name__=='__main__':unittest.main()
