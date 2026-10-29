# coding: utf-8
"""只检验文字网格的结构、原文和实际发送，不调用模型。"""
import json
from pathlib import Path
import unittest
from unittest.mock import patch

from 附注更新.表格文字 import grid_text, workbook_table_text
from 附注更新.Word原文材料 import word_source_material, word_text_material
from 测试.测试Word原文材料 import Word原文材料测试
from 测试 import 测试直接识别 as fixture
from 附注更新.分步流程 import run_step
from 附注更新.统一识别 import UnifiedSemanticEngine


class 表格文字结构测试(unittest.TestCase):
    def test_横纵合并与原空白和缺记录分别保留(self):
        values = {(1, 1): '期末余额', (2, 1): '账面余额', (2, 3): '账面价值',
                  (3, 1): '金额', (3, 2): '比例', (4, 1): 100, (4, 3): None}
        text = grid_text(range(1, 6), {1: 'A', 2: 'B', 3: 'C'}, values,
                         [(1, 1, 1, 3), (2, 1, 2, 2), (2, 3, 3, 3)],
                         lambda r, c: chr(64 + c) + str(r))
        rows = text.splitlines()
        self.assertIn('并入A1 | 并入A1', rows[2])
        self.assertIn('并入A2', rows[3])
        self.assertIn('并入C2', rows[4])
        self.assertIn('C4 null〔原格空白〕', rows[5])
        self.assertIn('未记录此位置的内容', rows[6])
        self.assertNotIn('〔原格空白〕', rows[6])

    def test_原文中的竖线换行与反斜杠可逐字取回(self):
        value = '原文 | 不拆列\n第二行\\路径'
        text = grid_text([1], {1: 'A'}, {(1, 1): value}, [], lambda r, c: 'A1')
        line = text.splitlines()[2]
        self.assertEqual(line.count(' | '), 1)
        encoded = line.split('A1 ', 1)[1].rsplit(' |', 1)[0].replace('\\|', '|')
        self.assertEqual(json.loads(encoded), value)

    def test_相同文字不同原格不去重(self):
        text = grid_text([1], {1: 'A', 2: 'B'}, {(1, 1): '金额', (1, 2): '金额'}, [],
                         lambda r, c: chr(64 + c) + str(r))
        self.assertIn('A1 "金额"', text)
        self.assertIn('B1 "金额"', text)

    def test_已有表区原样采用不判断标题或身份证(self):
        cells = {'A1': {'row': 1, 'column': 1, 'value': '期末余额'},
                 'A2': {'row': 2, 'column': 1, 'value': '甲'},
                 'B2': {'row': 2, 'column': 2, 'value': None}}
        source = {'names': {'原表': [{'sheet': '表', 'range': 'A1:B2'}]},
                  'sheets': [{'name': '表', 'cells': cells, 'merges': ['A1:B1']}]}
        text = workbook_table_text(source, [{'sheet': '表', 'range': 'B2'}])
        self.assertIn('原表区：原表；范围：A1:B2', text)
        self.assertIn('并入A1', text)
        self.assertNotIn('metric_id', text)
        self.assertNotIn('dimension_id', text)

    def test_Word段落顺序跨度和两套位置分别保留(self):
        case = Word原文材料测试()
        case.setUp()
        original = word_source_material(case.link, {'合成表!C12': '合成表!C20'})
        text = word_text_material(original)
        self.assertLess(text.index('表前说明'), text.index('Word表号：1'))
        self.assertLess(text.index('Word表号：1'), text.index('表后说明'))
        self.assertIn('并入第1行第1列', text)
        self.assertIn('并入第1行第2列', text)
        self.assertIn('行前省略1列', text)
        self.assertIn('原提取Excel=合成表!C12；当前Excel=合成表!C20', text)
        self.assertIn('原文 \\| 不拆开', text)
        self.assertIn('原格空白', text)

    def test_网格以真实换行发送且目标置于末尾(self):
        case = fixture.直接识别测试()
        work = Path(__file__).resolve().parents[1] / '输出' / '主线接续_20261027_213431' / '合成离线验证'
        with patch.object(fixture, 'WORK', work):
            case.setUp()
        engine = case.engine()
        grid = '| 行号 | A | B |\n| 1 | 合并起点 | 并入A1 |'
        target = {'载体': 'A', '位置': ['合成表!B3']}
        def receive(settings, folder, body, cancel):
            content = body['messages'][1]['content']
            self.assertIn('\n' + grid + '\n\n', content)
            self.assertNotIn('本次原表网格.txt', content.split('\n\n【', 1)[0])
            self.assertEqual(json.loads(content.rsplit('【本次作答位置.json】\n', 1)[1]), target)
            data = {'choices': [{'finish_reason': 'stop', 'message': {'content': '本格原答复'}}], 'usage': {}}
            return data, json.dumps(data).encode('utf-8')
        with patch('附注更新.Codex语义调用.call_codex', side_effect=receive):
            answer = engine.request_text('按原文识别', {'本次原表网格.txt': grid, '本次作答位置.json': target})
        self.assertEqual(answer, '本格原答复')

    def test_AB普通入口共用文字网格且默认不附图(self):
        case = fixture.直接识别测试()
        work = Path(__file__).resolve().parents[1] / '输出' / '主线接续_20261027_213431' / '合成离线验证'
        with patch.object(fixture, 'WORK', work):
            case.setUp()
        raw = '{"location":"合成表!B3","metric_id":"固定资产原值"}'
        a = None
        for side in ('a', 'b'):
            config = {**case.config, side + '_target_table':
                      {'id': '原表', 'sheet': '合成表', 'range': 'A2:C8', 'data_ranges': ['B3']}}
            if side == 'b':
                config['a_mapping'] = a
            def read(engine, prompt, materials):
                self.assertIn('| 行号 | A | B | C |', materials['本次原表网格.txt'])
                self.assertEqual(materials['本次作答位置.json']['位置'], ['合成表!B3'])
                self.assertNotIn('原表图片.json', materials)
                return raw
            with patch.object(UnifiedSemanticEngine, 'request_text', read):
                result = run_step('recognize_' + side, config)
            self.assertEqual(result['status'], 'complete', result)
            if side == 'a':
                a = result['files']['a_mapping']


    def test_B参考A完整网格及各自文件归属保留且基础字段不混A(self):
        from openpyxl import Workbook
        from openpyxl.workbook.defined_name import DefinedName
        from 附注更新.表格 import file_hash
        case = fixture.直接识别测试()
        work = Path(__file__).resolve().parents[1] / '输出' / '主线接续_20261027_213431' / '合成离线验证'
        with patch.object(fixture, 'WORK', work):
            case.setUp()
        a_source = case.folder / 'A合并表头.xlsx'
        b_source = case.folder / 'B不同原表.xlsx'
        for path, side in [(a_source, 'A'), (b_source, 'B')]:
            book = Workbook()
            sheet = book.active
            sheet.title = '合成表'
            sheet['A1'] = side + '本方原始共同说明'
            sheet['A2'] = '项目'
            sheet['B2'] = '期末余额'
            sheet['B3'] = '原值'
            sheet['C3'] = '减值准备'
            sheet['A4'] = '设备甲'
            if side == 'A':
                sheet.merge_cells('A2:A3')
                sheet.merge_cells('B2:C2')
            else:
                sheet['B4'] = 77
            book.defined_names.add(DefinedName(side + '完整表区', attr_text="'合成表'!$A$2:$C$4"))
            with path.open('xb') as stream:
                book.save(stream)
            book.close()
        target = {'id': '本格', 'sheet': '合成表', 'range': 'A2:C4', 'data_ranges': ['B4']}
        cfg = {**case.config, 'a_path': str(a_source), 'b_path': str(b_source),
               'a_target_table': target, 'b_target_table': target,
               'a_context_text': 'A完整表外说明及其例外必须保留',
               'a_required_ranges': [{'sheet': '合成表', 'range': 'B4'}]}
        a_raw = '{"location":"合成表!B4","metric_id":"固定资产原值","currency":"CNY"}'
        b_raw = '{"location":"合成表!B4","metric_id":"固定资产原值","currency":"USD"}'
        with patch.object(UnifiedSemanticEngine, 'request_text', return_value=a_raw):
            a_result = run_step('recognize_a', cfg)
        self.assertEqual(a_result['status'], 'complete', a_result)
        cfg['a_mapping'] = a_result['files']['a_mapping']
        with patch.object(UnifiedSemanticEngine, 'request_text', return_value=b_raw) as call:
            b_result = run_step('recognize_b', cfg)
        self.assertEqual(b_result['status'], 'complete', b_result)
        call.assert_called_once()
        materials = call.call_args.args[1]
        self.assertIn('A原表网格.txt', materials)
        grid = materials['A原表网格.txt']
        self.assertIn('A完整表区', grid)
        self.assertIn('$A$2:$C$4', grid)
        self.assertIn('并入B2', grid)
        self.assertIn('B4 null〔原格空白〕', grid)
        self.assertNotIn('B4 77', grid)
        self.assertIn('B4 77', materials['本次原表网格.txt'])
        self.assertIn('A完整表外说明及其例外必须保留', json.dumps(materials['A原表.json'], ensure_ascii=False))
        self.assertEqual(materials['A原表.json']['源文件'], {'path': str(a_source), 'sha256': file_hash(a_source)})
        own = materials.get('B原文.json', materials.get('原文.json'))
        self.assertEqual(own['载体'], 'B')
        self.assertEqual(own['源文件'], {'path': str(b_source), 'sha256': file_hash(b_source)})
        with patch.object(UnifiedSemanticEngine, 'request_text', return_value='{"location":"合成表!B4","currency":"USD"}') as call:
            basic = run_step('recognize_b', {**cfg, 'b_mapping': b_result['files']['b_mapping'],
                                            'b_field_scope': {'合成表!B4': ['currency']}})
        self.assertEqual(basic['status'], 'complete', basic)
        call.assert_called_once()
        for name in ('A原表.json', 'A需要.json', 'A原表网格.txt'):
            self.assertNotIn(name, call.call_args.args[1])


if __name__ == '__main__':
    unittest.main()
