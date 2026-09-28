"""分步业务流程契约测试；使用受控模型回执，真实读写Excel/Word，不代表模型准确率。"""
import copy
import hashlib
import importlib
import json
import shutil
import threading
import unittest
import uuid
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook, load_workbook
from 附注更新.语义 import SemanticEngine as RealSemanticEngine

ROOT = Path(__file__).resolve().parents[1]


def digest(path):
    with Path(path).open("rb") as handle:
        return hashlib.file_digest(handle, "sha256").hexdigest()


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


class StepEngine:
    """明确限于测试的回执，每次调用可观察；不访问任何外部服务。"""
    calls = []
    # 模拟模型回执，但来源范围身份仍使用正式实现，不能绕过原范围校验。
    _selection_identity = staticmethod(RealSemanticEngine._selection_identity)

    def __init__(self, settings, gold_path, work_dir, log=None, cancel=None):
        self.gold_hash = digest(gold_path)
        self.usage = {}

    def recognize(self, snapshot, reference=None, progress=None):
        type(self).calls.append({"path": snapshot["path"], "reference": copy.deepcopy(reference)})
        mappings, excluded = [], []
        for sheet in snapshot["sheets"]:
            for address, cell in sheet["cells"].items():
                if isinstance(cell.get("value"), (int, float)):
                    mappings.append({"sheet": sheet["name"], "cell": address,
                        "slot_id": "C-N087-T001-S001", "scope": "consolidated",
                        "dimensions": {"currency": "CNY", "unit": "元", "scale": 1,
                            "metric": "库存现金期末余额"},
                        "semantic_field": "库存现金期末余额", "value_type": "monetary",
                        "row_label": "库存现金", "column_label": "期末余额",
                        "reason": "测试回执", "reviewed": True})
                else:
                    excluded.append({"sheet": sheet["name"], "cell": address, "reason": "测试标题"})
        return {"source_hash": snapshot["sha256"], "gold_hash": self.gold_hash,
            "mappings": mappings, "excluded": excluded, "unresolved": [], "complete": True, "usage": {}}


class StepFlowTests(unittest.TestCase):
    def setUp(self):
        self.work = ROOT / "测试结果" / ("分步_" + uuid.uuid4().hex[:10])
        self.work.mkdir(parents=True)
        self.output = self.work / "分步成果"
        self.a = self.work / "表格A.xlsx"
        self.b = self.work / "表格B.xlsx"
        for path, rows in ((self.a, [["项目", "期末余额"], ["库存现金", 7]]),
                           (self.b, [["项目", "库存现金"], ["期末余额", 125]])):
            book = Workbook()
            for row in rows:
                book.active.append(row)
            book.save(path)
            book.close()
        self.gold = self.work / "测试金标准.jsonl"
        self.gold.write_text(json.dumps({"id":"C-N087-T001-S001","status":"active","scope":"consolidated",
            "note":{"id":"C-N087","name":"货币资金","meaning":"货币资金披露"},
            "table":{"id":"C-N087-T001","name":"余额","meaning":"货币资金余额"},
            "slot":{"name":"库存现金期末余额","role":"measure"},"value_type":"monetary",
            "dimensions":[{"name":"currency","required":True},{"name":"unit","required":True},{"name":"scale","required":True}]},ensure_ascii=False)+"\n", encoding="utf-8-sig")
        self.config = {"output_dir": str(self.output), "a_path": str(self.a), "b_path": str(self.b),
            "gold_path": str(self.gold), "base_url": "http://127.0.0.1:1/v1", "model": "test",
            "api_key": "", "review": True, "context_text": "测试合并附注，人民币元"}
        StepEngine.calls = []

    def api(self):
        try:
            return importlib.import_module("附注更新.分步流程")
        except ImportError:
            self.fail("尚未实现文件化分步流程")

    def test_snapshot_preserves_independent_cell_layout_evidence(self):
        api = self.api()
        source_hash = digest(self.b)
        proof = {"source_hash": source_hash, "table_range": "B2:B2",
            "note_ids": ["C-N087"], "scope": "consolidated", "table_semantic": "库存现金",
            "header_cells": ["B1", "A2"], "layout_batch_id": "本地测试原包",
            "layout_review": {"method": "仅用于本地契约测试", "rounds": [{"round": 1}, {"round": 2}]}}
        record = {"source_path": str(self.b), "source_hash": source_hash, "context": [],
            "confirmed_business_ranges": {"Sheet": [proof]}}
        before = copy.deepcopy(record)
        snapshot = api._snapshot(record)
        actual = snapshot["sheets"][0]["confirmed_business_ranges"]
        self.assertEqual(len(actual), 1)
        self.assertEqual(actual[0]["layout_batch_id"], proof["layout_batch_id"])
        self.assertEqual(actual[0]["layout_review"], proof["layout_review"])
        actual[0]["layout_review"]["rounds"][0]["round"] = 99
        self.assertEqual(record, before)
        self.assertEqual(digest(self.b), source_hash)

    def test_snapshot_rejects_cell_proof_expanded_to_rectangle(self):
        api = self.api()
        source_hash = digest(self.b)
        record = {"source_path": str(self.b), "source_hash": source_hash, "context": [],
            "confirmed_business_ranges": {"Sheet": [{"source_hash": source_hash,
                "table_range": "A1:B2", "note_ids": ["C-N087"], "scope": "consolidated",
                "table_semantic": "库存现金", "header_cells": ["B1"],
                "layout_batch_id": "本地测试原包", "layout_review": {"rounds": []}}]}}
        with self.assertRaises(ValueError):
            api._snapshot(record)

    def test_snapshot_rebuilds_source_grid_blanks_only_for_declared_policy(self):
        api = self.api()
        path = self.work / "无格式业务空白.xlsx"
        book = Workbook()
        book.active.append(["项目", "期末余额"])
        book.active.append(["库存现金"])
        book.save(path)
        book.close()
        record = {"source_path": str(path), "source_hash": digest(path), "context": [],
            "recognition_basis": {"layout_mode": "cells_v1", "candidate_blanks": "source_grid_v1"}}
        snapshot = api._snapshot(record)
        blank = snapshot["sheets"][0]["cells"]["B2"]
        self.assertIsNone(blank["value"])
        self.assertTrue(blank["derived_blank"])
        self.assertEqual(blank["blank_evidence"]["kind"], "source_grid")
        self.assertNotIn("confirmed_business_ranges", snapshot["sheets"][0])
        legacy = api._snapshot({k: v for k, v in record.items() if k != "recognition_basis"})
        self.assertNotIn("B2", legacy["sheets"][0]["cells"])
        self.assertEqual(digest(path), record["source_hash"])

    def test_only_update_data_engine_uses_cell_layout_with_separate_cache(self):
        api = self.api()
        original = copy.deepcopy(self.config)
        identities = [{"carrier": "A"}, {"carrier": "B"}, {"action": "structure"}]
        for identity in identities:
            before = copy.deepcopy(identity)
            with patch.object(api, "SemanticEngine") as factory:
                api._engine(self.config, self.gold, identity, None, None)
            settings, gold_path, cache_path, _, _ = factory.call_args.args
            expected = dict(identity)
            if identity.get("carrier") == "B":
                self.assertEqual(settings.get("layout_mode"), "cells_v1")
                expected["layout_mode"] = "cells_v1"
            else:
                self.assertNotIn("layout_mode", settings)
            digest = hashlib.sha256(api._json(expected).encode("utf-8")).hexdigest()[:24]
            self.assertEqual(Path(cache_path).name, digest)
            self.assertEqual(gold_path, str(self.gold))
            self.assertEqual(identity, before)
        self.assertEqual(self.config, original)

    def run_step(self, api, step, **config):
        merged = dict(self.config, **config)
        return api.run_step(step, merged, lambda text: None, threading.Event())

    def complete(self, result):
        self.assertEqual(result["status"], "complete", result)
        self.assertTrue(Path(result["output_dir"]).is_dir(), result)
        self.assertIsInstance(result["files"], dict)
        for path in result["files"].values():
            self.assertTrue(Path(path).is_file(), result)
        return result["files"]

    def recognized_pair(self, api):
        with patch.object(api, "SemanticEngine", StepEngine):
            a = self.complete(self.run_step(api, "recognize_a"))["a_mapping"]
            b = self.complete(self.run_step(api, "recognize_b", a_mapping=a))["b_mapping"]
        return a, b

    def mapping_copy(self, path, **changes):
        record = read_json(path)
        record.update(changes)
        target = self.work / ("映射_" + uuid.uuid4().hex[:8] + ".json")
        target.write_text(json.dumps(record, ensure_ascii=False), encoding="utf-8-sig")
        return str(target)


    def test_reviewed_existing_mapping_accepts_same_meaning_note_without_changing_layout(self):
        api=self.api();am,_=self.recognized_pair(api)
        record,snapshot=api._load_mapping(am,"A",complete=False)
        from 附注更新.表格 import materialize_business_blanks
        base=read_json(self.gold);peer=copy.deepcopy(base)
        peer.update(id="C-N088-T001-S001");peer["note"]["id"]="C-N088";peer["table"]["id"]="C-N088-T001"
        gold=self.work/"同名定义测试标准.jsonl"
        gold.write_text("\n".join(json.dumps(x,ensure_ascii=False) for x in (base,peer)),encoding="utf-8-sig")
        engine=RealSemanticEngine({},str(gold),str(self.work));sheet=snapshot["sheets"][0]
        table={"range":"A1:B2","note_ids":["C-N087"],"scope":"consolidated","table_semantic":"库存现金余额","header_cells":["A1","B1","A2"]}
        materialize_business_blanks(snapshot,sheet["name"],table)
        record["confirmed_business_ranges"]={sheet["name"]:copy.deepcopy(sheet["confirmed_business_ranges"])}
        raw=copy.deepcopy(record["mappings"][0]);raw.update(slot_id=peer["id"],evidence_cells=["A2","B1"],reason="相同科目含义下的现有余额指标")
        checked=engine._validate_classification({"mappings":[raw],"excluded":[],"unresolved":[]},["B2"],{peer["id"]},sheet,table)["B2"][1]
        checked.update(reviewed=True,review_method="existing_slot_proposal_and_independent_review",
                       semantic_review={"proposal":{"reason":"指标含义一致","evidence_cells":["A2","B1"]},
                                        "independent_review":{"accepted":True,"reason":"原始行列证据一致","evidence_cells":["A2","B1"]}})
        record.update(mappings=[],unresolved=[{"sheet":sheet["name"],"cell":"B2","reason":"待核实"}],complete=False,gold_path=str(gold),gold_hash=digest(gold))
        original=copy.deepcopy(record)
        learned={"source_hash":snapshot["sha256"],"gold_hash":engine.gold_hash,"reviewed_existing_mappings":[checked],"new_ids":[]}
        report=self.work/"同名科目独立审核.json";report.write_text(json.dumps(learned,ensure_ascii=False),encoding="utf-8-sig")
        learned.update(report_path=str(report),report_sha256=digest(report))
        with patch("urllib.request.build_opener",side_effect=AssertionError("禁止网络")):
            updated=api._apply_reviewed_existing(record,learned,engine,snapshot)
        self.assertEqual(record,original)
        self.assertEqual(updated["mappings"][0]["slot_id"],peer["id"])
        self.assertEqual(updated["confirmed_business_ranges"],original["confirmed_business_ranges"])

    def test_reviewed_existing_mapping_completes_only_original_pending_cells(self):
        api=self.api();am,_=self.recognized_pair(api)
        record,snapshot=api._load_mapping(am,"A",complete=False)
        from 附注更新.语义 import SemanticEngine
        from 附注更新.表格 import materialize_business_blanks
        engine=SemanticEngine({"report_scope":"consolidated","review":True},str(self.gold),str(self.work/"复核结果"))
        sheet=snapshot["sheets"][0]
        table={"range":"A1:B2","note_ids":["C-N087"],"scope":"consolidated","table_semantic":"库存现金余额","header_cells":["A1","B1","A2"]}
        materialize_business_blanks(snapshot,sheet["name"],table)
        record["confirmed_business_ranges"]={sheet["name"]:copy.deepcopy(sheet["confirmed_business_ranges"])}
        raw=copy.deepcopy(record["mappings"][0])
        raw.update(evidence_cells=["A2","B1"],reason="原行列明确为库存现金余额")
        response={"mappings":[raw],"excluded":[],"unresolved":[]}
        mapping=engine._validate_classification(response,["B2"],{raw["slot_id"]},sheet,table)["B2"][1]
        mapping.update(reviewed=True,review_method="existing_slot_proposal_and_independent_review",
                       semantic_review={"proposal":{"reason":"现有指标完整覆盖","evidence_cells":["A2","B1"]},
                                        "independent_review":{"accepted":True,"reason":"独立复核原行列、值类型及必要维度均相符","evidence_cells":["A2","B1"]}})
        record.update(mappings=[],unresolved=[{"sheet":sheet["name"],"cell":"B2","reason":"旧识别有分歧"}],complete=False)
        before=copy.deepcopy(record);source_before=digest(self.a)
        learned={"source_hash":snapshot["sha256"],"gold_hash":engine.gold_hash,
                 "reviewed_existing_mappings":[mapping],"new_ids":[]}
        report=self.work/"已有语义独立复核.json"
        report.write_text(json.dumps(learned,ensure_ascii=False),encoding="utf-8-sig")
        learned.update(report_path=str(report),report_sha256=digest(report))
        updated=api._apply_reviewed_existing(record,learned,engine,snapshot)
        self.assertEqual(record,before,"原识别结果保留，不覆盖旧证据")
        self.assertTrue(updated["complete"])
        self.assertEqual(len(updated["mappings"]),1)
        self.assertFalse(updated["unresolved"])
        self.assertEqual(updated["mappings"][0]["review_method"],"existing_slot_proposal_and_independent_review")
        self.assertEqual(digest(self.a),source_before)
        saved=read_json(updated["evidence_path"])
        self.assertEqual(saved["mappings"],updated["mappings"])
        self.assertEqual(updated["evidence_hash"],digest(updated["evidence_path"]))
        self.assertTrue(api._apply_reviewed_existing(record,{"reviewed_existing_mappings":[]},engine,snapshot) is record)
        for bad_kind in ("hash","dimension","scope","unreviewed","not_pending","duplicate"):
            with self.subTest(bad_kind=bad_kind):
                bad=copy.deepcopy(learned);bad_record=copy.deepcopy(record)
                bad["reviewed_existing_mappings"]=copy.deepcopy(learned["reviewed_existing_mappings"])
                item=bad["reviewed_existing_mappings"][0]
                if bad_kind=="hash":bad["source_hash"]="wrong"
                elif bad_kind=="dimension":item["dimensions"].pop("currency")
                elif bad_kind=="scope":item["scope"]="parent"
                elif bad_kind=="unreviewed":item["semantic_review"]["independent_review"]["accepted"]=False
                elif bad_kind=="not_pending":bad_record["unresolved"]=[]
                elif bad_kind=="duplicate":bad["reviewed_existing_mappings"].append(copy.deepcopy(item))
                report_path=self.work/("错误复核_"+bad_kind+".json")
                body={k:v for k,v in bad.items() if k not in {"report_path","report_sha256"}}
                report_path.write_text(json.dumps(body,ensure_ascii=False),encoding="utf-8-sig")
                bad.update(report_path=str(report_path),report_sha256=digest(report_path))
                original=copy.deepcopy(bad_record)
                with self.assertRaises(ValueError):api._apply_reviewed_existing(bad_record,bad,engine,snapshot)
                self.assertEqual(bad_record,original)

    def test_old_and_new_metric_records_match_without_rewriting_old_files(self):
        api=self.api();am,bm=self.recognized_pair(api)
        legacy=[]
        for path in (am,bm):
            record=read_json(path)
            self.assertNotIn("metric",record["mappings"][0]["dimensions"])
            record["mappings"][0]["dimensions"]["metric"]="库存现金期末余额"
            legacy.append(self.mapping_copy(path,mappings=record["mappings"]))
        originals={path:digest(path) for path in (am,bm,*legacy,str(self.a),str(self.b))}
        for path,carrier in zip(legacy,("A","B")):
            loaded,_=api._load_mapping(path,carrier)
            self.assertNotIn("metric",loaded["mappings"][0]["dimensions"])
        with patch.object(api,"SemanticEngine",side_effect=AssertionError("兼容读取与勾稽无需模型")):
            for a_path,b_path in ((am,bm),(legacy[0],legacy[1]),(legacy[0],bm),(am,legacy[1])):
                with self.subTest(a=a_path,b=b_path):
                    files=self.complete(self.run_step(api,"match",a_mapping=a_path,b_mapping=b_path))
                    plan=read_json(files["update_plan"])
                    self.assertEqual(len(plan["updates"]),1)
            written=self.complete(self.run_step(api,"write",update_plan=files["update_plan"]))
        book=load_workbook(written["a_prime"]);self.assertEqual(book.active["B2"].value,125);book.close()
        self.assertNotIn("metric",read_json(written["a_prime_mapping"])["mappings"][0]["dimensions"])
        self.assertEqual({path:digest(path) for path in originals},originals)
        wrong=read_json(am)["mappings"];wrong[0]["dimensions"]["metric"]="坏账准备"
        invalid=self.mapping_copy(am,mappings=wrong);before=digest(invalid)
        with self.assertRaisesRegex(ValueError,"metric|指标"):api._load_mapping(invalid,"A")
        self.assertEqual(digest(invalid),before)

    def test_progress_survives_cancel_before_gold_review_and_can_resume(self):
        api=self.api();book=load_workbook(self.a);book.active.append(["尚待确认",31]);book.save(self.a);book.close()
        events=[];calls=[]
        def reply(engine,system,payload):
            calls.append(copy.deepcopy(payload))
            if payload["task"]=="layout":return {"tables":[{"range":"A1:B3","note_ids":["C-N087"],"scope":"consolidated","table_semantic":"货币资金余额","header_cells":["A1","B1"]}]}
            if payload["task"]=="tables":return {"table_ids":["C-N087-T001"]}
            owned=payload["candidate_cells"]
            item={"cell":"B2","slot_id":"C-N087-T001-S001","scope":"consolidated","dimensions":{"currency":"CNY","unit":"元","scale":1},"value_type":"monetary","semantic_field":"库存现金期末余额","reason":"A2为库存现金，B1为期末余额"}
            return {"mappings":[item] if "B2" in owned else [],"excluded":[{"cell":a,"category":"label","reason":"原表文字标签"} for a in owned if a not in {"B2","B3"}],"unresolved":[{"cell":"B3","reason":"业务含义尚待确认"}] if "B3" in owned else []}
        def interrupted(*args,**kwargs):
            self.assertTrue(events,"进入金标准补充之前必须已发布可用映射")
            raise api.SemanticCancelled("测试正常取消")
        with patch.object(api.SemanticEngine,"_http",new=reply),patch("附注更新.语义补充.review_missing_semantics",side_effect=interrupted),patch.object(api,"_mapping_review",side_effect=AssertionError("每包不得写Excel")):
            result=api.run_step("recognize_a",self.config,progress=events.append)
        self.assertEqual(result["status"],"cancelled",result)
        path=result["files"]["a_mapping"];record,actual=api._load_mapping(path,"A",complete=False)
        self.assertEqual(len(record["mappings"]),1);self.assertFalse(record["complete"])
        self.assertEqual(len(record["unresolved"]),1)
        self.assertEqual(len(record["mappings"])+len(record["excluded"])+len(record["unresolved"]),6)
        self.assertEqual(events[-1]["files"]["a_mapping"],path)
        self.assertEqual(events[-1]["status"],"running");self.assertTrue(events[-1]["checkpoint"])
        self.assertEqual(events[-1]["counts"],{"mappings":1,"excluded":4,"unresolved":1})
        self.assertFalse(list(Path(result["output_dir"]).rglob("*.xlsx")))
        prior=digest(path);calls.clear()
        with patch.object(api.SemanticEngine,"_http",new=reply),patch("附注更新.语义补充.review_missing_semantics",return_value={"new_ids":[]}):
            resumed=api.run_step("recognize_a",{**self.config,"a_mapping":path})
        self.assertEqual(resumed["status"],"partial",resumed)
        classified={a for request in calls if request["task"]=="classify" for a in request["candidate_cells"]}
        self.assertNotIn("B2",classified);self.assertEqual(digest(path),prior)

    def test_progress_cancel_before_and_after_new_gold_recognition_keeps_actual_version(self):
        api=self.api();book=load_workbook(self.a);book.active.append(["新指标",31]);book.save(self.a);book.close()
        original_hash=digest(self.gold);new_gold=self.work/"新版本识别用标准.jsonl"
        old=json.loads(self.gold.read_text(encoding="utf-8-sig"));new=copy.deepcopy(old)
        new["id"]="C-N087-T001-S002";new["slot"]["name"]="新指标"
        new_gold.write_text("\n".join(json.dumps(item,ensure_ascii=False) for item in [old,new])+"\n",encoding="utf-8-sig")
        new_hash=digest(new_gold)
        for boundary in ("新轮开始前","新轮已有结果后","新轮已有结果后失败"):
            with self.subTest(boundary=boundary):
                events=[]
                class LearningProgressEngine(StepEngine):
                    def recognize(engine,snapshot,reference=None,previous_result=None,progress=None):
                        result=super().recognize(snapshot,reference)
                        if engine.gold_hash==original_hash:
                            result["mappings"]=[m for m in result["mappings"] if m["cell"]!="B3"]
                            result.update(complete=False,unresolved=[{"sheet":"Sheet","cell":"B3","reason":"新指标待独立核实"}])
                            return result
                        if boundary=="新轮开始前":raise api.SemanticCancelled("新版本尚未识别")
                        for m in result["mappings"]:
                            if m["cell"]=="B3":
                                m.update(slot_id="C-N087-T001-S002",semantic_field="新指标")
                                m["dimensions"].pop("metric",None)
                        progress(snapshot,result)
                        if boundary=="新轮已有结果后失败":raise ValueError("新版本已保存两格后普通失败")
                        raise api.SemanticCancelled("新版本已保存首批后取消")
                with patch.object(api,"SemanticEngine",LearningProgressEngine),patch("附注更新.语义补充.review_missing_semantics",return_value={"new_ids":[new["id"]],"path":str(new_gold)}):
                    result=api.run_step("recognize_a",self.config,progress=events.append)
                self.assertEqual(result["status"],"failed" if boundary.endswith("失败") else "cancelled",result)
                record,_=api._load_mapping(result["files"]["a_mapping"],"A",complete=False)
                expected=original_hash if boundary=="新轮开始前" else new_hash
                self.assertEqual(record["gold_hash"],expected)
                self.assertEqual(len(record["mappings"]),1 if boundary=="新轮开始前" else 2)
                self.assertEqual(digest(record["gold_path"]),expected)
                self.assertEqual(digest(record["selected_gold_path"]),expected)
                self.assertEqual(record["source_hash"],digest(self.a))
                self.assertEqual(result["files"]["gold_path"],str(new_gold))
                self.assertEqual(events[-1]["files"]["gold_path"],str(new_gold))
                self.assertNotIn("gold_path",events[0]["files"],"尚未发布新标准时不得改模型识别依据")
                self.assertEqual(record["latest_published_gold_path"],str(new_gold))
                self.assertEqual(record["latest_published_gold_hash"],new_hash)
                self.assertFalse(record["complete"])
                self.assertNotIn("evidence_hash",record)
                self.assertNotIn("evidence_path",record)
                self.assertEqual(events[-1]["file_hashes"][result["files"]["a_mapping"]],digest(result["files"]["a_mapping"]))
                self.assertEqual(len(list((Path(result["output_dir"])/"识别进度"/"识别依据").glob("原始实际结构.json"))),1)
        self.assertEqual(digest(self.gold),original_hash)

    def test_published_gold_survives_apply_failure_and_following_progress_save_failure(self):
        api=self.api();old_hash=digest(self.gold);new_gold=self.work/"复核刚发布的新标准.jsonl"
        old=json.loads(self.gold.read_text(encoding="utf-8-sig"));new=copy.deepcopy(old)
        new["id"]="C-N087-T001-S002";new["slot"]["name"]="新增指标"
        new_gold.write_text("\n".join(json.dumps(i,ensure_ascii=False) for i in [old,new])+"\n",encoding="utf-8-sig")
        class PendingEngine(StepEngine):
            def recognize(engine,snapshot,reference=None,progress=None):
                result=super().recognize(snapshot,reference)
                result.update(mappings=[],complete=False,unresolved=[{"sheet":"Sheet","cell":"B2","reason":"待核实指标"}])
                return result
        for boundary in ("应用复核失败","随后保存失败"):
            with self.subTest(boundary=boundary):
                events=[];original_save=api.save_json
                def save(path,value):
                    if boundary=="随后保存失败" and Path(path).suffix==".tmp" and any(e["files"].get("gold_path") for e in events):
                        raise OSError("测试发布后保存失败")
                    return original_save(path,value)
                def apply(result,*args):
                    if boundary=="应用复核失败":raise ValueError("测试已发布后应用失败")
                    return result
                with patch.object(api,"SemanticEngine",PendingEngine),patch("附注更新.语义补充.review_missing_semantics",return_value={"new_ids":[new["id"]],"path":str(new_gold)}),patch.object(api,"_apply_reviewed_existing",side_effect=apply),patch.object(api,"save_json",side_effect=save):
                    result=api.run_step("recognize_a",self.config,progress=events.append)
                self.assertEqual(result["status"],"failed",result)
                self.assertEqual(result["files"]["gold_path"],str(new_gold))
                self.assertEqual(events[-1]["files"]["gold_path"],str(new_gold))
                self.assertEqual(result["files"]["a_mapping"],events[0]["files"]["a_mapping"])
                record,_=api._load_mapping(result["files"]["a_mapping"],"A",complete=False)
                self.assertEqual(record["gold_hash"],old_hash)
                self.assertEqual(digest(result["files"]["a_mapping"]),events[0]["file_hashes"][result["files"]["a_mapping"]])
                self.assertNotIn("gold_path",events[0]["files"])
        self.assertEqual(digest(self.gold),old_hash)

    def test_progress_callback_failure_returns_last_saved_mapping(self):
        api=self.api();events=[]
        def fail(event):
            events.append(event)
            raise RuntimeError("进度通知故障")
        with patch.object(api,"SemanticEngine",StepEngine):
            result=api.run_step("recognize_b",self.config,progress=fail)
        self.assertEqual(result["status"],"failed",result)
        self.assertEqual(result["files"]["b_mapping"],events[-1]["files"]["b_mapping"])
        record,_=api._load_mapping(result["files"]["b_mapping"],"B",complete=False)
        self.assertFalse(record["complete"])

    def test_selected_formal_mapping_resumes_through_recognize_entry(self):
        api=self.api()
        def reply(engine,system,payload):
            if payload["task"]=="layout":return {"tables":[{"range":"A1:B2","note_ids":["C-N087"],"scope":"consolidated","table_semantic":"货币资金余额","header_cells":["A1","B1"]}]}
            if payload["task"]=="tables":return {"table_ids":["C-N087-T001"]}
            owned=payload["candidate_cells"]
            item={"cell":"B2","slot_id":"C-N087-T001-S001","scope":"consolidated","dimensions":{"currency":"CNY","unit":"元","scale":1},"value_type":"monetary","semantic_field":"库存现金期末余额","reason":"A2为库存现金，B1为期末余额"}
            return {"mappings":[item] if "B2" in owned else [],"excluded":[{"cell":a,"category":"label","reason":"原表文字标签"} for a in owned if a!="B2"],"unresolved":[]}
        with patch.object(api.SemanticEngine,"_http",new=reply):
            first=self.complete(self.run_step(api,"recognize_a"))["a_mapping"]
        previous_hash=digest(first)
        with patch.object(api.SemanticEngine,"_http",side_effect=AssertionError("正式入口应复用已选择成果")):
            result=self.complete(self.run_step(api,"recognize_a",a_mapping=first))
        saved=read_json(result["a_mapping"])
        self.assertEqual(saved["reuse"]["counts"]["mappings"],1)
        self.assertEqual(saved["reuse"]["evidence_hash"],previous_hash)
        self.assertEqual(digest(first),previous_hash)

    def test_empty_partial_mapping_retries_but_keeps_all_resume_checks(self):
        from 附注更新.语义 import SemanticError
        api=self.api()
        class EmptyEngine(StepEngine):
            def recognize(engine,snapshot,reference=None,progress=None):
                result=super().recognize(snapshot,reference)
                result["unresolved"]=[{"sheet":sheet["name"],"cell":a,"reason":"旧协议失败，尚未确认"} for sheet in snapshot["sheets"] for a in sheet["cells"]]
                result.update(mappings=[],excluded=[],complete=False)
                return result
        with patch.object(api,"SemanticEngine",EmptyEngine),patch("附注更新.语义补充.review_missing_semantics",return_value={"new_ids":[]}):
            first=self.run_step(api,"recognize_a")
        self.assertEqual(first["status"],"partial")
        previous=first["files"]["a_mapping"];original=digest(previous);old=read_json(previous)
        with self.assertRaisesRegex(ValueError,"未确认|没有已确认"):
            api._load_mapping(previous,"A",complete=True)
        claimed_complete=self.mapping_copy(previous,complete=True,unresolved=[])
        with self.assertRaisesRegex(ValueError,"没有已确认"):
            api._load_mapping(claimed_complete,"A",complete=True)
        calls=[]
        def reply(engine,system,payload):
            calls.append(copy.deepcopy(payload))
            if payload["task"]=="layout":return {"tables":[{"range":"A1:B2","note_ids":["C-N087"],"scope":"consolidated","table_semantic":"货币资金余额","header_cells":["A1","B1"]}]}
            if payload["task"]=="tables":return {"table_ids":["C-N087-T001"]}
            owned=payload["candidate_cells"]
            item={"cell":"B2","slot_id":"C-N087-T001-S001","scope":"consolidated","dimensions":{"currency":"CNY","unit":"元","scale":1},"value_type":"monetary","semantic_field":"库存现金期末余额","reason":"A2为库存现金，B1为期末余额"}
            return {"mappings":[item] if "B2" in owned else [],"excluded":[{"cell":a,"category":"label","reason":"原表文字标签"} for a in owned if a!="B2"],"unresolved":[]}
        with patch.object(api.SemanticEngine,"_http",new=reply):
            finished=self.complete(self.run_step(api,"recognize_a",a_mapping=previous))
        result=read_json(finished["a_mapping"])
        self.assertEqual(result["reuse"]["counts"]["mappings"],0)
        self.assertEqual(len(result["mappings"]),1)
        self.assertEqual({a for p in calls if p["task"]=="classify" for a in p["candidate_cells"]},{"A1","A2","B1","B2"})
        self.assertEqual(digest(previous),original)
        for fields in ({"gold_hash":"不匹配"},{"source_hash":"不匹配"},{"original_hash":"不匹配"},
                       {"context":["另一企业、另一期间"]},{"unresolved":old["unresolved"][:-1]}):
            with self.subTest(fields=tuple(fields)):
                altered=self.mapping_copy(previous,**fields)
                with patch.object(api.SemanticEngine,"_http",side_effect=AssertionError("校验失败不能请求模型")) as request:
                    rejected=self.run_step(api,"recognize_a",a_mapping=altered)
                self.assertEqual(rejected["status"],"failed",rejected)
                request.assert_not_called()

    def test_new_semantics_publish_and_resume_recognition(self):
        api=self.api(); original=digest(self.gold); source=digest(self.a)
        new_gold=self.work/"新增后完整金标准.jsonl"
        rows=[json.loads(self.gold.read_text(encoding="utf-8-sig"))]
        addition=copy.deepcopy(rows[0]);addition["id"]="C-N087-T001-S002";addition["slot"]["name"]="新增业务指标"
        new_gold.write_text("\n".join(json.dumps(x,ensure_ascii=False) for x in rows+[addition])+"\n",encoding="utf-8-sig")
        class LearningEngine(StepEngine):
            def recognize(engine,snapshot,reference=None,previous_result=None,progress=None):
                result=super().recognize(snapshot,reference)
                if engine.gold_hash==original:
                    result["unresolved"]=[{"sheet":"Sheet","cell":"B2","reason":"待核实新业务语义"}]
                    result["complete"]=False;result["mappings"]=[]
                else:
                    if previous_result is None:raise AssertionError("新金标准识别须明确带入前轮成果")
                    if previous_result["gold_hash"]!=original:raise AssertionError("前轮成果须保留旧金标准身份")
                    for m in result["mappings"]:
                        m["slot_id"]="C-N087-T001-S002"
                        m["semantic_field"]="新增业务指标"
                        m["dimensions"].pop("metric",None)
                return result
        learned={"path":str(new_gold),"new_ids":["C-N087-T001-S002"],"changes":[],"unresolved":[]}
        with patch.object(api,"SemanticEngine",LearningEngine),patch("附注更新.语义补充.review_missing_semantics",return_value=learned) as review:
            files=self.complete(self.run_step(api,"recognize_a"))
        self.assertEqual(review.call_count,1)
        self.assertEqual(len(StepEngine.calls),2)
        self.assertEqual(files["gold_path"],str(new_gold))
        result=read_json(files["a_mapping"])
        self.assertEqual(result["new_gold_ids"],["C-N087-T001-S002"])
        self.assertEqual(result["mappings"][0]["slot_id"],"C-N087-T001-S002")
        self.assertTrue(Path(files["a_gold_changes"]).is_file())
        self.assertEqual(digest(self.gold),original);self.assertEqual(digest(self.a),source)

    def test_failed_gold_review_keeps_recognized_files_and_pending_items(self):
        api=self.api();original=digest(self.a)
        class PendingEngine(StepEngine):
            def recognize(engine,snapshot,reference=None,previous_result=None,progress=None):
                result=super().recognize(snapshot,reference)
                pending=result["excluded"].pop(0)
                result["unresolved"]=[{**pending,"reason":"尚待核实的真实语义"}]
                result["complete"]=False
                return result
        with patch.object(api,"SemanticEngine",PendingEngine),patch("附注更新.语义补充.review_missing_semantics",side_effect=ValueError("模拟补充故障")):
            result=self.run_step(api,"recognize_a")
        self.assertEqual(result["status"],"partial",result)
        record=read_json(result["files"]["a_mapping"])
        self.assertEqual(len(record["mappings"]),1)
        self.assertEqual(len(record["unresolved"]),1)
        self.assertTrue(Path(result["files"]["a_review"]).is_file())
        changes=read_json(result["files"]["a_gold_changes"])
        self.assertEqual(changes["reviews"][0]["status"],"failed")
        self.assertEqual(digest(self.a),original)

    def test_failure_after_publication_keeps_old_mapping_and_new_gold(self):
        api=self.api();original=digest(self.gold);source=digest(self.a)
        rows=[json.loads(self.gold.read_text(encoding="utf-8-sig"))]
        addition=copy.deepcopy(rows[0]);addition["id"]="C-N087-T001-S002";addition["slot"]["name"]="新的真实业务指标"
        new_gold=self.work/"已发布的新金标准.jsonl"
        new_gold.write_text("\n".join(json.dumps(x,ensure_ascii=False) for x in rows+[addition])+"\n",encoding="utf-8-sig")
        class RetryFailureEngine(StepEngine):
            def recognize(engine,snapshot,reference=None,previous_result=None,progress=None):
                if engine.gold_hash!=original:
                    snapshot["sheets"][0]["cells"].clear()
                    raise ValueError("模拟发布后重新识别失败")
                result=super().recognize(snapshot,reference)
                pending=result["excluded"].pop(0)
                result["unresolved"]=[{**pending,"reason":"待核实的新业务"}];result["complete"]=False
                return result
        learned={"path":str(new_gold),"new_ids":["C-N087-T001-S002"],"changes":[],"unresolved":[]}
        with patch.object(api,"SemanticEngine",RetryFailureEngine),patch("附注更新.语义补充.review_missing_semantics",return_value=learned):
            result=self.run_step(api,"recognize_a")
        self.assertEqual(result["status"],"failed",result)
        self.assertEqual(result["files"]["gold_path"],str(new_gold))
        saved=read_json(result["files"]["a_mapping"])
        self.assertEqual(saved["gold_hash"],original)
        self.assertEqual(len(saved["mappings"]),1)
        self.assertTrue(saved["unresolved"])
        self.assertEqual(digest(self.a),source)
        self.assertEqual(saved["latest_published_gold_path"],str(new_gold))
        self.assertIn("模拟发布后重新识别失败",result["message"])

    def test_word_files_flow_through_five_separate_steps(self):
        api = self.api()
        word = self.work / "文档一.docx"
        doc = Document()
        doc.add_paragraph("六、财务报表重要项目的说明")
        doc.add_paragraph("合并附注，人民币元")
        table = doc.add_table(rows=2, cols=2)
        table.style = "Table Grid"
        for row, values in zip(table.rows, [["项目", "期末余额"], ["库存现金", "7.00"]]):
            for cell, value in zip(row.cells, values):
                cell.text = value
        doc.save(word)
        original_word = digest(word)
        with patch.object(api, "SemanticEngine", side_effect=AssertionError("提取不得调用模型")):
            extracted = self.complete(self.run_step(api, "extract_a", word_path=str(word)))
        self.assertEqual(set(extracted), {"a_path", "a_link"})
        self.assertEqual(StepEngine.calls, [])
        original_a = digest(extracted["a_path"])
        with patch.object(api, "SemanticEngine", StepEngine):
            a_map = self.complete(self.run_step(api, "recognize_a", **extracted))["a_mapping"]
            self.assertEqual(len(StepEngine.calls), 1)
            self.assertNotIn("b_mapping", read_json(a_map))
            b_map = self.complete(self.run_step(api, "recognize_b", a_mapping=a_map))["b_mapping"]
            self.assertEqual(len(StepEngine.calls), 2)
            self.assertTrue(StepEngine.calls[-1]["reference"])
            before_match = set(self.output.rglob("*.xlsx"))
            with patch.object(StepEngine, "recognize", side_effect=AssertionError("对应时不得重新识别")):
                matched = self.complete(self.run_step(api, "match", a_mapping=a_map, b_mapping=b_map,a_path=""))
        self.assertIn("update_plan", matched)
        self.assertIn("trace", matched)
        self.assertNotIn("a_prime", matched)
        self.assertNotIn("word", matched)
        self.assertEqual(set(self.output.rglob("*.xlsx")) - before_match, {Path(matched["trace"])})
        with patch.object(api, "SemanticEngine", side_effect=AssertionError("写入不得调用模型")):
            written = self.complete(self.run_step(api, "write", update_plan=matched["update_plan"],
                gold_path="", base_url="", model="", api_key="",a_path=""))
        self.assertIn("word", written)
        self.assertIn("word_link", written)
        self.assertEqual(Document(written["word"]).tables[0].cell(1, 1).text, "125.00")
        excel_only=self.complete(self.run_step(api,"write",update_plan=matched["update_plan"],a_link="",a_path=""))
        self.assertNotIn("word",excel_only)
        self.assertNotIn("a_link",read_json(excel_only["a_prime_mapping"]))
        self.assertEqual(digest(word), original_word)
        self.assertEqual(digest(extracted["a_path"]), original_a)
        updated_mapping = read_json(written["a_prime_mapping"])
        self.assertEqual(updated_mapping["source_hash"], digest(written["a_prime"]))
        self.assertEqual(Path(updated_mapping["source_path"]), Path(written["a_prime"]))
        self.assertEqual(updated_mapping["carrier"], "A")
        self.assertEqual(Path(updated_mapping["a_link"]), Path(written["word_link"]))

        # 下一轮直接复用A′映射与位置关联，只识别新的B。
        new_b=self.work/"第二轮B.xlsx";book=Workbook();book.active.append(["项目","库存现金"]);book.active.append(["期末余额",250]);book.save(new_b);book.close()
        with patch.object(api,"SemanticEngine",StepEngine):
            second_b=self.complete(self.run_step(api,"recognize_b",b_path=str(new_b),a_mapping=written["a_prime_mapping"]))["b_mapping"]
        self.assertEqual(len(StepEngine.calls),3)
        with patch.object(api,"SemanticEngine",side_effect=AssertionError("下一轮不应再次识别A")):
            second_plan=self.complete(self.run_step(api,"match",a_path=written["a_prime"],b_path=str(new_b),
                a_mapping=written["a_prime_mapping"],b_mapping=second_b))
            second_write=self.complete(self.run_step(api,"write",a_path=written["a_prime"],b_path=str(new_b),
                update_plan=second_plan["update_plan"]))
        self.assertEqual(Document(second_write["word"]).tables[0].cell(1,1).text,"250.00")
        self.assertEqual(Document(written["word"]).tables[0].cell(1,1).text,"125.00")


    def test_missing_chapter_never_falls_back_to_extracting_whole_word(self):
        api=self.api();word=self.work/"只有其他章节.docx"
        doc=Document();doc.add_paragraph("七、关联方关系")
        table=doc.add_table(rows=1,cols=2);table.cell(0,0).text="其他章节";table.cell(0,1).text="100"
        doc.save(word);before=digest(word)
        result=self.run_step(api,"extract_a",word_path=str(word))
        self.assertEqual(result["status"],"failed",result)
        self.assertNotIn("a_path",result["files"])
        self.assertEqual(digest(word),before)


    def test_B默认完整识别不读取已保留的关联和旧范围(self):
        api = self.api()
        for options in ({}, {"scope_prefilter":False}, {"scope_prefilter":True,"a_link":""}):
            with self.subTest(options=options), patch.object(api, "SemanticEngine", StepEngine), \
                 patch.object(api, "_load_link", side_effect=AssertionError("不应读取目标关联")) as link:
                files = self.complete(self.run_step(api, "recognize_b",
                    **{"a_link":"保留供后续回写.json","b_scope":"不存在的旧范围.json",**options}))
                link.assert_not_called()
                record = read_json(files["b_mapping"])
                self.assertFalse(record.get("scope_selection"))
                self.assertEqual({m["cell"] for m in record["mappings"]}, {"B2"})
                self.assertNotIn("b_scope", files)

    def test_B明确勾选且有Word关联才传递目标和已存范围(self):
        api = self.api();captured = []
        target = {"disclosures":["库存现金"],"table_business":["货币资金"]}
        old = self.work/"已有范围.json";old.write_text('{"范围":"旧范围"}',encoding="utf-8-sig")
        class ScopeEngine(StepEngine):
            def recognize(engine, snapshot, reference=None, **kwargs):
                captured.append(kwargs)
                return super().recognize(snapshot, reference)
        link = {"a_path":str(self.a),"a_hash":digest(self.a),"context":{"tables":[]}}
        with patch.object(api, "SemanticEngine", ScopeEngine), patch.object(api, "_load_link", return_value=link), \
             patch("附注更新.识别范围.build_target_meanings", return_value=target) as build:
            self.complete(self.run_step(api,"recognize_b",scope_prefilter=True,a_link="关联.json",b_scope=str(old)))
        build.assert_called_once()
        self.assertEqual(len(captured), 1)
        arguments = dict(captured[0])
        self.assertTrue(callable(arguments.pop("progress", None)))
        self.assertEqual(arguments, {"target_scope":target,"scope_selection":{"范围":"旧范围"}})

    def test_完整识别拒绝旧范围映射且未启动模型(self):
        api = self.api()
        previous = self.work/"旧范围映射.json"
        for scope in ({"scope_selection":{"target_hash":"旧目标"}}, {"scope_selection_path":"旧范围.json"}):
            previous.write_text(json.dumps({"schema":"附注语义映射-v1","carrier":"B",**scope}),encoding="utf-8-sig")
            with self.subTest(scope=scope), patch.object(api, "SemanticEngine", side_effect=AssertionError("不应启动模型")) as engine:
                result = self.run_step(api,"recognize_b",b_mapping=str(previous),a_link="后续回写仍保留.json")
            self.assertEqual(result["status"], "failed")
            self.assertIn("重新识别", result["message"])
            self.assertIn("完整更新数据", result["message"])
            engine.assert_not_called()

    def test_完整模式建立对应前拒绝旧范围映射(self):
        api = self.api();am,bm = self.recognized_pair(api)
        scoped = self.mapping_copy(bm,scope_selection={"target_hash":"旧目标"})
        a_record,a_snapshot = api._load_mapping(am,"A")
        b_record,b_snapshot = api._load_mapping(bm,"B")
        b_record["scope_selection"] = {"target_hash":"旧目标"}
        with patch.object(api,"_load_mapping",side_effect=[(a_record,a_snapshot),(b_record,b_snapshot)]), \
             patch.object(api,"build_update_plan",side_effect=AssertionError("不应建立对应")) as plan:
            result = self.run_step(api,"match",a_mapping=am,b_mapping=scoped)
        self.assertEqual(result["status"], "failed")
        self.assertIn("完整更新数据", result["message"])
        plan.assert_not_called()

    def test_recognize_b_is_independent_and_does_not_require_a(self):
        api = self.api()
        with patch.object(api, "SemanticEngine", StepEngine):
            files = self.complete(self.run_step(api, "recognize_b", a_path="", a_mapping=""))
        self.assertEqual(len(StepEngine.calls), 1)
        self.assertIsNone(StepEngine.calls[0]["reference"])
        record = read_json(files["b_mapping"])
        self.assertEqual(record["schema"], "附注语义映射-v1")
        self.assertEqual(record["carrier"], "B")
        self.assertEqual(record["source_hash"], digest(self.b))
        self.assertEqual(record["gold_hash"], digest(self.gold))
        self.assertNotEqual(Path(record["gold_path"]),self.gold)
        self.assertEqual(digest(record["gold_path"]),digest(self.gold))
        self.assertEqual(Path(record["source_path"]), self.b)
        self.assertNotIn("api_key", record)
        self.assertEqual(set(files), {"b_mapping","b_review"})
        book=load_workbook(files["b_review"])
        self.assertEqual(book["业务格语义"]["E2"].value,"C-N087-T001-S001")
        self.assertEqual(book["业务格语义"]["D2"].value,125)
        book.close()

    def test_saved_mappings_can_match_and_write_without_recognizing_again(self):
        api = self.api()
        a_map, b_map = self.recognized_pair(api)
        originals = {str(path): digest(path) for path in (self.a, self.b, Path(a_map), Path(b_map))}
        with patch.object(api, "SemanticEngine", side_effect=AssertionError("普通对应与写入无需模型")):
            matched = self.complete(self.run_step(api, "match", a_mapping=a_map, b_mapping=b_map))
            written = self.complete(self.run_step(api, "write", update_plan=matched["update_plan"]))
        self.assertEqual(load_workbook(written["a_prime"]).active["B2"].value, 125)
        self.assertEqual(read_json(written["a_prime_mapping"])["source_hash"], digest(written["a_prime"]))
        for path, expected in originals.items():
            self.assertEqual(digest(path), expected)


    def test_write_preserves_unprovided_original_blank_and_rechecks_record(self):
        api=self.api();source=self.work/"附注含原业务空白.xlsx"
        book=load_workbook(self.a);ws=book.active;ws["A3"]="预留对象";ws["B3"].number_format="0.00";book.save(source);book.close()
        self.config["a_path"]=str(source)
        class BlankEngine(StepEngine):
            def recognize(engine,snapshot,reference=None,previous_result=None,progress=None):
                result=super().recognize(snapshot,reference)
                if Path(snapshot["path"])==source:
                    mapped=copy.deepcopy(result["mappings"][0]);mapped["cell"]="B3"
                    mapped["dimensions"]["counterparty"]="预留对象"
                    result["mappings"].append(mapped)
                    result["excluded"]=[m for m in result["excluded"] if m["cell"]!="B3"]
                return result
        with patch.object(api,"SemanticEngine",BlankEngine):
            a_map=self.complete(self.run_step(api,"recognize_a"))["a_mapping"]
            b_map=self.complete(self.run_step(api,"recognize_b",a_mapping=a_map))["b_mapping"]
        matched=self.complete(self.run_step(api,"match",a_mapping=a_map,b_mapping=b_map))
        plan=read_json(matched["update_plan"])
        self.assertEqual(len(plan["updates"]),1);self.assertEqual(len(plan["retained_blanks"]),1)
        with patch.object(api,"SemanticEngine",side_effect=AssertionError("回写不调用模型")):
            result=self.run_step(api,"write",update_plan=matched["update_plan"])
        files=self.complete(result)
        self.assertEqual(result["updated_cells"],1);self.assertEqual(result["retained_blank_cells"],1)
        output=load_workbook(files["a_prime"])
        self.assertEqual(output.active["B2"].value,125);self.assertIsNone(output.active["B3"].value);output.close()
        review=load_workbook(files["trace"])
        self.assertEqual(review["原空白保留"].max_row,2);review.close()
        forged_dir=self.work/"错误保留记录";forged_dir.mkdir()
        forged=forged_dir/"更新清单.json";plan["retained_blanks"][0]["old_value"]=0
        api.save_json(forged,plan)
        api.save_json(forged_dir/"更新清单校验.json",{"path":str(forged),"sha256":digest(forged)})
        rejected=self.run_step(api,"write",update_plan=str(forged))
        self.assertNotEqual(rejected["status"],"complete",rejected)
        self.assertNotIn("a_prime",rejected["files"])

    def test_new_selected_b_cannot_silently_use_old_b_mapping(self):
        api=self.api();a_map,b_map=self.recognized_pair(api)
        newer=self.work/"另一个B.xlsx";book=Workbook();book.active.append(["项目","期末"]);book.active.append(["库存现金",999]);book.save(newer);book.close()
        result=self.run_step(api,"match",a_mapping=a_map,b_mapping=b_map,b_path=str(newer))
        self.assertNotEqual(result["status"],"complete",result)
        matched=self.complete(self.run_step(api,"match",a_mapping=a_map,b_mapping=b_map))
        result=self.run_step(api,"write",update_plan=matched["update_plan"],b_path=str(newer))
        self.assertNotEqual(result["status"],"complete",result)

    def test_saved_mapping_rechecks_active_gold_ids(self):
        api=self.api();a_map,b_map=self.recognized_pair(api)
        record=read_json(b_map);record["mappings"][0]["slot_id"]="不存在的ID"
        changed=self.mapping_copy(b_map,mappings=record["mappings"])
        a_record=read_json(a_map);a_record["mappings"][0]["slot_id"]="不存在的ID"
        changed_a=self.mapping_copy(a_map,mappings=a_record["mappings"])
        result=self.run_step(api,"match",a_mapping=changed_a,b_mapping=changed)
        self.assertNotEqual(result["status"],"complete",result)

    def test_match_requires_mapping_files(self):
        api = self.api()
        with patch.object(api, "SemanticEngine", side_effect=AssertionError("缺文件不得启动模型")):
            result = self.run_step(api, "match", a_mapping="", b_mapping=str(self.work / "不存在.json"))
        self.assertNotEqual(result["status"], "complete", result)
        self.assertNotIn("update_plan", result.get("files", {}))

    def test_changed_source_after_recognition_is_rejected(self):
        api = self.api()
        a_map, b_map = self.recognized_pair(api)
        backup_dir = Path(r"C:\Users\27651\BackUp") / ("附注自动更新系统_" + datetime.now().strftime("%Y%m%d_%H%M%S_%f"))
        target = backup_dir / self.a.relative_to(ROOT)
        target.parent.mkdir(parents=True)
        shutil.copy2(self.a, target)
        self.assertEqual(digest(self.a), digest(target))
        book = load_workbook(self.a)
        book.active["B2"] = 999
        book.save(self.a)
        book.close()
        result = self.run_step(api, "match", a_mapping=a_map, b_mapping=b_map)
        self.assertNotEqual(result["status"], "complete", result)
        self.assertNotIn("update_plan", result.get("files", {}))

    def test_different_gold_versions_are_rejected(self):
        api = self.api()
        a_map, b_map = self.recognized_pair(api)
        other_gold = self.work / "另一版本金标准.jsonl"
        other_gold.write_text('{"other":true}\n', encoding="utf-8-sig")
        b_map = self.mapping_copy(b_map, gold_path=str(other_gold), gold_hash=digest(other_gold))
        result = self.run_step(api, "match", a_mapping=a_map, b_mapping=b_map)
        self.assertNotEqual(result["status"], "complete", result)
        self.assertNotIn("update_plan", result.get("files", {}))

    def test_incomplete_mapping_cannot_be_used_for_write_plan(self):
        api = self.api()
        a_map, b_map = self.recognized_pair(api)
        b_map = self.mapping_copy(b_map, complete=False, unresolved=[{"cell": "B2", "reason": "口径未确认"}])
        result = self.run_step(api, "match", a_mapping=a_map, b_mapping=b_map)
        self.assertNotEqual(result["status"], "complete", result)
        self.assertNotIn("update_plan", result.get("files", {}))

    def test_write_rechecks_b_source_after_plan_is_saved(self):
        api = self.api()
        a_map, b_map = self.recognized_pair(api)
        matched = self.complete(self.run_step(api, "match", a_mapping=a_map, b_mapping=b_map))
        backup_dir = Path(r"C:\Users\27651\BackUp") / ("附注自动更新系统_" + datetime.now().strftime("%Y%m%d_%H%M%S_%f"))
        target = backup_dir / self.b.relative_to(ROOT)
        target.parent.mkdir(parents=True)
        shutil.copy2(self.b, target)
        self.assertEqual(digest(self.b), digest(target))
        book = load_workbook(self.b)
        book.active["B2"] = 888
        book.save(self.b)
        book.close()
        with patch.object(api, "SemanticEngine", side_effect=AssertionError("写入不能调用模型")):
            result = self.run_step(api, "write", update_plan=matched["update_plan"])
        self.assertNotEqual(result["status"], "complete", result)
        self.assertNotIn("a_prime", result.get("files", {}))
        self.assertEqual(load_workbook(self.a).active["B2"].value, 7)


    def test_supplement_is_separate_for_a_and_b(self):
        api=self.api()
        with patch.object(api,"SemanticEngine",StepEngine):
            a=self.complete(self.run_step(api,"recognize_a",a_context_text="A：2025年，人民币元",b_context_text="B：2026年，人民币万元"))
            b=self.complete(self.run_step(api,"recognize_b",a_context_text="A：2025年，人民币元",b_context_text="B：2026年，人民币万元"))
        self.assertEqual(read_json(a["a_mapping"])["additional_context"],"A：2025年，人民币元")
        self.assertEqual(read_json(b["b_mapping"])["additional_context"],"B：2026年，人民币万元")
        self.assertNotIn("A：2025年",str(read_json(b["b_mapping"])["context"]))

    def new_gold(self,change_used=False):
        slot=json.loads(self.gold.read_text(encoding="utf-8-sig").splitlines()[0])
        extra=copy.deepcopy(slot);extra["id"]="C-N087-T001-S002";extra["slot"]["name"]="银行存款期末余额"
        if change_used:slot["slot"]["name"]="含义已经修订"
        target=self.work/("新版金标准_"+uuid.uuid4().hex[:5]+".jsonl")
        target.write_text("\n".join(json.dumps(v,ensure_ascii=False) for v in (slot,extra))+"\n",encoding="utf-8-sig")
        return str(target)

    def test_added_gold_keeps_old_a_mapping_usable_for_new_b_and_next_round(self):
        api=self.api()
        with patch.object(api,"SemanticEngine",StepEngine):
            a_map=self.complete(self.run_step(api,"recognize_a"))["a_mapping"]
            newer=self.new_gold()
            b_map=self.complete(self.run_step(api,"recognize_b",a_mapping=a_map,gold_path=newer))["b_mapping"]
        with patch.object(api,"SemanticEngine",side_effect=AssertionError("复用旧A映射不得重新识别")):
            plan=self.complete(self.run_step(api,"match",a_mapping=a_map,b_mapping=b_map,gold_path=newer))
            output=self.complete(self.run_step(api,"write",update_plan=plan["update_plan"],gold_path=newer))
        self.assertEqual(len(StepEngine.calls),2)
        self.assertEqual(read_json(output["a_prime_mapping"])["gold_hash"],digest(newer))
        self.assertNotEqual(Path(read_json(output["a_prime_mapping"])["gold_path"]),Path(newer))
        self.assertEqual(digest(read_json(output["a_prime_mapping"])["gold_path"]),digest(newer))
        self.assertEqual(read_json(a_map)["gold_hash"],digest(self.gold))
        self.assertEqual(load_workbook(output["a_prime"]).active["B2"].value,125)

    def test_revised_used_definition_blocks_before_a_prime(self):
        api=self.api();a_map,b_map=self.recognized_pair(api)
        plan=self.complete(self.run_step(api,"match",a_mapping=a_map,b_mapping=b_map))
        changed=self.new_gold(change_used=True)
        result=self.run_step(api,"match",a_mapping=a_map,b_mapping=b_map,gold_path=changed)
        self.assertEqual(result["status"],"failed",result)
        result=self.run_step(api,"write",update_plan=plan["update_plan"],gold_path=changed)
        self.assertEqual(result["status"],"failed",result)
        self.assertFalse(list(Path(result["output_dir"]).glob("更新后附注表格*")))

    def test_unresolved_mapping_exports_real_source_package_without_inventing_standard(self):
        api=self.api()
        class UnresolvedEngine(StepEngine):
            def recognize(self,snapshot,reference=None,progress=None):
                result=super().recognize(snapshot,reference)
                source=result["mappings"].pop()
                result.update(complete=False,unresolved=[{"sheet":source["sheet"],"cell":source["cell"],"reason":"现有标准无法确认，需区分维度或缺项"}])
                return result
        with patch.object(api,"SemanticEngine",UnresolvedEngine):
            result=self.run_step(api,"recognize_b",b_context_text="这份B采用人民币元")
        self.assertEqual(result["status"],"partial",result)
        package=read_json(result["files"]["b_standard_review"])
        self.assertEqual(package["carrier"],"B")
        self.assertEqual(package["source_hash"],digest(self.b))
        self.assertEqual(package["gold_hash"],digest(self.gold))
        self.assertEqual(package["items"][0]["source_cell"]["value"],125)
        self.assertTrue(Path(package["structure_path"]).is_file())
        self.assertTrue(Path(package["instructions_path"]).is_file())
        self.assertEqual(package["proposed_changes"],[])
        self.assertEqual(digest(self.gold),package["gold_hash"])

    def test_b_can_be_recognized_while_a_still_has_unresolved_items(self):
        api=self.api()
        with patch.object(api,"SemanticEngine",StepEngine):
            a=self.complete(self.run_step(api,"recognize_a"))["a_mapping"]
            pending=self.mapping_copy(a,complete=False,unresolved=[{"sheet":"Sheet","cell":"B2","reason":"A仍需核实"}])
            b=self.complete(self.run_step(api,"recognize_b",a_mapping=pending))
        self.assertTrue(read_json(b["b_mapping"])["complete"])
        self.assertIsNone(StepEngine.calls[-1]["reference"])


    def test_word_omitted_position_stays_absent_across_saved_mappings_and_write(self):
        api=self.api()
        source=ROOT/"依赖"/"audit-notes-tool"/"核心测试"/"测试资料"/"Word编号"/"07_行首省略一列.docx"
        original_hash=digest(source)
        document=Document(source)
        heading=document.add_paragraph("六、财务报表重要项目的说明")
        document.tables[0]._tbl.addprevious(heading._p)
        target=document.tables[0]._tbl.tr_lst[1].tc_lst[0]
        texts=target.xpath(".//w:t")
        self.assertTrue(texts)
        texts[0].text="7.00"
        for text in texts[1:]:text.text=""
        word=self.work/"行首省略的财务附注.docx";document.save(word)
        extracted=self.complete(self.run_step(api,"extract_a",word_path=str(word)))
        context=read_json(extracted["a_link"])["context"]["tables"]
        self.assertEqual(context[0]["structural_omissions"][0]["range"],"B3")
        with patch.object(api,"SemanticEngine",StepEngine):
            a_map=self.complete(self.run_step(api,"recognize_a",**extracted))["a_mapping"]
            b_map=self.complete(self.run_step(api,"recognize_b",a_mapping=a_map))["b_mapping"]
        record=read_json(a_map)
        self.assertEqual([(m["sheet"],m["cell"]) for m in record["mappings"]],[("sheet1","C3")])
        self.assertNotIn("B3",api._snapshot(record)["sheets"][0]["cells"])
        with patch.object(api,"SemanticEngine",side_effect=AssertionError("普通对应和写入不需模型")):
            plan=self.complete(self.run_step(api,"match",a_mapping=a_map,b_mapping=b_map,a_path=""))
            written=self.complete(self.run_step(api,"write",update_plan=plan["update_plan"],a_path=""))
        updated=read_json(written["a_prime_mapping"])
        _,restored=api._load_mapping(written["a_prime_mapping"],"A")
        self.assertNotIn("B3",restored["sheets"][0]["cells"])
        self.assertEqual(restored["sheets"][0]["cells"]["C3"]["value"],125)
        self.assertEqual(updated["context"][0]["structural_omissions"][0]["range"],"B3")
        output=Document(written["word"]).tables[0]._tbl
        self.assertEqual(len(output.tr_lst[1].tc_lst),len(document.tables[0]._tbl.tr_lst[1].tc_lst))
        self.assertEqual("".join(t.text for t in output.tr_lst[1].tc_lst[0].xpath(".//w:t")),"125.00")
        self.assertEqual(digest(source),original_hash)

if __name__ == "__main__":
    unittest.main()
