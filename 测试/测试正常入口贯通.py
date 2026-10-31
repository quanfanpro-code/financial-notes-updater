# coding: utf-8
"""新入口的要求传递、范围接入和真实文件搬运；模拟答复不算业务识别验收。"""
import json
from pathlib import Path
import unittest
import uuid
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from 附注更新.正常入口识别 import read_scope_answer
from 附注更新.设置 import save_settings, load_settings
from 附注更新.分步流程 import run_step
from 附注更新.统一识别 import UnifiedSemanticEngine
from 附注更新.表格 import file_hash
from 测试.测试统一语义 import records


class 正常入口贯通测试(unittest.TestCase):
    def test_新增空名称行的两处来源使用已识别期间列(self):
        from 附注更新.A合集 import declared_named_row_relations,source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'names':{'表':[{'sheet':'原表','range':'B1:D5'}]},
           'sheets':[{'name':'原表','cells':{}}]},'outline':{'subtables':[{'sheet':'原表','metrics':[
           {'ranges':['C4'],'template_period_role':'closing'},{'ranges':['D4'],'template_period_role':'opening'}]}]}}
        b={'source_hash':'b','机械记录':{'sheets':[{'name':'来源甲','cells':{}},{'name':'来源乙','cells':{}}]}}
        text=('A表原始范围 原表!B1:D5，按要求于原小计4之前插入一条业务行，名称保留空白。\n'
              '新增行期末余额 → 来源甲!I15。\n新增行期初余额 → 来源乙!D27。\n'
              '原A位置 原表!B2:B3 的两行保留；名称依B原文更新为“项目甲”“项目乙”。')
        _,tables=declared_named_row_relations(a,b,text)
        self.assertEqual(len(tables),1)
        self.assertEqual(tables[0]['row_sources'],{'15':{'C':'来源甲!I15','D':'来源乙!D27'}})
        self.assertEqual(tables[0]['first_row_offset'],3)
        self.assertEqual(tables[0]['declared_labels'],{'15':{'value':None}})
        labels=source_relation(a,b,text,['原表!B2','原表!B3'])['label_actions']
        self.assertEqual([(x['target'],x['value']) for x in labels],[('原表!B2','项目甲'),('原表!B3','项目乙')])

    def test_来源位置列与依据列分开接入(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'sheets':[{'name':'原表'}]}}
        b={'source_hash':'b','机械记录':{'sheets':[{'name':'来源甲'},{'name':'来源乙'}]}}
        text=('| A位置及含义 | 最终采用的B位置 | B位置依据及实际期间 |\n|---|---|---|\n'
              '| 原表!C3、D3：两个期间 | 来源甲!I10、来源乙!D20 | 来源乙!B20属于另一期间，不采用 |')
        self.assertEqual(source_relation(a,b,text,['原表!C3','原表!D3'])['bindings'],
                         [{'target':'原表!C3','source':'来源甲!I10'},{'target':'原表!D3','source':'来源乙!D20'}])

    def test_原合计行号分别声明的新增成员复用整行列关系(self):
        from 附注更新.A合集 import declared_named_row_relations,source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'names':{
            '甲':[{'sheet':'原表','range':'B1:D5'}],'乙':[{'sheet':'原表','range':'B8:D12'}]},'sheets':[{'name':'原表'}]},
            'source_context':[{'range_name':'甲','title':'原标题'}]}
        b={'source_hash':'b','机械记录':{'sheets':[{'name':'来源'}]}}
        text=('A 原表!B5:D5 ← B 来源!A22:C22。\nA 原表!B12:D12 ← B 来源!A32:C32。\n'
              '于原A合计行 A5 之前另列新成员，取 B 来源!A20:C20；于原A合计行 A12 之前另列新成员，取 B 来源!A30:C30。')
        _,tables=declared_named_row_relations(a,b,text)
        self.assertEqual([(t['range_name'],t['source_rows'],t['columns']) for t in tables],
                         [('甲',[20],{'C':'B','D':'C'}),('乙',[30],{'C':'B','D':'C'})])
        dates='本表期末采用的实际日期为2030-12-31；期初采用的B实际日期为2029-01-01。'
        caption=source_relation(a,b,dates+'仅在本表局部期间说明处标明上述实际日期；不改其他表。',['原表!C3'])['caption_updates']
        self.assertEqual(caption,[{'range_name':'甲','old_text':'原标题','text':'原标题 '+dates}])

    def test_分列声明的年初来源与期末来源分别展开(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'sheets':[{'name':'原表'}]}}
        b={'source_hash':'b','机械记录':{'sheets':[{'name':'来源'}]},'b_source_ranges':[{'sheet':'来源','range':'A10:G11'}]}
        text=('A表的期末列 C/D/E 对应B的期末列 B/C/D；A表的期初列 F/G/H 对应B的年初列 E/F/G。\n'
              '- A 3“项目甲” ↔ B 11“项目甲”\n- A 4“项目乙” ↔ B 10“项目乙”')
        targets=['原表!'+col+str(row) for row in [3,4] for col in 'CDEFGH']
        result=source_relation(a,b,text,targets)
        self.assertEqual({x['target']:x['source'] for x in result['bindings']},
                         {f'原表!{left}{old}':f'来源!{right}{new}'
                          for old,new in [(3,11),(4,10)] for left,right in zip('CDEFGH','BCDEFG')})

    def test_分写期间来源使用已有列角色且不广播首个来源(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'sheets':[{'name':'原表'}]},
           'outline':{'subtables':[{'sheet':'原表','metrics':[
               {'ranges':['C7','C8'],'template_period_role':'closing'},
               {'ranges':['D7','D8'],'template_period_role':'opening'}]}]}}
        b={'source_hash':'b','机械记录':{'sheets':[{'name':'来源'}]}}
        reply=('原小计行 A 原表!C7:D7：期末 ← B 来源!B28，期初 ← B 来源!D28。\n'
               '原净额行 A 原表!C8:D8：期末 ← B 来源!F38，期初 ← B 来源!F47。')
        targets=['原表!'+col+str(row) for row in [7,8] for col in 'CD']
        result=source_relation(a,b,reply,targets)
        self.assertEqual({r['target']:r['source'] for r in result['bindings']},
                         dict(zip(targets,['来源!B28','来源!D28','来源!F38','来源!F47'])))
        a['outline']={}
        self.assertEqual(source_relation(a,b,reply,targets)['bindings'],[])

    def test_合计前另列来源成员使用已明确整行列对应(self):
        from 附注更新.A合集 import declared_named_row_relations
        a={'机械记录':{'names':{'表甲':[{'sheet':'原表','range':'B1:G5'}],
                                '表乙':[{'sheet':'原表','range':'B8:G12'}]},'sheets':[]}}
        text=('### 期末\nB新增成员排在原合计行之前，采用B 来源!A20:F20。\n'
              'A原 原表!C5:G5“合计” ↔ B 来源!B22:F22。\n'
              '### 期初\n输出副本在合计前另列B成员，来源为 来源!A30:F30。\n'
              'A原 原表!B12:G12“合计” ↔ B 来源!A32:F32。')
        _,tables=declared_named_row_relations(a,{'机械记录':{'sheets':[]}},text)
        self.assertEqual([t['source_rows'] for t in tables],[[20],[30]])
        self.assertTrue(all(t['columns']==dict(zip('CDEFG','BCDEF')) for t in tables))
        self.assertTrue(all(t['old_detail_count']==0 and t['labels']=={'B':'A'} and t['totals']==[] for t in tables))
        self.assertEqual([t['operations'][0]['index'] for t in tables],[4,4])

    def test_已有空名称行补跨表金额来源并保留一次插行及名称(self):
        from 附注更新.A合集 import declared_named_row_relations,append_declared_relation,source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'names':{'表':[{'sheet':'原表','range':'B1:D5'}]},
            'sheets':[{'name':'原表','cells':{}}]},'outline':{'subtables':[{'sheet':'原表','metrics':[
                {'ranges':['C4'],'template_period_role':'closing'},{'ranges':['D4'],'template_period_role':'opening'}]}]},
            'source_context':[{'range_name':'表','title':'账龄表'}]}
        old={'range_name':'表','source_sheet':'旧来源','source_rows':[15],'old_detail_count':0,'first_row_offset':3,
             'columns':{'C':'C'},'row_columns':{'15':{}},'labels':{'B':'B'},'totals':[],
             'operations':[{'axis':'row','action':'insert','index':3,'count':1,'template_index':2}]}
        b={'source_hash':'b','facts':[],'机械记录':{'sheets':[{'name':'来源'},{'name':'旧来源'}]},
           '已确认关系':[{'kind':'名单编排','tables':[old]}]}
        text=('原小计 A 原表!C4:D4。\n已有空名称业务行：期末金额 ← B 来源!B27，期初金额 ← B 来源!D27；不另新增一行。\n'
              'B说明标明期末为 2024-12-31、期初为 2024-01-01；A副本本表标题旁的局部期间文字相应更新。')
        _,tables=declared_named_row_relations(a,b,text)
        self.assertEqual(len(tables),1);self.assertEqual(tables[0]['source_rows'],[15])
        self.assertEqual(tables[0]['labels'],{'B':'B'});self.assertEqual(tables[0]['operations'],old['operations'])
        self.assertEqual(tables[0]['row_sources'],{'15':{'C':'来源!B27','D':'来源!D27'}})
        relations=b['已确认关系'];append_declared_relation(relations,{'kind':'名单编排','answer_hash':'new','tables':tables})
        self.assertEqual(len([t for r in relations for t in r.get('tables',[])]),1)
        caption=source_relation(a,b,text,['原表!C4'])['caption_updates']
        self.assertEqual(caption,[{'range_name':'表','old_text':'账龄表','text':'账龄表 期末为 2024-12-31、期初为 2024-01-01'}])

    def test_新成员期间追加时完整保留既有旧成员说明(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'names':{'表':[{'sheet':'原表','range':'B2:D5'}]},'sheets':[{'name':'原表','cells':{}}]},
           'source_context':[{'range_name':'表','title':'前五名'}]}
        b={'source_hash':'b','facts':[],'机械记录':{'sheets':[{'name':'来源','cells':{}}]},
           '已确认关系':[{'caption_updates':[{'range_name':'表','old_text':'前五名','text':'前五名（旧明细为2025年）'}]}]}
        raw=('新增行排在原独立合计之前：A 原表!B5:D5 ↔ B 来源!A10:C10。\n'
             '新增匿名行继承B的期末角色，实际日期为2024年12月31日；不改动A原有明细的期间。')
        result=source_relation(a,b,raw,['原表!B5'])['caption_updates']
        self.assertEqual(result,[{'range_name':'表','old_text':'前五名','text':'前五名（旧明细为2025年） '+raw.split('\n')[1]}])
        b['已确认关系'].append({'caption_updates':result})
        self.assertEqual(source_relation(a,b,raw,['原表!B5'])['caption_updates'],result)

    def test_新增成员占原合计位置但独立合计后移保留(self):
        from 附注更新.A合集 import declared_named_row_relations
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{
            'names':{'表':[{'sheet':'原表','range':'B2:D5'}]},'sheets':[{'name':'原表','cells':{}}]}}
        b={'source_hash':'b','facts':[],'机械记录':{'sheets':[{'name':'来源','cells':{}}]},'b_source_ranges':[{'sheet':'来源','range':'A10:C13'}]}
        raw=('新增行排在原明细之后、原独立合计之前：A 原表!B5:D5 ↔ B 来源!A10:C10；A B6:D6 ↔ B A11:C11。\n'
             '插入后，原独立合计移至第7行，沿用已确定的来源。')
        _,tables=declared_named_row_relations(a,b,raw)
        self.assertEqual(tables[0]['source_rows'],[10,11]);self.assertEqual(tables[0]['old_detail_count'],0)
        self.assertEqual(tables[0]['columns'],{'C':'B','D':'C'})
        self.assertEqual(tables[0]['target_row_offsets'],[3,4]);self.assertEqual(tables[0]['totals'],[])
        self.assertEqual(tables[0]['operations'],[{'axis':'row','action':'insert','index':3,'count':2,'template_index':2}])

    def test_局部来源期间说明保留来源限定并只作用指定表(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'names':{
            '表甲':[{'sheet':'原表','range':'$B$2:$D$5'}],'表乙':[{'sheet':'原表','range':'B8:D10'}]},'sheets':[{'name':'原表','cells':{}}]},
           'source_context':[{'range_name':'表甲','title':'名单甲'},{'range_name':'表乙','title':'名单乙'}]}
        b={'source_hash':'b','facts':[],'机械记录':{'sheets':[{'name':'来源','cells':{}}]}}
        raw='本次纳入的B相关附注表期初、期末分别为2024年1月1日、2024年12月31日，本期为2024年度；因此对这些目标表适用的A期间说明，应按角色改为“期初为2024年1月1日”。'
        captions=source_relation(a,b,raw,['原表!B2'])['caption_updates']
        self.assertEqual(captions,[{'range_name':'表甲','old_text':'名单甲','text':'名单甲 '+raw.split('；')[0]+'。'}])

    def test_原小计前新增空名行保留未采用金额列为空(self):
        from 附注更新.A合集 import declared_named_row_relations
        a={'机械记录':{'names':{'表':[{'sheet':'原表','range':'$B$2:$D$7'}]}},
           'outline':{'subtables':[{'sheet':'原表','metrics':[{'ranges':['C3'],'template_period_role':'closing'}]}]}}
        raw=('仅限 原表!B2:D7。在明细末尾、原小计之前增加一行。\n'
             'A B3:B4 分别对应 B 来源!B10:B11。新增A明细行的名称保持空白，对应 B B12；期末账面余额对应 B C12。')
        _,tables=declared_named_row_relations(a,{},raw)
        self.assertEqual(tables[0]['source_rows'],[12]);self.assertEqual(tables[0]['columns'],{'C':'C'})
        self.assertEqual(tables[0]['labels'],{'B':'B'})
        self.assertEqual(tables[0]['blank_values'],{'12':{'C':{'value':None},'D':{'value':None}}})
        self.assertEqual(tables[0]['operations'],[{'axis':'row','action':'insert','index':3,'count':1,'template_index':2}])

    def test_匿名新名单接在明确保留的旧明细之后(self):
        from 附注更新.A合集 import declared_named_row_relations
        a={'机械记录':{'names':{'表':[{'sheet':'原表','range':'B2:D5'}]}}}
        raw=('A 原表!B3:D4 的两个原具名明细，不能按行序对应B 来源!A10:C11 的匿名明细。'
             '按B更新名单时列示名称空白的业务行，A的各列分别对应B的各列。A原单位作为旧名单另行保留。\n'
             '合计仍独立：A 原表!B5:D5 对应B 来源!A12:C12。')
        _,tables=declared_named_row_relations(a,{},raw)
        self.assertEqual(tables[0]['first_row_offset'],3)
        self.assertEqual(tables[0]['target_row_offsets'],[3,4])
        self.assertEqual(tables[0]['old_detail_count'],0)
        self.assertEqual(tables[0]['columns'],{'C':'B','D':'C'})
        self.assertEqual(tables[0]['blank_values'],{'10':{'C':{'value':None},'D':{'value':None}},'11':{'C':{'value':None},'D':{'value':None}}})
        self.assertEqual(declared_named_row_relations(a,{},raw.replace('旧名单另行保留','')),( [],[]))

    def test_额外空名业务行沿用同表已有列关系(self):
        from 附注更新.A合集 import declared_named_row_relations
        a={'机械记录':{'names':{'表':[{'sheet':'原表','range':'B2:D5'}]}}}
        raw='## 账龄组合\nA 5“合计” ↔ B 12“合计”\nB 11 是未标明账龄名称的额外行；A表没有对应的未命名明细行。'
        b={'已确认关系':[{'content':raw,'bindings':[{'target':'原表!C5','source':'来源!B12'},{'target':'原表!D5','source':'来源!C12'}]}]}
        _,tables=declared_named_row_relations(a,b,raw)
        self.assertEqual(tables[0]['source_rows'],[11]);self.assertEqual(tables[0]['columns'],{'C':'B','D':'C'})
        self.assertEqual(tables[0]['declared_labels'],{'11':{'value':None}})
        self.assertEqual(tables[0]['operations'],[{'axis':'row','action':'insert','index':3,'count':1,'template_index':2}])
        self.assertEqual(declared_named_row_relations(a,{},raw),( [],[]))

    def test_明确来源行按已有答复的期间标签列展开(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'names':{},'sheets':[{'name':'原表','cells':{}}]}}
        b={'source_hash':'b','facts':[],'机械记录':{'sheets':[{'name':'来源','cells':{}}]},
           '合集来源':[{'识别答复':'C 列为期末，D 列为期初。B 的相应期间标签位置为来源 A2、C4:D4。'}]}
        raw='A 原表!C8期末、D8期初 → B“来源”第11行对应期末、期初栏。'
        self.assertEqual(source_relation(a,b,raw,['原表!C8','原表!D8'])['bindings'],[
            {'target':'原表!C8','source':'来源!C11'},{'target':'原表!D8','source':'来源!D11'}])
        b['合集来源']=[]
        self.assertEqual(source_relation(a,b,raw,['原表!C8','原表!D8'])['bindings'],[])

    def test_明细开放行范围按明说来源行数扩展且不压到首格(self):
        from 附注更新.A合集 import declared_named_row_relations,source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'names':{'表区':[{'sheet':'原表','range':'B2:D5'}]},'sheets':[{'name':'原表','cells':{}}]}}
        b={'source_hash':'b','gold_hash':'g','facts':[],'机械记录':{'sheets':[{'name':'来源','cells':{}}]}}
        raw=('## A 原表!B2:D5\nA原明细为 B3:B4，独立合计为 B5。\n'
            'A明细区 B3:D 对应B 来源!A10:C13，逐行按行号对应。\nA合计行对应B 来源!A14:C14。')
        self.assertEqual(source_relation(a,b,raw,['原表!B3'])['bindings'],[])
        _,tables=declared_named_row_relations(a,b,raw)
        self.assertEqual(tables[0]['source_rows'],[10,11,12,13]);self.assertEqual(tables[0]['columns'],{'C':'B','D':'C'})
        self.assertEqual(tables[0]['totals'],[{'target_row':5,'source_row':14}])
        self.assertEqual(tables[0]['operations'][0]['count'],2)

    def test_已明示账龄来源多一条无名称行沿用既有列角色(self):
        from 附注更新.A合集 import declared_named_row_relations
        a={'机械记录':{'names':{'表区':[{'sheet':'原表','range':'B2:D6'}]},'sheets':[]},
            'outline':{'subtables':[{'sheet':'原表','metrics':[
                {'ranges':['C3:C4'],'template_period_role':'期末'},
                {'ranges':['D3:D4'],'template_period_role':'期初'}]}]}}
        b={'机械记录':{'sheets':[{'name':'来源','cells':{}}]}}
        raw=('### A 原表!B2:D6\nA 各账龄的余额采用 B 来源!F10:F12（期末原值）及 H10:H12（期初原值）；B行名有一条无名称行。\n'
            'A C3:C4、D3:D4：各账龄余额。')
        bindings,tables=declared_named_row_relations(a,b,raw)
        self.assertEqual(bindings,[{'target':'原表!'+col+str(row),'source':'来源!'+src+str(row+7)} for row in (3,4) for col,src in [('C','F'),('D','H')]])
        self.assertEqual(tables[0]['source_rows'],[12]);self.assertEqual(tables[0]['columns'],{'C':'F','D':'H'})
        self.assertEqual(tables[0]['declared_labels']['12'],{'value':None})

    def test_明示原行名称定位插行且未映射新增列留空(self):
        from 附注更新.A合集 import declared_named_row_relations
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'names':{'表区':[{'sheet':'原表','range':'B2:D6'}]},
            'sheets':[{'name':'原表','cells':{'B3':{'value':'甲'},'B4':{'value':'乙'},'B5':{'value':'小计'},'B6':{'value':'合计'}}}]}}
        b={'source_hash':'b','gold_hash':'g','facts':[],'机械记录':{'sheets':[{'name':'来源','cells':{}}]}}
        raw=('### A 原表!B2:D6\nA C3:C4（期末）← B 来源!C10:C11。\n'
            'B B12 为空、C12 有业务值。原小计之前插入一行，期末业务格取 B C12；新增行期初格留空。\n'
            '后移后，A 小计行对应 B 来源!C13；合计行对应 B C14。')
        bindings,tables=declared_named_row_relations(a,b,raw)
        self.assertEqual(bindings,[{'target':'原表!C5','source':'来源!C13'},{'target':'原表!C6','source':'来源!C14'}])
        self.assertEqual(tables[0]['operations'][0],{'axis':'row','action':'insert','index':3,'count':1,'template_index':2})
        self.assertEqual(tables[0]['blank_values']['12']['D'],{'value':None})
        self.assertEqual(tables[0]['columns'],{'C':'C'})

    def test_相同原答复定向复用保留未请求来源和标签(self):
        from 附注更新.A合集 import append_declared_relation
        old={'kind':'取数对应','answer_hash':'同一原答复','bindings':[
            {'target':'原表!C3','source':'来源!B3'},{'target':'原表!C8','source':'来源!B8'}],
            'label_actions':[{'target':'原表!C7','action':'set','value':'另一表标签'}]}
        relations=[old]
        incoming={'kind':'取数对应','answer_hash':'同一原答复','requested_targets':['原表!C3'],
            'bindings':[{'target':'原表!C3','source':'来源!D3'}],'label_actions':[]}
        append_declared_relation(relations,incoming)
        self.assertEqual(relations[0]['bindings'],[old['bindings'][1],incoming['bindings'][0]])
        self.assertEqual(relations[0]['label_actions'],old['label_actions'])
        self.assertEqual(incoming['label_actions'],[])

    def test_表格列日期角色及上表行引用保持各自归属(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.test':{'dimension_refs':[{'id':'item'},{'id':'period'}]}},
            'dimension':{'item':{'value_type':'selection'},'period':{'value_type':'period'},'report_scope':{'allowed_values':['consolidated']}}}
        raw=('原表!C3:F4。\n| 行 | item 实际取值 | C、D 列：period / template_period_role | E、F 列：period / template_period_role |\n'
            '|---|---|---|---|\n| 3 | 指定对象甲 | 2025-12-31 / 期末 | 2024-12-31 / 期初 |\n'
            '| 4 | 指定对象乙 | 2025-12-31 / 期末 | 2024-12-31 / 期初 |\n'
            'C3:F4：metric.test；item 为上表对应行范围。\n'
            '原表!G8：report_scope 实际取值为 `consolidated`。相应报表范围是合并财务报表，不是母公司单体报表。')
        parsed=direct_answer(raw,standard,['原表'],field_scope={'原表!C3:F4':['metric_id','item','period','template_period_role'],'原表!G8':['report_scope']})
        periods=[r for r in parsed['dimensions'] if r['id']=='period' and r['value']=='2025-12-31']
        self.assertTrue(periods)
        self.assertEqual({loc for r in periods for loc in r['ranges']},{'原表!C3','原表!D3','原表!C4','原表!D4'})
        self.assertFalse(any(str(r['value']).startswith('上表') for r in parsed['dimensions']))
        self.assertEqual([r['value'] for r in parsed['dimensions'] if r['id']=='report_scope'],['consolidated'])

    def test_分别来源跨分号保留前述目标及各自日期(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'sheets':[{'name':'原表','cells':{}}]}}
        b={'source_hash':'b','gold_hash':'g','facts':[],'机械记录':{'sheets':[{'name':name,'cells':{}} for name in ('来源一','来源二')]}}
        raw='A 原表!C9:D9：已明确的同类余额，未限定单项；分别对应 B 来源一!I18 （2025-12-31）与 来源二!C29 （2024-12-31）。'
        self.assertEqual(source_relation(a,b,raw,['原表!C9','原表!D9'])['bindings'],[
            {'target':'原表!C9','source':'来源一!I18'},{'target':'原表!D9','source':'来源二!C29'}])



    def test_明确替换合并单元格文字落到承载格(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'sheets':[{'name':'原表','merges':['C6:G6']}]}}
        b={'source_hash':'b','机械记录':{'sheets':[{'name':'来源'}]}}
        result=source_relation(a,b,'将A `原表!C6:G6`合并单元格的文字替换为：**“新标签”**。',['原表!C6'])
        self.assertEqual([(r['target'],r['value']) for r in result['label_actions']],[('原表!C6','新标签')])

    def test_明示新增空名称行与合计后移接入(self):
        from 附注更新.A合集 import declared_named_row_relations
        a={'机械记录':{'names':{'表区':[{'sheet':'原表','range':'B2:H6'}]}}}
        b={'机械记录':{'sheets':[{'name':'来源'}]}}
        reply=('**A 原表!B2:H6**\n'
            '- A第4—5行分别对应B 来源!A10:A11。\n'
            '- （A列C、D、E）分别对应B 来源!B10:B11、C10:C11、D10:D11；（A列F、G、H）分别对应B 来源!E10:E11、F10:F11、G10:G11。\n'
            '- 新增的A第6行对应B第12行：名称取B 来源!A12（空白）。\n'
            '- 原A第6行的独立合计移至A第7行，对应B 来源!A13:G13。')
        _,tables=declared_named_row_relations(a,b,reply)
        self.assertEqual(tables[0]['source_rows'],[10,11,12])
        self.assertEqual(tables[0]['columns'],dict(zip('CDEFGH','BCDEFG')))
        self.assertEqual(tables[0]['totals'],[{'target_row':6,'source_row':13}])
        self.assertEqual(tables[0]['operations'],[{'axis':'row','action':'insert','index':4,'count':1,'template_index':2}])



    def test_两方合并表头引用落到实际保存文字的位置(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'sheets':[{'name':'目标','merges':['B1:D1']}]}}
        b={'source_hash':'b','机械记录':{'sheets':[{'name':'来源','merges':['A10:C10']}]}}
        result=source_relation(a,b,'表头：A B1:D1 → B 来源!A10:C10。',['目标!B3'])
        self.assertEqual(result['label_actions'],[{'action':'copy','target':'目标!B1','source':'来源!A10'}])

    def test_同组槽位只按明示列别名及原有同列引用展开(self):
        from 附注更新.A合集 import _plain_answer
        standard={'dimension':{},'metric':{key:{'dimension_refs':[]} for key in ('metric.amount','metric.share')}}
        raw=('## B 来源!B3:C4 与 B7:C8\n|B行|对应槽位|\n|---|---|\n'
             '|3／7|metric.amount（金额）、metric.share（占比）|\n|4／8|同一组槽位|\n'
             '列映射为：B金额→金额、C比例→比例。')
        mask={'来源!B3:C4':['metric_id'],'来源!B7:C8':['metric_id']}
        answer=_plain_answer(raw,standard,['来源'],{'来源!C7':'metric.share'},field_scope=mask)
        actual={cell:row['id'] for row in answer['metrics'] for cell in row['ranges'] if cell.startswith('来源!')}
        self.assertEqual(actual['来源!B3'],'metric.amount')
        self.assertEqual(actual['来源!C3'],'metric.share')
        self.assertNotIn('来源!C4',actual)
        raw=('## B 来源!B3:D5\n|B行|槽位|\n|---|---|\n'
             '|3–4|B金额、D金额→metric.amount；C比例→metric.share|\n|5|同一组槽位|')
        answer=_plain_answer(raw,standard,['来源'],field_scope={'来源!B3:D5':['metric_id']})
        actual={cell:row['id'] for row in answer['metrics'] for cell in row['ranges'] if cell.startswith('来源!')}
        self.assertEqual(actual['来源!D5'],'metric.amount')
        self.assertEqual(actual['来源!C5'],'metric.share')

    def test_自然文字分列和编号行对应不混入未声明行(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'sheets':[{'name':'目标'}]}}
        b={'source_hash':'b','机械记录':{'sheets':[{'name':'来源'}]}}
        raw=('## 表内关系\nA 目标!B10:H15 → B 来源!A20:G26。\n'
             'A 期末金额、比例、准备 C:D:E → B 期末组 B:C:D；A 期初金额、比例、准备 F:G:H → B 期初组 E:F:G。\n'
             'A 第13、14行分别对应 B 第23、24行；A 第15行“合计”对应 B 第26行。')
        targets=['目标!'+c+str(r) for r in [13,14,15] for c in 'CDEFGH']
        relation=source_relation(a,b,raw,targets)
        self.assertEqual(len(relation['bindings']),len(targets))
        self.assertEqual({r['target']:r['source'] for r in relation['bindings']},
            {'目标!'+a+str(ar):'来源!'+b+str(br) for ar,br in [(13,23),(14,24),(15,26)] for a,b in zip('CDEFGH','BCDEFG')})
        raw=('A 目标!B10:H15；B 来源!A20:G26。\nA 13←B 23；A 14←B 24；A 15采用B 26。\n'
             'A的C/D/E列分别采用B的B/C/D列；A的F/G/H列分别采用B的E/F/G列。')
        self.assertEqual(source_relation(a,b,raw,targets)['bindings'],relation['bindings'])
        raw=('## 整表及明细\nA 目标!B10:D14 → B 来源!A20:C24。\n'
             'A第13行“甲”对应B第24行；A第14行“乙”对应B第23行。')
        targets=['目标!'+c+str(r) for r in [13,14] for c in 'BCD']
        self.assertEqual(source_relation(a,b,raw,targets)['bindings'],[
            {'target':'目标!'+a+str(ar),'source':'来源!'+b+str(br)}
            for ar,br in [(13,24),(14,23)] for a,b in zip('BCD','ABC')])

    def test_分段表区内的列编号只接入本段请求位置(self):
        from 附注更新.A合集 import _plain_answer
        standard={'dimension':{},'metric':{key:{'dimension_refs':[]} for key in ('metric.balance','metric.allowance')}}
        raw=('## 第一表：来源!A1:C4\n- B列：metric.balance\n- C列：metric.allowance\n'
             '## 第二表：来源!A8:C11\n- B列：metric.allowance\n'
             '## 未说明表区\n- C列：metric.balance')
        fields={'来源!B3:C4':['metric_id'],'来源!B10:C11':['metric_id']}
        answer=_plain_answer(raw,standard,['来源'],field_scope=fields)
        actual={cell:row['id'] for row in answer['metrics'] for cell in row['ranges']}
        self.assertEqual(actual,{'来源!B3':'metric.balance','来源!B4':'metric.balance',
            '来源!C3':'metric.allowance','来源!C4':'metric.allowance',
            '来源!B10':'metric.allowance','来源!B11':'metric.allowance'})

    def test_共同日期声明不能补成本格期间(self):
        from 附注更新.A合集 import _per_cell_answer
        reply='期末实际日期为2025年12月31日。\n- 来源!B3：metric_id=m1。'
        saved=_per_cell_answer(reply,{'metric':{'m1':{}},'dimension':{'period':{}}},['来源'])
        self.assertEqual([r['id'] for r in saved['metrics']],['m1'])
        self.assertFalse(saved['dimensions'])

    def test_明确新名单引用及原成员原值单列保留(self):
        from 附注更新.A合集 import declared_named_row_relations,source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{
            'names':{'表':[{'sheet':'原表','range':'B1:D4'}]},'sheets':[{'name':'原表','merges':['C1:D1'],
                'cells':{'B2':{'value':'旧甲'},'B3':{'value':'原小计'},'B4':{'value':'合计'}}}]}}
        b={'source_hash':'b','facts':[],'机械记录':{'sheets':[{'name':'来源','merges':[]}]}}
        raw=('## A 原表!B1:D4\n本表期间标签：A C1:D1 采用 B 来源!B7 的期末标签。\n'
             '| A位置 | B位置 | 说明 |\n|---|---|---|\n'
             '| B2:D4 | 来源!A8:C10 | 明细及独立合计。原小计及其原值作为原有内容单列保留。 |')
        labels=source_relation(a,b,raw,['原表!C1'])['label_actions']
        self.assertIn({'action':'copy','target':'原表!C1','source':'来源!B7'},labels)
        _,tables=declared_named_row_relations(a,b,raw)
        self.assertEqual(tables[0]['source_rows'],[8,9])
        self.assertEqual(tables[0]['retained_members'][0],{'index':2,'cells':{col:{'source':'原表!'+col+'3'} for col in 'BCD'}})
        self.assertEqual(tables[0]['blank_values'],{'9':{'C':{'value':None},'D':{'value':None}}})

    def test_普通名单表复用已识别期间列并接入列简写(self):
        from 附注更新.A合集 import declared_named_row_relations
        a={'机械记录':{'names':{'名单':[{'sheet':'原表','range':'B2:D4'}]}},
           'outline':{'subtables':[{'sheet':'原表','metrics':[
             {'ranges':['C3'],'template_period_role':'current'},
             {'ranges':['D3'],'template_period_role':'prior'}]}]}}
        b={'机械记录':{'sheets':[{'name':'来源'}]}}
        raw=('A 原明细为 原表!B3:B3，合计为 原表!B4。采用 B《来源》完整表。\n'
             '| 更新后 A 排列 | A 位置关系 | 采用的 B 位置 | 名称及依据 |\n|---|---|---|---|\n'
             '| 甲 | A 原明细 | B20；本期/上期来源列 I/J20 | 采用原成员 |\n'
             '| 未命名业务位置 | 新增明细 | B21；I/J21 | 名称保留空白 |\n'
             '| 合计 | A 原合计 | B22；I/J22 | 独立父项 |')
        _,tables=declared_named_row_relations(a,b,raw)
        self.assertEqual(tables[0]['columns'],{'C':'I','D':'J'})
        self.assertEqual(tables[0]['source_rows'],[20,21])
        self.assertEqual(tables[0]['labels'],{'B':'B'})
        self.assertEqual(tables[0]['totals'],[{'target_row':4,'source_row':22}])
        a['outline']['subtables'][0]['metrics']=[]
        self.assertEqual(declared_named_row_relations(a,b,raw),([],[]))

    def test_名单更新后位置表按明示数量和已识别列展开(self):
        from 附注更新.A合集 import declared_named_row_relations
        a={'机械记录':{'names':{'名单':[{'sheet':'原表','range':'B8:D11'}]}},
           'outline':{'subtables':[{'sheet':'原表','metrics':[
             {'ranges':['C9'],'template_period_role':'closing'},
             {'ranges':['D9'],'template_period_role':'opening'}]}]}}
        b={'机械记录':{'sheets':[{'name':'来源'}]}}
        raw=('A《原表》原表!B9:B10原有两行明细，B11为合计。\n'
             'B实际列示3项，应将A原明细区扩为三行。\n'
             '| 更新后A位置 | 搬运的B来源及B成员名称 | B金额列 |\n|---|---|---|\n'
             '| B9 | B《来源》A20：甲 | B20期末；C20期初 |\n'
             '| B10 | A21：乙 | B21期末；C21期初 |\n'
             '| B11 | A22：丙 | B22期末；C22期初 |\n'
             '| B12 | A23：合计 | B23期末；C23期初 |')
        bindings,tables=declared_named_row_relations(a,b,raw)
        self.assertEqual(bindings,[{'target':'原表!B11','source':'来源!A23','label':True}])
        self.assertEqual(len(tables),1)
        self.assertEqual(tables[0]['columns'],{'C':'B','D':'C'})
        self.assertEqual(tables[0]['source_rows'],[20,21,22])
        self.assertEqual(tables[0]['totals'],[{'target_row':11,'source_row':23}])
        self.assertEqual(tables[0]['operations'],[{'axis':'row','action':'insert','index':3,'count':1,'template_index':1}])
        self.assertEqual(declared_named_row_relations(a,b,raw.replace('B实际列示3项，应将A原明细区扩为三行。','')),( [],[]))
        a['outline']['subtables'][0]['metrics']=[]
        self.assertEqual(declared_named_row_relations(a,b,raw),( [],[]))

    def test_既有名单只补明列来源并保留原排列(self):
        from 附注更新.A合集 import declared_named_row_relations
        a={'机械记录':{'names':{'范围':[{'sheet':'原表','range':'B2:D5'}]}}}
        table={'range_name':'范围','selection_dimension':'member','source_sheet':'来源','source_rows':[10,11],
            'columns':{'C':'B'},'totals':[{'target_row':5,'source_row':12}]}
        b={'机械记录':{'sheets':[{'name':'来源'},{'name':'选取'}]},'已确认关系':[{'tables':[table]}]}
        raw=('**名单（A：原表!B2:D5）**\nA的期初列D3:D4沿用已有成员。\n'
             '| B成员及来源位置 | A安排 | 依据 |\n|---|---|---|\n'
             '| 甲：选取!J20 | 原甲 | 来源C10引用J20。 |\n| 乙：选取!J21 | 原乙 | 来源C11引用J21。 |\n'
             '期初合计使用B的来源C12，名单排列保持。')
        bindings,tables=declared_named_row_relations(a,b,raw)
        self.assertEqual(table['columns'],{'C':'B'})
        self.assertEqual(tables[0]['source_rows'],[10,11])
        self.assertEqual(tables[0]['row_sources']['11'],{'D':'选取!J21'})
        self.assertEqual(bindings,[{'target':'原表!D5','source':'来源!C12'}])

    def test_四段来源矩阵保留原成员且合计各取本段来源(self):
        from 附注更新.A合集 import _declared_list_tables
        a={'机械记录':{'names':{'范围':[{'sheet':'原表','range':'B2:F5'}]},
            'sheets':[{'name':'原表','cells':{'B3':{'value':'甲'},'B4':{'value':'保留项'}}}]}}
        raw=('本次采用的B来源为 来源。A原明细为 原表!B3:F4，合计为第5行。\n'
             '| B行 | B项目名称 | 四段来源 |\n|---|---|---|\n'
             '| 20 | 甲 | I20 / I40 / I60 / I80 |\n| 21 | 乙 | I21 / I41 / I61 / I81 |\n'
             'A第3–4行排B的2项，A的 C/D/E/F 列分别沿用原含义。\n'
             'A第5行保留“保留项”及原值。\nA第6行为合计，四列来源 I22 / I42 / I62 / I82。')
        table=_declared_list_tables(a,raw,'member',[{'sheet':'来源','range':'A20:J82'}])[0]
        self.assertEqual(table['row_sources']['22'],{'C':'来源!I22','D':'来源!I42','E':'来源!I62','F':'来源!I82'})
        self.assertEqual(table['retained_members'][0]['cells']['F'],{'source':'原表!F4'})
        self.assertEqual(table['operations'],[{'axis':'row','action':'insert','index':3,'count':1,'template_index':1}])


    def test_原成员名称与明示来源行接入且新增空白位置不丢失(self):
        from 附注更新.A合集 import declared_named_row_relations
        a={'机械记录':{'names':{'范围':[{'sheet':'原表','range':'B2:D5'}]},
            'sheets':[{'name':'原表','cells':{'B3':{'value':'一、甲总体'},'B4':{'value':'旧甲'},'B5':{'value':'二、乙总体'}}}]}}
        b={'机械记录':{'sheets':[{'name':'来源','cells':{'A10':{'value':'一、甲总体'},'A14':{'value':'二、乙总体'}}}]},
            'b_source_ranges':[{'sheet':'来源','range':'A9:C17'}],'已确认关系':[]}
        raw=('**表｜A：原表!B2:D5**\n采用B 来源!A9:C17，各列依次对应，名单随B更新。\n'
             '- **甲段**：A“旧甲”对应B第11行；合计对应B第10行。在合计之后增列B第12行，并保留B第13行空白。\n'
             '- **乙段**：A原有父项对应B第14行；父项下排入B第15—16行，并保留B第17行空白。')
        bindings,tables=declared_named_row_relations(a,b,raw)
        self.assertIn({'target':'原表!C4','source':'来源!B11'},bindings)
        self.assertEqual([t['source_rows'] for t in tables],[[11,12,13],[15,16,17]])
        self.assertEqual([t['operations'][0]['index'] for t in tables],[3,4])
        self.assertTrue(all(t['selection_dimension'] is None for t in tables))
        self.assertEqual(tables[1]['operations'][0]['template_index'],2)
        self.assertEqual(tables[1]['target_row_offsets'],[6,7,8])
        self.assertNotIn('source_ranges',tables[0])
        self.assertEqual(declared_named_row_relations(a,b,raw.replace('各列依次对应','列关系未给出')),( [],[]))

    def test_同段名称和分段金额不把名称广播进金额列(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'sheets':[{'name':'原表','merges':[]}]}}
        b={'source_hash':'b','facts':[],'机械记录':{'sheets':[{'name':'来源','merges':[]}]}}
        raw='金额采用 B 来源 的 I列。\nA B8:F8 名称采用 B B30；金额依次取 I30、I50、I70、I90。'
        result=source_relation(a,b,raw,['原表!'+col+'8' for col in 'BCDEF'])
        self.assertEqual(result['bindings'],[{'target':'原表!'+col+'8','source':'来源!'+source} for col,source in zip('BCDEF',['B30','I30','I50','I70','I90'])])
        kept='A B8:F8父项对应 B B30。A B9:F9“旧项目”、B11:F11“另一旧项目”及原有金额均独立保留，不借给新增行。'
        targets=['原表!'+col+str(row) for row in [8,9,11] for col in 'BCDEF']
        result=source_relation(a,b,kept,targets)
        self.assertEqual(result['retained_targets'],['原表!'+col+str(row) for row in [9,11] for col in 'BCDEF'])

    def test_按已声明前后成员插入名单并保留四段独立来源(self):
        from 附注更新.A合集 import _declared_list_tables
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'names':{'表区':[{'sheet':'原表','range':'B5:F9'}]},
            'sheets':[{'name':'原表','merges':[],'cells':{'B6':{'value':'\t旧项'},'B7':{'value':'\t合计'}}}]},
            'outline':{'subtables':[{'sheet':'原表','range':'B5:F9','metrics':[{'ranges':['C6:F7']}]}]}}
        b={'source_hash':'b','facts':[],'机械记录':{'sheets':[{'name':'来源','merges':[],'cells':{}}]}}
        raw='金额采用 B 来源 的 I列。\n\n新增成员：沿用 member 名单维度，放在原旧项之后、原合计之前。\n- 新甲：B B30；金额 I30、I50、I70、I90。\n- 新乙：B B31；金额 I31、I51、I71、I91。'
        result=_declared_list_tables(a,raw,'member',[{'sheet':'来源','range':'B30:I91'}],b)
        self.assertEqual(result[0]['target_row_offsets'],[2,3])
        self.assertEqual(result[0]['row_sources']['31']['F'],'来源!I91')
        self.assertEqual(result[0]['operations'],[{'axis':'row','action':'insert','index':2,'count':2,'template_index':1}])


    def test_保留指令标题覆盖其后位置且不越过下一标题(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'sheets':[{'name':'原表','merges':[]}]}}
        b={'source_hash':'b','facts':[],'机械记录':{'sheets':[{'name':'来源','merges':[]}]}}
        raw='### 无独立来源，保留A原内容\n- A 原表!D3:D4，原值照留。\n### 采用来源\nA 原表!D5采用B来源!E7。'
        result=source_relation(a,b,raw,['原表!D3','原表!D4','原表!D5'])
        self.assertEqual(result['retained_targets'],['原表!D3','原表!D4'])
        self.assertEqual(result['bindings'],[{'target':'原表!D5','source':'来源!E7'}])

    def test_列角色期间文字仅替换明确列出的表头原文(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'names':{'T1':[{'sheet':'原表','range':'B2:D4'}],'T2':[{'sheet':'原表','range':'B8:D10'}]},
            'sheets':[{'name':'原表','merges':[],'cells':{'C2':{'value':'\t\t期初余额'},'D2':{'value':'期末余额'},'C8':{'value':'期初余额'},'D8':{'value':'期末余额'}}}]}}
        b={'source_hash':'b','facts':[],'机械记录':{'sheets':[{'name':'来源','merges':[]}]}}
        raw='## 期间\n在A原表!B2:D4及B8:D10，局部期间表头替换为：\n- C、D列“期初余额”“期末余额” → 期初余额（2023年12月31日）、期末余额（2024年12月31日）'
        result=source_relation(a,b,raw,['原表!C3'])
        self.assertEqual({r['target']:r['value'] for r in result['label_actions']},
                         {'原表!C2':'期初余额（2023年12月31日）','原表!D2':'期末余额（2024年12月31日）','原表!C8':'期初余额（2023年12月31日）','原表!D8':'期末余额（2024年12月31日）'})

    def test_答复明确的表头对应独立接入标签搬运(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'sheets':[{'name':'原表','merges':[]}]}}
        b={'source_hash':'b','facts':[],'机械记录':{'sheets':[{'name':'来源','merges':[]}]}}
        raw='本次统一采用 B 的 来源。\n表头：A C2:D2 → B B8:C8，分别为期初和期末。\nA C3:D3采用B B9:C9。'
        result=source_relation(a,b,raw,['原表!C3','原表!D3'])
        self.assertEqual(result['label_actions'],[{'action':'copy','target':'原表!C2','source':'来源!B8'},{'action':'copy','target':'原表!D2','source':'来源!C8'}])


    def test_名单位置表的期间来源沿用已识别列角色(self):
        from 附注更新.A合集 import _declared_list_tables
        a={'机械记录':{'names':{'TA_001':[{'sheet':'原表','range':'B10:D13'}]},'sheets':[]},
           'outline':{'subtables':[{'sheet':'原表','range':'B10:D13','metrics':[
               {'ranges':['C11:C13'],'template_period_role':'closing'},
               {'ranges':['D11:D13'],'template_period_role':'opening'}]}]}}
        raw=('A原表款项明细为 B11:B12，合计为 B13；沿用原明细样式。\n'
             '|A版式位置|A更新后的名称|B来源名称|期末金额来源（B）|期初金额来源（B）|\n'
             '|---|---|---|---|---|\n'
             '|B11|甲|来源!A20|来源!B20|来源!C20|\n'
             '|B12|乙|来源!A21|来源!B21|来源!C21|\n'
             '|B13|丙|来源!A22|来源!B22|来源!C22|\n'
             '|B14|合计|来源!A25|来源!B25|来源!C25|')
        result=_declared_list_tables(a,raw,'member',[{'sheet':'来源','range':'A20:C25'},{'sheet':'备选','range':'A1:D8'}])
        self.assertEqual(len(result),1)
        self.assertEqual(result[0]['source_rows'],[20,21,22])
        self.assertEqual(result[0]['old_detail_count'],2)
        self.assertEqual(result[0]['columns'],{'C':'B','D':'C'})
        self.assertEqual(result[0]['labels'],{'B':'A'})
        self.assertEqual(result[0]['totals'],[{'target_row':13,'source_row':25}])

    def test_明选来源不会被后文未采用表覆盖且反向关系可接入(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'sheets':[{'name':'原表','merges':[]}]}}
        b={'source_hash':'b','facts':[],'机械记录':{'sheets':[{'name':'来源甲','merges':[]},{'name':'来源乙','merges':[]}]}}
        raw=('本次采用的B来源为“来源甲”，不采用来源乙。\n'
             'B 来源乙!A1有另一说明，不用于本次搬运。\n'
             'A C3对应B I7。\n- B B8:C8 对应A C4:D4。\n'
             '- A C5、D5（原有明细）保留原内容。\n'
             '- A C6不保留原内容，采用B I9。')
        result=source_relation(a,b,raw,['原表!C3','原表!C4','原表!D4','原表!C5','原表!D5','原表!C6'])
        self.assertEqual({x['target']:x['source'] for x in result['bindings']},
                         {'原表!C3':'来源甲!I7','原表!C4':'来源甲!B8','原表!D4':'来源甲!C8','原表!C6':'来源甲!I9'})
        self.assertEqual(result['retained_targets'],['原表!C5','原表!D5'])
        alternate=raw.replace('本次采用的B来源为“来源甲”','本次统一采用 B 的 来源甲').replace('保留原内容','保留 A 原内容')
        self.assertEqual(source_relation(a,b,alternate,['原表!C3','原表!C4','原表!D4','原表!C5','原表!D5','原表!C6'])['bindings'],result['bindings'])

    def test_先列各行范围后列指标坐标不需要纠正字段名单(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.one':{'dimension_refs':[{'id':'member'}]}},'dimension':{'member':{'value_type':'selection'}}}
        raw=('范围原表!B2:D4。\n### 各行范围\n|行|member|\n|---|---|\n|3|类别甲|\n|4|类别乙|\n'
             '### 各列指标\n|坐标|槽位编号|\n|---|---|\n|C3:C4|metric.one|\n|D3:D4|metric.one|')
        answer=direct_answer(raw,standard,['原表'])
        values={loc:r['value'] for r in answer['dimensions'] if r['id']=='member' for loc in r['ranges']}
        self.assertEqual(values,{'原表!C3':'类别甲','原表!D3':'类别甲','原表!C4':'类别乙','原表!D4':'类别乙'})

    def test_数字行与列头组成字段归属并保留列内日期(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.one':{'dimension_refs':[{'id':'member'},{'id':'period'}]}},
                  'dimension':{'member':{'value_type':'selection'},'period':{'value_type':'period'}}}
        raw=('范围原表!B2:D4。所有金额对应metric.one。\n|行|B列原文|member|C列（期末）|D列（期初）|\n|---|---|---|---|---|\n'
             '|3|项目甲|成员甲|空白；2025-12-31|8；2024-12-31|\n|4|合计|全体|8；2025-12-31|9；2024-12-31|')
        answer=direct_answer(raw,standard,['原表'])
        fields={(r['id'],loc):r['value'] for r in answer['dimensions'] for loc in r['ranges']}
        self.assertEqual(fields[('member','原表!C3')],'成员甲')
        self.assertEqual(fields[('member','原表!D4')],'全体')
        self.assertEqual(fields[('period','原表!C3')],'2025-12-31')
        self.assertEqual(fields[('period','原表!D4')],'2024-12-31')
        self.assertNotIn(('member','原表!B3'),fields)

    def test_A同名表引用不能充当未声明B表默认引用(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'sheets':[{'name':'同名','merges':[]}]}}
        b={'source_hash':'b','facts':[],'机械记录':{'sheets':[{'name':'同名','merges':[]},{'name':'另一来源','merges':[]}]}}
        raw='A同名!B2说明。\n## 明细\nA C/D 分别对应B B/C。\nA 3“项目甲” ↔ B 8“项目甲”'
        self.assertEqual(source_relation(a,b,raw,['同名!C3','同名!D3'])['bindings'],[])

    def test_表区分组说明和引用块原标题替换直接接通(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],
           'source_context':[{'range_name':'TA_001','title':'原分类标题'}],
           '机械记录':{'sheets':[{'name':'原表','merges':[]}]}}
        b={'source_hash':'b','facts':[],'机械记录':{'sheets':[{'name':'来源','merges':[]}]}}
        raw=('**TA_001、TA_002：**原角色保留，建议在这两张表旁增加以下说明：\n\n'
             '> 模型已经明确的共同期间说明。\n\n'
             '**TA_003：**将原标题“原名单标题”替换为：\n\n'
             '> 模型明确的新名单标题。\n\n必要时在表旁补充：\n\n> 可选重复说明。')
        result=source_relation(a,b,raw,['原表!B3'])
        self.assertEqual(result['caption_updates'],[
            {'range_name':'TA_001','old_text':'原分类标题','text':'原分类标题 模型已经明确的共同期间说明。'},
            {'range_name':'TA_003','old_text':'原名单标题','text':'模型明确的新名单标题。'}])
        self.assertEqual(result['label_actions'],[])



    def test_分列分行来源引用和分组坐标直接接通(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'sheets':[{'name':'原表','merges':[]}]}}
        b={'source_hash':'b','facts':[],'机械记录':{'sheets':[{'name':'来源','merges':[]},{'name':'别表','merges':[]}]}}
        raw=('## 分类\nA表的期末列 C/D 对应B的期末列 B/C。\n'
             '- A 3“类别甲” ↔ B 8“类别甲”\n'
             '- A 4“合计” ↔ B 9“合计”\n'
             '## 变动\nA C/D 分别对应B B/C。\n- A 7“增加” ↔ B 12“增加”\n'
             '## 另一组\nA原表!B11对应B来源!A20。\n'
             'B期末两列为 B20/C20，期初对应 D20/E20；分别按A期末 C11/D11、期初 E11/F11 的列含义对应。')
        targets=['原表!'+c+str(r) for r in [3,4,7,11] for c in ('CDEF' if r==11 else 'CD')]
        relation=source_relation(a,b,raw,targets)
        expected={'原表!'+c+str(r):'来源!'+s+str(n) for r,n in [(3,8),(4,9),(7,12)] for c,s in [('C','B'),('D','C')]}
        expected.update({'原表!'+c+'11':'来源!'+s+'20' for c,s in zip('CDEF','BCDE')})
        self.assertEqual({r['target']:r['source'] for r in relation['bindings']},expected)
        ambiguous=raw.replace('## 分类','另有B别表!A1。\n## 分类')
        self.assertFalse(any(r['target']=='原表!C3' for r in source_relation(a,b,ambiguous,targets)['bindings']))

    def test_明示适用范围的日期完整保留而不裁定混写角色(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.one':{'dimension_refs':[{'id':'period'}]}},
                  'dimension':{'period':{'value_type':'period'}}}
        raw=('范围原表!C2:H4。\n表头C列为期末余额，适用于C3:E4，角色为本期、期末，实际日期2025年12月31日；'
             'F列为期初余额，适用于F3:H4，角色为上期、期初，实际日期2024年12月31日。')
        result=direct_answer(raw,standard,['原表'])
        self.assertTrue(any(r['id']=='period' and r['ranges']==['原表!C3:E4'] and r['value']=='2025-12-31' for r in result['dimensions']))
        self.assertTrue(any(r['id']=='period' and r['ranges']==['原表!F3:H4'] and r['value']=='2024-12-31' for r in result['dimensions']))

    def test_答复表格字段不能吞并后续栏且保留字段引用(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.one':{'dimension_refs':[{'id':'member'}]}},
                  'dimension':{'member':{'value_type':'selection'},'period':{'value_type':'period'}}}
        raw=('|坐标|槽位及维度实际取值|期间角色、表内标签及实际日期|\n|---|---|---|\n'
             '|B3|metric.one；member：类别甲。|B列：期末，实际日期2025-12-31。|\n'
             '|C3|metric.one；member同上。|C列：期初，实际日期2024-12-31。|\n'
             '|B4|metric.one；member：本表全部项目（表内合计范围）。|B列：期末，实际日期2025-12-31。|\n'
             '|C4|metric.one；member与合计范围相同。|C列：期初，实际日期2024-12-31。|')
        result=direct_answer(raw,standard,['原表'])
        values={loc:row['value'] for row in result['dimensions'] if row['id']=='member' for loc in row['ranges']}
        self.assertEqual(values['原表!B3'],'类别甲')
        self.assertEqual(values['原表!C3'],'类别甲')
        self.assertEqual(values['原表!B4'],'本表全部项目（表内合计范围）')
        self.assertEqual(values['原表!C4'],values['原表!B4'])

    def test_表头文字沿用同表已识别列角色(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],
           '机械记录':{'sheets':[{'name':'原表','merges':[]}]},
           'outline':{'subtables':[{'sheet':'原表','range':'A2:C4','metrics':[
               {'ranges':['B3:B4'],'template_period_role':'closing'},
               {'ranges':['C3:C4'],'template_period_role':'opening'}]}]}}
        b={'source_hash':'b','facts':[],'机械记录':{'sheets':[{'name':'来源','merges':[]}]}}
        raw=('期末角色的表头文字为“期末新文字”，期初角色为“期初新文字”：\n'
             '- A表性质列头原表!B2:C2，对应B性质表来源!F5:G5。')
        result=source_relation(a,b,raw,['原表!B3','原表!C3'])
        self.assertEqual({r['target']:r['value'] for r in result['label_actions']},
                         {'原表!B2':'期末新文字','原表!C2':'期初新文字'})

    def test_先列指标顺序再明示行关系接通列(self):
        from 附注更新.A合集 import direct_answer,source_relation
        standard={'metric':{key:{'dimension_refs':[]} for key in ['metric.one','metric.two']},'dimension':{}}
        raw=('## 分类\nB计量列与A列对应为：B金额列→A金额列；C比例列→A比例列。对应指标依次是metric.one、metric.two。\n'
             '- A表 原表!A3:D3 ← B 来源!A8:C8：所选类别。')
        answer=direct_answer(raw,standard,['来源'])
        found={cell:row['id'] for row in answer['metrics'] for cell in row['ranges']}
        self.assertEqual(found['来源!B8'],'metric.one');self.assertEqual(found['来源!C8'],'metric.two')
        a={'source_hash':'a','gold_hash':'g','facts':[
            {'record_type':'metric_fact','source_reference':{'location':'原表!C3'},'metric_id':'metric.one'},
            {'record_type':'metric_fact','source_reference':{'location':'原表!D3'},'metric_id':'metric.two'}],
           '机械记录':{'sheets':[{'name':'原表','merges':[]}]}}
        b={'source_hash':'b','facts':[
            {'record_type':'metric_fact','source_reference':{'location':'来源!B8'},'metric_id':'metric.one'},
            {'record_type':'metric_fact','source_reference':{'location':'来源!C8'},'metric_id':'metric.two'}],
           '机械记录':{'sheets':[{'name':'来源','merges':[]}]}}
        relation=source_relation(a,b,raw,['原表!C3','原表!D3'])
        self.assertEqual({r['target']:r['source'] for r in relation['bindings']},{'原表!C3':'来源!B8','原表!D3':'来源!C8'})

    def test_跨句来源及分组表头原文接入(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'sheets':[{'name':'原表','merges':[]}]}}
        b={'source_hash':'b','机械记录':{'sheets':[{'name':'来源','merges':[]}]}}
        raw=('A表 原表!C3:D3所需的是本表净额。原值列不作净额；直接同范围净额来源为B 来源!I8（期末）、J8（期初）。\n'
             '## 应搬入期间文字\n期末角色的表头文字为“期末（2024年末）”，期初角色为“期初（2024年初）”：\n'
             '- A表列头原表!C2:D2，对应来源的期末、期初列头。')
        relation=source_relation(a,b,raw,['原表!C3','原表!D3'])
        self.assertEqual(relation['bindings'],[{'target':'原表!C3','source':'来源!I8'},{'target':'原表!D3','source':'来源!J8'}])
        self.assertEqual({row['target']:row['value'] for row in relation['label_actions']},
                         {'原表!C2':'期末（2024年末）','原表!D2':'期初（2024年初）'})

    def test_名单名称和金额分表沿用已明示列角色(self):
        from 附注更新.A合集 import _declared_list_tables
        a={'机械记录':{'names':{'名单':[{'sheet':'原表','range':'$B$10:$D$13'}]}},
           'outline':{'subtables':[{'sheet':'原表','range':'B10:D13','metrics':[
               {'ranges':['C11:C13'],'template_period_role':'closing'},
               {'ranges':['D11:D13'],'template_period_role':'opening'}]}]}}
        raw=('### 性质表：A 原表!B10:D13\nA 原版式为2个明细行加1个合计行。\n'
             '名称及明细顺序取自 B 标签!E20:E22；金额采用对应的 B 金额!G30:G32（期末）、H30:H32（期初）。\n'
             'A 合计行采用 B 金额!G33:H33。')
        result=_declared_list_tables(a,raw,'member',[])[0]
        self.assertEqual(result['columns'],{'C':'G','D':'H'})
        self.assertEqual(result['source_rows'],[30,31,32])
        self.assertEqual(result['declared_labels']['32'],{'source':'标签!E22'})
        self.assertEqual(result['old_detail_count'],2)
        self.assertEqual(result['totals'],[{'target_row':13,'source_row':33}])


    def test_列期间声明旁列标签原文引用随明确金额范围保存(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.one':{'dimension_refs':[{'id':'period'}]}},
                  'dimension':{'period':{'value_type':'period'}}}
        raw=("范围：原表!B2:D4。\n"
             "列期间标签及实际期间（适用于以下明细和合计）：\n"
             "- C列，原表!C2，原文“期末余额”：period=2025-12-31，template_period_role=期末。\n"
             "- D列，原表!D2，原文“期初余额”：period=2024-12-31，template_period_role=期初。\n"
             "- C3:D4：metric.one。")
        answer=direct_answer(raw,standard,['原表'])
        labels={loc:set(row.get('label_references',row.get('evidence',[]))) for row in answer['dimensions']
                if row['id']=='period' for loc in row['ranges'] if row.get('label_references',row.get('evidence'))}
        self.assertEqual(labels['原表!C3'],{'原表!C2'})
        self.assertEqual(labels['原表!D4'],{'原表!D2'})

    def test_来源括注角色连接紧随的分别日期不混为一个值(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.one':{'dimension_refs':[{'id':'period'}]}},
                  'dimension':{'period':{'value_type':'period'}}}
        raw=("- [A位置]分别采用 B 来源!B3:B4（期末）及 来源!C3:C4（期初）。\n"
             "  金标准编号：metric.one。\n"
             "- B 期间取值：期末 2024-12-31，期初 2024-01-01。原值照源保留。")
        answer=direct_answer(raw,standard,['来源'])
        dates={loc:r['value'] for r in answer['dimensions'] if r['id']=='period' for loc in r['ranges']}
        self.assertEqual(dates,{'来源!B3':'2024-12-31','来源!B4':'2024-12-31',
                               '来源!C3':'2024-01-01','来源!C4':'2024-01-01'})

    def test_B采用范围承接紧随编号且A需求不污染B口径(self):
        from 附注更新.A合集 import _own_answer_text,direct_answer
        standard={'metric':{'metric.one':{'dimension_refs':[{'id':'report_scope'}]}},
                  'dimension':{'report_scope':{'value_type':'enum','allowed_values':['consolidated','standalone']}}}
        raw=("B 的 report_scope=consolidated；A 需求为单体，故 A 的 report_scope=standalone。\n"
             "- A 原表!C2:D5（四项）分别采用 B 来源!B3:B6（期末）及 来源!C3:C6（期初）。\n"
             "  金标准编号：metric.one。")
        snapshot={'sheets':[{'name':'来源'}]}
        own=_own_answer_text(raw,{'carrier':'B','A需要':{'需要':[{'location':'原表!C2'}]}},snapshot)
        self.assertNotIn('standalone',own)
        result=direct_answer(own,standard,['来源'])
        self.assertEqual(result['metrics'][0]['ranges'],['来源!B3:B6','来源!C3:C6'])
        self.assertEqual(next(r['value'] for r in result['dimensions'] if r['id']=='report_scope'),'consolidated')

    def test_旧说明跨表引用保留但不充当本表身份范围(self):
        from 附注更新.A合集 import _direct_table,direct_answer
        table={'sheet':'来源','range':'B3','data_ranges':['B3'],'metrics':[],'dimensions':[],
               'unresolved':[{'id':'entity','ranges':['另一表!C4'],'reason':'旧解析说明','original':'原答复引用另一表!C4'}]}
        saved=_direct_table(table,{'source_context':[]},{'metric':{},'dimension':{}})
        self.assertEqual(saved['unresolved'][0]['ranges'],[])
        self.assertEqual(saved['unresolved'][0]['referenced_ranges'],['另一表!C4'])
        self.assertEqual(saved['unresolved'][0]['original'],'原答复引用另一表!C4')
        standard={'metric':{'metric.one':{'dimension_refs':[]}},'dimension':{'entity':{'value_type':'text'},'currency':{'value_type':'text'}}}
        answer=direct_answer('B报告主体是甲公司，金额币种为 CNY。',standard,['来源'])
        self.assertEqual(next(r['value'] for r in answer['dimensions'] if r['id']=='entity'),'甲公司')

    def test_矩形分别采用两列来源保留列分组和括注之后的来源(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],
           '机械记录':{'sheets':[{'name':'原表','merges':[]}]}}
        b={'source_hash':'b','机械记录':{'sheets':[{'name':'来源','merges':[]}]}}
        raw='A 原表!C2:D5（四个项目）分别采用 B 来源!B10:B13（期末）及 来源!C10:C13（期初）。'
        targets=[f'原表!{col}{row}' for row in range(2,6) for col in ('C','D')]
        relation=source_relation(a,b,raw,targets)
        self.assertEqual({row['target']:row['source'] for row in relation['bindings']},
                         {f'原表!{col}{row}':f'来源!{src}{row+8}' for row in range(2,6)
                          for col,src in [('C','B'),('D','C')]})

    def test_B共同身份只用于答复已选来源不铺满候选(self):
        from 附注更新.A合集 import direct_answer,_own_answer_text,_answer_tables
        standard={'metric':{'metric.change':{'dimension_refs':[{'id':'member'},{'id':'currency'}]}},
                  'dimension':{'member':{'value_type':'selection','domain':'项目'},
                               'currency':{'value_type':'text'}}}
        raw=("B来源的共同身份取值：标准编号 metric.change；currency=CNY。\n"
             "可直接采用的 B 来源按行对应如下：\n"
             "| A坐标 | B来源（发生额表） | member |\n|---|---|---|\n"
             "| C2 | B3 | 类别甲 |\n| C3 | B4 | 类别乙 |")
        snapshot={'sheets':[{'name':name,'max_row':6,'max_column':4} for name in ('发生额表','余额表')]}
        areas=[{'id':name,'sheet':name,'range':'A1:D6','data_ranges':['B3:C4']}
               for name in ('发生额表','余额表')]
        mapping={'carrier':'B','答复指纹':'例','target_tables':areas,
                 'field_scope':{name+'!B3:C4':['metric_id','currency','member'] for name in ('发生额表','余额表')}}
        answer=direct_answer(_own_answer_text(raw,mapping,snapshot),standard,['发生额表','余额表'])
        tables=_answer_tables(answer,mapping,snapshot)
        selected={t['sheet']+'!'+cell for t in tables for row in t['metrics'] for cell in row['ranges']}
        self.assertEqual(selected,{'发生额表!B3','发生额表!B4'})
        self.assertFalse(tables[1]['dimensions'])
        self.assertEqual(tables[1]['data_ranges'],['B3:C4'])
        self.assertTrue(any(r['id']=='currency' and set(r['ranges'])=={'B3','B4'} for r in tables[0]['dimensions']))
        only_common={'metrics':[],'dimensions':[{'id':'currency','ranges':['*'],'value':'CNY'}],'unresolved':[]}
        common=_answer_tables(only_common,mapping,snapshot)
        self.assertTrue(all(t['dimensions'][0]['ranges']==['B3:C4'] for t in common))

    def test_末段明列列日期仍适用前面已列金额(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.one':{'dimension_refs':[{'id':'period'}]}},
                  'dimension':{'period':{'value_type':'period'}}}
        raw="""### 各金额格对应
C2:D3：metric.one。
上述金额格的实际期间取值依列相同：C列为2025-12-31、角色期末；D列为2024-12-31、角色期初。
"""
        parsed=direct_answer(raw,standard,['原表'])
        dates={loc:r['value'] for r in parsed['dimensions'] if r['id']=='period' for loc in r['ranges']}
        self.assertEqual(dates['原表!C2'],'2025-12-31')
        self.assertEqual(dates['原表!D3'],'2024-12-31')

    def test_答复列标题中的期间各归对应列(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.one':{'dimension_refs':[{'id':'period'}]}},
                  'dimension':{'period':{'value_type':'period'}}}
        answer="""范围：原表!C2:D3
metric_id（全部坐标）：metric.one。
| 坐标 | C列 | D列 |
|---|---|---|
| C2:D2 | 期末，2025-12-31 | 期初，2024-12-31 |
| C3:D3 | 期末，2025-12-31；C3为空 | 期初，2024-12-31 |
"""
        parsed=direct_answer(answer,standard,['原表'])
        dates={tuple(r['ranges']):r['value'] for r in parsed['dimensions'] if r['id']=='period'}
        self.assertEqual(dates[('原表!C2',)],'2025-12-31')
        self.assertEqual(dates[('原表!D2',)],'2024-12-31')
        self.assertEqual(dates[('原表!C3',)],'2025-12-31')
        self.assertEqual(dates[('原表!D3',)],'2024-12-31')
        self.assertNotIn(('*',),dates)

    def test_全部坐标的编号及分组范围含旁述引用仍完整接入(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.one':{'dimension_refs':[{'id':'selection'}]}},
                  'dimension':{'selection':{'value_type':'selection'}}}
        answer="""**原表!C2:D4**
- **metric_id（全部坐标）**：metric.one。
- **selection**：
  - C2:D2：members — 类别甲。
  - C3:D3：members — 类别乙；C3为空，不据此推定类别不存在。
  - C4:D4：predicate — 原表总体，不限于列出的明细。
"""
        answer=answer.replace('members —',chr(96)+'members'+chr(96)+' —')
        parsed=direct_answer(answer,standard,['原表'])
        self.assertTrue(any(row['id']=='metric.one' and row['ranges']==['原表!C2:D4'] for row in parsed['metrics']))
        values={tuple(row['ranges']):row['value'] for row in parsed['dimensions'] if row['id']=='selection'}
        self.assertEqual(values[('原表!C2:D2',)],{'mode':'members','members':['类别甲']})
        self.assertEqual(values[('原表!C3:D3',)],{'mode':'members','members':['类别乙；C3为空，不据此推定类别不存在']})
        self.assertEqual(values[('原表!C4:D4',)],{'mode':'predicate','predicate':'原表总体，不限于列出的明细'})

    def test_同名表标题只接明确目标表区且保留裸标签引用(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{
            'sheets':[{'name':'原表','merges':[]}],
            'names':{'TA_001':[{'sheet':'原表','range':'$A$2:$C$4'}],
                     'TA_002':[{'sheet':'原表','range':'$A$7:$C$9'}]}}}
        b={'source_hash':'b','机械记录':{'sheets':[{'name':'来源','merges':[]}]}}
        raw=('目标表：原表!A7:C9。不涉及另一张同名表原表!A2:C4。\n'
             '表头标题：原文“明细表”；替换为“明细表（原明细旧期，合计新期）”。\n'
             '合计标签：B9 原文“合计”；替换为“合计（新期）”。')
        relation=source_relation(a,b,raw,['原表!C9'])
        self.assertEqual(relation['caption_updates'],[
            {'range_name':'TA_002','old_text':'明细表','text':'明细表（原明细旧期，合计新期）'}])
        self.assertEqual([(r['target'],r['value']) for r in relation['label_actions']],[('原表!B9','合计（新期）')])

    def test_标签原文与替换文字用分号分开时直接接入(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'sheets':[{'name':'原表','merges':[]}]}}
        b={'source_hash':'b','机械记录':{'sheets':[{'name':'来源','merges':[]}]}}
        raw='- 分类表，原表!C2（表头）：原文“期初余额”；替换为“期初余额（2024年1月1日）”。'
        relation=source_relation(a,b,raw,['原表!C3'])
        self.assertEqual([(r['target'],r['value']) for r in relation['label_actions']],
                         [('原表!C2','期初余额（2024年1月1日）')])

    def test_分列行关系与分段共同日期完整接入(self):
        from 附注更新.A合集 import _own_answer_text, direct_answer, source_relation
        standard={'metric':{'metric.one':{'dimension_refs':[{'id':'selection'},{'id':'period'}]},
            'metric.ratio':{'dimension_refs':[{'id':'selection'},{'id':'period'}],
                'definition':{'ratio':{'numerator':{'dimension_bindings':{'selection':{'op':'dimension','key':'selection'}}}}}}},
            'dimension':{'selection':{'value_type':'selection'},'period':{'value_type':'period'}}}
        raw="""## 分类
对应B 来源!A5:D6。
| A行及项目 | B对应行及项目 | A列 ← B列 | 范围字段 |
|---|---|---|---|
| 原表!3 甲 | 来源!5 甲 | C←B：metric.one；D←C：metric.ratio | selection：predicate「类别甲」 |
| 原表!4 乙 | 来源!6 乙 | 同上 | selection：predicate「类别乙」 |
上表共同维度：period={date:2024-01-01, kind:instant}。
## 合计
A 原表!C9 ← B 来源!B8：metric.one；selection为predicate「合计总体」。
A 原表!D9 ← B 来源!C8：metric.ratio；分子范围同上。
A 原表!E9 ← B 来源!D8：metric.one；selection 同上。
共同维度：period={date:2024-12-31, kind:instant}。"""
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'sheets':[{'name':'原表','merges':[]}]}}
        b={'carrier':'B','source_hash':'b','b_source_ranges':[{'sheet':'来源'}],
            'A需要':{'需要':[{'location':'原表!C3'}]},'机械记录':{'sheets':[{'name':'来源','merges':[]}]}}
        targets=['原表!'+c for c in ('C3','D3','C4','D4','C9','D9','E9')]
        relation=source_relation(a,b,raw,targets)
        self.assertEqual({r['target']:r['source'] for r in relation['bindings']},
            dict(zip(targets,['来源!'+c for c in ('B5','C5','B6','C6','B8','C8','D8')])))
        parsed=direct_answer(_own_answer_text(raw,b,b['机械记录']),standard,['来源'],default_sheet='来源')
        def values(cell,key):
            return [r['value'] for r in parsed['dimensions'] if r['id']==key and '来源!'+cell in r['ranges']]
        self.assertEqual(values('B5','selection'),[{'mode':'predicate','predicate':'类别甲'}])
        self.assertEqual(values('C6','selection'),[{'mode':'predicate','predicate':'类别乙'}])
        self.assertEqual(values('B5','period'),[{'date':'2024-01-01','kind':'instant'}])
        self.assertEqual(values('B8','period'),[{'date':'2024-12-31','kind':'instant'}])
        self.assertEqual(values('C8','selection'),[{'mode':'predicate','predicate':'合计总体'}])
        self.assertEqual(values('D8','selection'),[{'mode':'predicate','predicate':'合计总体'}])

    def test_期间段落按明写实际日期和条件括号保存(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.one':{'dimension_refs':[{'id':'period'},{'id':'selection'}]}},
                  'dimension':{'period':{'value_type':'period'},'selection':{'value_type':'selection'}}}
        raw="""期间：下列余额均为时点数据，template_period_role=期末，实际日期为2025-12-31。
原表!C5：metric.one；selection=predicate[甲范围；含乙项]。"""
        result=direct_answer(raw,standard,['原表'])
        self.assertEqual([r['value'] for r in result['dimensions'] if r['id']=='period'],['2025-12-31'])
        self.assertEqual(next(r['value'] for r in result['dimensions'] if r['id']=='selection'),
                         {'mode':'predicate','predicate':'甲范围；含乙项'})

    def test_列声明的分母及行范围说明保留原文(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.one':{'dimension_refs':[{'id':'selection'},{'id':'test_denominator_selection'},{'id':'period'}]}},
                  'dimension':{'selection':{'value_type':'selection'},'test_denominator_selection':{'value_type':'selection'},'period':{'value_type':'period'}}}
        raw="""范围：原表!B5:C6。
期间：表头位置原表!B3:C4；本区各数据为期初，template_period_role=期初，实际日期 period=2024-12-31。
- C列比例：metric.one；分母范围为甲总体，不以明细替代总体。
| 行 | B列原文及对应 selection | 范围说明 |
|---|---|---|
| 5 | 甲 | 甲父项下明细 |
| 6 | 合计 | 本表甲总体 |"""
        result=direct_answer(raw,standard,['原表'])
        values=[r for r in result['dimensions'] if r['id']=='test_denominator_selection']
        self.assertEqual({loc for r in values for loc in r['ranges']},{'原表!C5','原表!C6'})
        self.assertTrue(all(r['value']=='甲总体，不以明细替代总体' for r in values))
        self.assertTrue(any(r['id']=='selection' and '本表甲总体' in r['value'] for r in result['dimensions']))
        self.assertTrue(any(r['id']=='period' and r['ranges']==['原表!B5:C6'] and r['value']=='2024-12-31' for r in result['dimensions']))

    def test_范围内分别声明的列指标沿用明示行范围(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.one':{'dimension_refs':[]},'metric.two':{'dimension_refs':[]}},'dimension':{}}
        raw="""**范围：原表!B5:E6。**
**列对应：**
- **C列“金额”**：metric.one。
- **E列“余额”**：metric.two。
| 行 | 名称 |
|---|---|
| 5 | 甲 |
| 6 | 乙 |"""
        result=direct_answer(raw,standard,['原表'])
        expected={'metric.one':{'原表!C5','原表!C6'},'metric.two':{'原表!E5','原表!E6'}}
        actual={key:{loc for row in result['metrics'] if row['id']==key for loc in row['ranges']} for key in expected}
        self.assertEqual(actual,expected)

    def test_行号和列字母分栏时接到明确交叉位置(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.one':{'dimension_refs':[{'id':'selection'}]},'metric.two':{'dimension_refs':[{'id':'selection'}]}},
                  'dimension':{'selection':{'value_type':'selection'}}}
        raw="""| 行（名称） | 列及原文列名 | 金标准编号及实际含义 |
|---|---|---|
| 5 甲 | C 金额 | metric.one；selection=甲对象 |
| 5 甲 | E 余额 | metric.two；selection=乙对象 |"""
        result=direct_answer(raw,standard,['原表'])
        self.assertEqual({r['id']:r['ranges'] for r in result['metrics']},
                         {'metric.one':['原表!C5'],'metric.two':['原表!E5']})
        self.assertEqual({tuple(r['ranges']):r['value'] for r in result['dimensions']},
                         {('原表!C5',):'甲对象',('原表!E5',):'乙对象'})

    def test_唯一需要采用来源标题并约束后续维度归属(self):
        from 附注更新.A合集 import direct_answer,source_relation
        raw="""- **选择的B来源：**来源!C6，即本次采用项目。
- **编号与维度：**
  - metric.test：测试金额
  - selection：甲范围
- **不选其他格：**来源!C7不是该范围。"""
        standard={'metric':{'metric.test':{'dimension_refs':[{'id':'selection'}]}},
                  'dimension':{'selection':{'value_type':'selection'}}}
        result=direct_answer(raw,standard,['来源'])
        self.assertEqual(result['metrics'][0]['ranges'],['来源!C6'])
        self.assertEqual(result['dimensions'][0]['ranges'],['来源!C6'])
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'sheets':[{'name':'原表'}]}}
        b={'source_hash':'b','gold_hash':'g','机械记录':{'sheets':[{'name':'来源'}]}}
        self.assertEqual(source_relation(a,b,raw,['原表!D8'])['bindings'],
                         [{'target':'原表!D8','source':'来源!C6'}])
        self.assertEqual(source_relation(a,b,raw,['原表!D8','原表!E8'])['bindings'],[])

    def test_中文紧接标准编号及对象后注释不丢失已明写取值(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.test':{'dimension_refs':[{'id':'selection'}]}},
                  'dimension':{'selection':{'value_type':'selection'}}}
        raw='- 来源 原表!C6；来源 原表!D6。采用metric.test；selection={domain:测试范围, mode:all}，即本范围全部项目。'
        result=direct_answer(raw,standard,['原表'])
        self.assertEqual(result['metrics'][0]['id'],'metric.test')
        self.assertEqual(result['metrics'][0]['ranges'],['原表!C6','原表!D6'])
        self.assertEqual(result['dimensions'][0]['value'],{'domain':'测试范围','mode':'all'})

    def test_项目分隔标题中的位置保留各自编号和维度(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.first':{'dimension_refs':[{'id':'selection'},{'id':'form'}]},
                            'metric.second':{'dimension_refs':[{'id':'selection'}]}},
                  'dimension':{'selection':{'value_type':'selection'},'form':{'value_type':'text'}}}
        raw="""- **原表!C6、D6｜甲项目**：metric_id=metric.first；selection=甲范围；form=未确定，原文未说明。
- **原表!C8、D8｜乙项目**：metric_id=metric.second；selection=乙范围。"""
        result=direct_answer(raw,standard,['原表'])
        metric_ranges={r['id']:r['ranges'] for r in result['metrics']}
        self.assertEqual(metric_ranges,{'metric.first':['原表!C6','原表!D6'],'metric.second':['原表!C8','原表!D8']})
        values={(r['id'],tuple(r['ranges'])):r['value'] for r in result['dimensions']}
        self.assertEqual(values[('selection',('原表!C6','原表!D6'))],'甲范围')
        self.assertEqual(values[('selection',('原表!C8','原表!D8'))],'乙范围')
        self.assertEqual(values[('form',('原表!C6','原表!D6'))],'未知')

    def test_两字段共用表头按原声明接入(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.test':{'dimension_refs':[{'id':'first'},{'id':'second'}]}},
                  'dimension':{'first':{'value_type':'selection'},'second':{'value_type':'selection'}}}
        raw="""metric.test；first 与 second 采用同一范围值。
| 位置 | first、second 的范围和值 |
|---|---|
| 原表!C6 | predicate：甲项及其紧邻调整的小计。 |
"""
        result=direct_answer(raw,standard,['原表'])
        values={r['id']:r['value'] for r in result['dimensions']}
        self.assertEqual(values['first'],values['second'])
        self.assertIn('甲项及其紧邻调整的小计',str(values['first']))

    def test_B表头所注明工作表用于身份与来源对应(self):
        from 附注更新.A合集 import _own_answer_text,direct_answer,source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'sheets':[{'name':'原表','cells':{}}]}}
        b={'carrier':'B','source_hash':'b','gold_hash':'g',
           'A需要':{'需要':[{'location':'原表!E6'}]},
           '机械记录':{'sheets':[{'name':'来源','cells':{}},{'name':'另表','cells':{}}]}}
        raw='metric_id=metric.test。\n| A坐标 | B来源（来源） | `selection` 实际取值 |\n|---|---|---|\n| E6 | C41 | 甲父项 |'
        own=_own_answer_text(raw,b,b['机械记录'])
        standard={'metric':{'metric.test':{'dimension_refs':[{'id':'selection'}]}},
                  'dimension':{'selection':{'value_type':'selection'}}}
        result=direct_answer(own,standard,['来源','另表'])
        self.assertTrue(any(r['ranges']==['来源!C41'] and r['value']=='甲父项' for r in result['dimensions']))
        self.assertEqual(source_relation(a,b,raw,['原表!E6'])['bindings'],
                         [{'target':'原表!E6','source':'来源!C41'}])
        changed=source_relation(a,b,raw+'\n撤销上次错误落到 A!E6、E7 的来源对应，按表中关系改接。',['原表!E6','原表!E7'])
        self.assertEqual(changed['retained_targets'],['原表!E6','原表!E7'])
        self.assertEqual(changed['bindings'],[{'target':'原表!E6','source':'来源!C41'}])


    def test_答复明说范围同时归属两个字段并沿用同上(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.test':{'dimension_refs':[{'id':'first'},{'id':'second'}]}},
                  'dimension':{'first':{'value_type':'selection'},'second':{'value_type':'selection'}}}
        raw="""下列“项目范围”取值同时归属 `first` 与 `second`，只按本行原声明保存。
| 位置 | 槽位及取值 |
|---|---|
| C6 | metric.test；项目范围：甲父项。 |
| D6 | metric.test；范围同上。 |
| C7:D7 | metric.test；范围限于“甲父项”下的其他项目；C、D 无适用标准。 |
"""
        result=direct_answer(raw,standard,['原表'])
        values={(r['id'],tuple(r['ranges'])):r['value'] for r in result['dimensions']}
        for key in ('first','second'):
            self.assertEqual(values[key,('原表!C6',)],'甲父项')
            self.assertEqual(values[key,('原表!D6',)],'甲父项')
            self.assertEqual(values[key,('原表!C7:D7',)],'“甲父项”下的其他项目')

    def test_否定示例不成为全表编号且逐列未知不吞肯定声明(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{m:{'dimension_refs':[]} for m in ['metric.one','metric.two']},'dimension':{}}
        raw="""不得把甲套成 `metric.one`；甲的 metric_id 均为未知／无适用标准，乙可用 `metric.two`。
| 位置 | 适用槽位 |
|---|---|
| C6 | metric_id 未知／无适用标准。 |
| D6 | metric_id=metric.two。 |
| C7:D7 | C 无适用标准，D 用 metric.two。 |
"""
        result=direct_answer(raw,standard,['原表'])
        self.assertEqual({(r['id'],loc) for r in result['metrics'] for loc in r['ranges']},
                         {('metric.two','原表!D6'),('metric.two','原表!D7')})
        unknown={loc for r in result['unresolved'] if r.get('id')=='metric_id' and r.get('value')=='未知' for loc in r['ranges']}
        self.assertEqual(unknown,{'原表!C6','原表!C7'})
        from 附注更新.A合集 import _merge_direct_table
        table={'id':'t','sheet':'原表','range':'B6:D7','data_ranges':['B6:D7'],
               'metrics':[{'id':'metric.one','ranges':['B6:D7'],'original':raw.splitlines()[0]}],
               'dimensions':[],'unresolved':[]}
        old={'facts':[],'source_context':[],'outline':{'subtables':[table]}}
        with patch('附注更新.A合集._standard',return_value=standard):
            merged,_,_=_merge_direct_table(old,raw,{loc:['metric_id'] for loc in ('B6','C6','D6','C7','D7')})
        self.assertEqual([r['ranges'] for r in merged['metrics'] if r['id']=='metric.one'],[['B7']])
        standard['dimension']={'currency':{'value_type':'text','domain':'币种'},'report_scope':{'value_type':'enum'},
                               'direction':{'value_type':'enum'}}
        grouped=direct_answer('共同取值：报表口径为**合并**；币种 **CNY**，单位倍率 **1**。\n'
            +raw+'\n上述 metric.two 格的 direction 均为“净变动”。',standard,['原表'])
        self.assertEqual(next(r['value'] for r in grouped['dimensions'] if r['id']=='currency'),'CNY')
        self.assertEqual(next(r['value'] for r in grouped['dimensions'] if r['id']=='report_scope'),'合并')
        self.assertEqual(set(next(r['ranges'] for r in grouped['dimensions'] if r['id']=='direction')),
                         {'原表!D6','原表!D7'})



    def test_共同金额列与逐行指标不扩成整段且日期各自保留(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{m:{'dimension_refs':[{'id':k} for k in ('period','selection','report_scope')]} for m in ['metric.one','metric.two']},
                  'dimension':{'period':{'value_type':'period'},'selection':{'value_type':'selection'},
                               'report_scope':{'value_type':'enum','allowed_values':['consolidated','standalone']}}}
        raw="""- **报表范围（report_scope）**：`consolidated`（合并）。原文在第679行表题写“调节表”，正文称“合并现金流量表”。
- C 列均为 **本期（template_period_role）**，实际期间为 2025年度；D 列均为 **上期**，实际期间为 2024年度。
每组所列槽位适用于该行的 C、D 两格；空白格保留。
| 行及原文项目（B列） | C、D格对应的金标准槽位及范围 |
|---|---|
| 6 项目一 | metric.one。 |
| 7 项目二 | metric.two；selection：现金，不含现金等价物。C7：期末，实际时点2025年12月31日；D7：期末，实际时点2024年12月31日。 |
"""
        actual=direct_answer(raw,standard,['原表'])
        self.assertEqual({r['id']:set(r['ranges']) for r in actual['metrics']},
                         {'metric.one':{'原表!C6','原表!D6'},'metric.two':{'原表!C7','原表!D7'}})
        periods={r['ranges'][0]:r['value'] for r in actual['dimensions'] if r['id']=='period'}
        self.assertEqual(periods['原表!C7'],'2025-12-31')
        self.assertEqual(periods['原表!D7'],'2024-12-31')
        roles={loc:r['value'] for r in actual['dimensions'] if r['id']=='template_period_role' for loc in r['ranges']}
        self.assertEqual(roles,{'原表!C6':'current','原表!D6':'prior','原表!C7':'closing','原表!D7':'closing'})
        common=next(r for r in actual['dimensions'] if r['id']=='report_scope')
        self.assertEqual(common['ranges'],['*']);self.assertEqual(common['value'],'consolidated')
        self.assertEqual({r['range'] for r in actual['non_business']},{'原表!B6','原表!B7'})
        self.assertEqual(next(r['value'] for r in actual['dimensions'] if r['id']=='selection'),'现金，不含现金等价物')

    def test_B列日期括注与范围说明不混成共同值或新成员(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.test':{'dimension_refs':[{'id':k} for k in ('period','selection','form')]}},
                  'dimension':{'period':{'value_type':'period'},'selection':{'value_type':'selection'},
                               'form':{'value_type':'enum','allowed_values':['gross','net']}}}
        raw="""A本期金额列对应B的C列、期间为2025年度；A上期金额列对应B的D列、期间为2024年度。
| 位置 | 金标准编号 | 维度 |
|---|---|---|
| C6；D6 | metric.test | selection=members（现金，不含现金等价物）；form=unknown（原文未说明）。 |
| C7（2024-12-31）；D7（2023-12-31） | metric.test | selection=members（现金等价物）。 |
来源!C8:D8对应金标准 metric.test。B实际事项范围（维度 **selection**）是“融资租入固定资产”，不是A所需的“新增资产”。
"""
        result=direct_answer(raw,standard,['来源'])
        periods={loc:r['value'] for r in result['dimensions'] if r['id']=='period' for loc in r['ranges']}
        for loc,value in {'来源!C6':'2025年度','来源!D6':'2024年度','来源!C7':'2024-12-31','来源!D7':'2023-12-31'}.items():
            self.assertEqual(periods[loc],value)
        selection=next(r['value'] for r in result['dimensions'] if r['id']=='selection' and '来源!C6' in r['ranges'])
        self.assertEqual(selection,'members（现金，不含现金等价物）')
        self.assertEqual(next(r['value'] for r in result['dimensions'] if r['id']=='selection' and '来源!C8:D8' in r['ranges']),'融资租入固定资产')
        self.assertEqual(next(r['value'] for r in result['dimensions'] if r['id']=='form'),'未知')


    def test_原句重接撤下误铺字段且保留范围外原声明(self):
        from 附注更新.A合集 import _merge_direct_table
        standard={'metric':{'metric.one':{'dimension_refs':[{'id':'first'}]},'metric.two':{'dimension_refs':[{'id':'second'}]}},
                  'dimension':{'first':{'value_type':'text'},'second':{'value_type':'text'}}}
        first='原表!B3：metric.one；first=甲。'
        second='原表!B4：metric.two；second=乙。'
        table={'id':'t','sheet':'原表','range':'B3:C4','data_ranges':['B3:C4'],
               'metrics':[{'id':'metric.one','ranges':['B3:C4'],'original':first},
                          {'id':'metric.two','ranges':['B3:B4'],'original':second}],
               'dimensions':[{'id':'first','ranges':['B3:C4'],'value':'甲','original':first},
                             {'id':'second','ranges':['B3:B4'],'value':'乙','original':second}],'unresolved':[]}
        old={'facts':[],'source_context':[],'outline':{'subtables':[table]}}
        with patch('附注更新.A合集._standard',return_value=standard):
            merged,_,_=_merge_direct_table(old,first+'\n'+second,{loc:['metric_id','first','second'] for loc in ('B3','B4')})
        def locations(kind,key):
            from 附注更新.A语义阅读 import addresses
            return {c for row in merged[kind] if row['id']==key for area in row['ranges'] for c in addresses(area)}
        self.assertEqual(locations('dimensions','first'),{'B3','C3','C4'})
        self.assertEqual(locations('dimensions','second'),{'B4'})
        self.assertEqual(locations('metrics','metric.one'),{'B3','C3','C4'})
        self.assertEqual(locations('metrics','metric.two'),{'B4'})

    def test_同一编号及选择范围沿用明确前项(self):
        from 附注更新.A合集 import direct_answer
        raw="""- 来源!C31：metric.test；selection=现金等价物小计。
- 来源!D31：同一编号及选择范围。

以上四格的共同维度值：currency=CNY；unit_scale=1，人民币元。
"""
        standard={'metric':{'metric.test':{'dimension_refs':[{'id':'selection'},{'id':'currency'},{'id':'unit_scale'}]}},
                  'dimension':{'selection':{'value_type':'selection'},'currency':{'value_type':'text'},'unit_scale':{'value_type':'number'}}}
        answer=direct_answer(raw,standard,['来源'])
        self.assertTrue(any(r['id']=='metric.test' and '来源!D31' in r['ranges'] for r in answer['metrics']))
        self.assertTrue(any(r['id']=='selection' and '来源!D31' in r['ranges'] and r['value']=='现金等价物小计' for r in answer['dimensions']))
        units=[r for r in answer['dimensions'] if r['id']=='unit_scale']
        self.assertEqual(units[0]['value'],1)
        self.assertEqual(set(units[0]['ranges']),{'来源!C31','来源!D31'})

    def test_上述来源共同身份覆盖明列来源且不混入A角色(self):
        from 附注更新.A合集 import _own_answer_text,direct_answer
        mapping={'carrier':'B'};snapshot={'sheets':[{'name':'来源'}]}
        raw="""- A C3 ← B 来源!C9；A D3 ← B 来源!D9。两格均为 metric.test。
- A C4 ← B 来源!C10；A D4 ← B 来源!D10。两格均为 metric.test。

上述B来源均为报告主体“甲”、currency=CNY。A列期间角色按需沿用：C列 template_period_role=closing，D列 template_period_role=opening；
A该格需按 metric.other 识别。
"""
        own=_own_answer_text(raw,mapping,snapshot)
        self.assertNotIn('template_period_role',own)
        self.assertNotIn('metric.other',own)
        standard={'metric':{'metric.test':{'dimension_refs':[{'id':'entity'},{'id':'currency'}]}},
                  'dimension':{'entity':{'value_type':'text'},'currency':{'value_type':'text'}}}
        answer=direct_answer(own,standard,['来源'])
        row=next(r for r in answer['dimensions'] if r['id']=='entity')
        self.assertEqual(set(row['ranges']),{'来源!C9','来源!D9','来源!C10','来源!D10'})

    def test_逐行逐列明示日期及同一范围不落到表头(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.test':{'dimension_refs':[{'id':'selection'},{'id':'period'}]}},
                  'dimension':{'selection':{'value_type':'selection'},'period':{'value_type':'period'}}}
        raw="""编号 metric.test。
selection：按各行所列现金、现金等价物分别识别；明细清单不替代合计总体。
期间栏位标签位于 C4、D4：原文分别为“期末余额”“期初余额”。
| 原表行 | B列项目原文 | C列：期末余额 | D列：期初余额 |
|---|---|---|---|
| 5 | 现金 | selection：现金小计，不含现金等价物；period：2025-12-31，角色“期末” | 同一现金小计；period：2024-12-31，角色“期初” |
| 6 | 库存现金 | 选择范围：库存现金；2025-12-31，“期末”；原格空白 | 同一选择范围；2024-12-31，“期初”；原格空白 |
"""
        actual=direct_answer(raw,standard,['原表'])
        rows={(r['ranges'][0],r['id']):r for r in actual['dimensions']}
        self.assertEqual(rows['原表!C5','selection']['value'],'现金小计，不含现金等价物')
        self.assertEqual(rows['原表!D5','selection']['value'],'现金小计，不含现金等价物')
        self.assertEqual(rows['原表!D6','selection']['value'],'库存现金')
        self.assertEqual(rows['原表!C6','period']['value'],'2025-12-31')
        self.assertEqual(rows['原表!D6','period']['value'],'2024-12-31')
        self.assertEqual(rows['原表!D6','period']['label_references'],['原表!D4'])
        self.assertEqual({r['range'] for r in actual['non_business']},{'原表!B5','原表!B6'})
        self.assertEqual(set(actual['data_ranges']),{'原表!C5','原表!D5','原表!C6','原表!D6'})

    def test_槽位维度取值列不误作标准编号(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.test':{'dimension_refs':[{'id':'selection'}]}},
                  'dimension':{'selection':{'value_type':'selection'}}}
        raw="""共同槽位：metric.test。
| 坐标 | 对象及范围 | 槽位维度取值 |
|---|---|---|
| C2、D2 | 甲项目 | selection＝members：甲 |
| C3、D3 | 合计 | selection＝all |
"""
        actual=direct_answer(raw,standard,['原表'])
        self.assertEqual(actual['unresolved'],[])
        self.assertEqual(actual['metrics'][0]['id'],'metric.test')
        self.assertEqual([(r['ranges'],r['value']) for r in actual['dimensions']],
                         [(['原表!C2','原表!D2'],{'mode':'members','members':['甲']}),
                          (['原表!C3','原表!D3'],'all')])

    def test_明示名单列与保留列不按同期间混用(self):
        from 附注更新.A合集 import _declared_list_tables
        def fact(sheet,cell,mode,member='甲'):
            return {'record_type':'metric_fact','metric_id':'metric.test','source_reference':{'location':sheet+'!'+cell},
                    'dimensions':{'selection':{'mode':mode,**({'members':[member]} if mode=='members' else {})}}}
        a={'facts':[fact('原表',c+str(r),'members' if r==3 else 'all') for c in 'CD' for r in [3,4]]+[fact('原表','B3','all')],
           '机械记录':{'names':{'名单':[{'sheet':'原表','range':'B2:E4'}]},
                       'sheets':[{'name':'原表','cells':{'B3':{'value':'甲'},'E3':{'value':9},'E4':{'value':9}}}]},
           'outline':{'subtables':[{'sheet':'原表','metrics':[{'ranges':['C3:C4','E3:E4'],'template_period_role':'current'},
                                                             {'ranges':['D3:D4'],'template_period_role':'prior'}]}]}}
        b={'facts':[fact('来源',c+str(r),'all' if r==9 else 'members','甲' if r==7 else '乙') for c in 'BC' for r in [7,8,9]],
           '机械记录':{'sheets':[{'name':'来源'}]},
           'outline':{'subtables':[{'sheet':'来源','metrics':[{'ranges':['B7:B9'],'template_period_role':'current'},
                                                             {'ranges':['C7:C9'],'template_period_role':'prior'}]}]}}
        raw="""A的C列接B本期，A的D列接B上期。
### 项目
金标准：metric.test；范围维度：selection。
B 来源!B7:C8 为明细，B9:C9为合计。
A沿用原样式，明细按B顺序排布于 B3:B4，合计移至5。
原有E列内容随原成员及合计保留，新成员E列留空。"""
        tables=_declared_list_tables(a,raw,'selection',[{'sheet':'来源','range':'A6:C9'}],b)
        self.assertEqual(len(tables),1)
        table=tables[0]
        self.assertEqual(table['columns'],{'C':'B','D':'C'})
        self.assertEqual(table['blank_values']['7']['E'],{'source':'原表!E3'})
        self.assertEqual(table['blank_values']['8']['E'],{'value':None})
        self.assertEqual(table['blank_values']['9']['E'],{'source':'原表!E4'})
        self.assertEqual(table['operations'],[{'axis':'row','action':'insert','index':2,'count':1,'template_index':1}])

    def test_逐行期间引用不串到同列其他表(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.test':{'dimension_refs':[{'id':'period'}]}},
                  'dimension':{'period':{'value_type':'period'}}}
        raw="""期间角色与日期：本期＝2025年度（2025-01-01至2025-12-31）；上期＝2024年度（2024-01-01至2024-12-31）。
| 位置 | metric_id | template_period_role；B期间标签 |
|---|---|---|
| 来源!B3:B4 | metric.test | current；B标签 来源!B2“本期发生额”。 |
| 来源!C8:C9 | 同上 | prior；B标签 来源!C7“上期发生额”。 |
"""
        rows=direct_answer(raw,standard,['来源'])['dimensions']
        dates=[r for r in rows if r['id']=='period']
        self.assertEqual([(r['ranges'],r['label_references'],r['value']) for r in dates],
                         [(['来源!B3:B4'],['来源!B2'],'2025-01-01至2025-12-31'),
                          (['来源!C8:C9'],['来源!C7'],'2024-01-01至2024-12-31')])

    def test_分节共同标准及依次成员保持重复名称(self):
        from 附注更新.A合集 import direct_answer
        from 附注更新.A语义阅读 import addresses
        standard={'metric':{'metric.test':{'dimension_refs':[{'id':'selection'}]}},
                  'dimension':{'selection':{'value_type':'selection'},'period':{'value_type':'period'}}}
        raw="""各表本期对应2025-01-01 至 2025-12-31，上期对应2024-01-01 至 2024-12-31。
### 项目
金标准：metric.test；范围维度：selection。

B 来源!B3:C5 为明细；B6:C6 为合计。明细按B原顺序各自取 members：甲、乙、甲。B6合计取 all。
B的B列是本期、C列是上期。
"""
        parsed=direct_answer(raw,standard,['来源'])
        rows={}
        for item in parsed['dimensions']:
            if item['id']=='selection':
                for area in item['ranges']:
                    for cell in addresses(area.rsplit('!',1)[-1]):rows[cell]=item['value']
        self.assertEqual(rows['B3'],{'mode':'members','members':['甲']})
        self.assertEqual(rows['B4'],{'mode':'members','members':['乙']})
        self.assertEqual(rows['B5'],{'mode':'members','members':['甲']})
        self.assertEqual(rows['B6'],{'mode':'all'})
        self.assertTrue(any(r['id']=='period' and r['ranges']==['来源!C5'] and
                            r['value'].replace(' ','')=='2024-01-01至2024-12-31' for r in parsed['dimensions']))

    def test_列分组及并列日期接入不混淆其中部分(self):
        from 附注更新.A合集 import direct_answer
        from 附注更新.A语义阅读 import addresses
        standard={'metric':{'metric.test':{'dimension_refs':[{'id':'selection'},{'id':'period'}]}},'dimension':{'selection':{'value_type':'selection'},'period':{'value_type':'period'}}}
        scope={'来源!C3:E4':['selection','period','template_period_role'],
               '来源!C8:D8':['selection','period','template_period_role'],
               '来源!C10:D11':['period','template_period_role']}
        raw="""C列“本期发生额”及E列“条件金额”均为 template_period_role=本期，期间为2025-01-01至2025-12-31；D列“上期发生额”为 template_period_role=上期，期间为2024-01-01至2024-12-31。
| 行及项目 | selection：C、D列 | selection：E列 |
|---|---|---|
| 3 甲 | members：甲 | predicate：甲中符合条件的部分 |
| 4 合计 | all：全部事项 | predicate：全部事项中符合条件的部分 |
| 行 | C列：本期发生额 | D列：上期发生额 | E列：条件金额 |
|---|---|---|---|
| 3 甲 | C3：原文内容 | D3：原文内容 | E3：空白 |
### 来源!C8:D8
| 坐标 | 事项 | period | template_period_role |
|---|---|---|---|
| C8、D8 | all：全部事项 | C8：2025年1月1日至2025年12月31日；D8：2024年1月1日至2024年12月31日 | C8：本期发生额；D8：上期发生额 |
- C10:D11中C列：period=2025年1月1日至2025年12月31日；template_period_role=本期发生额。
- C10:D11中D列：period=2024年1月1日至2024年12月31日；template_period_role=上期发生额。"""
        result=direct_answer(raw,standard,['来源'],field_scope=scope)
        def values(cell,key):
            return [r['value'] for r in result['dimensions'] if r['id']==key and any(
                cell in addresses(a.rsplit('!',1)[-1]) for a in r['ranges'] if a!='*')]
        self.assertEqual(values('C3','selection'),['members：甲'])
        self.assertEqual(values('E3','selection'),['predicate：甲中符合条件的部分'])
        self.assertEqual(values('E4','selection'),['predicate：全部事项中符合条件的部分'])
        self.assertTrue(values('C8','selection'))
        for cell,year in [('C3','2025'),('D3','2024'),('E3','2025'),('C8','2025'),('D8','2024'),('C10','2025'),('D11','2024')]:
            self.assertIn(year+'-01-01至'+year+'-12-31',values(cell,'period'),cell)
        self.assertNotIn('2024-01-01至2024-12-31',values('C10','period'))

    def test_上列各格维度指代及多列共同实际日期(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.test':{'dimension_refs':[{'id':'selection'},{'id':'period'}]}},
            'dimension':{'selection':{'value_type':'selection'},'period':{'value_type':'period'}}}
        raw=('C列“本期发生额”和E列“条件金额”：template_period_role=本期；实际期间为2025年度（2025-01-01至2025-12-31）。\n'
             'D列“上期发生额”：template_period_role=上期；实际期间为2024年度（2024-01-01至2024-12-31）。\n'
             '来源!C3：metric_id=metric.test；指标取值范围为 members：甲。\n'
             '来源!D3：metric_id=metric.test；指标取值范围为 members：甲。\n'
             '来源!E3：metric_id=metric.test；指标取值范围为 predicate：甲中符合原文条件的部分。\n'
             'selection按上列各格分别取值。')
        result=direct_answer(raw,standard,['来源'])
        values=[r for r in result['dimensions'] if r['id']=='selection']
        self.assertTrue(any(r['value']=={'mode':'members','members':['甲']} for r in values))
        self.assertTrue(any(r['value']=={'mode':'predicate','predicate':'甲中符合原文条件的部分'} for r in values))
        for col,year in [('C','2025'),('D','2024'),('E','2025')]:
            self.assertTrue(any(r['id']=='period' and r['ranges']==['来源!'+col+'3'] and r.get('value')==year+'-01-01至'+year+'-12-31' for r in result['dimensions']),col)

    def test_自然名单安排复用已识别期间并保留原成员(self):
        from 附注更新.A合集 import _declared_list_tables
        def fact(sheet,cell,mode):
            return {'record_type':'metric_fact','metric_id':'metric.test','source_reference':{'location':sheet+'!'+cell},
                'dimensions':{'selection':{'mode':mode,'members':['旧成员'] if mode=='members' else []}}}
        a={'facts':[fact('原表',c+str(r),'members' if r==3 else 'all') for c in 'BC' for r in [3,4]],
           '机械记录':{'names':{'名单':[{'sheet':'原表','range':'A2:C4'}]}},
           'outline':{'subtables':[{'sheet':'原表','metrics':[
               {'ranges':['B3:B4'],'template_period_role':'current'},{'ranges':['C3:C4'],'template_period_role':'prior'}]}]}}
        b={'facts':[fact('来源',c+str(r),'all' if r==9 else 'members') for c in 'BC' for r in [7,8,9]],
           '机械记录':{'sheets':[{'name':'来源'}]}}
        raw=('明细用维度selection、mode=members。\n'
             '- B 来源!A7 甲：metric.test；本期B7=0，上期C7 原文空白，不填金额。\n'
             '- B 来源!A8 乙：同一指标；本期B8=2，上期C8=0。\n'
             '- B 来源!A9 合计：同一指标；mode=all，本期B9=2，上期C9=0。\n'
             '按 A 原表样式，项目表更新为上述两项明细后接合计。本期、上期金额列分别取自 B 的 B、C 列。'
             'A原明细“旧成员”对应的原表!B3:C3值为1、2，B没有同名来源；这组A原内容应保留在原行/保留位。')
        tables=_declared_list_tables(a,raw,'selection',[{'sheet':'来源','range':'A6:C9'}],b)
        self.assertEqual(len(tables),1);table=tables[0]
        self.assertEqual(table['source_rows'],[7,8]);self.assertEqual(table['retained_rows'],[3])
        self.assertEqual(table['target_row_offsets'],[2,3])
        self.assertEqual(table['operations'],[{'axis':'row','action':'insert','index':2,'count':2,'template_index':1}])
        self.assertEqual(table['columns'],{'B':'B','C':'C'})
        self.assertEqual(table['blank_values'],{'7':{'C':{'value':None}}})

    def test_B逐项文字保留名称与金额来源的分工(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.test':{'dimension_refs':[{'id':'selection'},{'id':'period'}]}},
            'dimension':{'selection':{'value_type':'selection'},'period':{'value_type':'period'},
                'entity':{'value_type':'text'},'report_scope':{'value_type':'text'},'currency':{'value_type':'text'},'unit_scale':{'value_type':'number'}}}
        raw=('共同范围：两表均为示例公司合并报表，币种人民币（CNY），金额倍率 1。\n'
             '原表“本期／上期”对应2025年度／2024年度，分别保留期间：2025-01-01 至 2025-12-31、2024-01-01 至 2024-12-31。\n'
             '明细用维度 selection、mode=members，合计用同维度mode=all。\n'
             '- B 来源!A5 甲：metric.test；成员取值“甲”。本期值来源 来源!B5=0；上期 来源!C5 原文空白，不填金额。\n'
             '- B 来源!A6 乙：同一指标；成员取值“乙”。本期 B6=2，上期 C6=0。\n'
             '- B 来源!A7 合计：同一指标，范围mode=all，本期 B7=2，上期 C7=0。')
        result=direct_answer(raw,standard,['来源'])
        self.assertEqual({v['id'] for v in result['metrics']},{'metric.test'})
        self.assertEqual({loc for v in result['metrics'] for loc in v['ranges']},{'来源!'+c+str(r) for c in 'BC' for r in [5,6,7]})
        self.assertEqual({v['range'] for v in result['non_business']},{'来源!A5','来源!A6','来源!A7'})
        self.assertTrue(any(v['id']=='selection' and v['value']=={'mode':'members','members':['乙']} for v in result['dimensions']))
        self.assertTrue(any(v['id']=='period' and v['ranges']==['来源!C5'] and '2024-01-01' in v['value'] for v in result['dimensions']))
        self.assertTrue(any(v['id']=='entity' and v['value']=='示例公司' for v in result['dimensions']))

    def test_分期来源表及自然逐行名单按原答复接入(self):
        from 附注更新.A合集 import direct_answer, _declared_list_tables
        standard={'metric':{'metric.test':{'dimension_refs':[{'id':'selection'}]}},
            'dimension':{'selection':{'value_type':'selection'},'period':{'value_type':'period'}}}
        raw=('metric_id=metric.test。\n'
             'B来源“本期发生额”对应期间2025-01-01至2025-12-31；“上期发生额”对应2024-01-01至2024-12-31。\n'
             'A的B列“本期发生额”取B的B列，A的C列“上期发生额”取B的C列。\n'
             '| B来源项目及坐标 | 金标准选取范围 | 2025本期来源 | 2024上期来源 |\n|---|---|---|---|\n'
             '| 甲乙，来源!B5:C5 | selection：members；实际成员为“甲”和“乙”。合列，不拆分。 | B5：0 | C5：空白 |\n'
             '| 合计，来源!B6:C6 | selection：all，总体 | B6：0 | C6：空白 |\n'
             'A原表 原表!A2:C4 的明细为第3至3行，合计为第4行。\n'
             '## A表更新后的成员与排布\n'
             '- A第3行：甲乙，金额来源为B 来源!B5:C5。\n'
             '- A第4行：合计，金额来源为B 来源!B6:C6。')
        result=direct_answer(raw,standard,['来源'],default_sheet='来源')
        values=[r['value'] for r in result['dimensions'] if r['id']=='selection']
        self.assertIn({'mode':'members','members':['甲','乙']},values)
        self.assertIn('all',values)
        for cell,date,role in [('B5','2025-01-01至2025-12-31','current'),('C5','2024-01-01至2024-12-31','prior')]:
            self.assertTrue(any(r['id']=='period' and r['ranges']==['来源!'+cell] and r['value']==date for r in result['dimensions']))
            self.assertTrue(any(r['id']=='template_period_role' and r['ranges']==['来源!'+cell] and r['value']==role for r in result['dimensions']))
        a={'机械记录':{'names':{'名单':[{'sheet':'原表','range':'$A$2:$C$4'}]}}}
        tables=_declared_list_tables(a,raw,'selection',[{'sheet':'来源','range':'A4:C6'}])
        self.assertEqual(len(tables),1)
        self.assertEqual(tables[0]['source_rows'],[5])
        self.assertEqual(tables[0]['columns'],{'B':'B','C':'C'})
        self.assertEqual(tables[0]['declared_labels'],{'5':{'value':'甲乙'}})
        self.assertEqual(tables[0]['totals'],[{'target_row':4,'source_row':6}])

    def test_年度附日期及带引号的范围文字不被截断(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.test':{'dimension_refs':[{'id':'selection'},{'id':'period'}]}},
            'dimension':{'selection':{'value_type':'selection'},'period':{'value_type':'period'}}}
        raw=('来源!A3:C4：metric_id=metric.test。A列是项目名称。\n'
             'C列为 `2025年度（2025-01-01至2025-12-31）`，template_period_role=本期。\n'
             '| 坐标 | selection | C列：本期发生额 |\n|---|---|---|\n'
             '| A3、C3 | members：项目项下的“甲”。说明保留。 | C3：空白 |\n'
             '| A4、C4 | members：“乙”这一类别；保留从属范围，不作为总体。 | C4：空白 |')
        result=direct_answer(raw,standard,['来源'])
        values=[r['value'] for r in result['dimensions'] if r['id']=='selection']
        self.assertEqual(values,[{'mode':'members','members':['项目项下的“甲”']},{'mode':'members','members':['“乙”这一类别；保留从属范围，不作为总体']}])
        periods=[r for r in result['dimensions'] if r['id']=='period']
        self.assertTrue(periods);self.assertTrue(all(r['value']=='2025-01-01至2025-12-31' for r in periods))
        self.assertEqual({r['range'] for r in result['non_business']},{'来源!A3','来源!A4'})
        one=direct_answer('来源!C3：selection=members：“甲、乙”。',standard,['来源'])
        self.assertEqual(one['dimensions'][0]['value']['members'],['甲、乙'])

    def test_明示列范围的期间角色接同句标签(self):
        from 附注更新.A合集 import direct_answer
        raw='来源!B5:B8（本期），C5:C8（上期）；本期、上期分别由 来源!B4“本期发生额”、C4“上期发生额”标识。'
        result=direct_answer(raw,{'metric':{'metric.test':{}},'dimension':{'period':{'value_type':'period'}}},['来源'],field_scope={'来源!B5:C8':['period']})
        rows=[r for r in result['dimensions'] if r['id']=='period']
        self.assertEqual({tuple(r['label_references']) for r in rows},{('来源!B4',),('来源!C4',)})
        self.assertTrue(all('value' not in row for row in rows))

    def test_明示明细合计行采用期间标签原文(self):
        from 附注更新.A合集 import direct_answer
        raw=('来源!B3 原文：本期发生额；来源!C3 原文：上期发生额。\n'
             '这两个列标签分别适用于第4行明细和第5行合计。\n'
             '各格均对应metric.test。\n'
             '| 格子 | period及template_period_role |\n|---|---|\n'
             '| 来源!B4:B5 | 2025-01-01 至 2025-12-31；`本期` |\n'
             '| 来源!C4:C5 | 2024-01-01 至 2024-12-31；`上期` |')
        result=direct_answer(raw,{'metric':{'metric.test':{'dimension_refs':[{'id':'period'}]}},'dimension':{'period':{'value_type':'period'}}},['来源'])
        refs=[r for r in result['dimensions'] if r['id']=='period']
        self.assertEqual({tuple(r['label_references']) for r in refs},{('来源!B3',),('来源!C3',)})
        self.assertEqual({loc for r in refs for loc in r['ranges']},{'来源!B4','来源!C4','来源!B5','来源!C5'})

    def test_定向范围复用模型已选共同说明(self):
        from 附注更新.A语义阅读 import _with_declared_context
        mapping={'正常入口范围认识':{'ranges':[{'tables':[
            {'sheet':'来源','range':'A20:D30','context_references':[{'sheet':'来源','range':'A1:D1'}]},
            {'sheet':'无关表','range':'A20:D30','context_references':[{'sheet':'无关表','range':'A1:D1'}]}]}]}}
        scope=[{'sheet':'来源','range':'A22:C24'}]
        result=_with_declared_context(mapping,scope)
        self.assertEqual(result,[{'sheet':'来源','range':'A22:C24','context_references':[{'sheet':'来源','range':'A1:D1'}]}])
        self.assertEqual(scope,[{'sheet':'来源','range':'A22:C24'}])

    def test_B来源箭头表保留分列期间与名单安排(self):
        from 附注更新.A合集 import direct_answer, _own_answer_text, _declared_list_tables
        raw=('共同维度：currency、unit_scale 均未知。\n'
             'B表头“本期发生额”对应template_period_role=current、实际期间2025-01-01—2025-12-31；“上期发生额”对应template_period_role=prior、实际期间2024-01-01—2024-12-31。\n'
             '| B来源坐标 | selection 实际取值 | A采用行 |\n|---|---|---|\n'
             '| 来源!B5:C5 | `mode=members`；`domain=项目`；`members=[甲]` | A第3行 |\n'
             '| 来源!B6:C6 | `mode=members`；`domain=项目`；`members=[乙]` | 在 A 第4行新增 |\n'
             '| 来源!B7:C7 | `mode=all`；`domain=项目` | A第5行 |\n'
             'A 原明细起止为第 3—3 行，合计为第 4 行；名单更新后明细为第3—4行，合计为第5行。\n'
             '| 项目 | B本期来源格 → A本期格（current；2025-01-01—2025-12-31） | B上期来源格 → A上期格（prior；2024-01-01—2024-12-31） |\n|---|---|---|\n'
             '| 甲 | 来源!B5 → A B3 | 来源!C5 → A C3 |\n'
             '| 乙 | 来源!B6 → A 新增 B4 | 来源!C6 → A 新增 C4 |\n'
             '| 本表合计 | 来源!B7 → A B5 | 来源!C7 → A C5 |')
        standard={'metric':{'metric.test':{'dimension_refs':[{'id':'selection'}]}},'dimension':{
            'selection':{'value_type':'selection'},'period':{'value_type':'period'},'currency':{'value_type':'text'},'unit_scale':{'value_type':'number'}}}
        b={'carrier':'B','b_source_ranges':[{'sheet':'来源','range':'A4:C7'},{'sheet':'底稿','range':'A1:D8'}],
           '机械记录':{'sheets':[{'name':'来源'},{'name':'底稿'}]}}
        own=_own_answer_text(raw,b,b['机械记录']);result=direct_answer(own,standard,['来源','底稿'],default_sheet='来源')
        roles=[r for r in result['dimensions'] if r['id']=='template_period_role']
        self.assertTrue(roles);self.assertFalse(any(r['ranges']==['*'] for r in roles))
        self.assertTrue(any(r['ranges']==['来源!B5'] and r['value']=='current' for r in roles))
        self.assertTrue(any(r['ranges']==['来源!C5'] and r['value']=='prior' for r in roles))
        self.assertTrue(any(r['id']=='selection' and r['value'].get('members')==['甲'] for r in result['dimensions'] if isinstance(r['value'],dict)))
        self.assertEqual({r['id'] for r in result['dimensions'] if r['value']=='未知'},{'currency','unit_scale'})
        a={'机械记录':{'names':{'名单':[{'sheet':'原表','range':'$A$2:$C$4'}]}}}
        tables=_declared_list_tables(a,raw,'selection',b['b_source_ranges'],b)
        self.assertEqual(len(tables),1)
        self.assertEqual(tables[0]['source_rows'],[5,6]);self.assertEqual(tables[0]['columns'],{'B':'B','C':'C'})
        self.assertEqual(tables[0]['declared_labels'],{'5':{'value':'甲'},'6':{'value':'乙'}})
        self.assertEqual(tables[0]['totals'],[{'target_row':4,'source_row':7}])
        self.assertEqual(tables[0]['operations'],[{'axis':'row','action':'insert','index':2,'count':1,'template_index':1}])


    def test_相对名单顺序沿用原表并按已选公式引用取值(self):
        from 附注更新.A合集 import _declared_list_tables
        a={'source_hash':'a','gold_hash':'g','facts':[],'slot_bindings':[],
           '机械记录':{'names':{'名单':[{'sheet':'原表','range':'$A$1:$C$7'}]},'sheets':[{'name':'原表','merges':[],
              'cells':{f'A{row}':{'value':name} for row,name in [(2,'1．组甲'),(3,'甲'),(4,'旧乙'),(5,'组乙'),(6,'丁'),(7,'合计')]}}]}}
        b={'source_hash':'b','机械记录':{'sheets':[{'name':'来源','merges':[],
            'cells':{**{f'A{row}':{'value':name} for row,name in [(21,'甲'),(22,'丙'),(31,'丁'),(32,'戊')]},
                     'B21':{'formula':"='底稿'!I1"},'C21':{'formula':"='底稿'!J1"}}},
            {'name':'底稿','merges':[],'cells':{}}]},
           '答复接入来源':{'本次原答复':'| B位置与列义 | 采用的B来源 |\n|---|---|\n|B列|底稿!I1:I9|\n|C列|底稿!J1:J9|\nJ列为空白的来源保留为空白，不以A旧值替代。'}}
        raw=('将A的 A1:C7 按名单扩展排布。\n'
             '| B表位置 | 来源路径（B） | A采用位置 |\n|---|---|---|\n'
             '| B A20 | 2025：来源!B20；2024：C20 | A A2组甲；B2（2025）、C2（2024） |\n'
             '| B A30 | 2025：来源!B30；2024：C30 | B5（2025）、C5（2024） |\n'
             '| B A21 | 甲 | |\n| B A22 | 丙 | |\n| B A31 | 丁 | |\n| B A32 | 戊 | |\n'
             '- 甲段：组甲之后列B成员：甲、丙。随后单列保留A旧成员“旧乙”及其原有内容。\n'
             '- 乙段：组乙之后列B成员：丁、戊。')
        tables=_declared_list_tables(a,raw,'selection',[{'sheet':'来源','range':'A20:C32'}],b)
        self.assertEqual([t['source_rows'] for t in tables],[[21,22],[31,32]])
        self.assertEqual([t['old_detail_count'] for t in tables],[2,1])
        self.assertEqual(tables[0]['retained_members'][0]['cells']['B'],{'source':'原表!B4'})
        self.assertEqual(tables[0]['row_sources'],{'21':{'B':'底稿!I1','C':'底稿!J1'}})
        self.assertEqual(tables[0]['blank_values'],{'21':{'C':{'value':None}}})
        self.assertEqual([t['operations'][0]['count'] for t in tables],[1,1])


    def test_同表列限定与仅行号范围沿用答复已声明列(self):
        from 附注更新.A合集 import direct_answer
        from 附注更新.A语义阅读 import addresses
        standard={'metric':{m:{'dimension_refs':[{'id':'period'},{'id':'selection'}]} for m in ('metric.one','metric.two')},
                  'dimension':{'period':{'value_type':'period'},'selection':{'value_type':'selection'}}}
        raw=('| B位置与列义 | 金标准编号 | period | selection |\n|---|---|---|---|\n'
             '| 来源!B7:E9 中 B列甲 | metric.one | 2025年度 | 按所在行的范围，见下方逐段说明 |\n'
             '| 同表 C列乙 | metric.two | 2025年度 | 按所在行的范围，见下方逐段说明 |\n'
             '| 同表 D列甲 | metric.one | 2024年度 | 按所在行的范围，见下方逐段说明 |\n'
             '| 同表 E列乙 | metric.two | 2024年度 | 按所在行的范围，见下方逐段说明 |\n'
             '### `selection`及预留位置\n'
             '- 来源!7:8：明细，依次为甲、乙组合；丙。各行分别选取该名称对应的成员。\n'
             '- 来源!9：明细，依次为**丁**；“丁”保留原限定。')
        result=direct_answer(raw,standard,['来源'])
        found={m:set() for m in standard['metric']}
        for row in result['metrics']:
            for span in row['ranges']:found[row['id']].update(addresses(span.split('!',1)[-1]))
        self.assertEqual(found['metric.one'],{f'{c}{r}' for c in 'BD' for r in range(7,10)})
        self.assertEqual(found['metric.two'],{f'{c}{r}' for c in 'CE' for r in range(7,10)})
        periods={cell:r['value'] for r in result['dimensions'] if r['id']=='period'
                 for span in r['ranges'] for cell in addresses(span.split('!',1)[-1])}
        self.assertEqual(periods,{f'{c}{r}':('2025年度' if c in 'BC' else '2024年度') for c in 'BCDE' for r in range(7,10)})
        selections={cell:r['value'] for r in result['dimensions'] if r['id']=='selection'
                    for span in r['ranges'] if span!='*' for cell in addresses(span.split('!',1)[-1])}
        self.assertEqual(selections,{f'{c}{r}':{7:'甲、乙组合',8:'丙',9:'丁'}[r] for c in 'BCDE' for r in range(7,10)})

    def test_来源表中按明示期间分组且不把公式下层当作采用来源(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'slot_bindings':[],
           '机械记录':{'sheets':[{'name':'原表','merges':[]}]}}
        b={'source_hash':'b','机械记录':{'sheets':[{'name':name,'merges':[]} for name in ('展示','底稿')]},
           'b_source_ranges':[{'sheet':'展示'},{'sheet':'底稿'}]}
        raw=('| 业务范围 | B表位置及实际范围 | 来源路径（B） | A版式采用位置及处理 |\n'
             '|---|---|---|---|\n'
             '| 范围甲 | B A20 | 2025：展示!B20/C20 → 底稿!I30；2024：D20/E20 → J30 | '
             'A B7；分别放A C7/D7（2025）、E7/F7（2024） |')
        result=source_relation(a,b,raw,[f'原表!{c}7' for c in 'CDEF'])
        self.assertEqual(result['bindings'],[
            {'target':f'原表!{target}7','source':f'展示!{source}20'}
            for target,source in zip('CDEF','BCDE')])

    def test_同一指标并列列范围及年度角色完整接入(self):
        from 附注更新.A合集 import direct_answer
        from 附注更新.A语义阅读 import addresses
        standard={'metric':{m:{'dimension_refs':[{'id':'period'}]} for m in ('metric.income','metric.cost')},
                  'dimension':{'period':{'value_type':'period'}}}
        raw=('范围：原表!B6:F9。\n'
             'period及template_period_role：列C、D为本期，实际期间为2025年度，角色为current_period；列E、F为上期，实际期间为2024年度，角色为prior_period。\n'
             '- C列（C7:C9）、E列（E7:E9）：metric.income。\n'
             '- D列（D7:D9）、F列（F7:F9）：metric.cost。')
        result=direct_answer(raw,standard,['原表'])
        found={m:set() for m in standard['metric']}
        for row in result['metrics']:
            for span in row['ranges']:found[row['id']].update(addresses(span.split('!',1)[-1]))
        self.assertEqual(found['metric.income'],{f'{c}{r}' for c in 'CE' for r in range(7,10)})
        self.assertEqual(found['metric.cost'],{f'{c}{r}' for c in 'DF' for r in range(7,10)})
        values={r['ranges'][0].split('!',1)[-1]:(r['id'],r['value']) for r in result['dimensions'] if r['id']=='period'}
        self.assertTrue(any(r['value']=='2025年度' for r in result['dimensions'] if r['id']=='period'))
        self.assertTrue(any(r['value']=='2024年度' for r in result['dimensions'] if r['id']=='period'))

    def test_行标签说明连接答复明列财务位置且不扩展分母(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.amount':{'dimension_refs':[{'id':'selection'}]},
                           'metric.ratio':{'dimension_refs':[{'id':'selection'},{'id':'denominator'}]}},
                  'dimension':{k:{'value_type':'selection'} for k in ('selection','denominator')}}
        raw=('范围：原表!B7:D8。B列是类别，不是财务格；本次财务格为C7:D8。\n'
             'C列金额：metric.amount\nD列比例：metric.ratio\n'
             '- B7“父项”：selection 为父项全部。C7、D7 均为空白财务格，分别对应上述 C 至 D 列指标。D7 的 denominator 为本表合计。\n'
             '- B8“子项”：selection 为其中子类。C8、D8 分别对应上述 C 至 D 列指标。D8 的 denominator 为本表合计。')
        result=direct_answer(raw,standard,['原表'])
        selected={r['value']:r['ranges'] for r in result['dimensions'] if r['id']=='selection'}
        self.assertEqual(selected,{'父项全部':['原表!C7','原表!D7'],'其中子类':['原表!C8','原表!D8']})
        self.assertEqual([r['ranges'] for r in result['dimensions'] if r['id']=='denominator'],[['原表!D7'],['原表!D8']])

    def test_实际时点说明保留日期本身(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.amount':{'dimension_refs':[{'id':'period'}]}},'dimension':{'period':{'value_type':'period'}}}
        raw='原表!C7：metric.amount；period：实际时点为 **2024年12月31日**，template_period_role 为 **期初**。'
        result=direct_answer(raw,standard,['原表'])
        self.assertEqual([r['value'] for r in result['dimensions'] if r['id']=='period'],['2024-12-31'])

    def test_名单行表复用原答复并按明确名称保留旧值(self):
        from 附注更新.A合集 import _declared_list_tables
        from copy import deepcopy
        a={'facts':[],'slot_bindings':[{'sheet':'原表','cell':f'{col}{row}','partial_semantics':
            {'dimensions':{'selection':{'mode':'members','members':[name]}}}}
            for row,name in [(7,'甲'),(8,'旧乙')] for col in 'CD'],
            '机械记录':{'names':{'名单':[{'sheet':'原表','range':'B6:D9'}]}},
            'outline':{'subtables':[{'sheet':'原表','metrics':[
                {'ranges':['C7:C8'],'template_period_role':'current'},
                {'ranges':['D7:D8'],'template_period_role':'prior'}]}]}}
        sources=[{'sheet':'来源','range':'B20:J25'},{'sheet':'展示','range':'A40:C50'}]
        prior=('名称取 B 列、本期金额取 I 列。\n'
            '| B 来源行 | 搬至 A 行 | 成员 |\n|---|---|---|\n'
            '| 20 | 7 | 甲 |\n| 21 | 8 | 新丙 |\n| 22 | 9 | 新丁 |\n'
            '新增成员无对应旧值的保持空白。')
        b={'已确认关系':[{'kind':'名单编排','dimension':'selection','source_ranges':sources,'content':prior}],
            '合集来源':[{'识别答复':'A原明细为 原表!B7:B8，合计在 B9。'}]}
        reply=('名称相同的1项：原A上期值随同名成员保留在D列：甲→A7。\n'
            '无法确认同义的1项：旧乙，不与新丙合并。将它作为单列保留项接在3个B成员之后，依次放在A10:A10；名称和原A的C、D内容仍与原项目配对。\n'
            '沿用A第8行样式。合计移至A11。')
        before=deepcopy((a,b));tables=_declared_list_tables(a,reply,'selection',sources,b)
        self.assertEqual((a,b),before);self.assertEqual(len(tables),1)
        table=tables[0]
        self.assertEqual(table['source_rows'],[20,21,22]);self.assertEqual(table['columns'],{'C':'I'})
        self.assertEqual(table['labels'],{'B':'B'})
        self.assertEqual(table['blank_values']['20']['D'],{'source':'原表!D7'})
        self.assertEqual(table['blank_values']['21']['D'],{'value':None})
        self.assertEqual(table['retained_members'],[{'index':3,'cells':
            {'B':{'source':'原表!B8'},'C':{'source':'原表!C8'},'D':{'source':'原表!D8'}}}])
        self.assertEqual(table['operations'],[{'axis':'row','action':'insert','index':3,'count':2,'template_index':2}])

    def test_成员名称顺序及对应列引用保留到维度(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.amount':{'dimension_refs':[{'id':'selection'}]}},
                  'dimension':{'selection':{'value_type':'selection'}}}
        raw=('| 位置 | 原文对象及范围 | 金标准编号与维度取值 |\n|---|---|---|\n'
             '| C7:C8 | 各明细行，依次为甲、乙。各行仅指各自名称。 | metric.amount；`selection=members`，成员分别为上述各行名称。 |\n'
             '| D7:D8 | 与C列相同行名对应。 | metric.amount；`selection=members`，成员与对应C列相同。 |\n'
             '| E7:E8 | 未说明成员。 | metric.amount；selection=members。 |')
        result=direct_answer(raw,standard,['原表'])
        values={cell:row['value'] for row in result['dimensions'] if row['id']=='selection' for cell in row['ranges']}
        for cell,name in [('C7','甲'),('C8','乙'),('D7','甲'),('D8','乙')]:
            self.assertEqual(values['原表!'+cell],{'mode':'members','members':[name]})
        self.assertEqual(values['原表!E7:E8'],'members')

    def test_已声明财务范围内的列编号不扩成全部列(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{key:{'dimension_refs':[]} for key in ('metric.balance','metric.ratio')},'dimension':{}}
        raw=('本次财务格为 原表!C7:D8。\n'
             '- C列“金额”：metric.balance\n- D列“比例”：metric.ratio\n')
        result=direct_answer(raw,standard,['原表'])
        self.assertEqual({r['id']:r['ranges'] for r in result['metrics']},
            {'metric.balance':['原表!C7','原表!C8'],'metric.ratio':['原表!D7','原表!D8']})

    def test_行字段使用本行明列的数据位置而非上一行(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{key:{'dimension_refs':[{'id':'item'}]} for key in ('metric.balance','metric.ratio')},
                  'dimension':{'item':{'value_type':'text'}}}
        raw=('## 原表!B5:D9\n- C列“金额”：metric.balance\n- D列“比例”：metric.ratio\n'
             '| 行及B列原文 | item 实际范围 | C、D格 |\n|---|---|---|\n'
             '| 7“父项” | 父项全部 | C7、D7 分别对应上述两列指标。 |\n'
             '| 8“其中项” | 其中的子项 | C8、D8 为空白，仍分别对应上述两列指标。 |')
        result=direct_answer(raw,standard,['原表'])
        self.assertEqual({r['id']:r['ranges'] for r in result['metrics']},
            {'metric.balance':['原表!C7','原表!C8'],'metric.ratio':['原表!D7','原表!D8']})
        self.assertEqual([(r['value'],r['ranges']) for r in result['dimensions'] if r['id']=='item'],
            [('父项全部',['原表!C7','原表!D7']),('其中的子项',['原表!C8','原表!D8'])])

    def test_来源载体后的中英文冒号不丢失明确对应(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'gold','facts':[],
           '机械记录':{'sheets':[{'name':'原表','merges':[]}]}}
        b={'source_hash':'b','机械记录':{'sheets':[{'name':'来源','merges':[]}]}}
        raw=('| A需要的位置及字段 | 选定B来源位置 |\n|---|---|\n'
             '| A：原表!C3 | B：来源!B7 |\n| A：原表!D3 | B:来源!C7 |\n'
             '| A：原表!C4:C5 | B来源：来源!B8:B9 |')
        result=source_relation(a,b,raw,['原表!C3','原表!D3','原表!C4','原表!C5'])
        self.assertEqual(result['bindings'],[{'target':t,'source':s} for t,s in
            [('原表!C3','来源!B7'),('原表!D3','来源!C7'),('原表!C4','来源!B8'),('原表!C5','来源!B9')]])
        self.assertEqual(result['content'],raw)
        dates='| A需要坐标 | B来源坐标 |\n|---|---|\n| 原表!C6；D6 | 来源!B9（2025-12-31）；C9（2024-12-31） |'
        self.assertEqual(source_relation(a,b,dates,['原表!C6','原表!D6'])['bindings'],
                         [{'target':'原表!C6','source':'来源!B9'},{'target':'原表!D6','source':'来源!C9'}])

    def test_明示行关系和共用列关系机械展开且合并标签沿用原引用(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'gold','facts':[],
           '机械记录':{'sheets':[{'name':'原表','merges':['C2:D2']}]}}
        b={'source_hash':'b','机械记录':{'sheets':[{'name':'来源','merges':[]}]}}
        raw=('A 原表!C2:D2（合并表头“期末余额”）对应 2024年12月31日。\n'
             'A 原表!B3“甲” ↔ B 来源!A7“甲”\nA 原表!B4“乙” ↔ B 来源!A8“乙”\n'
             '上述每一行均沿用已确定的列对应：B 来源!B金额、C比例 ↔ A 原表!C、D。')
        targets=['原表!C3','原表!D3','原表!C4','原表!D4'];result=source_relation(a,b,raw,targets)
        self.assertEqual({r['target']:r['source'] for r in result['bindings']},
            {'原表!C3':'来源!B7','原表!D3':'来源!C7','原表!C4':'来源!B8','原表!D4':'来源!C8'})
        self.assertEqual([(r['target'],r['value']) for r in result['label_actions']],
                         [('原表!C2','期末余额（2024年12月31日）')])
        self.assertEqual(source_relation(a,b,raw.split('上述每一行')[0],targets)['bindings'],[])

    def test_B字段补充沿用直接公式材料不展开金额公式下层(self):
        from 附注更新.A语义阅读 import recognize_a,read_workbook,source_material
        engine=UnifiedSemanticEngine(self.config,self.gold,self.folder/'字段材料')
        snapshot=read_workbook(self.folder/'B.xlsx');snapshot['sheets'][0]['cells']['B3']['formula']='=1'
        a={'gold_hash':engine.gold_hash,'机械记录':read_workbook(self.folder/'A.xlsx'),'source_context':[]}
        a['机械记录']['names']={'表区1':[{'sheet':'原表','range':'A1:C3'}]}
        b={'recognition_method':'direct','source_hash':file_hash(self.folder/'B.xlsx'),
           'gold_hash':engine.gold_hash,'source_context':[],'金标准原文':'','机械记录':snapshot}
        metadata={'carrier':'B','normal_entry':True,'a_required_ranges':[{'sheet':'原表','range':'B3:C3'}],
                  'b_source_ranges':[{'sheet':'原表','range':'A2:C3'}],
                  'field_scope':{'原表!B3':['metric_id','period']}}
        previous=self.folder/'字段已有合集.json'
        previous.write_text(json.dumps(b,ensure_ascii=False),encoding='utf-8')
        with patch('附注更新.A语义阅读.source_material',wraps=source_material) as material, \
             patch('附注更新.A语义阅读.direct_reference',return_value={'需要':[
                 {'location':'原表!B3','metric_id':'固定资产原值','dimensions':{}},
                 {'location':'原表!C3','metric_id':'固定资产原值','dimensions':{}}],
                 'unresolved':[{'sheet':'原表','cell':'C3','reason':'未提供标准编号'}],
                 'source_identities':[{'source_reference':{'location':'原表!B3'},'原文':'已识别的含义'}]}), \
             patch('附注更新.统一勾稽.load_fact_mapping',return_value=b), \
             patch('附注更新.A合集.save_direct_answer',return_value={}), \
             patch.object(engine,'request_text',return_value='已有答复') as request:
            recognize_a(engine,self.folder/'B.xlsx',context=[],target_mapping=a,source_snapshot=snapshot,
                        previous_result=previous,artifact_metadata=metadata)
        self.assertEqual(material.call_args_list[0].args[0]['semantic_formula_sources'],4)
        self.assertNotIn('A原表.json',request.call_args.args[1])
        self.assertNotIn('A原表网格.txt',request.call_args.args[1])
        reference=request.call_args.args[1]['A需要.json']
        self.assertEqual([row['location'] for row in reference['需要']],['原表!B3','原表!C3'])
        self.assertTrue(all('locations' not in row for row in reference['需要']))
        self.assertNotIn('unresolved',reference)
        self.assertNotIn('source_identities',reference)

    def test_冒号行号和角色前置表头保持各自字段与期间(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.amount':{'dimension_refs':[{'id':'period'},{'id':'member'}]}},
                  'dimension':{'period':{'value_type':'period'},'member':{'value_type':'selection','domain':'对象'}}}
        raw=('C列各行（含合计）为期末余额，实际日期 **2025年12月31日**；D列各行（含合计）为期初余额，实际日期 **2024年12月31日**。\n'
             '| 行及原文 | 期末：C金额 | 期初：D金额 |\n|---|---|---|\n'
             '| 192：甲 | `metric.amount`；member=甲 | `metric.amount`；member=甲 |\n'
             '| 193：乙 | `metric.amount`；member=乙 | `metric.amount`；member=乙 |')
        answer=direct_answer(raw,standard,['原表'])
        members={loc:r['value'] for r in answer['dimensions'] if r['id']=='member' for loc in r['ranges']}
        self.assertEqual(members,{'原表!C192':'甲','原表!D192':'甲','原表!C193':'乙','原表!D193':'乙'})
        periods={loc:r['value'] for r in answer['dimensions'] if r['id']=='period' for loc in r['ranges']}
        for row in ('192','193'):
            self.assertEqual(periods['原表!C'+row],'2025-12-31')
            self.assertEqual(periods['原表!D'+row],'2024-12-31')


    def test_期间标签直接采用答复明列文字且不重建金额身份(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'gold','facts':[],
           '机械记录':{'sheets':[{'name':'原表','merges':[]}]}}
        b={'source_hash':'b','机械记录':{'sheets':[{'name':'来源','merges':[]}]}}
        raw=('**A原表 原表 第2行表头：**将B2“期末余额”改为“期末余额（某日期）”，将C2“期初余额”改为“期初余额（另一日期）”。\n'
             '**A原表 原表 第8行表头：**将B8、C8、D8、E8分别明确为“原文甲”“原文乙”“原文丙”“原文丁”。')
        relation=source_relation(a,b,raw,['原表!B3'])
        self.assertEqual(relation['bindings'],[])
        self.assertEqual({v['target']:v['value'] for v in relation['label_actions']},
            {'原表!B2':'期末余额（某日期）','原表!C2':'期初余额（另一日期）',
             '原表!B8':'原文甲','原表!C8':'原文乙','原表!D8':'原文丙','原表!E8':'原文丁'})
        self.assertEqual(relation['content'],raw)

    def test_旧定向来源入口只识别各格且保留完整原文(self):
        from 附注更新.A语义阅读 import recognize_a,read_workbook
        engine=UnifiedSemanticEngine(self.config,self.gold,self.folder/'关系材料')
        a={'gold_hash':engine.gold_hash,'机械记录':read_workbook(self.folder/'A.xlsx'),
           'source_context':['表外共同期间原文，不能在定向来源入口丢掉。',
                             {'title':'Word正文','text':'保存为原文记录的完整正文同样保留。'}]}
        b={'recognition_method':'direct','source_hash':file_hash(self.folder/'B.xlsx'),
           'gold_hash':engine.gold_hash,'facts':[],'slot_bindings':[],'机械记录':read_workbook(self.folder/'B.xlsx')}
        metadata={'carrier':'B','source_targets':['原表!B3'],
            'source_candidates':{'原表!B3':['原表!B3']},
            'a_required_ranges':[{'sheet':'原表','range':'A2:C3'}],
            'b_source_ranges':[{'sheet':'原表','range':'A2:C3'}],
            'field_scope':{'原表!B3':['period']}}
        previous=self.folder/'来源已有合集.json'
        previous.write_text(json.dumps(b,ensure_ascii=False),encoding='utf-8')
        def answer(prompt,materials):
            self.assertIn('各格分别是什么意思',prompt)
            self.assertNotIn('本次采用的必要来源及其他未采用候选',prompt)
            self.assertNotIn('A原表.json',materials)
            self.assertNotIn('A原表网格.txt',materials)
            self.assertNotIn('已保存的当前候选对应',materials)
            self.assertNotIn('本次需要确定来源的A位置',materials)
            self.assertIn('A需要.json',materials)
            return '原答复'
        with patch('附注更新.A语义阅读.direct_reference',return_value={}), \
             patch('附注更新.统一勾稽.load_fact_mapping',return_value=b), \
             patch('附注更新.A合集.save_direct_answer',return_value={}) as saved, \
             patch('附注更新.A合集.save_source_relations') as source_only, \
             patch.object(engine,'request_text',side_effect=answer) as asked:
            recognize_a(engine,self.folder/'B.xlsx',context=[],target_mapping=a,
                        previous_result=previous,artifact_metadata=metadata)
        self.assertEqual(asked.call_count,1)
        self.assertEqual(saved.call_args.args[0]['b_source_ranges'],metadata['b_source_ranges'])
        self.assertEqual(saved.call_args.args[0]['target_tables'],[])
        self.assertEqual(saved.call_args.kwargs['field_scope'],metadata['field_scope'])
        self.assertIs(saved.call_args.kwargs['target_mapping'],a)
        source_only.assert_not_called()
        self.assertNotIn('b_source_ranges',b)

    def test_正常设置回读保留定向范围与补充要求(self):
        cfg={**self.config,'a_required_ranges':[{'sheet':'原表','range':'A2:C3'}],
            'b_source_ranges':[{'sheet':'来源','range':'A2:D8'}],'b_source_targets':['原表!B3'],
            'b_allowed_metrics':['固定资产原值'],'b_additional_context':'只处理上述需要，复用其余结果。',
            'a_field_scope':{'原表!B3':['period']},'a_withdrawn':['旧解释'],
            'b_arranged_targets':False,'list_dimension':'测试名单'}
        path=self.folder/'定向设置.json';save_settings(cfg,path);loaded=load_settings(path)
        for key in ('a_required_ranges','b_source_ranges','b_source_targets','b_allowed_metrics',
                    'b_additional_context','a_field_scope','a_withdrawn','b_arranged_targets','list_dimension'):
            self.assertEqual(loaded[key],cfg[key],key)

    def test_来源答复明列名单行时沿用已识别标签位置(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'gold','facts':[],
            'excluded':[{'sheet':'原表','cell':cell,'category':'label'} for cell in ('A3','A4')],
            '机械记录':{'sheets':[{'name':'原表','merges':[]}]}}
        b={'source_hash':'b','gold_hash':'gold','机械记录':{'sheets':[{'name':'来源','merges':[]}]}}
        raw='A明细行按B有名称的两个成员排列：`新甲`、`新乙`；对应A行3—4。'
        result=source_relation(a,b,raw,['原表!B3','原表!B4'])
        self.assertEqual({r['target']:r['value'] for r in result['label_actions']},{'原表!A3':'新甲','原表!A4':'新乙'})
        a['excluded'].append({'sheet':'原表','cell':'C4','category':'label'})
        result=source_relation(a,b,raw,['原表!B3','原表!B4'])
        self.assertEqual([r['target'] for r in result['label_actions']],['原表!A3'])
        self.assertEqual(result['content'],raw)

    def test_粗体指标段落不撤销下级范围的共同字段(self):
        from 附注更新.A合集 import direct_answer,_answer_tables
        standard={'metric':{'metric.amount':{'dimension_refs':[{'id':'member'},{'id':'currency'}]}},
            'dimension':{'member':{'value_type':'selection','domain':'成员'},'currency':{'value_type':'text'}}}
        raw=('### 候选识别\n'
             '- **`metric.amount`**（余额）：currency=CNY。\n'
             '  - `原表!C3`：member=甲。\n'
             '  - `原表!C4`：member=乙。')
        answer=direct_answer(raw,standard,['原表'])
        snapshot={'sheets':[{'name':'原表','max_row':4,'max_column':3}]}
        saved=_answer_tables(answer,{'carrier':'B','答复指纹':'test'},snapshot)[0]
        self.assertTrue(any(r['id']=='metric.amount' and {'C3','C4'}<=set(r['ranges']) for r in saved['metrics']))
        self.assertTrue(any(r['id']=='currency' and r['value']=='CNY' for r in saved['dimensions']))

    def test_范围段落简写与旁述标签不混入金额身份(self):
        from 附注更新.A合集 import direct_answer,_answer_tables
        standard={'metric':{'metric.amount':{'dimension_refs':[{'id':'member'}]}},
            'dimension':{'member':{'value_type':'selection','domain':'成员'},
                         'currency':{'value_type':'text','domain':'币种'},'unit_scale':{'value_type':'number'}}}
        raw=('- **共同范围**\n- 币种及倍率：人民币元，`unit_scale=1`。\n'
             '- **逐行对应**（`member`保留各行的实际范围）\n'
             '  - `C3:D3`，B3“第一类”：`metric.amount`；member=第一类对象。C为期末，D为期初。\n'
             '  - `C4:D4`，B4“第二类”：`metric.amount`；范围为第二类对象。C为期末，D为期初。\n'
             '  - `C5:D5`，B5“小计”：`metric.amount`；对象为上述两类小计，不扩大为全体。C为期末，D为期初。\n'
             '- **另一段**\n- `C9`：`metric.amount`；范围为尚未指明维度的内容。\n'
             '- `B10`“原文”：`metric.amount`；member=明确目标。')
        answer=direct_answer(raw,standard,['原表'])
        self.assertEqual({r['ranges'][0]:r['value'] for r in answer['dimensions'] if r['id']=='member'},
            {'原表!C3:D3':'第一类对象','原表!C4:D4':'第二类对象',
             '原表!C5:D5':'上述两类小计，不扩大为全体','原表!B10':'明确目标'})
        self.assertTrue(any(r['id']=='currency' and r['value']=='人民币元' for r in answer['dimensions']))
        self.assertFalse(any('原表!B3' in r['ranges'] for r in answer['metrics']))
        self.assertTrue(any('原表!B10' in r['ranges'] for r in answer['metrics']))
        table={'id':'原表','sheet':'原表','range':'B2:D10','data_ranges':['B3:D10']}
        for mapping in ({'target_table':table},{'target_tables':[table]}):
            saved=_answer_tables(answer,mapping,{})[0]
            self.assertEqual({r['range'] for r in saved['non_business']},{'B3','B4','B5'})
            self.assertEqual(saved['data_ranges'],['B3:D10'])

    def test_名单分段公式取值沿用明确来源引用(self):
        from 附注更新.统一勾稽 import _calculation_formulas
        mapping={'facts':[],'机械记录':{'sheets':[{'name':'来源','cells':{cell:{'formula':'=1'} for cell in ('I4','I14','J31','A4','I41','I81','K99')}}]},
            '已确认关系':[{'kind':'名单编排','tables':[{'source_sheet':'来源','source_rows':[4],
                'columns':{'B':'I','C':'I'},'column_row_offsets':{'C':10},'labels':{'A':'A'},
                'row_sources':{'4':{'B':'来源!J31'},'8':{'B':'来源!I41','C':'来源!I81'}},'totals':[{'source_row':8}]}]}]}
        self.assertEqual(set(_calculation_formulas(mapping)),{'来源!J31','来源!I14','来源!A4','来源!I41','来源!I81'})


    def test_移除表头引用后的空引号不吞入期间解释(self):
        from 附注更新.A合集 import direct_answer,_direct_table
        standard={'metric':{'metric.cash':{'dimension_refs':[{'id':'period'}]}},
            'dimension':{'period':{'value_type':'period'}}}
        raw=('所选槽位 metric.cash，数据格为C3:D4。\n'
             '- `原表!C2` 原文“期末余额”：适用于 `C3:C4`，`template_period_role=期末`（本期期末），实际日期为 **2025年12月31日**。\n'
             '- `原表!D2` 原文“期初余额”：适用于 `D3:D4`，`template_period_role=期初`，实际日期为 **2024年12月31日**（上期期末）。')
        answer=direct_answer(raw,standard,['原表'])
        roles=[r for r in answer['dimensions'] if r['id']=='template_period_role']
        self.assertTrue(roles)
        self.assertLessEqual({r['value'] for r in roles},{'closing','opening','期末','期初'})
        self.assertTrue(all(r['value'] in ('closing','期末') for r in roles if r['ranges']==['原表!C3:C4']))
        self.assertTrue(all(r['value'] in ('opening','期初') for r in roles if r['ranges']==['原表!D3:D4']))
        periods={r['ranges'][0]:r['value'] for r in answer['dimensions'] if r['id']=='period'}
        self.assertEqual(periods['原表!C3:C4'],'2025-12-31')
        self.assertEqual(periods['原表!D3:D4'],'2024-12-31')
        table=_direct_table({'sheet':'原表','data_ranges':['C3:D4'],
            'metrics':[{'id':'metric.cash','ranges':['C3:D4']}],'dimensions':roles},{'source_context':[]},standard)
        self.assertFalse(table.get('unresolved'))
        self.assertEqual({cell:m['template_period_role'] for m in table['metrics'] for cell in m['ranges']},
                         {'C3':'closing','C4':'closing','D3':'opening','D4':'opening'})



    def test_共同说明不继承槽位且重复引用不阻断合集(self):
        from copy import deepcopy
        from 附注更新.A合集 import _answer_tables, combine_direct_collections, direct_answer, _records
        standard={'metric':{'metric.cash':{'dimension_refs':[{'id':'item'}]}},
                  'dimension':{'item':{'value_type':'text'},'entity':{'value_type':'text'}}}
        raw=('## 第一段\n来源!A1：entity=前段说明。\n## 第二段\n'
             '所选槽位：metric.cash。\n- **entity**：本段主体。\n'
             '来源!B3：item=现金。\n共同说明在 来源!A1，表头位置为 来源!B2。')
        answer=direct_answer(raw,standard,['来源'])
        metrics=[r for r in answer['metrics'] if r['id']=='metric.cash']
        self.assertEqual({p for r in metrics for p in r['ranges']},{'来源!B3'})
        entity=next(r for r in answer['dimensions'] if r['value']=='本段主体')
        self.assertEqual(entity['ranges'],['*'])
        snapshot={'sheets':[{'name':'来源','max_column':3,'max_row':3,'cells':{},'merges':[]}]}
        declared={'metrics':[{'id':'metric.cash','ranges':['*']}],
                  'dimensions':[{'id':'item','ranges':['来源!B3'],'value':'现金'}],
                  'unresolved':[{'ranges':['来源!A1'],'reason':'共同说明','original':'原引用'}]}
        table=_answer_tables(declared,{'carrier':'B','答复指纹':'first'},snapshot)[0]
        self.assertEqual(table['metrics'][0]['ranges'],['B3'])
        common={'carrier':'B','source_hash':'source','gold_hash':'gold','schema_version':2,
                'recognition_method':'direct','collection_version':2,'机械记录':snapshot,
                'facts':[],'owners':{},'完整性':{},'source_identities':[],
                '封存定义原文':[],'source_context':[],'coverage':{'financial_cells':1}}
        parts=[]
        for number,cell in enumerate(('B3','C3')):
            parts.append({**deepcopy(common),'识别答复':'原答复'+str(number),
                'outline':{'subtables':[{**deepcopy(table),'id':str(number)}]},
                'slot_bindings':[{'sheet':'来源','cell':cell,
                    'partial_semantics':{'metric_id':'metric.cash','dimensions':{'item':'现金'}}}],
                'unresolved':[{'sheet':'来源','cell':'A1','original':'原引用'+str(number)}]})
        before=deepcopy(parts);combined=combine_direct_collections(parts)
        self.assertEqual(parts,before)
        self.assertEqual(set(dict(_records(combined))),{'来源!B3','来源!C3'})
        self.assertEqual([r['original'] for r in combined['unresolved']],['原引用0','原引用1'])
        self.assertEqual(len(combined['合集来源']),2)

    def test_B工作表引号不会吞掉已有身份及来源对应(self):
        from 附注更新.A合集 import _own_answer_text, source_relation
        snapshot={'sheets':[{'name':'word附注'}]}
        mapping={'carrier':'B','source_hash':'B','机械记录':snapshot,'target_tables':[{'sheet':'word附注'}],
                 'A需要':{'需要':[{'location':'原表!C5'}]}}
        raw=('所选槽位 metric.cash。\n|A需要位置及角色|采用的B来源位置|item|\n'
             '|---|---|---|\n|原表!C5（期末）|「word附注」!B4|现金|\n'
             '本次选用 B「word附注」!B4；《word附注》!C4 为期初。\nB《word附注》`B4:C4` 已列明。')
        own=_own_answer_text(raw,mapping,snapshot)
        self.assertIn('|[A位置]|word附注!B4|现金|',own.replace(' ',''))
        self.assertIn('本次选用 word附注!B4',own)
        self.assertIn('word附注!C4 为期初',own)
        self.assertIn('word附注!B4:C4 已列明',own)
        target={'facts':[],'机械记录':{'sheets':[{'name':'原表'}]},'source_hash':'A','gold_hash':'gold'}
        relation=source_relation(target,mapping,raw,['原表!C5'])
        self.assertEqual(relation['bindings'],[{'target':'原表!C5','source':'word附注!B4'}])

    def test_B答复中的A简写位置不能归入B来源表(self):
        from 附注更新.A合集 import _own_answer_text
        mapping={'carrier':'B','target_tables':[{'sheet':'来源表'}],
                 'A需要':{'需要':[{'location':'原表!D219'}]}}
        raw=('A D219:D223、F219:F223 采用 来源表!B3:B7。\n'
             'A 原表!C219:C223、E219:E223 采用 来源表!C3:C7。\n'
             'A C224:D224 没有来源，metric.cash。\n'
             'A需要中的C324/C325对应B《来源表》B3。')
        result=_own_answer_text(raw,mapping,{'sheets':[{'name':'来源表'}]})
        self.assertNotIn('219',result);self.assertNotIn('224',result)
        self.assertNotIn('324',result);self.assertNotIn('325',result)
        self.assertIn('来源表!B3:B7',result);self.assertIn('来源表!C3:C7',result)

    def test_A说明之后同段的B明确内容仍保留(self):
        from 附注更新.A合集 import _own_answer_text
        mapping={'carrier':'B','target_tables':[{'sheet':'来源表'}],
                 'A需要':{'需要':[{'location':'原表!C3'}]}}
        raw=('A原标签在原表!C2，覆盖C3。A 的期初日期为2024-12-31。'
             'B 的期初角色对应日期为2025-01-01。')
        own=_own_answer_text(raw,mapping,{'sheets':[{'name':'来源表'}]})
        self.assertNotIn('2024-12-31',own);self.assertNotIn('C3',own)
        self.assertIn('B 的期初角色对应日期为2025-01-01',own)

    def test_明示未知及币种后的说明不冒充字段取值(self):
        from 附注更新.A合集 import _direct_table
        table={'sheet':'来源','data_ranges':['B4'],'metrics':[],'dimensions':[
            {'id':'report_scope','ranges':['B4'],'value':'原文未明确单体还是合并，保留未知'},
            {'id':'currency','ranges':['B4'],'value':'CNY；金额单位为人民币元'}]}
        standard={'dimension':{'report_scope':{'value_type':'text'},'currency':{'value_type':'text'}}}
        rows=_direct_table(table,{'source_context':[]},standard)['dimensions']
        self.assertEqual({r['id']:r['value'] for r in rows},{'report_scope':'未知','currency':'CNY'})
        self.assertTrue(all('original' in row for row in rows))

    def test_B表名共同声明与表内箭头保留来源而不混入A位置(self):
        from 附注更新.A合集 import _own_answer_text,direct_answer,source_relation
        mapping={'carrier':'B','target_tables':[{'sheet':'来源表'},{'sheet':'另表'}],
                 'A需要':{'需要':[{'location':'原表!C342'}]}}
        snapshot={'sheets':[{'name':'来源表'},{'name':'另表'}]}
        raw=('以下 B 坐标均属于 `来源表` 工作表。\n'
             '|A原表位置及需要|金标准槽位；实际范围|采用的B来源及期间角色|\n|---|---|---|\n'
             '|原表!C342、F342|metric.cash；item=all|I467 年初数 → C342；I497 年末数 → F342|')
        own=_own_answer_text(raw,mapping,snapshot)
        self.assertNotIn('342',own)
        self.assertIn('来源表!I467',own);self.assertIn('来源表!I497',own)
        standard={'metric':{'metric.cash':{'dimension_refs':[{'id':'item'}]}},'dimension':{'item':{'value_type':'text'}}}
        parsed=direct_answer(own,standard,['来源表','另表'])
        self.assertEqual({loc for row in parsed['metrics'] for loc in row['ranges']},{'来源表!I467','来源表!I497'})
        target={'facts':[],'机械记录':{'sheets':[{'name':'原表'}]},'source_hash':'A','gold_hash':'gold'}
        mapping.update(机械记录=snapshot,source_hash='B')
        relation=source_relation(target,mapping,raw,['原表!C342','原表!F342'])
        self.assertEqual(relation['bindings'],[{'target':'原表!C342','source':'来源表!I467'},
                                               {'target':'原表!F342','source':'来源表!I497'}])

    def test_自然文字列日期及共同标签直接接入(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.cash':{'dimension_refs':[{'id':'item'},{'id':'period'}]}},
                  'dimension':{'item':{'value_type':'text'},'period':{'value_type':'period'}}}
        raw=('金额列标签位于 来源表!B3:C3，分别为期末余额、期初余额。'
             '这些标签适用于明细行和合计行：B列为期末，日期 2024-12-31；C列为期初，日期 2024-01-01。\n'
             '以下各项均使用 metric.cash。\n|位置|item|\n|---|---|\n|来源表!B4|现金|\n|来源表!C4|现金|')
        result=direct_answer(raw,standard,['来源表'])
        for location,date,label in [('来源表!B4','2024-12-31','来源表!B3'),('来源表!C4','2024-01-01','来源表!C3')]:
            rows=[r for r in result['dimensions'] if r['id']=='period' and location in r['ranges']]
            self.assertEqual({r['value'] for r in rows},{date})
            self.assertTrue(any(label in r.get('label_references',[]) for r in rows))

    def test_B分组来源答复及已有候选补充共用接入(self):
        snapshot={'sheets':[{'name':'展示表'},{'name':'汇总表'}]}
        target={'机械记录':{'sheets':[{'name':'原表'}]}}
        standard={'metric':{'metric.cash':{},'metric.other':{}},'dimension':{}}
        original={'范围认识':{'A范围对应的B来源候选':[
            {'A范围':'原表!B3:D5','B来源候选':'现金：展示表!A2:D8、A11:D15；汇总表!B2:F9。'}],
            '表外共同原文引用':[{'引用':'展示表!A1:D1','说明':'表头'}]}}
        supplement='|A范围|相关标准|\n|---|---|\n|原表!B3:D5|metric.cash|'
        groups,unread=read_scope_answer(json.dumps([original,supplement],ensure_ascii=False),snapshot,standard,target)
        self.assertEqual(unread,[])
        self.assertEqual([(g['sheet'],g['range']) for g in groups],[('展示表','A2:D8'),('展示表','A11:D15'),('汇总表','B2:F9')])
        self.assertTrue(all(g['metric_ids']==['metric.cash'] and g['data_ranges']==[] for g in groups))
        self.assertTrue(all(g['context_references']==[{'sheet':'展示表','range':'A1:D1'}] for g in groups))
        without,_=read_scope_answer(json.dumps(original,ensure_ascii=False),snapshot,standard,target)
        self.assertTrue(all(not g['metric_ids'] for g in without))
        listed,_=read_scope_answer(json.dumps([original,'- `原表!B3:D5`：`metric.cash`。'],ensure_ascii=False),snapshot,standard,target)
        self.assertTrue(all(g['metric_ids']==['metric.cash'] for g in listed))

    def test_B只给候选表区时由实际答复列明数据位置(self):
        from 附注更新.A合集 import _answer_tables
        answer={'metrics':[{'id':'metric.cash','ranges':['来源!C4:C5']}],
                'dimensions':[{'id':'currency','ranges':['*'],'value':'CNY'}]}
        mapping={'carrier':'B','答复指纹':'原答复','target_tables':[{'id':'现金','sheet':'来源','range':'A2:A8','data_ranges':[]}]}
        snapshot={'sheets':[{'name':'来源','max_row':8,'max_column':6}]}
        tables=_answer_tables(answer,mapping,snapshot)
        self.assertEqual(tables[0]['data_ranges'],['C4','C5'])
        self.assertEqual(tables[0]['metrics'][0]['ranges'],['C4:C5'])
        self.assertEqual(tables[0]['dimensions'][0]['value'],'CNY')

    def test_B名称列范围仍呈现同行金额和表头原文(self):
        from 附注更新.A语义阅读 import source_material
        from 附注更新.表格 import read_workbook,materialize_candidate_blanks
        path=self.folder/'横向原文.xlsx';book=Workbook();sheet=book.active;sheet.title='来源'
        sheet.append(['项目','期末','期初']);sheet.append(['现金',80,None]);sheet.append(['其他科目',40,20])
        book.save(path);book.close();snapshot=read_workbook(path);materialize_candidate_blanks(snapshot)
        material,_=source_material(snapshot,[],[{'sheet':'来源','range':'A1:A2','data_ranges':[],'read_full_rows':True}])
        raw=material['sheets'][0]['rows']
        self.assertIn('B=期末',raw);self.assertIn('B=80',raw);self.assertIn('C=[空白]',raw)
        self.assertNotIn('其他科目',raw)

    def test_完整横向读取不能复用只有名称列的材料缓存(self):
        from 附注更新.正常入口识别 import scope_key
        old={'sheet':'来源','range':'A1:A3','data_ranges':[]}
        self.assertNotEqual(scope_key(old),scope_key({**old,'read_full_rows':True}))

    def test_编号承接已列范围及字段表保留行列限定(self):
        from 附注更新.A合集 import direct_answer
        from 附注更新.A语义阅读 import addresses
        standard={'metric':{'metric.change':{'dimension_refs':[{'id':'item'},{'id':'direction'}]}},
                  'dimension':{'item':{'value_type':'text'},'direction':{'value_type':'text'}}}
        raw=('**D3:E4：变动金额槽位及维度**\n\n这些格均采用金标准槽位 `metric.change`。\n\n'
             '| 行 | item | D列 direction | E列 direction |\n|---|---|---|---|\n'
             '| 3 | 设备 | 增加 | 减少 |\n| 4 | 其他设备 | 增加 | 减少 |\n\n'
             'E3:E4 原文为空白。\n\n| 类别行 | 适用区段 | item 实际取值 |\n|---|---|---|\n'
             '| 7 | C7:F8、C11:F12 | 房屋 |\n| 8 | 同上 | 设备 |')
        parsed=direct_answer(raw,standard,['原表'],field_scope={'原表!D3:E4':['metric_id','direction'],'原表!C7:F8':['item'],'原表!C11:F12':['item']})
        locations=lambda row:{cell for area in row['ranges'] for cell in addresses(area.rsplit('!',1)[-1])}
        self.assertEqual(set.union(*(locations(r) for r in parsed['metrics'])),{'D3','E3','D4','E4'})
        for cell,direction in [('D3','增加'),('E4','减少')]:
            self.assertEqual({r['value'] for r in parsed['dimensions'] if r['id']=='direction' and cell in locations(r)},{direction})
        self.assertEqual({r['value'] for r in parsed['dimensions'] if r['id']=='item' and 'F8' in locations(r)},{'设备'})
        self.assertFalse(any('C11' in locations(r) for r in parsed['dimensions'] if r['id']=='item'))

    def test_行列矩阵标题冒号和跨行编号保留列内字段(self):
        from 附注更新.A合集 import direct_answer
        from 附注更新.A语义阅读 import addresses
        standard={'metric':{'metric.balance':{'dimension_refs':[{'id':'member'}]},'metric.change':{'dimension_refs':[{'id':'member'},{'id':'direction'}]}},
                  'dimension':{'member':{'value_type':'text'},'direction':{'value_type':'text'}}}
        raw=('|行|member|C：期初余额|D：增加|E：减少|F：期末余额|\n|---|---|---|---|---|---|\n'
             '|3|设备|metric.balance|metric.change；direction=增加|metric.change；direction=减少|metric.balance|\n'
             '|4–6|其他设备|各行均为 metric.balance|各行均为 metric.change；direction=增加|各行同一变动指标；direction=减少|各行均为 metric.balance|')
        result=direct_answer(raw,standard,['原表'])
        slots={cell:row['id'] for row in result['metrics'] for area in row['ranges'] for cell in addresses(area.rsplit('!',1)[-1])}
        self.assertEqual(slots['C3'],'metric.balance');self.assertEqual(slots['E6'],'metric.change')
        for cell,value in [('D3','增加'),('E3','减少'),('D5','增加'),('E6','减少')]:
            values={r['value'] for r in result['dimensions'] if r['id']=='direction' and any(cell in addresses(a.rsplit('!',1)[-1]) for a in r['ranges'])}
            self.assertEqual(values,{value})

    def test_矩阵按明确依次声明搬运成员和列内方向(self):
        from 附注更新.A合集 import direct_answer
        from 附注更新.A语义阅读 import addresses
        standard={'metric':{'metric.change':{'dimension_refs':[{'id':'item'},{'id':'change_direction'},{'id':'asset_change_selection'}]}},
                  'dimension':{'item':{'value_type':'selection'},'change_direction':{'value_type':'text'},'asset_change_selection':{'value_type':'selection'}}}
        raw=('|行|item|D：增加|E：减少|\n|---|---|---|---|\n'
             '|4–5|各行分别为 members=[甲／乙]|各行均为 metric.change；方向增加，事项范围 `all`|各行同一变动指标；方向减少，事项范围 `all`|')
        answer=direct_answer(raw,standard,['原表'])
        for cell,member,direction in [('D4','甲','增加'),('E5','乙','减少')]:
            fields={r['id']:r['value'] for r in answer['dimensions'] if any(cell in addresses(a.rsplit('!',1)[-1]) for a in r['ranges'])}
            self.assertEqual(fields['item'],{'mode':'members','members':[member]})
            self.assertEqual(fields['change_direction'],direction)
            self.assertEqual(fields['asset_change_selection'],'all')

    def test_矩阵一列未给编号仍保存该列已给出的共同字段(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.amount':{'dimension_refs':[{'id':'item'},{'id':'period'}]}},
                  'dimension':{'item':{'value_type':'text'},'period':{'value_type':'period'}}}
        raw=('C列、D列：period=2025-12-31。\n|行|item|C：第一列|D：第二列|\n|---|---|---|---|\n'
             '|3|甲|metric.amount|未给编号|')
        answer=direct_answer(raw,standard,['原表'])
        self.assertFalse(any('原表!D3' in row['ranges'] for row in answer['metrics']))
        fields={row['id']:row['value'] for row in answer['dimensions'] if '原表!D3' in row['ranges']}
        self.assertEqual(fields['item'],'甲');self.assertEqual(fields['period'],'2025-12-31')

    def test_并列列名期间及已写明币种代码接入(self):
        from 附注更新.A合集 import direct_answer,_direct_table
        standard={'metric':{'metric.amount':{'dimension_refs':[{'id':'period'}]}},
                  'dimension':{'period':{'value_type':'period'},'currency':{'value_type':'text'}}}
        raw=('D列、E列：template_period_role=本期；period=2025年1月1日至2025年12月31日。\n'
             '原表!D3:E3：metric_id=metric.amount。')
        answer=direct_answer(raw,standard,['原表'])
        periods={cell:r['value'] for r in answer['dimensions'] if r['id']=='period' for cell in r['ranges']}
        self.assertEqual(periods['原表!D3'],'2025-01-01至2025-12-31')
        self.assertEqual(periods['原表!E3'],'2025-01-01至2025-12-31')
        table={'sheet':'原表','data_ranges':['D3'],'metrics':[],
               'dimensions':[{'id':'currency','ranges':['D3'],'value':'人民币，`CNY'}]}
        saved=_direct_table(table,{'source_context':[]},standard)
        self.assertEqual(saved['dimensions'][0]['value'],'CNY')

    def test_明说的范围字段简称和实际时点接入(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.amount':{'dimension_refs':[{'id':'member'},{'id':'period'},{'id':'unit_scale'}]}},
                  'dimension':{'member':{'value_type':'text'},'period':{'value_type':'period'},'unit_scale':{'value_type':'number'}}}
        raw=('member按各行实际范围取值，不用其他含义代替实际范围。\nunit_scale=1，原文说明金额单位为人民币元。\n'
             'B列为期末，template_period_role=期末，实际时点为2025年12月31日；标签在原表!B2。'
             'C列为期初，template_period_role=期初，实际时点为2024年12月31日；标签在原表!C2。这两个标签适用于以下明细行及合计行。\n'
             '原表!B3:C3：metric_id=metric.amount；范围为成员甲。')
        answer=direct_answer(raw,standard,['原表'])
        self.assertTrue(any(r['id']=='member' and r['value']=='成员甲' for r in answer['dimensions']))
        self.assertEqual({r['value'] for r in answer['dimensions'] if r['id']=='member'},{'成员甲'})
        self.assertTrue(any(r['id']=='unit_scale' and str(r['value'])=='1' for r in answer['dimensions']))
        self.assertTrue(any(r['id']=='period' and r['value']=='2024-12-31' and '原表!C3' in r['ranges'] and r.get('label_references')==['原表!C2'] for r in answer['dimensions']))

    def test_本期角色两种明确写法保存时不形成假冲突(self):
        from 附注更新.A合集 import _direct_table
        table={'sheet':'原表','data_ranges':['D3'],'metrics':[{'id':'metric.amount','ranges':['D3']}],
               'dimensions':[{'id':'template_period_role','ranges':['D3'],'value':value} for value in ('current','current_period')]}
        saved=_direct_table(table,{'source_context':[]},{'dimension':{}})
        self.assertEqual(saved['metrics'][0]['template_period_role'],'current')
        self.assertFalse(saved.get('unresolved'))

    def test_字段明确同某位置只沿用所指字段并保留说明(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.amount':{'dimension_refs':[{'id':'member'},{'id':'direction'}]}},
                  'dimension':{'member':{'value_type':'text'},'direction':{'value_type':'text'}}}
        raw=('|位置|编号|取值|\n|---|---|---|\n'
             '|C3|metric.amount|member=甲；direction=增加|\n'
             '|D3|metric.amount|member=同 `C3`；原格空白，不填写金额。direction=减少|')
        result=direct_answer(raw,standard,['原表'])
        fields={r['id']:r for r in result['dimensions'] if r['ranges']==['原表!D3']}
        self.assertEqual(fields['member']['value'],'甲')
        self.assertEqual(fields['direction']['value'],'减少')
        self.assertIn('原格空白',fields['member']['original']['field_reference'])

    def test_连续列声明和标签适用行保留原答复期间关系(self):
        from 附注更新.A合集 import direct_answer
        from 附注更新.A语义阅读 import addresses
        standard={'metric':{'metric.amount':{'dimension_refs':[{'id':'period'}]}},
                  'dimension':{'period':{'value_type':'period'}}}
        raw=('C:E列：template_period_role=本期期末余额，period=2025-12-31。对应 原表!C2 原文“期末余额”。\n'
             'F:H列：template_period_role=上期期初余额，period=2024-12-31。对应 原表!F2 原文“期初余额”。\n'
             '|坐标|编号|\n|---|---|\n|C3:H4|metric.amount|')
        result=direct_answer(raw,standard,['原表'])
        def values(parsed,cell,key):
            return [r for r in parsed['dimensions'] if r['id']==key and any(a!='*' and cell in addresses(a.rsplit('!',1)[-1]) for a in r['ranges'])]
        for cell,date,role,ref in [('D3','2025-12-31','closing','原表!C2'),('G4','2024-12-31','opening','原表!F2')]:
            self.assertTrue(any(r.get('value')==date and r.get('label_references')==[ref] for r in values(result,cell,'period')))
            self.assertTrue(any(r.get('value')==role for r in values(result,cell,'template_period_role')))
        raw=('数据格为 C5:E7。\nC5:E7：metric.amount。\n'
             '以上余额及比例的 period 均为实际日期 **2025-12-31**，template_period_role 为 **期末**。'
             '解释期间的原表标签位于合并单元格 C3:E3，原文为“期末余额”，适用于第5至7行。')
        result=direct_answer(raw,standard,['原表'])
        for cell in ('C5','D6','E7'):
            self.assertTrue(any(r.get('value')=='2025-12-31' and r.get('label_references')==['原表!C3:E3'] for r in values(result,cell,'period')))
            self.assertEqual({r['value'] for r in values(result,cell,'period') if r.get('value')},{'2025-12-31'})
            self.assertTrue(any(r.get('value')=='期末' for r in values(result,cell,'template_period_role')))

    def test_材料准备中断可接续且不覆盖已有文件(self):
        engine=UnifiedSemanticEngine(self.config,self.gold,self.folder/'中断接续')
        from 附注更新.语义 import _hash
        prompt='直接识别';materials={'原文.json':{'原文':'保留空白','值':None}}
        identity={'prompt':prompt,'materials':materials,'model':engine.settings.get('model'),
                  'endpoint':engine.settings.get('base_url'),'thinking':engine.settings.get('thinking_mode'),
                  'effort':engine.settings.get('model_reasoning_effort')}
        directory=engine.work_dir/'直接识别请求'/_hash(identity)
        material=directory/'材料';material.mkdir(parents=True)
        (material/'原文.json').write_bytes(b'partial')
        task=directory/'任务.json';task.write_bytes(b'original task')
        with patch.object(engine,'_http',return_value='已有含义') as send:
            self.assertEqual(engine.request_text(prompt,materials),'已有含义')
            self.assertEqual(engine.request_text(prompt,materials),'已有含义')
        self.assertEqual(send.call_count,1)
        self.assertEqual((material/'原文.json').read_bytes(),b'partial')
        self.assertEqual(task.read_bytes(),b'original task')
        saved=Path(send.call_args.args[1]['material_directory'])/'原文.json'
        self.assertEqual(json.loads(saved.read_text(encoding='utf-8-sig')),materials['原文.json'])

    def test_拒绝采用的标准提及不变成绑定且共同块保持作用范围(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.amount':{'dimension_refs':[{'id':'entity'},{'id':'currency'}]}},
                  'dimension':{'entity':{'value_type':'text'},'currency':{'value_type':'text'}}}
        rejected=direct_answer('本次范围 A2:A3。\n金标准中只有 metric.amount；它不适用于本表。因此，这些格没有可据以选择的金标准槽位编号。\n|坐标|含义|\n|---|---|\n|A2|税率|\n|A3|备注|',standard,['原表'])
        self.assertFalse(rejected['metrics'])
        raw=('固定目标为原表!A1:C3；本次待更新财务格为B3:C3。\n共同实际取值：\n'
             '- entity：甲公司\n- currency：人民币，代码 `CNY`\n'
             '|坐标|编号|\n|---|---|\n|B3:C3|metric.amount|')
        result=direct_answer(raw,standard,['原表'])
        self.assertTrue(any(row['id']=='entity' and row['ranges']==['*'] for row in result['dimensions']))
        self.assertTrue(any(row['id']=='currency' and row['value']=='CNY' for row in result['dimensions']))

    def test_括注维度共同取值及列标签适用声明原样接入(self):
        from 附注更新.A合集 import direct_answer
        ids=('entity','report_scope','currency','unit_scale','period','cash_selection')
        standard={'metric':{'metric.amount':{'dimension_refs':[{'id':key} for key in ids]}},
                  'dimension':{key:{'value_type':'selection' if key=='cash_selection' else 'text'} for key in ids}}
        raw=('数据格为 C3:D4；B列提供行项目。\n'
             '共同适用的金标准为 metric.amount。共同维度取值：报告主体（`entity`）为甲公司；实际报表范围（`report_scope`）为`consolidated`；金额币种（`currency`）为CNY；单位倍率（`unit_scale`）为1。\n'
             '列标签在 **sheet1!C2**，原文“期末余额”，对应 template_period_role=期末、period=2025年12月31日；在 **sheet1!D2**，原文“期初余额”，对应 template_period_role=期初、period=2024年12月31日。适用于下列全部明细和合计行。\n'
             '|财务格|对象及范围；维度取值|\n|---|---|\n'
             '|C3、D3（B3“现金”）|cash_selection=members(现金)|\n'
             '|C4、D4（B4“存款”）|cash_selection=members(存款)|')
        result=direct_answer(raw,standard,['sheet1'])
        common={r['id']:r['value'] for r in result['dimensions'] if r['ranges']==['*']}
        self.assertEqual({key:common[key] for key in ('entity','report_scope','currency','unit_scale')},
                         {'entity':'甲公司','report_scope':'consolidated','currency':'CNY','unit_scale':1})
        self.assertEqual(result['data_ranges'],['C3:D4'])
        selection=[r for r in result['dimensions'] if r['id']=='cash_selection']
        self.assertEqual(selection[0]['ranges'],['sheet1!C3','sheet1!D3'])
        self.assertEqual(selection[0]['value'],{'mode':'members','members':['现金']})
        self.assertEqual({cell for row in result['metrics'] for cell in row['ranges']},
                         {'sheet1!C3','sheet1!D3','sheet1!C4','sheet1!D4'})
        periods={cell:r for r in result['dimensions'] if r['id']=='period' for cell in r['ranges']}
        self.assertEqual(periods['sheet1!C3']['value'],'2025-12-31')
        self.assertEqual(periods['sheet1!D3']['label_references'],['sheet1!D2'])

    def test_范围箭头并列子表及明说相同候选直接接入(self):
        raw=('|子表（工作表!完整范围 → 数据范围）|相关标准候选|\n|---|---|\n'
             '|原表说明（sheet1!B2:D5 → C3:D5）；续表（B9:D12 → C10:D12）|metric.amount|\n'
             '|另一表（sheet1!B16:D19 → C17:D19）|与原表说明表相同的相关候选|')
        rows,unread=read_scope_answer(raw,{'sheets':[{'name':'sheet1'}]},
                                    {'metric':{'metric.amount':{}},'dimension':{}})
        self.assertEqual([r['range'] for r in rows],['B2:D5','B9:D12','B16:D19'])
        self.assertEqual([r['data_ranges'] for r in rows],[['C3:D5'],['C10:D12'],['C17:D19']])
        self.assertEqual([r['metric_ids'] for r in rows],[['metric.amount']]*3)
        self.assertFalse(unread)


    def test_表格同某格引用只继承未改写的字段(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.amount':{'dimension_refs':[{'id':'member'},{'id':'direction'}]}},
                  'dimension':{'member':{'value_type':'text'},'direction':{'value_type':'text'}}}
        raw=('|位置|编号|维度|\n|---|---|---|\n'
             '|D5|metric.amount|member=甲；direction=增加|\n'
             '|E5|metric.amount|同D5，但 direction=减少|')
        result=direct_answer(raw,standard,['原表'])
        values={(row['id'],cell):row['value'] for row in result['dimensions'] for cell in row['ranges']}
        self.assertEqual(values[('member','原表!E5')],'甲')
        self.assertEqual(values[('direction','原表!E5')],'减少')

    def test_分层范围各自保留编号行归属和列期间(self):
        from 附注更新.A合集 import direct_answer
        standard={'metric':{'metric.amount':{'dimension_refs':[{'id':'period'},{'id':'member'}]},
                            'metric.ratio':{'dimension_refs':[{'id':'period'}]}},
                  'dimension':{'period':{'value_type':'period'},'member':{'value_type':'text'}}}
        raw=(' - **C2:C4、G2:G4：金额**\n'
             '   - 槽位：`metric.amount`\n'
             '   - C列 period=2024-12-31，template_period_role=期初；G列 period=2025-12-31，template_period_role=期末。\n'
             '   - 第2行、第3行 member 分别为“甲”“乙”。\n'
             '   - 第4行 member=总体。\n'
             ' - **D2:D4、H2:H4：比例**\n'
             '   - 槽位：`metric.ratio`\n')
        result=direct_answer(raw,standard,['原表'])
        amounts=[r for r in result['metrics'] if r['id']=='metric.amount']
        ratios=[r for r in result['metrics'] if r['id']=='metric.ratio']
        self.assertEqual(amounts[0]['ranges'],['原表!C2:C4','原表!G2:G4'])
        self.assertEqual(ratios[0]['ranges'],['原表!D2:D4','原表!H2:H4'])
        periods=[r for r in result['dimensions'] if r['id']=='period']
        self.assertEqual({r['value']:r['ranges'] for r in periods},
            {'2024-12-31':['原表!C2','原表!C3','原表!C4'],'2025-12-31':['原表!G2','原表!G3','原表!G4']})
        self.assertTrue(any(r['id']=='member' and r['value']=='乙' and r['ranges']==['原表!C3','原表!G3'] for r in result['dimensions']))

    def test_来源省略号沿用同段已经声明的表名(self):
        from 附注更新.A合集 import source_relation
        base={'source_hash':'a','gold_hash':'g','facts':[],'slot_bindings':[],
              '机械记录':{'sheets':[{'name':'原表','merges':[]}]}}
        source={**base,'source_hash':'b','机械记录':{'sheets':[{'name':'来源表','merges':[]}]}}
        raw='A 原表!C2,D2 ← B 来源表!B4,C4\nA 原表!C3,D3 ← B …!B5,C5'
        targets=['原表!C2','原表!D2','原表!C3','原表!D3']
        result=source_relation(base,source,raw,targets)
        self.assertEqual(result['bindings'],[{'target':target,'source':'来源表!'+cell}
            for target,cell in zip(targets,['B4','C4','B5','C5'])])
        name=source_relation(base,source,'A 原表!B3 名称采用 B 来源表!A5：“带限定的名称”。仍使用原行，不新增行。',targets)
        self.assertEqual(name['label_actions'],[{'action':'copy','target':'原表!B3','source':'来源表!A5'}])

    def setUp(self):
        self.folder=Path(__file__).resolve().parents[1]/'测试结果'/('正常入口贯通_'+uuid.uuid4().hex)
        self.folder.mkdir(parents=True)
        self.gold=self.folder/'标准.jsonl'
        self.gold.write_text('\n'.join(json.dumps(row,ensure_ascii=False) for row in records()),encoding='utf-8')
        for side,year,value in [('A',2024,10),('B',2025,20)]:
            book=Workbook();ws=book.active;ws.title='原表'
            ws.append(['合成企业；单体；人民币；元'])
            ws.append(['项目',f'{year}年末',f'{year-1}年末'])
            ws.append(['设备',value,None]);book.save(self.folder/(side+'.xlsx'));book.close()
        self.config={'a_path':str(self.folder/'A.xlsx'),'b_path':str(self.folder/'B.xlsx'),
            'gold_path':str(self.gold),'output_dir':str(self.folder/'输出'),'base_url':'codex://local','model':'offline-test',
            'update_purpose':{'reason':'日期随B，原值搬运；缺来源保留'},
            'comparison_requirements':[{'维度ID':'entity','要求含义':'不参与比较','出处':'测试明确要求'}]}

    def test_明确要求经过设置保存回读不丢失(self):
        self.config['b_scope_answers']=['已保存范围原答复.txt','已保存候选原答复.txt']
        path=self.folder/'设置.json';save_settings(self.config,path);saved=load_settings(path)
        for key in ('update_purpose','comparison_requirements','b_scope_answers'):self.assertEqual(saved[key],self.config[key])

    def test_正常启动接续沿用分步入口已有成果而不修改旧记录(self):
        from 附注更新.无界面 import run_background
        config={**self.config,'word_path':self.config['a_path']}
        def first(step,settings,*args):
            return {'status':'complete' if step=='extract_a' else 'failed','message':'离线步骤',
                    'files':{'a_path':self.config['a_path']} if step=='extract_a' else {}}
        with patch('附注更新.无界面.run_step',side_effect=first):
            result=run_background(config)
        folder=Path(result['output_dir']);originals={p:p.read_bytes() for p in folder.iterdir() if p.is_file()}
        config.update(a_mapping='已保存A合集.json',b_mapping='已保存B合集.json',b_scope_answers=['范围原答复.txt'])
        called=[]
        def continued(step,settings,*args):
            called.append(step)
            self.assertEqual(settings['a_mapping'],config['a_mapping'])
            self.assertEqual(settings['b_mapping'],config['b_mapping'])
            self.assertEqual(settings['b_scope_answers'],config['b_scope_answers'])
            return {'status':'complete','message':'离线步骤','files':{}}
        with patch('附注更新.无界面.run_step',side_effect=continued):
            resumed=run_background(config,resume=folder)
        self.assertEqual(called,['recognize_a','recognize_b','match','write'])
        self.assertEqual(resumed['status'],'complete')
        self.assertTrue(all(p.read_bytes()==data for p,data in originals.items()))

    def test_正常启动已结束作业接入明确选入的新合集(self):
        from 附注更新.无界面 import run_background
        config={**self.config,'word_path':self.config['a_path']}
        def first(step,settings,*args):
            return {'status':'complete','message':'离线步骤','files':{'b_mapping':'原B合集.json'} if step=='recognize_b' else {}}
        with patch('附注更新.无界面.run_step',side_effect=first):
            finished=run_background(config)
        folder=Path(finished['output_dir']);old={p:p.read_bytes() for p in folder.iterdir() if p.is_file()}
        calls=[]
        def continued(step,settings,*args):
            calls.append(step);self.assertEqual(settings['b_mapping'],'明确选入的新B合集.json')
            return {'status':'complete','message':'离线接续','files':{}}
        with patch('附注更新.无界面.run_step',side_effect=continued):
            result=run_background({**config,'b_mapping':'明确选入的新B合集.json'},resume=folder)
        self.assertEqual(calls,['recognize_b','match','write'])
        self.assertEqual(result['status'],'complete')
        self.assertTrue(all(p.read_bytes()==data for p,data in old.items()))

    def test_按范围引用不重复塞入旧答复历史(self):
        from 附注更新.A语义阅读 import direct_reference
        record={'source_hash':'a','gold_hash':'g','facts':[],'slot_bindings':[],
            'unresolved':[{'sheet':'原表','cell':'B3','reason':'未答','original':{'reply':'历史全文'*1000}},
                          {'sheet':'原表','cell':'B4','reason':'范围外'}]}
        result=direct_reference(record,compact=True,ranges=[{'sheet':'原表','range':'B3'}])
        self.assertEqual(result['unresolved'],[{'sheet':'原表','cell':'B3','reason':'未答'}])
        self.assertIn('original',record['unresolved'][0])

    def test_收窄材料包含标准明说的相关定义且不选择身份(self):
        from 附注更新.A语义阅读 import standard_material
        standard={'dimension':{},'metric':{
            'metric.combined':{'id':'metric.combined','status':'active','dimension_refs':[],
                'definition':{'measurement_basis':'单类使用metric.single；旧项metric.old已停用。'}},
            'metric.single':{'id':'metric.single','status':'active','dimension_refs':[],
                'definition':{'measurement_basis':'组合使用metric.combined。'}},
            'metric.old':{'id':'metric.old','status':'retired','dimension_refs':[],'definition':{}}}}
        material=standard_material(['metric.combined'],standard)
        self.assertEqual([row['id'] for row in material['metrics']],['metric.combined','metric.single'])
        self.assertEqual(material['metrics'][0],standard['metric']['metric.combined'])
        self.assertNotIn('selected_metric',material)

    def test_行号表格只应用到明说的同一行各列(self):
        from 附注更新.A合集 import direct_answer
        engine=UnifiedSemanticEngine(self.config,self.gold,self.folder/'行号表')
        raw=('下表所列槽位及维度取值，对该行B、C两格都适用。\n'
             '|行及原文项目|金标准槽位与实际范围|\n|---|---|\n'
             '|2 “标题”|不是金额项目，无对应金额槽位。|\n'
             '|3 “旧项”|固定资产原值；asset_selection=members[旧项]。|\n'
             '|4 “新项”|同一指标；asset_selection=members[新项]。|')
        result=direct_answer(raw,engine.standard,['原表'])
        matched={loc for row in result['metrics'] if row['id']=='固定资产原值' for loc in row['ranges']}
        self.assertEqual(matched,{'原表!B3','原表!C3','原表!B4','原表!C4'})
        self.assertTrue(all(set(row['ranges'])<={'原表!B2','原表!C2'} for row in result['unresolved']))

    def test_否定说明提到的编号不冒充已选槽位且保留已说维度(self):
        from 附注更新.A合集 import direct_answer
        engine=UnifiedSemanticEngine(self.config,self.gold,self.folder/'否定编号')
        result=direct_answer('原表!B3:C3：metric_id=金标准无适用槽位；固定资产原值不是本格的计量。asset_selection=members[设备]。',
            engine.standard,['原表'])
        self.assertFalse(result['metrics'])
        self.assertTrue(any(row['id']=='asset_selection' for row in result['dimensions']))
        self.assertTrue(result['unresolved'])

    def test_行号列名矩阵保留槽位维度及分列日期(self):
        from 附注更新.A合集 import direct_answer
        engine=UnifiedSemanticEngine(self.config,self.gold,self.folder/'行列矩阵')
        raw=('period与template_period_role：B列为2025年12月31日（期末）；C列为2024年12月31日（期初）。\n'
             '| 行及原文对象 | asset_selection（范围） | B列 | C列 |\n|---|---|---|---|\n'
             '| 3 设备 | `members`：设备 | 固定资产原值 | 固定资产原值 |\n'
             '| 4 合计 | `all`：完整总体 | 固定资产原值 | 固定资产原值 |\n'
             'B各格的 `period` 均为2025年12月31日；C各格的 `period` 均为2024年12月31日。')
        result=direct_answer(raw,engine.standard,['原表'])
        self.assertEqual({loc for row in result['metrics'] for loc in row['ranges']},
                         {'原表!B3','原表!C3','原表!B4','原表!C4'})
        periods={loc:row['value'] for row in result['dimensions'] if row['id']=='period' for loc in row['ranges']}
        self.assertEqual(periods['原表!B3'],'2025-12-31');self.assertEqual(periods['原表!C3'],'2024-12-31')
        self.assertNotIn('*',periods)
        self.assertTrue(any(row['id']=='asset_selection' and row['value']=='all' for row in result['dimensions']))

    def test_同列分组编号与项目文字中的维度均接入(self):
        from 附注更新.A合集 import direct_answer
        engine=UnifiedSemanticEngine(self.config,self.gold,self.folder/'列分组')
        raw=('B至C列依次是期末余额、期初余额。期末时点为2025-12-31，期初时点为2024-12-31。\n'
             '|B行|项目及范围|B的金额列与标准编号|\n|---|---|---|\n'
             '|3|设备；`asset_selection=members：设备`|B、C：固定资产原值|')
        result=direct_answer(raw,engine.standard,['原表'])
        self.assertEqual({loc for row in result['metrics'] for loc in row['ranges']},{'原表!B3','原表!C3'})
        self.assertTrue(any(row['id']=='asset_selection' and '设备' in str(row['value']) for row in result['dimensions']))
        dates={loc:row['value'] for row in result['dimensions'] if row['id']=='period' for loc in row['ranges']}
        self.assertEqual(dates['原表!C3'],'2024-12-31')

    def test_同上只承接同列文字且日期仍各自保留(self):
        from 附注更新.A合集 import direct_answer
        engine=UnifiedSemanticEngine(self.config,self.gold,self.folder/'同上')
        raw=('|坐标|金标准槽位与范围|实际期间与角色|\n|---|---|---|\n'
             '|B3|“余额”；固定资产原值；asset_selection=members[设备]|2025-12-31；期末|\n'
             '|C3|同上|2024-12-31；期初|')
        result=direct_answer(raw,engine.standard,['原表'])
        self.assertEqual({loc for row in result['metrics'] for loc in row['ranges']},{'原表!B3','原表!C3'})
        dates={loc:row['value'] for row in result['dimensions'] if row['id']=='period' for loc in row['ranges']}
        self.assertEqual(dates,{'原表!B3':'2025-12-31','原表!C3':'2024-12-31'})
        self.assertFalse(any(row.get('id')=='metric_id' for row in result['unresolved']))
        ranged=direct_answer('B3:B5「余额」：固定资产原值。',engine.standard,['原表'])
        self.assertEqual({loc for row in ranged['metrics'] for loc in row['ranges']},{'原表!B3:B5'})

    def test_按行顺序成员与同行多项编号不串位(self):
        from 附注更新.A合集 import direct_answer
        standard={'dimension':{'period':{'value_type':'period'},'asset_selection':{'value_type':'selection','domain':'资产'},
                              'change_direction':{'value_type':'text','allowed_values':['增加','减少']}},
                  'metric':{'metric.balance':{'dimension_refs':[{'id':'period'},{'id':'asset_selection'}]},
                            'metric.change':{'dimension_refs':[{'id':'period'},{'id':'asset_selection'},{'id':'change_direction'}]}}}
        raw=('C列“本期增加”及D列“本期减少”=template_period_role:本期，期间2025-01-01至2025-12-31。\n'
             '|行坐标／项目|槽位及取值|\n|---|---|\n'
             '|A3:D3 合计|B3: metric.balance；C3: metric.change，change_direction=增加；D3: 同一变动槽位，方向=减少。行范围=asset_selection:all。|\n'
             '|A4:D4 设备；A5:D5 房屋|B=metric.balance；C、D=metric.change，方向分别为增加、减少。各行分别取 asset_selection=members[设备]、members[房屋]。|')
        result=direct_answer(raw,standard,['原表'])
        metrics={loc:row['id'] for row in result['metrics'] for loc in row['ranges']}
        self.assertEqual(metrics['原表!D3'],'metric.change');self.assertEqual(metrics['原表!B3'],'metric.balance')
        members={loc:row['value'] for row in result['dimensions'] if row['id']=='asset_selection' for loc in row['ranges']}
        self.assertEqual(members['原表!C4'],'members[设备]');self.assertEqual(members['原表!C5'],'members[房屋]')
        directions={loc:row['value'] for row in result['dimensions'] if row['id']=='change_direction' for loc in row['ranges']}
        self.assertEqual(directions['原表!D3'],'减少');self.assertEqual(directions['原表!D5'],'减少')

    def test_来源在前的整段名单及保留空白按原声明展开(self):
        from 附注更新.A合集 import _declared_list_tables
        a={'机械记录':{'names':{'表区':[{'sheet':'模板','range':'B10:D20'}]}}}
        raw=('来源!A4“父项”对应模板!B12。明细来源!A5:A8依次对应扩展后模板!B13:B16。新增的两行沿用原模板!B13的样式。\n'
             '来源!A9对应扩展后模板!B17。\n'
             '来源的B列“本期”搬至模板的C列“本期”；C列“上期”搬至模板的D列“上期”。\n'
             '来源!A3“调整”对应模板!B11，A的原空白保留。')
        result=_declared_list_tables(a,raw,'change_selection',[{'sheet':'来源','range':'A1:C10'}])
        self.assertEqual(len(result),1);table=result[0]
        self.assertEqual(table['source_rows'],[3,4,5,6,7,8,9])
        self.assertEqual(table['target_row_offsets'],[1,2,3,4,5,6,7])
        self.assertEqual(table['operations'],[{'axis':'row','action':'insert','index':5,'count':2,'template_index':3}])
        self.assertEqual(table['blank_values']['3']['C'],{'source':'模板!C11'})
        self.assertEqual(table['row_columns']['3'],{})

    def test_来源表格双向列序与名单名称金额同次接入(self):
        from 附注更新.A合集 import source_relation
        a={'facts':[],'slot_bindings':[],'source_hash':'a','gold_hash':'g','机械记录':{'sheets':[]}}
        b={'source_hash':'b','gold_hash':'g','机械记录':{'sheets':[{'name':'来源'}]},'b_source_ranges':[{'sheet':'来源','range':'A1:C9'}]}
        raw=('|B坐标|采用的A来源|\n|---|---|\n|B4|原表!C3|\n'
             '| A 目标行 | B 来源行 | 更新后的事项名称 | A 本期 ← B B列 |\n|---|---|---|---|\n'
             '| A3 | A4 | 新项 | C3 ← B4：0 |')
        relation=source_relation(a,b,raw,['原表!C3'])
        self.assertEqual(relation['bindings'],[{'target':'原表!C3','source':'来源!B4'}])
        self.assertEqual(relation['label_actions'],[{'action':'copy','target':'原表!A3','source':'来源!A4'}])
        raw='|A需要位置、项目|B来源位置|\n|---|---|\n|原表!A3 旧项|B4、C4|'
        relation=source_relation(a,b,raw,['原表!B3','原表!C3'])
        self.assertEqual(relation['bindings'],[{'target':'原表!B3','source':'来源!B4'},{'target':'原表!C3','source':'来源!C4'}])

    def test_未选A的普通B入口仍只认识目录(self):
        calls=[]
        def answer(engine,prompt,payload):
            calls.append(prompt);self.assertNotIn('金标准.jsonl',payload['materials']);return '仅目录认识。'
        with patch.object(UnifiedSemanticEngine,'_http',answer):result=run_step('recognize_b',self.config)
        self.assertEqual(result['status'],'saved');self.assertEqual(len(calls),1)

    def test_普通窗口保留要求并接收B同时更新的两个合集(self):
        import tkinter as tk
        from 附注更新 import 界面
        root=tk.Tk();root.withdraw()
        try:
            with patch.object(界面,'load_settings',return_value=self.config):app=界面.Application(root)
            self.assertEqual(app.config()['update_purpose'],self.config['update_purpose'])
            self.assertEqual(app.config()['comparison_requirements'],self.config['comparison_requirements'])
            app.update_step_files('recognize_b',{'a_mapping':'新A.json','b_mapping':'新B.json'})
            self.assertEqual(app.config()['a_mapping'],'新A.json')
            self.assertEqual(app.config()['b_mapping'],'新B.json')
        finally:
            for pending in root.tk.call('after','info'):root.tk.call('after','cancel',pending)
            root.destroy()

    def test_范围认识接自然文字与表格不补范围(self):
        standard={'metric':{'固定资产原值':{}},'dimension':{'asset_selection':{}}}
        snapshot={'sheets':[{'name':'原表'}]}
        for text in ['工作表：原表；完整范围：A1:C3；数据范围：B3:C3；相关标准：固定资产原值。',
                     '| 工作表 | 完整范围 | 数据范围 | 标准编号 |\n|---|---|---|---|\n|原表|A1:C3|B3:C3|固定资产原值|']:
            groups,unread=read_scope_answer(text,snapshot,standard)
            self.assertFalse(unread);self.assertEqual(groups[0]['data_ranges'],['B3:C3'])
            self.assertEqual(groups[0]['metric_ids'],['固定资产原值'])
        groups,unread=read_scope_answer('工作表：原表；完整范围：A1:C3；相关标准：固定资产原值',snapshot,standard)
        self.assertFalse(groups);self.assertTrue(unread)

    def test_整份范围材料保留单元格内换行(self):
        from 附注更新.A语义阅读 import scope_material
        from 附注更新.表格 import read_workbook
        source=self.folder/'换行原文.xlsx'
        book=Workbook();book.active['A1']='年末\n账面余额';book.save(source);book.close()
        material=scope_material(read_workbook(source),[])
        self.assertIn('年末\\n账面余额',material['sheets'][0]['rows'])

    def test_范围表头声明A工作表及B工作表中的空格括号(self):
        reply='| B子表 | 完整范围 | 待识别数据范围 | 对应A范围（sheet1!范围） | 相关标准候选 |\n|---|---|---|---|---|\n|货币资金|Z01 资产负债表(企财01表)!A4:D6|A5:D6|B2:D8|固定资产原值|'
        groups,unread=read_scope_answer(reply,{'sheets':[{'name':'Z01 资产负债表(企财01表)'}]},
            {'metric':{'固定资产原值':{}},'dimension':{}})
        self.assertFalse(unread)
        self.assertEqual(groups[0]['sheet'],'Z01 资产负债表(企财01表)')
        self.assertEqual(groups[0]['a_ranges'],[{'sheet':'sheet1','range':'B2:D8'}])

    def test_自然答复的说明编号列归属及成员列表(self):
        from 附注更新.A合集 import direct_answer
        engine=UnifiedSemanticEngine(self.config,self.gold,self.folder/'语义')
        reply=('原表!B3:C3：metric_id=固定资产原值。\n'
            'B 列为 `period=2025-12-31`、`template_period_role=期末`；C 列为 `period=2024-12-31`、`template_period_role=期初`。\n'
            '共同说明 T0 说金额以人民币元列示。\n'
            '| 原表行 | 财务格 | asset_selection |\n|---|---|---|\n|A3设备|B3、C3|members: ["设备"]|\n'
            'B3：period=2025-12-31；C3：period=2024-12-31。')
        result=direct_answer(reply,engine.standard,['原表'])
        selections=[r for r in result['dimensions'] if r['id']=='asset_selection']
        self.assertEqual(selections[0]['value']['members'],['设备'])
        self.assertEqual(set(selections[0]['ranges']),{'原表!B3','原表!C3'})
        self.assertFalse(any('T0' in area for rows in result.values() for r in rows for area in r.get('ranges',[])))

    def test_A共同期间角色可直接交给B(self):
        from 附注更新.A语义阅读 import direct_reference
        mapping={'source_hash':'a','gold_hash':'g','facts':[],'slot_bindings':[],
            'outline':{'subtables':[{'sheet':'原表','data_ranges':['B3:C3'],'dimensions':[
                {'id':'template_period_role','ranges':['*'],'value':'closing'}]}]}}
        result=direct_reference(mapping,compact=True)
        self.assertEqual(result['outline']['subtables'][0]['dimensions'][0]['ranges'],['B3:C3'])

    def test_期间标签在对应关系左侧并按表名分别归属(self):
        from 附注更新.A合集 import direct_answer
        engine=UnifiedSemanticEngine(self.config,self.gold,self.folder/'标签关系')
        raw='来源一：A2“项目”右侧，B2“期末余额”对应 B3:B5；C2“期初余额”对应 C3:C5。\n来源二：B2“期初余额”对应 B3:B7。'
        result=direct_answer(raw,engine.standard,['来源一','来源二'],field_scope={
            '来源一!B3:C5':['period'],'来源二!B3:B7':['period']})
        rows=[r for r in result['dimensions'] if r['id']=='period']
        self.assertEqual({tuple(r['label_references']) for r in rows},
            {('来源一!B2',),('来源一!C2',),('来源二!B2',)})
        self.assertTrue(all('value' not in r for r in rows))

    def test_补充逐格身份保留旧结果且不接额外位置对应(self):
        from 附注更新.A合集 import _records
        book=load_workbook(self.config['b_path']);book.active.append(['另一设备',30,None]);book.save(self.config['b_path']);book.close()
        cfg=dict(self.config)
        def answer(engine,prompt,payload):
            if '整份原文.json' in payload['materials']:
                return '工作表：原表；完整范围：A1:C3；数据范围：B3:C3；相关标准：固定资产原值；对应A：原表!A1:C3'
            return '\n'.join(f'原表!{col}3：metric_id=固定资产原值；entity=合成企业；report_scope=standalone；currency=CNY；unit_scale=1；period={{"kind":"instant","date":"2025-12-31"}}；asset_selection={{"mode":"members","members":["设备"]}}。' for col in 'BC')
        with patch.object(UnifiedSemanticEngine,'_http',answer):
            for step in ('recognize_a','recognize_b'):cfg.update(run_step(step,cfg)['files'])
        reply='\n'.join(f'原表!{col}4：metric_id=固定资产原值；entity=合成企业；report_scope=standalone；currency=CNY；unit_scale=1；period={{"kind":"instant","date":"2025-12-31"}}；asset_selection={{"mode":"members","members":["另一设备"]}}。' for col in 'BC')
        reply+='\nA 原表!B3 采用 B 原表!B4。'
        fields=['metric_id','entity','report_scope','currency','unit_scale','period','asset_selection']
        request={**cfg,'b_source_ranges':[{'sheet':'原表','range':'A1:C4'}],
                 'a_required_ranges':[{'sheet':'原表','range':'A1:C3'}],'b_field_scope':{'原表!B4:C4':fields},
                 'b_additional_context':'本次新来源需要与原有需要一并说明。'}
        with patch.object(UnifiedSemanticEngine,'request_text',return_value=reply) as call:
            result=run_step('recognize_b',request);self.assertEqual(result['status'],'complete',result);cfg.update(result['files'])
        self.assertNotIn('A原表.json',call.call_args.args[1]);call.assert_called_once()
        self.assertIn(request['b_additional_context'],call.call_args.args[0])
        with patch.object(UnifiedSemanticEngine,'_http',side_effect=AssertionError('不重复调用')):
            cfg.update(run_step('recognize_b',cfg)['files'])
        saved=json.loads(Path(cfg['b_mapping']).read_text(encoding='utf-8-sig'))
        self.assertEqual({loc for loc,row in _records(saved)},{'原表!B3','原表!C3','原表!B4','原表!C4'})
        self.assertEqual(saved['已确认关系'],[])

    def test_位置来源不能冒充逐格身份且原答复保留(self):
        from 附注更新.A合集 import save_direct_answer
        cfg=dict(self.config)
        def answer(engine,prompt,payload):
            if '整份原文.json' in payload['materials']:
                return '工作表：原表；完整范围：A1:C3；数据范围：B3:C3；相关标准：固定资产原值；对应A：原表!A1:C3'
            return '\n'.join(f'原表!{col}3：metric_id=固定资产原值；entity=合成企业；report_scope=standalone；currency=CNY；unit_scale=1；period={{"kind":"instant","date":"2025-12-31"}}；asset_selection={{"mode":"members","members":["设备"]}}。' for col in 'BC')
        with patch.object(UnifiedSemanticEngine,'_http',answer):
            for step in ('recognize_a','recognize_b'):cfg.update(run_step(step,cfg)['files'])
        target=json.loads(Path(cfg['a_mapping']).read_text('utf-8-sig'))
        changed=save_direct_answer(target,'原表!B3：entity=新主体。',self.folder/'来源请求前A.json',field_scope={'原表!B3':['entity']})
        cfg['a_mapping']=changed['mapping_path']
        before=json.loads(Path(cfg['b_mapping']).read_text('utf-8-sig'))
        request={**cfg,'b_source_ranges':[{'sheet':'原表','range':'A1:C3'}],
            'a_required_ranges':[{'sheet':'原表','range':'A1:C3'}],
            'b_source_targets':['原表!B3'],'b_field_scope':{'原表!B3':['period']}}
        with patch.object(UnifiedSemanticEngine,'_http',return_value='A 原表!B3 ← B 原表!B3。') as call:
            result=run_step('recognize_b',request)
        self.assertEqual(result['status'],'partial',result);call.assert_called_once()
        self.assertEqual(json.loads(Path(cfg['b_mapping']).read_text('utf-8-sig')),before)
        saved=json.loads(Path(result['files']['b_mapping']).read_text('utf-8-sig'))
        self.assertEqual(saved['本次逐格原答复'],'A 原表!B3 ← B 原表!B3。')
        self.assertFalse(saved['本次逐格接入']['complete'])
        self.assertTrue(any(json.loads(path.read_text(encoding='utf-8-sig')).get('text')=='A 原表!B3 ← B 原表!B3。'
                            for path in self.folder.rglob('答复.json')))

    def test_章节声明槽位及来源列保留维度值而不采用新A行作B位置(self):
        from 附注更新.A合集 import direct_answer
        engine=UnifiedSemanticEngine(self.config,self.gold,self.folder/'章节来源')
        raw=('设备使用 `固定资产原值`。\n本年数为 `2025-01-01` 至 `2025-12-31`。\n'
             '### 设备\n|更新后A行|项目|范围|B本年来源及金额|\n|---|---|---|---|\n'
             '|18|设备|`members=["设备"]`|B3：20|')
        result=direct_answer(raw,engine.standard,['原表'])
        self.assertTrue(any(r['id']=='固定资产原值' and r['ranges']==['原表!B3'] for r in result['metrics']))
        self.assertTrue(any(r['id']=='asset_selection' and '设备' in str(r['value']) for r in result['dimensions']))
        self.assertTrue(any(r['id']=='period' and r['value']=='2025-01-01至2025-12-31' for r in result['dimensions']))

    def test_自然名单排布表复用原明细总计和明示列来源(self):
        from 附注更新.A合集 import _declared_list_tables
        a={'机械记录':{'names':{'名单':[{'sheet':'原表','range':'A2:C5'}]}}}
        raw=('A旧明细为第3—4行、合计第5行。\n'
             '|更新后A行|项目|范围|B本年来源及金额|B上年来源及金额|\n|---|---|---|---|---|\n'
             '|3|设备甲|members|G7：0|H7：0|\n|4|设备乙|members|G8：10|H8：11|\n'
             '|5|设备丙|members|G9：2|H9：3|\n|6|合计|all|G6：12|H6：14|\n'
             'A原设备甲B3/C3取B G7/H7。')
        rows=_declared_list_tables(a,raw,'asset_selection',[{'sheet':'来源','range':'E1:H9'}])
        self.assertEqual(rows[0]['source_rows'],[7,8,9]);self.assertEqual(rows[0]['columns'],{'B':'G','C':'H'})
        self.assertEqual(rows[0]['totals'],[{'target_row':5,'source_row':6}])

    def test_名单续列复用旧来源并保留新增成员自己的列(self):
        from 附注更新.A合集 import _declared_list_tables
        a={'机械记录':{'names':{'名单':[{'sheet':'原表','range':'A2:C5'}]}}}
        table={'selection_dimension':'asset_selection','range_name':'名单','source_sheet':'来源','source_rows':[7,8],
            'first_row_offset':1,'old_detail_count':2,'columns':{'B':'C','C':'D'},'labels':{'A':'A'},'totals':[{'target_row':5,'source_row':6}]}
        b={'机械记录':{'sheets':[{'name':'来源'}]},'已确认关系':[{'dimension':'asset_selection','tables':[table]}]}
        reply=('前2项沿用已保存安排：B表第7–8行名称及C/D列金额，依序搬至A表第3–4行。\n'
               '名称从B表E5搬至A表A5；本年数从G5搬至B5，上年数从H5搬至C5。\n'
               '扩展明细行沿用A表第4行样式。')
        result=_declared_list_tables(a,reply,'asset_selection',[{'sheet':'来源','range':'A1:H8'}],b)[0]
        self.assertEqual(result['columns'],table['columns']);self.assertEqual(result['source_rows'],[7,8,5])
        self.assertEqual(result['row_columns'],{'5':{'B':'G','C':'H'}})
        self.assertEqual(result['declared_labels']['5'],{'source':'来源!E5'})
        self.assertEqual(table['source_rows'],[7,8])

    def test_名单旧范围带工作表名且金额列沿用已明说来源(self):
        from 附注更新.A合集 import _declared_list_tables
        a={'source_hash':'a','gold_hash':'g','facts':[],'slot_bindings':[],'机械记录':{'sheets':[{'name':'原表','cells':{}}],
            'names':{'名单':[{'sheet':'原表','range':'A2:C5'}]}}}
        b={'source_hash':'b','gold_hash':'g','机械记录':{'sheets':[{'name':'来源','cells':{}}]},'target_tables':[{'sheet':'来源'}]}
        raw=('明细为 原表!A3:A4，合计行为 原表!A5。\n'
             '|A需要坐标|本次采用的B来源|\n|---|---|\n|B3、C3|G7、H7|\n|B4、C4|G8、H8|\n|B5、C5|G6、H6|\n'
             '|更新后的 A 行|名称|期末金额来源|期初金额来源|\n|---|---|---|---|\n'
             '|3|甲|G7|H7|\n|4|乙|G8|H8|\n|5|丙|G9|H9|\n|6|合计|G6|H6|')
        table=_declared_list_tables(a,raw,'asset_selection',[{'sheet':'来源','range':'A1:H9'}],b)[0]
        self.assertEqual(table['columns'],{'B':'G','C':'H'});self.assertEqual(table['source_rows'],[7,8,9])

    def test_名单换序时原值跟随模型指定成员而不是旧行(self):
        from 附注更新.A合集 import _declared_list_tables
        a={'机械记录':{'names':{'名单':[{'sheet':'原表','range':'A2:C4'}]}}}
        raw=('按 B 表第 4–5 行顺序写入 A 表 A3:A4。B4→A3 新项；B5→A4 旧项。\n'
             '本期按 B 表 B4:B5 → A 表 B3:B4；上期按 B 表 C4:C5 → A 表 C3:C4。\n'
             '旧项原内容从 B3、C3 保留至 B4、C4；其他新增成员金额留空。\n'
             '合计：B6 名称 → A 表 A5；A原合计内容从 B4、C4 保留在合计位置 B5、C5。')
        table=_declared_list_tables(a,raw,'asset_selection',[{'sheet':'来源','range':'A1:C6'}])[0]
        self.assertEqual(table['blank_values']['5'],{'B':{'source':'原表!B3'},'C':{'source':'原表!C3'}})
        self.assertEqual(table['blank_values']['4'],{'B':{'value':None},'C':{'value':None}})
        self.assertEqual(table['totals'],[{'target_row':4,'source_row':6}])

    def test_旧名单简称只用该答复范围不受后续补充来源影响(self):
        from 附注更新.A合集 import _declared_list_tables
        a={'机械记录':{'names':{'名单':[{'sheet':'原表','range':'B2:D5'}]}}}
        raw=('|A位置|名称|B名称来源|金额来源|\n|---|---|---|---|\n'
             '|A3|甲|B!A6|B!B6、C6 → A!C3、D3|\n'
             '|A4|乙|B!A7|B!B7、C7 → A!C4、D4|\n'
             '沿用A原合计行C5:D5的值与空白，不搬入B8。')
        source=[{'sheet':'QCF52 来源','range':'A1:C8'}]
        expected=_declared_list_tables(a,raw,'asset_selection',source,{'b_source_ranges':source})
        later=_declared_list_tables(a,raw,'asset_selection',source,
            {'b_source_ranges':[{'sheet':'QCF135 后续另一科目','range':'A1:C8'}]})
        self.assertEqual(later,expected)
        self.assertEqual(later[0]['columns'],{'C':'B','D':'C'})

    def test_B答复A引用不会混成B身份及来源简称沿用(self):
        from 附注更新.A合集 import _own_answer_text,direct_answer,source_relation
        engine=UnifiedSemanticEngine(self.config,self.gold,self.folder/'来源语义')
        reply=('B来源表!B3：metric_id=固定资产原值。\n'
            'A!C3←B来源表!B3；A!D3←!C3。asset_selection：domain=固定资产，mode=all。解释文字。\n'
            'A!C4未找到来源。按A原文识别asset_selection=members[旧成员]。')
        b={'carrier':'B','A需要':{'需要':[{'location':'原表!C3'}]},'机械记录':{'sheets':[{'name':'B来源表'}]},
           'source_hash':'b','gold_hash':'g','target_tables':[{'sheet':'B来源表'}]}
        plain=_own_answer_text(reply,b,b['机械记录'])
        self.assertNotIn('旧成员',plain)
        parsed=direct_answer(plain,engine.standard,['B来源表'])
        row=next(r for r in parsed['dimensions'] if r['id']=='asset_selection')
        self.assertEqual(row['ranges'],['B来源表!B3','B来源表!C3']);self.assertEqual(row['value']['mode'],'all')
        a={'facts':[],'slot_bindings':[],'source_hash':'a','gold_hash':'g','机械记录':{'sheets':[]}}
        relation=source_relation(a,b,reply,['原表!C3','原表!D3'])
        self.assertEqual(relation['bindings'],[{'target':'原表!C3','source':'B来源表!B3'},{'target':'原表!D3','source':'B来源表!C3'}])
        aliases='A!C3←B!B3；A!D3←B!C3。'
        self.assertEqual(source_relation(a,b,aliases,['原表!C3','原表!D3'])['bindings'],relation['bindings'])
        self.assertIn('B来源表!B3',_own_answer_text(aliases,b,b['机械记录']))

    def test_字段解释与表格引用不混成字段取值(self):
        from 附注更新.A合集 import direct_answer
        engine=UnifiedSemanticEngine(self.config,self.gold,self.folder/'表格语义')
        reply=('B 的主体为该公司。\nentity=合成企业。\n'
            'B、C列 template_period_role=closing；D、E列 template_period_role=opening。\n'
            '| B坐标及原文 | 所选槽位 | asset_selection | 说明 |\n|---|---|---|---|\n'
            '|B3、D3|固定资产原值|members：设备|维度同B4。|\n'
            'B4：所选槽位为固定资产原值，含义说明。')
        result=direct_answer(reply,engine.standard,['原表'])
        self.assertEqual([r['value'] for r in result['dimensions'] if r['id']=='entity'],['合成企业'])
        row=next(r for r in result['dimensions'] if r['id']=='asset_selection')
        self.assertEqual(row['value']['members'],['设备'])
        self.assertEqual(row['ranges'],['原表!B3','原表!D3'])
        self.assertTrue(any(r['id']=='固定资产原值' and r['ranges']==['原表!B4'] for r in result['metrics']))
        roles=[r for r in result['dimensions'] if r['id']=='template_period_role']
        self.assertFalse(any('*' in r['ranges'] for r in roles))

    def test_自然来源表头及分句行列对应直接接入(self):
        from 附注更新.A合集 import source_relation
        a={'facts':[],'slot_bindings':[],'source_hash':'a','gold_hash':'g','机械记录':{'sheets':[]}}
        b={'source_hash':'b','gold_hash':'g','机械记录':{'sheets':[{'name':'来源表'}]},'target_tables':[{'sheet':'来源表'}]}
        reply=('| A需要坐标 | B采用来源 |\n|---|---|\n|A!C3、D3|B!B4、C4|\n'
               'A C5 对应B B6，A D5 对应B C6。\n'
               'A 行7至8分别对应 B 行10至11。\nA 的 C 列对应 B 的 B 列。\nA 的 D 列对应 B 的 C 列。')
        targets=['原表!'+c+str(r) for r in (3,5,7,8) for c in 'CD']
        relation=source_relation(a,b,reply,targets)
        self.assertEqual({r['target']:r['source'] for r in relation['bindings']},
            {'原表!'+c+str(r):'来源表!'+s+str(n) for r,n in ((3,4),(5,6),(7,10),(8,11)) for c,s in (('C','B'),('D','C'))})
        narrative='A的 C3:C6 对应这四项，C7是仅合计这四家单位的余额。A各项需要目前均无可采用的B搬运来源。'
        self.assertEqual(source_relation(a,b,narrative,['原表!C3','原表!C4','原表!C5','原表!C6','原表!C7'])['bindings'],[])

    def test_行号与列声明组成的位置保留指标和期间(self):
        from 附注更新.A合集 import direct_answer
        engine=UnifiedSemanticEngine(self.config,self.gold,self.folder/'行列语义')
        reply=('列 B 为 2024年12月31日、期初；列 C 为 2025年12月31日、期末。\n'
               '|行及原文项目|metric（槽位编号）|asset_selection|\n|---|---|---|\n'
               '|3「设备」|B、C：固定资产原值|members=["设备"]|')
        answer=direct_answer(reply,engine.standard,['原表'])
        self.assertEqual({a for row in answer['metrics'] for a in row['ranges']},{'原表!B3','原表!C3'})
        periods={r['ranges'][0]:r['value'] for r in answer['dimensions'] if r['id']=='period'}
        self.assertEqual(periods,{'原表!B3':'2024-12-31','原表!C3':'2025-12-31'})

    def test_标签和日期同句时均保留且标签不是金额目标(self):
        from 附注更新.A合集 import direct_answer
        engine=UnifiedSemanticEngine(self.config,self.gold,self.folder/'同句期间')
        answer=direct_answer('原表!B3:B4 对应表头 B2「期末余额」，template_period_role=期末，period=2025年12月31日。',engine.standard,['原表'])
        periods=[r for r in answer['dimensions'] if r['id']=='period' and 'value' in r]
        self.assertTrue(periods)
        self.assertEqual(periods[0]['value'],'2025-12-31')
        self.assertEqual(periods[0]['evidence'],['原表!B2'])
        self.assertFalse(any('B2' in a for r in periods for a in r['ranges']))

    def test_表头与适用范围分句明确时保留标签引用(self):
        from 附注更新.A合集 import direct_answer
        engine=UnifiedSemanticEngine(self.config,self.gold,self.folder/'分句标签')
        reply=('原表!B3:C4：metric_id=固定资产原值。\n表头原文位于 B2、C2：B2“期末余额”、C2“期初余额”。\n'
               'B3:B4：template_period_role=closing；period=2025-12-31。\n'
               'C3:C4：template_period_role=opening；period=2024-12-31。\n'
               '这些期间角色适用于各明细、合计。')
        answer=direct_answer(reply,engine.standard,['原表'])
        dates=[r for r in answer['dimensions'] if r['id']=='period' and r.get('label_references')]
        self.assertTrue(any(r['value']=='2025-12-31' and r['label_references']==['原表!B2'] for r in dates))
        self.assertTrue(any(r['value']=='2024-12-31' and r['label_references']==['原表!C2'] for r in dates))

    def test_共同列声明的反引号冒号及纯行号表格(self):
        from 附注更新.A合集 import direct_answer
        engine=UnifiedSemanticEngine(self.config,self.gold,self.folder/'列声明')
        reply=('B列：`period` = 2025-12-31，`template_period_role` = 期末。\n'
               'C列：`template_period_role=期初`；实际日期为2024-12-31。\n'
               '原表!B3:C4：metric_id=固定资产原值。\n'
               '|行|asset_selection|\n|---|---|\n|3|`members`：设备|\n|4|all|')
        answer=direct_answer(reply,engine.standard,['原表'])
        periods={loc:r['value'] for r in answer['dimensions'] if r['id']=='period' for loc in r['ranges']}
        self.assertEqual(periods['原表!B3'],'2025-12-31')
        self.assertEqual(periods['原表!C4'],'2024-12-31')
        rows=[r for r in answer['dimensions'] if r['id']=='asset_selection']
        self.assertTrue(any(r['ranges']==['原表!B3','原表!C3'] and r['value']=='members：设备' for r in rows))


    def test_同一需要的候选共同识别且保留各自位置(self):
        from 附注更新.正常入口识别 import group_source_ranges
        from 附注更新.A合集 import _answer_tables
        groups=[{'id':str(i),'sheet':sheet,'range':'A1:C3','data_ranges':['B3:C3'],
            'a_ranges':[{'sheet':'原表','range':span}],'subject':sheet,'metric_ids':['固定资产原值'],'list_dimensions':[]}
            for i,(sheet,span) in enumerate([('明细','B3'),('总表','B3:C3'),('附表','C3')])]
        grouped=group_source_ranges(groups)
        self.assertEqual(len(grouped),1)
        tables=_answer_tables({'metrics':[{'id':'固定资产原值','ranges':['明细!B3','总表!B3']}],
            'dimensions':[{'id':'entity','ranges':['*'],'value':'测试'}]},
            {'target_tables':grouped[0]['tables']},{})
        self.assertEqual({t['sheet']:len(t['metrics']) for t in tables},{'附表':0,'总表':1,'明细':1})
        self.assertTrue(all(t['data_ranges']==['B3:C3'] for t in tables))

    def test_正常入口自动收窄并把换期和空白接到副本(self):
        original={p:file_hash(p) for p in (self.folder/'A.xlsx',self.folder/'B.xlsx',self.gold)}
        calls=[]
        def answer(engine,prompt,payload):
            material=payload['materials'];calls.append(prompt)
            self.assertNotIn('返回JSON',prompt)
            if '整份原文.json' in material:
                self.assertEqual(material['明确更新要求.json'],self.config['update_purpose'])
                return '| 工作表 | 完整范围 | 数据范围 | 标准编号 | 对应A范围 |\n|---|---|---|---|---|\n|原表|A1:C3|B3:C3|固定资产原值|原表!A1:C3|'
            self.assertIn('本次明确更新要求.json',material)
            self.assertEqual(material['用户明确的比较排除.json'],self.config['comparison_requirements'])
            source=material.get('B原文.json',material.get('原文.json'))
            self.assertIn('原表!A1:C3',[r['sheet']+'!'+r['range'] for r in source['本次读取范围']])
            if 'A需要.json' in material:
                self.assertEqual(material['A需要.json']['载体'],'A')
                self.assertEqual(material['A需要.json']['目标表区'],[{'sheet':'原表','range':'A1:C3'}])
            year=2025 if 'A需要.json' in material else 2024
            reply='\n'.join(f'原表!{col}3：metric_id=固定资产原值；entity=合成企业；report_scope=standalone；currency=CNY；unit_scale=1；asset_selection={{"mode":"members","members":["设备"]}}；period={{"kind":"instant","date":"{date}-12-31"}}；template_period_role={role}；label_references={{"period":["原表!{col}2"]}}。'
                            for col,date,role in [('B',year,'closing'),('C',year-1,'opening')])
            return reply
        config=dict(self.config)
        with patch.object(UnifiedSemanticEngine,'_http',answer), \
             patch('附注更新.正常入口识别.answer_relations',side_effect=AssertionError('逐格识别不得另生成A/B对应')):
            for step in ('recognize_a','recognize_b','match','write'):
                result=run_step(step,config);self.assertEqual(result['status'],'complete',result)
                config.update(result['files'])
        self.assertEqual(len(calls),4)
        book=load_workbook(config['a_prime']);ws=book.active
        self.assertEqual(ws['B3'].value,20);self.assertIsNone(ws['C3'].value)
        # 未识别表头自身身份时，不从数据格的引用推导新日期。
        self.assertEqual(ws['B2'].value,'2024年末');book.close()
        plan=json.loads(Path(config['update_plan']).read_text(encoding='utf-8-sig'))
        self.assertEqual(plan['comparison_requirements'],self.config['comparison_requirements'])
        for key in ('a_mapping','b_mapping'):
            collection=json.loads(Path(config[key]).read_text(encoding='utf-8-sig'))
            self.assertEqual(collection['已确认关系'],[])
        self.assertEqual(original,{p:file_hash(p) for p in original})
        from 附注更新.A合集 import direct_digest
        prior=json.loads(Path(config['b_mapping']).read_text(encoding='utf-8-sig'))
        prior['封存校验']=direct_digest(prior)
        old_path=self.folder/'逐格身份复用.json'
        old_path.write_text(json.dumps(prior,ensure_ascii=False),encoding='utf-8')
        with patch.object(UnifiedSemanticEngine,'_http',side_effect=AssertionError('已有原答复不应重问')):
            replayed=run_step('recognize_b',{**config,'b_mapping':str(old_path)})
        self.assertEqual(replayed['status'],'complete',replayed)
        replayed_collection=json.loads(Path(replayed['files']['b_mapping']).read_text(encoding='utf-8-sig'))
        self.assertEqual(replayed_collection['已确认关系'],[])
        with patch.object(UnifiedSemanticEngine,'_http',answer):
            scoped=run_step('recognize_b',{**config,'b_source_ranges':[{'sheet':'原表','range':'A1:C3'}],
                'a_required_ranges':[{'sheet':'原表','range':'A1:C3'}],'b_allowed_metrics':['固定资产原值'],
                'b_field_scope':{'原表!B3':['metric_id']}})
        self.assertEqual(scoped['status'],'complete',scoped)
        with patch.object(UnifiedSemanticEngine,'_http',side_effect=AssertionError('已有输入与答复不应重问')):
            again=run_step('recognize_a',config)
            self.assertEqual(again['status'],'complete',again)
            config.update(again['files'])
            again=run_step('recognize_b',config)
            self.assertEqual(again['status'],'complete',again)
        from 附注更新.A合集 import save_direct_answer, _records
        a=json.loads(Path(config['a_mapping']).read_text(encoding='utf-8-sig'))
        corrected=save_direct_answer(a,'原表!B3：entity=已明确的新主体。',self.folder/'A纠正.json',field_scope={'原表!B3':['entity']})
        config['a_mapping']=corrected['mapping_path']
        with patch.object(UnifiedSemanticEngine,'_http',side_effect=AssertionError('已有纠正不应重问')):
            again=run_step('recognize_a',config)
        self.assertEqual(again['status'],'complete',again)
        restored=json.loads(Path(again['files']['a_mapping']).read_text(encoding='utf-8-sig'))
        self.assertEqual(dict(_records(restored))['原表!B3']['dimensions']['entity'],'已明确的新主体')
        b=json.loads(Path(config['b_mapping']).read_text(encoding='utf-8-sig'))
        prior=dict(_records(b))['原表!C3']['dimensions']['period']
        reply='原表!B3：period={"kind":"instant","date":"2026-12-31"}；template_period_role=closing。'
        saved=save_direct_answer(b,reply,self.folder/'B共同日期.json',field_scope={'原表!B3:C3':['period']})
        self.assertEqual(dict(_records(saved))['原表!B3']['dimensions']['period'],{'kind':'instant','date':'2026-12-31'})
        self.assertEqual(dict(_records(saved))['原表!C3']['dimensions']['period'],prior)

    def test_分段名单沿用已声明列对应和原样式行(self):
        from 附注更新.A合集 import _declared_list_tables
        a={'机械记录':{'names':{'资产表':[{'sheet':'原表','range':'B10:F18'}]}}}
        b={'合集来源':[{'识别答复':'来源表。A 表中项目列 B 对应 B 表项目列 A；A 表金额列 C、D、E、F 分别对应 B 表金额列 B、C、D、E。'}]}
        reply=('B4 为合计，对应 A11；B5:B8 的成员依次对应 A12:A15。\n'
               'B14 为合计，对应扩展后的 A16；B15:B18 依次对应 A17:A20。\n'
               '新增明细沿用 A 原样式行（原 A12、A16）。')
        tables=_declared_list_tables(a,reply,'asset_selection',[{'sheet':'来源表','range':'A1:E20'}],b)
        self.assertEqual([t['old_detail_count'] for t in tables],[3,3])
        self.assertEqual([t['first_row_offset'] for t in tables],[2,6])
        self.assertEqual(tables[1]['totals'],[{'target_row':15,'source_row':14}])
        self.assertEqual(tables[0]['columns'],{'C':'B','D':'C','E':'D','F':'E'})
        self.assertEqual(tables[0]['labels'],{'B':'A'})
        reply=('|分组|合计位置|B行 → A行（成员名称）|金额列搬运|\n'
               '|原值|11|B6→12 设备；B5→13 土地；B7→14 其他；B8→15 新设备|期初 B→C，期末 E→F|\n'
               '|折旧|16|B16→17 设备；B15→18 土地；B17→19 其他；B18→20 新设备|期初 B→C，期末 E→F|\n'
               '分别以12、16行的样式为参照。')
        tables=_declared_list_tables(a,reply,'asset_selection',[{'sheet':'来源表','range':'A1:E20'}],b)
        self.assertEqual(tables[0]['source_rows'],[6,5,7,8])
        self.assertEqual(tables[0]['declared_labels']['5'],{'value':'土地'})
        self.assertEqual(tables[0]['columns'],{'C':'B','F':'E'})

    def test_名单更新段落复用明确的范围和列对应(self):
        from 附注更新.A合集 import _declared_list_tables
        a={'机械记录':{'names':{'名单':[{'sheet':'原表','range':'A2:C5'}]}}}
        reply=('采用来源：B 来源表，不采用汇总表。\nB4:C4“新增成员”；B5:C5“原成员”；B6:C6“其他”。\n'
               'A原有明细为 A!A3:A4，合计为 A!A5；B实际项目名单更新为上述第4–6行，合计为第9行。\n'
               '金额列对应 A!B3:C3 → 来源表!B5:C5。')
        tables=_declared_list_tables(a,reply,'asset_selection',[{'sheet':'来源表','range':'A1:C9'},{'sheet':'汇总表','range':'A1:C9'}])
        self.assertEqual(tables[0]['source_rows'],[4,5,6])
        self.assertEqual(tables[0]['old_detail_count'],2)
        self.assertEqual(tables[0]['labels'],{'A':None})
        self.assertEqual(tables[0]['declared_labels']['4'],{'value':'新增成员'})
        self.assertEqual(tables[0]['columns'],{'B':'B','C':'C'})

    def test_名单箭头前后的表简称及指定样式原样接入(self):
        from 附注更新.A合集 import _declared_list_tables
        a={'机械记录':{'names':{'名单':[{'sheet':'原表','range':'A2:C5'}]}},
           '合集来源':[{'识别答复':'| A5 | 行项目“合计” | all |'}]}
        reply=('原合计行顺延；新增成员沿用 A 表明细行 4 的样式。\n'
               '|名称|名称来源|本期金额搬运|上期金额搬运|\n|---|---|---|---|\n'
               '|成员一|B 表 A4 → A 表 A3|B 表 B4 → A 表 B3|B 表 C4 → A 表 C3|\n'
               '|成员二|B 表 A5 → A 表 A4|B 表 B5 → A 表 B4|B 表 C5 → A 表 C4|\n'
               '|新增成员|B 表 A6 → A 表 A5|B 表 B6 → A 表 B5|B 表 C6 → A 表 C5|\n'
               '|合计|B 表 A7 → A 表 A6|B 表 B7 → A 表 B6|B 表 C7 → A 表 C6|')
        table=_declared_list_tables(a,reply,'asset_selection',[{'sheet':'来源','range':'A1:C7'}])[0]
        self.assertEqual(table['old_detail_count'],2)
        self.assertEqual(table['labels'],{'A':'A'})
        self.assertEqual(table['operations'],[{'axis':'row','action':'insert','index':3,'count':1,'template_index':2}])


    def test_答复用原文标题引用唯一来源页及后续省略页名(self):
        from 附注更新.A合集 import source_relation
        a={'facts':[],'slot_bindings':[],'source_hash':'a','gold_hash':'g','机械记录':{'sheets':[{'name':'原表'}]}}
        b={'source_hash':'b','机械记录':{'sheets':[{'name':'来源2','cells':{'A1':{'cached_value':'来源续表'}}}]},
           'target_tables':[{'sheet':'来源2'}]}
        reply=('B来源页名为“来源续表”。\n|A需要项目（A坐标）|B采用来源（B坐标）|\n|---|---|\n'
               '|项目一（C10、D10）|来源续表!C6、D6；“项目一”|\n'
               '|项目二（C11、D11）|!C8、D8；“项目二”|')
        targets=['原表!C10','原表!D10','原表!C11','原表!D11']
        result=source_relation(a,b,reply,targets)
        self.assertEqual(result['bindings'],[{'target':t,'source':'来源2!'+s}
            for t,s in zip(targets,['C6','D6','C8','D8'])])
        held=source_relation(a,b,reply.replace('“项目一”','本次不扩展搬运范围').replace('“项目二”','A原值保留'),targets)
        self.assertEqual(held['bindings'],[])
        self.assertEqual(held['retained_targets'],targets)
        b['机械记录']['sheets'].append({'name':'另一来源','cells':{'A1':{'cached_value':'来源续表'}}})
        b['target_tables'].append({'sheet':'另一来源'})
        self.assertEqual(source_relation(a,b,reply,targets)['bindings'],[])

    def test_来源竖向范围明确对应目标横向范围(self):
        from 附注更新.A合集 import source_relation
        a={'facts':[],'slot_bindings':[],'source_hash':'a','gold_hash':'g','机械记录':{'sheets':[{'name':'原表'}]}}
        b={'source_hash':'b','机械记录':{'sheets':[{'name':'来源表'}]},'target_tables':[{'sheet':'来源表'}]}
        reply='采用D3:D6的期初、本期增加、本期减少、期末角色对应A明细表的D9:G9期间角色。'
        result=source_relation(a,b,reply,['原表!'+c+'9' for c in 'DEFG'])
        self.assertEqual(result['bindings'],[{'target':'原表!'+c+'9','source':'来源表!D'+str(n)}
            for c,n in zip('DEFG',range(3,7))])

    def test_数字行序名单同时接入两组来源与旧成员保留(self):
        from 附注更新.A合集 import _declared_list_tables
        a={'机械记录':{'names':{'名单':[{'sheet':'原表','range':'A1:E8'}]},
            'sheets':[{'name':'原表','cells':{'A3':{'value':'旧甲'},'A6':{'value':'旧乙'}}}]}}
        reply=('样式沿用A主营明细行3、其他业务明细行5–7。\n'
            '| A位置 | 名称安排 | 四个金额列来源 |\n|---|---|---|\n'
            '| 2 | 小计，保留原位 | 原小计 |\n'
            '| 3–4 | 依次列B 来源 A10–A11 的两项名称 | A的B←B的B、A的C←B的C |\n'
            '| 5 | 保留A原“旧甲”独立明细及其原金额 | 无新来源 |\n'
            '| 6 | 小计，沿用A原小计样式行4 | 原小计 |\n'
            '| 7–10 | 依次列B 来源 A20–A23 的四项名称 | 同上 |\n'
            '| 11 | 保留A原“旧乙”独立明细及其原内容 | 无新来源 |\n'
            '| 12 | 合计，沿用A原合计样式行8 | 原合计 |')
        tables=_declared_list_tables(a,reply,'asset_selection',[{'sheet':'来源','range':'A1:E24'}])
        self.assertEqual([t['source_rows'] for t in tables],[[10,11],[20,21,22,23]])
        self.assertEqual([t['old_detail_count'] for t in tables],[1,3])
        self.assertEqual([t['operations'][0]['count'] for t in tables],[2,2])
        self.assertEqual(tables[0]['retained_members'][0]['cells']['B'],{'source':'原表!B3'})
        self.assertEqual(tables[1]['retained_members'][0]['index'],4)
        self.assertEqual(tables[0]['blank_values']['10']['B'],{'value':None})

    def test_分行标签替换和明确保留范围直接接入(self):
        from 附注更新.A合集 import source_relation
        a={'facts':[],'source_hash':'a','gold_hash':'g','机械记录':{'sheets':[{'name':'原表'}]}}
        b={'source_hash':'b','机械记录':{'sheets':[{'name':'来源'}]}}
        reply=('- TA_001，原表!C7:D7\n原文：C“本期”；D“上期”。\n替换为：C“本期（本年）”；D“上期（上年）”。\n'
               '原表!C8:D9均保留原空白，暂不取数。')
        result=source_relation(a,b,reply,['原表!C8','原表!D8','原表!C9','原表!D9'])
        self.assertEqual({r['target']:r['value'] for r in result['label_actions']},
                         {'原表!C7':'本期（本年）','原表!D7':'上期（上年）'})
        self.assertEqual(set(result['retained_targets']),{'原表!C8','原表!D8','原表!C9','原表!D9'})
        a['机械记录']['sheets'][0]['merges']=['C7:D7']
        merged=source_relation(a,b,'更新原表!C7:D7合并表头为“本期新标签”；子表头保留。',['原表!C8'])
        self.assertEqual(len(merged['label_actions']),1)
        self.assertEqual(merged['label_actions'][0]['target'],'原表!C7')
        self.assertEqual(merged['label_actions'][0]['value'],'本期新标签')


    def test_带数字名称的明说行关系优先于表区顺序(self):
        from 附注更新.A合集 import source_relation
        a={'facts':[],'slot_bindings':[],'source_hash':'a','gold_hash':'g','机械记录':{'sheets':[{'name':'原表'}]}}
        b={'source_hash':'b','机械记录':{'sheets':[{'name':'来源'}]}}
        text=('A 原表!C3:D4（本组列）对应B 来源!B10:C11。\n'
              'A第3行“3个月内” ↔ B第11行“3个月内”\nA第4行“4个月后” ↔ B第10行“4个月后”\n'
              'A列C/D分别对应B列B/C。')
        result=source_relation(a,b,text,['原表!'+c+str(r) for r in [3,4] for c in 'CD'])
        self.assertEqual({x['target']:x['source'] for x in result['bindings']},
                         {'原表!C3':'来源!B11','原表!D3':'来源!C11','原表!C4':'来源!B10','原表!D4':'来源!C10'})
        grouped=source_relation(a,b,'A第3—4行依次为项目一、项目二；分别对应B第10—11行。A列C/D分别对应B列B/C。',
                                ['原表!C3','原表!D3','原表!C4','原表!D4'])
        self.assertEqual(grouped['bindings'],[])
        b['target_tables']=[{'sheet':'来源'}]
        grouped=source_relation(a,b,'A第3—4行依次为项目一、项目二；分别对应B第10—11行。A列C/D分别对应B列B/C。',
                                ['原表!C3','原表!D3','原表!C4','原表!D4'])
        self.assertEqual(len(grouped['bindings']),4)

    def test_具名明细后新增空名称行按明说范围扩展(self):
        from 附注更新.A合集 import declared_named_row_relations
        a={'机械记录':{'names':{'表':[{'sheet':'原表','range':'B1:D5'}]},
            'sheets':[{'name':'原表','cells':{'B5':{'value':'合计'}}}]}}
        text=('A 原表!B3:D4 的具名行分别对应B 来源!A10:C11。\n'
              '在A表原合计行之前新增一条名称保持空白的业务行，对应B 来源!A12:C12。A原合计行对应B第13行。')
        _,tables=declared_named_row_relations(a,{'机械记录':{'sheets':[]}},text)
        self.assertEqual(tables[0]['source_rows'],[10,11,12])
        self.assertEqual(tables[0]['columns'],{'C':'B','D':'C'})
        self.assertEqual(tables[0]['totals'],[{'target_row':5,'source_row':13}])
        self.assertEqual(tables[0]['operations'][0]['index'],4)

    def test_编号顺序名单保留名称文字并展开两列来源(self):
        from 附注更新.A合集 import declared_named_row_relations
        a={'机械记录':{'names':{'表':[{'sheet':'原表','range':'B1:D4'}]},'sheets':[]}}
        reply=('A原表 原表!B1:D4 中，原有明细为 B2 甲、B3 乙，合计为 B4；期间列为 C期末余额、D期初余额。名单随B更新后，排布为：\n'
               '1. 甲：期末对应 来源!F10，期初对应 来源!H10。\n'
               '2. 乙：期末对应 来源!F11，期初对应 来源!H11。\n'
               '3. 新项：期末对应 来源!F12，期初对应 来源!H12。\n'
               '4. 合计：期末对应 来源!F13，期初对应 来源!H13。')
        _,tables=declared_named_row_relations(a,{'机械记录':{'sheets':[]}},reply)
        self.assertEqual(tables[0]['columns'],{'C':'F','D':'H'})
        self.assertEqual(tables[0]['source_rows'],[10,11,12])
        self.assertEqual(tables[0]['declared_labels']['12'],{'value':'新项'})
        self.assertEqual(tables[0]['operations'][0]['count'],1)

    def test_明确表旁说明按唯一表区接入且保存沿用范围过滤(self):
        from 附注更新.A合集 import source_relation,save_source_relations
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'names':{'TA_001':[{'sheet':'原表','range':'B2:D4'}]},
            'sheets':[{'name':'原表','merges':[]}]},'source_context':[{'range_name':'TA_001','title':'余额表'}]}
        b={'source_hash':'b','facts':[],'机械记录':{'sheets':[{'name':'来源','merges':[]}]}}
        raw='### 余额表（原表!B2:D4）\n如输出副本需在本表旁说明换期，可用以下完整文字；不改其他表：\n“本表使用已明确的新期间，原日期另存。”'
        parsed=source_relation(a,b,raw,['原表!C3'])
        self.assertEqual(parsed['caption_updates'],[{'range_name':'TA_001','old_text':'余额表','text':'余额表 本表使用已明确的新期间，原日期另存。'}])
        local='A表头 原表!C2、原表!D2 的“期末余额”“期初余额”角色保留。仅就A本表适用的期间说明，将“期末为旧日期”替换为“期末为新日期”；不据此修改其他表的说明。'
        self.assertEqual(source_relation(a,b,local,['原表!C3'])['caption_updates'],
                         [{'range_name':'TA_001','old_text':'余额表','text':'余额表 期末为新日期'}])
        with self.assertRaisesRegex(ValueError,'旧位置取数安排入口已停用'):
            save_source_relations(a,b,raw,['原表!C3'],self.folder/'不应保存旧安排.json')

    def test_明说来源及目标之间的公式引用不混为搬运来源(self):
        from 附注更新.A合集 import source_relation
        a={'source_hash':'a','gold_hash':'g','facts':[],'机械记录':{'sheets':[{'name':'原表'}]}}
        b={'source_hash':'b','机械记录':{'sheets':[{'name':'来源'}]},'b_source_ranges':[{'sheet':'来源','range':'A1:G20'}]}
        text='- B F10:G10 对应B汇总表 C50:I50：所列项目。对应A C3:D3。'
        result=source_relation(a,b,text,['原表!C3','原表!D3'])
        self.assertEqual(result['bindings'],[{'target':'原表!C3','source':'来源!F10'},{'target':'原表!D3','source':'来源!G10'}])

if __name__=='__main__':unittest.main()
