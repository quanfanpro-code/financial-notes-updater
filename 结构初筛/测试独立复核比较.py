# -*- coding: utf-8 -*-
"""独立复核发现的错漏：固定预期，使用真实比较函数；纯内存，无删除。"""
import unittest
from datetime import datetime
from copy import deepcopy
from collections import defaultdict
from openpyxl.utils import range_boundaries
from 比较新表 import 比较工作簿, _生成交接说明, _候选问题


def 参照(编号,名称,项目):
    内容={'A1':名称,'A2':'项目','C2':'期末余额','D2':'期初余额',
          'A3':项目[0],'C3':10,'D3':9,'A4':项目[1],'C4':20,'D4':19,'A5':'合计','C5':30,'D5':28}
    节点=[]
    for 坐标,原文 in 内容.items():
        列,行,_,_=range_boundaries(坐标)
        角色=('标题' if 行==1 else '表头' if 行==2 else '项目') if isinstance(原文,str) else '业务数值'
        n={'坐标':坐标,'行':行,'列':列,'原文':原文,'角色':角色,'查找文字':原文 if isinstance(原文,str) else None}
        if 角色=='业务数值':n['语义引用']={'representation_id':编号+':'+坐标,'semantic_id':编号+':'+坐标}
        节点.append(n)
    return {'小表编号':编号,'小表名称':名称+'明细','样式':1,'来源':{'工作表':名称},
            '节点':节点,'项目顺序':[],'合并':[],'上下文':[],'表头':[],
            '上层语义':{'完整':True,'业务路径':[{'附注':名称,'表义':名称+'明细'}],
                        '名称依据':[{'原文':名称}]}},内容


def 索引(记录,缺口=()):
    正文=defaultdict(list)
    for r in 记录:
        for n in r['节点']:
            if n['查找文字']:正文[n['查找文字']].append({'小表编号':r['小表编号'],'坐标':n['坐标'],'样式':1})
    return {'目录':'受控参照','规则版本':'受控参照','记录':记录,'小表':{r['小表编号']:r for r in 记录},
            '正文':dict(正文),'别名':{},'索引清单':{'缺口':list(缺口)}}


def 工作簿(名称,内容):
    cells={}
    for pos,value in 内容.items():
        c,r,_,_=range_boundaries(pos)
        cells[pos]={'row':r,'column':c,'value':value,'cached_value':value,'formula':None}
    return {'path':'受控新表.xlsx','sha256':'0'*64,'sheets':[{'name':名称,'cells':cells,
            'max_row':max(x['row'] for x in cells.values()),'max_column':max(x['column'] for x in cells.values()),'merges':[]}]}


class 独立复核比较(unittest.TestCase):
    def setUp(self):
        self.现金,self.内容=参照('cash','货币资金',['库存现金','银行存款'])
        self.应收,self.应收内容=参照('receivable','应收账款',['甲客户','乙客户'])

    def test_候选内部新增说明必须交接位置和原文(self):
        self.内容['B3']='其中：受限款项'
        结果=比较工作簿(工作簿('货币资金',self.内容),索引([self.现金]))
        self.assertTrue(_候选问题(结果['候选'][0]))
        说明=_生成交接说明(结果)
        self.assertNotIn('本次没有未解决问题',说明)
        self.assertIn('B3',说明)
        self.assertIn('其中：受限款项',说明)

    def test_候选内部未知数字不能自动当成已解释金额(self):
        self.内容['B3']=2026
        结果=比较工作簿(工作簿('货币资金',self.内容),索引([self.现金]))
        self.assertTrue(_候选问题(结果['候选'][0]))
        self.assertIn('B3',_生成交接说明(结果))

    def test_候选内部无缓存公式完整交接且不执行(self):
        数据=工作簿('货币资金',self.内容)
        数据['sheets'][0]['cells']['B3']={'row':3,'column':2,'value':'="其中：受限款项"',
            'formula':'="其中：受限款项"','cached_value':None}
        结果=比较工作簿(数据,索引([self.现金]))
        未对应=结果['候选'][0].get('未对应内容',[])
        公式=next((r for r in 未对应 if r['坐标']=='B3'),None)
        self.assertIsNotNone(公式)
        self.assertEqual(公式['公式'],'="其中：受限款项"')
        self.assertIsNone(公式['缓存'])
        self.assertIn('B3',_生成交接说明(结果))

    def test_候选内部日期与布尔值不消失(self):
        for 值 in (datetime(2026,12,31),True):
            with self.subTest(值=值):
                内容={**self.内容,'B3':值}
                结果=比较工作簿(工作簿('货币资金',内容),索引([self.现金]))
                未对应=结果['候选'][0].get('未对应内容',[])
                self.assertEqual([r['原文'] for r in 未对应 if r['坐标']=='B3'],[值])
                self.assertIn('B3',_生成交接说明(结果))

    def test_已对应业务格改为无缓存公式不增加识别任务(self):
        数据=工作簿('货币资金',self.内容)
        数据['sheets'][0]['cells']['C3'].update(value='=1+2',formula='=1+2',cached_value=None)
        结果=比较工作簿(数据,索引([self.现金]))
        self.assertEqual(结果['候选'][0].get('未对应内容',[]),[])
        self.assertEqual(_候选问题(结果['候选'][0]),[])

    def test_缺标题小表不能只因沿用旧工作表名而漏掉(self):
        del self.应收内容['A1']
        结果=比较工作簿(工作簿('货币资金',self.应收内容),索引([self.现金,self.应收]))
        候选=[c for c in 结果['候选'] if c['旧表编号']=='receivable']
        self.assertTrue(候选)
        self.assertTrue(候选[0]['上层语义核对']['待确认原因'])

    def test_同名项目的不同具体对应必须保留(self):
        self.内容.update({'A5':'银行存款','C5':7,'D5':6,'A6':'合计','C6':37,'D6':34})
        结果=比较工作簿(工作簿('货币资金',self.内容),索引([self.现金]))
        新位置={d['新坐标'] for c in 结果['候选'] for d in c['结构对应'] if d['旧坐标']=='C4'}
        self.assertTrue({'C4','C5'} <= 新位置)
        self.assertTrue(all(c['歧义'] for c in 结果['候选']))

    def test_仅一个明确项目加合计也不能被旧表名排除(self):
        self.应收['节点']=[n for n in self.应收['节点'] if n['行']!=4]
        内容={p:v for p,v in self.应收内容.items() if p not in ('A1','A4','C4','D4')}
        结果=比较工作簿(工作簿('货币资金',内容),索引([self.现金,self.应收]))
        self.assertTrue([c for c in 结果['候选'] if c['旧表编号']=='receivable'])

    def test_同一完整对应不因种子重复而膨胀(self):
        结果=比较工作簿(工作簿('货币资金',self.内容),索引([self.现金]))
        self.assertEqual(len(结果['候选']),1)
        self.assertEqual(_候选问题(结果['候选'][0]),[])

    def test_旧索引小表缺口传到候选和交接(self):
        缺口={'样式':1,'小表':'货币资金明细','坐标':'C列','问题':'该业务列缺表头归属依据'}
        结果=比较工作簿(工作簿('货币资金',self.内容),索引([self.现金],[缺口]))
        self.assertTrue(结果['缺口'])
        候选=结果['候选'][0]
        self.assertIn(缺口,候选.get('旧索引缺口',[]))
        说明=_生成交接说明(结果)
        self.assertIn('C列',说明)
        self.assertIn('该业务列缺表头归属依据',说明)
        业务={d['旧坐标'] for d in 候选['结构对应'] if d['角色']=='业务数值'}
        self.assertNotIn('C3',业务)
        self.assertIn('D3',业务)

    def test_不能定位小表的来源缺口仍须全局交接(self):
        缺口={'样式':1,'小表':None,'问题':'源文件中的一处结构归属缺依据'}
        结果=比较工作簿(工作簿('货币资金',self.内容),索引([self.现金],[缺口]))
        self.assertTrue(结果['缺口'])
        self.assertIn('源文件中的一处结构归属缺依据',_生成交接说明(结果))
        self.assertNotIn('本次没有未解决问题',_生成交接说明(结果))

    def test_表头冲突只阻止受影响业务列且保留原内容(self):
        self.现金['表头']=[{'坐标':p,'覆盖列':p[0],'查找文字':self.内容[p]}
                           for p in ('C2','D2')]
        self.内容['C2']='本期增加'
        结果=比较工作簿(工作簿('货币资金',self.内容),索引([self.现金]))
        self.assertTrue(结果['候选'])
        对应=[d for c in 结果['候选'] for d in c['结构对应'] if d['角色']=='业务数值']
        self.assertFalse([d for d in 对应 if d['旧坐标'].startswith('C')])
        self.assertIn(('D3','D3'),[(d['旧坐标'],d['新坐标']) for d in 对应])
        self.assertTrue(any(d['坐标']=='C3' and d['原文']==10
                            for c in 结果['候选'] for d in c['未对应内容']))

    def test_合并期间表头冲突不能被下层相同表头放过(self):
        内容={'A1':'存货','A2':'项目','C2':'期末余额','C3':'账面余额','D3':'跌价准备',
              'A4':'原材料','C4':10,'D4':2,'A5':'产成品','C5':20,'D5':3}
        旧=deepcopy(self.现金);旧['节点']=[];旧['合并']=['C2:D2'];旧['表头']=[]
        for p,v in 内容.items():
            c,r,_,_=range_boundaries(p)
            角色=('表头' if c>=3 else '标题' if r==1 else '项目') if isinstance(v,str) else '业务数值'
            旧['节点'].append({'坐标':p,'列':c,'行':r,'原文':v,'角色':角色,
                              '查找文字':v if isinstance(v,str) else None})
        for p,col in [('C2','C'),('C2','D'),('C3','C'),('D3','D')]:
            旧['表头'].append({'坐标':p,'覆盖列':col,'查找文字':内容[p]})
        for 情形 in ('相同','期间冲突','下层列互换','并排期间'):
            with self.subTest(情形=情形):
                新=deepcopy(内容)
                if 情形=='期间冲突':新['C2']='期初余额'
                if 情形=='下层列互换':
                    for row in (3,4,5):新[f'C{row}'],新[f'D{row}']=新[f'D{row}'],新[f'C{row}']
                数据=工作簿('货币资金',新);数据['sheets'][0]['merges']=['C2:D2']
                if 情形=='并排期间':
                    新['C2']='期初余额';新['F2']='期末余额'
                    for row in (3,4,5):新[f'F{row}'],新[f'G{row}']=新[f'C{row}'],新[f'D{row}']
                    数据=工作簿('货币资金',新);数据['sheets'][0]['merges']=['C2:D2','F2:G2']
                结果=比较工作簿(数据,索引([旧]));self.assertTrue(结果['候选'])
                对应={(d['旧坐标'],d['新坐标']) for c in 结果['候选'] for d in c['结构对应'] if d['角色']=='业务数值'}
                if 情形=='期间冲突':self.assertEqual(对应,set())
                elif 情形=='并排期间':self.assertEqual(对应,{('C4','F4'),('D4','G4'),('C5','F5'),('D5','G5')})
                else:
                    self.assertEqual(对应,{('C4','D4' if 情形=='下层列互换' else 'C4'),
                                          ('D4','C4' if 情形=='下层列互换' else 'D4'),
                                          ('C5','D5' if 情形=='下层列互换' else 'C5'),
                                          ('D5','C5' if 情形=='下层列互换' else 'D5')})

    def test_并排独立表不能混用项目列和期间列(self):
        from openpyxl.utils import get_column_letter
        for 左列,间距 in ((1,6),(4,9)):
            with self.subTest(左列=左列,间距=间距):
                def 地址(c,r):return f'{get_column_letter(c)}{r}'
                内容={};合并=[];旧=deepcopy(self.现金);旧['节点']=[];旧['表头']=[]
                for 起列,期间,项目 in ((左列,'甲期间',('甲对象','乙对象')),
                                      (左列+间距,'乙期间',('乙对象','甲对象'))):
                    局部={地址(起列,1):'工具明细',地址(起列,2):'工具类别',
                          地址(起列+1,2):期间,地址(起列+1,3):'数量',地址(起列+2,3):'价值'}
                    for r,名称 in enumerate(项目,4):
                        局部.update({地址(起列,r):名称,地址(起列+1,r):r,地址(起列+2,r):r*10})
                    内容.update(局部);合并.append(f'{地址(起列+1,2)}:{地址(起列+2,2)}')
                    if 起列!=左列:continue
                    for p,v in 局部.items():
                        c,r,_,_=range_boundaries(p)
                        角色=('表头' if c>起列 else '标题' if r==1 else '说明') if isinstance(v,str) else '业务数值'
                        旧['节点'].append({'坐标':p,'列':c,'行':r,'原文':v,'角色':角色,
                                          '查找文字':v if isinstance(v,str) else None})
                    for c in (起列+1,起列+2):
                        for p in (地址(起列+1,2),地址(c,3)):
                            旧['表头'].append({'坐标':p,'覆盖列':get_column_letter(c),'查找文字':局部[p]})
                旧['合并']=[合并[0]];数据=工作簿('工具明细',内容);数据['sheets'][0]['merges']=合并
                结果=比较工作簿(数据,索引([旧]))
                对应={(d['旧坐标'],d['新坐标']) for c in 结果['候选'] for d in c['结构对应'] if d['角色']=='业务数值'}
                self.assertEqual(对应,{(地址(c,r),地址(c,r)) for c in (左列+1,左列+2) for r in (4,5)})

    def test_不能跨过当前期间表头借用前段表头(self):
        内容={'A1':'资产','A2':'项目','C2':'期末余额','F2':'期初余额',
              'C3':'账面余额','D3':'减值准备','F3':'账面余额','G3':'减值准备',
              'A4':'甲项目','A5':'乙项目','C4':10,'D4':1,'F4':9,'G4':2,
              'C5':20,'D5':3,'F5':18,'G5':4}
        旧=deepcopy(self.现金);旧['节点']=[];旧['合并']=['C2:D2','F2:G2'];旧['表头']=[]
        for p,v in 内容.items():
            c,r,_,_=range_boundaries(p)
            角色=('表头' if c>=3 else '标题' if r==1 else '项目') if isinstance(v,str) else '业务数值'
            旧['节点'].append({'坐标':p,'列':c,'行':r,'原文':v,'角色':角色,
                              '查找文字':v if isinstance(v,str) else None})
        for p,col in [('C2','C'),('C2','D'),('F2','F'),('F2','G'),
                      ('C3','C'),('D3','D'),('F3','F'),('G3','G')]:
            旧['表头'].append({'坐标':p,'覆盖列':col,'查找文字':内容[p]})
        新={}
        for p,v in 内容.items():
            c,r,_,_=range_boundaries(p);新[p if r<3 else p[:-1]+str(r+3)]=v
        新.update({'A5':'项目','C5':'期初余额','F5':'期末余额'})
        数据=工作簿('货币资金',新)
        数据['sheets'][0]['merges']=['C2:D2','F2:G2','C5:D5','F5:G5']
        结果=比较工作簿(数据,索引([旧]))
        对应={(d['旧坐标'],d['新坐标']) for c in 结果['候选'] for d in c['结构对应'] if d['角色']=='业务数值'}
        self.assertEqual(对应,{(old+str(r),new+str(r+3)) for old,new in [('C','F'),('D','G'),('F','C'),('G','D')] for r in (4,5)})

if __name__=='__main__':unittest.main(verbosity=2)
