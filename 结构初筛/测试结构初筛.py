# -*- coding: utf-8 -*-
"""结构初筛受控测试。用法：python 测试结构初筛.py --组 输入|索引|查找|比较|交接|增量|全部

每个测试组先构造受控输入，再断言真实行为；预期值在运行被测程序前由构造规则固定。
"""
import argparse
import hashlib
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

本批目录 = Path(__file__).resolve().parent
sys.path.insert(0, str(本批目录))

成果库 = Path(r"C:\Users\27651\Desktop\附注自动更新系统\语义成果整理")
五样式目录 = Path(r"D:\BaiduSyncdisk\数字资产\审计数据自动搬运系统\附注样式示例")


def 指纹(路径):
    h = hashlib.sha256()
    with open(路径, "rb") as f:
        for 块 in iter(lambda: f.read(1 << 20), b""):
            h.update(块)
    return h.hexdigest()


class 输入组(unittest.TestCase):
    """清点与建库前的输入核验：缺失、指纹不符、同名冲突必须报明原因。"""

    def setUp(self):
        from 建立结构索引 import 清点
        self.临时 = Path(tempfile.mkdtemp(prefix="初筛输入测试_"))
        self.addCleanup(shutil.rmtree, self.临时, True)
        self.清单路径 = self.临时 / "输入清单.json"
        清点(str(成果库), str(五样式目录), str(self.清单路径))
        self.清单 = json.loads(self.清单路径.read_text(encoding="utf-8-sig"))

    def test_正确清单可读且覆盖五样式(self):
        用途 = [e["用途"] for e in self.清单["条目"]]
        for i in range(1, 6):
            self.assertIn(f"五样式原Excel:附注样式{i}.xlsx", 用途)
        self.assertIn("统一成果:表现记录", 用途)
        self.assertIn("统一成果:资产清单", 用途)
        self.assertIn("复用代码:机械读取模块", 用途)
        映射 = [e for e in self.清单["条目"] if e["用途"].startswith("五样式原映射:")]
        self.assertGreaterEqual(len(映射), 5)
        for e in self.清单["条目"]:
            self.assertEqual(指纹(e["路径"]), e["sha256"], e["路径"])

    def test_文件缺失必须报明(self):
        from 建立结构索引 import 核对输入清单
        self.清单["条目"][0]["路径"] = str(self.临时 / "不存在.xlsx")
        with self.assertRaises(Exception) as ctx:
            核对输入清单(self.清单)
        self.assertIn("不存在", str(ctx.exception))

    def test_指纹不符必须报明(self):
        from 建立结构索引 import 核对输入清单
        条目 = self.清单["条目"][0]
        条目["sha256"] = "0" * 64
        with self.assertRaises(Exception) as ctx:
            核对输入清单(self.清单)
        self.assertIn("指纹", str(ctx.exception))
        self.assertIn(条目["路径"], str(ctx.exception))

    def test_同名来源冲突必须报明(self):
        from 建立结构索引 import 核对输入清单
        副本 = dict(self.清单["条目"][0])
        副本["用途"] = self.清单["条目"][1]["用途"]
        self.清单["条目"].append(副本)
        with self.assertRaises(Exception) as ctx:
            核对输入清单(self.清单)
        self.assertIn("冲突", str(ctx.exception))

    def test_清点前后原件指纹不变(self):
        前 = {e["路径"]: e["sha256"] for e in self.清单["条目"]}
        from 建立结构索引 import 清点
        清单2 = self.临时 / "输入清单2.json"
        清点(str(成果库), str(五样式目录), str(清单2))
        for 路径, 摘要 in 前.items():
            self.assertEqual(指纹(路径), 摘要, 路径)


_索引目录 = None


def 取得测试索引():
    """整套测试共享一次真实建库结果；每次完整运行重新建立。"""
    global _索引目录
    if _索引目录 is None:
        from 建立结构索引 import 建立索引, 清点
        临时 = Path(tempfile.mkdtemp(prefix="初筛索引测试_"))
        清单 = 临时 / "输入清单.json"
        清点(str(成果库), str(五样式目录), str(清单))
        _索引目录 = 临时 / "索引" / "首版"
        建立索引(str(清单), str(_索引目录))
    return _索引目录


def 读小表记录(索引目录):
    return [json.loads(l) for l in (Path(索引目录) / "小表记录.jsonl").read_text(encoding="utf-8-sig").splitlines() if l.strip()]


class 索引组(unittest.TestCase):
    """旧小表记录：映射去向完整、语义引用可回查、原文与位置保留。"""

    @classmethod
    def setUpClass(cls):
        cls.目录 = 取得测试索引()
        cls.记录 = 读小表记录(cls.目录)
        cls.清单 = json.loads((cls.目录 / "索引清单.json").read_text(encoding="utf-8-sig"))

    def test_五种样式都有实际小表记录(self):
        按样式 = {}
        for r in self.记录:
            按样式.setdefault(r["样式"], 0)
            按样式[r["样式"]] += 1
        for i in range(1, 6):
            self.assertGreater(按样式.get(i, 0), 0, f"样式{i}没有小表记录")

    def test_样式1银行存款期末格定位与语义回查(self):
        # 事先固定的事实链：表现记录(语义C-N087-T001-S003, mapping_id S1-REVIEW-0003)
        # → 原映射(报表格式附注（填入）!C5 银行存款/期末余额)
        目标 = None
        for r in self.记录:
            if r["样式"] != 1:
                continue
            for n in r["节点"]:
                引用 = n.get("语义引用") or {}
                if 引用.get("mapping_id") == "S1-REVIEW-0003":
                    目标 = (r, n)
        self.assertIsNotNone(目标, "找不到映射S1-REVIEW-0003对应的小表节点")
        记录, 节点 = 目标
        self.assertEqual(记录["来源"]["工作表"], "报表格式附注（填入）")
        self.assertEqual(节点["坐标"], "C5")
        self.assertEqual(节点["角色"], "业务数值")
        引用 = 节点["语义引用"]
        self.assertEqual(引用["semantic_id"], "C-N087-T001-S003")
        self.assertTrue(引用["representation_id"].startswith("rep:"))
        self.assertEqual((引用.get("维度") or {}).get("period_role"), "closing")
        # 回查统一表现记录：同一表现编号应指向同一位置
        with open(成果库 / "数据" / "表现记录.jsonl", encoding="utf-8-sig") as f:
            表现 = next(json.loads(l) for l in f if l.strip() and json.loads(l)["representation_id"] == 引用["representation_id"])
        self.assertEqual(表现["source_locator"]["cell"], "C5")
        self.assertEqual(表现["source_locator"]["sheet"], "报表格式附注（填入）")

    def test_每条原映射有去向(self):
        去向 = self.清单["映射去向"]
        总数 = sum(去向.values())
        # 事先从原映射清点总数（含结构格排除）
        预期 = 0
        for e in json.loads((Path(self.目录).parent.parent / "输入清单.json").read_text(encoding="utf-8-sig"))["条目"]:
            if e["用途"].startswith("五样式原映射:"):
                with open(e["路径"], encoding="utf-8-sig") as f:
                    预期 += sum(1 for l in f if l.strip())
        self.assertEqual(总数, 预期, f"去向合计{总数}不等于原映射总数{预期}")
        self.assertEqual(去向.get("缺口", 0), 0, "存在未说明去向的映射缺口")

    def test_文字查找保留原名与位置(self):
        查找 = json.loads((self.目录 / "文字查找.json").read_text(encoding="utf-8-sig"))
        self.assertIn("银行存款", 查找["正文"])
        条目 = [e for e in 查找["正文"]["银行存款"] if e["样式"] == 1]
        self.assertTrue(条目, "样式1应有银行存款项目文字")
        self.assertTrue(any(e["坐标"] == "B5" for e in 条目))

    def test_公式表头保留公式与缓存(self):
        # 样式2货币资金C3：公式=[1]首页!B5，缓存“2025年度”，角色依据为结构格排除
        找到 = False
        for r in self.记录:
            if r["样式"] != 2:
                continue
            for n in r["节点"]:
                if n["坐标"] == "C3" and n.get("公式") == "=[1]首页!B5":
                    self.assertEqual(n.get("缓存"), "2025年度")
                    self.assertEqual(n["角色"], "表头")
                    找到 = True
        self.assertTrue(找到, "样式2公式表头C3未保留公式与缓存")


# ============ 受控变换工具：预期由变换规则事先固定，不用被测程序输出当答案 ============

def 变换复制(源路径, 输出路径, 行移=0, 列移=0, 插入行=None, 删除行=None, 改值=None, 新表名=None):
    """把源工作簿每个工作表的内容按规则复制到新工作簿。
    插入行：{在源行号前插入: [{列字母: 值}]}；删除行：源行号集合；改值：{源坐标: 新值}。
    返回 {(源工作表, 源坐标): (新工作表, 新坐标)} 的对应规则，即测试预期。"""
    from openpyxl import load_workbook, Workbook
    from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
    插入行 = 插入行 or {}
    删除行 = set(删除行 or ())
    改值 = dict(改值 or {})
    对应 = {}

    def 行偏移(r):
        return sum(len(v) for k, v in 插入行.items() if k <= r) - sum(1 for k in 删除行 if k < r)

    w = load_workbook(源路径, data_only=False)
    新 = Workbook()
    新.remove(新.active)
    for s in w:
        名 = 新表名 or s.title
        新表 = 新.create_sheet(名)
        新表.sheet_state = s.sheet_state
        for 位置, 内容 in sorted(插入行.items()):
            目标行 = 位置 + 行移 + sum(len(v) for k, v in 插入行.items() if k < 位置) - sum(1 for k in 删除行 if k < 位置)
            for 一组 in 内容:
                for 列字母, 值 in 一组.items():
                    新表.cell(目标行, column_index_from_string(列字母)).value = 值
                目标行 += 1
        for row in s.iter_rows():
            for c in row:
                if c.row in 删除行:
                    continue
                nr, nc = c.row + 行移 + 行偏移(c.row), c.column + 列移
                if c.value is not None:
                    值 = 改值.get(c.coordinate, c.value)
                    新表.cell(nr, nc).value = 值
                # 空格也登记对应：候选里的业务空位按同一平移规则核对
                对应[(s.title, c.coordinate)] = (名, f"{get_column_letter(nc)}{nr}")
        for m in s.merged_cells.ranges:
            左, 顶, 右, 底 = m.min_col, m.min_row, m.max_col, m.max_row
            if 顶 in 删除行:
                continue
            新表.merge_cells(start_row=顶 + 行移 + 行偏移(顶), start_column=左 + 列移,
                           end_row=底 + 行移 + 行偏移(底), end_column=右 + 列移)
    w.close()
    新.save(输出路径)
    return 对应


class 查找组(unittest.TestCase):
    """完整新表不预传表区，代码自己找候选；平移、混排、重复、弱线索分别验证。"""

    @classmethod
    def setUpClass(cls):
        cls.索引目录 = 取得测试索引()
        from 比较新表 import 载入索引
        cls.索引 = 载入索引(str(cls.索引目录))
        cls.临时 = Path(tempfile.mkdtemp(prefix="初筛查找测试_"))

    def _比较(self, 新表路径):
        from 建立结构索引 import 读取工作簿
        from 比较新表 import 比较工作簿
        return 比较工作簿(读取工作簿(str(新表路径)), self.索引)

    def _找候选(self, 结果, 小表名称, 样式=1):
        return [c for c in 结果["候选"] if c["小表名称"] == 小表名称 and c["样式来源"] == 样式]

    def test_原样查回自身(self):
        结果 = self._比较(五样式目录 / "附注样式1.xlsx")
        候选 = self._找候选(结果, "1、货币资金余额情况")
        self.assertEqual(len(候选), 1)
        对应 = {(d["旧工作表"], d["旧坐标"]): (d["新工作表"], d["新坐标"]) for d in 候选[0]["结构对应"]}
        self.assertEqual(对应[("报表格式附注（填入）", "C5")], ("报表格式附注（填入）", "C5"))
        self.assertEqual(候选[0]["覆盖计算"]["旧表文字覆盖"]["分子"], 候选[0]["覆盖计算"]["旧表文字覆盖"]["分母"])

    def test_整体平移找到货币资金和其他小表(self):
        新表 = self.临时 / "平移.xlsx"
        变换复制(五样式目录 / "附注样式1.xlsx", 新表, 行移=10, 列移=3, 新表名="新表")
        结果 = self._比较(新表)
        候选 = self._找候选(结果, "1、货币资金余额情况")
        self.assertEqual(len(候选), 1)
        对应 = {(d["旧工作表"], d["旧坐标"]): (d["新工作表"], d["新坐标"]) for d in 候选[0]["结构对应"]}
        self.assertEqual(对应[("报表格式附注（填入）", "C5")], ("新表", "F15"))
        其他 = [c for c in 结果["候选"] if "应收" in c["小表名称"]]
        self.assertTrue(其他, "平移后除第一个命中外还应找到其他小表")

    def test_混排两样式各自找到参照(self):
        from openpyxl import Workbook
        记录 = 读小表记录(self.索引目录)
        样1 = next(r for r in 记录 if r["样式"] == 1 and r["小表名称"] == "1、货币资金余额情况")
        样3候选 = [r for r in 记录 if r["样式"] == 3 and "应收账款" in r["小表名称"]]
        self.assertTrue(样3候选, "样式3应有应收账款小表")
        样3 = 样3候选[0]
        新 = Workbook()
        s = 新.active
        s.title = "混合"
        for n in 样1["节点"]:
            if n.get("原文") is not None:
                s.cell(n["行"] + 5, n["列"] + 2).value = n.get("公式") or n.get("原文")
        for n in 样3["节点"]:
            if n.get("原文") is not None:
                s.cell(n["行"] + 60, n["列"] + 1).value = n.get("公式") or n.get("原文")
        路径 = self.临时 / "混排.xlsx"
        新.save(路径)
        结果 = self._比较(路径)
        self.assertTrue(self._找候选(结果, "1、货币资金余额情况", 1), "混排中样式1货币资金应找到参照")
        self.assertTrue(self._找候选(结果, 样3["小表名称"], 3), "混排中样式3应收账款应找到参照")
        self.assertIn("1", str(结果["样式命中"]))
        self.assertIn("3", str(结果["样式命中"]))

    def test_重复小表保留多个候选(self):
        from openpyxl import Workbook
        记录 = 读小表记录(self.索引目录)
        样1 = next(r for r in 记录 if r["样式"] == 1 and r["小表名称"] == "1、货币资金余额情况")
        新 = Workbook()
        s = 新.active
        s.title = "重复"
        for 偏移 in (0, 30):
            for n in 样1["节点"]:
                if n.get("原文") is not None:
                    s.cell(n["行"] + 偏移, n["列"]).value = n.get("公式") or n.get("原文")
        路径 = self.临时 / "重复.xlsx"
        新.save(路径)
        结果 = self._比较(路径)
        候选 = self._找候选(结果, "1、货币资金余额情况")
        self.assertEqual(len(候选), 2, "同一旧小表的两处重复位置都应保留为候选")

    def test_单个通用词不构成候选(self):
        from openpyxl import Workbook
        新 = Workbook()
        s = 新.active
        s.title = "弱线索"
        s["B3"] = "合计"
        s["C3"] = 100
        路径 = self.临时 / "弱线索.xlsx"
        新.save(路径)
        结果 = self._比较(路径)
        self.assertEqual(结果["候选"], [], "仅有'合计'不应声称找到任何业务对应")
        self.assertTrue(结果["未匹配范围"], "未匹配内容必须列出")

    def test_候选筛选不漏基准(self):
        新表 = self.临时 / "基准.xlsx"
        变换复制(五样式目录 / "附注样式1.xlsx", 新表, 行移=4, 列移=1, 新表名="基准表")
        from 建立结构索引 import 读取工作簿
        from 比较新表 import 比较工作簿
        记录 = 读取工作簿(str(新表))
        快速 = 比较工作簿(记录, self.索引)
        基准 = 比较工作簿(记录, self.索引, 预筛选=False)
        快速集合 = {(c["旧表编号"], c["新表范围"]) for c in 快速["候选"]}
        基准集合 = {(c["旧表编号"], c["新表范围"]) for c in 基准["候选"]}
        遗漏 = 基准集合 - 快速集合
        self.assertEqual(遗漏, set(), f"快速筛选漏掉基准候选：{list(遗漏)[:3]}")
        self.assertLess(快速["运行统计"]["比较的小表数"], 基准["运行统计"]["比较的小表数"])


class 比较组(unittest.TestCase):
    """插行、删行、数值变化、表头变化、列互换、父标题变化与错误反例。"""

    @classmethod
    def setUpClass(cls):
        cls.索引目录 = 取得测试索引()
        from 比较新表 import 载入索引
        cls.索引 = 载入索引(str(cls.索引目录))
        cls.临时 = Path(tempfile.mkdtemp(prefix="初筛比较测试_"))

    def _货币资金候选(self, 结果):
        候选 = [c for c in 结果["候选"] if c["小表名称"] == "1、货币资金余额情况" and c["样式来源"] == 1]
        self.assertEqual(len(候选), 1)
        return 候选[0]

    def _比较(self, 新表路径):
        from 建立结构索引 import 读取工作簿
        from 比较新表 import 比较工作簿
        return 比较工作簿(读取工作簿(str(新表路径)), self.索引)

    def test_插入新项目行共同部分保留(self):
        新表 = self.临时 / "插入.xlsx"
        预期 = 变换复制(五样式目录 / "附注样式1.xlsx", 新表,
                     插入行={6: [{"B": "其中：存放财务公司款项"}]}, 新表名="新表")
        候选 = self._货币资金候选(self._比较(新表))
        对应 = {(d["旧工作表"], d["旧坐标"]): (d["新工作表"], d["新坐标"]) for d in 候选["结构对应"]}
        self.assertEqual(对应[("报表格式附注（填入）", "C5")], 预期[("报表格式附注（填入）", "C5")])
        self.assertEqual(对应[("报表格式附注（填入）", "B6")], 预期[("报表格式附注（填入）", "B6")])
        self.assertEqual(候选["新增项目原文"], ["其中：存放财务公司款项"])
        self.assertEqual(候选["缺失项目原文"], [])
        self.assertTrue(any("合计" in str(x) or "父" in str(x) for x in 候选["新增影响"]),
                        "新增'其中'行对父项/合计的影响必须列明")
        新覆盖 = 候选["覆盖计算"]["新表文字覆盖"]
        self.assertEqual(新覆盖["分母"] - 新覆盖["分子"], 1)

    def test_删除旧项目行标出缺失(self):
        新表 = self.临时 / "删除.xlsx"
        预期 = 变换复制(五样式目录 / "附注样式1.xlsx", 新表, 删除行={5}, 新表名="新表")
        结果 = self._比较(新表)
        候选 = self._货币资金候选(结果)
        对应 = {(d["旧工作表"], d["旧坐标"]): (d["新工作表"], d["新坐标"]) for d in 候选["结构对应"]}
        self.assertEqual(候选["缺失项目原文"], ["银行存款"])
        self.assertEqual(对应[("报表格式附注（填入）", "B6")], 预期[("报表格式附注（填入）", "B6")])
        self.assertEqual(对应[("报表格式附注（填入）", "B4")], 预期[("报表格式附注（填入）", "B4")])
        self.assertTrue(self._找候选存续(结果, "2、货币资金年末受限制情况"), "删除一张小表的行不影响其他小表")

    def _找候选存续(self, 结果, 小表名称):
        return [c for c in 结果["候选"] if c["小表名称"] == 小表名称]

    def test_整张小表删除后其他保留(self):
        新表 = self.临时 / "删表.xlsx"
        变换复制(五样式目录 / "附注样式1.xlsx", 新表, 删除行=set(range(2, 10)), 新表名="新表")
        结果 = self._比较(新表)
        self.assertEqual([c for c in 结果["候选"] if c["小表名称"] == "1、货币资金余额情况"], [])
        self.assertTrue([c for c in 结果["候选"] if c["小表名称"] == "2、货币资金年末受限制情况"])
        self.assertTrue(结果["未匹配范围"] is not None)

    def test_仅改业务数值结构不变(self):
        新表 = self.临时 / "改数值.xlsx"
        变换复制(五样式目录 / "附注样式1.xlsx", 新表, 改值={"C5": 999999.99, "D5": 1}, 新表名="新表")
        候选 = self._货币资金候选(self._比较(新表))
        对应 = {(d["旧工作表"], d["旧坐标"]) for d in 候选["结构对应"]}
        self.assertIn(("报表格式附注（填入）", "C5"), 对应)
        self.assertEqual(候选["冲突"], [])

    def test_结构数字变化留下差异(self):
        新表 = self.临时 / "改年份.xlsx"
        变换复制(五样式目录 / "附注样式1.xlsx", 新表, 改值={"C3": "末年余额"}, 新表名="新表")
        候选 = self._货币资金候选(self._比较(新表))
        缺失 = " ".join(候选["缺失项目原文"] + [str(x) for x in 候选["差异"]])
        self.assertIn("期末余额", 缺失, "表头文字变化必须留下差异记录")

    def test_期初期末列互换按文字重新对应(self):
        from openpyxl import load_workbook
        源 = self.临时 / "列互换_准备.xlsx"
        变换复制(五样式目录 / "附注样式1.xlsx", 源, 新表名="新表")
        w = load_workbook(源)
        s = w["新表"]
        from openpyxl.cell.cell import MergedCell
        换号 = {}
        for row in range(1, s.max_row + 1):
            左, 右 = s.cell(row, 3), s.cell(row, 4)
            if isinstance(左, MergedCell) or isinstance(右, MergedCell):
                continue
            c, d = 左.value, 右.value
            if c is not None or d is not None:
                左.value, 右.value = d, c
                换号[row] = True
        w.save(源)
        w.close()
        候选 = self._货币资金候选(self._比较(源))
        对应 = {(d["旧工作表"], d["旧坐标"]): (d["新工作表"], d["新坐标"]) for d in 候选["结构对应"]}
        self.assertEqual(对应.get(("报表格式附注（填入）", "C5")), ("新表", "D5"),
                         "列互换后应按表头文字重新对应，不得仍按旧列位置套用")
        self.assertTrue(any("列" in str(x) for x in 候选["关系核对"]["变化"]),
                        "列位置变化必须在关系核对中报告")

    def test_父标题变化报告冲突(self):
        新表 = self.临时 / "改父标题.xlsx"
        变换复制(五样式目录 / "附注样式1.xlsx", 新表, 改值={"B1": "其中：受限资金"}, 新表名="新表")
        候选 = self._货币资金候选(self._比较(新表))
        self.assertIn("上级标题变化", str(候选["上下文核对"]["冲突原因"]))

    def test_交换对应被预期对照检出(self):
        from 比较新表 import 核验预期对应
        新表 = self.临时 / "对照.xlsx"
        预期 = 变换复制(五样式目录 / "附注样式1.xlsx", 新表, 新表名="新表")
        结果 = self._比较(新表)
        核验预期对应(结果, 预期, 小表名称="1、货币资金余额情况")  # 不抛异常
        错误预期 = dict(预期)
        错误预期[("报表格式附注（填入）", "C5")] = 预期[("报表格式附注（填入）", "C4")]
        with self.assertRaises(AssertionError):
            核验预期对应(结果, 错误预期, 小表名称="1、货币资金余额情况")


class 交接组(unittest.TestCase):
    """完整交接材料：五文件齐全可回读、完整性检查能检出篡改、无未解决不强制细读。"""

    @classmethod
    def setUpClass(cls):
        cls.索引目录 = 取得测试索引()
        from 比较新表 import 载入索引
        cls.索引 = 载入索引(str(cls.索引目录))
        cls.临时 = Path(tempfile.mkdtemp(prefix="初筛交接测试_"))
        新表 = cls.临时 / "交接输入.xlsx"
        变换复制(五样式目录 / "附注样式1.xlsx", 新表, 行移=2, 新表名="新表")
        from 建立结构索引 import 读取工作簿
        from 比较新表 import 比较工作簿
        cls.记录 = 读取工作簿(str(新表))
        cls.新表路径 = 新表
        cls.结果 = 比较工作簿(cls.记录, cls.索引)

    def test_交接五文件齐全且可回读(self):
        from 比较新表 import 写出交接材料
        目录 = self.临时 / "交接输出"
        写出交接材料(self.记录, self.结果, str(目录))
        for 名 in ("工作簿结构.json", "初筛结果.json", "初筛报告.md", "LLM交接说明.md", "运行记录.json"):
            self.assertTrue((目录 / 名).is_file(), 名)
        结构 = json.loads((目录 / "工作簿结构.json").read_text(encoding="utf-8-sig"))
        self.assertEqual(结构["sha256"], 指纹(self.新表路径))
        self.assertTrue(结构["工作表"], "工作簿结构必须含完整工作表")
        结果 = json.loads((目录 / "初筛结果.json").read_text(encoding="utf-8-sig"))
        for 字段 in ("输入版本", "索引版本", "工作表", "候选", "未匹配范围"):
            self.assertIn(字段, 结果)
        for c in 结果["候选"]:
            for 字段 in ("新表范围", "旧表编号", "样式来源", "结构对应", "关系核对", "上下文核对", "覆盖计算", "歧义"):
                self.assertIn(字段, c, f"候选缺字段{字段}")
        记录 = json.loads((目录 / "运行记录.json").read_text(encoding="utf-8-sig"))
        self.assertEqual(记录["完整性结论"], "通过")
        self.assertEqual(记录["输入"]["sha256"], 指纹(self.新表路径))

    def test_输出目录已存在拒绝覆盖(self):
        from 比较新表 import 写出交接材料
        目录 = self.临时 / "交接输出"  # 上一测试已生成
        with self.assertRaises(Exception):
            写出交接材料(self.记录, self.结果, str(目录))

    def test_完整性检查检出篡改(self):
        import copy
        from 比较新表 import 完整性检查, 写出交接材料
        篡改 = copy.deepcopy(self.结果)
        篡改["候选"] = 篡改["候选"][:1]  # 删掉其余候选，其范围内容既不在候选也不在未匹配
        if len(self.结果["候选"]) > 1:
            with self.assertRaises(AssertionError):
                完整性检查(self.记录, 篡改)
            with self.assertRaises(AssertionError):
                写出交接材料(self.记录, 篡改, str(self.临时 / "篡改输出"))

    def test_无未解决时不强制细读(self):
        from 比较新表 import _生成交接说明
        空 = {"候选": [], "未匹配范围": []}
        self.assertIn("不强制", _生成交接说明(空))
        有 = {"候选": [], "未匹配范围": [{"工作表": "X", "坐标数": 3, "分段": []}]}
        说明 = _生成交接说明(有)
        self.assertIn("必须细读", 说明)
        self.assertIn("X", 说明)


class 增量组(unittest.TestCase):
    """增量建库：同输入重跑身份不变；行序无关；加无关小表旧身份保持；局部改动只更新受影响记录。"""

    @classmethod
    def setUpClass(cls):
        cls.首版 = 取得测试索引()
        cls.清单路径 = Path(cls.首版).parent.parent / "输入清单.json"
        cls.临时 = Path(tempfile.mkdtemp(prefix="初筛增量测试_"))

    def _改清单(self, 替换=None, 追加=None, 名="清单改版.json"):
        from 建立结构索引 import 文件指纹
        清单 = json.loads(self.清单路径.read_text(encoding="utf-8-sig"))
        for e in 清单["条目"]:
            if 替换 and e["用途"] in 替换:
                e["路径"] = str(替换[e["用途"]])
                e["sha256"] = 文件指纹(e["路径"])
                e["字节"] = Path(e["路径"]).stat().st_size
        for 用途, 路径 in (追加 or []):
            清单["条目"].append({"用途": 用途, "路径": str(路径),
                            "sha256": 文件指纹(路径), "字节": Path(路径).stat().st_size})
        路径 = self.临时 / 名
        路径.write_text(json.dumps(清单, ensure_ascii=False, indent=1), encoding="utf-8-sig")
        return 路径

    def test_同输入重跑逐字节一致(self):
        from 建立结构索引 import 建立索引
        目录2 = self.临时 / "再跑"
        建立索引(str(self.清单路径), str(目录2), 旧索引目录=str(self.首版))
        for 名 in ("小表记录.jsonl", "文字查找.json", "去向明细.jsonl"):
            self.assertEqual((目录2 / 名).read_bytes(), (Path(self.首版) / 名).read_bytes(), 名)
        清单 = json.loads((目录2 / "索引清单.json").read_text(encoding="utf-8-sig"))
        self.assertTrue(all(g["动作"] == "沿用旧记录" for g in 清单["沿革"]), "同输入重跑应全部沿用")

    def test_映射行序反转身份不变(self):
        from 建立结构索引 import 建立索引
        清单 = json.loads(self.清单路径.read_text(encoding="utf-8-sig"))
        条目 = next(e for e in 清单["条目"] if e["用途"].startswith("五样式原映射:") and "style1-" in e["用途"])
        原行 = Path(条目["路径"]).read_text(encoding="utf-8-sig").splitlines()
        反转目录 = self.临时 / "反转目录"
        反转目录.mkdir(exist_ok=True)
        反转 = 反转目录 / Path(条目["路径"]).name  # 同名文件只改行序
        反转.write_text("\n".join(reversed([l for l in 原行 if l.strip()])) + "\n", encoding="utf-8-sig")
        清单2 = self._改清单(替换={条目["用途"]: 反转}, 名="清单反转.json")
        目录 = self.临时 / "反转索引"
        建立索引(str(清单2), str(目录), 旧索引目录=str(self.首版))
        self.assertEqual((目录 / "小表记录.jsonl").read_bytes(),
                         (Path(self.首版) / "小表记录.jsonl").read_bytes(), "行序变化不得改变小表记录")

    def test_加无关小表旧身份保持(self):
        from 建立结构索引 import 建立索引
        清单 = json.loads(self.清单路径.read_text(encoding="utf-8-sig"))
        源条目 = next(e for e in 清单["条目"]
                   if e["用途"].startswith("五样式原映射:") and "style2-" in e["用途"] and "001" in e["用途"])
        新文件 = self.临时 / "999-测试新增-样式2-人工映射.jsonl"
        shutil.copy(源条目["路径"], 新文件)
        新用途 = "五样式原映射:work/excel-mapping-v1/style2-manual-review-v2/999-测试新增-样式2-人工映射.jsonl"
        清单2 = self._改清单(追加=[(新用途, 新文件)], 名="清单加表.json")
        目录 = self.临时 / "加表索引"
        建立索引(str(清单2), str(目录), 旧索引目录=str(self.首版))
        旧 = {r["小表编号"]: l for l, r in ((l, json.loads(l)) for l in
              (Path(self.首版) / "小表记录.jsonl").read_text(encoding="utf-8-sig").splitlines() if l.strip())}
        新 = {r["小表编号"]: l for l, r in ((l, json.loads(l)) for l in
              (目录 / "小表记录.jsonl").read_text(encoding="utf-8-sig").splitlines() if l.strip())}
        self.assertEqual(len(新), len(旧) + 1, "只应多出一张小表")
        for 编号, 行 in 旧.items():
            self.assertIn(编号, 新, f"旧编号消失：{编号}")
            self.assertEqual(行, 新[编号], f"无关小表加入后旧记录内容变化：{编号}")

    def test_改父标题只更新受影响记录(self):
        import zipfile, re
        from 建立结构索引 import 建立索引
        源 = 五样式目录 / "附注样式1.xlsx"
        改后 = self.临时 / "样式1改父标题.xlsx"
        with zipfile.ZipFile(源) as z:
            项 = {i.filename: z.read(i.filename) for i in z.infolist()}
        表xml = 项["xl/worksheets/sheet1.xml"].decode("utf-8")
        m = re.search(r'<c r="B1"([^>]*)>.*?</c>', 表xml)
        self.assertIsNotNone(m, "样式1 sheet1 应有B1单元格")
        新格 = '<c r="B1" s="2" t="inlineStr"><is><t>测试父标题改动</t></is></c>'
        项["xl/worksheets/sheet1.xml"] = (表xml[:m.start()] + 新格 + 表xml[m.end():]).encode("utf-8")
        with zipfile.ZipFile(改后, "w", zipfile.ZIP_DEFLATED) as z:
            for 名, 数据 in 项.items():
                z.writestr(名, 数据)
        清单2 = self._改清单(替换={"五样式原Excel:附注样式1.xlsx": 改后}, 名="清单改父标题.json")
        目录 = self.临时 / "改父标题索引"
        建立索引(str(清单2), str(目录), 旧索引目录=str(self.首版))
        清单新 = json.loads((目录 / "索引清单.json").read_text(encoding="utf-8-sig"))
        动作 = {g["样式"]: g["动作"] for g in 清单新["沿革"]}
        self.assertEqual(动作[1], "重新提取")
        for i in (2, 3, 4, 5):
            self.assertEqual(动作[i], "沿用旧记录")
        旧 = {r["小表编号"]: r for r in 读小表记录(self.首版)}
        新行 = {r["小表编号"]: l for l, r in ((l, json.loads(l)) for l in
              (目录 / "小表记录.jsonl").read_text(encoding="utf-8-sig").splitlines() if l.strip())}
        旧行 = {r["小表编号"]: l for l, r in ((l, json.loads(l)) for l in
              (Path(self.首版) / "小表记录.jsonl").read_text(encoding="utf-8-sig").splitlines() if l.strip())}
        self.assertEqual(set(新行), set(旧行), "改父标题不得改变小表身份集合")
        变化 = [k for k in 旧行 if 旧行[k] != 新行[k]]
        货币资金 = next(r for r in 旧.values() if r["样式"] == 1 and r["小表名称"] == "1、货币资金余额情况")
        self.assertIn(货币资金["小表编号"], 变化, "父标题所在的货币资金记录必须更新")
        新记录 = json.loads(新行[货币资金["小表编号"]])
        self.assertIn("测试父标题改动", json.dumps(新记录["上下文"], ensure_ascii=False))
        for k in 变化:
            self.assertEqual(k, 货币资金["小表编号"], f"未受影响的记录不应变化：{旧[k]['小表名称']}")


class 入口组(unittest.TestCase):
    """命令行入口：比较子命令端到端跑通，输出完整；坏输入退出非零。"""

    @classmethod
    def setUpClass(cls):
        cls.索引目录 = 取得测试索引()
        cls.临时 = Path(tempfile.mkdtemp(prefix="初筛入口测试_"))
        from openpyxl import Workbook
        记录 = 读小表记录(cls.索引目录)
        样1 = next(r for r in 记录 if r["样式"] == 1 and r["小表名称"] == "1、货币资金余额情况")
        新 = Workbook()
        s = 新.active
        s.title = "新表"
        for n in 样1["节点"]:
            if n.get("原文") is not None:
                s.cell(n["行"] + 3, n["列"] + 1).value = n.get("公式") or n.get("原文")
        cls.新表 = cls.临时 / "小样.xlsx"
        新.save(cls.新表)

    def _跑(self, *参数):
        import subprocess
        return subprocess.run([sys.executable, "-B", "-X", "utf8", str(本批目录 / "运行初筛.py"), *参数],
                              capture_output=True, text=True, timeout=300)

    def test_比较入口端到端(self):
        输出 = self.临时 / "入口输出"
        r = self._跑("比较", "--新表", str(self.新表), "--索引目录", str(self.索引目录),
                     "--输出", str(输出), "--基准对照")
        self.assertEqual(r.returncode, 0, r.stderr)
        for 名 in ("工作簿结构.json", "初筛结果.json", "初筛报告.md", "LLM交接说明.md", "运行记录.json"):
            self.assertTrue((输出 / 名).is_file(), 名)
        结果 = json.loads((输出 / "初筛结果.json").read_text(encoding="utf-8-sig"))
        self.assertTrue([c for c in 结果["候选"] if c["小表名称"] == "1、货币资金余额情况"])
        self.assertEqual(结果.get("基准对照", {}).get("结论"), "一致")

    def test_输入不存在退出非零(self):
        r = self._跑("比较", "--新表", str(self.临时 / "没有.xlsx"), "--索引目录", str(self.索引目录))
        self.assertEqual(r.returncode, 2)
        self.assertIn("不存在", r.stderr)

    def test_缺省输出建时间戳目录(self):
        r = self._跑("比较", "--新表", str(self.新表), "--索引目录", str(self.索引目录))
        self.assertEqual(r.returncode, 0, r.stderr)
        生成 = list((本批目录 / "运行结果").glob("运行_*"))
        self.assertTrue(生成, "缺省输出应在运行结果下建时间戳目录")
        最新 = max(生成, key=lambda p: p.stat().st_mtime)
        self.assertTrue((最新 / "初筛结果.json").is_file())

    def test_选择窗口取消不留成果(self):
        from unittest import mock
        import 运行初筛
        前 = set((本批目录 / "运行结果").glob("运行_*")) if (本批目录 / "运行结果").exists() else set()
        with mock.patch("tkinter.filedialog.askopenfilename", return_value=""):
            码 = 运行初筛.main(["选择文件", "--索引目录", str(self.索引目录)])
        self.assertEqual(码, 1)
        后 = set((本批目录 / "运行结果").glob("运行_*")) if (本批目录 / "运行结果").exists() else set()
        self.assertEqual(前, 后, "取消选择不得新建运行目录")

    def test_选择窗口选定真实文件结果可打开(self):
        from unittest import mock
        import 运行初筛
        真实 = r"C:\Users\27651\Desktop\附注自动更新系统\测试结果\真实数据验收_20260919_204300_181288\流程成果\提取附注表格_20260919_204300_52123b\附注表格.xlsx"
        if not Path(真实).is_file():
            self.skipTest("真实A端文件不在本机")
        with mock.patch("tkinter.filedialog.askopenfilename", return_value=真实):
            码 = 运行初筛.main(["选择文件", "--索引目录", str(self.索引目录)])
        self.assertEqual(码, 0)
        最新 = max((本批目录 / "运行结果").glob("运行_*"), key=lambda p: p.stat().st_mtime)
        结果 = json.loads((最新 / "初筛结果.json").read_text(encoding="utf-8-sig"))
        self.assertIn("候选", 结果)
        self.assertTrue((最新 / "初筛报告.md").read_text(encoding="utf-8-sig").strip())


class 修复回归组(unittest.TestCase):
    """独立复核发现问题的回归：共用表头不串表、重复位置不截断、缩进/新增合并被记录、
    有差异候选必须进细读清单、无参数入口进选择窗口、年份数字缓存变化留差异。"""

    @classmethod
    def setUpClass(cls):
        cls.索引目录 = 取得测试索引()
        from 比较新表 import 载入索引
        cls.索引 = 载入索引(str(cls.索引目录))
        cls.临时 = Path(tempfile.mkdtemp(prefix="初筛修复回归_"))

    def _比较(self, 新表路径):
        from 建立结构索引 import 读取工作簿
        from 比较新表 import 比较工作簿
        return 比较工作簿(读取工作簿(str(新表路径)), self.索引)

    @staticmethod
    def _恒等(候选):
        return [c for c in 候选
                if c["覆盖计算"]["旧表文字覆盖"]["分子"] == c["覆盖计算"]["旧表文字覆盖"]["分母"]
                and not [d for d in c["结构对应"]
                       if d["旧坐标"] != d["新坐标"] or d["旧工作表"] != d["新工作表"]]]

    def test_样式2全表原样查回不串表(self):
        """复核问题1：共用表头不得把小表串到全表共用行；样式2全库恒等查回。"""
        from openpyxl.utils import range_boundaries
        结果 = self._比较(五样式目录 / "附注样式2.xlsx")
        无身份 = []
        for 编号, t in self.索引["小表"].items():
            if t["样式"] != 2:
                continue
            恒等 = self._恒等([c for c in 结果["候选"] if c["旧表编号"] == 编号])
            if len(恒等) != 1:
                无身份.append(t["小表名称"])
        # 仅4张单行小表（唯一文字是通用词"检验"）结构上无法靠文字定位，属已登记覆盖限制
        self.assertEqual(sorted(无身份),
                         sorted(["交易性金融负债", "衍生金融负债", "债权投资-按类别列示", "其他债权投资-按类别列示"]),
                         f"样式2应除4张单行通用词小表外全部恒等查回，实际：{无身份}")
        # 应收股利：恒等候选必须在自己的位置，范围不得吞并其他科目
        编号 = next(k for k, v in self.索引["小表"].items() if v["样式"] == 2 and v["小表名称"] == "应收股利")
        恒等 = self._恒等([c for c in 结果["候选"] if c["旧表编号"] == 编号])
        self.assertEqual(len(恒等), 1)
        self.assertEqual(range_boundaries(恒等[0]["新表范围"]), (2, 238, 2, 239))
        self.assertTrue(恒等[0]["共用表头核对"]["一致"], "恒等位置上共用表头应核对一致")
        self.assertFalse(恒等[0]["共用表头核对"]["冲突"] or 恒等[0]["共用表头核对"]["缺依据"])

    def test_重复12处全部保留(self):
        """复核问题2：同一小表复制12处，12个独立位置都必须是候选，不得按数量截断。"""
        from openpyxl import load_workbook, Workbook
        源 = load_workbook(五样式目录 / "附注样式1.xlsx")["报表格式附注（填入）"]
        新 = Workbook()
        s = 新.active
        s.title = "新表"
        块 = [(r, c, 源.cell(r, c).value) for r in range(1, 10) for c in range(1, 6)
            if 源.cell(r, c).value is not None]
        期望范围 = []
        for k in range(12):
            位移 = 30 * k
            for r, c, v in 块:
                s.cell(r + 位移, c).value = v
            期望范围.append(f"B{2 + 位移}:D{8 + 位移}")
        路径 = self.临时 / "重复12张.xlsx"
        新.save(路径)
        结果 = self._比较(路径)
        范围集 = {c["新表范围"] for c in 结果["候选"]
                if c["小表名称"] == "1、货币资金余额情况" and c["样式来源"] == 1}
        self.assertEqual(范围集, set(期望范围), f"12处重复位置必须全部保留，实际{sorted(范围集)}")

    def test_缩进变化被记录(self):
        """复核问题4：项目缩进层级变化必须在差异或层级核对中留痕，不得满覆盖却无记录。"""
        新表 = self.临时 / "缩进.xlsx"
        变换复制(五样式目录 / "附注样式1.xlsx", 新表, 改值={"B5": "　　银行存款"}, 新表名="新表")
        候选 = [c for c in self._比较(新表)["候选"] if c["小表名称"] == "1、货币资金余额情况" and c["样式来源"] == 1]
        self.assertEqual(len(候选), 1)
        c = 候选[0]
        排版 = [d for d in c["差异"] if d["类型"] == "排版差异" and d["旧坐标"] == "B5"]
        层级 = [x for x in c["关系核对"]["层级"]["冲突"] if x["旧坐标"] == "B5"]
        self.assertTrue(排版 or 层级, f"缩进变化必须留痕：差异{c['差异']} 层级{c['关系核对']['层级']}")

    def test_新增合并被报告(self):
        """复核问题4：新表新增合并（旧表无对应）必须出现在关系核对。"""
        from openpyxl import load_workbook
        新表 = self.临时 / "新增合并.xlsx"
        变换复制(五样式目录 / "附注样式1.xlsx", 新表, 新表名="新表")
        w = load_workbook(新表)
        w["新表"].merge_cells("C4:D4")  # 现金行的空值格，新增合并不损内容
        w.save(新表)
        w.close()
        候选 = [c for c in self._比较(新表)["候选"] if c["小表名称"] == "1、货币资金余额情况" and c["样式来源"] == 1]
        self.assertEqual(len(候选), 1)
        合并 = 候选[0]["关系核对"]["合并"]
        self.assertTrue(any(m.get("新合并") == "C4:D4" and "新增" in m["结果"] for m in 合并),
                       f"新增合并C4:D4必须被报告：{合并}")

    def test_年份数字表头缓存变化留下差异(self):
        """复核三.4：样式2共用/自有表头的年份缓存被改（公式不动），必须有差异记录。"""
        import re
        import shutil as  shutil_
        import zipfile
        源 = 五样式目录 / "附注样式2.xlsx"
        改后 = self.临时 / "样式2改年份缓存.xlsx"
        with zipfile.ZipFile(源) as z:
            项 = {i.filename: z.read(i.filename) for i in z.infolist()}
        xml = 项["xl/worksheets/sheet1.xml"].decode("utf-8")
        m = re.search(r'<c r="C3"[^>]*>.*?</c>', xml)
        self.assertIsNotNone(m)
        新格 = m.group(0).replace("<v>2025年度</v>", "<v>2026年度</v>")
        self.assertNotEqual(新格, m.group(0))
        项["xl/worksheets/sheet1.xml"] = (xml[:m.start()] + 新格 + xml[m.end():]).encode("utf-8")
        with zipfile.ZipFile(改后, "w", zipfile.ZIP_DEFLATED) as z:
            for 名, 数据 in 项.items():
                z.writestr(名, 数据)
        结果 = self._比较(改后)
        候选 = [c for c in 结果["候选"] if c["小表名称"] == "货币资金" and c["样式来源"] == 2]
        self.assertTrue(候选)
        c = max(候选, key=lambda x: x["覆盖计算"]["旧表文字覆盖"]["分子"])
        痕迹 = json.dumps(c["差异"] + c["共用表头核对"]["冲突"] + c["上下文核对"]["冲突"], ensure_ascii=False)
        self.assertIn("2025年度", 痕迹, f"年份缓存从2025改2026必须留差异：{c['差异'][:5]}")
        self.assertIn("2026年度", 痕迹)

    def test_有差异候选必须进细读清单(self):
        """复核问题3：带差异的候选必须出现在交接说明待处理清单，不得报'没有未解决问题'。"""
        from 比较新表 import _生成交接说明
        新表 = self.临时 / "缺表头.xlsx"
        变换复制(五样式目录 / "附注样式1.xlsx", 新表, 改值={"C3": None}, 新表名="新表")
        结果 = self._比较(新表)
        候选 = [c for c in 结果["候选"] if c["小表名称"] == "1、货币资金余额情况" and c["样式来源"] == 1]
        self.assertEqual(len(候选), 1)
        self.assertTrue(候选[0]["差异"], "删表头必须留下差异")
        说明 = _生成交接说明(结果)
        self.assertIn("必须细读", 说明)
        self.assertIn("1、货币资金余额情况", 说明)
        self.assertNotIn("本次没有未解决问题", 说明)

    def test_无参数默认进入选择窗口(self):
        """复核问题5：无参数运行（双击.cmd同等）必须进入选择文件分支。"""
        from unittest import mock
        import 运行初筛
        with mock.patch("tkinter.filedialog.askopenfilename", return_value="") as 窗口:
            码 = 运行初筛.main([])
        self.assertEqual(码, 1)
        self.assertTrue(窗口.called, "无参数必须弹出文件选择窗口")

    def test_cmd入口转发参数(self):
        """复核问题5：实际执行启动初筛.cmd，参数转发到Python入口。"""
        import subprocess
        r = subprocess.run(["cmd", "/c", str(本批目录 / "启动初筛.cmd"), "比较",
                            "--新表", str(self.临时 / "不存在.xlsx"), "--索引目录", str(self.索引目录)],
                           capture_output=True, text=True, timeout=120, cwd=str(本批目录))
        self.assertEqual(r.returncode, 2, f"cmd应转发参数并返回参数检查失败码：{r.stdout}{r.stderr}")
        self.assertIn("不存在", r.stdout + r.stderr)


if __name__ == "__main__":
    解析 = argparse.ArgumentParser()
    解析.add_argument("--组", default="全部")
    参数 = 解析.parse_args()
    组表 = {"输入": 输入组, "索引": 索引组, "查找": 查找组, "比较": 比较组,
          "交接": 交接组, "增量": 增量组, "入口": 入口组, "修复回归": 修复回归组}
    if 参数.组 == "全部":
        套件 = unittest.TestLoader().loadTestsFromModule(sys.modules[__name__])
    else:
        套件 = unittest.TestSuite()
        for 名 in 参数.组.split(","):
            套件.addTests(unittest.TestLoader().loadTestsFromTestCase(组表[名.strip()]))
    结果 = unittest.TextTestRunner(verbosity=2).run(套件)
    sys.exit(0 if 结果.wasSuccessful() else 1)
