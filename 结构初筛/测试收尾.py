"""独立小样本回归：全部材料虚构，不依赖开发者电脑或收费模型。"""
import copy
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parent))
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from 建立结构索引 import 读取工作簿, 规范文字
from 比较新表 import 比较工作簿, _候选问题
import 运行初筛


def 虚构参照():
    表 = {'小表编号': '演示资金表', '小表名称': '演示货币资金', '样式': 1,
         '来源': {'工作表': '演示表'}, '节点': [], '合并': [], '项目顺序': [], '上下文': [], '表头': []}
    数据 = [('演示货币资金', None, None), ('项目', '期末余额', '期初余额'),
            ('现金', 100, 90), ('银行存款', 200, 180), ('合计', 300, 270)]
    for r, 行 in enumerate(数据, 1):
        for c, 值 in enumerate(行, 1):
            if 值 is None:
                continue
            文 = isinstance(值, str)
            n = {'坐标': f'{get_column_letter(c)}{r}', '行': r, '列': c, '原文': 值,
                 '角色': ('标题' if r == 1 else '表头' if r == 2 else '项目') if 文 else '业务数值'}
            if 文:
                n['查找文字'] = 规范文字(值)
            else:
                n['语义引用'] = {'说明': '完全虚构的演示位置，不是已批准的业务语义'}
            表['节点'].append(n)
    return 表


def 内存索引(表):
    正文 = {}
    for n in 表['节点']:
        if n.get('查找文字'):
            正文.setdefault(n['查找文字'], []).append({'小表编号': 表['小表编号']})
    return {'目录': '虚构演示', '规则版本': '演示', '正文': 正文, '别名': {}, '记录': [表],
            '小表': {表['小表编号']: 表}, '索引清单': {}}


class 收尾回归(unittest.TestCase):
    def setUp(self):
        self.目录 = Path(tempfile.mkdtemp(prefix='附注收尾验证_'))
        self.表 = 虚构参照()
        self.索引 = 内存索引(self.表)

    def 新表(self, offsets=((0, 0),), merge=None, sheets=('演示表',)):
        w = Workbook()
        for i, 名 in enumerate(sheets):
            s = w.active if i == 0 else w.create_sheet()
            s.title = 名
            for dr, dc in offsets:
                for n in self.表['节点']:
                    s.cell(n['行'] + dr, n['列'] + dc, n['原文'])
            if merge:
                s.merge_cells(merge)
        p = self.目录 / '虚构.xlsx'
        w.save(p)
        w.close()
        return p

    def test_左右并排两个位置都保留(self):
        r = 比较工作簿(读取工作簿(self.新表(((0, 0), (0, 5)))), self.索引)
        self.assertEqual({x['新表范围'] for x in r['候选']}, {'A1:C5', 'F1:H5'})

    def test_重复十二处全部保留(self):
        r = 比较工作簿(读取工作簿(self.新表(tuple((10*k, 0) for k in range(12)))), self.索引)
        self.assertEqual({x['新表范围'] for x in r['候选']}, {f'A{1+10*k}:C{5+10*k}' for k in range(12)})

    def test_跨出候选边界的合并仍需细读(self):
        r = 比较工作簿(读取工作簿(self.新表(merge='A1:E1')), self.索引)
        c = next(x for x in r['候选'] if x['种子偏移'] == [0, 0])
        self.assertTrue(any(x.get('新合并') == 'A1:E1' for x in c['关系核对']['合并']))
        self.assertTrue(_候选问题(c))

    def test_单词重复两行不能冒充两个识别词(self):
        for n in self.表['节点']:
            if n.get('查找文字'):
                n['原文'] = n['查找文字'] = '现金' if n['行'] in (3, 4) else '不存在' + n['坐标']
        self.索引 = 内存索引(self.表)
        w = Workbook(); s = w.active; s['A3'] = '现金'; s['A4'] = '现金'
        p = self.目录 / '单词.xlsx'; w.save(p); w.close()
        self.assertEqual(比较工作簿(读取工作簿(p), self.索引)['候选'], [])

    def test_基准跨工作表遗漏不能返回成功(self):
        p = self.新表(sheets=('甲表', '乙表'))
        真比较 = 比较工作簿
        def 故障注入(记录, 索引, 预筛选=True):
            r = 真比较(记录, 索引, 预筛选)
            if 预筛选:
                r['候选'] = [c for c in r['候选'] if c['新工作表'] == '甲表']
                cells = list(记录['sheets'][1]['cells'])
                r['未匹配范围'].append({'工作表': '乙表', '隐藏': False, '坐标数': len(cells),
                    '分段': [{'范围': 'A1:C5', '坐标数': len(cells), '坐标清单': cells, '原文摘录': []}]})
            return r
        输出 = self.目录 / '交接材料'
        with patch('比较新表.载入索引', return_value=self.索引), patch('比较新表.比较工作簿', side_effect=故障注入):
            码 = 运行初筛.执行比较(p, '虚构', 输出, True)
        r = json.loads((输出 / '初筛结果.json').read_text(encoding='utf-8-sig'))
        self.assertNotEqual(码, 0)
        self.assertEqual(r['基准对照']['基准候选数'], 2)
        self.assertEqual(len(r['基准对照']['遗漏']), 1)

    def test_一个词重复四次仍不构成部分候选(self):
        self.表['节点'] = [{'坐标': f'A{r}', '行': r, '列': 1, '角色': '项目',
            '原文': t, '查找文字': t} for r, t in enumerate(['现金']*4 + ['存款', '票据', '合计'], 1)]
        self.索引 = 内存索引(self.表)
        w = Workbook(); s = w.active
        for r in range(1, 5): s.cell(r, 1, '现金')
        p = self.目录 / '重复同词.xlsx'; w.save(p); w.close()
        self.assertEqual(比较工作簿(读取工作簿(p), self.索引)['候选'], [])

    def test_输入清单过期必须在整簿比较之前失败(self):
        import hashlib
        import 核验真实样本
        原件 = self.目录 / '虚构输入.txt'
        原件.write_text('只核对指纹，不读取为工作簿', encoding='utf-8')
        样本 = {'文件': str(原件), 'sha256': hashlib.sha256(原件.read_bytes()).hexdigest()}
        清单 = self.目录 / '固定样本.json'
        清单.write_text(json.dumps({'五样式原样查回': [样本], '真实A端': 样本}), encoding='utf-8')
        输入 = self.目录 / '过期输入.json'
        输入.write_text(json.dumps({'条目': [{'用途': '复用代码:机械读取模块', '路径': str(原件), 'sha256': '0'*64}]}), encoding='utf-8')
        with patch.object(核验真实样本, '读取工作簿') as 读取:
            with self.assertRaisesRegex(ValueError, '指纹不符'):
                核验真实样本.main(['--索引目录', str(self.目录/'未用索引'), '--样本清单', str(清单), '--输入清单', str(输入), '--输出', str(self.目录/'核验输出')])
            读取.assert_not_called()

    def test_同一通用词两行全对不得被预筛选挡住(self):
        self.表['节点'] = [{'坐标': f'A{r}', '行': r, '列': 1, '角色': '项目',
            '原文': '合计', '查找文字': '合计'} for r in (1, 2)]
        索引 = 内存索引(self.表)
        索引['正文']['合计'] = [{'小表编号': self.表['小表编号']}] * 11
        记录 = 读取工作簿(self.新表())
        self.assertEqual(len(比较工作簿(记录, 索引, False)['候选']), 1)
        self.assertEqual(len(比较工作簿(记录, 索引, True)['候选']), 1)

    def test_整表文字都不足以成立时不展开重复位置(self):
        self.表['节点'] = [{'坐标': f'A{r}', '行': r, '列': 1, '角色': '项目',
            '原文': t, '查找文字': t} for r, t in enumerate(['现金'] * 4 + ['存款', '票据', '合计'], 1)]
        w = Workbook(); s = w.active
        for r in range(1, 41): s.cell(r, 1, '现金')
        p = self.目录 / '只能命中一个词.xlsx'; w.save(p); w.close()
        r = 比较工作簿(读取工作簿(p), 内存索引(self.表), False)
        self.assertEqual(r['候选'], [])
        self.assertEqual(r['工作表'][0]['运行统计']['种子数'], 0)

    def test_重复合计旁的唯一项目不应反复扫描所有行(self):
        import 比较新表 as m
        self.表['节点'] = []
        w = Workbook(); s = w.active
        for r in range(1, 201):
            for c, t in ((1, '合计'), (2, f'唯一项目{r}')):
                s.cell(r, c, t)
                self.表['节点'].append({'坐标': f'{get_column_letter(c)}{r}', '行': r, '列': c,
                    '角色': '项目', '原文': t, '查找文字': t})
        p = self.目录 / '大量相同合计.xlsx'; w.save(p); w.close()
        数据 = m._工作表数据(读取工作簿(p)['sheets'][0])
        class 计数文字(dict):
            次数 = 0
            def get(self, *args):
                self.次数 += 1
                return super().get(*args)
        数据['显示'] = 计数文字(数据['显示'])
        with patch.object(m, '索引正文计数', 内存索引(self.表)['正文']):
            r = m._对齐候选(self.表, '演示表', 数据, (0, 0, 400))
        self.assertEqual(len(r['结构对应']), 400)
        self.assertTrue(all(d['旧坐标'] == d['新坐标'] for d in r['结构对应']))
        self.assertLess(数据['显示'].次数, 5000, '两百行对齐不得反复查询数万次文字')


class 上层语义回归(unittest.TestCase):
    """正文相同不等于业务相同；期望业务直接来自虚构的人工定义。"""

    def setUp(self):
        self.目录 = Path(tempfile.mkdtemp(prefix='工作表上层语义_'))
        self.表们 = [self.旧表('应收股利', '其他应收款', '应收股利分类'),
                    self.旧表('应收利息', '其他应收款', '应收利息分类')]

    def 旧表(self, 名, 附注, 表义):
        from 建立结构索引 import _提取小表
        记录 = self.新表(名)['sheets'][0]
        映射 = {'mapping_id': 名 + '-合计', 'cell': 'B3', '__文件': '虚构人工映射',
            'target_slot_id': 名 + '-余额', 'row_semantics': '合计', 'column_semantics': '期末余额',
            'target_slot': {'note': 附注, 'table': 表义}}
        表现 = {映射['mapping_id']: {'representation_id': 名, 'semantic_id': 名 + '-余额'}}
        return _提取小表(1, '虚构.xlsx', 名, 名, [映射], (1, 1, 2, 3), '虚构范围',
                      记录, 表现, 'A', [], [])

    def 新表(self, 名, 标题=None):
        内容 = {'A1': '项目', 'B1': '期末余额', 'A2': '第一项', 'B2': 1, 'A3': '合计', 'B3': 1}
        if 标题:
            内容['A5'] = 标题
        cells = {}
        from openpyxl.utils import range_boundaries
        for pos, value in 内容.items():
            c, r, _, _ = range_boundaries(pos)
            cells[pos] = {'row': r, 'column': c, 'value': value, 'cached_value': value, 'formula': None}
        return {'path': '虚构.xlsx', 'sha256': '0'*64, 'sheets': [
            {'name': 名, 'cells': cells, 'max_row': 5 if 标题 else 3, 'max_column': 2, 'merges': []}]}

    def 索引(self):
        rs = self.表们
        索引 = 内存索引(rs[0])
        索引['记录'] = rs
        索引['小表'] = {r['小表编号']: r for r in rs}
        索引['正文'] = {}
        for r in rs:
            for n in r['节点']:
                if n.get('查找文字'):
                    索引['正文'].setdefault(n['查找文字'], []).append({'小表编号': r['小表编号']})
        return 索引

    def test_相同合计必须继承各自工作表业务(self):
        for 快速 in (True, False):
            with self.subTest(快速=快速):
                r = 比较工作簿(self.新表('应收股利'), self.索引(), 快速)
                self.assertEqual([c['小表名称'] for c in r['候选']], ['应收股利'])
                self.assertEqual(r['运行统计']['比较的小表数'], 1)
                c = r['候选'][0]
                self.assertIn('上层语义核对', c)
                self.assertEqual(c['上层语义核对']['业务路径'][0]['表义'], '应收股利分类')

    def test_不同附注的同名表义不能混同(self):
        self.表们 = [self.旧表('应付利息', '其他应付款', '余额构成'),
                    self.旧表('交易性金融资产', '交易性金融资产', '余额构成')]
        r = 比较工作簿(self.新表('应付利息'), self.索引())
        self.assertEqual([c['小表名称'] for c in r['候选']], ['应付利息'])

    def test_未知名称不猜同义且保留所有候选(self):
        r = 比较工作簿(self.新表('本年明细'), self.索引())
        self.assertEqual(len(r['候选']), 2)
        self.assertTrue(all(c.get('上层语义核对', {}).get('待确认原因') for c in r['候选']))

    def test_工作表名与表内另一业务标题冲突必须放宽(self):
        r = 比较工作簿(self.新表('应收股利', '应收利息'), self.索引())
        self.assertEqual(len(r['候选']), 2)
        self.assertTrue(any('冲突' in p for c in r['候选'] for p in _候选问题(c)))

    def test_上级附注名必须包括其下多种表义(self):
        r = 比较工作簿(self.新表('其他应收款'), self.索引())
        self.assertEqual({c['小表名称'] for c in r['候选']}, {'应收股利', '应收利息'})
        self.assertTrue(all(c.get('上层语义核对', {}).get('业务路径') for c in r['候选']))

    def test_同一附注内混合业务只扩大到有依据的业务(self):
        self.表们.append(self.旧表('应付账款', '应付账款', '余额构成'))
        r = 比较工作簿(self.新表('应收股利', '应收利息'), self.索引())
        self.assertEqual({c['小表名称'] for c in r['候选']}, {'应收股利', '应收利息'})
        self.assertEqual(r['运行统计']['比较的小表数'], 2)

    def test_跨附注混合业务必须全范围保留(self):
        self.表们.append(self.旧表('应付账款', '应付账款', '余额构成'))
        r = 比较工作簿(self.新表('应收股利', '应付账款'), self.索引())
        self.assertEqual(len(r['候选']), 3)
        self.assertTrue(all(c['上层语义核对']['待确认原因'] for c in r['候选']))

    def test_表内更宽的父标题不能只因部分相交而忽略(self):
        r = 比较工作簿(self.新表('应收股利', '其他应收款'), self.索引())
        self.assertEqual({c['小表名称'] for c in r['候选']}, {'应收股利', '应收利息'})

    def test_旧索引缺上层信息仍可运行并明确缺依据(self):
        for r in self.表们:
            r.pop('上层语义', None)
        r = 比较工作簿(self.新表('应收股利'), self.索引())
        self.assertEqual(len(r['候选']), 2)
        self.assertTrue(all(c.get('上层语义核对', {}).get('待确认原因') for c in r['候选']))

    def test_新索引保存工作表父语义及映射依据(self):
        meta = self.表们[0].get('上层语义', {})
        self.assertTrue(meta.get('完整'))
        self.assertEqual(meta.get('业务路径'), [{'附注': '其他应收款', '表义': '应收股利分类',
                                               '映射依据': ['应收股利-合计']}])

    def test_旁边检验标签不能冒充表标题(self):
        from 建立结构索引 import _提取小表
        记录 = self.新表('应收股利')['sheets'][0]
        记录['cells'] = {'A2': {'value': '应收股利明细'}, 'H1': {'value': '余额检验'},
                         'A3': {'value': '合计'}, 'B3': {'value': 1}}
        映射 = {'mapping_id': '甲', 'target_slot_id': '股利', 'cell': 'B3', '__文件': '虚构',
                'target_slot': {'note': '其他应收款', 'table': '应收股利分类'}}
        表 = _提取小表(1, '虚构', '应收股利', '应收股利', [映射], (1, 1, 8, 3), '虚构',
            记录, {'甲': {'representation_id': '甲', 'semantic_id': '股利'}}, 'A', [], [])
        self.assertNotIn('余额检验', [e['原文'] for e in 表.get('上层语义', {}).get('名称依据', [])])

    def test_已登记的不同标题表达仍可找到同一业务(self):
        self.表们[0]['上层语义']['名称依据'].append({'原文': '股利应收明细', '来源': '表内标题', '坐标': 'A1'})
        r = 比较工作簿(self.新表('股利应收明细'), self.索引())
        self.assertEqual([c['小表名称'] for c in r['候选']], ['应收股利'])

    def test_旧表只有部分上层依据时不能被排除(self):
        self.表们[1]['上层语义']['完整'] = False
        r = 比较工作簿(self.新表('应收股利'), self.索引())
        self.assertEqual(len(r['候选']), 2)
        c = next(c for c in r['候选'] if c['小表名称'] == '应收利息')
        self.assertTrue(c['上层语义核对']['待确认原因'])

    def test_报告和交接保留父语义冲突及筛除依据(self):
        from 比较新表 import 写出交接材料
        数据 = self.新表('应收股利', '应收利息')
        r = 比较工作簿(数据, self.索引())
        输出 = self.目录 / '交接'
        写出交接材料(数据, r, 输出)
        for 名 in ('初筛报告.md', 'LLM交接说明.md'):
            文 = (输出 / 名).read_text(encoding='utf-8-sig')
            self.assertIn('上层', 文)
            self.assertIn('冲突', 文)
        saved = json.loads((输出 / '初筛结果.json').read_text(encoding='utf-8-sig'))
        self.assertTrue(saved['候选'][0]['上层语义核对']['旧业务路径'][0]['映射依据'])


if __name__ == '__main__':
    unittest.main(verbosity=2)
