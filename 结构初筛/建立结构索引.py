# -*- coding: utf-8 -*-
"""建立结构索引：只读装载资产、提取旧小表记录、生成文字查找目录。

接口约定（02_设计.md 第8节）：
    读取工作簿(路径) -> dict
    建立索引(输入清单, 输出目录, 旧索引目录=None) -> dict
另含清点入口：清点(成果库, 五样式目录, 输出) 生成本批输入清单.json。
"""
from __future__ import annotations

import hashlib
import importlib.util
import json
import re
import unicodedata
from collections import defaultdict
from pathlib import Path

规则版本 = "结构初筛规则v3"  # v3：结构格原文存显示值（公式格取缓存），排版差异可比对；v2：共用表头移出正文节点

项目根 = Path(__file__).resolve().parents[1]
机械读取模块路径 = 项目根 / "附注更新" / "表格.py"


def 文件指纹(路径):
    h = hashlib.sha256()
    with open(路径, "rb") as f:
        for 块 in iter(lambda: f.read(1 << 20), b""):
            h.update(块)
    return h.hexdigest()


def 装载机械读取模块(路径=None):
    """只读复用 附注更新/表格.py 的 read_workbook/file_hash；按文件装载，不触发包内其他模块。"""
    路径 = Path(路径 or 机械读取模块路径)
    if not 路径.is_file():
        raise ValueError(f"机械读取模块不存在：{路径}")
    规范 = importlib.util.spec_from_file_location("结构初筛_表格", 路径)
    模块 = importlib.util.module_from_spec(规范)
    规范.loader.exec_module(模块)
    return 模块


def 读jsonl(路径):
    with open(路径, encoding="utf-8-sig") as f:
        return [json.loads(行) for 行 in f if 行.strip()]


def 规范文字(值):
    """只作不改变含义的排版整理：NFKC、折叠空白、去首尾。其他类型原样返回。"""
    if not isinstance(值, str):
        return 值
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", 值)).strip()


def _登记(条目, 用途, 路径, 预期指纹=None):
    路径 = str(Path(路径))
    if not Path(路径).is_file():
        raise ValueError(f"清点失败，文件不存在：{用途} {路径}")
    实际 = 文件指纹(路径)
    if 预期指纹 and 实际 != 预期指纹:
        raise ValueError(f"清点失败，与资产清单登记指纹不符：{用途} {路径}")
    条目.append({"用途": 用途, "路径": 路径, "sha256": 实际, "字节": Path(路径).stat().st_size})


def 清点(成果库, 五样式目录, 输出):
    """固定本批实际读取的全部输入，输出路径+SHA-256清单；原件只读。"""
    成果库, 五样式目录 = Path(成果库), Path(五样式目录)
    清单路径 = 成果库 / "资产清单.json"
    if not 清单路径.is_file():
        raise ValueError(f"资产清单不存在：{清单路径}")
    资产清单 = json.loads(清单路径.read_text(encoding="utf-8-sig"))
    条目 = []
    _登记(条目, "统一成果:资产清单", 清单路径)

    工作簿登记 = {a["路径"]: a["sha256"] for a in 资产清单["资产"] if a.get("角色") == "原始工作簿"}
    for i in range(1, 6):
        路径 = 五样式目录 / f"附注样式{i}.xlsx"
        预期 = 工作簿登记.get(str(路径))
        if 预期 is None:
            raise ValueError(f"资产清单未登记原始工作簿：{路径}")
        _登记(条目, f"五样式原Excel:附注样式{i}.xlsx", 路径, 预期)

    映射登记 = set()
    for m in 资产清单["固定输入"]["五样式原映射"]:
        _登记(条目, f"五样式原映射:{m['相对路径']}", m["路径"], m["sha256"])
        映射登记.add(str(Path(m["路径"])))

    for 名称 in ("表现记录.jsonl", "统一语义记录.jsonl", "维度字典.jsonl", "上下文证据.json"):
        _登记(条目, f"统一成果:{名称.removesuffix('.jsonl')}", 成果库 / "数据" / 名称)
    for 名称 in ("旧新对应.jsonl", "来源示例.jsonl", "维度规范对应.json", "表现身份迁移.jsonl"):
        _登记(条目, f"统一成果:对应:{名称.removesuffix('.json') if 名称.endswith('.json') else 名称.removesuffix('.jsonl')}", 成果库 / "对应" / 名称)
    _登记(条目, "统一成果:查找索引", 成果库 / "索引" / "查找索引.json")

    映射根 = None
    for m in 资产清单["固定输入"]["五样式原映射"]:
        p = Path(m["路径"])
        映射根 = p.parents[1] if 映射根 is None else 映射根
    复核文件 = [
        ("复核文档:样式1逐格分类清单", 映射根 / "style1-manual-review-v2" / "样式1-全部可见内容逐格分类清单-待审批.jsonl"),
        ("复核文档:样式3逐表复核记录", 映射根 / "style3-manual-review-v2" / "样式3-逐表复核记录.md"),
        ("复核文档:样式4逐表复核记录", 映射根 / "style4-manual-review-v2" / "样式4-逐表复核记录.md"),
        ("复核文档:样式5逐表复核记录", 映射根 / "style5-manual-review-v2" / "样式5-逐表复核记录.md"),
    ]
    已登记 = set(映射登记)
    for 用途, 路径 in 复核文件:
        _登记(条目, 用途, 路径)
        已登记.add(str(路径))
    for 样式目录 in sorted(映射根.glob("style*-manual-review-v2")):
        for 文件 in sorted(样式目录.glob("*.jsonl")):
            if str(文件) not in 已登记:
                raise ValueError(f"映射目录中存在未登记文件，请核实来源：{文件}")

    _登记(条目, "复用代码:机械读取模块", 机械读取模块路径)
    清单 = {"说明": "结构索引与代码初筛本批输入清单；路径与指纹为开工实测", "条目": 条目}
    输出 = Path(输出)
    输出.parent.mkdir(parents=True, exist_ok=True)
    输出.write_text(json.dumps(清单, ensure_ascii=False, indent=1) + "\n", encoding="utf-8-sig")
    return 清单


def 核对输入清单(清单):
    """逐条回读核验；任何问题都报明原因并拒绝继续。"""
    问题 = []
    按用途 = defaultdict(list)
    按路径 = defaultdict(set)
    for e in 清单.get("条目", []):
        按用途[e["用途"]].append(e["路径"])
        按路径[e["路径"]].add(e["用途"])
    for 用途, 路径们 in 按用途.items():
        if len(路径们) != 1:
            问题.append(f"用途冲突，同一用途登记了{len(路径们)}个文件：{用途}")
    for 路径, 用途们 in 按路径.items():
        if len(用途们) != 1:
            问题.append(f"来源冲突，同一文件被登记为多个用途：{路径} -> {sorted(用途们)}")
    for e in 清单.get("条目", []):
        p = Path(e["路径"])
        if not p.is_file():
            问题.append(f"文件不存在：{e['用途']} {e['路径']}")
            continue
        if 文件指纹(p) != e["sha256"]:
            问题.append(f"指纹不符：{e['用途']} {e['路径']}")
    if 问题:
        raise ValueError("输入清单核验失败：\n" + "\n".join(问题))
    return True


# ============ 工作簿读取与小表提取 ============

from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

_表格模块缓存 = {}


def 表格模块(路径=None):
    键 = str(路径 or 机械读取模块路径)
    if 键 not in _表格模块缓存:
        _表格模块缓存[键] = 装载机械读取模块(路径)
    return _表格模块缓存[键]


def 读取工作簿(路径: str) -> dict:
    """机械读取完整工作簿：原值、公式及缓存、合并、隐藏、命名区域；不执行公式与宏。"""
    return 表格模块().read_workbook(str(路径))


def _json默认(o):
    if hasattr(o, "isoformat"):
        return o.isoformat()
    return str(o)


def _坐标(行, 列):
    return f"{get_column_letter(列)}{行}"


def _bbox(范围们):
    盒 = [range_boundaries(r) for r in 范围们]
    return (min(b[0] for b in 盒), min(b[1] for b in 盒), max(b[2] for b in 盒), max(b[3] for b in 盒))


def _解析样式3复核(路径):
    结果 = {}
    for m in re.finditer(r"^##\s*(\d+)．(.+?)（(\d+)—(\d+)行）\s*$",
                         Path(路径).read_text(encoding="utf-8-sig"), re.M):
        结果[int(m.group(1))] = {"名称": m.group(2).strip(), "顶": int(m.group(3)), "底": int(m.group(4))}
    return 结果


def _解析样式4复核(路径):
    结果 = {}
    当前 = None
    for 行 in Path(路径).read_text(encoding="utf-8-sig").splitlines():
        m = re.match(r"^##\s*(\d+)．(.+?)\s*$", 行)
        if m:
            当前 = m.group(2).strip()
            continue
        m2 = re.search(r"人工查看范围：([A-Z]+\d+:[A-Z]+\d+)", 行)
        if m2 and 当前 and 当前 not in 结果:
            结果[当前] = m2.group(1)
    return 结果


def _解析样式5复核(路径):
    结果 = []
    当前 = None
    for 行 in Path(路径).read_text(encoding="utf-8-sig").splitlines():
        m = re.match(r"^##\s*(\d+(?:—\d+)?)．(.+?)（(?:[^（）]*[,，])?起始行(\d+)）\s*$", 行)
        if m:
            if 当前 is not None and 当前.get("范围"):
                结果.append(当前)
            当前 = {"序号": m.group(1), "名称": m.group(2).strip(), "起始行": int(m.group(3)), "范围": []}
            continue
        if 当前 is not None:
            if re.match(r"^##\s", 行):
                if 当前.get("范围"):
                    结果.append(当前)
                当前 = None
                continue
            for m2 in re.finditer(r"已查看范围：((?:`[A-Z]+\d+:[A-Z]+\d+`、?)+)", 行):
                当前["范围"].extend(x.strip("`") for x in m2.group(1).split("、"))
    if 当前 is not None and 当前.get("范围"):
        结果.append(当前)
    return 结果


def _样式1复核范围(路径):
    按小表 = defaultdict(list)
    for r in 读jsonl(路径):
        for t in r.get("containing_review_tables") or []:
            if t.get("subtable_order") and t.get("source_range"):
                按小表[t["subtable_order"]].append(t["source_range"])
    return {k: _bbox(v) for k, v in 按小表.items()}


def _小表编号(样式, 工作簿名, 工作表名, 名称, 业务坐标):
    材料 = json.dumps({"样式": 样式, "工作簿": 工作簿名, "工作表": 工作表名,
                     "名称": 名称, "业务坐标": sorted(业务坐标)},
                    ensure_ascii=False, sort_keys=True)
    return "st:" + hashlib.sha256(材料.encode("utf-8")).hexdigest()[:16]


def _单元格(工作表记录, 行, 列):
    return 工作表记录["cells"].get(_坐标(行, 列))


def _取显示文字(格):
    """查找用文字：普通格取原值，公式格取缓存显示；其他返回None。"""
    if 格 is None:
        return None
    if 格.get("formula"):
        return 格.get("cached_value")
    return 格.get("value")


def _提取小表(样式, 工作簿名, 工作表名, 名称, 记录们, 边界, 依据, 工作表记录,
              表现按映射, 项目列字母, 缺口, 去向):
    左, 顶, 右, 底 = 边界
    业务映射 = {}
    结构格 = {}
    for r in 记录们:
        if r.get("target_slot_id"):
            业务映射[r["cell"]] = r
        else:
            结构格[r["cell"]] = r
    列语义文字 = {规范文字(r.get("column_semantics")) for r in 业务映射.values() if r.get("column_semantics")}
    行语义文字 = {规范文字(r.get("row_semantics")) for r in 业务映射.values() if r.get("row_semantics")}
    业务行 = set()
    业务列 = set()
    for coord in 业务映射:
        c, t, _, _ = range_boundaries(coord)
        业务行.add(t)
        业务列.add(c)
    项目列 = column_index_from_string(项目列字母) if 项目列字母 else None
    首业务行 = min(业务行) if 业务行 else 顶

    合并覆盖 = {}
    for m in 工作表记录.get("merges", []):
        ml, mt, mr, mb = range_boundaries(m)
        if ml <= 右 and mr >= 左 and mt <= 底 and mb >= 顶:
            合并覆盖[_坐标(mt, ml)] = m

    节点 = []
    空白位置 = []
    for 行 in range(顶, 底 + 1):
        for 列 in range(左, 右 + 1):
            坐标 = _坐标(行, 列)
            格 = _单元格(工作表记录, 行, 列)
            if 坐标 in 业务映射:
                r = 业务映射[坐标]
                表现 = 表现按映射.get(r["mapping_id"])
                if 表现 is None:
                    缺口.append({"小表": 名称, "坐标": 坐标, "问题": "映射在统一表现记录中不存在",
                               "mapping_id": r["mapping_id"]})
                    去向.append({"mapping_id": r["mapping_id"], "去向": "缺口"})
                    continue
                节点.append({"坐标": 坐标, "角色": "业务数值", "行": 行, "列": 列,
                           "原文": 格.get("value") if 格 else r.get("raw_value"),
                           "公式": (格 or {}).get("formula") or (r.get("raw_value") if r.get("is_formula") else None),
                           "缓存": (格 or {}).get("cached_value"),
                           "隐藏": (格 or {}).get("hidden", False),
                           "语义引用": {"mapping_id": r["mapping_id"],
                                     "representation_id": 表现["representation_id"],
                                     "semantic_id": 表现["semantic_id"],
                                     "维度": 表现.get("规范维度值"),
                                     "定义旧ID": 表现.get("定义旧ID"),
                                     "确认依据": 表现.get("确认依据")}})
                去向.append({"mapping_id": r["mapping_id"], "去向": "业务格进入索引"})
                continue
            if 坐标 in 结构格:
                r = 结构格[坐标]
                语义 = f"{r.get('row_semantics') or ''}{r.get('column_semantics') or ''}"
                角色 = "表头" if "标题" in 语义 else ("检验" if "检验" in 语义 else "结构排除")
                显示 = _取显示文字(格)
                节点.append({"坐标": 坐标, "角色": 角色, "行": 行, "列": 列,
                           "原文": 显示 if isinstance(显示, str) else (格 or {}).get("value"),
                           "查找文字": 规范文字(显示) if isinstance(显示, str) else None,
                           "公式": (格 or {}).get("formula"), "缓存": (格 or {}).get("cached_value"),
                           "隐藏": (格 or {}).get("hidden", False),
                           "依据": f"原映射结构格排除:{r.get('reason') or 语义}"})
                去向.append({"mapping_id": r["mapping_id"], "去向": "结构格记录"})
                continue
            显示 = _取显示文字(格)
            if 显示 is None or (isinstance(显示, str) and not 规范文字(显示)):
                if 格 is None or (格.get("value") is None and not 格.get("formula")):
                    空白位置.append(坐标)
                continue
            文字 = 规范文字(显示) if isinstance(显示, str) else None
            if 文字 is None:
                节点.append({"坐标": 坐标, "角色": "未分类数值", "行": 行, "列": 列,
                           "原文": 显示, "隐藏": 格.get("hidden", False), "依据": "表区内未映射数值"})
                continue
            if 文字 == 规范文字(名称):
                角色, 角色依据 = "标题", "与小表名称一致"
            elif 文字 in 列语义文字:
                角色, 角色依据 = "表头", "与映射列语义文字一致"
            elif 文字 in 行语义文字:
                角色, 角色依据 = "项目", "与映射行语义文字一致"
            elif 项目列 is not None and 列 == 项目列 and 行 in 业务行:
                角色, 角色依据 = "项目", "位于项目列且与业务格同行（补充结构依据）"
            elif 行 < 首业务行 and 列 in 业务列:
                角色, 角色依据 = "表头", "位于业务列首业务行上方（位置事实）"
            else:
                角色, 角色依据 = "说明", "表区内其他文字"
            节点.append({"坐标": 坐标, "角色": 角色, "行": 行, "列": 列,
                       "原文": 显示, "查找文字": 文字,
                       "公式": 格.get("formula"), "缓存": 格.get("cached_value"),
                       "隐藏": 格.get("hidden", False), "依据": 角色依据})

    项目顺序 = []
    上级 = None
    for n in sorted((n for n in 节点 if n["角色"] == "项目"), key=lambda n: (n["行"], n["列"])):
        原文 = str(n.get("原文") or "")
        缩进 = len(原文) - len(原文.lstrip(" 　"))
        是其中 = 规范文字(原文).startswith("其中")
        条目 = {"坐标": n["坐标"], "原文": 原文, "行": n["行"], "缩进": 缩进,
              "是否合计": 规范文字(原文).startswith("合计"), "上级": None}
        if 是其中 or 缩进 > 0:
            条目["上级"] = 上级
        elif not 条目["是否合计"]:
            上级 = n["坐标"]
        项目顺序.append(条目)

    表头 = []
    for 列 in sorted(业务列):
        覆盖 = [n for n in 节点 if n["角色"] == "表头" and n["列"] == 列 and n["行"] < 首业务行]
        if 覆盖:
            for n in 覆盖:
                表头.append({"坐标": n["坐标"], "原文": n.get("原文"), "查找文字": n.get("查找文字"),
                           "覆盖列": get_column_letter(列), "依据": n.get("依据")})
        else:
            缺口.append({"小表": 名称, "坐标": f"{get_column_letter(列)}列", "问题": "该业务列缺表头归属依据"})

    上下文 = []
    上一小表底 = 顶
    for 行 in range(顶 - 1, max(0, 顶 - 4), -1):
        行文字 = [(列, _单元格(工作表记录, 行, 列)) for 列 in range(1, max(右, 项目列 or 1) + 1)]
        行文字 = [(列, 格) for 列, 格 in 行文字 if 格 is not None and isinstance(_取显示文字(格), str) and 规范文字(_取显示文字(格))]
        if not 行文字:
            break
        for 列, 格 in 行文字:
            上下文.append({"坐标": _坐标(行, 列), "原文": _取显示文字(格),
                         "查找文字": 规范文字(_取显示文字(格)), "角色": "上级标题" if 行 >= 顶 - 1 else "前置说明",
                         "依据": "表区上方相邻文字（位置事实）"})
        上一小表底 = 行
    上下文.reverse()

    业务坐标 = sorted(业务映射.keys(), key=lambda a: (range_boundaries(a)[1], range_boundaries(a)[0]))
    return {
        "小表编号": _小表编号(样式, 工作簿名, 工作表名, 名称, 业务坐标),
        "样式": 样式,
        "来源": {"工作簿": 工作簿名, "工作表": 工作表名, "映射文件": sorted({r["__文件"] for r in 记录们})},
        "小表名称": 名称,
        "表区范围": {"范围": f"{_坐标(顶, 左)}:{_坐标(底, 右)}", "依据": 依据},
        "上下文": 上下文,
        "节点": 节点,
        "项目顺序": 项目顺序,
        "表头": 表头,
        "合并": sorted(set(合并覆盖.values())),
        "空白位置": 空白位置,
    }


def 建立索引(输入清单: str, 输出目录: str, 旧索引目录: str | None = None) -> dict:
    清单 = json.loads(Path(输入清单).read_text(encoding="utf-8-sig")) if isinstance(输入清单, (str, Path)) else 输入清单
    核对输入清单(清单)
    输出目录 = Path(输出目录)
    if 输出目录.exists():
        raise ValueError(f"输出目录已存在，拒绝覆盖：{输出目录}")

    按用途 = {e["用途"]: e for e in 清单["条目"]}

    def 取(前缀):
        return [e for 名, e in 按用途.items() if 名.startswith(前缀)]

    表现按映射 = {}
    for p in 读jsonl(按用途["统一成果:表现记录"]["路径"]):
        if p.get("source_kind") == "five_style":
            表现按映射[p["source_locator"]["mapping_id"]] = p

    语义别名 = {}
    for r in 读jsonl(按用途["统一成果:统一语义记录"]["路径"]):
        别名 = ((r.get("meaning") or {}).get("原表达") or {}).get("aliases") or {}
        文字们 = [str(t) for t in 别名.get("row_labels", []) + 别名.get("column_labels", [])]
        if 文字们:
            语义别名[r["semantic_id"]] = 文字们

    工作簿 = {i: 按用途[f"五样式原Excel:附注样式{i}.xlsx"] for i in range(1, 6)}
    映射文件 = defaultdict(list)
    for e in 取("五样式原映射:"):
        m = re.search(r"style(\d)-", e["用途"])
        映射文件[int(m.group(1))].append(e["路径"])

    旧记录 = {}
    旧指纹 = {}
    旧去向 = []
    旧缺口 = []
    沿革 = []
    if 旧索引目录:
        旧索引目录 = Path(旧索引目录)
        旧清单 = json.loads((旧索引目录 / "索引清单.json").read_text(encoding="utf-8-sig"))
        旧指纹 = 旧清单.get("样式指纹", {})
        for r in 读小表记录文件(旧索引目录 / "小表记录.jsonl"):
            旧记录.setdefault(r["样式"], []).append(r)
        旧去向 = 读jsonl(旧索引目录 / "去向明细.jsonl")
        旧缺口 = 旧清单.get("缺口", [])

    全部记录 = []
    缺口 = []
    去向 = []
    样式指纹 = {}
    for 样式 in range(1, 6):
        来源材料 = "|".join([工作簿[样式]["sha256"],
                          *[e["sha256"] for e in sorted(取("五样式原映射:"), key=lambda e: e["用途"]) if f"style{样式}-" in e["用途"]],
                          (按用途.get(f"复核文档:样式{样式}逐表复核记录") or 按用途.get("复核文档:样式1逐格分类清单") or {"sha256": ""})["sha256"] if 样式 != 2 else "",
                          按用途["统一成果:表现记录"]["sha256"], 按用途["统一成果:统一语义记录"]["sha256"], 规则版本])
        指纹 = hashlib.sha256(来源材料.encode("utf-8")).hexdigest()
        样式指纹[str(样式)] = 指纹
        if 旧指纹.get(str(样式)) == 指纹 and 样式 in 旧记录:
            全部记录.extend(旧记录[样式])
            去向.extend(d for d in 旧去向 if d.get("样式") == 样式)
            缺口.extend(g for g in 旧缺口 if g.get("样式") == 样式)
            沿革.append({"样式": 样式, "动作": "沿用旧记录", "原因": "来源指纹与规则版本未变"})
            continue
        if 旧索引目录:
            沿革.append({"样式": 样式, "动作": "重新提取", "原因": "来源内容或规则版本变化"})
        记录们, 去向们, 缺口们 = _提取样式(样式, 工作簿[样式]["路径"], 映射文件[样式], 按用途, 表现按映射)
        全部记录.extend(记录们)
        去向.extend(去向们)
        缺口.extend(缺口们)

    全部记录.sort(key=lambda r: r["小表编号"])
    编号计数 = defaultdict(int)
    for r in 全部记录:
        编号计数[r["小表编号"]] += 1
    重复 = [k for k, v in 编号计数.items() if v > 1]
    if 重复:
        raise ValueError(f"小表编号冲突：{重复[:5]}")

    查找 = defaultdict(list)
    for r in 全部记录:
        for n in r["节点"]:
            if n.get("查找文字") and n["角色"] in ("标题", "表头", "项目", "说明", "检验", "结构排除"):
                查找[n["查找文字"]].append({"小表编号": r["小表编号"], "样式": r["样式"],
                                        "工作表": r["来源"]["工作表"], "坐标": n["坐标"], "角色": n["角色"]})
    编号到小表 = {r["小表编号"]: r for r in 全部记录}
    语义到小表 = defaultdict(set)
    for r in 全部记录:
        for n in r["节点"]:
            引用 = n.get("语义引用")
            if 引用:
                语义到小表[引用["semantic_id"]].add(r["小表编号"])
    别名查找 = defaultdict(list)
    for sid, 文字们 in 语义别名.items():
        for 文字 in 文字们:
            规范 = 规范文字(文字)
            if not 规范:
                continue
            for 编号 in sorted(语义到小表.get(sid, ())):
                别名查找[规范].append({"小表编号": 编号, "semantic_id": sid,
                                  "依据": "统一语义记录原表达别名（已确认叫法）"})

    输出目录.mkdir(parents=True)
    with open(输出目录 / "小表记录.jsonl", "w", encoding="utf-8-sig", newline="\n") as f:
        for r in 全部记录:
            f.write(json.dumps(r, ensure_ascii=False, sort_keys=True, default=_json默认) + "\n")
    with open(输出目录 / "文字查找.json", "w", encoding="utf-8-sig", newline="\n") as f:
        json.dump({"规则版本": 规则版本,
                  "正文": {k: v for k, v in sorted(查找.items())},
                  "别名": {k: v for k, v in sorted(别名查找.items())}},
                 f, ensure_ascii=False, indent=1, sort_keys=True, default=_json默认)
    去向计数 = defaultdict(int)
    for d in 去向:
        去向计数[d["去向"]] += 1
    with open(输出目录 / "去向明细.jsonl", "w", encoding="utf-8-sig", newline="\n") as f:
        for d in sorted(去向, key=lambda d: (d.get("样式") or 0, d.get("mapping_id") or "")):
            f.write(json.dumps(d, ensure_ascii=False, sort_keys=True, default=_json默认) + "\n")
    索引清单 = {
        "规则版本": 规则版本,
        "输入清单": str(Path(输入清单).resolve()) if isinstance(输入清单, (str, Path)) else "（内存清单）",
        "输入指纹": {e["用途"]: e["sha256"] for e in 清单["条目"]},
        "实际读取": ["五样式原Excel", "五样式原映射", "复核文档", "统一成果:表现记录", "统一成果:统一语义记录", "复用代码:机械读取模块"],
        "样式指纹": 样式指纹,
        "小表数": len(全部记录),
        "按样式": {str(i): sum(1 for r in 全部记录 if r["样式"] == i) for i in range(1, 6)},
        "映射去向": dict(去向计数),
        "缺口": 缺口,
        "沿革": 沿革,
    }
    with open(输出目录 / "索引清单.json", "w", encoding="utf-8-sig", newline="\n") as f:
        json.dump(索引清单, f, ensure_ascii=False, indent=1, sort_keys=True, default=_json默认)
    return {"输出目录": str(输出目录), "小表数": len(全部记录), "映射去向": dict(去向计数), "缺口数": len(缺口)}


def 读小表记录文件(路径):
    return 读jsonl(路径)


def _提取样式(样式, 工作簿路径, 映射路径们, 按用途, 表现按映射):
    工作簿记录 = 读取工作簿(工作簿路径)
    记录们 = []
    for 路径 in 映射路径们:
        for r in 读jsonl(路径):
            r["__文件"] = str(Path(路径).name)
            记录们.append(r)
    缺口 = []
    去向 = []
    结果 = []

    if 样式 == 1:
        复核范围 = _样式1复核范围(按用途["复核文档:样式1逐格分类清单"]["路径"])
        分组 = defaultdict(list)
        for r in 记录们:
            分组[r["subtable_order"]].append(r)
        for 序号, 组 in 分组.items():
            工作表名 = 组[0]["worksheet"]
            表 = next(s for s in 工作簿记录["sheets"] if s["name"] == 工作表名)
            if 序号 not in 复核范围:
                盒 = _bbox([r["cell"] for r in 组])
                依据 = "复核清单无该小表范围，退用业务格外包矩形（缺表头边界依据）"
                缺口.append({"小表": 组[0].get("subtable"), "问题": "逐格分类清单缺少该小表范围", "subtable_order": 序号})
            else:
                盒 = 复核范围[序号]
                依据 = "样式1逐格分类清单containing_review_tables.source_range"
                越界 = [r["cell"] for r in 组 if not _在盒内(r["cell"], 盒)]
                if 越界:
                    缺口.append({"小表": 组[0].get("subtable"), "问题": f"映射格超出复核范围：{越界[:5]}"})
            结果.append(_提取小表(样式, "附注样式1.xlsx", 工作表名, 组[0].get("subtable") or f"小表{序号}",
                              组, 盒, 依据, 表, 表现按映射, "B", 缺口, 去向))
    elif 样式 == 2:
        分组 = defaultdict(list)
        for r in 记录们:
            分组[r["__文件"]].append(r)
        for 文件, 组 in sorted(分组.items()):
            工作表名 = 组[0]["worksheet"]
            表 = next(s for s in 工作簿记录["sheets"] if s["name"] == 工作表名)
            名称 = re.sub(r"^\d+-", "", 文件).removesuffix("-样式2-人工映射.jsonl")
            盒 = _bbox([r["cell"] for r in 组])
            左 = min(盒[0], 2)  # ponytail: 项目列B；列A为科目/其中标记，超出映射范围的列在缺口说明
            盒 = (左, 盒[1], 盒[2], 盒[3])
            上行 = _单元格(表, 盒[1] - 1, 1) if 盒[1] > 1 else None
            if 上行 is not None and isinstance(_取显示文字(上行), str) and 规范文字(_取显示文字(上行)):
                盒 = (盒[0], 盒[1] - 1, 盒[2], 盒[3])  # A列科目名行纳入表区（补充依据）
            依据 = "映射文件逐格记录（含结构格排除）边界，左扩至项目列B（执行agent回看原件补充，确认状态：已核实）"
            右外 = [c for c in range(盒[2] + 1, 盒[2] + 3)
                  for 行 in range(盒[1], 盒[3] + 1)
                  if (g := _单元格(表, 行, c)) is not None and _取显示文字(g) not in (None, "")]
            if 右外:
                缺口.append({"小表": 名称, "问题": "映射范围右侧存在未映射内容列（样式2原映射只覆盖部分年度列）",
                           "位置": f"{_坐标(盒[1], min(右外))}起{len(set(右外))}列"})
            结果.append(_提取小表(样式, "附注样式2.xlsx", 工作表名, 名称, 组, 盒, 依据, 表, 表现按映射, "B", 缺口, 去向))
            记录 = 结果[-1]
            if not 记录["表头"]:
                共用 = []
                for 列字母 in ("C", "D"):
                    格 = _单元格(表, 3, column_index_from_string(列字母))
                    if 格 is not None and (格.get("formula") or 格.get("value") is not None):
                        坐标 = f"{列字母}3"
                        共用.append({"坐标": 坐标, "原文": 格.get("value"), "查找文字": 规范文字(格.get("cached_value")),
                                   "覆盖列": 列字母, "共用": True,
                                   "依据": "样式2全表共用附注1第3行表头；001货币资金映射文件登记为列标题，其余小表共用（执行agent回看原件补充，确认状态：已核实）"})
                if 共用:
                    # 共用表头只作外部所属依据独立保存（进"表头"不进"节点"）：
                    # 它到小表正文之间的整段距离不是小表正文，不参与查找定位、覆盖分母与候选范围
                    记录["表头"].extend(共用)
                    记录["缺口_共用表头"] = "表头为全表共用行，E、F列（2023/2022年度）未纳入原映射"
    elif 样式 == 3:
        复核 = _解析样式3复核(按用途["复核文档:样式3逐表复核记录"]["路径"])
        分组 = defaultdict(list)
        for r in 记录们:
            分组[r["subtable_order"]].append(r)
        for 序号, 组 in 分组.items():
            工作表名 = 组[0]["worksheet"]
            表 = next(s for s in 工作簿记录["sheets"] if s["name"] == 工作表名)
            if 序号 in 复核:
                节 = 复核[序号]
                列盒 = _bbox([r["cell"] for r in 组])
                盒 = (2, 节["顶"] - 1, max(列盒[2], 2), 节["底"])  # 顶上一行为标题行（位置事实）
                依据 = f"样式3逐表复核记录第{序号}节（{节['顶']}—{节['底']}行），标题取上一行"
            else:
                盒 = _bbox([r["cell"] for r in 组])
                依据 = "复核记录缺该小表节，退用业务格外包矩形（缺边界依据）"
                缺口.append({"小表": 组[0].get("subtable"), "问题": "复核记录缺节", "subtable_order": 序号})
            结果.append(_提取小表(样式, "附注样式3.xlsx", 工作表名, 组[0].get("subtable") or f"小表{序号}",
                              组, 盒, 依据, 表, 表现按映射, "B", 缺口, 去向))
    elif 样式 == 4:
        复核 = _解析样式4复核(按用途["复核文档:样式4逐表复核记录"]["路径"])
        复核按表 = {}
        for 名称, 范围 in 复核.items():
            复核按表[名称] = 范围
        分组 = defaultdict(list)
        for r in 记录们:
            分组[r["worksheet"]].append(r)
        for 工作表名, 组 in sorted(分组.items()):
            表 = next(s for s in 工作簿记录["sheets"] if s["name"] == 工作表名)
            范围 = 复核按表.get(工作表名.strip()) or 复核按表.get(工作表名)
            if 范围:
                盒 = range_boundaries(范围)
                依据 = f"样式4逐表复核记录人工查看范围{范围}"
            else:
                盒 = _bbox([r["cell"] for r in 组])
                依据 = "复核记录缺对应节，退用业务格外包矩形（缺边界依据）"
                缺口.append({"小表": 工作表名, "问题": "复核记录缺对应节"})
            结果.append(_提取小表(样式, "附注样式4.xlsx", 工作表名, 工作表名.strip(), 组, 盒, 依据, 表, 表现按映射, "A", 缺口, 去向))
    elif 样式 == 5:
        节们 = _解析样式5复核(按用途["复核文档:样式5逐表复核记录"]["路径"])
        工作表名 = 记录们[0]["worksheet"]
        表 = next(s for s in 工作簿记录["sheets"] if s["name"] == 工作表名)
        盒们 = [(节, [range_boundaries(x) for x in 节["范围"]]) for 节 in 节们]
        分组 = defaultdict(list)
        重叠记录 = []
        for r in 记录们:
            命中 = [节 for 节, 盒列表 in 盒们 if any(_在盒内(r["cell"], b) for b in 盒列表)]
            if len(命中) == 1:
                分组[命中[0]["序号"]].append(r)
            elif len(命中) > 1:
                行号 = range_boundaries(r["cell"])[1]
                候选 = [节 for 节 in 命中 if 节["起始行"] <= 行号]
                属 = max(候选 or 命中, key=lambda 节: 节["起始行"])
                分组[属["序号"]].append(r)
                重叠记录.append(r["cell"])
            else:
                缺口.append({"小表": None, "坐标": r["cell"], "问题": "映射格落在0个复核范围内"})
                去向.append({"mapping_id": r["mapping_id"], "去向": "缺口"})
        if 重叠记录:
            缺口.append({"小表": None, "问题": "相邻复核范围重叠，按起始行归属较后一节（复核记录行号预留空行所致）",
                       "影响格数": len(重叠记录), "示例": 重叠记录[:5]})
        for 节 in 节们:
            组 = 分组.get(节["序号"])
            if not 组:
                continue
            盒 = _bbox(节["范围"])
            结果.append(_提取小表(样式, "附注样式5.xlsx", 工作表名, 节["名称"], 组, 盒,
                              f"样式5逐表复核记录第{节['序号']}节已查看范围{'、'.join(节['范围'])}",
                              表, 表现按映射, "A", 缺口, 去向))
    for d in 去向:
        d.setdefault("样式", 样式)
    for g in 缺口:
        g.setdefault("样式", 样式)
    return 结果, 去向, 缺口


def _在盒内(坐标, 盒):
    列, 行, _, _ = range_boundaries(坐标)
    return 盒[0] <= 列 <= 盒[2] and 盒[1] <= 行 <= 盒[3]
