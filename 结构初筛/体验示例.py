"""生成完全虚构的资金小表，演示离线初筛；每次保存到新目录。"""
import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from 测试收尾 import 虚构参照, 内存索引
from 运行初筛 import 执行比较
from 建立结构索引 import 索引文件指纹


def main():
    根 = Path(__file__).resolve().parent
    输出 = 根 / '运行结果' / ('虚构示例_' + datetime.now().strftime('%Y%m%d_%H%M%S_%f'))
    索引目录 = 输出 / '演示索引'
    索引目录.mkdir(parents=True, exist_ok=False)
    表 = 虚构参照()
    表['上层语义'] = {'完整': True, '业务路径': [{'附注': '虚构货币资金', '表义': '虚构余额构成',
        '映射依据': ['虚构演示依据，不是正式金标准']}],
        '名称依据': [{'原文': '演示表', '来源': '虚构工作表名称'},
                    {'原文': '演示货币资金', '来源': '虚构小表名称'}]}
    索引 = 内存索引(表)
    去向 = []
    正文 = {}
    for n in 表['节点']:
        if n['角色'] == '业务数值':
            身份 = '虚构演示:' + n['坐标']
            n['语义引用'].update(mapping_id=身份, representation_id=身份, semantic_id=身份,
                              维度={}, 确认依据={'说明': '虚构演示，未经业务语义批准'})
            去向.append({'样式': 表['样式'], 'mapping_id': 身份, '去向': '业务格进入索引'})
        if n.get('查找文字'):
            正文.setdefault(n['查找文字'], []).append({'小表编号': 表['小表编号'], '样式': 表['样式'],
                '工作表': 表['来源']['工作表'], '坐标': n['坐标'], '角色': n['角色']})
    (索引目录 / '小表记录.jsonl').write_text(json.dumps(表, ensure_ascii=False)+'\n', encoding='utf-8-sig')
    (索引目录 / '文字查找.json').write_text(json.dumps({'规则版本': '虚构演示', '正文': 正文, '别名': {}}, ensure_ascii=False), encoding='utf-8-sig')
    (索引目录 / '去向明细.jsonl').write_text(''.join(json.dumps(d, ensure_ascii=False)+'\n' for d in 去向), encoding='utf-8-sig')
    清单 = {'规则版本': '虚构演示', '说明': '演示只有一张虚构小表，不是正式714张索引',
          '小表数': 1, '按样式': {'1': 1}, '映射去向': {'业务格进入索引': len(去向)}, '缺口': [],
          '生成物指纹': 索引文件指纹(索引目录)}
    (索引目录 / '索引清单.json').write_text(json.dumps(清单, ensure_ascii=False), encoding='utf-8-sig')
    w = Workbook(); s = w.active; s.title = '演示表'
    for n in 表['节点']:
        s.cell(n['行'], n['列'], n['原文'])
    s['B3'] = 123
    s['B4'] = 456
    s['B5'] = 579
    新表 = 输出 / '虚构新表.xlsx'
    w.save(新表); w.close()
    码 = 执行比较(新表, 索引目录, 输出 / '初筛材料', True)
    print('以上金额完全虚构。业务数值变动不会改变结构文字的对应；初筛没有批准语义。')
    print('请打开结果目录中的“初筛报告.md”和“LLM交接说明.md”查看。')
    return 码


if __name__ == '__main__':
    sys.exit(main())
